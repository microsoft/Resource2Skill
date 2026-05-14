def create_slide(
    output_pptx_path: str,
    title_text: str = "V/S",
    subtitle_text: str = "COMPARISON",
    left_palette: list = [(0, 191, 255), (0, 160, 220), (0, 130, 190), (0, 100, 160)],
    right_palette: list = [(255, 20, 147), (220, 15, 120), (190, 10, 100), (160, 5, 80)],
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dual-Axis Split Infographic Comparison" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # 1. Set Slide Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(235, 235, 240)
    
    # --- HELPER FUNCTIONS FOR LXML INJECTIONS ---
    
    def add_shadow(shape):
        """Adds a subtle offset drop shadow to a shape."""
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="50800" dist="38100" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="20000"/>
                </a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        shape.element.spPr.append(parse_xml(shadow_xml))

    def make_pill_shape(shape):
        """Forces a rounded rectangle to have fully rounded (pill) ends."""
        avLst = shape.element.spPr.prstGeom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
        if avLst is None:
            avLst = parse_xml('<a:avLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
            shape.element.spPr.prstGeom.append(avLst)
        # 50000 = 50% radius
        gd = parse_xml('<a:gd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="adj" fmla="val 50000"/>')
        avLst.append(gd)

    def apply_gradient(shape, color1, color2):
        """Applies a linear gradient fill to a shape."""
        c1_hex = '%02x%02x%02x' % color1
        c2_hex = '%02x%02x%02x' % color2
        grad_xml = f"""
        <a:gradFill rotWithShape="1" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:gsLst>
                <a:gs pos="0"><a:srgbClr val="{c1_hex}"/></a:gs>
                <a:gs pos="100000"><a:srgbClr val="{c2_hex}"/></a:gs>
            </a:gsLst>
            <a:lin ang="5400000" scaled="0"/>
        </a:gradFill>
        """
        # Remove solidFill if exists, then append gradFill
        solid_fill = shape.element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        if solid_fill is not None:
            shape.element.spPr.remove(solid_fill)
        shape.element.spPr.append(parse_xml(grad_xml))

    # --- CENTER TEXT ---
    tb = slide.shapes.add_textbox(Inches(5.66), Inches(0.5), Inches(2.0), Inches(1.0))
    tf = tb.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = "Calibri"
    p.font.color.rgb = RGBColor(40, 50, 60)
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.name = "Calibri"
    p2.font.color.rgb = RGBColor(40, 50, 60)
    p2.alignment = PP_ALIGN.CENTER


    # --- GENERATOR ENGINE FOR SYMMETRICAL SIDES ---
    def create_side(side="left"):
        is_left = (side == "left")
        
        # Dimensions & Coordinates
        panel_w, panel_h = 3.0, 3.4
        circle_d = 2.0
        inner_d = 1.6
        bar_w, bar_h = 3.2, 0.4
        
        # Layout Math ensuring symmetry
        if is_left:
            panel_l = 1.5
            bars_l = 3.0
            letter = "A"
            colors = left_palette
            base_dark = (26, 45, 84)
            align = PP_ALIGN.LEFT
        else:
            panel_l = 13.333 - 1.5 - panel_w
            bars_l = 13.333 - 3.0 - bar_w
            letter = "B"
            colors = right_palette
            base_dark = (60, 20, 45)
            align = PP_ALIGN.RIGHT
            
        panel_t = 0.8
        circle_l = panel_l + (panel_w / 2) - (circle_d / 2)
        circle_t = panel_t + panel_h - 0.2
        
        # 1. Vertical Panel
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(panel_l), Inches(panel_t), Inches(panel_w), Inches(panel_h))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        panel.line.fill.background()
        add_shadow(panel)
        
        # Panel Text
        tf = panel.text_frame
        tf.margin_top = Inches(0.3)
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        p = tf.paragraphs[0]
        p.text = letter
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*colors[0])
        p.alignment = PP_ALIGN.CENTER
        
        p_desc = tf.add_paragraph()
        p_desc.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Fusce posuere, magna sed."
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = RGBColor(100, 100, 100)
        p_desc.alignment = PP_ALIGN.CENTER
        
        # 2. Horizontal Data Bars
        bar_start_y = circle_t + 0.2
        for i in range(4):
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, 
                Inches(bars_l), Inches(bar_start_y + (i * 0.5)), 
                Inches(bar_w), Inches(bar_h)
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(*colors[i])
            bar.line.fill.background()
            make_pill_shape(bar)
            add_shadow(bar)
            
            # Bar Text
            btf = bar.text_frame
            btf.clear()
            bp = btf.paragraphs[0]
            bp.text = f"0{i+1}   Add your detailed text here" if is_left else f"Add your detailed text here   0{i+1}"
            bp.font.size = Pt(10)
            bp.font.color.rgb = RGBColor(255, 255, 255)
            bp.alignment = PP_ALIGN.CENTER
            
        # 3. Outer Dark Badge
        outer_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(circle_l), Inches(circle_t), Inches(circle_d), Inches(circle_d))
        outer_circle.fill.solid()
        outer_circle.fill.fore_color.rgb = RGBColor(*base_dark)
        outer_circle.line.fill.background()
        add_shadow(outer_circle)
        
        # 4. Inner Gradient Badge
        inner_l = circle_l + ((circle_d - inner_d) / 2)
        inner_t = circle_t + ((circle_d - inner_d) / 2)
        inner_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(inner_l), Inches(inner_t), Inches(inner_d), Inches(inner_d))
        inner_circle.line.fill.background()
        apply_gradient(inner_circle, colors[0], colors[2])
        
        # Inner Circle Text/Icon
        ictf = inner_circle.text_frame
        icp = ictf.paragraphs[0]
        icp.text = "👤" # Placeholder icon
        icp.font.size = Pt(36)
        icp.alignment = PP_ALIGN.CENTER

    # Execute generators
    create_side("left")
    create_side("right")

    prs.save(output_pptx_path)
    return output_pptx_path
