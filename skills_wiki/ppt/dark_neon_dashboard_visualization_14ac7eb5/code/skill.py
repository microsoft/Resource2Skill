def create_slide(
    output_pptx_path: str,
    title_text: str = "Dark Neon Dashboard",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Dark Neon Dashboard visual effect.
    Generates a dark background with UI cards containing custom-styled neon charts.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Palette
    bg_color = RGBColor(20, 22, 28)
    card_color = RGBColor(30, 32, 38)
    card_border = RGBColor(50, 55, 65)
    tag_color = RGBColor(220, 60, 40)
    axis_text_color = RGBColor(150, 150, 160)
    neon_pink = RGBColor(255, 40, 130)
    neon_cyan = RGBColor(0, 220, 255)
    neon_lime = RGBColor(180, 255, 50)
    neon_purple = RGBColor(150, 50, 255)

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background() # No line

    # Grid configuration
    margin_x = Inches(0.5)
    margin_y = Inches(0.5)
    spacing = Inches(0.3)
    card_w = (prs.slide_width - (margin_x * 2) - spacing) / 2
    card_h = (prs.slide_height - (margin_y * 2) - spacing) / 2

    def draw_ui_card(slide, x, y, w, h, title):
        """Draws the dashboard container card and its floating tag."""
        # Main Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = card_color
        card.line.color.rgb = card_border
        
        # Tag Label (Pill shape slightly offset)
        tag_w, tag_h = Inches(1.2), Inches(0.3)
        tag = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            x - Inches(0.1), y + Inches(0.2), 
            tag_w, tag_h
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = tag_color
        tag.line.fill.background()
        
        tf = tag.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        
        return x + Inches(0.3), y + Inches(0.4), w - Inches(0.6), h - Inches(0.6)

    def style_axes(chart):
        """Applies dark UI styling to chart axes."""
        if chart.has_value_axis:
            val_axis = chart.value_axis
            val_axis.has_major_gridlines = True
            val_axis.major_gridlines.format.line.color.rgb = card_border
            val_axis.tick_labels.font.color.rgb = axis_text_color
            val_axis.tick_labels.font.size = Pt(10)
        
        if chart.has_category_axis:
            cat_axis = chart.category_axis
            cat_axis.tick_labels.font.color.rgb = axis_text_color
            cat_axis.tick_labels.font.size = Pt(10)
            
        chart.chart_area.fill.solid()
        chart.chart_area.fill.fore_color.rgb = card_color
        chart.plot_area.fill.solid()
        chart.plot_area.fill.fore_color.rgb = card_color
        
        if chart.has_legend:
            chart.legend.font.color.rgb = axis_text_color
            chart.legend.font.size = Pt(10)
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    # === Quadrant 1: Line Chart ===
    x, y = margin_x, margin_y
    cx, cy, cw, ch = draw_ui_card(slide, x, y, card_w, card_h, "Line Chart")
    
    cd1 = CategoryChartData()
    cd1.categories = ['2021', '2022', '2023', '2024']
    cd1.add_series('Growth', (2.5, 4.0, 3.2, 5.8))
    cd1.add_series('Metrics', (4.8, 3.1, 5.5, 4.2))
    
    chart1 = slide.shapes.add_chart(XL_CHART_TYPE.LINE, cx, cy, cw, ch, cd1).chart
    style_axes(chart1)
    
    chart1.series[0].format.line.color.rgb = neon_pink
    chart1.series[0].format.line.width = Pt(2.5)
    chart1.series[0].smooth = True
    chart1.series[1].format.line.color.rgb = neon_cyan
    chart1.series[1].format.line.width = Pt(2.5)
    chart1.series[1].smooth = True

    # === Quadrant 2: Column Chart ===
    x = margin_x + card_w + spacing
    cx, cy, cw, ch = draw_ui_card(slide, x, y, card_w, card_h, "Column Chart")
    
    cd2 = CategoryChartData()
    cd2.categories = ['2021', '2022', '2023', '2024']
    cd2.add_series('Alpha', (15, 22, 18, 35))
    cd2.add_series('Beta', (20, 15, 28, 25))
    
    chart2 = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, cx, cy, cw, ch, cd2).chart
    style_axes(chart2)
    
    chart2.series[0].format.fill.solid()
    chart2.series[0].format.fill.fore_color.rgb = neon_lime
    chart2.series[1].format.fill.solid()
    chart2.series[1].format.fill.fore_color.rgb = neon_purple

    # === Quadrant 3: Area Chart ===
    x, y = margin_x, margin_y + card_h + spacing
    cx, cy, cw, ch = draw_ui_card(slide, x, y, card_w, card_h, "Area Chart")
    
    cd3 = CategoryChartData()
    cd3.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    cd3.add_series('Volume', (100, 150, 120, 200))
    
    chart3 = slide.shapes.add_chart(XL_CHART_TYPE.AREA, cx, cy, cw, ch, cd3).chart
    style_axes(chart3)
    
    chart3.series[0].format.fill.solid()
    chart3.series[0].format.fill.fore_color.rgb = neon_cyan
    chart3.series[0].format.line.color.rgb = neon_cyan

    # === Quadrant 4: Doughnut Chart ===
    x = margin_x + card_w + spacing
    cx, cy, cw, ch = draw_ui_card(slide, x, y, card_w, card_h, "Doughnut")
    
    cd4 = CategoryChartData()
    cd4.categories = ['Mobile', 'Desktop', 'Tablet']
    cd4.add_series('Traffic', (55, 30, 15))
    
    chart4 = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, cx, cy, cw, ch, cd4).chart
    style_axes(chart4)
    chart4.has_legend = True
    
    # Assign specific colors to pie slices
    colors = [neon_pink, neon_cyan, neon_lime]
    for idx, point in enumerate(chart4.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = colors[idx % len(colors)]
        point.format.line.solid()
        point.format.line.color.rgb = card_color
        point.format.line.width = Pt(2)

    prs.save(output_pptx_path)
    return output_pptx_path
