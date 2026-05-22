def create_slide(
    output_pptx_path: str,
    title_text: str = "PROJECT OVERVIEW",
    body_text: str = "We pursue relationships based on transparency and mutual trust.",
    label_text: str = "CREATIVE PORTFOLIO",
    grad_hex_start: str = "FF0080",  # Hot Pink
    grad_hex_end: str = "00F0FF",    # Cyan
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Vibrant Geometric Framed Overlay' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Generation via PIL ===
    bg_img_path = "temp_dark_bg.png"
    try:
        # Fetch a random abstract/business image
        req = urllib.request.Request(
            "https://picsum.photos/1920/1080?grayscale", 
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req, timeout=5) as response:
            bg_img = Image.open(BytesIO(response.read())).convert("RGBA")
            bg_img = bg_img.resize((1920, 1080))
    except Exception:
        # Fallback if download fails
        bg_img = Image.new("RGBA", (1920, 1080), (40, 40, 45, 255))

    # Apply deep navy/black dark overlay for maximum contrast
    overlay = Image.new("RGBA", bg_img.size, (15, 20, 30, 200)) 
    final_bg = Image.alpha_composite(bg_img, overlay)
    final_bg.save(bg_img_path, format="PNG")

    # Insert background into slide
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Gradient Frame ===
    frame_width = Inches(8)
    frame_height = Inches(4.5)
    left = (prs.slide_width - frame_width) / 2
    top = (prs.slide_height - frame_height) / 2

    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, frame_width, frame_height)
    frame.fill.background()  # Make interior transparent
    frame.line.width = Pt(12) # Thick stroke

    # XML Injection for Gradient Stroke
    grad_xml = f"""
    <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">
        <a:gsLst>
            <a:gs pos="0"><a:srgbClr val="{grad_hex_start}"/></a:gs>
            <a:gs pos="100000"><a:srgbClr val="{grad_hex_end}"/></a:gs>
        </a:gsLst>
        <a:lin ang="2700000" scaled="1"/>
    </a:gradFill>
    """
    grad_element = parse_xml(grad_xml)
    ln = frame._element.spPr.get_or_add_ln()
    
    # Remove existing solidFill inside the line properties
    for child in list(ln):
        if child.tag.endswith('Fill'):
            ln.remove(child)
    ln.append(grad_element)

    # === Layer 3: Intersecting Label Badge ===
    badge_width = Inches(2.8)
    badge_height = Inches(0.4)
    badge_left = left + (frame_width - badge_width) / 2
    badge_top = top - (badge_height / 2)  # Overlap the top edge of the frame

    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, badge_left, badge_top, badge_width, badge_height)
    badge.fill.solid()
    # Color badge using the start of the gradient
    badge_r, badge_g, badge_b = int(grad_hex_start[:2], 16), int(grad_hex_start[2:4], 16), int(grad_hex_start[4:], 16)
    badge.fill.fore_color.rgb = RGBColor(badge_r, badge_g, badge_b)
    badge.line.fill.background() # No border for the badge

    badge_tf = badge.text_frame
    badge_tf.text = label_text.upper()
    badge_p = badge_tf.paragraphs[0]
    badge_p.alignment = PP_ALIGN.CENTER
    badge_p.font.name = "Arial Bold"
    badge_p.font.size = Pt(11)
    badge_p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 4: Core Text Formatting ===
    # Main Title
    tx_title = slide.shapes.add_textbox(left, top + Inches(1.2), frame_width, Inches(1.5))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text.upper()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial Black"
    p_title.font.size = Pt(54)
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Body
    if body_text:
        tx_body = slide.shapes.add_textbox(left + Inches(0.5), top + Inches(2.6), frame_width - Inches(1), Inches(1))
        tf_body = tx_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        p_body.text = body_text
        p_body.alignment = PP_ALIGN.CENTER
        p_body.font.name = "Arial"
        p_body.font.size = Pt(16)
        p_body.font.color.rgb = RGBColor(200, 200, 200) # Soft grey for secondary text

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
