def create_slide(
    output_pptx_path: str,
    title_text: str = "UNSTOPPABLE",
    accent_box_text: str = "U,",
    bg_color: tuple = (28, 31, 38),     # Dark charcoal
    line_color: tuple = (224, 49, 49),  # Tech Red
    box_color: tuple = (255, 212, 59),  # Cyber Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Tech-Industrial Registration Grid' title style.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Slide Dimensions ===
    w = prs.slide_width
    h = prs.slide_height

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # Helper function to create thin lines
    def add_line(left, top, width, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.fill.background() # No border
        return shape

    # === Layer 2: Tech Grid / Crop Marks ===
    line_thickness = Pt(2)
    mark_length = Inches(1.5)
    margin_x = Inches(2)
    margin_y = Inches(1.5)

    # Top Left Crop Mark
    add_line(margin_x, margin_y, mark_length, line_thickness, line_color) # Horizontal
    add_line(margin_x, margin_y, line_thickness, mark_length, line_color) # Vertical

    # Bottom Right Crop Mark
    add_line(w - margin_x - mark_length, h - margin_y, mark_length, line_thickness, line_color) # Horizontal
    add_line(w - margin_x, h - margin_y - mark_length, line_thickness, mark_length, line_color) # Vertical

    # Decorative "+" marks
    def add_plus_mark(left, top, color):
        size = Inches(0.2)
        add_line(left - size/2, top, size, Pt(1.5), color)
        add_line(left, top - size/2, Pt(1.5), size, color)

    add_plus_mark(w - margin_x, margin_y, line_color)
    add_plus_mark(margin_x, h - margin_y, line_color)


    # === Layer 3: Central Typography & Accent ===
    
    # 1. Main Title Text
    title_box_width = Inches(12)
    title_box_height = Inches(2)
    title_left = (w - title_box_width) / 2
    title_top = (h - title_box_height) / 2
    
    txBox = slide.shapes.add_textbox(title_left, title_top, title_box_width, title_box_height)
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(88)
    p.font.bold = True
    p.font.name = "Arial Black" # Fallback to a heavy standard font
    p.font.color.rgb = RGBColor(255, 255, 255)

    # 2. Yellow Accent Box (Placed above main text)
    accent_size = Inches(0.8)
    accent_left = (w - accent_size) / 2
    accent_top = title_top - Inches(0.5)

    accent_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, accent_left, accent_top, accent_size, accent_size
    )
    accent_shape.fill.solid()
    accent_shape.fill.fore_color.rgb = RGBColor(*box_color)
    accent_shape.line.fill.background()

    # 3. Text inside yellow accent box
    accent_tf = accent_shape.text_frame
    accent_tf.margin_left = Pt(0)
    accent_tf.margin_right = Pt(0)
    accent_p = accent_tf.paragraphs[0]
    accent_p.text = accent_box_text
    accent_p.alignment = PP_ALIGN.CENTER
    accent_p.font.size = Pt(28)
    accent_p.font.bold = True
    accent_p.font.name = "Arial Black"
    accent_p.font.color.rgb = RGBColor(20, 20, 20) # Dark text for contrast

    prs.save(output_pptx_path)
    return output_pptx_path
