def create_slide(
    output_pptx_path: str,
    title_text: str = "JING TIAN",
    subtitle_text: str = "HU GE / LEAD ACTOR",
    body_text: str = "“ The heavens and earth are vast, \nbut true happiness is the greatest. ”\n\nPossessing a unique talent for appraising ancient treasures, clever and quick-witted. Dreams of becoming a legendary swordsman.",
    bg_keyword: str = "cinematic portrait sword",
    base_color: tuple = (15, 20, 25),  # RGB for the solid blend (dark slate)
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Seamless Gradient Image Blend' poster effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageOps
    import urllib.request
    import io

    # 1. Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Dimensions in pixels (standard 1080p 16:9)
    WIDTH, HEIGHT = 1920, 1080

    # 2. Fetch and Prepare Background Image
    try:
        # Use Unsplash source to get an image matching the theme
        req = urllib.request.Request(
            f"https://source.unsplash.com/featured/1920x1080/?{urllib.parse.quote(bg_keyword)}",
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req, timeout=10) as response:
            bg_image_data = response.read()
        
        base_img = Image.open(io.BytesIO(bg_image_data)).convert("RGBA")
        # Ensure it's exactly 16:9 without warping
        base_img = ImageOps.fit(base_img, (WIDTH, HEIGHT), method=Image.Resampling.LANCZOS)
    except Exception as e:
        print(f"Failed to download image, using fallback background. Error: {e}")
        # Fallback: a dark grey box
        base_img = Image.new("RGBA", (WIDTH, HEIGHT), (40, 40, 45, 255))
    
    # Save base image to a buffer
    bg_buffer = io.BytesIO()
    base_img.convert("RGB").save(bg_buffer, format="JPEG", quality=90)
    bg_buffer.seek(0)

    # 3. Generate the Seamless Gradient Mask using PIL
    # We create a small mask (e.g., 200x100) and resize it for performance and ultra-smooth gradients
    mask_w, mask_h = 200, 100
    mask_base = Image.new("RGBA", (mask_w, mask_h))
    
    r, g, b = base_color
    
    # Left 40% is solid, 40% to 75% is gradient fade, right 25% is transparent
    solid_point = int(mask_w * 0.40)
    fade_point = int(mask_w * 0.75)
    
    for x in range(mask_w):
        if x <= solid_point:
            alpha = 255
        elif x <= fade_point:
            # Linear fade from 255 down to 0
            ratio = (x - solid_point) / (fade_point - solid_point)
            alpha = int(255 * (1 - ratio))
        else:
            alpha = 0
            
        # Paint the column
        for y in range(mask_h):
            mask_base.putpixel((x, y), (r, g, b, alpha))
            
    # Resize to full 1080p for smooth blending
    final_mask = mask_base.resize((WIDTH, HEIGHT), Image.Resampling.LANCZOS)
    mask_buffer = io.BytesIO()
    final_mask.save(mask_buffer, format="PNG")
    mask_buffer.seek(0)

    # 4. Insert Images into Slide
    # Add Base Picture
    slide.shapes.add_picture(bg_buffer, 0, 0, prs.slide_width, prs.slide_height)
    # Add Gradient Mask exactly over it
    slide.shapes.add_picture(mask_buffer, 0, 0, prs.slide_width, prs.slide_height)

    # 5. Add Typography (Left Side, in the solid colored zone)
    # Title
    tx_box_title = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.5), Inches(1.5))
    tf_title = tx_box_title.text_frame
    tf_title.clear()
    p = tf_title.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(64)
    run.font.bold = True
    run.font.name = "Georgia" # Serif for elegance
    run.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Accent text
    tx_box_sub = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.5), Inches(0.8))
    tf_sub = tx_box_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    run_sub = p_sub.add_run()
    run_sub.text = subtitle_text
    run_sub.font.size = Pt(18)
    run_sub.font.bold = True
    run_sub.font.name = "Arial"
    run_sub.font.color.rgb = RGBColor(212, 175, 55) # Gold accent

    # Divider Line
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(1.0), Inches(3.3), Inches(0.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(212, 175, 55)
    line.line.fill.background()

    # Body Text / Quote
    tx_box_body = slide.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(4.5), Inches(3.0))
    tf_body = tx_box_body.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    run_body = p_body.add_run()
    run_body.text = body_text
    run_body.font.size = Pt(14)
    run_body.font.name = "Arial"
    run_body.font.color.rgb = RGBColor(200, 200, 200) # Soft grey for readability

    # Save Presentation
    prs.save(output_pptx_path)
    return output_pptx_path
