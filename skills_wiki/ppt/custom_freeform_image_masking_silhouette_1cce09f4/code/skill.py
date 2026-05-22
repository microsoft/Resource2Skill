def create_slide(
    output_pptx_path: str,
    title_text: str = "Break The Grid",
    body_text: str = "By masking standard photographs into dynamic, custom polygons, we eliminate unnecessary background noise and create a bespoke, highly modern visual flow.",
    bg_palette: str = "architecture",  
    accent_color: tuple = (0, 120, 212),  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Custom Freeform Image Masking" visual effect.
    Uses PIL to perform a boolean 'intersect' mask with a custom polygon, and lxml for alpha-shadows.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw
    from lxml import etree
    import urllib.request
    import io
    import os

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Background Color (Off-white)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 247, 250)

    # 2. Fetch Image (Simulating the user's photo)
    try:
        url = f"https://source.unsplash.com/random/1200x800/?{bg_palette}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        response = urllib.request.urlopen(req, timeout=10)
        img = Image.open(io.BytesIO(response.read())).convert("RGBA")
    except Exception:
        # Fallback if download fails: Create a gradient image
        img = Image.new("RGBA", (1200, 800), (40, 40, 40, 255))
        draw = ImageDraw.Draw(img)
        for y in range(800):
            r = int(40 + (y / 800) * 100)
            g = int(40 + (y / 800) * 150)
            b = int(100 + (y / 800) * 155)
            draw.line([(0, y), (1200, y)], fill=(r, g, b, 255))

    # 3. Create the Custom Polygon Mask (The "Freeform Shape")
    # This simulates drawing a custom shape over the image to intersect it
    width, height = img.size
    
    # Define a dynamic, irregular polygonal shape (like an angled sleek crop)
    custom_shape_points = [
        (width * 0.1, 0),                 # Top slightly inset
        (width * 0.95, height * 0.05),    # Top right, angled down
        (width * 0.8, height * 0.95),     # Bottom right, angled in
        (0, height * 0.85),               # Bottom left, angled up
        (width * 0.05, height * 0.3)      # Mid left, indent
    ]

    # Create an empty alpha mask
    mask = Image.new("L", (width, height), 0)
    mask_draw = ImageDraw.Draw(mask)
    
    # Draw the custom polygon onto the mask
    mask_draw.polygon(custom_shape_points, fill=255)

    # Apply the mask to the original image (The "Intersect" action)
    img.putalpha(mask)

    # Save the custom-masked image
    temp_img_path = "temp_custom_mask.png"
    img.save(temp_img_path, "PNG")

    # 4. Insert the Masked Image into PowerPoint
    pic_left = Inches(1.0)
    pic_top = Inches(1.0)
    pic_width = Inches(5.5)
    
    pic = slide.shapes.add_picture(temp_img_path, pic_left, pic_top, width=pic_width)

    # 5. Add Custom Drop Shadow via lxml (Matches edge of the transparent mask)
    # This elevates the image and proves it is an isolated custom shape
    spPr = pic.element.xpath('.//p:spPr')[0]
    shadow_xml = """
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:outerShdw blurRad="150000" dist="80000" dir="2700000" algn="bl" rotWithShape="0">
            <a:srgbClr val="000000">
                <a:alpha val="20000"/>
            </a:srgbClr>
        </a:outerShdw>
    </a:effectLst>
    """
    spPr.append(etree.fromstring(shadow_xml))

    # 6. Add Typography balancing the shape
    # Headline
    tx_box = slide.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.5), Inches(1.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text.upper()
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42) # Dark Navy
    p.font.name = "Arial Black"

    # Decorative Line
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(7.1), Inches(3.6), Inches(1.0), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    line.line.fill.background()

    # Body Text
    body_box = slide.shapes.add_textbox(Inches(7.0), Inches(3.9), Inches(5.0), Inches(2.5))
    btf = body_box.text_frame
    btf.word_wrap = True
    bp = btf.add_paragraph()
    bp.text = body_text
    bp.font.size = Pt(18)
    bp.font.color.rgb = RGBColor(71, 85, 105) # Slate Gray
    bp.line_spacing = 1.4

    # Cleanup temp image
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)

    # Save
    prs.save(output_pptx_path)
    return output_pptx_path
