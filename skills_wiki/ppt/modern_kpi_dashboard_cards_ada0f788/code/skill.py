def create_slide(
    output_pptx_path: str,
    title_text: str = "Key Metrics",
    metrics: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a modern KPI dashboard card layout.

    Args:
        output_pptx_path: The path to save the generated PPTX file.
        title_text: The main title for the slide.
        metrics: A list of dictionaries, where each dictionary defines a KPI card.
                 Example: {'title': 'New Customers', 'value': '500', 'icon_url': '...'}

    Returns:
        Path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    import urllib.request
    import io
    from PIL import Image, ImageDraw

    # --- Helper function to get icons ---
    def get_icon_image(url: str, fallback_color: tuple = (220, 50, 50)):
        """Downloads an icon or creates a fallback PIL image."""
        try:
            with urllib.request.urlopen(url) as response:
                image_data = response.read()
                return io.BytesIO(image_data)
        except Exception:
            # Create a simple fallback shape if download fails
            img = Image.new('RGBA', (200, 200), (0, 0, 0, 0))
            draw = ImageDraw.Draw(img)
            draw.rectangle((50, 50, 150, 150), fill=fallback_color)
            byte_io = io.BytesIO()
            img.save(byte_io, format='PNG')
            byte_io.seek(0)
            return byte_io

    # --- Default data if none provided ---
    if metrics is None:
        metrics = [
            {
                'title': 'New Customers per month', 'value': '500',
                'icon_url': 'https://icon-icons.com/icons2/2248/PNG/512/group_users_icon_138862.png',
                'icon_color': (220, 50, 50)
            },
            {
                'title': 'On-time Delivery', 'value': '95%',
                'icon_url': 'https://icon-icons.com/icons2/238/PNG/256/delivery-truck_26830.png',
                'icon_color': (22, 160, 133)
            },
            {
                'title': 'Customer Satisfaction', 'value': '90%',
                'icon_url': 'https://icon-icons.com/icons2/933/PNG/512/like-symbol_icon-icons.com_72317.png',
                'icon_color': (241, 196, 15)
            },
            {
                'title': 'Lead Conversion Rate', 'value': '30%',
                'icon_url': 'https://icon-icons.com/icons2/2641/PNG/512/funnel_icon_159151.png',
                'icon_color': (52, 73, 94)
            },
            {
                'title': 'Customer Retention Rate', 'value': '80%',
                'icon_url': 'https://icon-icons.com/icons2/2242/PNG/512/save_user_customer_retention_icon_134763.png',
                'icon_color': (211, 84, 0)
            }
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background (White) ===
    # Slide background is white by default.

    # === Layer 2: Content ===
    # --- Title ---
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(5), Inches(0.75))
    title_tf = title_shape.text_frame
    p = title_tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(64, 64, 64)

    # --- KPI Cards ---
    num_metrics = len(metrics)
    total_width = prs.slide_width - Inches(1) # Total available width with margins
    padding = Inches(0.15)
    card_width = (total_width - (num_metrics - 1) * padding) / num_metrics
    start_left = Inches(0.5)

    for i, metric in enumerate(metrics):
        card_left = start_left + i * (card_width + padding)
        
        # --- Card Background ---
        if i % 2 != 0: # Alternating color
            background_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, Inches(1.2), card_width, Inches(5.5))
            background_shape.fill.solid()
            background_shape.fill.fore_color.rgb = RGBColor(242, 242, 242)
            background_shape.line.fill.background()

        # --- Icon ---
        icon_size = Inches(1.2)
        icon_left = card_left + (card_width - icon_size) / 2
        icon_top = Inches(1.8)
        icon_image_stream = get_icon_image(metric['icon_url'], metric.get('icon_color', (220, 50, 50)))
        slide.shapes.add_picture(icon_image_stream, icon_left, icon_top, height=icon_size)

        # --- Metric Value ---
        value_box = slide.shapes.add_textbox(card_left, Inches(3.2), card_width, Inches(1.5))
        value_tf = value_box.text_frame
        value_tf.word_wrap = False
        p_val = value_tf.paragraphs[0]
        p_val.text = metric['value']
        p_val.font.name = 'Calibri'
        p_val.font.size = Pt(72)
        p_val.font.bold = True
        p_val.font.color.rgb = RGBColor(64, 64, 64)
        p_val.alignment = PP_ALIGN.CENTER

        # --- Metric Title ---
        title_box = slide.shapes.add_textbox(card_left + Inches(0.1), Inches(4.7), card_width - Inches(0.2), Inches(0.8))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        p_title = title_tf.paragraphs[0]
        p_title.text = metric['title']
        p_title.font.name = 'Calibri'
        p_title.font.size = Pt(18)
        p_title.font.color.rgb = RGBColor(64, 64, 64)
        p_title.alignment = PP_ALIGN.CENTER

    # --- Decorative Sticky Note ---
    note_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), Inches(0.2), Inches(3), Inches(1.5))
    note_shape.fill.solid()
    note_shape.fill.fore_color.rgb = RGBColor(255, 242, 204) # Light yellow
    note_shape.line.fill.background()
    note_tf = note_shape.text_frame
    note_tf.margin_left = Inches(0.1)
    note_tf.margin_right = Inches(0.1)
    note_p = note_tf.paragraphs[0]
    note_p.text = "Monitor the performance on the basis of below mentioned parameters."
    note_p.font.name = 'Calibri'
    note_p.font.size = Pt(12)
    note_p.font.color.rgb = RGBColor(64, 64, 64)

    prs.save(output_pptx_path)
    return output_pptx_path
