def create_slide(
    output_pptx_path: str,
    title_text: str = "Joe Zeplin",
    body_text: str = "SOCIAL MEDIA MANAGER",
    bg_palette: str = "dark",  
    accent_color: tuple = (147, 112, 219),  # Medium Purple
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Mobile-Optimized Digital Business Card effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw
    import random

    # 1. Setup Presentation (Vertical 9:16 aspect ratio)
    prs = Presentation()
    # Standard widescreen is 13.333 x 7.5. Let's make it vertical (e.g., 5.625 x 10)
    prs.slide_width = Inches(5.625)
    prs.slide_height = Inches(10.0)
    
    # Use blank layout
    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)

    # 2. Background Generation
    # Add a dark solid background
    bg_color = RGBColor(30, 35, 45)
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg_color
    bg_shape.line.fill.background() # No line

    # Add a subtle accent shape in the background (geometric depth)
    accent_bg = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, Inches(-1), Inches(-1), Inches(4), Inches(4)
    )
    accent_bg.fill.solid()
    accent_bg.fill.fore_color.rgb = RGBColor(40, 46, 60)
    accent_bg.line.fill.background()
    
    accent_bg2 = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, Inches(3.5), Inches(7), Inches(5), Inches(5)
    )
    accent_bg2.fill.solid()
    accent_bg2.fill.fore_color.rgb = RGBColor(40, 46, 60)
    accent_bg2.line.fill.background()

    # 3. Process Profile Picture using PIL (Circular crop + border)
    avatar_url = "https://images.unsplash.com/photo-1560250097-0b93528c311a?q=80&w=400&auto=format&fit=crop"
    try:
        req = urllib.request.Request(avatar_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            avatar_img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception:
        # Fallback: create a dummy colored square if download fails
        avatar_img = Image.new("RGBA", (400, 400), (100, 150, 200, 255))

    # Make it square
    min_dim = min(avatar_img.size)
    left = (avatar_img.width - min_dim)/2
    top = (avatar_img.height - min_dim)/2
    avatar_img = avatar_img.crop((left, top, left+min_dim, top+min_dim))
    avatar_img = avatar_img.resize((400, 400), Image.Resampling.LANCZOS)

    # Create circular mask
    mask = Image.new('L', (400, 400), 0)
    draw = ImageDraw.Draw(mask)
    draw.ellipse((10, 10, 390, 390), fill=255)
    
    # Apply mask
    circular_avatar = Image.new('RGBA', (400, 400), (0, 0, 0, 0))
    circular_avatar.paste(avatar_img, (0, 0), mask)
    
    # Draw White Border Ring
    ring_draw = ImageDraw.Draw(circular_avatar)
    ring_draw.ellipse((10, 10, 390, 390), outline=(255, 255, 255, 255), width=15)

    avatar_stream = BytesIO()
    circular_avatar.save(avatar_stream, format='PNG')
    avatar_stream.seek(0)

    # Insert Avatar into slide
    avatar_size = Inches(2.5)
    avatar_left = (prs.slide_width - avatar_size) / 2
    avatar_top = Inches(1.2)
    slide.shapes.add_picture(avatar_stream, avatar_left, avatar_top, avatar_size, avatar_size)

    # 4. Typography (Name, Title, Contact Info)
    # Name
    name_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.9), Inches(4.625), Inches(0.6))
    name_tf = name_box.text_frame
    name_tf.clear()
    p = name_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title_text.upper()
    run.font.bold = True
    run.font.size = Pt(24)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(255, 255, 255)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.4), Inches(4.625), Inches(0.4))
    title_tf = title_box.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = body_text.upper()
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])

    # Divider Line
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2), Inches(4.9), Inches(1.625), Inches(0.02)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    divider.line.fill.background()

    # Contact Info
    contact_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(4.625), Inches(1.0))
    contact_tf = contact_box.text_frame
    contact_tf.clear()
    
    contacts = ["123-456-7890", "hello@reallygreatsite.com", "www.reallygreatsite.com"]
    for i, text in enumerate(contacts):
        p = contact_tf.add_paragraph() if i > 0 else contact_tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(11)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(220, 220, 220)

    # 5. Simulated QR Code Generation using PIL
    # Generates a convincing fake QR code to avoid external dependencies
    qr_size = 300
    qr_img = Image.new("RGB", (qr_size, qr_size), "white")
    qr_draw = ImageDraw.Draw(qr_img)
    
    # Draw position markers (the 3 large squares in the corners)
    def draw_marker(x, y):
        ms = 60 # marker size
        qr_draw.rectangle([x, y, x+ms, y+ms], fill="black")
        qr_draw.rectangle([x+10, y+10, x+ms-10, y+ms-10], fill="white")
        qr_draw.rectangle([x+20, y+20, x+ms-20, y+ms-20], fill="black")

    draw_marker(20, 20)           # Top Left
    draw_marker(qr_size-80, 20)   # Top Right
    draw_marker(20, qr_size-80)   # Bottom Left
    
    # Draw random data modules
    random.seed(42) # Fixed seed for reproducible layout
    grid_steps = 20
    step = qr_size // grid_steps
    for i in range(grid_steps):
        for j in range(grid_steps):
            # Skip marker areas
            if (i < 6 and j < 6) or (i > grid_steps-7 and j < 6) or (i < 6 and j > grid_steps-7):
                continue
            if random.random() > 0.5:
                qr_draw.rectangle([i*step, j*step, (i+1)*step, (j+1)*step], fill="black")

    # Add a small padding ring
    qr_padded = Image.new("RGB", (qr_size + 40, qr_size + 40), "white")
    qr_padded.paste(qr_img, (20, 20))

    qr_stream = BytesIO()
    qr_padded.save(qr_stream, format="PNG")
    qr_stream.seek(0)

    # Insert QR Code into slide
    qr_display_size = Inches(2.2)
    qr_left = (prs.slide_width - qr_display_size) / 2
    qr_top = Inches(6.8)
    
    slide.shapes.add_picture(qr_stream, qr_left, qr_top, qr_display_size, qr_display_size)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
