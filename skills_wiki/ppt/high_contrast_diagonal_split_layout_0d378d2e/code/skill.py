def create_slide(
    output_pptx_path: str,
    title_text: str = "V K COMPUTER",
    subtitle_text: str = "Graphic Designer",
    accent_color: tuple = (220, 20, 40),  # Vibrant Red
    dark_bg_color: tuple = (20, 20, 20),  # Deep Black
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "High-Contrast Diagonal Split Layout".
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    # Initialize Presentation (16:9 standard)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Define Colors
    c_accent = RGBColor(*accent_color)
    c_dark = RGBColor(*dark_bg_color)
    c_white = RGBColor(255, 255, 255)
    c_black = RGBColor(0, 0, 0)

    # === Layer 1: Background Diagonal Split ===
    
    # 1a. Dark Base Shape (Left Side Slant)
    ff_builder_dark = slide.shapes.build_freeform(Inches(0), Inches(0))
    ff_builder_dark.add_line_segments([
        (Inches(5.8), Inches(0)),     # Top right (starts the slant)
        (Inches(4.0), Inches(7.5)),   # Bottom right (ends the slant)
        (Inches(0), Inches(7.5)),     # Bottom left
        (Inches(0), Inches(0))        # Back to Top left
    ])
    shape_dark = ff_builder_dark.convert_to_shape()
    shape_dark.fill.solid()
    shape_dark.fill.fore_color.rgb = c_dark
    shape_dark.line.fill.background() # No outline

    # 1b. Accent Band Shape (Red Slant)
    ff_builder_accent = slide.shapes.build_freeform(Inches(5.8), Inches(0))
    ff_builder_accent.add_line_segments([
        (Inches(6.8), Inches(0)),     # Top right of red band
        (Inches(5.0), Inches(7.5)),   # Bottom right of red band
        (Inches(4.0), Inches(7.5)),   # Bottom left of red band (touches dark shape)
        (Inches(5.8), Inches(0))      # Back to Top left of red band
    ])
    shape_accent = ff_builder_accent.convert_to_shape()
    shape_accent.fill.solid()
    shape_accent.fill.fore_color.rgb = c_accent
    shape_accent.line.fill.background()

    # === Layer 2: Decorative Geometric Element ===
    
    # Hexagon
    hex_size = Inches(1.8)
    hexagon = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, 
        Inches(2.2), Inches(1.2), 
        hex_size, hex_size
    )
    hexagon.fill.solid()
    hexagon.fill.fore_color.rgb = c_accent
    hexagon.line.color.rgb = c_white
    hexagon.line.width = Pt(5)

    # === Layer 3: Left Side Content (Brand) ===
    
    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(4.0), Inches(1.0))
    tf_title = tx_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.bold = True
    p_title.font.size = Pt(44)
    p_title.font.color.rgb = c_white

    # Subtitle
    tx_sub = slide.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(4.0), Inches(0.5))
    p_sub = tx_sub.text_frame.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(24)
    p_sub.font.color.rgb = RGBColor(200, 200, 200)

    # Horizontal Divider Line
    div_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(1.2), Inches(5.2), Inches(2.5), Inches(0.08)
    )
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = c_accent
    div_line.line.fill.background()

    # === Layer 4: Right Side Content (Contact Info) ===
    
    contact_data = [
        {"icon": "📞", "text": "+1 (555) 123-4567\n+1 (555) 987-6543"},
        {"icon": "🌐", "text": "www.yourwebsite.com\nportfolio.design.com"},
        {"icon": "✉", "text": "contact@yourwebsite.com\nhello@design.com"}
    ]

    start_y = 3.0
    spacing_y = 1.3

    for index, item in enumerate(contact_data):
        current_y = start_y + (index * spacing_y)
        
        # Red Icon Background Box
        icon_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(7.5), Inches(current_y), 
            Inches(0.6), Inches(0.6)
        )
        icon_bg.fill.solid()
        icon_bg.fill.fore_color.rgb = c_accent
        icon_bg.line.fill.background()
        
        # Icon Text (Using Unicode symbols as a proxy for Wingdings)
        tx_icon = slide.shapes.add_textbox(
            Inches(7.5), Inches(current_y - 0.05), 
            Inches(0.6), Inches(0.6)
        )
        p_icon = tx_icon.text_frame.paragraphs[0]
        p_icon.text = item["icon"]
        p_icon.font.size = Pt(28)
        p_icon.font.color.rgb = c_white
        p_icon.alignment = PP_ALIGN.CENTER
        tx_icon.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Contact Info Text
        tx_info = slide.shapes.add_textbox(
            Inches(8.3), Inches(current_y), 
            Inches(4.0), Inches(0.8)
        )
        p_info = tx_info.text_frame.paragraphs[0]
        p_info.text = item["text"]
        p_info.font.size = Pt(18)
        p_info.font.color.rgb = c_black

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
