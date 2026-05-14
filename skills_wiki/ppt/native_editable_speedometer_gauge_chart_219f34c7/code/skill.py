def create_slide(
    output_pptx_path: str,
    title_text: str = "Performance Dashboard",
    score: float = 72.5,  # 0 to 100
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Native Editable Speedometer visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import math
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Colors
    c_red = RGBColor(226, 62, 87)
    c_yellow = RGBColor(244, 164, 44)
    c_green = RGBColor(81, 152, 114)
    c_dark = RGBColor(43, 43, 43)
    c_white = RGBColor(255, 255, 255)

    # Gauge Geometry parameters
    cx, cy = Inches(6.66), Inches(5.0)  # Center point
    r_outer = Inches(3.5)
    r_inner = Inches(2.2)

    # --- Helper Function to create an Arc Polygon ---
    def add_arc_zone(start_deg, end_deg, fill_color):
        steps = 40
        points = []
        # Draw outer curve
        for i in range(steps + 1):
            t = math.radians(start_deg + (end_deg - start_deg) * i / steps)
            # In screen coords, Y goes down, so we subtract to go UP
            x = cx + r_outer * math.cos(t)
            y = cy - r_outer * math.sin(t)
            points.append((x, y))
        # Draw inner curve (reverse direction)
        for i in range(steps + 1):
            t = math.radians(end_deg + (start_deg - end_deg) * i / steps)
            x = cx + r_inner * math.cos(t)
            y = cy - r_inner * math.sin(t)
            points.append((x, y))
            
        ff_builder = slide.shapes.build_freeform()
        ff_builder.add_line_segments(points, close=True)
        shape = ff_builder.convert_to_shape()
        
        # Style
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.color.rgb = c_white
        shape.line.width = Pt(3)
        return shape

    # === Layer 1: Draw the 3 Dial Zones ===
    # Math angles: 180 is left, 0 is right.
    # We want 3 equal zones of 60 degrees.
    add_arc_zone(180, 120, c_red)    # Poor
    add_arc_zone(120, 60, c_yellow)  # OK
    add_arc_zone(60, 0, c_green)     # Good

    # === Layer 2: Draw the Needle ===
    # Map score (0-100) to angle (180 to 0 degrees)
    score_clamped = max(0, min(100, score))
    needle_deg = 180 - (score_clamped / 100.0) * 180
    t = math.radians(needle_deg)
    
    needle_length = r_outer - Inches(0.2)
    base_r = Inches(0.15)  # Width of the needle base

    # Calculate 3 points of the isosceles triangle
    tip_x = cx + needle_length * math.cos(t)
    tip_y = cy - needle_length * math.sin(t)
    
    left_x = cx + base_r * math.cos(t + math.pi/2)
    left_y = cy - base_r * math.sin(t + math.pi/2)
    
    right_x = cx + base_r * math.cos(t - math.pi/2)
    right_y = cy - base_r * math.sin(t - math.pi/2)

    ff_needle = slide.shapes.build_freeform()
    ff_needle.add_line_segments([(tip_x, tip_y), (left_x, left_y), (right_x, right_y)], close=True)
    needle = ff_needle.convert_to_shape()
    needle.fill.solid()
    needle.fill.fore_color.rgb = c_dark
    needle.line.color.rgb = c_white
    needle.line.width = Pt(1.5)

    # === Layer 3: Center Pivot Cap ===
    cap_r = Inches(0.3)
    cap = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        cx - cap_r, cy - cap_r, 
        cap_r * 2, cap_r * 2
    )
    cap.fill.solid()
    cap.fill.fore_color.rgb = c_red
    cap.line.color.rgb = c_white
    cap.line.width = Pt(2)

    # === Layer 4: Text Labels ===
    # Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = c_dark
    p.alignment = PP_ALIGN.CENTER

    # Value Label (Bottom Center)
    val_box = slide.shapes.add_textbox(cx - Inches(1.5), cy + Inches(0.2), Inches(3), Inches(1))
    v_tf = val_box.text_frame
    v_p = v_tf.paragraphs[0]
    v_p.text = f"{score_clamped:.1f}%"
    v_p.font.size = Pt(36)
    v_p.font.bold = True
    v_p.font.color.rgb = c_dark
    v_p.alignment = PP_ALIGN.CENTER

    # Zone Labels
    def add_zone_label(text, x, y):
        tb = slide.shapes.add_textbox(x, y, Inches(1.5), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(120, 120, 120)
        p.alignment = PP_ALIGN.CENTER

    add_zone_label("POOR", cx - Inches(3.2), cy - Inches(0.5))
    add_zone_label("OK", cx - Inches(0.75), cy - Inches(3.0))
    add_zone_label("GOOD", cx + Inches(1.7), cy - Inches(0.5))

    prs.save(output_pptx_path)
    return output_pptx_path
