import os
import io
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.chart import XL_DATA_LABEL_POSITION
from PIL import Image, ImageDraw

def _create_pil_overlay(width_px, height_px, base_color, alpha_start, alpha_end, image_url=None):
    """
    Creates a PIL image with an optional background photo and a gradient overlay.
    """
    base_img = Image.new('RGBA', (width_px, height_px), (30, 30, 30, 255))
    
    if image_url:
        try:
            response = requests.get(image_url, timeout=5)
            if response.status_code == 200:
                downloaded_img = Image.open(io.BytesIO(response.content)).convert('RGBA')
                # Resize and crop to fit
                aspect_ratio = width_px / height_px
                img_aspect = downloaded_img.width / downloaded_img.height
                if img_aspect > aspect_ratio:
                    new_w = int(downloaded_img.height * aspect_ratio)
                    offset = (downloaded_img.width - new_w) // 2
                    downloaded_img = downloaded_img.crop((offset, 0, offset + new_w, downloaded_img.height))
                else:
                    new_h = int(downloaded_img.width / aspect_ratio)
                    offset = (downloaded_img.height - new_h) // 2
                    downloaded_img = downloaded_img.crop((0, offset, downloaded_img.width, offset + new_h))
                
                base_img = downloaded_img.resize((width_px, height_px), Image.Resampling.LANCZOS)
        except Exception as e:
            print(f"Failed to fetch image, using fallback. Error: {e}")

    # Create gradient overlay
    overlay = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    
    r, g, b = base_color
    for y in range(height_px):
        # Calculate alpha based on vertical position
        alpha = int(alpha_start + (alpha_end - alpha_start) * (y / height_px))
        draw.line([(0, y), (width_px, y)], fill=(r, g, b, alpha))
        
    final_img = Image.alpha_composite(base_img, overlay)
    
    # Save to bytes
    img_byte_arr = io.BytesIO()
    final_img.save(img_byte_arr, format='PNG')
    img_byte_arr.seek(0)
    return img_byte_arr

def create_slide(
    output_pptx_path: str,
    title_text: str = "BUSINESS PRESENTATION",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa. Fusce posuere, magna sed pulvinar ultricies, purus lectus malesuada libero, sit amet commodo magna eros quis urna.",
    accent_color: tuple = (0, 212, 255),  # Cyan
    cta_color: tuple = (255, 183, 3),     # Yellow
    **kwargs,
) -> str:
    
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Colors
    bg_dark = RGBColor(11, 19, 30)
    text_white = RGBColor(255, 255, 255)
    text_gray = RGBColor(160, 175, 190)
    cyan_rgb = RGBColor(*accent_color)
    yellow_rgb = RGBColor(*cta_color)
    chart_bg = RGBColor(30, 40, 50)

    # --- LAYER 1: Generate & Insert Background Overlays via PIL ---
    
    # Left Pane (30% width) - Portrait image with Teal overlay
    left_w_px, left_h_px = int(13.333 * 0.3 * 100), int(7.5 * 100)
    left_img_stream = _create_pil_overlay(
        left_w_px, left_h_px, 
        base_color=(0, 100, 120), 
        alpha_start=180, alpha_end=230,
        image_url="https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?w=800&q=80" # Portrait
    )
    slide.shapes.add_picture(left_img_stream, Inches(0), Inches(0), width=Inches(13.333 * 0.3), height=Inches(7.5))

    # Right Pane (70% width) - City image with Dark Navy overlay
    right_w_px, right_h_px = int(13.333 * 0.7 * 100), int(7.5 * 100)
    right_img_stream = _create_pil_overlay(
        right_w_px, right_h_px, 
        base_color=(11, 19, 30), 
        alpha_start=200, alpha_end=250,
        image_url="https://images.unsplash.com/photo-1449844908441-8829872d2607?w=1600&q=80" # City
    )
    slide.shapes.add_picture(right_img_stream, Inches(13.333 * 0.3), Inches(0), width=Inches(13.333 * 0.7), height=Inches(7.5))


    # --- LAYER 2: Left Pane Content ---
    
    # Left Heading
    tx_left_head = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(3), Inches(0.5))
    p = tx_left_head.text_frame.add_paragraph()
    p.text = "HEADING"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = text_white

    # Left Body
    tx_left_body = slide.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(3), Inches(1.5))
    tx_left_body.text_frame.word_wrap = True
    p = tx_left_body.text_frame.add_paragraph()
    p.text = "Pellentesque habitant morbi tristique senectus et netus et malesuada fames ac turpis."
    p.font.size = Pt(12)
    p.font.color.rgb = text_white

    # CTA Button
    btn = slide.shapes.add_shape(
        1, # msoShapeRectangle (Rounded would be 5, but simple rectangle works here)
        Inches(0.5), Inches(6.2), Inches(2), Inches(0.5)
    )
    btn.fill.solid()
    btn.fill.fore_color.rgb = yellow_rgb
    btn.line.fill.background()
    
    p = btn.text_frame.paragraphs[0]
    p.text = "LEARN MORE"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = bg_dark
    p.alignment = PP_ALIGN.CENTER


    # --- LAYER 3: Right Pane Content (Headers) ---
    
    # Super title
    tx_super = slide.shapes.add_textbox(Inches(4.5), Inches(1.0), Inches(8), Inches(0.4))
    p = tx_super.text_frame.add_paragraph()
    p.text = "CREATIVE VENUS"
    p.font.size = Pt(14)
    p.font.color.rgb = text_white
    
    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(4.4), Inches(1.3), Inches(8), Inches(1.0))
    p = tx_title.text_frame.add_paragraph()
    p.text = title_text
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = text_white

    # Body Text
    tx_body = slide.shapes.add_textbox(Inches(4.5), Inches(2.4), Inches(8), Inches(1.5))
    tx_body.text_frame.word_wrap = True
    p = tx_body.text_frame.add_paragraph()
    p.text = body_text
    p.font.size = Pt(14)
    p.font.color.rgb = text_gray


    # --- LAYER 4: Doughnut Charts ---
    chart_data_list = [
        {"name": "PARAMETER A", "val": 72},
        {"name": "PARAMETER B", "val": 79},
        {"name": "PARAMETER C", "val": 72}
    ]
    
    start_x = 4.5
    for i, data in enumerate(chart_data_list):
        cx = Inches(start_x + (i * 2.8))
        cy = Inches(4.5)
        cw = Inches(2.2)
        ch = Inches(2.2)

        # Chart Data
        chart_data = ChartData()
        chart_data.categories = ['Achieved', 'Remaining']
        chart_data.add_series('Series 1', (data["val"], 100 - data["val"]))

        # Add Chart
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, cx, cy, cw, ch, chart_data
        )
        chart = chart_shape.chart
        
        # Transparent Chart Background & Border
        chart_shape.fill.background()
        chart_shape.line.fill.background()

        # Format series colors
        series = chart.series[0]
        # First slice (Cyan)
        pt0 = series.points[0]
        pt0.format.fill.solid()
        pt0.format.fill.fore_color.rgb = cyan_rgb
        pt0.format.line.fill.background()
        
        # Second slice (Dark Track)
        pt1 = series.points[1]
        pt1.format.fill.solid()
        pt1.format.fill.fore_color.rgb = chart_bg
        pt1.format.line.fill.background()

        # Chart Label (Top)
        lbl = slide.shapes.add_textbox(cx, Inches(4.0), cw, Inches(0.4))
        p = lbl.text_frame.add_paragraph()
        p.text = data["name"]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = text_white
        p.alignment = PP_ALIGN.CENTER
        
        # Center Percentage Text
        pct = slide.shapes.add_textbox(cx, Inches(5.3), cw, Inches(0.6))
        p = pct.text_frame.add_paragraph()
        p.text = f"{data['val']}%"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = text_white
        p.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
