import io
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFont

def _create_circular_badge_stream(text: str, radius: int = 120, font_size: int = 24, color: tuple = (250, 246, 238, 255)) -> io.BytesIO:
    """
    Generates a PNG image of text wrapped in a circle using PIL.
    """
    size = radius * 2 + font_size * 3
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    
    # Attempt to load a smooth TTF font, fallback to default if unavailable
    try:
        font = ImageFont.truetype("arial.ttf", font_size)
    except IOError:
        try:
            font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", font_size) # MacOS fallback
        except IOError:
            font = ImageFont.load_default()

    total_chars = len(text)
    angle_step = 360 / total_chars
    center_x, center_y = size / 2, size / 2

    for i, char in enumerate(text):
        angle_deg = i * angle_step
        angle_rad = math.radians(angle_deg)

        # Create temporary canvas for the character to allow clean rotation
        char_canvas_size = font_size * 3
        char_img = Image.new('RGBA', (char_canvas_size, char_canvas_size), (0, 0, 0, 0))
        char_draw = ImageDraw.Draw(char_img)
        
        # Draw character centered in temp canvas
        bbox = char_draw.textbbox((0, 0), char, font=font)
        w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
        char_draw.text(((char_canvas_size - w) / 2, (char_canvas_size - h) / 2), char, font=font, fill=color)

        # Rotate character. -angle_deg to match circle direction, -90 to face outward
        rotated_char = char_img.rotate(-angle_deg - 90, resample=Image.BICUBIC, expand=True)
        
        # Calculate placement on the circle's circumference
        x = center_x + radius * math.cos(angle_rad)
        y = center_y + radius * math.sin(angle_rad)
        
        rw, rh = rotated_char.size
        img.paste(rotated_char, (int(x - rw / 2), int(y - rh / 2)), rotated_char)

    image_stream = io.BytesIO()
    img.save(image_stream, format='PNG')
    image_stream.seek(0)
    return image_stream

def _create_mock_product_stream(width: int = 400, height: int = 500) -> io.BytesIO:
    """
    Generates a placeholder "product cutout" with a drop shadow using PIL.
    Simulates a stylized 3D capsule/bottle.
    """
    img = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # 1. Drop Shadow
    shadow_rect = [width * 0.15, height * 0.85, width * 0.85, height * 0.98]
    draw.ellipse(shadow_rect, fill=(0, 0, 0, 60))
    
    # 2. Main Body (Capsule shape)
    body_rect = [width * 0.25, height * 0.1, width * 0.75, height * 0.9]
    draw.rounded_rectangle(body_rect, radius=width * 0.25, fill=(240, 235, 225, 255))
    
    # 3. Decorative / Branding lines on product
    draw.line([width * 0.25, height * 0.3, width * 0.75, height * 0.3], fill=(40, 40, 40, 255), width=8)
    draw.line([width * 0.25, height * 0.7, width * 0.75, height * 0.7], fill=(40, 40, 40, 255), width=8)
    
    # 4. Highlight for faux 3D volume
    highlight_rect = [width * 0.3, height * 0.15, width * 0.4, height * 0.85]
    draw.rounded_rectangle(highlight_rect, radius=width * 0.05, fill=(255, 255, 255, 150))

    image_stream = io.BytesIO()
    img.save(image_stream, format='PNG')
    image_stream.seek(0)
    return image_stream

def create_slide(
    output_pptx_path: str,
    title_text: str = "CRESCENTE",
    badge_text: str = "TYPICAL SICILIAN PRODUCTS • 100% ORGANIC • ",
    bg_color: tuple = (238, 93, 56),        # Vibrant Orange
    text_color: tuple = (250, 246, 238),    # Off-white
    accent_color: tuple = (20, 60, 40),     # Deep Green for accents
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Editorial Product Hero & Typography Rings' effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Solid Vibrant Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Massive Display Typography ===
    # Adding text behind the product to create depth
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.333), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(140)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(*text_color)

    # === Layer 3: Central Product Cutout ===
    # Generate the transparent product image
    product_stream = _create_mock_product_stream(width=500, height=600)
    # Center the product
    pic_left = Inches((13.333 - 4.5) / 2) # approx 4.5 inches wide
    pic_top = Inches((7.5 - 5.5) / 2)
    slide.shapes.add_picture(product_stream, pic_left, pic_top, width=Inches(4.5))

    # === Layer 4: Circular Typography Badge ===
    # Generate the circular text ring
    badge_stream = _create_circular_badge_stream(text=badge_text, radius=130, font_size=28, color=(*text_color, 255))
    # Position asymmetrically, overlapping the product and background
    badge_left = Inches(8.0)
    badge_top = Inches(4.0)
    slide.shapes.add_picture(badge_stream, badge_left, badge_top, width=Inches(3.5))

    # === Layer 5: Secondary Typography (Editorial touch) ===
    desc_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(4.0), Inches(1.0))
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = "Crafted with passion, transforming traditional specialties into new, high-quality experiences."
    p_desc.font.size = Pt(16)
    p_desc.font.name = "Georgia" # Serif for editorial contrast
    p_desc.font.italic = True
    p_desc.font.color.rgb = RGBColor(*text_color)

    # Add a small accent geometric detail
    accent_shape = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(1.0), Inches(5.8), Inches(0.5), Inches(0.05)
    )
    accent_shape.fill.solid()
    accent_shape.fill.fore_color.rgb = RGBColor(*text_color)
    accent_shape.line.fill.background() # No line

    prs.save(output_pptx_path)
    return output_pptx_path

