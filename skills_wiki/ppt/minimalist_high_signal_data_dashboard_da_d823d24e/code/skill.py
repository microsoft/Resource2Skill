import io
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "Q3 Executive Dashboard",
    body_text: str = "Performance overview sorted by impact, highlighting key metrics.",
    accent_color_hex: str = "#0078D4",  # Microsoft Blue
    alert_color_hex: str = "#D2222D",   # Red for peaks/alerts
    **kwargs,
) -> str:
    """
    Creates a minimalist, Tufte-style data dashboard reproducing the visual rules
    from the tutorial (No chart junk, sorted data, selective color, white space).
    """
    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Colors
    bg_color = RGBColor(250, 250, 250)
    text_color = RGBColor(60, 60, 60)
    muted_text = RGBColor(150, 150, 150)
    border_color = RGBColor(230, 230, 230)
    
    # 1. Slide Background (Very faint gray to make white cards pop)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # Helper: Add Text
    def add_text(slide, text, left, top, width, height, font_size, color, bold=False):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        return txBox

    # Add Titles
    add_text(slide, title_text, Inches(0.5), Inches(0.4), Inches(8), Inches(0.5), 28, text_color, bold=True)
    add_text(slide, body_text, Inches(0.5), Inches(0.9), Inches(8), Inches(0.4), 14, muted_text)

    # =========================================================
    # CHART 1: Minimalist Sorted Bar Chart (Rules: Sort, Color, Junk Removal)
    # =========================================================
    categories = ['Bikes', 'Components', 'Clothing', 'Accessories', 'Services']
    values = [1250000, 850000, 420000, 150000, 50000]
    
    fig, ax = plt.subplots(figsize=(6, 4), dpi=150)
    fig.patch.set_alpha(0.0) # Transparent background
    ax.patch.set_alpha(0.0)
    
    # Colors: Accent color for the top performer, gray for the rest
    colors = [accent_color_hex] + ['#CCCCCC'] * (len(categories) - 1)
    
    # Create horizontal bars (sorted visually top to bottom)
    bars = ax.barh(categories[::-1], values[::-1], color=colors[::-1], height=0.6)
    
    # Remove Chart Junk (Spines, Ticks, Gridlines)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.xaxis.set_visible(False) # Hide x axis completely
    ax.tick_params(axis='y', length=0, labelsize=12, colors='#555555') # Clean y labels
    
    # Direct Data Labels (Instead of axis)
    for bar in bars:
        width = bar.get_width()
        label_x_pos = width + 20000
        ax.text(label_x_pos, bar.get_y() + bar.get_height()/2, f'${width/1000:.0f}K', 
                va='center', ha='left', color='#555555', fontsize=10, fontweight='bold')

    plt.title("Sales by Category (Sorted)", loc='left', fontsize=14, color='#333333', pad=20)
    plt.tight_layout()
    
    # Save to memory
    bar_stream = io.BytesIO()
    plt.savefig(bar_stream, format='png', bbox_inches='tight')
    bar_stream.seek(0)
    plt.close(fig)

    # =========================================================
    # CHART 2: Minimalist Line Trend with Highlight (Rules: Color focus, Reduce noise)
    # =========================================================
    quarters = ['Q1', 'Q2', 'Q3', 'Q4', 'Q1 ', 'Q2 ', 'Q3 ']
    trend = [20, 25, 22, 35, 30, 45, 40]
    
    fig2, ax2 = plt.subplots(figsize=(6, 3), dpi=150)
    fig2.patch.set_alpha(0.0)
    ax2.patch.set_alpha(0.0)
    
    # Plot baseline in gray
    ax2.plot(quarters, trend, color='#CCCCCC', linewidth=3, marker='o', markersize=6)
    
    # Highlight the maximum point
    max_idx = trend.index(max(trend))
    ax2.plot(quarters[max_idx], trend[max_idx], color=alert_color_hex, marker='o', markersize=10)
    ax2.text(max_idx, trend[max_idx] + 3, f'{trend[max_idx]}', color=alert_color_hex, 
             fontsize=12, fontweight='bold', ha='center')

    # Remove Junk
    ax2.spines['top'].set_visible(False)
    ax2.spines['right'].set_visible(False)
    ax2.spines['left'].set_color('#EEEEEE')
    ax2.spines['bottom'].set_color('#EEEEEE')
    ax2.tick_params(axis='both', length=0, colors='#888888')
    ax2.grid(True, axis='y', color='#FAFAFA', linestyle='-', linewidth=1) # Barely visible grid
    
    plt.title("Revenue Trend (Peak Highlighted)", loc='left', fontsize=14, color='#333333', pad=15)
    plt.tight_layout()
    
    line_stream = io.BytesIO()
    plt.savefig(line_stream, format='png', bbox_inches='tight')
    line_stream.seek(0)
    plt.close(fig2)

    # =========================================================
    # Dashboard Layout (Rule: White Space & Subtle Borders)
    # =========================================================
    
    def create_card(slide, left, top, width, height):
        """Creates a subtle white card container using white space principles"""
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        # Adjust corner radius to be subtle
        card.adjustments[0] = 0.05 
        return card

    # Card 1: Bar Chart
    create_card(slide, Inches(0.5), Inches(1.8), Inches(6), Inches(5.2))
    slide.shapes.add_picture(bar_stream, Inches(0.7), Inches(2.3), width=Inches(5.5))

    # Card 2: KPI Area (Top Right)
    create_card(slide, Inches(6.8), Inches(1.8), Inches(6), Inches(1.5))
    add_text(slide, "Avg Days to Ship", Inches(7.0), Inches(2.0), Inches(2), Inches(0.3), 12, muted_text)
    add_text(slide, "1.76", Inches(7.0), Inches(2.3), Inches(2), Inches(0.8), 44, text_color, bold=True)
    
    add_text(slide, "Critical Incidents", Inches(10.0), Inches(2.0), Inches(2), Inches(0.3), 12, muted_text)
    add_text(slide, "0", Inches(10.0), Inches(2.3), Inches(2), Inches(0.8), 44, RGBColor(40, 167, 69), bold=True) # Green

    # Card 3: Line Chart (Bottom Right)
    create_card(slide, Inches(6.8), Inches(3.5), Inches(6), Inches(3.5))
    slide.shapes.add_picture(line_stream, Inches(7.0), Inches(3.8), width=Inches(5.5))

    # Save Presentation
    prs.save(output_pptx_path)
    return output_pptx_path
