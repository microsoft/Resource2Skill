import pandas as pd
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import matplotlib.pyplot as plt
import io

def generate_mock_data():
    """Generates a sample DataFrame mimicking the video's data structure."""
    years = range(2014, 2025)
    departments = {
        "Finance": ["Suja Mohanty", "Suraj Rajput", "Pramod Bhavsar", "Satsh Ojha", "Sintu Kumar"],
        "Operation": ["Harivansh Gautam", "Vini Saini", "Anand Singh Rajput", "Jaishri Saxena", "Virender Sroha"],
        "HR": ["Ekanshika Kalra", "Hemant Gusain", "Sarthak Bhagu", "Gopala Krishna", "Puran Singh"]
    }
    
    data = []
    for dept, employees in departments.items():
        for emp in employees:
            for year in years:
                sales = np.random.randint(200, 950)
                # Ensure target is reasonably close to sales
                target = int(sales * (1 + np.random.uniform(-0.15, 0.25)))
                data.append([dept, emp, year, sales, target])
    
    df = pd.DataFrame(data, columns=["Department", "Employee", "Year", "Sales", "Target"])
    return df

def create_slide(
    output_pptx_path: str,
    department: str = "Finance",
    employee_name: str = "Pramod Bhavsar",
    start_year: int = 2014,
    num_years_to_show: int = 7,
    **kwargs,
) -> str:
    """
    Creates a PPTX slide reproducing the KPI Performance Tracker dashboard.
    
    Note: The interactivity of Excel form controls is simulated. The output is a
    static slide representing the dashboard state for the given parameters.

    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background & Layout Shapes ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Header/Control panel background shape
    header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.4), Inches(12.333), Inches(1.5))
    header_fill = header_shape.fill
    header_fill.solid()
    header_fill.fore_color.rgb = RGBColor(70, 114, 196) # A professional blue
    header_shape.line.fill.background()

    # === Data Preparation (Simulating control selection) ===
    df = generate_mock_data()
    
    # Filter data for the selected employee and year range
    end_year = start_year + num_years_to_show - 1
    employee_data = df[
        (df['Employee'] == employee_name) &
        (df['Department'] == department) &
        (df['Year'].between(start_year, end_year))
    ].copy()

    # === Simulate Controls Visuals ===
    # Department selection visual
    for i, dept_name in enumerate(["Finance", "Operation", "HR"]):
        is_selected = (dept_name == department)
        box_color = RGBColor(0, 176, 80) if is_selected else RGBColor(146, 208, 80)
        
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + i * 2.5), Inches(0.6), Inches(2.2), Inches(0.6))
        box.fill.solid()
        box.fill.fore_color.rgb = box_color
        box.text = dept_name.upper()
        box.text_frame.paragraphs[0].font.bold = True
        box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        box.line.fill.background()

    # Employee selection visual
    emp_label = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(2), Inches(0.5))
    emp_label.text_frame.text = "EMPLOYEE NAME"
    p = emp_label.text_frame.paragraphs[0]
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.bold = True

    selected_emp_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(1.25), Inches(4), Inches(0.5))
    selected_emp_box.fill.solid()
    selected_emp_box.fill.fore_color.rgb = RGBColor(220, 230, 241)
    selected_emp_box.text = employee_name
    p = selected_emp_box.text_frame.paragraphs[0]
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(0, 0, 0)

    # === Layer 2: Chart Generation (matplotlib) ===
    fig, ax = plt.subplots(figsize=(12, 4.5))
    fig.patch.set_alpha(0.0)
    ax.patch.set_alpha(0.0)
    
    bar_width = 0.4
    index = np.arange(len(employee_data['Year']))
    
    # Bar for Sales
    bars = ax.bar(index, employee_data['Sales'], bar_width, label=employee_name, color='#4F81BD')
    # Line for Target
    line = ax.plot(index, employee_data['Target'], label='Target', color='#C0504D', marker='o', linewidth=2.5)
    
    ax.set_title(f"Employee-wise Revenue: {employee_name}", fontsize=18, weight='bold', pad=20)
    ax.set_ylabel('Revenue', fontsize=12)
    ax.set_xticks(index)
    ax.set_xticklabels(employee_data['Year'])
    ax.tick_params(axis='x', colors='grey')
    ax.tick_params(axis='y', colors='grey')
    
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('grey')
    ax.spines['bottom'].set_color('grey')
    ax.yaxis.grid(True, linestyle='--', which='major', color='grey', alpha=.25)

    for bar in bars:
        yval = bar.get_height()
        ax.text(bar.get_x() + bar.get_width()/2.0, yval, int(yval), ha='center', va='bottom', fontsize=9, color='white', weight='bold')

    ax.legend(frameon=False)
    plt.tight_layout()
    
    img_stream = io.BytesIO()
    plt.savefig(img_stream, format='png', dpi=300, transparent=True)
    img_stream.seek(0)
    plt.close(fig)

    slide.shapes.add_picture(img_stream, Inches(0.6), Inches(2.2), width=Inches(12.1))

    prs.save(output_pptx_path)
    return output_pptx_path

