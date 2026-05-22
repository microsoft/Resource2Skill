import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def create_gauge_image(percentage, color_rgb, output_path):
    """
    Generates a high-resolution transparent PNG of a semi-circular gauge.
    Uses Pillow to draw at 4x scale and downsamples for perfect anti-aliasing.
    """
    # 4x resolution for smooth anti-aliasing
    canvas_size = (2000, 1100)
    bbox = [100, 100, 1900, 1900] # Bounding box for a full circle
    line_width = 160
    
    # Create transparent image
    img = Image.new('RGBA', canvas_size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # Draw background track (light gray) from 180 to 360 degrees (top half)
    draw.arc(bbox, start=180, end=360, fill=(235, 235, 235, 255), width=line_width)
    
    # Calculate end angle based on percentage (180 base + (percentage * 180))
    end_angle = 180 + int(percentage * 180)
    
    # Draw foreground colored track
    draw.arc(bbox, start=180, end=end_angle, fill=color_rgb + (255,), width=line_width)
    
    # Downsample for smoothness
    img = img.resize((500, 275), Image.Resampling.LANCZOS)
    img.save(output_path, format="PNG")
    return output_path

def create_slide(
    output_pptx_path: str,
    title_text: str = "P E R C E N T A G E S",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Semi-Circular Gauge Array effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # === Master Title ===
    title_box = slide.shapes.add_textbox(Inches(1.66), Inches(0.4), Inches(10), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.name = "Century Gothic"
    p.font.color.rgb = RGBColor(150, 150, 150)
    p.alignment = PP_ALIGN.CENTER

    sub_p = tf.add_paragraph()
    sub_p.text = "This is a demo text you may write a brief text here to explain the title or if you think you do\nnot need this you may consider to delete the text box"
    sub_p.font.size = Pt(12)
    sub_p.font.name = "Century Gothic"
    sub_p.font.color.rgb = RGBColor(120, 120, 120)
    sub_p.alignment = PP_ALIGN.CENTER
    
    # === Data Definition ===
    gauges_data = [
        {"pct": 0.60, "title": "GRAPHIC DESIGN", "color": (232, 58, 114)},
        {"pct": 0.70, "title": "WEB DESIGN", "color": (0, 153, 143)},
        {"pct": 0.50, "title": "VIDEO EDITING", "color": (139, 195, 74)},
        {"pct": 0.90, "title": "UX DESIGN", "color": (26, 115, 232)}
    ]
    
    # === Layout Math ===
    num_items = len(gauges_data)
    section_width = prs.slide_width / num_items
    img_width = Inches(2.4)
    img_height = Inches(1.32)
    
    y_img = Inches(2.5)
    y_pct = Inches(4.0)
    y_title = Inches(4.8)
    y_body = Inches(5.2)

    # === Render Gauges ===
    for i, data in enumerate(gauges_data):
        center_x = (i * section_width) + (section_width / 2)
        
        # 1. Generate and insert Gauge Image
        img_path = f"temp_gauge_{i}.png"
        create_gauge_image(data["pct"], data["color"], img_path)
        
        pic_left = center_x - (img_width / 2)
        slide.shapes.add_picture(img_path, pic_left, y_img, width=img_width, height=img_height)
        
        # Cleanup temp image
        if os.path.exists(img_path):
            os.remove(img_path)

        # 2. Percentage Text
        pct_box = slide.shapes.add_textbox(center_x - Inches(1.25), y_pct, Inches(2.5), Inches(0.8))
        tf_pct = pct_box.text_frame
        p_pct = tf_pct.paragraphs[0]
        p_pct.text = f"{int(data['pct']*100)}%"
        p_pct.font.bold = True
        p_pct.font.size = Pt(36)
        p_pct.font.name = "Century Gothic"
        p_pct.font.color.rgb = RGBColor(*data["color"])
        p_pct.alignment = PP_ALIGN.CENTER

        # 3. Category Title
        title_box = slide.shapes.add_textbox(center_x - Inches(1.5), y_title, Inches(3.0), Inches(0.5))
        tf_t = title_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = data["title"]
        p_t.font.bold = True
        p_t.font.size = Pt(16)
        p_t.font.name = "Century Gothic"
        p_t.font.color.rgb = RGBColor(*data["color"])
        p_t.alignment = PP_ALIGN.CENTER

        # 4. Explanatory Body
        body_box = slide.shapes.add_textbox(center_x - Inches(1.5), y_body, Inches(3.0), Inches(1.0))
        tf_b = body_box.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = "Here You Should Add\nSome Brief Text to Explain\nMain Title"
        p_b.font.size = Pt(12)
        p_b.font.name = "Century Gothic"
        p_b.font.color.rgb = RGBColor(150, 150, 150)
        p_b.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
