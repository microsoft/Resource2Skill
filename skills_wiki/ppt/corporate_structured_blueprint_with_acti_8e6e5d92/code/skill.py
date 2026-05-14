def create_slide(
    output_pptx_path: str,
    title_text: str = "Visualize Effectively",
    total_steps: int = 7,
    current_step: int = 5,
    accent_color: tuple = (213, 0, 249),  # Vibrant Magenta
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Corporate Structured Blueprint with Active Navigator".
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    # Use standard 16:9 widescreen ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Colors ===
    DARK_TEXT = RGBColor(26, 26, 26)
    MUTED_GREY = RGBColor(150, 150, 150)
    LIGHT_BG = RGBColor(245, 245, 245)
    ACCENT = RGBColor(*accent_color)

    # === Layer 1: Header / Wayfinding ===
    
    # 1A. Active Navigator (e.g., 1 | 2 | 3 | 4 | 5 | 6 | 7)
    nav_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.5))
    tf_nav = nav_box.text_frame
    tf_nav.clear()
    p_nav = tf_nav.paragraphs[0]
    
    for i in range(1, total_steps + 1):
        run = p_nav.add_run()
        run.text = str(i)
        run.font.name = "Arial"
        
        # Highlight active step
        if i == current_step:
            run.font.color.rgb = ACCENT
            run.font.bold = True
            run.font.size = Pt(16)
        else:
            run.font.color.rgb = MUTED_GREY
            run.font.size = Pt(14)
            run.font.bold = False
            
        # Add separator
        if i < total_steps:
            sep = p_nav.add_run()
            sep.text = "  |  "
            sep.font.color.rgb = MUTED_GREY
            sep.font.size = Pt(14)

    # 1B. Corporate Mark / Logo Placeholder (Top Right)
    logo_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(12.033), Inches(0.3), Inches(0.8), Inches(0.6)
    )
    logo_shape.fill.solid()
    logo_shape.fill.fore_color.rgb = DARK_TEXT
    logo_shape.line.fill.background()  # Remove border
    
    tf_logo = logo_shape.text_frame
    tf_logo.text = "PM"
    tf_logo.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_logo.paragraphs[0].runs[0].font.color.rgb = ACCENT
    tf_logo.paragraphs[0].runs[0].font.bold = True
    tf_logo.paragraphs[0].runs[0].font.size = Pt(20)

    # === Layer 2: Main Context ===
    
    # 2A. Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(10), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.text = title_text
    tf_title.paragraphs[0].runs[0].font.size = Pt(36)
    tf_title.paragraphs[0].runs[0].font.bold = True
    tf_title.paragraphs[0].runs[0].font.color.rgb = DARK_TEXT

    # 2B. Structural Accent Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.8), Inches(1.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # === Layer 3: Content Blueprint (Simulated Layout) ===
    
    # 3A. Left Column: Bulleted Principles
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(6), Inches(4.5))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    
    header_p = tf_content.add_paragraph()
    header_p.text = "Core Principles:"
    header_p.font.size = Pt(24)
    header_p.font.bold = True
    header_p.font.color.rgb = DARK_TEXT

    bullets = [
        "Stick to your master layout and corporate design.",
        "Use modern icons or high-quality photography.",
        "Avoid outdated clip art and complex 3D effects.",
        "Limit messages to one key takeaway per slide."
    ]
    for bullet in bullets:
        p = tf_content.add_paragraph()
        p.text = bullet
        p.level = 1
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_TEXT
        # Add some vertical spacing
        p.space_before = Pt(14)

    # 3B. Right Column: Data/Visual Placeholder
    vis_placeholder = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.3), Inches(5.8), Inches(4.0)
    )
    vis_placeholder.fill.solid()
    vis_placeholder.fill.fore_color.rgb = LIGHT_BG
    vis_placeholder.line.color.rgb = MUTED_GREY
    
    tf_vis = vis_placeholder.text_frame
    tf_vis.text = "Visual Data Placeholder\n(Insert Chart or Infographic Here)"
    tf_vis.paragraphs[0].alignment = PP_ALIGN.CENTER
    for r in tf_vis.paragraphs[0].runs:
        r.font.color.rgb = MUTED_GREY
        r.font.size = Pt(16)

    # === Layer 4: Footer / Tracking ===
    
    # 4A. Source Note (Bottom Left)
    footer_left = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(5), Inches(0.4))
    tf_footer = footer_left.text_frame
    tf_footer.text = "Source: Corporate Presentation Mastery Guidelines"
    tf_footer.paragraphs[0].runs[0].font.size = Pt(10)
    tf_footer.paragraphs[0].runs[0].font.color.rgb = MUTED_GREY

    # 4B. Pagination (Bottom Right)
    footer_right = slide.shapes.add_textbox(Inches(12.0), Inches(7.0), Inches(0.8), Inches(0.4))
    tf_page = footer_right.text_frame
    tf_page.text = f"{current_step} / {total_steps}"
    tf_page.paragraphs[0].alignment = PP_ALIGN.RIGHT
    tf_page.paragraphs[0].runs[0].font.size = Pt(10)
    tf_page.paragraphs[0].runs[0].font.bold = True
    tf_page.paragraphs[0].runs[0].font.color.rgb = MUTED_GREY

    prs.save(output_pptx_path)
    return output_pptx_path
