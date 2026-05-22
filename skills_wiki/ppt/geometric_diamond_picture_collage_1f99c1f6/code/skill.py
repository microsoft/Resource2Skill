import os
import tempfile
import urllib.request
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def download_and_crop_square(url: str, filename: str) -> str:
    """
    Downloads an image and uses PIL to crop it to a perfect 1:1 square.
    This guarantees that when PPT applies 'stretch' picture fill into a 1:1 diamond shape,
    the image aspect ratio remains perfectly undistorted.
    """
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(filename, 'wb') as f:
                f.write(response.read())
        
        # Crop to square using PIL
        with Image.open(filename) as img:
            img = img.convert('RGB')
            w, h = img.size
            s = min(w, h)
            left = (w - s) // 2
            top = (h - s) // 2
            img_cropped = img.crop((left, top, left + s, top + s))
            img_cropped.save(filename)
            
    except Exception as e:
        print(f"Failed to download/crop {url}: {e}")
        # Create a fallback flat colored square
        img = Image.new('RGB', (800, 800), color=(200, 210, 220))
        img.save(filename)
        
    return filename

def create_slide(
    output_pptx_path: str,
    title_text: str = "Company Name",
    accent_color: tuple = (41, 53, 86),  # Deep Navy Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Geometric Diamond Picture Collage.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # 1. Setup temp directory for downloaded images
    tmpdir = tempfile.mkdtemp()
    
    # Using generic image placeholders that fit a fashion/apparel theme
    image_urls = [
        "https://images.unsplash.com/photo-1515886657613-9f3515b0c78f?w=800&q=80", # Top
        "https://images.unsplash.com/photo-1434389678369-182cb11b6510?w=800&q=80", # Bottom
        "https://images.unsplash.com/photo-1483985988355-763728e1935b?w=800&q=80", # Left
        "https://images.unsplash.com/photo-1520006403909-838d6b92c22e?w=800&q=80", # Right
        "https://images.unsplash.com/photo-1550614000-4b95d4ed7963?w=800&q=80"  # Portrait left
    ]
    
    img_paths = []
    for i, url in enumerate(image_urls):
        path = os.path.join(tmpdir, f"img_{i}.jpg")
        download_and_crop_square(url, path)
        img_paths.append(path)

    # === Layer 1: The Diamond Grid Collage ===
    
    # Math for the grid
    W = Inches(2.6) # Width/Height of the bounding box of each diamond
    gap = Inches(0.08) # Spacing parameter
    D = W/2 + gap # Distance from cluster center to individual diamond centers
    
    # Center point for the whole diamond cluster (shifted slightly right)
    CX = Inches(8.5)
    CY = Inches(4.2)
    
    # Positions: Top, Bottom, Left, Right
    positions = [
        (CX, CY - D), 
        (CX, CY + D), 
        (CX - D, CY), 
        (CX + D, CY)  
    ]
    
    for i, (cx, cy) in enumerate(positions):
        left = cx - W/2
        top = cy - W/2
        # Adding MSO_SHAPE.DIAMOND. By ensuring width == height, it becomes a 45deg rotated square.
        diamond = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, left, top, W, W)
        
        # Style the diamond
        diamond.line.color.rgb = RGBColor(255, 255, 255)
        diamond.line.width = Pt(3) # Nice white border to define the shape
        
        # Insert pre-cropped 1:1 image
        diamond.fill.user_picture(img_paths[i])

    # === Layer 2: Editorial Portrait Image (Left Side) ===
    # Adding a standalone vertical image to balance the composition
    left_img_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.5), Inches(2.5), Inches(4.0))
    left_img_shape.line.color.rgb = RGBColor(200, 200, 200)
    left_img_shape.line.width = Pt(1)
    left_img_shape.fill.user_picture(img_paths[4])


    # === Layer 3: Typography & Lines ===
    
    # Top Accent Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.33), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background() # No border
    
    # Header Title
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.3), Inches(10.33), Inches(0.8))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.name = 'Georgia' # Using a serif font for fashion/editorial feel
    p.font.color.rgb = RGBColor(33, 33, 33)
    
    # Small Contact/Detail Info Block (Bottom Left)
    info_box = slide.shapes.add_textbox(Inches(1.5), Inches(6.6), Inches(3.0), Inches(0.8))
    tf_info = info_box.text_frame
    p_info = tf_info.paragraphs[0]
    p_info.text = "Company Location\nCompany Email\nCompany Website"
    p_info.font.size = Pt(9)
    p_info.font.name = 'Arial'
    p_info.font.color.rgb = RGBColor(100, 100, 100)

    prs.save(output_pptx_path)
    return output_pptx_path
