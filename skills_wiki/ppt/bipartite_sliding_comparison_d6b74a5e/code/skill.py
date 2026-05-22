def create_bipartite_sliding_comparison(
    output_pptx_path: str = "bipartite_comparison.pptx",
    # --- Topic 1 (e.g., Republican) ---
    topic1_title: str = "REPUBLICAN",
    topic1_points: list = [
        ("CORE VALUE 1", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 2", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 3", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 4", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
    ],
    topic1_color: tuple = (222, 1, 0), # Red
    topic1_icon_url: str = "https://i.ibb.co/L5T9T2D/republican-logo-republican-party-logo-1.png",
    # --- Topic 2 (e.g., Democratic) ---
    topic2_title: str = "DEMOCRATIC",
    topic2_points: list = [
        ("CORE VALUE 1", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 2", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 3", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
        ("CORE VALUE 4", "Lorem ipsum dolor sit amet, consectetur adipiscing elit."),
    ],
    topic2_color: tuple = (30, 56, 136), # Blue
    topic2_icon_url: str = "https://i.ibb.co/mS6D4tF/democrat-logo-democratic-party-logo-1.png",
    # --- Core Visuals ---
    main_image_url: str = "https://images.unsplash.com/photo-1601815454228-21d1b0c97858",
    laptop_mockup_url: str = "https://i.ibb.co/h7gJj4h/laptop-mockup-transparent.png",
) -> str:
    """
    Creates a two-slide presentation reproducing the Bipartite Sliding Comparison effect.

    This function generates two static slides with the correct layouts.
    **MANUAL STEP REQUIRED**: To complete the effect, open the generated .pptx file,
    select both slides, go to the 'Transitions' tab, select 'Push', and in
    'Effect Options', choose 'From Right'. Set the duration to 1.75s.

    Args:
        output_pptx_path (str): Path to save the final presentation.
        All other arguments configure the content and appearance of the two topics.

    Returns:
        str: The path to the saved PPTX file.
    """
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image

    # --- Asset Downloading and Processing ---
    def get_image_from_url(url):
        try:
            with urllib.request.urlopen(url) as response:
                return Image.open(io.BytesIO(response.read())).convert("RGBA")
        except Exception as e:
            print(f"Warning: Could not download image from {url}. Error: {e}")
            return Image.new("RGBA", (800, 600), (128, 128, 128, 255))

    # Create the split-toned image using PIL
    main_img = get_image_from_url(main_image_url)
    
    # Create semi-transparent color overlays
    width, height = main_img.size
    red_overlay = Image.new("RGBA", main_img.size, topic1_color + (100,))
    blue_overlay = Image.new("RGBA", main_img.size, topic2_color + (120,))

    # Create a combined overlay (left half red, right half blue)
    split_overlay = Image.new("RGBA", main_img.size)
    split_overlay.paste(red_overlay.crop((0, 0, width // 2, height)), (0, 0))
    split_overlay.paste(blue_overlay.crop((width // 2, 0, width, height)), (width // 2, 0))

    # Composite the overlay onto the main image
    split_toned_img = Image.alpha_composite(main_img, split_overlay)
    
    split_img_bytes = io.BytesIO()
    split_toned_img.save(split_img_bytes, format='PNG')
    split_img_bytes.seek(0)

    laptop_img_bytes = io.BytesIO()
    get_image_from_url(laptop_mockup_url).save(laptop_img_bytes, format='PNG')
    laptop_img_bytes.seek(0)
    
    topic1_icon_bytes = io.BytesIO()
    get_image_from_url(topic1_icon_url).save(topic1_icon_bytes, format='PNG')
    topic1_icon_bytes.seek(0)
    
    topic2_icon_bytes = io.BytesIO()
    get_image_from_url(topic2_icon_url).save(topic2_icon_bytes, format='PNG')
    topic2_icon_bytes.seek(0)


    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SLIDE_WIDTH, SLIDE_HEIGHT = prs.slide_width, prs.slide_height
    blank_layout = prs.slide_layouts[6]

    # Common Colors
    BG_COLOR = RGBColor(242, 242, 242)
    TEXT_COLOR = RGBColor(89, 89, 89)
    ICON_BG_COLOR = RGBColor(51, 51, 51)
    
    # Common Dimensions
    LAPTOP_WIDTH = Inches(8.5)
    LAPTOP_HEIGHT = Inches(5.0)
    LAPTOP_Y_POS = (SLIDE_HEIGHT - LAPTOP_HEIGHT) / 2
    
    # Screen area inside mockup (approximate values, adjust if mockup changes)
    SCREEN_WIDTH_RATIO = 0.77
    SCREEN_HEIGHT_RATIO = 0.81
    SCREEN_X_OFFSET_RATIO = 0.115
    SCREEN_Y_OFFSET_RATIO = 0.08
    
    screen_width = LAPTOP_WIDTH * SCREEN_WIDTH_RATIO
    screen_height = LAPTOP_HEIGHT * SCREEN_HEIGHT_RATIO
    screen_x_offset = LAPTOP_WIDTH * SCREEN_X_OFFSET_RATIO
    screen_y_offset = LAPTOP_HEIGHT * SCREEN_Y_OFFSET_RATIO
    
    # --- Create Slide 1 ---
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.background.fill.solid()
    slide1.background.fill.fore_color.rgb = BG_COLOR

    # Laptop and Screen
    laptop_x1 = SLIDE_WIDTH - (LAPTOP_WIDTH / 2)
    screen_x1 = laptop_x1 + screen_x_offset
    screen_y1 = LAPTOP_Y_POS + screen_y_offset
    slide1.shapes.add_picture(split_img_bytes, screen_x1, screen_y1, width=screen_width, height=screen_height)
    slide1.shapes.add_picture(laptop_img_bytes, laptop_x1, LAPTOP_Y_POS, width=LAPTOP_WIDTH, height=LAPTOP_HEIGHT)

    # Icon
    icon_size = Inches(1.1)
    icon_bg_shape = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(2.2), icon_size, icon_size)
    icon_bg_shape.fill.solid()
    icon_bg_shape.fill.fore_color.rgb = ICON_BG_COLOR
    icon_bg_shape.line.fill.background()
    slide1.shapes.add_picture(topic1_icon_bytes, Inches(1.2) + Inches(0.1), Inches(2.2) + Inches(0.1), width=Inches(0.9), height=Inches(0.9))

    # Text Content
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(0.8), Inches(6), Inches(1))
    p = title_box.text_frame.paragraphs[0]
    p.text = topic1_title
    p.font.name = 'Arial Black'
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(*topic1_color)

    y_start = Inches(3.5)
    for title, body in topic1_points:
        sub_title_box = slide1.shapes.add_textbox(Inches(1.2), y_start - Inches(0.4), Inches(4), Inches(0.5))
        p = sub_title_box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = 'Arial Bold'
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(*topic1_color)
        
        body_box = slide1.shapes.add_textbox(Inches(1.2), y_start, Inches(4), Inches(0.5))
        p = body_box.text_frame.paragraphs[0]
        p.text = body
        p.font.name = 'Arial'
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_COLOR
        y_start += Inches(1)

    # --- Create Slide 2 (Mirror Image) ---
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = BG_COLOR
    
    # Reset image stream
    split_img_bytes.seek(0)
    laptop_img_bytes.seek(0)
    
    # Laptop and Screen
    laptop_x2 = -(LAPTOP_WIDTH / 2)
    screen_x2 = laptop_x2 + screen_x_offset
    screen_y2 = LAPTOP_Y_POS + screen_y_offset
    slide2.shapes.add_picture(split_img_bytes, screen_x2, screen_y2, width=screen_width, height=screen_height)
    slide2.shapes.add_picture(laptop_img_bytes, laptop_x2, LAPTOP_Y_POS, width=LAPTOP_WIDTH, height=LAPTOP_HEIGHT)
    
    # Icon
    icon2_x_pos = SLIDE_WIDTH - Inches(1.2) - icon_size
    icon_bg_shape2 = slide2.shapes.add_shape(MSO_SHAPE.OVAL, icon2_x_pos, Inches(2.2), icon_size, icon_size)
    icon_bg_shape2.fill.solid()
    icon_bg_shape2.fill.fore_color.rgb = ICON_BG_COLOR
    icon_bg_shape2.line.fill.background()
    slide2.shapes.add_picture(topic2_icon_bytes, icon2_x_pos + Inches(0.1), Inches(2.2) + Inches(0.1), width=Inches(0.9), height=Inches(0.9))
    
    # Text Content
    title_box2 = slide2.shapes.add_textbox(Inches(6.1), Inches(0.8), Inches(6), Inches(1))
    tf = title_box2.text_frame
    tf.paragraphs[0].text = topic2_title
    tf.paragraphs[0].font.name = 'Arial Black'
    tf.paragraphs[0].font.size = Pt(44)
    tf.paragraphs[0].font.color.rgb = RGBColor(*topic2_color)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    
    y_start = Inches(3.5)
    text_x_pos = SLIDE_WIDTH - Inches(1.2) - Inches(4)
    for title, body in topic2_points:
        sub_title_box2 = slide2.shapes.add_textbox(text_x_pos, y_start - Inches(0.4), Inches(4), Inches(0.5))
        tf = sub_title_box2.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.name = 'Arial Bold'
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = RGBColor(*topic2_color)
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT

        body_box2 = slide2.shapes.add_textbox(text_x_pos, y_start, Inches(4), Inches(0.5))
        tf = body_box2.text_frame
        tf.paragraphs[0].text = body
        tf.paragraphs[0].font.name = 'Arial'
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = TEXT_COLOR
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        y_start += Inches(1)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_bipartite_sliding_comparison()

