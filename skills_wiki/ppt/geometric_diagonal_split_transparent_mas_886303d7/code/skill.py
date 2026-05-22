def create_slide(
    output_pptx_path: str,
    title_text: str = "BUSINESS",
    title_accent: str = "PRESENTATION",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa. Fusce posuere, magna sed pulvinar ultricies, purus lectus malesuada libero, sit amet commodo nunc viverra imperdiet enim.",
    theme_color: tuple = (21, 96, 161),  # RGB for Corporate Blue
    accent_color: tuple = (255, 192, 0),  # RGB for Accent Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Geometric Diagonal Split" visual effect.
    """
    import io
    import urllib.request
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation (16:9 widescreen)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # ==========================================
    # Layer 1: Background Image
    # ==========================================
    image_url = "https://images.unsplash.com/photo-1556761175-5973dc0f32e7?q=80&w=1920&auto=format&fit=crop"
    
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            bg_img = Image.open(io.BytesIO(response.read())).convert("RGB")
            bg_img = bg_img.resize((1920, 1080))
    except Exception:
        # Fallback if download fails
        bg_img = Image.new("RGB", (1920, 1080), (80, 80, 80))
        draw_bg = ImageDraw.Draw(bg_img)
        draw_bg.text((800, 500), "BACKGROUND IMAGE (Download Failed)", fill=(200, 200, 200))

    # Save to buffer and insert
    bg_stream = io.BytesIO()
    bg_img.save(bg_stream, format="JPEG")
    bg_stream.seek(0)
    slide.shapes.add_picture(bg_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # ==========================================
    # Layer 2: Geometric PIL Alpha Overlay Mask
    # ==========================================
    # We generate a 1920x1080 RGBA image to lay exactly over the slide
    overlay = Image.new("RGBA", (1920, 1080), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")

    # Math for the slant: dx = 480px leftward shift from top to bottom
    # Base solid block
    draw.polygon([(0, 0), (1250, 0), (770, 1080), (0, 1080)], fill=(*theme_color, 255))
    
    # Stripe 1: Thicker, medium alpha
    draw.polygon([(1270, 0), (1390, 0), (910, 1080), (790, 1080)], fill=(*theme_color, 200))
    
    # Stripe 2: Thinner, lower alpha
    draw.polygon([(1410, 0), (1470, 0), (990, 1080), (930, 1080)], fill=(*theme_color, 128))
    
    # Stripe 3: Thinnest, lowest alpha (accent edge)
    draw.polygon([(1485, 0), (1505, 0), (1025, 1080), (1005, 1080)], fill=(*theme_color, 76))

    # Save overlay to buffer and insert
    overlay_stream = io.BytesIO()
    overlay.save(overlay_stream, format="PNG")
    overlay_stream.seek(0)
    slide.shapes.add_picture(overlay_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # ==========================================
    # Layer 3: Logo & Typography
    # ==========================================
    # Add a mock logo icon (4 small squares) top-left
    for row in range(2):
        for col in range(2):
            left = Inches(0.8 + col * 0.15)
            top = Inches(0.6 + row * 0.15)
            shp = slide.shapes.add_shape(5, left, top, Inches(0.12), Inches(0.12)) # 5 = rounded rectangle
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(255, 255, 255)
            shp.line.fill.background()

    # "LOGO" text next to the icon
    tb_logo = slide.shapes.add_textbox(Inches(1.15), Inches(0.6), Inches(2), Inches(0.4))
    run_logo = tb_logo.text_frame.paragraphs[0].add_run()
    run_logo.text = "LOGO"
    run_logo.font.bold = True
    run_logo.font.size = Pt(24)
    run_logo.font.color.rgb = RGBColor(255, 255, 255)

    # Main Text Block
    tb_main = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.0))
    tf_main = tb_main.text_frame
    tf_main.word_wrap = True
    
    # 1. Super-title
    p_super = tf_main.paragraphs[0]
    run_super = p_super.add_run()
    run_super.text = "POWERPOINT SHOW\n"
    run_super.font.size = Pt(14)
    run_super.font.color.rgb = RGBColor(200, 220, 255)  # Light blue-white
    
    # 2. Main Title - First line
    p_title1 = tf_main.add_paragraph()
    run_title1 = p_title1.add_run()
    run_title1.text = f"{title_text.upper()}\n"
    run_title1.font.bold = True
    run_title1.font.size = Pt(54)
    run_title1.font.color.rgb = RGBColor(255, 255, 255)
    
    # 3. Main Title - Second line (Accent)
    p_title2 = tf_main.add_paragraph()
    run_title2 = p_title2.add_run()
    run_title2.text = f"{title_accent.upper()}"
    run_title2.font.bold = True
    run_title2.font.size = Pt(54)
    run_title2.font.color.rgb = RGBColor(*accent_color)
    
    # 4. Spacer line
    p_space = tf_main.add_paragraph()
    p_space.add_run().font.size = Pt(14)
    
    # 5. Subtitle
    p_sub = tf_main.add_paragraph()
    run_sub = p_sub.add_run()
    run_sub.text = "ADD SUB TITLE HERE\n"
    run_sub.font.bold = True
    run_sub.font.size = Pt(16)
    run_sub.font.color.rgb = RGBColor(255, 255, 255)
    
    # 6. Body Text
    p_body = tf_main.add_paragraph()
    p_body.alignment = PP_ALIGN.LEFT
    run_body = p_body.add_run()
    run_body.text = body_text
    run_body.font.size = Pt(11)
    run_body.font.color.rgb = RGBColor(210, 230, 250)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
