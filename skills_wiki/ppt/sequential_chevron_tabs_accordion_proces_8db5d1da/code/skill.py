import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree

def create_slide(output_pptx_path: str):
    """
    Creates a presentation featuring the Sequential Chevron Tabs process layout.
    """
    prs = Presentation()
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Design configuration
    steps = 5
    base_width = prs.slide_width / steps
    point_extension = Inches(0.6)  # The chevron point extending to the right
    height = prs.slide_height

    # Monochromatic blue palette (from light to dark)
    palette = [
        RGBColor(135, 230, 245),
        RGBColor(105, 205, 235),
        RGBColor(70, 165, 205),
        RGBColor(45, 120, 170),
        RGBColor(25, 65, 105)
    ]

    letters = ["A", "B", "C", "D", "E"]
    icons = ["★", "⚙", "🚀", "💡", "📊"] # Unicode fallbacks for icons
    
    # Helper function to inject shadow via lxml
    def apply_right_shadow(shape):
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw')
        
        # 100000 EMUs is ~7.8 points. 
        outerShdw.set('blurRad', '150000') # Soft blur
        outerShdw.set('dist', '80000')     # Offset distance
        outerShdw.set('dir', '0')          # 0 degrees = rightwards
        outerShdw.set('algn', 'ctr')
        outerShdw.set('rotWithShape', '0')
        
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        srgbClr.set('val', '000000')
        alpha = etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
        alpha.set('val', '30000') # 30% opacity

    # ==========================================
    # LAYER 1: SHAPES (Drawn Back to Front)
    # ==========================================
    # We loop backward (4 down to 0) so the rightmost shape is at the bottom,
    # and the leftmost shape is on top, making the shadows cast correctly onto the shape next to it.
    
    for i in range(steps - 1, -1, -1):
        x_offset = i * base_width
        
        # Custom Chevron Tab Geometry
        ff_builder = slide.shapes.build_freeform()
        ff_builder.add_line_segments([
            (x_offset, 0),                                       # Top Left
            (x_offset + base_width, 0),                          # Top Right
            (x_offset + base_width + point_extension, height/2), # Center Right Point
            (x_offset + base_width, height),                     # Bottom Right
            (x_offset, height),                                  # Bottom Left
            (x_offset, 0)                                        # Close shape
        ], close=True)
        
        shape = ff_builder.convert_to_shape()
        
        # Styling
        shape.fill.solid()
        shape.fill.fore_color.rgb = palette[i]
        shape.line.fill.background() # No outline
        
        # Apply Depth
        apply_right_shadow(shape)

    # ==========================================
    # LAYER 2: TEXT & CONTENT (Drawn Left to Right)
    # ==========================================
    # Text is drawn afterwards so it sits on top of all shapes
    
    for i in range(steps):
        x_offset = i * base_width
        # Center of the rectangular part of the tab
        center_x = x_offset + (base_width / 2)
        
        # 1. Large Step Letter
        tx_letter = slide.shapes.add_textbox(center_x - Inches(0.5), Inches(0.8), Inches(1.0), Inches(1.0))
        p_letter = tx_letter.text_frame.paragraphs[0]
        p_letter.text = letters[i]
        p_letter.alignment = PP_ALIGN.CENTER
        p_letter.font.name = "Montserrat Black"
        p_letter.font.size = Pt(65)
        p_letter.font.bold = True
        p_letter.font.color.rgb = RGBColor(255, 255, 255)

        # 2. Subtitle
        tx_title = slide.shapes.add_textbox(x_offset + Inches(0.2), Inches(2.6), base_width - Inches(0.4), Inches(0.5))
        p_title = tx_title.text_frame.paragraphs[0]
        p_title.text = f"LOREM IPSUM"
        p_title.alignment = PP_ALIGN.CENTER
        p_title.font.name = "Montserrat"
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

        # 3. Body Text
        tx_body = slide.shapes.add_textbox(x_offset + Inches(0.2), Inches(3.2), base_width - Inches(0.4), Inches(2.0))
        tx_body.text_frame.word_wrap = True
        p_body = tx_body.text_frame.paragraphs[0]
        p_body.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore."
        p_body.alignment = PP_ALIGN.CENTER
        p_body.font.name = "Montserrat"
        p_body.font.size = Pt(10)
        p_body.font.color.rgb = RGBColor(255, 255, 255)

        # 4. Icon (Using Unicode characters for immediate rendering)
        tx_icon = slide.shapes.add_textbox(center_x - Inches(0.5), Inches(5.8), Inches(1.0), Inches(1.0))
        p_icon = tx_icon.text_frame.paragraphs[0]
        p_icon.text = icons[i]
        p_icon.alignment = PP_ALIGN.CENTER
        p_icon.font.name = "Segoe UI Emoji"
        p_icon.font.size = Pt(36)
        p_icon.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
