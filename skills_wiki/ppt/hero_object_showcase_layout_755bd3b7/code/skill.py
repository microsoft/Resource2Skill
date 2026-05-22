import os
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_slide(
    output_pptx_path: str,
    title_text: str = "The Lovely Parrot",
    image_keyword: str = "macaw,parrot",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Hero Object Showcase" static layout,
    ready for animation application.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Background ===
    # Default white background is sufficient for this clean style

    # === Layer 2: Title Typography ===
    # Create a text box at the top
    title_left = Inches(1)
    title_top = Inches(0.5)
    title_width = Inches(11.333)
    title_height = Inches(1.5)
    
    txBox = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    
    # Style the text to match the bold, punchy look in the video
    run = p.runs[0]
    run.font.name = 'Arial Black' # A common heavy font
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 30, 30) # Near black

    # === Layer 3: Hero Image ===
    # Download a sample image
    image_path = "temp_hero_image.jpg"
    try:
        url = f"https://source.unsplash.com/featured/800x600/?{image_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(image_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception as e:
        print(f"Image download failed: {e}. Creating a placeholder rectangle instead.")
        # Fallback if download fails: draw a colored rectangle
        shape = slide.shapes.add_shape(
            1, # msoShapeRectangle
            Inches(3.16), Inches(2.2), Inches(7), Inches(4.5)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(200, 50, 50)
        shape.line.fill.background()
        image_path = None

    if image_path and os.path.exists(image_path):
        # Insert and center the image
        # Assuming a target width of about 7 inches for a 13.333 wide slide
        target_width = Inches(7)
        pic = slide.shapes.add_picture(image_path, Inches(0), Inches(0), width=target_width)
        
        # Center horizontally
        pic.left = int((prs.slide_width - pic.width) / 2)
        # Position below the title
        pic.top = Inches(2.2)
        
        # Clean up temp file
        os.remove(image_path)

    # Save presentation
    prs.save(output_pptx_path)
    print(f"Slide saved to {output_pptx_path}. Remember to manually add Fade (Title) and Zoom (Image) animations!")
    return output_pptx_path

# Example usage:
# create_slide("hero_layout_ready_for_animation.pptx")
