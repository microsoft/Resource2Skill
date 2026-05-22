import os
import urllib.request
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "PROJECT\nDESCRIPTION",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Aenean commodo ligula eget dolor. Aenean massa. Cum sociis natoque penatibus et magnis dis parturient montes, nascetur ridiculus mus.",
    image_url: str = "https://images.unsplash.com/photo-1556761175-5973dc0f32d7?ixlib=rb-4.0.3&auto=format&fit=crop&w=800&q=80",
    accent_color: tuple = (213, 43, 82),  # Magenta/Red
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Geometric Offset Diamond Overlay' visual effect.
    """
    # 1. Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # 2. PIL Image Processing: Create the Diamond Mask
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception as e:
        print(f"Failed to download image, using fallback solid color. Error: {e}")
        img = Image.new("RGBA", (800, 800), (200, 200, 200, 255))

    # Make the image perfectly square via center crop
    w, h = img.size
    min_dim = min(w, h)
    left = (w - min_dim) // 2
    top = (h - min_dim) // 2
    img_square = img.crop((left, top, left + min_dim, top + min_dim))
    img_square = img_square.resize((1000, 1000), Image.Resampling.LANCZOS)

    # Create diamond (rhombus) mask
    mask = Image.new("L", (1000, 1000), 0)
    draw = ImageDraw.Draw(mask)
    draw.polygon([(500, 0), (1000, 500), (500, 1000), (0, 500)], fill=255)
    
    # Apply mask and save temporarily
    img_square.putalpha(mask)
    temp_img_path = "temp_diamond_mask.png"
    img_square.save(temp_img_path, format="PNG")

    # 3. Build Slide Layout - Right Side (The layered geometric effect)
    
    # Layer 1: The masked image
    img_width = Inches(6.0)
    img_height = Inches(6.0)
    img_left = Inches(6.5)
    img_top = Inches(0.75)
    slide.shapes.add_picture(temp_img_path, img_left, img_top, img_width, img_height)
    
    # Clean up temp image
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)

    # Layer 2: The offset hollow frame
    frame_left = Inches(7.2)  # deliberately offset to the right
    frame_top = Inches(0.5)   # deliberately offset upwards
    
    frame = slide.shapes.add_shape(
        MSO_SHAPE.DIAMOND, frame_left, frame_top, img_width, img_height
    )
    
    # LXML Hack: Remove fill completely so the image shows through
    frame.fill.solid() # Initialize spPr
    spPr = frame.element.spPr
    solidFill = spPr.find(qn('a:solidFill'))
    if solidFill is not None:
        spPr.remove(solidFill)
    spPr.append(parse_xml(r'<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
    
    # Apply thick accent outline
    frame.line.color.rgb = RGBColor(*accent_color)
    frame.line.width = Pt(8)


    # 4. Build Slide Layout - Left Side (Typography and Structure)
    
    # Kicker Text
    tx_kicker = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(2.0), Inches(0.5))
    p_kicker = tx_kicker.text_frame.paragraphs[0]
    p_kicker.text = "MORE."
    p_kicker.font.name = "Arial"
    p_kicker.font.size = Pt(12)
    p_kicker.font.color.rgb = RGBColor(150, 150, 150)

    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(5.0), Inches(1.5))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(40, 40, 45)

    # Accent Divider Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(4.3), Inches(1.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background() # No border

    # Body Paragraph
    tx_body = slide.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(4.5), Inches(1.5))
    tf_body = tx_body.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = RGBColor(120, 120, 120)
    p_body.line_spacing = 1.3

    # Pill Label ("PART ONE")
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(6.0), Inches(1.2), Inches(0.4)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(245, 245, 245)
    pill.line.color.rgb = RGBColor(180, 180, 180)
    pill.line.width = Pt(1)
    
    pill_text = pill.text_frame.paragraphs[0]
    pill_text.text = "PART ONE"
    pill_text.font.size = Pt(9)
    pill_text.font.bold = True
    pill_text.font.color.rgb = RGBColor(100, 100, 100)
    pill_text.alignment = PP_ALIGN.CENTER

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
