def create_slide(
    output_pptx_path: str,
    employee_name: str = "MEL STEAD",
    job_title: str = "Procurement Manager",
    quote: str = '"I innovate on how we can walk the\nfine line between keeping our\ninventory low without stocking out."',
    qa_pairs: list = None,
    portrait_url: str = "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?q=80&w=600&auto=format&fit=crop",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Split-Pane Employee Spotlight" visual effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # Default Q&A if none provided
    if not qa_pairs:
        qa_pairs = [
            {"q": "List five hashtags that describe your personality.", "a": "#whatisahastag #MelAttilaTheHun #bakingsnob\n#studentofancienthistory #sarcastic"},
            {"q": "If you could vacation anywhere in the world, where would you go?", "a": "Norway"},
            {"q": "What is your favorite part about working here?", "a": "I get to work with super smart and amazing people who are really fun and don't mind my odd humor."},
            {"q": "Outside of work, what activity can we find you doing?", "a": "Baking, eating, reading about baking, combing the internet for new ways to use my sous vide machine."},
            {"q": "Apple or Android?", "a": "Apple! Unless they don't support Sonos on Apple Music."}
        ]

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Color Palette
    color_left_bg = RGBColor(227, 232, 236)
    color_nameplate = RGBColor(44, 62, 80)
    color_accent = RGBColor(92, 158, 173)  # Teal
    color_text_dark = RGBColor(51, 51, 51)
    color_white = RGBColor(255, 255, 255)

    # === LEFT PANE ===
    left_pane_width = Inches(4.5)
    
    # 1. Left Background Block
    left_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, left_pane_width, prs.slide_height
    )
    left_bg.fill.solid()
    left_bg.fill.fore_color.rgb = color_left_bg
    left_bg.line.fill.background()

    # 2. Process and Add Portrait (with PIL bottom-fade to blend into background)
    img_path = "temp_portrait.png"
    try:
        req = urllib.request.Request(portrait_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert("RGBA")
            
            # Crop to standard portrait ratio (4:5)
            target_ratio = 4.0 / 5.0
            w, h = img.size
            if w / h > target_ratio:
                new_w = int(h * target_ratio)
                img = img.crop(((w - new_w) // 2, 0, (w + new_w) // 2, h))
            else:
                new_h = int(w / target_ratio)
                img = img.crop((0, 0, w, new_h))
                
            img = img.resize((600, 750), Image.Resampling.LANCZOS)
            
            # Create alpha mask to fade out the bottom seamlessly
            w, h = img.size
            mask = Image.new("L", (w, h), 255)
            draw = ImageDraw.Draw(mask)
            fade_height = int(h * 0.3)  # Bottom 30% fades
            
            for y in range(fade_height):
                alpha = int(255 * (1 - (y / fade_height)))
                y_pos = h - fade_height + y
                draw.line((0, y_pos, w, y_pos), fill=alpha)
                
            img.putalpha(mask)
            img.save(img_path, "PNG")
            
            # Insert into slide
            pic_top = Inches(0.4)
            pic_height = Inches(3.8)
            pic = slide.shapes.add_picture(img_path, Inches(0.5), pic_top, height=pic_height)
            # Center the picture in the pane
            pic.left = int((left_pane_width - pic.width) / 2)
    except Exception as e:
        print(f"Failed to process image: {e}")

    # 3. Nameplate Banner
    nameplate_top = Inches(4.3)
    nameplate = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, nameplate_top, left_pane_width, Inches(0.65)
    )
    nameplate.fill.solid()
    nameplate.fill.fore_color.rgb = color_nameplate
    nameplate.line.fill.background()
    
    tf_name = nameplate.text_frame
    tf_name.text = employee_name.upper()
    tf_name.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_name.paragraphs[0].font.color.rgb = color_white
    tf_name.paragraphs[0].font.size = Pt(24)
    tf_name.paragraphs[0].font.bold = True
    tf_name.vertical_anchor = MSO_ANCHOR.MIDDLE

    # 4. Job Title
    tx_title = slide.shapes.add_textbox(0, Inches(5.1), left_pane_width, Inches(0.5))
    tf_title = tx_title.text_frame
    tf_title.text = job_title
    tf_title.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_title.paragraphs[0].font.color.rgb = color_nameplate
    tf_title.paragraphs[0].font.size = Pt(16)
    tf_title.paragraphs[0].font.bold = True

    # 5. Quote
    tx_quote = slide.shapes.add_textbox(Inches(0.4), Inches(5.6), left_pane_width - Inches(0.8), Inches(1.5))
    tf_quote = tx_quote.text_frame
    tf_quote.word_wrap = True
    tf_quote.text = quote
    p_quote = tf_quote.paragraphs[0]
    p_quote.alignment = PP_ALIGN.CENTER
    p_quote.font.color.rgb = color_nameplate
    p_quote.font.size = Pt(12)
    p_quote.font.italic = True

    # === RIGHT PANE (Q&A LIST) ===
    # Start coordinates for the list
    start_y = Inches(0.6)
    spacing_y = Inches(1.3)
    
    for i, qa in enumerate(qa_pairs[:5]):  # limit to 5 to fit slide
        current_y = start_y + (i * spacing_y)
        
        # Number Badge (Circle)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(5.0), current_y, Inches(0.5), Inches(0.5)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = color_accent
        badge.line.fill.background()
        
        tf_badge = badge.text_frame
        tf_badge.text = str(i + 1)
        p_badge = tf_badge.paragraphs[0]
        p_badge.alignment = PP_ALIGN.CENTER
        p_badge.font.color.rgb = color_white
        p_badge.font.size = Pt(20)
        p_badge.font.bold = True
        tf_badge.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Question Text
        tx_q = slide.shapes.add_textbox(Inches(5.7), current_y - Inches(0.05), Inches(7.0), Inches(0.4))
        tf_q = tx_q.text_frame
        tf_q.word_wrap = True
        p_q = tf_q.paragraphs[0]
        p_q.text = qa['q']
        p_q.font.color.rgb = color_accent
        p_q.font.size = Pt(14)
        p_q.font.bold = True

        # Answer Text
        tx_a = slide.shapes.add_textbox(Inches(5.7), current_y + Inches(0.35), Inches(7.0), Inches(0.8))
        tf_a = tx_a.text_frame
        tf_a.word_wrap = True
        p_a = tf_a.paragraphs[0]
        p_a.text = qa['a']
        p_a.font.color.rgb = color_text_dark
        p_a.font.size = Pt(12)

    # === BRANDING ELEMENT (Optional Logo Badge) ===
    # Adds a small decorative badge spanning the column divide to anchor the aesthetic
    logo_badge = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(4.1), Inches(0.4), Inches(0.8), Inches(0.8)
    )
    logo_badge.fill.solid()
    logo_badge.fill.fore_color.rgb = RGBColor(132, 190, 65)  # Bright Green
    logo_badge.line.fill.background()
    tf_logo = logo_badge.text_frame
    tf_logo.text = "POS\nPORTAL"
    tf_logo.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_logo.paragraphs[0].font.color.rgb = color_white
    tf_logo.paragraphs[0].font.size = Pt(10)
    tf_logo.paragraphs[0].font.bold = True
    if len(tf_logo.paragraphs) > 1:
        tf_logo.paragraphs[1].alignment = PP_ALIGN.CENTER
        tf_logo.paragraphs[1].font.color.rgb = color_white
        tf_logo.paragraphs[1].font.size = Pt(10)
        tf_logo.paragraphs[1].font.bold = True
    tf_logo.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Cleanup temp image
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
