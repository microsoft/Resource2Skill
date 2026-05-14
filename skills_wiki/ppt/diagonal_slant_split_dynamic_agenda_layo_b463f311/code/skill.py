def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda",
    bg_palette: str = "modern building glass skyscraper",
    accent_color: tuple = (74, 144, 226),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Diagonal Slant Split / Dynamic Agenda layout.
    """
    import os
    import requests
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1 & 2: Background & Masked Image Image via PIL ===
    w_px, h_px = 4000, 2250 # High-res canvas corresponding to 13.333 x 7.5 inches
    base_img = Image.new('RGBA', (w_px, h_px), (0, 0, 0, 0))
    
    # Try downloading an architectural photo
    try:
        # Fallback static Unsplash architecture image to ensure it works
        url = "https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=4000&auto=format&fit=crop"
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        photo = Image.open(BytesIO(response.content)).convert("RGBA")
        
        # Center-crop to fit 16:9
        photo_ratio = photo.width / photo.height
        target_ratio = w_px / h_px
        if photo_ratio > target_ratio:
            new_w = int(photo.height * target_ratio)
            left = (photo.width - new_w) // 2
            photo = photo.crop((left, 0, left + new_w, photo.height))
        else:
            new_h = int(photo.width / target_ratio)
            top = (photo.height - new_h) // 2
            photo = photo.crop((0, top, photo.width, top + new_h))
        photo = photo.resize((w_px, h_px), Image.Resampling.LANCZOS)
    except Exception:
        # Fallback placeholder if network fails
        photo = Image.new("RGBA", (w_px, h_px), (40, 50, 60, 255))
        draw = ImageDraw.Draw(photo)
        for i in range(h_px):
            c = int(40 + (i/h_px)*30)
            draw.line([(0, i), (w_px, i)], fill=(c, c+10, c+20, 255))

    # Create the slant mask
    mask = Image.new("L", (w_px, h_px), 0)
    draw_mask = ImageDraw.Draw(mask)
    
    # Slant geometry (inches -> pixels)
    top_inch, bottom_inch = 8.5, 6.0
    top_px = int(w_px * (top_inch / 13.333))
    bottom_px = int(w_px * (bottom_inch / 13.333))
    
    # Draw right-aligned trapezoid
    draw_mask.polygon([(top_px, 0), (w_px, 0), (w_px, h_px), (bottom_px, h_px)], fill=255)
    base_img = Image.composite(photo, base_img, mask)
    
    # Save composite and add to slide
    img_stream = BytesIO()
    base_img.save(img_stream, format="PNG")
    img_stream.seek(0)
    slide.shapes.add_picture(img_stream, 0, 0, Inches(13.333), Inches(7.5))

    # === Layer 3: Diagonal Accent Line ===
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, 
        Inches(top_inch), Inches(0), 
        Inches(bottom_inch), Inches(7.5)
    )
    line.line.color.rgb = RGBColor(*accent_color)
    line.line.width = Pt(4.5)

    # === Layer 4: Title Typography ===
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(5), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(0, 0, 0)
    p.font.name = "Arial"

    # === Layer 5: Dynamic Agenda List ===
    start_y = 2.0
    item_spacing = 1.05
    
    # Mock data for agenda
    items = [
        {"title": "Welcome & Introduction", "desc": "Overview of today's key objectives and session guidelines."},
        {"title": "Q3 Performance Review", "desc": "Analyzing metrics, KPIs, and overall growth from the last quarter."},
        {"title": "Strategic Roadmap", "desc": "A look ahead at product milestones and marketing strategy."},
        {"title": "Team Restructuring", "desc": "Updates on department alignment and new management roles."},
        {"title": "Open Floor Q&A", "desc": "Dedicated time for questions, feedback, and open discussion."}
    ]

    for i, item in enumerate(items):
        y = start_y + (i * item_spacing)
        
        # Accent Circle
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(y + 0.1), Inches(0.35), Inches(0.35))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(*accent_color)
        circle.line.fill.background() # No outline
        
        # Text Content
        tb = slide.shapes.add_textbox(Inches(1.4), Inches(y), Inches(4.2), Inches(0.8))
        tf_item = tb.text_frame
        tf_item.word_wrap = True
        
        # Item Title
        p1 = tf_item.paragraphs[0]
        p1.text = f"Agenda / {item['title']}"
        p1.font.bold = True
        p1.font.size = Pt(16)
        p1.font.color.rgb = RGBColor(50, 50, 50)
        p1.font.name = "Arial"
        
        # Item Description
        p2 = tf_item.add_paragraph()
        p2.text = item['desc']
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(120, 120, 120)
        p2.font.name = "Arial"

    prs.save(output_pptx_path)
    return output_pptx_path
