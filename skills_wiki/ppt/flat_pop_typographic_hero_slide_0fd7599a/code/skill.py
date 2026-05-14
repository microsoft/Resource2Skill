def create_slide(
    output_pptx_path: str,
    title_text: str = "Nuts & Bolts\nSpeed Training",
    bg_color: tuple = (226, 35, 26),        # Vibrant Red
    text_color: tuple = (255, 255, 255),    # White
    shadow_color_hex: str = "4A0905",       # Dark Red/Brown (hex format for XML)
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Flat-Pop Typographic' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree
    from pptx.oxml.ns import qn

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Core Typography Setup ===
    # Calculate perfect center position
    tb_w = Inches(11)
    tb_h = Inches(5)
    tb_left = (prs.slide_width - tb_w) / 2
    tb_top = (prs.slide_height - tb_h) / 2
    
    tb = slide.shapes.add_textbox(tb_left, tb_top, tb_w, tb_h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = title_text
    
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.9  # Compress line spacing to create a cohesive text block
        
        # Font settings for the dynamic, retro look
        p.font.name = "Arial Black"  # Heavy weight sans-serif
        p.font.size = Pt(96)
        p.font.bold = True
        p.font.italic = True
        p.font.color.rgb = RGBColor(*text_color)

    # === Layer 3: XML Injection for Hard Shadow ===
    # The signature of this style is a 100% opacity, 0-blur drop shadow applied to the shape.
    # Because the shape has no background fill, PowerPoint applies this shadow directly to the text.
    def apply_hard_shadow(shape, color_hex, distance_pt=8, angle_deg=45):
        spPr = shape.element.spPr
        
        # Check for existing effect list, create if missing, clear if present
        effectLst = spPr.find(qn('a:effectLst'))
        if effectLst is None:
            effectLst = etree.SubElement(spPr, qn('a:effectLst'))
        else:
            for child in list(effectLst):
                effectLst.remove(child)
                
        # Create outer shadow element
        outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
        outerShdw.set('blurRad', '0')  # CRITICAL: 0 blur creates the hard pop-art edge
        
        # Convert pts to EMUs (1 pt = 12700 EMUs)
        dist_emu = int(distance_pt * 12700)
        outerShdw.set('dist', str(dist_emu))
        
        # Convert degrees to 60000ths of a degree (standard OpenXML format)
        dir_emu = int(angle_deg * 60000)
        outerShdw.set('dir', str(dir_emu))
        
        # Set shadow color
        srgbClr = etree.SubElement(outerShdw, qn('a:srgbClr'))
        srgbClr.set('val', color_hex)
        
        # Ensure 100% opacity
        alpha = etree.SubElement(srgbClr, qn('a:alpha'))
        alpha.set('val', '100000')

    # Apply the shadow to our text box
    apply_hard_shadow(tb, color_hex=shadow_color_hex, distance_pt=7, angle_deg=45)

    # Save and return
    prs.save(output_pptx_path)
    return output_pptx_path
