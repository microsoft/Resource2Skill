def create_slide(
    output_pptx_path: str,
    title_text: str = "Color Explorer",
    body_text: str = "Applying a strict 5-color palette ensures absolute visual harmony across your presentation. Every element is mapped to a specific role: background, primary text, secondary text, and accents.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Strict 5-Color Thematic Styling" effect.
    This programmatic approach simulates the "Eyedropper" technique from the tutorial.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # The 5-color palette extracted from the tutorial's chosen Adobe Color scheme
    palette = {
        "primary_dark": (24, 39, 55),    # Dark Navy
        "secondary": (61, 78, 97),       # Muted Slate Blue
        "neutral_bg": (226, 219, 208),   # Sand / Light Beige
        "accent_bright": (171, 35, 40),  # Crimson Red
        "accent_dark": (111, 22, 28)     # Deep Burgundy
    }

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    # Apply the neutral light color to the background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*palette["neutral_bg"])

    # === Layer 2: Accent Shape (Decorative Banner/Line) ===
    # Use the bright accent color for a visual anchor
    left = Inches(1.5)
    top = Inches(1.5)
    width = Inches(10.333)
    height = Inches(0.15)
    
    accent_bar = slide.shapes.add_shape(
        1,  # msoShapeRectangle
        left, top, width, height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = RGBColor(*palette["accent_bright"])
    accent_bar.line.fill.background() # No outline

    # === Layer 3: Typography & Content ===
    
    # Title Text (Primary Dark Color)
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p_title = title_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Open Sans"
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(*palette["primary_dark"])

    # Body Text (Secondary Muted Color)
    body_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(8.0), Inches(2.5))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    p_body = body_frame.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Open Sans"
    p_body.font.size = Pt(24)
    p_body.font.color.rgb = RGBColor(*palette["secondary"])
    
    # Small "Palette Swatch" Footer to demonstrate the 5 colors (like the tutorial)
    swatch_width = Inches(0.8)
    swatch_height = Inches(0.8)
    swatch_start_x = Inches(1.5)
    swatch_y = Inches(6.0)
    
    for i, (role, color) in enumerate(palette.items()):
        swatch = slide.shapes.add_shape(
            1, # Rectangle
            swatch_start_x + (i * (swatch_width + Inches(0.1))), 
            swatch_y, 
            swatch_width, 
            swatch_height
        )
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = RGBColor(*color)
        swatch.line.color.rgb = RGBColor(*palette["primary_dark"]) # Slight dark border
        swatch.line.width = Pt(1)

    prs.save(output_pptx_path)
    return output_pptx_path
