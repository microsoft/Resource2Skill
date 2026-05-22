import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "POWERPOINT",
    subtitle_text: str = "工作计划 / 汇报总结 / 年中总结 / 述职报告等",
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide reproducing the 'Prismatic Green Geometric' style.
    Generates a custom low-poly background via PIL and lays out hexagonal content nodes.
    
    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Generate & Apply Prismatic Background (PIL) ===
    bg_path = "temp_prismatic_bg.png"
    width, height = 1920, 1080
    img = Image.new("RGBA", (width, height), (142, 201, 146, 255))
    draw = ImageDraw.Draw(img)

    # Define polygons to simulate a 3D faceted low-poly surface
    polygons = [
        # (Coordinates list), (R, G, B)
        ([(0,0), (1200,0), (800, 600), (0, 800)], (154, 222, 178)),
        ([(0,800), (800, 600), (550, 1080), (0, 1080)], (94, 187, 137)),
        ([(1200,0), (1920,0), (1920, 500), (1450, 350)], (175, 235, 195)),
        ([(800, 600), (1200, 0), (1450, 350), (1600, 1080), (550, 1080)], (120, 204, 156)),
        ([(1450, 350), (1920, 500), (1920, 1080), (1600, 1080)], (76, 168, 118)),
        ([(400, 400), (800, 600), (550, 1080), (200, 900)], (105, 195, 145)), # Extra overlay facet
    ]

    for coords, color in polygons:
        draw.polygon(coords, fill=color)
    
    img.save(bg_path)
    
    # Insert background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Main Title & Subtitle Typography ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(6.0), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial Black"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    sub_box = slide.shapes.add_textbox(Inches(1.1), Inches(3.6), Inches(6.0), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = "Microsoft YaHei"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 3: Geometric Content Nodes (Hexagons) ===
    # Creates a 3-item list on the right side using hexagonal nodes
    content_items = [
        "Focus on professional PPT template development and corporate visual design.",
        "Provide customized solutions for internal logic and data visualization.",
        "Committed to enhancing the communication value of every slide presentation."
    ]
    
    start_y = Inches(2.0)
    spacing = Inches(1.5)
    
    for i, text in enumerate(content_items):
        y_pos = start_y + (i * spacing)
        
        # Add Hexagon Node
        hex_shape = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(7.5), y_pos, Inches(1.0), Inches(1.0))
        hex_shape.fill.solid()
        hex_shape.fill.fore_color.rgb = RGBColor(255, 255, 255) # White Hexagon
        hex_shape.line.fill.background() # No line
        
        # Number inside hexagon
        hex_tf = hex_shape.text_frame
        hex_p = hex_tf.paragraphs[0]
        hex_p.text = str(i + 1)
        hex_p.alignment = PP_ALIGN.CENTER
        hex_p.font.name = "Arial"
        hex_p.font.size = Pt(24)
        hex_p.font.bold = True
        hex_p.font.color.rgb = RGBColor(94, 187, 137) # Green text matching background
        
        # Add body text next to hexagon
        body_box = slide.shapes.add_textbox(Inches(8.7), y_pos + Inches(0.1), Inches(3.8), Inches(0.8))
        body_tf = body_box.text_frame
        body_tf.word_wrap = True
        body_p = body_tf.paragraphs[0]
        body_p.text = text
        body_p.font.name = "Microsoft YaHei"
        body_p.font.size = Pt(14)
        body_p.font.color.rgb = RGBColor(255, 255, 255)

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
