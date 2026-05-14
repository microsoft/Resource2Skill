def create_slide(
    output_pptx_path: str,
    title_text: str = "RAW\nMATERIALS",
    body_text: str = "WHAT DO YOU PRESENT? A LITTLE OF EVERYTHING",
    bg_palette: str = "architecture",
    accent_color: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Architectural Minimalist Chapter Title' visual effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # Extract optional kwargs
    super_title = kwargs.get("super_title", "TYPICAL CLIENT MEETING")

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # ==========================================
    # Layer 1 & 2: Background Image & PIL Overlay
    # ==========================================
    bg_path = "temp_bg_arch.jpg"
    try:
        # Attempt to grab a contextual background
        url = f"https://picsum.photos/seed/{bg_palette}/1920/1080"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            with open(bg_path, 'wb') as f:
                f.write(response.read())
    except Exception:
        # Fallback: Create a sleek architectural slate gradient if download fails
        img = Image.new('RGB', (1920, 1080))
        draw = ImageDraw.Draw(img)
        for y in range(1080):
            r = int(25 - (15 * y / 1080))
            g = int(28 - (15 * y / 1080))
            b = int(32 - (15 * y / 1080))
            draw.line([(0, y), (1920, y)], fill=(r, g, b))
        img.save(bg_path)
        
    slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Create and apply the semi-transparent dark overlay mask via PIL
    overlay_path = "temp_overlay_mask.png"
    # Deep charcoal with ~65% opacity (160/255)
    overlay = Image.new('RGBA', (1920, 1080), (15, 17, 20, 160))
    overlay.save(overlay_path)
    slide.shapes.add_picture(overlay_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # ==========================================
    # Layer 3: Architectural Typography
    # ==========================================
    left_margin = Inches(1.2)

    # 3a. Super Title (Small, Light Grey, Heavily Tracked)
    if super_title:
        # Simulate tracking (letter-spacing)
        words = super_title.upper().split()
        tracked_words = [" ".join(list(w)) for w in words]
        tracked_super = "   ".join(tracked_words)
        
        tx_super = slide.shapes.add_textbox(left_margin, Inches(2.0), Inches(10), Inches(0.5))
        p_super = tx_super.text_frame.paragraphs[0]
        p_super.text = tracked_super
        p_super.font.name = 'Arial'
        p_super.font.size = Pt(11)
        p_super.font.bold = True
        p_super.font.color.rgb = RGBColor(190, 190, 190)
        p_super.alignment = PP_ALIGN.LEFT

    # 3b. Main Title (Massive, Tight Line Spacing)
    # Slight negative left offset to visually align large text bounding box with elements below
    tx_main = slide.shapes.add_textbox(left_margin - Inches(0.04), Inches(2.3), Inches(11), Inches(2.5))
    tf_main = tx_main.text_frame
    tf_main.word_wrap = True
    
    paragraphs = title_text.upper().split('\n')
    for i, line in enumerate(paragraphs):
        p_main = tf_main.paragraphs[0] if i == 0 else tf_main.add_paragraph()
        p_main.text = line
        p_main.font.name = 'Arial'
        p_main.font.size = Pt(88)
        p_main.font.bold = True
        p_main.font.color.rgb = RGBColor(*accent_color)
        p_main.alignment = PP_ALIGN.LEFT
        # Critical for the style: tight line spacing
        p_main.line_spacing = 0.85
        p_main.space_after = Pt(0)

    # 3c. Micro-Geometric Accent Line
    num_lines = len(paragraphs)
    # Dynamically position line based on title lines
    line_y = Inches(2.4) + (num_lines * Inches(1.15))
    
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left_margin, line_y, Inches(0.6), Pt(2.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(*accent_color)
    accent.line.fill.background() # Remove border

    # 3d. Subtitle
    tx_sub = slide.shapes.add_textbox(left_margin, line_y + Inches(0.15), Inches(10), Inches(1))
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text.upper()
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(12)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(230, 230, 230)
    p_sub.alignment = PP_ALIGN.LEFT
    p_sub.line_spacing = 1.2

    # Cleanup temporary files
    if os.path.exists(bg_path):
        os.remove(bg_path)
    if os.path.exists(overlay_path):
        os.remove(overlay_path)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
