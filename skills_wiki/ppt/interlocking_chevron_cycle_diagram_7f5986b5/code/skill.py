def create_slide(
    output_pptx_path: str,
    title_text: str = "Continuous Cycle Flow",
    body_text: str = "",
    bg_palette: str = "light",
    accent_color: tuple = (41, 128, 185),
    num_segments: int = 6,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Interlocking Chevron Cycle Diagram.
    Generates native vector shapes using FreeformBuilder for perfect editability.
    
    Returns: path to the saved PPTX file.
    """
    import math
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)

    # --- Configuration ---
    cx, cy = prs.slide_width / 2, prs.slide_height / 2
    R_out = Inches(2.8)    # Outer radius
    R_in = Inches(1.5)     # Inner radius
    R_mid = (R_out + R_in) / 2
    
    gap_deg = 2.0          # Gap between segments
    chevron_depth = 18.0   # How deep the chevron point goes

    # Vibrant categorical color palette
    colors = [
        RGBColor(41, 128, 185),  # Dark Blue
        RGBColor(211, 84, 0),    # Orange
        RGBColor(241, 196, 15),  # Yellow
        RGBColor(39, 174, 96),   # Green
        RGBColor(142, 68, 173),  # Purple
        RGBColor(52, 152, 219),  # Light Blue
        RGBColor(231, 76, 60),   # Red
        RGBColor(52, 73, 94)     # Dark Slate
    ]

    # Helper function for polar to cartesian coordinates (in Emus)
    def pt(r, angle_deg):
        # Subtract 90 to make 0 degrees point straight up (12 o'clock)
        a = math.radians(angle_deg - 90)
        return int(cx + r * math.cos(a)), int(cy + r * math.sin(a))

    # --- Draw the Interlocking Segments ---
    segment_angle = 360 / num_segments
    arc_steps = 20  # Number of lines to approximate the curve
    
    for i in range(num_segments):
        base_angle = i * segment_angle
        
        # Start and end angles for the arcs, accounting for the gap
        theta1 = base_angle + (gap_deg / 2)
        theta2 = base_angle + segment_angle - (gap_deg / 2)
        
        # 1. Initialize FreeformBuilder
        builder = slide.shapes.build_freeform()
        
        # Start at the inner tail corner
        builder.move_to(*pt(R_in, theta1))
        
        # 2. Draw Tail Cut (Chevron indentation)
        builder.add_line_segments([pt(R_mid, theta1 + chevron_depth)])
        builder.add_line_segments([pt(R_out, theta1)])
        
        # 3. Draw Outer Arc (approximate with small lines)
        outer_arc_pts = []
        for step in range(1, arc_steps + 1):
            fraction = step / arc_steps
            angle = theta1 + (theta2 - theta1) * fraction
            outer_arc_pts.append(pt(R_out, angle))
        builder.add_line_segments(outer_arc_pts)
        
        # 4. Draw Head Point (Chevron protrusion)
        builder.add_line_segments([pt(R_mid, theta2 + chevron_depth)])
        builder.add_line_segments([pt(R_in, theta2)])
        
        # 5. Draw Inner Arc (approximate backwards)
        inner_arc_pts = []
        for step in range(1, arc_steps + 1):
            fraction = step / arc_steps
            angle = theta2 - (theta2 - theta1) * fraction
            inner_arc_pts.append(pt(R_in, angle))
        builder.add_line_segments(inner_arc_pts)
        
        # 6. Convert to shape and style
        shape = builder.convert_to_shape()
        
        # Style the shape
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = colors[i % len(colors)]
        shape.line.fill.background() # Remove outline
        
        # --- Add Text Overlay ---
        # Calculate visual center (offset by half the chevron depth)
        text_angle = (theta1 + theta2) / 2 + (chevron_depth / 2)
        tx, ty = pt(R_mid, text_angle)
        
        # Add textbox
        tb_width, tb_height = Inches(1.5), Inches(0.8)
        textbox = slide.shapes.add_textbox(
            tx - tb_width / 2, 
            ty - tb_height / 2, 
            tb_width, 
            tb_height
        )
        tf = textbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        
        # Step Number
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        run1 = p1.add_run()
        run1.text = f"{i + 1}."
        run1.font.bold = True
        run1.font.size = Pt(20)
        run1.font.color.rgb = RGBColor(255, 255, 255)
        
        # Title text
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = "Lorem\nIpsum"
        run2.font.size = Pt(12)
        run2.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path

