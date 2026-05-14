def create_slide(
    output_pptx_path: str,
    title_text: str = "The Taiwanese identity is growing stronger",
    subtitle_text: str = "Taiwan, % of respondents identifying as:",
    source_text: str = "Source: Election Study Centre; The Economist",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Editorial Data Journalism (Economist) style.
    Generates a minimalist matplotlib chart and overlays it with signature editorial PPT elements.
    """
    import os
    import io
    import numpy as np
    import matplotlib.pyplot as plt
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    
    # --- Color Palette ---
    COLOR_RED = RGBColor(227, 18, 11)      # Editorial Red
    COLOR_TEXT = RGBColor(38, 38, 38)      # Charcoal
    COLOR_GRAY_LINE = RGBColor(160, 160, 160)
    
    # --- 1. Generate Editorial Matplotlib Chart ---
    # Create synthetic data representative of the tutorial's style
    years = np.arange(1992, 2024)
    # Trend 1: Rising (Highlight)
    taiwanese = 20 + 1.5 * (years - 1992) + np.random.normal(0, 3, len(years))
    # Trend 2: Falling (Context)
    chinese = 30 - 0.8 * (years - 1992) + np.random.normal(0, 2, len(years))
    # Trend 3: Flat/Middle (Context)
    both = 45 - 0.4 * (years - 1992) + np.random.normal(0, 2, len(years))

    # Matplotlib setup
    fig, ax = plt.subplots(figsize=(10, 5), dpi=300)
    
    # Plot lines
    ax.plot(years, both, color='#A0A0A0', linewidth=2.5, label='Taiwanese and Chinese')
    ax.plot(years, chinese, color='#C0C0C0', linewidth=2.5, label='Chinese')
    # The Highlight Line
    ax.plot(years, taiwanese, color='#E3120B', linewidth=3.5, label='Taiwanese')

    # Editorial Styling: Remove spines
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_visible(False)
    ax.spines['bottom'].set_color('#333333')

    # Editorial Styling: Horizontal grid only
    ax.yaxis.grid(True, color='#E0E0E0', linestyle='-', linewidth=1)
    ax.set_axisbelow(True) # Put grid behind lines

    # Editorial Styling: Ticks and Labels
    ax.tick_params(axis='both', which='both', length=0) # Remove actual tick marks
    ax.yaxis.tick_right() # Move Y axis to the right
    ax.set_yticks([0, 20, 40, 60, 80])
    ax.set_yticklabels(['0', '20', '40', '60', '80'], fontsize=12, color='#333333', weight='bold')
    
    # Custom X ticks to mimic the simplified style
    ax.set_xticks([1992, 1995, 2000, 2005, 2010, 2015, 2020])
    ax.set_xticklabels(['1992', '95', '2000', '05', '10', '15', '2020'], fontsize=12, color='#333333')

    # Direct labeling instead of legends
    ax.text(2023.5, taiwanese[-1], 'Taiwanese', color='#E3120B', weight='bold', fontsize=12, va='center')
    ax.text(2023.5, both[-1], 'Taiwanese and Chinese', color='#808080', weight='bold', fontsize=11, va='center')
    ax.text(2023.5, chinese[-1], 'Chinese', color='#A0A0A0', weight='bold', fontsize=11, va='center')

    # Save chart to memory
    img_stream = io.BytesIO()
    plt.savefig(img_stream, format='png', transparent=True, bbox_inches='tight')
    plt.close(fig)
    img_stream.seek(0)

    # --- 2. Setup PPTX ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- 3. Add Signature Header Elements ---
    # Signature thick red line
    red_line = slide.shapes.add_shape(
        1, # rectangle
        Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.08)
    )
    red_line.fill.solid()
    red_line.fill.fore_color.rgb = COLOR_RED
    red_line.line.fill.background() # No border
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(10), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(10), Inches(0.4))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(14)
    p_sub.font.bold = False
    p_sub.font.color.rgb = COLOR_TEXT

    # --- 4. Insert Chart ---
    # Insert the matplotlib image
    pic = slide.shapes.add_picture(
        img_stream, 
        Inches(0.6), Inches(2.2), 
        width=Inches(11.5)
    )

    # --- 5. Add Source Text ---
    source_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(10), Inches(0.4))
    tf_source = source_box.text_frame
    p_source = tf_source.paragraphs[0]
    p_source.text = source_text
    p_source.font.name = 'Arial'
    p_source.font.size = Pt(10)
    p_source.font.color.rgb = COLOR_GRAY_LINE

    # Save
    prs.save(output_pptx_path)
    return output_pptx_path
