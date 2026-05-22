def create_slide(
    output_pptx_path: str,
    title_text: str = "Troubleshooting Flowchart",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Semantic Logic Flowchart effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.oxml import OxmlElement

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Slide Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 247, 250)

    # === Title Elements ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 58, 138)

    # === Helper Functions ===
    
    def add_shadow(shape):
        """Injects XML to create a soft drop shadow."""
        spPr = shape.element.spPr
        effectLst = spPr.find(qn('a:effectLst'))
        if effectLst is None:
            effectLst = OxmlElement('a:effectLst')
            extLst = spPr.find(qn('a:extLst'))
            if extLst is not None:
                extLst.addprevious(effectLst)
            else:
                spPr.append(effectLst)
        
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', '50800')  # 4pt
        outerShdw.set('dist', '38100')     # 3pt
        outerShdw.set('dir', '2700000')    # 45 degrees
        outerShdw.set('algn', 'tl')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '15000')          # 15% opacity
        srgbClr.append(alpha)
        
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)

    def create_node(slide, shape_type, text, cx, cy, w, h, bg_color):
        """Creates a central logic node."""
        left = cx - w/2
        top = cy - h/2
        shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*bg_color)
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(1.5)
        
        add_shadow(shape)
        
        tf = shape.text_frame
        tf.text = text
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.bold = True
        return shape

    def draw_arrow_line(slide, x1, y1, x2, y2, color=(160, 164, 171)):
        """Draws a straight connector line and injects an arrowhead at the destination (tailEnd)."""
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 
            Inches(x1), Inches(y1), 
            Inches(x2), Inches(y2)
        )
        line.line.color.rgb = RGBColor(*color)
        line.line.width = Pt(2)
        
        # XML Injection for Arrowhead
        ln = line.element.spPr.ln
        if ln is not None:
            tailEnd = ln.find(qn('a:tailEnd'))
            if tailEnd is not None:
                tailEnd.set('type', 'triangle')
                tailEnd.set('w', 'med')
                tailEnd.set('len', 'med')
            else:
                new_tail = OxmlElement('a:tailEnd')
                new_tail.set('type', 'triangle')
                new_tail.set('w', 'med')
                new_tail.set('len', 'med')
                extLst = ln.find(qn('a:extLst'))
                if extLst is not None:
                    extLst.addprevious(new_tail)
                else:
                    ln.append(new_tail)
        return line

    def create_label(slide, text, cx, cy, w=0.6, h=0.25):
        """Creates a small inline path label."""
        left = cx - w/2
        top = cy - h/2
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(160, 164, 171)
        shape.line.width = Pt(1)
        
        tf = shape.text_frame
        tf.text = text
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(70, 80, 90)
        p.font.bold = True
        return shape

    # === Layout & Coordinates ===
    W, H = 2.4, 0.8
    CX_TRUNK = 4.5
    CX_BRANCH = 9.5
    
    # Vertically distribute 4 rows on the slide (from Y=1.26 to Y=6.24)
    Y_ROWS = [1.26, 2.92, 4.58, 6.24]
    
    # Theme Colors
    CLR_ROOT = (30, 58, 138)
    CLR_DECISION = (249, 115, 22)
    CLR_ACTION = (14, 165, 233)
    CLR_SUCCESS = (34, 197, 94)
    CLR_ERROR = (239, 68, 68)

    # === 1. Draw Connectors First (so they sit behind nodes/labels) ===
    # Trunk lines (Vertical)
    draw_arrow_line(slide, CX_TRUNK, Y_ROWS[0] + H/2, CX_TRUNK, Y_ROWS[1] - H/2)
    draw_arrow_line(slide, CX_TRUNK, Y_ROWS[1] + H/2, CX_TRUNK, Y_ROWS[2] - H/2)
    draw_arrow_line(slide, CX_TRUNK, Y_ROWS[2] + H/2, CX_TRUNK, Y_ROWS[3] - H/2)
    # Branch lines (Horizontal)
    draw_arrow_line(slide, CX_TRUNK + W/2, Y_ROWS[1], CX_BRANCH - W/2, Y_ROWS[1])
    draw_arrow_line(slide, CX_TRUNK + W/2, Y_ROWS[2], CX_BRANCH - W/2, Y_ROWS[2])

    # === 2. Draw Nodes ===
    # Row 0
    create_node(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Lamp doesn't work", CX_TRUNK, Y_ROWS[0], W, H, CLR_ROOT)
    
    # Row 1
    create_node(slide, MSO_SHAPE.DIAMOND, "Is it plugged in?", CX_TRUNK, Y_ROWS[1], W, H, CLR_DECISION)
    create_node(slide, MSO_SHAPE.RECTANGLE, "Plug in lamp", CX_BRANCH, Y_ROWS[1], W, H, CLR_SUCCESS)
    
    # Row 2
    create_node(slide, MSO_SHAPE.DIAMOND, "Bulb burned out?", CX_TRUNK, Y_ROWS[2], W, H, CLR_DECISION)
    create_node(slide, MSO_SHAPE.RECTANGLE, "Replace bulb", CX_BRANCH, Y_ROWS[2], W, H, CLR_SUCCESS)
    
    # Row 3
    create_node(slide, MSO_SHAPE.RECTANGLE, "Repair lamp", CX_TRUNK, Y_ROWS[3], W, H, CLR_ERROR)

    # === 3. Draw Path Labels ===
    # Vertical Yes/No Labels
    create_label(slide, "YES", CX_TRUNK, (Y_ROWS[1] + Y_ROWS[2]) / 2)
    create_label(slide, "NO", CX_TRUNK, (Y_ROWS[2] + Y_ROWS[3]) / 2)
    
    # Horizontal Yes/No Labels
    create_label(slide, "NO", (CX_TRUNK + W/2 + CX_BRANCH - W/2) / 2, Y_ROWS[1])
    create_label(slide, "YES", (CX_TRUNK + W/2 + CX_BRANCH - W/2) / 2, Y_ROWS[2])

    prs.save(output_pptx_path)
    return output_pptx_path
