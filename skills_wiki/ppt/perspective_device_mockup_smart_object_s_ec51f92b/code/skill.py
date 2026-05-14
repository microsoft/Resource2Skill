def create_slide(
    output_pptx_path: str,
    title_text: str = "Digital Ecosystem",
    body_text: str = "Seamlessly deploy your solutions across all platforms with responsive, pixel-perfect accuracy.",
    bg_palette: str = "technology",
    accent_color: tuple = (0, 191, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Perspective Device Mockup (Smart Object) effect.
    Returns: path to the saved PPTX file.
    """
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    import urllib.request
    import io
    import os

    # 1. Setup Base Canvas
    W, H = 1280, 720
    base_img = Image.new('RGBA', (W, H), (245, 247, 250, 255))
    draw = ImageDraw.Draw(base_img)

    # 2. Draw Environment (Stylized Desk)
    # Gradient desk
    desk_start = H // 2 + 50
    for y in range(desk_start, H):
        ratio = (y - desk_start) / (H - desk_start)
        c = int(240 - ratio * 45)
        draw.line([(0, y), (W, y)], fill=(c, c+3, c+8, 255))
    
    # Subtle drop shadow for the laptop base
    draw.polygon([(100, 660), (900, 580), (1100, 680), (200, 760)], fill=(180, 185, 190, 120))

    # 3. Draw Laptop Geometry (Viewed from front-left angle)
    # Lid / Bezel Outer
    bezel_corners = [(150, 120), (800, 180), (800, 600), (150, 680)]
    draw.polygon(bezel_corners, fill=(40, 42, 45, 255))
    draw.line(bezel_corners + [bezel_corners[0]], fill=(180, 185, 190, 255), width=3)

    # Keyboard Base
    base_corners = [(150, 680), (800, 600), (1050, 680), (250, 780)]
    clipped_base = [(x, min(y, H)) for x, y in base_corners] # Clip to bottom
    draw.polygon(clipped_base, fill=(200, 205, 210, 255))
    
    # Keyboard Well (Darker inner polygon)
    kbd_corners = [(220, 680), (760, 615), (950, 650), (320, 730)]
    clipped_kbd = [(x, min(y, H)) for x, y in kbd_corners]
    draw.polygon(clipped_kbd, fill=(100, 105, 110, 255))

    # Screen Display Area (Inner corners)
    screen_corners = [(170, 145), (780, 200), (780, 580), (170, 650)]
    draw.polygon(screen_corners, fill=(10, 10, 10, 255))

    # 4. Fetch or Generate Content Image (The "Smart Object" Replacement)
    content_img = None
    try:
        url = f"https://images.unsplash.com/photo-1498050108023-c5249f4df085?auto=format&fit=crop&w=1000&q=80"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            content_img = Image.open(io.BytesIO(response.read())).convert('RGBA')
    except Exception:
        # Fallback: Create a vibrant gradient screen if network fails
        content_img = Image.new('RGBA', (1000, 600), (0, 0, 0, 255))
        cdraw = ImageDraw.Draw(content_img)
        for i in range(600):
            r = int(accent_color[0] * (i/600))
            g = int(accent_color[1] * (i/600))
            b = int(accent_color[2] + (255 - accent_color[2]) * (1 - i/600))
            cdraw.line([(0, i), (1000, i)], fill=(r, g, min(b,255), 255))

    # 5. Calculate Perspective Homography (Solving A*c = B)
    cw, ch = content_img.size
    pa = screen_corners  # Destination polygon on canvas
    pb = [(0, 0), (cw, 0), (cw, ch), (0, ch)]  # Source original corners

    matrix = []
    for p1, p2 in zip(pa, pb):
        matrix.append([p1[0], p1[1], 1, 0, 0, 0, -p2[0]*p1[0], -p2[0]*p1[1]])
        matrix.append([0, 0, 0, p1[0], p1[1], 1, -p2[1]*p1[0], -p2[1]*p1[1]])

    A = np.matrix(matrix, dtype=np.float64)
    B = np.array(pb).reshape(8)
    # Solve linear equations for the 8 perspective coefficients
    coeffs = np.linalg.solve(A, B).reshape(8).tolist()

    # 6. Apply Perspective Warp
    warped_img = content_img.transform((W, H), Image.PERSPECTIVE, coeffs, Image.BICUBIC)

    # 7. Mask and Composite
    mask = Image.new('L', (W, H), 0)
    ImageDraw.Draw(mask).polygon(screen_corners, fill=255)
    # Antialiasing the mask edges slightly
    mask = mask.filter(ImageFilter.GaussianBlur(0.5))
    base_img.paste(warped_img, (0, 0), mask)

    # 8. Add Glass Reflection (Gloss)
    glare = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    glare_corners = [
        screen_corners[0], 
        (450, screen_corners[1][1] - 10), 
        (450, 615), 
        screen_corners[3]
    ]
    ImageDraw.Draw(glare).polygon(glare_corners, fill=(255, 255, 255, 12))
    base_img = Image.alpha_composite(base_img, glare)

    # Save Composite Image
    img_path = "mockup_composite.png"
    base_img.save(img_path)

    # 9. Build PPTX Slide
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Insert background mockup
    slide.shapes.add_picture(img_path, Inches(0), Inches(0), Inches(13.333), Inches(7.5))

    # Add Typography on the right
    tx_box = slide.shapes.add_textbox(Inches(8.5), Inches(2.2), Inches(4.2), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(38)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(30, 35, 45)
    p_title.font.name = "Arial"

    p_body = tf.add_paragraph()
    p_body.text = "\n" + body_text
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(90, 100, 110)
    p_body.font.name = "Arial"

    # Add a sleek accent line under title
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(8.5), Inches(3.0), Inches(0.8), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    line.line.fill.background()

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(img_path):
        os.remove(img_path)
        
    return output_pptx_path
