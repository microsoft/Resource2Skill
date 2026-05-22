def create_slide(
    output_pptx_path: str,
    title_text: str = "Table of Contents",
    bg_palette: str = "interior architecture",  # Keyword for background image
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Glassmorphic Section Hub visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFilter
    import urllib.request
    import io

    # 1. Setup Presentation (16:9 Aspect Ratio)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Screen dimensions for PIL processing (matching 13.333x7.5 inches at 144 DPI)
    CANVAS_W, CANVAS_H = 1920, 1080
    DPI = 144

    # 2. Fetch or Generate Background Image
    try:
        url = f"https://source.unsplash.com/random/1920x1080/?{bg_palette.replace(' ', ',')}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img_sharp = Image.open(io.BytesIO(response.read())).convert("RGBA")
            img_sharp = img_sharp.resize((CANVAS_W, CANVAS_H), Image.Resampling.LANCZOS)
    except Exception:
        # Fallback Gradient Background if download fails
        img_sharp = Image.new("RGBA", (CANVAS_W, CANVAS_H), (20, 30, 40, 255))
        draw = ImageDraw.Draw(img_sharp)
        for y in range(CANVAS_H):
            r = int(20 + (40 * y / CANVAS_H))
            g = int(30 + (50 * y / CANVAS_H))
            b = int(40 + (30 * y / CANVAS_H))
            draw.line([(0, y), (CANVAS_W, y)], fill=(r, g, b, 255))

    # 3. Create Heavily Blurred Background for the Glass effect
    img_blurred = img_sharp.filter(ImageFilter.GaussianBlur(radius=40))

    # 4. Insert Sharp Background into Slide
    bg_stream = io.BytesIO()
    img_sharp.save(bg_stream, format="PNG")
    bg_stream.seek(0)
    slide.shapes.add_picture(bg_stream, 0, 0, prs.slide_width, prs.slide_height)

    # 5. Define Bubble Layout (Staggered Grid)
    # Row 1 (5 bubbles), Row 2 (4 bubbles)
    bubble_radius = 130
    sections = [
        {"title": "Our Mission", "cx": 360, "cy": 450},
        {"title": "Solution", "cx": 660, "cy": 450},
        {"title": "Business\nModel", "cx": 960, "cy": 450},
        {"title": "Competition", "cx": 1260, "cy": 450},
        {"title": "Our Team", "cx": 1560, "cy": 450},
        
        {"title": "Problem", "cx": 510, "cy": 750},
        {"title": "Market\nPotential", "cx": 810, "cy": 750},
        {"title": "Growth\nStrategy", "cx": 1110, "cy": 750},
        {"title": "Financials", "cx": 1410, "cy": 750},
    ]

    # 6. Generate and Insert Glass Bubbles
    for sec in sections:
        cx, cy, r = sec["cx"], sec["cy"], bubble_radius
        
        # a. Crop localized area from blurred background
        bbox = (cx - r, cy - r, cx + r, cy + r)
        bubble_bg = img_blurred.crop(bbox).convert("RGBA")
        
        # b. Create sharp circular mask
        mask = Image.new("L", (2*r, 2*r), 0)
        draw_mask = ImageDraw.Draw(mask)
        draw_mask.ellipse((0, 0, 2*r, 2*r), fill=255)
        
        # Apply mask to keep outside transparent
        bubble_bg.putalpha(mask)
        
        # c. Create Frost Layer (Inner Shadow)
        frost = Image.new("RGBA", (2*r, 2*r), (0, 0, 0, 0))
        draw_frost = ImageDraw.Draw(frost)
        # Draw thick semi-transparent white ring
        draw_frost.ellipse((0, 0, 2*r, 2*r), outline=(255, 255, 255, 160), width=int(r*0.2))
        # Blur it to create the inner glow/shadow effect
        frost = frost.filter(ImageFilter.GaussianBlur(int(r*0.15)))
        
        # d. Composite Frost over Blurred Background
        bubble = Image.alpha_composite(bubble_bg, frost)
        
        # e. Re-apply mask to clean up frost bleed outside the circle
        bubble.putalpha(mask)
        
        # f. Draw crisp outer border
        draw_final = ImageDraw.Draw(bubble)
        draw_final.ellipse((1, 1, 2*r-1, 2*r-1), outline=(255, 255, 255, 220), width=2)
        
        # g. Save bubble to memory and insert to PPTX
        bubble_stream = io.BytesIO()
        bubble.save(bubble_stream, format="PNG")
        bubble_stream.seek(0)
        
        # Calculate positioning in Inches based on DPI
        left = Inches((cx - r) / DPI)
        top = Inches((cy - r) / DPI)
        size = Inches((2 * r) / DPI)
        
        slide.shapes.add_picture(bubble_stream, left, top, size, size)
        
        # h. Overlay Editable Text Box
        txBox = slide.shapes.add_textbox(left, top, size, size)
        tf = txBox.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = sec["title"]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(255, 255, 255)

    # 7. Add Main Title
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.8), Inches(9.333), Inches(1.5))
    tf_title = title_box.text_frame
    tf_title.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial"
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
