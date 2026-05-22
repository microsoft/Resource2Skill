import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from lxml import etree

def add_drop_shadow(shape):
    """
    Injects Open XML to add a subtle drop shadow to a python-pptx shape.
    """
    spPr = shape.element.spPr
    effectLst = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    outerShdw = etree.SubElement(effectLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw")
    outerShdw.set("blurRad", "150000") # Blur radius
    outerShdw.set("dist", "100000")    # Distance
    outerShdw.set("dir", "2700000")    # Direction (angle)
    outerShdw.set("algn", "tl")
    
    srgbClr = etree.SubElement(outerShdw, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    srgbClr.set("val", "000000")       # Black shadow
    alpha = etree.SubElement(srgbClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
    alpha.set("val", "50000")          # 50% opacity

def create_slide(
    output_pptx_path: str,
    active_menu: str = "Go diagrams",
    code_snippet: str = 'with Diagram("Web Services"):\n    dns = Route53("dns")\n    lb = ELB("lb")\n    dns >> lb',
    accent_color: tuple = (138, 43, 226), # Neon Purple
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Dark Mode Split-View Topology" effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Colors ===
    bg_color = RGBColor(38, 38, 41)
    terminal_color = RGBColor(15, 15, 15)
    text_white = RGBColor(240, 240, 240)
    text_grey = RGBColor(160, 160, 160)
    keyword_color = RGBColor(255, 152, 0) # Orange for python keywords
    string_color = RGBColor(165, 214, 255) # Light blue for strings

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background() # No border

    # === Layer 2: Sidebar Navigation ===
    menus = ["Diagrams", "Go diagrams", "Mermaid", "PlantUML", "ASCII editors"]
    start_y = 2.0
    for menu in menus:
        # Create pill shape
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(0.5), Inches(start_y), Inches(2.0), Inches(0.6)
        )
        pill.adjustments[0] = 0.5 # Max roundness for pill shape
        
        # Style based on active state
        if menu == active_menu:
            pill.fill.solid()
            pill.fill.fore_color.rgb = RGBColor(*accent_color)
            pill.line.fill.background()
            text_color = text_white
        else:
            pill.fill.solid()
            pill.fill.fore_color.rgb = RGBColor(60, 60, 65)
            pill.line.fill.background()
            text_color = text_grey

        # Add text
        text_frame = pill.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = "   " + menu
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = text_color
        
        start_y += 0.8

    # === Layer 3: Terminal Window (Code) ===
    term_width = Inches(4.5)
    term_height = Inches(3.0)
    term = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(3.0), Inches(1.5), term_width, term_height
    )
    term.adjustments[0] = 0.05 # Slight rounding
    term.fill.solid()
    term.fill.fore_color.rgb = terminal_color
    term.line.fill.background()
    add_drop_shadow(term) # Apply lxml drop shadow

    # Add Code Text with basic highlighting simulation
    tf = term.text_frame
    tf.clear()
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    
    for line in code_snippet.split('\n'):
        p = tf.add_paragraph()
        p.space_after = Pt(4)
        
        # Extremely basic syntax highlighting logic for demonstration
        words = line.split(" ")
        for i, word in enumerate(words):
            run = p.add_run()
            # Restore spacing
            if i > 0: run.text = " " + word
            else: run.text = word
            
            run.font.name = "Consolas"
            run.font.size = Pt(12)
            
            if "with" in word or "from" in word or "import" in word:
                run.font.color.rgb = keyword_color
            elif '"' in word:
                run.font.color.rgb = string_color
            else:
                run.font.color.rgb = text_white

    # === Layer 4: Topology Diagram ===
    # Node 1: DNS
    n1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(3.5), Inches(1.2), Inches(1.2))
    n1.fill.solid()
    n1.fill.fore_color.rgb = RGBColor(138, 43, 226) # Purple
    n1.line.color.rgb = text_white
    n1.line.width = Pt(1.5)
    n1.text = "DNS"
    n1.text_frame.paragraphs[0].font.name = "Arial"
    n1.text_frame.paragraphs[0].font.size = Pt(14)

    # Node 2: Load Balancer
    n2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(3.5), Inches(1.2), Inches(1.2))
    n2.fill.solid()
    n2.fill.fore_color.rgb = RGBColor(65, 105, 225) # Royal Blue
    n2.line.color.rgb = text_white
    n2.line.width = Pt(1.5)
    n2.text = "LB"
    n2.text_frame.paragraphs[0].font.name = "Arial"
    n2.text_frame.paragraphs[0].font.size = Pt(14)

    # Connect Nodes
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.2), Inches(4.1), Inches(10.5), Inches(4.1))
    connector.line.color.rgb = text_white
    connector.line.width = Pt(2.5)
    
    # Add arrowhead via lxml to connector
    line_pr = connector.element.spPr.ln
    headEnd = etree.SubElement(line_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd")
    tailEnd = etree.SubElement(line_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tailEnd.set("type", "triangle")
    tailEnd.set("w", "med")
    tailEnd.set("len", "med")

    prs.save(output_pptx_path)
    return output_pptx_path

