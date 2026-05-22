def create_slide(
    output_pptx_path: str,
    title_text: str = "Focus Area: Newspaper Excerpt",
    body_text: str = "",
    bg_palette: str = "newspaper",
    accent_color: tuple = (220, 38, 38),  # Deep Red border
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Focal Zoom & Fade Overlay' visual effect.
    """
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # ---------------------------------------------------------
    # 1. Fetch Base Context Image (or generate fallback)
    # ---------------------------------------------------------
    try:
        # Attempt to get a realistic document/newspaper layout
        url = f"https://image.pollinations.ai/prompt/newspaper%20document%20columns?width=1600&height=900&nologo=true"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGB")
    except Exception:
        # Fallback: Generate a structural mockup of a newspaper
        base_img = Image.new('RGB', (1600, 900), color=(245, 245, 245))
        draw = ImageDraw.Draw(base_img)
        col_width = 450
        for i in range(3):
            x = 100 + i * (col_width + 50)
            draw.rectangle([x, 100, x + col_width, 150], fill=(80, 80, 80)) # Header
            y_start = 180
            if i == 1: # Middle column has an 'image' to zoom into
                draw.rectangle([x, 180, x + col_width, 420], fill=(150, 170, 190))
                y_start = 450
            for y in range(y_start, 800, 45): # Text lines
                line_width = col_width if y % 135 != 0 else col_width - 80
                draw.line([x, y, x + line_width, y], fill=(180, 180, 180), width=18)

    # ---------------------------------------------------------
    # 2. Add Base Image to Slide (Fit to slide dimensions)
    # ---------------------------------------------------------
    slide_w_emu = prs.slide_width
    slide_h_emu = prs.slide_height

    img_ratio = base_img.width / base_img.height
    slide_ratio = slide_w_emu / slide_h_emu

    if img_ratio > slide_ratio:
        fit_w = slide_w_emu
        fit_h = int(slide_w_emu / img_ratio)
    else:
        fit_h = slide_h_emu
        fit_w = int(slide_h_emu * img_ratio)

    pic_left = int((slide_w_emu - fit_w) / 2)
    pic_top = int((slide_h_emu - fit_h) / 2)

    # Save to BytesIO for pptx insertion
    base_io = BytesIO()
    base_img.save(base_io, format='JPEG', quality=90)
    base_io.seek(0)
    slide.shapes.add_picture(base_io, pic_left, pic_top, fit_w, fit_h)

    # ---------------------------------------------------------
    # 3. Add Semi-Transparent "Dimmer" Overlay
    # ---------------------------------------------------------
    # Creates the "fade" effect to mute the background
    overlay_img = Image.new('RGBA', (100, 100), (255, 255, 255, 170)) # ~66% opacity white
    overlay_io = BytesIO()
    overlay_img.save(overlay_io, format='PNG')
    overlay_io.seek(0)
    slide.shapes.add_picture(overlay_io, pic_left, pic_top, fit_w, fit_h)

    # ---------------------------------------------------------
    # 4. Crop, Scale, and Insert the Focal Zoom Area
    # ---------------------------------------------------------
    # Define ROI as percentages (targeting upper middle area)
    rx, ry, rw, rh = 0.35, 0.20, 0.30, 0.35 
    
    crop_box = (
        int(rx * base_img.width),
        int(ry * base_img.height),
        int((rx + rw) * base_img.width),
        int((ry + rh) * base_img.height)
    )
    cropped_img = base_img.crop(crop_box)

    # Scale the focal image up
    zoom_factor = 1.6
    zoom_w_emu = int(rw * fit_w * zoom_factor)
    zoom_h_emu = int(rh * fit_h * zoom_factor)

    crop_io = BytesIO()
    cropped_img.save(crop_io, format='PNG')
    crop_io.seek(0)

    # Calculate placement so it scales outward from its original center
    center_x = pic_left + (rx + rw/2) * fit_w
    center_y = pic_top + (ry + rh/2) * fit_h

    zoom_left = int(center_x - zoom_w_emu / 2)
    zoom_top = int(center_y - zoom_h_emu / 2)

    zoom_pic = slide.shapes.add_picture(crop_io, zoom_left, zoom_top, zoom_w_emu, zoom_h_emu)

    # ---------------------------------------------------------
    # 5. Apply Border and Shadow Formatting
    # ---------------------------------------------------------
    # Border
    zoom_pic.line.color.rgb = RGBColor(*accent_color)
    zoom_pic.line.width = Pt(5)

    # Deep Drop Shadow via lxml
    shadow_xml = """
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:outerShdw blurRad="250000" dist="100000" dir="5400000" algn="b" rotWithShape="0">
            <a:srgbClr val="000000">
                <a:alpha val="45000"/>
            </a:srgbClr>
        </a:outerShdw>
    </a:effectLst>
    """
    zoom_pic._element.spPr.append(parse_xml(shadow_xml))

    # ---------------------------------------------------------
    # 6. Add Explanatory Title (Optional Context)
    # ---------------------------------------------------------
    if title_text:
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(10), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 41, 59) # Dark slate

        # Add a subtle background to the title for readability over the image
        fill = title_box.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(255, 255, 255)
        title_box.line.color.rgb = RGBColor(*accent_color)
        title_box.line.width = Pt(2)

    prs.save(output_pptx_path)
    return output_pptx_path
