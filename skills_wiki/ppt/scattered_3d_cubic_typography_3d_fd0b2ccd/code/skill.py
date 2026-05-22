def create_slide(
    output_pptx_path: str,
    title_text: str = "翻轉立方",
    bg_color: tuple = (255, 214, 0),       # Bright Yellow background
    block_color: tuple = (255, 192, 0),    # Golden Yellow block
    text_color: tuple = (40, 40, 40),      # Dark text
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Scattered 3D Cubic Typography" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import random
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import OxmlElement
    
    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # 2. Set Background Color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)
    
    # Helper function to inject 3D XML into a shape
    def apply_3d_effect(shape, depth_pt=80, camera_preset="isometricTopUp", material="plastic"):
        spPr = shape.element.spPr
        
        # Define Scene 3D (Camera Angle & Lighting)
        scene3d = OxmlElement('a:scene3d')
        
        # Camera
        camera = OxmlElement('a:camera')
        camera.set('prst', camera_preset)
        scene3d.append(camera)
        
        # Lighting Rig
        lightRig = OxmlElement('a:lightRig')
        lightRig.set('rig', 'threePt')
        lightRig.set('dir', 't')
        scene3d.append(lightRig)
        
        spPr.append(scene3d)
        
        # Define Shape 3D (Extrusion / Depth)
        # 1 pt = 12700 EMUs
        extrusion_emu = int(depth_pt * 12700)
        sp3d = OxmlElement('a:sp3d')
        sp3d.set('extrusionH', str(extrusion_emu))
        sp3d.set('prstMaterial', material)
        
        # Add a subtle bevel to make edges catch light
        bevelT = OxmlElement('a:bevelT')
        bevelT.set('w', '38100') # 3pt
        bevelT.set('h', '38100')
        bevelT.set('prst', 'circle')
        sp3d.append(bevelT)
        
        spPr.append(sp3d)

    # 3. Define 3D Camera Presets for scattering effect
    camera_presets = [
        "isometricTopUp",
        "isometricRightUp",
        "isometricLeftUp",
        "perspectiveContrastingLeftFacing",
        "perspectiveContrastingRightFacing",
        "obliqueTopLeft",
        "obliqueTopRight"
    ]
    
    # 4. Generate 3D Blocks for each character
    num_chars = len(title_text)
    if num_chars == 0:
        title_text = "3D文字"
        num_chars = len(title_text)
        
    block_size = Inches(1.8)
    
    # Calculate starting position to roughly center the group
    start_x = (prs.slide_width - (num_chars * block_size * 1.2)) / 2
    base_y = Inches(3.0)
    
    for i, char in enumerate(title_text):
        # Add slight randomness to layout
        x_offset = start_x + (i * block_size * 1.2) + Inches(random.uniform(-0.2, 0.2))
        y_offset = base_y + Inches(random.uniform(-0.5, 0.5))
        
        # Create shape (Rounded Rectangle)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x_offset, y_offset, block_size, block_size
        )
        
        # Format Shape Color
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*block_color)
        shape.line.color.rgb = RGBColor(*block_color) # Match line to fill
        
        # Format Text
        text_frame = shape.text_frame
        text_frame.text = char
        text_frame.word_wrap = False
        p = text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        
        font = p.font
        font.name = 'Arial' # Best standard font for bold blocks
        font.size = Pt(64)
        font.bold = True
        font.color.rgb = RGBColor(*text_color)
        
        # Adjust text margins so it centers properly
        text_frame.margin_left = Inches(0)
        text_frame.margin_right = Inches(0)
        text_frame.margin_top = Inches(0)
        text_frame.margin_bottom = Inches(0)
        
        # Apply 3D Extrusion and Random Rotation
        preset = random.choice(camera_presets)
        apply_3d_effect(shape, depth_pt=80, camera_preset=preset, material="metal")

    # 5. Add a subtle secondary descriptive text box
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.33), Inches(0.5))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Generated 3D Typographic Elements using OXML Injection"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(100, 100, 100)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("3d_cubic_text.pptx", title_text="翻轉吧!")
