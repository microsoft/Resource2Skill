def create_slide(
    output_pptx_path: str,
    title_text: str = "Karen Walter",
    body_text: str = "Presentation Expert",
    bg_palette: str = "professional portrait", 
    accent_color: tuple = (0, 0, 0), 
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Minimalist Floating Identity Badge 
    and a final high-contrast Call-to-Action slide.

    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # ==========================================
    # SLIDE 1: Floating Identity Badge Profile
    # ==========================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # --- Layer 1: Background Image ---
    image_url = "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?q=80&w=1920&auto=format&fit=crop"
    img_path = "temp_bg_portrait.jpg"
    
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
        slide1.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception:
        # Fallback if download fails
        fallback_bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        fallback_bg.fill.solid()
        fallback_bg.fill.fore_color.rgb = RGBColor(80, 85, 90)
        fallback_bg.line.fill.background()

    # --- Layer 2: Floating Identity Badge ---
    badge_width = Inches(3.8)
    badge_height = Inches(0.8)
    badge_left = Inches(1.0)
    badge_top = Inches(4.5)
    
    badge = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        badge_left, badge_top, badge_width, badge_height
    )
    
    # Adjust corner radius to be slightly rounded (15%)
    badge.adjustments[0] = 0.15
    badge.line.fill.background() # Remove outline
    
    # Apply base color
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(248, 249, 250)
    
    # Inject XML for Transparency and Drop Shadow
    spPr = badge.element.spPr
    
    # Add Alpha (90% opacity)
    srgbClr = spPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    if srgbClr is not None:
        etree.SubElement(srgbClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", val="90000")
    
    # Add Drop Shadow
    effectLst = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    outerShdw = etree.SubElement(effectLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw", 
                                 blurRad="60000", dist="30000", dir="5400000", algn="tl", rotWithShape="0")
    shdwClr = etree.SubElement(outerShdw, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", val="000000")
    etree.SubElement(shdwClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", val="20000") # 20% opacity shadow

    # Add Name Text to Badge
    tf = badge.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title_text
    run.font.name = "Arial"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 30, 30)
    
    # --- Layer 3: Floating Subtitle (Role) ---
    if body_text:
        role_box = slide1.shapes.add_textbox(
            badge_left, badge_top + badge_height, 
            badge_width, Inches(0.6)
        )
        tf_role = role_box.text_frame
        tf_role.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_role = tf_role.paragraphs[0]
        run_role = p_role.add_run()
        run_role.text = body_text.upper()
        run_role.font.name = "Arial"
        run_role.font.size = Pt(13)
        run_role.font.bold = True
        run_role.font.color.rgb = RGBColor(255, 255, 255)
        
        # Add subtle text shadow to ensure subtitle pops against light backgrounds
        rPr = run_role._r.get_or_add_rPr()
        effectLst_txt = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        outerShdw_txt = etree.SubElement(effectLst_txt, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw", 
                                     blurRad="30000", dist="10000", dir="5400000")
        shdwClr_txt = etree.SubElement(outerShdw_txt, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", val="000000")
        etree.SubElement(shdwClr_txt, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", val="60000")

    # ==========================================
    # SLIDE 2: High-Contrast Call to Action
    # (Capturing the final frame of the video)
    # ==========================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Solid Navy Background
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(0, 0, 80) # Deep Navy Blue
    bg2.line.fill.background()
    
    # Centered CTA Text
    cta_box = slide2.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2.5))
    cta_tf = cta_box.text_frame
    cta_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p2 = cta_tf.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    
    run2a = p2.add_run()
    run2a.text = f"Book {title_text} for your event\n"
    run2a.font.name = "Arial"
    run2a.font.size = Pt(40)
    run2a.font.color.rgb = RGBColor(255, 255, 255)
    
    run2b = p2.add_run()
    run2b.text = "www.PowerfulPresentations.nl"
    run2b.font.name = "Arial"
    run2b.font.size = Pt(40)
    run2b.font.bold = True
    run2b.font.color.rgb = RGBColor(255, 255, 255)

    # Clean up temp files
    if os.path.exists(img_path):
        os.remove(img_path)
        
    prs.save(output_pptx_path)
    return output_pptx_path
