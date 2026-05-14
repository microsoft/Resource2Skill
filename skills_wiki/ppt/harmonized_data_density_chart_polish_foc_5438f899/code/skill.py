def create_slide(
    output_pptx_path: str,
    title_text: str = "Annual Performance Review: Data Highlights",
    body_text: str = "",
    bg_palette: str = "corporate",  
    accent_color: tuple = (255, 87, 34),  # Vibrant Orange for highlighting
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Harmonized Data Density chart styles.
    Features an optimized Column Chart (gap widths, contrast highlighting) 
    and an optimized Doughnut Chart (monochromatic gradient palette).
    """
    import os
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Slide Title ===
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 40, 50)
    
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(11), Inches(0.5))
    subtitle_frame = subtitle_box.text_frame
    p2 = subtitle_frame.paragraphs[0]
    p2.text = "Optimized spacing, contrast highlighting, and tonal harmony"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(120, 130, 140)

    # === Visual Effect 1: Optimized Column Chart ===
    # Strategy: Gap width 50%, subtle base color, single vibrant highlight point
    
    chart_data = CategoryChartData()
    chart_data.categories = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct']
    chart_data.add_series('Revenue', (40, 55, 45, 70, 95, 60, 50, 75, 80, 55))

    x, y, cx, cy = Inches(0.8), Inches(2.2), Inches(6.5), Inches(4.5)
    column_chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart

    # Format the Column Chart
    plot = column_chart.plots[0]
    plot.has_data_labels = True
    
    # CRITICAL: Adjust gap width to make bars dominant (Video rule: Gap = 1/2 Bar width)
    plot.gap_width = 50  

    series = plot.series[0]
    highlight_index = 4  # Highlight 'May'
    
    # Loop through points to apply Contrast Color Logic
    for idx, point in enumerate(series.points):
        fill = point.format.fill
        fill.solid()
        if idx == highlight_index:
            # Highlight Color
            fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
            point.data_label.font.bold = True
            point.data_label.font.size = Pt(14)
            point.data_label.font.color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
        else:
            # Subdued Base Color
            fill.fore_color.rgb = RGBColor(200, 210, 225)
            point.data_label.font.size = Pt(12)
            point.data_label.font.color.rgb = RGBColor(150, 160, 175)

    # Clean axes
    val_axis = column_chart.value_axis
    val_axis.has_major_gridlines = True
    val_axis.major_gridlines.format.line.color.rgb = RGBColor(230, 235, 240)
    val_axis.visible = False # Hide Y axis line, keep gridlines
    
    cat_axis = column_chart.category_axis
    cat_axis.tick_labels.font.size = Pt(12)
    cat_axis.tick_labels.font.color.rgb = RGBColor(100, 110, 120)
    cat_axis.major_tick_mark = XL_TICK_MARK.NONE

    column_chart.has_legend = False

    # === Visual Effect 2: Monochromatic Doughnut Chart ===
    # Strategy: Replace default colors with a calculated monochromatic gradient scale
    
    doughnut_data = CategoryChartData()
    doughnut_data.categories = ['North America', 'Europe', 'Asia', 'South America', 'Africa']
    doughnut_data.add_series('Market Share', (34, 30, 15, 15, 6))

    x2, y2, cx2, cy2 = Inches(7.8), Inches(2.2), Inches(4.5), Inches(4.5)
    doughnut_chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, x2, y2, cx2, cy2, doughnut_data
    ).chart

    doughnut_chart.has_legend = True
    doughnut_chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    doughnut_chart.legend.font.size = Pt(12)

    plot2 = doughnut_chart.plots[0]
    plot2.has_data_labels = True
    
    # Calculate Monochromatic Gradient Scale (e.g., Deep Teal to Light Mint)
    # Base: Teal (15, 100, 90) -> Lighten incrementally
    base_r, base_g, base_b = 15, 100, 90
    step_r, step_g, step_b = 40, 30, 30
    
    series2 = plot2.series[0]
    for idx, point in enumerate(series2.points):
        fill = point.format.fill
        fill.solid()
        
        # Calculate stepped color
        r = min(255, base_r + (step_r * idx))
        g = min(255, base_g + (step_g * idx))
        b = min(255, base_b + (step_b * idx))
        
        fill.fore_color.rgb = RGBColor(r, g, b)
        
        # Clean labels
        point.data_label.font.size = Pt(12)
        point.data_label.font.bold = True
        point.data_label.font.color.rgb = RGBColor(255, 255, 255) # White text for contrast
        # Note: PPTX engine handles label placement, usually fits inside large slices

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

