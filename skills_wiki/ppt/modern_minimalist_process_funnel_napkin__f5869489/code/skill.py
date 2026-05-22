def create_slide(
    output_pptx_path: str,
    title_text: str = "Process Methodology & Pipeline",
    body_text: str = "", # Not used directly in this custom layout
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Modern Minimalist Process Funnel (Napkin-Style).
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # --- Background ---
    # Clean white/off-white background for modern AI-tool aesthetic
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(250, 250, 250)

    # --- Add Slide Title ---
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = 'Arial'
    p.font.color.rgb = RGBColor(40, 40, 40)

    # --- Funnel Data & Palette ---
    stages = [
        {"title": "Phase 1: Broad Sourcing", "desc": "Initial data gathering and wide-net capture across all available channels.", "color": (255, 218, 185)}, # Peach
        {"title": "Phase 2: Qualification", "desc": "Filtering candidates based on baseline criteria and preliminary automated scoring.", "color": (230, 230, 250)}, # Lavender
        {"title": "Phase 3: Deep Evaluation", "desc": "In-depth analysis, structured interviews, and technical competency mapping.", "color": (152, 251, 152)}, # Mint
        {"title": "Phase 4: Final Selection", "desc": "Final review board alignment, offer generation, and successful onboarding.", "color": (135, 206, 235)},  # Sky Blue
    ]

    # --- Funnel Math & Geometry ---
    center_x = Inches(4.0)
    start_y = Inches(2.0)
    stage_height = Inches(1.1)
    
    top_width = Inches(5.0)
    bottom_width = Inches(1.5)
    
    num_stages = len(stages)
    width_decrement = (top_width - bottom_width) / num_stages

    # --- Draw Funnel & Text ---
    current_y = start_y
    current_width = top_width

    for i, stage in enumerate(stages):
        next_width = current_width - width_decrement
        
        # Calculate the 4 corners of the trapezoid for this specific stage
        # Top Left, Top Right, Bottom Right, Bottom Left
        x1 = center_x - (current_width / 2)
        y1 = current_y
        x2 = center_x + (current_width / 2)
        y2 = current_y
        x3 = center_x + (next_width / 2)
        y3 = current_y + stage_height
        x4 = center_x - (next_width / 2)
        y4 = current_y + stage_height

        # Build Custom Freeform Shape (guarantees perfect stacking)
        builder = slide.shapes.build_freeform(x1, y1)
        builder.add_line_segments((
            (x2, y2),
            (x3, y3),
            (x4, y4),
            (x1, y1) # Close the polygon
        ))
        funnel_segment = builder.convert_to_shape()
        
        # Style the Funnel Segment
        funnel_segment.fill.solid()
        funnel_segment.fill.fore_color.rgb = RGBColor(*stage["color"])
        funnel_segment.line.fill.background() # No border line

        # --- Add Explanatory Text Box ---
        text_x = center_x + Inches(3.0) # Anchor to the right
        text_y = current_y + (stage_height / 4)
        
        tx_box = slide.shapes.add_textbox(text_x, text_y, Inches(5.0), stage_height)
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        # Title Paragraph
        p_title = tf.paragraphs[0]
        p_title.text = stage["title"]
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = RGBColor(50, 50, 50)
        
        # Description Paragraph
        p_desc = tf.add_paragraph()
        p_desc.text = stage["desc"]
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = RGBColor(100, 100, 100)
        p_desc.space_before = Pt(4)

        # --- Draw Connecting Line ---
        # Draw a thin grey line from the middle-right edge of the funnel segment to the text box
        mid_right_x = center_x + ((current_width + next_width) / 4)
        mid_y = current_y + (stage_height / 2)
        
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 
            mid_right_x + Inches(0.2), mid_y,  # Start slightly away from the funnel
            text_x - Inches(0.2), mid_y        # End slightly before the text
        )
        connector.line.color.rgb = RGBColor(200, 200, 200)
        connector.line.width = Pt(1.5)

        # Advance Y and Width for next loop
        current_y += stage_height + Inches(0.05) # Add a tiny 0.05 gap for modern slice aesthetic
        current_width = next_width

    # Save the presentation
    prs.save(output_pptx_path)
    return output_pptx_path
