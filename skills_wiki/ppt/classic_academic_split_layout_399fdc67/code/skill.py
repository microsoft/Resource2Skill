def create_slide(
    output_pptx_path: str,
    title_text: str = "Red Delicious",
    body_text: str = "Season: Sep - Jan\nFlavor: Mild, sweet\nUses: Snacking\nCost: Low",
    bg_palette: str = "red apple",  # keyword for the right-pane image
    accent_color: tuple = (45, 30, 20),  # Deep Espresso text
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Classic Academic Split-Layout.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageOps
    import urllib.request
    import io
    import os

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 2. Setup Background (Warm Ivory)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(238, 235, 226)

    # 3. Create Left Pane Content (Text)
    # Title
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.0), Inches(1.5))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Georgia"  # Classic serif
    p_title.font.size = Pt(64)
    p_title.font.bold = False
    p_title.font.color.rgb = RGBColor(*accent_color)
    p_title.alignment = PP_ALIGN.LEFT

    # Body Details (Key-Value style)
    tx_body = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(5.0), Inches(3.0))
    tf_body = tx_body.text_frame
    tf_body.word_wrap = True
    
    lines = body_text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf_body.paragraphs[0]
        else:
            p = tf_body.add_paragraph()
        p.text = line
        p.font.name = "Georgia"
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(70, 60, 50)  # Slightly softer than title
        p.space_after = Pt(12)

    # 4. Generate Right Pane Image (50% Split)
    # Target dimensions: Width = 6.666", Height = 7.5"
    target_w_px = int(6.666 * 300) # 300 dpi approx
    target_h_px = int(7.5 * 300)
    
    img_path = "temp_split_img.jpg"
    try:
        # Fetch image based on keyword
        url = f"https://source.unsplash.com/featured/{target_w_px}x{target_h_px}/?{urllib.parse.quote(bg_palette)}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img_data = response.read()
            
        img = Image.open(io.BytesIO(img_data)).convert("RGB")
        # Crop exactly to the 50% split aspect ratio to prevent PPTX distortion
        img_cropped = ImageOps.fit(img, (target_w_px, target_h_px), Image.Resampling.LANCZOS)
        img_cropped.save(img_path)
    except Exception as e:
        # Fallback: Generate a textured/colored block if network fails
        print(f"Image fetch failed, generating fallback: {e}")
        img_fallback = Image.new("RGB", (target_w_px, target_h_px), (200, 190, 180))
        draw = ImageDraw.Draw(img_fallback)
        # Add subtle pattern
        for y in range(0, target_h_px, 40):
            draw.line([(0, y), (target_w_px, y)], fill=(210, 200, 190), width=2)
        img_fallback.save(img_path)

    # 5. Insert Image into Right Pane
    slide.shapes.add_picture(
        img_path, 
        Inches(6.666), 
        Inches(0), 
        width=Inches(6.667), 
        height=Inches(7.5)
    )

    # Cleanup temp file
    if os.path.exists(img_path):
        os.remove(img_path)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
