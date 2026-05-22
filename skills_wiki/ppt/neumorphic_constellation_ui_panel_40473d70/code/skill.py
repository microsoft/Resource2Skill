def create_slide(
    output_pptx_path: str,
    title_text: str = "EASY UI\nANIMATION",
    subtitle_text: str = "POWERPOINT",
    tagline_text: str = "MOTION GRAPHIC",
    bg_color: tuple = (28, 34, 55),
    accent_blue: tuple = (29, 78, 216),
    accent_red: tuple = (239, 68, 68),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neumorphic Constellation UI effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import OxmlElement
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Helper: Set Solid Fill
    def set_solid_fill(shape, rgb_tuple):
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*rgb_tuple)

    # Helper: Add Soft Shadow via lxml
    def add_soft_shadow(shape):
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        # 61pt blur, 18pt distance, 351 deg angle, 41% opacity
        outerShdw.set('blurRad', str(int(61 * 12700)))
        outerShdw.set('dist', str(int(18 * 12700)))
        outerShdw.set('dir', str(int(351 * 60000))) 
        outerShdw.set('algn', 'tl')
        outerShdw.set('rotWithShape', '0')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '41000') # 41%
        srgbClr.append(alpha)
        
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_solid_fill(bg, bg_color)
    bg.line.fill.background()

    # === Layer 2: Constellation & U-Arc ===
    # 2a. The "Subtracted" U-Arc illusion (Donut + Mask)
    arc_color = (48, 62, 90)
    donut = slide.shapes.add_shape(
        MSO_SHAPE.DONUT, Inches(5.5), Inches(1.5), Inches(5), Inches(5)
    )
    set_solid_fill(donut, bg_color) # Core is BG color
    donut.line.color.rgb = RGBColor(*arc_color)
    donut.line.width = Pt(15)
    
    # Masking rectangle to cover the top half of the donut (simulating boolean subtract)
    mask = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5.0), Inches(0), Inches(6), Inches(4.0)
    )
    set_solid_fill(mask, bg_color)
    mask.line.fill.background()

    # 2b. Constellation Points & Lines
    pts = [
        (7.5, 1.5), (9.0, 1.0), (8.2, 2.5), (10.0, 2.0),
        (7.0, 3.5), (8.5, 4.0), (9.8, 3.2), (11.0, 4.5),
        (10.5, 1.2), (7.8, 4.8)
    ]
    edges = [
        (0,1), (0,2), (1,2), (1,3), (2,3), (2,4), (2,5), 
        (3,6), (4,5), (5,6), (6,7), (3,8), (1,8), (5,9), (4,9)
    ]

    for (p1, p2) in edges:
        line = slide.shapes.add_connector(
            MSO_SHAPE.LINE, 
            Inches(pts[p1][0]), Inches(pts[p1][1]), 
            Inches(pts[p2][0]), Inches(pts[p2][1])
        )
        line.line.color.rgb = RGBColor(*arc_color)
        line.line.width = Pt(1)

    for pt in pts:
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(pt[0]-0.05), Inches(pt[1]-0.05), Inches(0.1), Inches(0.1)
        )
        set_solid_fill(node, arc_color)
        node.line.fill.background()

    # === Layer 3: Main Elevated Panel ===
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2.0), Inches(1.5), Inches(4.5), Inches(4.5)
    )
    set_solid_fill(panel, bg_color)
    panel.line.fill.background()
    add_soft_shadow(panel) # This creates the Neumorphic pop

    # === Layer 4: Typography & UI Accents ===
    # Text helper
    def add_styled_text(x, y, w, h, text, size_pt, color_tuple, is_italic=True, is_bold=True):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Segoe UI Black"
        p.font.size = Pt(size_pt)
        p.font.bold = is_bold
        p.font.italic = is_italic
        p.font.color.rgb = RGBColor(*color_tuple)
        return tb

    add_styled_text(Inches(2.5), Inches(2.2), Inches(4.0), Inches(1), title_text, 44, (255, 255, 255))
    add_styled_text(Inches(2.5), Inches(3.8), Inches(4.0), Inches(0.8), subtitle_text, 28, (255, 255, 255))
    
    # Tagline with a grey/slate color
    tag = add_styled_text(Inches(2.5), Inches(6.3), Inches(4.0), Inches(0.5), f"◆  {tagline_text}  ◆", 18, (156, 163, 175), is_italic=True)
    tag.text_frame.paragraphs[0].font.name = "Arial"

    # Follow Button
    btn = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(5.0), Inches(2.2), Inches(0.8)
    )
    set_solid_fill(btn, accent_blue)
    btn.line.fill.background()
    tf = btn.text_frame
    tf.text = "FOLLOW"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.name = "Segoe UI Black"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # Decorative Right Arrows
    for i in range(3):
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(7.0 + i*0.4), Inches(5.25), Inches(0.2), Inches(0.3)
        )
        arrow.rotation = 90
        set_solid_fill(arrow, (255, 255, 255))
        arrow.line.fill.background()

    # Decorative scattered diamonds
    diamonds = [
        (1.0, 1.0, 0.15, (255, 255, 255)),
        (1.8, 2.5, 0.1, accent_red),
        (8.5, 6.5, 0.1, accent_red),
        (11.0, 1.5, 0.1, accent_red),
        (11.5, 6.0, 0.15, (255, 255, 255)),
        (8.5, 0.5, 0.1, (255, 255, 255)),
    ]
    for (x, y, size, clr) in diamonds:
        d = slide.shapes.add_shape(
            MSO_SHAPE.DIAMOND, Inches(x), Inches(y), Inches(size), Inches(size)
        )
        set_solid_fill(d, clr)
        d.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
