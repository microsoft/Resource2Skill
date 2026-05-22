def create_slide(
    output_pptx_path: str,
    title_text: str = "System Flow Animation",
    body_text: str = "Visualizing continuous energy transfer through animated components",
    bg_color: tuple = (15, 23, 42),
    wire_color: tuple = (0, 200, 200),
    arrow_color: tuple = (0, 255, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Continuous Flow Animated Diagram.
    Generates seamless looping animated GIFs via PIL and embeds them.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from PIL import Image, ImageDraw, ImageFilter

    # --- Helper: Generate Seamless Looping GIF ---
    def create_arrow_gif(filename, w_px, h_px, spacing=100, frames=20, fps=25):
        imgs = []
        for i in range(frames):
            img = Image.new('RGBA', (w_px, h_px), (255, 255, 255, 0))
            glow = Image.new('RGBA', (w_px, h_px), (255, 255, 255, 0))
            sharp = Image.new('RGBA', (w_px, h_px), (255, 255, 255, 0))
            
            g_draw = ImageDraw.Draw(glow)
            s_draw = ImageDraw.Draw(sharp)
            
            # Shift calculates the micro-movement for this specific frame
            shift = (i / frames) * spacing
            
            # Draw enough arrows to cover canvas + margins for a seamless loop jump
            for j in range(-2, (w_px // spacing) + 2):
                x = j * spacing + shift
                y = h_px / 2
                # Arrow polygon geometry
                points = [(x, y - 12), (x + 20, y), (x, y + 12)]
                
                # Draw translucent glow and sharp core
                g_draw.polygon(points, fill=(arrow_color[0], arrow_color[1], arrow_color[2], 100))
                s_draw.polygon(points, fill=(arrow_color[0], arrow_color[1], arrow_color[2], 255))
                
            glow = glow.filter(ImageFilter.GaussianBlur(3))
            img = Image.alpha_composite(img, glow)
            img = Image.alpha_composite(img, sharp)
            imgs.append(img)
            
        # disposal=2 forces GIF to clear background each frame (critical for transparency)
        imgs[0].save(filename, save_all=True, append_images=imgs[1:], duration=1000//fps, loop=0, disposal=2)

    # 1. Prepare dynamic GIF assets tailored to our edge lengths
    gif_h = "temp_flow_h.gif"
    gif_v = "temp_flow_v.gif"
    # Widths must be strict multiples of 'spacing' (100) to ensure mathematical seamlessness
    create_arrow_gif(gif_h, 700, 60) # For horizontal lines
    create_arrow_gif(gif_v, 400, 60) # For vertical lines

    # 2. Base Presentation Setup
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*bg_color)
    bg.line.color.rgb = RGBColor(*bg_color) # Hide outline

    # Typography
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(1))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    p1.text = title_text
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(160, 170, 180)

    # 3. Circuit Pathway (Layer 2)
    cx, cy = Inches(13.333 / 2), Inches(4.25)
    w_circ, h_circ = Inches(8), Inches(5)
    left = cx - w_circ / 2
    top = cy - h_circ / 2

    wire = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w_circ, h_circ)
    wire.fill.background()
    wire.line.color.rgb = RGBColor(*wire_color)
    wire.line.width = Pt(3)

    # 4. Inject Flow Animations (Layer 3)
    # Top edge (Moves Right -> 0 deg)
    slide.shapes.add_picture(gif_h, cx - Inches(3.5), top - Inches(0.3), width=Inches(7), height=Inches(0.6))
    
    # Right edge (Moves Down -> 90 deg)
    pic_right = slide.shapes.add_picture(gif_v, left + w_circ - Inches(2), cy - Inches(0.3), width=Inches(4), height=Inches(0.6))
    pic_right.rotation = 90
    
    # Bottom edge (Moves Left -> 180 deg)
    pic_bottom = slide.shapes.add_picture(gif_h, cx - Inches(3.5), top + h_circ - Inches(0.3), width=Inches(7), height=Inches(0.6))
    pic_bottom.rotation = 180
    
    # Left edge (Moves Up -> 270 deg)
    pic_left = slide.shapes.add_picture(gif_v, left - Inches(2), cy - Inches(0.3), width=Inches(4), height=Inches(0.6))
    pic_left.rotation = 270

    # 5. Component Blocks (Layer 4 - Covers the GIFs so arrows seamlessly emerge)
    def add_block(x_c, y_c, text, outline_color):
        w, h = Inches(1.6), Inches(2.2)
        blk = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_c - w/2, y_c - h/2, w, h)
        blk.fill.solid()
        blk.fill.fore_color.rgb = RGBColor(20, 30, 50)
        blk.line.color.rgb = outline_color
        blk.line.width = Pt(2)
        
        btf = blk.text_frame
        btf.text = text
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.font.color.rgb = RGBColor(255, 255, 255)
        bp.font.bold = True
        bp.font.size = Pt(14)
        return blk

    add_block(left, cy, "BATTERY\nSOURCE", RGBColor(0, 255, 255))
    add_block(left + w_circ, cy, "SYSTEM\nLOAD", RGBColor(255, 50, 100))

    # Clean up temporary generation artifacts
    try:
        os.remove(gif_h)
        os.remove(gif_v)
    except OSError:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path
