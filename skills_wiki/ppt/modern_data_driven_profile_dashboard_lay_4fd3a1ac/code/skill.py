def create_slide(
    output_pptx_path: str,
    title_text: str = "Dashboard Overview",
    body_text: str = "A wonderful serenity has taken possession of my entire soul, optimizing our metrics for speed and massive scale.",
    bg_palette: str = "business,technology",
    accent_color: tuple = (0, 229, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Modern Data-Driven Dashboard Layout'.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Colors
    bg_color = (20, 24, 33)
    secondary_accent = (255, 50, 150)
    track_color = (50, 55, 65)

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.format.fill.solid()
    bg.format.fill.fore_color.rgb = RGBColor(*bg_color)
    bg.format.line.fill.background()

    # === Layer 2: Circular Hero Image (PIL Masking) ===
    img_size = 500
    temp_img_path = "temp_circular_hero.png"
    try:
        url = f"https://source.unsplash.com/random/800x800/?{bg_palette}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception:
        # Fallback if download fails
        img = Image.new('RGBA', (800, 800), (40, 50, 70, 255))
        draw = ImageDraw.Draw(img)
        draw.text((300, 400), "Image Offline", fill="white")

    # Center crop to square
    w, h = img.size
    min_dim = min(w, h)
    left = (w - min_dim)/2
    top = (h - min_dim)/2
    img = img.crop((left, top, left+min_dim, top+min_dim))
    img = img.resize((img_size, img_size), Image.Resampling.LANCZOS)

    # High-quality anti-aliased circular mask
    mask_size = (img_size * 3, img_size * 3)
    mask = Image.new('L', mask_size, 0)
    draw = ImageDraw.Draw(mask)
    draw.ellipse((0, 0, mask_size[0], mask_size[1]), fill=255)
    mask = mask.resize((img_size, img_size), Image.Resampling.LANCZOS)

    circular_img = Image.new('RGBA', (img_size, img_size), (0, 0, 0, 0))
    circular_img.paste(img, (0, 0), mask)
    circular_img.save(temp_img_path)

    # Insert circular image on the right side
    slide.shapes.add_picture(temp_img_path, Inches(7.5), Inches(1.25), width=Inches(5.0))

    # === Layer 3: Text & Content ===
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.5), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Body
    body_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.0), Inches(1.0))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(180, 180, 180)

    # === Layer 4: Data Widgets (Progress Rings) ===
    def add_progress_ring(x_pos, y_pos, size_in, percentage, color, label):
        # 1. Add Doughnut Chart
        chart_data = CategoryChartData()
        chart_data.categories = ['Value', 'Remaining']
        chart_data.add_series('Data', (percentage, 100 - percentage))

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, x_pos, y_pos, size_in, size_in, chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = False
        
        # Match background to slide to make it blend seamlessly
        chart.chart_area.format.fill.solid()
        chart.chart_area.format.fill.fore_color.rgb = RGBColor(*bg_color)
        chart.chart_area.format.line.fill.background()
        
        # Color the slices
        try:
            series = chart.series[0]
            pt1 = series.points[0]
            pt1.format.fill.solid()
            pt1.format.fill.fore_color.rgb = RGBColor(*color)
            
            pt2 = series.points[1]
            pt2.format.fill.solid()
            pt2.format.fill.fore_color.rgb = RGBColor(*track_color)
        except Exception:
            pass # Failsafe for version differences

        # 2. Add Center Text Box
        tx_box = slide.shapes.add_textbox(x_pos, y_pos, size_in, size_in)
        tf_ring = tx_box.text_frame
        tf_ring.text = f"{percentage}%"
        tf_ring.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_ring.margin_left = tf_ring.margin_right = tf_ring.margin_top = tf_ring.margin_bottom = 0
        p_ring = tf_ring.paragraphs[0]
        p_ring.alignment = PP_ALIGN.CENTER
        p_ring.font.size = Pt(28)
        p_ring.font.bold = True
        p_ring.font.color.rgb = RGBColor(255, 255, 255)

        # 3. Add Label below
        lbl_box = slide.shapes.add_textbox(x_pos, y_pos + size_in - Inches(0.2), size_in, Inches(0.5))
        tf_lbl = lbl_box.text_frame
        tf_lbl.text = label
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.font.size = Pt(14)
        p_lbl.font.color.rgb = RGBColor(180, 180, 180)

    # Add two progress rings side-by-side
    add_progress_ring(Inches(1.0), Inches(4.2), Inches(2.2), 76, accent_color, "Conversion")
    add_progress_ring(Inches(3.8), Inches(4.2), Inches(2.2), 42, secondary_accent, "Retention")

    # === Layer 5: Floating UI Card ===
    # Overlaps the image slightly to create depth
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(5.5), Inches(3.5), Inches(1.2))
    card.format.fill.solid()
    card.format.fill.fore_color.rgb = RGBColor(30, 35, 45) # Slightly lighter than BG
    card.format.line.solid()
    card.format.line.color.rgb = RGBColor(*accent_color)
    card.format.line.width = Pt(1.5)
    
    tf_card = card.text_frame
    tf_card.margin_left = Inches(0.2)
    p_card1 = tf_card.paragraphs[0]
    p_card1.text = "High Performance"
    p_card1.font.size = Pt(16)
    p_card1.font.bold = True
    p_card1.font.color.rgb = RGBColor(255, 255, 255)
    
    p_card2 = tf_card.add_paragraph()
    p_card2.text = "System optimized for scale."
    p_card2.font.size = Pt(12)
    p_card2.font.color.rgb = RGBColor(180, 180, 180)

    prs.save(output_pptx_path)
    
    # Cleanup temp file
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
        
    return output_pptx_path
