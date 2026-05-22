def create_slide(
    output_pptx_path: str,
    title_text: str = "Controls Explorer",
    header_color: tuple = (103, 58, 183),  # Deep Purple
    bg_color: tuple = (245, 245, 245),     # Light Gray
    **kwargs,
) -> str:
    """
    Creates a slide featuring a modern application dashboard UI grid layout with elevated cards.
    """
    import collections
    import collections.abc
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Helper Function: Add Subtle UI Drop Shadow ---
    def apply_ui_shadow(shape):
        spPr = shape._element.spPr
        shadow_xml = """
            <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <a:outerShdw blurRad="40000" dist="30000" dir="5400000" algn="b" rotWithShape="0">
                    <a:srgbClr val="000000">
                        <a:alpha val="10000"/>
                    </a:srgbClr>
                </a:outerShdw>
            </a:effectLst>
        """
        spPr.append(parse_xml(shadow_xml))

    # --- Layer 1: Slide Background ---
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(*bg_color)
    bg_shape.line.fill.background()

    # --- Layer 2: Header Ribbon ---
    header_height = Inches(1.2)
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, header_height
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(*header_color)
    header.line.fill.background()

    # Header Text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(10), Inches(0.8))
    tf = txBox.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Segoe UI" # Modern UI font
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # --- Layer 3: Card Grid Content ---
    # Sample data representing dashboard items
    card_data = [
        {"title": "Chart", "desc": "Plot over 30 chart types ranging from line charts to specialized financial charts.", "badge": "Updated"},
        {"title": "DataGrid", "desc": "Powerful grid control with advanced features like grouping, sorting, filtering and export to excel.", "badge": None},
        {"title": "PdfViewer", "desc": "High performance PDF Viewer component with features like search, zooming and text selection.", "badge": "New"},
        {"title": "ListView", "desc": "Advanced ListView component with features like grid layout, grouping, pull-to-refresh and selection.", "badge": None},
        {"title": "Schedule", "desc": "The Schedule control is used to schedule and manage the appointments through an intuitive user interface.", "badge": None},
        {"title": "ComboBox", "desc": "Allows users to type a value or choose an option from a list of predefined options.", "badge": None},
    ]

    # Grid settings
    cols = 3
    margin_x = Inches(0.6)
    margin_y = header_height + Inches(0.5)
    spacing_x = Inches(0.4)
    spacing_y = Inches(0.4)
    
    available_width = prs.slide_width - (margin_x * 2)
    card_width = (available_width - (spacing_x * (cols - 1))) / cols
    card_height = Inches(2.2)

    for i, data in enumerate(card_data):
        row = i // cols
        col = i % cols
        x = margin_x + col * (card_width + spacing_x)
        y = margin_y + row * (card_height + spacing_y)

        # 1. Card Container (White with Shadow)
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.fill.background()
        apply_ui_shadow(card)

        # 2. Simulated Icon Area (Light gray square)
        icon_size = Inches(0.5)
        icon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y + Inches(0.2), icon_size, icon_size)
        icon.fill.solid()
        icon.fill.fore_color.rgb = RGBColor(230, 230, 235) # Soft icon placeholder color
        icon.line.fill.background()
        
        # Adjust corner radius of icon via XML for modern look
        adjLst = icon._element.xpath('.//a:prstGeom/a:avLst')[0]
        parse_xml('<a:gd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="adj" fmla="val 20000"/>')
        adjLst.append(parse_xml('<a:gd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="adj" fmla="val 20000"/>'))

        # 3. Card Title
        text_x = x + Inches(0.9)
        title_box = slide.shapes.add_textbox(text_x, y + Inches(0.15), card_width - Inches(1.0), Inches(0.4))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(33, 33, 33)

        # 4. Card Description
        desc_box = slide.shapes.add_textbox(text_x, y + Inches(0.6), card_width - Inches(1.1), Inches(1.2))
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = data["desc"]
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = RGBColor(90, 90, 90)

        # 5. Status Badge (Optional)
        if data["badge"]:
            badge_width = Inches(0.8)
            badge_height = Inches(0.25)
            badge_x = x + card_width - badge_width - Inches(0.2)
            badge_y = y + Inches(0.2)
            
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, badge_x, badge_y, badge_width, badge_height)
            badge.fill.solid()
            badge.fill.fore_color.rgb = RGBColor(76, 175, 80) # Success Green
            badge.line.fill.background()
            
            # Badge text
            btf = badge.text_frame
            btf.margin_top = btf.margin_bottom = btf.margin_left = btf.margin_right = 0
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            bp.text = data["badge"]
            bp.font.name = "Segoe UI"
            bp.font.size = Pt(10)
            bp.font.bold = True
            bp.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
