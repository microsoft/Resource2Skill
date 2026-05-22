def create_slide(
    output_pptx_path: str,
    title_text: str = "Company Values",
    body_text: str = "Respect\n\nWhile we represent a significant range of suppliers, we recognise that each supplier relationship is unique and a privilege to uphold.\n\nService\n\nA cornerstone of our business and we must always strive for improvement.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Vibrant Geometric Color-Block Overlay' effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    
    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Add a blank slide
    slide_layout = prs.slide_layouts[6] # 6 is usually a completely blank layout
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Dark Slate Background ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(76, 92, 104)
    bg_shape.line.fill.background() # Remove border

    # === Layer 2: Mustard Yellow Skewed Block ===
    # Using FreeformBuilder to create the dynamic angled background on the right
    builder = slide.shapes.build_freeform(Inches(4.5), Inches(0))
    builder.add_line_segments([
        (Inches(13.333), Inches(0)),
        (Inches(13.333), Inches(7.5)),
        (Inches(2.5), Inches(7.5)),
        (Inches(4.5), Inches(0))
    ], close=True)
    yellow_block = builder.convert_to_shape()
    yellow_block.fill.solid()
    yellow_block.fill.fore_color.rgb = RGBColor(249, 188, 36)
    yellow_block.line.fill.background()

    # === Layer 3: Vibrant Teal Triangle ===
    # A large downward pointing triangle overlaying the background and yellow block
    builder_tri = slide.shapes.build_freeform(Inches(3.0), Inches(0))
    builder_tri.add_line_segments([
        (Inches(8.5), Inches(0)),
        (Inches(6.0), Inches(6.5)),
        (Inches(3.0), Inches(0))
    ], close=True)
    teal_tri = builder_tri.convert_to_shape()
    teal_tri.fill.solid()
    teal_tri.fill.fore_color.rgb = RGBColor(33, 178, 166)
    teal_tri.line.fill.background()

    # === Layer 4: Title Banner (Folded Ribbon Effect) ===
    # Small shadow/fold element
    fold = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(0.4), Inches(2.3), Inches(0.6), Inches(0.5)
    )
    fold.fill.solid()
    fold.fill.fore_color.rgb = RGBColor(50, 60, 70) # Darker shadow color
    fold.line.fill.background()
    
    # Main White Banner
    banner = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.2), Inches(4.2), Inches(1.2)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(255, 255, 255)
    banner.line.fill.solid()
    banner.line.fore_color.rgb = RGBColor(33, 178, 166)
    banner.line.width = Pt(3)

    # Title Text inside Banner
    title_tf = banner.text_frame
    title_tf.clear()
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text
    title_p.alignment = PP_ALIGN.CENTER
    title_font = title_p.font
    title_font.name = 'Arial' # Standard sans-serif
    title_font.size = Pt(36)
    title_font.bold = True
    title_font.color.rgb = RGBColor(76, 92, 104)

    # === Layer 5: Body Content ===
    # Placed dynamically over the yellow area
    content_box = slide.shapes.add_textbox(
        Inches(6.5), Inches(1.5), Inches(6.0), Inches(5.0)
    )
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    # Parse body text to simulate bold headers for paragraphs
    paragraphs = body_text.split('\n\n')
    for i, para_text in enumerate(paragraphs):
        p = text_frame.add_paragraph() if i > 0 else text_frame.paragraphs[0]
        
        # Simple heuristic: if it's a short single word/phrase, treat as header
        is_header = len(para_text.split()) <= 3 and not para_text.endswith('.')
        
        p.text = para_text
        p.font.name = 'Arial'
        p.font.size = Pt(24) if is_header else Pt(18)
        p.font.bold = is_header
        p.font.color.rgb = RGBColor(30, 30, 30) # Dark grey for readability
        
        if not is_header:
            p.space_after = Pt(14) # Add spacing after body paragraphs

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

