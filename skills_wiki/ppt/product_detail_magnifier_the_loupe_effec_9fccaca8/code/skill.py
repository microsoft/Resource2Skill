def create_slide(
    output_pptx_path: str,
    title_text: str = "PREMIUM CRAFTSMANSHIP",
    body_text: str = "Notice the intricate stitching and premium materials used in the armrest. Every detail is engineered for maximum comfort and durability, providing an unmatched user experience.",
    image_url: str = "https://images.unsplash.com/photo-1555041469-a586c61ea9bc?q=80&w=1920&auto=format&fit=crop", # Default: a nice sofa
    detail_center_x_pct: float = 0.25,  # X coordinate of the detail to zoom (0.0 to 1.0)
    detail_center_y_pct: float = 0.65,  # Y coordinate of the detail to zoom (0.0 to 1.0)
    zoom_factor: float = 2.0,           # How much to magnify the detail
    loupe_radius_px: int = 300,         # Size of the magnifying circle
    accent_color: tuple = (255, 255, 255, 255), # Border color of the loupe
):
    """
    Creates a PPTX file demonstrating the "Product Detail Magnifier (Loupe)" effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Pt, Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageEnhance, ImageDraw
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # 1. Fetch or generate background image
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception as e:
        print(f"Failed to download image, using fallback. Error: {e}")
        base_img = Image.new("RGBA", (1920, 1080), (40, 45, 55, 255))
        draw = ImageDraw.Draw(base_img)
        draw.rectangle([400, 400, 800, 800], fill=(80, 85, 95, 255))
        draw.text((450, 600), "Product Image Fallback", fill=(255, 255, 255, 255))

    # 2. Crop base image to exactly 16:9 to prevent distortion in PPT
    target_ratio = 13.333 / 7.5
    img_ratio = base_img.width / base_img.height
    if img_ratio > target_ratio:
        new_w = int(base_img.height * target_ratio)
        offset = (base_img.width - new_w) // 2
        base_img = base_img.crop((offset, 0, offset + new_w, base_img.height))
    elif img_ratio < target_ratio:
        new_h = int(base_img.width / target_ratio)
        offset = (base_img.height - new_h) // 2
        base_img = base_img.crop((0, offset, base_img.width, offset + new_h))
        
    base_img = base_img.resize((1920, 1080), Image.Resampling.LANCZOS)
    width, height = base_img.size
    
    # 3. Create the Darkened Background
    enhancer = ImageEnhance.Brightness(base_img)
    bg_dark = enhancer.enhance(0.35) # Darken to 35% brightness
    
    bg_dark_io = BytesIO()
    bg_dark.save(bg_dark_io, format="PNG")
    bg_dark_io.seek(0)
    
    # Insert dark background into slide
    slide.shapes.add_picture(bg_dark_io, 0, 0, width=prs.slide_width, height=prs.slide_height)
    
    # 4. Create the Magnified Loupe
    detail_px_x = int(width * detail_center_x_pct)
    detail_px_y = int(height * detail_center_y_pct)
    
    # Calculate crop box (before zoom)
    crop_radius = int(loupe_radius_px / zoom_factor)
    left = max(0, detail_px_x - crop_radius)
    top = max(0, detail_px_y - crop_radius)
    right = min(width, detail_px_x + crop_radius)
    bottom = min(height, detail_px_y + crop_radius)
    
    crop_img = base_img.crop((left, top, right, bottom))
    
    # Zoom it
    zoomed_img = crop_img.resize((loupe_radius_px * 2, loupe_radius_px * 2), Image.Resampling.LANCZOS)
    
    # Apply circular mask
    mask = Image.new("L", zoomed_img.size, 0)
    draw_mask = ImageDraw.Draw(mask)
    draw_mask.ellipse((0, 0, zoomed_img.width, zoomed_img.height), fill=255)
    zoomed_img.putalpha(mask)
    
    # Add Border
    border_overlay = Image.new("RGBA", zoomed_img.size, (0, 0, 0, 0))
    draw_border = ImageDraw.Draw(border_overlay)
    border_width = 8
    draw_border.ellipse(
        (border_width//2, border_width//2, zoomed_img.width - border_width//2, zoomed_img.height - border_width//2), 
        outline=accent_color, 
        width=border_width
    )
    final_loupe = Image.alpha_composite(zoomed_img, border_overlay)
    
    loupe_io = BytesIO()
    final_loupe.save(loupe_io, format="PNG")
    loupe_io.seek(0)
    
    # 5. Place Loupe and Elements on Slide
    # Determine placement (put loupe on the opposite side of the detail to balance)
    if detail_center_x_pct < 0.5:
        loupe_slide_cx = Inches(9.0) # Right side
    else:
        loupe_slide_cx = Inches(4.0) # Left side
        
    loupe_slide_cy = Inches(3.75) # Center vertically
    loupe_slide_radius = Inches(2.2) # ~4.4 inches diameter
    
    loupe_left = loupe_slide_cx - loupe_slide_radius
    loupe_top = loupe_slide_cy - loupe_slide_radius
    
    # Add connecting line FIRST so it goes behind the loupe
    detail_slide_x = prs.slide_width * detail_center_x_pct
    detail_slide_y = prs.slide_height * detail_center_y_pct
    
    connector = slide.shapes.add_connector(
        1, # Straight line
        detail_slide_x, detail_slide_y,
        loupe_slide_cx, loupe_slide_cy
    )
    connector.line.color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    connector.line.width = Pt(1.5)
    # Add a small dot at the origin
    dot = slide.shapes.add_shape(
        9, # msoShapeOval
        detail_slide_x - Pt(4), detail_slide_y - Pt(4),
        Pt(8), Pt(8)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    dot.line.fill.background()
    
    # Insert Loupe
    slide.shapes.add_picture(
        loupe_io, 
        loupe_left, loupe_top, 
        width=loupe_slide_radius*2, height=loupe_slide_radius*2
    )
    
    # 6. Add Typography
    # Place text near the loupe
    if detail_center_x_pct < 0.5:
        tx_left = Inches(7.5) # Next to loupe on right
        align = PP_ALIGN.RIGHT
    else:
        tx_left = Inches(1.5) # Next to loupe on left
        align = PP_ALIGN.LEFT
        
    tx_box = slide.shapes.add_textbox(tx_left, Inches(5.8), Inches(4.5), Inches(1.5))
    tf = tx_box.text_frame
    tf.clear()
    
    # Title
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = title_text
    run.font.bold = True
    run.font.size = Pt(18)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Body
    p2 = tf.add_paragraph()
    p2.alignment = align
    run2 = p2.add_run()
    run2.text = body_text
    run2.font.size = Pt(12)
    run2.font.name = "Arial"
    run2.font.color.rgb = RGBColor(180, 180, 180) # Light grey for readability
    
    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
