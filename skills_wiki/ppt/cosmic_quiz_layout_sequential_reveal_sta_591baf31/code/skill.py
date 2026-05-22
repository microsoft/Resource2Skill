def create_slide(
    output_pptx_path: str,
    title_text: str = "Who was the first person to walk on the moon?",
    answer_text: str = "Answer: Neil Armstrong",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Cosmic Quiz Layout.
    Generates a starry background and vector planet geometry.
    """
    import os
    import random
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Generate Deep Space Background with PIL ===
    bg_width, bg_height = 1920, 1080
    bg_img = Image.new('RGBA', (bg_width, bg_height), (13, 17, 28, 255))
    draw = ImageDraw.Draw(bg_img)
    
    # Generate random stars
    for _ in range(400):
        x = random.randint(0, bg_width)
        y = random.randint(0, bg_height)
        # Vary star sizes and opacities for depth
        size = random.choices([1, 2, 3], weights=[80, 15, 5])[0]
        opacity = random.randint(100, 255)
        draw.ellipse([x, y, x + size, y + size], fill=(255, 255, 255, opacity))
        
    bg_path = "temp_starfield.png"
    bg_img.save(bg_path)
    
    # Add background image to slide
    bg_shape = slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    # Send background to back (in python-pptx, shapes are ordered by creation, so it's already at the back)

    # === Layer 2: Vector Graphic (The Moon) ===
    # Main Moon Base (Positioned off-center to the right)
    moon = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        Inches(8.5), Inches(1), Inches(6.5), Inches(6.5)
    )
    moon.fill.solid()
    moon.fill.fore_color.rgb = RGBColor(220, 220, 220)
    moon.line.fill.background() # No line
    
    # Craters (scattered over the moon)
    craters_data = [
        (9.5, 2.5, 1.2, 1.2),
        (11.5, 1.5, 0.8, 0.8),
        (10.5, 4.5, 1.5, 1.5),
        (12.0, 5.0, 0.6, 0.6),
        (9.0, 5.5, 0.9, 0.9)
    ]
    
    for cx, cy, cw, ch in craters_data:
        crater = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx), Inches(cy), Inches(cw), Inches(ch)
        )
        crater.fill.solid()
        crater.fill.fore_color.rgb = RGBColor(180, 180, 180)
        crater.line.fill.background()

    # === Layer 3: Typography & Quiz Layout ===
    
    # 1. Question Box
    q_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(6.5), Inches(2.0))
    q_tf = q_box.text_frame
    q_tf.word_wrap = True
    q_p = q_tf.paragraphs[0]
    q_p.text = title_text
    q_p.font.size = Pt(44)
    q_p.font.color.rgb = RGBColor(255, 255, 255)
    q_p.font.name = "Arial"
    
    # 2. Answer Reveal Box (Styled to look like an animated reveal panel)
    ans_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(1.0), Inches(5.0), Inches(5.5), Inches(1.0)
    )
    
    # Subtle dark background for answer box to pop against the stars
    ans_box.fill.solid()
    ans_box.fill.fore_color.rgb = RGBColor(20, 25, 40)
    
    # Gold Outline to match tutorial's accent
    ans_box.line.color.rgb = RGBColor(255, 215, 0)
    ans_box.line.width = Pt(2)
    
    ans_tf = ans_box.text_frame
    ans_tf.vertical_anchor = PP_ALIGN.CENTER
    ans_p = ans_tf.paragraphs[0]
    ans_p.text = answer_text
    ans_p.font.size = Pt(28)
    ans_p.font.color.rgb = RGBColor(255, 215, 0) # Gold text
    ans_p.font.bold = True
    ans_p.font.name = "Arial"

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
