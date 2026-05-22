def create_slide(
    output_pptx_path: str,
    title_text: str = "CASEY SADLER",
    body_text: str = "SENIOR MANAGER OF\nPROFESSIONAL DEVELOPMENT",
    bg_theme: str = "portrait",
    accent_color: tuple = (111, 212, 228),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Hexagonal Portrait Spotlight' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import math
    import urllib.request
    from io import BytesIO
    import tempfile
    import os
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # --- Initialize Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- 1. Background Setup ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(91, 131, 166)

    # --- 2. Download Subject Image ---
    try:
        # Professional portrait photo
        url = "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?w=800&q=80"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception:
        # Fallback if download fails
        base_img = Image.new("RGBA", (800, 800), (200, 200, 200, 255))

    # --- 3. Generate Hexagonal Portrait via PIL (with supersampling for smooth edges) ---
    final_size = 800
    border_width = 30
    scale = 4  # 4x supersampling for perfect anti-aliasing
    super_size = final_size * scale

    # Crop to square
    w, h = base_img.size
    min_dim = min(w, h)
    left = (w - min_dim) / 2
    top = (h - min_dim) / 2
    img_cropped = base_img.crop((left, top, left + min_dim, top + min_dim))
    img_super = img_cropped.resize((super_size, super_size), Image.LANCZOS)

    out_img = Image.new("RGBA", (super_size, super_size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(out_img, "RGBA")
    cx, cy = super_size / 2, super_size / 2

    # Draw outer colored border hexagon
    R_outer = super_size / 2 - (scale * 2) # Slight padding
    points_outer = []
    for i in range(6):
        angle_rad = math.radians(60 * i - 30) # Pointy top orientation
        points_outer.append((cx + R_outer * math.cos(angle_rad), cy + R_outer * math.sin(angle_rad)))
    draw.polygon(points_outer, fill=accent_color + (255,))

    # Draw inner mask for the photo
    mask = Image.new("L", (super_size, super_size), 0)
    draw_mask = ImageDraw.Draw(mask)
    R_inner = R_outer - (border_width * scale)
    points_inner = []
    for i in range(6):
        angle_rad = math.radians(60 * i - 30)
        points_inner.append((cx + R_inner * math.cos(angle_rad), cy + R_inner * math.sin(angle_rad)))
    draw_mask.polygon(points_inner, fill=255)

    # Paste photo into mask
    out_img.paste(img_super, (0, 0), mask)

    # Scale down to final size to apply anti-aliasing
    final_img = out_img.resize((final_size, final_size), Image.LANCZOS)
    
    # Save temp image
    temp_img_fd, temp_img_path = tempfile.mkstemp(suffix=".png")
    os.close(temp_img_fd)
    final_img.save(temp_img_path)

    # --- 4. Insert Hexagon into PPTX ---
    img_size = Inches(3.8)
    pic_left = (prs.slide_width - img_size) / 2
    pic_top = Inches(1.2)
    slide.shapes.add_picture(temp_img_path, pic_left, pic_top, width=img_size, height=img_size)

    # --- 5. Add Decorative Radiating Lines ---
    line_len = Inches(0.3)
    center_x = prs.slide_width / 2
    center_y = pic_top + img_size / 2
    offset = img_size / 2 + Inches(0.15)
    
    # Draw 3 subtle lines radiating from the top right edge
    for angle_deg in [-70, -55, -40]:
        angle_rad = math.radians(angle_deg)
        start_x = center_x + offset * math.cos(angle_rad)
        start_y = center_y + offset * math.sin(angle_rad)
        end_x = center_x + (offset + line_len) * math.cos(angle_rad)
        end_y = center_y + (offset + line_len) * math.sin(angle_rad)
        
        # 1 = MSO_CONNECTOR.STRAIGHT
        line = slide.shapes.add_connector(1, start_x, start_y, end_x, end_y)
        line.line.color.rgb = RGBColor(255, 255, 255)
        line.line.width = Pt(2.5)

    # --- 6. Typography ---
    # Eyebrow Category
    txBox_cat = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(0.5))
    p_cat = txBox_cat.text_frame.paragraphs[0]
    p_cat.text = "EMPLOYEE SPOTLIGHT"
    p_cat.alignment = PP_ALIGN.CENTER
    p_cat.font.name = "Arial"
    p_cat.font.size = Pt(16)
    p_cat.font.bold = True
    p_cat.font.color.rgb = RGBColor(*accent_color)

    # Name
    txBox_name = slide.shapes.add_textbox(Inches(1), Inches(5.2), Inches(11.33), Inches(1))
    p_name = txBox_name.text_frame.paragraphs[0]
    p_name.text = title_text.upper()
    p_name.alignment = PP_ALIGN.CENTER
    p_name.font.name = "Arial"
    p_name.font.size = Pt(44)
    p_name.font.bold = True
    p_name.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Job Title
    txBox_title = slide.shapes.add_textbox(Inches(1), Inches(6.0), Inches(11.33), Inches(1))
    tf_title = txBox_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = body_text.upper()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial"
    p_title.font.size = Pt(18)
    p_title.font.bold = False
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Cleanup temp image and save
    prs.save(output_pptx_path)
    try:
        os.remove(temp_img_path)
    except OSError:
        pass
        
    return output_pptx_path
