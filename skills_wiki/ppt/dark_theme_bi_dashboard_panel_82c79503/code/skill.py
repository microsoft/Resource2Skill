def create_slide(
    output_pptx_path: str,
    title_text: str = "Quarterly Performance Dashboard",
    accent_color_main: tuple = (0, 191, 255),
    accent_color_pos: tuple = (46, 204, 113),
    accent_color_sec: tuple = (243, 156, 18),
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a Dark-Theme BI Dashboard Panel.

    Returns: Path to the saved PPTX file.
    """
    import io
    import numpy as np
    import matplotlib.pyplot as plt
    from matplotlib.patches import Wedge
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Color & Font Palette ---
    BG_COLOR = RGBColor(10, 20, 41)
    PANEL_COLOR = RGBColor(21, 35, 66)
    TEXT_COLOR = RGBColor(220, 220, 220)
    ACCENT_MAIN_RGB_FLOAT = tuple(c/255.0 for c in accent_color_main)
    ACCENT_POS_RGB_FLOAT = tuple(c/255.0 for c in accent_color_pos)
    ACCENT_SEC_RGB_FLOAT = tuple(c/255.0 for c in accent_color_sec)
    
    FONT_FAMILY = "Segoe UI Light"

    # --- Background ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    # --- Helper function to add a panel ---
    def add_panel(left, top, width, height):
        return slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )

    # --- Helper function to add text ---
    def add_text(shape, text, size, bold=False, align=PP_ALIGN.LEFT):
        text_frame = shape.text_frame
        p = text_frame.paragraphs[0]
        p.text = text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT_COLOR
        p.font.bold = bold
        p.alignment = align
        text_frame.margin_bottom = Inches(0.05)
        text_frame.margin_top = Inches(0.05)
        text_frame.margin_left = Inches(0.1)
        text_frame.margin_right = Inches(0.1)
        text_frame.vertical_anchor = "middle"

    # --- Matplotlib Chart Generators ---
    plt.style.use('dark_background')
    plt.rcParams['font.family'] = 'sans-serif'
    plt.rcParams['font.sans-serif'] = 'Segoe UI'

    def create_line_chart():
        fig, ax = plt.subplots(figsize=(6, 3))
        months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug']
        sales = np.random.randint(100, 500, size=8)
        sales[4] = sales[4] + 200 # peak
        ax.plot(months, sales, color=ACCENT_MAIN_RGB_FLOAT, linewidth=2.5, marker='o', markersize=8, markerfacecolor=ACCENT_MAIN_RGB_FLOAT, markeredgecolor='white')
        ax.fill_between(months, sales, color=ACCENT_MAIN_RGB_FLOAT, alpha=0.1)
        ax.tick_params(axis='x', colors='grey')
        ax.tick_params(axis='y', colors='grey')
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['bottom'].set_color('grey')
        ax.spines['left'].set_color('grey')
        ax.grid(axis='y', linestyle='--', alpha=0.2)
        fig.patch.set_alpha(0)
        ax.patch.set_alpha(0)
        
        buffer = io.BytesIO()
        plt.savefig(buffer, format='png', dpi=300, bbox_inches='tight', pad_inches=0.1)
        plt.close(fig)
        buffer.seek(0)
        return buffer

    def create_pie_chart():
        fig, ax = plt.subplots(figsize=(3, 3))
        labels = 'Product A', 'Product B', 'Product C'
        sizes = [45, 30, 25]
        colors = [ACCENT_MAIN_RGB_FLOAT, ACCENT_POS_RGB_FLOAT, ACCENT_SEC_RGB_FLOAT]
        ax.pie(sizes, labels=labels, autopct='%1.1f%%', startangle=90, colors=colors, wedgeprops={'edgecolor': 'black', 'linewidth': 2}, textprops={'color':"w"})
        ax.axis('equal')
        fig.patch.set_alpha(0)
        ax.patch.set_alpha(0)
        
        buffer = io.BytesIO()
        plt.savefig(buffer, format='png', dpi=300, bbox_inches='tight')
        plt.close(fig)
        buffer.seek(0)
        return buffer
    
    def create_gauge_chart(value=76): # value 0-100
        fig, ax = plt.subplots(figsize=(3, 1.5))
        
        # Background arc
        bg_arc = Wedge((0.5, 0.4), 0.4, 0, 180, width=0.15, facecolor='#2C3E50', edgecolor=None)
        ax.add_patch(bg_arc)

        # Foreground arc based on value
        angle = 180 * (value / 100.0)
        fg_arc = Wedge((0.5, 0.4), 0.4, 0, angle, width=0.15, facecolor=ACCENT_POS_RGB_FLOAT, edgecolor=None)
        ax.add_patch(fg_arc)

        # Needle
        theta = np.deg2rad(180 - angle)
        ax.arrow(0.5, 0.4, 0.35 * np.cos(theta), 0.35 * np.sin(theta),
                 width=0.01, head_width=0.0, head_length=0.0, fc=TEXT_COLOR.to_rgb(), ec=TEXT_COLOR.to_rgb())

        ax.text(0.5, 0.45, f'{value}%', ha='center', va='center', fontsize=20, color='white', weight='bold')
        
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 0.8)
        ax.set_aspect('equal', adjustable='box')
        ax.axis('off')
        fig.patch.set_alpha(0)
        ax.patch.set_alpha(0)

        buffer = io.BytesIO()
        plt.savefig(buffer, format='png', dpi=300, bbox_inches='tight')
        plt.close(fig)
        buffer.seek(0)
        return buffer

    # --- Build Slide Layout ---
    
    # Title
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(15), Inches(0.75))
    add_text(title_shape, title_text.upper(), 28, bold=True)
    
    # --- Main Content Area ---
    main_panel = add_panel(3.5, 1.25, 8.5, 4.5)
    main_panel.fill.solid()
    main_panel.fill.fore_color.rgb = PANEL_COLOR
    main_panel.line.fill.background()
    
    # Add Title to Main Panel
    main_title = slide.shapes.add_textbox(Inches(3.7), Inches(1.4), Inches(4), Inches(0.5))
    add_text(main_title, "Monthly Sales Trend", 16)
    
    # Add Line Chart to Main Panel
    line_chart_img = create_line_chart()
    slide.shapes.add_picture(line_chart_img, Inches(4.5), Inches(2.2), width=Inches(7))

    # --- KPI Panels ---
    kpi_panel_1 = add_panel(12.5, 1.25, 3, 2.1)
    kpi_panel_1.fill.solid(); kpi_panel_1.fill.fore_color.rgb = PANEL_COLOR; kpi_panel_1.line.fill.background()
    kpi_title_1 = slide.shapes.add_textbox(Inches(12.6), Inches(1.35), Inches(2.8), Inches(0.4))
    add_text(kpi_title_1, "Total Revenue", 12)
    kpi_value_1 = slide.shapes.add_textbox(Inches(12.6), Inches(1.8), Inches(2.8), Inches(1.0))
    add_text(kpi_value_1, "$2.8M", 44, bold=True)
    
    kpi_panel_2 = add_panel(12.5, 3.65, 3, 2.1)
    kpi_panel_2.fill.solid(); kpi_panel_2.fill.fore_color.rgb = PANEL_COLOR; kpi_panel_2.line.fill.background()
    kpi_title_2 = slide.shapes.add_textbox(Inches(12.6), Inches(3.75), Inches(2.8), Inches(0.4))
    add_text(kpi_title_2, "New Customers", 12)
    kpi_value_2 = slide.shapes.add_textbox(Inches(12.6), Inches(4.2), Inches(2.8), Inches(1.0))
    add_text(kpi_value_2, "1,245", 44, bold=True)

    # --- Bottom Panels ---
    bottom_panel_1 = add_panel(3.5, 6.1, 4.1, 2.5)
    bottom_panel_1.fill.solid(); bottom_panel_1.fill.fore_color.rgb = PANEL_COLOR; bottom_panel_1.line.fill.background()
    bottom_title_1 = slide.shapes.add_textbox(Inches(3.6), Inches(6.2), Inches(3.9), Inches(0.4))
    add_text(bottom_title_1, "Sales by Category", 14)
    pie_chart_img = create_pie_chart()
    slide.shapes.add_picture(pie_chart_img, Inches(4), Inches(6.5), height=Inches(1.8))
    
    bottom_panel_2 = add_panel(7.9, 6.1, 7.6, 2.5)
    bottom_panel_2.fill.solid(); bottom_panel_2.fill.fore_color.rgb = PANEL_COLOR; bottom_panel_2.line.fill.background()
    bottom_title_2 = slide.shapes.add_textbox(Inches(8), Inches(6.2), Inches(3.9), Inches(0.4))
    add_text(bottom_title_2, "Customer Satisfaction", 14)
    gauge_chart_img = create_gauge_chart(82)
    slide.shapes.add_picture(gauge_chart_img, Inches(9.8), Inches(6.5), height=Inches(1.8))

    # --- Slicer Panel (Left) ---
    slicer_panel = add_panel(0.5, 1.25, 2.5, 7.3)
    slicer_panel.fill.solid()
    slicer_panel.fill.fore_color.rgb = PANEL_COLOR
    slicer_panel.line.fill.background()
    
    slicer_title = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(2.3), Inches(0.4))
    add_text(slicer_title, "FILTERS", 14, bold=True)

    slicer_items = ["All Regions", "North America", "Europe", "Asia-Pacific", "South America"]
    for i, item in enumerate(slicer_items):
        top = 2.0 + i * 0.7
        is_active = (i == 1) # Highlight one item
        
        btn_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(top), Inches(2.1), Inches(0.5)
        )
        if is_active:
            btn_shape.fill.solid()
            btn_shape.fill.fore_color.rgb = RGBColor.from_string(accent_color_main.hex())
            btn_shape.line.fill.background()
        else:
            btn_shape.fill.solid()
            btn_shape.fill.fore_color.rgb = RGBColor(40, 60, 90)
            btn_shape.line.fill.background()
        
        add_text(btn_shape, item, 11, bold=is_active, align=PP_ALIGN.CENTER)

    # --- Save and Return ---
    prs.save(output_pptx_path)
    return output_pptx_path

