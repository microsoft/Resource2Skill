def create_slide(
    output_pptx_path: str,
    title_text: str = "李彦宏 Robin",
    role_text: str = "百度公司创始人 / 董事长兼首席执行官",
    bullet_points: list = None,
    accent_color: tuple = (41, 128, 185),  # Professional Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Geometric Anchored Character Profile" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw
    import urllib.request
    import io
    import os

    if bullet_points is None:
        bullet_points = [
            "大数据科学与产业研究院名誉院长",
            "纽约州立大学完成计算机科学硕士",
            "引领中国搜索引擎与人工智能发展",
            "提出并践行'技术改变世界'的理念"
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # ==========================================
    # Layer 1: Background Geometric Blocks
    # ==========================================
    
    # 1a. Solid Color Block (Left)
    block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(0), Inches(4.5), Inches(7.5)
    )
    block.fill.solid()
    block.fill.fore_color.rgb = RGBColor(*accent_color)
    block.line.fill.background() # No outline

    # 1b. Offset Wireframe Box (adds architectural depth)
    wireframe = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0.5), Inches(0.5), Inches(4.5), Inches(6.5)
    )
    wireframe.fill.background() # Transparent fill
    wireframe.line.color.rgb = RGBColor(20, 20, 20) # Dark outline
    wireframe.line.width = Pt(2)

    # ==========================================
    # Layer 2: Image Processing (PIL)
    # ==========================================
    
    # Fetch a placeholder portrait image
    portrait_url = "https://images.unsplash.com/photo-1560250097-0b93528c311a?q=80&w=800&auto=format&fit=crop"
    try:
        req = urllib.request.Request(portrait_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img_data = response.read()
            img = Image.open(io.BytesIO(img_data)).convert("RGBA")
    except Exception:
        # Fallback if download fails
        img = Image.new("RGBA", (800, 800), (200, 200, 200, 255))
        d = ImageDraw.Draw(img)
        d.text((300, 380), "Portrait", fill=(100, 100, 100, 255))

    # Crop to square
    width, height = img.size
    min_dim = min(width, height)
    left = (width - min_dim) / 2
    top = (height - min_dim) / 2
    img = img.crop((left, top, left + min_dim, top + min_dim))

    # Create circular mask
    mask = Image.new('L', img.size, 0)
    draw = ImageDraw.Draw(mask)
    draw.ellipse((0, 0, min_dim, min_dim), fill=255)
    
    # Apply mask
    circular_img = Image.new('RGBA', img.size, (0, 0, 0, 0))
    circular_img.paste(img, (0, 0), mask)
    
    # Add a white border to the circle
    border_width = int(min_dim * 0.03)
    draw_circle = ImageDraw.Draw(circular_img)
    draw_circle.ellipse((border_width/2, border_width/2, min_dim - border_width/2, min_dim - border_width/2), 
                        outline=(255, 255, 255, 255), width=border_width)

    # Save to buffer
    img_stream = io.BytesIO()
    circular_img.save(img_stream, format='PNG')
    img_stream.seek(0)

    # Insert Image onto Slide (Overlapping the color block boundary)
    pic_size = Inches(4.5)
    slide.shapes.add_picture(
        img_stream, 
        Inches(2.25), Inches(1.5), # Positioned to bridge the blue and white areas
        pic_size, pic_size
    )

    # ==========================================
    # Layer 3: Text & Typography
    # ==========================================
    
    # Name Text
    name_box = slide.shapes.add_textbox(Inches(7.2), Inches(1.5), Inches(5), Inches(1))
    name_tf = name_box.text_frame
    name_p = name_tf.paragraphs[0]
    name_p.text = title_text
    name_p.font.size = Pt(44)
    name_p.font.bold = True
    name_p.font.color.rgb = RGBColor(30, 30, 30)

    # Role / Title Text
    role_box = slide.shapes.add_textbox(Inches(7.2), Inches(2.3), Inches(5), Inches(0.5))
    role_tf = role_box.text_frame
    role_p = role_tf.paragraphs[0]
    role_p.text = role_text
    role_p.font.size = Pt(18)
    role_p.font.bold = True
    role_p.font.color.rgb = RGBColor(*accent_color)

    # Add a thin separator line
    sep_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(7.3), Inches(3.0), Inches(0.5), Pt(3)
    )
    sep_line.fill.solid()
    sep_line.fill.fore_color.rgb = RGBColor(*accent_color)
    sep_line.line.fill.background()

    # Bullet Points Details
    detail_box = slide.shapes.add_textbox(Inches(7.2), Inches(3.4), Inches(5.5), Inches(3))
    detail_tf = detail_box.text_frame
    detail_tf.word_wrap = True
    
    for point in bullet_points:
        p = detail_tf.add_paragraph()
        p.text = point
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(80, 80, 80)
        p.level = 0
        p.space_after = Pt(12)

    # ==========================================
    # Layer 4: Decorative Elements
    # ==========================================
    
    # Add a decorative triangle in the top right corner
    triangle = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE,
        Inches(12.0), Inches(0.5), Inches(0.8), Inches(0.8)
    )
    triangle.fill.background()
    triangle.line.color.rgb = RGBColor(*accent_color)
    triangle.line.width = Pt(2)
    
    # Save Presentation
    prs.save(output_pptx_path)
    return output_pptx_path
