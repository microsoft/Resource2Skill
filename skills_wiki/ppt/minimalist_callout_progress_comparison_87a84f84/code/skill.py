def create_slide(
    output_pptx_path: str,
    title_text: str = "ANIMATED COMPARISON SLIDE",
    subtitle_text: str = "Create a clean, data-driven horizontal layout",
    data: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Minimalist Callout Progress Comparison effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Set up presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set pure white background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # 1. Main Title
    title_tb = slide.shapes.add_textbox(Inches(2), Inches(0.5), Inches(9.333), Inches(0.6))
    tf = title_tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(50, 50, 50)

    # 2. Subtitle
    sub_tb = slide.shapes.add_textbox(Inches(2), Inches(1.1), Inches(9.333), Inches(0.5))
    tf = sub_tb.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(150, 150, 150)

    # Default data payload if none provided
    if data is None:
        data = [
            {
                "task": "TASK ONE", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above", 
                "value": 0.75, 
                "color": (216, 59, 65)
            },
            {
                "task": "TASK TWO", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above", 
                "value": 0.50, 
                "color": (75, 172, 143)
            },
            {
                "task": "TASK THREE", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above", 
                "value": 0.25, 
                "color": (63, 136, 197)
            },
        ]

    # Layout configuration
    start_y = Inches(2.2)
    row_height = Inches(1.7)
    
    txt_l = Inches(1.0)
    txt_w = Inches(3.0)
    
    track_l = Inches(4.5)
    track_w = Inches(7.5)
    track_h = Inches(0.15)

    # Generate rows
    for index, item in enumerate(data):
        row_y = start_y + (index * row_height)
        color_rgb = RGBColor(*item['color'])
        
        # --- Left Column: Text Content ---
        tb = slide.shapes.add_textbox(txt_l, row_y, txt_w, Inches(1.0))
        tf = tb.text_frame
        
        # Task Name
        p = tf.paragraphs[0]
        p.text = item['task']
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = color_rgb
        
        # Task Description
        p2 = tf.add_paragraph()
        p2.text = item['desc']
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(120, 120, 120)
        p2.space_before = Pt(5)

        # --- Right Column: Visualization ---
        track_t = row_y + Inches(0.35) # Vertically align with the text block
        
        # Track Background
        track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, track_l, track_t, track_w, track_h)
        track.fill.solid()
        track.fill.fore_color.rgb = RGBColor(235, 235, 235)
        track.line.fill.background()

        # Track Tick Marks (Visual Dividers)
        num_segments = 5
        for i in range(1, num_segments):
            tick_cx = track_l + (track_w * i / num_segments)
            tick = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 
                tick_cx - Inches(0.015), 
                track_t - Inches(0.05), 
                Inches(0.03), 
                track_h + Inches(0.1) # Extends slightly above and below the track
            )
            tick.fill.solid()
            tick.fill.fore_color.rgb = RGBColor(200, 200, 200)
            tick.line.fill.background()

        # Progress Bar Overlay
        prog_w = track_w * item['value']
        prog = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, track_l, track_t, prog_w, track_h)
        prog.fill.solid()
        prog.fill.fore_color.rgb = color_rgb
        prog.line.fill.background()

        # --- Data Callout (Box + Pointer) ---
        callout_cx = track_l + prog_w
        
        # Callout Box (Rounded Rectangle)
        box_w = Inches(0.8)
        box_h = Inches(0.4)
        box_l = callout_cx - box_w / 2
        box_t = track_t - box_h - Inches(0.08) # 0.08 leaves room for pointer
        
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_l, box_t, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = color_rgb
        box.line.fill.background()
        
        # Callout Pointer (Inverted Triangle)
        ptr_w = Inches(0.15)
        ptr_h = Inches(0.1)
        ptr_l = callout_cx - ptr_w / 2
        ptr_t = box_t + box_h - Inches(0.01) # Slight overlap to remove visual seam
        
        ptr = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, ptr_l, ptr_t, ptr_w, ptr_h)
        ptr.rotation = 180 # Point downwards
        ptr.fill.solid()
        ptr.fill.fore_color.rgb = color_rgb
        ptr.line.fill.background()

        # Callout Text
        tf_box = box.text_frame
        tf_box.text = f"{int(item['value'] * 100)}%"
        # Neutralize margins to perfectly center text in small shape
        tf_box.margin_left = 0
        tf_box.margin_right = 0
        tf_box.margin_top = Inches(0.05)
        tf_box.margin_bottom = 0
        
        p_box = tf_box.paragraphs[0]
        p_box.alignment = PP_ALIGN.CENTER
        p_box.font.bold = True
        p_box.font.size = Pt(14)
        p_box.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
