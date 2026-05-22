import os
import urllib.request
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image, ImageEnhance, ImageFilter

def apply_gradient(shape, color1, color2):
    """Injects OpenXML to apply a linear gradient fill to a shape."""
    shape.fill.solid()  # Initialize fill
    spPr = shape.element.spPr
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    
    # Remove existing fill elements
    for fill_type in ['solidFill', 'noFill', 'blipFill', 'pattFill', 'grpFill']:
        fill = spPr.find(f'.//{{{a_ns}}}{fill_type}')
        if fill is not None:
            spPr.remove(fill)
            
    gradFill = etree.SubElement(spPr, f'{{{a_ns}}}gradFill')
    gradFill.set("rotWithShape", "1")
    
    gsLst = etree.SubElement(gradFill, f'{{{a_ns}}}gsLst')
    
    gs1 = etree.SubElement(gsLst, f'{{{a_ns}}}gs')
    gs1.set("pos", "0")
    srgb1 = etree.SubElement(gs1, f'{{{a_ns}}}srgbClr')
    srgb1.set("val", f"{color1[0]:02X}{color1[1]:02X}{color1[2]:02X}")
    
    gs2 = etree.SubElement(gsLst, f'{{{a_ns}}}gs')
    gs2.set("pos", "100000")
    srgb2 = etree.SubElement(gs2, f'{{{a_ns}}}srgbClr')
    srgb2.set("val", f"{color2[0]:02X}{color2[1]:02X}{color2[2]:02X}")
    
    lin = etree.SubElement(gradFill, f'{{{a_ns}}}lin')
    lin.set("ang", "5400000")  # 90 degrees (Top to Bottom)
    lin.set("scaled", "1")

def apply_soft_edge(shape, radius_pt):
    """Injects OpenXML to apply a soft edge (blur) to a shape."""
    shape.line.fill.background()  # Remove outline
    spPr = shape.element.spPr
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    
    effectLst = spPr.find(f'.//{{{a_ns}}}effectLst')
    if effectLst is None:
        effectLst = etree.SubElement(spPr, f'{{{a_ns}}}effectLst')
        
    softEdge = etree.SubElement(effectLst, f'{{{a_ns}}}softEdge')
    softEdge.set('rad', str(int(radius_pt * 12700)))  # 1 pt = 12700 EMUs

def create_slide(
    output_pptx_path: str,
    title_text: str = "12 Points Agenda",
    body_text: str = "",
    bg_palette: str = "white minimalism",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 12-Point 3D Ribbon Agenda Grid.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # --- Background Generation ---
    try:
        bg_path = "temp_bg.jpg"
        url = f"https://image.pollinations.ai/prompt/abstract%20{bg_palette.replace(' ', '%20')}%20blur%20background?width=1920&height=1080&nologo=true"
        urllib.request.urlretrieve(url, bg_path)
        
        img = Image.open(bg_path).convert("RGBA")
        img = img.filter(ImageFilter.GaussianBlur(15))
        enhancer = ImageEnhance.Brightness(img)
        img = enhancer.enhance(1.4)
        
        # Apply a white wash for readability
        wash = Image.new('RGBA', img.size, (250, 252, 255, 210))
        img = Image.alpha_composite(img, wash)
        img.convert('RGB').save(bg_path)
        
        slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
        os.remove(bg_path)
    except Exception:
        # Fallback to solid soft gray
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(240, 243, 246)

    # Add Overall Title (Top Center)
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.2), Inches(9.333), Inches(0.6))
    tf = title_box.text_frame
    tf.text = title_text.upper()
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(40, 50, 70)

    # --- Palette: Base, Dark (fold/shadow), Light (gradient) ---
    colors = [
        ((220, 53, 69),  (140, 20, 30),  (250, 100, 110)),  # Red
        ((253, 126, 20), (180, 70, 0),   (255, 170, 80)),   # Orange
        ((255, 193, 7),  (180, 130, 0),  (255, 230, 100)),  # Yellow
        ((40, 167, 69),  (20, 100, 30),  (90, 210, 120)),   # Green
        ((32, 201, 151), (15, 130, 90),  (90, 230, 180)),   # Teal
        ((23, 162, 184), (10, 100, 120), (80, 200, 220)),   # Cyan
        ((0, 123, 255),  (0, 70, 160),   (100, 170, 255)),  # Blue
        ((102, 16, 242), (60, 10, 160),  (150, 80, 255)),   # Indigo
        ((111, 66, 193), (70, 30, 130),  (160, 110, 230)),  # Purple
        ((232, 62, 140), (150, 30, 90),  (255, 110, 180)),  # Pink
        ((73, 80, 87),   (40, 45, 50),   (130, 140, 150)),  # Dark Gray
        ((108, 117, 125),(60, 65, 70),   (160, 170, 180))   # Slate
    ]

    # Grid Settings
    card_w = 5.5
    card_h = 0.75
    row_height = 0.95
    y_start = 1.0

    for i in range(12):
        col = i // 6
        row = i % 6
        
        # Calculate Base Coordinates for the Tab's Top-Left
        cx = 0.7 if col == 0 else 6.8
        cy = y_start + (row * row_height)
        
        base_color, dark_color, light_color = colors[i]
        
        # 1. Soft Shadow (Drawn first to sit at the very back)
        shadow = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx + 0.6), Inches(cy + 0.3), 
            Inches(card_w - 0.7), Inches(0.55)
        )
        shadow.fill.solid()
        shadow.fill.fore_color.rgb = RGBColor(0, 0, 0)
        # Apply 15pt blur, opacity via shape fill not supported directly in simple API without lxml, 
        # but pure black with heavy blur acts as a perfect shadow.
        apply_soft_edge(shadow, 15) 

        # 2. 3D Fold Triangle
        ff_builder = slide.shapes.build_freeform()
        ff_builder.add_line_segments([
            (Inches(cx + 0.5), Inches(cy + card_h - 0.05)),
            (Inches(cx + 0.5), Inches(cy + card_h + 0.15)),
            (Inches(cx + 0.2), Inches(cy + card_h - 0.05)),
        ], close=True)
        fold = ff_builder.convert_to_shape()
        fold.fill.solid()
        fold.fill.fore_color.rgb = RGBColor(*dark_color)
        fold.line.fill.background()

        # 3. Main White Card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(cx + 0.5), Inches(cy), 
            Inches(card_w - 0.5), Inches(card_h)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.fill.background()

        # 4. Colored Ribbon/Tab (Left side rounded, right side flat)
        # We do this by combining a rounded rect and a normal rect
        tab_round = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(cx), Inches(cy), 
            Inches(1.0), Inches(card_h)
        )
        tab_round.fill.solid()
        tab_round.fill.fore_color.rgb = RGBColor(*base_color)
        tab_round.line.fill.background()
        
        tab_square = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(cx + 0.5), Inches(cy), 
            Inches(0.5), Inches(card_h)
        )
        tab_square.fill.solid()
        tab_square.fill.fore_color.rgb = RGBColor(*base_color)
        tab_square.line.fill.background()

        # 5. Background Circle (Depth)
        circ_back = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx + 0.35), Inches(cy + 0.25), 
            Inches(0.4), Inches(0.4)
        )
        circ_back.fill.solid()
        circ_back.fill.fore_color.rgb = RGBColor(*dark_color)
        circ_back.line.fill.background()

        # 6. Foreground Gradient Circle (The Badge)
        circ_front = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx + 0.3), Inches(cy + 0.2), 
            Inches(0.4), Inches(0.4)
        )
        circ_front.line.fill.background()
        apply_gradient(circ_front, light_color, base_color)
        
        # Add Step Number to Gradient Circle
        tf = circ_front.text_frame
        tf.text = f"{i+1:02d}"
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(12)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        # 7. "STEP" Text Label
        step_txt = slide.shapes.add_textbox(Inches(cx), Inches(cy - 0.05), Inches(1.0), Inches(0.3))
        tf = step_txt.text_frame
        tf.text = "STEP"
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        # 8. Main Title Text
        title_txt = slide.shapes.add_textbox(Inches(cx + 1.1), Inches(cy + 0.05), Inches(4.0), Inches(0.3))
        tf = title_txt.text_frame
        tf.text = f"TITLE OF AGENDA POINT {i+1}"
        p = tf.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.color.rgb = RGBColor(60, 60, 60)

        # 9. Body Description Text
        desc_txt = slide.shapes.add_textbox(Inches(cx + 1.1), Inches(cy + 0.3), Inches(4.2), Inches(0.4))
        tf = desc_txt.text_frame
        tf.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa."
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(140, 140, 140)

    prs.save(output_pptx_path)
    return output_pptx_path
