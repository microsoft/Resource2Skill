def create_slide(
    output_pptx_path: str,
    title_text: str = "a chisel",
    body_text: str = "a gene",
    bg_palette: str = "machine",  # Keyword for left image
    accent_color: tuple = (255, 255, 255),  # Ignored for this pure BW structural style
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Bifurcated Minimalist Dichotomy' visual effect.
    Creates a perfect 50/50 split screen with two contrasting images and museum-style placards.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    import urllib.request
    import io

    # Keywords for the contrasting images
    left_keyword = bg_palette
    right_keyword = kwargs.get("right_keyword", "biology")

    # Image fetching URLs
    left_img_url = f"https://source.unsplash.com/featured/1000x1200/?{left_keyword}"
    right_img_url = f"https://source.unsplash.com/featured/1000x1200/?{right_keyword}"

    # Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Canvas dimensions
    width = prs.slide_width
    height = prs.slide_height
    half_width = width / 2

    def fetch_and_crop_image(url, target_aspect_ratio=(8, 9), fallback_color=(30, 30, 30)):
        """Fetches an image, crops it to the target aspect ratio, and returns a BytesIO object."""
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as response:
                img_data = response.read()
            img = Image.open(io.BytesIO(img_data)).convert("RGB")
            
            # Crop to aspect ratio
            img_w, img_h = img.size
            target_w_ratio, target_h_ratio = target_aspect_ratio
            
            # Calculate new dimensions
            if (img_w / img_h) > (target_w_ratio / target_h_ratio):
                # Image is too wide
                new_w = int(img_h * (target_w_ratio / target_h_ratio))
                new_h = img_h
            else:
                # Image is too tall
                new_w = img_w
                new_h = int(img_w * (target_h_ratio / target_w_ratio))
                
            left = (img_w - new_w) / 2
            top = (img_h - new_h) / 2
            right = (img_w + new_w) / 2
            bottom = (img_h + new_h) / 2
            
            img_cropped = img.crop((left, top, right, bottom))
            
            img_io = io.BytesIO()
            img_cropped.save(img_io, format='PNG')
            img_io.seek(0)
            return img_io
        except Exception as e:
            print(f"Failed to fetch image: {e}. Generating fallback.")
            # Fallback solid color image
            img = Image.new('RGB', (800, 900), color=fallback_color)
            img_io = io.BytesIO()
            img.save(img_io, format='PNG')
            img_io.seek(0)
            return img_io

    # === Layer 1: Split Imagery ===
    # Target aspect ratio for a 6.666 x 7.5 inch half is exactly 8:9
    left_img_stream = fetch_and_crop_image(left_img_url, target_aspect_ratio=(8, 9), fallback_color=(20, 20, 25))
    right_img_stream = fetch_and_crop_image(right_img_url, target_aspect_ratio=(8, 9), fallback_color=(15, 30, 20))

    # Add Left Image
    slide.shapes.add_picture(left_img_stream, 0, 0, width=half_width, height=height)
    # Add Right Image
    slide.shapes.add_picture(right_img_stream, half_width, 0, width=half_width, height=height)


    # === Layer 2: The Divider ===
    # A stark black line separating the two realms
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, half_width - Pt(1.5), 0, Pt(3), height)
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(0, 0, 0)
    divider.line.fill.background() # No border


    # === Layer 3: Museum Placards (Text) ===
    def add_placard(x_center, text_str):
        if not text_str:
            return
            
        box_width = Inches(3.5)
        box_height = Inches(1.0)
        box_x = x_center - (box_width / 2)
        box_y = (height / 2) - (box_height / 2)
        
        # Add stark white box
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box_x, box_y, box_width, box_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.fill.background()  # No border
        
        # Apply minimalist typography
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p = text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text_str
        run.font.name = "Arial"
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    # Left placard center: 3.333 inches
    add_placard(half_width / 2, title_text)
    
    # Right placard center: 10.0 inches
    add_placard(half_width + (half_width / 2), body_text)

    # Save
    prs.save(output_pptx_path)
    return output_pptx_path
