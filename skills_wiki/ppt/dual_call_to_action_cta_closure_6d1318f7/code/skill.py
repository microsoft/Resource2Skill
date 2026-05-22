import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image

def create_slide(
    output_pptx_path: str,
    title_text: str = "Expert Academy",
    url_text: str = "www.expertacademy.be",
    cta_left_text: str = "FOLLOW US\nCLICK HERE",
    cta_right_text: str = "MORE INFO\nCLICK HERE",
    accent_color: tuple = (31, 78, 121),  # Corporate Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dual Call-to-Action Closure" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Gradient via PIL ===
    bg_path = "temp_bg_gradient.png"
    width, height = int(13.333 * 100), int(7.5 * 100) # scale for resolution
    base = Image.new('RGB', (width, height), (250, 250, 250))
    top = Image.new('RGB', (width, height), (220, 225, 230))
    mask = Image.new('L', (width, height))
    
    # Create a vertical linear gradient mask
    mask_data = []
    for y in range(height):
        val = int(255 * (y / height))
        mask_data.extend([val] * width)
    mask.putdata(mask_data)
    
    base.paste(top, (0, 0), mask)
    base.save(bg_path)
    
    # Add background to slide
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Main Typography ===
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.666), Inches(1.0), Inches(10), Inches(1.2))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.name = "Arial"
    p_title.font.color.rgb = RGBColor(0, 0, 0)
    p_title.alignment = PP_ALIGN.CENTER

    # Center URL
    url_box = slide.shapes.add_textbox(Inches(3.666), Inches(3.2), Inches(6), Inches(1.0))
    tf_url = url_box.text_frame
    p_url = tf_url.paragraphs[0]
    p_url.text = url_text
    p_url.font.size = Pt(28)
    p_url.font.bold = True
    p_url.font.name = "Arial"
    p_url.font.color.rgb = RGBColor(*accent_color)
    p_url.font.underline = True
    p_url.alignment = PP_ALIGN.CENTER

    # === Layer 3: Visual Vectors & Arrows ===
    
    # Left Arrow
    left_arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, 
        Inches(2.5), Inches(2.2), Inches(1.5), Inches(3.8)
    )
    left_arrow.fill.solid()
    left_arrow.fill.fore_color.rgb = RGBColor(*accent_color)
    left_arrow.line.color.rgb = RGBColor(*accent_color)

    # Right Arrow
    right_arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, 
        Inches(9.333), Inches(2.2), Inches(1.5), Inches(3.8)
    )
    right_arrow.fill.solid()
    right_arrow.fill.fore_color.rgb = RGBColor(*accent_color)
    right_arrow.line.color.rgb = RGBColor(*accent_color)

    # === Layer 4: Action Text ===
    def add_action_text(x, y, text):
        box = slide.shapes.add_textbox(x, y, Inches(3.0), Inches(1.0))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.name = "Arial"
        p.font.color.rgb = RGBColor(*accent_color)
        p.alignment = PP_ALIGN.CENTER

    # Place text directly below the arrow tips
    add_action_text(Inches(1.75), Inches(6.2), cta_left_text)
    add_action_text(Inches(8.583), Inches(6.2), cta_right_text)

    # Save presentation
    prs.save(output_pptx_path)
    
    # Cleanup temp file
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
