def create_slide(
    output_pptx_path: str,
    title_text: str = "REVENUE UPDATE",
    body_text: str = "Revenue plummeted due to declining sales. We need to pivot our strategy.",
    presenter_name: str = "Kevin Stratvert",
    presenter_image_url: str = "https://images.unsplash.com/photo-1560250097-0b93528c311a?q=80&w=800&auto=format&fit=crop",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Dynamic Presenter Layout' visual effect.
    """
    import os
    import io
    import urllib.request
    from PIL import Image
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(17, 21, 31)

    # === Layer 2: Title & Footer ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(1.0))
    tf_title = title_box.text_frame
    tf_title.text = title_text
    p = tf_title.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(12.333), Inches(0.5))
    tf_footer = footer_box.text_frame
    tf_footer.text = body_text
    p_foot = tf_footer.paragraphs[0]
    p_foot.font.size = Pt(14)
    p_foot.font.color.rgb = RGBColor(180, 180, 190)
    p_foot.alignment = PP_ALIGN.CENTER

    # === Layer 3: Context Chart (Matplotlib) ===
    # Generating a sleek dark-mode chart indicating decline
    plt.style.use('dark_background')
    fig, ax = plt.subplots(figsize=(8, 5))
    fig.patch.set_alpha(0.0)
    ax.patch.set_alpha(0.0)

    dates = ['2/1/2021', '3/1/2021', '4/1/2021', '5/1/2021', '6/1/2021', '7/1/2021', '8/1/2021']
    revenue = [70, 68, 65, 40, 35, 30, 25] 

    ax.plot(dates, revenue, color='#FF8C00', linewidth=4)
    ax.plot(dates[-1], revenue[-1], marker='o', markersize=12, color='#FF4500') # Explosion point

    ax.set_ylim(0, 80)
    ax.set_yticks([0, 10, 20, 30, 40, 50, 60, 70])
    ax.set_yticklabels([f"${y},000,000" for y in ax.get_yticks()], fontsize=8, color='#AAAAAA')
    ax.set_xticklabels(dates, fontsize=9, color='#AAAAAA')

    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#555555')
    ax.spines['bottom'].set_color('#555555')
    ax.tick_params(axis='both', colors='#555555')
    ax.grid(axis='y', color='#333333', linestyle='-', alpha=0.7)

    plt.tight_layout()
    chart_stream = io.BytesIO()
    plt.savefig(chart_stream, format='png', transparent=True, dpi=300)
    chart_stream.seek(0)
    plt.close(fig)

    slide.shapes.add_picture(chart_stream, Inches(0.5), Inches(1.6), width=Inches(7.5))

    # === Layer 4: Presenter Shape (Arrow Mockup) ===
    # 4a. Download and crop portrait exactly to arrow aspect ratio to prevent distortion
    target_w, target_h = 3.5, 4.5
    tmp_img_path = "tmp_presenter_crop.jpg"
    
    try:
        req = urllib.request.Request(presenter_image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(response).convert("RGB")
    except Exception:
        img = Image.new('RGB', (800, 1000), color=(80, 120, 160))
        
    target_ratio = target_w / target_h
    w, h = img.size
    curr_ratio = w / h
    
    if curr_ratio > target_ratio:
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        img = img.crop((left, 0, left + new_w, h))
    else:
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        img = img.crop((0, top, w, top + new_h))
        
    img.save(tmp_img_path)

    # 4b. Create Shape and fill with cropped picture
    presenter_shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, Inches(9.0), Inches(1.8), Inches(target_w), Inches(target_h)
    )
    presenter_shape.fill.user_picture(tmp_img_path)
    presenter_shape.line.color.rgb = RGBColor(255, 255, 255)
    presenter_shape.line.width = Pt(2)

    # 4c. Inject Glow Effect via lxml
    glow_color_hex = "FF4500" # Orange-Red
    glow_radius_emu = int(20 * 12700) # 20pt glow
    glow_xml = f"""
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:glow rad="{glow_radius_emu}">
            <a:srgbClr val="{glow_color_hex}"/>
        </a:glow>
    </a:effectLst>
    """
    effect_lst = parse_xml(glow_xml)
    spPr = presenter_shape.element.spPr
    existing_effects = spPr.xpath('./a:effectLst')
    if existing_effects:
        glow_el = effect_lst.find('{http://schemas.openxmlformats.org/drawingml/2006/main}glow')
        existing_effects[0].append(glow_el)
    else:
        spPr.append(effect_lst)

    # === Layer 5: Transparent Name Tag ===
    name_tag = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.75), Inches(6.0), Inches(4.0), Inches(0.5)
    )
    name_tag.fill.solid()
    name_tag.fill.fore_color.rgb = RGBColor(0, 0, 0)
    name_tag.line.color.rgb = RGBColor(100, 100, 100)
    name_tag.line.width = Pt(1)
    
    # Inject 60% opacity (40% transparent)
    alpha_xml = '<a:alpha val="60000" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    srgb_elements = name_tag.element.xpath('.//a:solidFill/a:srgbClr')
    if srgb_elements:
        srgb_elements[0].append(parse_xml(alpha_xml))

    # Add Name Tag Text
    tf_name = name_tag.text_frame
    tf_name.text = f"🎥  {presenter_name}"
    p_name = tf_name.paragraphs[0]
    p_name.font.bold = True
    p_name.font.size = Pt(16)
    p_name.font.color.rgb = RGBColor(255, 255, 255)
    p_name.alignment = PP_ALIGN.CENTER

    # Cleanup temporary image
    if os.path.exists(tmp_img_path):
        os.remove(tmp_img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
