def create_slide(
    output_pptx_path: str,
    target_percentage: int = 94,
    main_number_text: str = "94%",
    title_text: str = "觀看 YouTube",
    subtitle_text: str = "18-44 歲的網路使用者\n每週至少觀看一次",
    accent_color_hex: str = "#00BFFF",  # Cyan accent
    **kwargs,
) -> str:
    """
    Creates a Keynote-style high-contrast minimalist metric reveal slide.
    Combines PIL for the background, Matplotlib for the vector-style progress ring, 
    and python-pptx for crisp typography.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from PIL import Image, ImageDraw
    import matplotlib.pyplot as plt

    # Helpers for color conversion
    def hex_to_rgb(hex_color):
        hex_color = hex_color.lstrip('#')
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    
    accent_rgb = hex_to_rgb(accent_color_hex)

    # 1. Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # 2. Generate Deep Gradient Background using PIL
    bg_path = "temp_bg.png"
    bg_width, bg_height = 1920, 1080
    bg_img = Image.new('RGB', (bg_width, bg_height))
    draw = ImageDraw.Draw(bg_img)
    # Draw linear gradient from deep charcoal to pure black
    for y in range(bg_height):
        ratio = y / bg_height
        r = int(25 * (1 - ratio))
        g = int(25 * (1 - ratio))
        b = int(30 * (1 - ratio))
        draw.line([(0, y), (bg_width, y)], fill=(r, g, b))
    bg_img.save(bg_path)

    # Insert background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # 3. Generate Minimalist Progress Ring using Matplotlib
    chart_path = "temp_chart.png"
    fig, ax = plt.subplots(figsize=(4.5, 4.5), subplot_kw=dict(aspect="equal"))
    
    data = [target_percentage, 100 - target_percentage]
    colors = [accent_color_hex, '#333333'] # Accent color and dark gray track
    
    # Create pie chart with a hole (Donut)
    wedges, texts = ax.pie(
        data, 
        colors=colors, 
        startangle=90, 
        counterclock=False,
        wedgeprops=dict(width=0.12, edgecolor='none') # 0.12 defines the ring thickness
    )
    
    # Make figures perfectly transparent
    fig.patch.set_visible(False)
    plt.savefig(chart_path, transparent=True, dpi=300, format='png', bbox_inches='tight', pad_inches=0)
    plt.close(fig)

    # Insert Chart Picture into PPTX
    chart_size = Inches(4.5)
    chart_left = Inches(1.5)
    chart_top = Inches(1.5)
    slide.shapes.add_picture(chart_path, chart_left, chart_top, chart_size, chart_size)

    # 4. Add Massive Typography inside the Ring
    tx_box = slide.shapes.add_textbox(chart_left, chart_top, chart_size, chart_size)
    tf = tx_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.clear()
    p = tf.paragraphs[0]
    p.text = main_number_text
    p.font.size = Pt(80)
    p.font.bold = True
    p.font.name = "Arial" # Best cross-platform clean sans-serif fallback
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # 5. Add Right-side Typography (Headline)
    text_left = Inches(6.5)
    
    headline_box = slide.shapes.add_textbox(text_left, Inches(2.8), Inches(6), Inches(1))
    h_tf = headline_box.text_frame
    h_p = h_tf.paragraphs[0]
    h_p.text = title_text
    h_p.font.size = Pt(44)
    h_p.font.bold = True
    h_p.font.name = "Arial"
    h_p.font.color.rgb = RGBColor(255, 255, 255)

    # 6. Add Right-side Typography (Context/Subtitle)
    sub_box = slide.shapes.add_textbox(text_left, Inches(3.8), Inches(6), Inches(2))
    s_tf = sub_box.text_frame
    s_p = s_tf.paragraphs[0]
    s_p.text = subtitle_text
    s_p.font.size = Pt(24)
    s_p.font.name = "Arial"
    s_p.font.color.rgb = RGBColor(180, 180, 180) # Light Gray

    # Save output
    prs.save(output_pptx_path)

    # Cleanup temporary files
    if os.path.exists(bg_path): os.remove(bg_path)
    if os.path.exists(chart_path): os.remove(chart_path)

    return output_pptx_path
