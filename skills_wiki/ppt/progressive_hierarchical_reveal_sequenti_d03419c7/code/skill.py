def create_slide(
    output_pptx_path: str,
    title_text: str = "Key Deliverables",
    bullets_data: list = None,
    bg_color: tuple = (20, 25, 35),      # Deep slate
    text_color: tuple = (255, 255, 255), # White
    sub_color: tuple = (180, 180, 180),  # Light Grey
    accent_color: tuple = (0, 191, 255), # Cyan
    **kwargs,
) -> str:
    """
    Creates a PPTX file replicating the progressive bullet reveal effect using "Slide Builds".
    It generates a sequence of slides, adding one bullet/sub-bullet per slide.
    
    bullets_data should be a list of tuples: (level, text)
    e.g., [(0, "Main Point 1"), (1, "Sub Point 1"), (1, "Sub Point 2"), (0, "Main Point 2")]
    """
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    
    if not bullets_data:
        bullets_data = [
            (0, "Phase 1: Market Research"),
            (1, "Competitor analysis"),
            (1, "Customer demographic surveys"),
            (0, "Phase 2: Product Development"),
            (1, "Initial prototyping"),
            (1, "Iterative QA testing"),
            (1, "Final design lock"),
            (0, "Phase 3: Go-to-Market Strategy"),
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank layout
    
    # We will generate one slide for each progressive step
    for step in range(1, len(bullets_data) + 1):
        slide = prs.slides.add_slide(blank_layout)
        
        # 1. Background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*bg_color)
        
        # 2. Side Accent Bar
        accent_bar = slide.shapes.add_shape(
            1, # Rectangle
            Inches(0), Inches(0), Inches(0.15), prs.slide_height
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = RGBColor(*accent_color)
        accent_bar.line.fill.background()
        
        # 3. Static Title (Left aligned, vertically centered-ish)
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1), Inches(4), Inches(2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = title_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*text_color)
        p.font.name = "Arial"
        
        # Optional: Add a subtle divider line under title
        divider = slide.shapes.add_shape(
            1, Inches(0.8), Inches(2.2), Inches(1.5), Inches(0.05)
        )
        divider.fill.solid()
        divider.fill.fore_color.rgb = RGBColor(*accent_color)
        divider.line.fill.background()

        # 4. Bullet Points Container (Right side)
        # We draw only up to the current 'step'
        content_box = slide.shapes.add_textbox(Inches(5.5), Inches(1), Inches(7), Inches(5.5))
        content_tf = content_box.text_frame
        content_tf.word_wrap = True
        
        current_items = bullets_data[:step]
        
        for idx, (level, text) in enumerate(current_items):
            # First item modifies the default paragraph, subsequent ones add new
            p = content_tf.paragraphs[0] if idx == 0 else content_tf.add_paragraph()
            p.text = text
            p.level = level
            p.font.name = "Arial"
            
            # Styling based on level
            if level == 0:
                p.font.size = Pt(28)
                p.font.bold = True
                p.font.color.rgb = RGBColor(*text_color)
                # Adding some space before new main points (except the first)
                if idx > 0:
                    p.space_before = Pt(20)
            else:
                p.font.size = Pt(20)
                p.font.bold = False
                p.font.color.rgb = RGBColor(*sub_color)
                p.space_before = Pt(8)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("progressive_reveal.pptx")
