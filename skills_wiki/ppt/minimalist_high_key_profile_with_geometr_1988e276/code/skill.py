def create_slide(
    output_pptx_path: str,
    title_text: str = "Marcus Brotz",
    body_text: str = "Director of Business Development\nLegion Enterprises",
    bg_palette: str = "business portrait, clean background",
    accent_color: tuple = (44, 130, 240),
    **kwargs,
) -> str:
    """
    Creates a High-Key Profile slide with a custom geometric faceted logo.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import urllib.request
    import os

    # Initialize presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: High-Key White Background ===
    bg_shape = slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, prs.slide_height # 1 is msoShapeRectangle
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_shape.line.fill.background() # No outline

    # === Layer 2: Geometric Faceted Logo (The "Legion" Style) ===
    # Draw a faceted 'L' shape using FreeformBuilder
    def draw_polygon(points, color_rgb):
        builder = slide.shapes.build_freeform()
        # Convert points to inches
        scaled_points = [(Inches(x), Inches(y)) for x, y in points]
        builder.add_line_segments(scaled_points)
        builder.add_line_segments([scaled_points[0]]) # Close shape
        shape = builder.convert_to_shape()
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color_rgb)
        shape.line.fill.background() # No line to make them look seamlessly connected
        return shape

    # Logo Base Position
    start_x, start_y = 1.0, 1.0
    scale = 0.3 # Size of the geometric units

    # Define polygons for a faceted 'L' shape
    # Colors extracted/inspired from the video logo
    col_light = (186, 196, 238)
    col_mid = (44, 130, 240)
    col_dark = (18, 52, 102)

    polygons = [
        # Vertical stem
        ([(start_x, start_y), (start_x + scale, start_y), (start_x, start_y + scale*1.5)], col_light),
        ([(start_x + scale, start_y), (start_x + scale, start_y + scale*1.5), (start_x, start_y + scale*1.5)], col_mid),
        ([(start_x, start_y + scale*1.5), (start_x + scale, start_y + scale*1.5), (start_x, start_y + scale*3)], col_mid),
        ([(start_x + scale, start_y + scale*1.5), (start_x + scale, start_y + scale*3), (start_x, start_y + scale*3)], col_dark),
        # Horizontal base
        ([(start_x + scale, start_y + scale*2), (start_x + scale*2, start_y + scale*2), (start_x + scale, start_y + scale*3)], col_mid),
        ([(start_x + scale*2, start_y + scale*2), (start_x + scale*2, start_y + scale*3), (start_x + scale, start_y + scale*3)], col_light),
        ([(start_x + scale*2, start_y + scale*2), (start_x + scale*3, start_y + scale*2.5), (start_x + scale*2, start_y + scale*3)], col_mid)
    ]

    for pts, color in polygons:
        draw_polygon(pts, color)

    # === Layer 3: Typography ===
    # Name
    name_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(5.0), Inches(1.0))
    tf_name = name_box.text_frame
    p_name = tf_name.paragraphs[0]
    p_name.text = title_text
    p_name.font.name = 'Arial'
    p_name.font.size = Pt(48)
    p_name.font.bold = True
    p_name.font.color.rgb = RGBColor(30, 30, 30)

    # Title / Body
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(5.0), Inches(1.5))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = body_text
    p_title.font.name = 'Arial'
    p_title.font.size = Pt(22)
    p_title.font.color.rgb = RGBColor(120, 120, 120)

    # === Layer 4: Subject Image (Right Side) ===
    # Attempt to download a suitable placeholder image
    img_path = "temp_portrait.jpg"
    try:
        url = "https://images.unsplash.com/photo-1560250097-0b93528c311a?q=80&w=800&auto=format&fit=crop"
        urllib.request.urlretrieve(url, img_path)
    except Exception:
        # Fallback if download fails: Create a clean gray placeholder
        img = Image.new('RGB', (800, 1000), color=(230, 230, 230))
        d = ImageDraw.Draw(img)
        d.text((300, 480), "Portrait Image\nPlaceholder", fill=(150, 150, 150))
        img.save(img_path)

    # Calculate positioning for the image (Right aligned, full height approx)
    # We want it to occupy the right half.
    pic = slide.shapes.add_picture(img_path, Inches(7.5), Inches(0.5), height=Inches(6.5))
    
    # Optional: Crop the image to a consistent aspect ratio if needed, 
    # but add_picture with just height maintains aspect ratio.
    
    # Cleanup temp file
    if os.path.exists(img_path):
        os.remove(img_path)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
