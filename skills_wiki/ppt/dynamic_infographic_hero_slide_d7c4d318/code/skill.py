import os
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

def create_infographic_hero_slide(
    output_pptx_path: str,
    hero_image_path: str,
    background_image_path: str,
    title: str,
    subtitle: str,
    logo_path: str,
    career_data: list,
    achievements_data: list,
    accent_color: tuple = (253, 185, 39) # Lakers Gold
) -> str:
    """
    Creates a PPTX file reproducing the "Infographic Hero Slide" pattern.

    Args:
        output_pptx_path: Path to save the final .pptx file.
        hero_image_path: Path to a transparent PNG of the main subject.
        background_image_path: Path to the background image.
        title: Main title text (e.g., "LeBron James").
        subtitle: Secondary title text (e.g., "LeBron James").
        logo_path: Path to a small logo image (e.g., LeBron's personal logo).
        career_data: List of dicts for career timeline.
                     e.g., [{'period': '2003-2010', 'team': 'Cleveland Cavaliers', 'logo': 'cavs.png'}]
        achievements_data: List of dicts for achievements.
                           e.g., [{'award': 'MVP', 'count': 4, 'icon': 'mvp.png'}]
        accent_color: RGB tuple for the background overlay color.

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Skill 3: Full-Bleed Photographic Background ---
    bg_pic = slide.shapes.add_picture(background_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    # Move to back
    slide.shapes._spTree.remove(bg_pic._element)
    slide.shapes._spTree.insert(2, bg_pic._element)

    # --- Skill 4: Semi-Transparent Overlay Mask ---
    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill = overlay.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_rgb(*accent_color)
    fill.transparency = 0.15 # 15% transparent
    overlay.line.fill.background()
    # Move to just above background
    slide.shapes._spTree.remove(overlay._element)
    slide.shapes._spTree.insert(3, overlay._element)

    # --- Skill 2: The Frame-Breaking Cutout ---
    # Position hero to the left, slightly enlarged to break the frame
    hero_height = Inches(6.5)
    hero_pic = slide.shapes.add_picture(hero_image_path, Inches(-0.5), Inches(1.0), height=hero_height)

    # --- Skill 5: Deconstructed Text-to-Table Layout ---
    # Add gray rounded rectangle panels as backing for text
    panel_y = Inches(1.25)
    panel_height = Inches(5.0)
    
    # Career Panel
    career_panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), panel_y, Inches(3.2), panel_height)
    fill = career_panel.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)
    fill.transparency = 0.2
    career_panel.line.fill.background()
    
    # Achievements Panel
    ach_panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), panel_y, Inches(3.0), panel_height)
    fill = ach_panel.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)
    fill.transparency = 0.2
    ach_panel.line.fill.background()

    # --- Skill 1: Expressive Typography (Applied with modern hierarchy) ---
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(6.5), Inches(0.5), Inches(5), Inches(0.5))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = 'Impact'
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(0, 0, 0)

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(6.5), Inches(0.9), Inches(5), Inches(0.5))
    p = sub_box.text_frame.paragraphs[0]
    p.text = subtitle
    p.font.name = 'Arial'
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(80, 80, 80)
    
    # Logo
    slide.shapes.add_picture(logo_path, Inches(11.5), Inches(0.5), height=Inches(0.5))

    # Populate Career Data
    current_y = panel_y + Inches(0.3)
    for item in career_data:
        # Period
        tb = slide.shapes.add_textbox(Inches(6.7), current_y, Inches(1.5), Inches(0.3))
        tb.text_frame.paragraphs[0].text = item['period']
        tb.text_frame.paragraphs[0].font.size = Pt(12)
        tb.text_frame.paragraphs[0].font.bold = True
        # Team Name
        tb = slide.shapes.add_textbox(Inches(6.7), current_y + Inches(0.2), Inches(1.5), Inches(0.3))
        tb.text_frame.paragraphs[0].text = item['team']
        tb.text_frame.paragraphs[0].font.size = Pt(10)
        # Logo
        slide.shapes.add_picture(item['logo'], Inches(8.9), current_y + Inches(0.1), height=Inches(0.3))
        current_y += Inches(0.8)
        
    # Populate Achievements Data
    current_y = panel_y + Inches(0.3)
    for item in achievements_data:
        # Award Name
        tb = slide.shapes.add_textbox(Inches(10.2), current_y, Inches(1.5), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.text = item['award']
        p.font.size = Pt(14)
        p.font.bold = True
        
        # Icons
        icon_x = Inches(10.2)
        for i in range(item['count']):
            slide.shapes.add_picture(item['icon'], icon_x, current_y + Inches(0.3), height=Inches(0.5))
            icon_x += Inches(0.6)
            
        current_y += Inches(1.2)

    prs.save(output_pptx_path)
    return output_pptx_path

def download_asset(url, filename):
    if not os.path.exists(filename):
        print(f"Downloading {filename}...")
        urllib.request.urlretrieve(url, filename)
    return filename

# Example Usage:
if __name__ == '__main__':
    # Create dummy assets for demonstration
    from PIL import Image, ImageDraw
    
    # Dummy hero image (transparent)
    hero_img = Image.new('RGBA', (400, 600), (0, 0, 0, 0))
    draw = ImageDraw.Draw(hero_img)
    draw.ellipse((50, 50, 350, 550), fill='purple')
    hero_img.save('hero.png')

    # Dummy background
    bg_img = Image.new('RGB', (1280, 720), 'gray')
    bg_img.save('background.jpg')
    
    # Dummy logos and icons
    logo_img = Image.new('RGBA', (100, 100), 'white')
    logo_img.save('logo.png')
    
    team_logo = Image.new('RGBA', (100, 100), (255,0,0,128))
    team_logo.save('team_logo.png')
    
    icon_img = Image.new('RGBA', (100, 100), (0,255,0,128))
    icon_img.save('icon.png')

    # Data for the slide
    career = [
        {'period': '2003-2010', 'team': 'Cleveland Cavaliers', 'logo': 'team_logo.png'},
        {'period': '2010-2014', 'team': 'Miami Heat', 'logo': 'team_logo.png'},
        {'period': '2014-2018', 'team': 'Cleveland Cavaliers', 'logo': 'team_logo.png'},
        {'period': '2018-', 'team': 'Los Angeles Lakers', 'logo': 'team_logo.png'},
    ]

    achievements = [
        {'award': 'MVP', 'count': 4, 'icon': 'icon.png'},
        {'award': 'Finals MVP', 'count': 4, 'icon': 'icon.png'},
        {'award': 'Championships', 'count': 4, 'icon': 'icon.png'},
    ]
    
    # Generate the slide
    create_infographic_hero_slide(
        output_pptx_path='LeBron_James_Infographic.pptx',
        hero_image_path='hero.png',
        background_image_path='background.jpg',
        title='LeBron James',
        subtitle='King James',
        logo_path='logo.png',
        career_data=career,
        achievements_data=achievements
    )

    print("PPTX file 'LeBron_James_Infographic.pptx' created successfully.")

