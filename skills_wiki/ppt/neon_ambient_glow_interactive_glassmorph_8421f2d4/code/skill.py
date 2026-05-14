import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "Hypothesis",
    bg_color: tuple = (20, 11, 18),         # Dark Purple/Black (#140B12)
    glow_color: tuple = (0, 255, 179),      # Neon Green (#00FFB3)
    **kwargs,
) -> str:
    """
    Creates a PPTX slide reproducing the "Neon Ambient Glow & Glassmorphism Menu" style.
    Generates complex visual elements using PIL and composite them in python-pptx.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # --- Layer 1: Dark Background ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # --- Layer 2: Generate and Insert Ambient Glow (PIL) ---
    glow_filename = "temp_glow.png"
    img_size = 1000
    glow_img = Image.new('RGBA', (img_size, img_size), (0, 0, 0, 0))
    draw_glow = ImageDraw.Draw(glow_img)
    
    # Draw a solid circle in the center
    circle_radius = 180
    center = img_size // 2
    bbox = [center - circle_radius, center - circle_radius, 
            center + circle_radius, center + circle_radius]
    
    # Use the requested glow color, fully opaque
    draw_glow.ellipse(bbox, fill=(glow_color[0], glow_color[1], glow_color[2], 255))
    
    # Apply a heavy Gaussian blur to simulate the "Soft Edges" effect
    blur_radius = 120
    glow_img = glow_img.filter(ImageFilter.GaussianBlur(radius=blur_radius))
    glow_img.save(glow_filename)

    # Insert glow into the center of the slide
    # Slide is 13.333 x 7.5 inches. Image is 1000x1000. Let's make it 8x8 inches.
    glow_size = Inches(8)
    left = (prs.slide_width - glow_size) / 2
    top = (prs.slide_height - glow_size) / 2
    slide.shapes.add_picture(glow_filename, left, top, width=glow_size, height=glow_size)

    # --- Layer 3: Generate and Insert Glassmorphism Menu (PIL) ---
    menu_filename = "temp_menu.png"
    menu_w, menu_h = 600, 80
    menu_img = Image.new('RGBA', (menu_w, menu_h), (0, 0, 0, 0))
    draw_menu = ImageDraw.Draw(menu_img)
    
    # Pill shape parameters
    # Fill: White, very transparent (~6% opacity -> alpha 15)
    # Outline: White, slightly more visible (~20% opacity -> alpha 50)
    draw_menu.rounded_rectangle(
        [2, 2, menu_w - 2, menu_h - 2], 
        radius=38, 
        fill=(255, 255, 255, 15), 
        outline=(255, 255, 255, 50), 
        width=2
    )
    menu_img.save(menu_filename)

    # Insert menu at the top center
    menu_width_in = Inches(5.5)
    menu_height_in = Inches(0.75)
    menu_left = (prs.slide_width - menu_width_in) / 2
    menu_top = Inches(0.4)
    slide.shapes.add_picture(menu_filename, menu_left, menu_top, width=menu_width_in, height=menu_height_in)

    # Add pseudo-icons (text placeholders for icons) inside the menu
    # In a real scenario, you'd insert actual SVG/PNG icons here.
    icon_box = slide.shapes.add_textbox(menu_left, menu_top + Inches(0.1), menu_width_in, Inches(0.5))
    tf = icon_box.text_frame
    tf.text = "  Δ        🎯        ⚙️  "  # Using unicode characters as stand-in icons
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # --- Layer 4: Central Typography ---
    title_width = Inches(8)
    title_height = Inches(1.5)
    title_left = (prs.slide_width - title_width) / 2
    title_top = Inches(2.2)
    
    title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    
    p = tf_title.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial" # Fallback for Darker Grotesque
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Add an active indicator text under the title (like in the video)
    indicator_box = slide.shapes.add_textbox(title_left, title_top + Inches(1.2), title_width, Inches(0.5))
    tf_ind = indicator_box.text_frame
    p_ind = tf_ind.paragraphs[0]
    p_ind.text = "ACTIVE SECTION"
    p_ind.alignment = PP_ALIGN.CENTER
    p_ind.font.name = "Arial"
    p_ind.font.size = Pt(14)
    p_ind.font.color.rgb = RGBColor(255, 255, 255)

    # --- Cleanup temp files ---
    prs.save(output_pptx_path)
    if os.path.exists(glow_filename):
        os.remove(glow_filename)
    if os.path.exists(menu_filename):
        os.remove(menu_filename)
        
    return output_pptx_path
