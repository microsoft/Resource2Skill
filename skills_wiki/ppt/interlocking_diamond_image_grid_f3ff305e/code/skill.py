def create_slide(
    output_pptx_path: str,
    title_text: str = "New Product\nDetailed\nAnalysis",
    subtitle_text: str = "Your Company Name",
    theme_keyword: str = "corporate",
    accent_color: tuple = (142, 40, 54, 255),  # Maroon
    navy_color: tuple = (34, 52, 70, 255),     # Dark Navy
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Interlocking Diamond Image Grid effect.
    """
    import os
    import tempfile
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Helper: Convert RGBA tuple to RGBColor
    def to_rgb(color_tuple):
        return RGBColor(color_tuple[0], color_tuple[1], color_tuple[2])

    # === Helper Functions for Image Generation ===
    
    def fetch_image(url, fallback_color=(200, 200, 200)):
        """Fetch image from URL or return a solid color fallback."""
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            res = urllib.request.urlopen(req, timeout=5)
            return Image.open(BytesIO(res.read())).convert("RGB")
        except Exception:
            return Image.new("RGB", (800, 800), fallback_color)

    def create_diamond_mask_png(img, size_px, filepath):
        """Crops an image into a diamond shape and saves as transparent PNG."""
        w, h = img.size
        min_side = min(w, h)
        # Center crop to square
        img = img.crop(((w-min_side)//2, (h-min_side)//2, (w+min_side)//2, (h+min_side)//2))
        img = img.resize((size_px, size_px), Image.Resampling.LANCZOS)

        mask = Image.new("L", (size_px, size_px), 0)
        draw = ImageDraw.Draw(mask)
        # Draw diamond polygon
        draw.polygon([(size_px/2, 0), (size_px, size_px/2), (size_px/2, size_px), (0, size_px/2)], fill=255)

        out = Image.new("RGBA", (size_px, size_px), (0, 0, 0, 0))
        out.paste(img, (0, 0), mask)
        out.save(filepath)

    def create_hollow_diamond_png(size_px, filepath, color_rgba, thickness=8):
        """Creates a hollow diamond outline as a transparent PNG."""
        out = Image.new("RGBA", (size_px, size_px), (0, 0, 0, 0))
        draw = ImageDraw.Draw(out)
        inset = thickness
        poly = [
            (size_px/2, inset),
            (size_px - inset, size_px/2),
            (size_px/2, size_px - inset),
            (inset, size_px/2)
        ]
        draw.polygon(poly, outline=color_rgba, width=thickness)
        out.save(filepath)

    # === Build the Slide ===
    temp_dir = tempfile.mkdtemp()
    try:
        # --- 1. Typography & Lines (Left Side) ---
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6), Inches(3.5))
        tf = tx_box.text_frame
        
        # Split title into lines to handle them
        title_lines = title_text.split('\n')
        for i, text_line in enumerate(title_lines):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = text_line
            p.font.size = Pt(54)
            p.font.bold = True
            p.font.name = 'Arial'
            p.font.color.rgb = RGBColor(0, 0, 0)
            
        # Maroon divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.2), Inches(2.5), Pt(2))
        line.fill.solid()
        line.fill.fore_color.rgb = to_rgb(accent_color)
        line.line.fill.background()

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(5), Inches(1))
        p_sub = sub_box.text_frame.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(22)
        p_sub.font.italic = True
        p_sub.font.name = 'Arial'
        p_sub.font.color.rgb = RGBColor(100, 100, 100)

        # --- 2. Diamond Grid Generation (Right Side) ---
        # Grid parameters
        D = 2.8 # Diameter/Width of individual diamond in inches
        gap = 0.1 # Gap between edges
        delta = (D / 2) + gap # Distance offset for adjacent touching diamonds
        cx, cy = 9.5, 3.8 # Center coordinate of the main cluster
        
        # Download images
        img1 = fetch_image("https://images.unsplash.com/photo-1522071820081-009f0129c71c?ixlib=rb-4.0.3&auto=format&fit=crop&w=800&q=80", (200, 200, 200))
        img2 = fetch_image("https://images.unsplash.com/photo-1573164713988-8665fc963095?ixlib=rb-4.0.3&auto=format&fit=crop&w=800&q=80", (180, 180, 180))
        img3 = fetch_image("https://images.unsplash.com/photo-1498050108023-c5249f4df085?ixlib=rb-4.0.3&auto=format&fit=crop&w=800&q=80", (160, 160, 160))

        # Generate PNGs
        path_img1 = os.path.join(temp_dir, "img1.png")
        path_img2 = os.path.join(temp_dir, "img2.png")
        path_img3 = os.path.join(temp_dir, "img3.png")
        path_outline_bg = os.path.join(temp_dir, "out_bg.png")
        path_outline_fg = os.path.join(temp_dir, "out_fg.png")
        
        create_diamond_mask_png(img1, 800, path_img1)
        create_diamond_mask_png(img2, 800, path_img2)
        create_diamond_mask_png(img3, 800, path_img3)
        create_hollow_diamond_png(1000, path_outline_bg, accent_color, thickness=12) # Large Maroon
        create_hollow_diamond_png(800, path_outline_fg, (200, 200, 200, 255), thickness=6) # Center Grey

        # Place Elements (Order matters for Z-index)
        
        # Background Accent Outline (Large Maroon)
        w_bg = D * 1.5
        slide.shapes.add_picture(path_outline_bg, Inches(cx - delta/2 - w_bg/2), Inches(cy - delta/2 - w_bg/2), Inches(w_bg), Inches(w_bg))

        # Position TL: Image 1
        slide.shapes.add_picture(path_img1, Inches(cx - delta - D/2), Inches(cy - delta - D/2), Inches(D), Inches(D))
        
        # Position TR: Image 2
        slide.shapes.add_picture(path_img2, Inches(cx + delta - D/2), Inches(cy - delta - D/2), Inches(D), Inches(D))
        
        # Position BR: Image 3
        slide.shapes.add_picture(path_img3, Inches(cx + delta - D/2), Inches(cy + delta - D/2), Inches(D), Inches(D))

        # Position BL: Solid Navy Diamond with Icon
        navy_dia = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(cx - delta - D/2), Inches(cy + delta - D/2), Inches(D), Inches(D))
        navy_dia.fill.solid()
        navy_dia.fill.fore_color.rgb = to_rgb(navy_color)
        navy_dia.line.fill.background()
        
        # Add white icon inside the Navy Diamond
        icon_size = 0.8
        icon = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(cx - delta - icon_size/2), Inches(cy + delta - icon_size/2), Inches(icon_size), Inches(icon_size))
        icon.fill.background() # Hollow
        icon.line.color.rgb = RGBColor(255, 255, 255)
        icon.line.width = Pt(2)
        
        # Position Center: Hollow Grey Outline overlaying the cluster
        slide.shapes.add_picture(path_outline_fg, Inches(cx - D/2), Inches(cy - D/2), Inches(D), Inches(D))

    finally:
        # Cleanup temporary files
        import shutil
        shutil.rmtree(temp_dir, ignore_errors=True)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
