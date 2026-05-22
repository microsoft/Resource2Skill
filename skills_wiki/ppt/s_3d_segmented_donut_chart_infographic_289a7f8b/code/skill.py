def create_slide(
    output_pptx_path: str,
    title_text: str = "Pie Chart",
    body_text: str = "",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 3D Segmented Donut Chart visual effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageFilter
    from lxml import etree
    import math
    import os

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(235, 240, 245)

    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.5), Inches(9.333), Inches(1))
    tf = title_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].font.name = 'Century Gothic'
    tf.paragraphs[0].font.color.rgb = RGBColor(40, 40, 40)

    cx, cy = 6.666, 4.0  # Chart center coordinates

    # === Layer 2: PIL Ambient Drop Shadow ===
    shadow_path = "temp_ambient_shadow.png"
    shadow_img = Image.new('RGBA', (800, 400), (0, 0, 0, 0))
    draw = ImageDraw.Draw(shadow_img)
    # Draw highly squashed ellipse representing the perspective base
    draw.ellipse([200, 150, 600, 250], fill=(0, 0, 0, 80))
    shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(30))
    shadow_img.save(shadow_path)
    
    # Place shadow behind and slightly below the 3D object
    slide.shapes.add_picture(shadow_path, Inches(cx - 3.5), Inches(cy - 0.7), Inches(7.0), Inches(3.5))
    if os.path.exists(shadow_path):
        os.remove(shadow_path)

    # === XML Helper Functions ===
    def apply_3d_effect(shape, depth_pt=35):
        spPr = shape.element.find('.//a:spPr', namespaces=shape.element.nsmap)
        
        # Clear existing 3D properties if any
        for tag in ['a:scene3d', 'a:sp3d']:
            el = spPr.find(tag, namespaces=shape.element.nsmap)
            if el is not None: spPr.remove(el)
                
        # Inject Camera & Light Rig (Tilt 60 deg back, Spin 30 deg counter-clockwise)
        scene3d = etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}scene3d')
        camera = etree.SubElement(scene3d, '{http://schemas.openxmlformats.org/drawingml/2006/main}camera', prst="perspectiveRelaxed")
        etree.SubElement(camera, '{http://schemas.openxmlformats.org/drawingml/2006/main}rot', lat="18000000", lon="0", rev="19800000")
        etree.SubElement(scene3d, '{http://schemas.openxmlformats.org/drawingml/2006/main}lightRig', rig="threePt", dir="t")
        
        # Inject Extrusion (Depth) and Bevel
        sp3d = etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}sp3d', extrusionH=str(int(depth_pt * 12700)))
        etree.SubElement(sp3d, '{http://schemas.openxmlformats.org/drawingml/2006/main}bevelT', w="38100", h="38100", prst="circle")
        
        # Safely insert before extLst
        extLst = spPr.find('a:extLst', namespaces=shape.element.nsmap)
        if extLst is not None:
            extLst.addprevious(scene3d)
            extLst.addprevious(sp3d)
        else:
            spPr.append(scene3d)
            spPr.append(sp3d)

    def set_arc_geometry(shape, start_deg, end_deg, inner_pct=45):
        prstGeom = shape.element.find('.//a:prstGeom', namespaces=shape.element.nsmap)
        avLst = prstGeom.find('a:avLst', namespaces=shape.element.nsmap)
        if avLst is not None: prstGeom.remove(avLst)
        avLst = etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
        
        etree.SubElement(avLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd', name='adj1', fmla=f'val {int(start_deg * 60000)}')
        etree.SubElement(avLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd', name='adj2', fmla=f'val {int(end_deg * 60000)}')
        etree.SubElement(avLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd', name='adj3', fmla=f'val {int(inner_pct * 1000)}')
        prstGeom.insert(0, avLst)

    # === Layer 3: The 3D Chart ===
    radius = 2.25
    w = h = radius * 2
    left = cx - radius
    top = cy - radius

    colors = [
        (217, 35, 35),   # Red
        (38, 160, 217),  # Light Blue
        (42, 87, 140),   # Navy
        (242, 140, 56),  # Orange
        (69, 103, 52),   # Green
        (242, 194, 48)   # Yellow
    ]

    # Calculate properties and Z-order (draw back to front)
    slices_data = []
    for i in range(6):
        start_deg = i * 60
        mid_deg = start_deg + 30
        
        # Calculate apparent position factoring in the 30-degree Z-spin
        vis_mid_deg = (mid_deg - 30) % 360 
        y_rel = math.sin(math.radians(vis_mid_deg)) # Y dictates front/back occlusion
        
        slices_data.append({
            'start': start_deg, 'end': start_deg + 60,
            'vis_mid': vis_mid_deg, 'y_rel': y_rel, 'color': colors[i]
        })
        
    slices_data.sort(key=lambda x: x['y_rel']) # Process back shapes first

    hub_drawn = False
    for sd in slices_data:
        # Draw the center Hub right before we draw the mid/front slices
        if not hub_drawn and sd['y_rel'] >= -0.001:
            hub_radius = radius * 0.45
            hub = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - hub_radius), Inches(cy - hub_radius), Inches(hub_radius * 2), Inches(hub_radius * 2))
            hub.fill.solid()
            hub.fill.fore_color.rgb = RGBColor(210, 215, 220)
            hub.line.fill.background()
            apply_3d_effect(hub, depth_pt=35)
            hub_drawn = True
            
        # Draw fragment slice
        shape = slide.shapes.add_shape(MSO_SHAPE.BLOCK_ARC, Inches(left), Inches(top), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*sd['color'])
        shape.line.fill.background() # No Line
        set_arc_geometry(shape, sd['start'], sd['end'], inner_pct=45)
        apply_3d_effect(shape, depth_pt=35)

    # === Layer 4: Annotations ===
    for sd in slices_data:
        vis_mid = sd['vis_mid']
        squash = 0.5  # Cosine of 60 deg tilt
        
        # Connector anchor calculation
        rad = math.radians(vis_mid)
        sx = cx + 1.6 * math.cos(rad) # Start inside the 3D footprint
        sy = cy + 1.6 * squash * math.sin(rad)
        ex = cx + 3.2 * math.cos(rad) # End outside the 3D footprint
        ey = cy + 3.2 * squash * math.sin(rad)
        
        is_left = math.cos(rad) < -0.01
        
        # Line Segment
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(sx), Inches(sy), Inches(ex), Inches(ey))
        line.line.color.rgb = RGBColor(80, 80, 80)
        line.line.width = Pt(1.5)
        
        # Anchor Dot
        dot_r = 0.05
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ex - dot_r), Inches(ey - dot_r), Inches(dot_r*2), Inches(dot_r*2))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(80, 80, 80)
        dot.line.fill.background()
        
        # Text block Layout
        tw, th = 2.0, 1.0
        tx = ex - tw - 0.2 if is_left else ex + 0.2
        ty = ey - 0.4
        
        tb = slide.shapes.add_textbox(Inches(tx), Inches(ty), Inches(tw), Inches(th))
        
        # Header
        p = tb.text_frame.paragraphs[0]
        p.text = "Title Here"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.name = 'Century Gothic'
        p.alignment = PP_ALIGN.RIGHT if is_left else PP_ALIGN.LEFT
        
        # Body
        p2 = tb.text_frame.add_paragraph()
        p2.text = "Add details in 2-3 lines to describe the title. Lesser the content better it will look."
        p2.font.size = Pt(10)
        p2.font.name = 'Century Gothic'
        p2.font.color.rgb = RGBColor(100, 100, 100)
        p2.alignment = PP_ALIGN.RIGHT if is_left else PP_ALIGN.LEFT

    prs.save(output_pptx_path)
    return output_pptx_path
