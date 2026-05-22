def create_slide(
    output_pptx_path: str,
    title_text: str = "THE TRUE STORY\nOF DESIGN",
    body_text: str = "",
    bg_palette: str = "fashion",  # keyword for background image theme
    accent_color: tuple = (0, 174, 239),  # Vibrant Cyan
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Editorial Magazine Overlap Spread' visual effect.
    """
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from PIL import Image

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background Image ===
    # Fetch a thematic image to act as the right-side bleed
    img_width_px, img_height_px = 1280, 720
    url = f"https://loremflickr.com/{img_width_px}/{img_height_px}/{bg_palette.replace(' ', ',')}"
    
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            image_stream = io.BytesIO(response.read())
            img = Image.open(image_stream).convert("RGB")
            temp_img = io.BytesIO()
            img.save(temp_img, format="JPEG")
            temp_img.seek(0)
            slide.shapes.add_picture(temp_img, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    except Exception:
        # Fallback if image download fails
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(40, 40, 40)
        bg_shape.line.fill.background()

    # === Layer 2: Left Editorial Panel (White) ===
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(5.0), Inches(7.5))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    panel.line.fill.background()

    # === Layer 3: Vertical Watermark Typography ===
    watermark = slide.shapes.add_textbox(Inches(-2.0), Inches(3.25), Inches(5.0), Inches(1.0))
    watermark.rotation = -90
    tf_wm = watermark.text_frame
    p_wm = tf_wm.paragraphs[0]
    p_wm.text = "EDITORIAL"
    p_wm.font.name = 'Arial Black'
    p_wm.font.size = Pt(64)
    p_wm.font.color.rgb = RGBColor(242, 242, 242) # Very subtle light gray

    # === Layer 4: Overlapping Accent Header Block ===
    header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.5), Inches(8.0), Inches(1.8))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*accent_color)
    header_shape.line.fill.background()

    # Inject OpenXML to add a drop shadow to the header block for 3D depth
    try:
        from lxml import etree
        spPr = header_shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw',
                                     blurRad="100000", dist="50000", dir="5400000", algn="tl", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="30000")
    except ImportError:
        pass # Graceful degradation if lxml is missing

    # === Layer 5: Text Hierarchy ===
    # 5a. Headline inside accent block
    tf_head = header_shape.text_frame
    tf_head.word_wrap = True
    tf_head.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_head.margin_left = Inches(0.8) 
    
    p_head = tf_head.paragraphs[0]
    p_head.text = title_text.upper()
    p_head.font.name = 'Arial'
    p_head.font.size = Pt(36)
    p_head.font.bold = True
    p_head.font.color.rgb = RGBColor(255, 255, 255)

    # 5b. Eyebrow Tag
    eyebrow = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(4.0), Inches(0.5))
    tf_eye = eyebrow.text_frame
    p_eye = tf_eye.paragraphs[0]
    p_eye.text = "LIFESTYLE & DESIGN"
    p_eye.font.name = 'Arial'
    p_eye.font.size = Pt(10)
    p_eye.font.bold = True
    p_eye.font.color.rgb = RGBColor(*accent_color)

    # 5c. Divider Rule
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.6), Inches(1.5), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # 5d. Body Paragraph
    body_shape = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(3.8), Inches(2.5))
    tf_body = body_shape.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    default_body = "It is a long established fact that a reader will be distracted by the readable content of a page when looking at its layout. The point of using Lorem Ipsum is that it has a more-or-less normal distribution of letters, as opposed to using 'Content here, content here', making it look like readable English."
    p_body.text = body_text if body_text else default_body
    p_body.font.name = 'Arial'
    p_body.font.size = Pt(10)
    p_body.font.color.rgb = RGBColor(100, 100, 100)
    p_body.alignment = PP_ALIGN.JUSTIFY

    # 5e. Folio / Metadata
    folio = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(3.0), Inches(0.3))
    tf_folio = folio.text_frame
    p_folio = tf_folio.paragraphs[0]
    p_folio.text = "01 // MULTIPURPOSE TEMPLATE"
    p_folio.font.name = 'Arial'
    p_folio.font.size = Pt(8)
    p_folio.font.bold = True
    p_folio.font.color.rgb = RGBColor(180, 180, 180)

    prs.save(output_pptx_path)
    return output_pptx_path
