def create_slide(
    output_pptx_path: str,
    title_text: str = "HERE ARE ALL OF THE",
    subtitle_text: str = "UNDERSEA CABLES",
    highlight_text: str = "THAT POWER THE INTERNET",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Neon Global Network Overlay' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFilter
    import random
    import math
    import os

    # Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === PIL GENERATION: Background, Globe, and Glowing Cables ===
    width, height = 1920, 1080
    bg_color = (5, 8, 15, 255) # Deep space blue/black
    
    # Create base canvas
    canvas = Image.new("RGBA", (width, height), bg_color)
    draw = ImageDraw.Draw(canvas)

    # 1. Generate Starfield
    for _ in range(300):
        x = random.randint(0, width)
        y = random.randint(0, height)
        brightness = random.randint(50, 255)
        size = random.choice([1, 1, 1, 2]) # Mostly small, some slightly larger
        draw.ellipse([x, y, x+size, y+size], fill=(255, 255, 255, brightness))

    # 2. Draw the Globe (Base and Atmosphere)
    globe_center = (width // 2, height // 2)
    globe_radius = 450
    
    # Atmosphere Glow (Blurred larger circle)
    glow_img = Image.new("RGBA", (width, height), (0,0,0,0))
    glow_draw = ImageDraw.Draw(glow_img)
    glow_radius = globe_radius + 40
    glow_draw.ellipse(
        [globe_center[0] - glow_radius, globe_center[1] - glow_radius, 
         globe_center[0] + glow_radius, globe_center[1] + glow_radius],
        fill=(0, 150, 255, 60)
    )
    glow_img = glow_img.filter(ImageFilter.GaussianBlur(30))
    canvas.alpha_composite(glow_img)

    # Globe Solid Base
    draw.ellipse(
        [globe_center[0] - globe_radius, globe_center[1] - globe_radius, 
         globe_center[0] + globe_radius, globe_center[1] + globe_radius],
        fill=(10, 25, 45, 255) # Dark ocean blue
    )

    # 3. Draw Network Cables (Bezier Curves with Glow)
    # Helper to calculate quadratic bezier points
    def get_bezier_curve(p0, p1, p2, num_points=100):
        points = []
        for t in range(num_points + 1):
            t = t / num_points
            x = (1 - t)**2 * p0[0] + 2 * (1 - t) * t * p1[0] + t**2 * p2[0]
            y = (1 - t)**2 * p0[1] + 2 * (1 - t) * t * p1[1] + t**2 * p2[1]
            points.append((x, y))
        return points

    # Define cable routes (start, end, control point for curve) and colors
    cables = [
        # (Start, End, Control Point, Color)
        ((globe_center[0]-300, globe_center[1]-100), (globe_center[0]+100, globe_center[1]-250), (globe_center[0]-100, globe_center[1]-300), (0, 255, 255)), # Cyan
        ((globe_center[0]-250, globe_center[1]+150), (globe_center[0]+200, globe_center[1]-50), (globe_center[0], globe_center[1]+250), (255, 0, 255)),   # Magenta
        ((globe_center[0]+50, globe_center[1]+300), (globe_center[0]+350, globe_center[1]+100), (globe_center[0]+250, globe_center[1]+350), (50, 255, 50)),    # Lime
        ((globe_center[0]-350, globe_center[1]+50), (globe_center[0]-100, globe_center[1]+350), (globe_center[0]-300, globe_center[1]+250), (255, 80, 50)),  # Orange
        ((globe_center[0]-150, globe_center[1]-200), (globe_center[0]+300, globe_center[1]+150), (globe_center[0]+150, globe_center[1]-200), (0, 191, 255)), # Light Blue
    ]

    cables_layer = Image.new("RGBA", (width, height), (0,0,0,0))
    cables_draw = ImageDraw.Draw(cables_layer)

    for start, end, control, color in cables:
        curve_points = get_bezier_curve(start, control, end)
        
        # Draw outer glow (thick, semi-transparent)
        glow_color = (*color, 80)
        cables_draw.line(curve_points, fill=glow_color, width=12, joint="curve")
        
        # Draw core line (thin, opaque)
        cables_draw.line(curve_points, fill=(255, 255, 255, 255), width=3, joint="curve")
        
        # Draw End Nodes (glow + white core)
        for point in [start, end]:
            cables_draw.ellipse([point[0]-10, point[1]-10, point[0]+10, point[1]+10], fill=(*color, 150))
            cables_draw.ellipse([point[0]-4, point[1]-4, point[0]+4, point[1]+4], fill=(255, 255, 255, 255))

    # Blur the glow layer slightly for better effect, then composite
    cables_layer = cables_layer.filter(ImageFilter.GaussianBlur(1))
    canvas.alpha_composite(cables_layer)

    # Save PIL image
    temp_bg_path = "temp_globe_network_bg.png"
    canvas.save(temp_bg_path, "PNG")

    # === PPTX ASSEMBLY ===
    
    # 1. Add Background Image to Slide
    slide.shapes.add_picture(temp_bg_path, 0, 0, width=Inches(13.333), height=Inches(7.5))

    # 2. Add Typography
    # In the video, the text is centered or left-aligned, stacked, with specific coloring.
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    # Line 1: HERE ARE ALL OF THE
    p1 = tf.paragraphs[0]
    run1 = p1.add_run()
    run1.text = title_text + "\n"
    run1.font.name = 'Arial'
    run1.font.size = Pt(44)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(200, 200, 200) # Slightly off-white

    # Line 2: UNDERSEA CABLES
    run2 = p1.add_run()
    run2.text = subtitle_text + "\n"
    run2.font.name = 'Arial'
    run2.font.size = Pt(64)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(255, 255, 255) # Pure white

    # Line 3: THAT POWER THE INTERNET
    run3 = p1.add_run()
    run3.text = highlight_text
    run3.font.name = 'Arial'
    run3.font.size = Pt(44)
    run3.font.bold = True
    run3.font.color.rgb = RGBColor(0, 191, 255) # Cyan highlight

    # Clean up temp file
    if os.path.exists(temp_bg_path):
        os.remove(temp_bg_path)

    prs.save(output_pptx_path)
    return output_pptx_path
