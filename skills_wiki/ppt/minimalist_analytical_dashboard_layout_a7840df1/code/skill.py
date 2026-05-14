def create_slide(
    output_pptx_path: str,
    title_text: str = "Sales Dashboard",
    body_text: str = "",
    bg_palette: str = "light", 
    accent_color: tuple = (90, 150, 150),  # Muted Teal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Minimalist Analytical Dashboard layout.
    Generates custom clean charts via matplotlib and a transparent hero image via PIL.
    """
    import os
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    import matplotlib.pyplot as plt
    import numpy as np
    from PIL import Image, ImageDraw

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Colors ---
    BG_COLOR = RGBColor(248, 249, 250)
    DARK_TEXT = RGBColor(51, 51, 51)
    LIGHT_TEXT = RGBColor(136, 136, 136)
    ACCENT_RGB = RGBColor(*accent_color)
    ACCENT_HEX = '#%02x%02x%02x' % accent_color
    LIGHT_GRAY_HEX = '#E0E0E0'

    # --- Background ---
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.fill.background()

    # --- Helper: Add Text ---
    def add_text(x, y, w, h, text, font_size, color, bold=False, align=PP_ALIGN.LEFT):
        tx_box = slide.shapes.add_textbox(x, y, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Segoe UI"
        return tx_box

    # --- Header & Title ---
    add_text(Inches(0.5), Inches(0.3), Inches(3), Inches(0.5), title_text, 20, DARK_TEXT, bold=True)
    
    # Slicer Tabs (2019, 2020, 2021, 2022)
    years = ["2019", "2020", "2021", "2022"]
    start_x = 9.5
    for i, year in enumerate(years):
        is_active = (year == "2022")
        tab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(start_x + i*0.8), Inches(0.35), Inches(0.7), Inches(0.3))
        tab.fill.solid()
        tab.fill.fore_color.rgb = DARK_TEXT if is_active else BG_COLOR
        tab.line.color.rgb = DARK_TEXT if is_active else LIGHT_TEXT
        
        tf = tab.text_frame
        p = tf.paragraphs[0]
        p.text = year
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = is_active
        p.font.color.rgb = RGBColor(255, 255, 255) if is_active else LIGHT_TEXT

    # --- KPI Ribbon ---
    kpis = [
        ("746K", "SALE"),
        ("21.4%", "GROWTH"),
        ("614K", "PRE SALE"),
        ("473K", "BUDGET")
    ]
    
    for i, (val, label) in enumerate(kpis):
        x_pos = Inches(1.5 + i*2.8)
        add_text(x_pos, Inches(1.0), Inches(2), Inches(0.5), val, 26, DARK_TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(x_pos, Inches(1.4), Inches(2), Inches(0.3), label, 10, LIGHT_TEXT, bold=False, align=PP_ALIGN.CENTER)

    # Divider Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.8), Inches(11.333), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 220, 220)
    line.line.fill.background()

    # --- Hero Section (PIL Transparent Car Placeholder) ---
    def generate_car_placeholder():
        img = Image.new('RGBA', (800, 400), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        # Draw stylized sports car silhouette
        car_color = (220, 225, 230, 255)
        # Body
        draw.polygon([(150, 200), (300, 120), (500, 120), (650, 200), (700, 250), (700, 300), (100, 300), (100, 250)], fill=car_color)
        # Windows
        draw.polygon([(310, 130), (420, 130), (450, 190), (280, 190)], fill=(255, 255, 255, 150))
        draw.polygon([(435, 130), (490, 130), (580, 190), (465, 190)], fill=(255, 255, 255, 150))
        # Wheels
        draw.ellipse([180, 250, 280, 350], fill=(40, 40, 40, 255))
        draw.ellipse([520, 250, 620, 350], fill=(40, 40, 40, 255))
        draw.ellipse([210, 280, 250, 320], fill=(200, 200, 200, 255))
        draw.ellipse([550, 280, 590, 320], fill=(200, 200, 200, 255))
        return img

    car_img = generate_car_placeholder()
    car_io = io.BytesIO()
    car_img.save(car_io, format='PNG')
    car_io.seek(0)
    slide.shapes.add_picture(car_io, Inches(3.1), Inches(1.9), width=Inches(7.1))

    # Navigation Arrows
    add_text(Inches(0.5), Inches(3.0), Inches(0.5), Inches(1), "❮", 36, LIGHT_TEXT)
    add_text(Inches(12.3), Inches(3.0), Inches(0.5), Inches(1), "❯", 36, LIGHT_TEXT)

    # --- Charts Section (Matplotlib) ---
    plt.rcParams['font.family'] = 'sans-serif'
    
    # 1. Bar Chart (Left)
    fig1, ax1 = plt.subplots(figsize=(3.5, 2), dpi=150)
    fig1.patch.set_alpha(0)
    ax1.patch.set_alpha(0)
    categories = ['Replacement', 'Servicing', 'Other']
    values = [65, 85, 30]
    ax1.barh(categories, values, color=ACCENT_HEX, height=0.5)
    for spine in ax1.spines.values():
        spine.set_visible(False)
    ax1.set_xticks([])
    ax1.tick_params(axis='y', length=0, colors='#555555')
    ax1.invert_yaxis()
    
    buf1 = io.BytesIO()
    fig1.savefig(buf1, format='png', transparent=True, bbox_inches='tight', pad_inches=0)
    buf1.seek(0)
    slide.shapes.add_picture(buf1, Inches(0.8), Inches(5.2), height=Inches(1.8))
    plt.close(fig1)

    # 2. Donut Chart (Center)
    fig2, ax2 = plt.subplots(figsize=(2.5, 2.5), dpi=150)
    fig2.patch.set_alpha(0)
    ax2.patch.set_alpha(0)
    sizes = [45, 30, 25]
    colors = [ACCENT_HEX, LIGHT_GRAY_HEX, '#555555']
    wedges, _ = ax2.pie(sizes, colors=colors, startangle=90, wedgeprops=dict(width=0.3, edgecolor='none'))
    
    buf2 = io.BytesIO()
    fig2.savefig(buf2, format='png', transparent=True, bbox_inches='tight', pad_inches=0)
    buf2.seek(0)
    slide.shapes.add_picture(buf2, Inches(5.2), Inches(5.0), height=Inches(2.0))
    # Legend for Donut
    add_text(Inches(7.2), Inches(5.3), Inches(1.5), Inches(0.3), "● Third Party", 10, LIGHT_TEXT)
    add_text(Inches(7.2), Inches(5.6), Inches(1.5), Inches(0.3), "● Direct Sale", 10, ACCENT_RGB)
    add_text(Inches(7.2), Inches(5.9), Inches(1.5), Inches(0.3), "● Pre Book", 10, DARK_TEXT)
    plt.close(fig2)

    # 3. Column Chart (Right)
    fig3, ax3 = plt.subplots(figsize=(4, 2), dpi=150)
    fig3.patch.set_alpha(0)
    ax3.patch.set_alpha(0)
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    trend = [20, 25, 30, 22, 40, 45, 35, 50, 55, 48, 80, 65]
    colors_col = [ACCENT_HEX if i != 10 else '#333333' for i in range(len(months))] # Highlight Nov
    ax3.bar(months, trend, color=colors_col, width=0.6)
    for spine in ax3.spines.values():
        spine.set_visible(False)
    ax3.set_yticks([])
    ax3.tick_params(axis='x', length=0, colors='#888888', labelsize=8)
    
    buf3 = io.BytesIO()
    fig3.savefig(buf3, format='png', transparent=True, bbox_inches='tight', pad_inches=0)
    buf3.seek(0)
    slide.shapes.add_picture(buf3, Inches(9.0), Inches(5.2), width=Inches(3.8))
    plt.close(fig3)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
