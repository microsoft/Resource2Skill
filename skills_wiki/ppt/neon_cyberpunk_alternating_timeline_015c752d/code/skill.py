def create_slide(
    output_pptx_path: str,
    title_text: str = "TIMELINE SLIDE",
    milestones: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neon Cyberpunk Alternating Timeline effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
    
    # Default data if none provided
    if not milestones:
        milestones = [
            {"year": "2019", "text": "Put your text here. This is where your idea begins.", "hex": "FF2A2A"},  # Red
            {"year": "2020", "text": "Put your text here. This is where your idea begins.", "hex": "FFCC00"},  # Yellow
            {"year": "2021", "text": "Put your text here. This is where your idea begins.", "hex": "00FFFF"},  # Cyan
            {"year": "2022", "text": "Put your text here. This is where your idea begins.", "hex": "39FF14"},  # Green
            {"year": "2023", "text": "Put your text here. This is where your idea begins.", "hex": "8A2BE2"},  # Purple
        ]

    # --- XML Helper for Glow Effect ---
    def apply_shape_glow(shape, hex_color: str, radius_pt: int = 15, alpha_pct: int = 50):
        """Injects DrawingML XML to apply a glow effect to a shape."""
        spPr = shape._element.spPr
        alpha_val = int(alpha_pct * 1000) # 50000 = 50%
        glow_rad = int(radius_pt * 12700) # Convert points to EMUs
        
        # Build the effectLst XML
        glow_xml = f'''
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:glow rad="{glow_rad}">
                <a:srgbClr val="{hex_color}">
                    <a:alpha val="{alpha_val}"/>
                </a:srgbClr>
            </a:glow>
        </a:effectLst>
        '''
        effect_element = parse_xml(glow_xml)
        
        # Remove existing effectLst to avoid XML schema errors
        existing = spPr.find(qn('a:effectLst'))
        if existing is not None:
            spPr.remove(existing)
        spPr.append(effect_element)

    # --- Helper to convert Hex string to RGB tuple ---
    def hex_to_rgb(hex_str):
        hex_str = hex_str.lstrip('#')
        return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

    # --- Initialize Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # 1. Background Fill (Dark Charcoal)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 30, 32)

    # 2. Main Title
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.5), Inches(13.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial" # Fallback for Lemon Milk
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    # Apply subtle white glow to title
    apply_shape_glow(title_box, "FFFFFF", radius_pt=10, alpha_pct=30)

    # 3. Main Horizontal Axis Line
    center_y = Inches(3.75)
    margin_x = Inches(1.5)
    line_width = prs.slide_width - (margin_x * 2)
    
    main_line = slide.shapes.add_connector(MSO_SHAPE.LINE_INVERSE, margin_x, center_y, margin_x + line_width, center_y)
    main_line.line.color.rgb = RGBColor(255, 255, 255)
    main_line.line.width = Pt(1.5)

    # 4. Generate Timeline Nodes
    num_nodes = len(milestones)
    spacing = line_width / (num_nodes - 1) if num_nodes > 1 else 0
    
    branch_length = Inches(1.6)
    ring_radius = Inches(0.2)
    core_radius = Inches(0.08)
    base_dot_radius = Inches(0.06)

    for i, data in enumerate(milestones):
        x_center = margin_x + (i * spacing)
        
        # Alternating direction: Even index goes Down (1), Odd index goes Up (-1)
        # Reversing standard logic so first node (2019) goes down as seen in many alternating designs
        direction = 1 if i % 2 == 0 else -1 
        y_end = center_y + (direction * branch_length)

        rgb = hex_to_rgb(data["hex"])
        node_color = RGBColor(*rgb)

        # A. Base Dot (on main axis)
        base_dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x_center - base_dot_radius, center_y - base_dot_radius, base_dot_radius*2, base_dot_radius*2)
        base_dot.fill.solid()
        base_dot.fill.fore_color.rgb = node_color
        base_dot.line.fill.background() # No line
        apply_shape_glow(base_dot, data["hex"], radius_pt=8)

        # B. Vertical Branch Line
        branch = slide.shapes.add_connector(MSO_SHAPE.LINE_INVERSE, x_center, center_y, x_center, y_end)
        branch.line.color.rgb = node_color
        branch.line.width = Pt(1.5)

        # C. Hollow Ring (Outer Circle)
        ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, x_center - ring_radius, y_end - ring_radius, ring_radius*2, ring_radius*2)
        ring.fill.background() # Transparent
        ring.line.color.rgb = node_color
        ring.line.width = Pt(2.5)

        # D. Glowing Core (Inner Dot)
        core = slide.shapes.add_shape(MSO_SHAPE.OVAL, x_center - core_radius, y_end - core_radius, core_radius*2, core_radius*2)
        core.fill.solid()
        core.fill.fore_color.rgb = node_color
        core.line.fill.background()
        apply_shape_glow(core, data["hex"], radius_pt=18, alpha_pct=50)

        # E. Year Text (Matching Color)
        text_w = Inches(2)
        text_h = Inches(0.5)
        # Position text beyond the ring
        year_y = y_end + (direction * Inches(0.4)) - (text_h/2) if direction == 1 else y_end + (direction * Inches(0.4)) - (text_h/2)
        
        year_box = slide.shapes.add_textbox(x_center - (text_w/2), year_y, text_w, text_h)
        p_year = year_box.text_frame.paragraphs[0]
        p_year.text = data["year"]
        p_year.alignment = PP_ALIGN.CENTER
        p_year.font.name = "Arial"
        p_year.font.size = Pt(24)
        p_year.font.bold = True
        p_year.font.color.rgb = node_color

        # F. Description Text (White)
        desc_h = Inches(0.8)
        desc_y = year_y + (direction * Inches(0.4)) if direction == 1 else year_y + (direction * Inches(0.6))
        
        desc_box = slide.shapes.add_textbox(x_center - (text_w/2), desc_y, text_w, desc_h)
        desc_box.text_frame.word_wrap = True
        p_desc = desc_box.text_frame.paragraphs[0]
        p_desc.text = data["text"]
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = RGBColor(200, 200, 200)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
