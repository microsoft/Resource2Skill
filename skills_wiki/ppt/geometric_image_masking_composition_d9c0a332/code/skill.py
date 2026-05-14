import os
import urllib.request
from io import BytesIO
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def fetch_and_crop_image(url: str, target_width: int, target_height: int) -> Image.Image:
    """Fetches an image from a URL and center-crops it to the exact target dimensions."""
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception as e:
        print(f"Failed to download image: {e}. Generating fallback gradient.")
        # Fallback: Create a gradient image if network fails
        img = Image.new("RGBA", (target_width, target_height))
        draw = ImageDraw.Draw(img)
        for y in range(target_height):
            r = int(20 + (100 * y / target_height))
            g = int(50 + (150 * y / target_height))
            b = int(150 + (105 * y / target_height))
            draw.line([(0, y), (target_width, y)], fill=(r, g, b, 255))
        return img

    # Center Crop Logic
    img_ratio = img.width / img.height
    target_ratio = target_width / target_height
    
    if target_ratio > img_ratio:
        # Target is wider, crop top and bottom
        new_height = int(img.width / target_ratio)
        offset = (img.height - new_height) // 2
        crop_box = (0, offset, img.width, offset + new_height)
    else:
        # Target is taller, crop left and right
        new_width = int(img.height * target_ratio)
        offset = (img.width - new_width) // 2
        crop_box = (offset, 0, offset + new_width, img.height)
        
    img = img.crop(crop_box)
    img = img.resize((target_width, target_height), Image.Resampling.LANCZOS)
    return img

def apply_shape_mask(img: Image.Image, shape_type: str = "parallelogram") -> Image.Image:
    """Applies an alpha mask to the image based on the chosen shape."""
    mask = Image.new("L", img.size, 0)
    draw = ImageDraw.Draw(mask)
    w, h = img.size
    
    if shape_type == "parallelogram":
        # Create a slanted parallelogram
        # Top-left is indented, bottom-right is indented
        slant_offset = int(w * 0.25)
        polygon_points = [
            (slant_offset, 0), 
            (w, 0), 
            (w - slant_offset, h), 
            (0, h)
        ]
        draw.polygon(polygon_points, fill=255)
    elif shape_type == "circle":
        draw.ellipse((0, 0, w, h), fill=255)
    else:
        draw.rectangle((0, 0, w, h), fill=255)
        
    img.putalpha(mask)
    return img

def create_slide(
    output_pptx_path: str,
    title_text: str = "Marine Biology",
    subtitle_text: str = "An Introduction to Organisms In The Sea",
    theme_keyword: str = "ocean,reef",
    shape_style: str = "parallelogram", # "parallelogram" or "circle"
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Geometric Image Masking effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Slide Background ===
    # Dark slate/navy background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(15, 32, 45)

    # === Layer 2: Masked Image Generation ===
    # Convert inches to pixels for PIL (assuming 300 DPI for crispness)
    dpi = 300
    
    if shape_style == "parallelogram":
        img_w_inches, img_h_inches = 5.0, 7.5
        left_pos, top_pos = Inches(1.0), Inches(0)
    else: # Circle layout
        img_w_inches, img_h_inches = 5.5, 5.5
        left_pos, top_pos = Inches(7.0), Inches(1.0)
        
    target_px_w = int(img_w_inches * dpi)
    target_px_h = int(img_h_inches * dpi)

    # Fetch and process image
    img_url = f"https://source.unsplash.com/featured/{target_px_w}x{target_px_h}/?{theme_keyword}"
    img = fetch_and_crop_image(img_url, target_px_w, target_px_h)
    masked_img = apply_shape_mask(img, shape_type=shape_style)
    
    # Save temporarily
    temp_img_path = "temp_masked_shape.png"
    masked_img.save(temp_img_path, format="PNG")

    # Insert into PPTX
    slide.shapes.add_picture(
        temp_img_path, 
        left_pos, top_pos, 
        width=Inches(img_w_inches), 
        height=Inches(img_h_inches)
    )

    # === Layer 3: Typography ===
    if shape_style == "parallelogram":
        tx_left = Inches(6.5)
        tx_top = Inches(3.0)
        tx_width = Inches(6.0)
        tx_height = Inches(2.0)
    else:
        tx_left = Inches(1.0)
        tx_top = Inches(2.5)
        tx_width = Inches(5.5)
        tx_height = Inches(3.0)

    textbox = slide.shapes.add_textbox(tx_left, tx_top, tx_width, tx_height)
    text_frame = textbox.text_frame
    text_frame.word_wrap = True

    # Title
    p_title = text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.name = "Arial"
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Body
    p_sub = text_frame.add_paragraph()
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(24)
    p_sub.font.name = "Arial"
    p_sub.font.color.rgb = RGBColor(200, 210, 220)
    p_sub.space_before = Pt(10)
    
    # Clean up temp file
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
