def create_slide(
    output_pptx_path: str,
    bg_text_top: str = "KINETIC",
    bg_text_bottom: str = "TYPOGRAPHY",
    frame_1_text: str = "IDEAS",
    frame_2_text: str = "NEED",
    frame_3_text: str = "MOTION",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Cinematic Film Strip Typography effect.
    """
    import os
    import tempfile
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Solid Black Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)

    # === Layer 2: Massive Background Texture Text ===
    def add_bg_text(text, top_inches):
        tb = slide.shapes.add_textbox(0, Inches(top_inches), Inches(13.333), Inches(2.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial Black"
        p.font.size = Pt(140)
        p.font.color.rgb = RGBColor(30, 30, 30) # Very dark grey

    add_bg_text(bg_text_top, 0.2)
    add_bg_text(bg_text_bottom, 5.0)

    # === Layer 3: Generative Film Strip Mask (PIL) ===
    # We generate a 1920x1080 transparent PNG with white film borders and black sprocket holes
    temp_dir = tempfile.gettempdir()
    film_strip_path = os.path.join(temp_dir, "film_strip_mask.png")

    img = Image.new('RGBA', (1920, 1080), (0, 0, 0, 0)) # Transparent background
    draw = ImageDraw.Draw(img)
    
    strip_y1 = 350
    strip_y2 = 730
    border_h = 60
    
    # Draw horizontal white borders
    draw.rectangle([0, strip_y1, 1920, strip_y1 + border_h], fill=(255, 255, 255, 255))
    draw.rectangle([0, strip_y2 - border_h, 1920, strip_y2], fill=(255, 255, 255, 255))
    
    # Draw vertical white frame separators
    sep_w = 30
    frame_w = (1920 - 4 * sep_w) // 3 # Divide remaining space into 3 frames
    
    for i in range(4):
        x = i * (frame_w + sep_w)
        draw.rectangle([x, strip_y1, x + sep_w, strip_y2], fill=(255, 255, 255, 255))
        
    # Draw black sprocket holes
    hole_w, hole_h = 24, 30
    hole_gap = 16
    hole_y_top = strip_y1 + 15
    hole_y_bottom = strip_y2 - border_h + 15
    
    for x in range(10, 1920, hole_w + hole_gap):
        draw.rectangle([x, hole_y_top, x + hole_w, hole_y_top + hole_h], fill=(0, 0, 0, 255))
        draw.rectangle([x, hole_y_bottom, x + hole_w, hole_y_bottom + hole_h], fill=(0, 0, 0, 255))
        
    img.save(film_strip_path)

    # Insert the full-screen PIL mask into PPTX
    slide.shapes.add_picture(film_strip_path, 0, 0, Inches(13.333), Inches(7.5))

    # === Layer 4: Foreground Framed Text ===
    def add_frame_text(text, center_x_px):
        # Convert pixel coordinates from the 1920x1080 PIL image to PPTX inches
        center_x_inches = (center_x_px / 1920) * 13.333
        center_y_inches = (540 / 1080) * 7.5 # Vertically centered
        
        width = Inches(4.0)
        height = Inches(1.5)
        left = center_x_inches - (4.0 / 2)
        top = center_y_inches - (1.5 / 2)
        
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        # Use an expressive font (fallback to standard if missing on user's OS, but requested here)
        p.font.name = "Segoe Script" 
        p.font.size = Pt(54)
        p.font.color.rgb = RGBColor(255, 255, 255) # White text inside the dark frame
        tf.margin_top = 0
        tf.margin_bottom = 0

    # Calculate pixel centers for the 3 frames based on our PIL math
    frame_1_center = sep_w + (frame_w / 2)
    frame_2_center = sep_w + frame_w + sep_w + (frame_w / 2)
    frame_3_center = sep_w + frame_w + sep_w + frame_w + sep_w + (frame_w / 2)

    add_frame_text(frame_1_text, frame_1_center)
    add_frame_text(frame_2_text, frame_2_center)
    add_frame_text(frame_3_text, frame_3_center)

    prs.save(output_pptx_path)
    
    # Cleanup temp file
    if os.path.exists(film_strip_path):
        os.remove(film_strip_path)
        
    return output_pptx_path
