def create_slide(
    output_pptx_path: str,
    title_text: str = "WhatsApp's expansive reach and popularity make it an appealing option for connecting with customers",
    body_text: str = "", # Unused in this specific layout
    bg_palette: str = "light", 
    accent_color: tuple = (214, 40, 40),  # Red highlight as seen in the final video frame
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Intentional Data Highlight' consulting chart pattern.
    Features two side-by-side bar charts with stripped-down aesthetics and specific point highlighting.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.shapes import MSO_CONNECTOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Highlight Color (from args) and Neutral Color
    hl_color = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    neutral_color = RGBColor(210, 210, 210)
    text_main_color = RGBColor(40, 40, 40)
    text_sub_color = RGBColor(120, 120, 120)

    # === Layer 1: Title and Structural Separator ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = text_main_color

    # Separator Line
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, 
        Inches(0.5), Inches(1.3), Inches(12.833), Inches(1.3)
    )
    line.line.color.rgb = RGBColor(200, 200, 200)
    line.line.width = Pt(1)

    # === Layer 2: Chart Generation Helper ===
    def add_highlighted_bar_chart(x, y, w, h, data, target_category, main_title, sub_title):
        # 1. Add Subtitles above the chart
        tx_box = slide.shapes.add_textbox(x, y, w, Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = main_title
        p1.font.bold = True
        p1.font.size = Pt(14)
        p1.font.color.rgb = text_main_color
        
        p2 = tf.add_paragraph()
        p2.text = sub_title
        p2.font.size = Pt(11)
        p2.font.color.rgb = text_sub_color

        # 2. Prepare Chart Data (Sort ascending so largest renders at the top of the bar chart)
        sorted_data = sorted(data, key=lambda i: i[1])
        chart_data = CategoryChartData()
        chart_data.categories = [item[0] for item in sorted_data]
        chart_data.add_series('Data', [item[1] for item in sorted_data])

        # 3. Add Chart Shape
        chart_y = y + Inches(0.8)
        chart_h = h - Inches(0.8)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, x, chart_y, w, chart_h, chart_data
        ).chart

        # 4. Strip away Chart Junk
        chart.has_legend = False
        
        val_axis = chart.value_axis
        val_axis.visible = False # Hides x-axis numbers and line
        val_axis.has_major_gridlines = False
        
        cat_axis = chart.category_axis
        cat_axis.has_major_gridlines = False
        cat_axis.tick_labels.font.size = Pt(11)
        cat_axis.tick_labels.font.color.rgb = text_main_color

        # 5. Add Clean Data Labels
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(11)
        plot.data_labels.font.color.rgb = text_main_color

        # 6. Apply Intentional Highlighting (Color logic)
        series = chart.series[0]
        for idx, point in enumerate(series.points):
            fill = point.format.fill
            fill.solid()
            cat_name = sorted_data[idx][0]
            if cat_name == target_category:
                fill.fore_color.rgb = hl_color
            else:
                fill.fore_color.rgb = neutral_color

    # === Layer 3: Execute Chart Rendering ===
    
    # Dataset 1: Monthly Active Users
    data_1 = [
        ("WhatsApp", 2000), ("WeChat", 1263), ("Facebook Messenger", 988),
        ("QQ", 574), ("Snapchat", 557), ("Telegram", 550)
    ]
    
    add_highlighted_bar_chart(
        x=Inches(0.5), y=Inches(1.6), w=Inches(5.8), h=Inches(5.2),
        data=data_1,
        target_category="WhatsApp",
        main_title="WhatsApp leads the global messenger market",
        sub_title="Monthly active users (in millions), January 2022"
    )

    # Dataset 2: Global Downloads
    data_2 = [
        ("TikTok", 656), ("Instagram", 545), ("Facebook", 416),
        ("WhatsApp", 395), ("Telegram", 329), ("Snapchat", 327), ("Zoom", 300)
    ]

    add_highlighted_bar_chart(
        x=Inches(6.8), y=Inches(1.6), w=Inches(5.8), h=Inches(5.2),
        data=data_2,
        target_category="WhatsApp",
        main_title="Last year WhatsApp was downloaded 395 million times",
        sub_title="Total global downloads (in millions), 2021"
    )

    # === Layer 4: Footnote ===
    foot_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(10.0), Inches(0.5))
    p_foot = foot_box.text_frame.paragraphs[0]
    p_foot.text = "Sources: Company data, DataReportal, BusinessApps.com"
    p_foot.font.size = Pt(9)
    p_foot.font.color.rgb = RGBColor(150, 150, 150)

    prs.save(output_pptx_path)
    return output_pptx_path
