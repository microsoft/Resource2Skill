def create_slide(
    output_pptx_path: str,
    title_text: str = "Company Milestone",
    body_text: str = "Replace this sample text with your own explanation.",
    bg_palette: str = "technology",  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Continuous Film Strip Timeline" visual effect.
    Generates a 2-slide continuous sequence.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw
    from lxml import etree
    import urllib.request
    import io

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Generate the Film Strip Backdrop using PIL
    def generate_filmstrip_image():
        # High resolution for crisp rendering
        width, height = 4000, 1000
        img = Image.new('RGBA', (width, height), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        
        # Main strip body
        strip_color = (43, 43, 43, 255)
        draw.rectangle([0, 0, width, height], fill=strip_color)

        # Perforations (Sprocket holes)
        perf_w, perf_h = 40, 60
        gap = 30
        y_top = 40
        y_bottom = height - 40 - perf_h

        shadow_color = (20, 20, 20, 255)
        white_color = (255, 255, 255, 255)

        # Draw repeating perforations
        for x in range(15, width, perf_w + gap):
            # Top row - shadow then offset white for "cutout" depth illusion
            draw.rounded_rectangle([x, y_top, x+perf_w, y_top+perf_h], radius=10, fill=shadow_color)
            draw.rounded_rectangle([x+4, y_top+4, x+perf_w, y_top+perf_h], radius=10, fill=white_color)

            # Bottom row
            draw.rounded_rectangle([x, y_bottom, x+perf_w, y_bottom+perf_h], radius=10, fill=shadow_color)
            draw.rounded_rectangle([x+4, y_bottom+4, x+perf_w, y_bottom+perf_h], radius=10, fill=white_color)

        stream = io.BytesIO()
        img.save(stream, format='PNG')
        stream.seek(0)
        return stream

    filmstrip_stream = generate_filmstrip_image()

    # Helper: Fetch mock photography with solid color fallback
    def get_photo(index: int, fallback_rgb: tuple):
        try:
            url = f"https://picsum.photos/seed/{index + 10}/400/300"
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=5) as res:
                return io.BytesIO(res.read())
        except Exception:
            img = Image.new('RGB', (400, 300), fallback_rgb)
            stream = io.BytesIO()
            img.save(stream, format='JPEG')
            stream.seek(0)
            return stream

    # Layout Parameters
    num_frames_per_slide = 4
    frame_width = 2.4
    frame_height = 2.2
    strip_y = 2.0
    strip_height = 3.5
    gap = (13.333 - (num_frames_per_slide * frame_width)) / (num_frames_per_slide + 1)
    
    # 2. Build Slides (Create 2 to demonstrate the continuous transition)
    for slide_idx in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

        # Add Film Strip Backdrop
        slide.shapes.add_picture(
            filmstrip_stream, 
            left=Inches(0), 
            top=Inches(strip_y), 
            width=Inches(13.333), 
            height=Inches(strip_height)
        )

        # Add Frames (Alternating Photos and Empty Placeholders)
        for i in range(num_frames_per_slide):
            x = gap + i * (frame_width + gap)
            y = strip_y + (strip_height - frame_height) / 2
            global_idx = slide_idx * num_frames_per_slide + i

            if global_idx % 2 == 0:
                # Active Frame (Photo)
                photo_stream = get_photo(global_idx, (70, 130, 180))
                slide.shapes.add_picture(photo_stream, Inches(x), Inches(y), width=Inches(frame_width), height=Inches(frame_height))

                # Text Explanation Below Photo
                txBox = slide.shapes.add_textbox(Inches(x - 0.2), Inches(strip_y + strip_height + 0.2), Inches(frame_width + 0.4), Inches(1.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                
                # Title
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                run = p.add_run()
                run.text = f"{title_text} {global_idx + 1}"
                run.font.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(30, 30, 30)

                # Body
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                run2 = p2.add_run()
                run2.text = body_text
                run2.font.size = Pt(11)
                run2.font.color.rgb = RGBColor(100, 100, 100)
            else:
                # Neutral Placeholder Frame (Empty)
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(frame_width), Inches(frame_height))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(100, 100, 100)
                shape.line.color.rgb = RGBColor(80, 80, 80)
                shape.line.width = Pt(1)

        # 3. Inject OpenXML for "Push" Transition to create continuous scrolling effect
        # <p:transition spd="slow"><p:push dir="r"/></p:transition>
        xml_transition = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:push dir="r"/></p:transition>'
        p_transition = etree.fromstring(xml_transition)
        
        # Safely insert before p:timing or at the end of the slide XML
        timing_tag = slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
        if timing_tag is not None:
            timing_tag.addprevious(p_transition)
        else:
            slide._element.append(p_transition)

    prs.save(output_pptx_path)
    return output_pptx_path
