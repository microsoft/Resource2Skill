def create_slide(
    output_pptx_path: str,
    list_number: str = "5",
    title_part1: str = "Canva gives you access",
    title_part2: str = "to a lot more templates",
    bg_color: tuple = (23, 43, 61),      # Deep Navy
    accent_color: tuple = (253, 209, 39), # Vibrant Yellow
    text_color: tuple = (255, 255, 255),  # White
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "High-Contrast Oversized Number Divider" visual effect.
    
    Args:
        output_pptx_path: Filepath to save the .pptx file.
        list_number: The large digit/number to display in the side panel.
        title_part1: The first part of the statement (colored in accent color).
        title_part2: The second part of the statement (colored in white).
        bg_color: RGB tuple for the main dark background.
        accent_color: RGB tuple for the left panel and highlighted text.
        text_color: RGB tuple for the standard text.
        
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    
    # Initialize presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Extract color tuples to RGBColor objects
    c_bg = RGBColor(*bg_color)
    c_accent = RGBColor(*accent_color)
    c_text = RGBColor(*text_color)

    # === Layer 1: Background ===
    # Set the main slide background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = c_bg

    # === Layer 2: Visual Structure (The Split Panel) ===
    # Add the vertical accent panel on the left
    panel_width = Inches(3.5)
    left_panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        0, 0, panel_width, prs.slide_height
    )
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = c_accent
    left_panel.line.fill.background() # Remove border

    # === Layer 3: Typography ===
    
    # 1. The Oversized Anchor Number
    num_box = slide.shapes.add_textbox(
        Inches(0), Inches(2.25), panel_width, Inches(3.0)
    )
    num_frame = num_box.text_frame
    num_frame.clear() # clear default paragraphs
    num_p = num_frame.paragraphs[0]
    num_p.alignment = PP_ALIGN.CENTER
    num_run = num_p.add_run()
    num_run.text = str(list_number)
    
    # Styling the oversized number
    font = num_run.font
    font.name = "Arial Black" # Use a widely available heavy font
    font.size = Pt(220)
    font.color.rgb = c_bg # Number inherits the dark background color for contrast

    # 2. The Main Statement (Two-Tone Text)
    # Positioned with generous left padding away from the yellow panel
    text_box = slide.shapes.add_textbox(
        Inches(4.2), Inches(2.5), Inches(8.5), Inches(3.0)
    )
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.1

    # Run 1: Highlighted Text (Yellow)
    run1 = p.add_run()
    run1.text = title_part1 + "\n"
    run1.font.name = "Arial Black"
    run1.font.size = Pt(54)
    run1.font.color.rgb = c_accent
    
    # Run 2: Base Text (White)
    run2 = p.add_run()
    run2.text = title_part2
    run2.font.name = "Arial Black"
    run2.font.size = Pt(54)
    run2.font.color.rgb = c_text

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
