import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def create_slide(
    output_pptx_path: str,
    title_text: str = "AGENDA",
    agenda_items: list = None,
    bg_color: tuple = (32, 178, 170),      # Teal green
    node_fill_color: tuple = (64, 64, 64), # Dark gray
    accent_color: tuple = (255, 255, 255), # White
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Sequential Vertical Node Timeline.
    """
    if agenda_items is None:
        agenda_items = ["Introduction", "Services", "Clients", "Portfolio", "Contact Us"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Solid Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Main Title ===
    # Vertical accent line next to title
    title_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.0), Inches(0.05), Inches(1.0)
    )
    title_line.fill.solid()
    title_line.fill.fore_color.rgb = RGBColor(*accent_color)
    title_line.line.fill.background() # No outline

    # Title Text
    txBox = slide.shapes.add_textbox(Inches(1.7), Inches(2.2), Inches(4), Inches(1))
    tf = txBox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = 'Roboto' # Will fallback to Arial if missing
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*accent_color)
    p.alignment = PP_ALIGN.LEFT

    # === Layer 3: Timeline Construction ===
    num_items = len(agenda_items)
    axis_x = Inches(6.0) # Center line X position
    start_y = Inches(1.5)
    end_y = Inches(6.0)
    
    # Calculate spacing
    if num_items > 1:
        gap_y = (end_y - start_y) / (num_items - 1)
    else:
        gap_y = 0
        start_y = Inches(3.75) # Center vertically if only 1 item

    node_radius = Inches(0.35)
    line_weight = Pt(2.25)

    # Draw connectors FIRST (so they sit behind the circles)
    for i in range(num_items - 1):
        y1 = start_y + (i * gap_y)
        y2 = start_y + ((i + 1) * gap_y)
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, axis_x, y1, axis_x, y2
        )
        connector.line.color.rgb = RGBColor(*accent_color)
        connector.line.width = line_weight

    # Draw Nodes and Labels
    for i, item_text in enumerate(agenda_items):
        cy = start_y + (i * gap_y)
        
        # 1. Draw Circle Node
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            axis_x - node_radius, 
            cy - node_radius, 
            node_radius * 2, 
            node_radius * 2
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(*node_fill_color)
        circle.line.color.rgb = RGBColor(*accent_color)
        circle.line.width = line_weight
        
        # 2. Add Number inside circle
        tf_circle = circle.text_frame
        tf_circle.clear()
        tf_circle.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_num = tf_circle.paragraphs[0]
        p_num.text = f"{i+1:02d}" # "01", "02", etc.
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.name = 'Roboto'
        p_num.font.size = Pt(16)
        p_num.font.bold = True
        p_num.font.color.rgb = RGBColor(*accent_color)
        
        # 3. Add Label Text
        label_x = axis_x + node_radius + Inches(0.2)
        label_y = cy - Inches(0.25)
        label_box = slide.shapes.add_textbox(label_x, label_y, Inches(5), Inches(0.5))
        tf_label = label_box.text_frame
        tf_label.clear()
        tf_label.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_label = tf_label.paragraphs[0]
        p_label.text = item_text
        p_label.font.name = 'Roboto'
        p_label.font.size = Pt(24)
        p_label.font.color.rgb = RGBColor(*accent_color)
        p_label.alignment = PP_ALIGN.LEFT

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_slide("agenda_timeline.pptx")
