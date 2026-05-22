import os
import requests
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    program_title: str = "Virtual CFO",
    company_name: str = "By Decisions Plus Strategic",
    tagline: str = "Connect with us and make a difference",
    contact_email: str = "hello@decisionsplusstrategic.com",
    impact_story: str = "Each time we collaborate in a Virtual CFO Engagement, 500 women in Malawi receive business start-up and bookkeeping training to give them the skills they need to grow their own micro-business and lift their family and community out of poverty. We do this through our partnership with B1G1 and MicroLoan Australia.",
    image_url: str = "https://images.unsplash.com/photo-1605786184133-2558509a55a8",
    **kwargs,
) -> str:
    """
    Creates a PPTX file with a single slide reproducing the "Purpose-Driven CTA Slide" visual effect.

    Args:
        output_pptx_path (str): The path to save the generated .pptx file.
        program_title (str): The main title of the service or program.
        company_name (str): The name of the company.
        tagline (str): The aspirational tagline.
        contact_email (str): The primary contact email address.
        impact_story (str): The paragraph describing the social impact.
        image_url (str): URL for the representative photo.

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(29, 78, 185)

    # === Layer 2: Visuals (Image) ===
    try:
        response = requests.get(image_url)
        response.raise_for_status()
        image_stream = BytesIO(response.content)
        img_left = Inches(0.75)
        img_top = Inches(2.25)
        img_height = Inches(3.0)
        slide.shapes.add_picture(image_stream, img_left, img_top, height=img_height)
    except requests.exceptions.RequestException as e:
        print(f"Warning: Could not download image from {image_url}. Error: {e}. Skipping image.")
        # Fallback: Add a placeholder shape if image download fails
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(2.25), Inches(3.0), Inches(3.0)
        )

    # === Layer 3: Text & Content ===
    # Program Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = program_title
    p.font.name = "Calibri Light"
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(204, 218, 249)

    # Company Name
    p2 = txBox.text_frame.add_paragraph()
    p2.text = company_name
    p2.font.name = "Calibri"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(255, 255, 255)
    txBox.text_frame.margin_bottom = Inches(0)


    # Tagline
    txBox_tag = slide.shapes.add_textbox(Inches(7.8), Inches(0.5), Inches(5), Inches(1))
    p_tag = txBox_tag.text_frame.paragraphs[0]
    p_tag.text = tagline
    p_tag.font.name = "Calibri"
    p_tag.font.size = Pt(28)
    p_tag.font.color.rgb = RGBColor(191, 191, 191)
    p_tag.alignment = PP_ALIGN.RIGHT
    txBox_tag.text_frame.margin_bottom = Inches(0)

    # @ Symbol
    txBox_at = slide.shapes.add_textbox(Inches(4.2), Inches(3.0), Inches(1), Inches(1.5))
    p_at = txBox_at.text_frame.paragraphs[0]
    p_at.text = "@"
    p_at.font.name = "Calibri Light"
    p_at.font.size = Pt(72)
    p_at.font.color.rgb = RGBColor(255, 255, 255)
    p_at.alignment = PP_ALIGN.CENTER

    # Email Address
    txBox_email = slide.shapes.add_textbox(Inches(5.2), Inches(3.3), Inches(7.5), Inches(1))
    p_email = txBox_email.text_frame.paragraphs[0]
    p_email.text = contact_email
    p_email.font.name = "Calibri"
    p_email.font.bold = True
    p_email.font.size = Pt(32)
    p_email.font.color.rgb = RGBColor(255, 255, 255)
    p_email.alignment = PP_ALIGN.LEFT
    
    # Impact Story
    txBox_story = slide.shapes.add_textbox(Inches(4.2), Inches(4.7), Inches(8.5), Inches(2.5))
    p_story = txBox_story.text_frame.paragraphs[0]
    p_story.text = impact_story
    p_story.font.name = "Calibri"
    p_story.font.size = Pt(18)
    p_story.line_spacing = 1.5
    p_story.font.color.rgb = RGBColor(255, 255, 255)
    p_story.alignment = PP_ALIGN.LEFT

    # Save the presentation
    prs.save(output_pptx_path)
    return output_pptx_path
