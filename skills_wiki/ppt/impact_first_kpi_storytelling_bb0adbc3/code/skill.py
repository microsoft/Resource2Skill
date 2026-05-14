def create_presentation(
    output_pptx_path: str,
    hero_value: str = "$1,100,000",
    hero_subtitle: str = "weekly sales",
    growth_title: str = "Revenue is soaring 83%",
    start_value: str = "$600K",
    end_value: str = "$1.1M",
    funnel_kpis: list = None,
    logo_path: str = None
) -> str:
    """
    Creates a PPTX file reproducing the Impact-First KPI Storytelling style.

    Args:
        output_pptx_path (str): The path to save the generated .pptx file.
        hero_value (str): The main impact number for the title slide.
        hero_subtitle (str): The subtitle for the title slide.
        growth_title (str): The title for the growth chart slide.
        start_value (str): The starting value for the growth chart.
        end_value (str): The ending value for the growth chart.
        funnel_kpis (list): A list of dicts for the funnel, e.g., 
                            [{'label': 'social media engagements', 'value': '330,000'}, ...].
        logo_path (str): Optional path to a logo file to be placed on slides.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.shapes.freeform import FreeformBuilder

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Define Style ---
    BG_COLOR = RGBColor(13, 107, 180)
    ACCENT_COLOR = RGBColor(255, 204, 0)
    WHITE_COLOR = RGBColor(255, 255, 255)

    if funnel_kpis is None:
        funnel_kpis = [
            {'label': 'social media engagements', 'value': '330,000'},
            {'label': 'new website visitors', 'value': '44,000'},
            {'label': 'increase in inbound leads', 'value': '323%'}
        ]

    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_logo(slide):
        if logo_path:
            try:
                slide.shapes.add_picture(logo_path, Inches(0.25), Inches(0.25), height=Inches(0.3))
            except FileNotFoundError:
                print(f"Warning: Logo file not found at {logo_path}")

    # --- Slide 1: Hero Slide ---
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide1)
    add_logo(slide1)
    
    # Hero Value
    title_shape = slide1.shapes.add_textbox(Inches(0), Inches(2.5), prs.slide_width, Inches(2))
    title_p = title_shape.text_frame.paragraphs[0]
    title_p.text = hero_value
    title_p.font.name = 'Arial Black'
    title_p.font.size = Pt(128)
    title_p.font.color.rgb = ACCENT_COLOR
    title_p.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle_shape = slide1.shapes.add_textbox(Inches(0), Inches(4.25), prs.slide_width, Inches(1))
    subtitle_p = subtitle_shape.text_frame.paragraphs[0]
    subtitle_p.text = hero_subtitle
    subtitle_p.font.name = 'Arial'
    subtitle_p.font.size = Pt(36)
    subtitle_p.font.color.rgb = WHITE_COLOR
    subtitle_p.alignment = PP_ALIGN.CENTER

    # --- Slide 2: Growth Chart ---
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide2)
    add_logo(slide2)

    # Title
    growth_title_shape = slide2.shapes.add_textbox(Inches(0), Inches(0.5), prs.slide_width, Inches(1))
    gt_p = growth_title_shape.text_frame.paragraphs[0]
    gt_p.text = growth_title
    gt_p.font.name = 'Arial Bold'
    gt_p.font.size = Pt(48)
    gt_p.font.color.rgb = WHITE_COLOR
    gt_p.alignment = PP_ALIGN.CENTER
    # Highlight "Revenue"
    run1 = gt_p.runs[0]
    if "Revenue" in growth_title:
        parts = growth_title.split("Revenue")
        gt_p.text = parts[0]
        run1 = gt_p.add_run()
        run1.text = "Revenue"
        run1.font.color.rgb = ACCENT_COLOR
        run2 = gt_p.add_run()
        run2.text = parts[1]
        
    # Chart Drawing
    chart_y = Inches(5.5)
    start_x, end_x = Inches(2.5), Inches(10.833)
    
    # Baseline
    line = slide2.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, start_x - Inches(0.5), chart_y, end_x - start_x + Inches(1), 0)
    line.line.fill.solid()
    line.line.fill.fore_color.rgb = RGBColor(100, 150, 200)
    line.line.width = Pt(1.5)

    # Growth Curve
    path = FreeformBuilder.new(
        slide2.shapes, Emu(start_x), Emu(chart_y - Inches(1)), Emu(Pt(4)), ACCENT_COLOR
    ).add_cubic_bezier_segment(
        Emu(end_x), Emu(chart_y - Inches(3)), Emu(start_x + Inches(2)), Emu(chart_y - Inches(1)), Emu(end_x - Inches(2)), Emu(chart_y - Inches(3))
    ).close()
    path.line.end_cap = 3 # Round

    # Markers
    slide2.shapes.add_shape(MSO_SHAPE.OVAL, start_x - Pt(6), chart_y - Inches(1) - Pt(6), Pt(12), Pt(12)).fill.fore_color.rgb = ACCENT_COLOR
    slide2.shapes.add_shape(MSO_SHAPE.OVAL, end_x - Pt(6), chart_y - Inches(3) - Pt(6), Pt(12), Pt(12)).fill.fore_color.rgb = ACCENT_COLOR

    # Labels
    start_label = slide2.shapes.add_textbox(start_x - Inches(0.5), chart_y - Inches(0.8), Inches(1), Inches(0.5))
    start_label.text_frame.paragraphs[0].text = start_value
    start_label.text_frame.paragraphs[0].font.color.rgb = WHITE_COLOR
    start_label.text_frame.paragraphs[0].font.size = Pt(20)

    end_label = slide2.shapes.add_textbox(end_x - Inches(0.5), chart_y - Inches(3.8), Inches(1), Inches(0.5))
    end_label.text_frame.paragraphs[0].text = end_value
    end_label.text_frame.paragraphs[0].font.color.rgb = WHITE_COLOR
    end_label.text_frame.paragraphs[0].font.size = Pt(20)

    # --- Slide 3: Funnel ---
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide3)
    add_logo(slide3)
    
    funnel_center_x = prs.slide_width / 2
    
    # Funnel Segments
    segment_height = Inches(1.5)
    segment_gap = Inches(0.2)
    top_y = Inches(1.5)
    widths = [Inches(7), Inches(5), Inches(3)]
    
    for i in range(3):
        y_pos = top_y + i * (segment_height + segment_gap)
        width = widths[i]
        
        # Segment Shape (Trapezoid)
        path = FreeformBuilder.new(
            slide3.shapes, Emu(funnel_center_x - width/2), Emu(y_pos), Emu(Pt(0))
        ).add_line_segment(
            Emu(funnel_center_x + width/2), Emu(y_pos)
        ).add_line_segment(
            Emu(funnel_center_x + widths[i]/2 * 0.8), Emu(y_pos + segment_height)
        ).add_line_segment(
            Emu(funnel_center_x - widths[i]/2 * 0.8), Emu(y_pos + segment_height)
        ).close()
        path.fill.solid()
        path.fill.fore_color.rgb = RGBColor(18, 122, 200)

        # Value Text
        val_box = slide3.shapes.add_textbox(funnel_center_x - Inches(1.5), y_pos + Inches(0.1), Inches(3), Inches(1))
        val_p = val_box.text_frame.paragraphs[0]
        val_p.text = funnel_kpis[i]['value']
        val_p.font.name = 'Arial Black'
        val_p.font.size = Pt(40)
        val_p.font.color.rgb = ACCENT_COLOR
        val_p.alignment = PP_ALIGN.CENTER
        
        # Label Text
        lbl_box = slide3.shapes.add_textbox(funnel_center_x - Inches(4), y_pos - Inches(0.5), Inches(2.5), Inches(1.2))
        lbl_p = lbl_box.text_frame.paragraphs[0]
        lbl_p.text = funnel_kpis[i]['label']
        lbl_p.font.name = 'Arial'
        lbl_p.font.size = Pt(16)
        lbl_p.font.color.rgb = WHITE_COLOR
        lbl_p.alignment = PP_ALIGN.RIGHT
        
        # Connecting Line
        line_start_x = funnel_center_x - Inches(1.5)
        line = slide3.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, line_start_x, y_pos, 0, Inches(1.2))
        line.line.fill.solid()
        line.line.fill.fore_color.rgb = RGBColor(100, 150, 200)
        line.line.width = Pt(1)

        # Arrow
        if i < 2:
            arrow_y = y_pos + segment_height + segment_gap / 2
            arrow = slide3.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, funnel_center_x - Pt(15), arrow_y - Pt(10), Pt(30), Pt(20))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = WHITE_COLOR
            arrow.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_presentation("kpi_story_presentation.pptx")

