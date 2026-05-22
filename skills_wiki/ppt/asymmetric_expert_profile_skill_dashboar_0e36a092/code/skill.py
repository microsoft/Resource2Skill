def create_slide(
    output_pptx_path: str,
    expert_name: str = "Our expert",
    tagline: str = "Tagline of this slide",
    bio_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat.",
    timeline_data: list = [("Year 2018", "Little info on what\nshe actually did there"), ("Year 2020", "Little info on what\nshe actually did there"), ("Year 2023", "Little info on what\nshe actually did there")],
    skills_data: list = [("Photoshop", 0.85), ("Illustrator", 0.70), ("InDesign", 0.90)],
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Split-Panel Expert Profile visual effect.
    """
    import os
    import requests
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # --- Colors ---
    COLOR_DARK_PANEL = RGBColor(43, 54, 72)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_ACCENT = RGBColor(115, 38, 61) # Maroon
    COLOR_STARS = RGBColor(235, 186, 52) # Gold
    COLOR_TEXT_DARK = RGBColor(50, 50, 50)
    COLOR_TRACK_BG = RGBColor(220, 224, 229)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ==========================================
    # Layer 1: Split Background
    # ==========================================
    left_width = Inches(4.5)
    left_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, left_width, prs.slide_height)
    left_bg.fill.solid()
    left_bg.fill.fore_color.rgb = COLOR_DARK_PANEL
    left_bg.line.fill.background() # No border

    # ==========================================
    # Layer 2: Left Panel (Persona)
    # ==========================================
    # 1. Circular Avatar using PIL
    avatar_size_px = 400
    avatar_size_in = Inches(2.5)
    avatar_x = (left_width - avatar_size_in) / 2
    avatar_y = Inches(1.0)
    
    img_path = "temp_avatar.png"
    try:
        # Fetch a generic portrait
        url = "https://images.unsplash.com/photo-1494790108377-be9c29b29330?ixlib=rb-4.0.3&w=400&q=80"
        response = requests.get(url, timeout=5)
        img = Image.open(BytesIO(response.content)).convert("RGBA")
        
        # Crop to square
        min_dim = min(img.size)
        left = (img.width - min_dim)/2
        top = (img.height - min_dim)/2
        img = img.crop((left, top, left+min_dim, top+min_dim))
        img = img.resize((avatar_size_px, avatar_size_px), Image.Resampling.LANCZOS)
        
        # Create circular mask
        mask = Image.new('L', (avatar_size_px, avatar_size_px), 0)
        draw = ImageDraw.Draw(mask)
        draw.ellipse((0, 0, avatar_size_px, avatar_size_px), fill=255)
        img.putalpha(mask)
        img.save(img_path)
    except Exception as e:
        # Fallback: Create a solid color circle if network fails
        img = Image.new('RGBA', (avatar_size_px, avatar_size_px), (0,0,0,0))
        draw = ImageDraw.Draw(img)
        draw.ellipse((0, 0, avatar_size_px, avatar_size_px), fill=(100, 100, 100, 255))
        img.save(img_path)

    slide.shapes.add_picture(img_path, avatar_x, avatar_y, avatar_size_in, avatar_size_in)
    if os.path.exists(img_path): os.remove(img_path)

    # 2. Stars
    star_box = slide.shapes.add_textbox(0, Inches(3.7), left_width, Inches(0.6))
    star_frame = star_box.text_frame
    star_p = star_frame.paragraphs[0]
    star_p.alignment = PP_ALIGN.CENTER
    star_run = star_p.add_run()
    star_run.text = "★★★★★"
    star_run.font.size = Pt(32)
    star_run.font.color.rgb = COLOR_STARS

    # 3. Bio Text
    bio_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(3.5), Inches(2.5))
    bio_box.text_frame.word_wrap = True
    bio_p = bio_box.text_frame.paragraphs[0]
    bio_p.alignment = PP_ALIGN.CENTER
    bio_run = bio_p.add_run()
    bio_run.text = bio_text
    bio_run.font.size = Pt(12)
    bio_run.font.color.rgb = COLOR_WHITE
    bio_run.font.name = "Calibri"

    # ==========================================
    # Layer 3: Right Panel (Dashboard)
    # ==========================================
    right_x_start = left_width + Inches(0.5)
    
    # 1. Header
    title_box = slide.shapes.add_textbox(right_x_start, Inches(0.5), Inches(7.0), Inches(1.0))
    tf = title_box.text_frame
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = expert_name
    r1.font.size = Pt(44)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK_PANEL
    
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = tagline
    r2.font.size = Pt(18)
    r2.font.color.rgb = COLOR_ACCENT

    # 2. Timeline
    timeline_y = Inches(3.0)
    timeline_width = Inches(6.5)
    
    # Base Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x_start + Inches(0.5), timeline_y, timeline_width - Inches(1.0), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()

    # Nodes and Text
    num_nodes = len(timeline_data)
    node_spacing = (timeline_width - Inches(1.0)) / (num_nodes - 1) if num_nodes > 1 else 0
    node_size = Inches(0.2)

    for i, (year, desc) in enumerate(timeline_data):
        nx = right_x_start + Inches(0.5) + (i * node_spacing)
        
        # Node circle
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, nx - (node_size/2), timeline_y - (node_size/2) + Inches(0.03), node_size, node_size)
        node.fill.solid()
        node.fill.fore_color.rgb = COLOR_ACCENT
        node.line.fill.background()

        # Year Text (Above)
        yt_box = slide.shapes.add_textbox(nx - Inches(1.0), timeline_y - Inches(0.6), Inches(2.0), Inches(0.5))
        yt_p = yt_box.text_frame.paragraphs[0]
        yt_p.alignment = PP_ALIGN.CENTER
        yt_run = yt_p.add_run()
        yt_run.text = year
        yt_run.font.size = Pt(16)
        yt_run.font.color.rgb = COLOR_TEXT_DARK

        # Desc Text (Below)
        dt_box = slide.shapes.add_textbox(nx - Inches(1.0), timeline_y + Inches(0.2), Inches(2.0), Inches(1.0))
        dt_box.text_frame.word_wrap = True
        dt_p = dt_box.text_frame.paragraphs[0]
        dt_p.alignment = PP_ALIGN.CENTER
        dt_run = dt_p.add_run()
        dt_run.text = desc
        dt_run.font.size = Pt(11)
        dt_run.font.color.rgb = COLOR_TEXT_DARK

    # 3. Skill Bars
    skills_start_y = Inches(4.8)
    skill_spacing = Inches(0.8)
    bar_width = Inches(5.0)
    bar_height = Inches(0.18)

    for i, (skill, pct) in enumerate(skills_data):
        sy = skills_start_y + (i * skill_spacing)
        
        # Skill Name
        sk_box = slide.shapes.add_textbox(right_x_start + Inches(0.2), sy - Inches(0.35), Inches(3.0), Inches(0.4))
        sk_p = sk_box.text_frame.paragraphs[0]
        sk_run = sk_p.add_run()
        sk_run.text = skill
        sk_run.font.size = Pt(14)
        sk_run.font.color.rgb = COLOR_TEXT_DARK

        # Background Track Bar
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x_start + Inches(0.3), sy, bar_width, bar_height)
        track.fill.solid()
        track.fill.fore_color.rgb = COLOR_TRACK_BG
        track.line.fill.background()

        # Foreground Fill Bar
        fill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x_start + Inches(0.3), sy, bar_width * pct, bar_height)
        fill.fill.solid()
        fill.fill.fore_color.rgb = COLOR_DARK_PANEL
        fill.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
