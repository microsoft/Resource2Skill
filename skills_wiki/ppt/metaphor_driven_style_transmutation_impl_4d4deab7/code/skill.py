import os
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw

def _create_glass_card(width_in, height_in, dpi=300):
    """
    Helper function: Generates a translucent glassmorphism card using PIL.
    Returns a BytesIO object containing the PNG image.
    """
    # Convert inches to pixels based on DPI
    w_px = int(width_in * dpi)
    h_px = int(height_in * dpi)
    
    # Create transparent image
    img = Image.new('RGBA', (w_px, h_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # Glassmorphism properties
    radius = int(0.15 * dpi) # Corner radius
    fill_color = (255, 255, 255, 18)    # 7% opaque white
    outline_color = (255, 255, 255, 75) # 30% opaque white edge
    line_width = int(0.02 * dpi)
    
    # Draw rounded rectangle
    draw.rounded_rectangle(
        [(line_width, line_width), (w_px - line_width, h_px - line_width)],
        radius=radius,
        fill=fill_color,
        outline=outline_color,
        width=line_width
    )
    
    image_stream = io.BytesIO()
    img.save(image_stream, format='PNG')
    image_stream.seek(0)
    return image_stream

def create_slide(
    output_pptx_path: str,
    title_text: str = "市场运行的底层代码",
    subtitle_text: str = "解码驱动万物价格的无形力量",
    card_data: list = [
        {"title": "需求 (Demand)", "desc": "买家的心智与购买力"},
        {"title": "供给 (Supply)", "desc": "生产者的算盘与产能"},
        {"title": "均衡 (Equilibrium)", "desc": "看不见的手促成的握手"}
    ],
    bg_color: tuple = (13, 17, 28),
    accent_color: tuple = (0, 191, 255),
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the "Apple Keynote Tech Glassmorphism" style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    # Solid deep tech background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*bg_color)
    bg.line.fill.background() # No line

    # === Layer 2: Hero Glass Card (Top) ===
    hero_w, hero_h = 11.333, 2.0
    hero_left = Inches(1.0)
    hero_top = Inches(1.0)
    
    hero_card_stream = _create_glass_card(hero_w, hero_h)
    slide.shapes.add_picture(hero_card_stream, hero_left, hero_top, width=Inches(hero_w), height=Inches(hero_h))

    # Hero Text
    hero_tb = slide.shapes.add_textbox(hero_left, hero_top + Inches(0.3), Inches(hero_w), Inches(1.0))
    tf = hero_tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(*accent_color)

    # === Layer 3: Metric Glass Cards (Bottom Grid) ===
    num_cards = len(card_data)
    card_w = 3.5
    card_h = 2.5
    spacing = 0.416
    start_left = 1.0
    cards_top = Inches(3.8)

    for i, data in enumerate(card_data):
        left_pos = Inches(start_left + i * (card_w + spacing))
        
        # Insert PIL Glass Card
        card_stream = _create_glass_card(card_w, card_h)
        slide.shapes.add_picture(card_stream, left_pos, cards_top, width=Inches(card_w), height=Inches(card_h))
        
        # Accent Line (Top of card indicator)
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left_pos + Inches(0.2), cards_top + Inches(0.3), Inches(0.3), Inches(0.05)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = RGBColor(*accent_color)
        accent_line.line.fill.background()

        # Card Textbox
        tb = slide.shapes.add_textbox(left_pos + Inches(0.1), cards_top + Inches(0.5), Inches(card_w - 0.2), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Title
        p_title = tf.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)
        
        # Description
        p_desc = tf.add_paragraph()
        p_desc.text = "\n" + data["desc"]
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = RGBColor(200, 200, 200)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("apple_keynote_style.pptx")
