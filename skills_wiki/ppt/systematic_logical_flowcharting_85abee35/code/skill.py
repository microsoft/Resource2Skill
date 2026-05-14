def create_slide(
    output_pptx_path: str,
    title_text: str = "Office Attendance Flowchart",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Systematic Logical Flowcharting visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Color Palette ===
    COLOR_START = RGBColor(0, 112, 192)    # Blue
    COLOR_PROCESS = RGBColor(0, 176, 80)   # Green
    COLOR_DECISION = RGBColor(192, 0, 0)   # Red
    COLOR_TEXT = RGBColor(0, 0, 0)         # Black
    COLOR_LABEL = RGBColor(80, 80, 80)     # Dark Gray

    # === Helper: Create and Style Nodes ===
    def add_node(shape_type, text, x, y, w, h, border_color):
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        
        # Solid white fill to mask anything behind it
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        
        # Thick colored border
        shape.line.color.rgb = border_color
        shape.line.width = Pt(3)
        
        # Text styling
        shape.text_frame.word_wrap = True
        p = shape.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = text
        p.runs[0].font.color.rgb = COLOR_TEXT
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.bold = True
        return shape

    # === Helper: Add Connectors with Arrowheads via lxml ===
    def add_arrowed_connector(node_from, site_from, node_to, site_to):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Inches(0), Inches(0), Inches(1), Inches(1)
        )
        connector.begin_connect(node_from, site_from)
        connector.end_connect(node_to, site_to)
        
        # Line styling
        connector.line.color.rgb = RGBColor(0, 0, 0)
        connector.line.width = Pt(1.5)
        
        # Inject Arrowhead via Open XML
        spPr = connector.element.spPr
        ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is None:
            ln = OxmlElement('a:ln')
            spPr.append(ln)
        tailEnd = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd')
        if tailEnd is None:
            tailEnd = OxmlElement('a:tailEnd')
            ln.append(tailEnd)
        tailEnd.set('type', 'triangle')
        tailEnd.set('w', 'med')
        tailEnd.set('len', 'med')
        
        return connector

    # === Helper: Add Floating Labels ===
    def add_label(text, x, y):
        txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(1), Inches(0.5))
        p = txBox.text_frame.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(12)
        p.runs[0].font.color.rgb = COLOR_LABEL

    # === Build Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(5), Inches(0.8))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.runs[0].font.size = Pt(28)
    title_p.runs[0].font.bold = True

    # === Build Flowchart Nodes ===
    # Center X = 5.4 to make a 2.5 wide shape perfectly centered at 6.66
    n_start = add_node(MSO_SHAPE.ROUNDED_RECTANGLE, "Start", 5.4, 0.8, 2.5, 0.6, COLOR_START)
    n_punch = add_node(MSO_SHAPE.RECTANGLE, "Fingerprint Punching", 5.4, 2.1, 2.5, 0.8, COLOR_PROCESS)
    n_decision = add_node(MSO_SHAPE.DIAMOND, "Scan Valid?", 5.15, 3.6, 3.0, 1.2, COLOR_DECISION)
    n_attend = add_node(MSO_SHAPE.RECTANGLE, "Mark Attendance", 5.4, 5.5, 2.5, 0.8, COLOR_PROCESS)
    n_end = add_node(MSO_SHAPE.ROUNDED_RECTANGLE, "End", 5.4, 6.7, 2.5, 0.6, COLOR_START)
    
    # Lateral exception branch
    n_letter = add_node(MSO_SHAPE.PARALLELOGRAM, "Submit Letter", 9.5, 3.7, 2.5, 1.0, COLOR_START)

    # === Build Routing (Connectors) ===
    # Standard connection sites: 0=Top, 1=Left, 2=Bottom, 3=Right
    add_arrowed_connector(n_start, 2, n_punch, 0)
    add_arrowed_connector(n_punch, 2, n_decision, 0)
    
    # Branch: Yes
    add_arrowed_connector(n_decision, 2, n_attend, 0)
    add_label("Yes", 6.8, 4.9)
    
    # Branch: No
    add_arrowed_connector(n_decision, 3, n_letter, 1)
    add_label("No", 8.4, 3.9)
    
    # Reconvergence
    add_arrowed_connector(n_letter, 2, n_attend, 3)
    
    # Final step
    add_arrowed_connector(n_attend, 2, n_end, 0)

    prs.save(output_pptx_path)
    return output_pptx_path
