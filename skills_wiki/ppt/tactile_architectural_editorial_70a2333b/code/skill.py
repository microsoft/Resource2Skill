import os
import urllib.request
from PIL import Image, ImageDraw, ImageChops
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR

def _create_tactile_background(width, height, color_top, color_bottom, output_path):
    """
    Generates a rich gradient blended with a paper texture using PIL.
    """
    # 1. Create Base Vertical Gradient
    grad = Image.new('RGB', (1, int(height)))
    draw = ImageDraw.Draw(grad)
    for y in range(int(height)):
        f = y / float(height)
        r = int(color_top[0] + (color_bottom[0] - color_top[0]) * f)
        g = int(color_top[1] + (color_bottom[1] - color_top[1]) * f)
        b = int(color_top[2] + (color_bottom[2] - color_top[2]) * f)
        draw.point((0, y), (r, g, b))
    
    # Scale up smoothly
    try:
        resample_filter = Image.Resampling.LANCZOS
    except AttributeError:
        resample_filter = Image.LANCZOS
        
    grad = grad.resize((int(width), int(height)), resample_filter)
    
    # 2. Blend with a High-Res Paper Texture
    try:
        # Stable URL for a light, subtle watercolor paper texture
        tex_url = "https://images.unsplash.com/photo-1586075010923-2dd4570fb338?q=80&w=1920&auto=format&fit=crop"
        req = urllib.request.Request(tex_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            texture = Image.open(response).convert('L')
            
            # Crop to fill canvas aspect ratio
            tex_aspect = texture.width / texture.height
            canvas_aspect = width / height
            if tex_aspect > canvas_aspect:
                new_w = int(texture.height * canvas_aspect)
                left = (texture.width - new_w) // 2
                texture = texture.crop((left, 0, left + new_w, texture.height))
            else:
                new_h = int(texture.width / canvas_aspect)
                top = (texture.height - new_h) // 2
                texture = texture.crop((0, top, texture.width, top + new_h))
                
            texture = texture.resize((int(width), int(height)), resample_filter)
            
            # Convert texture to RGB for blending
            texture_rgb = Image.merge('RGB', (texture, texture, texture))
            
            # Multiply blend to darken the gradient based on paper crevices
            multiplied = ImageChops.multiply(grad, texture_rgb)
            
            # Final blend (40% texture intensity) to keep colors rich
            final_img = Image.blend(grad, multiplied, 0.4)
            final_img.save(output_path)
    except Exception as e:
        print(f"Texture blend failed, falling back to base gradient: {e}")
        grad.save(output_path)

def create_slide(
    output_pptx_path: str,
    title_text: str = "1. CONCEPT WRAPPING",
    subtitle_text: str = "Key Qualities of a Compelling Architecture Presentation:",
    bg_color_top: tuple = (180, 90, 50),     # Terracotta top
    bg_color_bottom: tuple = (90, 40, 20),   # Dark clay bottom
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Tactile Architectural Editorial' design style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Tactile Texture Background ===
    bg_img_path = "temp_tactile_bg.jpg"
    # Convert inches to pixels (approx 150 DPI for good quality)
    _create_tactile_background(
        width=13.333 * 150, 
        height=7.5 * 150, 
        color_top=bg_color_top, 
        color_bottom=bg_color_bottom, 
        output_path=bg_img_path
    )
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Editorial Typography ===

    # A. Subtitle (Classic Italic Serif)
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(2.7), Inches(11.333), Inches(0.8))
    sub_tf = sub_box.text_frame
    sub_tf.word_wrap = True
    p_sub = sub_tf.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = subtitle_text
    p_sub.font.name = 'Georgia'
    p_sub.font.size = Pt(20)
    p_sub.font.italic = True
    p_sub.font.color.rgb = RGBColor(235, 225, 215)  # Warm off-white

    # B. Divider Line (Subtle structural element)
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(6.0), Inches(3.55), Inches(7.333), Inches(3.55)
    )
    line.line.color.rgb = RGBColor(200, 150, 120)
    line.line.width = Pt(1)

    # C. Main Title (Wide-tracked Sans-Serif)
    # Simulate wide tracking (letter-spacing) by injecting spaces between characters
    tracked_title = "   ".join(list(title_text.upper()))
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.333), Inches(1.5))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    p_title = title_tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.CENTER
    p_title.text = tracked_title
    p_title.font.name = 'Arial'  # Clean geometric sans-serif
    p_title.font.size = Pt(36)
    p_title.font.bold = False
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
