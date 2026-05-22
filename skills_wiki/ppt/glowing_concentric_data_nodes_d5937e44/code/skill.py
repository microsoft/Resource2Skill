def create_slide(
    output_pptx_path: str,
    title_text: str = "SPINNING CIRCLES TUTORIAL",
    body_text: str = "A total of 3 circles labeled with A, B, and C respectively.",
    bg_keyword: str = "dark geometric",  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Glowing Concentric Data Nodes visual effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # --- Helper Functions for XML Injection (Effects) ---
    def add_glow(shape, color_hex="00CCFF", radius_pt=18):
        """Injects a glow effect into a shape's XML."""
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        glow = OxmlElement('a:glow')
        glow.set('rad', str(int(radius_pt * 12700))) # Convert pt to EMUs
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', color_hex)
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '50000') # 50% opacity
        srgbClr.append(alpha)
        glow.append(srgbClr)
        effectLst.append(glow)
        spPr.append(effectLst)

    def add_shadow(shape):
        """Injects a drop shadow effect to simulate 3D depth."""
        spPr = shape.element.spPr
        # Remove existing effectLst if present to avoid conflicts
        for elem in spPr.findall('.//a:effectLst', namespaces=spPr.nsmap):
            spPr.remove(elem)
            
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', '63500') # 5pt
        outerShdw.set('dist', '38100')    # 3pt
        outerShdw.set('dir', '2700000')   # 45 degrees
        outerShdw.set('algn', 'tl')
        outerShdw.set('rotWithShape', '0')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '60000') # 60% opacity
        srgbClr.append(alpha)
        
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    # --- Layer 1: Background ---
    # Try downloading a dark geometric background
    bg_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1600x900/?{urllib.parse.quote(bg_keyword)}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(bg_path, 'wb') as out_file:
            out_file.write(response.read())
        slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception:
        # Fallback to dark solid background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(15, 20, 25)

    # --- Layer 2: Main Title ---
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.5), Inches(9.333), Inches(1))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = "Georgia" # Serif font as in tutorial
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Add a subtle shadow to title
    add_shadow(title_box)

    # --- Layer 3: Concentric Nodes ---
    labels = ["A", "B", "C"]
    
    # Calculate horizontal distribution
    num_nodes = len(labels)
    total_width = prs.slide_width
    spacing = total_width / (num_nodes + 1)
    
    # Node Configuration
    outer_radius = Inches(1.3)
    mid_radius = Inches(1.0)
    inner_radius = Inches(0.7)
    cy = Inches(4.0) # Center Y coordinate

    colors = {
        "outer": RGBColor(17, 65, 136),   # Darkest Blue
        "mid": RGBColor(41, 108, 196),    # Medium Blue
        "inner": RGBColor(90, 155, 220)   # Light Blue
    }

    for i, label in enumerate(labels):
        cx = spacing * (i + 1) # Center X coordinate
        
        # 1. Outer Circle (Glow)
        outer_shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            cx - outer_radius, cy - outer_radius, 
            outer_radius * 2, outer_radius * 2
        )
        outer_shape.fill.solid()
        outer_shape.fill.fore_color.rgb = colors["outer"]
        outer_shape.line.color.rgb = RGBColor(0, 40, 80)
        outer_shape.line.width = Pt(2)
        add_glow(outer_shape, color_hex="00CCFF", radius_pt=18)

        # 2. Middle Circle (Shadow/Bevel simulation)
        mid_shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            cx - mid_radius, cy - mid_radius, 
            mid_radius * 2, mid_radius * 2
        )
        mid_shape.fill.solid()
        mid_shape.fill.fore_color.rgb = colors["mid"]
        mid_shape.line.color.rgb = RGBColor(100, 150, 255)
        mid_shape.line.width = Pt(1.5)
        add_shadow(mid_shape)

        # 3. Inner Circle (Text)
        inner_shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            cx - inner_radius, cy - inner_radius, 
            inner_radius * 2, inner_radius * 2
        )
        inner_shape.fill.solid()
        inner_shape.fill.fore_color.rgb = colors["inner"]
        inner_shape.line.color.rgb = RGBColor(200, 220, 255)
        inner_shape.line.width = Pt(1)
        add_shadow(inner_shape)

        # Text inside inner shape
        tf = inner_shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.name = "Arial"
        p.font.color.rgb = RGBColor(255, 255, 255)

    # --- Layer 4: Description Text ---
    desc_box = slide.shapes.add_textbox(Inches(2), Inches(6.0), Inches(9.333), Inches(0.8))
    tf = desc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(20)
    p.font.name = "Arial"
    p.font.color.rgb = RGBColor(220, 220, 220)

    # --- Save and Cleanup ---
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
