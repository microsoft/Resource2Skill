def create_slide(
    output_pptx_path: str,
    title_text: str = "GLOBAL MARKET & TECH",
    subtitle_text: str = "TREND ANALYSIS REPORT 2024",
    accent_color: tuple = (0, 210, 255),  # Cyan
    panel_color: tuple = (10, 15, 25),    # Dark Navy
    panel_alpha: int = 210,               # 0-255 transparency (approx 82% opaque)
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Layered Color Block & Frame" cover design style.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background (Image with PIL Fallback) ===
    bg_path = "temp_bg.jpg"
    try:
        # Try fetching a high-quality abstract background
        req = urllib.request.Request(
            "https://picsum.photos/1920/1080?blur=2", 
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req) as response, open(bg_path, 'wb') as out_file:
            out_file.write(response.read())
        slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception as e:
        print(f"Network image failed, generating PIL gradient fallback: {e}")
        # PIL Gradient Fallback
        base = Image.new('RGB', (1920, 1080), (10, 15, 30))
        top = Image.new('RGB', (1920, 1080), (0, 60, 80))
        mask = Image.new('L', (1920, 1080))
        mask_data = [int(255 * ((x + y) / 3000)) for y in range(1080) for x in range(1920)]
        mask.putdata(mask_data)
        base.paste(top, (0, 0), mask)
        base.save(bg_path)
        slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Semi-transparent Color Block (The "色块") ===
    panel_path = "temp_panel.png"
    # Create a 10x10 RGBA image (it will be stretched, maintaining flawless solid transparency)
    img = Image.new('RGBA', (10, 10), panel_color + (panel_alpha,))
    img.save(panel_path)
    
    panel_width, panel_height = Inches(10.5), Inches(3.8)
    panel_left = (prs.slide_width - panel_width) / 2
    panel_top = (prs.slide_height - panel_height) / 2
    slide.shapes.add_picture(panel_path, panel_left, panel_top, width=panel_width, height=panel_height)

    # === Layer 3: Accent Line Frame (The "线框") ===
    frame_width, frame_height = panel_width - Inches(0.3), panel_height - Inches(0.3)
    frame_left = panel_left + Inches(0.15)
    frame_top = panel_top + Inches(0.15)
    
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, frame_left, frame_top, frame_width, frame_height)
    frame.fill.background() # Transparent fill
    frame.line.color.rgb = RGBColor(*accent_color)
    frame.line.width = Pt(1.5)

    # Add Tech Corner Accents (Small squares at the corners of the frame)
    corner_size = Inches(0.06)
    corners = [
        (frame_left, frame_top),
        (frame_left + frame_width - corner_size, frame_top),
        (frame_left, frame_top + frame_height - corner_size),
        (frame_left + frame_width - corner_size, frame_top + frame_height - corner_size)
    ]
    for cx, cy in corners:
        dot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, corner_size, corner_size)
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(*accent_color)
        dot.line.fill.background()

    # === Layer 4: Typography & Accent Ribbon ===
    # 4a. Main Title
    title_box = slide.shapes.add_textbox(panel_left, panel_top + Inches(0.6), panel_width, Inches(1.2))
    tf = title_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    font = tf.paragraphs[0].font
    font.name = 'Arial'
    font.size = Pt(48)
    font.bold = True
    font.color.rgb = RGBColor(255, 255, 255)

    # 4b. Solid Accent Ribbon for Subtitle
    ribbon_width, ribbon_height = Inches(4.5), Inches(0.45)
    ribbon_left = (prs.slide_width - ribbon_width) / 2
    ribbon_top = panel_top + Inches(2.2)
    
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ribbon_left, ribbon_top, ribbon_width, ribbon_height)
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = RGBColor(*accent_color)
    ribbon.line.fill.background()

    # 4c. Subtitle Text (Overlapping the ribbon)
    sub_box = slide.shapes.add_textbox(ribbon_left, ribbon_top - Inches(0.04), ribbon_width, ribbon_height)
    tf_sub = sub_box.text_frame
    tf_sub.text = subtitle_text
    tf_sub.paragraphs[0].alignment = PP_ALIGN.CENTER
    font_sub = tf_sub.paragraphs[0].font
    font_sub.name = 'Arial'
    font_sub.size = Pt(16)
    font_sub.bold = True
    # Very dark color for high contrast against the bright accent ribbon
    font_sub.color.rgb = RGBColor(10, 15, 25)

    # Cleanup temp files
    try:
        if os.path.exists(bg_path): os.remove(bg_path)
        if os.path.exists(panel_path): os.remove(panel_path)
    except Exception:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path
