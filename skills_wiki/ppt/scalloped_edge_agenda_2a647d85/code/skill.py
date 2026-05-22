def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda Slide Design #1",
    agenda_items: list = None,
    image_keyword: str = "business meeting",
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with the "Scalloped Edge Agenda" design.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        title_text (str): The main title for the slide.
        agenda_items (list): A list of strings, where each string is an agenda item.
        image_keyword (str): A keyword to search for a background image on Unsplash.

    Returns:
        str: The path to the saved PPTX file.
    """
    import requests
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    if agenda_items is None:
        agenda_items = [
            "Introduction",
            "Why do we need the transformation?",
            "What do we need to transform?",
            "How to transform the organization?",
            "Who are those driving the transformation?",
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background (White) ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Image ===
    title_height = Inches(1.2)
    img_width = Inches(5.8)
    img_height = prs.slide_height - title_height
    img_left = Inches(0)
    img_top = title_height
    
    try:
        unsplash_url = f"https://source.unsplash.com/1600x900/?{image_keyword}"
        response = requests.get(unsplash_url, timeout=10)
        image_stream = BytesIO(response.content)
        pic = slide.shapes.add_picture(
            image_stream,
            img_left,
            img_top,
            width=img_width,
            height=img_height
        )
    except requests.exceptions.RequestException:
        # Fallback to a solid color rectangle if image download fails
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, img_left, img_top, img_width, img_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(128, 128, 128)
        shape.line.fill.background()

    # === Layer 3: Text & Content ===
    # --- Main Title ---
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    title_tf = title_shape.text_frame
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text
    title_p.font.name = "Arial"
    title_p.font.size = Pt(32)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(0, 0, 0)

    # --- Agenda Items and Scalloped Edge ---
    num_items = len(agenda_items)
    content_area_height = prs.slide_height - title_height
    item_v_spacing = content_area_height / num_items
    circle_diameter = Inches(0.8)

    for i, item_text in enumerate(agenda_items):
        v_center = title_height + (i + 0.5) * item_v_spacing

        # A. Add the white circle for the cutout effect
        circle_left = img_width - (circle_diameter / 2)
        circle_top = v_center - (circle_diameter / 2)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, circle_left, circle_top, circle_diameter, circle_diameter)
        
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
        circle.line.fill.background() # No outline

        # B. Add the number inside the circle
        tf_num = circle.text_frame
        tf_num.clear()
        p_num = tf_num.paragraphs[0]
        p_num.text = str(i + 1)
        p_num.font.name = "Arial"
        p_num.font.size = Pt(18)
        p_num.font.bold = True
        p_num.font.color.rgb = RGBColor(50, 50, 50)
        p_num.alignment = PP_ALIGN.CENTER
        tf_num.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
        
        # C. Add agenda item text
        text_left = img_width + Inches(0.5)
        text_width = prs.slide_width - text_left - Inches(0.5)
        text_height = item_v_spacing
        text_top = title_height + i * item_v_spacing
        
        txt_box = slide.shapes.add_textbox(text_left, text_top, text_width, text_height)
        tf_item = txt_box.text_frame
        tf_item.word_wrap = True
        tf_item.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
        p_item = tf_item.paragraphs[0]
        p_item.text = item_text
        p_item.font.name = "Arial"
        p_item.font.size = Pt(16)
        p_item.font.color.rgb = RGBColor(89, 89, 89)

    prs.save(output_pptx_path)
    return output_pptx_path
