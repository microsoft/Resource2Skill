def create_slide(
    output_pptx_path: str,
    title_text: str = "Chart Title",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Native Data-Driven Chart Integration.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    from pptx.enum.chart import XL_LEGEND_POSITION

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Define Chart Data ===
    # Reproducing the sample data structure seen in the tutorial's Excel sheet
    chart_data = CategoryChartData()
    chart_data.categories = ['Category 1', 'Category 2', 'Category 3', 'Category 4']
    
    # Adding multiple series to create the "Clustered" effect
    chart_data.add_series('Series 1', (4.3, 2.5, 3.5, 4.5))
    chart_data.add_series('Series 2', (2.4, 4.4, 1.8, 2.8))
    chart_data.add_series('Series 3', (2.0, 2.0, 3.0, 5.0))

    # === Define Positioning and Add Chart ===
    # Centered with good margins
    x = Inches(2.0)
    y = Inches(1.5)
    cx = Inches(9.333)
    cy = Inches(5.0)

    # Insert the chart
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    )
    chart = graphic_frame.chart

    # === Customize Chart Appearance ===
    
    # 1. Title
    chart.has_title = True
    chart.chart_title.text_frame.text = title_text
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(18)
    
    # 2. Legend
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    
    # 3. Axes formatting (optional refinement)
    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
