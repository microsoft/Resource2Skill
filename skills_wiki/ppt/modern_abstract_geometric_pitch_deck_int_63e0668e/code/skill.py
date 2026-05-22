def create_slide(
    output_pptx_path: str,
    title_text: str = "Awesome\nPresentation",
    kicker_text: str = "About Us",
    body_text: str = "Write a compelling caption here. This layout utilizes modern minimalist design principles, overlapping abstract geometry, and crisp typography to capture audience attention.",
    grad_color1: tuple = (224, 86, 253),  # Vibrant Pink
    grad_color2: tuple = (142, 68, 173),  # Deep Purple
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Modern Abstract Geometric Pitch Deck Intro' visual effect.
    """
    import io
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    # --- Helper: Generate Diagonal Gradient Image ---
    def create_diagonal_gradient(size=(400, 400), c1=(255,255,255), c2=(0,0,0)):
        img = Image.new('RGB', size)
        pixels = img.load()
        max_dist = size[0] + size[1]
        for x in range(size[0]):
            for y in range(size[1]):
                ratio = (x + y) / max_dist
                r = int(c1[0] + (c2[0] - c1[0]) * ratio)
                g = int(c1[1] + (c2[1] - c1[1]) * ratio)
                b = int(c1[2] + (c2[2] - c1[2]) * ratio)
                pixels[x, y] = (r, g, b)
        return img

    # --- Helper: Generate Dot Grid Pattern Image ---
    def create_dot_pattern(size=(400, 400), dot_color=(200, 200, 200, 255), bg_color=(250, 250, 250, 255), spacing=18, radius=1.5):
        img = Image.new('RGBA', size, bg_color)
        draw = ImageDraw.Draw(img)
        for x in range(0, size[0], spacing):
            for y in range(0, size[1], spacing):
                draw.ellipse([x-radius, y-radius, x+radius, y+radius], fill=dot_color)
        return img

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Set pure white background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Layer 1: Text & Typography (Left Side) ===
    
    # 1. Kicker Text
    tx_kicker = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(3.0), Inches(0.5))
    tf_kicker = tx_kicker.text_frame
    p_kicker = tf_kicker.paragraphs[0]
    p_kicker.text = kicker_text.upper()
    p_kicker.font.size = Pt(12)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = RGBColor(100, 100, 100)

    # 2. Colored Accent Line (Anchoring Kicker)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.85), Inches(0.4), Inches(0.03)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(grad_color1[0], grad_color1[1], grad_color1[2])
    line.line.fill.background()

    # 3. Main Oversized Title
    tx_title = slide.shapes.add_textbox(Inches(0.95), Inches(2.2), Inches(6.0), Inches(2.0))
    tf_title = tx_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(64)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(30, 30, 30)

    # 4. Body Copy
    tx_body = slide.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(4.8), Inches(1.5))
    tf_body = tx_body.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(14)
    p_body.font.color.rgb = RGBColor(110, 110, 110)

    # === Layer 2: Abstract Geometric Cluster (Right Side) ===

    # Generate Image Assets in memory
    img_grad = create_diagonal_gradient(size=(500, 500), c1=grad_color1, c2=grad_color2)
    img_pat = create_dot_pattern(size=(500, 500))

    # Pattern Square 1 (Back Layer)
    pat1_io = io.BytesIO()
    img_pat.save(pat1_io, format='PNG')
    pat1_io.seek(0)
    slide.shapes.add_picture(pat1_io, Inches(6.8), Inches(2.5), Inches(3.2), Inches(3.2))

    # Gradient Square (Middle Hero Layer)
    grad_io = io.BytesIO()
    img_grad.save(grad_io, format='PNG')
    grad_io.seek(0)
    # Add picture; placing it offset from Pat 1
    slide.shapes.add_picture(grad_io, Inches(8.5), Inches(1.2), Inches(3.5), Inches(3.5))

    # Pattern Square 2 (Front Layer)
    pat2_io = io.BytesIO()
    img_pat.save(pat2_io, format='PNG')
    pat2_io.seek(0)
    # Overlapping the bottom right corner of the gradient
    slide.shapes.add_picture(pat2_io, Inches(9.5), Inches(4.0), Inches(2.8), Inches(2.8))

    # Save output
    prs.save(output_pptx_path)
    return output_pptx_path
