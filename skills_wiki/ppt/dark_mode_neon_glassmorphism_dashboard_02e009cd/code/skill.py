def create_slide(
    output_pptx_path: str,
    title_text: str = "Portfolio Overview",
    metric_1_val: str = "$129,987.05",
    metric_2_val: str = "+ 14.5%",
    metric_3_val: str = "Healthy",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Dark Mode Neon Glassmorphism Dashboard' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml
    from PIL import Image, ImageDraw, ImageFilter
    import os

    # === Step 1: Generate Neon Aura Background via PIL ===
    # We use a small canvas for a fast, extreme blur, then upscale for buttery smoothness
    bg_img_path = "temp_neon_aura.png"
    base_canvas = Image.new('RGB', (320, 180), (10, 11, 16)) # Deep dark background
    draw = ImageDraw.Draw(base_canvas)
    
    # Draw vivid colorful orbs where the cards will roughly sit
    draw.ellipse([-50, 20, 150, 220], fill=(0, 200, 255))   # Cyan glow (Behind Left Card)
    draw.ellipse([180, -40, 380, 160], fill=(255, 0, 128))  # Magenta glow (Behind Right Cards)
    draw.ellipse([100, 100, 250, 250], fill=(112, 0, 255))  # Purple center blender
    
    # Extreme blur to create the aura
    blurred = base_canvas.filter(ImageFilter.GaussianBlur(radius=30))
    # Upscale to 1080p
    final_bg = blurred.resize((1920, 1080), Image.Resampling.LANCZOS)
    final_bg.save(bg_img_path)

    # === Step 2: Presentation Setup ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Apply background image
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # Helper function to create Glassmorphism cards using lxml
    def add_glass_card(left, top, width, height):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        # Set solid dark color
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(26, 30, 41)
        
        # Inject Alpha Transparency (85% opacity / 15% transparent) for Glass effect
        srgbClr = card.fill.fore_color._xClr
        alpha_xml = '<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="85000"/>'
        srgbClr.append(parse_xml(alpha_xml))
        
        # Add crisp border
        card.line.color.rgb = RGBColor(60, 65, 80)
        card.line.width = Pt(1.5)
        return card

    # Helper function for consistent text formatting
    def add_text(shape, text, pt_size, rgb_color, bold=False, top_margin=0.2):
        tb = shape.text_frame
        tb.margin_top = Inches(top_margin)
        tb.margin_left = Inches(0.4)
        p = tb.add_paragraph() if tb.text else tb.paragraphs[0]
        p.text = text
        p.font.size = Pt(pt_size)
        p.font.color.rgb = RGBColor(*rgb_color)
        p.font.bold = bold
        p.font.name = "Arial"

    # === Step 3: Main Slide Title ===
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(1))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.font.size = Pt(32)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    # === Step 4: Build Dashboard UI Grid ===

    # Card 1: Hero Metric (Left)
    hero_card = add_glass_card(Inches(1), Inches(1.5), Inches(5.5), Inches(5))
    add_text(hero_card, "TOTAL BALANCE", 12, (150, 160, 180), bold=True, top_margin=0.5)
    add_text(hero_card, metric_1_val, 48, (255, 255, 255), bold=True, top_margin=0)
    
    # Add a visual element (Progress Bar) inside Hero Card
    bar_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(4), Inches(4.7), Inches(0.3))
    bar_bg.fill.solid()
    bar_bg.fill.fore_color.rgb = RGBColor(15, 17, 24)
    bar_bg.line.fill.background()
    
    bar_fill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(4), Inches(3.2), Inches(0.3))
    bar_fill.fill.solid()
    bar_fill.fill.fore_color.rgb = RGBColor(0, 229, 255) # Cyan fill
    bar_fill.line.fill.background()

    # Card 2: Secondary Metric (Top Right)
    top_right_card = add_glass_card(Inches(7), Inches(1.5), Inches(5.333), Inches(2.3))
    add_text(top_right_card, "MONTHLY GROWTH", 12, (150, 160, 180), bold=True, top_margin=0.4)
    add_text(top_right_card, metric_2_val, 36, (255, 0, 128), bold=True, top_margin=0) # Magenta text

    # Card 3: Tertiary Metric (Bottom Right)
    bottom_right_card = add_glass_card(Inches(7), Inches(4.2), Inches(5.333), Inches(2.3))
    add_text(bottom_right_card, "SYSTEM STATUS", 12, (150, 160, 180), bold=True, top_margin=0.4)
    add_text(bottom_right_card, metric_3_val, 36, (255, 255, 255), bold=True, top_margin=0)

    # Clean up temp files
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
