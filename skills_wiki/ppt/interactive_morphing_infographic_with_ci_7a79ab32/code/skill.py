import os
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_FILL
from lxml import etree

def add_inner_shadow(shape, blur_radius=Pt(30), distance=0, direction=0, color=(255, 255, 255), transparency=0.5):
    """Applies an inner shadow effect to a shape by manipulating its XML properties."""
    spPr = shape.element.spPr
    effect_lst_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst"
    inner_shdw_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}innerShdw"
    srgb_clr_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
    alpha_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"

    effect_lst = spPr.find(effect_lst_tag)
    if effect_lst is None:
        effect_lst = etree.SubElement(spPr, effect_lst_tag)

    inner_shadow = etree.SubElement(effect_lst, inner_shdw_tag)
    inner_shadow.set("blurRad", str(int(blur_radius.emu)))
    inner_shadow.set("dist", str(Emu(distance)))
    inner_shadow.set("dir", str(int(direction * 60000)))

    srgb_clr = etree.SubElement(inner_shadow, srgb_clr_tag)
    srgb_clr.set("val", f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    
    alpha = etree.SubElement(srgb_clr, alpha_tag)
    alpha.set("val", str(int((1 - transparency) * 100000)))

def create_slide(
    output_pptx_path: str,
    title_text: str = "YOUR IDEAS",
    item_titles: list = None,
    item_texts: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Interactive Morphing Infographic with Circular Menu" visual style.
    This function generates one complete slide. The interactive morphing, hyperlinks, and trigger animations
    must be set up manually in PowerPoint by duplicating this slide and modifying content.

    Returns: path to the saved PPTX file.
    """
    if item_texts is None:
        item_texts = ["Insert some text\nhere if needed"] * 4
    if item_titles is None:
        item_titles = ["IDEA"] * 4

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background Gradient ===
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_type = 'linear'
    fill.gradient_angle = 90
    
    stop1 = fill.gradient_stops.add()
    stop1.position = 0.0
    stop1.color.rgb = RGBColor(7, 244, 158)
    
    stop2 = fill.gradient_stops.add()
    stop2.position = 0.80
    stop2.color.rgb = RGBColor(66, 4, 126)

    # === Layer 2: Main Title & Info Boxes (Glassmorphism Style) ===
    title_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(0.5), Inches(7), Inches(1))
    title_shape.adjustments[0] = 0.5
    title_shape.fill.background()
    title_shape.line.color.rgb = RGBColor(255, 255, 255)
    title_shape.line.width = Pt(0.75)
    add_inner_shadow(title_shape, blur_radius=Pt(30), transparency=0.4)
    
    text_frame = title_shape.text_frame
    text_frame.text = title_text
    p = text_frame.paragraphs[0]
    p.font.name = 'Montserrat SemiBold'
    p.font.size = Pt(32)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = 1

    positions = [(2.5, 3.0), (2.5, 5.5), (10.5, 3.0), (10.5, 5.5)]
    bulb_center = (Inches(8), Inches(4.5))

    for i, (x, y) in enumerate(positions):
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.5), Inches(0.5))
        circle.fill.background()
        circle.line.color.rgb = RGBColor(255, 255, 255)
        circle.line.width = Pt(0.75)
        tf = circle.text_frame
        p = tf.paragraphs[0]
        p.text = str(i + 1)
        p.font.name = 'Montserrat'; p.font.size = Pt(14); p.font.color.rgb = RGBColor(255, 255, 255); p.alignment = 1

        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 0.5), Inches(y + 0.6), Inches(3), Inches(1.2))
        rect.adjustments[0] = 0.2
        rect.fill.background()
        rect.line.color.rgb = RGBColor(255, 255, 255)
        rect.line.width = Pt(0.75)
        add_inner_shadow(rect, blur_radius=Pt(5), transparency=0.5)

        tf_rect = rect.text_frame
        p_title = tf_rect.paragraphs[0]; p_title.text = item_titles[i]; p_title.font.name = 'Montserrat SemiBold'
        p_title.font.size = Pt(20); p_title.font.color.rgb = RGBColor(255, 255, 255)
        p_body = tf_rect.add_paragraph(); p_body.text = item_texts[i]; p_body.font.name = 'Montserrat Light'
        p_body.font.size = Pt(9); p_body.font.color.rgb = RGBColor(255, 255, 255)

        line_shape = slide.shapes.add_connector(1, Inches(x + 0.25), Inches(y + 0.25), bulb_center[0], bulb_center[1])
        line_shape.line.color.rgb = RGBColor(255, 255, 255)
        line_shape.line.width = Pt(1)

    # === Layer 3: Central Illustration (Lightbulb) ===
    bulb_group = slide.shapes.add_group_shape()
    bulb_group.name = "!!main_illustration" # Name for Morph
    bulb_shape = bulb_group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.75), Inches(3), Inches(2.5), Inches(2.5))
    bulb_shape.fill.background()
    bulb_shape.line.color.rgb = RGBColor(255, 255, 255)
    bulb_shape.line.width = Pt(2)
    add_inner_shadow(bulb_shape, blur_radius=Pt(35), transparency=0.2)
    base_shape = bulb_group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.25), Inches(5.4), Inches(1.5), Inches(1))
    base_shape.fill.solid(); base_shape.fill.fore_color.rgb = RGBColor(255, 255, 255); base_shape.line.fill.background()
    filament = bulb_group.shapes.add_shape(MSO_SHAPE.ARC, Inches(7.25), Inches(4), Inches(1.5), Inches(1))
    filament.rotation = 90; filament.adjustments[0] = 270 * 60000; filament.adjustments[1] = 0
    filament.line.color.rgb = RGBColor(255, 255, 255); filament.line.width = Pt(4); filament.fill.background()

    # === Layer 4: Static Navigation Menu Placeholder ===
    # The interactive functionality (triggers, hyperlinks) must be added manually.
    menu_icon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(0.3), Inches(0.6), Inches(0.6))
    menu_icon.adjustments[0] = 0.5
    menu_icon.fill.background()
    menu_icon.line.color.rgb = RGBColor(255, 255, 255); menu_icon.line.width = Pt(1)
    add_inner_shadow(menu_icon, blur_radius=Pt(5), transparency=0.5)
    for i in range(3):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.45 + i*0.12), Inches(0.3), Inches(0.05))
        line.fill.solid(); line.fill.fore_color.rgb = RGBColor(255, 255, 255); line.line.fill.background()


    prs.save(output_pptx_path)
    return output_pptx_path

