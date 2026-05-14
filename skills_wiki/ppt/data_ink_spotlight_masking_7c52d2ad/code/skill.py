def create_slide(
    output_pptx_path: str,
    title_text: str = "森林保护监测覆盖\nFOREST PROTECTION MONITORING COVERAGE",
    body_text: str = (
        "对森林资源进行定点观测、监测和评估，以了解森林资源的数量、质量、结"
        "构和变化情况，为制定森林资源保护政策和管理措施提供科学依据。\n\n"
        "国家林业和草原局负责森林、草原、湿地动态监测工作，构建林草、湿地"
        "一网通监测体系，依法开展林草湿调查监测工作，着力推进国家和地"
        "方一体化调查监测。"
    ),
    bg_theme: str = "forest,river,aerial",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Data-Ink Spotlight Masking' effect.
    Returns: path to the saved PPTX file.
    """
    import os
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.chart.data import ChartData
    from pptx.oxml.ns import qn
    from pptx.oxml import parse_xml
    from PIL import Image, ImageDraw

    # === Initialize Presentation ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Dimensions in pixels (assuming 96 DPI)
    dpi = 96
    w_px, h_px = int(13.333 * dpi), int(7.5 * dpi)

    # === Layer 1: Fetch and Insert Background Image ===
    try:
        url = f"https://source.unsplash.com/random/1280x720/?{bg_theme}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            bg_image_bytes = response.read()
        bg_stream = io.BytesIO(bg_image_bytes)
    except Exception:
        # Fallback: Create a lush green gradient/solid image
        fallback_img = Image.new('RGB', (w_px, h_px), (34, 139, 34))
        bg_stream = io.BytesIO()
        fallback_img.save(bg_stream, format='JPEG')
        bg_stream.seek(0)

    slide.shapes.add_picture(bg_stream, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Create Cutout Mask using PIL ===
    # Dark forest overlay: R=15, G=45, B=35, Alpha=215 (~85% opacity)
    overlay = Image.new('RGBA', (w_px, h_px), (15, 45, 35, 215))
    
    # Calculate Hole geometry
    # Place on the right side: center X at 72%, center Y at 50%
    hole_cx_px = int(w_px * 0.72)
    hole_cy_px = int(h_px * 0.50)
    hole_r_px = int(h_px * 0.28) # radius is 28% of height
    
    # Create an alpha mask to punch the hole
    alpha_mask = Image.new('L', (w_px, h_px), 215)
    alpha_draw = ImageDraw.Draw(alpha_mask)
    alpha_draw.ellipse(
        (hole_cx_px - hole_r_px, hole_cy_px - hole_r_px, hole_cx_px + hole_r_px, hole_cy_px + hole_r_px), 
        fill=0  # 0 alpha = fully transparent hole
    )
    overlay.putalpha(alpha_mask)

    # Save mask to stream and insert into slide
    mask_stream = io.BytesIO()
    overlay.save(mask_stream, format='PNG')
    mask_stream.seek(0)
    slide.shapes.add_picture(mask_stream, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 3: Dashed Decorative Ring ===
    # Convert px to Inches for pptx
    hole_cx_in = hole_cx_px / dpi
    hole_cy_in = hole_cy_px / dpi
    hole_r_in = hole_r_px / dpi

    # Ring is slightly smaller/larger to frame it
    ring_r_in = hole_r_in * 1.05
    ring = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(hole_cx_in - ring_r_in), Inches(hole_cy_in - ring_r_in),
        Inches(ring_r_in * 2), Inches(ring_r_in * 2)
    )
    ring.fill.background() # Make shape transparent
    ring.line.color.rgb = RGBColor(255, 255, 255)
    ring.line.width = Pt(1)
    ring.line.dash_style = 4 # Dashed line

    # === Layer 4: Integrated Donut Chart ===
    # A standard donut chart has a hole size of ~50% of its bounding box.
    # To make the chart's inner hole perfectly match our PIL cutout, 
    # the chart bounding box radius must be roughly 2 * hole_r_in.
    chart_r_in = hole_r_in * 2.1 # Slight offset to look balanced
    
    chart_data = ChartData()
    chart_data.categories = ['Forest', 'Wetland', 'Grassland', 'Other']
    chart_data.add_series('Coverage', (0.40, 0.30, 0.15, 0.15))

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(hole_cx_in - chart_r_in), Inches(hole_cy_in - chart_r_in),
        Inches(chart_r_in * 2), Inches(chart_r_in * 2),
        chart_data
    )
    chart = chart_shape.chart
    chart.has_legend = False

    # Force Chart Background to be Transparent via lxml
    # Sometimes PPTX adds a white background to charts. This overrides it.
    chartSpace = chart._element
    spPr = chartSpace.find(qn('c:spPr'))
    if spPr is None:
        spPr = parse_xml(r'<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/></c:spPr>')
        chartSpace.insert(0, spPr)
    else:
        # remove existing fills, add noFill
        for child in spPr.xpath('.//a:solidFill | .//a:blipFill | .//a:gradFill', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}):
            spPr.remove(child)
        spPr.append(parse_xml(r'<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))

    # Format Chart Segments
    colors = [
        RGBColor(255, 192, 0),    # Yellow
        RGBColor(237, 125, 49),   # Orange
        RGBColor(112, 173, 71),   # Green
        RGBColor(91, 155, 213)    # Blue
    ]
    series = chart.series[0]
    for idx, point in enumerate(series.points):
        fill = point.format.fill
        fill.solid()
        fill.fore_color.rgb = colors[idx]
        
        # Add data labels
        point.data_label.has_text_frame = True
        point.data_label.text_frame.text = f"{int(chart_data.series[0].values[idx]*100)}%"
        point.data_label.font.color.rgb = RGBColor(255, 255, 255)
        point.data_label.font.bold = True
        point.data_label.font.size = Pt(14)

    # === Layer 5: Typography ===
    # Left Content Box
    tb_left = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.5), Inches(5.0))
    tf = tb_left.text_frame
    tf.word_wrap = True
    
    # Title
    p_title = tf.paragraphs[0]
    p_title.text = title_text.split('\n')[0]
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 215, 0) # Gold
    
    # Subtitle
    if '\n' in title_text:
        p_sub = tf.add_paragraph()
        p_sub.text = title_text.split('\n')[1]
        p_sub.font.size = Pt(14)
        p_sub.font.bold = True
        p_sub.font.color.rgb = RGBColor(169, 208, 142) # Light Sage Green
    
    # Body text
    p_body = tf.add_paragraph()
    p_body.text = "\n" + body_text
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = RGBColor(230, 230, 230)
    p_body.line_spacing = 1.5

    # Center Text inside the cutout
    tb_center = slide.shapes.add_textbox(Inches(hole_cx_in - 1.5), Inches(hole_cy_in - 0.5), Inches(3.0), Inches(1.0))
    p_center = tb_center.text_frame.paragraphs[0]
    p_center.text = "监测覆盖"
    p_center.font.size = Pt(28)
    p_center.font.bold = True
    p_center.font.color.rgb = RGBColor(255, 255, 255)
    # Adding a slight text shadow for visibility against image
    # Note: Using lxml to add shadow
    try:
        defPr = p_center._element.get_or_add_pPr().get_or_add_defRPr()
        shadow_xml = r'<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:outerShdw blurRad="50800" dist="38100" dir="2700000" algn="tl"><a:srgbClr val="000000"><a:alpha val="60000"/></a:srgbClr></a:outerShdw></a:effectLst>'
        defPr.append(parse_xml(shadow_xml))
    except Exception:
        pass # Graceful fail if shadow manipulation fails

    prs.save(output_pptx_path)
    return output_pptx_path
