def create_slide(
    output_pptx_path: str,
    title_text: str = "PERFORMANCE DASHBOARD",
    accent_cyan: tuple = (0, 255, 255),
    accent_magenta: tuple = (255, 20, 147),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Dark Mode Neon Dashboard Layout.
    Uses matplotlib to render glowing, minimalist data visualizations on the fly.
    """
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    import matplotlib.pyplot as plt
    import numpy as np

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Color Palette
    bg_color = RGBColor(13, 17, 28)
    card_bg_color = RGBColor(25, 32, 48)
    border_color = RGBColor(70, 80, 100)
    text_primary = RGBColor(255, 255, 255)
    text_secondary = RGBColor(160, 170, 190)
    
    cyan_rgb = RGBColor(*accent_cyan)
    cyan_hex = '#%02x%02x%02x' % accent_cyan
    magenta_hex = '#%02x%02x%02x' % accent_magenta

    # === Layer 1: Background ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg_color
    bg_shape.line.fill.background() # No line

    # === Layer 2: Dashboard Header ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.name = "Arial"
    p.font.color.rgb = text_primary

    # Dashboard sub-line
    p2 = tf.add_paragraph()
    p2.text = "Q3 METRICS & DATA VISUALIZATION"
    p2.font.size = Pt(12)
    p2.font.name = "Arial"
    p2.font.color.rgb = cyan_rgb
    p2.font.bold = True

    # === Helper Function: Create a Data Card ===
    def add_card(x, y, w, h):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = card_bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # === Layer 3: Top KPI Row (3 blocks) ===
    # Demonstrates Balance & Hierarchy
    kpi_y = 1.4
    kpi_w = 3.84
    kpi_h = 1.2
    gap = 0.4
    
    kpis = [
        {"label": "TOTAL REVENUE", "value": "$1.24M", "trend": "+14% vs Last Month"},
        {"label": "ACTIVE USERS", "value": "84,592", "trend": "+5% vs Last Month"},
        {"label": "DEFECT RATE", "value": "1.2%", "trend": "-2% vs Last Month"}
    ]

    for i, kpi in enumerate(kpis):
        x_pos = 0.5 + i * (kpi_w + gap)
        add_card(x_pos, kpi_y, kpi_w, kpi_h)
        
        # KPI Label
        tx = slide.shapes.add_textbox(Inches(x_pos + 0.1), Inches(kpi_y + 0.1), Inches(kpi_w-0.2), Inches(0.3))
        p = tx.text_frame.paragraphs[0]
        p.text = kpi["label"]
        p.font.size = Pt(11)
        p.font.color.rgb = text_secondary
        p.alignment = PP_ALIGN.CENTER
        
        # KPI Value
        tx2 = slide.shapes.add_textbox(Inches(x_pos + 0.1), Inches(kpi_y + 0.35), Inches(kpi_w-0.2), Inches(0.5))
        p2 = tx2.text_frame.paragraphs[0]
        p2.text = kpi["value"]
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = cyan_rgb if i < 2 else RGBColor(*accent_magenta)
        p2.alignment = PP_ALIGN.CENTER
        
        # KPI Trend
        tx3 = slide.shapes.add_textbox(Inches(x_pos + 0.1), Inches(kpi_y + 0.85), Inches(kpi_w-0.2), Inches(0.3))
        p3 = tx3.text_frame.paragraphs[0]
        p3.text = kpi["trend"]
        p3.font.size = Pt(9)
        p3.font.color.rgb = text_secondary
        p3.alignment = PP_ALIGN.CENTER

    # === Layer 4: Chart Row (2 blocks) ===
    chart_y = 3.0
    chart_h = 4.0
    chart_w = 5.96
    
    # Left Block: Line Chart (Trends over time)
    add_card(0.5, chart_y, chart_w, chart_h)
    
    # Create Matplotlib Glowing Line Chart
    fig1, ax1 = plt.subplots(figsize=(6, 3.5), dpi=150)
    fig1.patch.set_alpha(0.0) # Transparent bg
    ax1.set_facecolor('none')
    
    x_data = np.arange(10)
    y_data = np.cumsum(np.random.randn(10)) + 10
    
    # Core line
    ax1.plot(x_data, y_data, color=cyan_hex, linewidth=2.5)
    # Glow effect
    for n in range(1, 4):
        ax1.plot(x_data, y_data, color=cyan_hex, linewidth=2.5+(n*2.5), alpha=0.15)
        
    ax1.spines['top'].set_visible(False)
    ax1.spines['right'].set_visible(False)
    ax1.spines['bottom'].set_color('#465064')
    ax1.spines['left'].set_color('#465064')
    ax1.tick_params(colors='#A0AABE')
    
    # Save to memory
    img_stream1 = io.BytesIO()
    plt.savefig(img_stream1, format='png', bbox_inches='tight', transparent=True)
    plt.close(fig1)
    img_stream1.seek(0)
    
    # Insert chart image
    slide.shapes.add_picture(img_stream1, Inches(0.7), Inches(3.2), width=Inches(5.5))
    
    # Add title to left chart
    cx1 = slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(3), Inches(0.4))
    cp1 = cx1.text_frame.paragraphs[0]
    cp1.text = "REVENUE TRAJECTORY"
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = text_primary


    # Right Block: Bar Chart (Categorical Comparison)
    add_card(6.86, chart_y, chart_w, chart_h)
    
    # Create Matplotlib Categorical Bar Chart
    fig2, ax2 = plt.subplots(figsize=(6, 3.5), dpi=150)
    fig2.patch.set_alpha(0.0)
    ax2.set_facecolor('none')
    
    cats = ['Prod A', 'Prod B', 'Prod C', 'Prod D', 'Prod E']
    vals = [45, 60, 35, 75, 50]
    colors = [cyan_hex, magenta_hex, '#9400D3', cyan_hex, magenta_hex]
    
    bars = ax2.bar(cats, vals, color=colors, width=0.6)
    
    # Clean up axes
    ax2.spines['top'].set_visible(False)
    ax2.spines['right'].set_visible(False)
    ax2.spines['left'].set_visible(False) # No left spine for clean look
    ax2.spines['bottom'].set_color('#465064')
    ax2.tick_params(colors='#A0AABE', left=False)
    ax2.set_yticks([]) # Remove y ticks, use direct labels
    
    # Add value labels on top of bars
    for bar in bars:
        height = bar.get_height()
        ax2.annotate(f'{height}',
                    xy=(bar.get_x() + bar.get_width() / 2, height),
                    xytext=(0, 3),  # 3 points vertical offset
                    textcoords="offset points",
                    ha='center', va='bottom', color='white', fontweight='bold')
    
    # Save to memory
    img_stream2 = io.BytesIO()
    plt.savefig(img_stream2, format='png', bbox_inches='tight', transparent=True)
    plt.close(fig2)
    img_stream2.seek(0)
    
    # Insert chart image
    slide.shapes.add_picture(img_stream2, Inches(7.06), Inches(3.2), width=Inches(5.5))
    
    # Add title to right chart
    cx2 = slide.shapes.add_textbox(Inches(6.96), Inches(3.1), Inches(3), Inches(0.4))
    cp2 = cx2.text_frame.paragraphs[0]
    cp2.text = "PRODUCT DISTRIBUTION"
    cp2.font.size = Pt(12)
    cp2.font.bold = True
    cp2.font.color.rgb = text_primary

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
