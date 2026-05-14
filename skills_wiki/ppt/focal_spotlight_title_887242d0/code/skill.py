import requests
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "Alignment Principle",
    subtitle_text: str = "FOUR PRINCIPLES OF TYPOGRAPHY",
    bg_keyword: str = "forest,aerial",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a 'Focal Spotlight Title' design.

    This effect features a full-bleed background image with a semi-transparent,
    soft-edged dark circle in the center, containing the title and subtitle.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        title_text: The main text to display in the spotlight.
        subtitle_text: The smaller text to display below the main title.
        bg_keyword: A comma-separated keyword string for fetching a background
                    image from Unsplash (e.g., 'technology', 'nature,abstract').

    Returns:
        The path to the saved PPTX file.
    """
    SLIDE_WIDTH_PX = 1920
    SLIDE_HEIGHT_PX = 1080

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- 1. Create the Background with PIL ---
    try:
        # Download a background image from Unsplash
        image_url = f"https://source.unsplash.com/{SLIDE_WIDTH_PX}x{SLIDE_HEIGHT_PX}/?{bg_keyword}"
        response = requests.get(image_url, timeout=10)
        response.raise_for_status()
        bg_img = Image.open(io.BytesIO(response.content)).convert("RGBA")
    except requests.exceptions.RequestException:
        # Fallback to a solid dark background if image download fails
        bg_img = Image.new("RGBA", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), (30, 30, 30, 255))

    # --- 2. Create the Soft-Edged Spotlight Overlay with PIL ---
    # Create a new transparent layer for the spotlight
    overlay = Image.new("RGBA", bg_img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)

    # Define circle properties
    circle_radius = SLIDE_HEIGHT_PX * 0.28  # Approx 56% of slide height
    circle_center = (SLIDE_WIDTH_PX / 2, SLIDE_HEIGHT_PX / 2)
    circle_box = [
        circle_center[0] - circle_radius,
        circle_center[1] - circle_radius,
        circle_center[0] + circle_radius,
        circle_center[1] + circle_radius,
    ]
    
    # Draw a black, semi-transparent circle onto the overlay layer
    spotlight_alpha = 180  # Control transparency (0-255)
    draw.ellipse(circle_box, fill=(0, 0, 0, spotlight_alpha))

    # Apply a Gaussian blur to create the soft edge
    blur_radius = 50
    overlay = overlay.filter(ImageFilter.GaussianBlur(radius=blur_radius))

    # Composite the spotlight overlay onto the background image
    combined_img = Image.alpha_composite(bg_img, overlay)

    # --- 3. Add the Composited Image to the Slide ---
    with io.BytesIO() as output:
        combined_img.save(output, format="PNG")
        slide.shapes.add_picture(output, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    # --- 4. Add and Format Text ---
    # Main Title
    title_shape = slide.shapes.add_textbox(
        Inches(2), Inches(3), Inches(9.333), Inches(1)
    )
    title_tf = title_shape.text_frame
    title_tf.text = title_text
    title_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    title_font = title_tf.paragraphs[0].font
    title_font.name = 'Arial Black'
    title_font.size = Pt(44)
    title_font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle
    subtitle_shape = slide.shapes.add_textbox(
        Inches(2), Inches(3.9), Inches(9.333), Inches(0.6)
    )
    subtitle_tf = subtitle_shape.text_frame
    subtitle_tf.text = subtitle_text
    subtitle_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    subtitle_font = subtitle_tf.paragraphs[0].font
    subtitle_font.name = 'Arial'
    subtitle_font.size = Pt(16)
    subtitle_font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# if __name__ == '__main__':
#     create_slide(
#         output_pptx_path="focal_spotlight_title.pptx",
#         title_text="对 比 原 则",
#         subtitle_text="FOUR PRINCIPLES OF TYPOGRAPHY",
#         bg_keyword="mountain,snow"
#     )
#     create_slide(
#         output_pptx_path="focal_spotlight_title_2.pptx",
#         title_text="靠 近 原 则",
#         subtitle_text="FOUR PRINCIPLES OF TYPOGRAPHY",
#         bg_keyword="ocean,wave"
#     )

