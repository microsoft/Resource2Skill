def create_slide(
    output_pptx_path: str,
    title_text: str = "Every tool has an opinion about how you should work.\nThis one doesn't.",
    highlight_word: str = "doesn't.",
    overline_text: str = "PRODUCT PHILOSOPHY",
    body_text: str = "A workspace that adapts to you — not the other way around.",
    footer_text: str = "CHARMIQ — 2026",
    accent_color: tuple = (220, 38, 38),  # Vibrant Swiss Red
    bg_color: tuple = (250, 250, 250),    # Off-white
    text_color: tuple = (20, 20, 20)      # Charcoal black
) -> str:
    """
    Create a PPTX file reproducing the "Swiss International Typographic" style.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import re

    # Initialize presentation (16:9 aspect ratio)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Graphic Elements (The Anchor Bar) ===
    # A crisp, vertical geometric bar that grounds the typography
    bar_left = Inches(1.0)
    bar_top = Inches(1.5)
    bar_width = Inches(0.08)
    bar_height = Inches(4.5)
    
    accent_bar = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        bar_left, bar_top, bar_width, bar_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = RGBColor(*accent_color)
    accent_bar.line.fill.background()  # No outline

    # === Layer 3: Typography ===
    
    # Base positioning for text blocks (aligned exactly right of the accent bar)
    text_left = Inches(1.3)
    text_width = Inches(9.0)

    # 1. Overline (Category Label)
    overline_box = slide.shapes.add_textbox(text_left, Inches(1.4), text_width, Inches(0.5))
    overline_frame = overline_box.text_frame
    overline_frame.word_wrap = True
    p_over = overline_frame.paragraphs[0]
    p_over.text = overline_text.upper()
    p_over.font.name = "Arial"  # Fallback sans-serif mimicking Helvetica
    p_over.font.size = Pt(11)
    p_over.font.bold = True
    p_over.font.color.rgb = RGBColor(100, 100, 100)  # Mid-grey

    # 2. Headline (Massive, tightly spaced, with inline highlight)
    headline_box = slide.shapes.add_textbox(text_left, Inches(1.9), text_width, Inches(3.0))
    headline_frame = headline_box.text_frame
    headline_frame.word_wrap = True
    p_head = headline_frame.paragraphs[0]
    p_head.line_spacing = 0.95  # Tight line spacing for editorial feel
    
    # Logic to split text and colorize the highlighted word
    if highlight_word in title_text:
        # Split text but keep the highlight word intact
        parts = title_text.split(highlight_word)
        for i, part in enumerate(parts):
            # Add normal text
            if part:
                run = p_head.add_run()
                run.text = part
                run.font.name = "Arial"
                run.font.size = Pt(54)
                run.font.bold = True
                run.font.color.rgb = RGBColor(*text_color)
            
            # Add highlight word (between parts)
            if i < len(parts) - 1:
                run_highlight = p_head.add_run()
                run_highlight.text = highlight_word
                run_highlight.font.name = "Arial"
                run_highlight.font.size = Pt(54)
                run_highlight.font.bold = True
                run_highlight.font.color.rgb = RGBColor(*accent_color)
    else:
        # Fallback if highlight word not found
        run = p_head.add_run()
        run.text = title_text
        run.font.name = "Arial"
        run.font.size = Pt(54)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*text_color)

    # 3. Body Text (Minimalist supportive copy)
    body_box = slide.shapes.add_textbox(text_left, Inches(5.0), text_width, Inches(1.0))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    p_body = body_frame.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(80, 80, 80)

    # 4. Footer (Small alignment text)
    footer_box = slide.shapes.add_textbox(text_left, Inches(6.8), Inches(4.0), Inches(0.5))
    footer_frame = footer_box.text_frame
    p_foot = footer_frame.paragraphs[0]
    p_foot.text = footer_text
    p_foot.font.name = "Arial"
    p_foot.font.size = Pt(10)
    p_foot.font.bold = True
    p_foot.font.color.rgb = RGBColor(150, 150, 150)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
