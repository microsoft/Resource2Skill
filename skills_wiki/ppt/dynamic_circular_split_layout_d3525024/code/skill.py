def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda for Product\nLaunch Event",
    agenda_items: list = None,
    bg_palette: str = "corporate,meeting",
    accent_color: tuple = (0, 210, 181),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Dynamic Circular Split Layout.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml
    from PIL import Image

    if agenda_items is None:
        agenda_items = [
            "Engage key stakeholders in launch event",
            "Target audience to create event awareness",
            "Create strong social media presence",
            "Review execution budget and expected ROI"
        ]

    # Initialize presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Helper function for XML Gradient Injection
    def apply_gradient_fill(shape, hex_color1, hex_color2, angle=0):
        spPr = shape.element.spPr
        for fill_tag in ['.//a:solidFill', './/a:gradFill', './/a:pattFill', './/a:blipFill', './/a:noFill']:
            fills = spPr.xpath(fill_tag, namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
            for fill in fills:
                spPr.remove(fill)
        
        pptx_angle = int(angle * 60000)
        grad_xml = f"""
        <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">
            <a:gsLst>
                <a:gs pos="0"><a:srgbClr val="{hex_color1}"/></a:gs>
                <a:gs pos="100000"><a:srgbClr val="{hex_color2}"/></a:gs>
            </a:gsLst>
            <a:lin ang="{pptx_angle}" scaled="1"/>
        </a:gradFill>
        """
        spPr.append(parse_xml(grad_xml))

    # === Layer 1: Right-side Gradient Framing Curve ===
    # Draw a massive oval that acts as the colored border offset
    grad_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.0), Inches(0.0), Inches(7.5), Inches(7.5))
    grad_circle.line.fill.background()  # Remove outline
    apply_gradient_fill(grad_circle, "00D2B5", "0066CC", angle=60) # Teal to Blue

    # === Layer 2: Circular Masked Image ===
    img_path = "temp_agenda_image.jpg"
    img_url = f"https://images.unsplash.com/photo-1552664730-d307ca884978?q=80&w=800&auto=format&fit=crop"
    
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Fallback if download fails: create a subtle grey placeholder
        img = Image.new('RGB', (800, 800), color=(230, 235, 240))
        img.save(img_path)

    # Insert image slightly offset to the right of the gradient circle to create the crescent
    pic = slide.shapes.add_picture(img_path, Inches(6.5), Inches(0.5), Inches(6.5), Inches(6.5))
    
    # XML Injection: Force the rectangular picture bounding box to be an ellipse (circle)
    spPr = pic.element.spPr
    prstGeom = spPr.find('.//a:prstGeom', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if prstGeom is not None:
        prstGeom.set('prst', 'ellipse')

    # === Layer 3: Text Content & Formatting (Left side) ===
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(5.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Georgia"  # Elegant serif
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 40, 50)
    
    # Decorative line under title
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.6), Pt(4))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # Agenda List Items
    # Dynamic palette blending from Teal to Blue
    list_colors = [
        (0, 210, 181),   # Teal
        (0, 170, 190),
        (0, 130, 200),
        (0, 102, 204)    # Blue
    ]
    
    start_y = 3.0
    spacing = 0.95
    
    for i, text in enumerate(agenda_items):
        y_pos = start_y + (i * spacing)
        color = list_colors[i % len(list_colors)]
        
        # Pill shape for number
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pos), Inches(0.6), Inches(0.35))
        pill.adjustments[0] = 0.5  # Maximize roundness to create a pill shape
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(*color)
        pill.line.fill.background()
        
        pill.text = f"{i+1:02d}"
        pill_p = pill.text_frame.paragraphs[0]
        pill_p.font.size = Pt(12)
        pill_p.font.bold = True
        pill_p.font.color.rgb = RGBColor(255, 255, 255)
        pill_p.alignment = PP_ALIGN.CENTER
        
        # List text
        tb = slide.shapes.add_textbox(Inches(1.6), Inches(y_pos - 0.05), Inches(4.2), Inches(0.5))
        tp = tb.text_frame.paragraphs[0]
        tp.text = text
        tp.font.name = "Calibri"
        tp.font.size = Pt(14)
        tp.font.color.rgb = RGBColor(100, 100, 100)
        
        # Subtle separator line
        sep_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.6), Inches(y_pos + 0.45), Inches(4.0), Pt(1))
        sep_line.fill.solid()
        sep_line.fill.fore_color.rgb = RGBColor(230, 230, 230)
        sep_line.line.fill.background()

    # Cleanup temporary image
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
