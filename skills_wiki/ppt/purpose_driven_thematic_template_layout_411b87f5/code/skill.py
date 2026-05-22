def create_slide(
    output_pptx_path: str,
    title_text: str = "MARKETING PLAN",
    body_text: str = "Here is where your presentation begins",
    bg_palette: str = "eco",  # Options: "eco" or "corporate"
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Purpose-Driven Thematic Template Layout'.
    It generates a 3-column KPI dashboard showing how Themes (colors) 
    apply to Templates (structured layouts).
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import ChartData
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # === 1. Establish the "Theme" (Color Palette) ===
    if bg_palette == "eco":
        c_bg = RGBColor(249, 246, 238)       # Light beige
        c_primary = RGBColor(85, 107, 47)    # Dark olive green
        c_accent1 = RGBColor(143, 188, 143)  # Dark sea green
        c_accent2 = RGBColor(210, 180, 140)  # Tan
        c_text = RGBColor(40, 40, 40)        # Dark Charcoal
    else:
        # Corporate Blue theme fallback
        c_bg = RGBColor(240, 244, 248)       
        c_primary = RGBColor(16, 42, 67)     
        c_accent1 = RGBColor(98, 125, 152)   
        c_accent2 = RGBColor(217, 226, 236)  
        c_text = RGBColor(16, 42, 67)

    # === 2. Build the Foundation (Background) ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = c_bg
    bg_shape.line.fill.background()
    
    # === 3. Construct the Template Header ===
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(1.4), Inches(11.333), Inches(0.5))
    tf_sub = subtitle_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = c_text
    p_sub.font.name = "Arial"
    p_sub.alignment = PP_ALIGN.CENTER
    
    # Divider Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.166), Inches(2.1), Inches(1), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = c_accent1
    line.line.fill.background()
    
    # === 4. Generate the Template Layout (3 Columns of Data) ===
    metrics = [
        {"val": 25, "title": "Social Media", "desc": "Engagement via platforms experiencing targeted ad locations."},
        {"val": 50, "title": "Community", "desc": "Community engagement before, during, and after events."},
        {"val": 25, "title": "Incentives", "desc": "Discounts, giveaways, and gamified reward systems."}
    ]
    
    col_width = Inches(3.2)
    start_x = Inches(1.2)
    spacing = Inches(0.76)
    
    for i, item in enumerate(metrics):
        x = start_x + (i * (col_width + spacing))
        y_chart = Inches(2.7)
        
        # A. Native Doughnut Chart Placeholder
        chart_data = ChartData()
        chart_data.categories = ['Achieved', 'Remaining']
        chart_data.add_series('Data', (item["val"], 100 - item["val"]))
        
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, x, y_chart, col_width, col_width, chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = False
        
        # Remove background of chart area to blend with slide background
        chart_shape.fill.solid()
        chart_shape.fill.fore_color.rgb = c_bg
        chart_shape.line.fill.background()
        
        # Apply theme colors to chart points
        series = chart.series[0]
        pts = series.points
        
        # Highlight color for the achieved metric
        pts[0].format.fill.solid()
        pts[0].format.fill.fore_color.rgb = c_primary if i % 2 == 0 else c_accent1
        
        # Muted color for the remainder
        pts[1].format.fill.solid()
        pts[1].format.fill.fore_color.rgb = c_accent2
        
        # B. Center Metric Text overlay
        val_box = slide.shapes.add_textbox(x, y_chart + (col_width/2) - Inches(0.25), col_width, Inches(0.5))
        tf_val = val_box.text_frame
        p_val = tf_val.paragraphs[0]
        p_val.text = f"{item['val']}%"
        p_val.font.size = Pt(28)
        p_val.font.bold = True
        p_val.font.color.rgb = c_primary
        p_val.font.name = "Arial"
        p_val.alignment = PP_ALIGN.CENTER
        
        # C. Column Title
        y_text = y_chart + col_width + Inches(0.1)
        h_box = slide.shapes.add_textbox(x, y_text, col_width, Inches(0.4))
        tf_h = h_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = item["title"].upper()
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = c_primary
        p_h.font.name = "Arial"
        p_h.alignment = PP_ALIGN.CENTER
        
        # D. Column Description
        desc_box = slide.shapes.add_textbox(x, y_text + Inches(0.4), col_width, Inches(1.5))
        tf_d = desc_box.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = item["desc"]
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = c_text
        p_d.font.name = "Arial"
        p_d.alignment = PP_ALIGN.CENTER
        
    prs.save(output_pptx_path)
    return output_pptx_path
