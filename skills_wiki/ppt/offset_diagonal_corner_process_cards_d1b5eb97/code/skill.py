def create_slide(
    output_pptx_path: str,
    title_text: str = "Process Overview",
    body_text: str = "",
    bg_palette: str = "dark",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Offset Diagonal-Corner Process Cards' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # === Layer 1: Dark Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(15, 15, 18)  # Off-black for modern tech feel

    # Define color palettes for the 3 steps (Front Neon, Back Shadow)
    card_colors = [
        {"front": RGBColor(0, 255, 128), "back": RGBColor(0, 100, 50)},    # Green
        {"front": RGBColor(50, 100, 255), "back": RGBColor(20, 40, 120)},  # Blue
        {"front": RGBColor(255, 50, 50),  "back": RGBColor(120, 20, 20)}   # Red
    ]

    # Content data
    steps_data = [
        {"num": "01.", "title": "First Step Title", "desc": "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam."},
        {"num": "02.", "title": "Second Step Title", "desc": "Quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore."},
        {"num": "03.", "title": "Third Step Title", "desc": "Excepteur sint occaecat cupidatat non proident, sunt in culpa qui officia deserunt mollit anim id est laborum. Praesent elementum facilisis leo vel."}
    ]

    # Global dimensions for cards
    card_width = Inches(8.5)
    card_height = Inches(1.6)
    start_left = Inches(2.6)
    start_top = Inches(1.0)
    vertical_spacing = Inches(2.0)
    
    # Offset for the back "shadow" shape
    offset_x = Inches(-0.15)
    offset_y = Inches(0.15)

    # === Layer 2 & 3: Shapes and Text ===
    for index, (colors, data) in enumerate(zip(card_colors, steps_data)):
        current_top = start_top + (index * vertical_spacing)
        
        # 1. Back Shape (Shadow/Depth)
        back_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUND_2_DIAG_RECTANGLE, 
            start_left + offset_x, current_top + offset_y, 
            card_width, card_height
        )
        back_shape.fill.solid()
        back_shape.fill.fore_color.rgb = colors["back"]
        back_shape.line.fill.solid()
        back_shape.line.fill.fore_color.rgb = colors["back"]

        # 2. Front Shape (Main Card)
        front_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUND_2_DIAG_RECTANGLE, 
            start_left, current_top, 
            card_width, card_height
        )
        front_shape.fill.solid()
        front_shape.fill.fore_color.rgb = colors["front"]
        front_shape.line.fill.solid()
        front_shape.line.fill.fore_color.rgb = colors["front"]

        # 3. Number Text Box (Left side of card)
        num_box = slide.shapes.add_textbox(
            start_left + Inches(0.3), current_top + Inches(0.25), 
            Inches(1.5), Inches(1.0)
        )
        num_frame = num_box.text_frame
        num_frame.clear() # clear default paragraph
        p_num = num_frame.paragraphs[0]
        p_num.text = data["num"]
        p_num.font.size = Pt(48)
        p_num.font.bold = True
        p_num.font.name = "Arial Black"
        p_num.font.color.rgb = RGBColor(255, 255, 255)
        p_num.alignment = PP_ALIGN.LEFT

        # 4. Title Text Box (Right side of card)
        title_box = slide.shapes.add_textbox(
            start_left + Inches(2.2), current_top + Inches(0.1), 
            card_width - Inches(2.5), Inches(0.5)
        )
        title_frame = title_box.text_frame
        p_title = title_frame.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.name = "Arial"
        p_title.font.color.rgb = RGBColor(20, 20, 20)  # Dark text for high contrast

        # 5. Body Text Box (Below Title)
        body_box = slide.shapes.add_textbox(
            start_left + Inches(2.2), current_top + Inches(0.5), 
            card_width - Inches(2.6), Inches(1.0)
        )
        body_frame = body_box.text_frame
        body_frame.word_wrap = True
        p_body = body_frame.paragraphs[0]
        p_body.text = data["desc"]
        p_body.font.size = Pt(12)
        p_body.font.name = "Arial"
        p_body.font.color.rgb = RGBColor(40, 40, 40)
        p_body.line_spacing = 1.2

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
