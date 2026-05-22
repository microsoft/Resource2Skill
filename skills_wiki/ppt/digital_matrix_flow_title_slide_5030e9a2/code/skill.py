def create_slide(
    output_pptx_path: str,
    title_text: str = "Software Architecture",
    subtitle_text: str = "Diagramming Tools & Frameworks",
    primary_glow: tuple = (0, 150, 255),  # Deep Blue
    secondary_glow: tuple = (0, 255, 255), # Cyan
) -> str:
    """
    Creates a presentation with a generated "Digital Matrix Flow" background 
    and crisp, bottom-left aligned tech typography.
    """
    import os
    import random
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFilter

    # --- 1. Generate the Digital Matrix Background via PIL ---
    width, height = 1920, 1080
    bg_color = (8, 10, 15, 255)
    base_img = Image.new('RGBA', (width, height), bg_color)

    # Layer A: Glowing Vertical Streaks
    glow_layer = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    draw_glow = ImageDraw.Draw(glow_layer)
    
    # Draw several thick vertical bands
    bands = [
        (400, primary_glow, 120),
        (700, secondary_glow, 80),
        (1000, primary_glow, 150),
        (1300, secondary_glow, 60),
        (1600, primary_glow, 200)
    ]
    for x_pos, color, thickness in bands:
        r, g, b = color
        draw_glow.line([(x_pos, 0), (x_pos, height)], fill=(r, g, b, 140), width=thickness)
    
    # Apply massive blur to create light streaks
    glow_layer = glow_layer.filter(ImageFilter.GaussianBlur(150))
    base_img = Image.alpha_composite(base_img, glow_layer)

    # Layer B: Rigid Square Matrix
    grid_layer = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    draw_grid = ImageDraw.Draw(grid_layer)
    
    cols, rows = 16, 9
    cell_w = width / cols
    cell_h = height / rows
    max_sq_size = min(cell_w, cell_h) * 0.6
    
    random.seed(42) # Ensure consistent beautiful pattern
    
    for c in range(cols):
        for r in range(rows):
            cx = c * cell_w + cell_w / 2
            cy = r * cell_h + cell_h / 2
            
            # Left to right degradation logic
            progress = c / (cols - 1)  # 0.0 on left, 1.0 on right
            
            # Size decreases as we move right
            size = max_sq_size * (1.0 - (progress * 0.7))
            
            # Alpha decreases as we move right
            alpha = int(255 * (1.0 - progress))
            
            # Color logic: Solid white on left, tinting to cyan/blue on right
            red = int(255 * (1.0 - progress))
            green = int(255 * (1.0 - (progress * 0.2)))
            blue = 255
            
            # Random dropout (more dropouts on the right)
            if random.random() > (0.1 + progress * 0.6):
                left = cx - size / 2
                top = cy - size / 2
                right = cx + size / 2
                bottom = cy + size / 2
                
                # 20% chance to be an outlined square instead of filled for variety
                if random.random() > 0.8:
                    draw_grid.rectangle([left, top, right, bottom], outline=(red, green, blue, alpha), width=3)
                else:
                    draw_grid.rectangle([left, top, right, bottom], fill=(red, green, blue, alpha))

    base_img = Image.alpha_composite(base_img, grid_layer)
    
    # Save background image
    bg_img_path = "temp_digital_matrix_bg.png"
    base_img.convert('RGB').save(bg_img_path)

    # --- 2. Build the PowerPoint Slide ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Add the generated background
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Add Text Elements (Bottom-Left anchored)
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(8), Inches(0.5))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.add_paragraph()
    sub_p.text = subtitle_text.upper()
    sub_p.font.name = "Arial"
    sub_p.font.size = Pt(16)
    sub_p.font.bold = True
    sub_p.font.color.rgb = RGBColor(0, 200, 255) # Cyan

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(5.6), Inches(10), Inches(1.5))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    title_p = title_tf.add_paragraph()
    title_p.text = title_text
    title_p.font.name = "Arial"
    title_p.font.size = Pt(48)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    # Decorative Line above text
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(0.8), Inches(5.1), Inches(0.8), Pt(4)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 200, 255)
    line.line.fill.background()

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    return output_pptx_path
