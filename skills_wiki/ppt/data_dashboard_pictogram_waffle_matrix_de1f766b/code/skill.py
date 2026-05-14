import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def draw_person(slide, x, y, size, color, bg_color):
    """Draws a stylized 'Person' pictogram using native shapes."""
    head_size = int(size * 0.35)
    head_x = int(x + (size - head_size) / 2)
    head_y = int(y + size * 0.1)
    
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, head_x, head_y, head_size, head_size)
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.color.rgb = bg_color
    head.line.width = Pt(0.5)
    
    body_width = int(size * 0.6)
    body_height = int(size * 0.45)
    body_x = int(x + (size - body_width) / 2)
    body_y = int(y + size * 0.5)
    
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, body_x, body_y, body_width, body_height)
    body.fill.solid()
    body.fill.fore_color.rgb = color
    body.line.color.rgb = bg_color
    body.line.width = Pt(0.5)

def draw_building(slide, x, y, size, color, bg_color):
    """Draws a stylized 'Building' pictogram with window cutouts using native shapes."""
    b_width = int(size * 0.55)
    b_height = int(size * 0.8)
    b_x = int(x + (size - b_width) / 2)
    b_y = int(y + size * 0.15)
    
    bldg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, b_x, b_y, b_width, b_height)
    bldg.fill.solid()
    bldg.fill.fore_color.rgb = color
    bldg.line.color.rgb = bg_color
    bldg.line.width = Pt(0.5)
    
    # Add window cutouts
    w_width = int(size * 0.12)
    w_height = int(size * 0.12)
    for r in range(3):
        for c in range(2):
            wx = int(b_x + b_width * 0.2 + c * (b_width * 0.48))
            wy = int(b_y + b_height * 0.2 + r * (b_height * 0.25))
            win = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, wx, wy, w_width, w_height)
            win.fill.solid()
            win.fill.fore_color.rgb = bg_color  # Match background to look like a cutout
            win.line.color.rgb = bg_color
            win.line.width = Pt(0.5)

def draw_waffle_with_legend(slide, x, y, size, data, chart_title, shape_type, bg_color):
    """Renders the 10x10 Waffle Chart and vertically anchored legends."""
    # Chart Title
    title_box = slide.shapes.add_textbox(int(x), int(y - Inches(0.8)), int(size), int(Inches(0.5)))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = chart_title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(220, 220, 220)
    p.alignment = PP_ALIGN.CENTER
    
    # Flatten percentage data into 100 elements array
    elements = []
    for d in data:
        elements.extend([d] * int(d["pct"]))
    while len(elements) < 100:  # Pad any rounding decimals
        elements.append(data[-1])
        
    # Draw Grid
    cell_size = size / 10.0
    for i in range(100):
        # Anchor filling from bottom to top
        row = 9 - (i // 10)
        col = i % 10
        
        d = elements[i]
        left = x + col * cell_size
        top = y + row * cell_size
        
        padding = cell_size * 0.1
        icon_size = cell_size - 2 * padding
        icon_x = left + padding
        icon_y = top + padding
        
        if shape_type == "person":
            draw_person(slide, icon_x, icon_y, icon_size, d["color"], bg_color)
        elif shape_type == "building":
            draw_building(slide, icon_x, icon_y, icon_size, d["color"], bg_color)

    # Draw Data Legends
    for idx, d in enumerate(data):
        # Stagger vertical alignment based on data position (bottom or top)
        if len(data) == 2:
            box_y = int(y + size - Inches(1.3)) if idx == 0 else int(y + Inches(0.3))
        else:
            spacing = size / len(data)
            box_y = int(y + size - (idx + 1) * spacing + spacing / 2 - Inches(0.6))
            
        leg_box = slide.shapes.add_textbox(int(x - Inches(2.2)), int(box_y), int(Inches(2.0)), int(Inches(1.2)))
        tf = leg_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0; tf.margin_right = 0
        
        p1 = tf.paragraphs[0]
        p1.text = d["label"]
        p1.font.size = Pt(18)
        p1.font.color.rgb = d["color"]
        p1.font.name = "Arial"
        p1.alignment = PP_ALIGN.RIGHT
        
        p2 = tf.add_paragraph()
        p2.text = f"{d['pct']}%"
        p2.font.size = Pt(36)
        p2.font.bold = True
        p2.font.color.rgb = d["color"]
        p2.font.name = "Arial"
        p2.alignment = PP_ALIGN.RIGHT


def create_slide(
    output_pptx_path: str,
    title_text: str = "Corporate Analytics Dashboard",
    body_text: str = "Visualizing distribution metrics and key segment data",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Data Dashboard Pictogram Waffle Matrix.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Global Background
    bg_color = RGBColor(20, 24, 34)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # Slide Header
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(140, 155, 170)
    
    # Header Separator Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(Inches(0.5)), int(Inches(1.5)), int(Inches(12.333)), int(Inches(0.02)))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(50, 60, 80)
    line.line.fill.background()

    # Define Data
    # Note: data[0] is drawn first and fills the bottom of the grid. 
    data_1 = [
        {"label": "Remote Workers", "pct": 58, "color": RGBColor(45, 225, 195)},
        {"label": "In-Office", "pct": 42, "color": RGBColor(55, 65, 80)},
    ]
    
    data_2 = [
        {"label": "Enterprise", "pct": 73, "color": RGBColor(255, 130, 70)},
        {"label": "SMB & Startup", "pct": 27, "color": RGBColor(55, 65, 80)},
    ]

    # Render Charts
    # Chart 1 uses custom "Person" pictograms
    draw_waffle_with_legend(slide, Inches(2.8), Inches(2.5), Inches(4.0), data_1, "Workforce Distribution", "person", bg_color)
    
    # Chart 2 uses custom "Building" pictograms
    draw_waffle_with_legend(slide, Inches(8.8), Inches(2.5), Inches(4.0), data_2, "Revenue by Segment", "building", bg_color)

    prs.save(output_pptx_path)
    return output_pptx_path
