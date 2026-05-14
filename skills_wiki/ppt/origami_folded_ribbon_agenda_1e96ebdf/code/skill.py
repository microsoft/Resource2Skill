import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

def create_slide(
    output_pptx_path: str,
    title_text: str = "AGENDA SLIDE",
    subtitle_text: str = "Not everyone falls into\nsuccess with their first try.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Origami Folded Ribbon Agenda' visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(252, 252, 254)

    # === Layer 2: Left Column (Title & Subtitle) ===
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(3.0), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(44, 62, 80) # Dark Navy Gray
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(3.0), Inches(1.0))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(16)
    p_sub.font.name = "Calibri"
    p_sub.font.color.rgb = RGBColor(127, 140, 141)

    # === Layer 3: Data Banners (The Folded Ribbons) ===
    
    # Palette definition for the 4 rows
    row_data = [
        {"color": (59, 89, 152),  "num": "01", "title": "Apply your own style", "body": "Customize this infographic to make your presentation unique."},
        {"color": (166, 38, 57),  "num": "02", "title": "Deliver clear messages", "body": "Ensure your points are easily digestible and visually appealing."},
        {"color": (46, 139, 87),  "num": "03", "title": "Engage your audience", "body": "Keep your listeners focused with bold, structured layouts."},
        {"color": (102, 51, 153), "num": "04", "title": "Achieve your goals", "body": "Drive your main points home with professional formatting."}
    ]

    base_x = Inches(4.8)
    base_y = Inches(1.25)
    row_h = Inches(1.0)
    gap_y = Inches(0.35)
    
    # Geometry parameters for the origami effect
    arrow_depth = Inches(0.4)
    chevron_width = Inches(0.28)
    rect_width = Inches(7.0)

    for i, data in enumerate(row_data):
        y = base_y + i * (row_h + gap_y)
        rgb_tuple = data["color"]
        
        # 1. Main Rounded Rectangle (Body)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            base_x + arrow_depth, y, rect_width, row_h
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(*rgb_tuple)
        rect.line.fill.solid()
        rect.line.fill.fore_color.rgb = RGBColor(*rgb_tuple)
        rect.adjustments[0] = 0.15 # Subtle corner rounding

        # 2. Colored Triangle (Left Tip forming the arrow)
        # Perfectly aligns with the flat left edge of the rounded rectangle
        pts_tri = [
            (base_x + arrow_depth, y),           # Top right
            (base_x + arrow_depth, y + row_h),   # Bottom right
            (base_x, y + row_h / 2)              # Left tip
        ]
        ff_tri = slide.shapes.build_freeform(pts_tri[0][0], pts_tri[0][1])
        ff_tri.add_line_segments(pts_tri[1:], close=True)
        tri = ff_tri.convert_to_shape()
        tri.fill.solid()
        tri.fill.fore_color.rgb = RGBColor(*rgb_tuple)
        tri.line.fill.solid()
        tri.line.fill.fore_color.rgb = RGBColor(*rgb_tuple)

        # 3. Gray Chevron (The Origami Fold)
        # Wraps perfectly around the colored triangle tip
        pts_chev = [
            (base_x + arrow_depth, y),                             # Top inner
            (base_x + arrow_depth - chevron_width, y),             # Top outer
            (base_x - chevron_width, y + row_h / 2),               # Middle outer tip
            (base_x + arrow_depth - chevron_width, y + row_h),     # Bottom outer
            (base_x + arrow_depth, y + row_h),                     # Bottom inner
            (base_x, y + row_h / 2)                                # Middle inner tip
        ]
        ff_chev = slide.shapes.build_freeform(pts_chev[0][0], pts_chev[0][1])
        ff_chev.add_line_segments(pts_chev[1:], close=True)
        chev = ff_chev.convert_to_shape()
        chev.fill.solid()
        chev.fill.fore_color.rgb = RGBColor(210, 210, 215) # Neutral Silver
        chev.line.fill.solid()
        chev.line.fill.fore_color.rgb = RGBColor(210, 210, 215)

        # === Layer 4: Text Content inside Banner ===
        
        # Number Text Box
        num_box = slide.shapes.add_textbox(base_x + arrow_depth + Inches(0.1), y, Inches(0.8), row_h)
        num_tf = num_box.text_frame
        num_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_num = num_tf.paragraphs[0]
        p_num.text = data["num"]
        p_num.font.bold = True
        p_num.font.size = Pt(32)
        p_num.font.name = "Arial Black"
        p_num.font.color.rgb = RGBColor(255, 255, 255)
        p_num.alignment = PP_ALIGN.CENTER

        # Subtle Vertical Divider
        div = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            base_x + arrow_depth + Inches(1.1), y + Inches(0.2), Pt(1.5), row_h - Inches(0.4)
        )
        div.fill.solid()
        div.fill.fore_color.rgb = RGBColor(255, 255, 255)
        div.line.fill.background()

        # Text Title & Body
        txt_box = slide.shapes.add_textbox(base_x + arrow_depth + Inches(1.3), y + Inches(0.1), Inches(4.5), row_h - Inches(0.2))
        txt_tf = txt_box.text_frame
        txt_tf.word_wrap = True
        
        p_title = txt_tf.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.bold = True
        p_title.font.size = Pt(14)
        p_title.font.color.rgb = RGBColor(255, 255, 255)
        
        p_body = txt_tf.add_paragraph()
        p_body.text = data["body"]
        p_body.font.size = Pt(11)
        p_body.font.color.rgb = RGBColor(240, 240, 245)

        # === Layer 5: Decorative Accent (Icon stand-in) ===
        # Generates a clean, native minimalist vector geometric target instead of an external image
        cx = base_x + arrow_depth + rect_width - Inches(0.7)
        cy = y + Inches(0.25)
        
        ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, Inches(0.5), Inches(0.5))
        ring.fill.background()
        ring.line.color.rgb = RGBColor(255, 255, 255)
        ring.line.width = Pt(1.5)
        
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx + Inches(0.15), cy + Inches(0.15), Inches(0.2), Inches(0.2))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(255, 255, 255)
        dot.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
