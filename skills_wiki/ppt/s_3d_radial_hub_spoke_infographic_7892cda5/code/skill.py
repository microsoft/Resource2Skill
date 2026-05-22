def create_slide(
    output_pptx_path: str,
    title_text: str = "8 Step Business Infographic",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 3D Radial Hub & Spoke Infographic.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    import math
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Slide Background
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 245, 245)

    # 2. Dimensions & Coordinates
    cx, cy = 6.666, 3.75
    r_pie = 2.0
    r_center = 1.2
    r_small = 0.35
    w_box = 2.6
    h_box = 0.7
    y_offs = [-1.2, -0.4, 0.4, 1.2]

    # Harmonious 8-color palette mapped to the radial slices
    slice_colors = {
        1: (0, 114, 181),    # Mid-Bot Right (Blue)
        2: (0, 161, 112),    # Bot Right (Teal)
        3: (146, 168, 209),  # Bot Left (Light Blue)
        4: (136, 176, 75),   # Mid-Bot Left (Green)
        5: (239, 192, 80),   # Mid-Top Left (Yellow)
        6: (225, 93, 68),    # Top Left (Red/Orange)
        7: (107, 91, 149),   # Top Right (Purple)
        8: (52, 86, 139)     # Mid-Top Right (Dark Blue)
    }

    # Helper: Inject OOXML 3D Sphere Effect
    def apply_3d_sphere(shape):
        spPr = shape.element.spPr
        nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        
        # Clear existing fills/lines/effects
        for tag in ['solidFill', 'gradFill', 'ln', 'effectLst']:
            for elem in spPr.findall(f"a:{tag}", namespaces=nsmap):
                spPr.remove(elem)
                
        # Radial gradient shifted to top-left (30% in) for 3D light source
        grad_xml = """
        <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:gsLst>
                <a:gs pos="0"><a:srgbClr val="FFFFFF"/></a:gs>
                <a:gs pos="70000"><a:srgbClr val="E8E8E8"/></a:gs>
                <a:gs pos="100000"><a:srgbClr val="CCCCCC"/></a:gs>
            </a:gsLst>
            <a:path path="circle"><a:fillToRect l="30000" t="30000" r="70000" b="70000"/></a:path>
        </a:gradFill>
        """
        ln_xml = '<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:ln>'
        # Downward drop shadow
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="120000" dist="40000" dir="5400000" algn="ctr">
                <a:srgbClr val="000000"><a:alpha val="25000"/></a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        spPr.append(etree.fromstring(grad_xml))
        spPr.append(etree.fromstring(ln_xml))
        spPr.append(etree.fromstring(shadow_xml))

    # Helper: Draw Pie Slice
    def draw_pie_slice(cx_in, cy_in, radius_in, start_angle, end_angle, color):
        builder = slide.shapes.build_freeform(Inches(cx_in), Inches(cy_in))
        pts = []
        steps = 15
        for i in range(steps + 1):
            angle = start_angle + (end_angle - start_angle) * i / steps
            x = cx_in + radius_in * math.cos(math.radians(angle))
            y = cy_in + radius_in * math.sin(math.radians(angle))
            pts.append((Inches(x), Inches(y)))
        builder.add_line_segments(pts)
        builder.add_line_segments([(Inches(cx_in), Inches(cy_in))])
        shp = builder.convert_to_shape()
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*color)
        shp.line.fill.background()

    # --- LAYER 1: Text Box Pills (Backgrounds) ---
    left_map = [6, 5, 4, 3]
    right_map = [7, 8, 1, 2]
    left_positions, right_positions = [], []

    for i in range(4):
        y = cy + y_offs[i]
        x_off = math.sqrt(r_pie**2 - y_offs[i]**2)
        
        # Left Pill
        left_x = cx - x_off - w_box + r_small
        box_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_x), Inches(y - h_box/2), Inches(w_box), Inches(h_box))
        box_l.adjustments[0] = 0.5
        box_l.fill.solid()
        box_l.fill.fore_color.rgb = RGBColor(*slice_colors[left_map[i]])
        box_l.line.fill.background()
        left_positions.append((left_x, y - h_box/2))
        
        # Right Pill
        right_x = cx + x_off - r_small
        box_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(right_x), Inches(y - h_box/2), Inches(w_box), Inches(h_box))
        box_r.adjustments[0] = 0.5
        box_r.fill.solid()
        box_r.fill.fore_color.rgb = RGBColor(*slice_colors[right_map[i]])
        box_r.line.fill.background()
        right_positions.append((right_x, y - h_box/2))

    # --- LAYER 2: Segmented Pie Ring ---
    slice_angles = {1: (0,45), 2: (45,90), 3: (90,135), 4: (135,180), 5: (180,225), 6: (225,270), 7: (270,315), 8: (315,360)}
    for sid, angles in slice_angles.items():
        draw_pie_slice(cx, cy, r_pie, angles[0], angles[1], slice_colors[sid])

    # --- LAYER 3: Central 3D Sphere ---
    center_sphere = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r_center), Inches(cy - r_center), Inches(r_center*2), Inches(r_center*2))
    apply_3d_sphere(center_sphere)
    tf_center = center_sphere.text_frame
    tf_center.word_wrap = True
    p = tf_center.paragraphs[0]
    p.text = "INFOGRAPHIC\nTEMPLATE"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(60, 60, 60)
    p.alignment = PP_ALIGN.CENTER

    # --- LAYER 4: Small 3D Spheres (End-caps) ---
    def add_small_sphere(x_c, y_c, num):
        sph = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x_c - r_small), Inches(y_c - r_small), Inches(r_small*2), Inches(r_small*2))
        apply_3d_sphere(sph)
        tf = sph.text_frame
        tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = 0, 0, 0, 0
        p = tf.paragraphs[0]
        p.text = f"{num:02d}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(90, 90, 90)
        p.alignment = PP_ALIGN.CENTER

    for i in range(4):
        x_off = math.sqrt(r_pie**2 - y_offs[i]**2)
        add_small_sphere(cx - x_off, cy + y_offs[i], left_map[i])   # Left
        add_small_sphere(cx + x_off, cy + y_offs[i], right_map[i])  # Right

    # --- LAYER 5: Text Overlays ---
    def add_text_overlay(x, y, is_left, num):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w_box), Inches(h_box))
        tf = tb.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_top, tf.margin_bottom = 0, 0
        if is_left:
            tf.margin_right, tf.margin_left, align = Inches(0.55), Inches(0.1), PP_ALIGN.RIGHT
        else:
            tf.margin_left, tf.margin_right, align = Inches(0.55), Inches(0.1), PP_ALIGN.LEFT
            
        p = tf.paragraphs[0]
        p.text = f"INFODATA {num:02d}"
        p.font.size, p.font.bold, p.font.color.rgb = Pt(11), True, RGBColor(255, 255, 255)
        p.alignment = align
        p2 = tf.add_paragraph()
        p2.text = "Lorem ipsum dolor sit amet\nconsectetur adipiscing elit."
        p2.font.size, p.font.color.rgb = Pt(9), RGBColor(240, 240, 240)
        p2.alignment = align

    for i in range(4):
        add_text_overlay(left_positions[i][0], left_positions[i][1], True, left_map[i])
        add_text_overlay(right_positions[i][0], right_positions[i][1], False, right_map[i])

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.4), Inches(9.333), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size, p.font.bold, p.font.color.rgb, p.alignment = Pt(28), True, RGBColor(50, 50, 50), PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
