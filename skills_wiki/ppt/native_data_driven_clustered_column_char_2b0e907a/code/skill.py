def create_slide(
    output_pptx_path: str,
    title_text: str = "Chart created in PowerPoint",
    chart_title: str = "Units",
    categories: list = None,
    series_name: str = "Units",
    series_values: list = None,
    bar_color: tuple = (237, 125, 49),  # Default Office Orange (R, G, B)
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Data-Driven Clustered Column Chart visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    # Fallback to tutorial's default data if none provided
    if categories is None:
        categories = ['Region A', 'Region B', 'Region C', 'Region D', 'Region E']
    if series_values is None:
        series_values = [24, 65, 36, 48, 51]

    prs = Presentation()
    # Use standard widescreen aspect ratio as seen in modern PPTs
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Layout 5 is typically "Title Only" in standard Office templates
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    # === Layer 1: Slide Title ===
    if slide.shapes.title:
        title_shape = slide.shapes.title
        title_shape.text = title_text
        title_shape.text_frame.paragraphs[0].font.size = Pt(44)

    # === Layer 2: Chart Data Preparation ===
    # This acts as the programmatic equivalent of copying/pasting Excel data
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(series_name, series_values)

    # === Layer 3: Chart Instantiation & Positioning ===
    x = Inches(2.1)
    y = Inches(1.8)
    cx = Inches(9.1)
    cy = Inches(5.0)
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart

    # === Layer 4: Chart Formatting & Styling ===
    
    # Add chart title
    chart.has_title = True
    chart.chart_title.text_frame.text = chart_title
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(18)

    # Move legend to bottom
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    # Format the specific bar color to match the tutorial's orange
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor(*bar_color)

    # Clean up axes for a modern look
    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
