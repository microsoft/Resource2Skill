def create_slide(
    output_pptx_path: str,
    title_text: str = "INFOGRAPHIC DESIGN",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore.",
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the 'Circular Triforce Infographic' visual effect.
    """
    import math
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import OxmlElement
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Colors
    bg_color = RGBColor(244, 245, 247)
    node_bg = RGBColor(255, 255, 255)
    text_gray = RGBColor(100, 100, 100)
    
    # 0: Green, 1: Orange, 2: Blue
    accent_colors_rgb = [
        (40, 167, 69),
        (253, 126, 20),
        (0, 123, 255)
    ]
    accent_colors_pptx = [RGBColor(*rgb) for rgb in accent_colors_rgb]

    # --- Set Slide Background ---
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # --- Math & Geometry Setup ---
    cx, cy = 13.333 / 2, 7.5 / 2
    r_inches = 2.0
    node_radius = 0.9

    # Angles for the 3 main nodes (0 is right, clockwise)
    # Top: 270 (-90), Bottom Right: 30, Bottom Left: 150
    node_angles = [270, 30, 150]
    
    # Midpoint angles for the small dots
    # Top-Right (between 270 & 30 -> 330), Bottom (90), Top-Left (210)
    dot_angles = [330, 90, 210]

    # --- Layer 1: Generate High-Res Arcs via PIL ---
    # Draw at 2x scale for anti-aliasing
    canvas_w, canvas_h = 8000, 4500
    img = Image.new('RGBA', (canvas_w, canvas_h), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    
    center_px_x, center_px_y = canvas_w // 2, canvas_h // 2
    r_px = int(r_inches * (canvas_w / 13.333))
    arc_width = 32 # Thick arcs
    
    bbox = [
        center_px_x - r_px, center_px_y - r_px,
        center_px_x + r_px, center_px_y + r_px
    ]
    
    # Draw Green Arc (Top to Bottom-Right): 270 to 30. PIL handles this nicely if split:
    draw.arc(bbox, 270, 360, fill=accent_colors_rgb[0], width=arc_width)
    draw.arc(bbox, 0, 30, fill=accent_colors_rgb[0], width=arc_width)
    
    # Draw Orange Arc (Bottom-Right to Bottom-Left): 30 to 150
    draw.arc(bbox, 30, 150, fill=accent_colors_rgb[1], width=arc_width)
    
    # Draw Blue Arc (Bottom-Left to Top): 150 to 270
    draw.arc(bbox, 150, 270, fill=accent_colors_rgb[2], width=arc_width)
    
    # Downscale for smoothness
    img = img.resize((4000, 2250), Image.Resampling.LANCZOS)
    
    # Save to BytesIO and insert to slide
    img_stream = io.BytesIO()
    img.save(img_stream, format='PNG')
    img_stream.seek(0)
    slide.shapes.add_picture(img_stream, 0, 0, width=Inches(13.333))

    # --- Helper: Apply XML Native Drop Shadow ---
    def apply_shadow(shape):
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', '254000') # ~20pt blur
        outerShdw.set('dist', '88900')     # ~7pt distance
        outerShdw.set('dir', '5400000')    # 90 degrees (downward)
        outerShdw.set('algn', 'tl')
        outerShdw.set('rotWithShape', '0')
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '15000')          # 15% opacity
        srgbClr.append(alpha)
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    # --- Layer 2: Small Accent Dots (Midpoints) ---
    dot_radius = 0.15
    for i, angle_deg in enumerate(dot_angles):
        rad = math.radians(angle_deg)
        x = cx + r_inches * math.cos(rad)
        y = cy + r_inches * math.sin(rad)
        
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - dot_radius), Inches(y - dot_radius),
            Inches(dot_radius * 2), Inches(dot_radius * 2)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = accent_colors_pptx[i]
        dot.line.color.rgb = node_bg # White border
        dot.line.width = Pt(2)
        apply_shadow(dot)

    # --- Layer 3: Main White Nodes ---
    for i, angle_deg in enumerate(node_angles):
        rad = math.radians(angle_deg)
        x = cx + r_inches * math.cos(rad)
        y = cy + r_inches * math.sin(rad)
        
        # Add white circle
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - node_radius), Inches(y - node_radius),
            Inches(node_radius * 2), Inches(node_radius * 2)
        )
        node.fill.solid()
        node.fill.fore_color.rgb = node_bg
        node.line.color.rgb = node_bg # Invisible outline avoids PPTX rendering glitches
        apply_shadow(node)
        
        # Add Text inside circle
        tf = node.text_frame
        tf.clear()
        tf.word_wrap = False
        
        # Paragraph 1: Number
        p1 = tf.paragraphs[0]
        p1.text = f"0{i+1}"
        p1.font.size = Pt(26)
        p1.font.color.rgb = accent_colors_pptx[i]
        p1.font.bold = True
        p1.alignment = PP_ALIGN.CENTER
        
        # Paragraph 2: "STEP"
        p2 = tf.add_paragraph()
        p2.text = "STEP"
        p2.font.size = Pt(11)
        p2.font.color.rgb = text_gray
        p2.font.bold = True
        p2.alignment = PP_ALIGN.CENTER

    # --- Layer 4: Descriptive Text Boxes ---
    def add_desc_text(left_in, top_in, align=PP_ALIGN.LEFT, title_color=RGBColor(0,0,0)):
        tx = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(2.5), Inches(1.0))
        tf = tx.text_frame
        tf.clear()
        p1 = tf.paragraphs[0]
        p1.text = "Lorem Ipsum"
        p1.font.size = Pt(14)
        p1.font.color.rgb = title_color
        p1.font.bold = True
        p1.alignment = align
        
        p2 = tf.add_paragraph()
        p2.text = body_text
        p2.font.size = Pt(10)
        p2.font.color.rgb = text_gray
        p2.alignment = align
        
    # Top text
    add_desc_text(cx + 1.2, cy - 2.5, PP_ALIGN.LEFT, accent_colors_pptx[0])
    # Bottom Right text
    add_desc_text(cx + 1.8, cy + 1.0, PP_ALIGN.LEFT, accent_colors_pptx[1])
    # Bottom Left text
    add_desc_text(cx - 4.3, cy + 1.0, PP_ALIGN.RIGHT, accent_colors_pptx[2])

    # --- Layer 5: Central Title Text ---
    center_tx = slide.shapes.add_textbox(Inches(cx - 1.5), Inches(cy - 0.4), Inches(3.0), Inches(0.8))
    c_tf = center_tx.text_frame
    c_tf.clear()
    c_p = c_tf.paragraphs[0]
    c_p.text = title_text
    c_p.font.size = Pt(18)
    c_p.font.color.rgb = RGBColor(30, 40, 60)
    c_p.font.bold = True
    c_p.alignment = PP_ALIGN.CENTER

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
