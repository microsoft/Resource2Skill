def create_slide(
    output_pptx_path: str,
    title_text: str = "TECH VISION 2025",
    subtitle_text: str = "The Future of Digital Transformation",
    bg_keyword: str = "cyberpunk,city",
    accent_hex: str = "00BFFF",  # Cyber Blue
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Tech-Geometric Reveal & Neon Layered Typography' effect.
    """
    import os
    import requests
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Helper function to add alpha to a solid fill
    def set_shape_transparency(shape, alpha_percent):
        alpha_val = int((100 - alpha_percent) * 1000)
        alpha_xml = f'<a:alpha val="{alpha_val}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
        shape.fill.fore_color._xClr.append(parse_xml(alpha_xml))

    # Helper function to convert hex to RGB
    def hex_to_rgb(hex_str):
        hex_str = hex_str.lstrip('#')
        return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

    # --- Layer 1: Background Image ---
    try:
        url = f"https://source.unsplash.com/random/1920x1080/?{bg_keyword}"
        response = requests.get(url, timeout=10)
        image_stream = BytesIO(response.content)
        slide.shapes.add_picture(image_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception:
        # Fallback to dark grey background if network fails
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(20, 20, 25)
        bg.line.fill.background()

    # --- Layer 2: Dark Transparency Mask (Tutorial Trick 3) ---
    mask = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    mask.fill.solid()
    mask.fill.fore_color.rgb = RGBColor(0, 0, 0)
    mask.line.fill.background()
    set_shape_transparency(mask, 65) # 65% transparency

    # --- Layer 3: Geometric Mesh / Triangles (Tutorial Trick 2 adaptation) ---
    # Draw a cluster of tech triangles on the right side
    tri_coords = [
        (10.5, 1.0, 90), (10.5, 4.0, 90), 
        (8.5, 2.5, 90), (12.0, 2.5, -90)
    ]
    for left, top, rot in tri_coords:
        tri = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE, 
            Inches(left), Inches(top), Inches(2), Inches(3.5)
        )
        tri.rotation = rot
        tri.fill.solid()
        tri.fill.fore_color.rgb = RGBColor(255, 255, 255)
        set_shape_transparency(tri, 90) # Highly transparent glass look
        tri.line.color.rgb = RGBColor(255, 255, 255)
        tri.line.width = Pt(1)

    # --- Layer 4: Neon Layered Typography (Tutorial Trick 4) ---
    
    # 4a. Bottom Layer (Hollow + Neon Outline)
    left_margin = Inches(1.0)
    top_margin = Inches(2.5)
    
    box_bottom = slide.shapes.add_textbox(left_margin, top_margin, Inches(8), Inches(2))
    tf_bottom = box_bottom.text_frame
    p_bottom = tf_bottom.paragraphs[0]
    run_bottom = p_bottom.add_run()
    run_bottom.text = title_text
    run_bottom.font.size = Pt(80)
    run_bottom.font.name = "Arial Black"
    run_bottom.font.bold = True

    # Inject lxml to remove solid fill and add colored stroke
    rPr = run_bottom._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag.endswith('Fill') or child.tag.endswith('ln'):
            rPr.remove(child)
            
    # Add <a:noFill/>
    rPr.append(parse_xml('<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
    # Add <a:ln> with accent color
    line_width_emu = int(1.5 * 12700) # 1.5 Pt
    xml_ln = f'''
        <a:ln w="{line_width_emu}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:solidFill>
                <a:srgbClr val="{accent_hex}"/>
            </a:solidFill>
        </a:ln>
    '''
    rPr.append(parse_xml(xml_ln))

    # 4b. Top Layer (Solid White)
    # Offset slightly up and left to create the 3D/glitch pop
    offset = Inches(0.06) 
    box_top = slide.shapes.add_textbox(left_margin - offset, top_margin - offset, Inches(8), Inches(2))
    tf_top = box_top.text_frame
    p_top = tf_top.paragraphs[0]
    run_top = p_top.add_run()
    run_top.text = title_text
    run_top.font.size = Pt(80)
    run_top.font.name = "Arial Black"
    run_top.font.bold = True
    run_top.font.color.rgb = RGBColor(255, 255, 255)

    # --- Layer 5: Subtitle ---
    box_sub = slide.shapes.add_textbox(left_margin, top_margin + Inches(1.5), Inches(8), Inches(1))
    tf_sub = box_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    run_sub = p_sub.add_run()
    run_sub.text = subtitle_text
    run_sub.font.size = Pt(24)
    run_sub.font.name = "Arial"
    run_sub.font.color.rgb = RGBColor(200, 200, 200)

    # Add a glowing accent line under the subtitle
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        left_margin, top_margin + Inches(2.2), Inches(3), Pt(4)
    )
    line.fill.solid()
    r, g, b = hex_to_rgb(accent_hex)
    line.fill.fore_color.rgb = RGBColor(r, g, b)
    line.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
