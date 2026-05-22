def create_slide(
    output_pptx_path: str,
    title_text: str = "A Single Focused Message.",
    body_text: str = "",
    bg_palette: str = "business",  
    accent_color: tuple = (255, 107, 107),  # Coral accent
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Modern Pitch Deck "Rule of Three" Layout.
    Synthesizes Tips 1, 2, 4, 7, and 8 into a professional slide.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    import tempfile
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation (16:9 widescreen)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Define Theme Colors
    CHARCOAL = RGBColor(30, 41, 59)
    LIGHT_GRAY = RGBColor(100, 116, 139)
    ACCENT = RGBColor(*accent_color)
    
    # === Layer 1: Powerful Hero Photo (Tip 1) ===
    # Left 40% of the screen
    img_width = Inches(5.333)
    img_height = Inches(7.5)
    
    try:
        # Fetch a professional placeholder image
        url = f"https://picsum.photos/seed/{bg_palette}/800/1200"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img_data = response.read()
            
        with tempfile.NamedTemporaryFile(delete=False, suffix='.jpg') as tmp_img:
            tmp_img.write(img_data)
            tmp_img_path = tmp_img.name
            
        slide.shapes.add_picture(tmp_img_path, 0, 0, width=img_width, height=img_height)
        os.remove(tmp_img_path)
    except Exception:
        # Fallback if download fails: Solid dark shape
        fallback = slide.shapes.add_shape(
            1, 0, 0, img_width, img_height # 1 = msoShapeRectangle
        )
        fallback.fill.solid()
        fallback.fill.fore_color.rgb = CHARCOAL
        fallback.line.fill.background()

    # === Layer 2: Single Focused Message & Typography (Tip 4 & 8) ===
    # Placed in the upper section of the right 60%
    title_box = slide.shapes.add_textbox(Inches(6.33), Inches(0.8), Inches(6.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = title_text
    p.font.name = "Segoe UI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = CHARCOAL

    p2 = tf.add_paragraph()
    p2.text = "Clarity drives action. Our approach minimizes friction and maximizes impact across all key verticals."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(18)
    p2.font.color.rgb = LIGHT_GRAY
    p2.space_before = Pt(12)

    # === Layer 3: The Rule of Three (Tip 7) ===
    # Three distinct columns below the main text
    col_y = Inches(3.8)
    col_w = Inches(1.8)
    col_h = Inches(2.5)
    gap = Inches(0.3)
    start_x = Inches(6.33)

    columns_data = [
        ("01 / Discovery", "We uncover hidden value by deeply analyzing the market landscape and user behaviors."),
        ("02 / Strategy", "We build robust, scalable frameworks designed to endure market volatility and drive growth."),
        ("03 / Execution", "We deploy pixel-perfect solutions with agile methodologies to guarantee speed to market.")
    ]

    for i, (col_title, col_body) in enumerate(columns_data):
        current_x = start_x + (i * (col_w + gap))
        
        # Accent Line for Contrast/Visual Interest (Tip 2)
        accent_line = slide.shapes.add_shape(
            1, current_x, col_y, Inches(0.4), Inches(0.05)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = ACCENT
        accent_line.line.fill.background()

        # Column Text Box
        col_box = slide.shapes.add_textbox(current_x, col_y + Inches(0.2), col_w, col_h)
        col_tf = col_box.text_frame
        col_tf.word_wrap = True
        
        # Subheading
        cp1 = col_tf.add_paragraph()
        cp1.text = col_title
        cp1.font.name = "Segoe UI"
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = CHARCOAL
        
        # Body text
        cp2 = col_tf.add_paragraph()
        cp2.text = col_body
        cp2.font.name = "Segoe UI"
        cp2.font.size = Pt(12)
        cp2.font.color.rgb = LIGHT_GRAY
        cp2.space_before = Pt(8)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
