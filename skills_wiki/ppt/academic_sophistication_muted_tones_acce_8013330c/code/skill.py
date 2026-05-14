import colorsys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_academic_color_palette_slide(
    output_pptx_path: str = "academic_color_guide.pptx"
) -> str:
    """
    Creates a PowerPoint slide demonstrating the principles of academic color theory
    as taught in the LabGirls tutorial.

    The slide showcases:
    1. Transformation of default "bad" colors into professional "good" colors.
    2. The "Dark Base + Vibrant Accent" combination pattern.

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Helper function to modify colors
    def modify_color(rgb_255, sat_factor=1.0, light_factor=1.0):
        r, g, b = [x / 255.0 for x in rgb_255]
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        new_s = max(0, min(1, s * sat_factor))
        new_l = max(0, min(1, l * light_factor))
        new_r, new_g, new_b = colorsys.hls_to_rgb(h, new_l, new_s)
        return (int(new_r * 255), int(new_g * 255), int(new_b * 255))

    # Helper to add a colored swatch with a label
    def add_swatch(left, top, width, height, rgb_color, text, font_size=12, font_color=RGBColor(0,0,0)):
        shape = slide.shapes.add_shape(1, left, top, width, height) # 1 = rectangle
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*rgb_color)
        shape.line.fill.background()

        tb = shape.text_frame
        tb.clear()
        p = tb.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = font_color
        p.alignment = PP_ALIGN.CENTER
        tb.vertical_anchor = 3 # MSO_ANCHOR_MIDDLE

    # --- Slide Title ---
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(15), Inches(0.75))
    p = title_box.text_frame.paragraphs[0]
    p.text = "Academic Color Principles: From Default to Professional"
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(50, 50, 50)

    # --- Part 1: Transforming Base Colors ---
    section_title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(7), Inches(0.5))
    section_title_box.text_frame.paragraphs[0].text = "1. Mute & Deepen Default Colors"
    section_title_box.text_frame.paragraphs[0].font.size = Pt(24)

    default_colors = {
        "Red": (255, 0, 0),
        "Green": (0, 255, 0),
        "Blue": (0, 0, 255),
        "Yellow": (255, 255, 0),
        "Purple": (112, 48, 160)
    }

    start_top = Inches(1.7)
    swatch_h = Inches(0.6)
    swatch_w = Inches(1.5)
    gap = Inches(0.1)
    
    y_pos = start_top
    for name, color in default_colors.items():
        # Label for the row
        label_box = slide.shapes.add_textbox(Inches(0.5), y_pos, Inches(1), swatch_h)
        label_box.text_frame.paragraphs[0].text = name
        label_box.text_frame.paragraphs[0].font.size = Pt(14)

        # Default Color (The "Bad" one)
        add_swatch(Inches(1.8), y_pos, swatch_w, swatch_h, color, "Default")

        # Muted Color (Desaturated)
        muted_color = modify_color(color, sat_factor=0.6, light_factor=1.1)
        add_swatch(Inches(1.8) + (swatch_w + gap) * 1, y_pos, swatch_w, swatch_h, muted_color, "Muted", font_color=RGBColor(255,255,255) if sum(muted_color) < 300 else RGBColor(0,0,0))
        
        # Deep Color (Darkened)
        deep_color = modify_color(color, sat_factor=0.9, light_factor=0.5)
        add_swatch(Inches(1.8) + (swatch_w + gap) * 2, y_pos, swatch_w, swatch_h, deep_color, "Deep", font_color=RGBColor(255,255,255))
        
        # Professional Color (Both)
        prof_color = modify_color(color, sat_factor=0.7, light_factor=0.6)
        add_swatch(Inches(1.8) + (swatch_w + gap) * 3, y_pos, swatch_w, swatch_h, prof_color, "Professional", font_color=RGBColor(255,255,255))
        
        y_pos += swatch_h + Inches(0.2)

    # --- Part 2: Dark Base + Vibrant Accent ---
    section_title_box_2 = slide.shapes.add_textbox(Inches(8.5), Inches(1.0), Inches(7), Inches(0.5))
    section_title_box_2.text_frame.paragraphs[0].text = "2. Pattern: Dark Base + Vibrant Accent"
    section_title_box_2.text_frame.paragraphs[0].font.size = Pt(24)

    base_color = (25, 42, 86) # Deep Navy Blue
    accent_colors = {
        "Highlight Orange": (255, 136, 77),
        "Highlight Yellow": (253, 222, 84),
        "Highlight Cyan": (0, 191, 165),
    }

    # Base background
    base_shape = slide.shapes.add_shape(1, Inches(8.5), Inches(1.7), Inches(7), Inches(4))
    base_shape.fill.solid()
    base_shape.fill.fore_color.rgb = RGBColor(*base_color)
    base_shape.line.fill.background()
    base_shape.text_frame.paragraphs[0].text = "Use a dark, muted color for the main structure or background."
    base_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    base_shape.text_frame.paragraphs[0].font.size = Pt(16)
    
    # Accent swatches
    accent_y = Inches(3.0)
    accent_x_start = Inches(9.0)
    for i, (name, color) in enumerate(accent_colors.items()):
        accent_swatch_w = Inches(1.8)
        accent_swatch_h = Inches(1.2)
        accent_gap = Inches(0.2)
        add_swatch(accent_x_start + (accent_swatch_w + accent_gap) * i, accent_y, accent_swatch_w, accent_swatch_h, color, "Accent", font_size=14)

    accent_label = slide.shapes.add_textbox(Inches(9.0), Inches(4.5), Inches(6.0), Inches(1.0))
    p = accent_label.text_frame.paragraphs[0]
    p.text = "Then, use a single, bright accent color to highlight the key finding or element."
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.size = Pt(16)
    
    # --- Final branding/credit ---
    credit_box = slide.shapes.add_textbox(Inches(0.5), Inches(8.2), Inches(15), Inches(0.5))
    p = credit_box.text_frame.paragraphs[0]
    p.text = "Design Principles from LabGirls Tutorial"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(150, 150, 150)
    p.alignment = PP_ALIGN.RIGHT
    
    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
if __name__ == '__main__':
    saved_path = create_academic_color_palette_slide()
    print(f"Academic color guide slide saved to: {saved_path}")

