def create_slide(
    output_pptx_path: str,
    title_text: str = "NATURE",
    body_text: str = "LOREM IPSUM DOLOR SIT AMET,\nCONSECTETUER ADIPISCING ELIT. AENEAN\nCOMMODO LIGULA EGET DOLOR. AENEAN\nMASSA. CUM SOCIIS NATOQUE PENATIBUS\nET MAGNIS DIS PARTURIENT MONTES.",
    bg_keyword: str = "foggy pine forest landscape",
    accent_color: tuple = (45, 90, 60),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Geometric Cutout Overlay Reveal visual effect.
    """
    import os
    import math
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageFilter

    # --- Configuration ---
    W, H = 1280, 720  # Widescreen 16:9 canvas
    overlay_color = (250, 250, 250, 255)
    shadow_color = (0, 0, 0, 180)
    bg_image_path = "temp_bg.jpg"
    overlay_img_path = "temp_overlay.png"

    # --- Helper: Fetch Background Image ---
    def fetch_image(keyword, path):
        # Using Pollinations AI for reliable thematic placeholder images
        url = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(keyword)}?width={W}&height={H}&nologo=true"
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as response, open(path, 'wb') as f:
                f.write(response.read())
            return True
        except Exception as e:
            print(f"Image download failed: {e}. Falling back to generated gradient.")
            return False

    if not fetch_image(bg_keyword, bg_image_path):
        # Fallback background
        bg = Image.new('RGB', (W, H))
        draw = ImageDraw.Draw(bg)
        for y in range(H):
            r = int(20 + (y/H)*10)
            g = int(30 + (y/H)*30)
            b = int(25 + (y/H)*15)
            draw.line([(0, y), (W, y)], fill=(r, g, b))
        bg.save(bg_image_path)

    # --- Helper: Create Rotated Rounded Rectangle Mask ---
    def create_rotated_rect(width, height, radius, angle):
        diag = int(math.ceil(math.sqrt(width**2 + height**2))) + 40
        img = Image.new('L', (diag, diag), 0)
        draw = ImageDraw.Draw(img)
        x0 = (diag - width) // 2
        y0 = (diag - height) // 2
        draw.rounded_rectangle([x0, y0, x0+width, y0+height], radius, fill=255)
        # Resample=Image.BICUBIC for smooth anti-aliased edges
        return img.rotate(angle, resample=Image.BICUBIC, expand=False), diag

    # --- Step 1: Generate the Overlay Mask with Cutouts ---
    # mask: 255 = solid overlay, 0 = transparent hole
    mask = Image.new('L', (W, H), 255)
    
    # We simulate the complex geometric cutout using two large rotated rounded rectangles
    shape1, size1 = create_rotated_rect(800, 800, 150, 45)
    pos1 = (int(W * 0.65 - size1 // 2), int(H * 0.75 - size1 // 2))
    # Paste 0 (black/hole) using the shape itself as the alpha mask
    mask.paste(0, pos1, shape1)

    shape2, size2 = create_rotated_rect(650, 650, 120, 45)
    pos2 = (int(W * 0.85 - size2 // 2), int(H * 0.25 - size2 // 2))
    mask.paste(0, pos2, shape2)

    # --- Step 2: Generate the Drop Shadow ---
    # Create shadow layer: black color, alpha defined by mask
    shadow = Image.new('RGBA', (W, H), shadow_color)
    shadow.putalpha(mask)
    # Blur it. The blur will bleed *inward* into the 0-alpha (transparent) holes.
    shadow = shadow.filter(ImageFilter.GaussianBlur(25))

    # --- Step 3: Composite Overlay and Shadow ---
    final_overlay = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    # Paste shadow first (slightly offset for depth)
    final_overlay.alpha_composite(shadow, dest=(0, 5))
    
    # Create the solid white overlay layer
    overlay = Image.new('RGBA', (W, H), overlay_color)
    overlay.putalpha(mask)
    
    # Paste solid overlay on top
    final_overlay.alpha_composite(overlay, dest=(0, 0))
    final_overlay.save(overlay_img_path)

    # --- Step 4: Assemble PowerPoint Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Layer 1: Full-bleed background image
    slide.shapes.add_picture(bg_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Layer 2: Geometric transparent overlay
    slide.shapes.add_picture(overlay_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Layer 3: Title Text (Placed on the solid top-left area)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(5), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text.upper()
    p.font.name = 'Georgia' # Elegant serif
    p.font.size = Pt(64)
    p.font.color.rgb = RGBColor(*accent_color)
    
    # Layer 4: Body Text (Placed inside the cutout on the bottom-right over the image)
    body_box = slide.shapes.add_textbox(Inches(5.5), Inches(4.5), Inches(6.5), Inches(2.5))
    bf = body_box.text_frame
    bf.word_wrap = True
    bp = bf.add_paragraph()
    bp.text = body_text
    bp.font.name = 'Arial'
    bp.font.size = Pt(16)
    bp.font.color.rgb = RGBColor(255, 255, 255) # White text for contrast against dark forest
    bp.alignment = PP_ALIGN.CENTER

    # Save and clean up
    prs.save(output_pptx_path)
    
    try:
        os.remove(bg_image_path)
        os.remove(overlay_img_path)
    except OSError:
        pass
        
    return output_pptx_path
