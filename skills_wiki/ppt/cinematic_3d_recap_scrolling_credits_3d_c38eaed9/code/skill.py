import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw, ImageFont
from pptx.oxml.ns import qn

def create_slide(
    output_pptx_path: str,
    credits_data: dict = None,
    bg_center_color: tuple = (20, 35, 60),
    bg_edge_color: tuple = (5, 10, 15),
    **kwargs,
) -> str:
    """
    Creates a cinematic ending slide with a 3D perspective image, a reflection, 
    and dual-column scrolling-style credits.
    """
    if credits_data is None:
        credits_data = {
            "Special Thanks": "All Attendees",
            "Content Strategy": "Product Team",
            "Data Analysis": "Data Science Dept",
            "Slide Design": "Creative Studio",
            "Review & QA": "Management Board",
            "Final Production": "Media Group"
        }

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # ==========================================
    # Helper: Generate Radial Gradient Background
    # ==========================================
    def create_radial_bg(width, height, center_col, edge_col, filename):
        img = Image.new('RGB', (width, height))
        draw = ImageDraw.Draw(img)
        cx, cy = width / 2, height / 2
        max_dist = (cx**2 + cy**2)**0.5
        
        for y in range(height):
            for x in range(width):
                dist = ((x - cx)**2 + (y - cy)**2)**0.5
                ratio = min(1.0, dist / max_dist)
                # Interpolate
                r = int(center_col[0] * (1 - ratio) + edge_col[0] * ratio)
                g = int(center_col[1] * (1 - ratio) + edge_col[1] * ratio)
                b = int(center_col[2] * (1 - ratio) + edge_col[2] * ratio)
                draw.point((x, y), fill=(r, g, b))
        img.save(filename)
        return filename

    # ==========================================
    # Helper: Generate Dummy "Recap" Screen
    # ==========================================
    def create_dummy_screen(filename):
        img = Image.new('RGB', (1920, 1080), color=(13, 20, 35))
        draw = ImageDraw.Draw(img)
        # Draw some tech-looking UI elements
        draw.rectangle([100, 100, 1820, 980], outline=(0, 191, 255), width=10)
        draw.line([100, 300, 1820, 300], fill=(0, 191, 255), width=5)
        # Add a fake chart/wave
        for i in range(10):
            x1 = 200 + i * 150
            y1 = 800 - (i % 3) * 150
            x2 = 200 + (i+1) * 150
            y2 = 800 - ((i+1) % 4) * 120
            draw.line([x1, y1, x2, y2], fill=(0, 255, 150), width=8)
        img.save(filename)
        return filename

    # 1. Apply Background
    bg_path = "temp_cinematic_bg.jpg"
    create_radial_bg(1280, 720, bg_center_color, bg_edge_color, bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # 2. Add 3D Screen Image with Reflection
    screen_path = "temp_recap_screen.jpg"
    create_dummy_screen(screen_path)
    
    # Position on the left side
    pic_left = Inches(1.0)
    pic_top = Inches(2.0)
    pic_width = Inches(5.5)
    pic_height = Inches(3.1)
    
    pic = slide.shapes.add_picture(screen_path, pic_left, pic_top, pic_width, pic_height)

    # --- LXML MAGIC: Inject 3D Rotation and Reflection ---
    spPr = pic.element.spPr
    
    # Inject 3D Perspective (Perspective Right)
    scene3d_xml = """
    <a:scene3d xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:camera prst="perspectiveRight" fov="60000">
            <a:rot lat="0" lon="0" rev="0"/>
        </a:camera>
        <a:lightRig rig="threePt" dir="t">
            <a:rot lat="0" lon="0" rev="1200000"/>
        </a:lightRig>
    </a:scene3d>
    """
    scene3d = parse_xml(scene3d_xml)
    spPr.insert(0, scene3d) # Insert at beginning of spPr

    # Inject Reflection Effect
    effectLst_xml = """
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:reflection blurRad="38100" stA="50000" endA="300" endPos="35000" dist="12700" dir="5400000" sy="-100000" algn="bl" rotWithShape="0"/>
    </a:effectLst>
    """
    effectLst = parse_xml(effectLst_xml)
    spPr.append(effectLst)

    # 3. Add Credits Typography (Two Columns)
    # Column 1: Roles (Right-Aligned)
    role_box = slide.shapes.add_textbox(Inches(6.5), Inches(2.0), Inches(2.5), Inches(4.0))
    role_tf = role_box.text_frame
    role_tf.word_wrap = True
    
    # Column 2: Names (Left-Aligned)
    name_box = slide.shapes.add_textbox(Inches(9.2), Inches(2.0), Inches(3.0), Inches(4.0))
    name_tf = name_box.text_frame
    name_tf.word_wrap = True

    first_line = True
    for role, name in credits_data.items():
        # Role formatting
        p_role = role_tf.paragraphs[0] if first_line else role_tf.add_paragraph()
        p_role.text = role
        p_role.alignment = PP_ALIGN.RIGHT
        p_role.space_after = Pt(14)
        p_role.font.size = Pt(18)
        p_role.font.color.rgb = RGBColor(150, 200, 220)
        p_role.font.name = 'Arial'

        # Name formatting
        p_name = name_tf.paragraphs[0] if first_line else name_tf.add_paragraph()
        p_name.text = name
        p_name.alignment = PP_ALIGN.LEFT
        p_name.space_after = Pt(14)
        p_name.font.size = Pt(18)
        p_name.font.color.rgb = RGBColor(255, 255, 255)
        p_name.font.bold = True
        p_name.font.name = 'Arial'
        
        first_line = False

    # Cleanup temp files
    prs.save(output_pptx_path)
    if os.path.exists(bg_path): os.remove(bg_path)
    if os.path.exists(screen_path): os.remove(screen_path)

    return output_pptx_path
