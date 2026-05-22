def create_slide(
    output_pptx_path: str,
    title_text: str = "NINE-SQUARE\nMATRIX",
    body_text: str = "Transforming standard photography into geometric, structured design assets. This layout introduces visual rhythm and a modern editorial aesthetic.",
    bg_palette: str = "nature,aerial", 
    accent_color: tuple = (0, 191, 255),  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Nine-Square Matrix Reveal" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import urllib.request
    import io

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set background color to very light gray/off-white
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 247, 250)

    # 2. Fetch and Process Image via PIL
    try:
        # Download random Unsplash image based on keyword
        url = f"https://source.unsplash.com/random/1200x1200/?{bg_palette}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            img = Image.open(io.BytesIO(response.read())).convert("RGBA")
    except Exception as e:
        print(f"Image download failed, generating fallback: {e}")
        # Fallback: Generate a gradient-like colorful square
        img = Image.new('RGBA', (1200, 1200), (41, 128, 185, 255))
        draw = ImageDraw.Draw(img)
        for i in range(1200):
            draw.line([(0, i), (1200, i)], fill=(41 + int(i*0.1), 128 + int(i*0.05), 185, 255))

    # Center crop image to perfect 1:1 square
    width, height = img.size
    min_dim = min(width, height)
    left = (width - min_dim) / 2
    top = (height - min_dim) / 2
    right = (width + min_dim) / 2
    bottom = (height + min_dim) / 2
    square_img = img.crop((left, top, right, bottom))

    # 3. Calculate Grid and Insert into PPTX
    # We want a 3x3 grid. Total size 6x6 inches on the right side.
    total_grid_size = 6.0
    cell_size = total_grid_size / 3
    
    # Starting coordinates (center vertically, align right)
    start_x = Inches(6.5)
    start_y = Inches((7.5 - total_grid_size) / 2) # Center vertically: (7.5 - 6) / 2 = 0.75

    # PIL cropping dimensions
    pil_cell_size = min_dim / 3

    for row in range(3):
        for col in range(3):
            # Crop the specific cell from the PIL image
            c_left = col * pil_cell_size
            c_top = row * pil_cell_size
            c_right = c_left + pil_cell_size
            c_bottom = c_top + pil_cell_size
            
            cell_img = square_img.crop((c_left, c_top, c_right, c_bottom))
            
            # Save cropped cell to memory stream
            img_stream = io.BytesIO()
            cell_img.save(img_stream, format='PNG')
            img_stream.seek(0)
            
            # Calculate PPTX placement
            x_pos = start_x + Inches(col * cell_size)
            y_pos = start_y + Inches(row * cell_size)
            
            # Insert picture
            pic = slide.shapes.add_picture(
                img_stream, 
                x_pos, y_pos, 
                width=Inches(cell_size), 
                height=Inches(cell_size)
            )
            
            # Apply thick white border to recreate the "grid lines/gaps" effect
            pic.line.color.rgb = RGBColor(255, 255, 255)
            pic.line.width = Pt(4.5) # Thick border creates the gap aesthetic

    # 4. Add Text Layout (Left Side)
    # Title
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(4.5), Inches(1.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(20, 20, 20)
    p.font.name = "Arial Black"
    p.line_spacing = 1.1

    # Decorative Line
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(1.5), Inches(3.8), Inches(0.8), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # Body Text
    body_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.1), Inches(4.2), Inches(2.0))
    bf = body_box.text_frame
    bf.word_wrap = True
    bp = bf.paragraphs[0]
    bp.text = body_text
    bp.font.size = Pt(14)
    bp.font.color.rgb = RGBColor(100, 100, 100)
    bp.font.name = "Calibri"
    bp.line_spacing = 1.5

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
