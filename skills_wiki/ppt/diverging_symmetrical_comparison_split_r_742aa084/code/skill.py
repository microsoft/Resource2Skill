import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "Before And After",
    left_title: str = "Before",
    right_title: str = "After",
    color_left: tuple = (231, 76, 60),    # Deep Red
    color_right: tuple = (52, 152, 219),  # Deep Blue
    color_left_light: tuple = (253, 237, 236),
    color_right_light: tuple = (235, 245, 251),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Diverging Symmetrical Comparison (Before & After) effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Background Fill
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(248, 249, 250)

    # Coordinates
    center_x = Inches(6.666)
    center_y = Inches(4.15)
    ring_radius = Inches(1.0)

    # 1. Title
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.4), Inches(9.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(44, 62, 80)

    # Calculate row positions
    row_centers_y = [Inches(2.2), Inches(3.5), Inches(4.8), Inches(6.1)]

    # --- LAYER 1: CONNECTING LINES (Drawn first so they hide behind the center ring) ---
    for i, y in enumerate(row_centers_y):
        # LEFT side lines
        left_node_x = Inches(5.2)
        left_elbow_x = Inches(5.6)
        
        # Horizontal segment
        l1 = slide.shapes.add_shape(MSO_SHAPE.LINE, left_node_x, y, left_elbow_x, y)
        l1.line.color.rgb = RGBColor(*color_left)
        l1.line.width = Pt(2)
        # Diagonal segment converging to center
        l2 = slide.shapes.add_shape(MSO_SHAPE.LINE, left_elbow_x, y, center_x, center_y)
        l2.line.color.rgb = RGBColor(*color_left)
        l2.line.width = Pt(2)

        # RIGHT side lines
        right_node_x = Inches(8.133)
        right_elbow_x = Inches(7.733)
        
        r1 = slide.shapes.add_shape(MSO_SHAPE.LINE, right_node_x, y, right_elbow_x, y)
        r1.line.color.rgb = RGBColor(*color_right)
        r1.line.width = Pt(2)
        
        r2 = slide.shapes.add_shape(MSO_SHAPE.LINE, right_elbow_x, y, center_x, center_y)
        r2.line.color.rgb = RGBColor(*color_right)
        r2.line.width = Pt(2)

    # --- LAYER 2: THE CENTRAL SPLIT RING (PIL Generated) ---
    img_size = 600
    ring_img = Image.new('RGBA', (img_size, img_size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(ring_img)
    stroke_width = 70
    
    # PIL angles: 0 is 3 o'clock, sweeping clockwise. 
    # Left half: 90 (6 o'clock) to 270 (12 o'clock)
    draw.arc((stroke_width, stroke_width, img_size-stroke_width, img_size-stroke_width), 
             90, 270, fill=color_left + (255,), width=stroke_width)
    # Right half: 270 (12 o'clock) to 90 (6 o'clock)
    draw.arc((stroke_width, stroke_width, img_size-stroke_width, img_size-stroke_width), 
             270, 90, fill=color_right + (255,), width=stroke_width)
    
    ring_path = "temp_split_ring.png"
    ring_img.save(ring_path)
    
    # Insert ring
    slide.shapes.add_picture(ring_path, center_x - ring_radius, center_y - ring_radius, ring_radius * 2, ring_radius * 2)

    # Central Text Box
    ct_box = slide.shapes.add_shape(MSO_SHAPE.OVAL, center_x - Inches(0.7), center_y - Inches(0.7), Inches(1.4), Inches(1.4))
    ct_box.fill.solid()
    ct_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    ct_box.line.fill.background()
    ct_tf = ct_box.text_frame
    ct_tf.word_wrap = True
    ct_p = ct_tf.paragraphs[0]
    ct_p.text = "Before\n& After"
    ct_p.alignment = PP_ALIGN.CENTER
    ct_p.font.size = Pt(16)
    ct_p.font.bold = True
    ct_p.font.color.rgb = RGBColor(44, 62, 80)

    # --- LAYER 3: CONTENT CARDS & NODES ---
    for i, y in enumerate(row_centers_y):
        idx_str = f"0{i+1}"
        box_w, box_h = Inches(3.3), Inches(0.9)
        badge_w = Inches(0.7)
        top_y = y - (box_h / 2)
        
        # --- LEFT CARD ---
        # Main Content Box
        l_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), top_y, box_w, box_h)
        l_box.fill.solid()
        l_box.fill.fore_color.rgb = RGBColor(*color_left_light)
        l_box.line.fill.background()
        
        l_tf = l_box.text_frame
        l_tf.margin_left = Inches(0.2)
        l_tf.margin_right = Inches(0.2)
        l_p1 = l_tf.paragraphs[0]
        l_p1.text = left_title
        l_p1.font.bold = True
        l_p1.font.color.rgb = RGBColor(*color_left)
        l_p1.font.size = Pt(14)
        
        l_p2 = l_tf.add_paragraph()
        l_p2.text = "Lorem ipsum is simply dummy text of the typesetting industry."
        l_p2.font.color.rgb = RGBColor(80, 80, 80)
        l_p2.font.size = Pt(11)

        # Number Badge
        l_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), top_y, badge_w, box_h)
        l_badge.fill.solid()
        l_badge.fill.fore_color.rgb = RGBColor(*color_left)
        l_badge.line.fill.background()
        l_btf = l_badge.text_frame
        l_bp = l_btf.paragraphs[0]
        l_bp.text = idx_str
        l_bp.alignment = PP_ALIGN.CENTER
        l_bp.font.bold = True
        l_bp.font.size = Pt(18)
        l_bp.font.color.rgb = RGBColor(255, 255, 255)

        # Connection Node
        l_node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.2) - Inches(0.075), y - Inches(0.075), Inches(0.15), Inches(0.15))
        l_node.fill.solid()
        l_node.fill.fore_color.rgb = RGBColor(*color_left)
        l_node.line.fill.background()


        # --- RIGHT CARD ---
        # Number Badge
        r_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.133), top_y, badge_w, box_h)
        r_badge.fill.solid()
        r_badge.fill.fore_color.rgb = RGBColor(*color_right)
        r_badge.line.fill.background()
        r_btf = r_badge.text_frame
        r_bp = r_btf.paragraphs[0]
        r_bp.text = idx_str
        r_bp.alignment = PP_ALIGN.CENTER
        r_bp.font.bold = True
        r_bp.font.size = Pt(18)
        r_bp.font.color.rgb = RGBColor(255, 255, 255)

        # Main Content Box
        r_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.833), top_y, box_w, box_h)
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = RGBColor(*color_right_light)
        r_box.line.fill.background()
        
        r_tf = r_box.text_frame
        r_tf.margin_left = Inches(0.2)
        r_tf.margin_right = Inches(0.2)
        r_p1 = r_tf.paragraphs[0]
        r_p1.text = right_title
        r_p1.font.bold = True
        r_p1.font.color.rgb = RGBColor(*color_right)
        r_p1.font.size = Pt(14)
        
        r_p2 = r_tf.add_paragraph()
        r_p2.text = "Lorem ipsum is simply dummy text of the typesetting industry."
        r_p2.font.color.rgb = RGBColor(80, 80, 80)
        r_p2.font.size = Pt(11)

        # Connection Node
        r_node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.133) - Inches(0.075), y - Inches(0.075), Inches(0.15), Inches(0.15))
        r_node.fill.solid()
        r_node.fill.fore_color.rgb = RGBColor(*color_right)
        r_node.line.fill.background()

    prs.save(output_pptx_path)
    
    # Cleanup temporary image
    if os.path.exists(ring_path):
        os.remove(ring_path)
        
    return output_pptx_path
