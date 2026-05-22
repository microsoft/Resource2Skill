def create_slide(
    output_pptx_path: str,
    title_text: str = "Corporate Organization Structure",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing an automated Org Chart using a recursive layout algorithm.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Data Structure (Hierarchy) ===
    # Designed to fit beautifully on a standard 16:9 slide
    org_data = {
        "name": "CEO",
        "children": [
            {
                "name": "VP Operations",
                "children": [{"name": "Director of Logistics"}, {"name": "Director of Production"}]
            },
            {
                "name": "VP Marketing",
                "children": [{"name": "SEO Manager"}, {"name": "Email Manager"}, {"name": "Webmaster"}]
            },
            {
                "name": "VP Engineering",
                "children": [{"name": "Frontend Lead"}, {"name": "Backend Lead"}]
            }
        ]
    }

    # === Configuration & Styling ===
    config = {
        'node_w': Inches(1.4),
        'node_h': Inches(0.55),
        'h_gap': Inches(0.2),  # horizontal space between sibling leaves
        'v_gap': Inches(0.65), # vertical space between levels
        'margin_top': Inches(1.8),
        'colors': [
            RGBColor(32, 56, 100),   # Level 0
            RGBColor(47, 85, 151),   # Level 1
            RGBColor(68, 114, 196),  # Level 2
            RGBColor(143, 170, 220)  # Level 3+
        ],
        'line_color': RGBColor(130, 130, 130)
    }

    # === Layout Algorithm ===
    leaf_x_cursor = 0

    def calc_positions(node, depth):
        """Recursive post-order traversal to calculate dimensions and positions."""
        nonlocal leaf_x_cursor
        node['y'] = config['margin_top'] + depth * (config['node_h'] + config['v_gap'])
        node['depth'] = depth
        
        if not node.get('children'):
            # It's a leaf node. Assign current cursor x, then advance cursor.
            node['x'] = leaf_x_cursor
            leaf_x_cursor += config['node_w'] + config['h_gap']
        else:
            # It's a parent node. Process children first.
            for child in node['children']:
                calc_positions(child, depth + 1)
            # Center parent over children
            child_x_vals = [c['x'] for c in node['children']]
            node['x'] = sum(child_x_vals) / len(child_x_vals)

    # 1. Calculate relative positions
    calc_positions(org_data, 0)

    # 2. Center the entire chart horizontally on the slide
    chart_total_width = leaf_x_cursor - config['h_gap']
    shift_x = (prs.slide_width - chart_total_width) / 2

    def shift_positions(node):
        node['x'] += shift_x
        for c in node.get('children', []):
            shift_positions(c)

    shift_positions(org_data)

    # === Drawing Functions ===
    def add_shadow(shape):
        """Injects OOXML to add a professional drop shadow to the shape."""
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', '50800') # ~4pt blur
        outerShdw.set('dist', '38100')    # ~3pt distance
        outerShdw.set('dir', '2700000')   # 45 degrees
        outerShdw.set('algn', 'tl')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '30000')         # 30% opacity
        
        srgbClr.append(alpha)
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    def draw_line(x1, y1, x2, y2):
        """Draws a straight connector line."""
        connector = slide.shapes.add_connector(1, x1, y1, x2, y2) # 1 = MSO_CONNECTOR.STRAIGHT
        connector.line.color.rgb = config['line_color']
        connector.line.width = Pt(1.5)

    def draw_connectors_recursive(node):
        """Draw lines first so they sit behind the shapes."""
        if node.get('children'):
            px = node['x'] + config['node_w'] / 2
            py_bottom = node['y'] + config['node_h']
            py_mid = py_bottom + config['v_gap'] / 2

            # Vertical line down from parent
            draw_line(px, py_bottom, px, py_mid)

            children = node['children']
            if len(children) > 1:
                # Horizontal line spanning children
                min_cx = children[0]['x'] + config['node_w'] / 2
                max_cx = children[-1]['x'] + config['node_w'] / 2
                draw_line(min_cx, py_mid, max_cx, py_mid)

            # Vertical lines down to children
            for c in children:
                cx = c['x'] + config['node_w'] / 2
                cy_top = c['y']
                draw_line(cx, py_mid, cx, cy_top)
                draw_connectors_recursive(c)

    def draw_nodes_recursive(node):
        """Draw shapes and text."""
        # Create shape
        color = config['colors'][min(node['depth'], len(config['colors']) - 1)]
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            node['x'], node['y'], config['node_w'], config['node_h']
        )
        
        # Format shape
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background() # Remove outline
        add_shadow(shape)
        
        # Format text
        tf = shape.text_frame
        tf.text = node['name']
        tf.word_wrap = True
        tf.margin_left = Pt(5)
        tf.margin_right = Pt(5)
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.bold = True

        # Recurse
        for c in node.get('children', []):
            draw_nodes_recursive(c)

    # 3. Execute drawing logic
    draw_connectors_recursive(org_data)
    draw_nodes_recursive(org_data)

    # === Title Elements ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.name = "Arial"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(32, 56, 100)
    p.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
