def create_slide(
    output_pptx_path: str,
    title_text: str = "Thank you",
    speaker_name: str = "Jane Doe",
    email_address: str = "jane.doe@example.com",
    qr_url: str = "https://www.linkedin.com/in/example/",
    bg_color: tuple = (13, 33, 79),  # Deep Corporate Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Actionable Personalized Closing Slide".
    Features a clean layout, a simulated handwritten signature, and a dynamic QR code.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    import urllib.parse
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    # Set to 16:9 aspect ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Simulated Logo (Top Left) ===
    # Small white square
    logo_icon = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1), Inches(0.3), Inches(0.3)
    )
    logo_icon.fill.solid()
    logo_icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
    logo_icon.line.fill.background()

    # Company Name
    logo_text_box = slide.shapes.add_textbox(Inches(1.9), Inches(0.85), Inches(3), Inches(0.5))
    logo_p = logo_text_box.text_frame.add_paragraph()
    logo_p.text = "Microsoft Style Presentation"
    logo_p.font.size = Pt(16)
    logo_p.font.bold = True
    logo_p.font.color.rgb = RGBColor(255, 255, 255)
    logo_p.font.name = 'Segoe UI'

    # === Layer 3: Main "Thank You" Title ===
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(8), Inches(1.5))
    title_p = title_box.text_frame.add_paragraph()
    title_p.text = title_text
    title_p.font.size = Pt(64)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(255, 255, 255)
    title_p.font.name = 'Segoe UI'

    # === Layer 4: Signature Element ===
    # Using a script font and a slight rotation to simulate handwriting
    sig_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(5), Inches(1.2))
    sig_box.rotation = -5  # Slight counter-clockwise tilt
    sig_p = sig_box.text_frame.add_paragraph()
    sig_p.text = speaker_name
    sig_p.font.size = Pt(54)
    sig_p.font.color.rgb = RGBColor(255, 255, 255)
    sig_p.font.name = 'Segoe Script'  # Standard Windows cursive font (fallback to standard if missing, but code sets the metadata)

    # === Layer 5: Clickable Contact Info ===
    email_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.6), Inches(6), Inches(0.5))
    email_p = email_box.text_frame.add_paragraph()
    email_p.font.size = Pt(20)
    email_p.font.color.rgb = RGBColor(255, 255, 255)
    email_p.font.name = 'Segoe UI'
    
    # Add clickable mailto hyperlink as emphasized in the tutorial
    email_run = email_p.add_run()
    email_run.text = email_address
    email_run.hyperlink.address = f"mailto:{email_address}"

    # === Layer 6: Dynamic QR Code ===
    qr_image_path = "temp_qr_code.png"
    try:
        # Fetch real QR code from a public API
        encoded_url = urllib.parse.quote(qr_url)
        api_url = f"https://api.qrserver.com/v1/create-qr-code/?size=250x250&data={encoded_url}&color=000000&bgcolor=FFFFFF&margin=1"
        urllib.request.urlretrieve(api_url, qr_image_path)
        
        # Add to bottom right
        slide.shapes.add_picture(
            qr_image_path, Inches(10.3), Inches(4.5), Inches(1.5), Inches(1.5)
        )
    except Exception as e:
        print(f"Notice: Failed to download QR code. Using fallback shape. ({e})")
        # Fallback if no internet or API fails: Draw a white placeholder box
        qr_fallback = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(10.3), Inches(4.5), Inches(1.5), Inches(1.5)
        )
        qr_fallback.fill.solid()
        qr_fallback.fill.fore_color.rgb = RGBColor(255, 255, 255)
        qr_fallback.line.color.rgb = RGBColor(0, 0, 0)
        
        tf = qr_fallback.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = "QR CODE"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0, 0, 0)

    # Cleanup temporary image
    if os.path.exists(qr_image_path):
        os.remove(qr_image_path)

    prs.save(output_pptx_path)
    return output_pptx_path
