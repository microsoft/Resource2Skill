import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "Proprietary Tech",
    body_text: str = "What advantages do we have?",
    bg_palette: str = "light",
    accent_color: tuple = (224, 0, 163),  # Hot Magenta default
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Cinematic Minimalist Product Reveal (Apple style).
    Features crisp typography, a sleek device silhouette, and a dramatic PIL-generated neon under-glow.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank slide
    slide = prs.slides.add_slide(slide_layout)

    # Colors based on palette
    if bg_palette == "dark":
        bg_color = RGBColor(0, 0, 0)
        title_color = RGBColor(245, 245, 247)
        subtitle_color = RGBColor(134, 134, 139)
        device_color = RGBColor(30, 30, 30)
        device_line = RGBColor(80, 80, 80)
    else:
        bg_color = RGBColor(255, 255, 255)
        title_color = RGBColor(29, 29, 31)
        subtitle_color = RGBColor(134, 134, 139)
        device_color = RGBColor(40, 40, 42)
        device_line = RGBColor(100, 100, 105)

    # 1. Set Background Color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # 2. Generate Cinematic Neon Glow using PIL
    # We create a wide canvas to allow the blur to spread without clipping edges
    glow_filename = "temp_neon_glow.png"
    img_w, img_h = 1600, 600
    glow_img = Image.new('RGBA', (img_w, img_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(glow_img)
    
    # Draw the central neon "tube" that will be blurred
    neon_r, neon_g, neon_b = accent_color
    draw.rounded_rectangle(
        [300, 250, 1300, 300], 
        radius=25, 
        fill=(neon_r, neon_g, neon_b, 255)
    )
    
    # Apply heavy Gaussian blur to create the diffuse optical glow
    glow_img = glow_img.filter(ImageFilter.GaussianBlur(radius=70))
    glow_img.save(glow_filename)

    # Insert Glow into slide (centered horizontally, roughly midway vertically)
    slide.shapes.add_picture(
        glow_filename, 
        Inches(1.66), Inches(2.8), width=Inches(10), height=Inches(3.75)
    )

    # 3. Create the Device Silhouette (overlaps the top half of the glow)
    # This creates the illusion that the device is casting the light downwards
    phone_width = Inches(8)
    phone_height = Inches(1.2)
    phone_x = (prs.slide_width - phone_width) / 2
    phone_y = Inches(3.4)  # Overlapping the glow
    
    device_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, phone_x, phone_y, phone_width, phone_height
    )
    device_shape.fill.solid()
    device_shape.fill.fore_color.rgb = device_color
    device_shape.line.color.rgb = device_line
    device_shape.line.width = Pt(1.5)
    
    # Try to make the corners less aggressive (standard PPTX adjustment)
    try:
        device_shape.adjustments[0] = 0.15
    except:
        pass

    # 4. Typography (Top Left Anchor)
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.0), Inches(8), Inches(1))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.name = "Arial" # Fallback for Apple's SF Pro
    p.font.color.rgb = title_color
    
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.font.size = Pt(20)
    p2.font.name = "Arial"
    p2.font.color.rgb = subtitle_color
    p2.space_before = Pt(10)

    # 5. Bottom Icon Grid Placeholders (To emulate the 3 icons in the video)
    icon_y = Inches(5.5)
    icon_size = Inches(0.8)
    spacing = Inches(1.5)
    start_x = (prs.slide_width - (icon_size * 3 + spacing * 2)) / 2

    labels = ["One Button", "Multi-Touch", "OS X"]
    
    for i in range(3):
        x_pos = start_x + (i * (icon_size + spacing))
        
        # Icon Box
        icon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, icon_y, icon_size, icon_size)
        icon.fill.background()
        icon.line.color.rgb = RGBColor(200, 200, 205) if bg_palette != "dark" else RGBColor(80, 80, 80)
        icon.line.width = Pt(1.5)
        try:
            icon.adjustments[0] = 0.2
        except:
            pass
            
        # Label underneath
        label_box = slide.shapes.add_textbox(x_pos - Inches(0.5), icon_y + icon_size + Inches(0.1), icon_size + Inches(1), Inches(0.5))
        lbl_tf = label_box.text_frame
        lbl_tf.word_wrap = True
        lbl_p = lbl_tf.paragraphs[0]
        lbl_p.text = labels[i]
        lbl_p.alignment = PP_ALIGN.CENTER
        lbl_p.font.size = Pt(14)
        lbl_p.font.color.rgb = title_color

    prs.save(output_pptx_path)
    
    # Cleanup temporary image
    if os.path.exists(glow_filename):
        os.remove(glow_filename)
        
    return output_pptx_path

