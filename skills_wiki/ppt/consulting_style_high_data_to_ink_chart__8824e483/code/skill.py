def create_slide(
    output_pptx_path: str,
    title_text: str = "Exercise – How can you improve the visual appearance of this chart?",
    body_text: str = "Net Sales\nEUR mn, German Business Unit",
    bg_palette: str = "none",  # Not used in this minimalist style
    accent_color: tuple = (31, 73, 125),  # Consulting Navy
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Consulting-Style High Data-to-Ink Chart.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Colors
    c_accent = RGBColor(*accent_color)
    c_gray = RGBColor(217, 217, 217)
    c_text = RGBColor(64, 64, 64)
    c_black = RGBColor(0, 0, 0)

    # --- Add Slide Titles ---
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = c_text

    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(6), Inches(0.8))
    tf_sub = subtitle_box.text_frame
    tf_sub.text = body_text
    for p in tf_sub.paragraphs:
        p.font.size = Pt(14)
        p.font.name = "Arial"
        p.font.color.rgb = c_text

    # --- Data Definition ---
    # Exact data from the video tutorial
    data = {
        "2015": 3,
        "2016": 5,
        "2017": 9,
        "2018": 12,
        "2019": 8,
        "2020": 13,
        "2021": 15
    }
    
    # --- Chart Layout Parameters ---
    chart_left = 2.0  # inches
    chart_top = 3.5
    chart_width = 9.0
    chart_height = 3.0
    
    keys = list(data.keys())
    values = list(data.values())
    n = len(data)
    max_val = max(values)
    
    # Calculate widths and scaling
    gap_ratio = 0.4 # Gap is 40% of a column's width
    total_units = n + (n - 1) * gap_ratio
    unit_width = chart_width / total_units
    col_width = unit_width
    gap_width = unit_width * gap_ratio
    
    scale_y = chart_height / max_val if max_val > 0 else 1

    # --- Draw X-Axis Baseline ---
    baseline = slide.shapes.add_connector(
        MSO_SHAPE.LINE_INVERSE, 
        Inches(chart_left - 0.2), Inches(chart_top + chart_height),
        Inches(chart_left + chart_width + 0.2), Inches(chart_top + chart_height)
    )
    baseline.line.color.rgb = RGBColor(180, 180, 180)
    baseline.line.width = Pt(1.5)

    # --- Draw Columns & Labels ---
    col_centers = []
    col_tops = []
    
    for i, (cat, val) in enumerate(data.items()):
        # X positioning
        x = chart_left + i * (col_width + gap_width)
        col_centers.append(x + col_width / 2)
        
        # Y positioning and height
        h = val * scale_y
        y = chart_top + chart_height - h
        col_tops.append(y)
        
        # Draw Column
        fill_color = c_accent if i == n - 1 else c_gray
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(x), Inches(y), Inches(col_width), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.fill.background() # No outline
        
        # Draw Value Label (Above Column)
        val_box = slide.shapes.add_textbox(
            Inches(x), Inches(y - 0.4), Inches(col_width), Inches(0.4)
        )
        p = val_box.text_frame.paragraphs[0]
        p.text = str(val)
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(16)
        p.font.name = "Arial"
        p.font.color.rgb = c_text
        
        # Draw Category Label (Below X-Axis)
        cat_box = slide.shapes.add_textbox(
            Inches(x), Inches(chart_top + chart_height + 0.05), Inches(col_width), Inches(0.4)
        )
        p = cat_box.text_frame.paragraphs[0]
        p.text = str(cat)
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(14)
        p.font.name = "Arial"
        p.font.color.rgb = c_text
        if i == n - 1:
            p.font.bold = True

    # --- Draw the Difference Arrow (Think-cell style) ---
    start_idx = 0
    end_idx = n - 1
    
    x_start = col_centers[start_idx]
    y_start = col_tops[start_idx]
    x_end = col_centers[end_idx]
    y_end = col_tops[end_idx]
    
    # Calculate hovering height (safely above all columns in between)
    highest_point_in_range = min(col_tops[start_idx:end_idx+1])
    arrow_y = highest_point_in_range - 0.6  # 0.6 inches above the highest column
    
    # Draw bracket line using FreeformBuilder
    builder = slide.shapes.build_freeform(Inches(x_start), Inches(y_start - 0.15))
    builder.add_line_segments([
        (Inches(x_start), Inches(arrow_y)),
        (Inches(x_end), Inches(arrow_y)),
        (Inches(x_end), Inches(y_end - 0.25)) # Stop slightly above the target column
    ])
    bracket = builder.convert_to_shape()
    bracket.line.color.rgb = c_black
    bracket.line.width = Pt(1.5)
    
    # Draw custom arrowhead (Solid Triangle) pointing down
    # Coordinates for a downward triangle at (x_end, y_end - 0.15)
    head_w = 0.08
    head_h = 0.12
    tri_builder = slide.shapes.build_freeform(Inches(x_end - head_w), Inches(y_end - 0.15 - head_h))
    tri_builder.add_line_segments([
        (Inches(x_end + head_w), Inches(y_end - 0.15 - head_h)),
        (Inches(x_end), Inches(y_end - 0.15)),
        (Inches(x_end - head_w), Inches(y_end - 0.15 - head_h)) # close shape
    ])
    arrowhead = tri_builder.convert_to_shape()
    arrowhead.fill.solid()
    arrowhead.fill.fore_color.rgb = c_black
    arrowhead.line.fill.background()

    # --- Draw Difference Value Label ---
    # Calculate percentage increase
    val_start = values[start_idx]
    val_end = values[end_idx]
    pct_diff = ((val_end - val_start) / val_start) * 100
    diff_text = f"+{int(pct_diff)}%"
    
    # Position box exactly in the middle of the arrow's horizontal span
    diff_box_width = 1.0
    diff_box_x = x_start + (x_end - x_start) / 2 - (diff_box_width / 2)
    diff_box_y = arrow_y - 0.25
    
    # Add a white background shape to act as a cutout/mask over the line
    bg_mask = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(diff_box_x), Inches(diff_box_y), Inches(diff_box_width), Inches(0.5)
    )
    bg_mask.fill.solid()
    bg_mask.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_mask.line.fill.background()
    
    # Add the text label
    diff_box = slide.shapes.add_textbox(
        Inches(diff_box_x), Inches(diff_box_y + 0.05), Inches(diff_box_width), Inches(0.4)
    )
    p = diff_box.text_frame.paragraphs[0]
    p.text = diff_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = c_black

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
