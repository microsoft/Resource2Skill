def create_slide(
    output_pptx_path: str,
    quote_text: str = "A great option to have in your presentation toolkit is a custom layout. It saves you time when you need to put together a slide deck in a hurry.",
    author_text: str = "— Presentation Bootcamp",
    theme_color_start: tuple = (143, 223, 214),  # Light Teal
    theme_color_end: tuple = (43, 153, 145),     # Dark Teal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Branded Gradient Header & Watermark Quote Layout.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Base Background ===
    # Very light gray background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(242, 242, 242)
    bg.line.fill.background()  # No outline

    # === Layer 2: Gradient Header Bar (Using PIL) ===
    header_height_in = 0.8
    header_width_px = 1333
    header_height_px = 80
    grad_img_path = "temp_header_gradient.png"
    
    # Create horizontal linear gradient
    img = Image.new('RGB', (header_width_px, header_height_px))
    draw = ImageDraw.Draw(img)
    for x in range(header_width_px):
        r = int(theme_color_start[0] + (theme_color_end[0] - theme_color_start[0]) * (x / header_width_px))
        g = int(theme_color_start[1] + (theme_color_end[1] - theme_color_start[1]) * (x / header_width_px))
        b = int(theme_color_start[2] + (theme_color_end[2] - theme_color_start[2]) * (x / header_width_px))
        draw.line([(x, 0), (x, header_height_px)], fill=(r, g, b))
    
    img.save(grad_img_path)
    
    # Insert gradient bar at the top
    slide.shapes.add_picture(grad_img_path, 0, 0, prs.slide_width, Inches(header_height_in))

    # === Layer 3: Watermark Quote Graphics ===
    # Open Quote (Top Left)
    tb_q1 = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(2), Inches(2))
    p1 = tb_q1.text_frame.paragraphs[0]
    p1.text = "“"
    p1.font.size = Pt(250)
    p1.font.name = "Georgia"
    p1.font.color.rgb = RGBColor(215, 215, 215)  # Light gray acts as 80% transparency watermark

    # Close Quote (Bottom Right)
    tb_q2 = slide.shapes.add_textbox(Inches(10.5), Inches(4.0), Inches(2), Inches(2))
    p2 = tb_q2.text_frame.paragraphs[0]
    p2.text = "”"
    p2.font.size = Pt(250)
    p2.font.name = "Georgia"
    p2.font.color.rgb = RGBColor(215, 215, 215)

    # === Layer 4: Central Typography ===
    # Main Quote Text
    tb_main = slide.shapes.add_textbox(Inches(2.5), Inches(2.2), Inches(8.333), Inches(2.5))
    tb_main.text_frame.word_wrap = True
    p_main = tb_main.text_frame.paragraphs[0]
    p_main.text = quote_text
    p_main.font.size = Pt(36)
    p_main.font.italic = True
    p_main.font.name = "Arial"
    p_main.font.color.rgb = RGBColor(80, 80, 80)
    p_main.alignment = PP_ALIGN.CENTER
    
    # Author / Attribution Text
    tb_author = slide.shapes.add_textbox(Inches(2.5), Inches(5.0), Inches(8.333), Inches(1.0))
    p_author = tb_author.text_frame.paragraphs[0]
    p_author.text = author_text
    p_author.font.size = Pt(20)
    p_author.font.bold = True
    p_author.font.name = "Arial"
    p_author.font.color.rgb = RGBColor(theme_color_end[0], theme_color_end[1], theme_color_end[2]) # Accent color
    p_author.alignment = PP_ALIGN.CENTER

    # Cleanup temp files
    prs.save(output_pptx_path)
    if os.path.exists(grad_img_path):
        os.remove(grad_img_path)
        
    return output_pptx_path
