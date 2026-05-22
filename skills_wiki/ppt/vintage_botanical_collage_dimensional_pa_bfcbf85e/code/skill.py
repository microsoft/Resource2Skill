def create_slide(
    output_pptx_path: str,
    title_text: str = "Follow",
    connector_text: str = "- YOUR -",
    keyword_text: str = "HEART",
    bg_color: tuple = (235, 240, 235),  # Soft sage background
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Vintage Botanical Collage' papercraft effect.
    """
    import os
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageFilter

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Canvas dimensions for PIL
    W, H = 1920, 1080
    
    # Colors
    gold = (212, 175, 55, 255)
    cream = (255, 253, 245, 255)
    shadow_color = (0, 0, 0, 80)

    # 1. Base Image Layer
    base_img = Image.new('RGBA', (W, H), bg_color + (255,))
    draw = ImageDraw.Draw(base_img)

    # Subtle vintage wallpaper pattern (stripes)
    for x in range(0, W, 40):
        draw.line([(x, 0), (x, H)], fill=(255, 255, 255, 100), width=3)

    # 2. Ornate Gold Frames
    margin = 120
    # Outer thick frame
    draw.rectangle([margin, margin, W-margin, H-margin], outline=gold, width=6)
    # Inner thin frame
    draw.rectangle([margin+20, margin+20, W-margin-20, H-margin-20], outline=gold, width=2)
    
    # Corner flourishes (simple geometric)
    flourish_rad = 40
    for cx, cy in [(margin, margin), (W-margin, margin), (margin, H-margin), (W-margin, H-margin)]:
        draw.ellipse([cx-flourish_rad, cy-flourish_rad, cx+flourish_rad, cy+flourish_rad], outline=gold, width=4)
        draw.ellipse([cx-flourish_rad+10, cy-flourish_rad+10, cx+flourish_rad-10, cy+flourish_rad-10], outline=gold, width=1)

    # 3. Sentiment Plaque
    plaque_w, plaque_h = 800, 500
    px1, py1 = (W - plaque_w) // 2, (H - plaque_h) // 2
    px2, py2 = px1 + plaque_w, py1 + plaque_h

    # Plaque Drop Shadow (to simulate 3D foam tape)
    shadow_img = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    s_draw = ImageDraw.Draw(shadow_img)
    # Draw shadow slightly offset down and right
    s_draw.rounded_rectangle([px1+15, py1+15, px2+15, py2+15], radius=50, fill=shadow_color)
    shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(20))
    base_img = Image.alpha_composite(base_img, shadow_img)

    # Plaque Body
    draw = ImageDraw.Draw(base_img)
    draw.rounded_rectangle([px1, py1, px2, py2], radius=50, fill=cream, outline=gold, width=5)
    draw.rounded_rectangle([px1+15, py1+15, px2-15, py2-15], radius=35, outline=gold, width=1)

    # 4. Floral Embellishment Clusters
    def draw_floral_cluster(base, cx, cy, angle_offset=0):
        # Create a separate layer for the cluster so we can apply a drop shadow to the whole cluster
        cluster_layer = Image.new('RGBA', (W, H), (0, 0, 0, 0))
        c_draw = ImageDraw.Draw(cluster_layer)
        
        # Simulated watercolor petals/leaves (overlapping ellipses)
        # Leaves
        c_draw.ellipse([cx-80, cy-40, cx, cy+40], fill=(130, 160, 120, 220)) # Green
        c_draw.ellipse([cx-20, cy-80, cx+60, cy+20], fill=(140, 170, 130, 220))
        # Flowers
        c_draw.ellipse([cx-50, cy-50, cx+50, cy+50], fill=(220, 120, 150, 240)) # Deep Rose
        c_draw.ellipse([cx-30, cy-30, cx+30, cy+30], fill=(240, 160, 180, 255)) # Light Pink
        c_draw.ellipse([cx+10, cy-20, cx+70, cy+40], fill=(240, 210, 120, 230)) # Yellow/Gold accent
        c_draw.ellipse([cx+25, cy-5, cx+55, cy+25], fill=(255, 230, 150, 255))
        
        # Add a subtle gold foil center to the main flower
        c_draw.ellipse([cx-10, cy-10, cx+10, cy+10], fill=gold)

        # Apply drop shadow to the entire cluster to make it look like a die-cut sticker
        c_shadow = Image.new('RGBA', (W, H), (0, 0, 0, 0))
        cs_draw = ImageDraw.Draw(c_shadow)
        # Duplicate cluster drawing on shadow layer in black
        cs_draw.ellipse([cx-80+10, cy-40+10, cx+10, cy+40+10], fill=(0,0,0,100))
        cs_draw.ellipse([cx-20+10, cy-80+10, cx+60+10, cy+20+10], fill=(0,0,0,100))
        cs_draw.ellipse([cx-50+10, cy-50+10, cx+50+10, cy+50+10], fill=(0,0,0,100))
        cs_draw.ellipse([cx+10+10, cy-20+10, cx+70+10, cy+40+10], fill=(0,0,0,100))
        c_shadow = c_shadow.filter(ImageFilter.GaussianBlur(12))
        
        # Composite shadow then cluster
        merged = Image.alpha_composite(base, c_shadow)
        return Image.alpha_composite(merged, cluster_layer)

    # Place clusters on top-left and bottom-right corners of the plaque, overlapping the borders
    base_img = draw_floral_cluster(base_img, px1, py1)
    base_img = draw_floral_cluster(base_img, px2, py2)

    # Save background
    bg_img_path = "vintage_craft_bg.png"
    base_img.save(bg_img_path)

    # 5. Insert into PPTX
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # 6. Add Elegant Typography via python-pptx
    # Map pixel coordinates to Inches for PPTX
    left_inch = Inches(13.333 * (px1 / W))
    top_inch = Inches(7.5 * (py1 / H))
    width_inch = Inches(13.333 * (plaque_w / W))
    height_inch = Inches(7.5 * (plaque_h / H))

    txBox = slide.shapes.add_textbox(left_inch, top_inch, width_inch, height_inch)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = 3 # Middle

    # Paragraph 1: "Follow" (Elegant, italicized)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    run1 = p1.add_run()
    run1.text = title_text + "\n"
    run1.font.name = 'Georgia' # Safe elegant serif
    run1.font.size = Pt(44)
    run1.font.italic = True
    run1.font.color.rgb = RGBColor(80, 80, 80)

    # Paragraph 2: "- YOUR -" (Small, spaced)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = connector_text + "\n"
    run2.font.name = 'Georgia'
    run2.font.size = Pt(16)
    run2.font.color.rgb = RGBColor(120, 120, 120)

    # Paragraph 3: "HEART" (Large, Bold)
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    run3 = p3.add_run()
    run3.text = keyword_text
    run3.font.name = 'Georgia'
    run3.font.size = Pt(64)
    run3.font.bold = True
    # Match the deep rose color of the flowers
    run3.font.color.rgb = RGBColor(220, 100, 130)

    prs.save(output_pptx_path)
    
    # Cleanup
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
