def create_slide(
    output_pptx_path: str,
    title_text: str = "IT PRODUCT\nPRESENTATION",
    bg_theme: str = "technology,code", 
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Neon-Accented Asymmetric Geometric Layout' 
    visual effect, complete with custom shape masking and overlapping neon accents.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Color Palette ===
    COLOR_BG = RGBColor(44, 45, 53)
    COLOR_NEON_GREEN = RGBColor(196, 240, 66)
    COLOR_CYAN = RGBColor(0, 208, 197)
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_GREY = RGBColor(180, 180, 180)

    # === Layer 1: Background ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background() # Remove border

    # === Layer 2: Background Geometric Accents ===
    # Neon Green accent block (Behind Image)
    accent_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUND_2_DIAG_RECT, 
        Inches(7.5), Inches(0.8), Inches(4.5), Inches(3.5)
    )
    accent_bg.fill.solid()
    accent_bg.fill.fore_color.rgb = COLOR_NEON_GREEN
    accent_bg.line.fill.background()
    # Increase corner rounding if supported by the shape
    try:
        accent_bg.adjustments[0] = 0.25 
    except:
        pass

    # === Layer 3: Image Fetch & Masking ===
    image_stream = BytesIO()
    try:
        # Attempt to fetch a relevant high-quality image
        url = f"https://images.unsplash.com/featured/800x600/?{bg_theme}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            image_stream.write(response.read())
    except Exception:
        # Fallback to PIL generated placeholder if offline or blocked
        img = Image.new('RGB', (800, 600), color=(60, 65, 80))
        draw = ImageDraw.Draw(img)
        draw.line([(0,0), (800,600)], fill=(80, 85, 100), width=5)
        draw.line([(0,600), (800,0)], fill=(80, 85, 100), width=5)
        img.save(image_stream, format='PNG')
    
    image_stream.seek(0)

    # Add Picture and apply Asymmetric Crop (The core visual trick)
    pic = slide.shapes.add_picture(
        image_stream, 
        Inches(6.2), Inches(1.8), Inches(5.0), Inches(3.8)
    )
    # Apply the signature top-left/bottom-right rounded shape
    pic.auto_shape_type = MSO_SHAPE.ROUND_2_DIAG_RECT


    # === Layer 4: Foreground Geometric Accents ===
    # Bottom Right Cyan Donut intersection
    donut = slide.shapes.add_shape(
        MSO_SHAPE.DONUT, 
        Inches(5.0), Inches(4.5), Inches(2.2), Inches(2.2)
    )
    donut.fill.solid()
    donut.fill.fore_color.rgb = COLOR_CYAN
    donut.line.fill.background()
    try:
        donut.adjustments[0] = 0.35 # Make the ring thicker
    except:
        pass

    # Small Top Left Green Logo/Accent block
    logo_accent = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, 
        Inches(1.0), Inches(0.8), Inches(0.4), Inches(0.6)
    )
    logo_accent.fill.solid()
    logo_accent.fill.fore_color.rgb = COLOR_NEON_GREEN
    logo_accent.line.fill.background()


    # === Layer 5: Typography ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(5.0), Inches(2.0))
    tf = title_box.text_frame
    tf.clear() # Clear default paragraph
    
    lines = title_text.split('\n')
    
    # Line 1 (White, bold)
    p1 = tf.paragraphs[0]
    p1.text = lines[0] if len(lines) > 0 else "IT PRODUCT"
    p1.font.name = 'Arial' # Standard fallback for clean sans-serif
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE

    # Line 2 (Grey, bold)
    if len(lines) > 1:
        p2 = tf.add_paragraph()
        p2.text = lines[1]
        p2.font.name = 'Arial'
        p2.font.size = Pt(56)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_GREY

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("neon_tech_layout.pptx", title_text="IT PRODUCT\nPRESENTATION")
