def create_slide(
    output_pptx_path: str,
    title_text: str = "Dashboard Template",
    body_text: str = "",
    bg_palette: str = "ui_light", 
    accent_color: tuple = (0, 128, 59),  # Default Dark Green matching tutorial
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Modular Dashboard Container Wireframe effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation with 16:9 aspect ratio
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Base Colors
    color_bg = RGBColor(230, 230, 230)       # Deep soft gray
    color_panel = RGBColor(234, 234, 234)    # Light gray container
    color_card = RGBColor(255, 255, 255)     # Pure white cards
    color_accent = RGBColor(*accent_color)   # Dynamic accent (sidebar + icons)

    # === Set Canvas Background ===
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color_bg

    # === Helper Function for UI Cards ===
    def add_ui_card(x: float, y: float, w: float, h: float, bg_color: RGBColor, corner_radius: float = 0.05):
        """Adds a borderless rounded rectangle simulating a UI container."""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        
        # Remove outline
        shape.line.fill.background()
        
        # Tweak corner radius (adjustment values are relative to shape width/height)
        if len(shape.adjustments) > 0:
            shape.adjustments[0] = corner_radius
            
        return shape

    # === Layer 1: Main Structural Panels ===
    
    # Main Content Panel (Light Gray)
    # Starts slightly offset from the left, extending to the right margin
    add_ui_card(x=1.0, y=0.3, w=12.033, h=6.9, bg_color=color_panel, corner_radius=0.03)

    # Sidebar Panel (Accent Color)
    # Overlaps the main panel on the left edge
    add_ui_card(x=0.3, y=0.3, w=1.2, h=6.9, bg_color=color_accent, corner_radius=0.08)

    # === Layer 2: Grid Content Cards ===
    
    grid_x_start = 1.8
    grid_gap = 0.28
    
    # 1. Top Row (6 KPI Cards)
    kpi_y = 0.6
    kpi_w = 1.6
    kpi_h = 1.0
    # Unicode symbols to represent dashboard icons (Profit, Money, Update, Trend, Database, Users)
    icons = ["📈", "💵", "🔄", "📉", "🗄️", "👥"]
    
    for i in range(6):
        x_pos = grid_x_start + i * (kpi_w + grid_gap)
        card = add_ui_card(x=x_pos, y=kpi_y, w=kpi_w, h=kpi_h, bg_color=color_card, corner_radius=0.15)
        
        # Add Icon to card
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = icons[i]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(32)
        p.font.color.rgb = color_accent

    # 2. Middle Row (2 Large Chart Cards)
    mid_y = 1.9
    mid_w = 5.36  # Exactly spans 3 KPI cards + 2 gaps -> (1.6*3) + (0.28*2) = 5.36
    mid_h = 2.8
    
    add_ui_card(x=grid_x_start, y=mid_y, w=mid_w, h=mid_h, bg_color=color_card, corner_radius=0.05)
    add_ui_card(x=grid_x_start + mid_w + grid_gap, y=mid_y, w=mid_w, h=mid_h, bg_color=color_card, corner_radius=0.05)

    # 3. Bottom Row (3 Medium Detail Cards)
    bot_y = 5.0
    # Span is 11.0 total width. (11.0 - (2 * 0.28)) / 3 = 3.48 width per card
    bot_w = 3.48
    bot_h = 1.9
    
    add_ui_card(x=grid_x_start, y=bot_y, w=bot_w, h=bot_h, bg_color=color_card, corner_radius=0.07)
    add_ui_card(x=grid_x_start + bot_w + grid_gap, y=bot_y, w=bot_w, h=bot_h, bg_color=color_card, corner_radius=0.07)
    add_ui_card(x=grid_x_start + 2*(bot_w + grid_gap), y=bot_y, w=bot_w, h=bot_h, bg_color=color_card, corner_radius=0.07)

    # === Optional: Dashboard Title in Sidebar ===
    # Adding some simple text to complete the wireframe look
    sidebar_text = slide.shapes.add_textbox(Inches(0.35), Inches(0.5), Inches(1.1), Inches(1.0))
    tf_side = sidebar_text.text_frame
    tf_side.word_wrap = True
    p_side = tf_side.paragraphs[0]
    p_side.text = title_text
    p_side.alignment = PP_ALIGN.CENTER
    p_side.font.size = Pt(14)
    p_side.font.bold = True
    p_side.font.color.rgb = RGBColor(255, 255, 255)

    # Save the generated presentation
    prs.save(output_pptx_path)
    return output_pptx_path
