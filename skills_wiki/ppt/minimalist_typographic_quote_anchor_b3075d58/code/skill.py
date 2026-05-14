def create_slide(
    output_pptx_path: str,
    quote_text: str = "Design is not just what it looks like and feels like. Design is how it works.",
    author_text: str = "Steve Jobs",
    theme_color_hex: str = "2D3748",  # Slate gray
    bg_color_hex: str = "F8F9FA",    # Off-white
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Minimalist Typographic Quote Anchor" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # === Layer 1: Background ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor.from_string(bg_color_hex)
    bg_shape.line.fill.background()  # Remove default shape border
    
    # === Layer 2: Oversized Outlined Quote Mark ===
    # We use a massive text character combined with custom XML to emulate the vector outline
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3.5), Inches(6.5))
    tf = tx_box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.add_paragraph()
    p.text = "“"
    p.alignment = PP_ALIGN.RIGHT
    
    run = p.runs[0]
    run.font.name = "Georgia"
    run.font.size = Pt(400)
    
    # lxml injection for text outline
    rPr = run._r.get_or_add_rPr()
    
    # Remove any existing fill or line elements to strictly control rendering order
    for child in list(rPr):
        tag_name = child.tag.split('}')[-1]
        if tag_name in ('ln', 'noFill', 'solidFill', 'blipFill', 'gradFill', 'pattFill', 'grpFill'):
            rPr.remove(child)
            
    # Create 5pt outline
    ln_xml = f"""
    <a:ln w="{int(5 * 12700)}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:solidFill>
            <a:srgbClr val="{theme_color_hex}"/>
        </a:solidFill>
    </a:ln>
    """
    ln = parse_xml(ln_xml)
    
    # Create transparent fill (no fill)
    nofill_xml = '<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    fill_el = parse_xml(nofill_xml)
    
    # Insert at the beginning of rPr as required by the CT_TextCharacterProperties schema
    rPr.insert(0, ln)
    rPr.insert(1, fill_el)
    
    # === Layer 3: Main Quote Text ===
    quote_box = slide.shapes.add_textbox(Inches(4.5), Inches(2.2), Inches(7.5), Inches(2.5))
    q_tf = quote_box.text_frame
    q_tf.word_wrap = True
    q_tf.clear()
    
    q_p = q_tf.paragraphs[0]
    q_p.line_spacing = 1.2
    
    q_run = q_p.add_run()
    q_run.text = quote_text
    q_run.font.name = "Arial"
    q_run.font.size = Pt(36)
    q_run.font.color.rgb = RGBColor.from_string(theme_color_hex)
    
    # === Layer 4: Accent Divider Line ===
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(5.2), Inches(1.5), Pt(4)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string(theme_color_hex)
    line.line.fill.background()
    
    # === Layer 5: Author / Citation ===
    author_box = slide.shapes.add_textbox(Inches(4.5), Inches(5.5), Inches(7.5), Inches(1))
    a_tf = author_box.text_frame
    a_tf.clear()
    
    a_p = a_tf.paragraphs[0]
    a_run = a_p.add_run()
    a_run.text = author_text
    a_run.font.name = "Arial"
    a_run.font.size = Pt(20)
    a_run.font.bold = True
    a_run.font.color.rgb = RGBColor.from_string(theme_color_hex)
    
    prs.save(output_pptx_path)
    return output_pptx_path
