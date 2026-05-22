import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

def add_shadow_via_lxml(shape):
    """Injects a subtle outer drop shadow to a shape using lxml."""
    spPr = shape.element.spPr
    
    # Check if effectLst exists, if not, create it
    effectLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
    if effectLst is None:
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
    
    # Create outer shadow
    outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                 blurRad="40000", dist="30000", dir="2700000", algn="ctr", rotWithShape="0")
    
    # Shadow Color (Black with 25% opacity)
    srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
    etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="25000")


def generate_chevron_arc_points(cx, cy, r_inner, r_outer, start_angle, end_angle, notch_angle, steps=30):
    """
    Calculates the exact (x,y) vertices for a curved doughnut slice with a chevron point and notch.
    Angles are in degrees. 0 degrees is facing right (3 o'clock).
    """
    points = []
    
    # Outer arc
    for i in range(steps + 1):
        t = start_angle + (end_angle - start_angle) * (i / steps)
        rad = math.radians(t)
        points.append((cx + r_outer * math.cos(rad), cy + r_outer * math.sin(rad)))
        
    # The inward notch at the end angle
    rad_notch = math.radians(end_angle - notch_angle)
    r_mid = (r_inner + r_outer) / 2
    points.append((cx + r_mid * math.cos(rad_notch), cy + r_mid * math.sin(rad_notch)))
    
    # Inner arc (going backwards)
    for i in range(steps + 1):
        t = end_angle - (end_angle - start_angle) * (i / steps)
        rad = math.radians(t)
        points.append((cx + r_inner * math.cos(rad), cy + r_inner * math.sin(rad)))
        
    # The outward chevron tip at the start angle
    rad_tip = math.radians(start_angle - notch_angle)
    points.append((cx + r_mid * math.cos(rad_tip), cy + r_mid * math.sin(rad_tip)))
    
    return points


def create_slide(
    output_pptx_path: str,
    title_text: str = "INFOGRAPHIC\n6 OPTIONS",
    segments_data: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Circular Chevron Process Ring.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Default data if none provided
    if not segments_data:
        segments_data = [
            {"title": "Lorem Ipsum 1", "desc": "Dummy text of the printing industry."},
            {"title": "Lorem Ipsum 2", "desc": "Dummy text of the printing industry."},
            {"title": "Lorem Ipsum 3", "desc": "Dummy text of the printing industry."},
            {"title": "Lorem Ipsum 4", "desc": "Dummy text of the printing industry."},
            {"title": "Lorem Ipsum 5", "desc": "Dummy text of the printing industry."},
            {"title": "Lorem Ipsum 6", "desc": "Dummy text of the printing industry."},
        ]

    # Theme colors mapped to the video
    colors = [
        (89, 44, 130),   # Purple
        (45, 52, 112),   # Navy
        (66, 110, 169),  # Slate Blue
        (249, 187, 14),  # Yellow
        (235, 126, 31),  # Orange
        (186, 20, 76)    # Red
    ]

    # Geometry Setup
    num_segments = 6
    cx = Inches(6.666)  # Center X
    cy = Inches(3.75)   # Center Y
    r_outer = Inches(2.6)
    r_inner = Inches(1.5)
    notch_deg = 12      # How deep the chevron arrow pushes in
    angle_step = 360 / num_segments
    
    # 1. Draw the chevron segments
    for i in range(num_segments):
        start_angle = i * angle_step
        end_angle = (i + 1) * angle_step
        
        # Calculate vertices
        pts = generate_chevron_arc_points(cx, cy, r_inner, r_outer, start_angle, end_angle, notch_deg)
        
        # Build Freeform shape natively
        builder = slide.shapes.build_freeform(pts[0][0], pts[0][1])
        for pt in pts[1:]:
            builder.add_line_segments((pt,))
        
        # Convert and style
        shape = builder.convert_to_shape()
        color = RGBColor(*colors[i % len(colors)])
        
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        
        # Add a white border to make the interlocking clean
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(1.5)
        
        add_shadow_via_lxml(shape)
        
        # 2. Add numbers inside the segments
        mid_angle_rad = math.radians(start_angle + (angle_step / 2))
        r_mid = (r_inner + r_outer) / 2
        
        # Adjusted slightly backwards to account for the visual shift caused by the chevron arrow
        visual_mid_rad = math.radians(start_angle + (angle_step / 2) - (notch_deg/2))
        icon_x = cx + r_mid * math.cos(visual_mid_rad)
        icon_y = cy + r_mid * math.sin(visual_mid_rad)
        
        icon_box = slide.shapes.add_textbox(icon_x - Inches(0.25), icon_y - Inches(0.25), Inches(0.5), Inches(0.5))
        tf = icon_box.text_frame
        tf.text = str(i + 1)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(24)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # 3. Add text labels radially outside the ring
        r_text = Inches(3.2)
        txt_w = Inches(2.2)
        txt_h = Inches(1.0)
        
        tx = cx + r_text * math.cos(visual_mid_rad)
        ty = cy + r_text * math.sin(visual_mid_rad)
        
        # Adjust text box positioning based on which side of the circle it is
        align = PP_ALIGN.LEFT
        is_left_side = 90 < math.degrees(visual_mid_rad % (2 * math.pi)) < 270
        
        if is_left_side:
            tx -= txt_w  # Shift box to the left of the coordinate
            align = PP_ALIGN.RIGHT
            
        ty -= txt_h / 2 # Center vertically
            
        t_box = slide.shapes.add_textbox(tx, ty, txt_w, txt_h)
        tf = t_box.text_frame
        
        p = tf.add_paragraph()
        p.text = segments_data[i]["title"]
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = color
        p.alignment = align
        
        p2 = tf.add_paragraph()
        p2.text = segments_data[i]["desc"]
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(100, 100, 100)
        p2.alignment = align

    # 4. Central Title Text
    center_box = slide.shapes.add_textbox(cx - Inches(1.5), cy - Inches(0.8), Inches(3.0), Inches(1.6))
    tf_center = center_box.text_frame
    tf_center.word_wrap = True
    
    p_center = tf_center.paragraphs[0]
    p_center.text = title_text
    p_center.font.bold = True
    p_center.font.size = Pt(18)
    p_center.font.color.rgb = RGBColor(50, 50, 50)
    p_center.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
