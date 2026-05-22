def create_slide(
    output_pptx_path: str,
    title_text: str = "Glass Morphism\nin PowerPoint",
    body_text: str = "Follow me for more modern design tips.",
    bg_image_url: str = "https://images.unsplash.com/photo-1557672172-298e090bd0f1?w=1920&q=80",
) -> str:
    """
    Creates a PowerPoint slide featuring a high-fidelity Frosted Glassmorphism panel.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageFilter
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    # --- 1. Background & Glass Rendering (PIL) ---
    slide_width_px = 1920
    slide_height_px = 1080
    
    # Attempt to download the vibrant background, fallback to generating an abstract mesh gradient
    try:
        req = urllib.request.Request(bg_image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            bg_bytes = response.read()
        bg = Image.open(BytesIO(bg_bytes)).convert("RGBA")
        # Crop to 16:9 aspect ratio and resize
        target_ratio = slide_width_px / slide_height_px
        bg_ratio = bg.width / bg.height
        if bg_ratio > target_ratio:
            new_w = int(bg.height * target_ratio)
            offset = (bg.width - new_w) // 2
            bg = bg.crop((offset, 0, offset + new_w, bg.height))
        else:
            new_h = int(bg.width / target_ratio)
            offset = (bg.height - new_h) // 2
            bg = bg.crop((0, offset, bg.width, offset + new_h))
        bg = bg.resize((slide_width_px, slide_height_px), Image.Resampling.LANCZOS)
    except Exception as e:
        print(f"Image download failed ({e}). Generating fallback mesh gradient.")
        bg = Image.new("RGBA", (slide_width_px, slide_height_px), (30, 10, 60, 255))
        draw = ImageDraw.Draw(bg)
        # Draw colorful orbs to create a vibrant abstract background
        draw.ellipse((-200, -200, 1000, 1000), fill=(120, 30, 200, 255))
        draw.ellipse((1000, 400, 2400, 1400), fill=(40, 150, 255, 255))
        draw.ellipse((400, 600, 1500, 1700), fill=(255, 50, 150, 255))
        draw.ellipse((1300, -300, 2100, 500), fill=(255, 200, 50, 255))
        bg = bg.filter(ImageFilter.GaussianBlur(150)) # heavy blur creates the mesh

    # Define Glass Panel Dimensions and Position
    gx, gy = 250, 250
    gw, gh = 850, 500
    g_bounds = (gx, gy, gx + gw, gy + gh)
    corner_radius = 40

    # A) Extract and Blur the region exactly beneath the panel
    glass_blur = bg.crop(g_bounds).filter(ImageFilter.GaussianBlur(35))

    # B) Create Drop Shadow
    shadow_layer = Image.new("RGBA", (slide_width_px, slide_height_px), (0, 0, 0, 0))
    shadow_draw = ImageDraw.Draw(shadow_layer)
    # Draw black rounded rect offset slightly down and right
    shadow_draw.rounded_rectangle((gx+15, gy+20, gx+gw+15, gy+gh+20), radius=corner_radius, fill=(0, 0, 0, 90))
    shadow_layer = shadow_layer.filter(ImageFilter.GaussianBlur(25))
    bg = Image.alpha_composite(bg, shadow_layer)

    # C) Create an exact rounded rectangle mask for pasting the blurred region
    mask = Image.new("L", (gw, gh), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, gw, gh), radius=corner_radius, fill=255)
    
    # Paste the localized frosted blur back onto the background
    bg.paste(glass_blur, (gx, gy), mask)

    # D) Draw Glass Tint and Stroke (Border)
    glass_overlay = Image.new("RGBA", (slide_width_px, slide_height_px), (0, 0, 0, 0))
    overlay_draw = ImageDraw.Draw(glass_overlay)
    # Light white tint
    overlay_draw.rounded_rectangle(g_bounds, radius=corner_radius, fill=(255, 255, 255, 45))
    # Bright semi-transparent border edge
    overlay_draw.rounded_rectangle(g_bounds, radius=corner_radius, outline=(255, 255, 255, 160), width=3)
    
    bg = Image.alpha_composite(bg, glass_overlay)

    # Save composite background
    bg_img_path = "glass_composite_bg.png"
    bg.convert("RGB").save(bg_img_path)


    # --- 2. PowerPoint Assembly ---
    prs = Presentation()
    # 16:9 dimensions matching the 1920x1080 canvas (144 DPI)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Add the generated glassmorphism image as the full-bleed background
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # Coordinate mapping from px to inches (DPI = 144)
    dpi = 144.0
    # Add padding inside the glass box for text
    padding_x = 80
    padding_y = 100
    
    # Text box positioning
    tb_left = Inches((gx + padding_x) / dpi)
    tb_top = Inches((gy + padding_y) / dpi)
    tb_width = Inches((gw - (padding_x * 2)) / dpi)
    tb_height = Inches((gh - (padding_y * 2)) / dpi)

    textbox = slide.shapes.add_textbox(tb_left, tb_top, tb_width, tb_height)
    text_frame = textbox.text_frame
    text_frame.word_wrap = True

    # Title Styling
    p1 = text_frame.paragraphs[0]
    p1.text = title_text
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.name = "Segoe UI"
    p1.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle Styling
    p2 = text_frame.add_paragraph()
    p2.text = f"\n{body_text}"
    p2.font.size = Pt(28)
    p2.font.bold = False
    p2.font.name = "Segoe UI Light"
    p2.font.color.rgb = RGBColor(240, 240, 240)

    prs.save(output_pptx_path)
    
    # Cleanup temporary image file (optional, leaving it for verification)
    # if os.path.exists(bg_img_path): os.remove(bg_img_path)
    
    return output_pptx_path
