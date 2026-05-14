def create_slide(
    output_pptx_path: str,
    title_text: str = "FASHION",
    body_text: str = "show",
    bg_palette: str = "fashion runway", 
    accent_color: tuple = (226, 184, 178),  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Geometric Quarter-Circle Collage' visual effect.
    """
    import os
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageOps

    # --- Setup PPTX ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Geometry & Canvas Constants ---
    # We use 1920x1080 to match 16:9
    W, H = 1920, 1080
    S = 360 # Size of each grid cell. 3 rows * 360 = 1080.
    offset_x = 480 # Leaves 480px on the left for text (4 columns * 360 = 1440)
    offset_y = 0

    # Color definitions
    color1 = (*accent_color, 255) # Dusty pink
    color2 = (140, 120, 115, 210) # Warm transparent taupe
    
    # --- Step 1: Download Background Image ---
    base_img = Image.new('RGB', (W, H), (240, 240, 240))
    try:
        # High-fashion minimalist image
        url = "https://images.unsplash.com/photo-1490481651871-ab68de25d43d?w=1920&h=1080&fit=crop&q=80"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            base_img = Image.open(io.BytesIO(response.read())).convert('RGBA')
    except Exception as e:
        print(f"Could not download image, using fallback. Error: {e}")
    
    base_img = ImageOps.fit(base_img, (W, H), Image.Resampling.LANCZOS).convert("RGBA")

    # --- Step 2: Prepare Masks and Layers ---
    # image_mask controls where the photo is visible
    image_mask = Image.new('L', (W, H), 0)
    img_draw = ImageDraw.Draw(image_mask)
    
    # color_layer holds the solid pastel shapes
    color_layer = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    col_draw = ImageDraw.Draw(color_layer)

    # Helper function to draw perfect geometric shapes based on PIL pieslice angles
    # PIL Angles: 0 is 3 o'clock, 90 is 6 o'clock, 180 is 9 o'clock, 270 is 12 o'clock
    def draw_shape(draw_obj, shape_type, x, y, S, fill_val):
        if shape_type == 'circle':
            draw_obj.ellipse([x, y, x+S, y+S], fill=fill_val)
        elif shape_type == 'arc_tl': # Center at bottom-right
            draw_obj.pieslice([x, y, x+2*S, y+2*S], 180, 270, fill=fill_val)
        elif shape_type == 'arc_tr': # Center at bottom-left
            draw_obj.pieslice([x-S, y, x+S, y+2*S], 270, 360, fill=fill_val)
        elif shape_type == 'arc_bl': # Center at top-right
            draw_obj.pieslice([x, y-S, x+2*S, y+S], 90, 180, fill=fill_val)
        elif shape_type == 'arc_br': # Center at top-left
            draw_obj.pieslice([x-S, y-S, x+S, y+S], 0, 90, fill=fill_val)
        elif shape_type == 'half_t':
            draw_obj.pieslice([x, y, x+S, y+S], 180, 360, fill=fill_val)
        elif shape_type == 'half_b':
            draw_obj.pieslice([x, y, x+S, y+S], 0, 180, fill=fill_val)
        elif shape_type == 'half_l':
            draw_obj.pieslice([x, y, x+S, y+S], 90, 270, fill=fill_val)
        elif shape_type == 'half_r':
            draw_obj.pieslice([x, y, x+S, y+S], 270, 360, fill=fill_val)
            draw_obj.pieslice([x, y, x+S, y+S], 0, 90, fill=fill_val)

    # Grid Mapping: 3 rows x 4 columns
    # Defines the modular collage structure
    layout = [
        # Row 0
        [('arc_br', 'img'), ('half_b', 'img'), ('circle', 'img'), ('arc_bl', 'color1')],
        # Row 1
        [('half_r', 'img'), ('circle', 'color2'), ('arc_tr', 'img'), ('arc_tl', 'img')],
        # Row 2
        [('arc_tr', 'color1'), ('circle', 'img'), ('arc_tl', 'color2'), ('half_l', 'img')]
    ]

    # Render layout to masks
    for row_idx, row in enumerate(layout):
        for col_idx, (shape, content_type) in enumerate(row):
            cx = offset_x + col_idx * S
            cy = offset_y + row_idx * S
            
            if content_type == 'img':
                draw_shape(img_draw, shape, cx, cy, S, 255)
            elif content_type == 'color1':
                draw_shape(col_draw, shape, cx, cy, S, color1)
            elif content_type == 'color2':
                draw_shape(col_draw, shape, cx, cy, S, color2)

    # --- Step 3: Composite Final Image ---
    # Base is pure white canvas
    final_canvas = Image.new('RGBA', (W, H), (255, 255, 255, 255))
    # Paste the photo only where the image_mask dictates
    final_canvas.paste(base_img, (0, 0), image_mask)
    # Overlay the solid color shapes
    final_canvas = Image.alpha_composite(final_canvas, color_layer)

    # Save to disk
    bg_img_path = output_pptx_path.replace('.pptx', '_bg.png')
    final_canvas.save(bg_img_path, format="PNG")

    # Add as background to PPTX
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # --- Step 4: Add Typography ---
    # Title Text
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4), Inches(2))
    tf = tx_box.text_frame
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(64)
    p.font.color.rgb = RGBColor(70, 70, 70)
    
    # Subtitle Text (overlapping / right underneath)
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.font.name = "Georgia"
    p2.font.italic = True
    p2.font.size = Pt(44)
    p2.font.color.rgb = RGBColor(*accent_color)
    
    # Optional small side text (metadata vibe)
    tx_meta = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(2), Inches(1))
    tf_meta = tx_meta.text_frame
    p_meta = tf_meta.paragraphs[0]
    p_meta.text = "WARDROBE\nSpring/Summer\nCollection"
    p_meta.font.name = "Arial"
    p_meta.font.size = Pt(12)
    p_meta.font.color.rgb = RGBColor(150, 150, 150)

    # Save presentation
    prs.save(output_pptx_path)
    
    # Clean up the temp image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    return output_pptx_path
