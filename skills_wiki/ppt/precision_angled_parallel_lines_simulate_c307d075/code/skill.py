import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

def create_slide(
    output_pptx_path: str,
    title_text: str = "Precise Angled Parallel Lines",
    body_text: str = "Simulated Ruler Output",
    angle_degrees: float = 45.0,
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the manual Ruler ink-drawing effect.
    Generates a stylized vector ruler and mathematically aligned parallel lines.
    
    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Background ===
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Slide center coordinates
    cx = Inches(13.333 / 2)
    cy = Inches(7.5 / 2)

    # Calculate radians for layout math
    angle_rad = math.radians(angle_degrees)
    perp_angle_rad = math.radians(angle_degrees + 90)

    # === Layer 1: Simulated Draft Ruler ===
    ruler_w = Inches(12)
    ruler_h = Inches(1.5)
    ruler_left = cx - ruler_w / 2
    ruler_top = cy - ruler_h / 2

    # Main ruler body
    ruler = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ruler_left, ruler_top, ruler_w, ruler_h)
    ruler.rotation = angle_degrees
    ruler.fill.solid()
    ruler.fill.fore_color.rgb = RGBColor(245, 245, 245)
    ruler.line.color.rgb = RGBColor(180, 180, 180)
    ruler.line.width = Pt(1)
    
    # Generate Ruler Tick Marks along the bottom edge
    num_ticks = 60
    tick_length = Inches(0.1)
    
    # Find the rotated bottom edge center
    ruler_bottom_offset = ruler_h / 2
    edge_cx = cx + ruler_bottom_offset * math.cos(perp_angle_rad)
    edge_cy = cy + ruler_bottom_offset * math.sin(perp_angle_rad)
    
    dx_ruler = (ruler_w / 2) * math.cos(angle_rad)
    dy_ruler = (ruler_w / 2) * math.sin(angle_rad)
    
    edge_start_x = edge_cx - dx_ruler
    edge_start_y = edge_cy - dy_ruler
    
    for i in range(num_ticks + 1):
        t = i / num_ticks
        pos_x = edge_start_x + t * (dx_ruler * 2)
        pos_y = edge_start_y + t * (dy_ruler * 2)
        
        # Vary tick lengths for realism (10th, 5th, and standard marks)
        actual_tick_length = tick_length * 2 if i % 10 == 0 else (tick_length * 1.5 if i % 5 == 0 else tick_length)
        
        # Extend inwards into the ruler
        tick_end_x = pos_x - actual_tick_length * math.cos(perp_angle_rad)
        tick_end_y = pos_y - actual_tick_length * math.sin(perp_angle_rad)
        
        tick = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pos_x, pos_y, tick_end_x, tick_end_y)
        tick.line.color.rgb = RGBColor(160, 160, 160)
        tick.line.width = Pt(1)

    # Angle Indicator Badge (Center)
    circle_size = Inches(1.2)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - circle_size/2, cy - circle_size/2, circle_size, circle_size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
    circle.line.color.rgb = RGBColor(200, 200, 200)
    circle.line.width = Pt(1.5)
    
    # Text inside badge (remains perfectly upright by bypassing shape rotation)
    text_frame = circle.text_frame
    text_frame.text = f"{int(angle_degrees)}°"
    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(120, 120, 120)
    p.font.bold = True

    # === Layer 2: Precision Ink Lines ===
    num_lines = 3
    spacing = Inches(0.5)
    length = Inches(8)
    color_rgb = (235, 64, 52)  # Crimson Red ink simulation
    line_width_pt = 2.5

    # Line vector differences
    dx = (length / 2) * math.cos(angle_rad)
    dy = (length / 2) * math.sin(angle_rad)
    
    for i in range(num_lines):
        # Calculate perpendicular offset to lay out parallel lines
        dist = ruler_bottom_offset + Inches(0.15) + (i * spacing)
        
        current_offset_x = dist * math.cos(perp_angle_rad)
        current_offset_y = dist * math.sin(perp_angle_rad)
        
        line_cx = cx + current_offset_x
        line_cy = cy + current_offset_y
        
        start_x = line_cx - dx
        start_y = line_cy - dy
        end_x = line_cx + dx
        end_y = line_cy + dy
        
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y
        )
        connector.line.color.rgb = RGBColor(*color_rgb)
        connector.line.width = Pt(line_width_pt)
        
    # === Layer 3: Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(40, 40, 40)
    
    prs.save(output_pptx_path)
    return output_pptx_path
