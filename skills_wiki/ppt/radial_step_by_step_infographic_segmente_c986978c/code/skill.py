def create_slide(
    output_pptx_path: str,
    title_text: str = "8 STEP CIRCULAR INFOGRAPHIC",
    body_text: str = "",
    bg_palette: str = "light abstract clean",
    **kwargs,
) -> str:
    import math
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Helper 1: Shadow Injection ===
    def add_shadow(shape, blur="40000", dist="20000", alpha="20000"):
        spPr = shape.element.spPr
        effectLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        if effectLst is None:
            effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad=blur, dist=dist, dir="2700000", algn="ctr", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val=alpha)

    # === Helper 2: Draw Custom Segment ===
    def create_pie_arrow_segment(cx_emu, cy_emu, r_in_emu, r_out_emu, r_tip_emu, angle_start, angle_end, arrow_width_deg, color):
        ff_builder = slide.shapes.build_freeform()
        a_start = math.radians(angle_start)
        a_end = math.radians(angle_end)
        a_mid = math.radians((angle_start + angle_end) / 2.0)
        a_arrow_half = math.radians(arrow_width_deg / 2.0)
        
        # 1. Inner arc
        steps = 20
        a_step = (a_end - a_start) / steps
        start_x = int(round(cx_emu + r_in_emu * math.cos(a_start)))
        start_y = int(round(cy_emu + r_in_emu * math.sin(a_start)))
        ff_builder.move_to(start_x, start_y)
        for i in range(1, steps + 1):
            a = a_start + i * a_step
            ff_builder.line_to(int(round(cx_emu + r_in_emu * math.cos(a))), int(round(cy_emu + r_in_emu * math.sin(a))))
            
        # 2. Line to outer radius
        ff_builder.line_to(int(round(cx_emu + r_out_emu * math.cos(a_end))), int(round(cy_emu + r_out_emu * math.sin(a_end))))
        
        # 3. Outer arc right side
        a_arrow_right = a_mid + a_arrow_half
        steps_out1 = max(3, int(20 * (a_end - a_arrow_right) / (a_end - a_start)))
        out_step1 = (a_arrow_right - a_end) / steps_out1
        for i in range(1, steps_out1 + 1):
            a = a_end + i * out_step1
            ff_builder.line_to(int(round(cx_emu + r_out_emu * math.cos(a))), int(round(cy_emu + r_out_emu * math.sin(a))))
            
        # 4. Line to tip
        ff_builder.line_to(int(round(cx_emu + r_tip_emu * math.cos(a_mid))), int(round(cy_emu + r_tip_emu * math.sin(a_mid))))
        
        # 5. Line to arrow base left
        a_arrow_left = a_mid - a_arrow_half
        ff_builder.line_to(int(round(cx_emu + r_out_emu * math.cos(a_arrow_left))), int(round(cy_emu + r_out_emu * math.sin(a_arrow_left))))
        
        # 6. Outer arc left side
        steps_out2 = max(3, int(20 * (a_arrow_left - a_start) / (a_end - a_start)))
        out_step2 = (a_start - a_arrow_left) / steps_out2
        for i in range(1, steps_out2 + 1):
            a = a_arrow_left + i * out_step2
            ff_builder.line_to(int(round(cx_emu + r_out_emu * math.cos(a))), int(round(cy_emu + r_out_emu * math.sin(a))))
            
        # Close
        ff_builder.line_to(start_x, start_y)
        shape = ff_builder.convert_to_shape()
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(2)
        add_shadow(shape)
        return shape

    # === Layer 1: Background ===
    bg_path = "radial_bg.jpg"
    try:
        url = "https://images.unsplash.com/photo-1550684848-fac1c5b4e853?w=1920&q=80"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            with open(bg_path, 'wb') as f:
                f.write(response.read())
    except:
        img = Image.new('RGB', (1920, 1080))
        draw = ImageDraw.Draw(img)
        for y in range(1080):
            r, g, b = int(245 + (10 * y / 1080)), int(247 + (8 * y / 1080)), int(250 + (5 * y / 1080))
            draw.line([(0, y), (1920, y)], fill=(r, g, b))
        img.save(bg_path)

    pic = slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    pic_elem = pic._element
    pic_elem.getparent().remove(pic_elem)
    slide.shapes._spTree.insert(2, pic_elem)  # Send to back safely

    # === Configuration & Palette ===
    cx, cy = Inches(13.333 / 2), Inches(4.2)
    r_in, r_out, r_tip = Inches(1.1), Inches(2.1), Inches(2.5)
    
    colors = [
        (230, 57, 70),   # Red
        (244, 162, 97),  # Orange
        (233, 196, 106), # Yellow
        (42, 157, 143),  # Green
        (0, 150, 199),   # Light Blue
        (2, 62, 138),    # Dark Blue
        (114, 9, 183),   # Purple
        (247, 37, 133)   # Pink/Magenta
    ]

    # === Layer 2: Core Graphics & Elements ===
    # Draw segments
    for i in range(8):
        base_angle = -90 + i * 45
        a_start, a_end = base_angle + 2, base_angle + 45 - 2
        a_mid_rad = math.radians(base_angle + 22.5)
        
        # 1. Main Polygon
        create_pie_arrow_segment(cx, cy, r_in, r_out, r_tip, a_start, a_end, 15.0, colors[i])
        
        # 2. Number inside polygon
        r_num = Inches(1.6)
        tx_inner = cx + r_num * math.cos(a_mid_rad) - Inches(0.2)
        ty_inner = cy + r_num * math.sin(a_mid_rad) - Inches(0.2)
        tb_inner = slide.shapes.add_textbox(int(round(tx_inner)), int(round(ty_inner)), Inches(0.4), Inches(0.4))
        tb_inner.margin_left = tb_inner.margin_right = tb_inner.margin_top = tb_inner.margin_bottom = 0
        p_inner = tb_inner.text_frame.paragraphs[0]
        p_inner.text = f"0{i+1}"
        p_inner.alignment = PP_ALIGN.CENTER
        p_inner.font.size, p_inner.font.bold = Pt(16), True
        p_inner.font.color.rgb = RGBColor(50, 50, 50) if i == 2 else RGBColor(255, 255, 255) # Contrast for yellow

        # 3. Outer Text Box Anchor Math
        r_text_anchor = Inches(2.8)
        anchor_x = cx + r_text_anchor * math.cos(a_mid_rad)
        anchor_y = cy + r_text_anchor * math.sin(a_mid_rad)
        tb_w, tb_h = Inches(1.8), Inches(0.8)

        if math.cos(a_mid_rad) > 0.1:  # Right side placement
            tx, align = anchor_x + Inches(0.1), PP_ALIGN.LEFT
        else:                          # Left side placement
            tx, align = anchor_x - tb_w - Inches(0.1), PP_ALIGN.RIGHT
        ty = anchor_y - (tb_h / 2)

        # 4. Connecting Line & Node
        r_line_start = r_tip + Inches(0.05) # Tiny gap
        line_start_x = cx + r_line_start * math.cos(a_mid_rad)
        line_start_y = cy + r_line_start * math.sin(a_mid_rad)
        
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(line_start_x), int(line_start_y), int(anchor_x), int(anchor_y))
        conn.line.color.rgb, conn.line.width = RGBColor(*colors[i]), Pt(1.5)
        
        n_rad = Inches(0.04)
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(anchor_x - n_rad), int(anchor_y - n_rad), int(n_rad*2), int(n_rad*2))
        node.fill.solid()
        node.fill.fore_color.rgb = RGBColor(*colors[i])
        node.line.fill.background()

        # 5. Outer Text Block
        tb_out = slide.shapes.add_textbox(int(round(tx)), int(round(ty)), int(tb_w), int(tb_h))
        tf_out = tb_out.text_frame
        tf_out.word_wrap = True
        
        p_out = tf_out.paragraphs[0]
        p_out.text = f"Step {i+1} Target"
        p_out.font.size, p_out.font.bold = Pt(14), True
        p_out.font.color.rgb, p_out.alignment = RGBColor(*colors[i]), align
        
        p_out_body = tf_out.add_paragraph()
        p_out_body.text = "Add concise descriptive text here to explain this step of the process."
        p_out_body.font.size, p_out_body.font.color.rgb = Pt(10), RGBColor(100, 100, 100)
        p_out_body.alignment = align

    # === Layer 3: Central Core & Titles ===
    # Central White Circle
    core_rad = Inches(1.0)
    core = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx - core_rad), int(cy - core_rad), int(core_rad*2), int(core_rad*2))
    core.fill.solid()
    core.fill.fore_color.rgb = RGBColor(255, 255, 255)
    core.line.color.rgb, core.line.width = RGBColor(240, 240, 240), Pt(1)
    add_shadow(core, blur="30000", alpha="15000")

    # Center Text
    tb_c = slide.shapes.add_textbox(int(cx - Inches(0.8)), int(cy - Inches(0.4)), Inches(1.6), Inches(0.8))
    p_c = tb_c.text_frame.paragraphs[0]
    p_c.text = "CORE\nCYCLE"
    p_c.alignment, p_c.font.bold, p_c.font.size = PP_ALIGN.CENTER, True, Pt(14)
    p_c.font.color.rgb = RGBColor(60, 60, 60)

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(11.333), Inches(1.2))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.alignment, p_title.font.size, p_title.font.bold = PP_ALIGN.CENTER, Pt(28), True
    p_title.font.color.rgb = RGBColor(40, 40, 40)
    
    p_sub = tf_title.add_paragraph()
    p_sub.text = "A versatile 8-step radial infographic for structured processes"
    p_sub.alignment, p_sub.font.size = PP_ALIGN.CENTER, Pt(14)
    p_sub.font.color.rgb = RGBColor(120, 120, 120)

    prs.save(output_pptx_path)
    return output_pptx_path
