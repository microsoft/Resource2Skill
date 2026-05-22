def create_slide(
    output_pptx_path: str,
    title_text: str = "THE TITLE OF YOUR\nRESEARCH PAPER / THESIS",
    body_text: str = "Your Name\nUniversity Name\nProgram Title\nName Of Advisor",
    bg_palette: str = "office,work", 
    accent_color: tuple = (85, 86, 115),  # Slate Purple from the video
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Modern Diagonal Split Hero" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import urllib.request
    from io import BytesIO
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Background Image ===
    # Download and perfectly crop image to 16:9 using PIL
    img_url = f"https://source.unsplash.com/featured/1920x1080/?{urllib.parse.quote(bg_palette)}"
    
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read()))
            
            # Crop to exactly 16:9
            target_ratio = 16 / 9
            img_ratio = img.width / img.height
            if img_ratio > target_ratio:
                new_w = int(img.height * target_ratio)
                offset = (img.width - new_w) // 2
                img = img.crop((offset, 0, offset + new_w, img.height))
            else:
                new_h = int(img.width / target_ratio)
                offset = (img.height - new_h) // 2
                img = img.crop((0, offset, img.width, offset + new_h))
            
            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG', quality=90)
            img_byte_arr.seek(0)
            
            # Insert full bleed
            slide.shapes.add_picture(img_byte_arr, 0, 0, prs.slide_width, prs.slide_height)
    except Exception as e:
        # Fallback if download fails: Light gray rectangle
        print(f"Image download failed, using fallback. Error: {e}")
        bg_shape = slide.shapes.add_shape(
            1, 0, 0, prs.slide_width, prs.slide_height # 1 is msoShapeRectangle
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(230, 230, 230)
        bg_shape.line.fill.background()

    # === Layer 2: Diagonal Polygon Split ===
    # We draw a polygon on the left side. Top edge spans 65% width, bottom edge 45%.
    w, h = prs.slide_width, prs.slide_height
    start_x, start_y = 0, 0
    builder = slide.shapes.build_freeform(start_x, start_y)
    builder.add_line_segments([
        (0, h),                  # Bottom Left
        (w * 0.45, h),           # Bottom Right (45% of width)
        (w * 0.65, 0),           # Top Right (65% of width - creates the diagonal)
        (0, 0)                   # Back to Top Left
    ])
    
    polygon = builder.convert_to_shape()
    polygon.fill.solid()
    polygon.fill.fore_color.rgb = RGBColor(*accent_color)
    # Hide the border by making it match the fill color
    polygon.line.color.rgb = RGBColor(*accent_color)

    # === Layer 3: Typography & Lines ===
    
    # 1. Main Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.5), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = 'Arial'
    
    # 2. Geometric Divider Line
    line = slide.shapes.add_connector(
        1, Inches(1.0), Inches(3.8), Inches(3.5), Inches(3.8) # 1 is msoConnectorStraight
    )
    line.line.color.rgb = RGBColor(255, 255, 255)
    line.line.width = Pt(2.5)

    # 3. Subtitle / Body Text
    body_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.2), Inches(5.0), Inches(2.0))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(18)
    p_body.font.color.rgb = RGBColor(220, 220, 230) # Slightly dimmed white for hierarchy
    p_body.font.name = 'Arial'
    
    # Set spacing for body text to look like a neat list
    p_body.space_before = Pt(6)
    p_body.line_spacing = 1.3

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
