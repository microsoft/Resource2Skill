def create_slide(
    output_pptx_path: str,
    title_text: str = "Our Core Team",
    card_color: tuple = (41, 128, 185, 255),    # Blue base
    accent_color: tuple = (236, 240, 241, 255), # Light grey avatar bg
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Dynamic Pop-out Profile Matrix'.
    Uses PIL to generate custom card bases and transparent cut-out avatars to demonstrate the layering effect.
    """
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ---------------------------------------------------------
    # Helper 1: Generate Custom Card Base (Curved top, flat bottom)
    # ---------------------------------------------------------
    def generate_card_base(width_px=400, height_px=600, color=(41, 128, 185, 255)):
        img = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        
        # Draw base rectangle
        draw.rectangle([0, 100, width_px, height_px], fill=color)
        
        # Draw a sweeping curve at the top
        # Simulating the boolean subtraction seen in the tutorial
        draw.pieslice([-50, -50, width_px + 50, 200], 0, 180, fill=color)
        
        # Cutout a curve on the top right to make it asymmetrical and dynamic
        draw.ellipse([width_px*0.5, -100, width_px*1.5, 150], fill=(0, 0, 0, 0))
        
        stream = io.BytesIO()
        img.save(stream, format='PNG')
        stream.seek(0)
        return stream

    # ---------------------------------------------------------
    # Helper 2: Generate Transparent "Cut-out" Avatar
    # ---------------------------------------------------------
    def generate_avatar(width_px=400, height_px=500):
        img = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        
        # Colors
        skin = (255, 219, 172, 255)
        shirt = (52, 73, 94, 255)
        
        # Shoulders (breaks out of the bottom)
        draw.ellipse([-50, 250, width_px+50, 750], fill=shirt)
        
        # Neck
        draw.rectangle([width_px/2 - 30, 200, width_px/2 + 30, 280], fill=skin)
        
        # Head
        head_radius = 80
        cx, cy = width_px/2, 160
        draw.ellipse([cx - head_radius, cy - head_radius, cx + head_radius, cy + head_radius], fill=skin)
        
        stream = io.BytesIO()
        img.save(stream, format='PNG')
        stream.seek(0)
        return stream

    # ---------------------------------------------------------
    # Layout Setup
    # ---------------------------------------------------------
    num_profiles = 4
    card_width = 2.2
    card_height = 3.0
    spacing = 0.8
    total_width = (card_width * num_profiles) + (spacing * (num_profiles - 1))
    start_x = (13.333 - total_width) / 2
    
    card_y = 3.5 # Position of the geometric base
    avatar_y = 1.6 # Position of avatar (overlaps the top of the card)
    text_y = 6.6 # Position of text below card

    # Data for the team
    team_data = [
        {"name": "Gao Ya Qi", "title": "Project Director", "desc": "Led major implementations. Exceptional team management."},
        {"name": "Zhang Meng Ting", "title": "PR Manager", "desc": "Expert in establishing extensive social networks and PR."},
        {"name": "Lin Yu Hang", "title": "Strategy Consultant", "desc": "Deep insights into corporate strategy and financial mgmt."},
        {"name": "Chen Ya Ting", "title": "Marketing VP", "desc": "Rich experience in marketing. Highly successful product launches."}
    ]

    # Insert global background curve (a huge soft blue oval at the bottom)
    bg_shape = slide.shapes.add_shape(
        9, # msoShapeOval
        Inches(-2), Inches(6.5), Inches(17.33), Inches(4)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(235, 245, 251) # Very light blue
    bg_shape.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(44, 62, 80)
    p.alignment = PP_ALIGN.CENTER

    # Loop to create each profile column
    for i in range(num_profiles):
        current_x = start_x + i * (card_width + spacing)
        
        # 1. Insert Custom Card Base
        card_stream = generate_card_base(color=card_color)
        slide.shapes.add_picture(card_stream, Inches(current_x), Inches(card_y), width=Inches(card_width), height=Inches(card_height))
        
        # 2. Insert Pop-out Avatar (Overlapping)
        avatar_stream = generate_avatar()
        slide.shapes.add_picture(avatar_stream, Inches(current_x), Inches(avatar_y), width=Inches(card_width))
        
        # 3. Add Name Text
        name_box = slide.shapes.add_textbox(Inches(current_x - 0.2), Inches(text_y), Inches(card_width + 0.4), Inches(0.4))
        name_tf = name_box.text_frame
        p_name = name_tf.paragraphs[0]
        p_name.text = team_data[i]["name"]
        p_name.font.size = Pt(16)
        p_name.font.bold = True
        p_name.font.color.rgb = RGBColor(44, 62, 80)
        p_name.alignment = PP_ALIGN.CENTER
        
        # 4. Add Title/Description Text
        desc_box = slide.shapes.add_textbox(Inches(current_x - 0.2), Inches(text_y + 0.35), Inches(card_width + 0.4), Inches(1.0))
        desc_tf = desc_box.text_frame
        desc_tf.word_wrap = True
        
        # Role Title (Bold, Blue)
        p_title = desc_tf.paragraphs[0]
        p_title.text = team_data[i]["title"] + "\n"
        p_title.font.size = Pt(11)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(*card_color[:3])
        p_title.alignment = PP_ALIGN.CENTER
        
        # Description Text
        p_desc = desc_tf.add_paragraph()
        p_desc.text = team_data[i]["desc"]
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = RGBColor(127, 140, 141)
        p_desc.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
