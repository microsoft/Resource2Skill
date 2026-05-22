def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda",
    bg_keyword: str = "skyscraper",
    theme_colors: list = ["673AB7", "3F51B5", "009688", "F44336", "FF9800"],
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "3D Folded Fading Ribbon Agenda" effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from PIL import Image, ImageEnhance

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Slide Background
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F5F5")

    # 2. Prepare and Insert Anchor Image via PIL
    img_path = "temp_anchor.jpg"
    img_width_in, img_height_in = 3.5, 7.5
    try:
        url = "https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1000"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(img_path, 'wb') as out_file:
                out_file.write(response.read())
        img = Image.open(img_path)
    except Exception:
        # Fallback if download fails
        img = Image.new('RGB', (800, 1500), color='#2C3E50')

    # Crop to exact aspect ratio (3.5 : 7.5)
    target_ratio = img_width_in / img_height_in
    img_ratio = img.width / img.height
    if img_ratio > target_ratio:
        new_width = int(img.height * target_ratio)
        left = (img.width - new_width) // 2
        img = img.crop((left, 0, left + new_width, img.height))
    else:
        new_height = int(img.width / target_ratio)
        top = (img.height - new_height) // 2
        img = img.crop((0, top, img.width, top + new_height))

    # Darken image slightly so white text pops
    img = ImageEnhance.Brightness(img).enhance(0.65)
    img.save(img_path)

    # Insert Image
    slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(img_width_in), height=Inches(img_height_in))

    # 3. Add Vertical Rotated Title on Image
    title_box = slide.shapes.add_textbox(Inches(0.2), Inches(3.0), Inches(4.0), Inches(1.0))
    title_box.rotation = 270  # Rotate to read bottom-to-top
    tf = title_box.text_frame
    tf.text = title_text.upper()
    p = tf.paragraphs[0]
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Helper: Darken Hex Color for Shadows
    def darken_hex(hex_str, factor=0.5):
        r, g, b = tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))
        return f"{int(r*factor):02X}{int(g*factor):02X}{int(b*factor):02X}"

    # Helper: Apply XML Gradient and remove outline
    def apply_fading_gradient(shape, hex_color):
        spPr = shape.element.spPr
        for fill_type in ['a:solidFill', 'a:gradFill', 'a:noFill']:
            for elem in spPr.xpath(f'./{fill_type}'):
                spPr.remove(elem)
        for elem in spPr.xpath('./a:ln'):
            spPr.remove(elem)
        
        # Linear gradient: 100% opacity to 0% opacity
        grad_xml = f"""
        <a:gradFill rotWithShape="1" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:gsLst>
                <a:gs pos="0"><a:srgbClr val="{hex_color}"><a:alpha val="100000"/></a:srgbClr></a:gs>
                <a:gs pos="100000"><a:srgbClr val="{hex_color}"><a:alpha val="0"/></a:srgbClr></a:gs>
            </a:gsLst>
            <a:lin ang="0" scaled="1"/>
        </a:gradFill>
        """
        spPr.append(parse_xml(grad_xml))
        spPr.append(parse_xml('<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:ln>'))

    # Helper: Add shadow to circle
    def apply_shadow(shape):
        spPr = shape.element.spPr
        effect_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="40000" dist="20000" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000"><a:alpha val="40000"/></a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        spPr.append(parse_xml(effect_xml))

    # 4. Generate List Items
    num_items = 5
    start_y = 0.8
    y_step = 1.3
    edge_x = 3.5

    for i in range(num_items):
        item_y = start_y + (i * y_step)
        color_hex = theme_colors[i % len(theme_colors)]
        
        # A. Fading Ribbon (Chevron to hide left edge behind circle)
        ribbon = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(edge_x), Inches(item_y), Inches(6.0), Inches(0.7))
        apply_fading_gradient(ribbon, color_hex)

        # B. 3D Fold (Triangle)
        fold = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(edge_x), Inches(item_y + 0.7), Inches(0.25), Inches(0.25))
        fold.flip_v = True # Move 90-deg angle to Top-Left
        fold.fill.solid()
        fold.fill.fore_color.rgb = RGBColor.from_string(darken_hex(color_hex))
        fold.line.fill.background()

        # C. Anchor Circle
        radius = 0.35
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(edge_x - radius), Inches(item_y), Inches(radius*2), Inches(radius*2))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor.from_string(color_hex)
        circle.line.fill.background()
        apply_shadow(circle)

        # Circle Number
        tf_circ = circle.text_frame
        tf_circ.text = f"{i+1:02d}"
        p_circ = tf_circ.paragraphs[0]
        p_circ.font.bold = True
        p_circ.font.size = Pt(18)
        p_circ.font.color.rgb = RGBColor(255, 255, 255)

        # D. Text Content Block
        tbox = slide.shapes.add_textbox(Inches(edge_x + 0.6), Inches(item_y - 0.05), Inches(5.0), Inches(0.7))
        tf_text = tbox.text_frame
        
        # Item Title
        p_title = tf_text.paragraphs[0]
        p_title.text = f"Agenda Topic {i+1}"
        p_title.font.bold = True
        p_title.font.size = Pt(18)
        p_title.font.color.rgb = RGBColor.from_string(color_hex)
        
        # Item Description
        p_desc = tf_text.add_paragraph()
        p_desc.text = "Provide brief details about your agenda here, in max 2-3 lines. With presentation, lesser the better."
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = RGBColor.from_string("555555")

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(img_path):
        os.remove(img_path)
        
    return output_pptx_path
