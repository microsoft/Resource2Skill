def create_slide(
    output_pptx_path: str,
    org_data: list = None,
    title_text: str = "Company Organization Chart",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing an Algorithmic Dynamic Org Chart.
    
    :param org_data: List of dicts with 'id', 'name', 'title', and 'parent' keys.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from lxml import etree

    # Default organizational hierarchy mirroring the tutorial context
    if not org_data:
        org_data = [
            {"id": "1", "name": "Patti Fernandez", "title": "President", "parent": None},
            {"id": "2", "name": "Kevin Stratvert", "title": "VP Marketing", "parent": "1"},
            {"id": "3", "name": "Miriam Graham", "title": "VP Operations", "parent": "1"},
            {"id": "4", "name": "Lee Gu", "title": "VP Engineering", "parent": "1"},
            {"id": "5", "name": "Megan Bowen", "title": "Marketing Manager", "parent": "2"},
            {"id": "6", "name": "Alex Wilber", "title": "PR Lead", "parent": "2"},
            {"id": "7", "name": "Lidia Holloway", "title": "Logistics", "parent": "3"},
            {"id": "8", "name": "Diego Siciliani", "title": "HR Manager", "parent": "3"},
            {"id": "9", "name": "Grady Archie", "title": "Lead Developer", "parent": "4"},
            {"id": "10", "name": "Johanna Lorenz", "title": "QA Lead", "parent": "4"},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Apply soft background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 247, 250)

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 30, 30)

    # --- Tree Data Structure & Build ---
    class Node:
        def __init__(self, data):
            self.id = data['id']
            self.name = data['name']
            self.title = data['title']
            self.parent_id = data['parent']
            self.children = []
            self.level = 0
            self.subtree_width = 0
            self.x = 0
            self.y = 0
            self.shape = None

    nodes_by_id = {str(d['id']): Node(d) for d in org_data}
    root_nodes = []
    
    for node in nodes_by_id.values():
        if node.parent_id is not None and str(node.parent_id) in nodes_by_id:
            parent = nodes_by_id[str(node.parent_id)]
            parent.children.append(node)
        else:
            root_nodes.append(node)

    def assign_levels(node, level):
        node.level = level
        for child in node.children:
            assign_levels(child, level + 1)

    for root in root_nodes:
        assign_levels(root, 0)

    # --- Algorithmic Layout Calculation ---
    base_w = Inches(2.0)
    base_h = Inches(0.8)
    base_h_gap = Inches(0.3)
    base_v_gap = Inches(0.8)

    def calc_width(node, scale=1.0):
        actual_w = base_w * scale
        actual_h_gap = base_h_gap * scale
        if not node.children:
            node.subtree_width = actual_w + actual_h_gap
        else:
            node.subtree_width = sum(calc_width(c, scale) for c in node.children)
        node.subtree_width = max(node.subtree_width, actual_w + actual_h_gap)
        return node.subtree_width

    def calc_depth(node):
        if not node.children:
            return 1
        return 1 + max(calc_depth(c) for c in node.children)

    # Determine required scaling to fit slide
    total_w = sum(calc_width(r, 1.0) for r in root_nodes)
    max_d = max(calc_depth(r) for r in root_nodes) if root_nodes else 1
    
    margin_x = Inches(0.5)
    start_y = Inches(1.5)
    max_w_avail = prs.slide_width - margin_x * 2
    max_h_avail = prs.slide_height - start_y - Inches(0.5)
    
    total_h = max_d * base_h + (max_d - 1) * base_v_gap

    scale = 1.0
    if total_w > max_w_avail:
        scale = min(scale, max_w_avail / total_w)
    if total_h > max_h_avail:
        scale = min(scale, max_h_avail / total_h)

    # Re-calculate geometry with final scale
    total_w_scaled = sum(calc_width(r, scale) for r in root_nodes)
    actual_w = base_w * scale
    actual_h = base_h * scale
    actual_v_gap = base_v_gap * scale

    current_x = (prs.slide_width - total_w_scaled) / 2

    def set_pos(node, x_start, y_pos):
        center_x = x_start + node.subtree_width / 2
        node.x = center_x - actual_w / 2
        node.y = y_pos

        child_x = x_start
        for child in node.children:
            set_pos(child, child_x, y_pos + actual_h + actual_v_gap)
            child_x += child.subtree_width

    for root in root_nodes:
        set_pos(root, current_x, start_y)
        current_x += root.subtree_width

    # --- Drawing Elements ---
    colors = [
        RGBColor(29, 66, 138),   # Level 0 (Exec)
        RGBColor(0, 128, 128),   # Level 1 (Director)
        RGBColor(230, 115, 0),   # Level 2 (Manager)
        RGBColor(103, 58, 183),  # Level 3 (Lead)
        RGBColor(46, 125, 50),   # Level 4
        RGBColor(194, 24, 91)    # Level 5+
    ]

    def add_shadow(shape):
        spPr = shape.element.spPr
        ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        effectLst = etree.SubElement(spPr, f'{{{ns}}}effectLst')
        outerShdw = etree.SubElement(effectLst, f'{{{ns}}}outerShdw',
                                     blurRad="40000", dist="30000", dir="5400000", algn="ctr", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, f'{{{ns}}}srgbClr', val="000000")
        etree.SubElement(srgbClr, f'{{{ns}}}alpha', val="30000")

    def draw_tree(node, slide):
        # Add shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, node.x, node.y, actual_w, actual_h)
        node.shape = shape
        
        # Style Box
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = colors[min(node.level, len(colors)-1)]
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(1.5)
        add_shadow(shape)
        
        # Text Configuration
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = node.name
        p.font.bold = True
        p.font.size = Pt(max(7, 11 * scale))
        p.font.color.rgb = RGBColor(255, 255, 255)
        
        if node.title:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.text = node.title
            p2.font.size = Pt(max(6, 9.5 * scale))
            p2.font.color.rgb = RGBColor(240, 240, 240)
            
        # Draw and Connect Children
        for child in node.children:
            child_shape = draw_tree(child, slide)
            # Add dynamic elbow connector
            connector = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
            connector.begin_connect(shape, 2)     # Anchored to Bottom of Parent
            connector.end_connect(child_shape, 0) # Anchored to Top of Child
            connector.line.color.rgb = RGBColor(170, 175, 180)
            connector.line.width = Pt(1.5)
            
        return shape

    # Execute Drawing
    for root in root_nodes:
        draw_tree(root, slide)

    prs.save(output_pptx_path)
    return output_pptx_path
