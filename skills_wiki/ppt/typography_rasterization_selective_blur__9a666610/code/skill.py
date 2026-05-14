import os
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from lxml import etree

def generate_text_image(text, font_size=250, text_color=(0, 0, 0, 255)):
    """Helper to render text to a transparent PIL Image with tight cropping."""
    # Attempt to load a heavy font, fallback to default
    try:
        font = ImageFont.truetype("arialbd.ttf", font_size)
    except IOError:
        try:
            font = ImageFont.truetype("impact.ttf", font_size)
        except IOError:
            font = ImageFont.load_default()
            
    # Create a temporary large canvas
    temp_img = Image.new("RGBA", (2000, 1000), (255, 255, 255, 0))
    draw = ImageDraw.Draw(temp_img)
    
    # Draw text
    draw.text((100, 100), text, font=font, fill=text_color)
    
    # Get bounding box and crop tightly
    bbox = temp_img.getbbox()
    if bbox:
        return temp_img.crop(bbox)
    return temp_img

def apply_3d_rotation(shape, preset="isometricTopUp"):
    """
    Injects Open XML to apply native PowerPoint 3D Isometric rotation to a picture shape.
    """
    spPr = shape.element.spPr
    # Create scene3d element
    scene3d = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}scene3d')
    # Add camera with specific preset (e.g., isometricTopUp, isometricRightUp)
    camera = etree.SubElement(scene3d, '{http://schemas.openxmlformats.org/drawingml/2006/main}camera', prst=preset)
    # Add standard light rig
    lightRig = etree.SubElement(scene3d, '{http://schemas.openxmlformats.org/drawingml/2006/main}lightRig', rig="threePt", dir="t")


def create_slide(
    output_pptx_path: str,
    title_text: str = "FOCUS",
    bend_text: str = "BEND",
    accent_color: tuple = (255, 0, 127, 255),  # Hot Pink
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Typography Blur and 3D Split effects.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ==========================================
    # SLIDE 1: Selective Blur Effect ("FOCUS")
    # ==========================================
    slide_blur = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 1. Generate base text image
    img_sharp = generate_text_image(title_text, font_size=300)
    
    # 2. Create blurred version
    img_blur = img_sharp.copy().filter(ImageFilter.GaussianBlur(radius=15))
    
    # 3. Composite: Center sharp, edges blurred. 
    # We'll cut the middle out of the blurred image and replace it with the sharp image.
    width, height = img_sharp.size
    composite = img_blur.copy()
    
    # Define a sharp region in the middle (approx 40% of the word)
    left_bound = int(width * 0.3)
    right_bound = int(width * 0.7)
    
    sharp_crop = img_sharp.crop((left_bound, 0, right_bound, height))
    composite.paste(sharp_crop, (left_bound, 0), sharp_crop)
    
    blur_img_path = "temp_focus_blur.png"
    composite.save(blur_img_path)
    
    # Add to slide
    pic_width = Inches(10)
    pic_height = pic_width * (height / width)
    left = (prs.slide_width - pic_width) / 2
    top = (prs.slide_height - pic_height) / 2
    slide_blur.shapes.add_picture(blur_img_path, left, top, pic_width, pic_height)


    # ==========================================
    # SLIDE 2: 3D Bisected Bend Effect ("BEND")
    # ==========================================
    slide_bend = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 1. Generate base text image
    img_base = generate_text_image(bend_text, font_size=400)
    bw, bh = img_base.size
    
    # 2. Split image horizontally
    mid_y = int(bh / 2)
    img_top = img_base.crop((0, 0, bw, mid_y))
    img_bot = img_base.crop((0, mid_y, bw, bh))
    
    # 3. Colorize bottom half using accent color
    # Extract alpha channel, create a solid color block, and mask it with the alpha
    alpha_bot = img_bot.getchannel("A")
    color_block = Image.new("RGBA", img_bot.size, accent_color)
    color_block.putalpha(alpha_bot)
    img_bot_colored = color_block
    
    # Save partial images
    top_path = "temp_bend_top.png"
    bot_path = "temp_bend_bot.png"
    img_top.save(top_path)
    img_bot_colored.save(bot_path)
    
    # Add Top Half to Slide and apply 3D
    base_pic_width = Inches(8)
    # Proportional heights for the halves
    top_pic_height = base_pic_width * (img_top.size[1] / bw)
    bot_pic_height = base_pic_width * (img_bot_colored.size[1] / bw)
    
    center_x = (prs.slide_width - base_pic_width) / 2
    center_y = (prs.slide_height - (top_pic_height + bot_pic_height)) / 2
    
    # Position logic (adjusted slightly to allow for 3D overlap visual)
    pic_top = slide_bend.shapes.add_picture(top_path, center_x - Inches(0.5), center_y - Inches(0.5), base_pic_width, top_pic_height)
    apply_3d_rotation(pic_top, preset="isometricTopUp")
    
    pic_bot = slide_bend.shapes.add_picture(bot_path, center_x + Inches(0.5), center_y + Inches(0.2), base_pic_width, bot_pic_height)
    apply_3d_rotation(pic_bot, preset="isometricRightUp")

    # Save presentation
    prs.save(output_pptx_path)
    
    # Cleanup temp files
    for p in [blur_img_path, top_path, bot_path]:
        if os.path.exists(p):
            os.remove(p)
            
    return output_pptx_path

if __name__ == "__main__":
    create_slide("Typography_Manipulation_Effects.pptx")
