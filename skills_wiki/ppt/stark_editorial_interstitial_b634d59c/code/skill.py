def create_slide(
    output_pptx_path: str,
    title_text: str = "AGENDA",
    eyebrow_text: str = "TYPICAL CLIENT MEETING",
    subtitle_text: str = "A GUIDE TO ENSURE YOU DON'T LEAVE ANY QUESTION UNANSWERED",
    bg_keyword: str = "architecture",
) -> str:
    """
    Create a PPTX file reproducing the 'Stark Editorial Interstitial' visual effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageFilter, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Helper function to inject letter spacing (tracking) via lxml
    def add_tracking(run, pt_spacing):
        """Add architectural letter spacing to a text run."""
        rPr = run._r.get_or_add_rPr()
        # OOXML spc attribute is in 1/100ths of a point
        spc_val = int(pt_spacing * 100)
        rPr.set('spc', str(spc_val))

    # === Layer 1: Cinematic Background (PIL) ===
    bg_path = "temp_cinematic_bg.jpg"
    try:
        # Fetch a relevant background image
        url = f"https://source.unsplash.com/random/1920x1080/?{bg_keyword},interior"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read()))
    except Exception:
        # Fallback to a solid dark base if download fails
        img = Image.new('RGB', (1920, 1080), (30, 35, 40))

    # Apply heavy cinematic blur
    img = img.filter(ImageFilter.GaussianBlur(radius=12))
    
    # Apply dark overlay for text contrast (Alpha compositing)
    img = img.convert("RGBA")
    dark_overlay = Image.new('RGBA', img.size, (10, 12, 15, 140)) # Deep dark grey, ~55% opacity
    img = Image.alpha_composite(img, dark_overlay)
    img = img.convert("RGB")
    
    img.save(bg_path, format="JPEG", quality=90)
    
    # Insert processed background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Stark Typography Layout ===
    
    left_margin = Inches(1.5)
    
    # 1. Eyebrow Text
    tx_eyebrow = slide.shapes.add_textbox(left_margin, Inches(2.3), Inches(10), Inches(0.5))
    p_eye = tx_eyebrow.text_frame.paragraphs[0]
    run_eye = p_eye.add_run()
    run_eye.text = eyebrow_text.upper()
    run_eye.font.name = "Arial"
    run_eye.font.size = Pt(12)
    run_eye.font.bold = True
    run_eye.font.color.rgb = RGBColor(180, 180, 180)
    add_tracking(run_eye, 5) # Heavy tracking for technical feel

    # 2. Main Title
    tx_title = slide.shapes.add_textbox(left_margin - Inches(0.05), Inches(2.6), Inches(11), Inches(1.5))
    p_title = tx_title.text_frame.paragraphs[0]
    run_title = p_title.add_run()
    run_title.text = title_text.upper()
    run_title.font.name = "Arial"
    run_title.font.size = Pt(88)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255)
    # No tracking on the main title, keep it heavy and dense

    # 3. Subtitle
    tx_sub = slide.shapes.add_textbox(left_margin, Inches(4.3), Inches(10), Inches(0.5))
    p_sub = tx_sub.text_frame.paragraphs[0]
    run_sub = p_sub.add_run()
    run_sub.text = subtitle_text.upper()
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = False
    run_sub.font.color.rgb = RGBColor(210, 210, 210)
    add_tracking(run_sub, 3) # Moderate tracking

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
