def create_slide(
    output_pptx_path: str,
    title_text: str = "Project Roadmap",
    steps_data: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Precision Corporate Alternating Vertical Timeline.
    """
    import copy
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree

    # Default data if none provided
    if not steps_data:
        steps_data = [
            {"title": "Project Kickoff", "desc": "Define the vision, core objectives, and assemble the primary team."},
            {"title": "Market Analysis", "desc": "Conduct competitor research and identify key target demographics."},
            {"title": "Product Development", "desc": "Iterative design and engineering phases to build the MVP."},
            {"title": "Beta Launch", "desc": "Release to a closed group of early adopters for feedback and QA."}
        ]

    # Corporate Color Palette (Teal, Coral, Green, Blue)
    colors = [
        RGBColor(0, 168, 143),
        RGBColor(242, 108, 79),
        RGBColor(136, 195, 64),
        RGBColor(0, 114, 188)
    ]
    
    charcoal = RGBColor(50, 50, 50)
    gray = RGBColor(200, 200, 200)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Helper function to inject XML drop shadows
    def add_premium_shadow(shape):
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="50800" dist="38100" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="25000"/>
                </a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        shadow_element = etree.fromstring(shadow_xml)
        shape.element.spPr.append(shadow_element)

    # --- Title ---
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text.upper()
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = charcoal
    p.alignment = PP_ALIGN.CENTER

    # --- Layout Calculations ---
    center_x = 13.333 / 2
    top_margin = 1.8
    bottom_margin = 6.8
    total_height = bottom_margin - top_margin
    num_steps = len(steps_data)
    
    # Calculate spacing (avoid division by zero if 1 step)
    y_spacing = total_height / (num_steps - 1) if num_steps > 1 else 0
    node_radius = 0.35

    # --- Draw Central Spine ---
    spine = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(center_x - 0.02), Inches(top_margin), 
        Inches(0.04), Inches(total_height)
    )
    spine.fill.solid()
    spine.fill.fore_color.rgb = gray
    spine.line.fill.background()

    # --- Draw Steps ---
    for i, step in enumerate(steps_data):
        is_left = i % 2 == 0
        current_y = top_margin + (i * y_spacing)
        color = colors[i % len(colors)]

        # 1. Connecting Stem (Horizontal line from spine to node)
        stem_width = 0.8
        stem_start_x = center_x - stem_width if is_left else center_x
        stem = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(stem_start_x), Inches(current_y - 0.015),
            Inches(stem_width), Inches(0.03)
        )
        stem.fill.solid()
        stem.fill.fore_color.rgb = gray
        stem.line.fill.background()

        # 2. Text Box
        box_width = 4.5
        box_height = 1.2
        box_x = center_x - stem_width - box_width - 0.2 if is_left else center_x + stem_width + 0.2
        box_y = current_y - (box_height / 2) + 0.1 # slight optical adjustment
        
        tx_box = slide.shapes.add_textbox(Inches(box_x), Inches(box_y), Inches(box_width), Inches(box_height))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        # Title paragraph
        p_title = tf.paragraphs[0]
        p_title.text = step["title"]
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = color
        p_title.alignment = PP_ALIGN.RIGHT if is_left else PP_ALIGN.LEFT
        
        # Description paragraph
        p_desc = tf.add_paragraph()
        p_desc.text = step["desc"]
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = charcoal
        p_desc.alignment = PP_ALIGN.RIGHT if is_left else PP_ALIGN.LEFT

        # 3. Circular Node (Drawn last so it sits on top of lines)
        node_x = center_x - stem_width - node_radius if is_left else center_x + stem_width - node_radius
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(node_x), Inches(current_y - node_radius), 
            Inches(node_radius * 2), Inches(node_radius * 2)
        )
        
        # Node Styling
        node.fill.solid()
        node.fill.fore_color.rgb = RGBColor(255, 255, 255)
        node.line.color.rgb = color
        node.line.width = Pt(4)
        add_premium_shadow(node) # Inject XML shadow

        # Node Text (Number)
        node_tf = node.text_frame
        node_tf.margin_left = 0
        node_tf.margin_right = 0
        node_tf.margin_top = 0
        node_tf.margin_bottom = 0
        p_num = node_tf.paragraphs[0]
        p_num.text = f"{i+1:02d}"
        p_num.font.size = Pt(16)
        p_num.font.bold = True
        p_num.font.color.rgb = color
        p_num.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
