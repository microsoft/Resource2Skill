def create_slide(
    output_pptx_path: str,
    title_text: str = "Creative Presentations",
    body_text: str = "",
    bg_palette: str = "technology",  
    accent_color: tuple = (30, 30, 30),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Realistic Typewriter Rhythm Morph' visual effect.
    Generates a sequence of slides with micro-transitions to simulate human typing.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from lxml import etree
    import random

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # === Step 1: Generate the Human Typing Script ===
    # Format: (Current String, Cursor Visible Boolean, Delay to Next Slide in ms)
    script = []
    
    # Initial cursor blinking sequence
    script.append(("", True, 400))
    script.append(("", False, 400))
    script.append(("", True, 400))
    
    current_text = ""
    random.seed(42) # Seed for predictable reproducibility, but allows organic logic
    
    for i, char in enumerate(title_text):
        # 15% chance to make a human typo (but not on spaces or the very first character)
        if i > 0 and char != " " and random.random() < 0.15:
            wrong_char = random.choice("abcdefghijklmnopqrstuvwxyz")
            if wrong_char != char.lower():
                script.append((current_text + wrong_char, True, 180)) # Type wrong char
                script.append((current_text + wrong_char, True, 600)) # Pause: realizing mistake
                script.append((current_text, True, 250))              # Backspace correction
                script.append((current_text, True, 300))              # Pause before continuing
                
        current_text += char
        
        # Variable typing speed (spaces take slightly longer)
        delay = random.randint(80, 180)
        if char == " ":
            delay += 120
            
        script.append((current_text, True, delay))
        
    # Final sequence: Stop typing, blink a few times, and hold
    script.append((current_text, False, 500))
    script.append((current_text, True, 500))
    script.append((current_text, False, 500))
    script.append((current_text, False, 3000)) # Final hold
    
    # XML Namespace for Slide Manipulation
    nsmap = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    
    # === Step 2: Build Frames (Slides) based on Script ===
    for text_state, cursor_visible, delay_ms in script:
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
        
        # 1. Canvas Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(250, 250, 250)
        
        # 2. Main Search Bar Outline
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(2.5), Inches(3.25), Inches(8.33), Inches(1.0)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(255, 255, 255)
        bar.line.color.rgb = RGBColor(50, 50, 50)
        bar.line.width = Pt(2)
        
        # 3. Right Side "Search" Button 
        # Using overlapping shapes to achieve rounded right corners with flat left edges
        btn_rrect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(9.5), Inches(3.25), Inches(1.33), Inches(1.0)
        )
        btn_rrect.fill.solid()
        btn_rrect.fill.fore_color.rgb = RGBColor(*accent_color)
        btn_rrect.line.color.rgb = RGBColor(*accent_color)
        
        btn_mask = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(9.5), Inches(3.25), Inches(0.5), Inches(1.0)
        )
        btn_mask.fill.solid()
        btn_mask.fill.fore_color.rgb = RGBColor(*accent_color)
        btn_mask.line.color.rgb = RGBColor(*accent_color)
        
        # 4. Search Icon Assembly (Circle + Connector Line)
        cx, cy = 10.165, 3.75
        radius = 0.15
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx - radius), Inches(cy - radius), Inches(radius*2), Inches(radius*2)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(*accent_color) # Match button BG to look hollow
        circle.line.color.rgb = RGBColor(255, 255, 255)
        circle.line.width = Pt(3)
        
        handle = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(cx + radius * 0.7), Inches(cy + radius * 0.7),
            Inches(cx + radius * 1.6), Inches(cy + radius * 1.6)
        )
        handle.line.color.rgb = RGBColor(255, 255, 255)
        handle.line.width = Pt(3)
        
        # 5. Typewriter Text Setup
        display_text = text_state + ("|" if cursor_visible else "")
        text_box = slide.shapes.add_textbox(
            Inches(2.8), Inches(3.35), Inches(6.5), Inches(0.8)
        )
        tf2 = text_box.text_frame
        tf2.word_wrap = False
        tf2.text = display_text
        p2 = tf2.paragraphs[0]
        p2.font.name = "Consolas" # Crucial monospace font for typewriter feel
        p2.font.size = Pt(28)
        p2.font.color.rgb = RGBColor(*accent_color)
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        # 6. XML Injection: Slide Transition Auto-Advance
        # This writes `<p:transition spd="fast" advTm="{delay_ms}"/>` to drive the animation
        sld = slide._element
        transition = etree.Element('{http://schemas.openxmlformats.org/presentationml/2006/main}transition')
        transition.set('advTm', str(delay_ms))
        transition.set('spd', 'fast')
        
        timing = sld.find('./p:timing', namespaces=nsmap)
        extLst = sld.find('./p:extLst', namespaces=nsmap)
        
        if timing is not None:
            timing.addprevious(transition)
        elif extLst is not None:
            extLst.addprevious(transition)
        else:
            sld.append(transition)
            
    prs.save(output_pptx_path)
    return output_pptx_path
