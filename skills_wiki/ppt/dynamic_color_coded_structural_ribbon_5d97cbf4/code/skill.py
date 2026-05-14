def create_slide(
    output_pptx_path: str,
    active_section_idx: int = 2,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Dynamic Color-Coded Structural Ribbon' visual effect.
    
    Args:
        output_pptx_path: Path to save the PPTX file.
        active_section_idx: Integer index of the currently active section (0-based).
        
    Returns: 
        path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Customizable section data: (Title, Background RGB, Text RGB)
    sections = kwargs.get("sections", [
        ("Intro", (176, 180, 184), (50, 50, 50)),
        ("Status Quo", (235, 87, 110), (255, 255, 255)),
        ("Product", (66, 133, 244), (255, 255, 255)),
        ("Market", (52, 168, 83), (255, 255, 255)),
        ("Why Us", (251, 188, 5), (50, 50, 50)),
        ("Ask", (32, 51, 72), (255, 255, 255))
    ])
    
    # Clamp active index
    active_section_idx = min(max(active_section_idx, 0), len(sections) - 1)
    active_name, active_bg, active_fg = sections[active_section_idx]

    # --- Layer 1: Background ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 247)  # Ultra-light premium gray

    # --- XML Shadow Helper ---
    def add_shadow(shape, blur_rad="300000", dist="40000", direction="5400000", alpha="10000"):
        spPr = shape._element.spPr
        # Remove any existing effect list
        for e in spPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst"):
            spPr.remove(e)
            
        effectLst = etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        outerShdw = etree.SubElement(effectLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw",
                                     blurRad=blur_rad, dist=dist, dir=direction, algn="ctr", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", val="000000")
        etree.SubElement(srgbClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", val=alpha)
        
        # Insert before extLst if present, otherwise append
        extLst = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}extLst")
        if extLst is not None:
            extLst.addprevious(effectLst)
        else:
            spPr.append(effectLst)

    # --- Layer 2: Main Canvas Content ---
    # Section Header Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.33), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = active_name.upper()
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*active_bg)

    # Elevated White Content Card
    content_rect = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.2)
    )
    content_rect.fill.solid()
    content_rect.fill.fore_color.rgb = RGBColor(255, 255, 255)
    content_rect.line.color.rgb = RGBColor(255, 255, 255)  # Hide border
    add_shadow(content_rect, blur_rad="300000", dist="40000", direction="5400000", alpha="12000")

    # Content Placeholder Text
    tf = content_rect.text_frame
    tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.6)
    p = tf.paragraphs[0]
    p.text = f"{active_name} Overview"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(40, 40, 40)
    p.alignment = PP_ALIGN.LEFT

    # Mock Data Bullets
    bullets = [
        "First primary directive or metric belonging to this section.", 
        "Second piece of vital structural context.",
        "Smooth transition data point leading into the specific details."
    ]
    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(120, 120, 120)
        p.level = 1
        p.space_before = Pt(14)

    # --- Layer 3: The Bottom Navigation Ribbon ---
    N = len(sections)
    W_emu = int(prs.slide_width)
    
    # Calculate widths in EMUs to prevent float rounding gaps
    # The active tab is dynamically sized to be 2x wider than inactive tabs
    w_in_emu = int(W_emu / (N + 1.0))
    w_act_emu = W_emu - (N - 1) * w_in_emu 

    current_x_emu = 0
    active_shape = None

    for i, (name, bg, fg) in enumerate(sections):
        is_active = (i == active_section_idx)
        width_emu = w_act_emu if is_active else w_in_emu
        height_emu = int(Inches(0.8)) if is_active else int(Inches(0.5))
        y_emu = int(prs.slide_height) - height_emu
        
        # Generate the ribbon segment
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            current_x_emu, y_emu, width_emu, height_emu
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*bg)
        shape.line.color.rgb = RGBColor(*bg) # Seamless border matching
        
        # Configure text
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        tf.margin_left = 0
        tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = name
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = RGBColor(*fg)
        p.font.size = Pt(14) if is_active else Pt(11)
        p.font.bold = is_active
        
        if is_active:
            active_shape = shape
            
        current_x_emu += width_emu

    # Make the active ribbon tab pop outwards with an omnidirectional glow/shadow
    if active_shape:
        add_shadow(active_shape, blur_rad="150000", dist="0", direction="0", alpha="25000")

    prs.save(output_pptx_path)
    return output_pptx_path
