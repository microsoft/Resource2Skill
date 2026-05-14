import random
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_slide(
    output_pptx_path: str,
    title_text: str = "ISTA 2016",
    subtitle_text: str = "LEARN. INSPIRE. GEEK OUT.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Fragmented Data Stream' (Tech Glitch Bar) visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Tech Conference Palette (Purple, Magenta, Orange, Navy, Light Gray)
    palette = [
        RGBColor(81, 35, 120),   # Purple
        RGBColor(216, 27, 96),   # Magenta
        RGBColor(255, 152, 0),   # Orange
        RGBColor(26, 35, 126),   # Navy
        RGBColor(158, 158, 158)  # Gray accent
    ]

    # === Layer 1: Typography (Top Center) ===
    # Title
    title_box = slide.shapes.add_textbox(Inches(2), Inches(2.0), Inches(9.333), Inches(1.5))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.name = "Arial Black"
    run.font.size = Pt(88)
    run.font.color.rgb = RGBColor(81, 35, 120)  # Brand Purple

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(2), Inches(3.5), Inches(9.333), Inches(0.5))
    tf_sub = sub_box.text_frame
    tf_sub.clear()
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.alignment = PP_ALIGN.CENTER
    run_sub = p_sub.runs[0]
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(20)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    # === Layer 2: The Fragmented Data Stream Bar (Bottom anchor) ===
    # Parameters for the procedural generation
    band_start_y = 4.8  # Start Y in inches
    row_height = 0.15   # Height of each rectangle row in inches
    num_rows = 9        # Number of horizontal tracks
    
    base_start_x = 3.0  # Where the cluster roughly begins
    base_end_x = 10.333 # Where the cluster roughly ends

    for row in range(num_rows):
        current_y = band_start_y + (row * (row_height + 0.05))  # 0.05 is vertical gap
        
        # Add random start/end jitter to make the edges look jagged
        row_start_x = base_start_x + random.uniform(-1.0, 1.0)
        row_end_x = base_end_x + random.uniform(-1.0, 1.0)
        
        current_x = row_start_x
        
        while current_x < row_end_x:
            # Generate random width for this specific fragment
            rect_width = random.uniform(0.3, 2.5)
            
            # Ensure it doesn't vastly overshoot the bounds
            if current_x + rect_width > row_end_x + 0.5:
                rect_width = row_end_x - current_x
                if rect_width <= 0:
                    break
            
            # Create the rectangle
            rect = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(current_x),
                Inches(current_y),
                Inches(rect_width),
                Inches(row_height)
            )
            
            # Style the rectangle (No border, random palette fill)
            rect.fill.solid()
            rect.fill.fore_color.rgb = random.choice(palette)
            rect.line.fill.background()  # Make outline transparent
            
            # Advance X position. 
            # Random factor between -0.1 (overlap) and 0.4 (gap)
            current_x += rect_width + random.uniform(-0.1, 0.4)

    # === Layer 3: Extra floating "Glitch" fragments for dynamic effect ===
    for _ in range(15):
        # Place these randomly around the main block
        extra_x = random.uniform(2.0, 11.0)
        extra_y = random.uniform(band_start_y - 0.5, band_start_y + (num_rows * 0.2) + 0.5)
        extra_width = random.uniform(0.1, 0.8)
        
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(extra_x),
            Inches(extra_y),
            Inches(extra_width),
            Inches(0.08) # Thinner accent lines
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = random.choice(palette)
        rect.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
