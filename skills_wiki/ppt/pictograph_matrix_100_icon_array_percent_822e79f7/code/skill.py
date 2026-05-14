import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    percentage: int = 35,
    title_text: str = "OF CAMPERS\nDON'T LIKE\nSMORES*",
    disclaimer_text: str = "*According to a recent survey",
    bg_color: tuple = (97, 244, 222),      # Mint Turquoise
    highlight_color: tuple = (13, 31, 45), # Deep Navy
    base_color: tuple = (255, 255, 255),   # Pure White
) -> str:
    """
    Creates a PPTX file reproducing the Pictograph Matrix (100-Icon Array) Data Visualization.
    
    Returns: path to the saved PPTX file.
    """
    # 1. Initialize Presentation (16:9 widescreen)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 2. Set Background Color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # 3. Generate the 10x10 Icon Matrix using PIL
    # We will draw a clean, minimalist "Person" vector-style icon programmatically.
    grid_size = 10
    cell_size = 100
    img_size = grid_size * cell_size
    
    # Create transparent canvas
    matrix_img = Image.new("RGBA", (img_size, img_size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(matrix_img)
    
    # Ensure percentage is bounded
    percentage = max(0, min(100, int(percentage)))
    
    for row in range(grid_size):
        for col in range(grid_size):
            # Calculate index to determine color
            # The video shows icons colored top-to-bottom, left-to-right, but 
            # standard infographics fill left-to-right, top-to-bottom. We use standard.
            idx = row * grid_size + col
            current_color = highlight_color if idx < percentage else base_color
            color_rgba = current_color + (255,) # Add full alpha
            
            # Calculate base coordinates for this cell
            cx = col * cell_size + (cell_size / 2)
            cy = row * cell_size + (cell_size / 2)
            
            # Draw "Person" Icon
            # Head (Circle)
            head_radius = 14
            head_y = cy - 25
            draw.ellipse(
                [cx - head_radius, head_y - head_radius, cx + head_radius, head_y + head_radius],
                fill=color_rgba
            )
            
            # Body (Trapezoid / rounded block approximation)
            body_top = head_y + head_radius + 4
            body_bottom = cy + 30
            body_width_top = 26
            body_width_bottom = 34
            
            draw.polygon([
                (cx - body_width_top/2, body_top),
                (cx + body_width_top/2, body_top),
                (cx + body_width_bottom/2, body_bottom),
                (cx - body_width_bottom/2, body_bottom)
            ], fill=color_rgba)

    # Save temporary image
    temp_img_path = "temp_pictograph_matrix.png"
    matrix_img.save(temp_img_path)

    # 4. Insert Matrix Image into PPTX
    # Placed on the left side, vertically centered
    pic_left = Inches(1.5)
    pic_top = Inches(1.0)
    pic_size = Inches(5.5)
    slide.shapes.add_picture(temp_img_path, pic_left, pic_top, width=pic_size, height=pic_size)

    # 5. Add Typography
    # 5a. The Giant Percentage Number
    tx_number = slide.shapes.add_textbox(Inches(7.5), Inches(1.5), Inches(5.0), Inches(1.5))
    tf_number = tx_number.text_frame
    p_num = tf_number.paragraphs[0]
    p_num.text = f"{percentage}%"
    p_num.font.size = Pt(110)
    p_num.font.bold = True
    p_num.font.name = "Arial Black" # Fallback to a universally available bold sans-serif
    p_num.font.color.rgb = RGBColor(*highlight_color)
    
    # 5b. The Context Text
    tx_context = slide.shapes.add_textbox(Inches(7.6), Inches(3.5), Inches(5.0), Inches(2.5))
    tf_context = tx_context.text_frame
    tf_context.word_wrap = True
    p_ctx = tf_context.paragraphs[0]
    p_ctx.text = title_text
    p_ctx.font.size = Pt(36)
    p_ctx.font.bold = True
    p_ctx.font.name = "Arial Black"
    p_ctx.font.color.rgb = RGBColor(*highlight_color)
    # Tweak line spacing for punchy visual effect
    p_ctx.line_spacing = 0.9 

    # 5c. The Disclaimer Text
    tx_disc = slide.shapes.add_textbox(Inches(7.6), Inches(6.0), Inches(5.0), Inches(0.5))
    tf_disc = tx_disc.text_frame
    p_disc = tf_disc.paragraphs[0]
    p_disc.text = disclaimer_text
    p_disc.font.size = Pt(12)
    p_disc.font.italic = True
    p_disc.font.name = "Arial"
    p_disc.font.color.rgb = RGBColor(*highlight_color)

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
        
    return output_pptx_path

# Example execution:
# create_slide("pictograph_percentage.pptx", percentage=35)
