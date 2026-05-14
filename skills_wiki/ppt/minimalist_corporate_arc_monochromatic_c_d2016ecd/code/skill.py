def create_slide(
    output_pptx_path: str,
    title_text: str = "CREATIVE BUSINESS PLAN",
    subtitle_text: str = "Your Favorite Topic to Discuss",
    bg_keyword: str = "architecture",
    accent_color: tuple = (186, 178, 172),  # Warm Stone/Beige
    dark_color: tuple = (38, 38, 42),       # Charcoal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Minimalist Corporate Arc" visual effect.
    Generates a 2-slide deck: a Title Arc slide and a 50/50 Split Content slide.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageEnhance

    # Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Helper: Download and process image (Grayscale + Darken)
    def get_processed_bg(keyword, filename, darken_factor=0.4):
        url = f"https://source.unsplash.com/random/1920x1080/?{keyword}"
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req) as response:
                with open(filename, 'wb') as f:
                    f.write(response.read())
            
            # Process with PIL
            with Image.open(filename) as img:
                img = img.convert('L')  # Convert to Grayscale
                enhancer = ImageEnhance.Brightness(img)
                img = enhancer.enhance(darken_factor)  # Darken
                img.save(filename)
            return True
        except Exception as e:
            print(f"Failed to fetch image: {e}")
            # Create a solid dark fallback image if download fails
            img = Image.new('RGB', (1920, 1080), dark_color)
            img.save(filename)
            return True

    # Colors
    c_accent = RGBColor(*accent_color)
    c_dark = RGBColor(*dark_color)
    c_white = RGBColor(255, 255, 255)

    # ==========================================
    # SLIDE 1: Title with Sweeping Bottom Arc
    # ==========================================
    slide_1 = prs.slides.add_slide(blank_layout)
    bg_file_1 = "temp_bg_1.jpg"
    get_processed_bg(bg_keyword, bg_file_1, darken_factor=0.35)

    # 1. Background Image
    slide_1.shapes.add_picture(bg_file_1, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # 2. Bottom Sweeping Arc
    # To create a gentle wave, we use a massive oval pushed below the slide edge
    arc_width = Inches(28)
    arc_height = Inches(10)
    arc_left = (prs.slide_width - arc_width) / 2
    arc_top = prs.slide_height - Inches(1.8) # How much the arc peeks above the bottom
    
    arc = slide_1.shapes.add_shape(MSO_SHAPE.OVAL, arc_left, arc_top, arc_width, arc_height)
    arc.fill.solid()
    arc.fill.fore_color.rgb = c_accent
    arc.line.color.rgb = c_white
    arc.line.width = Pt(3)

    # 3. Hexagon Logo Mark
    logo_size = Inches(0.8)
    logo_left = (prs.slide_width - logo_size) / 2
    logo_top = Inches(1.5)
    logo = slide_1.shapes.add_shape(MSO_SHAPE.HEXAGON, logo_left, logo_top, logo_size, logo_size)
    logo.fill.background() # transparent
    logo.line.color.rgb = c_white
    logo.line.width = Pt(2)
    
    logo_tf = logo.text_frame
    logo_p = logo_tf.paragraphs[0]
    logo_p.text = "b"
    logo_p.alignment = PP_ALIGN.CENTER
    logo_p.font.size = Pt(24)
    logo_p.font.color.rgb = c_white
    logo_p.font.bold = True

    # 4. Typography
    # Title
    title_box = slide_1.shapes.add_textbox(0, Inches(2.8), prs.slide_width, Inches(1))
    title_tf = title_box.text_frame
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text.upper()
    title_p.alignment = PP_ALIGN.CENTER
    title_p.font.size = Pt(44)
    title_p.font.color.rgb = c_white
    title_p.font.name = "Arial"
    title_p.font.bold = True

    # Subtitle
    sub_box = slide_1.shapes.add_textbox(0, Inches(3.7), prs.slide_width, Inches(0.5))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = subtitle_text
    sub_p.alignment = PP_ALIGN.CENTER
    sub_p.font.size = Pt(20)
    sub_p.font.color.rgb = c_white
    sub_p.font.name = "Arial"

    # ==========================================
    # SLIDE 2: 50/50 Monochromatic Split
    # ==========================================
    slide_2 = prs.slides.add_slide(blank_layout)
    bg_file_2 = "temp_bg_2.jpg"
    get_processed_bg("cityscape", bg_file_2, darken_factor=0.6)

    # 1. Split Left (Beige)
    left_pane = slide_2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width / 2, prs.slide_height)
    left_pane.fill.solid()
    left_pane.fill.fore_color.rgb = c_accent
    left_pane.line.fill.background()

    # 2. Split Right (Image)
    slide_2.shapes.add_picture(bg_file_2, prs.slide_width / 2, 0, width=prs.slide_width / 2, height=prs.slide_height)

    # 3. Left Content (Dark Text on Beige)
    content_box = slide_2.shapes.add_textbox(Inches(1), Inches(2), Inches(4.5), Inches(3))
    content_tf = content_box.text_frame
    
    cp1 = content_tf.paragraphs[0]
    cp1.text = "Our vision"
    cp1.font.size = Pt(36)
    cp1.font.color.rgb = c_dark
    cp1.font.name = "Arial"
    
    cp2 = content_tf.add_paragraph()
    cp2.text = "OVERVIEW\n"
    cp2.font.size = Pt(12)
    cp2.font.color.rgb = c_dark
    
    cp3 = content_tf.add_paragraph()
    cp3.text = "Therefore it may be the case that issues arise, i.e. issues with national postal services, international customs issues etc."
    cp3.font.size = Pt(14)
    cp3.font.color.rgb = RGBColor(80, 80, 85)

    # 4. Floating Circular Accent on Image
    circle_size = Inches(2.5)
    circle = slide_2.shapes.add_shape(MSO_SHAPE.OVAL, prs.slide_width / 2 + Inches(2), prs.slide_height / 2 - Inches(1.25), circle_size, circle_size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = c_dark
    circle.line.fill.background()

    # Cleanup temporary images
    prs.save(output_pptx_path)
    
    if os.path.exists(bg_file_1): os.remove(bg_file_1)
    if os.path.exists(bg_file_2): os.remove(bg_file_2)

    return output_pptx_path
