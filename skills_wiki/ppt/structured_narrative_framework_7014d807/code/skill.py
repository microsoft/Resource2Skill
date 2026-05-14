def create_structured_narrative_deck(
    output_pptx_path: str,
    title_text: str = "Your Presentation Title",
    subtitle_text: str = "Engaging and memorable subtitle",
    presenter_name: str = "Your Name",
    opening_hook: str = "Start with a powerful question or a startling statistic to grab attention.",
    section_titles: list = ["First Key Point", "Second Key Point", "Third Key Point"],
    closing_message: str = "End with a strong, actionable take-home message.",
    **kwargs,
) -> str:
    """
    Creates a PPTX file based on the Structured Narrative Framework,
    embodying the design principles of clarity, structure, and narrative flow.

    Args:
        output_pptx_path: The path to save the generated .pptx file.
        title_text: The main title for the presentation.
        subtitle_text: The subtitle for the title slide.
        presenter_name: The name of the presenter.
        opening_hook: The text for the strong opening slide.
        section_titles: A list of strings for the main section header slides.
        closing_message: The final, memorable message for the closing slide.

    Returns:
        The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    # --- Color Palette ---
    BG_COLOR = RGBColor(34, 40, 49)
    TEXT_COLOR = RGBColor(238, 238, 238)
    ACCENT_COLOR = RGBColor(0, 173, 181)

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    blank_layout = prs.slide_layouts[6]

    # --- Helper to set background ---
    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        return

    # --- Slide 1: Title Slide ---
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide)

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(14), Inches(2))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.font.bold = True
    title_p.font.size = Pt(54)
    title_p.font.color.rgb = TEXT_COLOR
    title_p.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(14), Inches(1))
    subtitle_p = subtitle_box.text_frame.paragraphs[0]
    subtitle_p.text = subtitle_text
    subtitle_p.font.size = Pt(24)
    subtitle_p.font.color.rgb = TEXT_COLOR
    subtitle_p.alignment = PP_ALIGN.CENTER
    
    presenter_box = slide.shapes.add_textbox(Inches(1), Inches(7.5), Inches(14), Inches(1))
    presenter_p = presenter_box.text_frame.paragraphs[0]
    presenter_p.text = presenter_name
    presenter_p.font.size = Pt(18)
    presenter_p.font.color.rgb = TEXT_COLOR
    presenter_p.alignment = PP_ALIGN.CENTER

    # --- Slide 2: The Opening Hook ---
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide)
    
    hook_box = slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(13), Inches(3))
    hook_p = hook_box.text_frame.paragraphs[0]
    hook_p.text = opening_hook
    hook_p.font.italic = True
    hook_p.font.size = Pt(32)
    hook_p.font.color.rgb = ACCENT_COLOR
    hook_p.alignment = PP_ALIGN.CENTER

    # --- Content Slides ---
    for title in section_titles:
        slide = prs.slides.add_slide(blank_layout)
        set_background(slide)

        header_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(14), Inches(1.0))
        header_p = header_box.text_frame.paragraphs[0]
        header_p.text = title
        header_p.font.bold = True
        header_p.font.size = Pt(36)
        header_p.font.color.rgb = ACCENT_COLOR
        
        line = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(1), Inches(1.5), Inches(5), Inches(0))
        line.line.color.rgb = ACCENT_COLOR
        line.line.width = Pt(2)
        
        content_box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(14), Inches(5.5))
        content_tf = content_box.text_frame
        content_tf.word_wrap = True
        
        p1 = content_tf.paragraphs[0]
        p1.text = "Use concise language here."
        p1.font.size = Pt(22)
        p1.font.color.rgb = TEXT_COLOR
        p1.level = 0
        
        p2 = content_tf.add_paragraph()
        p2.text = "Slides should support your message, not replace it."
        p2.font.size = Pt(20)
        p2.font.color.rgb = TEXT_COLOR
        p2.level = 1

    # --- Slide N: The Closing Message ---
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide)
    
    closing_box = slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(13), Inches(3))
    closing_p = closing_box.text_frame.paragraphs[0]
    closing_p.text = closing_message
    closing_p.font.bold = True
    closing_p.font.size = Pt(32)
    closing_p.font.color.rgb = TEXT_COLOR
    closing_p.alignment = PP_ALIGN.CENTER
    
    # --- Final Slide: Thank You / Q&A ---
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide)
    
    ty_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(13), Inches(2))
    ty_p = ty_box.text_frame.paragraphs[0]
    ty_p.text = "Thank You"
    ty_p.font.bold = True
    ty_p.font.size = Pt(48)
    ty_p.font.color.rgb = ACCENT_COLOR
    ty_p.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path

