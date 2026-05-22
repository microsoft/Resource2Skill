def create_slide(
    output_pptx_path: str,
    title_text: str = "Strategic Process Flow",
    body_text: str = "Visualizing the trajectory of our financial lifecycle.",
    accent_color: tuple = (0, 191, 255),  # Cyan
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Custom Motion Path effect.
    Draws a visual path and animates a glowing orb along it.
    """
    import os
    import random
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageFilter
    from lxml import etree
    from pptx.oxml.ns import qn

    # --- Helper 1: Generate PIL Background ---
    def create_background(filepath, width=1920, height=1080):
        bg = Image.new("RGBA", (width, height), (13, 17, 28, 255))
        draw = ImageDraw.Draw(bg)
        # Add subtle ambient color blobs
        draw.ellipse([-300, -300, 800, 800], fill=(0, 60, 100, 120))
        draw.ellipse([1200, 400, 2200, 1400], fill=(60, 20, 80, 100))
        bg = bg.filter(ImageFilter.GaussianBlur(150))
        bg.save(filepath)
        return filepath

    # --- Helper 2: Generate Glowing Orb ---
    def create_glowing_orb(filepath, size_px=150, color=(0, 191, 255)):
        img = Image.new("RGBA", (size_px, size_px), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        center = size_px // 2
        
        # Draw radial glow
        for r in range(center, 10, -4):
            alpha = int(255 * (1 - r/center)**2.5)
            draw.ellipse([center-r, center-r, center+r, center+r], fill=color + (alpha,))
            
        # Draw solid inner core
        draw.ellipse([center-15, center-15, center+15, center+15], fill=(255, 255, 255, 255))
        img.save(filepath)
        return filepath

    # --- Helper 3: Calculate PPTX Path String ---
    def generate_path_string(points, slide_w, slide_h):
        """Converts absolute slide inches to relative percentage path string for animMotion."""
        start_x, start_y = points[0]
        path_parts = ["M 0 0"]
        for x, y in points[1:]:
            dx = (x - start_x) / slide_w
            dy = (y - start_y) / slide_h
            path_parts.append(f"L {dx:.5f} {dy:.5f}")
        return " ".join(path_parts)

    # --- Helper 4: Inject XML Animation ---
    def inject_motion_path(slide, shape_id, path_str, duration_ms=4000):
        # 1. Ensure <p:timing> tree exists
        timing = slide.element.find(qn('p:timing'))
        if timing is None:
            timing_xml = """
            <p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:tnLst>
                <p:par>
                  <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
                    <p:childTnLst>
                      <p:seq concurrent="1" nextAc="seek">
                        <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                          <p:childTnLst/>
                        </p:cTn>
                      </p:seq>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:tnLst>
            </p:timing>
            """
            timing = etree.fromstring(timing_xml)
            extLst = slide.element.find(qn('p:extLst'))
            if extLst is not None:
                extLst.addprevious(timing)
            else:
                slide.element.append(timing)

        # 2. Find the main sequence child list
        ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
        main_seq_lst = timing.xpath('.//p:cTn[@nodeType="mainSeq"]/p:childTnLst', namespaces=ns)[0]

        base_id = random.randint(10000, 90000)

        # 3. Create the animation node (Plays automatically with slide)
        anim_xml = f"""
        <p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
          <p:cTn id="{base_id}" fill="hold" presetID="1" presetClass="path" presetSubtype="0">
            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
            <p:childTnLst>
              <p:par>
                <p:cTn id="{base_id+1}" fill="hold">
                  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                  <p:childTnLst>
                    <p:animMotion pathEditMode="relative" path="{path_str}">
                      <p:cTn id="{base_id+2}" dur="{duration_ms}" fill="hold"/>
                      <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                    </p:animMotion>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
            </p:childTnLst>
          </p:cTn>
        </p:par>
        """
        anim_node = etree.fromstring(anim_xml)
        main_seq_lst.append(anim_node)


    # === Slide Setup ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    bg_path = "temp_bg.png"
    create_background(bg_path, width=1920, height=1080)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Typography ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(8), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8), Inches(0.5))
    p_sub = sub_box.text_frame.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(180, 190, 200)

    # === Layer 3: Visual Path Construction ===
    # Define the coordinates of the process journey
    path_points_inches = [
        (1.5, 5.0),   # Start
        (4.5, 3.2),   # Stage 2
        (8.0, 5.8),   # Stage 3
        (11.5, 3.5)   # End
    ]

    # Draw the dashed trajectory line
    builder = slide.shapes.build_freeform(start_x=Inches(path_points_inches[0][0]), 
                                          start_y=Inches(path_points_inches[0][1]))
    builder.add_line_segments([(Inches(x), Inches(y)) for x, y in path_points_inches[1:]])
    path_shape = builder.convert_to_shape()
    path_shape.line.color.rgb = RGBColor(*accent_color)
    # Using 4 for dashed line style (MSO_LINE.DASH = 4 but varies by environment, 4 is safe default)
    path_shape.line.dash_style = 4  
    path_shape.line.width = Pt(2)

    # Add Stage Markers & Labels
    for i, (cx, cy) in enumerate(path_points_inches):
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.15), Inches(cy - 0.15), Inches(0.3), Inches(0.3))
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(13, 17, 28)
        marker.line.color.rgb = RGBColor(*accent_color)
        marker.line.width = Pt(2)

        lbl = slide.shapes.add_textbox(Inches(cx - 1), Inches(cy + 0.25), Inches(2), Inches(0.5))
        lbl_p = lbl.text_frame.paragraphs[0]
        lbl_p.text = f"Stage {i+1}"
        lbl_p.alignment = PP_ALIGN.CENTER
        lbl_p.font.size = Pt(12)
        lbl_p.font.color.rgb = RGBColor(200, 200, 200)
        lbl_p.font.bold = True

    # === Layer 4: The Traveling Object (Animated) ===
    orb_path = "temp_orb.png"
    orb_size_in = 1.0
    create_glowing_orb(orb_path, size_px=150, color=accent_color)
    
    # Place orb exactly centered on the first point
    orb_shape = slide.shapes.add_picture(
        orb_path,
        Inches(path_points_inches[0][0] - orb_size_in/2),
        Inches(path_points_inches[0][1] - orb_size_in/2),
        Inches(orb_size_in),
        Inches(orb_size_in)
    )

    # === Layer 5: Inject Motion Path Animation ===
    # Calculate the relative VML path string automatically based on coordinates
    ppt_path_string = generate_path_string(
        path_points_inches, 
        slide_w=13.333, 
        slide_h=7.5
    )
    
    # Apply the animation to the orb
    inject_motion_path(slide, orb_shape.shape_id, ppt_path_string, duration_ms=4500)

    # Save and clean up
    prs.save(output_pptx_path)
    
    if os.path.exists(bg_path): os.remove(bg_path)
    if os.path.exists(orb_path): os.remove(orb_path)
    
    return output_pptx_path
