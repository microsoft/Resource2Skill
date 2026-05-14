def create_slide(
    output_pptx_path: str,
    title_text: str = "LIQUID NEON",
    body_text: str = "Abstract fluid backgrounds powered by algorithmic curves.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Abstract Liquid Glow Background visual effect.
    Returns: path to the saved PPTX file.
    """
    import math
    import os
    from PIL import Image, ImageDraw, ImageFilter
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    # === Layer 1 & 2: Background & Visual Effect (PIL) ===
    # Set up a high-res canvas (4K-ish for smoothness)
    w, h = 3840, 2160
    bg_color = (13, 9, 36) # Deep dark purple
    img = Image.new('RGB', (w, h), bg_color)

    # Define the fluid waves to mimic the "Liquify" drag paths
    # Format: (amplitude_y, freq_x, phase, color_rgba, stroke_width, y_offset)
    waves = [
        # Broad ambient background sweep (Deep Magenta)
        (400, 1500, 0, (150, 10, 80, 255), 600, h/2 + 200),
        # Mid-level wave (Vibrant Pink/Purple)
        (500, 1800, 1.5, (255, 50, 150, 255), 300, h/2 - 100),
        # Core bright highlight (Glowing Orange/Yellow)
        (550, 1800, 1.6, (255, 200, 50, 255), 100, h/2 - 120),
        # Secondary counter-curve for complexity (Cool Purple)
        (350, 1300, 3.14, (70, 30, 200, 220), 400, h/2 + 400)
    ]

    for amp_y, freq_x, phase, color, width, y_off in waves:
        # Create a transparent layer for each wave to prevent blur bleeding artifacts
        layer = Image.new('RGBA', (w, h), (0, 0, 0, 0))
        draw = ImageDraw.Draw(layer)
        
        points = []
        # Generate points spanning beyond the canvas edges
        for x in range(-500, w+500, 50):
            # Combine two sine waves with different frequencies to create an organic, unpredictable S-curve
            y = y_off + math.sin(x/freq_x + phase) * amp_y + math.cos(x/(freq_x*0.5) + phase*1.3) * (amp_y*0.5)
            points.append((x, y))

        # Draw the thick sweeping line
        draw.line(points, fill=color, width=width, joint='curve')

        # Apply extreme Gaussian blur to simulate the soft brush and stretched liquify look
        # The blur radius scales with the stroke width for realistic light falloff
        layer = layer.filter(ImageFilter.GaussianBlur(radius=width*0.45))

        # Alpha composite the glowing wave onto the main image
        img.paste(layer, (0,0), layer)

    # Save the generated background image temporarily
    bg_path = "liquid_glow_temp_bg.png"
    img.save(bg_path)

    # === Layer 3: Text & Content (python-pptx) ===
    prs = Presentation()
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Insert the PIL-generated background
    slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Add Sleek Typography over the dark area
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2.5))
    tf = txBox.text_frame
    
    # Title
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(72)
    p_title.font.bold = True
    p_title.font.name = "Arial"
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle/Body
    p_body = tf.add_paragraph()
    p_body.text = body_text
    p_body.font.size = Pt(28)
    p_body.font.name = "Arial"
    p_body.font.color.rgb = RGBColor(200, 200, 215) # Slightly cool gray to match the purple base

    # Clean up the temporary image and save the presentation
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
