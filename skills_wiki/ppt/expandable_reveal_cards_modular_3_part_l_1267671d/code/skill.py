def create_slide(
    output_pptx_path: str,
    title_text: str = "Your Slide Title Text Here",
    body_text: str = "You can replace this sample text with your own text. Add more text if required.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Expandable Reveal Cards' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set slide background to very light grey so the white card footers pop out
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 240, 240)

    # Add dark Top Banner for Slide Title
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.4), prs.slide_width, Inches(0.9))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(70, 70, 70)
    banner.line.color.rgb = RGBColor(70, 70, 70)
    
    title_tf = banner.text_frame
    title_tf.word_wrap = True
    p = title_tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.LEFT
    banner.text_frame.margin_left = Inches(0.8)

    # Color Palette for the 4 Cards
    colors = [
        RGBColor(57, 160, 173),   # Teal
        RGBColor(230, 115, 0),    # Orange
        RGBColor(100, 75, 125),   # Purple
        RGBColor(120, 165, 60)    # Green
    ]

    # Layout Parameters
    num_cards = 4
    card_w = Inches(2.2)
    header_h = Inches(0.8)
    content_h = Inches(2.3)
    footer_h = Inches(0.8)
    gap = Inches(0.5)

    total_width = (num_cards * card_w) + ((num_cards - 1) * gap)
    start_x = (prs.slide_width - total_width) / 2
    start_y = Inches(2.3)

    # Generate Cards
    for i in range(num_cards):
        x = start_x + i * (card_w + gap)
        y = start_y

        color = colors[i]

        # ==========================================
        # 1. HEADER (Top-Corners Rounded)
        # Trick: Layer a rounded rect and a standard rect to hide the bottom rounded corners
        # ==========================================
        h_shape_rnd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, header_h)
        h_shape_sqr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + (header_h/2), card_w, header_h/2)

        for shp in (h_shape_rnd, h_shape_sqr):
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            shp.line.color.rgb = color # Hides the seam perfectly

        # Header Text Box
        h_txt = slide.shapes.add_textbox(x, y, card_w, header_h)
        h_tf = h_txt.text_frame
        h_tf.word_wrap = True
        h_p = h_tf.add_paragraph()
        h_p.text = "Subtitle Text"
        h_p.font.size = Pt(20)
        h_p.font.bold = True
        h_p.font.color.rgb = RGBColor(255, 255, 255)
        h_p.alignment = PP_ALIGN.CENTER
        # Vertical centering simulation
        h_txt.text_frame.margin_top = Inches(0.2)

        # ==========================================
        # 2. BODY CONTENT (Dark Grey)
        # ==========================================
        m_y = y + header_h
        m_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, m_y, card_w, content_h)
        m_shape.fill.solid()
        m_shape.fill.fore_color.rgb = RGBColor(45, 52, 54) # Dark Slate Grey
        m_shape.line.color.rgb = RGBColor(45, 52, 54)

        # Content Text Box
        m_txt = slide.shapes.add_textbox(x, m_y, card_w, content_h)
        m_tf = m_txt.text_frame
        m_tf.word_wrap = True
        m_p = m_tf.add_paragraph()
        m_p.text = body_text
        m_p.font.size = Pt(13)
        m_p.font.color.rgb = RGBColor(240, 240, 240)
        m_p.alignment = PP_ALIGN.CENTER
        
        m_txt.text_frame.margin_left = Inches(0.15)
        m_txt.text_frame.margin_right = Inches(0.15)
        m_txt.text_frame.margin_top = Inches(0.3)

        # ==========================================
        # 3. FOOTER (Bottom-Corners Rounded)
        # Trick: Reverse of the header, white color
        # ==========================================
        f_y = m_y + content_h
        f_shape_rnd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, f_y, card_w, footer_h)
        f_shape_sqr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, f_y, card_w, footer_h/2)

        for shp in (f_shape_rnd, f_shape_sqr):
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(255, 255, 255)
            shp.line.color.rgb = RGBColor(255, 255, 255)

        # Footer Icon (Using Unicode geometric shapes as placeholders)
        icon_txt = slide.shapes.add_textbox(x, f_y, card_w, footer_h)
        i_tf = icon_txt.text_frame
        i_p = i_tf.add_paragraph()
        
        icons = ["\u2699", "\u2605", "\u2714", "\u2690"] # Gear, Star, Checkmark, Flag
        i_p.text = icons[i]
        i_p.font.size = Pt(32)
        i_p.font.color.rgb = color # Match icon color to header color
        i_p.alignment = PP_ALIGN.CENTER
        icon_txt.text_frame.margin_top = Inches(0.1)

    prs.save(output_pptx_path)
    return output_pptx_path
