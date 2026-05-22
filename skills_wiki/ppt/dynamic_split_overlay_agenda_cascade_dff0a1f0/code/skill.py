import os
import urllib.request
from io import BytesIO
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_word_1: str = "MEETING",
    title_word_2: str = "AGENDA",
    bg_keyword: str = "office,desk,meeting",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dynamic Split-Overlay Agenda Cascade" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background & Split Overlay via PIL ===
    # Dimensions for 13.333x7.5 at 120dpi
    base_w, base_h = 1600, 900 
    
    # Try downloading a thematic background image, fallback to a solid gray image
    img = None
    try:
        url = f"https://images.unsplash.com/featured/1600x900/?{bg_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        res = urllib.request.urlopen(req, timeout=5)
        img = Image.open(BytesIO(res.read())).convert('RGBA')
        img = img.resize((base_w, base_h), Image.LANCZOS)
    except Exception:
        img = Image.new('RGBA', (base_w, base_h), color=(220, 220, 230, 255))

    # Create the transparent split overlay mask
    overlay = Image.new('RGBA', (base_w, base_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    
    # Calculate pixel width corresponding to 3.5 inches
    split_inch = 3.5
    split_x = int(base_w * (split_inch / 13.333))
    
    # Dark panel on the left (alpha 235/255)
    draw.rectangle([0, 0, split_x, base_h], fill=(45, 45, 45, 235))
    # Light panel on the right (alpha 225/255)
    draw.rectangle([split_x, 0, base_w, base_h], fill=(255, 255, 255, 225))

    # Composite the overlay onto the background
    composite = Image.alpha_composite(img, overlay)
    bg_path = "temp_agenda_bg.png"
    composite.save(bg_path)

    # Insert composite image as full slide background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    os.remove(bg_path)  # Cleanup temp file

    # === Layer 2: Vertical Sidebar Title ===
    # A 4-inch wide textbox rotated 270 degrees. 
    # To perfectly center it horizontally in the 3.5-inch sidebar, left = (3.5/2) - (4/2) = -0.25
    tx_box = slide.shapes.add_textbox(Inches(-0.25), Inches(3.25), Inches(4.0), Inches(1.0))
    tx_box.rotation = 270
    tf = tx_box.text_frame
    tf.word_wrap = False
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    
    run1 = p.add_run()
    run1.text = f"{title_word_1} "
    run1.font.bold = True
    run1.font.size = Pt(44)
    run1.font.color.rgb = RGBColor(255, 255, 255)

    run2 = p.add_run()
    run2.text = title_word_2
    run2.font.bold = False
    run2.font.size = Pt(44)
    run2.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 3: Agenda Cascades (Geometric Markers and Text) ===
    colors = [
        RGBColor(76, 175, 80),   # Green
        RGBColor(156, 39, 176),  # Purple
        RGBColor(192, 160, 32),  # Mustard
        RGBColor(3, 169, 244),   # Blue
        RGBColor(0, 188, 212)    # Cyan
    ]

    start_y = 1.0
    spacing = 1.25
    diamond_size = 0.65

    for i, color in enumerate(colors):
        y_ctr = start_y + i * spacing
        
        # 3.1 Diamond Shape placed exactly on the split-overlay edge
        diamond = slide.shapes.add_shape(
            MSO_SHAPE.DIAMOND,
            Inches(split_inch - diamond_size/2), 
            Inches(y_ctr - diamond_size/2),
            Inches(diamond_size), 
            Inches(diamond_size)
        )
        diamond.fill.solid()
        diamond.fill.fore_color.rgb = color
        diamond.line.color.rgb = RGBColor(255, 255, 255)
        diamond.line.width = Pt(2)
        
        # Numeric text inside the diamond
        df = diamond.text_frame
        dp = df.paragraphs[0]
        dp.text = str(i + 1)
        dp.font.bold = True
        dp.font.size = Pt(16)
        dp.font.color.rgb = RGBColor(255, 255, 255)
        dp.alignment = PP_ALIGN.CENTER

        # 3.2 Dynamic Colored Topic Title
        title_box = slide.shapes.add_textbox(
            Inches(split_inch + 0.6), 
            Inches(y_ctr - 0.35), 
            Inches(8.0), 
            Inches(0.4)
        )
        tp = title_box.text_frame.add_paragraph()
        tr = tp.add_run()
        tr.text = f"YOUR AGENDA TOPIC {i+1}"
        tr.font.bold = True
        tr.font.size = Pt(16)
        tr.font.color.rgb = color

        # 3.3 Subtle Gray Details Text
        body_box = slide.shapes.add_textbox(
            Inches(split_inch + 0.6), 
            Inches(y_ctr + 0.05), 
            Inches(8.0), 
            Inches(0.6)
        )
        body_box.text_frame.word_wrap = True
        bp = body_box.text_frame.add_paragraph()
        br = bp.add_run()
        br.text = "Your detailed agenda of discussion inserts here. Provide a brief overview of the topics that will be covered to keep the meeting on track."
        br.font.size = Pt(12)
        br.font.color.rgb = RGBColor(90, 90, 90)

    prs.save(output_pptx_path)
    return output_pptx_path
