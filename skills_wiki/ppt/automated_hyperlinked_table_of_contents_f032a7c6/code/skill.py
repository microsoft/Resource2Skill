import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from lxml import etree

def create_hyperlinked_toc_in_presentation(prs: Presentation, toc_slide_index: int = 1) -> Presentation:
    """
    Inserts a hyperlinked Table of Contents into an existing Presentation object.

    Args:
        prs (Presentation): The presentation object to modify.
        toc_slide_index (int): The index where the ToC slide should be inserted.

    Returns:
        Presentation: The modified presentation object.
    """
    if not (0 <= toc_slide_index <= len(prs.slides)):
        raise ValueError(f"toc_slide_index must be between 0 and {len(prs.slides)}")

    slide_layout = prs.slide_layouts[5]  # Title Only layout
    toc_slide = prs.slides.add_slide(slide_layout)
    
    # Reorder slides to place the new ToC slide at the correct index
    slides = list(prs.slides)
    xml_slides = prs.slides._sldIdLst
    slides_in_order = slides[:toc_slide_index] + [slides[-1]] + slides[toc_slide_index:-1]
    xml_slides.clear()
    for slide in slides_in_order:
        xml_slides.append(slide._element)

    # --- Design ToC Slide ---
    title_shape = toc_slide.shapes.title
    title_shape.text = "Table of Contents"
    title_shape.text_frame.paragraphs[0].font.size = Pt(44)
    title_shape.text_frame.paragraphs[0].font.bold = True
    
    left, top, width, height = Inches(1), Inches(1.5), prs.slide_width - Inches(2), prs.slide_height - Inches(2.5)
    textbox = toc_slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    tf.margin_bottom = Inches(0.1)
    tf.margin_left = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    # LXML Magic: Set the text box to have two columns
    bodyPr = tf._txBody.get_or_add_bodyPr()
    bodyPr.set("numCol", "2")
    bodyPr.set("spcCol", "360000")  # Spacing between columns in EMUs (0.4 inches)

    # Populate ToC with hyperlinked slide titles
    slides_to_list = [s for s in slides_in_order if s.slide_id != toc_slide.slide_id]

    for i, slide in enumerate(slides_to_list):
        slide_title = f"Slide {i + 1}"
        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text
        
        p = tf.add_paragraph()
        p.text = slide_title
        p.font.size = Pt(16)
        
        # Create hyperlink
        run = p.runs[0]
        run.hyperlink.address = None
        run.hyperlink._hlinkClick.rId = toc_slide.part.relate_to(
            slide.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
        ).rId
        
    return prs

def create_slide(
    output_pptx_path: str = "Automated_ToC_Presentation.pptx",
    **kwargs
) -> str:
    """
    Generates a sample presentation with an automated, hyperlinked Table of Contents.
    This demonstrates the ability to programmatically create a navigable ToC.

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_titles = [
        "Project Kick-off: A New Beginning",
        "Phase 1: Research & Discovery",
        "Understanding the Market Landscape",
        "Core Product Strategy",
        "Feature Deep Dive: Part A",
        "Feature Deep Dive: Part B",
        "Design & User Experience Mockups",
        "Technical Architecture Overview",
        "Go-to-Market Plan",
        "Financial Projections & KPIs",
        "Timeline & Key Milestones",
        "Our Talented Team",
        "Appendix: Supporting Data",
        "Next Steps & Q&A"
    ]
    
    # Create the title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Comprehensive Business Plan"
    subtitle.text = "Project Phoenix"

    # Create the content slides
    for title_text in slide_titles:
        content_layout = prs.slide_layouts[5] # Title Only
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = title_text

    # --- Insert the Table of Contents ---
    prs = create_hyperlinked_toc_in_presentation(prs, toc_slide_index=1)
    
    prs.save(output_pptx_path)
    return output_pptx_path
