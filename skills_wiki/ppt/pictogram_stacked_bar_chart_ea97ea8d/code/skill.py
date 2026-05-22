def create_slide(
    output_pptx_path: str,
    title_text: str = "NEXT 品牌服飾",
    subtitle_text: str = "2019 Q3 各品類營收業績",
    theme_color: tuple = (0, 225, 255), # Cyan for icons
    **kwargs,
) -> str:
    """
    Creates a slide featuring a custom pictogram stacked bar chart.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Data Definition ---
    data = [
        {"category": "內衣 (Innerwear)", "value": 600},
        {"category": "褲子 (Pants)", "value": 1500},
        {"category": "洋裝 (Dress)", "value": 1200},
        {"category": "上衣 (Top)", "value": 800},
        {"category": "外套 (Jacket)", "value": 3000},
    ]
    max_value = max(item["value"] for item in data)
    unit_per_icon = 200 # Each icon represents 200 units
    
    # --- Helper 1: Draw a simple T-Shirt Icon ---
    def create_tshirt_icon(color_rgb, size=40):
        img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        w, h = size, size
        # Stylized T-shirt polygon coordinates
        points = [
            (w*0.3, h*0.1), (w*0.7, h*0.1),       # Neck top
            (w*0.95, h*0.3),                      # Right shoulder outer
            (w*0.85, h*0.45),                     # Right sleeve bottom
            (w*0.75, h*0.35),                     # Right armpit
            (w*0.75, h*0.9),                      # Right bottom
            (w*0.25, h*0.9),                      # Left bottom
            (w*0.25, h*0.35),                     # Left armpit
            (w*0.15, h*0.45),                     # Left sleeve bottom
            (w*0.05, h*0.3)                       # Left shoulder outer
        ]
        color_rgba = color_rgb + (255,)
        draw.polygon(points, fill=color_rgba)
        return img

    # --- Helper 2: Create the Stacked Bar Image ---
    def create_pictogram_bar(icon_img, value, unit_val, spacing=4):
        total_icons = value / unit_val
        full_icons = int(total_icons)
        partial_ratio = total_icons - full_icons
        
        icon_w, icon_h = icon_img.size
        # Calculate total width needed
        total_w = int((full_icons + (1 if partial_ratio > 0 else 0)) * (icon_w + spacing))
        if total_w == 0: total_w = 1
        
        bar_img = Image.new('RGBA', (total_w, icon_h), (0, 0, 0, 0))
        current_x = 0
        
        # Paste full icons
        for _ in range(full_icons):
            bar_img.paste(icon_img, (current_x, 0))
            current_x += icon_w + spacing
            
        # Paste partial icon by cropping
        if partial_ratio > 0:
            crop_w = int(icon_w * partial_ratio)
            partial_img = icon_img.crop((0, 0, crop_w, icon_h))
            bar_img.paste(partial_img, (current_x, 0))
            
        return bar_img

    # --- Step 1: Create Tinted Background ---
    bg_img_path = "temp_bg.png"
    try:
        # Download a fashion-related image
        req = urllib.request.Request(
            "https://images.unsplash.com/photo-1445205170230-053b83016050?w=1600&q=80",
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req) as response:
            bg_base = Image.open(BytesIO(response.read())).convert("RGBA")
            bg_base = bg_base.resize((1920, 1080))
    except Exception:
        # Fallback to solid image if download fails
        bg_base = Image.new('RGBA', (1920, 1080), (50, 50, 50, 255))

    # Apply deep slate blue overlay mask
    overlay = Image.new('RGBA', bg_base.size, (50, 60, 100, 210)) # Tint color
    bg_final = Image.alpha_composite(bg_base, overlay)
    bg_final.save(bg_img_path)

    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # --- Step 2: Add Main Titles (Left Panel) ---
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.5), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(255, 255, 255)
    
    p3 = tf.add_paragraph()
    p3.text = "\n單位 萬 NT$"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(200, 200, 200)

    # --- Step 3: Generate and Layout the Pictogram Chart (Right Panel) ---
    start_y = Inches(1.5)
    step_y = Inches(0.85)
    chart_x = Inches(6.0)
    
    base_icon = create_tshirt_icon(theme_color, size=35)
    
    for i, item in enumerate(data):
        current_y = start_y + (i * step_y)
        
        # 3a. Add Category Label
        label_box = slide.shapes.add_textbox(Inches(4.5), current_y - Inches(0.1), Inches(1.3), Inches(0.5))
        tf_label = label_box.text_frame
        p_label = tf_label.paragraphs[0]
        p_label.text = item["category"]
        p_label.font.size = Pt(16)
        p_label.font.color.rgb = RGBColor(255, 255, 255)
        p_label.alignment = PP_ALIGN.RIGHT
        
        # 3b. Generate and Place Pictogram Bar
        bar_img = create_pictogram_bar(base_icon, item["value"], unit_per_icon)
        temp_bar_path = f"temp_bar_{i}.png"
        bar_img.save(temp_bar_path)
        
        # Calculate scaling to fit PPT dimensions neatly
        pt_width = bar_img.width * 0.75 # scale factor
        pt_height = bar_img.height * 0.75
        
        pic = slide.shapes.add_picture(temp_bar_path, chart_x, current_y, height=Pt(pt_height))
        
        # 3c. Add Value Label at the end of the bar
        val_x = chart_x + Pt(pt_width) + Inches(0.1)
        val_box = slide.shapes.add_textbox(val_x, current_y - Inches(0.1), Inches(1.5), Inches(0.5))
        tf_val = val_box.text_frame
        p_val = tf_val.paragraphs[0]
        p_val.text = str(item["value"])
        p_val.font.size = Pt(24)
        p_val.font.bold = True
        p_val.font.color.rgb = RGBColor(*theme_color) # Match icon color
        
        # Cleanup temp bar image
        if os.path.exists(temp_bar_path):
            os.remove(temp_bar_path)

    # Cleanup bg image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
