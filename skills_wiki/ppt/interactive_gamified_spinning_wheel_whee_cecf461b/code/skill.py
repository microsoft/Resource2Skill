def create_slide(
    output_pptx_path: str,
    title_text: str = "Who's Next?",
    names: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file featuring a visual 'Spinning Wheel of Names' setup.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    
    if names is None:
        names = ["Liam", "Olivia", "Noah", "Emma", "Oliver", "Ava", 
                 "Elijah", "Sophia", "James", "Isabella", "William", "Mia"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Constants for layout
    cx = prs.slide_width / 2
    cy = prs.slide_height / 2
    wheel_diameter = Inches(6.0)

    # === Layer 1: The Pie Chart (The Wheel) ===
    chart_data = CategoryChartData()
    chart_data.categories = names
    # Provide an equal value (1.0) for every slice so they are uniform
    chart_data.add_series('Names', [1.0] * len(names))

    x = cx - (wheel_diameter / 2)
    y = cy - (wheel_diameter / 2)
    
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, x, y, wheel_diameter, wheel_diameter, chart_data
    )
    chart = graphic_frame.chart

    # Clean up chart UI
    chart.has_legend = False
    chart.has_title = False

    # Format Data Labels
    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_category_name = True
    data_labels.show_value = False
    data_labels.show_percentage = False
    data_labels.font.size = Pt(14)
    data_labels.font.bold = True
    data_labels.font.color.rgb = RGBColor(0, 0, 0)

    # Vibrant Game-Show Color Palette
    palette = [
        RGBColor(255, 89, 94),   # Bright Red
        RGBColor(255, 202, 58),  # Sun Yellow
        RGBColor(138, 201, 38),  # Lime Green
        RGBColor(25, 130, 196),  # Vivid Blue
        RGBColor(106, 76, 147),  # Royal Purple
        RGBColor(255, 146, 76),  # Orange
        RGBColor(50, 205, 50),   # Bright Green
        RGBColor(0, 191, 255),   # Deep Sky Blue
    ]

    # Apply colors to individual slices
    for i, point in enumerate(chart.series[0].points):
        fill = point.format.fill
        fill.solid()
        fill.fore_color.rgb = palette[i % len(palette)]

    # === Layer 2: UI Elements (Pointer and Button) ===

    # Center Spin Button (Circle)
    btn_diameter = Inches(1.2)
    btn_x = cx - (btn_diameter / 2)
    btn_y = cy - (btn_diameter / 2)
    
    btn = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, btn_x, btn_y, btn_diameter, btn_diameter
    )
    btn.fill.solid()
    btn.fill.fore_color.rgb = RGBColor(255, 255, 255)  # White fill
    btn.line.color.rgb = RGBColor(50, 50, 50)          # Dark border
    btn.line.width = Pt(3)
    
    btn_tf = btn.text_frame
    btn_tf.text = "SPIN!"
    btn_p = btn_tf.paragraphs[0]
    btn_p.alignment = PP_ALIGN.CENTER
    btn_p.font.bold = True
    btn_p.font.size = Pt(20)
    btn_p.font.color.rgb = RGBColor(0, 0, 0)
    btn_tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Top Pointer (Triangle)
    tri_width = Inches(0.5)
    tri_height = Inches(0.7)
    tri_x = cx - (tri_width / 2)
    # Positioned slightly overlapping the top of the wheel
    tri_y = y - Inches(0.2) 
    
    pointer = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, tri_x, tri_y, tri_width, tri_height
    )
    pointer.rotation = 180  # Point downward
    pointer.fill.solid()
    pointer.fill.fore_color.rgb = RGBColor(220, 20, 60) # Crimson Red
    pointer.line.color.rgb = RGBColor(255, 255, 255)
    pointer.line.width = Pt(1.5)

    # Optional: Top Pointer Anchor (Small Circle over the triangle base)
    anchor_diam = Inches(0.3)
    anchor_x = cx - (anchor_diam / 2)
    anchor_y = tri_y - (anchor_diam / 2) + Inches(0.05)
    
    anchor = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, anchor_x, anchor_y, anchor_diam, anchor_diam
    )
    anchor.fill.solid()
    anchor.fill.fore_color.rgb = RGBColor(50, 50, 50)
    anchor.line.color.rgb = RGBColor(255, 255, 255)
    anchor.line.width = Pt(1)

    # Add a title at the top left
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(1))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.name = "Arial"
    
    p2 = tf.add_paragraph()
    p2.text = "Click the wheel in PPT and add 'Spin' Animation!"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(120, 120, 120)

    prs.save(output_pptx_path)
    return output_pptx_path
