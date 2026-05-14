def create_slide(
    output_pptx_path: str,
    hero_text: str = "X",
    subtitle_text: str = "Hello, future",
    **kwargs,
) -> str:
    """
    Creates a 2-slide PPTX reproducing the 'Cinematic Dark Mode Product Reveal' aesthetic.
    Slide 1: Gradient Hero Text Teaser
    Slide 2: Hardware Reveal with feature callouts
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml
    from PIL import Image, ImageDraw

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Helper Function: Set Background to Black ---
    def set_black_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 0, 0)

    # --- Helper Function: Apply Gradient to Text via LXML ---
    def apply_gradient_to_run(run):
        # Angle 2700000 is 45 degrees in 1/60000ths of a degree
        gradient_xml = """
        <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:gsLst>
                <a:gs pos="0"><a:srgbClr val="FF1493"/></a:gs>     <!-- Deep Pink -->
                <a:gs pos="40000"><a:srgbClr val="FF4500"/></a:gs> <!-- Vibrant Orange -->
                <a:gs pos="100000"><a:srgbClr val="00BFFF"/></a:gs> <!-- Cyan -->
            </a:gsLst>
            <a:lin ang="2700000" scaled="1"/>
        </a:gradFill>
        """
        grad_fill_element = parse_xml(gradient_xml)
        rPr = run._r.get_or_add_rPr()
        rPr.append(grad_fill_element)

    # --- Helper Function: Create Dummy Phone Mockup using PIL ---
    def create_phone_mockup(filename):
        # Create a transparent image
        img = Image.new('RGBA', (300, 600), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        
        # Outer bezel (dark gray with slightly lighter edge)
        draw.rounded_rectangle([0, 0, 300, 600], radius=45, fill=(20, 20, 20, 255), outline=(100, 100, 100, 255), width=3)
        # Inner screen (pure black)
        draw.rounded_rectangle([12, 12, 288, 588], radius=35, fill=(0, 0, 0, 255))
        # The iconic "Notch"
        draw.rounded_rectangle([80, 12, 220, 45], radius=15, fill=(20, 20, 20, 255))
        # Screen glow/reflection (subtle diagonal polygon)
        draw.polygon([(12, 12), (288, 150), (12, 300)], fill=(255, 255, 255, 5))
        
        img.save(filename)
        return filename

    # ==========================================
    # SLIDE 1: The Teaser (Gradient Hero Text)
    # ==========================================
    slide_1 = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    set_black_background(slide_1)

    # Add Hero Text box
    hero_box = slide_1.shapes.add_textbox(Inches(3.66), Inches(2.0), Inches(6.0), Inches(3.0))
    tf = hero_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = hero_text
    run.font.name = "Arial"
    run.font.size = Pt(220)
    run.font.bold = True
    
    # Inject Gradient XML
    apply_gradient_to_run(run)

    # Add Subtitle
    sub_box = slide_1.shapes.add_textbox(Inches(3.66), Inches(5.2), Inches(6.0), Inches(1.0))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER
    run_sub = p_sub.add_run()
    run_sub.text = subtitle_text
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(32)
    run_sub.font.bold = False
    run_sub.font.color.rgb = RGBColor(255, 255, 255) # Pure white

    # ==========================================
    # SLIDE 2: The Hardware Reveal
    # ==========================================
    slide_2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_black_background(slide_2)

    # Generate and insert phone mockup
    mockup_path = "temp_mockup.png"
    create_phone_mockup(mockup_path)
    slide_2.shapes.add_picture(mockup_path, Inches(2.0), Inches(1.0), height=Inches(5.5))

    # Add Reveal Feature Text
    feat_title_box = slide_2.shapes.add_textbox(Inches(6.0), Inches(2.5), Inches(6.0), Inches(1.0))
    tf_ft = feat_title_box.text_frame
    p_ft = tf_ft.paragraphs[0]
    run_ft = p_ft.add_run()
    run_ft.text = "All screen."
    run_ft.font.name = "Arial"
    run_ft.font.size = Pt(48)
    run_ft.font.bold = True
    apply_gradient_to_run(run_ft) # Use the same gradient for consistency

    feat_desc_box = slide_2.shapes.add_textbox(Inches(6.0), Inches(3.5), Inches(6.0), Inches(2.0))
    tf_fd = feat_desc_box.text_frame
    
    features = ["5.8-inch OLED Display", "Super Retina Tech", "Revolutionary Sensors"]
    for idx, feat in enumerate(features):
        p = tf_fd.add_paragraph() if idx > 0 else tf_fd.paragraphs[0]
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = feat
        run.font.name = "Arial"
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(170, 170, 170) # Light gray for secondary text

    # Cleanup temp image and save
    prs.save(output_pptx_path)
    if os.path.exists(mockup_path):
        os.remove(mockup_path)
        
    return output_pptx_path
