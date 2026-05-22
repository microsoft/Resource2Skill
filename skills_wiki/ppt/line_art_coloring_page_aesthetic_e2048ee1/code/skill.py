def create_slide(
    output_pptx_path: str,
    title_text: str = "YOU ARE PURE MAGIC",
    body_text: str = "",
    **kwargs,
) -> str:
    """
    Creates a printable Quote Coloring Page in an 8.5 x 11 Portrait layout.
    Uses lxml injection to create massive text outlines and native shapes for fillers.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from lxml import etree

    prs = Presentation()
    
    # Set to Standard Letter Portrait (Printable Coloring Page format)
    prs.slide_width = Inches(8.5)
    prs.slide_height = Inches(11.0)
    
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Clean any residual placeholders
    for shape in slide.shapes:
        sp = shape._element
        sp.getparent().remove(sp)

    # === Layer 1: Page Frame ===
    margin = Inches(0.25)
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        margin, margin, 
        prs.slide_width - (margin * 2), 
        prs.slide_height - (margin * 2)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor(0, 0, 0)
    frame.line.width = Pt(6)

    # === Layer 2: Decorative Filler Shapes ===
    # Tuple format: (Shape Type, Left, Top, Width, Height)
    fillers = [
        # Corner anchors
        (MSO_SHAPE.CLOUD, Inches(0.6), Inches(0.6), Inches(2.2), Inches(1.5)),
        (MSO_SHAPE.HEART, Inches(6.0), Inches(0.8), Inches(1.8), Inches(1.6)),
        (MSO_SHAPE.STAR_5_POINT, Inches(0.7), Inches(8.5), Inches(2.0), Inches(2.0)),
        (MSO_SHAPE.SUN, Inches(5.8), Inches(8.3), Inches(2.2), Inches(2.2)),
        
        # Side framing
        (MSO_SHAPE.MOON, Inches(6.6), Inches(4.5), Inches(1.2), Inches(1.8)),
        (MSO_SHAPE.DIAMOND, Inches(0.6), Inches(4.8), Inches(1.0), Inches(1.0)),
        
        # Sparkles and Dots (to fill negative space)
        (MSO_SHAPE.OVAL, Inches(3.2), Inches(1.2), Inches(0.4), Inches(0.4)),
        (MSO_SHAPE.STAR_4_POINT, Inches(4.8), Inches(2.0), Inches(0.6), Inches(0.6)),
        (MSO_SHAPE.OVAL, Inches(1.4), Inches(3.2), Inches(0.3), Inches(0.3)),
        (MSO_SHAPE.STAR_4_POINT, Inches(7.0), Inches(3.0), Inches(0.5), Inches(0.5)),
        (MSO_SHAPE.OVAL, Inches(1.2), Inches(6.8), Inches(0.4), Inches(0.4)),
        (MSO_SHAPE.STAR_4_POINT, Inches(6.5), Inches(6.8), Inches(0.7), Inches(0.7)),
        (MSO_SHAPE.OVAL, Inches(3.5), Inches(9.2), Inches(0.4), Inches(0.4)),
        (MSO_SHAPE.STAR_4_POINT, Inches(4.8), Inches(8.8), Inches(0.5), Inches(0.5)),
    ]

    for shape_type, left, top, width, height in fillers:
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(0, 0, 0)
        shape.line.width = Pt(4) # Thick black outline

    # === Layer 3: Stacked Typography ===
    # Logic to split text into chunky, short blocks
    words = title_text.upper().split()
    lines = []
    current_line = []
    for word in words:
        if len(" ".join(current_line + [word])) > 9 and current_line:
            lines.append(" ".join(current_line))
            current_line = [word]
        else:
            current_line.append(word)
    if current_line:
        lines.append(" ".join(current_line))

    # Add central text box
    txBox = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(6.5), Inches(6.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, line_text in enumerate(lines):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        
        # Compress vertical spacing to make text blocks interlock slightly
        p.line_spacing = 0.85
        
        run = p.add_run()
        run.text = line_text
        run.font.name = "Arial Black"
        # Alternate font sizes slightly to give an uneven, organic drawn feel
        run.font.size = Pt(85 if i % 2 == 0 else 70)
        
        # Text Fill = Pure White
        run.font.fill.solid()
        run.font.fill.fore_color.rgb = RGBColor(255, 255, 255)
        
        # --- LXML INJECTION: Heavy Text Outline ---
        # Grabs the Run Properties element
        rPr = run._r.get_or_add_rPr()
        
        # Inject the Line properties (<a:ln>)
        ln = etree.SubElement(rPr, qn('a:ln'))
        ln.set('w', '50800') # 4 pt stroke width (1 pt = 12700 EMUs)
        ln.set('cap', 'rnd') # Rounded caps for a friendly look
        
        solidFill = etree.SubElement(ln, qn('a:solidFill'))
        srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
        srgbClr.set('val', '000000') # Black stroke

    prs.save(output_pptx_path)
    return output_pptx_path
