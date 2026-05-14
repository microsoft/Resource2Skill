def create_slide(
    output_pptx_path: str,
    title_text: str = "公司组织架构图",
    bg_color: tuple = (13, 17, 28),
    node_main_color: tuple = (0, 102, 204),
    node_sub_color: tuple = (24, 144, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Tech-Style Organizational Hierarchy.
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Helper 1: Create Tech Background ===
    bg_img_path = "temp_tech_bg.png"
    try:
        width, height = 1920, 1080
        img = Image.new('RGB', (width, height), bg_color)
        draw = ImageDraw.Draw(img)
        # Create subtle radial gradient (center highlight)
        center_x, center_y = width // 2, height // 3
        max_radius = 800
        for r in range(max_radius, 0, -5):
            alpha = int(40 * (1 - r / max_radius))
            # Blend highlight (slightly lighter blue)
            r_c = min(255, bg_color[0] + alpha)
            g_c = min(255, bg_color[1] + alpha)
            b_c = min(255, bg_color[2] + alpha * 2)
            draw.ellipse(
                [center_x - r, center_y - r, center_x + r, center_y + r],
                fill=(r_c, g_c, b_c)
            )
        img.save(bg_img_path)
        slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)
        os.remove(bg_img_path)
    except Exception as e:
        print(f"Warning: Could not generate PIL background. {e}")

    # === Add Title ===
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.4), prs.slide_width, Inches(0.8))
    tf = title_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    tf.paragraphs[0].font.name = "Microsoft YaHei"

    # === Helper 2: Draw Shapes & Lines ===
    def add_node(x, y, w, h, text, fill_rgb):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
        shape.line.fill.background() # No border
        
        # Adjust corner radius (smaller radius looks more professional)
        shape.adjustments[0] = 0.15 
        
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.name = "Microsoft YaHei"
        return shape

    def add_v_line(x, y, h, color=(150, 150, 150), thickness=0.02):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - thickness/2), Inches(y), Inches(thickness), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.fill.background()

    def add_h_line(x, y, w, color=(150, 150, 150), thickness=0.02):
        if w < 0:
            x = x + w
            w = abs(w)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y - thickness/2), Inches(w), Inches(thickness))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.fill.background()

    # === Define Organization Structure Coordinates ===
    # Using absolute coordinate mapping for a perfect layout (1 Executive, 3 Managers, 9 Sub-groups)
    node_w, node_h = 1.8, 0.6
    cx = 13.333 / 2  # Center X

    # Level 1: Executive
    l1_y = 1.5
    add_node(cx - node_w/2, l1_y, node_w, node_h, "执行董事", node_main_color)

    # Level 2: Managers (Standard Layout)
    l2_y = 2.8
    l2_xs = [cx - 3.5, cx, cx + 3.5]
    l2_names = ["市场总监", "总经理", "财务总监"]
    
    # Draw L1 to L2 Connector
    add_v_line(cx, l1_y + node_h, l2_y - (l1_y + node_h) - 0.2) # Stem down
    add_h_line(l2_xs[0], l2_y - 0.2, l2_xs[2] - l2_xs[0]) # Main horizontal bus
    
    for x, name in zip(l2_xs, l2_names):
        add_v_line(x, l2_y - 0.2, 0.2) # Drop down to L2
        add_node(x - node_w/2, l2_y, node_w, node_h, name, node_main_color)

    # Level 3: Sub-departments (Hanging Layout)
    l3_names = [
        ["信息调研组", "产品策划组", "广告投放组"],
        ["销售一区", "销售二区", "销售三区"],
        ["审计组", "会计组", "出纳组"]
    ]
    
    # Draw hanging nodes under each L2 parent
    for i, parent_cx in enumerate(l2_xs):
        names = l3_names[i]
        
        # Determine the spine X (shifted to the left of the parent)
        spine_x = parent_cx - (node_w/2) + 0.2
        
        # Draw spine down from parent
        total_spine_h = (len(names) * 0.8)
        add_v_line(spine_x, l2_y + node_h, total_spine_h)
        
        # Draw children
        for j, name in enumerate(names):
            child_y = l2_y + node_h + 0.4 + (j * 0.8)
            # Drop horizontal connector to child
            add_h_line(spine_x, child_y + node_h/2, 0.3)
            # Add child node
            child_x = spine_x + 0.3
            add_node(child_x, child_y, node_w * 0.9, node_h, name, node_sub_color)

    prs.save(output_pptx_path)
    return output_pptx_path
