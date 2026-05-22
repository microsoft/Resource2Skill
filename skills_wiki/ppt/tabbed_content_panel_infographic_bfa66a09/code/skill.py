import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

def create_slide(
    output_pptx_path: str,
    title_text: str = "INFOGRAPHIC",
    sample_texts: list = None,
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with a Tabbed Content Panel Infographic.

    Args:
        output_pptx_path: The path to save the generated PPTX file.
        title_text: The main title for the infographic panel.
        sample_texts: A list of 5 dictionaries, each with 'title' and 'body' keys.
        **kwargs: Not used, but included for compatibility.

    Returns:
        The path to the saved PPTX file.
    """

    # Helper for Open XML manipulation
    def qn(tag):
        nsmap = {
            'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
            'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
            'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
        }
        prefix, tagroot = tag.split(':')
        uri = nsmap[prefix]
        return f'{{{uri}}}{tagroot}'

    def add_shadow_to_shape(shape):
        sp = shape._element
        spPr = sp.get_or_add_spPr()
        
        effect_lst = etree.SubElement(spPr, qn('a:effectLst'))
        outer_shadow = etree.SubElement(effect_lst, qn('a:outerShdw'))
        outer_shadow.set('blurRad', '101600')
        outer_shadow.set('dist', '76200')
        outer_shadow.set('dir', '2700000') # 45 degrees
        outer_shadow.set('algn', 'br')
        outer_shadow.set('rotWithShape', '0')
        
        srgb_clr = etree.SubElement(outer_shadow, qn('a:srgbClr'))
        srgb_clr.set('val', '000000')
        alpha = etree.SubElement(srgb_clr, qn('a:alpha'))
        alpha.set('val', '35000') # 35% opacity

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(128, 128, 128)

    # === Layer 2: Visual Elements ===
    # Define colors
    PALETTE = [
        RGBColor(255, 192, 0),  # Yellow
        RGBColor(146, 208, 80),  # Green
        RGBColor(0, 176, 240),   # Blue
        RGBColor(112, 48, 160),  # Purple
        RGBColor(244, 112, 38)   # Orange
    ]
    TEXT_LIGHT_COLOR = RGBColor(255, 255, 255)
    TEXT_DARK_COLOR = RGBColor(0, 0, 0)
    TEXT_BODY_COLOR = RGBColor(89, 89, 89)

    # Main container
    container_left = Inches(0.5)
    container_top = Inches(0.5)
    container_width = Inches(15)
    container_height = Inches(8)
    
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, container_left, container_top, container_width, container_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.fill.background()
    
    # Adjust corner radius (16667 is a common value for a gentle curve)
    shape.adjustments[0] = 0.16667
    
    # Add shadow using lxml
    add_shadow_to_shape(shape)

    # Right side colored tabs and text
    tab_area_left = container_left + Inches(9.5)
    tab_area_width = Inches(5)
    tab_height = (container_height / 5) - Inches(0.2)
    
    for i, color in enumerate(PALETTE):
        top = container_top + Inches(0.7) + (i * (tab_height + Inches(0.2)))
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, tab_area_left, top, tab_area_width, tab_height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = color
        rect.line.fill.background()

        # Add Option Text
        txBox = slide.shapes.add_textbox(tab_area_left, top, tab_area_width, tab_height)
        p = txBox.text_frame.paragraphs[0]
        p.text = f"OPTION\n{i+1:02d}"
        p.font.name = 'Arial Black'
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_LIGHT_COLOR
        p.alignment = PP_ALIGN.CENTER
        txBox.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Left side content
    if sample_texts is None:
        sample_texts = [
            {'title': 'Sample Text', 'body': 'This is a sample text. Insert your desired text here.'}
        ] * 5

    content_area_left = container_left + Inches(0.5)
    content_area_width = Inches(8.5)
    row_height = container_height / 5
    
    for i in range(5):
        row_top = container_top + (i * row_height)
        
        # Icon placeholder
        slide.shapes.add_shape(
            MSO_SHAPE.OVAL, content_area_left, row_top + Inches(0.35), Inches(0.6), Inches(0.6)
        )

        # Title text
        txBox = slide.shapes.add_textbox(content_area_left + Inches(0.8), row_top + Inches(0.2), content_area_width, Inches(0.5))
        p = txBox.text_frame.paragraphs[0]
        p.text = sample_texts[i]['title']
        p.font.name = 'Arial'
        p.font.bold = True
        p.font.size = Pt(20)
        p.font.color.rgb = PALETTE[i]
        
        # Body text
        txBox_body = slide.shapes.add_textbox(content_area_left + Inches(0.8), row_top + Inches(0.55), content_area_width, Inches(0.5))
        p_body = txBox_body.text_frame.paragraphs[0]
        p_body.text = sample_texts[i]['body']
        p_body.font.name = 'Arial'
        p_body.font.size = Pt(12)
        p_body.font.color.rgb = TEXT_BODY_COLOR
        
        # Separator line
        if i < 4:
            line_top = container_top + ((i + 1) * row_height)
            slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, container_left + Inches(0.2), line_top, Inches(9.1), 0)

    # Main Title on the right panel
    title_box = slide.shapes.add_textbox(Inches(10.5), Inches(1), Inches(4), Inches(0.5))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial Black'
    p.font.size = Pt(24)
    p.font.color.rgb = TEXT_DARK_COLOR
    
    # Underline for title
    line = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(10.5), Inches(1.5), Inches(2), 0)
    line.line.color.rgb = TEXT_DARK_COLOR
    line.line.width = Pt(1.5)

    # Sample text on the right panel
    st_box = slide.shapes.add_textbox(Inches(10.5), Inches(1.8), Inches(4), Inches(2))
    p = st_box.text_frame.paragraphs[0]
    p.text = "Sample Text\n\nThis is a sample text. Insert your desired text here. This is a sample text."
    p.font.name = 'Arial'
    p_run = p.runs[0]
    p_run.font.bold = True
    p_run.font.size = Pt(18)
    p_run.font.color.rgb = TEXT_DARK_COLOR

    # Set font properties for the rest of the text
    st_box.text_frame.paragraphs[1].font.size = Pt(12)
    st_box.text_frame.paragraphs[1].font.color.rgb = TEXT_BODY_COLOR

    prs.save(output_pptx_path)
    return output_pptx_path
