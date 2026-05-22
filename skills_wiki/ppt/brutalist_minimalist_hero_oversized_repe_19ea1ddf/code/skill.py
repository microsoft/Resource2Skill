def create_slide(
    output_pptx_path: str,
    title_text: str = "HEY, I'M HUY",
    bg_color: tuple = (242, 242, 240),      # Warm off-white
    text_color: tuple = (26, 26, 26),       # Stark black
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Brutalist Minimalist Hero web design effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml import OxmlElement
    from pptx.oxml.ns import qn
    from PIL import Image, ImageDraw
    import os

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Helper: Convert RGB tuple to Hex string for lxml
    def rgb_to_hex(rgb):
        return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"

    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 1: Abstract Geometric Background via PIL ===
    # Simulates the subtle fractal/folded shapes in the video
    bg_img_path = "temp_abstract_bg.png"
    img_size = (800, 800)
    img = Image.new("RGBA", img_size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    
    # Draw overlapping translucent geometric shapes
    base_r, base_g, base_b = bg_color
    draw.polygon([(0, 0), (800, 200), (400, 800)], fill=(base_r-20, base_g-20, base_b-20, 100))
    draw.polygon([(800, 0), (800, 800), (100, 600)], fill=(base_r-15, base_g-15, base_b-10, 130))
    draw.polygon([(0, 400), (600, 0), (800, 800)], fill=(base_r-30, base_g-25, base_b-20, 80))
    img.save(bg_img_path)

    # Insert abstract background in the center
    pic = slide.shapes.add_picture(bg_img_path, Inches(4.16), Inches(1.25), width=Inches(5.0), height=Inches(5.0))
    
    # === Layer 2: Web UI Accents ===
    # Top Left Logo
    tx_logo = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(2), Inches(0.5))
    p_logo = tx_logo.text_frame.paragraphs[0]
    p_logo.text = "huyng*"
    p_logo.font.bold = True
    p_logo.font.size = Pt(16)
    p_logo.font.color.rgb = RGBColor(*text_color)
    
    # Top Right Nav
    tx_nav = slide.shapes.add_textbox(Inches(8.5), Inches(0.4), Inches(4.3), Inches(0.5))
    p_nav = tx_nav.text_frame.paragraphs[0]
    p_nav.text = "about    services    projects    Let's Talk."
    p_nav.font.size = Pt(10)
    p_nav.font.bold = True
    p_nav.font.color.rgb = RGBColor(*text_color)
    p_nav.alignment = PP_ALIGN.RIGHT

    # Bottom Right Scroll Indicator
    tx_scroll = slide.shapes.add_textbox(Inches(12.3), Inches(6.0), Inches(1.5), Inches(0.5))
    tx_scroll.rotation = 90
    p_scroll = tx_scroll.text_frame.paragraphs[0]
    p_scroll.text = "Scroll —"
    p_scroll.font.size = Pt(10)
    p_scroll.font.color.rgb = RGBColor(*text_color)

    # === Layer 3: Giant Repeating Hero Typography ===
    # Helper to apply outline via lxml
    def apply_text_outline(run, fill_hex, outline_hex, width_emu=12700):
        rPr = run._r.get_or_add_rPr()
        
        # Add solidFill (Matches background to look transparent/hollow)
        fill = OxmlElement('a:solidFill')
        srgb = OxmlElement('a:srgbClr')
        srgb.set('val', fill_hex)
        fill.append(srgb)
        rPr.append(fill)
        
        # Add ln (Outline Stroke)
        ln = OxmlElement('a:ln')
        ln.set('w', str(width_emu))
        ln_fill = OxmlElement('a:solidFill')
        ln_srgb = OxmlElement('a:srgbClr')
        ln_srgb.set('val', outline_hex)
        ln_fill.append(ln_srgb)
        ln.append(ln_fill)
        rPr.append(ln)

    # Create a large central text box
    tx_hero = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11.333), Inches(4))
    tf = tx_hero.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Line 1: Solid Text (Top)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.line_spacing = Pt(100) # Extremely tight spacing
    run1 = p1.add_run()
    run1.text = title_text.upper()
    run1.font.name = "Arial Black"
    run1.font.size = Pt(110)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(*text_color)

    # Line 2: Outlined Text (Middle)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.line_spacing = Pt(100)
    run2 = p2.add_run()
    run2.text = title_text.upper()
    run2.font.name = "Arial Black"
    run2.font.size = Pt(110)
    run2.font.bold = True
    # Apply lxml outline effect
    apply_text_outline(run2, rgb_to_hex(bg_color), rgb_to_hex(text_color), width_emu=15000)

    # Line 3: Outlined Text (Bottom)
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.line_spacing = Pt(100)
    run3 = p3.add_run()
    run3.text = title_text.upper()
    run3.font.name = "Arial Black"
    run3.font.size = Pt(110)
    run3.font.bold = True
    # Apply lxml outline effect
    apply_text_outline(run3, rgb_to_hex(bg_color), rgb_to_hex(text_color), width_emu=15000)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
