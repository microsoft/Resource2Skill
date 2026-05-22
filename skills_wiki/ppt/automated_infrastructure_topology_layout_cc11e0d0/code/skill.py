def create_slide(
    output_pptx_path: str,
    title_text: str = "Enterprise Network Architecture",
    body_text: str = "High-level topology mapping and security boundaries",
    bg_palette: str = "light",
    accent_color: tuple = (41, 128, 185),  # Core network blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing an Automated Infrastructure Topology Layout.
    Generates a hierarchical IT diagram with automated elbow connectors and security zones.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Colors ===
    c_text_dark = RGBColor(44, 62, 80)
    c_zone_bg = RGBColor(248, 250, 252)
    c_zone_line = RGBColor(189, 195, 199)
    
    c_cloud = RGBColor(236, 240, 241)
    c_firewall = RGBColor(231, 76, 60)      # Red for security
    c_switch = RGBColor(*accent_color)      # Primary accent for core routing
    c_server = RGBColor(127, 140, 141)      # Gray for compute
    c_client = RGBColor(173, 216, 230)      # Light blue for clients

    # === Header Content ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_text_dark

    if body_text:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(10), Inches(0.4))
        p_sub = sub_box.text_frame.paragraphs[0]
        p_sub.text = body_text
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = RGBColor(127, 140, 141)

    # === Layer 1: Security Zone (Background Plate) ===
    # Encompasses the firewall, switch, and all internal servers/clients
    zone = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.6), Inches(10.33), Inches(4.5))
    zone.fill.solid()
    zone.fill.fore_color.rgb = c_zone_bg
    zone.line.color.rgb = c_zone_line
    zone.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    zone.line.width = Pt(1.5)

    ztb = slide.shapes.add_textbox(Inches(1.6), Inches(2.7), Inches(3), Inches(0.4))
    ztb_p = ztb.text_frame.paragraphs[0]
    ztb_p.text = "Secure Internal Network (VPC)"
    ztb_p.font.size = Pt(11)
    ztb_p.font.color.rgb = RGBColor(149, 165, 166)
    ztb_p.font.bold = True

    # === Layer 2: Network Nodes ===
    nodes = {}

    def add_node(node_id, name, shape_type, cx_in, cy_in, w_in, h_in, fill_rgb, outline_rgb=RGBColor(149, 165, 166), text_inside=False, font_color=RGBColor(255,255,255)):
        cx, cy, w, h = Inches(cx_in), Inches(cy_in), Inches(w_in), Inches(h_in)
        left, top = cx - w/2, cy - h/2
        
        shape = slide.shapes.add_shape(shape_type, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
        shape.line.color.rgb = outline_rgb
        shape.line.width = Pt(1.5)
        
        if text_inside:
            shape.text_frame.text = name
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = shape.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(11)
            p.font.color.rgb = font_color
            p.font.bold = True
        else:
            # Add detached label below the shape
            tb = slide.shapes.add_textbox(Inches(cx_in - 1.5), Inches(cy_in + h_in/2 + 0.05), Inches(3.0), Inches(0.5))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text = name
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(11)
            p.font.color.rgb = c_text_dark
            p.font.bold = True
            
        nodes[node_id] = shape

    # Top Level: Internet
    add_node('internet', "Public Internet", MSO_SHAPE.CLOUD, 6.66, 1.4, 2.5, 1.5, c_cloud, text_inside=True, font_color=c_text_dark)
    
    # Mid Level: Security & Routing
    add_node('fw', "Edge Firewall", MSO_SHAPE.ROUNDED_RECTANGLE, 6.66, 3.2, 1.8, 0.7, c_firewall, text_inside=True)
    add_node('switch', "Core Switch / Router", MSO_SHAPE.RECTANGLE, 6.66, 4.5, 6.0, 0.5, c_switch, text_inside=True)
    
    # Bottom Level: Compute & Clients (Y = 6.0)
    add_node('app_srv', "Application Server", MSO_SHAPE.CAN, 3.0, 6.0, 1.0, 1.2, c_server)
    add_node('db_srv', "Primary Database", MSO_SHAPE.CAN, 5.0, 6.0, 1.0, 1.2, c_server)
    add_node('pc1', "Workstation A", MSO_SHAPE.ROUNDED_RECTANGLE, 8.3, 6.0, 1.2, 0.9, c_client, text_inside=False)
    add_node('pc2', "Workstation B", MSO_SHAPE.ROUNDED_RECTANGLE, 10.3, 6.0, 1.2, 0.9, c_client, text_inside=False)

    # === Layer 3: Orthogonal Connectors ===
    def connect_nodes(n1_id, n2_id, site1=2, site2=0):
        # site 2 = Bottom edge, site 0 = Top edge
        connector = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(0), Inches(0), Inches(1), Inches(1))
        connector.begin_connect(nodes[n1_id], site1)
        connector.end_connect(nodes[n2_id], site2)
        connector.line.color.rgb = RGBColor(127, 140, 141)
        connector.line.width = Pt(1.5)

    # Vertical trunk
    connect_nodes('internet', 'fw', site1=2, site2=0)
    connect_nodes('fw', 'switch', site1=2, site2=0)
    
    # Bus distribution to compute layer
    connect_nodes('switch', 'app_srv', site1=2, site2=0)
    connect_nodes('switch', 'db_srv', site1=2, site2=0)
    connect_nodes('switch', 'pc1', site1=2, site2=0)
    connect_nodes('switch', 'pc2', site1=2, site2=0)

    prs.save(output_pptx_path)
    return output_pptx_path
