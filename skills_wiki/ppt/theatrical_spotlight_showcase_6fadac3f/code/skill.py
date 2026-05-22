AMBIENT_CAPABLE = True

def create_slide(
    output_pptx_path: str,
    title_text: str = "Employee Spotlight",
    name_text: str = "Kelli Schwartz",
    role_text: str = "Public Guardian - Conservator",
    bg_color: tuple = (97, 24, 138),  # Deep purple
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Theatrical Spotlight Showcase' visual effect.
    Returns: path to the saved PPTX file.
    """
    import urllib.request
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement
    from PIL import Image

    # Optionally import shell helpers for ambient motion if available
    try:
        from _shell_helpers import add_pulse_loop
    except ImportError:
        def add_pulse_loop(*args, **kwargs):
            pass

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Helper function to inject transparency into a shape's solid fill
    def set_shape_transparency(shape, alpha_percent):
        spPr = shape.element.spPr
        solidFill = spPr.find('.//a:solidFill', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
        if solidFill is not None:
            srgbClr = solidFill.find('.//a:srgbClr', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
            if srgbClr is not None:
                alpha = OxmlElement('a:alpha')
                # Alpha val in OOXML is 0 to 100000 (e.g., 30000 = 30% opaque)
                alpha.set('val', str(int(alpha_percent * 1000)))
                srgbClr.append(alpha)

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*bg_color)
    bg.line.fill.background()

    # === Layer 2: Stage Elements (Truss Structure) ===
    # Top and bottom cords of the truss
    top_cord = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.2), Inches(13.333), Inches(0.1))
    top_cord.fill.solid()
    top_cord.fill.fore_color.rgb = RGBColor(0, 0, 0)
    top_cord.line.fill.background()

    bottom_cord = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.8), Inches(13.333), Inches(0.1))
    bottom_cord.fill.solid()
    bottom_cord.fill.fore_color.rgb = RGBColor(0, 0, 0)
    bottom_cord.line.fill.background()

    # Draw diagonal truss supports
    num_zigzags = 12
    step = 13.333 / num_zigzags
    for i in range(num_zigzags):
        x_start = i * step
        x_mid = x_start + (step / 2)
        x_end = x_start + step
        
        # Downward stroke
        line1 = slide.shapes.add_connector(
            MSO_SHAPE.LINE_INVERSE, Inches(x_start), Inches(0.3), Inches(x_mid), Inches(0.8)
        )
        line1.line.color.rgb = RGBColor(0, 0, 0)
        line1.line.width = Pt(3)
        
        # Upward stroke
        line2 = slide.shapes.add_connector(
            MSO_SHAPE.LINE_INVERSE, Inches(x_mid), Inches(0.8), Inches(x_end), Inches(0.3)
        )
        line2.line.color.rgb = RGBColor(0, 0, 0)
        line2.line.width = Pt(3)

    # === Layer 3: Floor Oval ===
    floor_width = Inches(8)
    floor_height = Inches(1.8)
    floor_left = (prs.slide_width - floor_width) / 2
    floor_top = Inches(6.0)
    floor = slide.shapes.add_shape(MSO_SHAPE.OVAL, floor_left, floor_top, floor_width, floor_height)
    floor.fill.solid()
    floor.fill.fore_color.rgb = RGBColor(255, 255, 255)
    floor.line.fill.background()
    # Add an ambient glow pulse to the stage floor
    add_pulse_loop(slide, floor, duration_ms=2500, scale_pct=102)

    # === Layer 4: Lights & Beams ===
    # Configuration for the 4 spotlights: (x_center, rotation, target_left_x, target_right_x)
    light_configs = [
        (Inches(1.8), 25, floor_left + Inches(0.5), floor_left + Inches(2.5)),
        (Inches(4.8), 10, floor_left + Inches(2.5), floor_left + Inches(4.5)),
        (Inches(8.5), -10, floor_left + Inches(3.5), floor_left + Inches(5.5)),
        (Inches(11.5), -25, floor_left + Inches(5.5), floor_left + Inches(7.5)),
    ]

    for x_c, rot, t_left, t_right in light_configs:
        # Draw Light Beam (Polygon)
        beam_builder = slide.shapes.build_freeform(x_c, Inches(1.1))
        beam_builder.add_line_segments([
            (t_right, Inches(6.9)),
            (t_left, Inches(6.9)),
            (x_c, Inches(1.1))
        ])
        beam = beam_builder.convert_to_shape()
        beam.fill.solid()
        beam.fill.fore_color.rgb = RGBColor(255, 255, 255)
        beam.line.fill.background()
        set_shape_transparency(beam, 30) # 30% opacity (70% transparent)

        # Draw Light Fixture (Trapezoid & Base)
        base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_c - Inches(0.15), Inches(0.9), Inches(0.3), Inches(0.2))
        base.fill.solid()
        base.fill.fore_color.rgb = RGBColor(0, 0, 0)
        base.line.fill.background()
        
        lamp = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, x_c - Inches(0.3), Inches(1.1), Inches(0.6), Inches(0.6))
        lamp.fill.solid()
        lamp.fill.fore_color.rgb = RGBColor(0, 0, 0)
        lamp.line.fill.background()
        lamp.rotation = rot

    # === Layer 5: Content Placement ===
    # 1. Top Title Text
    tx_box = slide.shapes.add_textbox(Inches(3), Inches(1.5), Inches(7.333), Inches(1.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = RGBColor(255, 255, 255)

    # 2. Subject Photo
    img_url = "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?crop=faces&fit=crop&w=600&h=600&q=80"
    img_width = Inches(2.8)
    img_height = Inches(2.8)
    img_left = (prs.slide_width - img_width) / 2
    img_top = Inches(2.6)

    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img_data = io.BytesIO(response.read())
            slide.shapes.add_picture(img_data, img_left, img_top, img_width, img_height)
    except Exception:
        # Fallback empty rectangle if download fails
        fallback = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, img_left, img_top, img_width, img_height)
        fallback.fill.solid()
        fallback.fill.fore_color.rgb = RGBColor(200, 200, 200)

    # 3. Subject Name and Role
    name_box = slide.shapes.add_textbox(Inches(2.66), Inches(5.6), Inches(8), Inches(1.0))
    tf_name = name_box.text_frame
    tf_name.word_wrap = True
    p_name = tf_name.paragraphs[0]
    p_name.text = name_text
    p_name.alignment = PP_ALIGN.CENTER
    p_name.font.size = Pt(40)
    p_name.font.bold = True
    p_name.font.color.rgb = RGBColor(0, 0, 0) # Dark text over white stage oval

    role_box = slide.shapes.add_textbox(Inches(2.66), Inches(6.4), Inches(8), Inches(0.8))
    tf_role = role_box.text_frame
    tf_role.word_wrap = True
    p_role = tf_role.paragraphs[0]
    p_role.text = role_text
    p_role.alignment = PP_ALIGN.CENTER
    p_role.font.size = Pt(24)
    p_role.font.color.rgb = RGBColor(50, 50, 50)

    prs.save(output_pptx_path)
    return output_pptx_path
