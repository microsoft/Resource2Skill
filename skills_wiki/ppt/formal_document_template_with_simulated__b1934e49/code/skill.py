import datetime
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_meeting_minutes_template(
    output_pptx_path: str,
    meeting_title: str = "ＯＯＯＯＯ會議",
    chairperson: str = "ＯＯＯＯＯ",
    location: str = "501 會議室",
    default_recorder: str = "請點擊選取",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a formal meeting minutes template,
    simulating dynamic date and dropdown fields.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        meeting_title: The title of the meeting.
        chairperson: The name of the chairperson.
 хрониlocation: The location of the meeting.
        default_recorder: The default text for the simulated recorder dropdown.

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Content & Simulated Fields ===

    # --- Title ---
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.5), prs.slide_width, Inches(1))
    p = title_box.text_frame.paragraphs[0]
    p.text = "會議記錄"
    p.font.name = "標楷體"
    p.font.size = Pt(36)
    p.alignment = PP_ALIGN.CENTER

    # --- Meeting Name ---
    meeting_name_box = slide.shapes.add_textbox(Inches(0), Inches(1.5), prs.slide_width, Inches(0.5))
    p = meeting_name_box.text_frame.paragraphs[0]
    p.text = f"「{meeting_title}」"
    p.font.name = "標楷體"
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER

    # --- Field Definitions ---
    fields = [
        ("壹、時    間：", ""), # Placeholder for date
        ("貳、開會地點：", location),
        ("參、主    席：", chairperson),
        ("肆、出席人員：", "(詳簽到單)"),
        ("伍、主席致詞：", "¶"),
        ("陸、報告事項：", "¶"),
        ("柒、會議決議：", "¶"),
    ]
    
    start_top = Inches(2.5)
    line_height = Inches(0.5)
    label_left = Inches(1.0)
    label_width = Inches(2.5)
    content_left = Inches(3.5)
    content_width = Inches(8.8)

    # --- Generate and Place Date Field ---
    # Calculate ROC year and format date string in Chinese
    today = datetime.date.today()
    roc_year = today.year - 1911
    weekday_map = {0: '一', 1: '二', 2: '三', 3: '四', 4: '五', 5: '六', 6: '日'}
    weekday_str = weekday_map[today.weekday()]
    date_str = f"{roc_year}年{today.month}月{today.day}日 (星期{weekday_str})"
    fields[0] = (fields[0][0], date_str)

    for i, (label_text, content_text) in enumerate(fields):
        # Label
        tx_box = slide.shapes.add_textbox(label_left, start_top + i * line_height, label_width, line_height)
        p = tx_box.text_frame.paragraphs[0]
        p.text = label_text
        p.font.name = "標楷體"
        p.font.bold = True
        p.font.size = Pt(16)
        
        # Content
        content_box = slide.shapes.add_textbox(content_left, start_top + i * line_height, content_width, line_height)
        p_content = content_box.text_frame.paragraphs[0]
        p_content.text = content_text if content_text != "¶" else ""
        p_content.font.name = "標楷體"
        p_content.font.size = Pt(16)
        
        if i == 0: # Style the date field
            fill = p_content.font.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(0, 0, 0)
            # This is a bit of a trick; there's no direct background setting for text runs.
            # A shape fill behind the text would be another way.
    
    # --- Simulated Dropdown for Recorder ---
    recorder_label_left = Inches(9.0)
    recorder_label_width = Inches(1.2)
    recorder_field_left = Inches(10.2)
    recorder_field_width = Inches(2.0)
    recorder_top = start_top + 2 * line_height # Align with "主席" field

    # Recorder Label
    tx_box = slide.shapes.add_textbox(recorder_label_left, recorder_top, recorder_label_width, line_height)
    p = tx_box.text_frame.paragraphs[0]
    p.text = "記錄："
    p.font.name = "標楷體"
    p.font.bold = True
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.RIGHT

    # Recorder Field (visual simulation)
    recorder_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, recorder_field_left, recorder_top, recorder_field_width, Inches(0.4))
    recorder_shape.fill.solid()
    recorder_shape.fill.fore_color.rgb = RGBColor(240, 240, 240)
    recorder_shape.line.fill.solid()
    recorder_shape.line.fill.fore_color.rgb = RGBColor(200, 200, 200)
    recorder_shape.line.width = Pt(0.5)
    
    tf = recorder_shape.text_frame
    tf.margin_left = Emu(91440) # Standard margin
    p_rec = tf.paragraphs[0]
    p_rec.text = default_recorder
    p_rec.font.name = "標楷體"
    p_rec.font.size = Pt(14)
    
    # Dropdown Arrow
    arrow_left = recorder_field_left + recorder_field_width - Inches(0.3)
    arrow_top = recorder_top + Inches(0.1)
    arrow_size = Inches(0.2)
    arrow = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, arrow_left, arrow_top, arrow_size, arrow_size)
    arrow.rotation = 180.0
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(100, 100, 100)
    arrow.line.fill.background()
    
    # Placeholder lines for details
    placeholder_text = "（一）" + "○" * 45
    placeholder_top = start_top + 7 * line_height + Inches(0.2)
    for i in range(5):
        ph_box = slide.shapes.add_textbox(Inches(1.5), placeholder_top + i * line_height, Inches(10), line_height)
        p_ph = ph_box.text_frame.paragraphs[0]
        p_ph.text = placeholder_text
        p_ph.font.name = "標楷體"
        p_ph.font.size = Pt(16)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
if __name__ == '__main__':
    create_meeting_minutes_template(
        "meeting_minutes_template.pptx",
        meeting_title="NDC ODF 工具導入研討會",
        chairperson="陳飛亨 老師",
        location="台北市電腦公會 501 會議室",
        default_recorder="陳小亨"
    )

