def create_slide(
    output_pptx_path: str,
    title_text: str = "DUOTONE",
    subtitle_text: str = "Aesthetic Design Pattern",
    bg_keyword: str = "portrait,fashion",  # Unsplash keyword
    color_shadows: tuple = (42, 12, 148),  # Navy Blue for darks
    color_highlights: tuple = (252, 0, 88), # Hot Pink for lights
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Pop-Art Duotone Aesthetic visual effect.
    """
    import io
    import requests
    from PIL import Image, ImageOps, ImageEnhance, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Download & Process Background Image ===
    try:
        # Fetch image from Unsplash
        url = f"https://images.unsplash.com/featured/?{bg_keyword}"
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        img = Image.open(io.BytesIO(response.content)).convert("RGB")
    except Exception as e:
        print(f"Failed to download image, using fallback: {e}")
        # Fallback: create a dummy gradient/noise image
        img = Image.new('RGB', (1280, 720), color=(100, 100, 100))
        draw = ImageDraw.Draw(img)
        for i in range(0, 720, 20):
            draw.line([(0, i), (1280, i)], fill=(150, 150, 150), width=10)

    # Crop to fit 16:9 aspect ratio exactly
    target_ratio = 16.0 / 9.0
    img_ratio = img.width / img.height

    if img_ratio > target_ratio:
        # Image is wider, crop width
        new_width = int(target_ratio * img.height)
        offset = (img.width - new_width) / 2
        img = img.crop((offset, 0, img.width - offset, img.height))
    elif img_ratio < target_ratio:
        # Image is taller, crop height
        new_height = int(img.width / target_ratio)
        offset = (img.height - new_height) / 2
        img = img.crop((0, offset, img.width, img.height - offset))

    # Resize to standardize processing (optional, helps with memory)
    img = img.resize((1920, 1080), Image.Resampling.LANCZOS)

    # Apply Duotone Effect using PIL
    # 1. Convert to grayscale (Luminance)
    img_gray = img.convert("L")
    
    # 2. Boost contrast slightly so the duotone pops harder
    enhancer = ImageEnhance.Contrast(img_gray)
    img_gray = enhancer.enhance(1.3)
    
    # 3. Apply Gradient Map (colorize: black goes to color_shadows, white goes to color_highlights)
    img_duotone = ImageOps.colorize(img_gray, black=color_shadows, white=color_highlights)

    # Save to memory stream
    img_stream = io.BytesIO()
    img_duotone.save(img_stream, format='PNG')
    img_stream.seek(0)

    # Add to slide
    slide.shapes.add_picture(
        img_stream, 
        0, 0, 
        width=prs.slide_width, 
        height=prs.slide_height
    )

    # === Layer 2: Text Overlay ===
    # Add strong, bold text overlay to complete the look
    tx_box = slide.shapes.add_textbox(
        Inches(1), Inches(4.5), 
        Inches(11.333), Inches(2.5)
    )
    tf = tx_box.text_frame
    tf.word_wrap = True

    # Title
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(88)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.LEFT

    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.bold = True
    p2.font.size = Pt(32)
    p2.font.name = "Arial"
    p2.font.color.rgb = RGBColor(255, 255, 255)
    p2.alignment = PP_ALIGN.LEFT

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
