def create_slide(
    output_pptx_path: str,
    chapter_chars: list = ["C", "H", "1"],
    chapter_title: str = "EXECUTIVE SUMMARY",
    chapter_desc: str = "This section provides an overview of our strategic performance, key milestones, and the overarching vision for the upcoming fiscal year.",
    accent_color: tuple = (238, 77, 77),
    image_keyword: str = "office",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Crimson & Grayscale Split Screen Divider" visual effect.
    """
    import os
    import io
    import urllib.request
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.3333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Left Solid Background ===
    left_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(0), 
        Inches(6.6666), Inches(7.5)
    )
    left_rect.fill.solid()
    left_rect.fill.fore_color.rgb = RGBColor(*accent_color)
    left_rect.line.fill.background() # Hide border

    # === Layer 2: Right Grayscale Image (via PIL) ===
    # Download and process the right-half image
    try:
        url = f"https://picsum.photos/seed/{image_keyword}/800/900?grayscale"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(io.BytesIO(response.read())).convert('L') # Force grayscale
    except Exception:
        # Fallback to a solid gray image if download fails
        img = Image.new('L', (800, 900), color=100)

    # Crop precisely to the aspect ratio of the half-slide (6.6667 / 7.5)
    target_ratio = 6.6667 / 7.5
    img_ratio = img.width / img.height

    if img_ratio > target_ratio:
        new_w = int(img.height * target_ratio)
        left_crop = (img.width - new_w) / 2
        img = img.crop((left_crop, 0, left_crop + new_w, img.height))
    else:
        new_h = int(img.width / target_ratio)
        top_crop = (img.height - new_h) / 2
        img = img.crop((0, top_crop, img.width, top_crop + new_h))

    img_path = "temp_split_bg.png"
    img.save(img_path)
    # Insert right image
    slide.shapes.add_picture(img_path, Inches(6.6666), Inches(0), width=Inches(6.6667), height=Inches(7.5))
    os.remove(img_path)

    # === Layer 3: Central Floating Circles ===
    cy = 3.75
    radius = 1.0
    spacing = 2.2 # Distance between circle centers
    centers_x = [6.6666 - spacing, 6.6666, 6.6666 + spacing]

    for i, cx in enumerate(centers_x):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx - radius), Inches(cy - radius), 
            Inches(radius * 2), Inches(radius * 2)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
        circle.line.fill.background()
        
        # Inject XML Drop Shadow for the "floating" effect
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="150000" dist="40000" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="20000"/>
                </a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        circle.element.spPr.append(parse_xml(shadow_xml))
        
        # Add character text to circle
        tf = circle.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = chapter_chars[i] if i < len(chapter_chars) else ""
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(70, 70, 70)
        p.font.name = 'Arial'

    # === Layer 4: Typography & Information (Left side) ===
    # Small top label
    tx_top = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(4.5), Inches(0.5))
    p_top = tx_top.text_frame.paragraphs[0]
    p_top.text = "SECTION REPORT"
    p_top.font.size = Pt(10)
    p_top.font.color.rgb = RGBColor(255, 255, 255)
    p_top.font.bold = True
    p_top.font.name = 'Arial'

    # Bottom Title and Description
    tx_desc = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(5.0), Inches(2.0))
    tf_desc = tx_desc.text_frame
    tf_desc.word_wrap = True
    
    p_title = tf_desc.paragraphs[0]
    p_title.text = chapter_title
    p_title.font.bold = True
    p_title.font.size = Pt(28)
    p_title.font.color.rgb = RGBColor(255, 255, 255)
    p_title.font.name = 'Arial'
    
    p_desc = tf_desc.add_paragraph()
    p_desc.text = chapter_desc
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = RGBColor(255, 255, 255)
    p_desc.font.name = 'Arial'
    p_desc.space_before = Pt(12)

    prs.save(output_pptx_path)
    return output_pptx_path
