import os
import random
import urllib.request
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from lxml import etree

def create_slide(
    output_pptx_path: str,
    title_text: str = "MODIFIED IN 20 MINUTES",
    subtitle_text: str = "USING PYTHON-PPTX",
    bg_keyword: str = "cityscape",
    accent_color: tuple = (255, 213, 0),  # Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dynamic 2.5D Cutout Spotlight" visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Helper function to add drop shadow via lxml
    def add_drop_shadow(shape):
        spPr = shape.element.spPr
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="50800" dist="50800" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="50000"/>
                </a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        effectLst = etree.fromstring(shadow_xml)
        spPr.append(effectLst)

    # === Layer 1: Background Image ===
    bg_img_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1920x1080/?{bg_keyword},landscape"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(bg_img_path, 'wb') as f:
                f.write(response.read())
        slide.shapes.add_picture(bg_img_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
    except Exception:
        # Fallback to dark solid background if download fails
        bg = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(40, 45, 55)
        bg.line.fill.background()

    # === Layer 2: Subject Cutout (Generated Mockup via PIL) ===
    cutout_path = "temp_cutout.png"
    # Create a stylized "person silhouette" placeholder looking over the city
    cutout_img = Image.new('RGBA', (600, 800), (0, 0, 0, 0))
    draw = ImageDraw.Draw(cutout_img)
    # Draw silhouette (dark jacket)
    draw.polygon([(250, 200), (100, 350), (50, 800), (550, 800), (480, 380), (350, 200)], fill=(30, 35, 40, 255))
    # Draw head
    draw.ellipse([210, 50, 390, 230], fill=(220, 180, 150, 255)) # Skin tone head
    # Jacket details/shading
    draw.line([(250, 200), (300, 800)], fill=(20, 25, 30, 255), width=10)
    cutout_img.save(cutout_path)
    
    # Insert cutout bottom-right to create parallax depth
    cutout_shape = slide.shapes.add_picture(
        cutout_path, 
        Inches(8), Inches(2.5), 
        width=Inches(4.5), height=Inches(5.0)
    )
    add_drop_shadow(cutout_shape)

    # === Layer 3: Staggered Highlight Text Blocks ===
    
    # 3a. Main Title Box (Yellow)
    title_box = slide.shapes.add_shape(
        1, # MSO_SHAPE.RECTANGLE
        Inches(1.5), Inches(3.5), 
        width=Inches(7.0), height=Inches(1.2)
    )
    title_box.fill.solid()
    title_box.fill.fore_color.rgb = RGBColor(*accent_color)
    title_box.line.fill.background() # No line
    add_drop_shadow(title_box)

    text_frame = title_box.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = f" {title_text.upper()} "
    run.font.name = "Arial Black" # Fallback to standard bold font
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 30, 30)

    # 3b. Subtitle Box (White)
    sub_box = slide.shapes.add_shape(
        1, 
        Inches(1.5), Inches(4.7), 
        width=Inches(5.0), height=Inches(0.7)
    )
    sub_box.fill.solid()
    sub_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    sub_box.line.fill.background()
    add_drop_shadow(sub_box)

    text_frame_sub = sub_box.text_frame
    text_frame_sub.clear()
    p_sub = text_frame_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.LEFT
    run_sub = p_sub.add_run()
    run_sub.text = f" {subtitle_text.upper()} "
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(28)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(40, 90, 180) # Deep blue

    # === Layer 4: Atmospheric Particles (Snow) via PIL ===
    particles_path = "temp_particles.png"
    part_img = Image.new('RGBA', (1920, 1080), (0, 0, 0, 0))
    part_draw = ImageDraw.Draw(part_img)
    
    # Generate random overlapping particles
    for _ in range(120):
        x = random.randint(0, 1920)
        y = random.randint(0, 1080)
        radius = random.randint(2, 6)
        alpha = random.randint(50, 220)
        part_draw.ellipse([x, y, x + radius, y + radius], fill=(255, 255, 255, alpha))
        
    part_img.save(particles_path)
    
    # Overlay particles over the entire slide
    slide.shapes.add_picture(
        particles_path, 
        Inches(0), Inches(0), 
        width=prs.slide_width, height=prs.slide_height
    )

    # Cleanup temp files and save
    prs.save(output_pptx_path)
    
    for tmp_file in [bg_img_path, cutout_path, particles_path]:
        if os.path.exists(tmp_file):
            os.remove(tmp_file)
            
    return output_pptx_path
