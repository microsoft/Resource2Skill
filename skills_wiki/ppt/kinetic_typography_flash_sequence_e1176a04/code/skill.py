def create_slide(
    output_pptx_path: str,
    title_text: str = "IGNORED", # Kept for signature compatibility
    body_text: str = "IGNORED",  # Kept for signature compatibility
    bg_palette: str = "kinetic", 
    accent_color: tuple = (246, 215, 87),  # Mustard Yellow Background
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Kinetic Typography Flash Sequence" effect.
    Generates multiple auto-advancing slides to create a lyric video effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Define the kinetic sequence: (Text, Duration in milliseconds)
    # This simulates the rapid pacing seen in the tutorial.
    kinetic_sequence = [
        ("5 AM", 1000),
        ("AND", 400),
        ("WE ARE", 600),
        ("IN", 250),
        ("TROUBLE.", 1500),
        ("BUT", 400),
        ("WE DON'T", 800),
        ("REALLY", 600),
        ("CARE.", 1500)
    ]

    text_color = RGBColor(15, 15, 15)  # Near black
    bg_color_rgb = RGBColor(*accent_color)

    # XML Namespace for PowerPoint
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    nsmap = {'p': p_ns}

    for word, duration_ms in kinetic_sequence:
        # Add blank slide
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Set Solid Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color_rgb

        # Create centered text box
        # Spanning full width to ensure long words fit, vertically centered
        box_height = Inches(2)
        top_pos = (prs.slide_height - box_height) / 2
        
        txBox = slide.shapes.add_textbox(0, top_pos, prs.slide_width, box_height)
        tf = txBox.text_frame
        tf.word_wrap = False
        
        p = tf.paragraphs[0]
        p.text = word
        p.alignment = PP_ALIGN.CENTER
        
        # Style text (Heavy, wide sans-serif)
        font = p.font
        font.name = 'Arial Black' # Safe fallback for heavy font
        font.size = Pt(90)
        font.bold = True
        font.color.rgb = text_color

        # ==========================================
        # XML INJECTION: Slide Auto-Advance Timing
        # ==========================================
        # We need to add/modify <p:transition advTm="duration_ms"/>
        
        # Get the underlying XML element for the slide
        sld_xml = slide.element
        
        # Look for existing transition element
        transition = sld_xml.find('.//p:transition', namespaces=nsmap)
        
        if transition is None:
            # Create the transition element if it doesn't exist
            # It must be inserted in a specific order in the XML schema,
            # usually before <p:timing> or <p:extLst>. For simplicity, appending 
            # to the end of the slide element usually works for modern PPTX engines.
            transition = etree.SubElement(sld_xml, f"{{{p_ns}}}transition")
            
            # Set transition type to "None" (instant cut)
            # You can add <p:none/> or <p:fade/> as a child if needed, 
            # but an empty transition tag defaults to None/Cut.
            
        # advTm is the advance time in milliseconds
        transition.set('advTm', str(duration_ms))
        # advClick="0" disables advancing purely on click, forcing the timer
        # (Though keeping it default is safer so users can still click through if stuck)
        
    prs.save(output_pptx_path)
    return output_pptx_path
