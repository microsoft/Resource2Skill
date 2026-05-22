def create_slide(
    output_pptx_path: str,
    title_text: str = "SUB TITLE HERE",
    body_text: str = "Type your detailed text here. This is a place to enter your brochure text and describe the features in detail.",
    bg_palette: str = "cityscape",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Trifold Panoramic Gradient Blend' visual effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    # --- Configuration ---
    SLIDE_WIDTH_IN = 13.333
    SLIDE_HEIGHT_IN = 7.5
    CANVAS_W = 1920
    CANVAS_H = 1080
    
    COLOR_YELLOW = RGBColor(244, 180, 26)
    COLOR_YELLOW_TUPLE = (244, 180, 26)
    COLOR_GREEN = RGBColor(30, 107, 82)
    COLOR_DARK = RGBColor(51, 51, 51)
    
    # === Layer 1: Background Image Compositing via PIL ===
    # 1. Create Base Columns
    bg_img = Image.new('RGB', (CANVAS_W, CANVAS_H), (255, 255, 255))
    draw = ImageDraw.Draw(bg_img)
    # Center Yellow Panel
    draw.rectangle([CANVAS_W // 3, 0, 2 * CANVAS_W // 3, CANVAS_H], fill=COLOR_YELLOW_TUPLE)
    
    # 2. Fetch Cityscape Image
    target_h = 800
    try:
        url = "https://images.unsplash.com/photo-1512453979798-5ea266f8880c?q=80&w=1920&auto=format&fit=crop"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            city_img = Image.open(BytesIO(response.read())).convert('RGBA')
    except Exception:
        # Fallback to a solid blue-grey if network fails
        city_img = Image.new('RGBA', (CANVAS_W, target_h), (70, 90, 120, 255))
        
    # Resize and crop city image
    aspect = city_img.width / city_img.height
    new_w = max(CANVAS_W, int(target_h * aspect))
    city_img = city_img.resize((new_w, target_h), Image.LANCZOS)
    left_crop = (new_w - CANVAS_W) // 2
    city_img = city_img.crop((left_crop, 0, left_crop + CANVAS_W, target_h))
    
    # Full canvas size layer for the city
    city_layer = Image.new('RGBA', (CANVAS_W, CANVAS_H), (0, 0, 0, 0))
    city_layer.paste(city_img, (0, CANVAS_H - target_h))
    
    # 3. Create the Arch Mask with Gradient Fade
    mask = Image.new('L', (CANVAS_W, CANVAS_H), 0)
    mask_draw = ImageDraw.Draw(mask)
    
    # Draw large sweeping arch (ellipse)
    arch_top = 250
    mask_draw.ellipse([-300, arch_top, CANVAS_W + 300, CANVAS_H + 800], fill=255)
    
    # Apply vertical gradient fade to the mask to blend the top of the arch
    fade_start = arch_top
    fade_end = arch_top + 450
    pixels = mask.load()
    for y in range(CANVAS_H):
        for x in range(CANVAS_W):
            if pixels[x, y] > 0:
                if y < fade_end:
                    factor = max(0, min(1.0, (y - fade_start) / (fade_end - fade_start)))
                    pixels[x, y] = int(255 * factor)
                else:
                    pixels[x, y] = 255
                    
    # Composite the faded city over the columns
    bg_img.paste(city_layer, (0, 0), mask)
    
    # Save temp background
    bg_path = "temp_trifold_bg.png"
    bg_img.save(bg_path, format="PNG")

    # === Layer 2: PPTX Assembly ===
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Apply generated background
    slide.shapes.add_picture(bg_path, 0, 0, width=Inches(SLIDE_WIDTH_IN), height=Inches(SLIDE_HEIGHT_IN))
    
    col_width = SLIDE_WIDTH_IN / 3
    
    # === Layer 3: Vector Accents ===
    # Top Right Yellow Box (extends the yellow motif)
    tr_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_width * 2), Inches(0), Inches(col_width), Inches(2.2))
    tr_box.fill.solid()
    tr_box.fill.fore_color.rgb = COLOR_YELLOW
    tr_box.line.fill.background()
    
    # Bottom Left Dark Green Accent
    bl_shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.0), Inches(5.0), Inches(3.5), Inches(3.5))
    bl_shape.fill.solid()
    bl_shape.fill.fore_color.rgb = COLOR_GREEN
    bl_shape.line.fill.background()
    
    # Bottom Right Dark Green Accent
    br_shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(4.5), Inches(5.0), Inches(5.0))
    br_shape.fill.solid()
    br_shape.fill.fore_color.rgb = COLOR_GREEN
    br_shape.line.fill.background()
    
    # === Layer 4: Content Construction ===
    
    # -- Left Column (Bullets) --
    bullet_y_positions = [2.0, 3.2, 4.4]
    for i, y in enumerate(bullet_y_positions):
        # Bullet circle
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(y), Inches(0.4), Inches(0.4))
        circle.fill.solid()
        circle.fill.fore_color.rgb = COLOR_GREEN
        circle.line.fill.background()
        
        # Bullet text
        tb = slide.shapes.add_textbox(Inches(1.2), Inches(y - 0.1), Inches(2.8), Inches(1.0))
        p_title = tb.text_frame.add_paragraph()
        p_title.text = f"Your Title {i+1:02d}"
        p_title.font.bold = True
        p_title.font.size = Pt(14)
        p_title.font.color.rgb = COLOR_DARK
        
        p_desc = tb.text_frame.add_paragraph()
        p_desc.text = "Type your detail text here. Enter text here to explain."
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = RGBColor(100, 100, 100)

    # -- Center Column (Main Info) --
    center_tb = slide.shapes.add_textbox(Inches(col_width + 0.5), Inches(1.2), Inches(col_width - 1.0), Inches(2.0))
    center_p1 = center_tb.text_frame.add_paragraph()
    center_p1.text = title_text.upper()
    center_p1.font.bold = True
    center_p1.font.size = Pt(20)
    center_p1.font.color.rgb = COLOR_DARK
    center_p1.alignment = PP_ALIGN.CENTER
    
    center_p2 = center_tb.text_frame.add_paragraph()
    center_p2.text = body_text
    center_p2.font.size = Pt(11)
    center_p2.font.color.rgb = RGBColor(80, 80, 80)
    center_p2.alignment = PP_ALIGN.CENTER
    
    # -- Right Column (Header & Footer) --
    right_header = slide.shapes.add_textbox(Inches(col_width * 2 + 0.3), Inches(0.4), Inches(col_width - 0.6), Inches(1.5))
    rh_p1 = right_header.text_frame.add_paragraph()
    rh_p1.text = "BROCHURE"
    rh_p1.font.bold = True
    rh_p1.font.size = Pt(28)
    rh_p1.font.color.rgb = COLOR_DARK
    rh_p1.alignment = PP_ALIGN.RIGHT
    
    rh_p2 = right_header.text_frame.add_paragraph()
    rh_p2.text = "DESIGN | 2024"
    rh_p2.font.size = Pt(14)
    rh_p2.font.color.rgb = COLOR_DARK
    rh_p2.alignment = PP_ALIGN.RIGHT
    
    right_footer = slide.shapes.add_textbox(Inches(col_width * 2 + 0.5), Inches(6.5), Inches(col_width - 1.0), Inches(1.0))
    rf_p = right_footer.text_frame.add_paragraph()
    rf_p.text = "COMPANY\nNAME"
    rf_p.font.bold = True
    rf_p.font.size = Pt(16)
    rf_p.font.color.rgb = COLOR_DARK
    rf_p.alignment = PP_ALIGN.RIGHT

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
