def create_slide(
    output_pptx_path: str,
    title_text: str = "YOUR TITLE",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa. Fusce posuere, magna sed pulvinar ultricies, purus lectus.",
    accent_color: tuple = (71, 201, 157),  # Seafoam green
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Radial Origami Infographic Pivot' visual effect.
    """
    import os
    import math
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.line import MSO_LINE, MSO_ARROWHEAD
    from PIL import Image, ImageDraw, ImageFilter

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Constants & Coordinates
    cx = Inches(6.666)
    cy = Inches(3.75)
    r_base_inch = 1.6  # Base circle radius
    
    # ---------------------------------------------------------
    # 1. Native PPTX: Background & Base Structural Elements
    # ---------------------------------------------------------
    # Add a subtle off-white background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(250, 250, 250)
    bg.line.fill.background()

    # Dark Blue Base Circle
    base_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        cx - Inches(r_base_inch), cy - Inches(r_base_inch), 
        Inches(r_base_inch * 2), Inches(r_base_inch * 2)
    )
    base_circle.fill.solid()
    base_circle.fill.fore_color.rgb = RGBColor(43, 96, 113)
    base_circle.line.fill.background()

    # Base Circle Intersecting Grid Lines
    grid_angles = [0, 45, 90, 135]
    for a in grid_angles:
        rad = math.radians(a)
        start_x = cx + Inches(r_base_inch) * math.cos(rad)
        start_y = cy + Inches(r_base_inch) * math.sin(rad)
        end_x = cx - Inches(r_base_inch) * math.cos(rad)
        end_y = cy - Inches(r_base_inch) * math.sin(rad)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y)
        line.line.color.rgb = RGBColor(30, 70, 85)
        line.line.width = Pt(1.5)

    # ---------------------------------------------------------
    # 2. PIL: Generate the 3D Folds Overlay (Pixel-Perfect Shadows)
    # ---------------------------------------------------------
    def generate_folds_overlay(filepath):
        size = 1000  # 1000px = 4.0 Inches when inserted (250 ppi)
        canvas = Image.new('RGBA', (size, size), (0, 0, 0, 0))

        # Helper to generate a diagonal gradient
        def make_grad(c1, c2, angle):
            grad = Image.new('RGBA', (2000, 1))
            d = ImageDraw.Draw(grad)
            for i in range(2000):
                r = int(c1[0] + (c2[0] - c1[0]) * i / 1999)
                g = int(c1[1] + (c2[1] - c1[1]) * i / 1999)
                b = int(c1[2] + (c2[2] - c1[2]) * i / 1999)
                d.point((i, 0), fill=(r, g, b, 255))
            return grad.resize((2000, 2000)).rotate(angle).crop((500, 500, 1500, 1500))

        # The base circle in PIL corresponds to 800px diameter (400px radius)
        # Bounding box: [100, 100, 900, 900]
        
        # Fold 1: Main Top Wrap (3/4 Pie shape)
        pie_mask = Image.new('L', (size, size), 0)
        # Pillow draws clockwise. 180 to 90 covers top-left, top-right, bottom-right.
        ImageDraw.Draw(pie_mask).pieslice([100, 100, 900, 900], start=180, end=90, fill=255)
        
        # Shadow for pie
        shadow_pie = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        shadow_pie.paste((0, 0, 0, 80), (15, 15), pie_mask.filter(ImageFilter.GaussianBlur(18)))
        canvas.alpha_composite(shadow_pie)
        
        # Gradient for pie
        pie_grad = make_grad((71, 201, 157), (200, 240, 230), 45)
        pie_layer = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        pie_layer.paste(pie_grad, (0, 0), pie_mask)
        canvas.alpha_composite(pie_layer)

        # Fold 2: Bottom-Left Triangle Flap
        tri_mask = Image.new('L', (size, size), 0)
        ImageDraw.Draw(tri_mask).polygon([(500, 500), (100, 500), (500, 900)], fill=255)
        
        # Shadow for triangle
        shadow_tri = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        shadow_tri.paste((0, 0, 0, 90), (12, 12), tri_mask.filter(ImageFilter.GaussianBlur(15)))
        canvas.alpha_composite(shadow_tri)
        
        # Gradient for triangle
        tri_grad = make_grad((35, 140, 105), (71, 201, 157), 135)
        tri_layer = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        tri_layer.paste(tri_grad, (0, 0), tri_mask)
        canvas.alpha_composite(tri_layer)

        # Center Motif: Bright Inner Circle
        center_mask = Image.new('L', (size, size), 0)
        ImageDraw.Draw(center_mask).ellipse([280, 280, 720, 720], fill=255)
        
        # Shadow for center circle
        shadow_center = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        shadow_center.paste((0, 0, 0, 60), (0, 10), center_mask.filter(ImageFilter.GaussianBlur(15)))
        canvas.alpha_composite(shadow_center)
        
        # Gradient for center circle
        center_grad = make_grad((220, 250, 240), (255, 255, 255), 90)
        center_layer = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        center_layer.paste(center_grad, (0, 0), center_mask)
        canvas.alpha_composite(center_layer)

        canvas.save(filepath)

    temp_image_path = "temp_folds_overlay.png"
    generate_folds_overlay(temp_image_path)

    # ---------------------------------------------------------
    # 3. Native PPTX: Connectors (Arrows under the overlay)
    # ---------------------------------------------------------
    angles_deg = [135, 90, 45, 0, -45]
    length_inch = 2.8

    for i, a in enumerate(angles_deg):
        rad = math.radians(a)
        # Start at center (will be hidden by overlay)
        end_x = cx + Inches(length_inch) * math.cos(rad)
        end_y = cy - Inches(length_inch) * math.sin(rad)
        
        arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, cy, end_x, end_y)
        arrow.line.color.rgb = RGBColor(100, 100, 100)
        arrow.line.width = Pt(1.5)
        arrow.line.dash_style = MSO_LINE.DASH
        arrow.line.end_arrowhead = MSO_ARROWHEAD.STEALTH

    # ---------------------------------------------------------
    # 4. Native PPTX: Insert Folds Overlay & Center Text
    # ---------------------------------------------------------
    # The overlay is 4x4 inches, perfectly centered over our 3.2x3.2 inch base circle
    overlay_size = Inches(4.0)
    slide.shapes.add_picture(temp_image_path, cx - overlay_size/2, cy - overlay_size/2, overlay_size, overlay_size)
    os.remove(temp_image_path)  # Cleanup

    # Center Text "5"
    tx_center = slide.shapes.add_textbox(cx - Inches(0.5), cy - Inches(0.5), Inches(1.0), Inches(1.0))
    tf_center = tx_center.text_frame
    tf_center.text = "5"
    tf_center.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_center.paragraphs[0].font.size = Pt(64)
    tf_center.paragraphs[0].font.bold = True
    tf_center.paragraphs[0].font.color.rgb = RGBColor(*accent_color)

    # ---------------------------------------------------------
    # 5. Native PPTX: Surrounding Content & Typography
    # ---------------------------------------------------------
    # Manual offsets to perfectly align text blocks to the end of the arrows
    text_offsets = [
        (-Inches(1.8), -Inches(0.6)),  # 135
        (-Inches(0.75), -Inches(0.9)), # 90
        (Inches(0.2), -Inches(0.6)),   # 45
        (Inches(0.2), -Inches(0.3)),   # 0
        (Inches(0.2), Inches(0.1))     # -45
    ]

    for i, a in enumerate(angles_deg):
        rad = math.radians(a)
        end_x = cx + Inches(length_inch) * math.cos(rad)
        end_y = cy - Inches(length_inch) * math.sin(rad)
        
        tx_box = slide.shapes.add_textbox(
            end_x + text_offsets[i][0], 
            end_y + text_offsets[i][1], 
            Inches(1.5), Inches(0.8)
        )
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        # Option Title
        p_title = tf.paragraphs[0]
        p_title.text = f"OPTION {i+1}"
        p_title.font.bold = True
        p_title.font.size = Pt(14)
        p_title.font.color.rgb = RGBColor(60, 60, 60)
        p_title.alignment = PP_ALIGN.CENTER if a == 90 else (PP_ALIGN.LEFT if a in [45, 0, -45] else PP_ALIGN.RIGHT)
        
        # Body text
        p_body = tf.add_paragraph()
        p_body.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit."
        p_body.font.size = Pt(9)
        p_body.font.color.rgb = RGBColor(120, 120, 120)
        p_body.alignment = p_title.alignment

    # Main Title (Bottom Left)
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.8), Inches(4.0), Inches(2.0))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    
    # "YOUR" (Accent color)
    p_main = tf_title.paragraphs[0]
    run1 = p_main.add_run()
    run1.text = title_text.split(" ")[0] + " "
    run1.font.bold = True
    run1.font.size = Pt(28)
    run1.font.color.rgb = RGBColor(200, 70, 70)  # Red Accent
    
    # "TITLE" (Dark)
    if len(title_text.split(" ")) > 1:
        run2 = p_main.add_run()
        run2.text = " ".join(title_text.split(" ")[1:])
        run2.font.bold = True
        run2.font.size = Pt(28)
        run2.font.color.rgb = RGBColor(50, 50, 50)

    # Accent divider line under title
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(5.35), Inches(2.0), Pt(2))
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(200, 70, 70)
    divider.line.fill.background()

    # Main Body Text
    p_desc = tf_title.add_paragraph()
    p_desc.text = "\n" + body_text
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = RGBColor(100, 100, 100)

    prs.save(output_pptx_path)
    return output_pptx_path
