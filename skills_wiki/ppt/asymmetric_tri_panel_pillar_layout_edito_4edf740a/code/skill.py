import os
import urllib.request
from io import BytesIO
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "A SLIDE,\nBEAUTIFUL\nAS THIS!",
    accent_text: str = "JUST ONE",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam.",
    bg_keyword: str = "urban,street",
    accent_color: tuple = (217, 119, 54),  # Burnt Orange
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Asymmetric Tri-Panel Pillar Layout'.
    """
    # 1. Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # Theme Colors
    color_dark_text = RGBColor(51, 51, 51)
    color_accent = RGBColor(*accent_color)
    color_pillar_light = RGBColor(248, 248, 248)
    color_pillar_dark = RGBColor(45, 45, 45)
    
    # Layout Measurements
    slide_w_in = 13.333
    slide_h_in = 7.5
    left_area_w = 5.0
    right_area_w = slide_w_in - left_area_w
    pillar_w = right_area_w / 3.0  # Approx 2.777 inches
    
    # ==========================================
    # LAYER 1: THE LEFT CONTENT AREA
    # ==========================================
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(3.5), Inches(2.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    
    run_main = p.add_run()
    run_main.text = title_text + "\n"
    run_main.font.size = Pt(36)
    run_main.font.color.rgb = color_dark_text
    run_main.font.name = "Arial"
    
    run_accent = p.add_run()
    run_accent.text = accent_text
    run_accent.font.size = Pt(36)
    run_accent.font.color.rgb = color_accent
    run_accent.font.name = "Arial"
    
    # Body Box
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(3.5), Inches(2.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    
    p_body = tf_body.add_paragraph()
    p_body.alignment = PP_ALIGN.CENTER
    p_body.text = body_text
    p_body.font.size = Pt(14)
    p_body.font.color.rgb = RGBColor(120, 120, 120)
    
    # ==========================================
    # LAYER 2: THE THREE RIGHT PILLARS
    # ==========================================
    
    # Helper function to generate Pillow image with overlay for Pillar 1
    def generate_pillar1_image(width_in, height_in, keyword):
        dpi = 150
        px_w, px_h = int(width_in * dpi), int(height_in * dpi)
        
        # Fallback solid image
        base_img = Image.new('RGBA', (px_w, px_h), (100, 100, 100, 255))
        
        try:
            url = f"https://images.unsplash.com/photo-1517462964-21fdcec3f25b?ixlib=rb-4.0.3&auto=format&fit=crop&w={px_w}&h={px_h}&q=80"
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as response:
                img_data = response.read()
                fetched = Image.open(BytesIO(img_data)).convert('RGBA')
                # Resize/Crop to exact dimensions
                fetched_aspect = fetched.width / fetched.height
                target_aspect = px_w / px_h
                if fetched_aspect > target_aspect:
                    new_w = int(fetched.height * target_aspect)
                    offset = (fetched.width - new_w) // 2
                    fetched = fetched.crop((offset, 0, offset + new_w, fetched.height))
                else:
                    new_h = int(fetched.width / target_aspect)
                    offset = (fetched.height - new_h) // 2
                    fetched = fetched.crop((0, offset, fetched.width, offset + new_h))
                base_img = fetched.resize((px_w, px_h), Image.Resampling.LANCZOS)
        except Exception as e:
            print(f"Warning: Image download failed, using fallback. ({e})")
        
        # Create gradient overlay
        overlay = Image.new('RGBA', (px_w, px_h))
        draw = ImageDraw.Draw(overlay)
        for y in range(px_h):
            # Gradient from Dark Gray (alpha 180) to Accent Color (alpha 150)
            ratio = y / px_h
            r = int(45 + (accent_color[0] - 45) * ratio)
            g = int(45 + (accent_color[1] - 45) * ratio)
            b = int(45 + (accent_color[2] - 45) * ratio)
            a = int(180 - (30 * ratio)) # Opacity decreases slightly going down
            draw.line([(0, y), (px_w, y)], fill=(r, g, b, a))
            
        final_img = Image.alpha_composite(base_img, overlay)
        temp_path = "temp_pillar1.png"
        final_img.save(temp_path, "PNG")
        return temp_path

    # --- PILLAR 1 (Image + Overlay) ---
    p1_x = left_area_w
    p1_img_path = generate_pillar1_image(pillar_w, slide_h_in, bg_keyword)
    slide.shapes.add_picture(p1_img_path, Inches(p1_x), Inches(0), Inches(pillar_w), Inches(slide_h_in))
    
    # --- PILLAR 2 (Light Gray) ---
    p2_x = p1_x + pillar_w
    shape2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(p2_x), Inches(0), Inches(pillar_w), Inches(slide_h_in))
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = color_pillar_light
    shape2.line.fill.background() # No outline
    
    # --- PILLAR 3 (Dark Gray) ---
    p3_x = p2_x + pillar_w
    shape3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(p3_x), Inches(0), Inches(pillar_w), Inches(slide_h_in))
    shape3.fill.solid()
    shape3.fill.fore_color.rgb = color_pillar_dark
    shape3.line.fill.background()

    # ==========================================
    # LAYER 3: PILLAR CONTENT (Icons, Text, Nums)
    # ==========================================
    
    def add_pillar_content(x_offset, text_color, number, icon_shape):
        # 1. Top Icon (Simulated with PPTX shape)
        icon_size = 0.5
        icon_x = x_offset + (pillar_w - icon_size) / 2
        icon = slide.shapes.add_shape(icon_shape, Inches(icon_x), Inches(1.0), Inches(icon_size), Inches(icon_size))
        icon.fill.background() # Transparent fill
        icon.line.color.rgb = text_color
        icon.line.width = Pt(2)
        
        # 2. Middle Text
        tb = slide.shapes.add_textbox(Inches(x_offset + 0.2), Inches(2.2), Inches(pillar_w - 0.4), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore."
        p.font.size = Pt(13)
        p.font.color.rgb = text_color
        p.font.name = "Arial"
        
        # 3. Bottom Number
        num_tb = slide.shapes.add_textbox(Inches(x_offset), Inches(6.5), Inches(pillar_w), Inches(0.5))
        num_tf = num_tb.text_frame
        num_p = num_tf.add_paragraph()
        num_p.alignment = PP_ALIGN.CENTER
        num_p.text = number
        num_p.font.size = Pt(14)
        num_p.font.color.rgb = text_color
        num_p.font.name = "Arial"

    # Add content to the 3 pillars
    add_pillar_content(p1_x, RGBColor(255, 255, 255), "0 1", MSO_SHAPE.HEART)       # Pillar 1: White text on image
    add_pillar_content(p2_x, color_dark_text, "0 2", MSO_SHAPE.SUN)                 # Pillar 2: Dark text on light
    add_pillar_content(p3_x, RGBColor(255, 255, 255), "0 3", MSO_SHAPE.LIGHTNING_BOLT) # Pillar 3: White text on dark

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(p1_img_path):
        os.remove(p1_img_path)
        
    return output_pptx_path

# Example execution:
# create_slide("split_panel_editorial.pptx")
