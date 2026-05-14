def create_slide(
    output_pptx_path: str,
    title_text: str = "THANK YOU",
    sub_text: str = "DO YOU HAVE ANY QUESTIONS?",
    bg_keyword: str = "sunset,landscape",
    mask_color: tuple = (230, 230, 230, 255), # Light grey
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Typography Cutout Mask effect.
    Uses PIL to generate a solid layer with a transparent text hole.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw, ImageFont

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Helper: Download Background Image ---
    bg_img_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1920x1080/?{bg_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(bg_img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception as e:
        print(f"Image download failed, generating fallback: {e}")
        fallback = Image.new('RGB', (1920, 1080), color=(50, 100, 150))
        fallback.save(bg_img_path)

    # --- Helper: Download Thick Font for Mask ---
    font_path = "Montserrat-Black.ttf"
    if not os.path.exists(font_path):
        try:
            # Download a heavy font to ensure the cutout effect works well
            font_url = "https://github.com/google/fonts/raw/main/ofl/montserrat/Montserrat-Black.ttf"
            urllib.request.urlretrieve(font_url, font_path)
        except Exception:
            pass # Will fallback to default in PIL if download fails

    # --- Create the PIL Mask (The Core Effect) ---
    mask_img_path = "temp_mask.png"
    width, height = 1920, 1080
    
    # Create the solid foreground image
    img = Image.new('RGBA', (width, height), color=mask_color)
    
    # Create an alpha mask (255 = opaque foreground, 0 = transparent text hole)
    alpha_mask = Image.new('L', (width, height), color=255)
    draw = ImageDraw.Draw(alpha_mask)
    
    try:
        font = ImageFont.truetype(font_path, 360)
    except IOError:
        font = ImageFont.load_default()

    # Calculate text position (Centered)
    title_text = title_text.upper()
    try:
        bbox = draw.textbbox((0, 0), title_text, font=font)
        text_w = bbox[2] - bbox[0]
        text_h = bbox[3] - bbox[1]
    except AttributeError:
        # Fallback for older PIL versions
        text_w, text_h = draw.textsize(title_text, font=font)
        
    x = (width - text_w) / 2
    y = (height - text_h) / 2 - 50 # Slightly above true center
    
    # Draw text in black (0) on the alpha mask. This creates the "hole"
    draw.text((x, y), title_text, fill=0, font=font)
    
    # Apply the alpha mask to the solid image
    img.putalpha(alpha_mask)
    img.save(mask_img_path)

    # --- Assemble the Slide ---
    
    # 1. Background Image (Bottom Layer)
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    
    # 2. Text Cutout Mask (Middle Layer)
    slide.shapes.add_picture(mask_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # 3. Accents and Secondary Text (Top Layer)
    
    # Intersecting Vertical Line
    line_x = prs.slide_width / 2 - Inches(3) # Offset to the left
    line_y = Inches(1.5)
    line_h = Inches(4.5)
    line_w = Inches(0.06)
    shape_line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        line_x, line_y, line_w, line_h
    )
    shape_line.fill.solid()
    shape_line.fill.fore_color.rgb = RGBColor(20, 20, 20)
    shape_line.line.color.rgb = RGBColor(255, 255, 255)
    shape_line.line.width = Pt(1.5)

    # Secondary Text
    sub_text_box = slide.shapes.add_textbox(
        Inches(1), Inches(5.2), Inches(11.333), Inches(1)
    )
    tf = sub_text_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = sub_text.upper()
    p.alignment = PP_ALIGN.CENTER
    
    # Format Subtext
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(60, 60, 60)
    
    # Note: python-pptx doesn't natively support letter spacing/tracking
    # We simulate it by adding spaces between characters for the subtext
    spaced_sub_text = "  ".join(list(sub_text.upper()))
    p.text = spaced_sub_text

    # --- Save and Cleanup ---
    prs.save(output_pptx_path)
    
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
    if os.path.exists(mask_img_path):
        os.remove(mask_img_path)

    return output_pptx_path
