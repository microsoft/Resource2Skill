def create_slide(
    output_pptx_path: str,
    title_text: str = "Process Flowchart",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Structured Process Flowchart visual effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Helper Functions ===

    def add_drop_shadow(shape):
        """Injects Office Open XML to add a soft drop shadow to a shape."""
        shadow_xml = """
        <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:outerShdw blurRad="40000" dist="38100" dir="2700000" algn="tl" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="25000"/>
                </a:srgbClr>
            </a:outerShdw>
        </a:effectLst>
        """
        shape.element.spPr.append(parse_xml(shadow_xml))

    def create_node(slide, shape_type, text, x, y, w, h, bg_color):
        """Creates a flowchart node with styling, text, and shadow."""
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        
        # Fill and Outline
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*bg_color)
        shape.line.color.rgb = RGBColor(*bg_color)
        
        # Text Formatting
        shape.text = text
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.bold = True
            paragraph.font.size = Pt(13)
            paragraph.font.name = "Arial"
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            
        add_drop_shadow(shape)
        return shape

    def create_connector(slide, start_shape, end_shape, start_idx, end_idx):
        """Creates an elbow connector between two shape connection sites."""
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Inches(0), Inches(0), Inches(1), Inches(1)
        )
        connector.begin_connect(start_shape, start_idx)
        connector.end_connect(end_shape, end_idx)
        
        # Styling the line
        connector.line.color.rgb = RGBColor(160, 160, 160)
        connector.line.width = Pt(2.5)
        connector.line.end_arrowhead = MSO_LINE.ARROWHEAD_TRIANGLE
        return connector

    # === Layout & Palette Definitions ===
    
    # Palette based on the tutorial
    C_START = (66, 133, 244)   # Blue
    C_DECIS = (234, 67, 53)    # Red/Pink
    C_PROCS = (52, 168, 83)    # Green
    C_OUTPT = (251, 188, 5)    # Yellow/Orange

    # Dimensions
    w, h = 1.7, 0.8
    y_center = 3.35 # (7.5 height / 2) - (0.8 / 2) roughly
    y_top = 1.6
    
    # X coordinates (Columns)
    col1 = 0.8  # Start
    col2 = 3.2  # Decision
    col3 = 5.8  # Processes
    col4 = 8.4  # Output
    col5 = 11.0 # End

    # Connection site mappings (Typical for standard PPT autoshapes)
    SITE_TOP = 0
    SITE_LEFT = 1
    SITE_BOTTOM = 2
    SITE_RIGHT = 3

    # === Layer 1: Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(50, 50, 50)

    # === Layer 2: Shape Generation ===
    node_start = create_node(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "START", col1, y_center, w, h, C_START)
    node_decision = create_node(slide, MSO_SHAPE.DIAMOND, "DECISION", col2, y_center, w, h+0.2, C_DECIS)
    node_proc1 = create_node(slide, MSO_SHAPE.RECTANGLE, "PROCESS 1", col3, y_top, w, h, C_PROCS)
    node_proc2 = create_node(slide, MSO_SHAPE.RECTANGLE, "PROCESS 2", col3, y_center+0.1, w, h, C_PROCS)
    # Parallelogram shape logic requires slightly wider w for text fit due to slant
    node_output = create_node(slide, MSO_SHAPE.PARALLELOGRAM, "OUTPUT", col4, y_center+0.1, w+0.2, h, C_OUTPT)
    node_end = create_node(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "END", col5, y_center+0.1, w, h, C_START)

    # === Layer 3: Connectors ===
    # Note: If arrows look slightly misrouted upon first opening in PPT, simply 
    # select them and click 'Reroute Connectors' - this is a standard behavior 
    # of the Office drawing engine calculating initial bounding boxes.
    
    # Start (Right) -> Decision (Left)
    create_connector(slide, node_start, node_decision, SITE_RIGHT, SITE_LEFT)
    
    # Decision (Top) -> Process 1 (Left)
    create_connector(slide, node_decision, node_proc1, SITE_TOP, SITE_LEFT)
    
    # Decision (Right) -> Process 2 (Left)
    create_connector(slide, node_decision, node_proc2, SITE_RIGHT, SITE_LEFT)
    
    # Process 1 (Right) -> Output (Top)
    create_connector(slide, node_proc1, node_output, SITE_RIGHT, SITE_TOP)
    
    # Process 2 (Right) -> Output (Left)
    create_connector(slide, node_proc2, node_output, SITE_RIGHT, SITE_LEFT)
    
    # Output (Right) -> End (Left)
    create_connector(slide, node_output, node_end, SITE_RIGHT, SITE_LEFT)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
