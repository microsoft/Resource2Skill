import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pandas as pd
import io

def create_slide(
    output_pptx_path: str,
    chart_data: dict = None,
    title_text: str = "Small Multiples – Panel Bar Chart – Percentage Values",
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with a Small Multiples Comparative Panel Bar Chart.

    This function generates a panel of horizontal bar charts using matplotlib, saves it
    as an image, and inserts it into a new PowerPoint slide. This is a highly
    effective way to compare multiple metrics across the same set of categories.

    Args:
        output_pptx_path: Path to save the generated PPTX file.
        chart_data: A dictionary representing the data for the charts.
                    Example:
                    {
                        "Category": ["Product 1", "Product 2", "Product 3", "Product 4", "Product 5"],
                        "Parameter A": [35, 57, 17, 26, 45],
                        "Parameter B": [65, 29, 42, 12, 32],
                        "Parameter C": [23, 37, 18, 20, 12]
                    }
        title_text: The main title for the slide.

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Add a title to the slide
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.8))
    text_frame = title_shape.text_frame
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 64, 128)

    # --- Use default data if none is provided ---
    if chart_data is None:
        chart_data = {
            "Category": ["Product 1", "Product 2", "Product 3", "Product 4", "Product 5"],
            "Parameter A": [35, 57, 17, 26, 45],
            "Parameter B": [65, 29, 42, 12, 32],
            "Parameter C": [23, 37, 18, 20, 12]
        }
    df = pd.DataFrame(chart_data).set_index("Category")

    # --- Chart Generation with Matplotlib ---
    # Define colors inspired by the video
    colors = {
        'Parameter A': '#2d3e50',  # Dark Slate Gray
        'Parameter B': '#3498db',  # Professional Blue
        'Parameter C': '#95a5a6'   # Muted Gray-Blue
    }

    # Reverse the DataFrame to plot from top to bottom
    df = df.iloc[::-1]

    num_parameters = len(df.columns)

    # Create a figure and a set of subplots
    fig, axes = plt.subplots(
        nrows=1,
        ncols=num_parameters,
        figsize=(12, 5),
        sharey=True # Share the Y-axis labels and ticks
    )
    fig.patch.set_facecolor('white')

    for i, (param_name, param_series) in enumerate(df.items()):
        ax = axes[i]
        
        # Plot the horizontal bars
        bars = ax.barh(param_series.index, param_series, color=colors.get(param_name, '#34495e'))
        
        # Set title for each subplot (panel)
        ax.set_title(param_name, fontsize=14, pad=15, color='#333333', weight='bold')
        
        # Remove all spines (borders) for a cleaner look
        ax.spines[['top', 'right', 'bottom', 'left']].set_visible(False)
        
        # Remove x-axis ticks and labels
        ax.xaxis.set_ticks_position('none')
        ax.set_xticks([])
        
        # Remove y-axis ticks but keep labels on the first chart
        ax.tick_params(axis='y', length=0)

        # Set a consistent x-axis limit (e.g., 0 to 100 for percentages)
        ax.set_xlim(0, max(100, df.max().max() * 1.1))

        # Add data labels inside the bars
        for bar in bars:
            width = bar.get_width()
            label_x_pos = width - 3 # Position text inside the bar
            ax.text(
                label_x_pos,
                bar.get_y() + bar.get_height() / 2,
                f'{int(width)}%',
                ha='right',
                va='center',
                color='white',
                fontsize=11,
                fontweight='bold'
            )

    # Style the shared y-axis labels on the first subplot
    axes[0].tick_params(axis='y', labelsize=12, pad=10, labelcolor='#333333')

    # Use a tight layout to automatically adjust spacing
    plt.tight_layout(pad=3.0)

    # Save the plot to a memory buffer
    img_stream = io.BytesIO()
    plt.savefig(img_stream, format='png', dpi=300, bbox_inches='tight', facecolor='white')
    img_stream.seek(0)
    plt.close(fig)

    # --- Add the generated image to the slide ---
    left = Inches(0.5)
    top = Inches(1.2)
    pic = slide.shapes.add_picture(img_stream, left, top, width=Inches(12.33))

    prs.save(output_pptx_path)
    return output_pptx_path
