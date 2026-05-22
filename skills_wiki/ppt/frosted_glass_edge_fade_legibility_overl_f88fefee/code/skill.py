import io
import requests
from PIL import Image, ImageFilter, ImageEnhance
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_slide(
    output_pptx_path: str,
    title_text: str = "KYOTO",
    body_text: str = "Kyoto, located in the Kansai region of Japan, is a city steeped in history and tradition. It served as the imperial capital for over a millennium, making it a cultural and historical treasure trove.\n\nKnown for its beautifully preserved temples, shrines, and traditional architecture, Kyoto offers a glimpse into Japan's rich past.",
    bg_keyword: str = "japan city street",
    overlay_width_ratio: float = 0.45,  # Overlay covers left 45% of slide
):
    """
    Create a PPTX file reproducing the 'Frosted Glass Legibility Overlay' effect.
    Downloads a background image, processes a blurred/darkened partial overlay using PIL,
    and inserts crisp typography on top.
    """
    # 1. Setup Presentation (16:9 Aspect Ratio)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Constants for pixel dimensions (assuming 150 DPI for image processing)
    # 13.333 inches * 150 = ~2000px, 7.5 inches * 150 = 1125px
    target_width = 2000
    target_height = 1125

    # 2. Fetch and Prepare Background Image
    print(f"Downloading background image for '{bg_keyword}'...")
    try:
        url = f"https://source.unsplash.com/random/{target_width}x{target_height}/?{bg_keyword}"
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        base_img = Image.open(io.BytesIO(response.content)).convert("RGBA")
    except Exception as e:
        print(f"Failed to download image ({e}). Using solid dark fallback.")
        base_img = Image.new("RGBA", (target_width, target_height), (30, 30, 40, 255))

    # Resize/Crop base image to exactly match 16:9 target dimensions
    # This prevents PPTX from stretching the image and ruining the overlay alignment
    bg_w, bg_h = base_img.size
    aspect_ratio = target_width / target_height
    if bg_w / bg_h > aspect_ratio:
        # Image is too wide
        new_w = int(bg_h * aspect_ratio)
        offset = (bg_w - new_w) // 2
        base_img = base_img.crop((offset, 0, offset + new_w, bg_h))
    else:
        # Image is too tall
        new_h = int(bg_w / aspect_ratio)
        offset = (bg_h - new_h) // 2
        base_img = base_img.crop((0, offset, bg_w, offset + new_h))
    
    base_img = base_img.resize((target_width, target_height), Image.Resampling.LANCZOS)

    # Save base image to bytes and insert into slide
    bg_bytes = io.BytesIO()
    base_img.convert("RGB").save(bg_bytes, format="JPEG", quality=90)
    bg_bytes.seek(0)
    slide.shapes.add_picture(bg_bytes, 0, 0, prs.slide_width, prs.slide_height)

    # 3. Create the Frosted Glass Overlay (Method 4 from tutorial)
    print("Generating frosted glass overlay...")
    overlay_px_width = int(target_width * overlay_width_ratio)
    
    # Crop the exact area from the base image
    overlay_img = base_img.crop((0, 0, overlay_px_width, target_height))
    
    # Apply Gaussian Blur
    overlay_img = overlay_img.filter(ImageFilter.GaussianBlur(radius=30))
    
    # Darken the image (Brightness < 1.0 makes it darker)
    enhancer = ImageEnhance.Brightness(overlay_img)
    overlay_img = enhancer.enhance(0.4)  # Reduce brightness by 60%

    # Save overlay to bytes and insert into slide perfectly aligned
    overlay_bytes = io.BytesIO()
    overlay_img.save(overlay_bytes, format="PNG")
    overlay_bytes.seek(0)
    
    overlay_width_inches = Inches(13.333 * overlay_width_ratio)
    slide.shapes.add_picture(
        overlay_bytes, 
        0, 0, 
        width=overlay_width_inches, 
        height=prs.slide_height
    )

    # 4. Add Typography
    print("Adding typography...")
    margin_left = Inches(0.8)
    margin_top = Inches(1.5)
    text_width = overlay_width_inches - Inches(1.2) # Keep text inside the frosted area

    # Title Box
    title_box = slide.shapes.add_textbox(margin_left, margin_top, text_width, Inches(1))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = "Arial Black"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Body Text Box
    body_box = slide.shapes.add_textbox(margin_left, margin_top + Inches(1.2), text_width, Inches(4))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(240, 240, 240)  # Slightly off-white for body
    # Adjust line spacing for elegance (approx 1.2 lines)
    p_body.line_spacing = 1.2

    # Save the presentation
    prs.save(output_pptx_path)
    print(f"Presentation saved successfully to: {output_pptx_path}")
    return output_pptx_path

if __name__ == "__main__":
    create_slide("frosted_glass_overlay.pptx")
