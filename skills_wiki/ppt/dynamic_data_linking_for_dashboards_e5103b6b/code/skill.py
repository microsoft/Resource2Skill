import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, BarChart, LineChart, Reference

def create_source_excel(excel_path="Business_Dashboard_Source.xlsx"):
    """Creates a sample Excel file with data and charts to be linked."""
    if os.path.exists(excel_path):
        # To ensure the demo works correctly, remove the old file
        # In a real-world scenario, you might just update it
        os.remove(excel_path)

    wb = Workbook()
    
    # --- Traffic Usage Sheet ---
    ws_traffic = wb.active
    ws_traffic.title = "Traffic Usage"
    ws_traffic.sheet_view.showGridLines = False
    traffic_data = [
        ('Description', 'Traffic Usage', 'Percentage'),
        ('Already Used', 40.56, 0.67),
        ('Remaining', 20.36, 0.33)
    ]
    for row in traffic_data:
        ws_traffic.append(row)
    
    chart_t = DoughnutChart()
    labels_t = Reference(ws_traffic, min_col=1, min_row=2, max_row=3)
    data_t = Reference(ws_traffic, min_col=2, min_row=1, max_row=3)
    chart_t.add_data(data_t, titles_from_data=True)
    chart_t.set_categories(labels_t)
    chart_t.title = "Traffic Usage"
    ws_traffic.add_chart(chart_t, "G5")

    # --- Payment Sheet ---
    ws_payment = wb.create_sheet("Payment")
    ws_payment.sheet_view.showGridLines = False
    payment_data = [
        ('Month', 'Amount (K)'), ('Jan', 25), ('Feb', 20), ('Mar', 42), ('Apr', 36),
        ('May', 38), ('Jun', 35), ('Jul', 33), ('Aug', 37), ('Sep', 35), ('Oct', 45), ('Nov', 35)
    ]
    for row in payment_data:
        ws_payment.append(row)
    chart_p = BarChart()
    data_p = Reference(ws_payment, min_col=2, min_row=1, max_row=12)
    cats_p = Reference(ws_payment, min_col=1, min_row=2, max_row=12)
    chart_p.add_data(data_p, titles_from_data=True)
    chart_p.set_categories(cats_p)
    chart_p.title = "Monthly Payments"
    chart_p.legend = None
    ws_payment.add_chart(chart_p, "F2")

    # --- Expense Trend ---
    ws_expense = wb.create_sheet("Expense Trend")
    ws_expense.sheet_view.showGridLines = False
    expense_data = [
        ('Month', 'Expense (K)'), ('Jan', 3.4), ('Feb', 3.9), ('Mar', 3.5), ('Apr', 3.8), ('May', 3.2),
        ('Jun', 3.4), ('Jul', 4.1), ('Aug', 2.9), ('Sep', 4.2), ('Oct', 4.5), ('Nov', 3.2), ('Dec', 4.0)
    ]
    for row in expense_data:
        ws_expense.append(row)
    chart_e = LineChart()
    chart_e.title = "Expense Trend"
    data_e = Reference(ws_expense, min_col=2, min_row=1, max_row=13)
    cats_e = Reference(ws_expense, min_col=1, min_row=2, max_row=13)
    chart_e.add_data(data_e, titles_from_data=True)
    chart_e.set_categories(cats_e)
    ws_expense.add_chart(chart_e, "F2")

    wb.save(excel_path)
    print(f"Source Excel file created at: {excel_path}")

def create_slide(
    output_pptx_path: str = "Linked_Excel_Dashboard.pptx",
    title_text: str = "Dashboard",
    ceo_name: str = "JOHN STANLEY",
    **kwargs,
) -> str:
    """
    Creates a PPTX file with a dashboard layout, linking to widgets from an external Excel file.
    This function reproduces the "Dynamic Excel Dashboard Linking" pattern by injecting OLE links.
    """
    from PIL import Image

    # 1. Setup Files and Presentation
    source_excel_name = "Business_Dashboard_Source.xlsx"
    source_excel_path = os.path.abspath(source_excel_name)
    create_source_excel(source_excel_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 2. Create Dashboard Background and Layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(18, 18, 18)

    # Main Dashboard Title
    dash_title = slide.shapes.add_textbox(Inches(2.7), Inches(0.2), Inches(10), Inches(0.8))
    p = dash_title.text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri Light'
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(255, 255, 255)

    # 3. Embed Linked Excel Objects using lxml
    placeholder_img_path = "placeholder_ole.png"
    Image.new('RGBA', (1, 1), (0, 0, 0, 0)).save(placeholder_img_path)

    def add_linked_ole(slide, x_emu, y_emu, cx_emu, cy_emu, excel_path, sheet_name, cell_range):
        slide_part = slide.part
        r_id = slide_part.relate_to(
            excel_path,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
            is_external=True
        )
        img_r_id = slide_part.relate_to(
            placeholder_img_path,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
        )
        
        id1 = slide_part.next_id
        id2 = id1 + 1
        name = f"Object {id1}"
        
        # XML structure for a linked OLE object
        xml_str = f"""
        <p:graphicFrame {nsdecls('p', 'a')}>
          <p:nvGraphicFramePr>
            <p:cNvPr id="{id1}" name="{name}"/>
            <p:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr>
            <p:nvPr>
              <p:extLst>
                <p:ext uri="{{63837DB3-8F32-4F43-9BE5-756155938360}}">
                  <p14:creationId xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" id="{{{os.urandom(16).hex().upper()}}}"/>
                </p:ext>
              </p:extLst>
            </p:nvPr>
          </p:nvGraphicFramePr>
          <p:xfrm>
            <a:off x="{x_emu}" y="{y_emu}"/>
            <a:ext cx="{cx_emu}" cy="{cy_emu}"/>
          </p:xfrm>
          <a:graphic>
            <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/ole">
              <p:oleObj spid="{id2}" name="{sheet_name}" r:id="{r_id}" imgW="{cx_emu}" imgH="{cy_emu}" progId="Excel.Sheet.12">
                <p:embed>
                   <p:link updateAutomatic="true">{sheet_name}!{cell_range}</p:link>
                </p:embed>
                 <p:pic>
                    <p:nvPicPr>
                        <p:cNvPr id="0" name=""/>
                        <p:cNvPicPr><a:picLocks noChangeAspect="true" noChangeArrowheads="true"/></p:cNvPicPr>
                        <p:nvPr><a:hlinkClick r:id="" action="ppaction://hlinksldjump"/></p:nvPr>
                    </p:nvPicPr>
                    <p:blipFill>
                        <a:blip r:embed="{img_r_id}"/>
                        <a:stretch><a:fillRect/></a:stretch>
                    </p:blipFill>
                    <p:spPr bwMode="auto">
                        <a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>
                        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                        <a:noFill/>
                        <a:ln><a:noFill/></a:ln>
                    </p:spPr>
                </p:pic>
              </p:oleObj>
            </a:graphicData>
          </a:graphic>
        </p:graphicFrame>"""
        
        tree = slide.shapes._spTree
        graphic_frame = etree.fromstring(xml_str)
        tree.append(graphic_frame)

    # Define widget positions and their corresponding Excel ranges
    widgets_to_link = [
        {'pos': (Inches(2.7), Inches(1.2), Inches(3.5), Inches(2.5)), 'sheet': 'Traffic Usage', 'range': 'A1:K17'},
        {'pos': (Inches(6.4), Inches(1.2), Inches(3.2), Inches(2.5)), 'sheet': 'Payment', 'range': 'A1:K18'},
        {'pos': (Inches(9.8), Inches(1.2), Inches(3.2), Inches(2.5)), 'sheet': 'Expense Trend', 'range': 'A1:K18'}, # Placeholder
        {'pos': (Inches(2.7), Inches(3.9), Inches(6.9), Inches(3.2)), 'sheet': 'Expense Trend', 'range': 'A1:N20'},
        {'pos': (Inches(9.8), Inches(3.9), Inches(3.2), Inches(3.2)), 'sheet': 'Payment', 'range': 'A1:L20'} # Placeholder
    ]

    for widget in widgets_to_link:
        l, t, w, h = widget['pos']
        add_linked_ole(slide, Emu(l), Emu(t), Emu(w), Emu(h), source_excel_path, widget['sheet'], widget['range'])

    # 4. Save and Cleanup
    prs.save(output_pptx_path)
    os.remove(placeholder_img_path)
    print(f"Presentation saved to: {output_pptx_path}")
    return output_pptx_path

