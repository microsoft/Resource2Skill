import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import OxmlElement
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda",
    date_text: str = "12-Sep-202X",
    intro_text: str = "We provide the best resources to get free feedback on your presentation. Based on the feedback, you can fix issues and create a wonderful deck.",
    agenda_items: list = None,
    accent_color: tuple = (68, 114, 196)
) -> str:
    """
    Create a PPTX file reproducing the Split-Panel Geometric Agenda visual effect.
    """
    if agenda_items is None:
        # Default mock items if none provided
        agenda_items = [
            {"title": "Agenda / Topic Name", "desc": "This is the best section to detail out your presentation topic and context."},
            {"title": "Agenda / Topic Name", "desc": "This is the best section to detail out your presentation topic and context."},
            {"title": "Agenda / Topic Name", "desc": "This is the best section to detail out your presentation topic and context."},
            {"title": "Agenda / Topic Name", "desc": "This is the best section to detail out your presentation topic and context."},
            {"title": "Agenda / Topic Name", "desc": "This is the best section to detail out your presentation topic and context."}
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # ==========================================
    # Layer 1: Generate & Apply Background
    # ==========================================
    bg_img_path = "temp_agenda_bg.png"
    bg_width, bg_height = 1920, 1080
    bg_img = Image.new("RGB", (bg_width, bg_height))
    draw = ImageDraw.Draw(bg_img)
    # White to light-gray linear gradient (top to bottom)
    for y in range(bg_height):
        c = int(255 - (20 * (y / bg_height))) # 255 down to 235
        draw.line([(0, y), (bg_width, y)], fill=(c, c, c))
    bg_img.save(bg_img_path)
    
    # Insert as background picture
    slide.shapes.add_picture(bg_img_path, 0, 0, width=Inches(13.333), height=Inches(7.5))

    # ==========================================
    # Layer 2: Left Curved Geometric Panel
    # ==========================================
    # We use an oversized circle shifted to the left to create the semi-circle curve
    circle_radius = 6.0
    circle_diameter = circle_radius * 2
    left_x = Inches(-5.5) # Shifted heavily left
    top_y = Inches((7.5 - circle_diameter) / 2) # Centered vertically
    
    left_panel = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left_x, top_y, Inches(circle_diameter), Inches(circle_diameter)
    )
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = RGBColor(*accent_color)
    left_panel.line.fill.background() # Remove outline

    # ==========================================
    # Layer 3: Left Panel Typography
    # ==========================================
    left_text_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(3.5), Inches(3.0))
    tf = left_text_box.text_frame
    tf.clear()
    
    # Main Title
    p_title = tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)
    
    # Date Subtitle
    p_date = tf.add_paragraph()
    p_date.text = f"\n{date_text}\n"
    p_date.font.size = Pt(16)
    p_date.font.color.rgb = RGBColor(255, 255, 255)
    
    # Intro Paragraph
    p_intro = tf.add_paragraph()
    p_intro.text = intro_text
    p_intro.font.size = Pt(13)
    p_intro.font.color.rgb = RGBColor(240, 240, 240)
    p_intro.font.name = "Calibri"

    # ==========================================
    # Layer 4: Right Agenda List Items
    # ==========================================
    def apply_shadow(shape):
        """Helper to inject an outer shadow effect via lxml."""
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', '50800')  # ~4pt blur
        outerShdw.set('dist', '38100')     # ~3pt distance
        outerShdw.set('dir', '2700000')    # 45 degrees
        outerShdw.set('algn', 'tl')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')       # Black shadow
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '30000')          # 30% opacity
        
        srgbClr.append(alpha)
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    start_y = 1.0
    spacing_y = 1.2
    
    for idx, item in enumerate(agenda_items):
        current_y = start_y + (idx * spacing_y)
        
        # 1. Indicator Circle
        bullet_shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(6.8), Inches(current_y + 0.1), Inches(0.55), Inches(0.55)
        )
        bullet_shape.fill.solid()
        bullet_shape.fill.fore_color.rgb = RGBColor(*accent_color)
        bullet_shape.line.fill.background()
        apply_shadow(bullet_shape) # Add 3D depth to the flat UI
        
        # 2. Text Content Box
        item_tb = slide.shapes.add_textbox(Inches(7.6), Inches(current_y), Inches(5.0), Inches(1.0))
        item_tf = item_tb.text_frame
        item_tf.clear()
        
        # Item Title
        p_item_title = item_tf.paragraphs[0]
        p_item_title.text = item["title"]
        p_item_title.font.size = Pt(18)
        p_item_title.font.bold = True
        p_item_title.font.color.rgb = RGBColor(60, 60, 60)
        
        # Item Description
        p_item_desc = item_tf.add_paragraph()
        p_item_desc.text = item["desc"]
        p_item_desc.font.size = Pt(12)
        p_item_desc.font.color.rgb = RGBColor(120, 120, 120)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
