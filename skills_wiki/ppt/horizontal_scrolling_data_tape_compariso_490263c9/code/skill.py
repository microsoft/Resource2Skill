def create_slide(
    output_pptx_path: str,
    title_text: str = "Comparison Video\nTemplate",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Horizontal Scrolling Data Tape effect.
    Generates an overflowing tape of data cards and a static front-overlay mask.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    import urllib.request
    import io
    from PIL import Image

    # Helper function to fetch or generate placeholder images
    def get_image_stream(url, fallback_color=(150, 150, 150)):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=3) as response:
                return io.BytesIO(response.read())
        except Exception:
            img = Image.new('RGB', (200, 200), color=fallback_color)
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format='JPEG')
            img_byte_arr.seek(0)
            return img_byte_arr

    # Dummy data for the comparison tape
    data_items = [
        {"val": "10B", "name": "Item One"},
        {"val": "12B", "name": "Item Two"},
        {"val": "15B", "name": "Item Three"},
        {"val": "22B", "name": "Item Four"},
        {"val": "35B", "name": "Item Five"},
        {"val": "41B", "name": "Item Six"},
        {"val": "50B", "name": "Item Seven"},
        {"val": "75B", "name": "Item Eight"},
        {"val": "100B", "name": "Item Nine"},
        {"val": "150B", "name": "Item Ten"},
    ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Background Tracks ===
    # Dark top section
    bg_top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(3.5))
    bg_top.fill.solid()
    bg_top.fill.fore_color.rgb = RGBColor(40, 44, 52)
    bg_top.line.color.rgb = RGBColor(40, 44, 52)

    # Grey middle section (names)
    bg_mid = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.5), Inches(13.333), Inches(1.0))
    bg_mid.fill.solid()
    bg_mid.fill.fore_color.rgb = RGBColor(150, 155, 160)
    bg_mid.line.color.rgb = RGBColor(150, 155, 160)

    # Darker bottom section (images)
    bg_bot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(4.5), Inches(13.333), Inches(3.0))
    bg_bot.fill.solid()
    bg_bot.fill.fore_color.rgb = RGBColor(70, 75, 80)
    bg_bot.line.color.rgb = RGBColor(70, 75, 80)

    # === Layer 2: The Data Tape (Overflows to the right) ===
    item_width = 2.0
    item_gap = 0.2
    start_x = 0.5

    for i, item in enumerate(data_items):
        current_x = start_x + (i * (item_width + item_gap))

        # 1. Ribbon Rectangle (holds text)
        rect_y = 0.5
        rect_h = 1.6
        ribbon_color = RGBColor(220, 50, 32)
        
        ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(current_x), Inches(rect_y), Inches(item_width), Inches(rect_h))
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = ribbon_color
        ribbon.line.fill.background()
        
        tf = ribbon.text_frame
        tf.text = f"{item['val']}\nUSD"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        if len(tf.paragraphs) > 1:
            tf.paragraphs[1].font.size = Pt(14)
            tf.paragraphs[1].alignment = PP_ALIGN.CENTER

        # 2. Ribbon Triangle (pointing down, attached to bottom of rectangle)
        tri_y = rect_y + rect_h
        tri_h = 0.5
        triangle = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(current_x), Inches(tri_y), Inches(item_width), Inches(tri_h))
        triangle.fill.solid()
        triangle.fill.fore_color.rgb = ribbon_color
        triangle.line.fill.background()
        triangle.rotation = 180 # Point downwards

        # 3. Label Box
        lbl_y = 3.5
        lbl_h = 1.0
        label = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(current_x), Inches(lbl_y), Inches(item_width), Inches(lbl_h))
        label.fill.background() # transparent fill to let grey background show
        label.line.fill.background()
        
        lbl_tf = label.text_frame
        lbl_tf.text = item['name']
        lbl_tf.paragraphs[0].font.size = Pt(18)
        lbl_tf.paragraphs[0].font.bold = True
        lbl_tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        lbl_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # 4. Image
        img_y = 4.7
        img_h = 2.0
        img_stream = get_image_stream(f"https://picsum.photos/seed/{i+10}/200/200")
        slide.shapes.add_picture(img_stream, Inches(current_x), Inches(img_y), Inches(item_width), Inches(img_h))

    # === Layer 3: Static Front Overlay Panel ===
    # This acts as the viewport mask on the right side
    panel_w = 3.5
    panel_x = 13.333 - panel_w
    
    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(panel_x), Inches(0), Inches(panel_w), Inches(7.5))
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(30, 35, 40)
    overlay.line.color.rgb = RGBColor(0, 0, 0)

    # Overlay Title Text
    title_box = slide.shapes.add_textbox(Inches(panel_x + 0.2), Inches(1.0), Inches(panel_w - 0.4), Inches(2.0))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    p = title_tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Decorative Timer/Logo circle on overlay
    circle = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(panel_x + 1.25), Inches(4.5), Inches(1.0), Inches(1.0))
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
    circle.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
