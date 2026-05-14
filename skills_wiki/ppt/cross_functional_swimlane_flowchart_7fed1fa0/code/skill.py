def create_slide(
    output_pptx_path: str,
    title_text: str = "Expense Reimbursement Process",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Cross-Functional Swimlane Flowchart effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from lxml import etree
    from pptx.oxml.ns import qn

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- XML Injection Helpers ---
    def set_arrowhead(connector):
        """Inject XML to add a triangle arrowhead to a line."""
        spPr = connector.element.spPr
        ln = spPr.find(qn('a:ln'))
        if ln is None:
            ln = etree.SubElement(spPr, qn('a:ln'))
        headEnd = ln.find(qn('a:headEnd'))
        if headEnd is None:
            headEnd = etree.SubElement(ln, qn('a:headEnd'))
        headEnd.set('type', 'triangle')
        headEnd.set('w', 'med')
        headEnd.set('len', 'med')

    def add_shadow(shape):
        """Inject XML to add a subtle drop shadow to nodes."""
        spPr = shape.element.spPr
        effectLst = spPr.find(qn('a:effectLst'))
        if effectLst is None:
            effectLst = etree.SubElement(spPr, qn('a:effectLst'))
        outerShdw = effectLst.find(qn('a:outerShdw'))
        if outerShdw is None:
            outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
        outerShdw.set('blurRad', '38100') # Blur radius
        outerShdw.set('dist', '38100')    # Distance
        outerShdw.set('dir', '2700000')   # 45 degrees
        outerShdw.set('algn', 'tl')
        srgbClr = etree.SubElement(outerShdw, qn('a:srgbClr'))
        srgbClr.set('val', '000000')
        alpha = etree.SubElement(srgbClr, qn('a:alpha'))
        alpha.set('val', '15000')         # 15% opacity

    def style_node(shape, text, font_size=11):
        """Apply uniform styling to flowchart nodes."""
        shape.text = text
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(60, 60, 60)
                run.font.bold = True
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(120, 120, 120)
        shape.line.width = Pt(1.5)
        add_shadow(shape)

    # --- Title Setup ---
    tb = slide.shapes.add_textbox(Inches(1.86), Inches(0.4), Inches(9.6), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = title_text
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(24)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(50, 50, 50)

    # --- Swimlane Setup ---
    # Total width = 9.6 inches (3.2 * 3). Centered horizontally on a 13.333 slide.
    lanes = [
        {"title": "Applicant", "x": 1.86, "w": 3.2, "bg": RGBColor(242, 247, 252), "hdr": RGBColor(190, 215, 235)},
        {"title": "Manager", "x": 5.06, "w": 3.2, "bg": RGBColor(254, 248, 242), "hdr": RGBColor(245, 215, 185)},
        {"title": "Finance", "x": 8.26, "w": 3.2, "bg": RGBColor(248, 243, 252), "hdr": RGBColor(225, 205, 240)}
    ]

    for lane in lanes:
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(lane["x"]), Inches(1.5), Inches(lane["w"]), Inches(5.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = lane["bg"]
        bg.line.fill.background()
        
        # Header
        hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(lane["x"]), Inches(1.5), Inches(lane["w"]), Inches(0.5))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = lane["hdr"]
        hdr.line.color.rgb = RGBColor(255, 255, 255)
        hdr.line.width = Pt(1.0)
        
        hdr.text = lane["title"]
        p = hdr.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(80, 80, 80)

    # Centers for nodes in each lane
    c1_x, c2_x, c3_x = 3.46, 6.66, 9.86

    # --- Add Flowchart Nodes ---
    n1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c1_x - 0.6), Inches(2.3), Inches(1.2), Inches(0.4))
    style_node(n1, "Start")

    n2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c1_x - 0.8), Inches(3.2), Inches(1.6), Inches(0.6))
    style_node(n2, "Fill\nForm")

    n3 = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(c2_x - 1.0), Inches(4.2), Inches(2.0), Inches(1.0))
    style_node(n3, "Manager\nReview")

    n4 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c3_x - 0.8), Inches(5.6), Inches(1.6), Inches(0.6))
    style_node(n4, "Process\nPayment")

    n5 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c3_x - 0.6), Inches(6.6), Inches(1.2), Inches(0.4))
    style_node(n5, "End")

    n6 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c1_x - 0.6), Inches(4.5), Inches(1.2), Inches(0.4))
    style_node(n6, "End")

    # --- Add Connectors ---
    def draw_straight_arrow(x1, y1, x2, y2):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        conn.line.color.rgb = RGBColor(120, 120, 120)
        conn.line.width = Pt(1.5)
        set_arrowhead(conn)
        return conn

    def draw_elbow_down_right(x1, y1, x2, y2):
        """Draws a predictable down -> right -> down elbow path to bypass pptx auto-routing bugs."""
        mid_y = y1 + (y2 - y1) / 2
        l1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x1), Inches(mid_y))
        l2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(mid_y), Inches(x2), Inches(mid_y))
        l3 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x2), Inches(mid_y), Inches(x2), Inches(y2))
        for l in [l1, l2, l3]:
            l.line.color.rgb = RGBColor(120, 120, 120)
            l.line.width = Pt(1.5)
        set_arrowhead(l3)

    # Route 1: Start -> Fill Form
    draw_straight_arrow(c1_x, 2.7, c1_x, 3.2)
    # Route 2: Fill Form -> Manager Review (Elbow)
    draw_elbow_down_right(c1_x, 3.8, c2_x, 4.2)
    # Route 3: Manager Review -> Process Payment (Elbow)
    draw_elbow_down_right(c2_x, 5.2, c3_x, 5.6)
    # Route 4: Manager Review -> End (Reject)
    draw_straight_arrow(c2_x - 1.0, 4.7, c1_x + 0.6, 4.7)
    # Route 5: Process Payment -> End (Success)
    draw_straight_arrow(c3_x, 6.2, c3_x, 6.6)

    # --- Add Branch Labels ---
    def add_branch_label(text, x, y):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(0.8), Inches(0.3))
        tb.text_frame.margin_left = 0
        tb.text_frame.margin_top = 0
        tb.text_frame.margin_right = 0
        tb.text_frame.margin_bottom = 0
        tb.text_frame.text = text
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(120, 120, 120)
        p.runs[0].font.bold = True

    add_branch_label("Pass", 6.8, 5.25)
    add_branch_label("Reject", 4.6, 4.45)

    prs.save(output_pptx_path)
    return output_pptx_path
