import os
import urllib.request
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "MAGICAL\nSCOTLAND",
    subtitle_text: str = "LOREM IPSUM DOLOR",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do\neiusmod tempor incididunt ut labore et dolore magna aliqua.",
    image_url: str = "https://images.unsplash.com/photo-1506377247377-2a5b3b417ebb?q=80&w=1920&auto=format&fit=crop",
    accent_color: tuple = (255, 215, 0),  # Yellow
) -> str:
    """
    Create a PPTX file reproducing the 'Diagonal Matte Reveal' design pattern.
    """
    # ---------------------------------------------------------
    # 1. SETUP PRESENTATION
    # ---------------------------------------------------------
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    width, height = 1920, 1080  # HD canvas for PIL rendering

    # ---------------------------------------------------------
    # 2. GENERATE BACKGROUND IMAGE (PIL Compositing)
    # ---------------------------------------------------------
    # Fetch or generate base image
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            bg_img = Image.open(BytesIO(response.read())).convert("RGBA")
            
        # Scale and crop to exactly 1920x1080
        img_ratio = bg_img.width / bg_img.height
        target_ratio = width / height
        if img_ratio > target_ratio:
            new_h = height
            new_w = int(new_h * img_ratio)
        else:
            new_w = width
            new_h = int(new_w / img_ratio)
        
        bg_img = bg_img.resize((new_w, new_h), Image.Resampling.LANCZOS)
        left = (new_w - width) / 2
        top = (new_h - height) / 2
        bg_img = bg_img.crop((left, top, left + width, top + height))
    except Exception as e:
        print(f"Failed to download image, using fallback. Error: {e}")
        bg_img = Image.new("RGBA", (width, height), (30, 80, 60, 255))

    # Create the dark matte overlay (Vertical Gradient)
    overlay = Image.new("RGBA", (width, height))
    draw_overlay = ImageDraw.Draw(overlay)
    for y in range(height):
        # Dark slate `#282A36` to `#14161C`
        r = int(40 - (40 - 20) * (y / height))
        g = int(42 - (42 - 22) * (y / height))
        b = int(54 - (54 - 28) * (y / height))
        draw_overlay.line([(0, y), (width, y)], fill=(r, g, b, 255))

    # Create the alpha mask for the matte (White = Matte visible, Black = Hole/Image visible)
    mask = Image.new("L", (width, height), 255)

    def punch_hole(center, size, radius, angle):
        """Helper to punch rotated rounded rectangles into the mask."""
        cx, cy = center
        w, h = size
        # Create a temp canvas, draw the white shape
        temp = Image.new("L", (width, height), 0)
        draw = ImageDraw.Draw(temp)
        draw.rounded_rectangle([cx - w/2, cy - h/2, cx + w/2, cy + h/2], radius=radius, fill=255)
        # Rotate in place
        temp = temp.rotate(angle, center=center, resample=Image.Resampling.BICUBIC)
        # Paste 0 (black) into the mask wherever temp is white
        mask.paste(0, (0, 0), temp)

    def punch_circle(center, radius):
        """Helper to punch circles into the mask."""
        cx, cy = center
        temp = Image.new("L", (width, height), 0)
        draw = ImageDraw.Draw(temp)
        draw.ellipse([cx - radius, cy - radius, cx + radius, cy + radius], fill=255)
        mask.paste(0, (0, 0), temp)

    # Define the composition of diagonal cutouts (-45 degrees)
    punch_hole((1300, 200), (300, 1600), radius=150, angle=-45)  # Main slash
    punch_hole((1750, 700), (250, 1600), radius=125, angle=-45)  # Secondary right slash
    punch_hole((900, 950), (160, 900), radius=80, angle=-45)     # Small bottom slash
    punch_hole((850, -150), (220, 800), radius=110, angle=-45)   # Small top slash
    
    # Add floating circles to break up the parallel lines
    punch_circle((1600, 150), radius=110)
    punch_circle((1180, 850), radius=90)

    # Create the inner shadow effect
    # By blurring the mask (where holes are black and matte is white)
    # The edges of the holes fade from 255 (white) to 0 (black) inside the hole
    shadow_mask = mask.filter(ImageFilter.GaussianBlur(25))
    shadow_layer = Image.new("RGBA", (width, height), (0, 0, 0, 220)) # Deep black shadow
    
    # Apply shadow to the background image using the blurred mask
    # This darkens the edges just inside the hole
    bg_img.paste(shadow_layer, (0, 0), shadow_mask)

    # Finally, composite the dark matte over the background using the crisp mask
    bg_img.paste(overlay, (0, 0), mask)

    # Save to disk
    bg_path = "temp_slide_bg.png"
    bg_img.save(bg_path, format="PNG")

    # Set as slide background
    slide_bg = slide.background
    fill = slide_bg.fill
    fill.solid()
    # To place it seamlessly, we add a picture shape covering the entire slide
    pic = slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    
    # Optional watermark (like the globe/plane in the tutorial)
    # Simulated by placing text with a very dark color that matches the gradient closely
    watermark_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(4))
    tf_w = watermark_box.text_frame
    p_w = tf_w.paragraphs[0]
    p_w.text = "✈"
    p_w.font.size = Pt(250)
    p_w.font.color.rgb = RGBColor(60, 62, 74) # Slightly lighter than BG to simulate 75% opacity

    # ---------------------------------------------------------
    # 3. ADD TYPOGRAPHY & CONTENT
    # ---------------------------------------------------------
    # Title Text
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(6.5), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial Black"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Subtitle Text (Accent Color)
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(6.0), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(*accent_color)
    
    # Body Text
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.7), Inches(5.5), Inches(1.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = RGBColor(180, 182, 190)

    # Save presentation
    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
