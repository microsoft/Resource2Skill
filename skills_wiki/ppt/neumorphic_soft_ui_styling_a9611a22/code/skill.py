def create_slide(
    output_pptx_path: str,
    title_text: str = "NEUMORPHIC PRESENTATION",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Neumorphic Soft UI design style.
    Generates dynamic dual-shadow assets using PIL and overlays text via python-pptx.
    """
    import os
    import tempfile
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFilter

    # --- 1. Colors & Settings ---
    BG_COLOR = (238, 240, 245)
    DARK_SHADOW = (200, 205, 215, 255)
    LIGHT_SHADOW = (255, 255, 255, 255)
    ACCENT_COLOR = RGBColor(133, 193, 184)
    TEXT_COLOR = RGBColor(122, 133, 153)
    
    DPI = 300 # High resolution for PIL generation

    # --- 2. Helper: Generate Neumorphic Shape ---
    def create_neumorphic_asset(width_in, height_in, radius_in, offset_px=15, blur_px=25):
        """Generates a soft UI card as a transparent PNG."""
        w_px = int(width_in * DPI)
        h_px = int(height_in * DPI)
        r_px = int(radius_in * DPI)
        
        # Calculate padding needed so blur doesn't cut off
        padding = blur_px * 2 + offset_px
        img_w = w_px + padding * 2
        img_h = h_px + padding * 2
        
        # Base transparent image
        base = Image.new('RGBA', (img_w, img_h), (0, 0, 0, 0))
        shape_bounds = [padding, padding, padding + w_px, padding + h_px]
        
        # Light shadow (Top Left)
        light_layer = Image.new('RGBA', (img_w, img_h), (0, 0, 0, 0))
        light_draw = ImageDraw.Draw(light_layer)
        light_bounds = [
            padding - offset_px, padding - offset_px, 
            padding + w_px - offset_px, padding + h_px - offset_px
        ]
        light_draw.rounded_rectangle(light_bounds, radius=r_px, fill=LIGHT_SHADOW)
        light_layer = light_layer.filter(ImageFilter.GaussianBlur(blur_px))
        
        # Dark shadow (Bottom Right)
        dark_layer = Image.new('RGBA', (img_w, img_h), (0, 0, 0, 0))
        dark_draw = ImageDraw.Draw(dark_layer)
        dark_bounds = [
            padding + offset_px, padding + offset_px, 
            padding + w_px + offset_px, padding + h_px + offset_px
        ]
        dark_draw.rounded_rectangle(dark_bounds, radius=r_px, fill=DARK_SHADOW)
        dark_layer = dark_layer.filter(ImageFilter.GaussianBlur(blur_px))
        
        # Actual Shape (Center)
        shape_layer = Image.new('RGBA', (img_w, img_h), (0, 0, 0, 0))
        shape_draw = ImageDraw.Draw(shape_layer)
        # Using RGBA format of BG color
        shape_fill = BG_COLOR + (255,)
        shape_draw.rounded_rectangle(shape_bounds, radius=r_px, fill=shape_fill)
        
        # Composite layers
        final = Image.alpha_composite(base, light_layer)
        final = Image.alpha_composite(final, dark_layer)
        final = Image.alpha_composite(final, shape_layer)
        
        # Save to temp file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        final.save(temp_file.name)
        
        # Return path and the alignment offset in inches
        return temp_file.name, (padding / DPI)

    # --- 3. Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Set Slide Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(BG_COLOR[0], BG_COLOR[1], BG_COLOR[2])

    # --- 4. Create & Insert Visuals ---
    temp_files = [] # Track to clean up later

    # Top Pill (Title)
    pill_w, pill_h = 4.0, 0.7
    pill_path, pill_pad = create_neumorphic_asset(pill_w, pill_h, radius_in=0.35, offset_px=8, blur_px=15)
    temp_files.append(pill_path)
    
    # Position Top Pill (Centered, top)
    pill_x = (13.333 - pill_w) / 2
    pill_y = 0.8
    # Insert picture accounting for the blur padding
    slide.shapes.add_picture(
        pill_path, 
        Inches(pill_x - pill_pad), Inches(pill_y - pill_pad), 
        width=Inches(pill_w + pill_pad*2), height=Inches(pill_h + pill_pad*2)
    )

    # Add Text to Top Pill
    tx_box = slide.shapes.add_textbox(Inches(pill_x), Inches(pill_y), Inches(pill_w), Inches(pill_h))
    tf = tx_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].alignment = 2 # Center
    tf.paragraphs[0].font.name = "Arial"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ACCENT_COLOR

    # Create 3 Neumorphic Cards
    card_w, card_h = 2.5, 2.5
    card_path, card_pad = create_neumorphic_asset(card_w, card_h, radius_in=0.4, offset_px=12, blur_px=22)
    temp_files.append(card_path)

    gap = 1.2
    total_w = (card_w * 3) + (gap * 2)
    start_x = (13.333 - total_w) / 2
    card_y = 2.8

    for i in range(3):
        cur_x = start_x + (i * (card_w + gap))
        
        # Insert Card Image
        slide.shapes.add_picture(
            card_path, 
            Inches(cur_x - card_pad), Inches(card_y - card_pad), 
            width=Inches(card_w + card_pad*2), height=Inches(card_h + card_pad*2)
        )
        
        # Add Card Number (Accent)
        num_box = slide.shapes.add_textbox(Inches(cur_x), Inches(card_y + 0.6), Inches(card_w), Inches(0.5))
        tf_num = num_box.text_frame
        tf_num.text = f"0{i+1}"
        tf_num.paragraphs[0].alignment = 2
        tf_num.paragraphs[0].font.name = "Arial"
        tf_num.paragraphs[0].font.size = Pt(36)
        tf_num.paragraphs[0].font.bold = True
        tf_num.paragraphs[0].font.color.rgb = ACCENT_COLOR

        # Add Card Subtext
        desc_box = slide.shapes.add_textbox(Inches(cur_x + 0.2), Inches(card_y + 1.4), Inches(card_w - 0.4), Inches(0.8))
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        p = tf_desc.paragraphs[0]
        p.text = "SKILL LEVEL\nDescription of process."
        p.alignment = 2
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_COLOR

    # --- 5. Save and Cleanup ---
    prs.save(output_pptx_path)
    
    # Cleanup temporary files
    for path in temp_files:
        if os.path.exists(path):
            os.remove(path)

    return output_pptx_path
