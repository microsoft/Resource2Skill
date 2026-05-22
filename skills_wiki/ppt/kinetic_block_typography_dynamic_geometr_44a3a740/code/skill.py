import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

def create_slide(
    output_pptx_path: str,
    headline_top: str = "BEST",
    headline_mid: str = "WAY TO",
    block_word: str = "LEARN",
    side_text: str = "THE",
    accent_word: str = "SO...",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Kinetic Block Typography & Dynamic Overlay style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 10, 10)

    # === Layer 1: Dashed Kinetic Circle Accent (Top Right) ===
    # Draw a large circle extending slightly off-canvas
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6), Inches(6))
    circle.fill.background()  # Transparent/Matches background
    circle.line.color.rgb = RGBColor(255, 255, 255)
    circle.line.width = Pt(6)
    
    # Inject large dashed line style via lxml
    ln = circle.line._linePr
    prstDash = OxmlElement('a:prstDash')
    prstDash.set('val', 'lgDash')
    ln.append(prstDash)

    # Add accent text inside the circle
    tx_circle = slide.shapes.add_textbox(Inches(9.5), Inches(0.5), Inches(3), Inches(2))
    tf = tx_circle.text_frame
    p = tf.paragraphs[0]
    p.text = accent_word
    p.font.size = Pt(65)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Main Typography ===
    
    # "THE" - Rotated vertical structural text
    tx_side = slide.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(2), Inches(1))
    tx_side.rotation = 270
    p = tx_side.text_frame.paragraphs[0]
    p.text = side_text.upper()
    p.font.size = Pt(45)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # "BEST" - Massive Yellow Highlight
    tx_top = slide.shapes.add_textbox(Inches(2.5), Inches(0.2), Inches(8), Inches(2))
    p = tx_top.text_frame.paragraphs[0]
    p.text = headline_top.upper()
    p.font.size = Pt(140)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(255, 192, 0)

    # "WAY TO" - White connecting text
    tx_mid = slide.shapes.add_textbox(Inches(2.5), Inches(2.2), Inches(8), Inches(1.5))
    p = tx_mid.text_frame.paragraphs[0]
    p.text = headline_mid.upper()
    p.font.size = Pt(85)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 3: Colorful Letter Blocks ===
    # Define a high-contrast palette typical for kinetic typography
    block_colors = [
        RGBColor(220, 20, 60),   # Crimson Red
        RGBColor(0, 112, 192),   # Bright Blue
        RGBColor(0, 176, 80),    # Emerald Green
        RGBColor(112, 48, 160),  # Purple
        RGBColor(255, 102, 0)    # Orange
    ]
    
    # Alternating rotation angles for a dynamic "bouncing" feel
    rotations = [-8, 6, -11, 8, -5, 10, -7]
    
    block_size = 1.3  # inches
    gap = 0.15        # inches
    start_x = 2.6     # inches
    start_y = 3.9     # inches

    word = block_word.upper()[:10] # Limit to 10 chars for safety
    for i, char in enumerate(word):
        x = start_x + (i * (block_size + gap))
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(x), Inches(start_y), 
            Inches(block_size), Inches(block_size)
        )
        
        # Style the block
        rect.fill.solid()
        rect.fill.fore_color.rgb = block_colors[i % len(block_colors)]
        rect.line.color.rgb = RGBColor(255, 255, 255)
        rect.line.width = Pt(3)
        rect.rotation = rotations[i % len(rotations)]

        # Add the letter
        tf = rect.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = char
        p.font.size = Pt(75)
        p.font.bold = True
        p.font.name = "Arial Black"
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

    # === Layer 4: Cinematic Filmstrip Overlay ===
    # Creates a full-width film strip graphic intersecting the bottom of the layout
    strip_y_inch = 5.8
    strip_h_inch = 1.7
    
    # Film background
    film_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(strip_y_inch), 
        Inches(13.333), Inches(strip_h_inch)
    )
    film_bg.fill.solid()
    film_bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
    film_bg.line.fill.background()
    
    # Programmatically draw the film perforations (white squares)
    hole_size = 0.15
    hole_step = 0.35
    current_x = 0.1
    
    while current_x < 13.333:
        # Top perforation
        h_top = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(current_x), Inches(strip_y_inch + 0.15), 
            Inches(hole_size), Inches(hole_size)
        )
        h_top.fill.solid()
        h_top.fill.fore_color.rgb = RGBColor(255, 255, 255)
        h_top.line.fill.background()

        # Bottom perforation
        h_bot = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(current_x), Inches(strip_y_inch + strip_h_inch - 0.3), 
            Inches(hole_size), Inches(hole_size)
        )
        h_bot.fill.solid()
        h_bot.fill.fore_color.rgb = RGBColor(255, 255, 255)
        h_bot.line.fill.background()
        
        current_x += hole_step

    prs.save(output_pptx_path)
    return output_pptx_path
