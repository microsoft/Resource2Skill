def create_slide(
    output_pptx_path: str,
    title_text: str = "Photography Fundamentals",
    highlight_text: str = "Fundamentals",
    color_60: tuple = (242, 242, 242),  # Light Gray (Primary/Background)
    color_30: tuple = (91, 68, 106),    # Plum/Dark Purple (Secondary/Structure)
    color_10: tuple = (76, 145, 115),   # Dark Green (Accent/Highlight)
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 60-30-10 Proportional Color Harmony effect.
    This creates a feature slide that inherently uses the rule, and includes 
    the educational color proportion bar as seen in the tutorial.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    # Set 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Convert tuples to RGBColor objects
    c60 = RGBColor(*color_60)
    c30 = RGBColor(*color_30)
    c10 = RGBColor(*color_10)

    # === Layer 1: Background (The 60%) ===
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = c60
    background.line.fill.background() # No line

    # === Layer 2: The 60-30-10 Demonstration Bar ===
    # A literal representation of the 1000px bar from the video (scaled to 10 inches wide)
    bar_width_total = 10.0
    bar_height = 0.5
    start_x = (13.333 - bar_width_total) / 2
    start_y = Inches(6.0)

    # 60% block
    block_60 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, start_x * Inches(1), start_y, Inches(bar_width_total * 0.6), Inches(bar_height)
    )
    block_60.fill.solid()
    block_60.fill.fore_color.rgb = c60
    block_60.line.color.rgb = c30 # Give it a border so it's visible against the identical background

    # 30% block
    block_30 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(start_x + (bar_width_total * 0.6)), start_y, Inches(bar_width_total * 0.3), Inches(bar_height)
    )
    block_30.fill.solid()
    block_30.fill.fore_color.rgb = c30
    block_30.line.fill.background()

    # 10% block
    block_10 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(start_x + (bar_width_total * 0.9)), start_y, Inches(bar_width_total * 0.1), Inches(bar_height)
    )
    block_10.fill.solid()
    block_10.fill.fore_color.rgb = c10
    block_10.line.fill.background()


    # === Layer 3: Typography & Content (The 30% and 10% mix) ===
    
    # Title Text Box
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.333), Inches(1.5))
    tf = title_box.text_frame
    tf.clear() # Clear default
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    # Split title to apply accent color to the highlight word
    parts = title_text.split(highlight_text)
    
    if len(parts) > 1:
        # Before highlight
        run1 = p.add_run()
        run1.text = parts[0]
        run1.font.size = Pt(48)
        run1.font.bold = True
        run1.font.color.rgb = c30  # Secondary color (30%)
        
        # The highlight
        run2 = p.add_run()
        run2.text = highlight_text
        run2.font.size = Pt(48)
        run2.font.bold = True
        run2.font.color.rgb = c10  # Accent color (10%)
        
        # After highlight
        run3 = p.add_run()
        run3.text = parts[1]
        run3.font.size = Pt(48)
        run3.font.bold = True
        run3.font.color.rgb = c30  # Secondary color (30%)
    else:
        # Fallback if highlight word isn't in title
        run = p.add_run()
        run.text = title_text
        run.font.size = Pt(48)
        run.font.bold = True
        run.font.color.rgb = c30

    # Add feature layout (4 columns) to represent the 30% structure
    features = ["Lighting", "Focus", "Framing", "Equipment"]
    cols = 4
    spacing = 2.5
    start_col_x = (13.333 - (spacing * (cols - 1))) / 2

    for i, feature in enumerate(features):
        center_x = start_col_x + (i * spacing)
        
        # Icon placeholder (Circle) - uses 30% secondary color
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(center_x - 0.75), Inches(2.8), Inches(1.5), Inches(1.5)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = c30
        circle.line.fill.background()

        # Inner icon detail (to add some depth using the 60% color)
        inner_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(center_x - 0.5), Inches(3.05), Inches(1.0), Inches(1.0)
        )
        inner_circle.fill.solid()
        inner_circle.fill.fore_color.rgb = c60
        inner_circle.line.color.rgb = c30
        inner_circle.line.width = Pt(2)

        # Feature label - uses 30% secondary color
        label_box = slide.shapes.add_textbox(Inches(center_x - 1.5), Inches(4.5), Inches(3), Inches(0.5))
        lbl_tf = label_box.text_frame
        lbl_p = lbl_tf.paragraphs[0]
        lbl_p.alignment = PP_ALIGN.CENTER
        lbl_run = lbl_p.add_run()
        lbl_run.text = feature
        lbl_run.font.size = Pt(24)
        lbl_run.font.color.rgb = c30
        lbl_run.font.bold = True

        # Tiny accent dot under each label - uses 10% accent color
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(center_x - 0.1), Inches(5.1), Inches(0.2), Inches(0.2)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = c10
        dot.line.fill.background()

    # Footer/Author line
    footer = slide.shapes.add_textbox(Inches(1), Inches(6.7), Inches(11.333), Inches(0.5))
    ft_tf = footer.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.alignment = PP_ALIGN.CENTER
    ft_run = ft_p.add_run()
    ft_run.text = "Photography for Beginners | Color Proportion Demo"
    ft_run.font.size = Pt(14)
    ft_run.font.color.rgb = c30

    prs.save(output_pptx_path)
    return output_pptx_path
