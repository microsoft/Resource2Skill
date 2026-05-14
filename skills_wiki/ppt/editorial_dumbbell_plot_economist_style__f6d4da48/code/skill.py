def create_slide(
    output_pptx_path: str,
    title_text: str = "Engagement scores have shifted significantly",
    subtitle_text: str = "Employee Engagement Score by Department, 2023 vs 2024 (%)",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Editorial Dumbbell Plot" style.
    Generates a completely native, vector-based connected dot plot using python-pptx shapes.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # --- Colors ---
    COLOR_RED_ACCENT = RGBColor(227, 26, 28)
    COLOR_TEXT_MAIN = RGBColor(51, 51, 51)
    COLOR_TEXT_MUTED = RGBColor(120, 120, 120)
    COLOR_LINE_GRID = RGBColor(220, 220, 220)
    COLOR_DOT_BASE = RGBColor(154, 172, 184)  # 2023 - Muted Blue-Gray
    COLOR_DOT_NEW = RGBColor(216, 83, 73)     # 2024 - Punchy Orange-Red

    # --- Sample Data ---
    data = [
        {"cat": "Production", "val1": 20, "val2": 35},
        {"cat": "Finance", "val1": 60, "val2": 70},
        {"cat": "Customer Service", "val1": 80, "val2": 70},
        {"cat": "Operations", "val1": 85, "val2": 78},
        {"cat": "IT", "val1": 45, "val2": 55},
        {"cat": "Procurement", "val1": 32, "val2": 55},
        {"cat": "Sales", "val1": 60, "val2": 65},
        {"cat": "HR", "val1": 75, "val2": 80},
    ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- 1. Editorial Chrome (Top Left Accent & Text) ---
    
    # Red editorial accent block
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.4), Inches(0.5), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_RED_ACCENT
    accent.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.6), Inches(12.0), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.font.name = "Arial"

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.45), Inches(1.15), Inches(12.0), Inches(0.4))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.add_paragraph()
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.font.name = "Arial"

    # --- 2. Chart Layout Parameters ---
    chart_left = Inches(2.5)
    chart_top = Inches(2.2)
    chart_width = Inches(9.0)
    chart_height = Inches(4.5)
    
    max_val = 100.0
    x_scale = chart_width / max_val
    num_items = len(data)
    row_height = chart_height / num_items

    # --- 3. Draw Legend ---
    legend_top = Inches(1.7)
    
    # Legend Base Dot
    slide.shapes.add_shape(MSO_SHAPE.OVAL, chart_left, legend_top, Inches(0.15), Inches(0.15)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLOR_DOT_BASE
    slide.shapes[-1].line.fill.background()
    lbl1 = slide.shapes.add_textbox(chart_left + Inches(0.2), legend_top - Inches(0.05), Inches(1), Inches(0.3))
    lbl1.text_frame.paragraphs[0].text = "2023"
    lbl1.text_frame.paragraphs[0].font.size = Pt(12)
    lbl1.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MUTED

    # Legend New Dot
    legend_new_x = chart_left + Inches(1.0)
    slide.shapes.add_shape(MSO_SHAPE.OVAL, legend_new_x, legend_top, Inches(0.15), Inches(0.15)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLOR_DOT_NEW
    slide.shapes[-1].line.fill.background()
    lbl2 = slide.shapes.add_textbox(legend_new_x + Inches(0.2), legend_top - Inches(0.05), Inches(1), Inches(0.3))
    lbl2.text_frame.paragraphs[0].text = "2024"
    lbl2.text_frame.paragraphs[0].font.size = Pt(12)
    lbl2.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MUTED

    # --- 4. Draw X-Axis Guides (Vertical Lines) ---
    for x_val in range(0, 101, 20):
        x_pos = chart_left + (x_val * x_scale)
        # Vertical grid line
        grid_line = slide.shapes.add_connector(MSO_SHAPE.LINE_INVERSE, x_pos, chart_top - Inches(0.2), x_pos, chart_top + chart_height)
        grid_line.line.color.rgb = COLOR_LINE_GRID
        grid_line.line.width = Pt(0.5)
        # Axis label
        lbl_box = slide.shapes.add_textbox(x_pos - Inches(0.3), chart_top - Inches(0.5), Inches(0.6), Inches(0.3))
        p = lbl_box.text_frame.paragraphs[0]
        p.text = f"{x_val}%"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # --- 5. Draw Dumbbell Data Elements ---
    dot_size = Inches(0.16)
    half_dot = dot_size / 2

    for i, row in enumerate(data):
        y_center = chart_top + (i * row_height) + (row_height / 2)
        
        # Category Label (Y-Axis)
        cat_box = slide.shapes.add_textbox(Inches(0.5), y_center - Inches(0.2), chart_left - Inches(0.6), Inches(0.4))
        p_cat = cat_box.text_frame.paragraphs[0]
        p_cat.text = row["cat"]
        p_cat.font.size = Pt(12)
        p_cat.font.color.rgb = COLOR_TEXT_MAIN
        p_cat.alignment = PP_ALIGN.RIGHT
        cat_box.text_frame.margin_right = 0

        # Calculate X positions
        x1 = chart_left + (row["val1"] * x_scale)
        x2 = chart_left + (row["val2"] * x_scale)
        
        # Draw Connector Line
        conn = slide.shapes.add_connector(MSO_SHAPE.LINE_INVERSE, x1, y_center, x2, y_center)
        conn.line.color.rgb = COLOR_TEXT_MAIN
        conn.line.width = Pt(1.5)

        # Draw Base Dot (val1)
        dot1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, x1 - half_dot, y_center - half_dot, dot_size, dot_size)
        dot1.fill.solid()
        dot1.fill.fore_color.rgb = COLOR_DOT_BASE
        dot1.line.fill.background()

        # Draw Target Dot (val2)
        dot2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, x2 - half_dot, y_center - half_dot, dot_size, dot_size)
        dot2.fill.solid()
        dot2.fill.fore_color.rgb = COLOR_DOT_NEW
        dot2.line.fill.background()

        # Draw Value Labels (floating near dots)
        # Offset labels slightly so they don't overlap dots. If value is going up, 2023 on left, 2024 on right.
        if row["val2"] >= row["val1"]:
            val1_align, val1_offset = PP_ALIGN.RIGHT, -Inches(0.55)
            val2_align, val2_offset = PP_ALIGN.LEFT, Inches(0.1)
        else:
            val1_align, val1_offset = PP_ALIGN.LEFT, Inches(0.1)
            val2_align, val2_offset = PP_ALIGN.RIGHT, -Inches(0.55)

        # Label 1
        l1 = slide.shapes.add_textbox(x1 + val1_offset, y_center - Inches(0.2), Inches(0.45), Inches(0.4))
        p_l1 = l1.text_frame.paragraphs[0]
        p_l1.text = str(row["val1"])
        p_l1.alignment = val1_align
        p_l1.font.size = Pt(10)
        p_l1.font.color.rgb = COLOR_DOT_BASE
        
        # Label 2
        l2 = slide.shapes.add_textbox(x2 + val2_offset, y_center - Inches(0.2), Inches(0.45), Inches(0.4))
        p_l2 = l2.text_frame.paragraphs[0]
        p_l2.text = str(row["val2"])
        p_l2.alignment = val2_align
        p_l2.font.size = Pt(10)
        p_l2.font.bold = True
        p_l2.font.color.rgb = COLOR_DOT_NEW

    prs.save(output_pptx_path)
    return output_pptx_path
