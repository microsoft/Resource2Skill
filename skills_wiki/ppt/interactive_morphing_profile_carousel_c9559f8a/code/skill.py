def create_slide(output_pptx_path: str, **kwargs) -> str:
    """
    Create a PPTX file reproducing the Interactive Morphing Profile Carousel.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw
    from pptx.oxml.ns import qn
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Define team profiles with vibrant color gradients
    profiles = [
        {"name": "Jake", "role": "Creative Director", "c1": (230, 57, 70), "c2": (180, 20, 40)},
        {"name": "Marie", "role": "Lead Designer", "c1": (244, 162, 97), "c2": (210, 110, 30)},
        {"name": "Angela", "role": "Project Manager", "c1": (155, 93, 229), "c2": (100, 40, 180)},
        {"name": "Monica", "role": "UI/UX Specialist", "c1": (0, 187, 249), "c2": (0, 100, 200)},
        {"name": "Jennifer", "role": "Frontend Dev", "c1": (0, 245, 212), "c2": (0, 160, 140)}
    ]

    # Step 1: Generate high-quality placeholder profile cards using PIL
    img_paths = []
    for i, p in enumerate(profiles):
        path = f"profile_temp_{i}.png"
        w, h = 400, 800
        img = Image.new('RGB', (w, h))
        
        # Draw vertical gradient background
        for y in range(h):
            r = int(p["c1"][0] + (p["c2"][0] - p["c1"][0]) * y / h)
            g = int(p["c1"][1] + (p["c2"][1] - p["c1"][1]) * y / h)
            b = int(p["c1"][2] + (p["c2"][2] - p["c1"][2]) * y / h)
            ImageDraw.Draw(img).line([(0, y), (w, y)], fill=(r, g, b))
        
        # Add a subtle, translucent geometric "person" overlay
        overlay = Image.new('RGBA', (w, h), (0,0,0,0))
        draw = ImageDraw.Draw(overlay)
        draw.ellipse([(120, 200), (280, 360)], fill=(255, 255, 255, 120))
        draw.rounded_rectangle([(70, 420), (330, 900)], radius=70, fill=(255, 255, 255, 120))
        
        img = Image.alpha_composite(img.convert('RGBA'), overlay).convert('RGB')
        img.save(path)
        img_paths.append(path)

    # Step 2: Initialize layout topology (1 Base slide + 5 Detail slides)
    layout = prs.slide_layouts[6] # Blank slide
    base_slide = prs.slides.add_slide(layout)
    detail_slides = [prs.slides.add_slide(layout) for _ in range(5)]
    all_slides = [base_slide] + detail_slides

    # Helper: Inject Morph Transition via OOXML
    def apply_morph(slide):
        sld = slide.element
        transition = parse_xml(r'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="fast"><p:morph/></p:transition>')
        # Remove any existing transition elements safely
        for existing in sld.xpath('./p:transition'):
            sld.remove(existing)
        # Insert transition in correct XML sequence order (after cSld / clrMapOvr)
        insert_idx = 0
        for idx, child in enumerate(sld):
            if child.tag in (qn('p:cSld'), qn('p:clrMapOvr')):
                insert_idx = idx + 1
        sld.insert(insert_idx, transition)

    # Helper: Mathematical layout distributor
    def get_position(idx, active_idx):
        y_center = 3.75
        
        # State A: Base Gallery (Evenly spaced)
        if active_idx is None:
            w, h = 1.8, 3.5
            gap = (13.333 - (5 * w)) / 6
            x = gap + (w / 2) + idx * (w + gap)
            return x, y_center, w, h
        
        # State B: Active Image in Detail View (Hero sizing)
        if idx == active_idx:
            w, h = 3.6, 5.5
            return 13.333/2, 3.25, w, h
        
        # State C: Inactive Images in Detail View (Pushed to margins)
        w, h = 1.1, 2.2
        if idx < active_idx: # Distribute to left margin
            n_left = active_idx
            x = 2.5 if n_left == 1 else 0.8 + idx * ((4.2 - 0.8) / (n_left - 1))
        else: # Distribute to right margin
            n_right = 5 - active_idx - 1
            local_i = idx - active_idx - 1
            x = 13.333 - 2.5 if n_right == 1 else 9.133 + local_i * ((12.533 - 9.133) / (n_right - 1))
        
        return x, y_center, w, h

    # Step 3: Populate all slides ensuring rigid shape sequence for Morph matching
    for slide_idx, slide in enumerate(all_slides):
        apply_morph(slide)
        active_idx = None if slide_idx == 0 else slide_idx - 1
        
        # MUST add shapes in the exact same 0-4 sequence on every slide
        for i in range(5):
            x, y, w, h = get_position(i, active_idx)
            
            # Interactive wiring: clicking the active image goes back to base; otherwise, jump to detail
            target_slide = base_slide if i == active_idx else detail_slides[i]
            
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, 
                Inches(x - w/2), Inches(y - h/2), Inches(w), Inches(h)
            )
            # Dynamic corner radius adjustment (keeps pixel-radius consistent despite varying widths)
            shape.adjustments[0] = min(0.25, 0.15 / w)
            
            # Stylize and bind
            shape.fill.user_picture(img_paths[i])
            shape.line.color.rgb = RGBColor(255, 255, 255)
            shape.line.width = Pt(1.5)
            shape.click_action.target_slide = target_slide
            
        # Add metadata text box ONLY on detail slides
        if active_idx is not None:
            txBox = slide.shapes.add_textbox(Inches(13.333/2 - 2), Inches(6.2), Inches(4), Inches(1))
            tf = txBox.text_frame
            
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = profiles[active_idx]["name"]
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(50, 50, 50)
            
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.text = profiles[active_idx]["role"]
            p2.font.size = Pt(14)
            p2.font.color.rgb = RGBColor(120, 120, 120)

    # Step 4: Save and cleanup
    prs.save(output_pptx_path)
    
    for path in img_paths:
        if os.path.exists(path):
            os.remove(path)
            
    return output_pptx_path
