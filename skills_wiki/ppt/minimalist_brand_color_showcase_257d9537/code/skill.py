def create_slide(
    output_pptx_path: str,
    title_text: str = "Boho Color Palette",
    colors: list = None,
    bg_color_hex: str = "EFEAD5",
    text_color_hex: str = "40303A",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Minimalist Brand Color Showcase visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    def hex_to_rgb(hex_str):
        hex_str = str(hex_str).lstrip('#')
        return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

    # Default Boho Palette if none provided
    if colors is None:
        colors = [
            {"hex": "E38150", "name": "Rust"},
            {"hex": "FFB3A0", "name": "Peach"},
            {"hex": "EFC16D", "name": "Mustard"},
            {"hex": "B6CDBD", "name": "Mint"},
            {"hex": "829E8B", "name": "Sage"}
        ]
        
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    r_bg, g_bg, b_bg = hex_to_rgb(bg_color_hex)
    fill.fore_color.rgb = RGBColor(r_bg, g_bg, b_bg)

    # Calculate global text color
    r_txt, g_txt, b_txt = hex_to_rgb(text_color_hex)
    
    # === Layer 2: Title ===
    txBox = slide.shapes.add_textbox(Inches(0), Inches(1.2), Inches(13.333), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(54)
    run.font.bold = True
    run.font.name = "Arial Black"  # Heavy font for high contrast impact
    run.font.color.rgb = RGBColor(r_txt, g_txt, b_txt)

    # === Layer 3: Color Swatches & Labels ===
    num_colors = len(colors)
    circle_size = 2.4  # Diameter in inches
    center_dist = 2.0  # Distance between centers (creates a 0.4 inch overlap)
    
    # Calculate starting X to perfectly center the clustered shapes
    total_width = circle_size + (num_colors - 1) * center_dist
    start_x = (13.333 - total_width) / 2
    
    y_pos = 3.2  # Top Y coordinate for the circles
    y_center = y_pos + (circle_size / 2)

    for i, color_dict in enumerate(colors):
        x_pos = start_x + i * center_dist
        x_center = x_pos + (circle_size / 2)

        # Draw overlapping circle
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x_pos), Inches(y_pos), Inches(circle_size), Inches(circle_size)
        )
        fill = shape.fill
        fill.solid()
        r, g, b = hex_to_rgb(color_dict['hex'])
        fill.fore_color.rgb = RGBColor(r, g, b)
        
        # Match line color to fill to seamlessly hide the border
        shape.line.color.rgb = RGBColor(r, g, b)

        # Label text box (positioned securely below each circle's exact center)
        txBox_hex = slide.shapes.add_textbox(
            Inches(x_center - 1.5), Inches(y_center + circle_size / 2 + 0.4), Inches(3), Inches(1.0)
        )
        tf_hex = txBox_hex.text_frame
        p_hex = tf_hex.paragraphs[0]
        p_hex.alignment = PP_ALIGN.CENTER
        
        # Insert Color Name (Sans-Serif)
        if 'name' in color_dict and color_dict['name']:
            run_name = p_hex.add_run()
            run_name.text = color_dict['name'].upper() + "\n"
            run_name.font.size = Pt(13)
            run_name.font.bold = True
            run_name.font.name = "Arial"
            run_name.font.color.rgb = RGBColor(r_txt, g_txt, b_txt)
            
        # Insert Hex Code (Monospace)
        run_hex = p_hex.add_run()
        run_hex.text = "#" + color_dict['hex'].upper()
        run_hex.font.size = Pt(16)
        run_hex.font.name = "Courier New"
        run_hex.font.color.rgb = RGBColor(r_txt, g_txt, b_txt)

    prs.save(output_pptx_path)
    return output_pptx_path
