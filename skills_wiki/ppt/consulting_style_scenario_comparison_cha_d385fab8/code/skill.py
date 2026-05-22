def create_slide(
    output_pptx_path: str,
    title_text: str = "Scenario Analysis",
    bau_label: str = "BAU",
    accel_label: str = "Acceleration",
    callout_text: str = "Additional 100K Investment\nin Product Marketing",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a consulting-style scenario comparison chart.

    This chart visually contrasts a 'Business as Usual' (BAU) scenario with an
    'Acceleration' scenario, using color, line styles, and an annotation
    to justify a business decision.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.shapes.freeform import FreeformBuilder

    # --- Presentation and Slide Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # --- Chart Data ---
    chart_data = CategoryChartData()
    chart_data.categories = ['2021', '2022', '2023', '2024', '2025', '2026']
    # Series 1: Business as Usual (BAU)
    chart_data.add_series(bau_label, (1500, 2500, 3000, 8000, 10000, 10000))
    # Series 2: Acceleration
    chart_data.add_series(accel_label, (None, None, 3000, 11000, 12500, 15000)) # Start from divergence

    # --- Chart Creation ---
    x, y, cx, cy = Inches(1), Inches(1.5), Inches(11), Inches(5)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, x, y, cx, cy, chart_data
    ).chart

    # --- Chart Formatting ---
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.chart_title.text_frame.text = title_text
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(24)

    # --- Value Axis (Y-axis) Formatting ---
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor(217, 217, 217)
    value_axis.minimum_scale = 0.0
    value_axis.maximum_scale = 15000.0
    value_axis.tick_labels.font.size = Pt(12)
    value_axis.format.line.color.rgb = RGBColor(255, 255, 255) # Hide axis line

    # --- Category Axis (X-axis) Formatting ---
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(12)
    category_axis.format.line.color.rgb = RGBColor(180, 180, 180)
    category_axis.major_tick_mark = XL_TICK_MARK.OUTSIDE
    
    # --- Series Formatting ---
    # Series 1: BAU (Blue, Solid)
    bau_series = chart.series[0]
    bau_series.format.line.color.rgb = RGBColor(0, 112, 192)
    bau_series.format.line.width = Pt(2.5)

    # Series 2: Acceleration (Red, Dashed)
    accel_series = chart.series[1]
    accel_series.format.line.color.rgb = RGBColor(192, 0, 0)
    accel_series.format.line.width = Pt(2.5)
    accel_series.format.line.dash_style = MSO_LINE.DASH

    # --- Callout Annotation ---
    # 1. The Text Box
    callout_box = slide.shapes.add_textbox(
        Inches(2.5), Inches(2.2), Inches(2.0), Inches(0.8)
    )
    callout_box.text_frame.text = callout_text
    p = callout_box.text_frame.paragraphs[0]
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(192, 0, 0)
    
    fill = callout_box.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 235, 235)
    
    line = callout_box.line
    line.color.rgb = RGBColor(192, 0, 0)
    line.width = Pt(1)

    # 2. The Leader Line (using FreeformBuilder)
    # Positions are estimated to point to the '2023' data point
    shapes = slide.shapes
    x1, y1 = Inches(4.5), Inches(2.6)  # Start from edge of box
    x2, y2 = Inches(5.2), Inches(2.6)  # Horizontal segment
    x3, y3 = Inches(5.2), Inches(5.2)  # Vertical segment pointing to chart
    
    with FreeformBuilder(shapes, x1, y1) as builder:
        builder.add_line_segment(x2, y2)
        builder.add_line_segment(x3, y3)
    freeform_shape = builder.close()
    
    leader_line = freeform_shape.line
    leader_line.color.rgb = RGBColor(192, 0, 0)
    leader_line.width = Pt(1.5)

    # --- Save Presentation ---
    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_slide("scenario_chart.pptx")

