import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw, ImageFont, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "GET FIT",
    subtitle_text: str = "DOWNLOAD TODAY",
    bg_color: tuple = (13, 17, 28),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Isometric Glassmorphism UI Stack' effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Helper to load a generic sans-serif font
    def get_font(size):
        try:
            return ImageFont.truetype("arial.ttf", size)
        except IOError:
            return ImageFont.load_default()

    # ==========================================
    # 1. Generate Ambient Background
    # ==========================================
    bg_path = "temp_bg.png"
    bg_img = Image.new("RGBA", (1920, 1080), bg_color + (255,))
    
    # Create glowing orbs
    orb1 = Image.new("RGBA", (800, 800), (0, 0, 0, 0))
    ImageDraw.Draw(orb1).ellipse([0, 0, 800, 800], fill=(255, 0, 128, 140)) # Magenta
    orb1 = orb1.filter(ImageFilter.GaussianBlur(180))
    
    orb2 = Image.new("RGBA", (800, 800), (0, 0, 0, 0))
    ImageDraw.Draw(orb2).ellipse([0, 0, 800, 800], fill=(0, 229, 255, 120)) # Cyan
    orb2 = orb2.filter(ImageFilter.GaussianBlur(180))
    
    bg_img.paste(orb1, (200, -100), orb1)
    bg_img.paste(orb2, (900, 400), orb2)
    bg_img.save(bg_path)
    
    # Apply to slide background
    slide_bg = slide.background
    fill = slide_bg.fill
    fill.solid()
    # Setting an image as background requires a shape workaround in python-pptx
    bg_shape = slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    slide.shapes._spTree.remove(bg_shape._element)
    slide.shapes._spTree.insert(2, bg_shape._element)

    # ==========================================
    # 2. Generate Base Phone UI Layer (Flat)
    # ==========================================
    base_path = "temp_phone_base.png"
    # We use a large square canvas so all layers share the same 3D pivot point
    canvas_size = 1200
    base_img = Image.new("RGBA", (canvas_size, canvas_size), (0, 0, 0, 0))
    draw_base = ImageDraw.Draw(base_img)
    
    # Phone Chassis
    phone_box = [350, 150, 850, 1050]
    draw_base.rounded_rectangle(phone_box, radius=45, fill=(26, 27, 38, 255), outline=(85, 85, 85, 255), width=4)
    # Notch
    draw_base.rounded_rectangle([520, 150, 680, 190], radius=15, fill=(15, 15, 20, 255))
    
    # UI Element: Progress Ring
    ring_box = [480, 280, 720, 520]
    draw_base.arc(ring_box, start=135, end=405, fill=(0, 229, 255, 255), width=25)
    draw_base.arc(ring_box, start=405, end=495, fill=(40, 40, 60, 255), width=25)
    
    # UI Element: Inner Text
    fnt_large = get_font(50)
    fnt_small = get_font(25)
    draw_base.text((600, 380), "5678", font=fnt_large, fill=(255, 255, 255, 255), anchor="mm")
    draw_base.text((600, 430), "STEPS", font=fnt_small, fill=(180, 180, 180, 255), anchor="mm")
    
    # UI Elements: Bottom Grid Buttons
    btn_y1, btn_y2 = 620, 760
    buttons = [
        ([400, btn_y1, 580, btn_y1+120], "HEART", (255, 0, 128)),
        ([620, btn_y1, 800, btn_y1+120], "WATER", (0, 229, 255)),
        ([400, btn_y2, 580, btn_y2+120], "SLEEP", (150, 0, 255)),
        ([620, btn_y2, 800, btn_y2+120], "DIET", (0, 255, 150))
    ]
    for box, text, color in buttons:
        draw_base.rounded_rectangle(box, radius=20, fill=(30, 32, 45, 255), outline=(60, 60, 80, 255), width=2)
        draw_base.ellipse([box[0]+20, box[1]+40, box[0]+60, box[1]+80], fill=color)
        draw_base.text((box[0]+120, box[1]+60), text, font=fnt_small, fill=(200, 200, 200, 255), anchor="mm")

    base_img.save(base_path)

    # ==========================================
    # 3. Generate Floating Glass Layer (Flat)
    # ==========================================
    glass_path = "temp_phone_glass.png"
    glass_img = Image.new("RGBA", (canvas_size, canvas_size), (0, 0, 0, 0))
    draw_glass = ImageDraw.Draw(glass_img)
    
    # Floating Glass Panel over the ring
    glass_box1 = [430, 370, 770, 560]
    draw_glass.rounded_rectangle(glass_box1, radius=25, fill=(255, 255, 255, 15), outline=(255, 255, 255, 100), width=2)
    draw_glass.text((600, 510), "GOAL: 10,000", font=fnt_small, fill=(255, 255, 255, 255), anchor="mm")
    
    glass_img.save(glass_path)

    # ==========================================
    # 4. Insert into PPTX and Inject 3D LXML
    # ==========================================
    # Define exact same placement for both layers so their 3D pivot aligns perfectly
    pic_left = Inches(1.0)
    pic_top = Inches(0.0)
    pic_height = Inches(7.5) # Scale the 1200px canvas to slide height
    
    pic_base = slide.shapes.add_picture(base_path, pic_left, pic_top, height=pic_height)
    pic_glass = slide.shapes.add_picture(glass_path, pic_left, pic_top, height=pic_height)

    # LXML Helper to apply 3D properties
    def apply_3d_to_pic(pic, prst, z_offset_emu, depth_emu):
        spPr = pic.element.spPr
        scene3d_xml = f'''
        <a:scene3d xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:camera prst="{prst}"/>
            <a:lightRig rig="threePt" dir="t"/>
        </a:scene3d>
        '''
        # z = distance from ground (levitation), extrusionH = physical thickness
        sp3d_xml = f'''
        <a:sp3d xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" z="{int(z_offset_emu)}" extrusionH="{int(depth_emu)}">
            <a:extrusionClr><a:srgbClr val="1A1B26"/></a:extrusionClr>
        </a:sp3d>
        '''
        spPr.append(parse_xml(scene3d_xml))
        spPr.append(parse_xml(sp3d_xml))

    # Apply Perspective Right to both. 
    # Base layer gets physical depth (extrusion). Glass gets levitation (z-offset).
    preset = "perspectiveContrastingRight"
    apply_3d_to_pic(pic_base, preset, z_offset_emu=0, depth_emu=200000)
    apply_3d_to_pic(pic_glass, preset, z_offset_emu=500000, depth_emu=0) # Float ~0.5 inches out

    # ==========================================
    # 5. Add Right-side Typography
    # ==========================================
    tb = slide.shapes.add_textbox(Inches(7.5), Inches(3.0), Inches(5.0), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 229, 255)
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(32)
    p2.font.color.rgb = RGBColor(200, 200, 200)

    # Cleanup temp files (optional but good practice)
    for f in [bg_path, base_path, glass_path]:
        if os.path.exists(f):
            os.remove(f)

    prs.save(output_pptx_path)
    return output_pptx_path
