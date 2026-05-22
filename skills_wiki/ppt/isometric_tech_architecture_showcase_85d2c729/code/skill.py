import io
from lxml import etree
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

# Helper to add namespace prefixes to XML tags
def qn(tag):
    """
    Stands for 'qualified name', a utility function to turn a namespace prefixed
    tag name into a Clark-notation qualified tag name for lxml. For example,
    qn('p:cSld') returns '{http://schemas.openxmlformats.org/presentationml/2006/main}cSld'.
    """
    ns = {
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    }
    prefix, tagroot = tag.split(':')
    uri = ns[prefix]
    return f'{{{uri}}}{tagroot}'

def add_text_shadow(run):
    """Applies a soft outer shadow to a text run using lxml."""
    rPr = run._r.get_or_add_rPr()
    effect_lst = OxmlElement('a:effectLst')
    outer_shadow = OxmlElement('a:outerShdw')
    outer_shadow.set('blurRad', '40000')
    outer_shadow.set('dist', '20000')
    outer_shadow.set('dir', '2700000')
    outer_shadow.set('algn', 'bl')
    srgb_clr = OxmlElement('a:srgbClr')
    srgb_clr.set('val', '000000')
    alpha = OxmlElement('a:alpha')
    alpha.set('val', '50000') # 50% opacity
    srgb_clr.append(alpha)
    outer_shadow.append(srgb_clr)
    effect_lst.append(outer_shadow)
    rPr.append(effect_lst)

def create_isometric_server_png(width, height, depth, color1=(40, 40, 40), color2=(20, 20, 20)):
    """Creates a transparent PNG of an isometric server rack."""
    img_width = width + depth // 2
    img_height = height + depth // 2
    image = Image.new("RGBA", (img_width + 50, img_height + 50), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    
    # Coordinates
    x0, y0 = 25, 25 + depth // 2
    top = [(x0, y0), (x0 + width, y0), (x0 + width + depth // 2, y0 - depth // 2), (x0 + depth // 2, y0 - depth // 2)]
    front = [(x0, y0), (x0 + width, y0), (x0 + width, y0 + height), (x0, y0 + height)]
    side = [(x0 + width, y0), (x0 + width + depth // 2, y0 - depth // 2), (x0 + width + depth // 2, y0 + height - depth // 2), (x0 + width, y0 + height)]

    # Draw faces with gradients
    draw.polygon(top, fill=(60, 60, 60))
    for i in range(height):
        ratio = i / height
        inter_color = tuple(int(c1 + (c2 - c1) * ratio) for c1, c2 in zip(color1, color2))
        draw.line([(front[0][0], front[0][1] + i), (front[1][0], front[1][1] + i)], fill=inter_color)
    for i in range(depth // 2):
        ratio = i / (depth // 2)
        inter_color = tuple(int(c1 + (c2 - c1) * ratio) for c1, c2 in zip((50,50,50), (10,10,10)))
        draw.line([(side[0][0] + i, side[0][1] - i), (side[3][0] + i, side[3][1] - i)], fill=inter_color)

    # Add subtle details
    for i in range(1, 5):
        y_pos = y0 + (height // 5) * i
        draw.line([(x0 + 5, y_pos), (x0 + width - 5, y_pos)], fill=(70, 70, 70))

    # Create shadow
    shadow_img = Image.new("RGBA", image.size, (0, 0, 0, 0))
    shadow_draw = ImageDraw.Draw(shadow_img)
    shadow_draw.polygon(front, fill=(0, 0, 0, 80))
    shadow_draw.polygon(side, fill=(0, 0, 0, 80))
    shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(radius=8))

    final_image = Image.alpha_composite(shadow_img, image)
    return final_image

def create_isometric_database_png(radius, height, color1=(50, 50, 50), color2=(25, 25, 25)):
    """Creates a transparent PNG of an isometric database cylinder."""
    ellipse_height = radius // 2
    img_width = radius * 2 + 20
    img_height = height + ellipse_height + 20
    image = Image.new("RGBA", (img_width, img_height), (0, 0, 0, 0))
    
    # Create body with gradient
    body = Image.new("RGB", (radius * 2, height))
    for x in range(radius * 2):
        for y in range(height):
            ratio = x / (radius * 2)
            inter_color = tuple(int(c1 + (c2 - c1) * ratio) for c1, c2 in zip(color1, color2))
            body.putpixel((x, y), inter_color)
    image.paste(body, (10, 10 + ellipse_height))

    # Draw top ellipse
    draw = ImageDraw.Draw(image)
    top_bbox = (10, 10, 10 + radius * 2, 10 + ellipse_height)
    draw.ellipse(top_bbox, fill=(80, 80, 80), outline=(100, 100, 100))

    # Add glossy highlight to top
    highlight_bbox = (10 + radius // 2, 12, 10 + radius * 2 - radius // 2, 10 + ellipse_height - 4)
    draw.ellipse(highlight_bbox, fill=(120, 120, 120))
    
    # Shadow
    shadow_img = Image.new("RGBA", image.size, (0, 0, 0, 0))
    shadow_draw = ImageDraw.Draw(shadow_img)
    shadow_draw.ellipse((5, img_height - 20, img_width - 5, img_height-5), fill=(0, 0, 0, 100))
    shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(radius=6))
    
    final_image = Image.alpha_composite(shadow_img, image)
    return final_image

def create_slide(output_pptx_path: str):
    """
    Creates a PPTX file reproducing the Isometric Tech-Architecture Showcase effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = RGBColor(245, 245, 245)
    fill.gradient_stops[0].position = 0.0
    fill.gradient_stops[1].color.rgb = RGBColor(220, 220, 220)
    fill.gradient_stops[1].position = 1.0

    # === Layer 2: Visual Components (Generated with PIL) ===
    # Generate assets
    server_img = create_isometric_server_png(120, 150, 100)
    db_img = create_isometric_database_png(60, 90)
    desktop_img = create_isometric_server_png(80, 70, 60, color1=(50,50,50), color2=(30,30,30)) # Re-use for PC look

    # Convert PIL images to byte streams
    server_stream, db_stream, desktop_stream = io.BytesIO(), io.BytesIO(), io.BytesIO()
    server_img.save(server_stream, format="PNG")
    db_img.save(db_stream, format="PNG")
    desktop_img.save(desktop_stream, format="PNG")

    # Place images on slide
    pic_server1 = slide.shapes.add_picture(server_stream, Inches(3), Inches(1.5), height=Inches(2.5))
    pic_server2 = slide.shapes.add_picture(server_stream, Inches(7), Inches(1.5), height=Inches(2.5))
    pic_db1 = slide.shapes.add_picture(db_stream, Inches(2.7), Inches(4.5), height=Inches(1.8))
    pic_db2 = slide.shapes.add_picture(db_stream, Inches(5), Inches(4.5), height=Inches(1.8))
    pic_desktop = slide.shapes.add_picture(desktop_stream, Inches(12), Inches(3.5), height=Inches(1.5))
    
    # === Layer 3: Connectors & Text ===
    accent_color = RGBColor(204, 0, 0)
    line_width = Pt(3)

    # Connector from Server 2 to Desktop
    connector1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.2), Inches(2.7), Inches(12), Inches(4))
    line1 = connector1.line
    line1.color.rgb = accent_color
    line1.width = line_width
    
    # Connector from DB1 to Server 1
    connector2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3.5), Inches(4.5), Inches(3.8), Inches(3.8))
    line2 = connector2.line
    line2.color.rgb = accent_color
    line2.width = line_width

    # Labels with shadows
    def add_label(text, left, top, width, height):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        p = txBox.text_frame.paragraphs[0]
        p.font.name = 'Calibri Light'
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(255, 255, 255)
        run = p.add_run()
        run.text = text
        add_text_shadow(run)
    
    add_label("MVS\nMAINFRAME", 3, 1, 2, 0.5)
    add_label("PAYMENT\nPROCESSING", 7, 1, 2, 0.5)
    add_label("REGISTRATION DB", 2.5, 6.4, 2, 0.5)
    add_label("DATA\nWAREHOUSE", 4.8, 6.4, 2, 0.5)
    add_label("DESKTOP", 12.3, 5.1, 2, 0.5)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
if __name__ == '__main__':
    output_file = "isometric_tech_diagram.pptx"
    create_slide(output_file)
    print(f"Slide saved to {output_file}")
