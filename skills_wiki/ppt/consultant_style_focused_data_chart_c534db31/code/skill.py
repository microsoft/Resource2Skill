def create_slide(
    output_pptx_path: str,
    title_text: str = "Amidst lackluster sales for most tree species, sales of Kwanzaan Cherry have grown steadily since 2020",
    subtitle_text: str = "Trees Sold by Species (Aspira Nursery, 2020-2023)",
    highlight_series_name: str = "Kwanzen Cherry",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a "Consultant-Style Focused Data Chart".

    This function reproduces the data visualization best practice of using selective
    emphasis (color, direct labels) to tell a clear story, while de-emphasizing
    contextual data.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR

    # --- Data from the video tutorial ---
    chart_data_dict = {
        'Categories': ['2020', '2021', '2022', '2023'],
        'Series': [
            {'name': 'Thundercloud Plum', 'values': [83, 81, 80, 77]},
            {'name': 'Golden Rain Tree', 'values': [71, 75, 69, 68]},
            {'name': 'Kwanzen Cherry', 'values': [45, 51, 75, 103]},
            {'name': 'Tina Sargent Crabapple', 'values': [78, 71, 66, 55]}
        ]
    }

    # --- Color Palette ---
    COLOR_HIGHLIGHT = RGBColor(70, 130, 180)  # Steel Blue
    COLOR_MUTED = RGBColor(192, 192, 192)      # Light Gray
    COLOR_TEXT_DARK = RGBColor(30, 30, 30)
    COLOR_TEXT_LIGHT = RGBColor(128, 128, 128)
    COLOR_BACKGROUND = RGBColor(255, 255, 255)

    # --- Presentation and Slide Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Set slide background color
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = COLOR_BACKGROUND

    # --- Add Chart ---
    chart_data = ChartData()
    chart_data.categories = chart_data_dict['Categories']
    for series in chart_data_dict['Series']:
        chart_data.add_series(series['name'], series['values'])

    x, y, cx, cy = Inches(1), Inches(2.0), Inches(11.33), Inches(4.5)
    chart_graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, x, y, cx, cy, chart_data
    )
    chart = chart_graphic_frame.chart

    # --- Format Chart Elements ---
    chart.has_legend = False

    # Format value axis (Y-axis)
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor(230, 230, 230)
    value_axis.format.line.fill.background() # Effectively removes the axis line
    value_axis.tick_labels.font.size = Pt(10)
    value_axis.tick_labels.font.color.rgb = COLOR_TEXT_LIGHT

    # Format category axis (X-axis)
    category_axis = chart.category_axis
    category_axis.format.line.color.rgb = COLOR_MUTED
    category_axis.tick_labels.font.size = Pt(11)
    category_axis.tick_labels.font.color.rgb = COLOR_TEXT_DARK
    category_axis.tick_mark = XL_TICK_MARK.OUTSIDE

    # --- Format Data Series and Labels ---
    for i, series in enumerate(chart.series):
        is_highlight = series.name == highlight_series_name

        # Set line colors
        line_format = series.format.line
        line_format.width = Pt(2.5)
        line_format.color.rgb = COLOR_HIGHLIGHT if is_highlight else COLOR_MUTED

        # Add and format data labels
        data_labels = series.data_labels
        data_labels.show_category_name = False
        data_labels.show_value = False
        data_labels.show_series_name = True
        
        # We can only apply label settings to the whole series, but can format individual points
        # Apply settings to the last point's label
        last_point_label = data_labels.get_label(len(series.points) - 1)
        last_point_label.position = XL_LABEL_POSITION.RIGHT
        last_point_label.font.size = Pt(11)
        last_point_label.font.color.rgb = COLOR_HIGHLIGHT if is_highlight else COLOR_TEXT_LIGHT
        if is_highlight:
            last_point_label.font.bold = True
        
        # Hide labels for all other points
        for j in range(len(series.points) - 1):
             data_labels.get_label(j).show_series_name = False

    # Remove chart border
    chart.chart_style = 2 # A style with minimal formatting
    chart.format.line.fill.background()

    # --- Add Title and Subtitle ---
    title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(1.0))
    title_tf = title_shape.text_frame
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text
    title_p.font.name = 'Calibri'
    title_p.font.size = Pt(24)
    title_p.font.bold = True
    title_p.font.color.rgb = COLOR_TEXT_DARK

    subtitle_shape = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(0.5))
    subtitle_tf = subtitle_shape.text_frame
    subtitle_p = subtitle_tf.paragraphs[0]
    subtitle_p.text = subtitle_text
    subtitle_p.font.name = 'Calibri'
    subtitle_p.font.size = Pt(12)
    subtitle_p.font.color.rgb = COLOR_TEXT_LIGHT

    prs.save(output_pptx_path)
    return output_pptx_path
