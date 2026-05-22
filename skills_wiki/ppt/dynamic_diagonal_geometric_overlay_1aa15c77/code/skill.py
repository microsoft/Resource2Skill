def create_slide(
    output_pptx_path: str,
    title_text: str = "Product Launch\nGo-To Market Strategy",
    body_text: str = "A comprehensive roadmap for introducing your new product to the market successfully and maximizing initial impact.",
    bg_palette: str = "business,office",
    accent_color: tuple = (0, 168, 150),  # Teal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Dynamic Diagonal Geometric Overlay' visual effect.
    """
    import os
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Constants
    SLIDE_WIDTH_PX = 1920
    SLIDE_HEIGHT_PX = 1080
    
    # Colors based on video style
    color_base_navy = (34, 49, 63, 220)       # Dark slate/navy, high opacity
    color_accent_primary = accent_color + (230,) # Teal, high opacity
    color_accent_light = (123, 192, 227, 180) # Light blue, lower opacity

    # === Layer 1: Background Image ===
    bg_img_path = "temp_bg.jpg"
    try:
        # Try fetching a relevant Unsplash image
        url = f"https://source.unsplash.com/random/1920x1080/?{bg_palette}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(bg_img_path, 'wb') as f:
                f.write(response.read())
    except Exception:
        # Fallback: Create a solid neutral gray background if download fails
        fallback_bg = Image.new('RGB', (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), (230, 230, 230))
        fallback_bg.save(bg_img_path)

    # Insert Background Image
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Geometric Overlay via PIL ===
    # Create a transparent RGBA canvas
    overlay_img = Image.new('RGBA', (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay_img, 'RGBA')

    # Polygon 1: Base Dark Navy Sweep (Covers most of bottom left)
    # Coordinates: Top-Left edge down to Bottom-Right edge
    poly1 = [
        (0, SLIDE_HEIGHT_PX * 0.1), 
        (SLIDE_WIDTH_PX * 0.8, SLIDE_HEIGHT_PX), 
        (0, SLIDE_HEIGHT_PX)
    ]
    draw.polygon(poly1, fill=color_base_navy)

    # Polygon 2: Primary Accent (Teal) Sweep
    poly2 = [
        (0, SLIDE_HEIGHT_PX * 0.45), 
        (SLIDE_WIDTH_PX * 0.65, SLIDE_HEIGHT_PX), 
        (0, SLIDE_HEIGHT_PX)
    ]
    draw.polygon(poly2, fill=color_accent_primary)

    # Polygon 3: Secondary Accent (Light Blue) Triangle at bottom corner
    poly3 = [
        (0, SLIDE_HEIGHT_PX * 0.75), 
        (SLIDE_WIDTH_PX * 0.35, SLIDE_HEIGHT_PX), 
        (0, SLIDE_HEIGHT_PX)
    ]
    draw.polygon(poly3, fill=color_accent_light)

    # Save and insert the overlay mask
    overlay_path = "temp_overlay.png"
    overlay_img.save(overlay_path, format="PNG")
    slide.shapes.add_picture(overlay_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 3: Text & Content ===
    # Title Box (positioned over the thickest part of the dark overlay on the left)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(7.0), Inches(1.5))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    p_title = title_tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255) # White text
    p_title.font.name = "Arial"

    # Subtitle/Body Box
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(6.0), Inches(1.2))
    body_tf = body_box.text_frame
    body_tf.word_wrap = True
    p_body = body_tf.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(18)
    p_body.font.color.rgb = RGBColor(240, 240, 240) # Off-white for slight hierarchy
    p_body.font.name = "Arial"

    # Add a small decorative accent line above the title
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(0.85), Inches(3.7), Inches(0.8), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    line.line.color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])

    # Save presentation
    prs.save(output_pptx_path)

    # Cleanup temporary files
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
    if os.path.exists(overlay_path):
        os.remove(overlay_path)

    return output_pptx_path
