def create_slide(output_pptx_path: str, **kwargs) -> str:
    """
    Create a PPTX file reproducing the "Bilateral Gradient Comparison Bars" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from lxml import etree

    # --- Data Configuration ---
    slide_title = kwargs.get("title", "Product Comparison: Product A vs. Product B")
    
    # Palette
    bg_color = (30, 36, 48)            # Dark Charcoal
    center_pill_bg = (45, 55, 72)      # Lighter Slate
    
    # Left (Product A) Gradient: Light Green to Dark Green
    color_a_light = "68D391"
    color_a_dark = "276749"
    
    # Right (Product B) Gradient: Dark Blue to Light Blue
    color_b_dark = "2B6CB0"
    color_b_light = "63B3ED"

    # Comparison Data (Values 0.0 to 1.0)
    data = [
        {"metric": "Target Market Penetration", "A": 0.50, "B": 0.95},
        {"metric": "Market Share", "A": 0.75, "B": 0.65},
        {"metric": "Customer Acquisition Cost", "A": 1.00, "B": 0.50},
        {"metric": "Average Revenue Per User", "A": 0.40, "B": 0.50},
        {"metric": "Customer Lifetime Value", "A": 0.85, "B": 1.00},
    ]

    # --- Helper Functions (lxml XML Injection) ---
    def make_pill_shape(shape):
        """Forces a rounded rectangle to have maximum roundness (pill shape)."""
        prstGeom = shape.element.spPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
        if prstGeom is not None:
            avLst = prstGeom.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
            if avLst is None:
                avLst = etree.SubElement(prstGeom, '{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
            for gd in avLst.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}gd'):
                avLst.remove(gd)
            etree.SubElement(avLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd', name="adj", fmla="val 50000")

    def apply_gradient(shape, hex_start, hex_end, angle):
        """Applies a linear gradient fill to a shape via openXML."""
        spPr = shape.element.spPr
        for elem in spPr.xpath('.//a:solidFill', namespaces=spPr.nsmap):
            spPr.remove(elem)
        
        gradFill = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill', rotWithShape="1")
        gsLst = etree.SubElement(gradFill, '{http://schemas.openxmlformats.org/drawingml/2006/main}gsLst')
        
        # Start color (pos 0)
        gs1 = etree.SubElement(gsLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gs', pos="0")
        etree.SubElement(gs1, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val=hex_start)
        
        # End color (pos 100000)
        gs2 = etree.SubElement(gsLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gs', pos="100000")
        etree.SubElement(gs2, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val=hex_end)
        
        # Direction
        etree.SubElement(gradFill, '{http://schemas.openxmlformats.org/drawingml/2006/main}lin', ang=str(int(angle * 60000)), scaled="1")

    def apply_shadow(shape):
        """Applies a subtle drop shadow to elevate the shape."""
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad="50800", dist="38100", dir="5400000", algn="ctr", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="40000")

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Background Fill
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # 2. Main Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = slide_title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # 3. Headers (Product A & B)
    # Product A Badge & Text
    badge_a = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(1.2), Inches(0.6), Inches(0.6))
    badge_a.fill.solid()
    badge_a.fill.fore_color.rgb = RGBColor(104, 211, 145)  # Matches color_a_light
    badge_a.line.fill.background()
    badge_a_tf = badge_a.text_frame
    badge_a_tf.text = "A"
    badge_a_tf.paragraphs[0].font.bold = True
    badge_a_tf.paragraphs[0].font.size = Pt(24)
    badge_a_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    txt_a = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(2.8), Inches(0.8))
    tf_a = txt_a.text_frame
    p_a1 = tf_a.add_paragraph()
    p_a1.text = "Product A\n"
    p_a1.font.bold = True
    p_a1.font.size = Pt(20)
    p_a1.font.color.rgb = RGBColor(104, 211, 145)
    p_a1.alignment = PP_ALIGN.RIGHT
    p_a2 = tf_a.add_paragraph()
    p_a2.text = "Premium high-performance model."
    p_a2.font.size = Pt(12)
    p_a2.font.color.rgb = RGBColor(200, 200, 200)
    p_a2.alignment = PP_ALIGN.RIGHT

    # Product B Badge & Text
    badge_b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.2), Inches(1.2), Inches(0.6), Inches(0.6))
    badge_b.fill.solid()
    badge_b.fill.fore_color.rgb = RGBColor(99, 179, 237)  # Matches color_b_light
    badge_b.line.fill.background()
    badge_b_tf = badge_b.text_frame
    badge_b_tf.text = "B"
    badge_b_tf.paragraphs[0].font.bold = True
    badge_b_tf.paragraphs[0].font.size = Pt(24)
    badge_b_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    txt_b = slide.shapes.add_textbox(Inches(10.0), Inches(1.15), Inches(2.8), Inches(0.8))
    tf_b = txt_b.text_frame
    p_b1 = tf_b.add_paragraph()
    p_b1.text = "Product B\n"
    p_b1.font.bold = True
    p_b1.font.size = Pt(20)
    p_b1.font.color.rgb = RGBColor(99, 179, 237)
    p_b1.alignment = PP_ALIGN.LEFT
    p_b2 = tf_b.add_paragraph()
    p_b2.text = "Budget-friendly everyday model."
    p_b2.font.size = Pt(12)
    p_b2.font.color.rgb = RGBColor(200, 200, 200)
    p_b2.alignment = PP_ALIGN.LEFT

    # --- Layout Geometry Calculations ---
    y_start = Inches(2.5)
    y_gap = Inches(0.9)
    center_x = Inches(13.333 / 2)
    metric_w = Inches(2.8)
    metric_h = Inches(0.6)
    metric_x = center_x - (metric_w / 2)
    max_bar_w = Inches(4.0)
    tuck_in = Inches(0.4)  # How much the bar hides under the center pill

    # We must draw ALL Progress Bars FIRST so they sit behind the Center Metric Pills
    center_shapes_data = []

    for idx, item in enumerate(data):
        y = y_start + idx * y_gap
        
        # Calculate widths
        w_a = (max_bar_w * item["A"]) + tuck_in
        w_b = (max_bar_w * item["B"]) + tuck_in
        
        x_a = metric_x - w_a + tuck_in
        x_b = metric_x + metric_w - tuck_in

        # --- Draw Product A Bar (Left) ---
        bar_a = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_a, y, w_a, metric_h)
        bar_a.line.fill.background()
        make_pill_shape(bar_a)
        # Angle 0 is left-to-right. We want Light on Left, Dark on Right.
        apply_gradient(bar_a, color_a_light, color_a_dark, 0)
        
        tf_ba = bar_a.text_frame
        tf_ba.margin_left = Inches(0.2)
        p_ba = tf_ba.add_paragraph()
        p_ba.text = f"{int(item['A']*100)}%"
        p_ba.font.bold = True
        p_ba.font.size = Pt(16)
        p_ba.font.color.rgb = RGBColor(255, 255, 255)
        p_ba.alignment = PP_ALIGN.LEFT

        # --- Draw Product B Bar (Right) ---
        bar_b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_b, y, w_b, metric_h)
        bar_b.line.fill.background()
        make_pill_shape(bar_b)
        # Angle 0 is left-to-right. We want Dark on Left, Light on Right.
        apply_gradient(bar_b, color_b_dark, color_b_light, 0)
        
        tf_bb = bar_b.text_frame
        tf_bb.margin_right = Inches(0.2)
        p_bb = tf_bb.add_paragraph()
        p_bb.text = f"{int(item['B']*100)}%"
        p_bb.font.bold = True
        p_bb.font.size = Pt(16)
        p_bb.font.color.rgb = RGBColor(255, 255, 255)
        p_bb.alignment = PP_ALIGN.RIGHT
        
        # Save coordinates to draw center shapes next
        center_shapes_data.append((y, item["metric"]))

    # --- Draw Center Metric Pills (Drawn last to sit on top of the bars) ---
    for y, metric_text in center_shapes_data:
        center_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, metric_x, y, metric_w, metric_h)
        make_pill_shape(center_pill)
        
        # Styling
        center_pill.fill.solid()
        center_pill.fill.fore_color.rgb = RGBColor(*center_pill_bg)
        center_pill.line.fill.background()
        apply_shadow(center_pill)
        
        # Text
        tf_c = center_pill.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.add_paragraph()
        p_c.text = metric_text
        p_c.font.bold = True
        p_c.font.size = Pt(14)
        p_c.font.color.rgb = RGBColor(255, 255, 255)
        p_c.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
