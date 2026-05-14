def create_slide(
    output_pptx_path: str,
    title_text: str = "TITLE",
    option1_text: str = "OPTION 1",
    option2_text: str = "OPTION 2",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the '3D Floral Wreath S-Curve Morph' visual effect.
    Returns: path to the saved PPTX file.
    """
    import math
    import os
    from PIL import Image, ImageDraw, ImageFilter
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_CONNECTOR

    # ==========================================
    # 1. PIL Image Generation: The Leaf Wreath
    # ==========================================
    leaf_size = 400
    leaf_mask = Image.new('L', (leaf_size, leaf_size), 0)
    draw = ImageDraw.Draw(leaf_mask)

    # Draw a symmetrical plump leaf pointing right (0 degrees)
    pts = []
    # Upper curve
    for i in range(101):
        t = i / 100.0
        x = 100 + 200 * t
        y = 200 - 60 * math.sin(t * math.pi)
        pts.append((x, y))
    # Lower curve
    for i in range(100, -1, -1):
        t = i / 100.0
        x = 100 + 200 * t
        y = 200 + 60 * math.sin(t * math.pi)
        pts.append((x, y))
    draw.polygon(pts, fill=255)

    # Apply linear gradient to the leaf
    gradient_leaf = Image.new('RGBA', (leaf_size, leaf_size))
    for x in range(leaf_size):
        ratio = x / leaf_size
        r = int(0x06 + (0x00 - 0x06) * ratio)
        g = int(0x38 + (0x9A - 0x38) * ratio)
        b = int(0x08 + (0x16 - 0x08) * ratio)
        for y in range(leaf_size):
            gradient_leaf.putpixel((x, y), (r, g, b, 255))
    gradient_leaf.putalpha(leaf_mask)

    # Create drop shadow
    shadow_mask = leaf_mask.filter(ImageFilter.GaussianBlur(12))
    shadow = Image.new('RGBA', (leaf_size, leaf_size), (0, 0, 0, 0))
    shadow.putalpha(shadow_mask)
    # Reduce shadow opacity to 45%
    shadow_data = shadow.getdata()
    new_data = [(0, 0, 0, int(item[3] * 0.45)) for item in shadow_data]
    shadow.putdata(new_data)

    # Canvas for the S-curve
    canvas_size = 2000
    canvas = Image.new('RGBA', (canvas_size, canvas_size), (0, 0, 0, 0))

    # Calculate points for Top Loop (CCW: 40 deg to 230 deg)
    points_top = []
    for deg in range(40, 231, 12):
        rad = math.radians(deg)
        x = 1000 + 300 * math.cos(rad)
        y = 700 - 300 * math.sin(rad)
        points_top.append((x, y))

    # Calculate points for Bottom Loop (CCW: -140 deg to 130 deg)
    points_bottom = []
    for deg in range(-140, 131, 12):
        rad = math.radians(deg)
        x = 1000 + 300 * math.cos(rad)
        y = 1300 - 300 * math.sin(rad)
        points_bottom.append((x, y))

    def calculate_tangents(points):
        leaves = []
        for i in range(len(points)):
            x, y = points[i]
            if i < len(points) - 1:
                nx, ny = points[i+1]
                angle = math.degrees(math.atan2(ny - y, nx - x))
            else:
                px, py = points[i-1]
                angle = math.degrees(math.atan2(y - py, x - px))
            # PIL rotation is counter-clockwise, screen Y is down -> negate angle
            leaves.append((x, y, -angle))
        return leaves

    all_leaves = calculate_tangents(points_top) + calculate_tangents(points_bottom)

    # Paste leaves sequentially to create overlapping effect
    for x, y, angle in all_leaves:
        rot_leaf = gradient_leaf.rotate(angle, resample=Image.BICUBIC, expand=True)
        rot_shadow = shadow.rotate(angle, resample=Image.BICUBIC, expand=True)
        
        # Center coordinates
        lx = int(x - rot_leaf.width / 2)
        ly = int(y - rot_leaf.height / 2)
        
        # Shadow offset (+12, +12)
        sx = int(x - rot_shadow.width / 2 + 12)
        sy = int(y - rot_shadow.height / 2 + 12)
        
        canvas.alpha_composite(rot_shadow, (sx, sy))
        canvas.alpha_composite(rot_leaf, (lx, ly))

    # Crop blank space
    bbox = canvas.getbbox()
    canvas = canvas.crop(bbox)
    
    img_path = "s_wreath_temp.png"
    canvas.save(img_path)

    # ==========================================
    # 2. PowerPoint Construction
    # ==========================================
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background Color (Soft Off-White)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 248, 248)

    # Insert Graphic
    pic = slide.shapes.add_picture(img_path, 0, 0)
    pic.height = Inches(6.5)
    pic.left = (prs.slide_width - pic.width) / 2
    pic.top = (prs.slide_height - pic.height) / 2

    # Clean up temp file
    if os.path.exists(img_path):
        os.remove(img_path)

    # Center Title in the Gap
    tb_title = slide.shapes.add_textbox(Inches(5.66), Inches(3.5), Inches(2), Inches(0.5))
    tb_title.text_frame.word_wrap = True
    p = tb_title.text_frame.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(50, 50, 50)

    # Option 1 Text (Left)
    tb_opt1 = slide.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(3), Inches(1.5))
    p1 = tb_opt1.text_frame.paragraphs[0]
    p1.text = option1_text
    p1.alignment = PP_ALIGN.RIGHT
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(50, 50, 50)
    
    p1_desc = tb_opt1.text_frame.add_paragraph()
    p1_desc.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa."
    p1_desc.alignment = PP_ALIGN.RIGHT
    p1_desc.font.size = Pt(10)
    p1_desc.font.bold = False
    p1_desc.font.color.rgb = RGBColor(120, 120, 120)

    # Option 2 Text (Right)
    tb_opt2 = slide.shapes.add_textbox(Inches(9.3), Inches(3.2), Inches(3), Inches(1.5))
    p2 = tb_opt2.text_frame.paragraphs[0]
    p2.text = option2_text
    p2.alignment = PP_ALIGN.LEFT
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(50, 50, 50)
    
    p2_desc = tb_opt2.text_frame.add_paragraph()
    p2_desc.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa."
    p2_desc.alignment = PP_ALIGN.LEFT
    p2_desc.font.size = Pt(10)
    p2_desc.font.bold = False
    p2_desc.font.color.rgb = RGBColor(120, 120, 120)

    # Connecting Lines
    line1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.2), Inches(3.8), Inches(4.8), Inches(3.8))
    line1.line.color.rgb = RGBColor(180, 180, 180)
    
    line2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.5), Inches(3.8), Inches(9.1), Inches(3.8))
    line2.line.color.rgb = RGBColor(180, 180, 180)

    prs.save(output_pptx_path)
    return output_pptx_path
