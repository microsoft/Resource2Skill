def create_slide(
    output_pptx_path: str,
    title_text: str = "Quarterly Sales Summary",
    subtitle_text: str = "Performance metrics and key highlights across all regions",
    accent_color: tuple = (0, 176, 240),   # Corporate Teal
    secondary_color: tuple = (31, 73, 125), # Corporate Navy
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Corporate Flat-Design KPI Dashboard Layout' effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide layout

    # Helper function to inject premium drop shadow via lxml
    def apply_shadow(shape):
        spPr = shape.element.spPr
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        effectLst = etree.SubElement(spPr, f'{{{a}}}effectLst')
        # 80000 = ~6pt blur, 40000 = ~3pt distance, 5400000 = 90 degrees (straight down)
        outerShdw = etree.SubElement(effectLst, f'{{{a}}}outerShdw', 
                                     blurRad="80000", dist="40000", dir="5400000", algn="tl", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, f'{{{a}}}srgbClr', val="000000")
        etree.SubElement(srgbClr, f'{{{a}}}alpha', val="12000") # 12% opacity

    # Set background to very light gray for contrast against white cards
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(250, 250, 250)

    # Add Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.6), Inches(10), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.name = "Segoe UI"
    p.font.bold = True
    p.font.color.rgb = RGBColor(*secondary_color)

    # Add Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(10), Inches(0.5))
    tf_sub = subtitle_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(18)
    p_sub.font.name = "Segoe UI"
    p_sub.font.color.rgb = RGBColor(120, 120, 120)

    # Add Header Dividing Line
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1.9), Inches(12.333), Inches(1.9)
    )
    line.line.color.rgb = RGBColor(*accent_color)
    line.line.width = Pt(2)

    # Card Content Data
    metrics = ["$2.3M", "$8.8M", "$8.4M", "$9.2M"]
    quarters = ["Q1", "Q2", "Q3", "Q4"]
    descriptions = [
        "Initial growth phase with new product line rollouts.",
        "Strong performance driven by targeted marketing campaigns.",
        "Steady revenue stream from recurring enterprise clients.",
        "Aggressive year-end push resulting in record-breaking sales."
    ]

    # Grid Math Calculation
    margin = 1.0
    usable_width = 13.333 - (margin * 2)
    card_width_in = 2.5
    num_cards = 4
    total_gaps = num_cards - 1
    gap_in = (usable_width - (card_width_in * num_cards)) / total_gaps
    
    card_height_in = 4.2
    card_y_in = 2.5

    for i in range(num_cards):
        card_x_in = margin + i * (card_width_in + gap_in)

        # 1. Main Card Background
        card_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(card_x_in), Inches(card_y_in), Inches(card_width_in), Inches(card_height_in)
        )
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = RGBColor(255, 255, 255) # Pure White
        card_bg.line.fill.background() # No border
        apply_shadow(card_bg)

        # 2. Top Accent Bar
        bar_height_in = 0.15
        accent_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(card_x_in), Inches(card_y_in), Inches(card_width_in), Inches(bar_height_in)
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = RGBColor(*accent_color)
        accent_bar.line.fill.background()

        # 3. Label text (e.g., Q1)
        q_box = slide.shapes.add_textbox(Inches(card_x_in), Inches(card_y_in + 0.4), Inches(card_width_in), Inches(0.6))
        q_tf = q_box.text_frame
        q_p = q_tf.paragraphs[0]
        q_p.text = quarters[i]
        q_p.alignment = PP_ALIGN.CENTER
        q_p.font.size = Pt(20)
        q_p.font.name = "Segoe UI"
        q_p.font.bold = True
        q_p.font.color.rgb = RGBColor(160, 160, 160)

        # 4. Main Metric text (e.g., $2.3M)
        m_box = slide.shapes.add_textbox(Inches(card_x_in), Inches(card_y_in + 0.9), Inches(card_width_in), Inches(0.8))
        m_tf = m_box.text_frame
        m_p = m_tf.paragraphs[0]
        m_p.text = metrics[i]
        m_p.alignment = PP_ALIGN.CENTER
        m_p.font.size = Pt(40)
        m_p.font.name = "Segoe UI"
        m_p.font.bold = True
        m_p.font.color.rgb = RGBColor(*secondary_color)
        
        # 5. Inner Divider Line
        div_x_in = card_x_in + 0.5
        div_y_in = card_y_in + 1.9
        div_width_in = card_width_in - 1.0
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(div_x_in), Inches(div_y_in), Inches(div_width_in), Pt(2))
        div.fill.solid()
        div.fill.fore_color.rgb = RGBColor(*accent_color)
        div.line.fill.background()

        # 6. Description Text
        d_box = slide.shapes.add_textbox(Inches(card_x_in + 0.15), Inches(card_y_in + 2.2), Inches(card_width_in - 0.3), Inches(1.5))
        d_tf = d_box.text_frame
        d_tf.word_wrap = True
        d_p = d_tf.paragraphs[0]
        d_p.text = descriptions[i]
        d_p.alignment = PP_ALIGN.CENTER
        d_p.font.size = Pt(13)
        d_p.font.name = "Segoe UI"
        d_p.font.color.rgb = RGBColor(100, 100, 100)
        
        # 7. Directional Flow Arrow (between cards)
        if i < num_cards - 1:
            arrow_x_in = card_x_in + card_width_in + (gap_in / 2) - 0.1
            arrow_y_in = card_y_in + (card_height_in / 2) - 0.15
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(arrow_x_in), Inches(arrow_y_in), Inches(0.2), Inches(0.3)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(210, 210, 210)
            arrow.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
