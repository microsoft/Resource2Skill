def create_slide(
    output_pptx_path: str,
    presentation_title: str = "FASHION X AI",
    presentation_subtitle: str = "Visionary Female Fashion Prompts & Professional Reviews",
    concepts: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Editorial Fashion Split-Screen Lookbook' style.
    Generates a Dark Title slide, followed by Split-Screen Content slides.
    """
    import os
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # Default data if none provided
    if not concepts:
        concepts = [
            {
                "id": "01",
                "title": "AVANT-GARDE MINIMALISM",
                "image_keyword": "minimalist+fashion+female",
                "elements": "Structural oversized blazer, architectural trousers, monochromatic palette.",
                "fabric": "Heavyweight silk-wool blend, matte finish, seamless construction.",
                "details": "Hidden plackets, sharp lapels, asymmetric hemlines.",
                "review": "This concept masterfully redefines elegance through subtractive design, where the silhouette itself becomes the primary narrative."
            }
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- Helper Function: Download & Crop Image to Exact 8:9 Ratio (Half Slide) ---
    def get_half_slide_image(keyword):
        # Target size for half of a 13.333 x 7.5 inch slide at 150 DPI
        target_w, target_h = int((13.333 / 2) * 150), int(7.5 * 150) 
        
        try:
            url = f"https://source.unsplash.com/featured/?{keyword}"
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as response:
                img = Image.open(io.BytesIO(response.read())).convert("RGB")
                
            # Crop to aspect ratio
            img_ratio = img.width / img.height
            target_ratio = target_w / target_h
            
            if img_ratio > target_ratio:
                # Image is wider than needed, crop sides
                new_w = int(target_ratio * img.height)
                left = (img.width - new_w) / 2
                img = img.crop((left, 0, left + new_w, img.height))
            else:
                # Image is taller than needed, crop top/bottom
                new_h = int(img.width / target_ratio)
                top = (img.height - new_h) / 2
                img = img.crop((0, top, img.width, top + new_h))
                
            img = img.resize((target_w, target_h), Image.Resampling.LANCZOS)
        except Exception as e:
            # Fallback: Create a gradient dummy image if network fails
            img = Image.new('RGB', (target_w, target_h))
            draw = ImageDraw.Draw(img)
            for y in range(target_h):
                r = int(20 + (100 * y / target_h))
                g = int(20 + (100 * y / target_h))
                b = int(30 + (120 * y / target_h))
                draw.line([(0, y), (target_w, y)], fill=(r, g, b))
                
        img_path = f"temp_split_{keyword.replace('+', '_')}.jpg"
        img.save(img_path, quality=90)
        return img_path

    # ==========================================================
    # SLIDE 1: DARK MODE TITLE SLIDE
    # ==========================================================
    slide_title = prs.slides.add_slide(blank_layout)
    
    # Background
    bg = slide_title.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height) # shape 1 is rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(20, 20, 20)
    bg.line.fill.background()

    # Title Text
    tb_title = slide_title.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.333), Inches(1.5))
    tf_title = tb_title.text_frame
    p = tf_title.add_paragraph()
    p.text = presentation_title.upper()
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = RGBColor(255, 204, 51) # Gold
    p.alignment = PP_ALIGN.CENTER

    # Subtitle Text
    p2 = tf_title.add_paragraph()
    p2.text = presentation_subtitle
    p2.font.size = Pt(20)
    p2.font.name = "Arial"
    p2.font.color.rgb = RGBColor(220, 220, 220)
    p2.alignment = PP_ALIGN.CENTER

    # ==========================================================
    # SLIDE 2+: SPLIT SCREEN CONTENT SLIDES
    # ==========================================================
    for concept in concepts:
        slide = prs.slides.add_slide(blank_layout)
        
        # Left Background (Off-white)
        left_bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width/2, prs.slide_height)
        left_bg.fill.solid()
        left_bg.fill.fore_color.rgb = RGBColor(248, 248, 248)
        left_bg.line.fill.background()

        # Right Background (Full Bleed Image)
        img_path = get_half_slide_image(concept.get("image_keyword", "fashion"))
        slide.shapes.add_picture(img_path, prs.slide_width/2, 0, prs.slide_width/2, prs.slide_height)
        os.remove(img_path) # Cleanup

        # --- Left Pane Content ---
        margin_left = Inches(0.8)
        content_width = Inches(5.0)

        # Concept Number
        tb_num = slide.shapes.add_textbox(margin_left, Inches(0.8), content_width, Inches(0.5))
        p_num = tb_num.text_frame.paragraphs[0]
        p_num.text = f"CONCEPT {concept.get('id', '01')}"
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = RGBColor(100, 100, 100)

        # Main Title
        tb_head = slide.shapes.add_textbox(margin_left, Inches(1.1), content_width, Inches(1.0))
        p_head = tb_head.text_frame.paragraphs[0]
        p_head.text = concept.get("title", "UNTITLED")
        p_head.font.size = Pt(28)
        p_head.font.bold = True
        p_head.font.color.rgb = RGBColor(20, 20, 20)

        # The Styled Data Table
        table_top = Inches(2.2)
        shape_table = slide.shapes.add_table(rows=3, cols=2, left=margin_left, top=table_top, width=content_width, height=Inches(2.5))
        tbl = shape_table.table
        
        # Adjust column widths
        tbl.columns[0].width = Inches(1.5)
        tbl.columns[1].width = Inches(3.5)

        row_data = [
            ("Key Elements", concept.get("elements", ""), RGBColor(30, 45, 60)),     # Navy
            ("Fabric & Patterns", concept.get("fabric", ""), RGBColor(100, 120, 140)), # Slate
            ("Garment Details", concept.get("details", ""), RGBColor(200, 200, 200))   # Silver
        ]

        for row_idx, (label, text, bg_color) in enumerate(row_data):
            # Left Cell (Label)
            cell_left = tbl.cell(row_idx, 0)
            cell_left.fill.solid()
            cell_left.fill.fore_color.rgb = bg_color
            p_left = cell_left.text_frame.paragraphs[0]
            p_left.text = label
            p_left.font.size = Pt(11)
            p_left.font.bold = True
            p_left.font.color.rgb = RGBColor(255, 255, 255) if row_idx < 2 else RGBColor(40, 40, 40)
            cell_left.margin_top = Pt(12)

            # Right Cell (Data)
            cell_right = tbl.cell(row_idx, 1)
            cell_right.fill.solid()
            cell_right.fill.fore_color.rgb = RGBColor(240, 240, 240)
            p_right = cell_right.text_frame.paragraphs[0]
            p_right.text = text
            p_right.font.size = Pt(10)
            p_right.font.color.rgb = RGBColor(60, 60, 60)
            cell_right.margin_top = Pt(12)
            cell_right.margin_left = Pt(12)

        # Professional Review Block
        tb_rev = slide.shapes.add_textbox(margin_left, Inches(5.2), content_width, Inches(1.5))
        tb_rev.text_frame.word_wrap = True
        
        p_rev_head = tb_rev.text_frame.paragraphs[0]
        p_rev_head.text = "Professional Review: "
        p_rev_head.font.size = Pt(11)
        p_rev_head.font.bold = True
        p_rev_head.font.italic = True
        p_rev_head.font.color.rgb = RGBColor(20, 20, 20)

        p_rev_body = tb_rev.text_frame.add_paragraph()
        p_rev_body.text = f"\"{concept.get('review', '')}\""
        p_rev_body.font.size = Pt(11)
        p_rev_body.font.italic = True
        p_rev_body.font.color.rgb = RGBColor(80, 80, 80)

    prs.save(output_pptx_path)
    return output_pptx_path
