def create_slide(
    output_pptx_path: str,
    title_text: str = "Unified Design Strategy",
    body_text: str = "Applying Thematic Motifs Across Multiple Slides",
    bg_palette: str = "dark", 
    accent_color: tuple = (42, 186, 171),  # Bright Teal
    **kwargs,
) -> str:
    """
    Creates a 4-slide PPTX demonstrating the "Thematic Motif Unification & Variation" skill.
    Slide 1-3: Motif Cropping/Fragmentation for visual coherence.
    Slide 4: The 5 Variation Rules (Cropping, Symmetry, Layering, Repetition, Color).
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    bg_color = RGBColor(15, 15, 15) # Dark charcoal background
    motif_color = RGBColor(*accent_color)
    
    # Helper to set background color
    def set_solid_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper to add transparent motif (requires lxml)
    def add_transparent_motif(slide, shape_type, left, top, width, height, color, alpha=100000):
        """Adds a shape and applies transparency via XML (alpha: 0 to 100000)"""
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background() # No line
        
        # Inject Alpha via lxml
        fill_element = shape.fill._xPr.solidFill
        srgbClr = fill_element.find('.//a:srgbClr', namespaces=fill_element.nsmap)
        if srgbClr is not None:
            alpha_elem = etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
            alpha_elem.set('val', str(int(alpha))) # 100000 is 100% opaque, 50000 is 50%
        return shape

    # Helper to add standard text
    def add_slide_text(slide, title, subtitle, top_offset=Inches(3)):
        # Title
        tb1 = slide.shapes.add_textbox(Inches(1), top_offset, Inches(6), Inches(1))
        p1 = tb1.text_frame.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(44)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(255, 255, 255)
        
        # Subtitle
        tb2 = slide.shapes.add_textbox(Inches(1), top_offset + Inches(1), Inches(6), Inches(0.5))
        p2 = tb2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = RGBColor(180, 180, 180)

    # ==========================================
    # SLIDE 1: Right-Edge Mega Fragment
    # ==========================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_solid_bg(slide1, bg_color)
    
    # Motif: Large circle off-screen to the right
    add_transparent_motif(slide1, MSO_SHAPE.OVAL, Inches(10), Inches(-1), Inches(8), Inches(8), motif_color)
    add_slide_text(slide1, "01. Right Anchoring", "Massive off-screen shape creates dynamic tension.")

    # ==========================================
    # SLIDE 2: Top-Left Header Fragment
    # ==========================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_solid_bg(slide2, bg_color)
    
    # Motif: Smaller circle off-screen to the top left
    add_transparent_motif(slide2, MSO_SHAPE.OVAL, Inches(-1.5), Inches(-1.5), Inches(4), Inches(4), motif_color)
    add_slide_text(slide2, "02. Corner Branding", "The same motif scaled down, anchoring the layout.", top_offset=Inches(3.5))

    # ==========================================
    # SLIDE 3: Top-Center Symmetrical Drop
    # ==========================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_solid_bg(slide3, bg_color)
    
    # Motif: Circle emerging from the top center
    add_transparent_motif(slide3, MSO_SHAPE.OVAL, Inches(5.166), Inches(-2), Inches(3), Inches(3), motif_color)
    
    # Center-aligned text for this slide
    tb3 = slide3.shapes.add_textbox(Inches(3.66), Inches(3), Inches(6), Inches(2))
    p1 = tb3.text_frame.paragraphs[0]
    p1.text = "03. Symmetrical Balance"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.alignment = 2 # Center
    
    p2 = tb3.text_frame.add_paragraph()
    p2.text = "Using the same anchor to establish a center axis."
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(180, 180, 180)
    p2.alignment = 2

    # ==========================================
    # SLIDE 4: The 5 Variation Techniques
    # ==========================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_solid_bg(slide4, bg_color)
    add_slide_text(slide4, "04. Element Variations", "Same shape, different treatments (Cropping, Symmetry, Layering, Repetition, Color)", top_offset=Inches(0.5))

    base_y = Inches(2.5)
    
    # 1. Cropping (遮挡)
    add_transparent_motif(slide4, MSO_SHAPE.RECTANGLE, Inches(1), base_y, Inches(1.5), Inches(1.5), motif_color)
    # Block it to simulate cropping
    mask = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), base_y + Inches(0.75), Inches(1.5), Inches(0.75))
    mask.fill.solid()
    mask.fill.fore_color.rgb = bg_color
    mask.line.fill.background()

    # 2. Symmetry (对称)
    add_transparent_motif(slide4, MSO_SHAPE.RIGHT_TRIANGLE, Inches(3.5), base_y, Inches(0.7), Inches(1.5), motif_color)
    tri2 = add_transparent_motif(slide4, MSO_SHAPE.RIGHT_TRIANGLE, Inches(4.3), base_y, Inches(0.7), Inches(1.5), motif_color)
    tri2.rotation = 180 # Flip to create symmetry

    # 3. Layering (层叠) - Using lxml transparency helper
    add_transparent_motif(slide4, MSO_SHAPE.USER, Inches(6.5), base_y, Inches(1.2), Inches(1.5), motif_color, alpha=100000)
    add_transparent_motif(slide4, MSO_SHAPE.USER, Inches(6.2), base_y, Inches(1.2), Inches(1.5), motif_color, alpha=50000)
    add_transparent_motif(slide4, MSO_SHAPE.USER, Inches(5.9), base_y, Inches(1.2), Inches(1.5), motif_color, alpha=20000)

    # 4 & 5. Repetition & Color (重复 & 变色)
    grid_start_x = Inches(8.5)
    for row in range(3):
        for col in range(5):
            x = grid_start_x + (col * Inches(0.7))
            y = base_y + (row * Inches(0.7))
            
            # Apply color variation to one specific item
            if row == 1 and col == 2:
                c = motif_color
                alpha = 100000
            else:
                c = RGBColor(50, 50, 50) # Inactive pattern color
                alpha = 100000
                
            add_transparent_motif(slide4, MSO_SHAPE.USER, x, y, Inches(0.5), Inches(0.5), c, alpha)

    # Add Labels for the variations
    labels = ["Cropping", "Symmetry", "Layering", "Repetition & Color"]
    x_positions = [1, 3.5, 6, 8.5]
    for lbl, x_pos in zip(labels, x_positions):
        tb = slide4.shapes.add_textbox(Inches(x_pos), base_y + Inches(2), Inches(2), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = lbl
        p.font.size = Pt(14)
        p.font.color.rgb = motif_color

    prs.save(output_pptx_path)
    return output_pptx_path
