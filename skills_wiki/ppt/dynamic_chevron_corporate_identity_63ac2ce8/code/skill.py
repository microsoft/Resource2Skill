def create_slide(
    output_pptx_path: str,
    title_text: str = "Customer Service Review",
    body_text: str = "Annual Performance & Strategic Roadmap",
    bg_palette: str = "office,corporate",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Dynamic Chevron Corporate Identity' visual effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageOps
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    # === Presentation Setup ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Color Palette ===
    COLOR_TEAL = RGBColor(73, 142, 141)
    COLOR_GOLD = RGBColor(186, 163, 110)
    COLOR_CORAL = RGBColor(193, 125, 125)
    COLOR_SAGE = RGBColor(142, 186, 141)

    # === Geometric Configuration (Inches) ===
    PATTERN_H = 6.0
    IMG_W = 5.0
    POINT_EXT = 2.0
    BAND_W = 1.2
    GAP = 0.2

    # === Layer 1: Base Image with PIL Polygon Mask ===
    # Convert dimensions to pixels (using 300 DPI for crispness)
    DPI = 300
    PIL_W = int((IMG_W + POINT_EXT) * DPI)
    PIL_H = int(PATTERN_H * DPI)

    # Create the transparent polygon mask
    mask = Image.new("L", (PIL_W, PIL_H), 0)
    draw = ImageDraw.Draw(mask)
    # Points: Top-Left, Top-Right, Right-Point, Bottom-Right, Bottom-Left
    poly_points = [
        (0, 0),
        (int(IMG_W * DPI), 0),
        (int((IMG_W + POINT_EXT) * DPI), int((PATTERN_H / 2) * DPI)),
        (int(IMG_W * DPI), int(PATTERN_H * DPI)),
        (0, int(PATTERN_H * DPI))
    ]
    draw.polygon(poly_points, fill=255)

    # Fetch and format the image
    try:
        # High quality corporate image URL
        img_url = "https://images.unsplash.com/photo-1497366216548-37526070297c?auto=format&fit=crop&w=2100&q=80"
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGBA")
        
        # Handle Pillow version differences for Lanczos
        try:
            resample_filter = Image.Resampling.LANCZOS
        except AttributeError:
            resample_filter = Image.LANCZOS
            
        base_img = ImageOps.fit(base_img, (PIL_W, PIL_H), resample_filter)
    except Exception as e:
        print(f"Image download failed: {e}. Using geometric fallback.")
        # Fallback if download fails
        base_img = Image.new("RGBA", (PIL_W, PIL_H), (40, 50, 60, 255))
        draw_fallback = ImageDraw.Draw(base_img)
        draw_fallback.line([(0,0), (PIL_W, PIL_H)], fill=(60, 70, 80, 255), width=30)

    # Apply mask and save temporarily
    base_img.putalpha(mask)
    tmp_image_path = "tmp_masked_hero.png"
    base_img.save(tmp_image_path)

    # Insert Image into Slide
    slide.shapes.add_picture(tmp_image_path, 0, 0, Inches(IMG_W + POINT_EXT), Inches(PATTERN_H))

    # === Layer 2: Vector Chevron Bands (python-pptx Freeform) ===
    def add_chevron(slide, xs, ys, w, h, p, rgb_color):
        """Helper to draw parallel chevron shapes native to PowerPoint."""
        # Calculate vertices for the chevron band
        pts = [
            (xs + w, ys),                    # Top-Right
            (xs + w + p, ys + h/2),          # Outer-Point
            (xs + w, ys + h),                # Bottom-Right
            (xs, ys + h),                    # Bottom-Left
            (xs + p, ys + h/2)               # Inner-Point
        ]
        
        # Build freeform starting at Top-Left (xs, ys)
        builder = slide.shapes.build_freeform(int(xs), int(ys))
        builder.add_line_segments([(int(x), int(y)) for x, y in pts], close=True)
        shape = builder.convert_to_shape()
        
        # Apply solid color formatting
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb_color
        shape.line.fill.background()  # Hide outline natively

    # Calculate starting positions and render bands
    x_pos1 = Inches(IMG_W + GAP)
    add_chevron(slide, x_pos1, 0, Inches(BAND_W), Inches(PATTERN_H), Inches(POINT_EXT), COLOR_GOLD)

    x_pos2 = x_pos1 + Inches(BAND_W + GAP)
    add_chevron(slide, x_pos2, 0, Inches(BAND_W), Inches(PATTERN_H), Inches(POINT_EXT), COLOR_CORAL)

    x_pos3 = x_pos2 + Inches(BAND_W + GAP)
    add_chevron(slide, x_pos3, 0, Inches(BAND_W), Inches(PATTERN_H), Inches(POINT_EXT), COLOR_SAGE)


    # === Layer 3: Typographic Foundation ===
    # Horizontal Anchor Line separating visuals from text
    h_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(PATTERN_H), Inches(13.333), Inches(0.08))
    h_line.fill.solid()
    h_line.fill.fore_color.rgb = COLOR_TEAL
    h_line.line.fill.background()

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.0), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.clear()
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text.upper()
    p_title.font.name = "Arial"
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEAL

    # Subtitle
    if body_text:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.0), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.clear()
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = body_text
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(20)
        p_sub.font.color.rgb = RGBColor(120, 120, 120)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(tmp_image_path):
        os.remove(tmp_image_path)
        
    return output_pptx_path
