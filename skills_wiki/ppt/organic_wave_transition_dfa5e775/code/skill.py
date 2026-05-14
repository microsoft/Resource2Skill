def create_slide(
    output_pptx_path: str,
    title_text: str = "ARISTOTLE'S MODEL\nOF PERSUASION",
    subtitle_text: str = "PART #1",
    bg_top_color: tuple = (12, 16, 42),
    bg_bottom_color: tuple = (25, 30, 80),
    wave_back_color: tuple = (25, 28, 77),
    wave_mid_color: tuple = (70, 150, 130),
    wave_front_color: tuple = (100, 210, 180),
    text_color: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Organic Wave Transition' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import math
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background & Organic Waves (via PIL) ===
    # Draw at 2x resolution for perfect anti-aliasing
    w, h = 3840, 2160
    img = Image.new("RGBA", (w, h))
    draw = ImageDraw.Draw(img)

    # 1. Background Gradient
    for y in range(h):
        blend = y / h
        r = int(bg_top_color[0] * (1 - blend) + bg_bottom_color[0] * blend)
        g = int(bg_top_color[1] * (1 - blend) + bg_bottom_color[1] * blend)
        b = int(bg_top_color[2] * (1 - blend) + bg_bottom_color[2] * blend)
        draw.line([(0, y), (w, y)], fill=(r, g, b, 255))

    # 2. Wave 1 (Dark Blue Backing - Peaks on the right)
    points1 = [(0, h)]
    for x in range(w + 1):
        y1 = 1200 - 600 * math.sin((x / 3840) * 2 * math.pi + 1.08 * math.pi)
        points1.append((x, y1))
    points1.append((w, h))
    draw.polygon(points1, fill=wave_back_color + (255,))

    # 3. Wave 2 (Dark Teal Middle - Subtle depth layer)
    points_mid = [(0, h)]
    for x in range(w + 1):
        y_mid = 1350 - 450 * math.sin((x / 3840) * 2 * math.pi + 0.08 * math.pi)
        points_mid.append((x, y_mid))
    points_mid.append((w, h))
    draw.polygon(points_mid, fill=wave_mid_color + (255,))

    # 4. Wave 3 (Vibrant Teal Front - High hill on left, sweeping down)
    points2 = [(0, h)]
    for x in range(w + 1):
        y2 = 1300 - 500 * math.sin((x / 3840) * 2 * math.pi - 0.02 * math.pi)
        points2.append((x, y2))
    points2.append((w, h))
    draw.polygon(points2, fill=wave_front_color + (255,))

    # Resize image down for antialiasing and save to memory stream
    img_resized = img.resize((1920, 1080), Image.Resampling.LANCZOS)
    image_stream = io.BytesIO()
    img_resized.save(image_stream, format='PNG')
    image_stream.seek(0)

    # Insert background image
    slide.shapes.add_picture(image_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Text & Typography ===
    
    # Subtitle (Eyebrow text)
    txBox_sub = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11.33), Inches(0.8))
    tf_sub = txBox_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.alignment = PP_ALIGN.CENTER
    font_sub = p_sub.font
    font_sub.name = 'Arial'
    font_sub.size = Pt(28)
    font_sub.bold = True
    font_sub.color.rgb = RGBColor(*text_color)

    # Main Title
    txBox_title = slide.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.33), Inches(2.5))
    tf_title = txBox_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.alignment = PP_ALIGN.CENTER
    font_title = p_title.font
    font_title.name = 'Arial'  # Clean sans-serif
    font_title.size = Pt(54)
    font_title.bold = True
    font_title.color.rgb = RGBColor(*text_color)

    prs.save(output_pptx_path)
    return output_pptx_path
