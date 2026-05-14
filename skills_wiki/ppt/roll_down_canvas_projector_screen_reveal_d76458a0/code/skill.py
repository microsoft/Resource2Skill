import io
from typing import Tuple
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw

def _create_metallic_cylinder(width_px: int, height_px: int) -> io.BytesIO:
    """
    Creates a PIL image of a horizontal cylinder with a 3D metallic gradient.
    """
    img = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # 3-stop gradient for a horizontal cylinder (light hits the top-middle)
    color_top = (70, 70, 75, 255)
    color_mid = (230, 230, 235, 255)
    color_bot = (40, 40, 45, 255)
    
    mid_point = int(height_px * 0.4) # Highlight is slightly above center
    
    for y in range(height_px):
        if y <= mid_point:
            # Interpolate top to mid
            ratio = y / mid_point
            r = int(color_top[0] + (color_mid[0] - color_top[0]) * ratio)
            g = int(color_top[1] + (color_mid[1] - color_top[1]) * ratio)
            b = int(color_top[2] + (color_mid[2] - color_top[2]) * ratio)
        else:
            # Interpolate mid to bottom
            ratio = (y - mid_point) / (height_px - mid_point)
            r = int(color_mid[0] + (color_bot[0] - color_mid[0]) * ratio)
            g = int(color_mid[1] + (color_bot[1] - color_mid[1]) * ratio)
            b = int(color_mid[2] + (color_bot[2] - color_mid[2]) * ratio)
            
        draw.line([(0, y), (width_px, y)], fill=(r, g, b, 255))
        
    img_stream = io.BytesIO()
    img.save(img_stream, format='PNG')
    img_stream.seek(0)
    return img_stream

def _create_pull_handle(width_px: int, height_px: int) -> io.BytesIO:
    """
    Creates a PIL image of a pull string and ring.
    """
    img = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    color = (40, 40, 45, 255)
    center_x = width_px // 2
    
    # Dimensions
    ring_radius = int(width_px * 0.4)
    ring_thickness = int(width_px * 0.15)
    string_width = int(width_px * 0.15)
    
    string_bottom = height_px - (ring_radius * 2)
    
    # Draw string (vertical line)
    draw.rectangle(
        [center_x - string_width//2, 0, center_x + string_width//2, string_bottom],
        fill=color
    )
    
    # Draw ring (donut shape)
    bbox_outer = [center_x - ring_radius, string_bottom, center_x + ring_radius, string_bottom + ring_radius * 2]
    bbox_inner = [
        bbox_outer[0] + ring_thickness, bbox_outer[1] + ring_thickness,
        bbox_outer[2] - ring_thickness, bbox_outer[3] - ring_thickness
    ]
    
    draw.ellipse(bbox_outer, fill=color)
    draw.ellipse(bbox_inner, fill=(0, 0, 0, 0)) # Hollow out the center
    
    img_stream = io.BytesIO()
    img.save(img_stream, format='PNG')
    img_stream.seek(0)
    return img_stream

def create_slide(
    output_pptx_path: str,
    title_text: str = "Key Principle",
    body_text: str = '"Design is not just what it looks like and feels like. Design is how it works."\n— Steve Jobs',
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Roll-Down Canvas Reveal visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Background - clean white/light grey
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Add Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(50, 50, 60)
    
    # ==========================================
    # MEASUREMENTS & LAYOUT
    # ==========================================
    canvas_w = 8.5
    canvas_h = 4.0
    canvas_l = (13.333 - canvas_w) / 2
    canvas_t = 2.0
    
    bar_w = canvas_w + 0.4 # Overhangs 0.2 on each side
    bar_h = 0.35
    bar_l = canvas_l - 0.2
    
    # ==========================================
    # LAYER 1: THE CANVAS
    # ==========================================
    canvas = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(canvas_l), Inches(canvas_t), 
        Inches(canvas_w), Inches(canvas_h)
    )
    canvas.fill.solid()
    canvas.fill.fore_color.rgb = RGBColor(240, 240, 245) # Light silver
    canvas.line.color.rgb = RGBColor(200, 200, 205) # Subtle border
    canvas.line.width = Pt(1)

    # ==========================================
    # LAYER 2: THE TEXT ON CANVAS
    # ==========================================
    text_box = slide.shapes.add_textbox(
        Inches(canvas_l + 0.5), Inches(canvas_t + 0.5), 
        Inches(canvas_w - 1.0), Inches(canvas_h - 1.0)
    )
    tf_canvas = text_box.text_frame
    tf_canvas.word_wrap = True
    tf_canvas.text = body_text
    p_canvas = tf_canvas.paragraphs[0]
    p_canvas.font.size = Pt(28)
    p_canvas.font.italic = True
    p_canvas.font.name = 'Georgia' # Classic serif for quotes
    p_canvas.font.color.rgb = RGBColor(40, 40, 45)
    p_canvas.alignment = PP_ALIGN.CENTER
    
    # If multiple paragraphs (like quote author), format them
    if len(tf_canvas.paragraphs) > 1:
        p_author = tf_canvas.paragraphs[1]
        p_author.font.size = Pt(20)
        p_author.font.italic = False
        p_author.font.bold = True
        p_author.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # LAYER 3: 3D HARDWARE (Top Bar, Bottom Bar, Handle)
    # ==========================================
    # Generate high-res assets via PIL
    dpi_scale = 300 # rendering DPI for sharpness
    bar_img_stream = _create_metallic_cylinder(int(bar_w * dpi_scale), int(bar_h * dpi_scale))
    handle_img_stream = _create_pull_handle(int(0.6 * dpi_scale), int(1.2 * dpi_scale))
    
    # Top Bar (Anchors the canvas)
    slide.shapes.add_picture(
        bar_img_stream,
        Inches(bar_l), Inches(canvas_t - (bar_h/2)),
        width=Inches(bar_w), height=Inches(bar_h)
    )
    
    # Bottom Bar (Pulls the canvas down)
    slide.shapes.add_picture(
        bar_img_stream,
        Inches(bar_l), Inches(canvas_t + canvas_h - (bar_h/2)),
        width=Inches(bar_w), height=Inches(bar_h)
    )
    
    # Pull Handle (Attached to bottom bar)
    slide.shapes.add_picture(
        handle_img_stream,
        Inches(13.333/2 - 0.3), Inches(canvas_t + canvas_h),
        width=Inches(0.6), height=Inches(1.2)
    )

    prs.save(output_pptx_path)
    return output_pptx_path
