def create_slide(
    output_pptx_path: str,
    title_text: str = "AutoFit Options: Split Text",
    body_text: str = "",
    title_color: tuple = (237, 125, 49),  # Coral Orange
    text_color: tuple = (0, 112, 192),    # Medium Blue
    **kwargs,
) -> str:
    """
    Create a multi-slide PPTX file demonstrating Programmatic Sidebar Pagination.
    If the text overflows the algorithmic threshold, it splits cleanly across slides.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    # 1. Generate fallback text if none provided (simulating the video's lorem/rand)
    if not body_text:
        body_text = "\n".join([
            f"{i}. The quick brown fox jumps over the lazy dog. PowerPoint creates a new slide with the same title as the current slide, and divides the text approximately evenly."
            for i in range(1, 12)
        ])

    # 2. Generate the dual-pane background image using PIL
    bg_path = "pagination_bg_template.png"
    bg_w, bg_h = 1280, 720
    bg_img = Image.new('RGB', (bg_w, bg_h), (255, 255, 255))
    draw = ImageDraw.Draw(bg_img)
    
    # Calculate 30% split
    divider_x = int(bg_w * 0.30)
    
    # Right panel background (Light Peach)
    right_panel_color = (253, 240, 235) 
    draw.rectangle([divider_x, 0, bg_w, bg_h], fill=right_panel_color)
    
    # Divider Line (Coral Orange)
    draw.line([divider_x, 0, divider_x, bg_h], fill=title_color, width=4)
    bg_img.save(bg_path)

    # 3. Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Chunk text by paragraphs
    paragraphs = [p.strip() for p in body_text.split('\n') if p.strip()]

    # Layout configurations
    col_max_chars = 1100  # Threshold for triggering a slide split
    slide_count = 1

    def add_new_slide(is_first=True):
        nonlocal slide_count
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
        
        # Add background
        slide.shapes.add_picture(bg_path, 0, 0, width=Inches(13.333), height=Inches(7.5))
        
        # Add Sidebar Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(3.0), Inches(4.5))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        
        # Append (Cont.) for subsequent slides
        display_title = title_text if is_first else f"{title_text}\n\n(Cont.)"
        p.text = display_title
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*title_color)
        
        # Add Slide Indicator at bottom of sidebar
        page_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(3.0), Inches(0.5))
        page_tf = page_box.text_frame
        page_p = page_tf.paragraphs[0]
        page_p.text = f"Slide {slide_count}"
        page_p.font.size = Pt(12)
        page_p.font.color.rgb = RGBColor(150, 150, 150)
        
        slide_count += 1
        return slide

    def render_text_column(slide, text_list):
        if not text_list: return
        # Position inside the right-hand panel
        txBox = slide.shapes.add_textbox(Inches(4.5), Inches(0.8), Inches(8.0), Inches(6.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        for i, para_text in enumerate(text_list):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = para_text
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(*text_color)
            p.space_after = Pt(16) # Good breathing room between paragraphs

    # 4. Pagination Engine
    current_slide = add_new_slide(is_first=True)
    current_col_text = []
    current_col_chars = 0
    
    for p in paragraphs:
        # If adding this paragraph exceeds our safety limit, flush and paginate
        if current_col_chars + len(p) > col_max_chars and current_col_text:
            render_text_column(current_slide, current_col_text)
            current_slide = add_new_slide(is_first=False)
            current_col_text = []
            current_col_chars = 0
            
        current_col_text.append(p)
        current_col_chars += len(p)
        
    # Flush any remaining text to the final slide
    if current_col_text:
        render_text_column(current_slide, current_col_text)

    # Clean up temp background image
    if os.path.exists(bg_path):
        os.remove(bg_path)

    prs.save(output_pptx_path)
    return output_pptx_path
