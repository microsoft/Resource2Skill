import os
from PIL import Image, ImageDraw
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE

def create_slide(
    output_pptx_path: str,
    title_text: str = "Steps of Manufacturing in Production Plant",
    body_text: str = "The following slide highlights the key steps of the production plan illustrating initial planning, product development, prototype production, and commercial evaluation.",
    bg_palette: str = "industrial",
    accent_color: tuple = (38, 64, 72),  # Deep slate teal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Alternating Horizontal Process Flow style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Namespace map for lxml operations
    nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

    # Helper: Add drop shadow to a shape
    def add_drop_shadow(shape, alpha=15000, blur=40000, dist=30000, dir=5400000):
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad=str(blur), dist=str(dist), dir=str(dir), algn="ctr", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val=str(alpha))

    # Helper: Add right-facing arrow to the end of a line
    def add_line_tail_arrow(shape):
        ln_elements = shape.element.xpath('.//a:ln', namespaces=nsmap)
        if ln_elements:
            ln = ln_elements[0]
            etree.SubElement(ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', type="triangle", w="med", len="lrg")

    # === Layer 1: Background (PIL Gradient) ===
    bg_path = "temp_bg_gradient.png"
    img = Image.new('RGB', (1920, 1080), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    # Subtle vertical gradient from white to light cool-grey
    for i in range(1080):
        factor = i / 1080.0
        r = int(255 * (1 - factor) + 245 * factor)
        g = int(255 * (1 - factor) + 247 * factor)
        b = int(255 * (1 - factor) + 248 * factor)
        draw.line([(0, i), (1920, i)], fill=(r, g, b))
    img.save(bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    os.remove(bg_path) # cleanup

    # === Layer 2: Top Accent Bar & Typography ===
    primary_rgb = RGBColor(*accent_color)
    secondary_rgb = RGBColor(160, 180, 190)

    # Top accent line
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    top_bar.fill.solid(); top_bar.fill.fore_color.rgb = primary_rgb
    top_bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28); p.font.bold = True; p.font.name = "Calibri"
    p.font.color.rgb = primary_rgb
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.size = Pt(12); p_sub.font.name = "Calibri"
    p_sub.font.color.rgb = RGBColor(100, 100, 100)
    tf_sub.word_wrap = True

    # === Layer 3: Central Spine ===
    axis_y = Inches(4.5)
    axis_start_x = Inches(0.8)
    axis_end_x = Inches(12.5)
    
    axis_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, axis_start_x, axis_y, axis_end_x, axis_y)
    axis_line.line.color.rgb = secondary_rgb
    axis_line.line.width = Pt(4)
    add_line_tail_arrow(axis_line)

    # === Layer 4: Alternating Nodes and Cards ===
    steps_data = kwargs.get('steps_data', [
        {"title": "Initial Planning", "bullets": ["Develop product plan", "Estimate costs", "Resource allocation"]},
        {"title": "Product Dev", "bullets": ["Assess feasibility", "Select machinery", "Define requirements"]},
        {"title": "Prototype", "bullets": ["Produce initial units", "Evaluate quality", "Iterate design"]},
        {"title": "Commercial", "bullets": ["Scale production", "Monitor efficiency", "Quality assurance"]},
        {"title": "Inspection", "bullets": ["Final manual checks", "Package sorting", "Dispatch readiness"]},
    ])

    num_steps = len(steps_data)
    available_width = axis_end_x - axis_start_x - Inches(1.5)
    step_spacing = available_width / (num_steps - 1)
    start_x = axis_start_x + Inches(0.75)
    
    card_w = Inches(2.0)
    card_h_header = Inches(0.4)
    card_h_body = Inches(1.4)
    card_total_h = card_h_header + card_h_body

    for i, step in enumerate(steps_data):
        x = start_x + i * step_spacing
        is_top = (i % 2 == 0)

        # 1. Determine Y coordinates based on alternation
        if is_top:
            card_top_y = axis_y - Inches(0.6) - card_total_h
            line_start_y = card_top_y + card_total_h
            line_end_y = axis_y - Inches(0.125)
        else:
            card_top_y = axis_y + Inches(0.6)
            line_start_y = axis_y + Inches(0.125)
            line_end_y = card_top_y

        # 2. Draw vertical dashed connector
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, line_start_y, x, line_end_y)
        connector.line.color.rgb = secondary_rgb
        connector.line.width = Pt(2)
        connector.line.dash_style = MSO_LINE.DASH

        # 3. Draw Axis Node (Circle) over the connector
        node_size = Inches(0.24)
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - node_size/2, axis_y - node_size/2, node_size, node_size)
        node.fill.solid(); node.fill.fore_color.rgb = RGBColor(255, 255, 255)
        node.line.color.rgb = primary_rgb; node.line.width = Pt(2.5)

        # 4. Draw Card Body
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x - card_w/2, card_top_y + card_h_header, card_w, card_h_body)
        body.fill.solid(); body.fill.fore_color.rgb = RGBColor(255, 255, 255)
        body.line.color.rgb = primary_rgb; body.line.width = Pt(1)
        add_drop_shadow(body, alpha=12000, blur=35000, dist=25000)

        # 5. Populate Card Body Text
        tf_body = body.text_frame
        tf_body.word_wrap = True
        tf_body.margin_top = Pt(8); tf_body.margin_left = Pt(10); tf_body.margin_right = Pt(8)
        for j, bullet in enumerate(step['bullets']):
            p_b = tf_body.add_paragraph() if j > 0 else tf_body.paragraphs[0]
            p_b.text = f"•  {bullet}"
            p_b.font.size = Pt(10.5); p_b.font.name = "Calibri"
            p_b.font.color.rgb = RGBColor(70, 70, 70)
            p_b.space_after = Pt(4)

        # 6. Draw Card Header
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x - card_w/2, card_top_y, card_w, card_h_header)
        header.fill.solid(); header.fill.fore_color.rgb = primary_rgb
        header.line.color.rgb = primary_rgb; header.line.width = Pt(1)
        
        # 7. Populate Card Header Text (Center aligned)
        tf_h = header.text_frame
        tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_h = tf_h.paragraphs[0]
        p_h.text = step['title']
        p_h.font.size = Pt(11); p_h.font.bold = True; p_h.font.name = "Calibri"
        p_h.font.color.rgb = RGBColor(255, 255, 255)
        p_h.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
