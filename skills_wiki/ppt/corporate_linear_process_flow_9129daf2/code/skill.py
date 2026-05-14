def create_slide(
    output_pptx_path: str,
    title_text: str = "Review Methods for Evaluating Performance",
    subtitle_text: str = "A comprehensive 4-step approach to staff assessment.",
    accent_color: tuple = (244, 180, 26),  # Mustard Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Corporate Linear Process Flow' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from lxml.etree import ElementBase
    from pptx.oxml import parse_xml

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Color Palette
    bg_color = RGBColor(245, 245, 240)
    text_dark = RGBColor(44, 62, 80)
    text_light = RGBColor(100, 110, 120)
    acc_rgb = RGBColor(*accent_color)
    white = RGBColor(255, 255, 255)

    # --- Set Background Color ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # --- Helper Function for Shadows ---
    def apply_subtle_shadow(shape):
        """Inject OOXML to add a soft drop shadow to a shape."""
        spPr = shape.element.spPr
        shadow_xml = f"""
            <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <a:outerShdw blurRad="150000" dist="40000" dir="5400000" algn="b">
                    <a:srgbClr val="000000">
                        <a:alpha val="10000"/>
                    </a:srgbClr>
                </a:outerShdw>
            </a:effectLst>
        """
        effectLst = parse_xml(shadow_xml)
        spPr.append(effectLst)

    # --- Header Texts ---
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.33), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = text_dark
    p.alignment = PP_ALIGN.LEFT

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.33), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(16)
    p_sub.font.name = "Arial"
    p_sub.font.color.rgb = text_light
    p_sub.alignment = PP_ALIGN.LEFT

    # --- Process Flow Data ---
    steps = [
        {"num": "01", "title": "Self Evaluation", "desc": "Employee evaluating their own performance against pre-defined criteria."},
        {"num": "02", "title": "Checklist", "desc": "Performing structured checklist evaluation to assess specific employee behaviors."},
        {"num": "03", "title": "360-Degree Feedback", "desc": "Taking a holistic look at staff performance including peer and management feedback."},
        {"num": "04", "title": "Objective Management", "desc": "Setting attainable goals to be achieved by the employee in a specified timeframe."}
    ]

    # --- Layout Geometry ---
    num_steps = len(steps)
    margin_x = 1.0
    usable_width = 13.333 - (2 * margin_x)
    step_spacing = usable_width / (num_steps - 1)
    
    line_y = 2.8 # Y-coordinate for the main connecting line
    card_y = 3.6 # Y-coordinate for the top of the cards
    card_width = 2.4
    card_height = 2.8
    circle_radius = 0.45

    # 1. Draw Connecting Line (goes behind everything)
    connector = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(margin_x), Inches(line_y - 0.02), 
        Inches(usable_width), Inches(0.04)
    )
    connector.fill.solid()
    connector.fill.fore_color.rgb = RGBColor(210, 210, 205)
    connector.line.fill.background()

    # 2. Iterate and Draw Steps
    for i, step in enumerate(steps):
        center_x = margin_x + (i * step_spacing)
        
        # Vertical connector stem (line from circle to card)
        stem = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(center_x - 0.02), Inches(line_y),
            Inches(0.04), Inches(card_y - line_y + 0.2)
        )
        stem.fill.solid()
        stem.fill.fore_color.rgb = RGBColor(210, 210, 205)
        stem.line.fill.background()

        # Content Card (Rounded Rectangle)
        card_x = center_x - (card_width / 2)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(card_x), Inches(card_y),
            Inches(card_width), Inches(card_height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = white
        card.line.fill.background() # No border
        # Adjust corner radius to be subtle
        card.adjustments[0] = 0.05
        apply_subtle_shadow(card)

        # Card Text - Title
        c_title_box = slide.shapes.add_textbox(
            Inches(card_x + 0.1), Inches(card_y + 0.2), 
            Inches(card_width - 0.2), Inches(0.6)
        )
        c_tf = c_title_box.text_frame
        c_tf.word_wrap = True
        p_ct = c_tf.paragraphs[0]
        p_ct.text = step["title"]
        p_ct.font.size = Pt(16)
        p_ct.font.bold = True
        p_ct.font.name = "Arial"
        p_ct.font.color.rgb = text_dark
        p_ct.alignment = PP_ALIGN.CENTER

        # Card Text - Description
        c_desc_box = slide.shapes.add_textbox(
            Inches(card_x + 0.15), Inches(card_y + 0.8), 
            Inches(card_width - 0.3), Inches(1.8)
        )
        d_tf = c_desc_box.text_frame
        d_tf.word_wrap = True
        p_cd = d_tf.paragraphs[0]
        p_cd.text = step["desc"]
        p_cd.font.size = Pt(12)
        p_cd.font.name = "Arial"
        p_cd.font.color.rgb = text_light
        p_cd.alignment = PP_ALIGN.CENTER

        # Node Circle (Drawn last so it sits on top of lines)
        circle_x = center_x - circle_radius
        circle_y = line_y - circle_radius
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(circle_x), Inches(circle_y),
            Inches(circle_radius * 2), Inches(circle_radius * 2)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = acc_rgb
        # White border to make it pop off the line
        circle.line.color.rgb = white
        circle.line.width = Pt(3)
        apply_subtle_shadow(circle)

        # Number inside circle
        circle_tf = circle.text_frame
        circle_tf.word_wrap = False
        circle_tf.margin_top = circle_tf.margin_bottom = circle_tf.margin_left = circle_tf.margin_right = 0
        p_num = circle_tf.paragraphs[0]
        p_num.text = step["num"]
        p_num.font.size = Pt(18)
        p_num.font.bold = True
        p_num.font.name = "Arial"
        p_num.font.color.rgb = white
        p_num.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
