def create_slide(
    output_pptx_path: str,
    title_text: str = "ORDER NOW",
    link_text: str = "https://yourwebsite.com",
    bg_theme: str = "office,laptop",
    accent_color: tuple = (203, 175, 142),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Glassmorphic Concentric Pill Overlay pattern.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # --- HELPER FUNCTIONS FOR LXML XML INJECTION ---
    def set_shape_transparency(shape, opacity_percent):
        """Injects alpha value into shape's solid fill (opacity_percent: 0-100)"""
        opacity_val = int(opacity_percent * 1000)
        spPr = shape.element.spPr
        solidFill = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        if solidFill is not None:
            color_element = solidFill[0]
            for child in color_element:
                if child.tag.endswith('alpha'):
                    color_element.remove(child)
            alpha = parse_xml(f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{opacity_val}"/>')
            color_element.append(alpha)

    def set_line_transparency(shape, opacity_percent):
        """Injects alpha value into shape border's solid fill"""
        opacity_val = int(opacity_percent * 1000)
        spPr = shape.element.spPr
        ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is not None:
            solidFill = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
            if solidFill is not None:
                color_element = solidFill[0]
                for child in color_element:
                    if child.tag.endswith('alpha'):
                        color_element.remove(child)
                alpha = parse_xml(f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{opacity_val}"/>')
                color_element.append(alpha)

    def remove_fill(shape):
        """Removes fill entirely"""
        spPr = shape.element.spPr
        for child in list(spPr):
            if child.tag.endswith('Fill'):
                spPr.remove(child)
        noFill = parse_xml('<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
        spPr.insert(0, noFill)

    def remove_line(shape):
        """Removes outline entirely"""
        spPr = shape.element.spPr
        ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is not None:
            spPr.remove(ln)

    # === LAYER 1: Background Image ===
    bg_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1920x1080/?{bg_theme}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(bg_path, 'wb') as out_file:
            out_file.write(response.read())
        slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception:
        # Fallback to dark solid rectangle if download fails
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(40, 40, 40)
        remove_line(bg)

    # === LAYER 2: PIL Gradient Overlay Mask ===
    overlay_path = "temp_overlay.png"
    overlay = Image.new('RGBA', (1920, 1080))
    draw = ImageDraw.Draw(overlay)
    
    r1, g1, b1 = (60, 60, 60)         # Top: Dark Gray
    a1 = int((100 - 42) / 100 * 255)  # 42% trans -> 58% opacity
    r2, g2, b2 = accent_color         # Bottom: Peach/Accent
    a2 = int((100 - 20) / 100 * 255)  # 20% trans -> 80% opacity

    for y in range(1080):
        ratio = y / 1080
        r = int(r1 + (r2 - r1) * ratio)
        g = int(g1 + (g2 - g1) * ratio)
        b = int(b1 + (b2 - b1) * ratio)
        a = int(a1 + (a2 - a1) * ratio)
        draw.line([(0, y), (1920, y)], fill=(r, g, b, a))
        
    overlay.save(overlay_path)
    slide.shapes.add_picture(overlay_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === LAYER 3: Main Typographic CTA ===
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11.333), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(80)
    p.font.bold = True
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(50, 50, 50)

    # === LAYER 4: Glassmorphic Concentric Pill Button ===
    cx = 13.333 / 2
    cy = 5.2

    # 1. Outer Pill
    w_out, h_out = 6.2, 1.5
    outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - w_out/2), Inches(cy - h_out/2), Inches(w_out), Inches(h_out))
    outer.adjustments[0] = 0.5  # Max rounded (pill)
    remove_fill(outer)
    outer.line.fill.solid()
    outer.line.color.rgb = RGBColor(255, 255, 255)
    outer.line.width = Pt(1)
    set_line_transparency(outer, 52) # 48% transparency -> 52% opacity

    # 2. Middle Pill
    w_mid, h_mid = 6.0, 1.3
    middle = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - w_mid/2), Inches(cy - h_mid/2), Inches(w_mid), Inches(h_mid))
    middle.adjustments[0] = 0.5
    remove_fill(middle)
    middle.line.fill.solid()
    middle.line.color.rgb = RGBColor(255, 255, 255)
    middle.line.width = Pt(2)
    set_line_transparency(middle, 52)

    # 3. Inner Pill
    w_in, h_in = 5.8, 1.1
    inner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - w_in/2), Inches(cy - h_in/2), Inches(w_in), Inches(h_in))
    inner.adjustments[0] = 0.5
    inner.fill.solid()
    inner.fill.fore_color.rgb = RGBColor(255, 255, 255)
    set_shape_transparency(inner, 75) # 25% transparency -> 75% opacity
    remove_line(inner)

    # 4. URL Link Text Box
    link_box = slide.shapes.add_textbox(Inches(cx - w_in/2), Inches(cy - h_in/2 + 0.1), Inches(w_in), Inches(h_in))
    tf_link = link_box.text_frame
    p_link = tf_link.paragraphs[0]
    p_link.text = link_text
    p_link.alignment = PP_ALIGN.CENTER
    p_link.font.size = Pt(28)
    p_link.font.color.rgb = RGBColor(60, 60, 60)
    p_link.font.underline = True

    # === LAYER 5: Top Hanging Tab (Logo Placement) ===
    # Draw a rounded rectangle bleeding off the top edge to create a bottom-rounded tab
    w_tab, h_tab = 3.6, 2.0
    tab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - w_tab/2), Inches(-1.0), Inches(w_tab), Inches(h_tab))
    tab.adjustments[0] = 0.3 # softer curve
    tab.fill.solid()
    tab.fill.fore_color.rgb = RGBColor(255, 255, 255)
    set_shape_transparency(tab, 85) # 15% trans -> 85% opacity
    remove_line(tab)

    # Add dummy branding text to the visible part of the tab
    brand_box = slide.shapes.add_textbox(Inches(cx - w_tab/2), Inches(0.1), Inches(w_tab), Inches(0.8))
    p_brand = brand_box.text_frame.paragraphs[0]
    p_brand.text = "★ PREMIUM DESIGN ★"
    p_brand.alignment = PP_ALIGN.CENTER
    p_brand.font.size = Pt(14)
    p_brand.font.bold = True
    p_brand.font.color.rgb = RGBColor(100, 100, 100)

    # Cleanup temp files
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
    if os.path.exists(overlay_path):
        os.remove(overlay_path)

    return output_pptx_path
