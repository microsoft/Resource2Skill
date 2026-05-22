def create_slide(
    output_pptx_path: str,
    title_text: str = "Strategic Framework Analysis",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Modular Framework Alignment" (STP style) effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    
    prs = Presentation()
    # Use standard 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Data & Theme Definition ===
    # Defining the nodes with their specific colors (Dark for border/text, Light for background)
    nodes = [
        {
            "header": "S",
            "body": "Segmenting the market based on demographics, psychographics, and behavior.\n\nEnsure data is backed by reliable sources.",
            "color_main": RGBColor(146, 208, 80),  # Green
            "color_light": RGBColor(226, 240, 217)
        },
        {
            "header": "T",
            "body": "Targeting specific segments that offer the most value and align with business objectives.\n\nFocus on high-yield customer profiles.",
            "color_main": RGBColor(155, 194, 230),  # Blue
            "color_light": RGBColor(222, 235, 247)
        },
        {
            "header": "P",
            "body": "Positioning the brand in the minds of the consumers to differentiate from competitors.\n\nEstablish a strong, unique value proposition.",
            "color_main": RGBColor(244, 177, 131),  # Orange
            "color_light": RGBColor(252, 228, 214)
        }
    ]

    # === Geometric Calculations for "Distribute Horizontally" ===
    num_nodes = len(nodes)
    card_width = Inches(3.0)
    card_height = Inches(4.0)
    badge_radius = Inches(0.7)  # Diameter will be 1.4 inches
    
    # Calculate spacing
    total_canvas_width = prs.slide_width
    total_cards_width = card_width * num_nodes
    # Space remaining to be distributed (edges + between cards)
    total_empty_space = total_canvas_width - total_cards_width
    # We want margins on left/right and gaps between. Number of gaps = num_nodes + 1
    gap_width = total_empty_space / (num_nodes + 1)
    
    y_top_rect = Inches(2.5) # Y position for the main text box

    # Add Slide Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), prs.slide_width - Inches(2), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(50, 50, 50)
    p.alignment = PP_ALIGN.CENTER

    # === Create the Modular Cards ===
    for i, node in enumerate(nodes):
        # Calculate X position for this specific card
        current_x = gap_width + (i * (card_width + gap_width))
        
        # 1. Create Body Rectangle (Background)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            current_x, y_top_rect, card_width, card_height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = node["color_light"]
        rect.line.fill.background() # No border
        
        # Adjust rounded corner radius (magic adjustment value for pptx)
        rect.adjustments[0] = 0.05 
        
        # Add body text
        text_frame = rect.text_frame
        text_frame.word_wrap = True
        text_frame.margin_left = Inches(0.3)
        text_frame.margin_right = Inches(0.3)
        text_frame.margin_top = Inches(1.0) # Leave space for the overlapping badge
        
        p = text_frame.paragraphs[0]
        p.text = node["body"]
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(60, 60, 60)
        p.alignment = PP_ALIGN.LEFT
        
        # 2. Create Header Badge (Overlapping Circle)
        # Center of circle aligns with center of rectangle; middle of circle aligns with top of rectangle
        circle_x = current_x + (card_width / 2) - badge_radius
        circle_y = y_top_rect - badge_radius
        
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            circle_x, circle_y, badge_radius * 2, badge_radius * 2
        )
        
        # Style the Badge: White fill, Thick colored border (The core skill from the video)
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
        circle.line.color.rgb = node["color_main"]
        circle.line.width = Pt(6) # THICK border
        
        # Add Header Text
        badge_tf = circle.text_frame
        badge_p = badge_tf.paragraphs[0]
        badge_p.text = node["header"]
        badge_p.font.size = Pt(48)
        badge_p.font.bold = True
        badge_p.font.color.rgb = node["color_main"]
        badge_p.alignment = PP_ALIGN.CENTER
        
        # Center text vertically within the circle
        badge_tf.vertical_anchor = 3 # Middle

    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("framework_alignment.pptx")
