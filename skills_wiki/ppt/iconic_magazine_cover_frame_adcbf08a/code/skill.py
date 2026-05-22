def create_slide(
    output_pptx_path: str,
    title_text: str = "GLOBAL EXPEDITION",
    body_text: str = "INTO THE UNKNOWN\nDiscovering the secrets of the world's most isolated habitats.",
    bg_palette: str = "nature",
    accent_color: tuple = (255, 204, 0),  # Classic yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Iconic Magazine Cover Frame visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout

    # === Layer 1: Background Image ===
    # Download a high-quality wildlife/nature placeholder image
    img_url = "https://images.unsplash.com/photo-1540573133985-87b6da6d54a9?q=80&w=1920&auto=format&fit=crop"
    img_path = "temp_bg.jpg"
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            bg_img = Image.open(BytesIO(response.read())).convert("RGB")
            
            # Crop to 16:9 aspect ratio
            target_ratio = 16 / 9
            w, h = bg_img.size
            img_ratio = w / h
            if img_ratio > target_ratio:
                new_w = int(h * target_ratio)
                left = (w - new_w) / 2
                bg_img = bg_img.crop((left, 0, left + new_w, h))
            else:
                new_h = int(w / target_ratio)
                top = (h - new_h) / 2
                bg_img = bg_img.crop((0, top, w, top + new_h))
            
            bg_img.save(img_path)
    except Exception:
        # Fallback if download fails
        bg_img = Image.new("RGB", (1920, 1080), (30, 35, 40))
        bg_img.save(img_path)

    # Insert background
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: PIL Gradient Overlay & Frame ===
    frame_path = "temp_frame.png"
    dpi = 300
    sw, sh = 13.333, 7.5
    w_px, h_px = int(sw * dpi), int(sh * dpi)
    
    # Create empty transparent canvas
    overlay = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    
    # 1. Draw top and bottom gradients to ensure text readability over bright photos
    grad_height = int(2.5 * dpi)
    for y in range(grad_height):
        # Top gradient fading down
        alpha_top = int(180 * (1 - (y / grad_height)))
        draw.line([(0, y), (w_px, y)], fill=(0, 0, 0, alpha_top))
        
        # Bottom gradient fading up
        alpha_bot = int(220 * (1 - (y / grad_height)))
        draw.line([(0, h_px - y), (w_px, h_px - y)], fill=(0, 0, 0, alpha_bot))
        
    # 2. Draw the iconic inset border
    margin_px = int(0.5 * dpi) # 0.5 inch inset margin
    thickness_px = int(0.15 * dpi) # ~11pt line thickness
    
    draw.rectangle(
        [margin_px, margin_px, w_px - margin_px, h_px - margin_px],
        outline=accent_color + (255,), # Apply RGBA accent color
        width=thickness_px
    )
    
    overlay.save(frame_path)
    slide.shapes.add_picture(frame_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 3: Typography ===
    
    # Tiny Date / Metadata
    date_box = slide.shapes.add_textbox(Inches(0), Inches(0.5), prs.slide_width, Inches(0.5))
    dtf = date_box.text_frame
    dp = dtf.add_paragraph()
    dp.text = "VOL. 245 • SPECIAL EDITION"
    dp.alignment = PP_ALIGN.CENTER
    dfont = dp.font
    dfont.name = 'Arial'
    dfont.size = Pt(11)
    dfont.bold = True
    dfont.color.rgb = RGBColor(*accent_color)

    # Main Magazine Title
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.8), prs.slide_width, Inches(1.5))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    font = p.font
    font.name = 'Times New Roman'
    font.size = Pt(64)
    font.bold = True
    font.color.rgb = RGBColor(255, 255, 255)

    # Parse body text lines
    lines = body_text.split('\n')
    headline_text = lines[0] if len(lines) > 0 else "HEADLINE GOES HERE"
    sub_text = lines[1] if len(lines) > 1 else "Subtitle or description text goes here."

    # Article Headline
    hl_box = slide.shapes.add_textbox(Inches(1), Inches(5.3), Inches(11.333), Inches(1))
    htf = hl_box.text_frame
    hp = htf.add_paragraph()
    hp.text = headline_text
    hp.alignment = PP_ALIGN.CENTER
    hfont = hp.font
    hfont.name = 'Times New Roman'
    hfont.size = Pt(36)
    hfont.bold = True
    hfont.color.rgb = RGBColor(255, 255, 255)

    # Article Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(6.0), Inches(11.333), Inches(0.8))
    stf = sub_box.text_frame
    sp = stf.add_paragraph()
    sp.text = sub_text
    sp.alignment = PP_ALIGN.CENTER
    sfont = sp.font
    sfont.name = 'Arial'
    sfont.size = Pt(16)
    sfont.color.rgb = RGBColor(*accent_color)

    prs.save(output_pptx_path)

    # Cleanup temp files
    if os.path.exists(img_path):
        os.remove(img_path)
    if os.path.exists(frame_path):
        os.remove(frame_path)

    return output_pptx_path
