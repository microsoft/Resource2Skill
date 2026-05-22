import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFilter
from lxml import etree

def create_slide(
    output_pptx_path: str,
    title_text: str = "Name of Presentation",
    subtitle_text: str = "Company Name",
    slide_title: str = "Key Business Metrics",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Corporate Flow Master Template' visual effect.
    """
    prs = Presentation()
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    width_px, height_px = 1920, 1080
    bg_image_path = "temp_corporate_bg.png"

    # ==========================================
    # Layer 1: Generate Master Background (PIL)
    # ==========================================
    bg_img = Image.new('RGB', (width_px, height_px), (255, 255, 255))
    draw = ImageDraw.Draw(bg_img, 'RGBA')

    # 1. Top Header Banner (Gradient)
    banner_height = int(height_px * 0.16)
    for y in range(banner_height):
        ratio = y / banner_height
        r = int(230 - (230 - 160) * ratio)
        g = int(240 - (240 - 200) * ratio)
        b = int(250 - (250 - 230) * ratio)
        draw.line([(0, y), (width_px, y)], fill=(r, g, b, 255))
    
    # Solid bottom border for banner
    draw.line([(0, banner_height), (width_px, banner_height)], fill=(90, 140, 200, 255), width=6)

    # 2. Abstract Flowing Waves
    wave_layer = Image.new('RGBA', (width_px, height_px), (0, 0, 0, 0))
    w_draw = ImageDraw.Draw(wave_layer)

    # Use large, off-screen ellipses to create sweeping arcs
    # Light thick swoosh
    w_draw.ellipse([-width_px*0.1, height_px*0.3, width_px*0.8, height_px*1.5], outline=(100, 150, 220, 60), width=35)
    # Dark thin intertwining swoosh
    w_draw.ellipse([-width_px*0.05, height_px*0.1, width_px*0.9, height_px*1.4], outline=(30, 80, 180, 110), width=10)
    # Right-side balancing swoosh
    w_draw.ellipse([width_px*0.4, -height_px*0.2, width_px*1.4, height_px*1.1], outline=(50, 120, 200, 80), width=25)
    # Accent sharp line
    w_draw.ellipse([width_px*0.2, height_px*0.5, width_px*1.2, height_px*1.8], outline=(10, 50, 150, 140), width=5)

    # Soften the curves
    wave_layer = wave_layer.filter(ImageFilter.GaussianBlur(8))
    bg_img.paste(wave_layer, (0, 0), wave_layer)
    bg_img.save(bg_image_path)


    # ==========================================
    # Helper: Apply Gradient Fill to Shapes (lxml)
    # ==========================================
    def apply_gradient(shape, hex_color1, hex_color2):
        spPr = shape.element.spPr
        nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        
        # Remove any existing solid fills
        for solid_fill in spPr.xpath('./a:solidFill', namespaces=nsmap):
            spPr.remove(solid_fill)
            
        ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        gradFill = etree.SubElement(spPr, f"{{{ns}}}gradFill")
        # 5400000 = 90 degrees (top to bottom gradient)
        lin = etree.SubElement(gradFill, f"{{{ns}}}lin", ang="5400000", scaled="1") 
        gsLst = etree.SubElement(gradFill, f"{{{ns}}}gsLst")
        
        # Top Color (Lighter)
        gs1 = etree.SubElement(gsLst, f"{{{ns}}}gs", pos="0")
        etree.SubElement(gs1, f"{{{ns}}}srgbClr", val=hex_color1)
        
        # Bottom Color (Darker)
        gs2 = etree.SubElement(gsLst, f"{{{ns}}}gs", pos="100000")
        etree.SubElement(gs2, f"{{{ns}}}srgbClr", val=hex_color2)

    # ==========================================
    # Slide 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    slide1.shapes.add_picture(bg_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Title in the banner
    tx_box = slide1.shapes.add_textbox(Inches(1), Inches(0.3), Inches(11.333), Inches(1))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = 'Arial'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 65, 120)
    p.alignment = PP_ALIGN.RIGHT

    # Subtitle / Company name at bottom
    sub_box = slide1.shapes.add_textbox(Inches(6), Inches(6.5), Inches(6.333), Inches(0.8))
    tf2 = sub_box.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = subtitle_text
    p2.font.name = 'Arial'
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(100, 100, 100)
    p2.alignment = PP_ALIGN.RIGHT


    # ==========================================
    # Slide 2: Content Slide (4-Box Grid)
    # ==========================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_picture(bg_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Content Title in banner
    tx_box_2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10), Inches(0.8))
    tf3 = tx_box_2.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = slide_title
    p3.font.name = 'Arial'
    p3.font.size = Pt(32)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(15, 65, 120)

    # Create 4 Glossy/Gradient Blocks mimicking the tutorial style
    box_width = Inches(4.5)
    box_height = Inches(2.0)
    start_x = Inches(2.0)
    start_y = Inches(2.2)
    gap = Inches(0.4)

    blocks = [
        {"x": start_x, "y": start_y, "c1": "4FC3F7", "c2": "0288D1", "label": "Analysis"},
        {"x": start_x + box_width + gap, "y": start_y, "c1": "FFB74D", "c2": "F57C00", "label": "Strategy"},
        {"x": start_x, "y": start_y + box_height + gap, "c1": "81C784", "c2": "388E3C", "label": "Execution"},
        {"x": start_x + box_width + gap, "y": start_y + box_height + gap, "c1": "E57373", "c2": "D32F2F", "label": "Reporting"}
    ]

    for block in blocks:
        # Add shape
        shape = slide2.shapes.add_shape(
            1, # msoShapeRectangle
            block["x"], block["y"], box_width, box_height
        )
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(1.5)
        
        # Apply custom XML gradient
        apply_gradient(shape, block["c1"], block["c2"])

        # Add formatted text
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = block["label"]
        p.font.name = 'Arial'
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        
        # Add subtext
        p_sub = tf.add_paragraph()
        p_sub.text = "Click to add supporting descriptive text here. This area holds detailed information."
        p_sub.font.size = Pt(12)
        p_sub.font.bold = False
        p_sub.alignment = PP_ALIGN.CENTER

    # Cleanup temp image
    try:
        prs.save(output_pptx_path)
    finally:
        if os.path.exists(bg_image_path):
            os.remove(bg_image_path)

    return output_pptx_path
