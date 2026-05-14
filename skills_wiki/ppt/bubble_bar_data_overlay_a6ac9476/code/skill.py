import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN

def create_slide(
    output_pptx_path: str,
    chart_data: dict,
    title_text: str = "Bubble Bar Chart",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Bubble Bar Data Overlay visual effect.

    The function generates a slide containing a horizontal bar chart where one series
    (e.g., Profit) is overlaid on another (e.g., Revenue). To the right of each bar,
    a bubble is placed, with its size and label corresponding to a third metric (e.g., Profit Rate).

    Args:
        output_pptx_path: Path to save the output .pptx file.
        chart_data: A dictionary containing the data for the chart.
                    Expected format: {'Category': [...], 'Revenue': [...], 'Profit': [...]}.
                    'ProfitRate' will be calculated from this data.
        title_text: The main title for the chart.

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Data Preparation ===
    df = pd.DataFrame(chart_data)
    df['ProfitRate'] = df['Profit'] / df['Revenue']
    df = df.sort_values(by='Revenue', ascending=True) # Sort for visual order in chart
    categories = df['Category'].tolist()
    
    chart_data_obj = ChartData()
    chart_data_obj.categories = categories
    chart_data_obj.add_series('Revenue', df['Revenue'].tolist())
    chart_data_obj.add_series('Profit', df['Profit'].tolist())

    # === Layer 2: Bar Chart ===
    x, y, cx, cy = Inches(1.0), Inches(1.25), Inches(8), Inches(5.5)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, x, y, cx, cy, chart_data_obj
    )
    chart = graphic_frame.chart

    # Format plot area
    chart.plot_area.format.fill.background()
    chart.plot_area.format.line.fill.background()
    
    # Format category axis (Y-axis)
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(12)
    category_axis.format.line.fill.background()
    category_axis.reverse_order = True

    # Format value axis (X-axis)
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = False
    value_axis.visible = False

    # Format series
    series_revenue = chart.series[0]
    series_revenue.fill.solid()
    series_revenue.fill.fore_color.rgb = RGBColor(155, 194, 230)
    series_revenue.has_data_labels = True
    data_labels_revenue = series_revenue.data_labels
    data_labels_revenue.position = XL_DATA_LABEL_POSITION.INSIDE_END
    data_labels_revenue.font.size = Pt(11)
    data_labels_revenue.font.color.rgb = RGBColor(89, 89, 89)

    series_profit = chart.series[1]
    series_profit.fill.solid()
    series_profit.fill.fore_color.rgb = RGBColor(47, 117, 181)
    series_profit.has_data_labels = True
    data_labels_profit = series_profit.data_labels
    data_labels_profit.position = XL_DATA_LABEL_POSITION.CENTER
    data_labels_profit.font.size = Pt(11)
    data_labels_profit.font.color.rgb = RGBColor(255, 255, 255)
    data_labels_profit.font.bold = True

    # Overlap series and set gap width
    chart.plot_area.cat_groupings[0].overlap = 100
    chart.plot_area.cat_groupings[0].gap_width = 80

    # Format Legend and Title
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    chart.chart_title.text_frame.text = title_text
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(18)
    
    # === Layer 3: Bubbles ===
    num_categories = len(categories)
    plot_height_estimate = cy - Inches(0.5) # Estimate plot area height
    plot_y_start_estimate = y + Inches(0.25) # Estimate plot area top
    
    bubble_x_pos = x + cx + Inches(0.3)
    
    min_rate, max_rate = df['ProfitRate'].min(), df['ProfitRate'].max()
    min_bubble_size, max_bubble_size = Inches(0.4), Inches(0.9)

    for i, row in df.reset_index().iterrows():
        # Calculate Y position for the bubble, aligning with the bar center
        category_band_height = plot_height_estimate / num_categories
        bar_center_y = plot_y_start_estimate + (i * category_band_height) + (category_band_height / 2)
        
        rate = row['ProfitRate']
        scale_factor = (rate - min_rate) / (max_rate - min_rate) if (max_rate - min_rate) > 0 else 0.5
        bubble_diameter = min_bubble_size + (scale_factor * (max_bubble_size - min_bubble_size))
        bubble_y_pos = bar_center_y - (bubble_diameter / 2)
        
        bubble = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, bubble_x_pos, bubble_y_pos, bubble_diameter, bubble_diameter
        )
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = RGBColor(79, 178, 222)
        bubble.line.fill.background()
        
        text_frame = bubble.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.text = f"{rate:.1%}"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

    prs.save(output_pptx_path)
    return output_pptx_path

if __name__ == '__main__':
    # Example data from the tutorial
    tutorial_data = {
        'Category': ['产品 F', '产品 G', '产品 C', '产品 E', '产品 D', '产品 A', '产品 B'],
        'Revenue': [1916, 1843, 1839, 1820, 1614, 1612, 1304],
        'Profit': [188, 123, 385, 177, 163, 369, 200],
    }
    create_slide(
        output_pptx_path="bubble_bar_chart_reproduction.pptx",
        chart_data=tutorial_data,
        title_text="泡泡组合条形图类型"
    )

