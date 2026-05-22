def create_slide(
    output_pptx_path: str,
    title_text: str = "Comparison Of 2 Ideas",
    **kwargs
) -> str:
    """
    Creates a PPTX file reproducing the 'Mirrored Concentric Arc Comparison' layout.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw
    import io

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Main Slide Title ===
    title_box = slide.shapes.add_textbox(Inches(2.0), Inches(0.4), Inches(9.333), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.name = 'Arial'
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # === Layer 2: PIL Background Arcs Generator ===
    # We use PIL to generate perfect half-donuts and half-circles, avoiding PPTX rendering quirks
    size = 1600
    left_color = (137, 186, 227, 255)  # Light Blue
    right_color = (150, 150, 150, 255) # Light Gray

    # 1. Outer Rings
    outer = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    d_outer = ImageDraw.Draw(outer)
    # Pieslices for the main bulk (PIL angles: 0=Right, 90=Bottom)
    d_outer.pieslice([size//2 - 550, size//2 - 550, size//2 + 550, size//2 + 550], 90, 270, fill=left_color)
    d_outer.pieslice([size//2 - 550, size//2 - 550, size//2 + 550, size//2 + 550], -90, 90, fill=right_color)

    # 2. Cut central hole in the outer rings using an alpha mask
    mask = outer.split()[3]
    d_mask = ImageDraw.Draw(mask)
    d_mask.ellipse([size//2 - 400, size//2 - 400, size//2 + 400, size//2 + 400], fill=0)
    outer.putalpha(mask)

    # 3. Inner Solid Rings
    inner = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    d_inner = ImageDraw.Draw(inner)
    d_inner.pieslice([size//2 - 250, size//2 - 250, size//2 + 250, size//2 + 250], 90, 270, fill=left_color)
    d_inner.pieslice([size//2 - 250, size//2 - 250, size//2 + 250, size//2 + 250], -90, 90, fill=right_color)

    # 4. Composite Layers
    final_img = Image.alpha_composite(outer, inner)
    
    img_stream = io.BytesIO()
    final_img.save(img_stream, format='PNG')
    img_stream.seek(0)
    
    # Insert PIL image into PPTX, sized to exactly 6 inches and perfectly centered
    img_size = 6.0
    img_left = (13.333 / 2) - (img_size / 2)
    img_top = 1.0 # Lowered to account for title space
    slide.shapes.add_picture(img_stream, Inches(img_left), Inches(img_top), Inches(img_size), Inches(img_size))

    # Center anchor coordinates (Matches the visual center of the inserted image)
    cx, cy = 6.666, 4.0

    # === Layer 3: Central Dividing Line ===
    line = slide.shapes.add_connector(
        1, # MSO_CONNECTOR.STRAIGHT
        Inches(cx), Inches(1.5), Inches(cx), Inches(6.8)
    )
    line.line.dash_style = MSO_LINE.DASH
    line.line.width = Pt(2.5)
    line.line.color.rgb = RGBColor(100, 100, 100)

    # === Layer 4: Inner Core Labels ===
    def add_center_label(cx_shift, title, subtitle):
        tx = slide.shapes.add_textbox(Inches(cx + cx_shift), Inches(cy - 0.5), Inches(1.0), Inches(1.0))
        tf = tx.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(255, 255, 255)
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(255, 255, 255)
        p2.alignment = PP_ALIGN.CENTER

    add_center_label(-1.1, "Idea 1", "Title Here")
    add_center_label(0.1, "Idea 2", "Title Here")

    # === Layer 5: Data Nodes & Text Block Overlays ===
    # Math derivation: Outer arc spans R=400 to R=550px in 1600px canvas. 
    # The exact center of the arc is 475px. Scaling to 6 inches -> R = 1.78 inches.
    
    # Cos(45), Sin(45) applied for top/bottom nodes: 1.78 * 0.707 = 1.258
    node_coords = [
        # X, Y, Number, Color, Is_Left_Side
        (cx - 1.258, cy - 1.258, "1", RGBColor(244, 144, 71), True),  # Orange
        (cx - 1.78,  cy,         "2", RGBColor(67, 183, 110), True),  # Green
        (cx - 1.258, cy + 1.258, "3", RGBColor(74, 144, 226), True),  # Blue
        
        (cx + 1.258, cy - 1.258, "1", RGBColor(192, 57, 43), False),  # Red
        (cx + 1.78,  cy,         "2", RGBColor(142, 68, 173), False), # Purple
        (cx + 1.258, cy + 1.258, "3", RGBColor(241, 196, 15), False), # Yellow
    ]

    for x, y, num, color, is_left in node_coords:
        # Draw Native Editable Node Circle
        r = 0.22 
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(r*2), Inches(r*2)
        )
        node.fill.solid()
        node.fill.fore_color.rgb = color
        node.line.color.rgb = RGBColor(255, 255, 255)
        node.line.width = Pt(2.5)
        
        # Center Number inside Node
        tf = node.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = num
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        # Draw associated Descriptive Text Box
        tw, th = 2.5, 1.0
        tx_left = (x - tw - 0.4) if is_left else (x + 0.4)
        tx_top = y - 0.4

        txBox = slide.shapes.add_textbox(Inches(tx_left), Inches(tx_top), Inches(tw), Inches(th))
        tf_box = txBox.text_frame
        
        p_title = tf_box.paragraphs[0]
        p_title.text = "TITLE HERE"
        p_title.font.bold = True
        p_title.font.size = Pt(12)
        p_title.font.color.rgb = RGBColor(60, 60, 60)
        if is_left: p_title.alignment = PP_ALIGN.RIGHT
        
        p_body = tf_box.add_paragraph()
        p_body.text = "Add details in 2-3 lines to describe the title. Lesser the content better it will look like."
        p_body.font.size = Pt(10)
        p_body.font.color.rgb = RGBColor(120, 120, 120)
        if is_left: p_body.alignment = PP_ALIGN.RIGHT
        
    prs.save(output_pptx_path)
    return output_pptx_path
