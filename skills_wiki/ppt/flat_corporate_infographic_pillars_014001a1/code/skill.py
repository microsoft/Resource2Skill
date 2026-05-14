def create_slide(
    output_pptx_path: str,
    title_text: str = "COMPLETED, OUTSTANDING AND NEW ITEMS",
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the 'Flat Corporate Infographic Pillars' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.shapes.freeform import FreeformBuilder
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Define Theme Colors ===
    c_red = RGBColor(211, 49, 65)         # Strong corporate red
    c_slate = RGBColor(70, 75, 80)        # Dark slate gray
    c_light_bg = RGBColor(245, 245, 245)  # Light gray for text cards
    c_border = RGBColor(220, 220, 220)    # Soft border gray
    c_white = RGBColor(255, 255, 255)
    c_text_dark = RGBColor(50, 50, 50)
    c_text_light = RGBColor(120, 120, 120)

    # === Slide Header ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(10), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_slate

    # Thin decorative accent line under title
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.1), Inches(12.33), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = c_red
    line.line.fill.background() # No border

    # === Pillar Configuration ===
    pillar_w = Inches(3.4)
    spacing = Inches(0.8)
    start_x = Inches(0.85)
    start_y = Inches(1.8)

    # Content structure
    cards = [
        {"title": "COMPLETED ITEMS", "color": c_slate, "num": "01", "icon_char": "✔"},
        {"title": "OUTSTANDING ITEMS", "color": c_red, "num": "02", "icon_char": "!"},
        {"title": "NEW ITEMS", "color": c_slate, "num": "03", "icon_char": "+"}
    ]

    for i, card in enumerate(cards):
        x = start_x + i * (pillar_w + spacing)

        # 1. Body Block (Drawn first so it sits underneath the header tab)
        body_h = Inches(4.0)
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, start_y + Inches(1.2), pillar_w, body_h)
        body.fill.solid()
        body.fill.fore_color.rgb = c_light_bg
        body.line.color.rgb = c_border
        body.line.width = Pt(1)

        # 2. Custom Freeform Tabbed Header
        hdr_h = Inches(1.2)
        arrow_w = Inches(0.6)
        arrow_h = Inches(0.3)

        # Building a rectangle with a downward-pointing triangle at the bottom center
        builder = FreeformBuilder(slide.shapes)
        builder.move_to(x, start_y)                                       # Top Left
        builder.line_to(x + pillar_w, start_y)                            # Top Right
        builder.line_to(x + pillar_w, start_y + hdr_h)                    # Bottom Right
        builder.line_to(x + pillar_w/2 + arrow_w/2, start_y + hdr_h)      # Tab Right Base
        builder.line_to(x + pillar_w/2, start_y + hdr_h + arrow_h)        # Tab Point (Arrow tip)
        builder.line_to(x + pillar_w/2 - arrow_w/2, start_y + hdr_h)      # Tab Left Base
        builder.line_to(x, start_y + hdr_h)                               # Bottom Left
        builder.line_to(x, start_y)                                       # Back to Top Left

        hdr = builder.convert_to_shape()
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = card["color"]
        hdr.line.fill.background() # No border

        # 3. Icon Placeholder (Hollow Circle in Header)
        icon_size = Inches(0.5)
        icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + pillar_w/2 - icon_size/2, start_y + Inches(0.25), icon_size, icon_size)
        icon.fill.background() # Transparent fill
        icon.line.color.rgb = c_white
        icon.line.width = Pt(1.5)
        
        # Add basic symbol inside icon
        itf = icon.text_frame
        ip = itf.paragraphs[0]
        ip.text = card["icon_char"]
        ip.alignment = PP_ALIGN.CENTER
        ip.font.size = Pt(18)
        ip.font.bold = True
        ip.font.color.rgb = c_white
        itf.vertical_anchor = MSO_ANCHOR.MIDDLE

        # 4. Text Content within Body Block
        tf = body.text_frame
        tf.margin_top = Inches(0.6)  # Leave visual breathing room for the overlapping arrow tab
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)

        p1 = tf.paragraphs[0]
        p1.text = card["title"]
        p1.alignment = PP_ALIGN.CENTER
        p1.font.bold = True
        p1.font.size = Pt(16)
        p1.font.color.rgb = c_text_dark

        p2 = tf.add_paragraph()
        p2.text = "This slide section is entirely editable. Adapt this space to detail the strategic tasks, summarize quarterly progress, and capture your audience's attention."
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(12)
        p2.font.color.rgb = c_text_light
        p2.space_before = Pt(14)

        # 5. Overlapping Number Badge (Visual Anchor at bottom)
        badge_size = Inches(0.6)
        badge_y = start_y + Inches(1.2) + body_h - badge_size/2
        
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + pillar_w/2 - badge_size/2, badge_y, badge_size, badge_size)
        badge.fill.solid()
        badge.fill.fore_color.rgb = c_white
        badge.line.color.rgb = card["color"]
        badge.line.width = Pt(2.5)

        btf = badge.text_frame
        bp = btf.paragraphs[0]
        bp.text = card["num"]
        bp.alignment = PP_ALIGN.CENTER
        bp.font.size = Pt(14)
        bp.font.bold = True
        bp.font.color.rgb = card["color"]
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE

    prs.save(output_pptx_path)
    return output_pptx_path
