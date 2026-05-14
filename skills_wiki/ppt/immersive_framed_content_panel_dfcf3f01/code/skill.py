import os
import urllib.request
from io import BytesIO
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

def _create_fallback_gradient(filepath: str, width: int = 1920, height: int = 1080):
    """Creates a dark tech-themed gradient background if image download fails."""
    img = Image.new('RGB', (width, height))
    draw = ImageDraw.Draw(img)
    color_top = (13, 17, 28)    # Dark Navy
    color_bottom = (40, 20, 60) # Deep Purple
    
    for y in range(height):
        r = int(color_top[0] + (color_bottom[0] - color_top[0]) * y / height)
        g = int(color_top[1] + (color_bottom[1] - color_top[1]) * y / height)
        b = int(color_top[2] + (color_bottom[2] - color_top[2]) * y / height)
        draw.line([(0, y), (width, y)], fill=(r, g, b))
        
    img.save(filepath)

def _add_drop_shadow(shape):
    """Injects OOXML to add a subtle drop shadow to a shape."""
    spPr = shape.element.spPr
    effectLst = etree.SubElement(spPr, qn('a:effectLst'))
    outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
    
    # Shadow parameters
    outerShdw.set('blurRad', '400000')  # Blur radius
    outerShdw.set('dist', '350000')     # Distance
    outerShdw.set('dir', '2700000')     # Angle (45 degrees)
    outerShdw.set('algn', 'tl')         # Alignment
    outerShdw.set('rotWithShape', '0')
    
    srgbClr = etree.SubElement(outerShdw, qn('a:srgbClr'))
    srgbClr.set('val', '000000')        # Black shadow
    alpha = etree.SubElement(srgbClr, qn('a:alpha'))
    alpha.set('val', '40000')           # 40% Opacity

def create_slide(
    output_pptx_path: str,
    title_text: str = "Common Elements of a Pitch",
    body_text: str = "• What do you do?\n• Team\n• Traction\n• Unique Insights\n• Market Size\n• Ask",
    bg_keyword: str = "technology,network,dark",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Immersive Framed Content Panel effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Thematic Background ===
    bg_img_path = "temp_bg_framed_panel.jpg"
    bg_url = f"https://images.unsplash.com/photo-1451187580459-43490279c0fa?q=80&w=1920&auto=format&fit=crop"
    
    try:
        req = urllib.request.Request(bg_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(bg_img_path, 'wb') as out_file:
                out_file.write(response.read())
    except Exception as e:
        print(f"Background download failed, using PIL gradient fallback. Error: {e}")
        _create_fallback_gradient(bg_img_path)

    # Insert background covering the entire slide
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: The White Content Panel ===
    # Create the card leaving a generous margin to frame it with the background
    card_width = Inches(10.0)
    card_height = Inches(6.0)
    card_left = (prs.slide_width - card_width) / 2
    card_top = (prs.slide_height - card_height) / 2
    
    card = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        card_left, card_top, card_width, card_height
    )
    
    # Style the card: pure white, no outline, with shadow
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.fill.background() # No line
    
    _add_drop_shadow(card)

    # === Layer 3: Typography ===
    # Title Box
    title_box = slide.shapes.add_textbox(
        card_left + Inches(0.8), 
        card_top + Inches(0.6), 
        card_width - Inches(1.6), 
        Inches(1.0)
    )
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p_title = title_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Arial'
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(0, 0, 0) # Stark Black

    # Body / Bullets Box
    body_box = slide.shapes.add_textbox(
        card_left + Inches(0.8), 
        card_top + Inches(1.8), 
        card_width - Inches(1.6), 
        card_height - Inches(2.4)
    )
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    
    for i, line in enumerate(body_text.split('\n')):
        p = body_frame.paragraphs[i] if i == 0 else body_frame.add_paragraph()
        p.text = line.replace('• ', '') # Remove manual bullet if exists to apply formatting
        p.font.name = 'Arial'
        p.font.size = Pt(28)
        p.font.color.rgb = RGBColor(80, 80, 80) # Dark Grey for readability
        p.space_after = Pt(14)
        p.level = 0
        
        # Add XML for native bullet points
        pPr = p._pPr
        buFont = etree.SubElement(pPr, qn('a:buFont'))
        buFont.set('typeface', 'Arial')
        buChar = etree.SubElement(pPr, qn('a:buChar'))
        buChar.set('char', '•')

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
