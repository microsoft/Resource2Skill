def create_slide(
    output_pptx_path: str,
    title_text: str = "Serverless Event Processing Architecture",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Dynamic Cloud Architecture Flow Diagram.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.line import MSO_LINE_DASH_STYLE
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    # === Color Palette ===
    COLOR_AWS_NAVY = RGBColor(35, 47, 62)
    COLOR_DATA_CENTER_BG = RGBColor(249, 250, 251)
    COLOR_DATA_CENTER_LINE = RGBColor(107, 114, 128)
    COLOR_FLOW_CYAN = RGBColor(0, 191, 255)
    
    # Domain specific colors
    COLOR_COMPUTE = RGBColor(216, 102, 19)   # Orange
    COLOR_NETWORK = RGBColor(140, 79, 255)   # Purple
    COLOR_DB = RGBColor(51, 85, 218)         # Blue
    COLOR_STORAGE = RGBColor(63, 134, 36)    # Green

    # === Helper Functions ===

    def add_arrowhead(connector):
        """Inject drawingML to add a triangle arrowhead to a line/connector."""
        namespaces = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        a_ln = connector.element.xpath('.//a:ln', namespaces=namespaces)
        if not a_ln:
            return
        a_ln = a_ln[0]
        
        # Create arrowhead element
        tailEnd = etree.Element(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", 
            type="triangle", w="med", len="med"
        )
        
        # Update if exists, otherwise safely insert into schema order
        existing = a_ln.xpath('./a:tailEnd', namespaces=namespaces)
        if existing:
            existing[0].set('type', 'triangle')
        else:
            extLst = a_ln.xpath('./a:extLst', namespaces=namespaces)
            if extLst:
                extLst[0].addprevious(tailEnd)
            else:
                a_ln.append(tailEnd)

    def add_flow_connector(slide, start_x, start_y, end_x, end_y, is_elbow=False):
        """Draws a bright, dashed flow line with an arrowhead."""
        ctype = MSO_CONNECTOR.ELBOW if is_elbow else MSO_CONNECTOR.STRAIGHT
        connector = slide.shapes.add_connector(
            ctype, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y)
        )
        connector.line.color.rgb = COLOR_FLOW_CYAN
        connector.line.width = Pt(2.5)
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        add_arrowhead(connector)
        return connector

    def add_group_box(slide, x, y, w, h, title, line_color, bg_color=None, dashed=False):
        """Draws logical boundary boxes (e.g., VPC, Cloud)."""
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        
        if bg_color:
            rect.fill.solid()
            rect.fill.fore_color.rgb = bg_color
        else:
            rect.fill.background() # Transparent
            
        rect.line.color.rgb = line_color
        rect.line.width = Pt(1.5)
        if dashed:
            rect.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            
        # Top-left Title
        txBox = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.1), Inches(w - 0.2), Inches(0.4))
        p = txBox.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = line_color
        return rect

    def add_service_node(slide, x, y, name, color_rgb):
        """Constructs a vector proxy card for a cloud service icon."""
        w, h = 1.2, 1.4
        
        # Base Card
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(255, 255, 255)
        rect.line.color.rgb = color_rgb
        rect.line.width = Pt(2)
        
        # Proxy Graphic (Circle)
        icon_size = 0.5
        icon_x = x + (w - icon_size) / 2
        icon_y = y + 0.2
        icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(icon_x), Inches(icon_y), Inches(icon_size), Inches(icon_size))
        icon.fill.solid()
        icon.fill.fore_color.rgb = color_rgb
        icon.line.fill.background()
        
        # Text Label
        txBox = slide.shapes.add_textbox(Inches(x), Inches(y + 0.8), Inches(w), Inches(0.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_AWS_NAVY
        p.alignment = PP_ALIGN.CENTER
        return rect

    # === Slide Setup ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # 1. Slide Title & Legend
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(10), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_AWS_NAVY

    legend_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(10), Inches(0.5))
    p_leg = legend_box.text_frame.paragraphs[0]
    p_leg.text = "----▶  Dashed cyan lines denote active data flow paths and logical connections."
    p_leg.font.size = Pt(11)
    p_leg.font.italic = True
    p_leg.font.color.rgb = COLOR_DATA_CENTER_LINE

    # 2. Grouping Boundaries
    # Data Center on Left
    add_group_box(slide, 0.5, 1.5, 2.5, 5.0, "Corporate Data Center", COLOR_DATA_CENTER_LINE, COLOR_DATA_CENTER_BG, dashed=True)
    # AWS Cloud on Right
    add_group_box(slide, 3.5, 1.5, 9.0, 5.0, "AWS Cloud", COLOR_AWS_NAVY, dashed=False)

    # 3. Insert Service Nodes
    add_service_node(slide, 1.15, 3.25, "Client App", COLOR_DATA_CENTER_LINE)
    add_service_node(slide, 4.2, 3.25, "API Gateway", COLOR_NETWORK)
    add_service_node(slide, 6.7, 3.25, "AWS Lambda", COLOR_COMPUTE)
    add_service_node(slide, 9.5, 1.9, "Amazon S3", COLOR_STORAGE)
    add_service_node(slide, 9.5, 4.6, "DynamoDB", COLOR_DB)

    # 4. Draw Flow Connectors
    # Y-center for middle row = 3.25 + (1.4/2) = 3.95
    
    # Client -> API Gateway (Straight)
    add_flow_connector(slide, 2.35, 3.95, 4.2, 3.95, is_elbow=False)
    
    # API Gateway -> Lambda (Straight)
    add_flow_connector(slide, 5.4, 3.95, 6.7, 3.95, is_elbow=False)
    
    # Lambda -> S3 (Elbow)
    # Target Y = 1.9 + 0.7 = 2.6
    add_flow_connector(slide, 7.9, 3.95, 9.5, 2.6, is_elbow=True)
    
    # Lambda -> DynamoDB (Elbow)
    # Target Y = 4.6 + 0.7 = 5.3
    add_flow_connector(slide, 7.9, 3.95, 9.5, 5.3, is_elbow=True)

    # Save and return
    prs.save(output_pptx_path)
    return output_pptx_path
