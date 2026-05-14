import io
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageOps, ImageDraw

def create_filled_icon_image(
    icon_data: bytes,
    percentage: float,
    size: tuple = (300, 300),
    bg_color: tuple = (217, 217, 217),
    fill_color: tuple = (79, 129, 189)
) -> io.BytesIO:
    """
    Creates a partially filled icon image using PIL.

    Args:
        icon_data: The byte data of the source icon (PNG with transparency).
        percentage: The fill percentage (0.0 to 1.0).
        size: The output size of the icon.
        bg_color: RGB tuple for the background (100% total) part of the icon.
        fill_color: RGB tuple for the filled part of the icon.

    Returns:
        A BytesIO object containing the generated PNG image.
    """
    # 1. Load the base icon and ensure it has an alpha channel
    icon = Image.open(io.BytesIO(icon_data)).convert("RGBA").resize(size, Image.LANCZOS)
    
    # 2. Create the background (gray) version
    bg_icon = Image.new("RGBA", icon.size, bg_color)
    bg_icon.putalpha(icon.getchannel('A'))

    # 3. Create the fill (colored) version
    fill_icon_layer = Image.new("RGBA", icon.size, fill_color)
    fill_icon_layer.putalpha(icon.getchannel('A'))

    # 4. Calculate the crop height based on the percentage
    # We want to keep the bottom part of the image, so we calculate the top cutoff point
    height = icon.size[1]
    cutoff_y = int(height * (1 - percentage))

    # 5. Create a mask to clip the top portion of the fill icon
    mask = Image.new("L", icon.size, 0)
    draw = ImageDraw.Draw(mask)
    draw.rectangle([(0, cutoff_y), (icon.size[0], height)], fill=255)

    # 6. Composite the final image
    # Start with the background icon, then paste the masked fill icon on top
    final_image = bg_icon.copy()
    final_image.paste(fill_icon_layer, (0, 0), mask)

    # 7. Save to a byte buffer
    img_buffer = io.BytesIO()
    final_image.save(img_buffer, format='PNG')
    img_buffer.seek(0)
    
    return img_buffer

def create_slide(
    output_pptx_path: str,
    title_text: str = "Image Percentage Fill Chart",
    data: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with Pictogram Percentage Fill charts.

    Returns: path to the saved PPTX file.
    """
    if data is None:
        data = [
            {"label": "Metric A", "value": 0.80, "color": (247, 150, 70)},
            {"label": "Metric B", "value": 0.65, "color": (79, 129, 189)},
            {"label": "Metric C", "value": 0.52, "color": (155, 187, 89)},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set a light gray background for contrast
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 240, 240)

    # Add a title
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.33), Inches(1))
    title_tf = title_shape.text_frame
    title_tf.text = title_text
    p = title_tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(89, 89, 89)
    p.alignment = PP_ALIGN.CENTER
    
    # Download a default icon (human silhouette)
    try:
        icon_url = "https://www.flaticon.com/download/icon/3135715?format=png&size=512"
        # Flaticon requires a proper User-Agent header
        req = urllib.request.Request(icon_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            icon_data = response.read()
    except Exception as e:
        print(f"Failed to download icon, using a placeholder. Error: {e}")
        # Create a simple circle as a fallback
        img = Image.new('RGBA', (512, 512), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.ellipse((10, 10, 502, 502), fill=(128, 128, 128, 255))
        fallback_buffer = io.BytesIO()
        img.save(fallback_buffer, format='PNG')
        fallback_buffer.seek(0)
        icon_data = fallback_buffer.read()

    # --- Create and place pictograms ---
    total_items = len(data)
    total_width = Inches(12)
    item_width = total_width / total_items
    start_left = (prs.slide_width - total_width) / 2
    
    for i, item in enumerate(data):
        # Generate the filled icon image
        filled_icon_buffer = create_filled_icon_image(
            icon_data=icon_data,
            percentage=item["value"],
            fill_color=item["color"],
            bg_color=(220, 220, 220)
        )
        
        # Add the image to the slide
        icon_size = Inches(2.5)
        left_pos = start_left + (i * item_width) + (item_width - icon_size) / 2
        top_pos = Inches(2.5)
        slide.shapes.add_picture(filled_icon_buffer, left_pos, top_pos, height=icon_size)

        # Add the percentage label
        label_box = slide.shapes.add_textbox(left_pos, top_pos + icon_size + Inches(0.2), icon_size, Inches(0.5))
        label_tf = label_box.text_frame
        label_tf.text = f'{item["value"]:.0%}'
        p_label = label_tf.paragraphs[0]
        p_label.font.size = Pt(32)
        p_label.font.bold = True
        p_label.font.color.rgb = RGBColor.from_string("595959")
        p_label.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
