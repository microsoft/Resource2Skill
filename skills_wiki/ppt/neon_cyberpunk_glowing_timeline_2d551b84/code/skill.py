def create_slide(
    output_pptx_path: str,
    title_text: str = "TIMELINE SLIDE",
    milestones: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neon Cyberpunk Glowing Timeline effect.
    """
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    import copy

    # Helper function to inject XML Glow Effect
    def apply_glow_effect(shape, hex_color: str, radius_pt: int = 15, alpha_pct: int = 60):
        spPr = shape.element.spPr
        
        # Check if effectLst exists, if not, create it
        effectLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        if effectLst is None:
            effectLst = OxmlElement('a:effectLst')
            spPr.append(effectLst)
            
        # Create glow element
        glow = OxmlElement('a:glow')
        glow.set('rad', str(int(radius_pt * 12700))) # Convert points to EMUs
        
        # Create color element
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', hex_color.replace('#', ''))
        
        # Create opacity (alpha) element
        alpha = OxmlElement('a:alpha')
        alpha.set('val', str(int(alpha_pct * 1000))) # 60% = 60000
        
        srgbClr.append(alpha)
        glow.append(srgbClr)
        effectLst.append(glow)

    # Helper to convert hex to RGB tuple
    def hex_to_rgb(hex_str):
        hex_str = hex_str.lstrip('#')
        return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

    # Default Data
    if not milestones:
        milestones = [
            {"year": "2019", "text": "Put your text here. This is where\nyour idea begins to take shape.", "color": "#FF3B30"}, # Red
            {"year": "2020", "text": "Put your text here. This is where\nyour idea begins to take shape.", "color": "#FFCC00"}, # Yellow
            {"year": "2021", "text": "Put your text here. This is where\nyour idea begins to take shape.", "color": "#33CCFF"}, # Cyan
            {"year": "2022", "text": "Put your text here. This is where\nyour idea begins to take shape.", "color": "#34C759"}, # Green
            {"year": "2023", "text": "Put your text here. This is where\nyour idea begins to take shape.", "color": "#AF52DE"}, # Purple
        ]

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # 1. Slide Background (Dark Charcoal)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 30, 30)

    # 2. Main Title (Glowing)
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.5), Inches(9.333), Inches(1))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = RGBColor(255, 255, 255)
    # Apply glow to title text box shape (gives a subtle backing glow)
    apply_glow_effect(title_box, "FFFFFF", radius_pt=10, alpha_pct=30)

    # 3. Horizontal Axis Line
    axis_y = Inches(3.75)
    axis_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), axis_y, Inches(11.333), Pt(1.5))
    axis_line.fill.solid()
    axis_line.fill.fore_color.rgb = RGBColor(255, 255, 255)
    axis_line.line.fill.background() # No border

    # 4. Generate Timeline Nodes
    num_nodes = len(milestones)
    start_x = Inches(2.0)
    end_x = Inches(11.333)
    step_x = (end_x - start_x) / (num_nodes - 1) if num_nodes > 1 else 0

    stem_length = Inches(1.2)
    inner_radius = Inches(0.08)
    outer_radius = Inches(0.25)

    for i, node in enumerate(milestones):
        cx = start_x + (i * step_x)
        is_top = (i % 2 == 0) # Alternating logic
        
        node_color_hex = node["color"]
        node_color_rgb = RGBColor(*hex_to_rgb(node_color_hex))
        
        # Stem vertical position
        node_cy = axis_y - stem_length if is_top else axis_y + stem_length
        
        # A. Vertical Stem Line
        stem_top = axis_y - stem_length if is_top else axis_y
        stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx - Pt(0.75), stem_top, Pt(1.5), stem_length)
        stem.fill.solid()
        stem.fill.fore_color.rgb = node_color_rgb
        stem.line.fill.background()
        
        # B. Outer Hollow Circle
        outer_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            cx - outer_radius, node_cy - outer_radius, 
            outer_radius * 2, outer_radius * 2
        )
        outer_circle.fill.background() # Transparent fill
        outer_circle.line.color.rgb = node_color_rgb
        outer_circle.line.width = Pt(2)
        
        # C. Inner Glowing Core
        inner_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            cx - inner_radius, node_cy - inner_radius, 
            inner_radius * 2, inner_radius * 2
        )
        inner_circle.fill.solid()
        inner_circle.fill.fore_color.rgb = node_color_rgb
        inner_circle.line.fill.background() # No border
        apply_glow_effect(inner_circle, node_color_hex, radius_pt=12, alpha_pct=50)

        # D. Text Boxes
        text_width = Inches(2.2)
        
        # Calculate Y positions based on Top/Bottom placement
        if is_top:
            year_y = node_cy - Inches(0.5)
            desc_y = year_y - Inches(0.7)
        else:
            year_y = node_cy + Inches(0.2)
            desc_y = year_y + Inches(0.3)
            
        # Year Text
        year_box = slide.shapes.add_textbox(cx - (text_width/2), year_y, text_width, Inches(0.4))
        p_year = year_box.text_frame.paragraphs[0]
        p_year.text = node["year"]
        p_year.alignment = PP_ALIGN.CENTER
        p_year.font.size = Pt(20)
        p_year.font.bold = True
        p_year.font.name = "Arial"
        p_year.font.color.rgb = node_color_rgb
        
        # Description Text
        desc_box = slide.shapes.add_textbox(cx - (text_width/2), desc_y, text_width, Inches(0.8))
        desc_box.text_frame.word_wrap = True
        p_desc = desc_box.text_frame.paragraphs[0]
        p_desc.text = node["text"]
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.font.size = Pt(11)
        p_desc.font.name = "Arial"
        p_desc.font.color.rgb = RGBColor(230, 230, 230)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
