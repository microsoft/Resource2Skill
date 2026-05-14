def create_slide(
    output_pptx_path: str,
    title_text: str = "Net Sales grew by +400% from 2015 to 2021",
    chart_data: dict = None,
    highlight_index: int = -1,
    accent_color_rgb: tuple = (47, 82, 143),
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a single, clean, consulting-style column chart,
    highlighting a specific data point.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        title_text (str): The main title for the slide (action title).
        chart_data (dict): Data for the chart, e.g., {'2015': 3, '2016': 5, ...}.
        highlight_index (int): The zero-based index of the data point to highlight.
                               Defaults to -1 (the last point).
        accent_color_rgb (tuple): The RGB tuple for the highlighted bar.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_THEME_COLOR

    # --- Default Data if none provided ---
    if chart_data is None:
        chart_data = {
            '2015': 3, '2016': 5, '2017': 9, '2018': 12,
            '2019': 8, '2020': 13, '2021': 15
        }

    # --- Presentation and Slide Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # --- Add Slide Title (Action Title) ---
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.8))
    title_frame = title_shape.text_frame
    p = title_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(89, 89, 89)

    # --- Chart Data Preparation ---
    categories = list(chart_data.keys())
    values = list(chart_data.values())

    chart_data_obj = ChartData()
    chart_data_obj.categories = categories
    chart_data_obj.add_series('Net Sales', values)

    # --- Add and Position Chart ---
    x, y, cx, cy = Inches(0.5), Inches(1.5), Inches(12), Inches(5)
    chart_graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data_obj
    )
    chart = chart_graphic_frame.chart

    # --- Style the Chart (Data-Ink Maximization) ---
    plot = chart.plots[0]
    series = plot.series[0]

    # 1. Remove Chart Junk
    chart.has_legend = False
    chart.value_axis.visible = False
    if chart.value_axis.has_major_gridlines:
        chart.value_axis.major_gridlines.format.line.fill.background() # Effectively makes it invisible

    # 2. Style Category Axis (X-axis)
    cat_axis = chart.category_axis
    cat_axis.format.line.solid()
    cat_axis.format.line.color.rgb = RGBColor(191, 191, 191)
    cat_axis.format.line.width = Pt(0.75)
    cat_axis.tick_labels.font.size = Pt(12)
    cat_axis.tick_labels.font.color.rgb = RGBColor(89, 89, 89)

    # 3. Add and Style Data Labels
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    data_labels.font.size = Pt(12)
    data_labels.font.color.rgb = RGBColor(89, 89, 89)

    # 4. Style Bars (Default and Highlight)
    default_fill_color = RGBColor(217, 217, 217)
    highlight_fill_color = RGBColor.from_rgb(*accent_color_rgb)

    # Handle negative index for "last item"
    if highlight_index < 0:
        highlight_index = len(values) + highlight_index

    for i, point in enumerate(series.points):
        fill = point.format.fill
        if i == highlight_index:
            fill.solid()
            fill.fore_color.rgb = highlight_fill_color
        else:
            fill.solid()
            fill.fore_color.rgb = default_fill_color
        # Remove outline from bars
        point.format.line.fill.background()
        
    # 5. Add Growth Arrow Annotation
    first_val = values[0]
    last_val = values[-1]
    if first_val > 0:
        growth_pct = (last_val / first_val - 1) * 100
        growth_text = f"+{growth_pct:.0f}%"
        
        # Manually create an arrow and text (emulating think-cell's difference arrow)
        # Position it relative to the chart frame
        arrow_start_y = y + cy - Inches(0.5)
        arrow_end_y = arrow_start_y - Inches(1.5)
        arrow_x = x + cx - Inches(0.5)

        arrow_shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, arrow_x, arrow_end_y, Inches(0.2), Inches(1.5))
        arrow_shape.rotation = 180.0
        fill = arrow_shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(128, 128, 128)
        line = arrow_shape.line
        line.fill.background()

        label_shape = slide.shapes.add_textbox(arrow_x - Inches(0.4), arrow_start_y - Inches(0.9), Inches(1.0), Inches(0.4))
        label_frame = label_shape.text_frame
        label_frame.clear()
        p = label_frame.paragraphs[0]
        p.text = growth_text
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(89, 89, 89)
        p.font.bold = True

    # --- Save Presentation ---
    prs.save(output_pptx_path)
    return output_pptx_path
