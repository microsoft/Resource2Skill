import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "ANIMATED COMPARISON SLIDE DESIGN",
    data_points: list = None,
    **kwargs
) -> str:
    """
    Creates a sleek, minimalist infographic progress dashboard slide.
    
    :param output_pptx_path: Path to save the presentation.
    :param title_text: Main slide title.
    :param data_points: List of dictionaries containing 'title', 'desc', 'pct', and 'color' (RGB tuple).
    """
    # Default data if none provided
    if not data_points:
        data_points = [
            {
                "title": "TASK ONE", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above.", 
                "pct": 75, 
                "color": (214, 40, 40)   # Red
            },
            {
                "title": "TASK TWO", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above.", 
                "pct": 50, 
                "color": (42, 157, 143)  # Teal
            },
            {
                "title": "TASK THREE", 
                "desc": "You can write some text here\nto describe any information\nabout the heading above.", 
                "pct": 25, 
                "color": (0, 119, 182)   # Blue
            }
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Constants for layout
    SLATE_BLACK = RGBColor(51, 51, 51)
    GREY_TEXT = RGBColor(128, 128, 128)
    TRACK_COLOR = RGBColor(230, 230, 230)
    TICK_COLOR = RGBColor(180, 180, 180)
    
    # 1. Main Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Montserrat" # Will fallback to Arial if not installed
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = SLATE_BLACK

    # Coordinates and dimensions for the bars
    start_y = 2.0
    spacing_y = 1.6
    
    text_x = Inches(1.0)
    text_w = Inches(3.0)
    
    track_x = Inches(4.5)
    track_w_inches = 7.5
    track_w = Inches(track_w_inches)
    track_h = Inches(0.12)
    
    # Generate the 3 tasks
    for i, data in enumerate(data_points):
        current_y = Inches(start_y + (i * spacing_y))
        accent_color = RGBColor(*data["color"])
        
        # --- Left Side: Text ---
        # Category Title
        cat_box = slide.shapes.add_textbox(text_x, current_y - Inches(0.2), text_w, Inches(0.4))
        cat_tf = cat_box.text_frame
        cat_p = cat_tf.paragraphs[0]
        cat_p.text = data["title"]
        cat_p.font.name = "Montserrat"
        cat_p.font.size = Pt(20)
        cat_p.font.bold = True
        cat_p.font.color.rgb = accent_color
        
        # Category Description
        desc_box = slide.shapes.add_textbox(text_x, current_y + Inches(0.15), text_w, Inches(0.8))
        desc_tf = desc_box.text_frame
        desc_tf.word_wrap = True
        desc_p = desc_tf.paragraphs[0]
        desc_p.text = data["desc"]
        desc_p.font.name = "Arial"
        desc_p.font.size = Pt(11)
        desc_p.font.color.rgb = GREY_TEXT
        
        # --- Right Side: The Data Track ---
        # 1. Background Track
        track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, track_x, current_y + Inches(0.15), track_w, track_h)
        track.fill.solid()
        track.fill.fore_color.rgb = TRACK_COLOR
        track.line.color.rgb = TRACK_COLOR
        
        # 2. Tick marks (0%, 25%, 50%, 75%, 100%)
        num_segments = 4
        segment_w = track_w_inches / num_segments
        for j in range(num_segments + 1):
            tick_x = track_x + Inches(j * segment_w)
            tick = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, tick_x, current_y + Inches(0.1), Inches(0.03), Inches(0.22))
            tick.fill.solid()
            tick.fill.fore_color.rgb = TICK_COLOR
            tick.line.color.rgb = TICK_COLOR
            
        # 3. Colored Fill Bar
        fill_w = Inches(track_w_inches * (data["pct"] / 100.0))
        # Ensure it has a tiny minimum width so it renders even at 0%
        fill_w = max(fill_w, Inches(0.05))
        
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, track_x, current_y + Inches(0.15), fill_w, track_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent_color
        bar.line.color.rgb = accent_color
        
        # 4. Data Label (The Custom "Pin" Callout)
        pin_x_center = track_x + fill_w
        
        # Pin Triangle (Pointer)
        tri_w = Inches(0.2)
        tri_h = Inches(0.15)
        tri = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE, 
            pin_x_center - (tri_w/2), 
            current_y - Inches(0.02), 
            tri_w, 
            tri_h
        )
        tri.rotation = 180 # Point downwards
        tri.fill.solid()
        tri.fill.fore_color.rgb = accent_color
        tri.line.color.rgb = accent_color
        
        # Pin Rectangle (Body)
        rect_w = Inches(0.8)
        rect_h = Inches(0.4)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            pin_x_center - (rect_w/2), 
            current_y - Inches(0.4), 
            rect_w, 
            rect_h
        )
        rect.adjustments[0] = 0.2 # Slight rounding
        rect.fill.solid()
        rect.fill.fore_color.rgb = accent_color
        rect.line.color.rgb = accent_color
        
        # Text inside the Pin
        rect_tf = rect.text_frame
        rect_tf.margin_left = 0
        rect_tf.margin_right = 0
        rect_tf.margin_top = 0
        rect_tf.margin_bottom = 0
        rect_p = rect_tf.paragraphs[0]
        rect_p.alignment = PP_ALIGN.CENTER
        rect_p.text = f"{data['pct']}%"
        rect_p.font.name = "Arial"
        rect_p.font.size = Pt(12)
        rect_p.font.bold = True
        rect_p.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("animated_comparison_dashboard.pptx")
