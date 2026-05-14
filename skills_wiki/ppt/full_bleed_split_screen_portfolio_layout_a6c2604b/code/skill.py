def create_slide(
    output_pptx_path: str,
    title_text: str = "Balloon Dog",
    body_text: str = "The Balloon Dog is one of Jeff Koons' most famous sculptures. Completed in the 1990s, the ten-foot-tall piece is made out of mirror-polished stainless steel with a transparent color coating. It represents childhood innocence and the joy of simple celebrations.",
    image_keyword: str = "art,sculpture",
    panel_color: tuple = (40, 44, 52),      # Dark Slate RGB
    text_color: tuple = (255, 255, 255),    # White RGB
    accent_color: tuple = (255, 107, 107),  # Coral Red RGB
    **kwargs,
) -> str:
    """
    Creates a PPTX file featuring a Full-Bleed Split-Screen layout.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw

    # Initialize Presentation (16:9 Widescreen)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Left Content Panel ===
    left_panel_width = Inches(5.0)
    
    # Solid background block
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, left_panel_width, prs.slide_height
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(*panel_color)
    panel.line.fill.background() # Remove border

    # === Layer 2: Typography & Details ===
    # Decorative accent line above title
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.3), Inches(0.6), Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(*accent_color)
    accent.line.fill.background()

    # Title Box
    title_box = slide.shapes.add_textbox(
        Inches(0.4), Inches(1.5), Inches(4.2), Inches(1.0)
    )
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(*text_color)
    p_title.font.name = "Arial"

    # Body Text Box
    body_box = slide.shapes.add_textbox(
        Inches(0.4), Inches(2.8), Inches(4.0), Inches(4.0)
    )
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(*text_color)
    p_body.font.name = "Arial"

    # === Layer 3: Right Hero Image (with auto-crop-to-fill) ===
    img_path = "temp_split_img.jpg"
    target_w = prs.slide_width - left_panel_width
    target_h = prs.slide_height

    # 1. Acquire Image (Download or Fallback)
    try:
        url = f"https://picsum.photos/1200/800?random=1"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Generate synthetic fallback image if network fails
        img = Image.new('RGB', (1200, 800), color=(220, 224, 230))
        draw = ImageDraw.Draw(img)
        # Draw some abstract aesthetic shapes
        draw.ellipse([-200, -200, 600, 600], fill=(200, 205, 215))
        draw.rectangle([800, 400, 1400, 1000], fill=(205, 210, 220))
        img.save(img_path)

    # 2. Calculate Crop Percentages to prevent distortion
    with Image.open(img_path) as img:
        img_w, img_h = img.size
    
    img_ar = img_w / img_h
    target_ar = target_w / target_h

    # Insert picture without constraints initially
    pic = slide.shapes.add_picture(img_path, left_panel_width, 0)

    # Apply mathematically calculated crops
    if img_ar > target_ar:
        # Image is wider than target area -> crop sides
        crop_fraction = 1.0 - (target_ar / img_ar)
        pic.crop_left = crop_fraction / 2
        pic.crop_right = crop_fraction / 2
    elif img_ar < target_ar:
        # Image is taller than target area -> crop top/bottom
        crop_fraction = 1.0 - (img_ar / target_ar)
        pic.crop_top = crop_fraction / 2
        pic.crop_bottom = crop_fraction / 2

    # 3. Force the final cropped bounding box to fit the right pane perfectly
    pic.width = int(target_w)
    pic.height = int(target_h)

    # Cleanup temporary file
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
