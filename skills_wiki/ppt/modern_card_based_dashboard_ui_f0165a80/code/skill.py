import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from lxml import etree

def create_slide(
    output_pptx_path: str,
    title_text: str = "Performance Tracking 2024",
    bg_color: tuple = (15, 17, 26),     # Deep navy/grey background
    card_color: tuple = (30, 34, 51),   # Elevated card color
    text_main: tuple = (255, 255, 255), # White text
    text_muted: tuple = (150, 160, 180),# Muted grey text
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Modern Card-Based Dashboard UI" 
    using a strict mathematical grid, rounded corners, and XML-injected custom shadows.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # --- Step 5a: Color Theme & Canvas Background ---
    # Set slide background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # --- Helper Function: Step 3 (Rounded Corners) & Step 4 (Shadows) ---
    def add_dashboard_card(left, top, width, height):
        """Creates a rounded rectangle card with a custom soft shadow."""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        
        # Step 5b: Card Color & Remove Borders
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*card_color)
        shape.line.fill.background() # No line
        
        # Step 3: Rounded Corners (adjusting the handle)
        # The adjustment value determines radius. Smaller number = smaller radius.
        shape.adjustments[0] = 0.05 
        
        # Step 4: Borders & Shadows (via lxml injection)
        # We inject a custom outer shadow: black, 15% opacity, soft blur, downward angle
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        outerShdw = etree.SubElement(
            effectLst, 
            "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw",
            blurRad="150000",   # Soft blur
            dist="40000",       # Slight offset
            dir="5400000",      # 90 degrees (downward)
            algn="tl", 
            rotWithShape="0"
        )
        srgbClr = etree.SubElement(
            outerShdw, 
            "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", 
            val="000000"        # Black shadow
        )
        # 15% opacity (alpha val is in 1/1000th of a percent, so 15000 = 15%)
        etree.SubElement(
            srgbClr, 
            "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", 
            val="15000"
        )
        
        return shape

    def add_card_content(card_shape, value_text, label_text):
        """Adds text to the card demonstrating Step 2 (White Space/Padding)"""
        text_frame = card_shape.text_frame
        text_frame.clear()
        # Internal padding (Step 2: White space)
        text_frame.margin_left = Inches(0.2)
        text_frame.margin_top = Inches(0.15)
        text_frame.vertical_anchor = 3 # Middle
        
        # Value
        p = text_frame.paragraphs[0]
        p.text = value_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*text_main)
        
        # Label
        p2 = text_frame.add_paragraph()
        p2.text = label_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(*text_muted)

    # --- Step 1: Alignment & Sizing (Mathematical Grid Setup) ---
    slide_w = 13.333
    slide_h = 7.5
    margin = 0.5
    gutter = 0.2

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(margin), Inches(margin), Inches(5), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*text_main)

    # --- Grid Calculations ---
    # Row 1: KPI Cards (4 columns)
    row1_top = margin + 0.8
    row1_height = 1.2
    cols_r1 = 4
    card_w_r1 = (slide_w - (2 * margin) - ((cols_r1 - 1) * gutter)) / cols_r1

    kpis = [
        ("$95M", "Total Sales"),
        ("1,535", "Total Insured"),
        ("1,481", "Total Uninsured"),
        ("46 Years", "Average Age")
    ]

    for i in range(cols_r1):
        left = margin + (i * (card_w_r1 + gutter))
        card = add_dashboard_card(Inches(left), Inches(row1_top), Inches(card_w_r1), Inches(row1_height))
        add_card_content(card, kpis[i][0], kpis[i][1])

    # Row 2: Medium Charts (3 columns)
    row2_top = row1_top + row1_height + gutter
    row2_height = 2.0
    cols_r2 = 3
    card_w_r2 = (slide_w - (2 * margin) - ((cols_r2 - 1) * gutter)) / cols_r2

    for i in range(cols_r2):
        left = margin + (i * (card_w_r2 + gutter))
        card = add_dashboard_card(Inches(left), Inches(row2_top), Inches(card_w_r2), Inches(row2_height))
        add_card_content(card, f"Chart Area {i+1}", "Data visualization placeholder")

    # Row 3: Large Charts (2 columns)
    row3_top = row2_top + row2_height + gutter
    row3_height = slide_h - row3_top - margin
    cols_r3 = 2
    card_w_r3 = (slide_w - (2 * margin) - ((cols_r3 - 1) * gutter)) / cols_r3

    for i in range(cols_r3):
        left = margin + (i * (card_w_r3 + gutter))
        card = add_dashboard_card(Inches(left), Inches(row3_top), Inches(card_w_r3), Inches(row3_height))
        add_card_content(card, f"Detailed Analysis {i+1}", "Secondary data view placeholder")

    prs.save(output_pptx_path)
    return output_pptx_path
