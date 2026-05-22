def create_slide(
    output_pptx_path: str,
    title_text: str = "This is a sample",
    bg_palette: str = "technology",  # Kept for signature compatibility
    accent_color: tuple = (50, 205, 50),  # Vibrant Green
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Split-Panel Object Reveal layout.
    Provides the visual staging used for simultaneous animation tutorials.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)

    # === Layer 1: Left Text Block ===
    # Draw a vibrant green container on the left side
    left_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.0), Inches(1.75), Inches(5.5), Inches(4.0)
    )
    left_shape.fill.solid()
    left_shape.fill.fore_color.rgb = RGBColor(*accent_color)
    left_shape.line.fill.background()  # Remove border

    tf = left_shape.text_frame
    tf.word_wrap = True
    
    # Add repeated bold text pattern as seen in the tutorial
    for i in range(3):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = title_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial Black"
        p.font.size = Pt(40)
        p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Right Subject Image ===
    # Attempt to download a transparent PNG (airplane silhouette)
    image_url = "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Airplane_silhouette.svg/512px-Airplane_silhouette.svg.png"
    image_path = "temp_subject_image.png"
    
    try:
        # Request with headers to avoid basic scraping blocks
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(image_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Fallback to PIL: Generate a stylized paper plane with a transparent background
        img = Image.new("RGBA", (600, 400), (255, 255, 255, 0))
        draw = ImageDraw.Draw(img)
        
        # Main wing body
        draw.polygon([(100, 200), (500, 100), (300, 350), (250, 250)], fill=(100, 150, 255, 255))
        # Shadow/lower wing flap
        draw.polygon([(500, 100), (250, 250), (200, 300)], fill=(70, 100, 180, 255))
        img.save(image_path)

    # Insert the transparent image on the right side of the split layout
    slide.shapes.add_picture(image_path, Inches(7.5), Inches(2.0), width=Inches(4.5))

    # Clean up temporary asset
    if os.path.exists(image_path):
        os.remove(image_path)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
