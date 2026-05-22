import io
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_transparent_silhouette() -> io.BytesIO:
    """
    Generates a stylized RGBA image of a person (silhouette) with a transparent background.
    This simulates the "cutout" image effect without relying on external URLs.
    """
    img = Image.new("RGBA", (800, 1000), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    
    # Draw a stylized person (head and shoulders/blazer)
    # Blazer color to represent the "clothing" that dictates the slide theme
    blazer_color = (200, 100, 120, 255) 
    skin_color = (240, 210, 190, 255)
    
    # Body / Blazer (Slanted shoulders)
    draw.polygon([(100, 1000), (300, 500), (500, 500), (700, 1000)], fill=blazer_color)
    # Head/Neck
    draw.rectangle([(360, 400), (440, 550)], fill=skin_color)
    draw.ellipse([(300, 200), (500, 450)], fill=skin_color)
    
    # Hair style
    draw.ellipse([(280, 180), (520, 350)], fill=(220, 180, 120, 255))
    
    # Shadow/Fold on blazer for depth
    draw.polygon([(300, 500), (400, 800), (500, 500)], fill=(180, 80, 100, 255))

    image_stream = io.BytesIO()
    img.save(image_stream, format='PNG')
    image_stream.seek(0)
    return image_stream

def create_slide(
    output_pptx_path: str = "Dynamic_Portrait_Profile.pptx",
    name_text: str = "Jenny Davis",
    role_text: str = "CREATIVE DIRECTOR",
    base_color: tuple = (226, 169, 184),   # Soft Mauve/Pink
    accent_shape_color: tuple = (205, 137, 153), # Darker Mauve
    text_dark: tuple = (50, 50, 50),
    text_light: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dynamic Subject-Harmonized Portrait" visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Solid Harmonized Background ===
    bg_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_rect.fill.solid()
    bg_rect.fill.fore_color.rgb = RGBColor(*base_color)
    bg_rect.line.fill.background() # No line

    # === Layer 2: Dynamic Diagonal Shapes ===
    # Adds visual rhythm and motion behind the subject
    diag1 = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(4), Inches(-1), Inches(4), Inches(10)
    )
    diag1.rotation = 15
    diag1.fill.solid()
    diag1.fill.fore_color.rgb = RGBColor(*accent_shape_color)
    diag1.line.fill.background()

    diag2 = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(7), Inches(-1), Inches(1), Inches(10)
    )
    diag2.rotation = 15
    diag2.fill.solid()
    diag2.fill.fore_color.rgb = RGBColor(*accent_shape_color)
    diag2.line.fill.background()

    # === Layer 3: Subject Image (Transparent Cutout) ===
    # Simulating the background removal by inserting an RGBA PNG
    img_stream = create_transparent_silhouette()
    slide.shapes.add_picture(
        img_stream, Inches(0.5), Inches(1.5), width=Inches(4.5)
    )

    # === Layer 4: Typography and Details ===
    
    # Name (Strong Contrast)
    name_box = slide.shapes.add_textbox(Inches(6.0), Inches(1.2), Inches(5), Inches(1))
    tf_name = name_box.text_frame
    p = tf_name.paragraphs[0]
    p.text = name_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*text_dark)

    # Role (Weak Contrast / Secondary)
    role_box = slide.shapes.add_textbox(Inches(6.05), Inches(1.9), Inches(5), Inches(0.5))
    tf_role = role_box.text_frame
    p2 = tf_role.paragraphs[0]
    p2.text = role_text
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(*text_light)

    # Decorative Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.05), Inches(2.5), Inches(4), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*text_light)
    line.line.fill.background()

    # Detail List (Bullet points with custom circles)
    details = [
        "Graduated from Columbia University",
        "Winner of IF Design Gold Award",
        "Served 50+ Fortune 500 companies",
        "Published multiple personal design portfolios"
    ]
    
    start_y = 3.0
    for i, detail in enumerate(details):
        y_pos = start_y + (i * 0.6)
        
        # Custom concentric circle bullet
        outer = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.05), Inches(y_pos + 0.05), Inches(0.15), Inches(0.15))
        outer.fill.background() # transparent
        outer.line.color.rgb = RGBColor(*text_light)
        outer.line.width = Pt(1)
        
        inner = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.08), Inches(y_pos + 0.08), Inches(0.09), Inches(0.09))
        inner.fill.solid()
        inner.fill.fore_color.rgb = RGBColor(*text_light)
        inner.line.fill.background()

        # Text
        det_box = slide.shapes.add_textbox(Inches(6.3), Inches(y_pos - 0.05), Inches(6), Inches(0.4))
        p_det = det_box.text_frame.paragraphs[0]
        p_det.text = detail
        p_det.font.size = Pt(14)
        p_det.font.color.rgb = RGBColor(*text_dark)
        
    # Logo Placeholder Top Right
    logo = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(12.0), Inches(0.5), Inches(0.4), Inches(0.4))
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor(*text_dark)
    logo.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path

if __name__ == "__main__":
    create_slide()
