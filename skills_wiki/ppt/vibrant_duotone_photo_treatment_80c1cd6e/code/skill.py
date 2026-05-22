import os
import urllib.request
from io import BytesIO
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn

def add_drop_shadow(shape):
    """
    Injects Open XML to add a subtle drop shadow to a text box/shape.
    Ensures text readability over the duotone background.
    """
    spPr = shape.element.spPr
    effectLst = etree.SubElement(spPr, qn('a:effectLst'))
    outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
    outerShdw.set('blurRad', '101600')     # 8 pt blur
    outerShdw.set('dist', '38100')         # 3 pt distance
    outerShdw.set('dir', '2700000')        # 45 degrees angle
    outerShdw.set('algn', 'tl')
    
    srgbClr = etree.SubElement(outerShdw, qn('a:srgbClr'))
    srgbClr.set('val', '000000')           # Black shadow
    alpha = etree.SubElement(srgbClr, qn('a:alpha'))
    alpha.set('val', '40000')              # 40% opacity


def create_slide(
    output_pptx_path: str,
    title_text: str = "VIBRANT DUOTONE",
    body_text: str = "Unify your photography with bold, branded color treatments.",
    theme_keyword: str = "fashion,portrait",
    shadow_color: tuple = (29, 0, 168),     # Dark Blue/Violet
    highlight_color: tuple = (255, 226, 0), # Bright Yellow
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Duotone Photo Treatment visual effect.
    """
    
    # 1. Download source image
    # Using a reliable Unsplash image ID for a high-quality portrait (similar to tutorial)
    img_url = "https://images.unsplash.com/photo-1517841905240-472988babdf9?q=80&w=1920&auto=format&fit=crop"
    
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert("RGB")
    except Exception as e:
        print(f"Image download failed: {e}. Generating fallback gradient.")
        # Fallback: Create a smooth gradient image so the duotone effect is still visible
        img = Image.new('RGB', (1920, 1080))
        draw = ImageDraw.Draw(img)
        for i in range(1080):
            val = int(255 * (i / 1080))
            draw.line([(0, i), (1920, i)], fill=(val, val, val))

    # 2. Process Image: Apply the Duotone Effect
    # Convert image to grayscale to map luminance values
    gray_img = img.convert('L')
    
    # colorize maps darkest pixels to 'black' param, and lightest to 'white' param
    # This directly mimics the two solid color adjustment layers from the tutorial
    duotone_img = ImageOps.colorize(gray_img, black=shadow_color, white=highlight_color)
    
    temp_img_path = "temp_duotone_bg.jpg"
    duotone_img.save(temp_img_path, quality=95)

    # 3. Assemble Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add Duotone Background (Full Bleed)
    slide.shapes.add_picture(temp_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # 4. Add Typography
    # Title Text
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
    add_drop_shadow(title_box) # Apply lxml shadow
    
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial Black"
    p.font.size = Pt(72)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.bold = True

    # Subtitle Text
    sub_box = slide.shapes.add_textbox(Inches(2), Inches(4.2), Inches(9.333), Inches(1))
    add_drop_shadow(sub_box) # Apply lxml shadow
    
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(24)
    p_sub.font.color.rgb = RGBColor(255, 255, 255)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
        
    return output_pptx_path
