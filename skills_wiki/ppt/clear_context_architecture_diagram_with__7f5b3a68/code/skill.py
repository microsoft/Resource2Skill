def create_slide(
    output_pptx_path: str,
    title_text: str = "Story-Driven Architecture Diagram",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the clear-context architecture diagram style
    featuring logical node groups, standard shapes, connection flows, and a legend.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.line import MSO_LINE

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Colors
    c_bg = RGBColor(250, 252, 255)
    c_node_fill = RGBColor(235, 243, 250)
    c_node_border = RGBColor(50, 90, 140)
    c_boundary = RGBColor(180, 190, 200)
    c_text = RGBColor(30, 40, 50)
    c_alert = RGBColor(220, 60, 60)

    # Set background color
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = c_bg

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10), Inches(0.8))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_text

    # Helper function to create standard nodes
    def create_node(shape_type, text, left, top, width, height, fill_color=c_node_fill, border_color=c_node_border):
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        # Configure text
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = c_text
        p.alignment = PP_ALIGN.CENTER
        return shape

    # Helper function to create connectors
    def create_connector(pt1, pt2, label=None, label_offset_y=0.2, color=c_node_border, is_alert=False):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, pt1[0], pt1[1], pt2[0], pt2[1]
        )
        connector.line.color.rgb = color
        connector.line.width = Pt(2)
        connector.line.end_arrowhead = True
        
        if label:
            # Calculate midpoint
            mid_x = (pt1[0] + pt2[0]) / 2
            mid_y = (pt1[1] + pt2[1]) / 2
            
            # Add label box
            lbl_box = slide.shapes.add_textbox(
                mid_x - Inches(0.5), mid_y - Inches(label_offset_y), Inches(1), Inches(0.4)
            )
            lp = lbl_box.text_frame.paragraphs[0]
            lp.text = label
            lp.font.size = Pt(11)
            lp.font.italic = True
            lp.font.color.rgb = color if not is_alert else c_alert
            lp.alignment = PP_ALIGN.CENTER

    # === Layer 1: Context Boundaries (Groupings) ===
    # Draw a dashed bounding box for the "Core Processing" zone first so it sits behind nodes
    boundary = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(3.5), Inches(1.5), Inches(6.5), Inches(4.5)
    )
    boundary.fill.background()  # Transparent
    boundary.line.color.rgb = c_boundary
    boundary.line.width = Pt(2)
    boundary.line.dash_style = MSO_LINE.DASH
    
    # Boundary Label
    b_lbl = slide.shapes.add_textbox(Inches(3.6), Inches(1.6), Inches(2), Inches(0.4))
    b_tf = b_lbl.text_frame.paragraphs[0]
    b_tf.text = "Core Processing Group"
    b_tf.font.size = Pt(11)
    b_tf.font.bold = True
    b_tf.font.color.rgb = c_boundary

    # === Layer 2: Architecture Nodes ===
    # Node 1: Ingestion (Outside boundary)
    n1 = create_node(MSO_SHAPE.RECTANGLE, "Data Ingestion", Inches(1.0), Inches(3.25), Inches(1.8), Inches(0.8))
    
    # Node 2: Decision (Inside boundary)
    n2 = create_node(MSO_SHAPE.DIAMOND, "Decision", Inches(4.2), Inches(2.9), Inches(1.6), Inches(1.5))
    
    # Node 3: Database (Inside boundary)
    n3 = create_node(MSO_SHAPE.CAN, "Database", Inches(7.5), Inches(2.0), Inches(1.6), Inches(1.4))
    
    # Node 4: Terminator (Inside boundary)
    n4 = create_node(MSO_SHAPE.RECTANGLE, "Terminator", Inches(7.5), Inches(4.5), Inches(1.6), Inches(0.8))
    
    # Node 5: Display (Outside boundary)
    n5 = create_node(MSO_SHAPE.RECTANGLE, "Display", Inches(10.5), Inches(2.3), Inches(1.8), Inches(0.8))

    # === Layer 3: Connectors ===
    # Connect N1 to N2
    create_connector((Inches(2.8), Inches(3.65)), (Inches(4.2), Inches(3.65)))
    
    # Connect N2 to N3 (Yes path)
    create_connector((Inches(5.0), Inches(2.9)), (Inches(7.5), Inches(2.7)), label="Yes", label_offset_y=0.3)
    
    # Connect N2 to N4 (No path - Error/Alert)
    create_connector((Inches(5.0), Inches(4.4)), (Inches(7.5), Inches(4.9)), label="No", label_offset_y=0.1, color=c_alert, is_alert=True)
    
    # Connect N3 to N5
    create_connector((Inches(9.1), Inches(2.7)), (Inches(10.5), Inches(2.7)))


    # === Layer 4: The Legend (Crucial tip from the tutorial) ===
    # Legend Box Background
    leg_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(6.2), Inches(12.333), Inches(1.0))
    leg_box.fill.solid()
    leg_box.fill.fore_color.rgb = RGBColor(245, 245, 245)
    leg_box.line.color.rgb = c_boundary
    
    # Legend Title
    l_txt = slide.shapes.add_textbox(Inches(0.6), Inches(6.3), Inches(1), Inches(0.5))
    lp = l_txt.text_frame.paragraphs[0]
    lp.text = "Legend:"
    lp.font.size = Pt(14)
    lp.font.bold = True
    lp.font.color.rgb = c_text

    # Legend Item 1: Process
    create_node(MSO_SHAPE.RECTANGLE, "", Inches(2.0), Inches(6.45), Inches(0.6), Inches(0.4))
    l1 = slide.shapes.add_textbox(Inches(2.7), Inches(6.4), Inches(1.5), Inches(0.4))
    l1.text_frame.paragraphs[0].text = "= Standard Process"
    l1.text_frame.paragraphs[0].font.size = Pt(11)

    # Legend Item 2: Database
    create_node(MSO_SHAPE.CAN, "", Inches(4.5), Inches(6.4), Inches(0.5), Inches(0.5))
    l2 = slide.shapes.add_textbox(Inches(5.1), Inches(6.4), Inches(1.5), Inches(0.4))
    l2.text_frame.paragraphs[0].text = "= Storage / DB"
    l2.text_frame.paragraphs[0].font.size = Pt(11)

    # Legend Item 3: Alert/Error
    create_connector((Inches(7.0), Inches(6.65)), (Inches(7.8), Inches(6.65)), color=c_alert)
    l3 = slide.shapes.add_textbox(Inches(7.9), Inches(6.4), Inches(2.0), Inches(0.4))
    l3.text_frame.paragraphs[0].text = "= Exception / Inefficient"
    l3.text_frame.paragraphs[0].font.size = Pt(11)
    l3.text_frame.paragraphs[0].font.color.rgb = c_alert

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
