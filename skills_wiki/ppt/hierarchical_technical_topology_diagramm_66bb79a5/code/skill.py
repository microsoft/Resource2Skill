def create_slide(
    output_pptx_path: str,
    title_text: str = "Cloud Architecture Topology",
    bg_color: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing a hierarchical technical architecture diagram.
    Uses native python-pptx shapes to build nested containers (Region, VCN, Subnet)
    and places node elements within them.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Set solid background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # Palette
    COLOR_REGION = RGBColor(80, 80, 80)        # Dark Gray
    COLOR_AD_LINE = RGBColor(150, 150, 150)    # Light Gray
    COLOR_AD_FILL = RGBColor(248, 248, 248)    # Very Light Gray
    COLOR_NET = RGBColor(199, 70, 52)          # Brick Red
    COLOR_NODE = RGBColor(255, 255, 255)       # White
    COLOR_NODE_LINE = RGBColor(60, 60, 60)     # Dark Gray

    def add_group_box(x, y, w, h, title, line_color, dash_style=None, fill_color=None):
        """Helper to create a styled container box with a top-left label."""
        # Main container shape
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        
        # Fill styling
        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
        else:
            shape.fill.background() # Transparent fill
            
        # Line styling
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
        if dash_style:
            shape.line.dash_style = dash_style

        # Label Textbox (Top Left)
        txBox = slide.shapes.add_textbox(Inches(x), Inches(y - 0.1), Inches(w), Inches(0.5))
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = line_color
        return shape

    def add_node(x, y, label, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
        """Helper to create a service node (representing an icon/component)."""
        w, h = 1.2, 0.8
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_NODE
        shape.line.color.rgb = COLOR_NODE_LINE
        shape.line.width = Pt(1)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_NODE_LINE
        p.alignment = PP_ALIGN.CENTER
        return shape

    # === Build Topology Back-to-Front (Managing Z-Order visually) ===

    # 1. Outer Region
    add_group_box(0.5, 1.0, 12.3, 6.0, "CLOUD REGION", COLOR_REGION)

    # 2. Availability Domains (Physical boundaries)
    add_group_box(0.8, 1.5, 5.7, 5.2, "Availability Domain 1", COLOR_AD_LINE, MSO_LINE_DASH_STYLE.DASH, COLOR_AD_FILL)
    add_group_box(6.8, 1.5, 5.7, 5.2, "Availability Domain 2", COLOR_AD_LINE, MSO_LINE_DASH_STYLE.DASH, COLOR_AD_FILL)

    # 3. Virtual Cloud Network (Spans across ADs)
    add_group_box(1.1, 2.0, 11.1, 4.4, "Virtual Cloud Network (VCN)\n10.0.0.0/16", COLOR_NET, MSO_LINE_DASH_STYLE.SQUARE_DOT)

    # 4. Subnets (Within ADs, inside VCN)
    add_group_box(1.3, 2.7, 4.7, 3.4, "Public Subnet A\n10.0.1.0/24", COLOR_NET, MSO_LINE_DASH_STYLE.DASH)
    add_group_box(7.3, 2.7, 4.7, 3.4, "Private Subnet B\n10.0.2.0/24", COLOR_NET, MSO_LINE_DASH_STYLE.DASH)

    # 5. Service Nodes
    node_lb = add_node(1.5, 3.5, "Load\nBalancer", MSO_SHAPE.HEXAGON)
    node_web = add_node(3.5, 3.5, "Web\nServer")
    node_db = add_node(8.5, 3.5, "Database\nSystem", MSO_SHAPE.CAN)
    
    # Node outside VCN but in Region
    node_igw = add_node(1.1, 1.2, "Internet\nGateway", MSO_SHAPE.CLOUD)

    # 6. Add Connectors (Simple relationship lines)
    # Connect LB to Web Server
    connector1 = slide.shapes.add_connector(MSO_SHAPE.LINE_CALLOUT_2, Inches(2.7), Inches(3.9), Inches(3.5), Inches(3.9))
    connector1.line.color.rgb = COLOR_REGION
    
    # Connect Web Server to DB
    connector2 = slide.shapes.add_connector(MSO_SHAPE.LINE_CALLOUT_2, Inches(4.7), Inches(3.9), Inches(8.5), Inches(3.9))
    connector2.line.color.rgb = COLOR_REGION
    connector2.line.dash_style = MSO_LINE_DASH_STYLE.DASH # Indicate cross-AD traffic

    # === Add Main Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(10), Inches(0.6))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_REGION

    prs.save(output_pptx_path)
    return output_pptx_path
