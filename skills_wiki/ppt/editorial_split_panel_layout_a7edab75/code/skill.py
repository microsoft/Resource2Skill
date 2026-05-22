def create_slide(
    output_pptx_path: str,
    title_text: str = "A SLIDE,\nBEAUTIFUL\nAS THIS!\nJUST ONE",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam.",
    bg_theme_keyword: str = "architecture",
    accent_color: tuple = (210, 140, 70),  # RGB for the orange/gold
    dark_color: tuple = (50, 50, 50),      # RGB for dark charcoal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Editorial Split-Panel Layout' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    import requests
    from PIL import Image, ImageDraw
    from io import BytesIO
    import os

    # Helper function to emulate font tracking (letter-spacing)
    def add_tracking(text):
        words = text.split(" ")
        tracked_words = [" ".join(list(word)) for word in words]
        return "   ".join(tracked_words)

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Force white slide background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Pillar Setup Variables ===
    col_width = 3.0
    left_starts = [4.333, 7.333, 10.333]
    
    # === Layer 1: Construct the Image Pillar (Col 2) with PIL Overlay ===
    dpi = 150
    w_px, h_px = int(col_width * dpi), int(7.5 * dpi)
    img_url = f"https://picsum.photos/seed/{bg_theme_keyword}/{w_px}/{h_px}"
    
    try:
        response = requests.get(img_url, timeout=5)
        response.raise_for_status()
        base_img = Image.open(BytesIO(response.content)).convert("RGBA")
    except Exception as e:
        print(f"Image download failed, using fallback. Error: {e}")
        base_img = Image.new("RGBA", (w_px, h_px), (150, 150, 150, 255))
        
    # Draw gradient overlay (Dark -> Accent)
    overlay = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    
    # Top color: dark, mostly opaque. Bottom color: accent, slightly more opaque.
    top_rgba = (dark_color[0], dark_color[1], dark_color[2], 180)
    bottom_rgba = (accent_color[0], accent_color[1], accent_color[2], 210)
    
    for y in range(h_px):
        ratio = y / h_px
        r = int(top_rgba[0] * (1 - ratio) + bottom_rgba[0] * ratio)
        g = int(top_rgba[1] * (1 - ratio) + bottom_rgba[1] * ratio)
        b = int(top_rgba[2] * (1 - ratio) + bottom_rgba[2] * ratio)
        a = int(top_rgba[3] * (1 - ratio) + bottom_rgba[3] * ratio)
        draw.line([(0, y), (w_px, y)], fill=(r, g, b, a))
        
    final_img = Image.alpha_composite(base_img, overlay).convert("RGB")
    temp_img_path = "temp_pillar_bg.jpg"
    final_img.save(temp_img_path, quality=95)

    # Insert Image Pillar
    slide.shapes.add_picture(temp_img_path, Inches(left_starts[0]), Inches(0), width=Inches(col_width), height=Inches(7.5))

    # === Layer 2: Construct Solid Pillars (Col 3 & Col 4) ===
    # Column 3 (White - added explicitly to overwrite anything underneath)
    rect3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_starts[1]), Inches(0), Inches(col_width), Inches(7.5))
    rect3.fill.solid()
    rect3.fill.fore_color.rgb = RGBColor(255, 255, 255)
    rect3.line.color.rgb = RGBColor(255, 255, 255) # Borderless

    # Column 4 (Dark)
    rect4 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_starts[2]), Inches(0), Inches(col_width), Inches(7.5))
    rect4.fill.solid()
    rect4.fill.fore_color.rgb = RGBColor(*dark_color)
    rect4.line.color.rgb = RGBColor(*dark_color) # Borderless

    # === Layer 3: Left Primary Column Content (Title & Main Body) ===
    title_tb = slide.shapes.add_textbox(Inches(0.4), Inches(1.2), Inches(3.5), Inches(3.0))
    tf = title_tb.text_frame
    tf.word_wrap = True
    
    lines_input = title_text.strip().split('\n')
    for i, line_text in enumerate(lines_input):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = add_tracking(line_text)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(28)
        # Highlight the last line with the accent color
        color_tup = accent_color if i == len(lines_input) - 1 else dark_color
        p.font.color.rgb = RGBColor(*color_tup)

    body_main_tb = slide.shapes.add_textbox(Inches(0.4), Inches(4.5), Inches(3.5), Inches(2.5))
    tf_body = body_main_tb.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text * 2 # Make text block thicker
    p_body.alignment = PP_ALIGN.CENTER
    p_body.font.name = "Arial"
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = RGBColor(130, 130, 130)

    # === Layer 4: Populate the 3 Secondary Pillars ===
    icons = ["✎", "💡", "📖"]
    c_white = RGBColor(255, 255, 255)
    c_dark = RGBColor(*dark_color)
    pillar_text_colors = [c_white, c_dark, c_white]

    for i in range(3):
        left = left_starts[i]
        font_color = pillar_text_colors[i]
        
        # 1. Icon
        icon_tb = slide.shapes.add_textbox(Inches(left), Inches(1.0), Inches(col_width), Inches(1.0))
        p_ic = icon_tb.text_frame.paragraphs[0]
        p_ic.text = icons[i]
        p_ic.alignment = PP_ALIGN.CENTER
        p_ic.font.size = Pt(36)
        p_ic.font.color.rgb = font_color
        
        # 2. Body Text
        pb_tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(2.5), Inches(2.4), Inches(3.0))
        tf_pb = pb_tb.text_frame
        tf_pb.word_wrap = True
        p_pb = tf_pb.paragraphs[0]
        p_pb.text = body_text
        p_pb.alignment = PP_ALIGN.CENTER
        p_pb.font.name = "Arial"
        p_pb.font.size = Pt(11)
        p_pb.font.color.rgb = font_color
        
        # 3. Number anchors
        num_tb = slide.shapes.add_textbox(Inches(left), Inches(6.5), Inches(col_width), Inches(0.5))
        p_num = num_tb.text_frame.paragraphs[0]
        p_num.text = add_tracking(f"0{i+1}")
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.name = "Arial"
        p_num.font.size = Pt(12)
        p_num.font.color.rgb = font_color

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
        
    return output_pptx_path
