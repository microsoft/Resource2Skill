import os
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def _generate_arch_placeholder(width_in, height_in, diagram_type="hero"):
    """
    Generates a minimalist, architectural-style diagram using PIL.
    Ensures the code works perfectly offline and matches the domain aesthetic.
    """
    dpi = 150
    w_px = int(width_in * dpi)
    h_px = int(height_in * dpi)
    
    img = Image.new('RGBA', (w_px, h_px), (245, 245, 245, 255))
    draw = ImageDraw.Draw(img)
    
    # Draw border
    border_color = (200, 200, 200, 255)
    draw.rectangle([0, 0, w_px-1, h_px-1], outline=border_color, width=2)
    
    if diagram_type == "hero":
        # Create a subtle perspective/isometric grid feel
        grid_color = (230, 230, 230, 255)
        for i in range(0, w_px, 40):
            draw.line([(i, 0), (i, h_px)], fill=grid_color, width=1)
        for i in range(0, h_px, 40):
            draw.line([(0, i), (w_px, i)], fill=grid_color, width=1)
            
        # Draw a "massing model" geometric shape
        accent = (80, 100, 120, 255)
        draw.polygon([(w_px*0.2, h_px*0.8), (w_px*0.5, h_px*0.3), (w_px*0.8, h_px*0.6), (w_px*0.8, h_px*0.9), (w_px*0.2, h_px*0.9)], fill=accent)
        draw.polygon([(w_px*0.2, h_px*0.8), (w_px*0.5, h_px*0.3), (w_px*0.4, h_px*0.8)], fill=(60, 80, 100, 255))

    elif diagram_type == "plan":
        # Top-down floor plan look
        wall_color = (50, 50, 50, 255)
        room_color = (220, 220, 220, 255)
        draw.rectangle([w_px*0.1, h_px*0.1, w_px*0.9, h_px*0.9], outline=wall_color, width=6, fill=room_color)
        draw.rectangle([w_px*0.1, h_px*0.4, w_px*0.5, h_px*0.4], outline=wall_color, width=4) # Inner wall
        draw.rectangle([w_px*0.6, h_px*0.1, w_px*0.6, h_px*0.6], outline=wall_color, width=4) # Inner wall

    elif diagram_type == "section":
        # Section cut look
        earth_color = (180, 180, 180, 255)
        structure_color = (30, 30, 30, 255)
        # Ground line
        draw.rectangle([0, h_px*0.7, w_px, h_px], fill=earth_color)
        # Building section
        draw.rectangle([w_px*0.25, h_px*0.3, w_px*0.75, h_px*0.7], outline=structure_color, width=5)
        # Roof
        draw.line([(w_px*0.2, h_px*0.3), (w_px*0.8, h_px*0.3)], fill=structure_color, width=8)

    img_stream = io.BytesIO()
    img.save(img_stream, format='PNG')
    img_stream.seek(0)
    return img_stream

def create_slide(
    output_pptx_path: str,
    title_text: str = "URBAN RENEWAL PAVILION",
    body_text: str = "The concept explores the intersection of brutalist massing and permeable public spaces. By elevating the primary structure, the ground plane is liberated for community interaction, while the strict geometric grid organizes the analytical programmatic spaces above. The design prioritizes natural light, material honesty, and clear structural hierarchy.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Minimalist Architectural Grid Presentation effect.
    """
    prs = Presentation()
    # Use standard widescreen but treat it like a landscape presentation board
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Colors
    charcoal = RGBColor(40, 40, 40)
    light_grey = RGBColor(120, 120, 120)

    # ==========================================
    # GRID DEFINITION (Measurements in Inches)
    # ==========================================
    margin = 0.5
    gutter = 0.3 # Space between grid elements
    
    # Left Column (Hero)
    left_col_x = margin
    left_col_w = 7.5
    
    # Right Column (Analytical)
    right_col_x = left_col_x + left_col_w + gutter
    right_col_w = prs.slide_width.inches - right_col_x - margin

    # ==========================================
    # LAYER 1: TYPOGRAPHY & TEXT HIERARCHY
    # ==========================================
    # Title Box
    title_y = margin
    title_h = 1.2
    title_box = slide.shapes.add_textbox(Inches(left_col_x), Inches(title_y), Inches(left_col_w), Inches(title_h))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = 'Arial'
    p.font.color.rgb = charcoal

    p2 = tf.add_paragraph()
    p2.text = "PROCESS & ANALYTICAL DIAGRAMS"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.name = 'Arial'
    p2.font.color.rgb = light_grey

    # Right Column Text Box (Concept Statement)
    concept_y = margin
    concept_h = 1.8
    concept_box = slide.shapes.add_textbox(Inches(right_col_x), Inches(concept_y), Inches(right_col_w), Inches(concept_h))
    ctf = concept_box.text_frame
    ctf.word_wrap = True
    
    cp1 = ctf.paragraphs[0]
    cp1.text = "CONCEPTUAL FRAMEWORK"
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = charcoal
    
    cp2 = ctf.add_paragraph()
    cp2.text = body_text
    cp2.font.size = Pt(10)
    cp2.font.color.rgb = charcoal
    cp2.alignment = PP_ALIGN.JUSTIFY

    # ==========================================
    # LAYER 2: GRIDDED IMAGERY
    # ==========================================
    # Hero Image (Left Column)
    hero_y = title_y + title_h + gutter
    hero_h = prs.slide_height.inches - hero_y - margin
    hero_img_stream = _generate_arch_placeholder(left_col_w, hero_h, "hero")
    slide.shapes.add_picture(hero_img_stream, Inches(left_col_x), Inches(hero_y), Inches(left_col_w), Inches(hero_h))

    # Analytical Image 1 (Right Column, Top)
    img1_y = concept_y + concept_h + gutter
    img_h = (prs.slide_height.inches - img1_y - margin - gutter) / 2
    img1_stream = _generate_arch_placeholder(right_col_w, img_h, "plan")
    slide.shapes.add_picture(img1_stream, Inches(right_col_x), Inches(img1_y), Inches(right_col_w), Inches(img_h))
    
    # Add subtle caption for Image 1
    cap1 = slide.shapes.add_textbox(Inches(right_col_x), Inches(img1_y - 0.25), Inches(right_col_w), Inches(0.25))
    cap1.text_frame.text = "FIG 1: GROUND FLOOR PLAN"
    cap1.text_frame.paragraphs[0].font.size = Pt(8)
    cap1.text_frame.paragraphs[0].font.color.rgb = light_grey

    # Analytical Image 2 (Right Column, Bottom)
    img2_y = img1_y + img_h + gutter
    img2_stream = _generate_arch_placeholder(right_col_w, img_h, "section")
    slide.shapes.add_picture(img2_stream, Inches(right_col_x), Inches(img2_y), Inches(right_col_w), Inches(img_h))
    
    # Add subtle caption for Image 2
    cap2 = slide.shapes.add_textbox(Inches(right_col_x), Inches(img2_y - 0.25), Inches(right_col_w), Inches(0.25))
    cap2.text_frame.text = "FIG 2: TRANSVERSE SECTION"
    cap2.text_frame.paragraphs[0].font.size = Pt(8)
    cap2.text_frame.paragraphs[0].font.color.rgb = light_grey

    prs.save(output_pptx_path)
    return output_pptx_path
