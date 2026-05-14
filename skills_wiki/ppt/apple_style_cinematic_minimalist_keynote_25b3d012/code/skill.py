import os
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def _create_cinematic_vignette(output_path: str, width: int = 1920, height: int = 1080):
    """
    Helper function: Generates a premium dark radial gradient background using PIL.
    Simulates a subtle stage spotlight (Apple keynote style).
    """
    # Colors: Deep charcoal center, pitch black edges
    center_color = (35, 40, 45)
    edge_color = (0, 0, 0)
    
    # Base image (edges)
    base = Image.new('RGB', (width, height), edge_color)
    
    # Create a mask for the radial spotlight
    mask = Image.new('L', (width, height), 0)
    draw = ImageDraw.Draw(mask)
    
    # Draw a large ellipse in the center
    ellipse_w, ellipse_h = width * 1.2, height * 1.5
    left = (width - ellipse_w) / 2
    top = (height - ellipse_h) / 2
    draw.ellipse((left, top, left + ellipse_w, top + ellipse_h), fill=255)
    
    # Massively blur the ellipse to create a smooth gradient
    mask = mask.filter(ImageFilter.GaussianBlur(250))
    
    # Image for the center color
    center_img = Image.new('RGB', (width, height), center_color)
    
    # Composite the images
    bg = Image.composite(center_img, base, mask)
    bg.save(output_path)
    return output_path

def create_slide(
    output_pptx_path: str,
    title_text: str = "Less is More.",
    subtitle_text: str = "让观点更有力量",
    accent_color: tuple = (255, 255, 255),  # Pure white for maximum contrast
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Apple-Style Minimalist Keynote slide.
    
    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    # Set to 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Cinematic Vignette Background ===
    bg_img_path = "temp_cinematic_bg.png"
    _create_cinematic_vignette(bg_img_path)
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Massive Core Typography ===
    # Calculate perfectly centered text box dimensions
    tb_width = Inches(11)
    tb_height = Inches(4)
    left = (prs.slide_width - tb_width) / 2
    top = (prs.slide_height - tb_height) / 2

    txBox = slide.shapes.add_textbox(left, top, tb_width, tb_height)
    tf = txBox.text_frame
    tf.clear()  # Clear default formatting
    
    # Vertical centering within the text box
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # --- Title Configuration ---
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.CENTER
    run_title = p_title.add_run()
    run_title.text = title_text
    
    # Apply tutorial principles: Sans-serif, High Contrast, Massive Size
    font_title = run_title.font
    font_title.name = 'Arial'  # Safe sans-serif fallback
    font_title.size = Pt(110)
    font_title.bold = True
    font_title.color.rgb = RGBColor(*accent_color)

    # --- Subtitle Configuration (Optional) ---
    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.alignment = PP_ALIGN.CENTER
        run_sub = p_sub.add_run()
        run_sub.text = subtitle_text
        
        font_sub = run_sub.font
        font_sub.name = 'Microsoft YaHei'  # Safe sans-serif for Chinese
        font_sub.size = Pt(28)
        font_sub.bold = False
        # Lower contrast for subtitle to maintain hierarchy
        font_sub.color.rgb = RGBColor(150, 160, 170)

    # Clean up temp files
    try:
        os.remove(bg_img_path)
    except OSError:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("apple_style_keynote.pptx", title_text="Hello.", subtitle_text="This is a high-contrast minimalist slide.")
