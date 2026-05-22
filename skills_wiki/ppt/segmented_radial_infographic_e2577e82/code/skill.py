def create_slide(
    output_pptx_path: str,
    title_text: str = "Segmented Radial Infographic",
    body_text: str = "",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Segmented Radial Infographic visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    import math

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 248)

    # === Slide Title ===
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 30, 30)

    # === Geometry & Layout Configurations ===
    cx = int(prs.slide_width / 2)
    cy = int(prs.slide_height / 2 + Inches(0.4))
    r_out = Inches(2.2)
    r_in = Inches(1.3)
    
    # 4-Segment Palette and Alignment
    segments = [
        {"color": (218, 56, 50), "align": "right"}, # Red
        {"color": (19, 41, 75),  "align": "right"}, # Navy
        {"color": (0, 114, 206), "align": "left"},  # Blue
        {"color": (255, 192, 0), "align": "left"}   # Yellow
    ]

    # === Helper: Draw Arc Segment ===
    def add_arc_segment(slide_obj, cx_val, cy_val, radius_out, radius_in, start_deg, end_deg, color_rgb):
        steps = max(20, int((end_deg - start_deg) / 2))
        pts = []
        # Outer arc
        for i in range(steps + 1):
            ang = math.radians(start_deg + (end_deg - start_deg) * i / steps)
            pts.append((cx_val + radius_out * math.cos(ang), cy_val + radius_out * math.sin(ang)))
        # Inner arc (reversed)
        for i in range(steps, -1, -1):
            ang = math.radians(start_deg + (end_deg - start_deg) * i / steps)
            pts.append((cx_val + radius_in * math.cos(ang), cy_val + radius_in * math.sin(ang)))
            
        builder = slide_obj.shapes.build_freeform(pts[0][0], pts[0][1])
        builder.add_line_segments(pts[1:], close=True)
        shape = builder.convert_to_shape()
        
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color_rgb)
        # White border creates the visual "gap" between segments
        shape.line.color.rgb = RGBColor(245, 245, 248) # Match background
        shape.line.width = Pt(4)
        return shape

    # === Helper: Add Radial Labels & Connectors ===
    def add_radial_label(slide_obj, cx_val, cy_val, radius_out, radius_in, mid_deg, num, title, body, align_type, color_rgb):
        rad = math.radians(mid_deg)
        
        # 1. Floating Number inside the arc
        r_mid = (radius_out + radius_in) / 2 
        nx = cx_val + r_mid * math.cos(rad)
        ny = cy_val + r_mid * math.sin(rad)
        bs = Inches(0.6)
        nb = slide_obj.shapes.add_textbox(int(nx - bs/2), int(ny - bs/2), int(bs), int(bs))
        nb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_num = nb.text_frame.paragraphs[0]
        p_num.text = f"0{num}"
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.size = Pt(22)
        p_num.font.bold = True
        p_num.font.color.rgb = RGBColor(255, 255, 255)
        
        # 2. Connector Line
        x_start = cx_val + radius_out * math.cos(rad)
        y_start = cy_val + radius_out * math.sin(rad)
        line_len = Inches(0.5)
        
        w = Inches(2.6)
        h = Inches(1.2)
        
        if align_type == 'right':
            x_end = x_start - line_len
            left = x_end - w
            text_align = PP_ALIGN.RIGHT
        else:
            x_end = x_start + line_len
            left = x_end
            text_align = PP_ALIGN.LEFT
            
        conn = slide_obj.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, int(x_start), int(y_start), int(x_end), int(y_start)
        )
        conn.line.color.rgb = RGBColor(*color_rgb)
        conn.line.width = Pt(2)
        
        # 3. Text Box
        top = y_start - h / 2
        tb = slide_obj.shapes.add_textbox(int(left), int(top), int(w), int(h))
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p_title = tb.text_frame.paragraphs[0]
        p_title.text = title
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = RGBColor(*color_rgb)
        p_title.alignment = text_align
        
        p_body = tb.text_frame.add_paragraph()
        p_body.text = body
        p_body.font.size = Pt(11)
        p_body.font.color.rgb = RGBColor(100, 100, 100)
        p_body.alignment = text_align

    # === Build Segments ===
    # Total span 270 degrees, starting at bottom-left (135) to bottom-right (405)
    start_angle = 135
    span = 270 / len(segments)
    
    for i, seg in enumerate(segments):
        end_angle = start_angle + span
        mid_angle = (start_angle + end_angle) / 2
        
        # Draw the vector pie segment
        add_arc_segment(slide, cx, cy, r_out, r_in, start_angle, end_angle, seg["color"])
        
        # Add the connected text elements
        add_radial_label(
            slide, cx, cy, r_out, r_in, mid_angle, i+1, 
            f"Phase 0{i+1} Heading", 
            "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore.", 
            seg["align"], seg["color"]
        )
        
        start_angle = end_angle

    # === Central Hub Graphic ===
    cr = r_in - Inches(0.1)
    center_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, int(cx - cr), int(cy - cr), int(cr * 2), int(cr * 2)
    )
    center_circle.fill.solid()
    center_circle.fill.fore_color.rgb = RGBColor(25, 25, 35) # Dark contrasting hub
    center_circle.line.fill.background()
    
    ctb = slide.shapes.add_textbox(int(cx - cr), int(cy - cr), int(cr*2), int(cr*2))
    ctb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    cp = ctb.text_frame.paragraphs[0]
    cp.text = "💡" # Represents the core idea
    cp.alignment = PP_ALIGN.CENTER
    cp.font.size = Pt(50)

    prs.save(output_pptx_path)
    return output_pptx_path
