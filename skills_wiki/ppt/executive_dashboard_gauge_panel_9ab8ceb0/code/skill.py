def create_slide(output_pptx_path: str, **kwargs) -> str:
    """
    Creates a PowerPoint slide featuring a 3-column Dashboard Gauge Panel.
    Uses PIL to geometrically draw the speedometers and python-pptx for layout.
    """
    import os
    import math
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # --- Data & Color Configuration ---
    metrics = [
        {"label": "REACH", "value": 0.80, "color": (91, 135, 177)},
        {"label": "ENGAGEMENT", "value": 0.55, "color": (216, 151, 116)},
        {"label": "AWARENESS", "value": 0.75, "color": (142, 185, 147)}
    ]
    
    slide_title = "EFFICIENCY"
    bg_color = (255, 255, 255)

    # --- PIL Image Generation Function ---
    def generate_gauge_image(percentage: float, main_color: tuple, filename: str):
        # Create an image canvas (white background for clean anti-aliasing)
        size = 600
        # Height is half the size + some padding for the needle base
        img = Image.new('RGB', (size, int(size/2 + 20)), bg_color)
        draw = ImageDraw.Draw(img)

        cx, cy = size / 2, size / 2
        radius = size / 2 - 20
        thickness = 80
        
        # Calculate light color for the background track (blend with white)
        light_c = tuple(int(c + (255 - c) * 0.7) for c in main_color)
        
        # 1. Draw Background Track (using pieslice cutout method for clean edges)
        bbox = [cx - radius, cy - radius, cx + radius, cy + radius]
        draw.pieslice(bbox, 180, 360, fill=light_c)
        
        # 2. Draw Value Arc
        end_angle = 180 + (180 * percentage)
        draw.pieslice(bbox, 180, end_angle, fill=main_color)
        
        # 3. Inner Cutout to create the "Donut" arc
        inner_r = radius - thickness
        inner_bbox = [cx - inner_r, cy - inner_r, cx + inner_r, cy + inner_r]
        draw.pieslice(inner_bbox, 180, 360, fill=bg_color)
        
        # 4. Draw Segment Separators (White lines radiating from center)
        for i in range(1, 10):
            ang_rad = math.radians(180 + i * 18)
            lx_in = cx + inner_r * math.cos(ang_rad)
            ly_in = cy + inner_r * math.sin(ang_rad)
            lx_out = cx + radius * math.cos(ang_rad)
            ly_out = cy + radius * math.sin(ang_rad)
            draw.line([(lx_in, ly_in), (lx_out, ly_out)], fill=bg_color, width=10)
            
        # 5. Draw the Needle
        needle_angle_rad = math.radians(end_angle)
        needle_length = radius - 15
        
        # Needle Tip
        tip_x = cx + needle_length * math.cos(needle_angle_rad)
        tip_y = cy + needle_length * math.sin(needle_angle_rad)
        
        # Needle Base (perpendicular to create a triangle)
        base_width = 15
        perp_rad = needle_angle_rad + math.pi / 2
        bx1 = cx + base_width * math.cos(perp_rad)
        by1 = cy + base_width * math.sin(perp_rad)
        bx2 = cx - base_width * math.cos(perp_rad)
        by2 = cy - base_width * math.sin(perp_rad)
        
        draw.polygon([(tip_x, tip_y), (bx1, by1), (bx2, by2)], fill=main_color)
        
        # Center Pin
        pin_r = 20
        draw.ellipse([cx - pin_r, cy - pin_r, cx + pin_r, cy + pin_r], fill=main_color)
        draw.ellipse([cx - pin_r/2, cy - pin_r/2, cx + pin_r/2, cy + pin_r/2], fill=bg_color)
        
        img.save(filename)
        return filename

    # --- PPTX Generation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Add Slide Title
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.5), prs.slide_width, Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = slide_title
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(60, 64, 72)

    # Layout dimensions
    col_width = prs.slide_width / 3
    img_width = Inches(3.5)
    img_y = Inches(2.5)
    
    temp_files = []

    # Iterate through metrics and place them in 3 columns
    for idx, metric in enumerate(metrics):
        # 1. Generate local image
        tmp_filename = f"temp_gauge_{idx}.png"
        generate_gauge_image(metric["value"], metric["color"], tmp_filename)
        temp_files.append(tmp_filename)
        
        # Calculate X position (Center of the specific column)
        col_center_x = (idx * col_width) + (col_width / 2)
        
        # 2. Insert Gauge Image
        img_x = col_center_x - (img_width / 2)
        slide.shapes.add_picture(tmp_filename, img_x, img_y, width=img_width)
        
        # 3. Add Metric Label
        label_y = img_y + Inches(2.0)
        label_box = slide.shapes.add_textbox(col_center_x - Inches(2), label_y, Inches(4), Inches(0.6))
        lp = label_box.text_frame.paragraphs[0]
        lp.text = metric["label"]
        lp.alignment = PP_ALIGN.CENTER
        lp.font.name = "Arial"
        lp.font.size = Pt(20)
        lp.font.bold = True
        lp.font.color.rgb = RGBColor(*metric["color"])
        
        # 4. Add Metric Percentage
        val_y = label_y + Inches(0.4)
        val_box = slide.shapes.add_textbox(col_center_x - Inches(2), val_y, Inches(4), Inches(1))
        vp = val_box.text_frame.paragraphs[0]
        vp.text = f"{int(metric['value'] * 100)}%"
        vp.alignment = PP_ALIGN.CENTER
        vp.font.name = "Arial"
        vp.font.size = Pt(44)
        vp.font.color.rgb = RGBColor(*metric["color"])

    prs.save(output_pptx_path)
    
    # Cleanup temporary images
    for f in temp_files:
        if os.path.exists(f):
            os.remove(f)

    return output_pptx_path
