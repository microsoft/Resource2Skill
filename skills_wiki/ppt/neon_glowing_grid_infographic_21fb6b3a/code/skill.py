import os
import tempfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFilter

def _create_glow_orb(color_rgba, radius=150, blur_amount=80):
    """Generates a soft, glowing spherical PNG with alpha transparency."""
    # Create an image with padding to avoid clipping the blur
    size = (radius * 2) + (blur_amount * 4)
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    center = size // 2
    # Draw solid circle
    draw.ellipse(
        [center - radius, center - radius, center + radius, center + radius],
        fill=color_rgba
    )
    # Apply heavy blur to create the fuzzy glow effect
    img = img.filter(ImageFilter.GaussianBlur(blur_amount))
    
    temp_path = tempfile.mktemp(suffix=".png")
    img.save(temp_path, format="PNG")
    return temp_path

def _add_text_outline(run, hex_color="BCFF01", width_pt=1.5):
    """Injects OOXML to add a stroke (outline) to a specific text run."""
    rPr = run._r.get_or_add_rPr()
    
    # Create outline element <a:ln>
    ln = OxmlElement('a:ln')
    ln.set('w', str(int(width_pt * 12700))) # Convert Pt to EMUs (1 pt = 12700 EMUs)
    
    # Set outline solid fill
    solidFill = OxmlElement('a:solidFill')
    srgbClr = OxmlElement('a:srgbClr')
    srgbClr.set('val', hex_color)
    solidFill.append(srgbClr)
    ln.append(solidFill)
    
    # Add to run properties
    rPr.append(ln)

def create_slide(
    output_pptx_path: str,
    title_text: str = "PERCENTAGE CHART",
    percentage: int = 75,
    bg_color: tuple = (25, 34, 13),      # Very dark green
    accent_color: tuple = (188, 255, 1), # Neon lime
    dark_fill: tuple = (56, 80, 0),      # Dark muted green
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neon Glowing Grid Infographic.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # === Layer 1: Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Glowing Atmospheric Orbs ===
    rgba_accent = (accent_color[0], accent_color[1], accent_color[2], 120)
    
    orb1_path = _create_glow_orb(rgba_accent, radius=180, blur_amount=90)
    orb2_path = _create_glow_orb(rgba_accent, radius=100, blur_amount=60)
    
    # Position orbs subtly behind the main focus areas
    slide.shapes.add_picture(orb1_path, Inches(1.0), Inches(4.5), Inches(5), Inches(5))
    slide.shapes.add_picture(orb2_path, Inches(3.5), Inches(1.0), Inches(3), Inches(3))

    # === Layer 3: Typography & Lines ===
    
    # 1. Subtitle Text
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(4), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*accent_color)
    
    # 2. Subtitle Accent Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.7), Inches(2.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # 3. Main Percentage Value
    val_box = slide.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(4.5), Inches(2.5))
    val_tf = val_box.text_frame
    p_val = val_tf.paragraphs[0]
    
    # Integer part
    run_int = p_val.add_run()
    run_int.text = str(percentage)
    run_int.font.name = "Arial Black"
    run_int.font.size = Pt(130)
    run_int.font.bold = True
    run_int.font.color.rgb = RGBColor(*accent_color)
    
    # Percentage symbol part (Hollow / Neon Stroke)
    run_sym = p_val.add_run()
    run_sym.text = "%"
    run_sym.font.name = "Arial Black"
    run_sym.font.size = Pt(130)
    run_sym.font.bold = True
    run_sym.font.color.rgb = RGBColor(*dark_fill)
    
    # Apply custom XML to outline the % sign
    accent_hex = f"{accent_color[0]:02X}{accent_color[1]:02X}{accent_color[2]:02X}"
    _add_text_outline(run_sym, hex_color=accent_hex, width_pt=1.5)

    # === Layer 4: The 10x10 Grid Infographic ===
    
    grid_start_x = Inches(6.5)
    grid_start_y = Inches(1.25)
    dot_spacing = Inches(0.55)
    dot_size = Inches(0.42)
    
    # Ensure percentage is clamped between 0 and 100
    safe_percent = max(0, min(100, percentage))
    
    for row in range(10):
        for col in range(10):
            idx = (row * 10) + col
            
            x = grid_start_x + (col * dot_spacing)
            y = grid_start_y + (row * dot_spacing)
            
            oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, dot_size, dot_size)
            
            # Decide if circle is "filled" or "empty"
            if idx < safe_percent:
                # Active circle
                oval.fill.solid()
                oval.fill.fore_color.rgb = RGBColor(*accent_color)
                oval.line.fill.background() # No line
            else:
                # Inactive circle
                oval.fill.solid()
                oval.fill.fore_color.rgb = RGBColor(*dark_fill)
                oval.line.color.rgb = RGBColor(*accent_color)
                oval.line.width = Pt(1.5)

    # Cleanup temp image files
    try:
        os.remove(orb1_path)
        os.remove(orb2_path)
    except:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path
