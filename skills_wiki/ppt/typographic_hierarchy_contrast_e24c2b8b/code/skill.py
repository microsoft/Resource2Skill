import urllib.request
import io
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "ROBOTO",
    subtitle_text: str = "I'M A SUB-HEADING",
    body_text: str = "And now I'm a bunch of body text that finishes off the pairing of these two typefaces. Font pairing is important because it helps to create visual hierarchy and improve readability in design projects. The right combination of fonts can emphasize the intended message, create contrast and bring balance to a design.",
    bg_keyword: str = "dark abstract texture",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide demonstrating the Typographic Hierarchy & Contrast skill,
    pairing a strong Slab Serif (Roboto Slab) with a clean Sans-Serif (Mulish).

    NOTE: This code assumes the fonts 'Roboto Slab' and 'Mulish' are installed
    on the system where the presentation is viewed.

    Returns: Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    try:
        # Use a more reliable API for images, like Pexels or Unsplash with an API key if available
        # For this example, we'll use a placeholder service that might be unstable.
        image_url = f"https://source.unsplash.com/1280x720/?{urllib.parse.quote(bg_keyword)}"
        response = requests.get(image_url, stream=True, timeout=10)
        response.raise_for_status()
        image_stream = io.BytesIO(response.content)
        slide.shapes.add_picture(image_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except (requests.exceptions.RequestException, IOError) as e:
        print(f"Warning: Could not download background image. Using fallback. Error: {e}")
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(20, 20, 30) # Dark charcoal fallback

    # === Layer 2: Decorative Elements ===
    line_left = Inches(1.5)
    line_top = Inches(1.8)
    line_height = Inches(2.5)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_left, line_top, Inches(0.1), line_height)
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(255, 255, 255)
    line.line.fill.background()

    # === Layer 3: Text & Content ===
    # Headline (Roboto Slab)
    txBox = slide.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(10), Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Roboto Slab'
    p.font.size = Pt(88)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle (Mulish)
    txBox_sub = slide.shapes.add_textbox(Inches(1.8), Inches(3.2), Inches(10), Inches(1))
    p_sub = txBox_sub.text_frame.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = 'Mulish'
    p_sub.font.size = Pt(28)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(255, 255, 255)

    # Body Text (Mulish)
    txBox_body = slide.shapes.add_textbox(Inches(1.8), Inches(4.2), Inches(8), Inches(2.5))
    tf_body = txBox_body.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = 'Mulish'
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(255, 255, 255)
    p_body.line_spacing = 1.5

    prs.save(output_pptx_path)
    return output_pptx_path

