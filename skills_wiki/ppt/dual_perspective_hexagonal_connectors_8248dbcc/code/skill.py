from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_LINE
from pptx.enum.dml import MSO_THEME_COLOR, MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

def create_pros_cons_slide(output_pptx_path: str, **kwargs) -> str:
    """
    Creates a PowerPoint slide with a 'Pros and Cons' infographic using
    dual-perspective hexagonal connectors.

    Returns: Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Define color palettes
    positive_main_color = RGBColor(46, 172, 60)
    positive_light_color = RGBColor(155, 213, 162)
    negative_main_color = RGBColor(192, 0, 0)
    negative_light_color = RGBColor(217, 106, 107)
    text_color = RGBColor(89, 89, 89)

    # --- Helper function to draw one side of the infographic ---
    def draw_connector_side(side: str):
        is_positive = side == 'positive'
        
        # Determine position and colors based on side
        center_x = prs.slide_width / 2
        offset_x = Inches(0.2)
        start_x = center_x - offset_x if is_positive else center_x + offset_x
        
        main_color = positive_main_color if is_positive else negative_main_color
        light_color = positive_light_color if is_positive else negative_light_color
        
        # --- Define shape geometry ---
        cap_width = Inches(1.2)
        cap_height = Inches(1.4)
        body_length = Inches(4.5)
        
        # Multiplier for horizontal direction (-1 for left, 1 for right)
        direction = -1 if is_positive else 1
        
        # --- 1. Draw the Solid Cap ---
        # Vertices for the thick hexagonal cap
        cap_path = [
            (start_x, Emu(Inches(3.05))),
            (start_x, Emu(Inches(4.45))),
            (start_x + direction * cap_width * 0.5, Emu(Inches(5.15))),
            (start_x + direction * cap_width, Emu(Inches(4.45))),
            (start_x + direction * cap_width, Emu(Inches(3.05))),
            (start_x + direction * cap_width * 0.5, Emu(Inches(2.35))),
        ]
        cap_shape = slide.shapes.add_freeform_shape(Emu(cap_path[0][0]), Emu(cap_path[0][1]), cap_path)
        cap_shape.fill.solid()
        cap_shape.fill.fore_color.rgb = main_color
        cap_shape.line.fill.background() # No line

        # --- 2. Draw the Hollow Body ---
        line_thickness = Pt(4)
        # Top line
        top_line_path = [
            (start_x, Emu(Inches(3.05))),
            (start_x - direction * body_length, Emu(Inches(3.05))),
            (start_x - direction * (body_length + cap_width * 0.25), Emu(Inches(2.70)))
        ]
        top_line = slide.shapes.add_freeform_shape(Emu(top_line_path[0][0]), Emu(top_line_path[0][1]), top_line_path)
        top_line.line.color.rgb = light_color
        top_line.line.width = line_thickness
        top_line.fill.background()
        
        # Bottom line
        bottom_line_path = [
             (start_x, Emu(Inches(4.45))),
             (start_x - direction * body_length, Emu(Inches(4.45))),
             (start_x - direction * (body_length + cap_width * 0.25), Emu(Inches(4.80)))
        ]
        bottom_line = slide.shapes.add_freeform_shape(Emu(bottom_line_path[0][0]), Emu(bottom_line_path[0][1]), bottom_line_path)
        bottom_line.line.color.rgb = light_color
        bottom_line.line.width = line_thickness
        bottom_line.fill.background()
        
        # --- 3. Draw Decorative Elements (Dots and Arrows) ---
        dot_size = Inches(0.2)
        dot_left = start_x - dot_size/2 if is_positive else start_x + direction * cap_width - dot_size/2
        dot_top = Emu(Inches(2.35)) - dot_size/2
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, dot_left, dot_top, dot_size, dot_size)
        dot.fill.solid()
        dot.fill.fore_color.rgb = main_color
        dot.line.fill.background()

        arrow_start_x = top_line_path[1][0] if is_positive else bottom_line_path[1][0]
        arrow_start_y = top_line_path[1][1] if is_positive else bottom_line_path[1][1]
        arrow = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 
            arrow_start_x - direction * Inches(0.5), arrow_start_y,
            arrow_start_x - direction * Inches(0.5), arrow_start_y - Inches(0.5)
        )
        arrow.line.color.rgb = main_color
        arrow.line.width = Pt(3)
        arrow.line.end_arrow_type = MSO_LINE.ARROW_TRIANGLE
        if not is_positive: # Flip arrow for negative side
            arrow.rotation = 180

        # --- 4. Add Icons and Text ---
        icon_size = Pt(60)
        icon_left = start_x + direction * cap_width/2 - Emu(Pt(icon_size/2))
        icon_top = Emu(Inches(3.75)) - Emu(Pt(icon_size/2))
        icon_box = slide.shapes.add_textbox(icon_left, icon_top, Emu(icon_size), Emu(icon_size))
        p_icon = icon_box.text_frame.paragraphs[0]
        run_icon = p_icon.add_run()
        run_icon.text = "☺" if is_positive else "☹"
        run_icon.font.size = icon_size
        run_icon.font.color.rgb = light_color

        # Title text
        title_left = Inches(1) if is_positive else Inches(8)
        title_box = slide.shapes.add_textbox(title_left, Inches(2), Inches(4.5), Inches(0.5))
        p_title = title_box.text_frame.paragraphs[0]
        run_title = p_title.add_run()
        run_title.text = "POSITIVES" if is_positive else "NEGATIVES"
        run_title.font.name = 'Arial Black'
        run_title.font.size = Pt(28)
        run_title.font.bold = True
        run_title.font.color.rgb = main_color

        # List items
        list_items = [
            f"Add {'Positive' if is_positive else 'Negative'} Line here"
        ] * 4
        
        list_start_y = Inches(3.2)
        list_start_x = Inches(1.5) if is_positive else Inches(7.8)
        for i, item in enumerate(list_items):
            y_pos = list_start_y + i * Inches(0.6)
            
            # Check/Cross icon
            check_box = slide.shapes.add_textbox(list_start_x, y_pos, Inches(0.4), Inches(0.4))
            p_check = check_box.text_frame.paragraphs[0]
            run_check = p_check.add_run()
            run_check.text = "✔" if is_positive else "✖"
            run_check.font.color.rgb = main_color
            run_check.font.size = Pt(20)

            # Item text
            item_box = slide.shapes.add_textbox(list_start_x + Inches(0.4), y_pos, Inches(3.5), Inches(0.4))
            p_item = item_box.text_frame.paragraphs[0]
            run_item = p_item.add_run()
            run_item.text = item
            run_item.font.name = 'Open Sans'
            run_item.font.size = Pt(16)
            run_item.font.color.rgb = text_color


    # --- Draw both sides ---
    draw_connector_side('positive')
    draw_connector_side('negative')
    
    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
if __name__ == '__main__':
    create_pros_cons_slide("pros_and_cons_slide.pptx")

