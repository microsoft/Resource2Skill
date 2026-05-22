def create_slide(
    output_pptx_path: str,
    title_text: str = "Awesome Text",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. \nSed do eiusmod tempor incididunt ut labore et dolore magna aliqua. \nUt enim ad minim veniam, quis nostrud exercitation.",
    brush_color: tuple = (25, 28, 35, 255),  # Deep charcoal ink
    **kwargs,
) -> str:
    """
    Creates a presentation slide featuring an organic brush stroke accent 
    with high-contrast typography pairings (Serif Heading + Sans Body).
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import random
    import os

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Set background to light gray/off-white
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 247)

    # 2. Generate the Organic Brush Stroke using PIL
    # We simulate a brush stroke by drawing many overlapping, slightly randomized thick lines
    img_width, img_height = 1600, 600
    brush_img = Image.new('RGBA', (img_width, img_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(brush_img)

    base_x1, base_y1 = img_width * 0.1, img_height * 0.5
    base_x2, base_y2 = img_width * 0.9, img_height * 0.5

    # Seed for consistent procedural generation
    random.seed(42)

    # Draw bristles
    for _ in range(300):
        # Randomize start and end points to create jagged edges
        x1 = base_x1 + random.randint(-80, 80)
        y1 = base_y1 + random.randint(-120, 120)
        x2 = base_x2 + random.randint(-80, 80)
        y2 = base_y2 + random.randint(-120, 120)
        
        # Vary line thickness for texture
        w = random.randint(10, 45)
        
        # Add slight transparency variation to some bristles
        alpha = random.randint(200, 255)
        current_color = (brush_color[0], brush_color[1], brush_color[2], alpha)
        
        draw.line([(x1, y1), (x2, y2)], fill=current_color, width=w)
        # Round the ends of the bristles
        draw.ellipse([x1-w//2, y1-w//2, x1+w//2, y1+w//2], fill=current_color)
        draw.ellipse([x2-w//2, y2-w//2, x2+w//2, y2+w//2], fill=current_color)

    brush_path = "temp_brush_stroke.png"
    brush_img.save(brush_path)

    # 3. Add Brush Image to Slide
    # Place it centrally but slightly elevated
    pic_width = Inches(10)
    pic_height = Inches(3.75)
    pic_left = (prs.slide_width - pic_width) / 2
    pic_top = Inches(1.5)
    slide.shapes.add_picture(brush_path, pic_left, pic_top, pic_width, pic_height)

    # 4. Add Typography - The core lesson of the tutorial
    
    # Heading: Large, Bold, Serif (Georgia), White
    title_box = slide.shapes.add_textbox(pic_left, pic_top + Inches(1), pic_width, Inches(1.5))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.clear()
    
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.alignment = PP_ALIGN.CENTER
    
    run_title = p_title.runs[0]
    run_title.font.name = 'Georgia'  # As recommended in the video
    run_title.font.size = Pt(60)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255) # High contrast against dark brush

    # Body: Smaller, Clean, Sans-Serif (Arial), Dark
    body_top = pic_top + pic_height + Inches(0.2)
    body_box = slide.shapes.add_textbox(Inches(2.5), body_top, prs.slide_width - Inches(5), Inches(2))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    tf_body.clear()
    
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.alignment = PP_ALIGN.CENTER
    p_body.line_spacing = 1.5
    
    run_body = p_body.runs[0]
    run_body.font.name = 'Arial'  # Pairing Serif heading with Sans-Serif body
    run_body.font.size = Pt(18)
    run_body.font.color.rgb = RGBColor(60, 60, 60)

    # 5. Save and Cleanup
    prs.save(output_pptx_path)
    
    if os.path.exists(brush_path):
        os.remove(brush_path)
        
    return output_pptx_path
