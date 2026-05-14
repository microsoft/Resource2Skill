def create_slide(
    output_pptx_path: str,
    title_line1: str = "Visitors rate their experience most",
    title_line2: str = "welcoming and fun!",
    data: dict = None,
    accent_color: tuple = (38, 172, 198),  # Cyan
    base_color: tuple = (142, 149, 160),   # Cool Grey
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Data-Driven Pill-Shape Bar Chart.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR

    # Default data matching the tutorial
    if data is None:
        data = {
            "Welcoming Atmosphere": 0.90,
            "Fun Place to Be": 0.86,
            "Quality Time with Friends & Family": 0.82,
            "Educational Experience": 0.80
        }

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.0), Inches(1.5))
    title_tf = title_box.text_frame
    title_p = title_tf.paragraphs[0]
    
    # First line (Grey)
    r1 = title_p.add_run()
    r1.text = title_line1 + "\n"
    r1.font.size = Pt(36)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(120, 120, 120)
    
    # Second line (Accent Color)
    r2 = title_p.add_run()
    r2.text = title_line2
    r2.font.size = Pt(40)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(*accent_color)

    # === Layer 2: Mathematical Layout & Bar Drawing ===
    start_y = Inches(2.8)
    bar_height = Inches(0.6)
    gap = Inches(0.35)
    start_x = Inches(3.5)
    max_bar_width = Inches(8.0)

    for i, (label, value) in enumerate(data.items()):
        y = start_y + i * (bar_height + gap)
        
        # Highlight top 2 items
        color = accent_color if i < 2 else base_color
        
        # Calculate visual width based on percentage
        tw = max_bar_width * value
        
        # 1. Bar Body (Standard Rectangle)
        # We subtract half the height so the rectangle ends exactly at the center of the oval
        rect = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            start_x, y, tw - (bar_height / 2), bar_height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(*color)
        rect.line.fill.solid()
        rect.line.fill.fore_color.rgb = RGBColor(*color) # Match fill to hide border natively
        
        # 2. Bar Tip (Perfect Circle / Oval)
        oval = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            start_x + tw - bar_height, y, bar_height, bar_height
        )
        oval.fill.solid()
        oval.fill.fore_color.rgb = RGBColor(*color)
        oval.line.fill.solid()
        oval.line.fill.fore_color.rgb = RGBColor(*color)
        
        # 3. Y-Axis Label
        lbl_box = slide.shapes.add_textbox(Inches(0.5), y, start_x - Inches(0.7), bar_height)
        lbl_tf = lbl_box.text_frame
        lbl_tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
        lbl_tf.margin_left = 0
        lbl_tf.margin_right = 0
        
        lbl_p = lbl_tf.paragraphs[0]
        lbl_p.text = label
        lbl_p.alignment = PP_ALIGN.RIGHT
        lbl_p.font.size = Pt(13)
        lbl_p.font.color.rgb = RGBColor(100, 100, 100)
        
        # 4. Data Value Label (Centered perfectly inside the Oval tip)
        val_box = slide.shapes.add_textbox(start_x + tw - bar_height, y, bar_height, bar_height)
        val_tf = val_box.text_frame
        val_tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
        # Remove margins so text centers perfectly in the small box
        val_tf.margin_left = 0
        val_tf.margin_right = 0
        val_tf.margin_top = 0
        val_tf.margin_bottom = 0
        
        val_p = val_tf.paragraphs[0]
        val_p.text = f"{value:.0%}"
        val_p.alignment = PP_ALIGN.CENTER
        val_p.font.size = Pt(14)
        val_p.font.bold = True
        val_p.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
