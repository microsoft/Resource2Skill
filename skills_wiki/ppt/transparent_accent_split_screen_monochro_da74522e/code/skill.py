import io
import urllib.request
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import OxmlElement
from pptx.oxml.ns import qn

def apply_shape_transparency(shape, alpha_percent):
    """
    Injects XML into a python-pptx shape to set its transparency.
    alpha_percent: float between 0 and 100 (e.g., 85 for 15% transparency).
    """
    # Convert percentage to PowerPoint's alpha value (100000 = 100% opaque)
    alpha_val = int((alpha_percent / 100.0) * 100000)
    
    spPr = shape.element.spPr
    solidFill = spPr.find(qn('a:solidFill'))
    if solidFill is not None:
        srgbClr = solidFill.find(qn('a:srgbClr'))
        if srgbClr is not None:
            # Check if alpha already exists
            alpha = srgbClr.find(qn('a:alpha'))
            if alpha is None:
                alpha = OxmlElement('a:alpha')
                srgbClr.append(alpha)
            alpha.set('val', str(alpha_val))

def create_slide(
    output_pptx_path: str,
    title_text: str = "New Product\nLaunch Tools\n& Techniques",
    circle_text: str = "Our Mission",
    footer_text: str = "Download this and 100,000s of other templates at example.com",
    bg_image_url: str = "https://images.unsplash.com/photo-1497366216548-37526070297c?q=80&w=1600&auto=format&fit=crop",
    accent_color: tuple = (136, 191, 165), # Soft Mint Green
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Transparent Accent Split-Screen design pattern.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Define color palette
    clr_accent = RGBColor(*accent_color)
    clr_text_dark = RGBColor(51, 51, 51)
    clr_footer = RGBColor(255, 230, 0)

    # === Layer 1: Background Split (65% Image / 35% White) ===
    split_ratio = 0.65
    img_width_in = prs.slide_width.inches * split_ratio
    img_height_in = prs.slide_height.inches

    # Fetch and process background image via PIL
    try:
        req = urllib.request.Request(bg_image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(io.BytesIO(response.read()))
            
        # Crop to target aspect ratio
        tgt_ratio = img_width_in / img_height_in
        w, h = img.size
        img_ratio = w / h
        
        if img_ratio > tgt_ratio:
            new_w = int(h * tgt_ratio)
            left = (w - new_w) / 2
            img = img.crop((left, 0, left + new_w, h))
        else:
            new_h = int(w / tgt_ratio)
            top = (h - new_h) / 2
            img = img.crop((0, top, w, top + new_h))
            
        # Convert to Grayscale to make the accent color pop
        img = img.convert('L').convert('RGB')
        
    except Exception as e:
        # Fallback if download fails: generate a dark grey placeholder
        print(f"Image download failed, using fallback: {e}")
        img = Image.new('RGB', (int(img_width_in*100), int(img_height_in*100)), color=(100, 100, 100))

    # Save processed image to memory and insert
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format='JPEG', quality=90)
    img_byte_arr.seek(0)
    
    slide.shapes.add_picture(img_byte_arr, Inches(0), Inches(0), width=Inches(img_width_in), height=Inches(img_height_in))

    # === Layer 2: Transparent Accent Circle ===
    circle_size = Inches(5.0)
    # Position: Vertically centered, horizontally bridging the split line
    circle_left = Inches(4.2) 
    circle_top = Inches((7.5 - 5.0) / 2)

    circle = slide.shapes.add_shape(
        9, # MSO_SHAPE.OVAL
        circle_left, circle_top, circle_size, circle_size
    )
    
    # Style the circle
    circle.fill.solid()
    circle.fill.fore_color.rgb = clr_accent
    circle.line.fill.background() # Remove outline
    
    # Inject XML for 85% opacity (15% transparency)
    apply_shape_transparency(circle, 85)

    # Add text inside the circle
    text_frame = circle.text_frame
    text_frame.text = circle_text
    text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    font = text_frame.paragraphs[0].font
    font.size = Pt(32)
    font.name = "Calibri"
    font.bold = True
    font.color.rgb = clr_text_dark

    # === Layer 3: Main Typography (Right Side) ===
    title_box = slide.shapes.add_textbox(
        left=Inches(8.8), top=Inches(2.5), width=Inches(4.0), height=Inches(2.5)
    )
    t_frame = title_box.text_frame
    t_frame.word_wrap = True
    p = t_frame.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.LEFT
    
    title_font = p.font
    title_font.size = Pt(44)
    title_font.name = "Calibri"
    title_font.bold = True
    title_font.color.rgb = clr_text_dark

    # === Layer 4: Footer Accent Bar ===
    footer_height = Inches(0.35)
    footer = slide.shapes.add_shape(
        1, # MSO_SHAPE.RECTANGLE
        Inches(0), Inches(7.5) - footer_height, Inches(13.333), footer_height
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = clr_footer
    footer.line.fill.background()
    
    # Footer Text
    f_frame = footer.text_frame
    f_frame.text = footer_text
    f_p = f_frame.paragraphs[0]
    f_p.alignment = PP_ALIGN.CENTER
    f_font = f_p.font
    f_font.size = Pt(11)
    f_font.name = "Calibri"
    f_font.color.rgb = clr_text_dark

    prs.save(output_pptx_path)
    return output_pptx_path
