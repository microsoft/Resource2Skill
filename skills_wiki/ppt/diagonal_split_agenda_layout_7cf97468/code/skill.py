def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda",
    bg_theme: str = "modern glass office building",
    accent_color: tuple = (66, 133, 244),  # Vibrant Corporate Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Diagonal Split Agenda Layout' visual effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    # Initialize presentation (16:9 widescreen)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Colors
    color_accent = RGBColor(*accent_color)
    color_white = RGBColor(255, 255, 255)
    color_text_dark = RGBColor(51, 51, 51)
    color_text_light = RGBColor(119, 119, 119)

    # === Layer 1: Background Image ===
    # Download a contextual background image
    img_url = "https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1600&auto=format&fit=crop"
    img_path = "temp_agenda_bg.jpg"
    
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
        
        # Place image to cover the right side (and overflow left, which will be masked)
        # We stretch slightly to ensure full coverage
        pic = slide.shapes.add_picture(img_path, Inches(3.0), 0, Inches(10.333), Inches(7.5))
    except Exception as e:
        print(f"Image download failed, using fallback fill: {e}")
        # Fallback: Draw a gray rectangle if image fails
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.0), 0, Inches(10.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(200, 200, 200)
        bg_shape.line.fill.background()

    # === Layer 2: The Diagonal Mask (White Freeform Polygon) ===
    # Define points for a polygon that covers the left side with a diagonal right edge
    top_x = 7.8
    bottom_x = 6.2
    
    builder = slide.shapes.build_freeform()
    builder.add_line_segments([
        (0, 0),                           # Top Left
        (0, Inches(7.5)),                 # Bottom Left
        (Inches(bottom_x), Inches(7.5)),  # Bottom Right (angled)
        (Inches(top_x), 0),               # Top Right (angled)
        (0, 0)                            # Close shape
    ], close=True)
    
    mask_shape = builder.convert_to_shape()
    mask_shape.fill.solid()
    mask_shape.fill.fore_color.rgb = color_white
    mask_shape.line.fill.background() # Remove outline

    # === Layer 3: Diagonal Accent Line ===
    accent_line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, 
        Inches(top_x), 0, 
        Inches(bottom_x), Inches(7.5)
    )
    accent_line.line.color.rgb = color_accent
    accent_line.line.width = Pt(4.5)

    # === Layer 4: Content Layout ===
    
    # 1. Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(4.0), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = color_text_dark

    # 2. Agenda Items (Loop to create list)
    start_y = 1.8
    spacing_y = 1.0
    
    agenda_items = [
        "Introduction & Overview",
        "Market Analysis & Trends",
        "Financial Projections Q3",
        "Strategic Initiatives",
        "Q&A Session"
    ]

    for i, topic in enumerate(agenda_items):
        current_y = start_y + (i * spacing_y)
        
        # A. Bullet Circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(1.0), Inches(current_y + 0.1), 
            Inches(0.35), Inches(0.35)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = color_accent
        circle.line.fill.background()
        
        # B. Text Block (Title + Description)
        tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(current_y), Inches(4.5), Inches(0.8))
        tx_frame = tx_box.text_frame
        tx_frame.word_wrap = True
        
        # Topic Title
        p_title = tx_frame.paragraphs[0]
        p_title.text = f"Agenda / {topic}"
        p_title.font.name = "Arial"
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = color_text_dark
        
        # Topic Description
        p_desc = tx_frame.add_paragraph()
        p_desc.text = "Review key metrics, website updates, and gather feedback to improve our ongoing resume formatting strategies."
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = color_text_light
        p_desc.space_before = Pt(3)

    # Cleanup temp image
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
