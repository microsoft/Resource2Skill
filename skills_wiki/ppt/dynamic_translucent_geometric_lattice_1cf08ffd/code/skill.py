def create_slide(
    output_pptx_path: str,
    title_text: str = "OFFICE TEMPLATE",
    subtitle_text: str = "FREE PPT TEMPLATES\nInsert the Subtitle of Your Presentation",
    accent_r: int = 242,
    accent_g: int = 153,
    accent_b: int = 74,
    **kwargs,
) -> str:
    """
    Creates a presentation slide featuring a Dynamic Translucent Geometric Lattice.
    Uses lxml to inject alpha transparency into native rotated shapes.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml
    
    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Helper: Apply precise alpha transparency to a shape using OpenXML
    def apply_transparency(shape, r, g, b, alpha_percent):
        # 100% alpha_percent means fully opaque, 0% means fully transparent in this context
        # Actually, let's define it as opacity_percent (0-100)
        opacity_val = int(alpha_percent * 1000) # DrawingML alpha val: 100000 is 100%
        
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(r, g, b)
        
        solid_fill = shape.element.spPr.solidFill
        if solid_fill is not None:
            srgbClr = solid_fill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
            if srgbClr is not None:
                # Remove any existing alpha tags
                for child in list(srgbClr):
                    if child.tag.endswith('alpha'):
                        srgbClr.remove(child)
                # Inject new alpha tag
                alpha_elem = parse_xml(f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{opacity_val}"/>')
                srgbClr.append(alpha_elem)
                
    # Helper: Create an oversized rotated rectangle that bleeds off the slide
    def add_angled_bar(cx, cy, width, height, rotation, r, g, b, opacity):
        # Calculate top-left based on center points (Inches)
        left = cx - (width / 2)
        top = cy - (height / 2)
        
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        shape.rotation = rotation
        shape.line.fill.background() # Remove outline
        
        # Apply color and opacity
        apply_transparency(shape, r, g, b, opacity)
        return shape

    # === Layer 1: Geometric Lattice Background ===
    
    # Color Palette generation based on accent color
    # Lighter variants for background shapes, darker for foreground/intersections
    c_light = (min(255, accent_r + 30), min(255, accent_g + 50), min(255, accent_b + 50))
    c_base = (accent_r, accent_g, accent_b)
    c_dark = (max(0, accent_r - 40), max(0, accent_g - 60), max(0, accent_b - 40))
    
    # Top Right Cluster (45 degrees)
    add_angled_bar(cx=10, cy=0, width=15, height=3.5, rotation=45, r=c_light[0], g=c_light[1], b=c_light[2], opacity=30)
    add_angled_bar(cx=12, cy=3, width=15, height=2.0, rotation=45, r=c_base[0], g=c_base[1], b=c_base[2], opacity=40)
    add_angled_bar(cx=9.5, cy=4.5, width=15, height=1.2, rotation=45, r=c_base[0], g=c_base[1], b=c_base[2], opacity=25)
    
    # Perpendicular Intersectors (-45 degrees / 135 degrees)
    add_angled_bar(cx=12, cy=1.5, width=15, height=1.5, rotation=135, r=c_base[0], g=c_base[1], b=c_base[2], opacity=35)
    add_angled_bar(cx=14, cy=6, width=15, height=2.5, rotation=135, r=c_light[0], g=c_light[1], b=c_light[2], opacity=20)
    
    # Bottom Left Accent Cluster
    add_angled_bar(cx=1, cy=8, width=10, height=2.0, rotation=45, r=c_dark[0], g=c_dark[1], b=c_dark[2], opacity=50)
    add_angled_bar(cx=3, cy=9, width=10, height=1.0, rotation=45, r=c_base[0], g=c_base[1], b=c_base[2], opacity=30)


    # === Layer 2: Typography & Content ===
    
    # Kicker / Small top category
    kicker_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5), Inches(0.5))
    kb_tf = kicker_box.text_frame
    p0 = kb_tf.paragraphs[0]
    p0.text = "LOGO / BRAND NAME"
    p0.font.name = "Arial"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(c_base[0], c_base[1], c_base[2])
    
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(8), Inches(1))
    tb_tf = title_box.text_frame
    p1 = tb_tf.paragraphs[0]
    p1.text = title_text
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.color.rgb = RGBColor(180, 130, 110) # Soft brownish gray
    
    # Primary Subtitle / Main Hook
    subtitle_lines = subtitle_text.split('\n')
    if len(subtitle_lines) > 0:
        sub1_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.8), Inches(8), Inches(1.5))
        s1_tf = sub1_box.text_frame
        
        p2 = s1_tf.paragraphs[0]
        p2.text = subtitle_lines[0]
        p2.font.name = "Arial"
        p2.font.size = Pt(44)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(30, 30, 30) # Dark Slate
        
        if len(subtitle_lines) > 1:
            p3 = s1_tf.add_paragraph()
            p3.text = "— " + subtitle_lines[1]
            p3.font.name = "Arial"
            p3.font.size = Pt(18)
            p3.font.color.rgb = RGBColor(100, 100, 100) # Medium Gray
            
    # Small Decorative Metadata badges (bottom left)
    meta_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(3), Inches(0.5))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = RGBColor(250, 240, 235)
    m_tf = meta_box.text_frame
    m_tf.word_wrap = False
    p4 = m_tf.paragraphs[0]
    p4.text = "REPORT : Q3 BUSINESS"
    p4.font.name = "Arial"
    p4.font.size = Pt(10)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(c_dark[0], c_dark[1], c_dark[2])

    prs.save(output_pptx_path)
    return output_pptx_path
