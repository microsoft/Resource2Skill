def create_slide(
    output_pptx_path: str,
    title_text: str = "CREATIVITY",
    body_text: str = "IS INVENTING,\nEXPERIMENTING,\nGROWING,\nTAKING RISKS,\nBREAKING RULES,\nAND HAVING FUN.",
    author_text: str = "ANGELO BREWING",
    bg_keyword: str = "abstract texture dark",
    accent_color: tuple = (255, 255, 255),  # Default to white
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Immersive Typographic Quote Poster' visual effect.
    Mixes hollow (outlined) text with solid bold text over an abstract background.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Background Image ===
    # Attempt to download an abstract background from Unsplash
    bg_image_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1600x900/?{bg_keyword.replace(' ', ',')}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(bg_image_path, "wb") as f:
                f.write(response.read())
        
        slide.shapes.add_picture(bg_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception as e:
        print(f"Failed to download image: {e}. Falling back to dark solid background.")
        bg = slide.shapes.add_shape(
            1, 0, 0, prs.slide_width, prs.slide_height  # 1 = msoShapeRectangle
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(20, 20, 30)
        bg.line.fill.background()

    # === Layer 2: Dark Overlay for Readability ===
    overlay = slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, prs.slide_height
    )
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(0, 0, 0)
    # Set transparency using lxml on the shape's fill
    fill_pr = overlay.fill._xPr
    solidFill = fill_pr.find('.//a:solidFill', namespaces=fill_pr.nsmap)
    if solidFill is not None:
        alpha = parse_xml('<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="60000"/>') # 60% opacity
        solidFill.find('.//a:srgbClr', namespaces=solidFill.nsmap).append(alpha)
    overlay.line.fill.background()

    # === Layer 3: Typography ===
    
    # Text Box dimensions
    left = Inches(1.5)
    top = Inches(1.0)
    width = Inches(10.0)
    height = Inches(5.5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.clear()

    # 1. The Hook (Hollow Text)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    run1 = p1.add_run()
    run1.text = title_text.upper() + "\n"
    run1.font.name = "Arial Black" # Fallback for thick font
    run1.font.size = Pt(85)
    run1.font.bold = True

    # --- LXML MAGIC: Apply Hollow Text Effect ---
    rPr = run1._r.get_or_add_rPr()
    
    # Convert accent color to hex for XML
    hex_color = f"{accent_color[0]:02X}{accent_color[1]:02X}{accent_color[2]:02X}"
    
    # 1. Remove solid fill
    noFill = parse_xml('<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
    rPr.append(noFill)
    
    # 2. Add outline (ln) - roughly 1.5pt width (19050 EMUs)
    ln_xml = f"""
    <a:ln w="19050" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:solidFill>
            <a:srgbClr val="{hex_color}"/>
        </a:solidFill>
    </a:ln>
    """
    ln = parse_xml(ln_xml)
    rPr.append(ln)

    # 2. The Body (Solid Text)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    p2.line_spacing = 0.9 # Tight line spacing characteristic of posters
    run2 = p2.add_run()
    run2.text = body_text.upper() + "\n"
    run2.font.name = "Arial Black"
    run2.font.size = Pt(55)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(255, 255, 255) # Always white for body

    # 3. The Author/Attribution
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.LEFT
    p3.space_before = Pt(20) # Add a little gap
    run3 = p3.add_run()
    # Simulate tracking/letter-spacing by joining with spaces
    tracked_author = "  ".join(list(author_text.upper()))
    run3.text = tracked_author
    run3.font.name = "Arial"
    run3.font.size = Pt(16)
    run3.font.bold = True
    run3.font.color.rgb = RGBColor(*accent_color)

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(bg_image_path):
        os.remove(bg_image_path)
        
    return output_pptx_path
