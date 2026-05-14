def create_slide(
    output_pptx_path: str,
    title_text: str = "Welcome back to\nmy YouTube channel",
    dark_color: tuple = (45, 45, 45),    # RGB for the main polygon
    light_color: tuple = (125, 125, 125),  # RGB for the background slice
    text_color: tuple = (255, 255, 255),   # RGB for the text
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Modern Diagonal Split Hero Layout" visual effect.
    This generates the custom freeform geometry and typographic layout.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation with standard 16:9 widescreen dimensions
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a completely blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background (Light Gray Base) ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*light_color)

    # === Layer 2: Visual Effect (Diagonal Dark Polygon) ===
    # Draw a custom polygon that covers the left side and cuts diagonally on the right
    ff_builder = slide.shapes.build_freeform()
    ff_builder.add_line_segments([
        (Inches(0), Inches(0)),         # Top Left
        (Inches(10.5), Inches(0)),      # Top Right (extends ~78% across)
        (Inches(7.0), Inches(7.5)),       # Bottom Right (angles back to ~52% across)
        (Inches(0), Inches(7.5)),       # Bottom Left
        (Inches(0), Inches(0))          # Close path back to Top Left
    ])
    
    diagonal_shape = ff_builder.convert_to_shape()
    
    # Style the polygon
    diagonal_shape.fill.solid()
    diagonal_shape.fill.fore_color.rgb = RGBColor(*dark_color)
    # Remove the border line to keep it clean
    diagonal_shape.line.fill.solid()
    diagonal_shape.line.fill.fore_color.rgb = RGBColor(*dark_color)

    # === Layer 3: Text & Content ===
    # Place text within the "safe zone" of the dark polygon
    left_margin = Inches(1.5)
    top_margin = Inches(2.5)
    width = Inches(7.0)
    height = Inches(2.5)

    txBox = slide.shapes.add_textbox(left_margin, top_margin, width, height)
    text_frame = txBox.text_frame
    text_frame.word_wrap = True
    
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.LEFT
    
    # Style the text to be bold, white, and highly legible
    font = p.font
    font.name = 'Arial'
    font.size = Pt(48)
    font.bold = True
    font.color.rgb = RGBColor(*text_color)

    # Note: To fully match the video, open the resulting PPTX, select the text box,
    # go to Animations -> Wipe -> From Left -> Effect Options -> Animate Text: By Word (10% delay).

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
