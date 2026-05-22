def create_slide(
    output_pptx_path: str,
    title_text: str = "5-Year Strategic Timeline",
    timeline_data: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Segmented Chevron Progression Timeline.
    
    Args:
        output_pptx_path: Path to save the file.
        title_text: The main slide title.
        timeline_data: List of dicts with 'year', 'title', and 'desc'. 
                       If None, default placeholder data is generated.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    
    # 1. Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Default data if none provided
    if not timeline_data:
        timeline_data = [
            {"year": "2020", "title": "Foundation", "desc": "Establish core infrastructure and finalize initial team build-out."},
            {"year": "2021", "title": "Alpha Launch", "desc": "Release initial product to closed beta testers and gather feedback."},
            {"year": "2022", "title": "Market Entry", "desc": "Public launch with targeted marketing campaigns in key regions."},
            {"year": "2023", "title": "Scaling Up", "desc": "Expand server capacity, introduce tier 2 features, and grow user base."},
            {"year": "2024", "title": "Global Reach", "desc": "Open international offices and localize product for major global markets."}
        ]

    # Pre-defined vibrant modern palette
    colors = [
        RGBColor(226, 54, 75),   # Red
        RGBColor(46, 204, 113),  # Green
        RGBColor(52, 152, 219),  # Blue
        RGBColor(155, 89, 182),  # Purple
        RGBColor(243, 156, 18),  # Orange
        RGBColor(52, 73, 94)     # Dark Blue/Gray (fallback)
    ]

    # 2. Add main title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(33, 33, 33)

    # 3. Layout calculations
    num_nodes = len(timeline_data)
    total_width = Inches(11.333) # Leaving 1 inch margin on each side
    start_left = Inches(1.0)
    
    # We introduce a slight gap to create a segmented modern look
    gap = Inches(0.05)
    shape_width = (total_width - (gap * (num_nodes - 1))) / num_nodes
    shape_height = Inches(1.2)
    shape_top = Inches(3.0) # Vertical center for the shapes

    # 4. Generate the Timeline
    for i, data in enumerate(timeline_data):
        current_left = start_left + (i * (shape_width + gap))
        current_color = colors[i % len(colors)]
        
        # --- Add Chevron Shape ---
        # The last shape is traditionally a pentagon (block arrow) to cap the timeline, 
        # but chevrons throughout look highly consistent. We will use chevrons for all.
        chevron = slide.shapes.add_shape(
            MSO_SHAPE.CHEVRON, 
            current_left, 
            shape_top, 
            shape_width, 
            shape_height
        )
        
        # Format shape
        chevron.fill.solid()
        chevron.fill.fore_color.rgb = current_color
        chevron.line.fill.background() # No outline
        
        # Add year text inside chevron
        tf = chevron.text_frame
        tf.text = data["year"]
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

        # --- Add Content Text Box Below ---
        # Position slightly below the chevron
        text_box_top = shape_top + shape_height + Inches(0.3)
        # Make the text box slightly wider than the shape to prevent ugly text wrapping,
        # but keep it centered relative to the chevron.
        box_width = shape_width * 1.2 
        box_left = current_left - ((box_width - shape_width) / 2)
        
        text_box = slide.shapes.add_textbox(box_left, text_box_top, box_width, Inches(2.0))
        text_box.text_frame.word_wrap = True
        
        # Phase Title (Colored to match shape)
        p_title = text_box.text_frame.paragraphs[0]
        p_title.text = data["title"]
        p_title.alignment = PP_ALIGN.CENTER
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = current_color
        
        # Description Text (Gray)
        p_desc = text_box.text_frame.add_paragraph()
        p_desc.text = data["desc"]
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = RGBColor(100, 100, 100)
        # Add slight space before description
        p_desc.space_before = Pt(6)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
