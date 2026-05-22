def create_slide(
    output_pptx_path: str,
    brand_name: str = "MY BUSINESS",
    person_name: str = "JOHN DOE",
    person_title: str = "Business Owner",
    brand_color: tuple = (200, 16, 46),  # Bold Red
    accent_color: tuple = (0, 0, 0),     # Black
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Dynamic Angled Split Layout effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # ======================================================
    # Layer 1: Left Branding Block (Angled Red Polygon)
    # ======================================================
    # Geometry: Starts at x=0, goes to x=4.0 at top, slants to x=5.5 at bottom
    ff_red = slide.shapes.build_freeform(0, 0)
    ff_red.add_line_segments([
        (Inches(4.0), 0),
        (Inches(5.5), Inches(7.5)),
        (0, Inches(7.5)),
        (0, 0)
    ])
    red_shape = ff_red.convert_to_shape()
    red_shape.fill.solid()
    red_shape.fill.fore_color.rgb = RGBColor(*brand_color)
    red_shape.line.fill.background() # Remove border

    # ======================================================
    # Layer 2: Top Right Header Block (Angled Black Polygon)
    # ======================================================
    # Geometry: Parallel to the red shape, creating a 0.2 inch visual gap.
    # Red edge equation: x = 4.0 + (1.5/7.5)*y = 4.0 + 0.2*y
    # Black edge start: x = 4.2 + 0.2*y
    # At y=0, x=4.2. At y=2.5, x=4.2 + 0.5 = 4.7
    ff_black = slide.shapes.build_freeform(Inches(4.2), 0)
    ff_black.add_line_segments([
        (Inches(13.333), 0),
        (Inches(13.333), Inches(2.5)),
        (Inches(4.7), Inches(2.5)),
        (Inches(4.2), 0)
    ])
    black_shape = ff_black.convert_to_shape()
    black_shape.fill.solid()
    black_shape.fill.fore_color.rgb = RGBColor(*accent_color)
    black_shape.line.fill.background()

    # ======================================================
    # Layer 3: Branding & Logo (Left Side)
    # ======================================================
    # Logo Placeholder
    logo = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5)
    )
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor(255, 255, 255)
    logo.line.fill.background()
    logo.text_frame.text = "LOGO"
    logo.text_frame.paragraphs[0].font.color.rgb = RGBColor(*brand_color)
    logo.text_frame.paragraphs[0].font.bold = True
    logo.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    logo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Brand Name Text
    tb_brand = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(3.5), Inches(1.5))
    p_brand = tb_brand.text_frame.paragraphs[0]
    p_brand.text = brand_name
    p_brand.font.bold = True
    p_brand.font.size = Pt(40)
    p_brand.font.color.rgb = RGBColor(255, 255, 255)
    p_brand.alignment = PP_ALIGN.CENTER

    # ======================================================
    # Layer 4: Personal / Section Info (Top Right)
    # ======================================================
    tb_name = slide.shapes.add_textbox(Inches(6.0), Inches(0.5), Inches(6.8), Inches(1.5))
    
    p_name = tb_name.text_frame.paragraphs[0]
    p_name.text = person_name
    p_name.font.bold = True
    p_name.font.size = Pt(36)
    p_name.font.color.rgb = RGBColor(255, 255, 255)
    p_name.alignment = PP_ALIGN.RIGHT

    p_title = tb_name.text_frame.add_paragraph()
    p_title.text = person_title
    p_title.font.size = Pt(20)
    p_title.font.color.rgb = RGBColor(200, 200, 200)
    p_title.alignment = PP_ALIGN.RIGHT

    # ======================================================
    # Layer 5: Contact Information List (Bottom Right)
    # ======================================================
    contact_data = [
        ("123 Business Road, Corporate District, 90210", "A"), # Address
        ("+1 (555) 123-4567", "P"),                           # Phone
        ("contact@mybusiness.com", "E"),                      # Email
        ("www.mybusiness.com", "W")                           # Web
    ]
    
    start_y = 3.2
    spacing = 0.85
    
    for i, (text, icon) in enumerate(contact_data):
        y_pos = Inches(start_y + (i * spacing))
        
        # Hexagon Icon Container
        hex_shape = slide.shapes.add_shape(
            MSO_SHAPE.HEXAGON, Inches(12.0), y_pos, Inches(0.6), Inches(0.6)
        )
        hex_shape.fill.solid()
        hex_shape.fill.fore_color.rgb = RGBColor(*brand_color)
        hex_shape.line.fill.background()
        
        # Simple text acting as icon placeholder inside Hexagon
        tf = hex_shape.text_frame
        tf.text = icon
        p_icon = tf.paragraphs[0]
        p_icon.font.color.rgb = RGBColor(255, 255, 255)
        p_icon.font.bold = True
        p_icon.font.size = Pt(14)
        p_icon.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Contact Details Text Box (Right aligned next to hexagon)
        tb_contact = slide.shapes.add_textbox(Inches(5.5), y_pos + Inches(0.05), Inches(6.3), Inches(0.5))
        p_contact = tb_contact.text_frame.paragraphs[0]
        p_contact.text = text
        p_contact.font.size = Pt(18)
        p_contact.font.color.rgb = RGBColor(40, 40, 40)
        p_contact.alignment = PP_ALIGN.RIGHT

    prs.save(output_pptx_path)
    return output_pptx_path
