def create_slide(
    output_pptx_path: str,
    title_text: str = "Neo-Dark Analytics Dashboard",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neo-Dark Analytics Dashboard visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Colors ---
    BG_COLOR = RGBColor(18, 18, 18)
    PANEL_COLOR = RGBColor(30, 30, 36)
    STROKE_COLOR = RGBColor(51, 51, 51)
    TEXT_MUTED = RGBColor(150, 150, 150)
    
    NEON_COLORS = [
        RGBColor(255, 0, 127),    # Neon Pink
        RGBColor(0, 240, 255),    # Neon Cyan
        RGBColor(0, 255, 136),    # Neon Green
        RGBColor(255, 140, 0)     # Neon Orange
    ]

    # --- Slide Background ---
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG_COLOR

    # --- Helper Functions ---
    def create_panel(slide, x, y, w, h, label_text):
        # Base panel
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = PANEL_COLOR
        panel.line.color.rgb = STROKE_COLOR
        panel.line.width = Pt(1)
        
        # Rotated Edge Tag
        tag_w, tag_h = Inches(1.8), Inches(0.35)
        tag_x = x - tag_w / 2
        tag_y = y + Inches(1.0) - tag_h / 2
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tag_x, tag_y, tag_w, tag_h)
        tag.rotation = -90
        
        # XML Injection for Tag Gradient Fill
        spPr = tag.element.spPr
        for elem in list(spPr):
            if elem.tag.endswith('Fill'):
                spPr.remove(elem)
        gradFill_xml = """
        <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
          <a:gsLst>
            <a:gs pos="0"><a:srgbClr val="FF4B2B"/></a:gs>
            <a:gs pos="100000"><a:srgbClr val="FF416C"/></a:gs>
          </a:gsLst>
          <a:lin ang="5400000" scaled="1"/>
        </a:gradFill>
        """
        spPr.append(parse_xml(gradFill_xml))
        tag.line.color.rgb = PANEL_COLOR
        tag.line.width = Pt(0)
        
        # Tag Text
        tf = tag.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = label_text
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.name = "Arial"
        p.alignment = PP_ALIGN.CENTER

    def style_chart(chart, is_circular=False):
        # Match chart background to panel background to simulate transparency
        chart.chart_area.format.fill.solid()
        chart.chart_area.format.fill.fore_color.rgb = PANEL_COLOR
        chart.chart_area.format.line.color.rgb = PANEL_COLOR
        chart.plot_area.format.fill.solid()
        chart.plot_area.format.fill.fore_color.rgb = PANEL_COLOR
        
        if not is_circular:
            val_axis = chart.value_axis
            val_axis.has_major_gridlines = True
            val_axis.major_gridlines.format.line.color.rgb = RGBColor(60, 60, 60)
            val_axis.format.line.color.rgb = RGBColor(80, 80, 80)
            val_axis.tick_labels.font.color.rgb = TEXT_MUTED
            val_axis.tick_labels.font.size = Pt(10)

            cat_axis = chart.category_axis
            cat_axis.format.line.color.rgb = RGBColor(80, 80, 80)
            cat_axis.tick_labels.font.color.rgb = TEXT_MUTED
            cat_axis.tick_labels.font.size = Pt(10)
        
        chart.has_legend = True
        chart.legend.include_in_layout = False
        chart.legend.font.color.rgb = TEXT_MUTED
        chart.legend.font.size = Pt(10)
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    # --- Layout Definitions ---
    p1_x, p1_y, p1_w, p1_h = Inches(1.3), Inches(0.8), Inches(7.3), Inches(2.8)
    p2_x, p2_y, p2_w, p2_h = Inches(8.9), Inches(0.8), Inches(3.5), Inches(2.8)
    p3_x, p3_y, p3_w, p3_h = Inches(1.3), Inches(3.9), Inches(7.3), Inches(2.8)
    p4_x, p4_y, p4_w, p4_h = Inches(8.9), Inches(3.9), Inches(3.5), Inches(2.8)

    # --- Title ---
    title_box = slide.shapes.add_textbox(Inches(1.3), Inches(0.2), Inches(10.0), Inches(0.4))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title_text.upper()
    title_p.font.size = Pt(16)
    title_p.font.bold = True
    title_p.font.color.rgb = TEXT_MUTED
    title_p.font.name = "Arial"

    # ==========================================
    # Panel 1: Line Chart
    # ==========================================
    create_panel(slide, p1_x, p1_y, p1_w, p1_h, "Line Chart")
    
    cd1 = CategoryChartData()
    cd1.categories = ['2021', '2022', '2023', '2024']
    cd1.add_series('Product A', (3.5, 5.2, 2.8, 6.1))
    cd1.add_series('Product B', (2.1, 4.8, 5.9, 4.2))

    line_chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, p1_x + Inches(0.4), p1_y + Inches(0.1), p1_w - Inches(0.6), p1_h - Inches(0.3), cd1
    ).chart
    style_chart(line_chart)
    
    for i, series in enumerate(line_chart.series):
        series.format.line.color.rgb = NEON_COLORS[i % len(NEON_COLORS)]
        series.format.line.width = Pt(2.5)
        # XML Injection: Smooth Lines
        smooth_xml = parse_xml('<c:smooth val="1" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
        series._element.append(smooth_xml)

    # ==========================================
    # Panel 2: Doughnut Chart
    # ==========================================
    create_panel(slide, p2_x, p2_y, p2_w, p2_h, "Doughnut")
    
    cd2 = CategoryChartData()
    cd2.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    cd2.add_series('Sales', (25, 30, 20, 25))

    doughnut = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, p2_x + Inches(0.3), p2_y + Inches(0.1), p2_w - Inches(0.5), p2_h - Inches(0.2), cd2
    ).chart
    style_chart(doughnut, is_circular=True)
    
    for i, pt in enumerate(doughnut.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = NEON_COLORS[i % len(NEON_COLORS)]

    # XML Injection: Adjust Doughnut Hole Size
    doughnut_el = doughnut._chartSpace.find(qn('c:chart')).find(qn('c:plotArea')).find(qn('c:doughnutChart'))
    if doughnut_el is not None:
        doughnut_el.append(parse_xml('<c:holeSize val="70" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>'))

    # ==========================================
    # Panel 3: Column Chart
    # ==========================================
    create_panel(slide, p3_x, p3_y, p3_w, p3_h, "Column Chart")
    
    cd3 = CategoryChartData()
    cd3.categories = ['2021', '2022', '2023', '2024']
    cd3.add_series('Metric 1', (4, 6, 3, 7))
    cd3.add_series('Metric 2', (3, 2, 5, 4))

    col_chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, p3_x + Inches(0.4), p3_y + Inches(0.1), p3_w - Inches(0.6), p3_h - Inches(0.3), cd3
    ).chart
    style_chart(col_chart)
    
    for i, series in enumerate(col_chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = NEON_COLORS[i % len(NEON_COLORS)]

    # ==========================================
    # Panel 4: Data Analysis Grid
    # ==========================================
    create_panel(slide, p4_x, p4_y, p4_w, p4_h, "Data Analysis")
    
    stats = [
        ("50K+", "Users", NEON_COLORS[0]),
        ("72K+", "Views", NEON_COLORS[1]),
        ("60K+", "Sales", NEON_COLORS[2]),
        ("99K+", "Leads", NEON_COLORS[3])
    ]

    start_x = p4_x + Inches(0.2)
    start_y = p4_y + Inches(0.3)
    
    for i, stat in enumerate(stats):
        col, row = i % 2, i // 2
        bx, by = start_x + (col * Inches(1.5)), start_y + (row * Inches(1.1))
        
        # Big Number
        tb_val = slide.shapes.add_textbox(bx, by, Inches(1.4), Inches(0.5))
        tf_val = tb_val.text_frame
        for attr in ['margin_left', 'margin_right', 'margin_top', 'margin_bottom']: setattr(tf_val, attr, 0)
        p_val = tf_val.paragraphs[0]
        p_val.text = stat[0]
        p_val.font.size = Pt(32)
        p_val.font.bold = True
        p_val.font.color.rgb = stat[2]
        p_val.alignment = PP_ALIGN.CENTER
        
        # Label
        tb_lbl = slide.shapes.add_textbox(bx, by + Inches(0.45), Inches(1.4), Inches(0.3))
        tf_lbl = tb_lbl.text_frame
        for attr in ['margin_left', 'margin_right', 'margin_top', 'margin_bottom']: setattr(tf_lbl, attr, 0)
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.text = stat[1]
        p_lbl.font.size = Pt(12)
        p_lbl.font.color.rgb = TEXT_MUTED
        p_lbl.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
