import os
import urllib.request
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda",
    bg_keyword: str = "skyscraper",
    **kwargs
) -> str:
    """
    Creates a PPTX file reproducing the Split-Panel Edge-Bleed Agenda design.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # --- 1. Background Image Panel (Left 35%) ---
    img_width = Inches(4.66)
    img_height = Inches(7.5)
    img_path = "temp_bg_image.jpg"
    
    # Try fetching a high-quality contextual image
    url = f"https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1000&auto=format&fit=crop"
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(img_path, 'wb') as f:
                f.write(response.read())
    except Exception:
        # Fallback to a solid dark-blue PIL image if download fails
        img = Image.new('RGB', (800, 1200), (20, 30, 50))
        img.save(img_path)

    # Insert the background picture
    slide.shapes.add_picture(img_path, Inches(0), Inches(0), img_width, img_height)

    # --- 2. Title Block (Semi-transparent Overlay) ---
    overlay_path = "temp_overlay.png"
    # Create a 60% opaque black rectangle (alpha=153 out of 255)
    Image.new('RGBA', (100, 100), (0, 0, 0, 153)).save(overlay_path)
    
    box_top = Inches(4.5)
    box_height = Inches(1.5)
    slide.shapes.add_picture(overlay_path, Inches(0), box_top, img_width, box_height)

    # Title Text
    tb = slide.shapes.add_textbox(Inches(0.5), box_top, img_width - Inches(0.5), box_height)
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.add_paragraph()
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(255, 255, 255)

    # --- 3. Agenda Items (Nodes, Lines, and Text) ---
    agenda_items = kwargs.get("agenda_items", [
        ("01 Organization Structure", "Review current team alignments and proposed reporting lines."),
        ("02 Action Items", "Review deliverables from Q2 and assign owners for Q3 OKRs."),
        ("03 Key Updates", "High-level overview of recent product launches and market response."),
        ("04 Head Count", "Current capacity, hiring freeze updates, and open critical roles."),
        ("05 Project Updates", "Status reports on Alpha, Beta, and Gamma strategic initiatives.")
    ])

    palette = [
        (89, 43, 145),   # Deep Purple
        (52, 101, 164),  # Professional Blue
        (115, 210, 22),  # Vibrant Green
        (204, 0, 0),     # Bold Red
        (237, 125, 49)   # Accent Orange
    ]

    num_items = len(agenda_items)
    start_y = Inches(1.2)
    end_y = Inches(6.3)
    spacing = (end_y - start_y) / (num_items - 1) if num_items > 1 else Inches(1)
    
    circle_radius = Inches(0.2)
    # Circle center sits exactly on the image edge (4.66 inches)
    circle_left = img_width - circle_radius

    for i, (item_title, item_desc) in enumerate(agenda_items):
        y_center = start_y + i * spacing
        color = palette[i % len(palette)]
        
        # Line Connector (drawn first so it goes behind the circle slightly)
        line_start_x = img_width + Inches(0.1)
        line_end_x = Inches(12.5)
        # 1 = MSO_CONNECTOR.STRAIGHT
        connector = slide.shapes.add_connector(1, line_start_x, y_center, line_end_x, y_center)
        connector.line.color.rgb = RGBColor(*color)
        connector.line.width = Pt(1.5)

        # Boundary Node (Circle)
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            circle_left, y_center - circle_radius,
            circle_radius * 2, circle_radius * 2
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(*color)
        # Thick white border creates a beautiful "cutout" effect over the image
        circle.line.color.rgb = RGBColor(255, 255, 255)
        circle.line.width = Pt(2.5)

        # Item Header (Sits above the line)
        title_box = slide.shapes.add_textbox(
            line_start_x, y_center - Inches(0.45), Inches(7.5), Inches(0.4)
        )
        tf_title = title_box.text_frame
        p_title = tf_title.add_paragraph()
        p_title.text = item_title
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(*color)

        # Item Description (Sits below the line)
        desc_box = slide.shapes.add_textbox(
            line_start_x, y_center + Inches(0.05), Inches(7.5), Inches(0.5)
        )
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        p_desc = tf_desc.add_paragraph()
        p_desc.text = item_desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = RGBColor(100, 100, 100)

    # Cleanup temp files
    prs.save(output_pptx_path)
    
    if os.path.exists(img_path):
        os.remove(img_path)
    if os.path.exists(overlay_path):
        os.remove(overlay_path)
        
    return output_pptx_path
