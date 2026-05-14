import os
import io
import urllib.request
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

def create_slide(
    output_pptx_path: str,
    title_text: str = "ABOUT US",
    primary_color: tuple = (216, 0, 50),     # Red
    secondary_color: tuple = (0, 43, 91),    # Navy Blue
    image_url: str = "https://images.unsplash.com/photo-1441986300917-64674bd600d8?auto=format&fit=crop&q=80&w=800&h=600",
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the 'Asymmetric Color-Block Overlap' visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: White Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Graphic Cluster (Left Side) ===
    
    # Primary Color Block (Red) - acts as the base layer of the cluster
    red_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.5), Inches(2.0), Inches(4.5), Inches(4.5)
    )
    red_rect.fill.solid()
    red_rect.fill.fore_color.rgb = RGBColor(*primary_color)
    red_rect.line.fill.background()

    # Process and Insert Main Image (Overlaps the Red Block)
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            image_data = response.read()
        img = Image.open(io.BytesIO(image_data)).convert("RGB")
    except Exception:
        # Fallback image generation if network fails
        img = Image.new("RGB", (800, 600), (230, 230, 230))
        draw = ImageDraw.Draw(img)
        draw.line((0, 0, 800, 600), fill=(180, 180, 180), width=3)
        draw.line((0, 600, 800, 0), fill=(180, 180, 180), width=3)

    # Crop image to exact specific dimensions (4.5in x 3.5in at 300dpi)
    target_w, target_h = int(4.5 * 300), int(3.5 * 300)
    img_cropped = ImageOps.fit(img, (target_w, target_h), Image.Resampling.LANCZOS)
    img_path = "temp_overlap_img.png"
    img_cropped.save(img_path)

    # Insert Image (Offset slightly to the top-left of the red block)
    slide.shapes.add_picture(
        img_path,
        Inches(0.8), Inches(1.0), Inches(4.5), Inches(3.5)
    )
    if os.path.exists(img_path):
        os.remove(img_path)

    # Secondary Accent Block (Navy) - Overlaps the bottom right of the red block
    navy_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.0), Inches(5.5), Inches(1.5), Inches(1.5)
    )
    navy_rect.fill.solid()
    navy_rect.fill.fore_color.rgb = RGBColor(*secondary_color)
    navy_rect.line.fill.background()

    # Decorative Text inside the Red Block
    id_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(3.5), Inches(0.5))
    tf = id_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Company Identity"
    p.font.name = 'Georgia'
    p.font.italic = True
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 3: Structured Typography (Right Side) ===

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.0), Inches(5.0), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = 'Arial Black'
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(20, 20, 20)

    # Separator Line
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(6.8), Inches(2.1), Inches(11.8), Inches(2.1)
    )
    line.line.color.rgb = RGBColor(*secondary_color)
    line.line.width = Pt(2)

    # Content Sections helper function
    def add_text_section(top_in, heading, text):
        # Heading
        h_box = slide.shapes.add_textbox(Inches(6.8), Inches(top_in), Inches(5.0), Inches(0.5))
        p = h_box.text_frame.paragraphs[0]
        p.text = heading
        p.font.name = 'Arial'
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(*secondary_color)

        # Body Text
        b_box = slide.shapes.add_textbox(Inches(6.8), Inches(top_in + 0.4), Inches(5.0), Inches(1.5))
        tf = b_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = 'Arial'
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(90, 90, 90)
        p.line_spacing = 1.3

    add_text_section(
        2.5, 
        "VISION", 
        "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris."
    )
    
    add_text_section(
        4.2, 
        "MISSION", 
        "Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur. Excepteur sint occaecat cupidatat non proident, sunt in culpa qui officia deserunt mollit anim."
    )

    prs.save(output_pptx_path)
    return output_pptx_path
