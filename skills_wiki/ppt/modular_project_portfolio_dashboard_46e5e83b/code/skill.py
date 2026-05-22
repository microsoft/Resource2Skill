import os
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

def create_donut_chart(percentage: int, color: tuple, size: tuple = (200, 200), track_color: tuple = (233, 236, 239), text_color: tuple = (33, 37, 41)) -> io.BytesIO:
    """
    Creates a donut chart as a transparent PNG image in a byte stream.
    """
    image = Image.new("RGBA", size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    
    bbox = (10, 10, size[0] - 10, size[1] - 10)
    width = 20

    # Draw the background track
    draw.arc(bbox, start=-90, end=270, fill=track_color, width=width)
    
    # Draw the foreground arc
    end_angle = -90 + (percentage / 100.0) * 360
    if percentage > 0:
        draw.arc(bbox, start=-90, end=end_angle, fill=color, width=width)

    # Draw the text in the center
    try:
        font = ImageFont.truetype("arialbd.ttf", 48)
    except IOError:
        font = ImageFont.load_default()

    text = f"{percentage}%"
    text_bbox = draw.textbbox((0, 0), text, font=font)
    text_width = text_bbox[2] - text_bbox[0]
    text_height = text_bbox[3] - text_bbox[1]
    
    text_x = (size[0] - text_width) / 2
    text_y = (size[1] - text_height) / 2
    draw.text((text_x, text_y), text, font=font, fill=text_color)
    
    img_byte_arr = io.BytesIO()
    image.save(img_byte_arr, format='PNG')
    img_byte_arr.seek(0)
    return img_byte_arr

def create_slide(
    output_pptx_path: str,
    title_text: str = "Project KPI Dashboard",
    project_data: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Modular Project Portfolio Dashboard visual effect.

    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Slide Title
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.33), Inches(0.5))
    p = title_shape.text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial Black'
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(33, 37, 41)
    
    # Default data if none provided
    if project_data is None:
        project_data = [
            {"name": "Project A", "color": (88, 114, 255), "summary": "On track for Q3 launch. Key milestones for UI/UX are completed.", "status": 50, "risk": 30, "resources": 90},
            {"name": "Project B", "color": (40, 167, 69), "summary": "Budget review pending. Experiencing minor delays in backend integration.", "status": 70, "risk": 50, "resources": 90},
            {"name": "Project C", "color": (23, 162, 184), "summary": "Awaiting final stakeholder feedback. All development tasks are 90% complete.", "status": 90, "risk": 15, "resources": 90}
        ]
        
    # --- Sidebar ---
    sidebar_left = Inches(11.2)
    sidebar_width = Inches(1.8)
    sidebar_top = Inches(1.0)
    
    sidebar_data = [
        {"value": "08", "title": "Ongoing Projects"},
        {"value": "$25.3M", "title": "Allocated Budget"},
        {"value": "50", "title": "Team Members"},
        {"value": "102", "title": "Task Pending"},
    ]
    
    y_pos = sidebar_top
    for item in sidebar_data:
        tb = slide.shapes.add_textbox(sidebar_left, y_pos, sidebar_width, Inches(0.8))
        tf = tb.text_frame
        tf.clear()

        p_val = tf.paragraphs[0]
        p_val.text = item["value"]
        p_val.font.name = 'Arial Black'
        p_val.font.size = Pt(28)
        p_val.font.color.rgb = RGBColor(33, 37, 41)
        
        p_title = tf.add_paragraph()
        p_title.text = item["title"]
        p_title.font.name = 'Arial'
        p_title.font.size = Pt(11)
        p_title.font.color.rgb = RGBColor(108, 117, 125)
        
        y_pos += Inches(1.2)

    # --- Project Rows ---
    start_y = Inches(1.0)
    row_height = Inches(2.0)
    row_gutter = Inches(0.2)
    
    col_starts = [Inches(0.5), Inches(2.2), Inches(5.8), Inches(7.5), Inches(9.2)]
    col_widths = [Inches(1.5), Inches(3.4), Inches(1.5), Inches(1.5), Inches(1.5)]

    for i, project in enumerate(project_data):
        current_y = start_y + i * (row_height + row_gutter)
        
        # Project Name Card
        proj_card = slide.shapes.add_textbox(col_starts[0], current_y, col_widths[0], row_height)
        proj_card.fill.solid()
        proj_card.fill.fore_color.rgb = RGBColor(*project["color"])
        tf = proj_card.text_frame
        tf.vertical_anchor = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = project["name"]
        p.font.name = 'Arial Bold'
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        kpi_defs = [
            {"title": "Summary", "type": "text", "value": project["summary"]},
            {"title": "Status Track", "type": "donut", "value": project["status"], "color": (255, 193, 7)},
            {"title": "Risk Analysis", "type": "donut", "value": project["risk"], "color": (253, 126, 20)},
            {"title": "Resources", "type": "donut", "value": project["resources"], "color": (220, 53, 69)}
        ]

        for j in range(1, 5):
            kpi = kpi_defs[j-1]
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col_starts[j], current_y, col_widths[j], row_height)
            card.shadow.inherit = False
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(248, 249, 250)
            card.line.fill.background()

            title_box = slide.shapes.add_textbox(col_starts[j] + Inches(0.1), current_y + Inches(0.1), col_widths[j] - Inches(0.2), Inches(0.3))
            p = title_box.text_frame.paragraphs[0]
            p.text = kpi['title']
            p.font.name = 'Arial Bold'
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(108, 117, 125)

            if kpi['type'] == 'text':
                body_box = slide.shapes.add_textbox(col_starts[j] + Inches(0.1), current_y + Inches(0.5), col_widths[j] - Inches(0.2), row_height - Inches(0.6))
                tf = body_box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = kpi['value']
                p.font.name = 'Arial'
                p.font.size = Pt(11)
                p.font.color.rgb = RGBColor(33, 37, 41)
            elif kpi['type'] == 'donut':
                donut_img_stream = create_donut_chart(kpi['value'], kpi['color'])
                img_size = Inches(1.2)
                img_left = col_starts[j] + (col_widths[j] - img_size) / 2
                img_top = current_y + (row_height - img_size) / 2 + Inches(0.1)
                slide.shapes.add_picture(donut_img_stream, img_left, img_top, width=img_size)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
if __name__ == '__main__':
    output_path = "modular_project_dashboard.pptx"
    create_slide(output_path)
    print(f"Slide saved to {output_path}")
    if os.name == 'nt':
        os.startfile(output_path)
