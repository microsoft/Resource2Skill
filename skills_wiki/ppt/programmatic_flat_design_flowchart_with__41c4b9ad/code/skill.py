def create_slide(
    output_pptx_path: str,
    title_text: str = "Process Flowchart",
    steps: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing a flat-design, perfectly aligned process flowchart 
    with a retro-styled footer band.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    if steps is None:
        steps = [
            "Step 1\nInitiate", 
            "Step 2\nProcess", 
            "Step 3\nReview", 
            "Step 4\nDeploy"
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    # Terracotta Red
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(214, 88, 70)

    # === Layer 2: Retro Footer Stripes ===
    stripe_colors = [
        RGBColor(34, 40, 49),   # Navy
        RGBColor(255, 193, 7),  # Gold
        RGBColor(142, 40, 40),  # Maroon
        RGBColor(235, 126, 61)  # Orange
    ]
    stripe_height = 0.15
    start_y = 7.5 - (len(stripe_colors) * stripe_height)
    
    for i, color in enumerate(stripe_colors):
        y_pos = start_y + (i * stripe_height)
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(0), Inches(y_pos), Inches(13.333), Inches(stripe_height)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.fill.background() # No line

    # === Layer 3: Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 4: Flowchart Generation ===
    num_nodes = len(steps)
    node_size = 1.8  # Size in inches (width and height for circle)
    
    # Calculate spacing
    total_width = 13.333
    total_node_width = num_nodes * node_size
    gap = (total_width - total_node_width) / (num_nodes + 1)
    
    center_y = 3.5 # Vertical center for nodes
    
    nodes = []
    
    # Draw Nodes
    for i, text in enumerate(steps):
        x_pos = gap + i * (node_size + gap)
        
        # Add Circle
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(x_pos), Inches(center_y - (node_size/2)), 
            Inches(node_size), Inches(node_size)
        )
        
        # Style Shape
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(38, 166, 154) # Teal
        shape.line.fill.background() # No border
        
        # Add Text
        shape.text_frame.text = text
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(18)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            
        nodes.append({
            "right_edge": x_pos + node_size,
            "left_edge": x_pos,
            "center_y": center_y
        })

    # Draw Connector Lines
    for i in range(num_nodes - 1):
        start_x = nodes[i]["right_edge"] + 0.1 # slight padding
        end_x = nodes[i+1]["left_edge"] - 0.1  # slight padding
        line_y = nodes[i]["center_y"]
        
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(start_x), Inches(line_y),
            Inches(end_x), Inches(line_y)
        )
        
        # Style Line
        connector.line.color.rgb = RGBColor(255, 255, 255)
        connector.line.width = Pt(4)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
