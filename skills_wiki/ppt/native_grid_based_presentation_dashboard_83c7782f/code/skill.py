def create_slide(
    output_pptx_path: str,
    title_text: str = "Executive Pivot Dashboard",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Native Grid-Based Presentation Dashboard visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Dashboard Header & UI Slicers ===
    # Main Dashboard Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.5))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(80, 80, 80)

    # Simulated Excel Slicers (UI context filters)
    slicers = ["Accessories", "Bikes", "Clothing", "Components"]
    start_x = Inches(4.8)
    for i, text in enumerate(slicers):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            start_x + Inches(i * 1.8), Inches(0.35), Inches(1.6), Inches(0.4)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(245, 245, 245)
        shape.line.color.rgb = RGBColor(200, 200, 200)
        
        tf_slicer = shape.text_frame
        tf_slicer.word_wrap = False
        tf_slicer.margin_left = Pt(0)
        tf_slicer.margin_right = Pt(0)
        
        p_slicer = tf_slicer.paragraphs[0]
        p_slicer.text = text
        p_slicer.font.size = Pt(12)
        p_slicer.font.bold = True
        p_slicer.font.color.rgb = RGBColor(100, 100, 100)
        p_slicer.alignment = PP_ALIGN.CENTER

    # === Layer 3: Data Visualization Grid ===
    
    # Chart 1: Bar Chart (Sales by Category) - Left Column
    chart_data1 = CategoryChartData()
    chart_data1.categories = ['Bikes', 'Accessories', 'Clothing', 'Components']
    chart_data1.add_series('Sales Volume', (150000, 80000, 45000, 120000))
    
    chart1 = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, 
        Inches(0.5), Inches(1.2), Inches(4), Inches(5.8), 
        chart_data1
    ).chart
    chart1.has_title = True
    chart1.chart_title.text_frame.text = 'Sales by Category'
    if chart1.has_legend:
        chart1.legend.position = XL_LEGEND_POSITION.BOTTOM

    # Chart 2: Line Chart (Sales Trend) - Top Right Panel
    chart_data2 = CategoryChartData()
    chart_data2.categories = ['2017', '2018', '2019', '2020']
    chart_data2.add_series('Revenue Trend', (200000, 250000, 220000, 300000))
    
    chart2 = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, 
        Inches(4.8), Inches(1.2), Inches(8), Inches(2.8), 
        chart_data2
    ).chart
    chart2.has_title = True
    chart2.chart_title.text_frame.text = 'Sales Trend (2017-2020)'
    if chart2.has_legend:
        chart2.legend.position = XL_LEGEND_POSITION.BOTTOM

    # Chart 3: Column Chart (Year on Year Change) - Bottom Right Panel
    chart_data3 = CategoryChartData()
    chart_data3.categories = ['2018', '2019', '2020']
    chart_data3.add_series('YoY % Change', (25, -12, 36))
    
    chart3 = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 
        Inches(4.8), Inches(4.2), Inches(8), Inches(2.8), 
        chart_data3
    ).chart
    chart3.has_title = True
    chart3.chart_title.text_frame.text = 'Year on Year Change (%)'
    
    # Add Data Labels to the Column Chart for a more dashboard-like feel
    plot3 = chart3.plots[0]
    plot3.has_data_labels = True
    for series in plot3.series:
        for point in series.points:
            point.data_label.font.size = Pt(10)
            point.data_label.font.color.rgb = RGBColor(60, 60, 60)
            
    if chart3.has_legend:
        chart3.legend.position = XL_LEGEND_POSITION.BOTTOM

    prs.save(output_pptx_path)
    return output_pptx_path
