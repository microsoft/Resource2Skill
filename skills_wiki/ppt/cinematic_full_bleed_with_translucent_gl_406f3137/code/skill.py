import os
import urllib.request
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "THE QUTUB MINAR",
    body_text: str = "Minar means a 'place of fire'. They are one of the popular symbols of Islam. They are the oldest form in Islamic architecture according to Muslim tradition.",
    bg_keyword: str = "tower,sky,architecture",
    panel_color_rgba: tuple = (10, 80, 160, 200),  # Deep bright blue, 78% opacity
    text_color_rgb: tuple = (255, 255, 255)
) -> str:
    """
    Creates a cinematic full-bleed presentation slide with a translucent content panel.
    """
    # Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    width_px, height_px = 1920, 1080

    # === Layer 1: Fetch/Generate Background ===
    base_img = None
    try:
        # Attempt to download a relevant high-res image
        url = f"https://images.unsplash.com/featured/1920x1080/?{bg_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            img_data = response.read()
            base_img = Image.open(BytesIO(img_data)).convert("RGBA")
            # Ensure it fits exactly 1920x1080
            base_img = base_img.resize((width_px, height_px), Image.Resampling.LANCZOS)
    except Exception as e:
        print(f"Image download failed, using fallback gradient. Error: {e}")
        # Fallback: Create a gradient sky-blue background
        base_img = Image.new("RGBA", (width_px, height_px), (135, 206, 235, 255))
        draw_base = ImageDraw.Draw(base_img)
        for y in range(height_px):
            r = int(135 + (20 * (y / height_px)))
            g = int(206 - (50 * (y / height_px)))
            b = int(235 - (80 * (y / height_px)))
            draw_base.line([(0, y), (width_px, y)], fill=(r, g, b, 255))

    # === Layer 2: Translucent Overlay Panel ===
    # Create an overlay layer for the semi-transparent panel
    overlay = Image.new("RGBA", (width_px, height_px), (0, 0, 0, 0))
    draw_overlay = ImageDraw.Draw(overlay)
    
    # We place the panel on the right side, occupying 55% of the screen width
    panel_left = int(width_px * 0.45)
    draw_overlay.rectangle(
        [(panel_left, 0), (width_px, height_px)],
        fill=panel_color_rgba
    )

    # Composite the overlay onto the base image
    final_bg = Image.alpha_composite(base_img, overlay)

    # Save to BytesIO to insert into PPTX
    bg_stream = BytesIO()
    final_bg.convert("RGB").save(bg_stream, format="PNG")
    bg_stream.seek(0)

    # Insert background into slide
    slide.shapes.add_picture(bg_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 3: Typography & Content ===
    
    # Setup coordinates for text relative to the translucent panel
    margin_left = Inches(13.333 * 0.45 + 0.5) # Start slightly inside the panel
    margin_top_title = Inches(1.5)
    text_width = Inches((13.333 * 0.55) - 1.0) # Width of panel minus padding
    
    # Add Title
    title_box = slide.shapes.add_textbox(margin_left, margin_top_title, text_width, Inches(1.5))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text.upper() # Force uppercase for that cinematic look
    p_title.font.name = 'Segoe UI'
    p_title.font.size = Pt(64)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(*text_color_rgb)
    p_title.alignment = PP_ALIGN.LEFT

    # Add Body Text
    body_box = slide.shapes.add_textbox(margin_left, margin_top_title + Inches(1.8), text_width, Inches(4.0))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = 'Segoe UI'
    p_body.font.size = Pt(28)
    p_body.font.color.rgb = RGBColor(*text_color_rgb)
    p_body.alignment = PP_ALIGN.LEFT
    
    # Apply soft line spacing (1.2 lines) for readability
    p_body.line_spacing = 1.2

    # Save the presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_slide("modern_cinematic_slide.pptx", title_text="What is a 'Minar'?", bg_keyword="monument,sky")
