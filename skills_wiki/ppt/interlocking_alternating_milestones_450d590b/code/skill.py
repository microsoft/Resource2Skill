def create_slide(
    output_pptx_path: str,
    title_text: str = "Infographic Timeline",
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Interlocking Alternating Milestones visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
    from lxml import etree
    import urllib.request
    from PIL import Image, ImageDraw
    import os

    # 1. Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 2. Background Generation (PIL Washout)
    bg_path = "temp_bg_infographic.png"
    width, height = 1920, 1080
    image = Image.new('RGB', (width, height))
    draw = ImageDraw.Draw(image)
    # Extremely subtle radial gradient (white to 3% gray)
    for r in range(width, 0, -10):
        c = int(255 - (r/width)*8)
        draw.ellipse(
            [(width/2 - r, height/2 - r), (width/2 + r, height/2 + r)], 
            fill=(c, c, c)
        )
    image.save(bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # 3. Add Slide Title
    title_box = slide.shapes.add_textbox(Inches(2.0), Inches(0.5), Inches(8), Inches(1))
    tf_title = title_box.text_frame
    tf_title.text = title_text
    tf_title.paragraphs[0].font.size = Pt(36)
    tf_title.paragraphs[0].font.bold = True
    tf_title.paragraphs[0].font.color.rgb = RGBColor(60, 60, 60)

    # 4. Helper to inject soft shadows
    def add_subtle_shadow(shape):
        spPr = shape.element.spPr
        effectLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        if effectLst is None:
            effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        else:
            effectLst.clear()
        
        # Soft shadow pointing straight down
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad="50800", dist="38100", dir="5400000", algn="tl", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="15000")

    # 5. Node Drawing Logic
    def draw_node(cx, cy, is_top_pointing, color, number_str, year_str, next_cx=None):
        R = Inches(0.5)
        r = Inches(0.35)
        L = Inches(1.0)
        d = Inches(0.08)
        
        # Layer 1: Horizontal connecting line to next node
        if next_cx:
            line = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, cx + R, cy, next_cx - R, cy
            )
            line.line.color.rgb = color
            line.line.width = Pt(6)
            
        # Layer 2: Outer Arc (Drawn over the connector for a clean seam)
        left, top = cx - R, cy - R
        arc = slide.shapes.add_shape(MSO_SHAPE.ARC, left, top, R*2, R*2)
        arc.line.color.rgb = color
        arc.line.width = Pt(6)
        try:
            # Leave a ~30 degree gap at top or bottom
            if is_top_pointing:
                arc.adjustments[0], arc.adjustments[1] = 285, 255
            else:
                arc.adjustments[0], arc.adjustments[1] = 105, 75
        except Exception:
            pass
            
        # Layer 3: Vertical Branch Line
        v_start_y = (cy - r + Inches(0.05)) if is_top_pointing else (cy + r - Inches(0.05))
        v_end_y = (cy - R - L) if is_top_pointing else (cy + R + L)
        v_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, v_start_y, cx, v_end_y)
        v_line.line.color.rgb = color
        v_line.line.width = Pt(3)
        
        # Layer 4: Inner Solid Core (Drawn over the vertical line start to hide the seam)
        inner_left, inner_top = cx - r, cy - r
        inner_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, inner_left, inner_top, r*2, r*2)
        inner_circle.fill.solid()
        inner_circle.fill.fore_color.rgb = color
        inner_circle.line.color.rgb = color  # Avoid default outline
        add_subtle_shadow(inner_circle)
        
        # Text inside core
        tf = inner_circle.text_frame
        tf.text = number_str
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        # Layer 5: Terminal Dot
        dot_left, dot_top = cx - d, v_end_y - d
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, dot_left, dot_top, d*2, d*2)
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.color.rgb = color
        
        # Layer 6: Thin Horizontal Separator Line
        thin_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, v_end_y, cx + Inches(1.5), v_end_y)
        thin_line.line.color.rgb = color
        thin_line.line.width = Pt(1.5)
        
        # Layer 7: Text Blocks
        text_width = Inches(1.8)
        text_left = cx + Inches(0.05)
        
        def configure_tf(textbox, v_anchor):
            tframe = textbox.text_frame
            tframe.vertical_anchor = v_anchor
            tframe.word_wrap = True
            tframe.margin_left = tframe.margin_right = tframe.margin_top = tframe.margin_bottom = 0
            return tframe

        if is_top_pointing:
            # Block Above Line (Year + Subtitle)
            tb_top = slide.shapes.add_textbox(text_left, v_end_y - Inches(0.72), text_width, Inches(0.7))
            tf_top = configure_tf(tb_top, MSO_VERTICAL_ANCHOR.BOTTOM)
            
            p1 = tf_top.paragraphs[0]
            p1.text = year_str
            p1.font.size, p1.font.bold, p1.font.color.rgb = Pt(18), True, RGBColor(80, 80, 80)
            p1.space_after = Pt(0)
            
            p2 = tf_top.add_paragraph()
            p2.text = "Lorem Ipsum"
            p2.font.size, p2.font.bold, p2.font.color.rgb = Pt(12), True, color
            p2.space_before = Pt(0)
            
            # Block Below Line (Paragraph)
            tb_bot = slide.shapes.add_textbox(text_left, v_end_y + Inches(0.02), text_width, Inches(1.0))
            tf_bot = configure_tf(tb_bot, MSO_VERTICAL_ANCHOR.TOP)
            p3 = tf_bot.paragraphs[0]
            p3.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor."
            p3.font.size, p3.font.color.rgb = Pt(9), RGBColor(120, 120, 120)
            
        else:
            # Block Above Line (Paragraph)
            tb_top = slide.shapes.add_textbox(text_left, v_end_y - Inches(1.02), text_width, Inches(1.0))
            tf_top = configure_tf(tb_top, MSO_VERTICAL_ANCHOR.BOTTOM)
            p1 = tf_top.paragraphs[0]
            p1.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor."
            p1.font.size, p1.font.color.rgb = Pt(9), RGBColor(120, 120, 120)
            
            # Block Below Line (Subtitle + Year)
            tb_bot = slide.shapes.add_textbox(text_left, v_end_y + Inches(0.02), text_width, Inches(0.7))
            tf_bot = configure_tf(tb_bot, MSO_VERTICAL_ANCHOR.TOP)
            
            p2 = tf_bot.paragraphs[0]
            p2.text = "Lorem Ipsum"
            p2.font.size, p2.font.bold, p2.font.color.rgb = Pt(12), True, color
            p2.space_after = Pt(0)
            
            p3 = tf_bot.add_paragraph()
            p3.text = year_str
            p3.font.size, p3.font.bold, p3.font.color.rgb = Pt(18), True, RGBColor(80, 80, 80)
            p3.space_before = Pt(0)

    # 6. Build the Timeline Data
    years = ["2017", "2018", "2019", "2020", "2021"]
    colors = [
        RGBColor(42, 75, 124),   # Dark Blue
        RGBColor(61, 178, 211),  # Cyan
        RGBColor(242, 156, 56),  # Orange
        RGBColor(217, 78, 52),   # Red
        RGBColor(139, 168, 75)   # Green
    ]

    start_cx = Inches(2.0)
    spacing = Inches(2.333)
    cy = Inches(4.0)

    for i in range(5):
        cx = start_cx + i * spacing
        is_top = (i % 2 == 0)
        next_cx = (start_cx + (i + 1) * spacing) if i < 4 else None
        
        draw_node(
            cx=cx, cy=cy, 
            is_top_pointing=is_top, 
            color=colors[i], 
            number_str=f"0{i+1}", 
            year_str=years[i], 
            next_cx=next_cx
        )

    # 7. Save and Cleanup
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
