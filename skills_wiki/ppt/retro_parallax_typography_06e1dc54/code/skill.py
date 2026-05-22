def create_slide(
    output_pptx_path: str,
    title_text: str = "Parallax",
    body_text: str = "",
    bg_palette: str = "retro",
    accent_color: tuple = (142, 93, 75),  # Deep brown
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Retro Parallax Typography" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Color Palette
    bg_color = RGBColor(249, 244, 240)      # Cream/Beige
    shadow_color = RGBColor(142, 93, 75)    # Deep Brown
    main_color = RGBColor(26, 26, 26)       # Off-Black
    outline_color = RGBColor(255, 255, 255) # White
    
    # 1. Set Slide Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # Helper function to inject Slant Up WordArt transform
    def apply_slant_transform(shape):
        bodyPr = shape.text_frame._bodyPr
        warp_xml = r'''
        <a:prstTxWarp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" prst="slantUp">
            <a:avLst/>
        </a:prstTxWarp>
        '''
        warp = parse_xml(warp_xml)
        bodyPr.append(warp)

    # Helper function to create a base text box
    def create_text_layer(x, y, text, font_name="Brush Script MT", font_size=160):
        tx_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(10), Inches(4))
        tf = tx_box.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = font_name
        run.font.size = Pt(font_size)
        apply_slant_transform(tx_box)
        return run

    # Center-ish anchor coordinates
    base_x, base_y = 1.6, 1.5

    # ==========================================
    # LAYER 1: Bottom Drop Shadow (Solid Brown)
    # ==========================================
    offset_1 = 0.25
    run1 = create_text_layer(base_x + offset_1, base_y + offset_1, title_text)
    run1.font.color.rgb = shadow_color

    # ==========================================
    # LAYER 2: Middle Striped Shadow
    # ==========================================
    offset_2 = 0.12
    run2 = create_text_layer(base_x + offset_2, base_y + offset_2, title_text)
    
    # Inject Narrow Horizontal Pattern Fill
    rPr2 = run2._r.get_or_add_rPr()
    # Remove default solid fill if present to avoid conflicts
    for fill_tag in rPr2.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill'):
        rPr2.remove(fill_tag)
        
    pattern_xml = f'''
    <a:pattFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" prst="narHorz">
        <a:fgClr><a:srgbClr val="{main_color}"/></a:fgClr>
        <a:bgClr><a:srgbClr val="{bg_color}"/></a:bgClr>
    </a:pattFill>
    '''
    rPr2.append(parse_xml(pattern_xml))

    # ==========================================
    # LAYER 3: Top Main Text (Black + White Stroke + Highlight)
    # ==========================================
    run3 = create_text_layer(base_x, base_y, title_text)
    run3.font.color.rgb = main_color
    rPr3 = run3._r.get_or_add_rPr()

    # 3a. Inject Thick White Outline (63500 EMUs = 5pt stroke)
    outline_xml = f'''
    <a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="63500">
        <a:solidFill><a:srgbClr val="{outline_color}"/></a:solidFill>
    </a:ln>
    '''
    rPr3.append(parse_xml(outline_xml))

    # 3b. Inject Inner Shadow to simulate the cutout highlight on the left edge
    # dir="10800000" is 180 degrees (Left), dist="38100" is 3pt distance
    highlight_xml = f'''
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:innerShdw blurRad="0" dist="38100" dir="10800000">
            <a:srgbClr val="{bg_color}"/>
        </a:innerShdw>
    </a:effectLst>
    '''
    rPr3.append(parse_xml(highlight_xml))

    prs.save(output_pptx_path)
    return output_pptx_path
