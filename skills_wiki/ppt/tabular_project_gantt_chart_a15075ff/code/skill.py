import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.table import MSO_VERTICAL_ANCHOR
from lxml import etree

def _set_cell_properties(cell, text, bold=False, font_size=10, font_color=RGBColor(0, 0, 0), bg_color=None, align=PP_ALIGN.CENTER, v_align=MSO_VERTICAL_ANCHOR.MIDDLE):
    """Helper to style a table cell."""
    cell.text = text
    cell.vertical_anchor = v_align
    
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color
        
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    if not p.runs:
        p.add_run()
    run = p.runs[0]
    run.font.bold = bold
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color

def _set_no_outline(shape):
    """Remove the outline from a shape by manipulating its XML properties."""
    tree = shape.element
    ln = tree.xpath('.//a:ln', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if ln:
        ln[0].getparent().remove(ln[0])

def create_slide(
    output_pptx_path: str,
    chart_title: str = "Gantt Chart (3 Months Activity)",
    tasks: list = None,
    timeline_data: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file with a professional Gantt chart.

    Args:
        output_pptx_path: Path to save the generated PPTX file.
        chart_title: The main title for the Gantt chart slide.
        tasks: A list of task name strings for the rows.
        timeline_data: A list of dictionaries defining the timeline bars.
                       Each dict should contain: 'row', 'start_week', 'duration', 'color', and optional 'milestone'.

    Returns:
        The path to the saved PPTX file.
    """
    # Use default data if none is provided
    if tasks is None:
        tasks = [
            "Market Research", "Specifications", "Planning", "Design",
            "Development", "Training", "Assessment", "Documentation"
        ]
    if timeline_data is None:
        timeline_data = [
            {'row': 0, 'start_week': 1, 'duration': 3, 'color': (47, 82, 143)},
            {'row': 1, 'start_week': 2, 'duration': 4, 'color': (47, 82, 143)},
            {'row': 2, 'start_week': 5, 'duration': 2, 'color': (237, 125, 49), 'milestone': 7},
            {'row': 3, 'start_week': 7, 'duration': 3, 'color': (112, 48, 160)},
            {'row': 4, 'start_week': 8, 'duration': 4, 'color': (237, 125, 49), 'milestone': 11},
            {'row': 5, 'start_week': 9, 'duration': 3, 'color': (47, 82, 143)},
            {'row': 6, 'start_week': 10, 'duration': 2, 'color': (112, 48, 160), 'milestone': 12},
            {'row': 6, 'start_week': 11, 'duration': 2, 'color': (112, 48, 160)},
            {'row': 7, 'start_week': 12, 'duration': 1, 'color': (112, 48, 160)},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === 1. Chart Title ===
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.33), Inches(0.75))
    p = title_shape.text_frame.paragraphs[0]
    p.text = chart_title
    p.font.bold = True
    p.font.size = Pt(28)
    p.alignment = PP_ALIGN.LEFT

    # === 2. Create and Style the Table Grid ===
    num_tasks = len(tasks)
    rows, cols = 2 + num_tasks, 13
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.0), Inches(12.33), Inches(0.5 * rows))
    table = table_shape.table

    table.columns[0].width = Inches(3.0)
    week_col_width = Inches((12.33 - 3.0) / 12)
    for i in range(1, cols):
        table.columns[i].width = week_col_width

    # Month Headers
    month_colors = [RGBColor(23, 54, 93), RGBColor(237, 125, 49), RGBColor(112, 173, 71)]
    for i in range(3):
        cell_start = table.cell(0, 1 + i * 4)
        cell_end = table.cell(0, 4 + i * 4)
        cell_start.merge(cell_end)
        _set_cell_properties(cell_start, f"Month 0{i+1}", bold=True, font_size=16, font_color=RGBColor(255, 255, 255), bg_color=month_colors[i])

    # Week Headers
    for i in range(12):
        _set_cell_properties(table.cell(1, i + 1), f"Week 0{(i % 4) + 1}", font_size=10, bg_color=RGBColor(222, 235, 247))

    # Task Column
    for i, task_name in enumerate(tasks):
        cell = table.cell(i + 2, 0)
        _set_cell_properties(cell, task_name, bold=True, font_size=14, font_color=RGBColor(255, 255, 255), bg_color=RGBColor(127, 127, 127), align=PP_ALIGN.LEFT)
        cell.text_frame.margin_left = Inches(0.1)

    # === 3. Draw Timeline and Milestone Markers ===
    table_left, table_top = table_shape.left, table_shape.top
    task_col_width = table.columns[0].width
    header_height = table.rows[0].height + table.rows[1].height
    row_height = table.rows[2].height

    for item in timeline_data:
        task_row_idx, start_week, duration, color_rgb = item['row'], item['start_week'], item['duration'], item['color']
        milestone_week = item.get('milestone')

        marker_left = table_left + task_col_width + ((start_week - 1) * week_col_width)
        marker_top = table_top + header_height + (task_row_idx * row_height) + (row_height * 0.2)
        marker = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, marker_left, marker_top, duration * week_col_width, row_height * 0.6)
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(*color_rgb)
        _set_no_outline(marker)
        
        tf = marker.text_frame
        p = tf.paragraphs[0]
        p.text = "Enter Date Here"
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.LEFT
        tf.margin_left, tf.vertical_anchor = Inches(0.05), MSO_VERTICAL_ANCHOR.MIDDLE

        if milestone_week:
            milestone_left = table_left + task_col_width + ((milestone_week - 1) * week_col_width) + (week_col_width * 0.7)
            flag = slide.shapes.add_shape(MSO_SHAPE.WAVE, milestone_left, marker_top - (row_height * 0.1), row_height * 0.5, row_height * 0.5)
            flag.rotation = 90
            flag.fill.solid()
            flag.fill.fore_color.rgb = RGBColor(112, 173, 71)
            _set_no_outline(flag)

    # === 4. Add Legend ===
    legend_items = {"Team A": (47, 82, 143), "Team B": (237, 125, 49), "Team C": (112, 48, 160), "Milestone": (112, 173, 71)}
    legend_x, legend_y = Inches(0.5), table_top + table_shape.height + Inches(0.3)
    for label, color in legend_items.items():
        if label == "Milestone":
            shape = slide.shapes.add_shape(MSO_SHAPE.WAVE, legend_x, legend_y + Inches(0.05), Inches(0.2), Inches(0.2))
            shape.rotation = 90
        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, legend_x, legend_y, Inches(0.5), Inches(0.25))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        _set_no_outline(shape)
        
        p = slide.shapes.add_textbox(legend_x + Inches(0.6), legend_y - Inches(0.05), Inches(1.5), Inches(0.3)).text_frame.paragraphs[0]
        p.text = label
        p.font.size = Pt(12)
        legend_x += Inches(2.2)

    prs.save(output_pptx_path)
    return output_pptx_path

