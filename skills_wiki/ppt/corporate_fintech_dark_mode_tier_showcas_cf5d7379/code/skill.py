import os
import requests
from io import BytesIO
from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "CASH FOREX GROUP",
    body_text: str = "Select your Academy Pack to begin your journey.",
    bg_keyword: str = "laptop,dark",
    accent_color: tuple = (150, 210, 50),  # CFX Neon Green
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Corporate Fintech Dark Mode & Tier Showcase' style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Generation via PIL ===
    # Attempt to download a thematic background image, fallback to dark gray
    bg_img_path = "temp_bg.jpg"
    try:
        response = requests.get(f"https://source.unsplash.com/1920x1080/?{bg_keyword}", timeout=5)
        response.raise_for_status()
        img = Image.open(BytesIO(response.content)).convert("RGB")
    except Exception:
        # Fallback if network fails
        img = Image.new("RGB", (1920, 1080), (30, 30, 35))

    # Darken the image to create the "Financial Dark Mode" aesthetic
    enhancer = ImageEnhance.Brightness(img)
    img_dark = enhancer.enhance(0.25)  # Reduce brightness to 25%
    img_dark.save(bg_img_path, format="JPEG", quality=90)

    # Insert background
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Titles and Header ===
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*accent_color)

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(220, 220, 220)

    # === Layer 3: Tier Cards ===
    # Configuration for the 3 pricing tiers
    tiers = [
        {"name": "ELEMENTAL", "price": "300 PV", "color": (170, 175, 180)},  # Silver
        {"name": "SUPREME", "price": "500 PV", "color": (50, 150, 220)},     # Blue
        {"name": "ADVANCED", "price": "1K PV", "color": (220, 50, 60)}       # Red
    ]

    card_width = Inches(3.0)
    card_height = Inches(4.5)
    spacing = Inches(1.0)
    start_x = (prs.slide_width - (card_width * 3 + spacing * 2)) / 2
    start_y = Inches(2.2)

    for i, tier in enumerate(tiers):
        x = start_x + i * (card_width + spacing)
        
        # Draw Card Base (Dark Gray with matching border)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, start_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(40, 40, 45)
        card.line.color.rgb = RGBColor(*tier["color"])
        card.line.width = Pt(2)
        
        # Tier Header Shape (Colored top section)
        header_height = Inches(1.0)
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, start_y, card_width, header_height)
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor(*tier["color"])
        header.line.fill.background()
        
        # Tier Name Text
        h_tf = header.text_frame
        h_tf.vertical_anchor = PP_ALIGN.CENTER
        h_p = h_tf.paragraphs[0]
        h_p.text = tier["name"]
        h_p.alignment = PP_ALIGN.CENTER
        h_p.font.name = "Arial"
        h_p.font.size = Pt(20)
        h_p.font.bold = True
        h_p.font.color.rgb = RGBColor(255, 255, 255)

        # Price / Value Text
        price_box = slide.shapes.add_textbox(x, start_y + Inches(1.5), card_width, Inches(1.0))
        p_tf = price_box.text_frame
        p_p = p_tf.paragraphs[0]
        p_p.text = tier["price"]
        p_p.alignment = PP_ALIGN.CENTER
        p_p.font.name = "Arial"
        p_p.font.size = Pt(36)
        p_p.font.bold = True
        p_p.font.color.rgb = RGBColor(*tier["color"])

        # Divider Line
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.5), start_y + Inches(2.5), Inches(2.0), Pt(1))
        div.fill.solid()
        div.fill.fore_color.rgb = RGBColor(100, 100, 100)
        div.line.fill.background()

        # Features Text
        feat_box = slide.shapes.add_textbox(x + Inches(0.2), start_y + Inches(2.7), card_width - Inches(0.4), Inches(1.5))
        f_tf = feat_box.text_frame
        f_tf.word_wrap = True
        features = ["✓ Academy Program", "✓ Trading Pool", "✓ Leadership Points"]
        for idx, feature in enumerate(features):
            fp = f_tf.paragraphs[idx] if idx == 0 else f_tf.add_paragraph()
            fp.text = feature
            fp.font.name = "Arial"
            fp.font.size = Pt(14)
            fp.font.color.rgb = RGBColor(200, 200, 200)

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
