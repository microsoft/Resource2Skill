def create_slide(
    output_pptx_path: str,
    title_text: str = "BUSINESS INFOGRAPHIC",
    **kwargs,
) -> str:
    """
    Creates a symmetrical 10-step neural/circuit node infographic using native custom geometry.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # --- Color Palette ---
    colors_left = [
        RGBColor(156, 39, 176),   # Purple
        RGBColor(63, 81, 181),    # Indigo
        RGBColor(0, 188, 212),    # Cyan
        RGBColor(76, 175, 80),    # Green
        RGBColor(139, 195, 74)    # Light Green
    ]
    colors_right = [
        RGBColor(233, 30, 99),    # Pink
        RGBColor(244, 67, 54),    # Red
        RGBColor(255, 152, 0),    # Orange
        RGBColor(255, 193, 7),    # Amber
        RGBColor(255, 235, 59)    # Yellow
    ]
    center_color = RGBColor(38, 50, 56) # Dark Slate
    text_color_dark = RGBColor(50, 50, 50)
    text_color_light = RGBColor(120, 120, 120)

    # --- Coordinates & Measurements ---
    cx, cy = 13.333 / 2, 7.5 / 2
    y_offsets = [-2.3, -1.15, 0, 1.15, 2.3]  # Vertical spread
    
    rect_w, rect_h = 1.3, 0.7
    left_rect_x = 3.0
    right_rect_x = 13.333 - 3.0 - rect_w
    
    ribbon_thickness = 0.2
    
    # Helper to create text blocks
    def add_text_block(x, y, w, h, index, color, alignment):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y - 0.2), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Heading Paragraph
        p1 = tf.paragraphs[0]
        p1.alignment = alignment
        
        run_num = p1.add_run()
        run_num.text = f"{index:02d} "
        run_num.font.size = Pt(16)
        run_num.font.bold = True
        run_num.font.color.rgb = color
        
        run_head = p1.add_run()
        run_head.text = "HEADING"
        run_head.font.size = Pt(16)
        run_head.font.bold = True
        run_head.font.color.rgb = text_color_dark
        
        # Body Paragraph
        p2 = tf.add_paragraph()
        p2.alignment = alignment
        p2.space_before = Pt(5)
        run_body = p2.add_run()
        run_body.text = "Some text goes here. Some text goes here. Some text goes here."
        run_body.font.size = Pt(10)
        run_body.font.color.rgb = text_color_light

    # --- Draw Left Nodes ---
    for i, y_off in enumerate(y_offsets):
        node_y = cy + y_off
        color = colors_left[i]
        
        # 1. Custom Branching Ribbon (Freeform Polygon)
        # We start under the center circle, angle out, and seamlessly overlap the rounded rectangle.
        ffb = slide.shapes.build_freeform(Inches(cx - 0.8), Inches(cy - (ribbon_thickness/2)))
        ffb.add_line_segments([
            (Inches(cx - 1.6), Inches(node_y - (ribbon_thickness/2))),
            (Inches(left_rect_x + rect_w - 0.2), Inches(node_y - (ribbon_thickness/2))),
            (Inches(left_rect_x + rect_w - 0.2), Inches(node_y + (ribbon_thickness/2))),
            (Inches(cx - 1.6), Inches(node_y + (ribbon_thickness/2))),
            (Inches(cx - 0.8), Inches(cy + (ribbon_thickness/2))),
            (Inches(cx - 0.8), Inches(cy - (ribbon_thickness/2))) # Close
        ])
        ribbon = ffb.convert_to_shape()
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = color
        ribbon.line.fill.background() # No border
        
        # 2. Rounded Rectangle Endpoint
        rect_y = node_y - (rect_h / 2)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(left_rect_x), Inches(rect_y), Inches(rect_w), Inches(rect_h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        
        # 3. Icon Placeholder (White inner circle)
        icon_size = 0.4
        icon_x = left_rect_x + 0.15
        icon_y = node_y - (icon_size / 2)
        icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(icon_x), Inches(icon_y), Inches(icon_size), Inches(icon_size))
        icon.fill.solid()
        icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
        icon.line.fill.background()
        
        # 4. Text Block
        add_text_block(left_rect_x - 2.6, rect_y, 2.4, 1.0, i + 1, color, PP_ALIGN.RIGHT)

    # --- Draw Right Nodes ---
    for i, y_off in enumerate(y_offsets):
        node_y = cy + y_off
        color = colors_right[i]
        
        # 1. Custom Branching Ribbon
        ffb = slide.shapes.build_freeform(Inches(cx + 0.8), Inches(cy - (ribbon_thickness/2)))
        ffb.add_line_segments([
            (Inches(cx + 1.6), Inches(node_y - (ribbon_thickness/2))),
            (Inches(right_rect_x + 0.2), Inches(node_y - (ribbon_thickness/2))),
            (Inches(right_rect_x + 0.2), Inches(node_y + (ribbon_thickness/2))),
            (Inches(cx + 1.6), Inches(node_y + (ribbon_thickness/2))),
            (Inches(cx + 0.8), Inches(cy + (ribbon_thickness/2))),
            (Inches(cx + 0.8), Inches(cy - (ribbon_thickness/2)))
        ])
        ribbon = ffb.convert_to_shape()
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = color
        ribbon.line.fill.background()
        
        # 2. Rounded Rectangle Endpoint
        rect_y = node_y - (rect_h / 2)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(right_rect_x), Inches(rect_y), Inches(rect_w), Inches(rect_h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        
        # 3. Icon Placeholder
        icon_x = right_rect_x + rect_w - icon_size - 0.15
        icon_y = node_y - (icon_size / 2)
        icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(icon_x), Inches(icon_y), Inches(icon_size), Inches(icon_size))
        icon.fill.solid()
        icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
        icon.line.fill.background()
        
        # 4. Text Block
        add_text_block(right_rect_x + rect_w + 0.2, rect_y, 2.4, 1.0, i + 6, color, PP_ALIGN.LEFT)

    # --- Draw Central Hub ---
    # The hub is drawn last so it sits on top of all the radiating ribbons, hiding their origin points
    hub_radius = 1.0
    hub_shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        Inches(cx - hub_radius), Inches(cy - hub_radius), 
        Inches(hub_radius * 2), Inches(hub_radius * 2)
    )
    hub_shape.fill.solid()
    hub_shape.fill.fore_color.rgb = center_color
    
    # White border around hub for contrast
    hub_shape.line.color.rgb = RGBColor(255, 255, 255)
    hub_shape.line.width = Pt(4)
    
    # Center Hub Text
    tf_hub = hub_shape.text_frame
    tf_hub.word_wrap = True
    p_hub = tf_hub.paragraphs[0]
    p_hub.alignment = PP_ALIGN.CENTER
    
    run_hub = p_hub.add_run()
    parts = title_text.split()
    if len(parts) >= 2:
        run_hub.text = parts[0] + "\n" + " ".join(parts[1:])
    else:
        run_hub.text = title_text
        
    run_hub.font.bold = True
    run_hub.font.size = Pt(16)
    run_hub.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
