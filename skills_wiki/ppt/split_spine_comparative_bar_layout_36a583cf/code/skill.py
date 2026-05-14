def create_slide(
    output_pptx_path: str,
    title_text: str = "Comparison Slide Title / Pros Cons",
    body_text: str = "",
    bg_palette: str = "technology",
    accent_color: tuple = (40, 40, 40),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Split-Spine Comparative Bar' layout.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from lxml import etree

    # Utility to inject shadows for depth
    def add_shadow(shape):
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad="150000", dist="30000", dir="5400000", algn="ctr")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="20000")

    def clear_margins(tf):
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(245, 245, 248)
    bg.line.fill.background()

    # Title & Headers
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.2), Inches(11.33), Inches(0.8))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.font.size = Pt(32)
    title_p.font.color.rgb = RGBColor(*accent_color)
    title_p.alignment = PP_ALIGN.CENTER

    headers = [
        {"text": "OPTION 1", "x": 1.5, "color": (218, 62, 82)},
        {"text": "OPTION 2", "x": 8.3, "color": (46, 204, 113)}
    ]
    for h in headers:
        hbox = slide.shapes.add_textbox(Inches(h["x"]), Inches(1.0), Inches(3.5), Inches(0.4))
        hp = hbox.text_frame.paragraphs[0]
        hp.text = h["text"]
        hp.font.size = Pt(16)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(*h["color"])
        hp.alignment = PP_ALIGN.CENTER
        
        sbox = slide.shapes.add_textbox(Inches(h["x"] - 0.5), Inches(1.4), Inches(4.5), Inches(0.5))
        sbox.text_frame.word_wrap = True
        sp = sbox.text_frame.paragraphs[0]
        sp.text = f"Describe {h['text'].lower()} here in 2-3 lines so everybody gets more idea about it."
        sp.font.size = Pt(11)
        sp.font.color.rgb = RGBColor(100, 100, 100)
        sp.alignment = PP_ALIGN.CENTER

    # Data for the 4 rows
    data_rows = [
        {"val1": 40, "val2": 30, "color": (218, 62, 82), "icon": "$"},
        {"val1": 20, "val2": 50, "color": (46, 204, 113), "icon": "★"},
        {"val1": 40, "val2": 20, "color": (52, 152, 219), "icon": "♥"},
        {"val1": 50, "val2": 60, "color": (155, 89, 182), "icon": "✔"},
    ]

    # === Layer 2: Colored Bars & Text ===
    for i, row in enumerate(data_rows):
        bar_y = 2.5 + i * 1.1
        color = RGBColor(*row["color"])
        
        # --- LEFT BAR ---
        left_width = 4.2 + (row["val1"] / 100.0) * 1.0
        left_x = 5.8 - left_width # Anchored under the left spine
        
        l_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_x), Inches(bar_y), Inches(left_width), Inches(0.8))
        l_bar.fill.solid()
        l_bar.fill.fore_color.rgb = color
        l_bar.line.fill.solid()
        l_bar.line.color.rgb = color
        l_bar.adjustments[0] = 0.5 # Capsule shape

        # Left Percentage
        lp_box = slide.shapes.add_textbox(Inches(left_x + 0.15), Inches(bar_y + 0.1), Inches(0.8), Inches(0.6))
        lp_tf = lp_box.text_frame
        clear_margins(lp_tf)
        lp_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lp_p = lp_tf.paragraphs[0]
        lp_p.text = f"{row['val1']}%"
        lp_p.alignment = PP_ALIGN.CENTER
        lp_p.font.size = Pt(22)
        lp_p.font.bold = True
        lp_p.font.color.rgb = RGBColor(255, 255, 255)

        # Left Divider Line
        l_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(left_x + 1.05), Inches(bar_y + 0.15), Inches(left_x + 1.05), Inches(bar_y + 0.65))
        l_line.line.color.rgb = RGBColor(255, 255, 255)
        l_line.line.width = Pt(1.5)

        # Left Description (Dynamically sized to stay flush with central column)
        desc_l_x = left_x + 1.15
        desc_l_w = 5.0 - desc_l_x 
        ld_box = slide.shapes.add_textbox(Inches(desc_l_x), Inches(bar_y + 0.1), Inches(desc_l_w), Inches(0.6))
        ld_tf = ld_box.text_frame
        clear_margins(ld_tf)
        ld_tf.word_wrap = True
        ld_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ld_p = ld_tf.paragraphs[0]
        ld_p.text = "The reason to achieve this target are mentioned here in 2-3 lines at least."
        ld_p.alignment = PP_ALIGN.RIGHT
        ld_p.font.size = Pt(10)
        ld_p.font.color.rgb = RGBColor(255, 255, 255)

        # --- RIGHT BAR ---
        right_width = 4.2 + (row["val2"] / 100.0) * 1.0
        right_x = 7.5 # Anchored under the right spine
        
        r_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(right_x), Inches(bar_y), Inches(right_width), Inches(0.8))
        r_bar.fill.solid()
        r_bar.fill.fore_color.rgb = color
        r_bar.line.fill.solid()
        r_bar.line.color.rgb = color
        r_bar.adjustments[0] = 0.5
        
        bar_right_edge = right_x + right_width

        # Right Description
        desc_r_x = 8.3 
        desc_r_w = (bar_right_edge - 1.15) - desc_r_x
        rd_box = slide.shapes.add_textbox(Inches(desc_r_x), Inches(bar_y + 0.1), Inches(desc_r_w), Inches(0.6))
        rd_tf = rd_box.text_frame
        clear_margins(rd_tf)
        rd_tf.word_wrap = True
        rd_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rd_p = rd_tf.paragraphs[0]
        rd_p.text = "The reason to achieve this target are mentioned here in 2-3 lines at least."
        rd_p.alignment = PP_ALIGN.LEFT
        rd_p.font.size = Pt(10)
        rd_p.font.color.rgb = RGBColor(255, 255, 255)

        # Right Divider Line
        r_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bar_right_edge - 1.05), Inches(bar_y + 0.15), Inches(bar_right_edge - 1.05), Inches(bar_y + 0.65))
        r_line.line.color.rgb = RGBColor(255, 255, 255)
        r_line.line.width = Pt(1.5)

        # Right Percentage
        rp_box = slide.shapes.add_textbox(Inches(bar_right_edge - 0.95), Inches(bar_y + 0.1), Inches(0.8), Inches(0.6))
        rp_tf = rp_box.text_frame
        clear_margins(rp_tf)
        rp_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rp_p = rp_tf.paragraphs[0]
        rp_p.text = f"{row['val2']}%"
        rp_p.alignment = PP_ALIGN.CENTER
        rp_p.font.size = Pt(22)
        rp_p.font.bold = True
        rp_p.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 3: Central White Spine Columns (Masks the inner bar edges) ===
    left_col = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(2.1), Inches(1.2), Inches(4.5))
    left_col.fill.solid()
    left_col.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left_col.line.fill.background()
    left_col.adjustments[0] = 0.5
    add_shadow(left_col)

    right_col = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.1), Inches(1.2), Inches(4.5))
    right_col.fill.solid()
    right_col.fill.fore_color.rgb = RGBColor(255, 255, 255)
    right_col.line.fill.background()
    right_col.adjustments[0] = 0.5
    add_shadow(right_col)

    # === Layer 4: Central Icons ===
    for i, row in enumerate(data_rows):
        bar_y = 2.5 + i * 1.1
        cx = 6.65
        cy = bar_y + 0.4
        
        bg_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.25), Inches(cy - 0.25), Inches(0.5), Inches(0.5))
        bg_circle.fill.solid()
        bg_circle.fill.fore_color.rgb = RGBColor(40, 40, 40)
        bg_circle.line.fill.solid()
        bg_circle.line.color.rgb = RGBColor(40, 40, 40)
        
        icon_box = slide.shapes.add_textbox(Inches(cx - 0.25), Inches(cy - 0.25), Inches(0.5), Inches(0.5))
        icon_tf = icon_box.text_frame
        clear_margins(icon_tf)
        icon_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        icon_p = icon_tf.paragraphs[0]
        icon_p.text = row["icon"]
        icon_p.alignment = PP_ALIGN.CENTER
        icon_p.font.size = Pt(18)
        icon_p.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
