def create_slide(
    output_pptx_path: str,
    title_text: str = "Mobile Experience",
    body_text: str = "Seamlessly integrated designs that look perfect on any device.",
    bg_theme: str = "minimalist desk",
    content_theme: str = "app interface, ui",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide featuring a programmatically generated smartphone mockup
    composited over an aesthetic background, reproducing a premium product showcase.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFilter, ImageOps

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Constants for dimensions (Slide is 1280x720 at 96dpi)
    SLIDE_W, SLIDE_H = 1280, 720
    PHONE_W, PHONE_H = 340, 680
    BEZEL = 16
    CORNER_RADIUS_PHONE = 40
    CORNER_RADIUS_SCREEN = 30

    # --- Helper: Image Downloader with Fallback ---
    def get_image(keyword, width, height, fallback_color1, fallback_color2):
        try:
            # Using a reliable placeholder service with keywords
            url = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(keyword)}?width={width}&height={height}&nologo=true"
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as response:
                img = Image.open(BytesIO(response.read())).convert("RGBA")
                return img.resize((width, height), Image.Resampling.LANCZOS)
        except Exception as e:
            print(f"Image download failed, using gradient fallback. Error: {e}")
            # Fallback: Create a gradient image
            base = Image.new('RGBA', (width, height), fallback_color1)
            top = Image.new('RGBA', (width, height), fallback_color2)
            mask = Image.new('L', (width, height))
            mask_data = []
            for y in range(height):
                mask_data.extend([int(255 * (y / height))] * width)
            mask.putdata(mask_data)
            return Image.composite(base, top, mask)

    # --- Step 1: Prepare Background ---
    bg_img = get_image(bg_theme, SLIDE_W, SLIDE_H, (240, 245, 250, 255), (200, 210, 220, 255))
    
    # Optional: Lighten background slightly to ensure text legibility
    overlay = Image.new('RGBA', (SLIDE_W, SLIDE_H), (255, 255, 255, 100))
    bg_img = Image.alpha_composite(bg_img, overlay)

    # --- Step 2: Prepare Phone Mockup ---
    # 2a. Download content image for the screen
    screen_w = PHONE_W - (BEZEL * 2)
    screen_h = PHONE_H - (BEZEL * 2)
    screen_img = get_image(content_theme, screen_w, screen_h, (0, 191, 255, 255), (13, 17, 28, 255))
    
    # 2b. Create screen mask for rounded corners
    screen_mask = Image.new("L", (screen_w, screen_h), 0)
    draw_screen_mask = ImageDraw.Draw(screen_mask)
    draw_screen_mask.rounded_rectangle((0, 0, screen_w, screen_h), radius=CORNER_RADIUS_SCREEN, fill=255)
    
    # Crop content image to rounded corners
    rounded_screen = Image.new("RGBA", (screen_w, screen_h), (0,0,0,0))
    rounded_screen.paste(screen_img, (0,0), mask=screen_mask)

    # 2c. Build Phone Body
    phone = Image.new("RGBA", (PHONE_W, PHONE_H), (0,0,0,0))
    draw_phone = ImageDraw.Draw(phone)
    # Bezel (Dark Gray)
    draw_phone.rounded_rectangle((0, 0, PHONE_W, PHONE_H), radius=CORNER_RADIUS_PHONE, fill=(30, 30, 32, 255))
    # Inner border line (slight highlight)
    draw_phone.rounded_rectangle((2, 2, PHONE_W-2, PHONE_H-2), radius=CORNER_RADIUS_PHONE-2, outline=(80, 80, 85, 255), width=2)
    
    # Paste screen onto phone
    phone.paste(rounded_screen, (BEZEL, BEZEL), mask=rounded_screen)
    
    # Add a top notch/camera (classic modern phone look)
    notch_w, notch_h = 100, 25
    notch_x = (PHONE_W - notch_w) // 2
    draw_phone.rounded_rectangle((notch_x, BEZEL-5, notch_x+notch_w, BEZEL+notch_h), radius=10, fill=(30, 30, 32, 255))

    # --- Step 3: Create Shadow and Composite ---
    # Shadow canvas needs to be larger than the phone
    padding = 150
    shadow = Image.new("RGBA", (PHONE_W + padding*2, PHONE_H + padding*2), (0,0,0,0))
    draw_shadow = ImageDraw.Draw(shadow)
    # Draw dark shape for shadow, slightly offset downwards
    shadow_offset_y = 20
    draw_shadow.rounded_rectangle(
        (padding, padding + shadow_offset_y, padding + PHONE_W, padding + PHONE_H + shadow_offset_y),
        radius=CORNER_RADIUS_PHONE, 
        fill=(0, 0, 0, 140)
    )
    # Blur the shadow
    shadow = shadow.filter(ImageFilter.GaussianBlur(30))
    
    # Paste phone onto shadow
    shadow.paste(phone, (padding, padding), mask=phone)

    # Paste the whole mockup (shadow + phone) onto the right side of the background
    mockup_x = SLIDE_W - PHONE_W - padding - 80 # 80px margin from right
    mockup_y = (SLIDE_H - PHONE_H) // 2 - padding
    bg_img.paste(shadow, (mockup_x, mockup_y), mask=shadow)

    # Save final composite to temp file
    temp_bg_path = "temp_mockup_bg.png"
    bg_img.save(temp_bg_path)

    # --- Step 4: Add to PPTX ---
    # Set the composited image as the background of the slide
    slide.shapes.add_picture(temp_bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Add Typography on the left side
    text_margin_left = Inches(1.0)
    text_width = Inches(5.5)

    # Title Box
    title_box = slide.shapes.add_textbox(text_margin_left, Inches(2.8), text_width, Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.name = "Arial" # Standard clean font
    p.font.color.rgb = RGBColor(20, 20, 20)

    # Body Text Box
    body_box = slide.shapes.add_textbox(text_margin_left, Inches(4.0), text_width, Inches(1.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(20)
    p_body.font.name = "Arial"
    p_body.font.color.rgb = RGBColor(80, 80, 80)

    # Save and clean up
    prs.save(output_pptx_path)
    if os.path.exists(temp_bg_path):
        os.remove(temp_bg_path)

    return output_pptx_path
