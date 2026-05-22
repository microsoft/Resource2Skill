def create_slide(
    output_pptx_path: str,
    title_text: str = "Five Design Principles in Practice",
    subtitle_text: str = "Applying Hierarchy, Space, Alignment, Proximity, and Balance.",
    accent_color: tuple = (0, 160, 176),  # Vibrant Teal
    **kwargs,
) -> str:
    """
    Creates a PPTX file demonstrating a professional consulting-style grid layout.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # --- Color Palette ---
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_TITLE = RGBColor(20, 30, 50)      # Dark Navy
    COLOR_SUBTITLE = RGBColor(120, 130, 140) # Medium Gray
    COLOR_BODY = RGBColor(90, 100, 110)      # Dark Gray
    COLOR_ACCENT = RGBColor(*accent_color)
    COLOR_LINE = RGBColor(220, 225, 230)     # Light Gray

    # --- 1. Space & Alignment: Define the Master Grid ---
    MARGIN = Inches(0.8)
    SLIDE_W = prs.slide_width
    SLIDE_H = prs.slide_height
    CONTENT_W = SLIDE_W - (2 * MARGIN)

    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

    # --- 2. Hierarchy: Header Section ---
    # Main Title
    title_box = slide.shapes.add_textbox(MARGIN, MARGIN, CONTENT_W, Inches(0.8))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    p = title_tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial' # Standard sans-serif
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE

    # Subtitle
    sub_box = slide.shapes.add_textbox(MARGIN, MARGIN + Inches(0.8), CONTENT_W, Inches(0.5))
    sub_tf = sub_box.text_frame
    sub_tf.word_wrap = True
    p2 = sub_tf.paragraphs[0]
    p2.text = subtitle_text
    p2.font.name = 'Arial'
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_SUBTITLE

    # Divider Line (Separates header from content)
    line_y = MARGIN + Inches(1.5)
    divider = slide.shapes.add_connector(
        MSO_SHAPE.LINE_CALLOUT_1, 
        MARGIN, line_y, SLIDE_W - MARGIN, line_y
    )
    divider.line.color.rgb = COLOR_LINE
    divider.line.width = Pt(1.5)

    # --- 3. Proximity, Alignment & Balance: 3-Column Content ---
    NUM_COLS = 3
    GAP = Inches(0.6)
    COL_W = (CONTENT_W - (GAP * (NUM_COLS - 1))) / NUM_COLS
    START_Y = line_y + Inches(0.6)

    col_data = [
        {"icon": "1", "title": "Establish Hierarchy", "body": "Guide the reader's eye using scale, weight, and color. Ensure the most critical takeaway is immediately obvious upon first glance."},
        {"icon": "2", "title": "Respect Whitespace", "body": "Use generous margins and padding. Do not fill every empty pixel. Space creates elegance and reduces cognitive overload."},
        {"icon": "3", "title": "Group & Align", "body": "Place related elements in close proximity. Align everything to a strict invisible grid to establish order and professional trust."}
    ]

    for i, data in enumerate(col_data):
        col_x = MARGIN + (i * (COL_W + GAP))
        
        # Proximity: Icon/Accent Graphic
        icon_size = Inches(0.6)
        icon = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            col_x, START_Y, icon_size, icon_size
        )
        icon.fill.solid()
        icon.fill.fore_color.rgb = COLOR_ACCENT
        icon.line.fill.background() # No border
        
        # Number inside icon
        icon_tf = icon.text_frame
        icon_p = icon_tf.paragraphs[0]
        icon_p.text = data["icon"]
        icon_p.alignment = PP_ALIGN.CENTER
        icon_p.font.bold = True
        icon_p.font.size = Pt(18)
        icon_p.font.color.rgb = COLOR_BG

        # Proximity: Column Title (Close to icon)
        col_title_box = slide.shapes.add_textbox(
            col_x, START_Y + icon_size + Inches(0.15), COL_W, Inches(0.5)
        )
        ct_tf = col_title_box.text_frame
        ct_tf.word_wrap = True
        ct_p = ct_tf.paragraphs[0]
        ct_p.text = data["title"]
        ct_p.font.name = 'Arial'
        ct_p.font.size = Pt(20)
        ct_p.font.bold = True
        ct_p.font.color.rgb = COLOR_TITLE

        # Proximity: Column Body (Close to title)
        col_body_box = slide.shapes.add_textbox(
            col_x, START_Y + icon_size + Inches(0.7), COL_W, Inches(2.0)
        )
        cb_tf = col_body_box.text_frame
        cb_tf.word_wrap = True
        cb_p = cb_tf.paragraphs[0]
        cb_p.text = data["body"]
        cb_p.font.name = 'Arial'
        cb_p.font.size = Pt(14)
        cb_p.font.color.rgb = COLOR_BODY
        # Optional: increase line spacing slightly for readability
        cb_p.line_spacing = 1.2

    # --- 4. Balance: Asymmetrical Footer / Takeaway ---
    # To balance the heavy top-left title, we add a subtle callout box at the bottom right.
    footer_w = CONTENT_W * 0.6
    footer_h = Inches(0.8)
    footer_x = SLIDE_W - MARGIN - footer_w
    footer_y = SLIDE_H - MARGIN - footer_h

    footer_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, footer_x, footer_y, footer_w, footer_h
    )
    footer_box.fill.solid()
    # Light tint of accent color for background
    footer_box.fill.fore_color.rgb = RGBColor(
        int(COLOR_ACCENT[0] + (255 - COLOR_ACCENT[0]) * 0.9),
        int(COLOR_ACCENT[1] + (255 - COLOR_ACCENT[1]) * 0.9),
        int(COLOR_ACCENT[2] + (255 - COLOR_ACCENT[2]) * 0.9)
    )
    footer_box.line.fill.background()

    footer_tf = footer_box.text_frame
    footer_tf.word_wrap = True
    f_p = footer_tf.paragraphs[0]
    f_p.text = "Key Takeaway: Design is not just about making things 'pretty'; it's about structuring communication to minimize friction for the audience."
    f_p.font.name = 'Arial'
    f_p.font.size = Pt(12)
    f_p.font.bold = True
    f_p.font.color.rgb = COLOR_TITLE
    # Add some internal padding
    footer_tf.margin_left = Inches(0.2)
    footer_tf.margin_right = Inches(0.2)
    footer_tf.margin_top = Inches(0.15)

    prs.save(output_pptx_path)
    return output_pptx_path
