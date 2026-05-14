def create_slide(
    output_pptx_path: str,
    title_text: str = "How was your\nonboarding?",
    bg_color_left: tuple = (255, 0, 102),     # Hot Pink
    bg_color_right: tuple = (65, 30, 225),    # Electric Purple
    card_color_front: tuple = (255, 240, 0),  # Bright Yellow
    card_color_back: tuple = (255, 255, 255), # White
    text_color_top: tuple = (255, 255, 255),  # White
    text_color_shadow: tuple = (220, 20, 110),# Deep Magenta
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Pop-Art Angled Title Card visual effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Split Background ===
    
    # Base background (Right side)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*bg_color_right)

    # Diagonal Polygon (Left side)
    ff_builder = slide.shapes.build_freeform()
    # Draw a polygon covering the left-to-bottom area
    ff_builder.add_line_segments([
        (0, 0),
        (Inches(9), 0),
        (Inches(3), Inches(7.5)),
        (0, Inches(7.5))
    ], close=True)
    
    bg_left = ff_builder.convert_to_shape()
    bg_left.fill.solid()
    bg_left.fill.fore_color.rgb = RGBColor(*bg_color_left)
    # Safely remove border by matching fill color
    bg_left.line.fill.solid()
    bg_left.line.fill.fore_color.rgb = RGBColor(*bg_color_left)

    # === Layer 2: Tilted Card Stack ===
    
    card_w = Inches(9.5)
    card_h = Inches(5.5)
    cx, cy = Inches(13.333) / 2, Inches(7.5) / 2
    left, top = cx - card_w / 2, cy - card_h / 2
    rotation_angle = -4

    # Back Card (White offset)
    card_back = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, card_w, card_h)
    card_back.fill.solid()
    card_back.fill.fore_color.rgb = RGBColor(*card_color_back)
    card_back.line.fill.solid()
    card_back.line.fill.fore_color.rgb = RGBColor(*card_color_back)
    card_back.rotation = rotation_angle

    # Front Card (Yellow main)
    offset_front = Inches(0.25)
    card_front = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left - offset_front, top - offset_front, card_w, card_h)
    card_front.fill.solid()
    card_front.fill.fore_color.rgb = RGBColor(*card_color_front)
    card_front.line.fill.solid()
    card_front.line.fill.fore_color.rgb = RGBColor(*card_color_front)
    card_front.rotation = rotation_angle

    # === Layer 3: Solid 3D Extruded Typography ===
    
    text_x = left - offset_front + Inches(0.8)
    text_y = top - offset_front + Inches(0.8)
    text_w = card_w - Inches(1.6)
    text_h = card_h - Inches(1.6)

    # To create a solid, hard-edged 3D shadow, we stack multiple text boxes
    extrusion_depth = 8
    step_size = Inches(0.015)

    # Shadow Layers (Deep Magenta)
    for i in range(extrusion_depth, 0, -1):
        tx_box = slide.shapes.add_textbox(text_x + (i * step_size), text_y + (i * step_size), text_w, text_h)
        tx_box.rotation = rotation_angle
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.alignment = PP_ALIGN.LEFT
        p.font.name = "Arial Black"
        p.font.size = Pt(72)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*text_color_shadow)

    # Top Text Layer (White)
    tx_top = slide.shapes.add_textbox(text_x, text_y, text_w, text_h)
    tx_top.rotation = rotation_angle
    tf_top = tx_top.text_frame
    tf_top.word_wrap = True
    p_top = tf_top.paragraphs[0]
    p_top.text = title_text
    p_top.alignment = PP_ALIGN.LEFT
    p_top.font.name = "Arial Black"
    p_top.font.size = Pt(72)
    p_top.font.bold = True
    p_top.font.color.rgb = RGBColor(*text_color_top)

    prs.save(output_pptx_path)
    return output_pptx_path
