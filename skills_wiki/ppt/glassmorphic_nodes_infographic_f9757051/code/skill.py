def create_slide(
    output_pptx_path: str,
    title_text: str = "CORE ARCHITECTURE",
    body_text: str = "Four fundamental pillars driving the\nnext generation platform ecosystem.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Glassmorphic Nodes Infographic" visual effect.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from PIL import Image, ImageDraw, ImageFilter

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg_color = (15, 23, 42) # Slate 900
    
    # ==========================================
    # Helper: Create soft glowing gradient circle
    # ==========================================
    def create_gradient_circle(radius, color1, color2):
        size = int(radius * 2)
        gradient = Image.new('RGBA', (1, size))
        for y in range(size):
            factor = y / (size - 1) if size > 1 else 0
            r = int(color1[0] * (1 - factor) + color2[0] * factor)
            g = int(color1[1] * (1 - factor) + color2[1] * factor)
            b = int(color1[2] * (1 - factor) + color2[2] * factor)
            gradient.putpixel((0, y), (r, g, b, 255))
        gradient = gradient.resize((size, size))
        
        mask = Image.new('L', (size, size), 0)
        # 10px inset to allow blur without hard clipping
        ImageDraw.Draw(mask).ellipse((10, 10, size-10, size-10), fill=255) 
        gradient.putalpha(mask)
        return gradient.filter(ImageFilter.GaussianBlur(8))

    # ==========================================
    # Layer 1: PIL Masked Background Image
    # ==========================================
    img = Image.new('RGBA', (1920, 1080), bg_color)
    radius = 220
    
    # 4 node colors
    colors = [
        {"color1": (56, 189, 248), "color2": (2, 132, 199),   "rgb": RGBColor(56, 189, 248)}, # Cyan
        {"color1": (167, 139, 250), "color2": (124, 58, 237), "rgb": RGBColor(167, 139, 250)}, # Purple
        {"color1": (250, 204, 21), "color2": (234, 88, 12),   "rgb": RGBColor(250, 204, 21)},  # Yellow
        {"color1": (74, 222, 128), "color2": (22, 163, 74),   "rgb": RGBColor(74, 222, 128)}   # Green
    ]
    
    circles = [create_gradient_circle(radius, c["color1"], c["color2"]) for c in colors]
    
    # Center rounded rectangle dimensions
    rect_w, rect_h = 700, 700
    center_x, center_y = 960, 540
    rect_left = center_x - rect_w // 2
    rect_top = center_y - rect_h // 2
    rect_right = center_x + rect_w // 2
    rect_bottom = center_y + rect_h // 2
    
    centers = [
        (rect_left, rect_top),        # Top Left
        (rect_right, rect_top),       # Top Right
        (rect_left, rect_bottom),     # Bottom Left
        (rect_right, rect_bottom)     # Bottom Right
    ]
    
    # Paste circles onto background
    for circle, (cx, cy) in zip(circles, centers):
        img.paste(circle, (int(cx - radius), int(cy - radius)), circle)
        
    # ** THE ILLUSION **: Draw a shape filled with bg_color to mask the inner parts of the circles
    mask_draw = ImageDraw.Draw(img)
    mask_draw.rounded_rectangle([rect_left, rect_top, rect_right, rect_bottom], radius=150, fill=bg_color)
    
    bg_path = "temp_glass_bg.png"
    img.save(bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    
    # ==========================================
    # Layer 2: True Glassmorphism XML Injection
    # ==========================================
    glass_left, glass_top = rect_left / 144, rect_top / 144
    glass_width, glass_height = rect_w / 144, rect_h / 144
    
    # Place perfectly over the masked area
    glass = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(glass_left), Inches(glass_top),
        Inches(glass_width), Inches(glass_height)
    )
    glass.adjustments[0] = 150 / min(rect_w, rect_h) # Match PIL corner radius
    glass.line.color.rgb = RGBColor(255, 255, 255)
    glass.line.width = Pt(1.5)
    
    # Inject OpenXML for glass (gradients, alphas, shadows)
    spPr = glass.element.spPr
    for tag in ['{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill',
                '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst']:
        elem = spPr.find(tag)
        if elem is not None:
            spPr.remove(elem)
            
    gradFill = OxmlElement('a:gradFill')
    gsLst = OxmlElement('a:gsLst')
    
    # Glass Stop 1
    gs1 = OxmlElement('a:gs')
    gs1.set('pos', '0')
    srgbClr1 = OxmlElement('a:srgbClr')
    srgbClr1.set('val', 'FFFFFF')
    alpha1 = OxmlElement('a:alpha')
    alpha1.set('val', '20000') # 20% Alpha
    srgbClr1.append(alpha1)
    gs1.append(srgbClr1)
    
    # Glass Stop 2
    gs2 = OxmlElement('a:gs')
    gs2.set('pos', '100000')
    srgbClr2 = OxmlElement('a:srgbClr')
    srgbClr2.set('val', 'FFFFFF')
    alpha2 = OxmlElement('a:alpha')
    alpha2.set('val', '2000') # 2% Alpha
    srgbClr2.append(alpha2)
    gs2.append(srgbClr2)
    
    gsLst.append(gs1)
    gsLst.append(gs2)
    lin = OxmlElement('a:lin')
    lin.set('ang', '5400000') # 90 degrees
    lin.set('scaled', '1')
    gradFill.append(gsLst)
    gradFill.append(lin)
    
    # Soft Shadow
    effectLst = OxmlElement('a:effectLst')
    outerShdw = OxmlElement('a:outerShdw')
    outerShdw.set('blurRad', '200000')
    outerShdw.set('dist', '40000')
    outerShdw.set('dir', '5400000')
    shdwClr = OxmlElement('a:srgbClr')
    shdwClr.set('val', '000000')
    shdwAlpha = OxmlElement('a:alpha')
    shdwAlpha.set('val', '40000') # 40% shadow opacity
    shdwClr.append(shdwAlpha)
    outerShdw.append(shdwClr)
    effectLst.append(outerShdw)

    # Apply fills and make line translucent
    ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
    if ln is not None:
        ln.addprevious(gradFill)
        ln.addnext(effectLst)
        solidFill = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        if solidFill is not None:
            for c in solidFill.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                alpha = OxmlElement('a:alpha')
                alpha.set('val', '50000') # 50% opacity border
                c.append(alpha)
    else:
        spPr.append(gradFill)
        spPr.append(effectLst)

    # ==========================================
    # Layer 3: Text & Layout
    # ==========================================
    # Main Title (Center)
    txBox = slide.shapes.add_textbox(Inches(glass_left + 0.2), Inches(glass_top + 1.8), Inches(glass_width - 0.4), Inches(1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
    
    # Outer Nodes Layout
    node_data = [
        {"title": "01  SECURITY", "align": PP_ALIGN.RIGHT, "color": colors[0]["rgb"],
         "tx_x": 0.5, "tx_y": glass_top - 0.1, "line_sx": 3.6, "line_ex": glass_left - 0.1, "line_y": glass_top},
         
        {"title": "INTEGRATION  02", "align": PP_ALIGN.LEFT, "color": colors[1]["rgb"],
         "tx_x": glass_left + glass_width + 0.8, "tx_y": glass_top - 0.1, "line_sx": glass_left + glass_width + 0.1, "line_ex": glass_left + glass_width + 0.8, "line_y": glass_top},
         
        {"title": "03  ANALYTICS", "align": PP_ALIGN.RIGHT, "color": colors[2]["rgb"],
         "tx_x": 0.5, "tx_y": glass_top + glass_height - 0.4, "line_sx": 3.6, "line_ex": glass_left - 0.1, "line_y": glass_top + glass_height},
         
        {"title": "PERFORMANCE  04", "align": PP_ALIGN.LEFT, "color": colors[3]["rgb"],
         "tx_x": glass_left + glass_width + 0.8, "tx_y": glass_top + glass_height - 0.4, "line_sx": glass_left + glass_width + 0.1, "line_ex": glass_left + glass_width + 0.8, "line_y": glass_top + glass_height}
    ]
    
    for nd in node_data:
        tb = slide.shapes.add_textbox(Inches(nd["tx_x"]), Inches(nd["tx_y"]), Inches(3.0), Inches(1))
        t_frame = tb.text_frame
        
        p = t_frame.paragraphs[0]
        p.text = nd["title"]
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = nd["color"]
        p.alignment = nd["align"]
        
        p2 = t_frame.add_paragraph()
        p2.text = "Operational parameters and module logic details described here."
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(148, 163, 184)
        p2.alignment = nd["align"]
        
        # Connectors
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 
            Inches(nd["line_sx"]), Inches(nd["line_y"]), 
            Inches(nd["line_ex"]), Inches(nd["line_y"])
        )
        line.line.color.rgb = nd["color"]
        line.line.width = Pt(1.5)
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '40000') # 40%
        line.line.color._xClr.append(alpha)

    # Small decorative background dots
    dots = [
        (glass_left - 1.0, glass_top - 1.0, colors[0]["rgb"]),
        (glass_left + glass_width + 1.0, glass_top - 0.5, colors[1]["rgb"]),
        (glass_left - 0.5, glass_top + glass_height + 0.8, colors[2]["rgb"]),
        (glass_left + glass_width + 0.5, glass_top + glass_height + 1.0, colors[3]["rgb"])
    ]
    for dx, dy, color in dots:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(dx), Inches(dy), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
