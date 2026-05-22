def create_slide(
    output_pptx_path: str,
    title_text: str = "Production volume ahead of budget (tons)",
    data: list = None,
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Message-Driven Variance Infographic.
    Replaces a complex chart with a clean, shape-based visual that highlights variance.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Default data simulating the tutorial's transformation
    if data is None:
        data = [
            {"label": "Europe", "value": 330, "variance": 70, "bar_color": (112, 173, 71)},  # Highlighted positive
            {"label": "US", "value": 165, "variance": 15, "bar_color": (166, 166, 166)},      # Neutral
            {"label": "Asia", "value": 65, "variance": -55, "bar_color": (166, 166, 166)},    # Neutral
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Title ===
    tb_title = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p = tb_title.text_frame.add_paragraph()
    p.text = title_text
    p.font.size = Pt(36)
    p.font.name = "Calibri"
    p.font.color.rgb = RGBColor(64, 64, 64)

    # Subtle horizontal line under title to separate header from data
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.6), Inches(11.333), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(217, 217, 217)
    line.line.fill.background()

    # === Layer 2: Chart Area ===
    base_x = 3.0
    max_w = 7.0
    start_y = 2.5
    y_spacing = 1.2
    bar_height = 0.6

    max_val = max(d["value"] for d in data)
    scale = max_w / max_val if max_val > 0 else 1.0

    for i, row in enumerate(data):
        y = start_y + i * y_spacing

        # Y-Axis Label
        tb_label = slide.shapes.add_textbox(Inches(0.5), Inches(y + 0.05), Inches(2.3), Inches(0.5))
        p_label = tb_label.text_frame.paragraphs[0]
        p_label.text = row["label"]
        p_label.font.size = Pt(24)
        p_label.font.bold = True
        p_label.font.color.rgb = RGBColor(89, 89, 89)
        p_label.alignment = PP_ALIGN.RIGHT

        # Main Bar
        bar_w = row["value"] * scale
        bar_w = max(bar_w, 0.1)  # Ensure minimum width for visibility
        r, g, b = row["bar_color"]
        
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(base_x), Inches(y), Inches(bar_w), Inches(bar_height))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(r, g, b)
        bar.line.fill.background()

        # Actual Value Text (inside bar if wide enough, outside if too short)
        if bar_w > 1.2:
            val_x = base_x + bar_w - 1.0
            val_color = RGBColor(255, 255, 255)
            val_align = PP_ALIGN.RIGHT
        else:
            val_x = base_x + bar_w + 0.1
            val_color = RGBColor(89, 89, 89)
            val_align = PP_ALIGN.LEFT

        tb_val = slide.shapes.add_textbox(Inches(val_x), Inches(y + 0.05), Inches(0.9), Inches(0.5))
        p_val = tb_val.text_frame.paragraphs[0]
        p_val.text = str(row["value"])
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = val_color
        p_val.alignment = val_align

        # Variance Indicator Badge
        var_val = row["variance"]
        # Shift X further right if the actual value text was pushed outside the bar
        var_x = base_x + bar_w + (0.2 if bar_w > 1.2 else 1.2)
        
        var_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(var_x), Inches(y + 0.1), Inches(1.0), Inches(0.4))
        var_bg.fill.solid()
        var_bg.line.fill.background()
        
        # Apply semantic colors based on performance
        if var_val > 0:
            var_bg.fill.fore_color.rgb = RGBColor(226, 239, 218)  # Light Green
            txt_color = RGBColor(84, 130, 53)                     # Dark Green
            var_text = f"+{var_val}"
        else:
            var_bg.fill.fore_color.rgb = RGBColor(252, 228, 214)  # Light Red
            txt_color = RGBColor(192, 0, 0)                       # Dark Red
            var_text = str(var_val)

        p_var = var_bg.text_frame.paragraphs[0]
        p_var.text = var_text
        p_var.font.size = Pt(16)
        p_var.font.bold = True
        p_var.font.color.rgb = txt_color
        p_var.alignment = PP_ALIGN.CENTER

    # === Layer 3: Global Summary Footer ===
    global_y = start_y + len(data) * y_spacing + 0.2
    
    # Separation Line
    sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(global_y - 0.2), Inches(11.333), Inches(0.02))
    sep.fill.solid()
    sep.fill.fore_color.rgb = RGBColor(217, 217, 217)
    sep.line.fill.background()

    # Global Label
    tb_g_label = slide.shapes.add_textbox(Inches(0.5), Inches(global_y), Inches(2.3), Inches(0.5))
    p_g_label = tb_g_label.text_frame.paragraphs[0]
    p_g_label.text = "Global"
    p_g_label.font.size = Pt(24)
    p_g_label.font.bold = True
    p_g_label.font.color.rgb = RGBColor(89, 89, 89)
    p_g_label.alignment = PP_ALIGN.RIGHT

    # Global Actual Value
    total_val = sum(d["value"] for d in data)
    tb_g_val = slide.shapes.add_textbox(Inches(3.0), Inches(global_y), Inches(2.0), Inches(0.5))
    p_g_val = tb_g_val.text_frame.paragraphs[0]
    p_g_val.text = f"{total_val} t"
    p_g_val.font.size = Pt(24)
    p_g_val.font.bold = True
    p_g_val.font.color.rgb = RGBColor(64, 64, 64)

    # Global Total Variance
    total_var = sum(d["variance"] for d in data)
    var_color = RGBColor(84, 130, 53) if total_var > 0 else RGBColor(192, 0, 0)
    
    # Up/Down Arrow for global variance
    arrow_char = "↑" if total_var > 0 else "↓"
    
    tb_g_var = slide.shapes.add_textbox(Inches(4.5), Inches(global_y), Inches(2.0), Inches(0.5))
    p_g_var = tb_g_var.text_frame.paragraphs[0]
    p_g_var.text = f"{arrow_char} {abs(total_var)} t"
    p_g_var.font.size = Pt(24)
    p_g_var.font.bold = True
    p_g_var.font.color.rgb = var_color

    prs.save(output_pptx_path)
    return output_pptx_path
