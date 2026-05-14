import os
import urllib.request
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "INNOVATION\n& FUTURE",
    body_text: str = "Leveraging artificial intelligence to build modern, structured, and visually compelling presentations in seconds.",
    panel_color: tuple = (13, 17, 28),      # Dark Navy Blue
    text_color: tuple = (255, 255, 255),    # White text
    accent_color: tuple = (0, 191, 255),    # Vivid Cyan accent line
    image_keyword: str = "technology"       # Used for fetching dynamic image
) -> str:
    """
    Create a PPTX file reproducing the 'Modern High-Contrast Split-Panel' layout.
    """
    prs = Presentation()
    # Set to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ==========================================
    # Layer 1: Background Image (Right Side)
    # ==========================================
    img_path = "temp_bg_image.jpg"
    
    # Attempt to download a dynamic image mimicking Canva's photo library
    try:
        url = f"https://picsum.photos/seed/{image_keyword}/1200/800"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception as e:
        # Fallback: Use PIL to generate a placeholder image if network fails
        print(f"Image download failed, generating fallback. Error: {e}")
        img = Image.new('RGB', (1200, 800), color=(200, 200, 200))
        draw = ImageDraw.Draw(img)
        # Draw a simple pattern to make it look like a placeholder
        for i in range(0, 1200, 40):
            draw.line([(i, 0), (i, 800)], fill=(220, 220, 220), width=2)
        img.save(img_path)

    # Add image to slide (Positioned on the right, taking up ~65% of width)
    # Left = 4.5 inches, Width = 8.833 inches
    slide.shapes.add_picture(img_path, Inches(4.5), Inches(0), width=Inches(8.833), height=Inches(7.5))

    # ==========================================
    # Layer 2: Geometric Color Block Panel (Left Side)
    # ==========================================
    # Panel spans from x=0 to x=5.5 (Overlaps the image slightly to create depth)
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(0), Inches(5.5), Inches(7.5)
    )
    # Remove outline and set solid color
    panel.line.fill.background()
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(*panel_color)

    # ==========================================
    # Layer 3: Decorative Accent Element
    # ==========================================
    # Small line or box to anchor the design (Very common in modern templates)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.5), Inches(0.5), Inches(0.08)
    )
    accent.line.fill.background()
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(*accent_color)

    # ==========================================
    # Layer 4: Text Content
    # ==========================================
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.2), Inches(2.0))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.name = 'Arial' # Universally available clean sans-serif
    p_title.font.color.rgb = RGBColor(*text_color)
    p_title.alignment = PP_ALIGN.LEFT

    # Body Box
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(4.0), Inches(3.0))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(16)
    p_body.font.name = 'Arial'
    
    # Calculate a slightly dimmed text color for the body (simulate opacity/hierarchy)
    dim_color = tuple(int(c * 0.8) for c in text_color)
    p_body.font.color.rgb = RGBColor(*dim_color)
    p_body.alignment = PP_ALIGN.LEFT

    # Cleanup temporary image
    try:
        os.remove(img_path)
    except:
        pass

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution (uncomment to run locally)
# create_slide("canva_style_split_layout.pptx")
