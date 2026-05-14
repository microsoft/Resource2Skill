def create_slide(
    output_pptx_path: str,
    title_text: str = "Market Expansion",
    body_text: str = "Our new sector integration represents a 25% growth opportunity in the coming fiscal year. The highlighted segment demonstrates our core focus.",
    callout_text: str = "25%",
    bg_palette: str = "business,architecture", 
    accent_color: tuple = (0, 191, 255),  # Cyan
    **kwargs,
) -> str:
    """
    Creates a PPTX file featuring the "Exploded Geometric Reveal" design style.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Set background color
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(13, 17, 28)

    # 2. Image Processing via PIL (Fragmenting Shapes & Cropping)
    canvas_size = 1200
    
    # Try fetching a photo, fallback to gradient if offline
    try:
        url = f"https://source.unsplash.com/random/1200x1200/?{bg_palette}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGBA")
            base_img = base_img.resize((canvas_size, canvas_size), Image.LANCZOS)
    except Exception:
        # Fallback Gradient
        base_img = Image.new("RGBA", (canvas_size, canvas_size))
        draw_bg = ImageDraw.Draw(base_img)
        for i in range(canvas_size):
            r = int(20 + (50 - 20) * (i / canvas_size))
            g = int(40 + (100 - 40) * (i / canvas_size))
            b = int(80 + (180 - 80) * (i / canvas_size))
            draw_bg.line([(0, i), (canvas_size, i)], fill=(r, g, b, 255))

    # Create Mask for Main Shape (0 to 270 degrees -> leaves Top-Right empty)
    # PIL pieslice angles: 0 is 3 o'clock, going clockwise. 
    # Top-right quadrant is 270 to 360. 
    # Therefore, we want to keep 0 to 270.
    main_mask = Image.new("L", (canvas_size, canvas_size), 0)
    draw_mask = ImageDraw.Draw(main_mask)
    draw_mask.pieslice([0, 0, canvas_size, canvas_size], 0, 270, fill=255)
    
    base_img.putalpha(main_mask)
    main_img_path = "temp_main_fragment.png"
    base_img.save(main_img_path)

    # Create Exploded Accent Shape (Top-Right quadrant: 270 to 360 degrees)
    accent_img = Image.new("RGBA", (canvas_size, canvas_size), (0, 0, 0, 0))
    draw_accent = ImageDraw.Draw(accent_img)
    draw_accent.pieslice([0, 0, canvas_size, canvas_size], 270, 360, fill=(accent_color[0], accent_color[1], accent_color[2], 255))
    accent_img_path = "temp_accent_fragment.png"
    accent_img.save(accent_img_path)

    # 3. Insert Elements into PowerPoint
    base_size = Inches(5.5)
    base_x = Inches(1.5)
    base_y = Inches(1.0)
    
    # Explode Offset (Moving top-right)
    offset_x = Inches(0.3)
    offset_y = Inches(0.3)

    # Add Main Fragment
    slide.shapes.add_picture(main_img_path, base_x, base_y, base_size, base_size)

    # Add Accent Fragment
    pic_accent = slide.shapes.add_picture(accent_img_path, base_x + offset_x, base_y - offset_y, base_size, base_size)

    # Apply Shadow to Accent via lxml
    spPr = pic_accent.element.spPr
    effectLst = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    outerShdw = etree.SubElement(effectLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw",
                                 blurRad="350000", dist="150000", dir="2700000", algn="tl", rotWithShape="0")
    srgbClr = etree.SubElement(outerShdw, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", val="000000")
    etree.SubElement(srgbClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", val="60000")

    # 4. Text Layout
    # Callout Metric directly next to the exploded slice
    tx_callout = slide.shapes.add_textbox(base_x + base_size + offset_x - Inches(1.0), base_y, Inches(3), Inches(1.5))
    tf_callout = tx_callout.text_frame
    p_callout = tf_callout.paragraphs[0]
    p_callout.text = callout_text
    p_callout.font.size = Pt(72)
    p_callout.font.bold = True
    p_callout.font.color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])

    # Callout Subtext
    p_sub = tf_callout.add_paragraph()
    p_sub.text = "Key Metric Highlight"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(180, 190, 210)

    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(7.5), Inches(2.5), Inches(5), Inches(1))
    tf_title = tx_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Body Text
    tx_body = slide.shapes.add_textbox(Inches(7.5), Inches(3.5), Inches(4.5), Inches(3))
    tf_body = tx_body.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(200, 210, 225)
    p_body.line_spacing = 1.4

    # Cleanup temp files
    prs.save(output_pptx_path)
    if os.path.exists(main_img_path): os.remove(main_img_path)
    if os.path.exists(accent_img_path): os.remove(accent_img_path)

    return output_pptx_path
