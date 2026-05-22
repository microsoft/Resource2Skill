import random
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

def create_dataviz_color_strategy_slide(output_pptx_path: str) -> str:
    """
    Creates a PPTX slide demonstrating four key data visualization color strategies:
    Sequential, Divergent, Categorical, and Highlight.

    Args:
        output_pptx_path: The path to save the generated PPTX file.

    Returns:
        The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Helper function for linear interpolation of colors
    def interpolate_color(start_rgb, end_rgb, factor):
        return RGBColor(
            int(start_rgb[0] + (end_rgb[0] - start_rgb[0]) * factor),
            int(start_rgb[1] + (end_rgb[1] - start_rgb[1]) * factor),
            int(start_rgb[2] + (end_rgb[2] - start_rgb[2]) * factor)
        )

    # Helper function to draw a bar chart demo
    def draw_chart(shapes, left, top, width, height, colors, title):
        num_bars = len(colors)
        bar_width = width / num_bars
        bar_spacing = bar_width * 0.2
        drawable_bar_width = bar_width - bar_spacing

        for i, color in enumerate(colors):
            bar_left = left + (i * bar_width) + (bar_spacing / 2)
            bar_height = height * random.uniform(0.25, 1.0)
            bar_top = top + (height - bar_height)
            
            shape = shapes.add_shape(1, bar_left, bar_top, drawable_bar_width, bar_height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = color
            shape.line.fill.background()

        # Add title
        title_box = shapes.add_textbox(left, top + height + Inches(0.1), width, Inches(0.5))
        p = title_box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = 'Calibri'
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(89, 89, 89)
        p.alignment = PP_ALIGN.CENTER

    # --- Define Palettes ---
    # 1. Sequential Palette (Light to Dark Orange)
    seq_start = (254, 235, 218)
    seq_end = (230, 126, 34)
    sequential_colors = [interpolate_color(seq_start, seq_end, i / 6) for i in range(7)]

    # 2. Divergent Palette (Blue -> Neutral -> Orange)
    div_start = (36, 113, 163)
    div_mid = (240, 240, 240)
    div_end = (212, 85, 0)
    divergent_colors = [interpolate_color(div_start, div_mid, i / 5) for i in range(5)] + \
                       [interpolate_color(div_mid, div_end, i / 5) for i in range(1, 6)]

    # 3. Categorical Palette (5 distinct colors)
    categorical_colors = [
        RGBColor(41, 128, 185), RGBColor(39, 174, 96), RGBColor(241, 196, 15),
        RGBColor(230, 126, 34), RGBColor(142, 68, 173)
    ]

    # 4. Highlight Palette (Gray with Teal highlight)
    highlight_color = RGBColor(22, 160, 133)
    neutral_color = RGBColor(208, 211, 212)
    highlight_colors = [neutral_color] * 2 + [highlight_color] + [neutral_color] * 2

    # --- Draw the four charts on the slide ---
    chart_width = Inches(3.5)
    chart_height = Inches(3)
    total_width = chart_width * 4 + Inches(0.5) * 3
    start_left = (prs.slide_width - total_width) / 2
    y_pos = (prs.slide_height - chart_height) / 2 - Inches(0.25)

    draw_chart(slide.shapes, start_left, y_pos, chart_width, chart_height, sequential_colors, "SEQUENTIAL")
    draw_chart(slide.shapes, start_left + chart_width + Inches(0.5), y_pos, chart_width, chart_height, divergent_colors, "DIVERGENT")
    draw_chart(slide.shapes, start_left + (chart_width + Inches(0.5)) * 2, y_pos, chart_width, chart_height, categorical_colors, "CATEGORICAL")
    draw_chart(slide.shapes, start_left + (chart_width + Inches(0.5)) * 3, y_pos, chart_width, chart_height, highlight_colors, "HIGHLIGHT")

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_dataviz_color_strategy_slide("dataviz_color_strategies.pptx")
