import subprocess
import tempfile
import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

def create_slide(
    output_pptx_path: str,
    mermaid_definitions: dict = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file with slides containing diagrams rendered from Mermaid code.
    
    This function requires the Mermaid CLI (mmdc) to be installed and in the system's PATH.
    Install it via npm: `npm install -g @mermaid-js/mermaid-cli`

    Args:
        output_pptx_path (str): The path to save the generated .pptx file.
        mermaid_definitions (dict): A dictionary where keys are slide titles (str) 
                                    and values are Mermaid code blocks (str).

    Returns:
        str: The path to the saved PPTX file.
    """
    # --- Dependency Check ---
    if not shutil.which("mmdc"):
        raise EnvironmentError(
            "Mermaid CLI (mmdc) not found in your system's PATH. "
            "This skill requires it to render diagrams. \n"
            "Please install it globally by running: 'npm install -g @mermaid-js/mermaid-cli'"
        )

    # --- Default example if no definitions are provided ---
    if mermaid_definitions is None:
        mermaid_definitions = {
            "Example Flowchart": """
            flowchart TD
                A[Start] --> B{Decision};
                B -- Yes --> C[Process 1];
                B -- No --> D[Process 2];
                C --> E[End];
                D --> E[End];
            """,
            "Example Gantt Chart": """
            gantt
                title Project Development Timeline
                dateFormat YYYY-MM-DD
                section Planning Phase
                Market Research :done, 2024-09-01, 7d
                Initial Design :active, 2024-09-08, 5d
                section Development Phase
                Core Feature Dev : 2024-09-15, 15d
                Testing & QA : 2024-10-01, 10d
            """,
            "Example Pie Chart": """
            pie
                title Market Share
                "Apple" : 45
                "Samsung" : 25
                "Huawei" : 15
                "Other" : 15
            """
        }

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # --- Process Each Diagram ---
    for title, code in mermaid_definitions.items():
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Add a title to the slide
        title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), prs.slide_width - Inches(1), Inches(0.75))
        tf = title_shape.text_frame
        tf.text = title
        p = tf.paragraphs[0]
        p.font.size = Pt(32)
        p.font.bold = True

        # Use temporary files for processing
        with tempfile.NamedTemporaryFile(delete=False, mode='w', suffix='.mmd', encoding='utf-8') as mmd_file:
            mmd_file.write(code)
            mmd_file_path = mmd_file.name
        
        png_file_path = mmd_file_path.replace('.mmd', '.png')

        try:
            # Run the Mermaid CLI to convert the .mmd file to a .png image
            # The -b transparent flag makes the image background transparent for better integration
            subprocess.run(
                ["mmdc", "-i", mmd_file_path, "-o", png_file_path, "-b", "transparent"],
                check=True,
                capture_output=True, # Use capture_output to hide stdout/stderr unless there's an error
                text=True,
            )

            # Add the generated image to the slide, centered
            with Image.open(png_file_path) as img:
                img_width, img_height = img.size

            # Scale image to fit slide while maintaining aspect ratio
            max_width = prs.slide_width - Inches(1)
            max_height = prs.slide_height - Inches(1.5)
            
            ratio = min(max_width / img_width, max_height / img_height)
            
            display_width = int(img_width * ratio)
            display_height = int(img_height * ratio)

            left = int((prs.slide_width - display_width) / 2)
            top = Inches(1.0) # Place below the title

            slide.shapes.add_picture(png_file_path, left, top, width=display_width, height=display_height)

        except subprocess.CalledProcessError as e:
            error_message = f"Error rendering Mermaid diagram '{title}'.\nmmdc stderr: {e.stderr}"
            print(error_message)
            # Add an error message to the slide for user feedback
            err_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(4))
            err_box.text_frame.text = error_message
            
        finally:
            # Clean up the temporary files
            if os.path.exists(mmd_file_path):
                os.remove(mmd_file_path)
            if os.path.exists(png_file_path):
                os.remove(png_file_path)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to call the function:
# if __name__ == '__main__':
#     # The function uses default examples if mermaid_definitions is not provided.
#     # You can provide your own dictionary of diagrams.
#     my_diagrams = {
#         "Morning Routine": """
#         graph TD
#             A(Start) --> B(Alarm rings);
#             B --> C{Ready to wake up?};
#             C -- No --> D(Snooze - 5 minutes);
#             D --> B;
#             C -- Yes --> E(Get out of bed);
#             E --> F(End);
#         """
#     }
#     create_slide(output_pptx_path="my_mermaid_presentation.pptx", mermaid_definitions=my_diagrams)
#     print("Presentation saved to my_mermaid_presentation.pptx")
