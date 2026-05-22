import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR

def create_slide(
    output_pptx_path: str,
    data: list,
    colors: list,
    slide_title: str = "Preparing Animated Presentation Slide",
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with a Serpentine Process Timeline infographic.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        data (list): A list of dictionaries, where each dictionary represents a stage.
                     Each dict should have keys: 'title', 'description', 'icon' (path to image).
        colors (list): A list of RGB tuples for styling each stage.
        slide_title (str): The main title for the slide.

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Slide Title ---
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), prs.slide_width - Inches(1), Inches(0.75))
    title_tf = title_shape.text_frame
    p = title_tf.paragraphs[0]
    p.text = slide_title
    p.font.bold = True
    p.font.size = Pt(32)
    p.font.color.rgb = RGBColor(68, 84, 106)
    p.alignment = PP_ALIGN.CENTER

    # --- Layout Parameters ---
    num_stages = len(data)
    total_width = Inches(12.5)
    start_x = (prs.slide_width - total_width) / 2
    
    stage_width = total_width / (num_stages - 0.5 if num_stages > 1 else 1)
    box_width = Inches(1.8)
    box_height = Inches(1.2)
    gap = stage_width - box_width

    y_center = prs.slide_height / 2 + Inches(0.2)
    arc_thickness = Pt(4)

    # --- Create Stages ---
    for i in range(num_stages):
        stage_data = data[i]
        color_rgb = RGBColor(*colors[i % len(colors)])
        
        x_pos = start_x + i * stage_width

        # --- Main Icon Container ---
        container = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, y_center - box_height/2, box_width, box_height
        )
        container.fill.background()
        container.line.color.rgb = color_rgb
        container.line.width = arc_thickness
        container.shadow.inherit = False

        # --- Icon ---
        if 'icon' in stage_data and stage_data['icon']:
            try:
                icon_size = Inches(0.6)
                slide.shapes.add_picture(
                    stage_data['icon'], 
                    x_pos + (box_width - icon_size) / 2, 
                    y_center - icon_size / 2, 
                    width=icon_size
                )
            except:
                # Fallback if icon path is invalid
                pass

        # --- Content Box ---
        content_box_height = Inches(0.35)
        content_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            x_pos, y_center + box_height/2 - content_box_height, 
            box_width, content_box_height
        )
        content_box.fill.solid()
        content_box.fill.fore_color.rgb = color_rgb
        content_box.line.fill.background()
        content_box.shadow.inherit = False
        tf_content = content_box.text_frame
        tf_content.text = "Content"
        tf_content.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        tf_content.paragraphs[0].font.bold = True
        tf_content.paragraphs[0].font.size = Pt(12)
        tf_content.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf_content.margin_bottom = 0
        tf_content.margin_top = 0

        # --- Title and Description ---
        title_box_height = Inches(0.3)
        desc_height = Inches(0.6)
        title_gap = Inches(0.3)

        if i % 2 != 0:  # Odd stages (1, 3, 5...) - Title Below
            title_y = y_center + box_height/2 + title_gap
            desc_y = title_y + title_box_height
        else:  # Even stages (0, 2, 4...) - Title Above
            title_y = y_center - box_height/2 - title_gap - title_box_height
            desc_y = title_y - desc_height

        # Title Box
        title_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            x_pos, title_y, 
            box_width, title_box_height
        )
        title_box.fill.solid()
        title_box.fill.fore_color.rgb = color_rgb
        title_box.line.fill.background()
        title_box.shadow.inherit = False
        tf_title = title_box.text_frame
        tf_title.text = stage_data.get('title', f"Title {i+1}")
        tf_title.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        tf_title.paragraphs[0].font.bold = True
        tf_title.paragraphs[0].font.size = Pt(12)
        tf_title.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf_title.margin_bottom = 0
        tf_title.margin_top = 0

        # Description Box
        desc_box = slide.shapes.add_textbox(x_pos - Inches(0.1), desc_y, box_width + Inches(0.2), desc_height)
        tf_desc = desc_box.text_frame
        tf_desc.text = stage_data.get('description', "Description for this stage.")
        tf_desc.paragraphs[0].font.size = Pt(10)
        tf_desc.paragraphs[0].font.color.rgb = RGBColor(89, 89, 89)
        tf_desc.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf_desc.word_wrap = True

        # --- Connecting Arcs ---
        if i < num_stages - 1:
            arc_color_rgb = RGBColor(*colors[(i + 1) % len(colors)])
            arc_radius = gap * 0.7
            arc_size = arc_radius * 2
            arc_x = x_pos + box_width

            if i % 2 == 0:  # Bottom arc
                arc_y = y_center + box_height/2 - arc_radius
                arc = slide.shapes.add_shape(MSO_SHAPE.ARC, arc_x, arc_y, arc_size, arc_size)
                arc.rotation = 0
                arc.adjustments[0] = Emu(270 * 60000)
                arc.adjustments[1] = Emu(90 * 60000)
            else:  # Top arc
                arc_y = y_center - box_height/2 - arc_radius
                arc = slide.shapes.add_shape(MSO_SHAPE.ARC, arc_x, arc_y, arc_size, arc_size)
                arc.rotation = 180
                arc.adjustments[0] = Emu(270 * 60000)
                arc.adjustments[1] = Emu(90 * 60000)
            
            arc.fill.background()
            arc.line.color.rgb = arc_color_rgb
            arc.line.width = arc_thickness
            arc.shadow.inherit = False

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
if __name__ == '__main__':
    # Define the data for each stage
    stage_data = [
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
        {"title": "Title", "description": "I hope and I believe that this template will save your time and money.", "icon": None},
    ]

    # Define the color palette
    stage_colors = [
        (46, 61, 73),
        (46, 117, 182),
        (112, 173, 71),
        (255, 192, 0),
        (192, 0, 0),
        (112, 48, 160)
    ]

    create_slide("serpentine_timeline.pptx", data=stage_data, colors=stage_colors)

