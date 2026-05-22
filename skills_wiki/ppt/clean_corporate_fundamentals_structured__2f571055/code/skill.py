import os
import urllib.request
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def add_fade_transition(slide):
    """
    Injects OOXML to add a Fade transition to a slide.
    """
    # Create the transition element structure
    # <p:transition xmlns:p="..."><p:fade/></p:transition>
    nsmap = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    transition = etree.Element('{http://schemas.openxmlformats.org/presentationml/2006/main}transition', nsmap=nsmap)
    fade = etree.SubElement(transition, '{http://schemas.openxmlformats.org/presentationml/2006/main}fade')
    
    # Append to the slide's XML element
    slide.element.append(transition)

def create_slide(
    output_pptx_path: str,
    title_text: str = "My Presentation",
    body_text: str = "Keep text to a minimum\nFocus on impactful points\nEnsure legibility for the back row",
    accent_color: tuple = (0, 112, 192),  # Corporate Blue
    text_color: tuple = (64, 64, 64),     # Dark Charcoal
    **kwargs,
) -> str:
    """
    Creates a foundational corporate presentation with 3 slides, 
    demonstrating text hierarchy, lists, media insertion, and fade transitions.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Convert color tuples
    acc_rgb = RGBColor(*accent_color)
    txt_rgb = RGBColor(*text_color)

    # === Slide 1: Title Slide ===
    slide1 = prs.slides.add_slide(prs.slide_layouts[0]) # Title layout
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]

    title1.text = title_text.upper()
    title1.text_frame.paragraphs[0].font.bold = True
    title1.text_frame.paragraphs[0].font.color.rgb = txt_rgb
    title1.text_frame.paragraphs[0].font.name = 'Calibri'
    
    subtitle1.text = "A Foundational Corporate Layout"
    subtitle1.text_frame.paragraphs[0].font.color.rgb = acc_rgb

    # === Slide 2: Bulleted List Slide ===
    slide2 = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content layout
    title2 = slide2.shapes.title
    body2 = slide2.placeholders[1]

    title2.text = "Key Principles"
    title2.text_frame.paragraphs[0].font.bold = True
    title2.text_frame.paragraphs[0].font.color.rgb = txt_rgb

    # Add bullets
    bullets = body_text.split('\n')
    body2.text = bullets[0]
    for bullet in bullets[1:]:
        p = body2.text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
    
    # Format all bullet text
    for paragraph in body2.text_frame.paragraphs:
        paragraph.font.size = Pt(28)
        paragraph.font.color.rgb = txt_rgb

    # === Slide 3: Media & Shapes Slide ===
    slide3 = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only layout
    title3 = slide3.shapes.title
    title3.text = "Process Flow"
    title3.text_frame.paragraphs[0].font.bold = True
    title3.text_frame.paragraphs[0].font.color.rgb = txt_rgb

    # 1. Insert an Image (Download from web with fallback)
    img_path = "temp_chart.jpg"
    try:
        urllib.request.urlretrieve("https://images.unsplash.com/photo-1551288049-bebda4e38f71?auto=format&fit=crop&w=800&q=80", img_path)
        pic = slide3.shapes.add_picture(img_path, Inches(1), Inches(2), width=Inches(6))
    except Exception as e:
        # Fallback to a placeholder rectangle if download fails
        fallback_shape = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(6), Inches(4))
        fallback_shape.fill.solid()
        fallback_shape.fill.fore_color.rgb = RGBColor(200, 200, 200)
        fallback_shape.text = "Image Download Failed"
    
    # 2. Insert a Shape (Block Arrow)
    arrow = slide3.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, 
        Inches(7.5), Inches(3.5), Inches(2), Inches(1)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = acc_rgb
    arrow.line.color.rgb = txt_rgb

    # 3. Add a Text Box
    txBox = slide3.shapes.add_textbox(Inches(10), Inches(3.25), Inches(2.5), Inches(1.5))
    tf = txBox.text_frame
    tf.text = "Next\nSteps"
    for p in tf.paragraphs:
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = acc_rgb
        p.alignment = PP_ALIGN.CENTER

    # === Apply Transitions ===
    for slide in prs.slides:
        add_fade_transition(slide)

    # Save and cleanup
    prs.save(output_pptx_path)
    if os.path.exists(img_path):
        os.remove(img_path)
        
    return output_pptx_path
