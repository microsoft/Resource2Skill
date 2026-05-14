def create_slide(
    output_pptx_path: str,
    title_text: str = "AGENDA",
    agenda_items: list = None,
    bg_color_rgb: tuple = (26, 188, 156),     # Teal
    node_color_rgb: tuple = (52, 73, 94),     # Dark Charcoal
    line_color_rgb: tuple = (255, 255, 255),  # White
    text_color_rgb: tuple = (255, 255, 255),  # White
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Connected Vertical Flow Agenda" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.dml import MSO_THEME_COLOR

    if agenda_items is None:
        agenda_items = ["Introduction", "Our Services", "Key Clients", "Project Portfolio", "Contact Us"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Layer 1: Background ---
    # Create a full slide rectangle for background to ensure consistent coloring
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(*bg_color_rgb)
    bg_shape.line.fill.background() # No line

    # --- Layer 2: Left Side Title ---
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(3.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.name = 'Segoe UI' # Good modern default
    p.font.color.rgb = RGBColor(*text_color_rgb)
    p.alignment = PP_ALIGN.LEFT

    # Add a small vertical accent line next to the title
    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(4.2), Inches(0.05), Inches(0.8)
    )
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = RGBColor(*text_color_rgb)
    accent_line.line.fill.background()

    # --- Layer 3: Vertical Track & Nodes ---
    num_items = len(agenda_items)
    track_x = Inches(6.0)
    start_y = Inches(1.2)
    end_y = Inches(6.3)
    
    # Calculate spacing
    if num_items > 1:
        y_step = (end_y - start_y) / (num_items - 1)
    else:
        y_step = 0

    # Draw the main vertical connecting line
    track_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        track_x - Pt(1.5),  # Center the line
        start_y, 
        Pt(3),              # Thickness
        end_y - start_y
    )
    track_line.fill.solid()
    track_line.fill.fore_color.rgb = RGBColor(*line_color_rgb)
    track_line.line.fill.background()

    # Draw nodes and text
    node_radius = Inches(0.25)
    
    for i, item_text in enumerate(agenda_items):
        current_y = start_y + (i * y_step)
        
        # 1. Add Node (Circle)
        # Position is calculated from top-left, so subtract radius to center it on track_x/current_y
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            track_x - node_radius, 
            current_y - node_radius, 
            node_radius * 2, 
            node_radius * 2
        )
        node.fill.solid()
        node.fill.fore_color.rgb = RGBColor(*node_color_rgb)
        node.line.color.rgb = RGBColor(*line_color_rgb)
        node.line.width = Pt(2)
        
        # Add number inside node
        node_tf = node.text_frame
        node_tf.margin_left = 0
        node_tf.margin_right = 0
        node_tf.margin_top = 0
        node_tf.margin_bottom = 0
        node_p = node_tf.paragraphs[0]
        node_p.text = f"{i+1:02d}" # "01", "02", etc.
        node_p.font.size = Pt(14)
        node_p.font.bold = True
        node_p.font.name = 'Segoe UI'
        node_p.font.color.rgb = RGBColor(*text_color_rgb)
        node_p.alignment = PP_ALIGN.CENTER

        # 2. Add Agenda Item Text
        text_x = track_x + node_radius + Inches(0.3)
        # Shift Y up slightly so text aligns visually with the center of the circle
        text_y = current_y - Inches(0.2) 
        
        item_box = slide.shapes.add_textbox(text_x, text_y, Inches(5.0), Inches(0.5))
        item_tf = item_box.text_frame
        item_p = item_tf.paragraphs[0]
        item_p.text = item_text
        item_p.font.size = Pt(20)
        item_p.font.name = 'Segoe UI'
        item_p.font.color.rgb = RGBColor(*text_color_rgb)
        item_p.alignment = PP_ALIGN.LEFT

    prs.save(output_pptx_path)
    return output_pptx_path
