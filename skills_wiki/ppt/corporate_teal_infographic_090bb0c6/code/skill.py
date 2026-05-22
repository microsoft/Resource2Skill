def create_slide(
    output_pptx_path: str,
    title_text: str = "MARKETING PLAN",
    subtitle_text: str = "Go to market & growth strategy development. We help you scope and plan next steps.",
    main_header: str = "Supporting the founder",
    main_body: str = "Monitor every step, documenting and creating a playbook for the conversion process. Provide right-sized sales management support to move quickly toward success.",
    accent_color_rgb: tuple = (0, 169, 157),
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Corporate Teal Infographic' design style,
    specifically modeling the 'Marketing Plan' slide.

    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw, ImageFont
    import io

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Define Colors
    TEAL = RGBColor.from_rgb(*accent_color_rgb)
    DARK_GRAY = RGBColor(64, 64, 64)
    LIGHT_GRAY = RGBColor(242, 242, 242)
    WHITE = RGBColor(255, 255, 255)
    
    # --- Helper function to generate an icon ---
    def create_icon(icon_type, size=(100, 100), color=(0, 169, 157)):
        img = Image.new("RGBA", size, (255, 255, 255, 0))
        draw = ImageDraw.Draw(img)
        pen_width = 6
        
        if icon_type == "vision": # Telescope
            draw.line([(25, 75), (50, 50)], fill=color, width=pen_width+4)
            draw.line([(50, 50), (75, 25)], fill=color, width=pen_width+4)
            draw.ellipse([(65, 15), (85, 35)], outline=color, width=pen_width)
        elif icon_type == "stabilize": # Gears
            for i in range(8):
                angle = i * 45
                import math
                x1 = 50 + 35 * math.cos(math.radians(angle))
                y1 = 50 + 35 * math.sin(math.radians(angle))
                x2 = 50 + 45 * math.cos(math.radians(angle))
                y2 = 50 + 45 * math.sin(math.radians(angle))
                draw.line([(x1, y1), (x2, y2)], fill=color, width=pen_width+2)
            draw.ellipse([(15, 15), (85, 85)], outline=color, width=pen_width)
            draw.ellipse([(35, 35), (65, 65)], outline=color, width=pen_width)
        elif icon_type == "assess": # Document with check
            draw.rectangle([(20, 15), (70, 85)], outline=color, width=pen_width)
            draw.line([(30, 30), (60, 30)], fill=color, width=pen_width)
            draw.line([(30, 45), (60, 45)], fill=color, width=pen_width)
            draw.line([(55, 60), (65, 75)], fill=color, width=pen_width+2)
            draw.line([(65, 75), (85, 50)], fill=color, width=pen_width+2)
        elif icon_type == "coach": # Chart
            draw.line([(20, 80), (80, 80)], fill=color, width=pen_width)
            draw.line([(20, 80), (20, 20)], fill=color, width=pen_width)
            draw.rectangle([(30, 50), (40, 79)], fill=color)
            draw.rectangle([(50, 30), (60, 79)], fill=color)
            draw.rectangle([(70, 60), (80, 79)], fill=color)

        image_stream = io.BytesIO()
        img.save(image_stream, format="PNG")
        image_stream.seek(0)
        return image_stream

    # === Layer 1: Background & Main Panes ===
    # Left Teal Pane
    left_pane = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(3.5), Inches(7.5))
    left_pane.fill.solid()
    left_pane.fill.fore_color.rgb = TEAL
    left_pane.line.fill.background()

    # === Layer 2: Content ===
    # --- Left Pane Text ---
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2.5), Inches(2))
    p = title_box.text_frame.paragraphs[0]
    p.text = "\n".join(title_text.split()) # Stack words
    p.font.name = 'Montserrat ExtraBold'
    p.font.size = Pt(36)
    p.font.color.rgb = WHITE
    title_box.text_frame.margin_left = 0
    title_box.text_frame.margin_right = 0
    
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(2.7), Inches(1.5))
    p = subtitle_box.text_frame.paragraphs[0]
    p.text = subtitle_text
    p.font.name = 'Montserrat'
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE

    # --- Right Pane Content ---
    # Top Gray Box
    gray_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(0.5), Inches(9), Inches(2.0))
    gray_box.fill.solid()
    gray_box.fill.fore_color.rgb = LIGHT_GRAY
    gray_box.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(4.25), Inches(0.7), Inches(8.5), Inches(0.5))
    p = header_box.text_frame.paragraphs[0]
    p.text = main_header
    p.font.name = 'Montserrat SemiBold'
    p.font.size = Pt(18)
    p.font.color.rgb = DARK_GRAY

    body_box = slide.shapes.add_textbox(Inches(4.25), Inches(1.2), Inches(8.5), Inches(1.2))
    p = body_box.text_frame.paragraphs[0]
    p.text = main_body
    p.font.name = 'Montserrat'
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_GRAY
    body_box.text_frame.word_wrap = True

    # Process Flow Items
    process_items = [
        {"icon": "vision", "title": "Vision & Objective setting", "desc": "(Why & What?)"},
        {"icon": "stabilize", "title": "Stabilize product teams", "desc": "with revised roles & responsibilities"},
        {"icon": "assess", "title": "Introduce & Assess", "desc": "(continuous discovery & delivery)"},
        {"icon": "coach", "title": "Do & Train / Coach", "desc": ""},
    ]

    start_x = Inches(4.2)
    item_width = Inches(2.2)
    
    for i, item in enumerate(process_items):
        x_pos = start_x + (i * item_width)
        
        # Circle for icon
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x_pos, Inches(3.0), Inches(1.5), Inches(1.5))
        circle.fill.background()
        circle.line.color.rgb = TEAL
        circle.line.width = Pt(2.5)
        
        # Icon
        icon_stream = create_icon(item["icon"], color=accent_color_rgb)
        slide.shapes.add_picture(icon_stream, x_pos + Inches(0.25), Inches(3.25), Inches(1.0), Inches(1.0))
        
        # Text Box
        item_text_box = slide.shapes.add_textbox(x_pos - Inches(0.15), Inches(4.8), Inches(2.0), Inches(1.5))
        item_text_box.text_frame.text = f"{item['title']}\n{item['desc']}"
        
        # Title formatting
        p_title = item_text_box.text_frame.paragraphs[0]
        p_title.font.name = 'Montserrat SemiBold'
        p_title.font.size = Pt(12)
        p_title.font.color.rgb = DARK_GRAY
        p_title.alignment = PP_ALIGN.CENTER
        
        # Description formatting
        if len(item_text_box.text_frame.paragraphs) > 1:
            p_desc = item_text_box.text_frame.paragraphs[1]
            p_desc.font.name = 'Montserrat'
            p_desc.font.size = Pt(10)
            p_desc.font.color.rgb = DARK_GRAY
            p_desc.alignment = PP_ALIGN.CENTER
        
        item_text_box.text_frame.word_wrap = True

    # Page Number
    page_num_box = slide.shapes.add_textbox(Inches(12.8), Inches(7.0), Inches(0.5), Inches(0.5))
    p = page_num_box.text_frame.paragraphs[0]
    p.text = "3"
    p.font.name = 'Montserrat'
    p.font.size = Pt(10)
    p.font.color.rgb = DARK_GRAY
    p.alignment = PP_ALIGN.RIGHT

    prs.save(output_pptx_path)
    return output_pptx_path
