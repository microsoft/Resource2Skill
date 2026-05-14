AMBIENT_CAPABLE = True

def create_slide(
    output_pptx_path: str,
    title_text: str = "Season's Greetings",
    body_text: str = "",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Festive Serial Lights visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import io
    import math

    # Ambient helper import for continuous motion
    try:
        from _shell_helpers import add_pulse_loop
    except ImportError:
        def add_pulse_loop(*args, **kwargs):
            pass

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Dark Canvas ===
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(30, 34, 42)
    bg.line.fill.background()

    # === Layer 2: Glow Asset Generation ===
    def create_glow_image(color_rgb, radius=60):
        """Generate a radial transparent glow using PIL."""
        img = Image.new('RGBA', (radius*2, radius*2), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        for i in range(radius, 0, -1):
            factor = 1 - (i / radius)
            # Quadratic falloff for a softer, more realistic bloom
            alpha = int(255 * (factor ** 2))
            draw.ellipse((radius-i, radius-i, radius+i, radius+i), fill=color_rgb + (alpha,))
        
        img_io = io.BytesIO()
        img.save(img_io, format='PNG')
        return img_io

    glow_colors = [
        (255, 60, 60),   # Red
        (60, 255, 60),   # Green
        (60, 150, 255),  # Blue
        (255, 210, 50)   # Yellow
    ]
    glow_streams = [create_glow_image(c) for c in glow_colors]

    # === Layer 3: Wire & Bulbs Generation ===
    def draw_light_string(base_y, phase_offset=0, amplitude=Inches(0.6)):
        num_wire_points = 120
        wire_pts = []
        width = prs.slide_width
        
        # Calculate organic draped curve
        for i in range(num_wire_points):
            x = int(i * (width / (num_wire_points - 1)))
            # Fundamental wave + a harmonic wave to make the drape feel natural and slightly slack
            y_float = base_y + math.sin(x / width * 2 * math.pi * 1.5 + phase_offset) * amplitude \
                             + math.sin(x / width * 2 * math.pi * 3) * (amplitude * 0.3)
            wire_pts.append((x, int(y_float)))
        
        # Draw the physical wire
        ff_builder = slide.shapes.build_freeform()
        for i, (x, y) in enumerate(wire_pts):
            if i == 0:
                ff_builder.add_nodes(1, 0, x, y)
            else:
                ff_builder.add_nodes(0, 0, x, y)
        wire = ff_builder.convert_to_shape()
        wire.line.color.rgb = RGBColor(220, 220, 220)
        wire.line.width = Pt(1.5)

        # Place the bulbs and glowing halos
        num_bulbs = 18
        for i in range(num_bulbs):
            x = int(i * (width / (num_bulbs - 1)))
            y_float = base_y + math.sin(x / width * 2 * math.pi * 1.5 + phase_offset) * amplitude \
                             + math.sin(x / width * 2 * math.pi * 3) * (amplitude * 0.3)
            y = int(y_float)
            
            # --- Place Glow (Behind the actual bulb) ---
            color_idx = i % len(glow_streams)
            glow_stream = glow_streams[color_idx]
            glow_stream.seek(0)
            
            glow_size = int(Inches(1.2))
            glow = slide.shapes.add_picture(
                glow_stream, 
                x - int(glow_size/2), 
                y - int(glow_size/2), 
                glow_size, 
                glow_size
            )
            
            # Add infinite pulsing/blinking animation to the glow
            # Duration staggered by index to create out-of-sync flashing logic
            duration = 800 + (i % 3) * 300 
            add_pulse_loop(slide, glow, duration_ms=duration, scale_pct=25)
            
            # --- Place Physical Bulb (On top) ---
            bulb_size = int(Pt(10))
            bulb = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                x - int(bulb_size/2),
                y - int(bulb_size/2),
                bulb_size,
                bulb_size
            )
            bulb.fill.solid()
            bulb.fill.fore_color.rgb = RGBColor(255, 255, 255)
            bulb.line.fill.background() # No border

    # Execute draped lighting strings at top and bottom of slide
    draw_light_string(int(Inches(1.2)), phase_offset=0)
    draw_light_string(int(prs.slide_height - Inches(1.2)), phase_offset=math.pi)

    # === Layer 4: Content Setup ===
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2.5))
    tf = textbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Georgia"
    p.font.size = Pt(64)
    p.font.italic = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
