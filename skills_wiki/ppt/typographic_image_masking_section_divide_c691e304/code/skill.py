def create_slide(
    output_pptx_path: str,
    section_number: str = "01",
    title_text: str = "SECTION TITLE",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor congue massa. Fusce posuere, magna sed pulvinar ultricies, purus lectus malesuada libero.",
    accent_color: tuple = (139, 179, 45),  # Lime green
    bg_color: tuple = (26, 27, 32),        # Dark charcoal
    image_url: str = "https://images.unsplash.com/photo-1533038590840-1cbea976a55e?ixlib=rb-4.0.3&auto=format&fit=crop&w=800&q=80"
) -> str:
    """
    Create a PPTX file reproducing the 'Typographic Image Masking' section divider effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    
    # --- 1. Helper: Generate Masked Image Text using PIL ---
    def generate_masked_text_image(text, img_url, output_img_path):
        # Fetch image or use gradient fallback
        try:
            req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req) as response:
                base_img = Image.open(BytesIO(response.read())).convert("RGBA")
        except Exception as e:
            print(f"Image fetch failed ({e}). Using fallback gradient.")
            base_img = Image.new('RGBA', (800, 800), accent_color)
            top = Image.new('RGBA', (800, 800), (0, 100, 0, 255))
            mask = Image.new('L', (800, 800))
            mask.putdata([int(255 * (y / 800)) for y in range(800) for _ in range(800)])
            base_img.paste(top, (0, 0), mask)

        # Scale image to sufficient size
        base_img = base_img.resize((800, 800), Image.Resampling.LANCZOS)
        
        # Try to load a heavy font (Fallback chain for different OS)
        font_options = [
            "impact.ttf", "Arial Black.ttf", "arialbd.ttf", 
            "Trebuchet MS Bold.ttf", "DejaVuSans-Bold.ttf"
        ]
        font = None
        for f in font_options:
            try:
                font = ImageFont.truetype(f, 600)  # Massive font size
                break
            except IOError:
                continue
        if not font:
            font = ImageFont.load_default()
            print("Warning: Heavy fonts not found. Using default font.")

        # Create alpha mask for text
        text_mask = Image.new('L', base_img.size, 0)
        draw = ImageDraw.Draw(text_mask)
        
        # Draw text exactly in the center of the mask canvas
        draw.text((400, 400), text, fill=255, font=font, anchor="mm")
        
        # Apply mask to base image to extract the picture-filled text
        final_img = Image.new('RGBA', base_img.size, (0, 0, 0, 0))
        final_img.paste(base_img, (0, 0), text_mask)
        
        # Crop the transparent borders down to just the text box bounds
        bbox = text_mask.getbbox()
        if bbox:
            final_img = final_img.crop(bbox)
            
        final_img.save(output_img_path, format="PNG")
        return output_img_path

    # Prepare masked text asset
    temp_img_path = "temp_text_mask.png"
    generate_masked_text_image(section_number, image_url, temp_img_path)

    # --- 2. Initialize Presentation & Slide Layout ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # --- 3. Build the Slide Layers ---
    
    # Layer 1: Dark Background Rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*bg_color)
    bg.line.fill.background()

    # Layer 2: Typographic Image Mask ("01")
    # Positioned on the left side of the golden ratio split
    try:
        pic = slide.shapes.add_picture(temp_img_path, Inches(1.8), Inches(2.2), height=Inches(3.2))
    except Exception as e:
        print(f"Failed to place image mask: {e}")
    finally:
        if os.path.exists(temp_img_path):
            os.remove(temp_img_path)

    # Layer 3: Accent Divider Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(2.2), Inches(0.06), Inches(3.2))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # Layer 4: Section Title Text Box
    title_box = slide.shapes.add_textbox(Inches(5.5), Inches(2.05), Inches(6.0), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text.upper()
    p_title.font.name = "Arial"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(*accent_color)

    # Layer 5: Body Content Text Box
    body_box = slide.shapes.add_textbox(Inches(5.5), Inches(2.7), Inches(6.0), Inches(2.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(14)
    p_body.font.color.rgb = RGBColor(200, 200, 200)

    # Save to file
    prs.save(output_pptx_path)
    return output_pptx_path
