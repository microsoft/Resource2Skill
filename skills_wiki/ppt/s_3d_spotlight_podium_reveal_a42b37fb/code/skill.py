import os
import math
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "POWERPOINT\nUNIVERSITY",
    podium_color: tuple = (226, 0, 116),   # Magenta/Pink
    podium_shadow: tuple = (139, 0, 70),   # Darker Crimson
    bg_color_center: tuple = (173, 216, 230), # Light Blue
    bg_color_edge: tuple = (23, 42, 70)       # Dark Blue
) -> str:
    """
    Creates a slide with a cinematic 3D podium, leading ramp, and spotlights.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ---------------------------------------------------------
    # 1. GENERATE BACKGROUND (Radial Gradient via PIL)
    # ---------------------------------------------------------
    bg_width, bg_height = 1920, 1080
    bg_img = Image.new("RGB", (bg_width, bg_height))
    draw = ImageDraw.Draw(bg_img)
    
    max_radius = math.hypot(bg_width/2, bg_height/2)
    for r in range(int(max_radius), 0, -2):
        ratio = r / max_radius
        # Interpolate between edge and center colors
        r_col = int(bg_color_edge[0] * ratio + bg_color_center[0] * (1 - ratio))
        g_col = int(bg_color_edge[1] * ratio + bg_color_center[1] * (1 - ratio))
        b_col = int(bg_color_edge[2] * ratio + bg_color_center[2] * (1 - ratio))
        
        bbox = [
            bg_width/2 - r, bg_height/2 - r,
            bg_width/2 + r, bg_height/2 + r
        ]
        draw.ellipse(bbox, fill=(r_col, g_col, b_col))
    
    bg_stream = BytesIO()
    bg_img.save(bg_stream, format="PNG")
    bg_stream.seek(0)
    slide.shapes.add_picture(bg_stream, 0, 0, prs.slide_width, prs.slide_height)

    # ---------------------------------------------------------
    # 2. GENERATE AND PLACE SPOTLIGHTS (Soft Blur via PIL)
    # ---------------------------------------------------------
    beam_width, beam_height = 800, 1200
    beam_img = Image.new("RGBA", (beam_width, beam_height), (0,0,0,0))
    beam_draw = ImageDraw.Draw(beam_img)
    # Draw a cone shape
    beam_draw.polygon(
        [(beam_width/2 - 50, 0), (beam_width/2 + 50, 0), (beam_width, beam_height), (0, beam_height)],
        fill=(255, 255, 255, 70)
    )
    beam_img = beam_img.filter(ImageFilter.GaussianBlur(40))
    
    beam_stream = BytesIO()
    beam_img.save(beam_stream, format="PNG")
    
    # Left Spotlight
    beam_stream.seek(0)
    left_light = slide.shapes.add_picture(beam_stream, Inches(-1), Inches(-1.5), Inches(5), Inches(8))
    left_light.rotation = -35.0
    
    # Right Spotlight
    beam_stream.seek(0)
    right_light = slide.shapes.add_picture(beam_stream, Inches(9.333), Inches(-1.5), Inches(5), Inches(8))
    right_light.rotation = 35.0

    # ---------------------------------------------------------
    # 3. CONSTRUCT 3D PODIUM (Layered Vector Math)
    # ---------------------------------------------------------
    center_x = 13.333 / 2
    podium_y = 4.5  # Base anchor Y
    podium_w = 6.0
    podium_h = 1.5  # Squished oval height
    extrusion = 0.4 # Depth of the 3D edge
    
    def add_shape_filled(shape_type, left, top, width, height, r, g, b):
        sp = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor(r, g, b)
        sp.line.color.rgb = RGBColor(r, g, b)
        return sp

    # A) Podium Base Bottom Arc (Darker Grey for shadow)
    add_shape_filled(MSO_SHAPE.OVAL, center_x - podium_w/2, podium_y + extrusion, podium_w, podium_h, 200, 200, 200)
    
    # B) Podium Base Body (Rectangle bridging top and bottom ovals)
    add_shape_filled(MSO_SHAPE.RECTANGLE, center_x - podium_w/2, podium_y + podium_h/2, podium_w, extrusion, 230, 230, 230)
    
    # C) Podium Top Rim (White)
    add_shape_filled(MSO_SHAPE.OVAL, center_x - podium_w/2, podium_y, podium_w, podium_h, 255, 255, 255)

    # ---------------------------------------------------------
    # 4. DRAW RAMP (Freeform Polygons for Perspective)
    # ---------------------------------------------------------
    ramp_top_w = 3.5
    ramp_bottom_w = 10.0
    ramp_start_y = podium_y + (podium_h / 2) + 0.1 # Start from mid-point of top rim
    ramp_end_y = 7.6 # Extend slightly past bottom
    
    # Ramp Thickness (Shadow/Extrusion)
    ramp_ext = 0.2
    ff_shadow = slide.shapes.build_freeform()
    ff_shadow.add_line_segments([
        (Inches(center_x - ramp_top_w/2), Inches(ramp_start_y + ramp_ext)),
        (Inches(center_x + ramp_top_w/2), Inches(ramp_start_y + ramp_ext)),
        (Inches(center_x + ramp_bottom_w/2), Inches(ramp_end_y + ramp_ext)),
        (Inches(center_x - ramp_bottom_w/2), Inches(ramp_end_y + ramp_ext))
    ], close=True)
    shadow_shape = ff_shadow.convert_to_shape()
    shadow_shape.fill.solid()
    shadow_shape.fill.fore_color.rgb = RGBColor(*podium_shadow)
    shadow_shape.line.color.rgb = RGBColor(*podium_shadow)

    # Ramp Surface (Main)
    ff_main = slide.shapes.build_freeform()
    ff_main.add_line_segments([
        (Inches(center_x - ramp_top_w/2), Inches(ramp_start_y)),
        (Inches(center_x + ramp_top_w/2), Inches(ramp_start_y)),
        (Inches(center_x + ramp_bottom_w/2), Inches(ramp_end_y)),
        (Inches(center_x - ramp_bottom_w/2), Inches(ramp_end_y))
    ], close=True)
    main_shape = ff_main.convert_to_shape()
    main_shape.fill.solid()
    main_shape.fill.fore_color.rgb = RGBColor(*podium_color)
    main_shape.line.color.rgb = RGBColor(*podium_color)

    # ---------------------------------------------------------
    # 5. PODIUM TOP SURFACE (Overlaps the ramp to hide the top edge)
    # ---------------------------------------------------------
    surface_w = podium_w - 0.5
    surface_h = podium_h - 0.4
    surface_x = center_x - surface_w/2
    surface_y = podium_y + 0.2
    add_shape_filled(MSO_SHAPE.OVAL, surface_x, surface_y, surface_w, surface_h, *podium_color)

    # ---------------------------------------------------------
    # 6. CENTERPIECE (Logo Placeholder / Icon / Text)
    # ---------------------------------------------------------
    # Add a glowing back-plate for the logo
    glow_plate = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.75), Inches(surface_y - 0.5), Inches(1.5), Inches(1.5))
    glow_plate.fill.solid()
    glow_plate.fill.fore_color.rgb = RGBColor(255, 255, 255)
    glow_plate.line.fill.background()
    
    # Add Text on top of the plate
    txBox = slide.shapes.add_textbox(Inches(center_x - 1.5), Inches(surface_y - 0.35), Inches(3), Inches(1))
    tf = txBox.text_frame
    tf.text = "LOGO"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.color.rgb = RGBColor(*podium_color)
    tf.paragraphs[0].alignment = 2 # Center

    prs.save(output_pptx_path)
    return output_pptx_path
