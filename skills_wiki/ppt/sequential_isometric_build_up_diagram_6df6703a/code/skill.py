def create_slide(
    output_pptx_path: str,
    title_text: str = "Mixed-use Buildings",
    body_text: str = "Urban Planning Diagram",
    base_color: tuple = (240, 240, 240, 255),
    building_color: tuple = (255, 255, 255, 255),
    greenery_color: tuple = (190, 230, 200, 200),
    line_color: tuple = (100, 100, 100, 255),
    **kwargs,
) -> str:
    """
    Creates a multi-slide presentation simulating an isometric build-up diagram.
    Uses PIL to mathematically generate isometric polygons.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import io
    import math

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Slide dimensions in pixels for PIL
    W, H = 1920, 1080
    ORIGIN_X, ORIGIN_Y = W // 2, int(H * 0.65)
    SCALE = 50

    # Isometric Math Helper
    def iso_project(x, y, z):
        """Converts 3D grid coordinates to 2D isometric screen coordinates"""
        # Standard isometric projection angles (30 degrees)
        angle = math.radians(30)
        sx = ORIGIN_X + (x - y) * math.cos(angle) * SCALE
        sy = ORIGIN_Y + (x + y) * math.sin(angle) * SCALE - (z * SCALE)
        return (sx, sy)

    def draw_prism(draw, x, y, z, w, d, h, fill_color, edge_color):
        """Draws a 3D rectangular prism in isometric view"""
        # Calculate the 8 corners
        p1 = iso_project(x, y, z)
        p2 = iso_project(x+w, y, z)
        p3 = iso_project(x+w, y+d, z)
        p4 = iso_project(x, y+d, z)
        
        p5 = iso_project(x, y, z+h)
        p6 = iso_project(x+w, y, z+h)
        p7 = iso_project(x+w, y+d, z+h)
        p8 = iso_project(x, y+d, z+h)

        # Base color variations for 3D shading
        r, g, b, a = fill_color
        top_color = fill_color
        left_color = (int(r*0.9), int(g*0.9), int(b*0.9), a)
        right_color = (int(r*0.8), int(g*0.8), int(b*0.8), a)

        # Draw Left face (if w > 0 and h > 0)
        draw.polygon([p4, p3, p7, p8], fill=left_color, outline=edge_color, width=2)
        # Draw Right face (if d > 0 and h > 0)
        draw.polygon([p2, p3, p7, p6], fill=right_color, outline=edge_color, width=2)
        # Draw Top face
        draw.polygon([p5, p6, p7, p8], fill=top_color, outline=edge_color, width=2)

    def draw_flat_polygon(draw, points_3d, fill_color, edge_color):
        """Draws a flat polygon on a specific Z plane"""
        pts_2d = [iso_project(x, y, z) for x, y, z in points_3d]
        draw.polygon(pts_2d, fill=fill_color, outline=edge_color, width=2)

    # === GENERATE IMAGE LAYERS ===
    
    # Layer 1: Base Grid
    img_base = Image.new('RGBA', (W, H), (255, 255, 255, 0))
    draw_base = ImageDraw.Draw(img_base)
    draw_prism(draw_base, -6, -6, -0.5, 12, 12, 0.5, base_color, line_color)
    # Draw some grid lines on the base
    for i in range(-5, 6, 2):
        draw_flat_polygon(draw_base, [(i, -6, 0), (i+0.1, -6, 0), (i+0.1, 6, 0), (i, 6, 0)], line_color, line_color)
        draw_flat_polygon(draw_base, [(-6, i, 0), (-6, i+0.1, 0), (6, i+0.1, 0), (6, i, 0)], line_color, line_color)
    base_io = io.BytesIO()
    img_base.save(base_io, format='PNG')

    # Layer 2: Greenery / Environment
    img_env = Image.new('RGBA', (W, H), (255, 255, 255, 0))
    draw_env = ImageDraw.Draw(img_env)
    # A central park area
    green_pts = [(-2, -2, 0.05), (4, -2, 0.05), (4, 4, 0.05), (0, 4, 0.05), (0, 0, 0.05), (-2, 0, 0.05)]
    draw_flat_polygon(draw_env, green_pts, greenery_color, line_color)
    env_io = io.BytesIO()
    img_env.save(env_io, format='PNG')

    # Layer 3: Buildings (Hovering State)
    img_bld_hover = Image.new('RGBA', (W, H), (255, 255, 255, 0))
    draw_bld_hover = ImageDraw.Draw(img_bld_hover)
    HOVER_Z = 4  # Z offset to simulate falling
    draw_prism(draw_bld_hover, -4, 2, HOVER_Z, 3, 3, 4, building_color, line_color) # Main building
    draw_prism(draw_bld_hover, 1, 2, HOVER_Z, 2, 2, 2, building_color, line_color)  # Small building
    draw_prism(draw_bld_hover, -4, -4, HOVER_Z, 2, 3, 1.5, building_color, line_color) # Warehouse
    bld_hover_io = io.BytesIO()
    img_bld_hover.save(bld_hover_io, format='PNG')

    # Layer 4: Buildings (Landed State)
    img_bld_land = Image.new('RGBA', (W, H), (255, 255, 255, 0))
    draw_bld_land = ImageDraw.Draw(img_bld_land)
    LAND_Z = 0
    draw_prism(draw_bld_land, -4, 2, LAND_Z, 3, 3, 4, building_color, line_color)
    draw_prism(draw_bld_land, 1, 2, LAND_Z, 2, 2, 2, building_color, line_color)
    draw_prism(draw_bld_land, -4, -4, LAND_Z, 2, 3, 1.5, building_color, line_color)
    bld_land_io = io.BytesIO()
    img_bld_land.save(bld_land_io, format='PNG')

    # === BUILD PPTX SLIDES ===
    # Helper to add full-slide image
    def add_layer(slide, image_io):
        image_io.seek(0)
        slide.shapes.add_picture(image_io, 0, 0, width=Inches(13.333), height=Inches(7.5))

    def add_title(slide, text):
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.333), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(32)
        p.font.bold = True
        
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Base Only
    slide1 = prs.slides.add_slide(blank_layout)
    add_layer(slide1, base_io)
    add_title(slide1, "Platform Established")

    # Slide 2: Base + Hovering Buildings
    slide2 = prs.slides.add_slide(blank_layout)
    add_layer(slide2, base_io)
    add_layer(slide2, bld_hover_io)
    add_title(slide2, "Zoning & Massing")

    # Slide 3: Base + Landed Buildings
    slide3 = prs.slides.add_slide(blank_layout)
    add_layer(slide3, base_io)
    add_layer(slide3, bld_land_io)
    add_title(slide3, "Structures Settled")

    # Slide 4: Base + Greenery + Landed Buildings (Final)
    slide4 = prs.slides.add_slide(blank_layout)
    add_layer(slide4, base_io)
    add_layer(slide4, env_io)
    add_layer(slide4, bld_land_io)
    add_title(slide4, title_text)

    # Add descriptive floating labels to final slide
    labels = [
        ("Green Space", Inches(8.5), Inches(4.5)),
        ("Residential Block", Inches(3.0), Inches(2.5)),
        ("Commercial Hub", Inches(8.0), Inches(2.0))
    ]
    for text, x, y in labels:
        tx = slide4.shapes.add_textbox(x, y, Inches(2), Inches(0.5))
        p = tx.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True

    prs.save(output_pptx_path)
    return output_pptx_path
