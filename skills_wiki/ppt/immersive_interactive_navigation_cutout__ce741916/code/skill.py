import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "Presentation Title",  # Not strictly used in this visual pattern
    body_text: str = "",
    bg_palette: str = "business", 
    accent_color: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Creates an interactive PPTX with a Web-Style Top Navigation and a 
    Vertical Agenda featuring a PIL-generated "Cutout" highlight overlay.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Agenda/Menu items based on the tutorial
    menu_items = [
        "Accelerate our cultural transformation",
        "Strategic enterprise accounts",
        "Customer acquisition",
        "Customer retention and growth"
    ]
    
    # Web Top Nav items
    web_nav_items = ["Home", "Works", "Services", "About Us"]
    
    # Colors
    bg_color = (13, 110, 204) # Deep Azure Blue from tutorial
    overlay_color = (0, 40, 100, 180) # Semi-transparent dark blue overlay
    
    # We need to create the slides first to link them later
    slides = []
    for _ in range(len(menu_items)):
        slides.append(prs.slides.add_slide(prs.slide_layouts[6])) # Blank layout
        
    # Standard dimensions
    slide_w_px, slide_h_px = 1280, 720
    
    # Layout constants for the vertical agenda
    start_x = Inches(2.0)
    start_y = Inches(2.5)
    item_spacing = Inches(0.8)
    item_width = Inches(6.0)
    item_height = Inches(0.5)
    
    for i, slide in enumerate(slides):
        # === Layer 1: Base Background ===
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(*bg_color)
        bg_shape.line.fill.background()
        
        # Add Agenda Title
        title_box = slide.shapes.add_textbox(start_x, Inches(1.2), Inches(4), Inches(0.8))
        tf = title_box.text_frame
        p = tf.add_paragraph()
        p.text = "会议主题 (Meeting Agenda)"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        
        # === Layer 2: Text Content (Under the overlay) ===
        for j, item in enumerate(menu_items):
            y_pos = start_y + (j * item_spacing)
            
            # Number icon
            num_box = slide.shapes.add_textbox(start_x - Inches(0.6), y_pos, Inches(0.5), item_height)
            num_tf = num_box.text_frame
            num_p = num_tf.add_paragraph()
            num_p.text = f"{j+1}"
            num_p.font.size = Pt(16)
            num_p.font.color.rgb = RGBColor(255, 255, 255)
            
            # Text item
            tb = slide.shapes.add_textbox(start_x, y_pos, item_width, item_height)
            tf = tb.text_frame
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(255, 255, 255)

        # === Layer 3: PIL Masking Overlay (The "Cutout") ===
        # Calculate pixel coordinates for the cutout based on the active item (i)
        # 1 inch = 96 pixels approximately in PIL standard rendering
        dpi = 96
        active_y_px = int((start_y.inches + (i * item_spacing.inches) - 0.1) * dpi)
        active_h_px = int((item_height.inches + 0.2) * dpi)
        
        # Create a full-slide image with alpha channel
        overlay_img = Image.new('RGBA', (slide_w_px, slide_h_px), overlay_color)
        draw = ImageDraw.Draw(overlay_img)
        
        # Punch a transparent hole (cutout) for the active item
        bbox = [
            (int((start_x.inches - 0.8) * dpi), active_y_px), 
            (int((start_x.inches + 6.0) * dpi), active_y_px + active_h_px)
        ]
        draw.rectangle(bbox, fill=(0, 0, 0, 0)) # Transparent window
        
        mask_path = f"temp_mask_slide_{i}.png"
        overlay_img.save(mask_path)
        
        # Insert the overlay
        slide.shapes.add_picture(mask_path, 0, 0, prs.slide_width, prs.slide_height)
        os.remove(mask_path) # Cleanup
        
        # === Layer 4: Web Navigation Top Menu ===
        nav_start_x = Inches(4.5)
        nav_y = Inches(0.3)
        nav_spacing = Inches(1.5)
        
        for k, nav_item in enumerate(web_nav_items):
            x_pos = nav_start_x + (k * nav_spacing)
            
            # Create clickable text box
            nav_box = slide.shapes.add_textbox(x_pos, nav_y, Inches(1.2), Inches(0.5))
            tf = nav_box.text_frame
            p = tf.add_paragraph()
            p.text = nav_item
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
            
            # Add Interactivity: Link to respective slides (if within range)
            if k < len(slides):
                # Apply click action to jump to slide
                nav_box.click_action.target_slide = slides[k]
            
            # Add Underline Indicator for the active web menu item
            # To sync with the agenda, we assume Web Nav 1 = Agenda 1, etc.
            if k == i:
                line = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, 
                    x_pos + Inches(0.2), nav_y + Inches(0.45), 
                    Inches(0.8), Pt(3)
                )
                line.fill.solid()
                line.fill.fore_color.rgb = RGBColor(255, 255, 255)
                line.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
