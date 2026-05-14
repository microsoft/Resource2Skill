def create_slide(
    output_pptx_path: str,
    title_text: str = "Our global reach",
    body_text: str = "overview",
    bg_palette: str = "architecture", 
    accent_color: tuple = (218, 213, 206),  # Corporate Taupe
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Monochrome Corporate Node Panels' visual effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageEnhance

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Colors
    charcoal = RGBColor(43, 43, 45)
    taupe = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    white = RGBColor(255, 255, 255)
    
    # 2. Generate Moody Background via PIL
    img_path = "temp_bg_desaturated.jpg"
    try:
        # Fetch image
        url = f"https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1600&auto=format&fit=crop"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read())).convert('RGB')
        
        # Desaturate and Darken
        enhancer = ImageEnhance.Color(img)
        img = enhancer.enhance(0.0)  # Make grayscale
        enhancer = ImageEnhance.Brightness(img)
        img = enhancer.enhance(0.35) # Darken significantly for text contrast
    except Exception:
        # Fallback if download fails
        img = Image.new('RGB', (1600, 900), color=(30, 30, 30))
        
    img.save(img_path)
    
    # 3. Add Split Background to Slide
    # Top Photo (Top 60%)
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=Inches(4.5))
    
    # Bottom Taupe Block (Bottom 40%)
    bottom_block = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(4.5), prs.slide_width, Inches(3.0)
    )
    bottom_block.fill.solid()
    bottom_block.fill.fore_color.rgb = taupe
    bottom_block.line.fill.background()
    
    # Clean up temp image
    if os.path.exists(img_path):
        os.remove(img_path)

    # 4. Add Header Typography
    tx_box = slide.shapes.add_textbox(0, Inches(0.4), prs.slide_width, Inches(1.0))
    tf = tx_box.text_frame
    
    # Main Title
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.color.rgb = white
    p.font.name = 'Arial' # Fallback for clean sans-serif
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(180, 180, 180)
    
    # Small divider line under title
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.416), Inches(1.5), Inches(0.5), Inches(0.03)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = white
    divider.line.fill.background()

    # 5. Create the 3 Bridging Panels
    panel_width = 3.2
    panel_height = 3.8
    y_pos = 2.2
    # Calculate perfect spacing
    gap = (13.333 - (3 * panel_width)) / 4

    for i in range(3):
        x_pos = gap + i * (panel_width + gap)
        
        # Add White Panel
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(y_pos), Inches(panel_width), Inches(panel_height)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = white
        panel.line.fill.background()
        
        # Add Panel Content
        content_box = slide.shapes.add_textbox(
            Inches(x_pos + 0.2), Inches(y_pos + 0.4), Inches(panel_width - 0.4), Inches(3.0)
        )
        ctf = content_box.text_frame
        ctf.word_wrap = True
        
        # Panel Header
        cp = ctf.paragraphs[0]
        cp.text = "CONTACT\nINFORMATION"
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = charcoal
        
        # Panel Body
        cp2 = ctf.add_paragraph()
        cp2.text = "\nTherefore it may be the case that issues arise, i.e. issues with national postal services, international customs issues etc."
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = RGBColor(110, 110, 110)
        
        # Add Overlapping Circular Node at the bottom edge
        circ_size = 0.7
        cx = x_pos + (panel_width / 2) - (circ_size / 2)
        cy = y_pos + panel_height - (circ_size / 2) # Intersects the bottom border
        
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(circ_size), Inches(circ_size)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = charcoal
        # Add a thick white border to the circle to make it "cut" into the panel
        circle.line.color.rgb = white
        circle.line.width = Pt(4)
        
        # Number inside circle
        circ_tf = circle.text_frame
        # Center text vertically & horizontally
        circ_tf.word_wrap = False
        circ_p = circ_tf.paragraphs[0]
        circ_p.text = str(i + 1)
        circ_p.alignment = PP_ALIGN.CENTER
        circ_p.font.size = Pt(14)
        circ_p.font.bold = True
        circ_p.font.color.rgb = white
        
        # Nudge text down slightly since pptx vertical centering can be finicky
        circ_tf.margin_top = Inches(0.12)

    # Save
    prs.save(output_pptx_path)
    return output_pptx_path
