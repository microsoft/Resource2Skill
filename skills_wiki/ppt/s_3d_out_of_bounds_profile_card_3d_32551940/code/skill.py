def create_slide(
    output_pptx_path: str,
    title_text: str = "SANJAY",
    body_text: str = "Project Founder & CEO",
    bg_palette: str = "corporate",
    accent_color: tuple = (28, 53, 94),  # Deep Navy Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the '3D Out-of-Bounds Profile Card' visual effect.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # 1. Setup Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Define colors
    COLOR_ACCENT = RGBColor(*accent_color)
    COLOR_TEXT_DARK = RGBColor(30, 30, 30)
    COLOR_TEXT_MUTED = RGBColor(120, 120, 120)
    COLOR_WATERMARK = RGBColor(240, 243, 245)

    # 2. Add Background Watermark (BIOGRAPHY)
    tx_watermark = slide.shapes.add_textbox(Inches(0), Inches(0.5), Inches(8), Inches(2))
    tf_watermark = tx_watermark.text_frame
    p = tf_watermark.paragraphs[0]
    p.text = "BIOGRAPHY"
    p.font.size = Pt(96)
    p.font.name = "Arial Black"
    p.font.color.rgb = COLOR_WATERMARK
    
    # 3. Create the Geometric Anchor (The underlying circle)
    circle_size = Inches(4.5)
    circle_left = Inches(1.5)
    circle_top = Inches(2.0)
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, circle_left, circle_top, circle_size, circle_size
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = COLOR_ACCENT
    circle.line.fill.background() # No line

    # 4. Generate a Dummy Transparent Portrait (Simulating remove.bg result)
    # We generate a transparent PNG of a person's upper body silhouette to demonstrate the 3D effect
    img_path = "temp_portrait.png"
    portrait_width, portrait_height = 800, 1000
    img = Image.new("RGBA", (portrait_width, portrait_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # Draw silhouette (suit and head)
    # Shoulders/Suit
    draw.pieslice([-200, 400, 1000, 1600], 180, 360, fill=(150, 160, 180, 255))
    # Head
    draw.ellipse([250, 100, 550, 500], fill=(220, 200, 190, 255))
    # V-neck / Tie simulation for business feel
    draw.polygon([(400, 400), (350, 600), (450, 600)], fill=(255, 255, 255, 255))
    draw.polygon([(400, 450), (380, 700), (420, 700)], fill=(50, 60, 80, 255))
    
    img.save(img_path, "PNG")

    # 5. Insert Cutout Portrait (The 3D Pop-Out)
    # Key concept: Height is larger than the circle, positioned so the head breaks the top boundary
    pic_height = Inches(5.8)
    pic_top = Inches(0.8) # Starts higher than the circle's top (2.0)
    pic_left = Inches(1.1)
    
    portrait_pic = slide.shapes.add_picture(img_path, pic_left, pic_top, height=pic_height)
    
    # Clean up temp image
    if os.path.exists(img_path):
        os.remove(img_path)

    # 6. Typography and Right-Side Layout
    text_start_left = Inches(6.5)
    
    # Name
    tx_name = slide.shapes.add_textbox(text_start_left, Inches(1.5), Inches(5), Inches(1))
    p_name = tx_name.text_frame.paragraphs[0]
    p_name.text = title_text.upper()
    p_name.font.size = Pt(44)
    p_name.font.bold = True
    p_name.font.color.rgb = COLOR_ACCENT
    p_name.font.name = "Arial"

    # Title
    tx_title = slide.shapes.add_textbox(text_start_left, Inches(2.2), Inches(5), Inches(0.5))
    p_title = tx_title.text_frame.paragraphs[0]
    p_title.text = body_text
    p_title.font.size = Pt(14)
    p_title.font.color.rgb = COLOR_TEXT_MUTED
    
    # Decorative Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, text_start_left, Inches(2.7), Inches(0.5), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()

    # Biography Sections (Simulating the structured layout from the tutorial)
    sections = [
        ("INTRODUCTION", "• Founder of multiple tech initiatives.\n• Expert in UI/UX and visual storytelling.\n• Passionate about crossing the boundaries of design and tech."),
        ("EXPERIENCE", "• 10+ years in corporate presentation design.\n• Consulted for Fortune 500 companies.\n• Authored best-selling design frameworks.")
    ]

    current_top = 3.2
    for title, content in sections:
        # Section Title Box
        tx_sec = slide.shapes.add_textbox(text_start_left, Inches(current_top), Inches(2), Inches(0.4))
        p_sec = tx_sec.text_frame.paragraphs[0]
        p_sec.text = title
        p_sec.font.size = Pt(12)
        p_sec.font.bold = True
        p_sec.font.color.rgb = COLOR_ACCENT
        
        # Section Content
        tx_content = slide.shapes.add_textbox(text_start_left, Inches(current_top + 0.4), Inches(5.5), Inches(1))
        tx_content.text_frame.word_wrap = True
        p_content = tx_content.text_frame.paragraphs[0]
        p_content.text = content
        p_content.font.size = Pt(10.5)
        p_content.font.color.rgb = COLOR_TEXT_DARK
        p_content.line_spacing = 1.5
        
        current_top += 1.5

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
