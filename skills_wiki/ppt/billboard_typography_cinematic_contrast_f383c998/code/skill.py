def create_slide(
    output_pptx_path: str,
    topic_number: str = "#1",
    main_phrase: str = "GOOD PRESENTATION\nSLIDES ARE CLEAR",
    use_blurred_photo: bool = True,
    bg_color_start: tuple = (20, 30, 48),  # Deep Navy
    bg_color_end: tuple = (36, 59, 85),    # Lighter Navy
    text_color: tuple = (255, 255, 255)
) -> str:
    """
    Create a PPTX file reproducing the 'Billboard Typography & Cinematic Contrast' effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageFilter, ImageDraw

    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    bg_path = "temp_cinematic_bg.png"
    width, height = 1920, 1080
    bg_success = False

    # === Layer 1: Background Generation via PIL ===
    if use_blurred_photo:
        raw_bg_path = "raw_bg.jpg"
        try:
            # Download a high-res abstract/business image
            req = urllib.request.Request(
                "https://picsum.photos/1920/1080?blur=5", 
                headers={'User-Agent': 'Mozilla/5.0'}
            )
            with urllib.request.urlopen(req) as response, open(raw_bg_path, 'wb') as out_file:
                out_file.write(response.read())
            
            with Image.open(raw_bg_path) as img:
                img = img.convert("RGBA")
                # Apply heavy Gaussian blur to remove distracting details
                img = img.filter(ImageFilter.GaussianBlur(radius=20))
                
                # Apply a dark semi-transparent wash (alpha 140) to ensure text pops
                overlay = Image.new("RGBA", img.size, (15, 20, 25, 140))
                final_bg = Image.alpha_composite(img, overlay)
                final_bg.save(bg_path)
                bg_success = True
        except Exception as e:
            print(f"Image fetch failed: {e}. Falling back to premium gradient.")
            bg_success = False
            if os.path.exists(raw_bg_path):
                os.remove(raw_bg_path)

    # Fallback / Solid Gradient Background
    if not bg_success:
        base = Image.new("RGB", (width, height), bg_color_start)
        draw = ImageDraw.Draw(base)
        for y in range(height):
            # Calculate gradient color
            r = int(bg_color_start[0] + (bg_color_end[0] - bg_color_start[0]) * y / height)
            g = int(bg_color_start[1] + (bg_color_end[1] - bg_color_start[1]) * y / height)
            b = int(bg_color_start[2] + (bg_color_end[2] - bg_color_start[2]) * y / height)
            draw.line([(0, y), (width, y)], fill=(r, g, b))
        base.save(bg_path)

    # Insert the background into the slide
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Text Integration (Billboard Style) ===
    
    # 1. Topic Number (Smaller, top center)
    if topic_number:
        txBox_num = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11.333), Inches(1))
        tf_num = txBox_num.text_frame
        tf_num.text = topic_number
        tf_num.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        font_num = tf_num.paragraphs[0].runs[0].font
        font_num.name = 'Arial'
        font_num.size = Pt(40)
        font_num.bold = True
        # Slightly muted color for secondary information
        font_num.color.rgb = RGBColor(190, 200, 210) 

    # 2. Main Phrase (Massive, dead center)
    # Position shifted slightly down to accommodate the number
    txBox_main = slide.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.333), Inches(3.5))
    tf_main = txBox_main.text_frame
    tf_main.word_wrap = True
    
    p_main = tf_main.paragraphs[0]
    p_main.alignment = PP_ALIGN.CENTER
    run_main = p_main.add_run()
    run_main.text = main_phrase
    
    font_main = run_main.font
    font_main.name = 'Arial'
    font_main.size = Pt(68) # Massive typography
    font_main.bold = True
    font_main.color.rgb = RGBColor(text_color[0], text_color[1], text_color[2])

    # Save output
    prs.save(output_pptx_path)

    # Cleanup temporary images
    if os.path.exists(bg_path):
        os.remove(bg_path)
    if os.path.exists("raw_bg.jpg"):
        os.remove("raw_bg.jpg")

    return output_pptx_path
