import os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.shapes.freeform import FreeformBuilder
from pptx.oxml import parse_xml

def _inject_alpha_gradient(shape, color, alpha_top=0, alpha_bottom=60000):
    """
    Injects a vertical linear gradient into a shape with varying alpha transparency.
    alpha values: 0 is fully transparent, 100000 is fully opaque.
    angle 5400000 is 90 degrees (top to bottom).
    """
    hex_color = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    xml = f"""
    <a:gradFill rotWithShape="1" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <a:gsLst>
        <a:gs pos="0">
          <a:srgbClr val="{hex_color}">
            <a:alpha val="{alpha_top}"/>
          </a:srgbClr>
        </a:gs>
        <a:gs pos="100000">
          <a:srgbClr val="{hex_color}">
            <a:alpha val="{alpha_bottom}"/>
          </a:srgbClr>
        </a:gs>
      </a:gsLst>
      <a:lin ang="5400000" scaled="1"/>
    </a:gradFill>
    """
    # Ensure shape has a fill block to replace
    shape.fill.solid()
    spPr = shape.element.spPr
    for child in spPr:
        if child.tag.endswith('Fill'):
            spPr.remove(child)
            break
    spPr.append(parse_xml(xml))

def _generate_thumbnail(filepath, color):
    """Generates a simple, beautiful placeholder thumbnail image to avoid network dependency."""
    img = Image.new('RGB', (200, 120), color)
    draw = ImageDraw.Draw(img)
    # Draw a stylized landscape/mountain abstraction
    draw.polygon([(0, 120), (60, 50), (120, 120)], fill=(255, 255, 255, 100))
    draw.polygon([(80, 120), (150, 30), (200, 120)], fill=(255, 255, 255, 80))
    draw.ellipse((140, 20, 170, 50), fill=(255, 255, 255, 150))
    img.save(filepath)

def create_slide(
    output_pptx_path: str,
    title_text: str = "Corporate Journey",
    body_text: str = "",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 3D Perspective Journey Timeline visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 1. Background (White)
    bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    # 2. Add Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(80, 80, 80)

    # 3. Create the Perspective Road
    # Coordinates for a sweeping path going from bottom-right to top-left
    road_pts = [
        (Inches(1.5), Inches(3.2)),
        (Inches(2.0), Inches(2.7)),
        (Inches(14.0), Inches(7.5)),
        (Inches(6.0), Inches(7.5))
    ]
    fb = FreeformBuilder(slide.shapes)
    fb.add_line_segments(road_pts, close=True)
    road = fb.convert_to_shape()
    road.fill.solid()
    road.fill.fore_color.rgb = RGBColor(230, 230, 230)
    road.line.fill.background()

    # 4. Define Milestones
    milestones = [
        {"year": "2018", "title": "MARKET EXPANSION", "color": (146, 208, 80), "align": "right"},
        {"year": "2017", "title": "NEW PARTNERSHIPS", "color": (255, 192, 0), "align": "left"},
        {"year": "2016", "title": "GLOBAL RELEASE", "color": (0, 112, 192), "align": "right"},
        {"year": "2015", "title": "SERIES B FUNDING", "color": (0, 176, 240), "align": "left"},
        {"year": "2014", "title": "COMPANY FOUNDED", "color": (112, 48, 160), "align": "right"},
    ]

    # Path interpolation boundaries
    start_x, start_y = 10.5, 6.7  # Front (closest)
    end_x, end_y = 3.0, 3.4       # Back (furthest)
    scale_front = 1.0
    scale_back = 0.55

    n = len(milestones)
    
    # Render from back to front to ensure proper z-indexing overlapping
    for i in range(n - 1, -1, -1):
        ms = milestones[i]
        
        # Calculate interpolation factor t (0 is front, 1 is back)
        t = i / (n - 1)
        x = start_x + (end_x - start_x) * t
        y = start_y + (end_y - start_y) * t
        scale = scale_front + (scale_back - scale_front) * t
        color_rgb = RGBColor(*ms["color"])
        
        # Base dimensions (Scaled)
        outer_w, outer_h = 1.4 * scale, 0.45 * scale
        inner_w, inner_h = 0.5 * scale, 0.16 * scale
        pin_h = 2.2 * scale
        pin_w = 0.04 * scale
        shadow_w = 0.18 * scale
        
        # --- A. Flat 3D Rings ---
        # Outer Ring
        outer_ring = slide.shapes.add_shape(9, Inches(x - outer_w/2), Inches(y - outer_h/2), Inches(outer_w), Inches(outer_h))
        outer_ring.fill.background() # transparent
        outer_ring.line.color.rgb = color_rgb
        outer_ring.line.width = Pt(3 * scale)
        
        # Inner Ring
        inner_ring = slide.shapes.add_shape(9, Inches(x - inner_w/2), Inches(y - inner_h/2), Inches(inner_w), Inches(inner_h))
        inner_ring.fill.solid()
        inner_ring.fill.fore_color.rgb = color_rgb
        inner_ring.line.fill.background()

        # --- B. Vertical Pin Shadow (Glow) ---
        # Placed slightly behind the exact center
        shadow_oval = slide.shapes.add_shape(9, Inches(x - shadow_w/2), Inches(y - pin_h), Inches(shadow_w), Inches(pin_h))
        shadow_oval.line.fill.background()
        _inject_alpha_gradient(shadow_oval, ms["color"], alpha_top=0, alpha_bottom=70000)

        # --- C. Vertical Solid Pin ---
        pin = slide.shapes.add_shape(1, Inches(x - pin_w/2), Inches(y - pin_h), Inches(pin_w), Inches(pin_h))
        pin.fill.solid()
        pin.fill.fore_color.rgb = color_rgb
        pin.line.fill.background()

        # --- D. Thumbnail Image ---
        img_w, img_h = 1.6 * scale, 1.0 * scale
        img_offset = 0.4 * scale
        
        img_x = x + img_offset if ms["align"] == "right" else x - img_offset - img_w
        img_y = y - pin_h - img_h + (0.2 * scale)
        
        tmp_img = f"temp_thumb_{i}.jpg"
        _generate_thumbnail(tmp_img, ms["color"])
        slide.shapes.add_picture(tmp_img, Inches(img_x), Inches(img_y), Inches(img_w), Inches(img_h))
        os.remove(tmp_img)

        # --- E. Text Block ---
        txt_w, txt_h = 1.6 * scale, 1.2 * scale
        txt_x = img_x
        txt_y = img_y + img_h + 0.05
        
        tb = slide.shapes.add_textbox(Inches(txt_x), Inches(txt_y), Inches(txt_w), Inches(txt_h))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Alignment logic
        align_enum = PP_ALIGN.LEFT if ms["align"] == "right" else PP_ALIGN.RIGHT

        # Year
        p1 = tf.paragraphs[0]
        p1.text = ms["year"]
        p1.font.size = Pt(int(24 * scale))
        p1.font.bold = True
        p1.font.color.rgb = color_rgb
        p1.alignment = align_enum

        # Title
        p2 = tf.add_paragraph()
        p2.text = ms["title"]
        p2.font.size = Pt(int(10 * scale))
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(100, 100, 100)
        p2.alignment = align_enum
        p2.space_before = Pt(2)

        # Subtitle/Body
        p3 = tf.add_paragraph()
        p3.text = "Type your text here. Type your subtitle here to explain the timeline."
        p3.font.size = Pt(int(8 * scale))
        p3.font.color.rgb = RGBColor(150, 150, 150)
        p3.alignment = align_enum
        p3.space_before = Pt(2)

    prs.save(output_pptx_path)
    return output_pptx_path
