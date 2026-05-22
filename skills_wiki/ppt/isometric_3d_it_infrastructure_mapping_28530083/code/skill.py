def create_slide(
    output_pptx_path: str,
    title_text: str = "ISOMETRIC IT INFRASTRUCTURE",
    body_text: str = "",
    bg_palette: str = "technology",
    accent_color: tuple = (230, 57, 70),  # Crimson Red
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Isometric 3D Network Topology Map.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Colors ===
    COLOR_BG = RGBColor(244, 245, 247)
    COLOR_FLOOR = RGBColor(235, 237, 240)
    COLOR_FLOOR_LINE = RGBColor(220, 222, 225)
    COLOR_PATH = RGBColor(*accent_color)
    COLOR_TEXT = RGBColor(43, 45, 66)
    
    # Node Colors
    COLOR_SERVER = "2B2D42"  # Dark Slate
    COLOR_DB = "8D99AE"      # Cool Gray
    COLOR_FW = "E63946"      # Crimson Accent

    # Apply Background
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = COLOR_BG

    # === Helper: Safely Append XML to spPr ===
    def append_to_spPr(spPr, elem):
        """Safely inserts 3D tags before a:extLst to prevent file corruption."""
        extLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}extLst')
        if extLst is not None:
            extLst.addprevious(elem)
        else:
            spPr.append(elem)

    # === Helper: Generate Darker Shade for 3D Extrusion ===
    def darken_hex(hex_color, factor=0.6):
        r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
        return f"{int(r * factor):02X}{int(g * factor):02X}{int(b * factor):02X}"

    # === Helper: Generate 3D Native Shape ===
    def add_3d_node(slide, cx, cy, shape_type, width, height, depth_pt, hex_color):
        # Visually offset so the 'bottom' of the 3D shape rests on the coordinate
        offset_y = (depth_pt / 72.0) * 0.8
        top = cy - offset_y - (height / 2)
        left = cx - (width / 2)
        
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.line.fill.background()
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)
        
        spPr = shape.element.spPr
        
        # Add 3D Camera (Isometric Top Up)
        scene3d = parse_xml(
            '<a:scene3d xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:camera prst="isoTopUp"/>'
            '<a:lightRig rig="threePt" dir="t"/>'
            '</a:scene3d>'
        )
        append_to_spPr(spPr, scene3d)
        
        # Add 3D Extrusion Depth and Bevel
        if depth_pt > 0:
            depth_emu = int(depth_pt * 12700)
            dark_hex = darken_hex(hex_color)
            sp3d = parse_xml(
                f'<a:sp3d xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" extrusionH="{depth_emu}">'
                '<a:bevelT w="25400" h="12700" prst="circle"/>'
                f'<a:extrusionClr><a:srgbClr val="{dark_hex}"/></a:extrusionClr>'
                '</a:sp3d>'
            )
            append_to_spPr(spPr, sp3d)
        return shape

    # === Helper: Draw Isometric Flat Path ===
    def draw_path(slide, start, end, color):
        ff = slide.shapes.build_freeform(start[0], start[1])
        ff.add_line_segments([(end[0], end[1])])
        shape = ff.convert_to_shape()
        shape.line.color.rgb = color
        shape.line.width = Pt(4)

    # === Isometric Coordinates Math ===
    # Using 30 degree angles: dy = dx * tan(30) = dx * 0.57735
    A = (Inches(2.0), Inches(4.0))                                      # Mobile Client
    B = (Inches(4.5), Inches(4.0 + (2.5 * 0.57735)))                    # Firewall
    C = (Inches(7.5), Inches(5.443 - (3.0 * 0.57735)))                  # Switch/App Farm
    D = (Inches(10.5), Inches(3.711 + (3.0 * 0.57735)))                 # Operational DB
    E = (Inches(10.5), Inches(3.711 - (3.0 * 0.57735)))                 # Web Server

    # === Layer 1: Isometric Floor Plane ===
    ff_floor = slide.shapes.build_freeform(Inches(6.5), Inches(0.536))
    ff_floor.add_line_segments([
        (Inches(12.5), Inches(4.0)), 
        (Inches(6.5), Inches(7.464)), 
        (Inches(0.5), Inches(4.0)), 
        (Inches(6.5), Inches(0.536))
    ])
    floor = ff_floor.convert_to_shape()
    floor.fill.solid()
    floor.fill.fore_color.rgb = COLOR_FLOOR
    floor.line.color.rgb = COLOR_FLOOR_LINE
    floor.line.width = Pt(1)

    # === Layer 2: Network Paths (Drawn back-to-front underneath objects) ===
    draw_path(slide, A, B, COLOR_PATH)
    draw_path(slide, B, C, COLOR_PATH)
    draw_path(slide, C, D, COLOR_PATH)
    draw_path(slide, C, E, COLOR_PATH)

    # === Layer 3: 3D Nodes ===
    # Node E: Web Server (Tall Block)
    add_3d_node(slide, E[0], E[1], MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.0), depth_pt=50, hex_color=COLOR_SERVER)
    
    # Node D: Operational DB (Cylinder)
    add_3d_node(slide, D[0], D[1], MSO_SHAPE.OVAL, Inches(1.2), Inches(1.2), depth_pt=60, hex_color=COLOR_DB)
    
    # Node C: App Server Farm (3 servers offset to form a cluster)
    # Drawn back-to-front based on isometric Y sorting
    add_3d_node(slide, C[0] + Inches(0.6), C[1] - Inches(0.346), MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), depth_pt=35, hex_color=COLOR_SERVER)
    add_3d_node(slide, C[0], C[1], MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), depth_pt=35, hex_color=COLOR_SERVER)
    add_3d_node(slide, C[0] - Inches(0.6), C[1] + Inches(0.346), MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), depth_pt=35, hex_color=COLOR_SERVER)
    
    # Node B: Firewall (Red Block)
    add_3d_node(slide, B[0], B[1], MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(0.6), depth_pt=20, hex_color=COLOR_FW)
    
    # Node A: Mobile Client (Thin Block)
    add_3d_node(slide, A[0], A[1], MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), depth_pt=8, hex_color=COLOR_DB)

    # === Layer 4: Text Labels ===
    def add_label(x, y, text):
        txBox = slide.shapes.add_textbox(x - Inches(1), y, Inches(2), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT
        p.alignment = PP_ALIGN.CENTER

    add_label(A[0], A[1] + Inches(0.3), "Mobile Client")
    add_label(B[0], B[1] + Inches(0.3), "Firewall")
    add_label(C[0], C[1] + Inches(0.6), "App Server Farm")
    add_label(D[0], D[1] + Inches(0.5), "Operational DB")
    add_label(E[0], E[1] + Inches(0.3), "Web Server")

    # === Title Block ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1))
    p_title = title_box.text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT

    prs.save(output_pptx_path)
    return output_pptx_path
