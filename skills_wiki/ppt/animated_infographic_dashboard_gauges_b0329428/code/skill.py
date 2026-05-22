import io
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw

def create_gauge_graphic(
    value: int,
    start_color: tuple,
    end_color: tuple,
    needle_color: tuple,
    size: int = 800
) -> Image:
    """
    Generates a single static gauge graphic using PIL.
    The needle is pre-rotated to the final value.

    Args:
        value (int): The percentage value (0-100) to display.
        start_color (tuple): The light RGBA color for the gradient.
        end_color (tuple): The dark RGBA color for the gradient.
        needle_color (tuple): The RGBA color for the needle.
        size (int): The canvas size (width and height).

    Returns:
        PIL.Image: An RGBA image of the gauge.
    """
    # Create a transparent canvas
    im = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(im)

    # Gauge properties
    center = (size // 2, size // 2)
    radius = size // 2 * 0.9
    thickness = size * 0.15
    num_segments = 5

    # Draw the segmented arc
    for i in range(num_segments):
        start_angle = 180 + i * (180 / num_segments)
        end_angle = 180 + (i + 1) * (180 / num_segments)

        # Interpolate color for the gradient
        ratio = i / (num_segments - 1)
        r = int(start_color[0] + ratio * (end_color[0] - start_color[0]))
        g = int(start_color[1] + ratio * (end_color[1] - start_color[1]))
        b = int(start_color[2] + ratio * (end_color[2] - start_color[2]))
        segment_color = (r, g, b)

        draw.arc(
            [center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius],
            start=start_angle,
            end=end_angle,
            fill=segment_color,
            width=int(thickness)
        )

    # Create the needle graphic on a separate canvas for rotation
    needle_im = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    needle_draw = ImageDraw.Draw(needle_im)

    # Define needle shape (pointing upwards)
    needle_length = radius * 0.9
    needle_base_radius = thickness * 0.5
    
    p1 = (center[0], center[1] - needle_length)
    p2 = (center[0] - needle_base_radius * 0.6, center[1])
    p3 = (center[0] + needle_base_radius * 0.6, center[1])
    
    # Draw needle triangle and circle base
    needle_draw.polygon([p1, p2, p3], fill=needle_color)
    needle_draw.ellipse(
        [
            center[0] - needle_base_radius, center[1] - needle_base_radius,
            center[0] + needle_base_radius, center[1] + needle_base_radius
        ],
        fill=needle_color
    )

    # Rotate the needle to the correct angle
    # The gauge spans from -90 (left) to +90 (right) in standard angle terms
    # Our gauge is 180 (left) to 0 (right).
    # Rotation angle: 0% -> 0 degrees, 100% -> 180 degrees
    rotation_angle = (value / 100.0) * 180
    rotated_needle = needle_im.rotate(90 - rotation_angle, resample=Image.BICUBIC, center=center)

    # Composite the rotated needle onto the gauge
    im.paste(rotated_needle, (0, 0), rotated_needle)

    # Crop the image to be a semi-circle
    cropped_im = im.crop((0, 0, size, size // 2 + int(needle_base_radius)))
    return cropped_im


def create_slide(
    output_pptx_path: str,
    title_text: str = "EFFICIENCY",
    **kwargs,
) -> str:
    """
    Creates a PPTX file with three static dashboard gauges.
    Animation must be added manually in PowerPoint.

    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set a plain white background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # --- Slide Title ---
    title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(1))
    title_tf = title_shape.text_frame
    title_tf.word_wrap = False
    p = title_tf.add_paragraph()
    p.text = title_text
    p.font.name = 'Arial Black'
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(30, 55, 89)
    p.alignment = 1 # PP_ALIGN.CENTER

    # --- Gauge Data and Colors ---
    gauges_data = [
        {"label": "REACH", "value": 80, "colors": {"start": (173, 216, 230), "end": (70, 130, 180), "needle": (30, 55, 89)}},
        {"label": "ENGAGEMENT", "value": 55, "colors": {"start": (251, 213, 181), "end": (238, 142, 72), "needle": (176, 85, 23)}},
        {"label": "AWARENESS", "value": 75, "colors": {"start": (197, 224, 180), "end": (112, 173, 71), "needle": (47, 83, 22)}},
    ]

    # --- Create and place gauges ---
    num_gauges = len(gauges_data)
    total_width = Inches(12)
    gauge_width = total_width / num_gauges
    start_left = (prs.slide_width - total_width) / 2
    top = Inches(2.0)
    gauge_img_width = Inches(3.5)

    for i, data in enumerate(gauges_data):
        # Generate the gauge graphic
        gauge_image = create_gauge_graphic(
            value=data["value"],
            start_color=data["colors"]["start"],
            end_color=data["colors"]["end"],
            needle_color=data["colors"]["needle"]
        )
        
        # Convert PIL image to a stream to add to pptx
        image_stream = io.BytesIO()
        gauge_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        
        # Add image to slide
        left_pos = start_left + (i * gauge_width) + (gauge_width - gauge_img_width) / 2
        pic = slide.shapes.add_picture(image_stream, left_pos, top, width=gauge_img_width)

        # Add metric label
        label_top = top + Inches(gauge_img_width.inches / 2 * (gauge_image.height / gauge_image.width)) + Inches(0.2)
        label_box = slide.shapes.add_textbox(left_pos, label_top, gauge_img_width, Inches(0.5))
        label_tf = label_box.text_frame
        p_label = label_tf.add_paragraph()
        p_label.text = data["label"]
        p_label.font.name = 'Arial'
        p_label.font.size = Pt(18)
        p_label.font.color.rgb = RGBColor(*data["colors"]["needle"])
        p_label.alignment = 1 # PP_ALIGN.CENTER

        # Add value label
        value_top = label_top + Inches(0.4)
        value_box = slide.shapes.add_textbox(left_pos, value_top, gauge_img_width, Inches(0.5))
        value_tf = value_box.text_frame
        p_value = value_tf.add_paragraph()
        p_value.text = f"{data['value']}%"
        p_value.font.name = 'Arial Black'
        p_value.font.size = Pt(28)
        p_value.font.color.rgb = RGBColor(*data["colors"]["needle"])
        p_value.alignment = 1 # PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
