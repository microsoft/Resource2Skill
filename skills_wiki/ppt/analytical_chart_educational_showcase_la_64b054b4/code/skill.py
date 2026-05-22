def create_slide(
    output_pptx_path: str,
    chart_type: str = "scatter",  # options: "scatter", "histogram"
    category_text: str = "RELATIONSHIPS",
    title_text: str = "SCATTERPLOT",
    subtitle_text: str = '"Scatterplot demonstrates the relationship between two variables (X, Y)"',
    accent_color: tuple = (219, 68, 55),  # Red/Orange
    text_color: tuple = (80, 80, 80),     # Dark Gray
) -> str:
    """
    Creates an Educational Chart Reference slide, combining native PPTX text layout
    with a dynamically generated, cleanly styled matplotlib chart.
    """
    import io
    import numpy as np
    import matplotlib.pyplot as plt
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Helper function for text boxes
    def add_text_box(slide, text, left, top, width, height, font_size, is_bold, color_rgb, alignment=PP_ALIGN.CENTER, is_italic=False):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = alignment
        p.font.size = Pt(font_size)
        p.font.bold = is_bold
        p.font.italic = is_italic
        p.font.color.rgb = RGBColor(*color_rgb)
        p.font.name = "Arial"
        return txBox

    # === Layer 1: Typography & Layout ===
    
    # Category Header (Top Left)
    add_text_box(slide, category_text, Inches(0.5), Inches(0.3), Inches(5), Inches(0.5), 
                 font_size=18, is_bold=True, color_rgb=(100, 100, 100), alignment=PP_ALIGN.LEFT)

    # Chart Title (Centered, Accent Color)
    add_text_box(slide, title_text, Inches(1), Inches(1.2), Inches(11.333), Inches(0.8), 
                 font_size=36, is_bold=True, color_rgb=accent_color)

    # Subtitle / Definition (Centered, Gray, Italic)
    add_text_box(slide, subtitle_text, Inches(1.5), Inches(1.8), Inches(10.333), Inches(0.6), 
                 font_size=20, is_bold=False, is_italic=True, color_rgb=text_color)


    # === Layer 2: Dynamic Chart Generation ===
    
    # Matplotlib styling for modern flat look
    plt.style.use('default')
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=300)
    
    # Remove top and right borders (spines) for a clean look
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['bottom'].set_color('#cccccc')
    ax.spines['left'].set_color('#cccccc')
    ax.tick_params(colors='#666666')
    
    np.random.seed(42) # For reproducibility

    if chart_type.lower() == "scatter":
        # Generate Scatter Data
        x = np.random.rand(100)
        y = x + np.random.normal(0, 0.2, 100)
        
        ax.scatter(x, y, s=100, c='#00BFFF', alpha=0.6, edgecolors='white', linewidths=0.5)
        ax.set_xlabel("Variable X", color='#666666', fontsize=10)
        ax.set_ylabel("Variable Y", color='#666666', fontsize=10)
        
    elif chart_type.lower() == "histogram":
        # Generate Histogram Data
        data = np.random.normal(20, 3, 1000)
        
        ax.hist(data, bins=30, color='#7B68EE', edgecolor='white', linewidth=1.2)
        ax.set_xlabel("Value Range", color='#666666', fontsize=10)
        ax.set_ylabel("Probability / Frequency", color='#666666', fontsize=10)

    # Save chart to a memory buffer
    image_stream = io.BytesIO()
    plt.tight_layout()
    plt.savefig(image_stream, format='png', transparent=True)
    plt.close(fig)
    image_stream.seek(0)

    # === Layer 3: Insert Chart into PPTX ===
    
    # Center the chart on the slide, below the text
    chart_left = Inches(2.66) # Centered (13.333 - 8) / 2
    chart_top = Inches(2.5)
    chart_width = Inches(8.0)
    
    slide.shapes.add_picture(image_stream, chart_left, chart_top, width=chart_width)

    # Save Presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_slide("Scatterplot_Reference.pptx", chart_type="scatter", title_text="SCATTERPLOT", subtitle_text='"Scatterplot demonstrates the relationship between two variables (X, Y)"')
# create_slide("Histogram_Reference.pptx", chart_type="histogram", title_text="HISTOGRAM", category_text="DISTRIBUTIONS", subtitle_text='"Demonstrates the frequency distribution of continuous data."')
