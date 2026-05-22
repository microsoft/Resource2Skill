def create_slide(
    output_pptx_path: str,
    name: str = "Shinichi Kudo",
    bio_text: str = (
        "High school detective, originally studying in Class B, Year 2 at Teitan High School. "
        "He is the childhood friend of Ran Mouri, the only son of Yusaku Kudo and Yukiko Kudo. "
        "Possessing first-class deductive reasoning skills, he is known as the 'Savior of the Japanese Police' "
        "and the 'Sherlock Holmes of the Heisei Era'. He is also the main helper of Inspector Megure from the "
        "Tokyo Metropolitan Police Department, who highly appreciates his abilities."
    ),
    stats: dict = None,
    accent_color: tuple = (93, 156, 227),  # Soft Corporate Blue
    body_color: tuple = (89, 89, 89),      # Dark Grey (reduces visual pressure)
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Text-Driven Grid Profile" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    if stats is None:
        stats = {
            "Age": "17",
            "Gender": "Male",
            "Height": "174 cm",
            "Weight": "58 kg",
            "DOB": "May 4th",
            "Nationality": "Japan"
        }

    prs = Presentation()
    # Set to 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a completely blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    accent_rgb = RGBColor(*accent_color)
    body_rgb = RGBColor(*body_color)

    # ==========================================
    # 1. H1 Header: Name Box (Top Left)
    # ==========================================
    name_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.0), Inches(5.0), Inches(1.0))
    name_tf = name_box.text_frame
    name_tf.word_wrap = True
    
    p_name = name_tf.paragraphs[0]
    p_name.text = name
    p_name.font.size = Pt(44)
    p_name.font.bold = True
    p_name.font.name = "Arial" # Fallback clean sans-serif
    p_name.font.color.rgb = accent_rgb

    # ==========================================
    # 2. Narrative Block: Bio (Bottom Left)
    # ==========================================
    bio_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(6.0), Inches(4.5))
    bio_tf = bio_box.text_frame
    bio_tf.word_wrap = True # Crucial for paragraph blocks
    
    p_bio = bio_tf.paragraphs[0]
    p_bio.text = bio_text
    p_bio.font.size = Pt(18)
    p_bio.font.name = "Arial"
    p_bio.font.color.rgb = body_rgb
    
    # Tutorial Key Insight: Justified alignment & 1.5 Line Spacing
    p_bio.alignment = PP_ALIGN.JUSTIFY
    p_bio.line_spacing = 1.5 

    # ==========================================
    # 3. Structured Data Block: Stats (Right)
    # ==========================================
    stats_box = slide.shapes.add_textbox(Inches(8.5), Inches(2.2), Inches(4.0), Inches(4.5))
    stats_tf = stats_box.text_frame
    stats_tf.word_wrap = True

    # Iterate through dictionary to build the key-value pairs
    first_paragraph = True
    for key, value in stats.items():
        if first_paragraph:
            p_stat = stats_tf.paragraphs[0]
            first_paragraph = False
        else:
            p_stat = stats_tf.add_paragraph()
            
        # Tutorial Key Insight: 1.5 Line Spacing for readability
        p_stat.line_spacing = 1.5
        
        # Add the Key (Bold, Accent Color)
        run_key = p_stat.add_run()
        # Add spaces for pseudo-column alignment between keys and values
        run_key.text = f"{key}    " 
        run_key.font.size = Pt(18)
        run_key.font.bold = True
        run_key.font.name = "Arial"
        run_key.font.color.rgb = accent_rgb
        
        # Add the Value (Regular, Dark Grey)
        run_val = p_stat.add_run()
        run_val.text = str(value)
        run_val.font.size = Pt(18)
        run_val.font.bold = False
        run_val.font.name = "Arial"
        run_val.font.color.rgb = body_rgb

    prs.save(output_pptx_path)
    return output_pptx_path
