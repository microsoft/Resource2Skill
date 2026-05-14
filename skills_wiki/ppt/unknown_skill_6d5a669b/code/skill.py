import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml

def apply_text_gradient(run, hex_color1: str, hex_color2: str, angle_deg: float = 0):
    """
    Injects Open XML to apply a linear gradient fill directly to a text run.
    This preserves text editability while achieving a vector-shape aesthetic.
    """
    rPr = run._r.get_or_add_rPr()
    
    # Remove existing solid fill if present to avoid conflicts
    for child in rPr:
        if 'solidFill' in child.tag:
            rPr.remove(child)

    # PowerPoint XML angles are measured in 1/60,000ths of a degree
    ang_val = int(angle_deg * 60000)
    
    # Clean hex strings
    c1 = hex_color1.lstrip('#')
    c2 = hex_color2.lstrip('#')

    gradFill_xml = f"""
    <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:gsLst>
            <a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>
            <a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>
        </a:gsLst>
        <a:lin ang="{ang_val}" scaled="0"/>
    </a:gradFill>
    """
    rPr.append(parse_xml(gradFill_xml))

def add_dynamic_text(slide, text, left, top, font_size, hex1, hex2, angle=0):
    """Helper to create zero-margin, gradient-filled italic text."""
    # Arbitrary large width/height, the zero margins will handle tightness
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(8), Inches(2))
    tf = txBox.text_frame
    
    # Remove internal margins for precise typographic alignment
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    
    # Use a universally available heavy font
    run.font.name = 'Arial Black'
    run.font.size = Pt(font_size)
    run.font.italic = True
    
    # Apply the gradient
    apply_text_gradient(run, hex1, hex2, angle)
    return txBox

def create_slide(
    output_pptx_path: str,
    number_text: str = "10",
    text_line1: str = "PROJECTS",
    text_line2: str = "POWERPOINT",
    text_line3: str = "MOTIONS",
    **kwargs,
) -> str:
    """
    Creates a slide demonstrating the Dynamic Geometric Gradient Typography effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Setup ===
    # Soft off-white background to make the white geometrical masks pop
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 246, 250)

    # === Layer 2: Geometric "Masking" Slashes ===
    # These represent the physical masking blocks used in the video for the entrance animation.
    # We render them as aesthetic background shapes intersecting the composition.
    
    # Left Slash
    slash1 = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(0.5), Inches(-1), Inches(3.5), Inches(9.5))
    slash1.fill.solid()
    slash1.fill.fore_color.rgb = RGBColor(255, 255, 255)
    slash1.line.fill.background() # No outline
    
    # Right Slash
    slash2 = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(7.0), Inches(2.5), Inches(5.0), Inches(6.0))
    slash2.fill.solid()
    slash2.fill.fore_color.rgb = RGBColor(255, 255, 255)
    slash2.line.fill.background()
    
    # Small Accent Slash
    slash3 = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(4.5), Inches(0.5), Inches(3.0), Inches(1.5))
    slash3.fill.solid()
    slash3.fill.fore_color.rgb = RGBColor(230, 235, 240)
    slash3.line.fill.background()

    # === Layer 3: Typography & Gradients ===
    # Gradient Palettes mimicking the vibrant tech/agency aesthetic
    pink_orange = ("E94057", "F27121")  # Vibrant Red-Pink to Orange
    purple_pink = ("8A2387", "E94057")  # Deep Purple to Pink

    # Add text boxes in a precise typographic lockup
    # 1. The massive anchor number
    add_dynamic_text(
        slide=slide, 
        text=number_text, 
        left=1.2, top=2.2, 
        font_size=160, 
        hex1=pink_orange[0], hex2=pink_orange[1], 
        angle=90
    )
    
    # 2. Top sub-heading (smaller)
    add_dynamic_text(
        slide=slide, 
        text=text_line1, 
        left=3.8, top=2.0, 
        font_size=36, 
        hex1=purple_pink[0], hex2=purple_pink[1], 
        angle=90
    )
    
    # 3. Middle main heading
    add_dynamic_text(
        slide=slide, 
        text=text_line2, 
        left=3.7, top=2.8, 
        font_size=82, 
        hex1=purple_pink[0], hex2=purple_pink[1], 
        angle=90
    )
    
    # 4. Bottom main heading
    add_dynamic_text(
        slide=slide, 
        text=text_line3, 
        left=3.6, top=4.0, 
        font_size=82, 
        hex1=pink_orange[0], hex2=pink_orange[1], 
        angle=90
    )

    prs.save(output_pptx_path)
    return output_pptx_path
