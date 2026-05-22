import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

def draw_arc_shape(slide, xc, yc, r_out, r_in, start_angle, end_angle, fill_color=None, line_color=None, bg_color=None):
    """
    Builds a custom block arc shape using FreeformBuilder for perfect precision.
    Angles are provided in radians.
    """
    num_segments = 100
    # Create vertices for outer arc
    t_vals = [start_angle + (end_angle - start_angle) * i / num_segments for i in range(num_segments + 1)]
    points = []
    
    # Outer curve
    for t in t_vals:
        points.append((xc + r_out * math.cos(t), yc - r_out * math.sin(t)))
    
    # Inner curve (reversed)
    for t in reversed(t_vals):
        points.append((xc + r_in * math.cos(t), yc - r_in * math.sin(t)))
    
    # Close shape
    points.append(points[0])

    builder = slide.shapes.build_freeform(points[0][0], points[0][1])
    builder.add_line_segments(points[1:])
    shape = builder.convert_to_shape()

    if fill_color:
        # This is a filled data arc
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        # Hide outline by blending with fill
        shape.line.color.rgb = fill_color 
    else:
        # This is the background track outline
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.0)
        
    return shape

def add_connector_and_label(slide, x_tip, y_tip, label_y, pct_val, subtitle, color):
    """Draws a line from the arc tip to a label, adding a dot marker and text box."""
    # Vertical connector line
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, 
        x_tip, y_tip, x_tip, label_y
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(1.0)

    # Dot marker at the tip of the arc
    marker_size = Inches(0.12)
    marker = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        x_tip - marker_size/2, y_tip - marker_size/2, 
        marker_size, marker_size
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = color
    marker.line.fill.background()

    # Label box at the top of the line
    tb = slide.shapes.add_textbox(x_tip - Inches(1.0), label_y - Inches(0.8), Inches(2.0), Inches(0.8))
    tf = tb.text_frame
    
    p = tf.paragraphs[0]
    p.text = f"{int(pct_val * 100)}%"
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.name = "Arial"
    p2.font.size = Pt(11)
    p2.font.color.rgb = color
    p2.alignment = PP_ALIGN.CENTER

def create_slide(
    output_pptx_path: str,
    title_text: str = "Main Title",
    body_text: str = "Dynamic gauge visualization\nreproduced via precise vector geometry.",
    bg_color_rgb: tuple = (5, 10, 60), 
    **kwargs,
) -> str:
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg_color = RGBColor(*bg_color_rgb)
    
    # === Layer 1: Background ===
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg_color

    # === Layer 2: Main Title ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(4.0), Inches(2.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Georgia"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf.add_paragraph()
    p2.text = body_text
    p2.font.name = "Arial"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(180, 190, 220)

    # === Layer 3: Concentric Arc Gauge ===
    center_x = Inches(8.5)
    center_y = Inches(6.5)

    # Data defined as: [Percentage, Inner Radius, Outer Radius, RGB Color, Label Y-height]
    tiers = [
        {"pct": 0.93, "r_in": Inches(3.0), "r_out": Inches(3.6), "color": RGBColor(0, 229, 255), "label_y": Inches(2.0), "name": "Tier 01"},
        {"pct": 0.75, "r_in": Inches(2.2), "r_out": Inches(2.8), "color": RGBColor(0, 255, 133), "label_y": Inches(1.0), "name": "Tier 02"},
        {"pct": 0.50, "r_in": Inches(1.4), "r_out": Inches(2.0), "color": RGBColor(100, 200, 255), "label_y": Inches(2.5), "name": "Tier 03"},
    ]

    for tier in tiers:
        # Draw background track (full 180 degrees -> pi to 0)
        draw_arc_shape(
            slide, center_x, center_y, tier["r_out"], tier["r_in"], 
            start_angle=math.pi, end_angle=0, 
            bg_color=bg_color, line_color=tier["color"]
        )

        # Draw filled data arc
        # 0% is at math.pi (left), 100% is at 0 (right).
        data_end_angle = math.pi * (1.0 - tier["pct"])
        draw_arc_shape(
            slide, center_x, center_y, tier["r_out"], tier["r_in"], 
            start_angle=math.pi, end_angle=data_end_angle, 
            fill_color=tier["color"]
        )

        # Add connection line and label
        r_mid = (tier["r_out"] + tier["r_in"]) / 2
        x_tip = center_x + r_mid * math.cos(data_end_angle)
        y_tip = center_y - r_mid * math.sin(data_end_angle)

        add_connector_and_label(
            slide, x_tip, y_tip, tier["label_y"], 
            tier["pct"], tier["name"], tier["color"]
        )

    prs.save(output_pptx_path)
    return output_pptx_path
