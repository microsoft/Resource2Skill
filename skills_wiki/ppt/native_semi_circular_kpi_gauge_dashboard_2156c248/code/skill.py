def create_slide(
    output_pptx_path: str,
    title_text: str = "Customer Service Team Benchmarking Dashboard",
    metrics: list = None,
    **kwargs,
) -> str:
    """
    Creates a slide featuring a 4-column semi-circular KPI Gauge Dashboard.
    Uses native Donut Charts manipulated via OXML to create editable gauges.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    # Default metrics if none provided
    if metrics is None:
        metrics = [
            {"value": 35, "label": "Response Rate", "desc": "First contact within 1hr", "color": (56, 163, 165)},
            {"value": 50, "label": "Resolution", "desc": "Solved on first call", "color": (128, 237, 153)},
            {"value": 72, "label": "CSAT Score", "desc": "Positive feedback ratio", "color": (87, 204, 153)},
            {"value": 98, "label": "Uptime", "desc": "System availability", "color": (34, 87, 122)},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(250, 250, 250) # Very light off-white

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1.0))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = RGBColor(60, 60, 60)

    # Layout dimensions for 4 columns
    total_cols = 4
    margin_x = Inches(1.0)
    available_width = prs.slide_width - (margin_x * 2)
    col_width = available_width / total_cols
    
    chart_size = Inches(2.2)
    chart_y = Inches(2.5)

    # Function to strip borders via OXML
    def remove_borders(series_element):
        for pt in series_element.xpath('.//c:dPt'):
            spPr = pt.find(qn('c:spPr'))
            if spPr is None:
                spPr = parse_xml(r'<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
                pt.append(spPr)
            ln = spPr.find(qn('a:ln'))
            if ln is not None:
                spPr.remove(ln)
            spPr.append(parse_xml(r'<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:ln>'))

    # Build each gauge
    for idx, metric in enumerate(metrics[:4]):
        col_center_x = margin_x + (idx * col_width) + (col_width / 2)
        chart_x = col_center_x - (chart_size / 2)

        # 1. Create Chart Data (Value, Remaining, Hidden Bottom Half)
        val = metric["value"]
        chart_data = CategoryChartData()
        chart_data.categories = ['Achieved', 'Remaining', 'Hidden']
        chart_data.add_series('Data', (val, 100 - val, 100))

        # 2. Add Donut Chart
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, chart_x, chart_y, chart_size, chart_size, chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = False
        chart.has_title = False

        # 3. Format Slices
        series = chart.series[0]
        
        # Achieved (Color)
        pt1_fill = series.points[0].format.fill
        pt1_fill.solid()
        pt1_fill.fore_color.rgb = RGBColor(*metric["color"])
        
        # Remaining (Light Grey)
        pt2_fill = series.points[1].format.fill
        pt2_fill.solid()
        pt2_fill.fore_color.rgb = RGBColor(220, 220, 220)

        # Hidden (Transparent via OXML)
        pt3_element = series.points[2].format._element
        spPr = pt3_element.get_or_add_spPr()
        # clear existing fills
        for child in list(spPr):
            if child.tag.endswith('Fill'):
                spPr.remove(child)
        spPr.append(parse_xml(r'<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))

        # Remove borders
        remove_borders(series._element)

        # 4. Modify OXML for Gauge Angle (270) and Hole Size (70%)
        plot_elem = chart.plots[0]._element
        
        # Set angle
        fsa = plot_elem.find(qn('c:firstSliceAng'))
        if fsa is None:
            fsa = parse_xml(r'<c:firstSliceAng xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" val="270"/>')
            plot_elem.insert(0, fsa)
        else:
            fsa.set('val', '270')
            
        # Set hole size
        hs = plot_elem.find(qn('c:holeSize'))
        if hs is None:
            hs = parse_xml(r'<c:holeSize xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" val="70"/>')
            plot_elem.insert(1, hs)
        else:
            hs.set('val', '70')

        # 5. Add Typography
        # Visual center of the arch hole is at Y + (height/2)
        visual_baseline = chart_y + (chart_size / 2)
        
        # Percentage Text (Inside the hole, sitting on the baseline)
        txt_height = Inches(0.8)
        pct_box = slide.shapes.add_textbox(chart_x, visual_baseline - txt_height + Inches(0.1), chart_size, txt_height)
        tf = pct_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"{val}%"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(80, 80, 80)
        p.font.name = "Arial"

        # Label Text (Below the baseline)
        lbl_box = slide.shapes.add_textbox(chart_x - Inches(0.5), visual_baseline + Inches(0.2), chart_size + Inches(1.0), Inches(0.5))
        tf = lbl_box.text_frame
        p = tf.paragraphs[0]
        p.text = metric["label"]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*metric["color"])
        p.font.name = "Arial"

        # Description Text (Below label)
        desc_box = slide.shapes.add_textbox(chart_x - Inches(0.5), visual_baseline + Inches(0.6), chart_size + Inches(1.0), Inches(0.5))
        tf = desc_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = metric["desc"]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(150, 150, 150)
        p.font.name = "Arial"

    prs.save(output_pptx_path)
    return output_pptx_path
