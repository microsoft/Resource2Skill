def create_slide(
    output_pptx_path: str,
    title_text: str = "Unique Selling Points",
    accent_color: tuple = (255, 192, 0),  # Default: Mustard Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Intersecting Venn-Style Feature Cluster' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Colors
    accent_rgb = RGBColor(*accent_color)
    text_dark = RGBColor(64, 64, 64)
    line_gray = RGBColor(200, 200, 200)
    white = RGBColor(255, 255, 255)

    # === Step 1: Slide Title & Decorative Accent ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(5.0), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = text_dark

    # Small decorative accent next to the title
    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(0.5), Inches(0.7), Inches(0.3), Inches(0.4)
    )
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = accent_rgb
    accent_line.line.fill.background()

    # === Step 2: The Intersecting Venn Cluster (Center) ===
    center_x = 13.333 / 2
    center_y = 7.5 / 2
    circle_radius = 1.1 # 2.2 inches diameter
    
    # Calculate circle centers to form a triangle
    # Top circle
    c1_x = center_x
    c1_y = center_y - 0.6
    # Bottom Left circle
    c2_x = center_x - 0.7
    c2_y = center_y + 0.6
    # Bottom Right circle
    c3_x = center_x + 0.7
    c3_y = center_y + 0.6

    circle_coords = [(c1_x, c1_y), (c2_x, c2_y), (c3_x, c3_y)]

    for cx, cy in circle_coords:
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(cx - circle_radius), Inches(cy - circle_radius), 
            Inches(circle_radius * 2), Inches(circle_radius * 2)
        )
        circle.fill.background()  # Transparent fill
        circle.line.color.rgb = line_gray
        circle.line.width = Pt(2)

    # === Step 3: Central Icons ===
    icon_size = 0.6
    icon_shapes = [MSO_SHAPE.DIAMOND, MSO_SHAPE.LIGHTNING_BOLT, MSO_SHAPE.HEART]
    
    for i, (cx, cy) in enumerate(circle_coords):
        icon = slide.shapes.add_shape(
            icon_shapes[i],
            Inches(cx - icon_size/2), Inches(cy - icon_size/2),
            Inches(icon_size), Inches(icon_size)
        )
        icon.fill.solid()
        icon.fill.fore_color.rgb = accent_rgb
        icon.line.fill.background()

    # === Step 4: Flanking Feature Cards ===
    # Helper to build a two-part card (Solid Header + Text Body)
    def add_feature_card(slide_obj, x, y, title, desc, align="left"):
        card_w, header_h, body_h = 3.0, 0.45, 1.2
        
        # Header Rectangle
        header = slide_obj.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(card_w), Inches(header_h)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = accent_rgb
        header.line.fill.background()
        
        htf = header.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.text = title
        hp.alignment = PP_ALIGN.CENTER
        hp.font.bold = True
        hp.font.size = Pt(14)
        hp.font.color.rgb = white

        # Body Text Box
        body = slide_obj.shapes.add_textbox(Inches(x), Inches(y + header_h), Inches(card_w), Inches(body_h))
        btf = body.text_frame
        btf.word_wrap = True
        bp = btf.paragraphs[0]
        bp.text = desc
        bp.font.size = Pt(12)
        bp.font.color.rgb = text_dark
        
        if align == "right":
            bp.alignment = PP_ALIGN.RIGHT
        elif align == "center":
            bp.alignment = PP_ALIGN.CENTER
        else:
            bp.alignment = PP_ALIGN.LEFT

    # Data for the cards
    features = [
        {"title": "Solution 1: Precision", "desc": "Advanced metrics tracking providing real-time accuracy and minimizing operational drift effectively."},
        {"title": "Solution 2: Scalability", "desc": "Elastic architecture designed to grow seamlessly with your business demands without bottlenecks."},
        {"title": "Solution 3: Integration", "desc": "Unified API structure allowing frictionless connectivity with your existing ecosystem."}
    ]

    # Place Card 1 (Left)
    add_feature_card(slide, center_x - 4.5, center_y - 0.8, features[0]["title"], features[0]["desc"], align="right")
    
    # Place Card 2 (Right)
    add_feature_card(slide, center_x + 1.5, center_y - 0.8, features[1]["title"], features[1]["desc"], align="left")
    
    # Place Card 3 (Bottom Center)
    add_feature_card(slide, center_x - 1.5, center_y + 2.0, features[2]["title"], features[2]["desc"], align="center")

    prs.save(output_pptx_path)
    return output_pptx_path
