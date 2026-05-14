import matplotlib.pyplot as plt
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "Annual City PM2.5 Average Ranking",
    chart_data: dict = None,
    highlight_category: str = "City D",
    compare_categories: tuple = ("City A", "City B"),
    accent_color: tuple = (247, 148, 29),  # RGB for Orange
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a dashboard-style highlighted bar chart.

    This function uses matplotlib to generate a high-fidelity chart image
    and places it onto a dark-themed slide.

    Args:
        output_pptx_path (str): The path to save the generated .pptx file.
        title_text (str): The main title for the slide.
        chart_data (dict): Data for the chart, e.g., {'Category A': 80, 'Category B': 100}.
                           If None, default data is used.
        highlight_category (str): The category to highlight with the accent color.
        compare_categories (tuple): A tuple of two categories to draw comparison lines for.
        accent_color (tuple): The (R, G, B) accent color for highlights.

    Returns:
        str: The path to the saved PPTX file.
    """
    # --- Data Setup ---
    if chart_data is None:
        chart_data = {
            'City A': 80,
            'City B': 122.6,
            'City C': 127.2,
            'City D': 145.6
        }
    categories = list(chart_data.keys())
    values = list(chart_data.values())

    # --- Matplotlib Chart Generation ---
    plt.style.use('dark_background')
    fig, ax = plt.subplots(figsize=(10, 5.5))
    
    # Define colors
    accent_rgb_float = [c / 255.0 for c in accent_color]
    default_color = '#AAAAAA'
    bar_colors = [accent_rgb_float if cat == highlight_category else default_color for cat in categories]

    # Plot bars
    bars = ax.bar(categories, values, color=bar_colors, width=0.6)

    # Style the chart
    fig.set_facecolor('#1F1F1F')
    ax.set_facecolor('#1F1F1F')
    
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#555555')
    ax.spines['bottom'].set_color('#555555')

    ax.tick_params(axis='x', colors='white', length=0)
    ax.tick_params(axis='y', colors='white', length=0)
    ax.yaxis.grid(True, color='#444444', linestyle='-', linewidth=0.5)
    ax.set_axisbelow(True)

    # Add data labels
    for bar in bars:
        yval = bar.get_height()
        label_color = accent_rgb_float if bar.get_facecolor() == tuple(accent_rgb_float + [1.0]) else 'white'
        ax.text(bar.get_x() + bar.get_width()/2.0, yval + 5, f'{yval:.1f}', 
                ha='center', va='bottom', color=label_color, fontsize=12, weight='bold')

    # Add comparison annotation
    if len(compare_categories) == 2 and all(c in categories for c in compare_categories):
        idx1 = categories.index(compare_categories[0])
        idx2 = categories.index(compare_categories[1])
        val1 = values[idx1]
        val2 = values[idx2]
        
        ax.hlines(y=val1, xmin=-0.5, xmax=idx1, color=default_color, linestyle='--', linewidth=1)
        ax.hlines(y=val2, xmin=-0.5, xmax=idx2, color=default_color, linestyle='--', linewidth=1)
        
        ax.text(-0.6, val1, str(round(val1)), ha='right', va='center', color='white', fontsize=10)
        ax.text(-0.6, val2, str(round(val2)), ha='right', va='center', color='white', fontsize=10)
        
        diff = abs(val1 - val2)
        mid_y = (val1 + val2) / 2
        ax.plot([-0.55, -0.55], [val1, val2], color=accent_rgb_float, linewidth=2)
        ax.text(-0.7, mid_y, f"{diff:.1f}", ha='right', va='center', color=accent_rgb_float, fontsize=12, weight='bold')

    # Add trend arrow annotation (sticker)
    if highlight_category in categories:
        highlight_idx = categories.index(highlight_category)
        highlight_val = values[highlight_idx]
        ax.annotate('', xy=(highlight_idx, highlight_val * 1.15), xytext=(highlight_idx - 1, highlight_val * 0.9),
                    arrowprops=dict(arrowstyle='->, head_length=0.6, head_width=0.4',
                                    connectionstyle="arc3,rad=.2",
                                    ls='dotted', color='white'))

    plt.tight_layout()

    # Save plot to a BytesIO buffer
    img_buffer = io.BytesIO()
    plt.savefig(img_buffer, format='png', dpi=300, facecolor=fig.get_facecolor(), transparent=True)
    img_buffer.seek(0)
    plt.close(fig)

    # --- PowerPoint Slide Creation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(31, 31, 31) # #1F1F1F

    # Add title
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(1))
    title_tf = title_shape.text_frame
    p = title_tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Add chart image to slide
    chart_width = Inches(12)
    chart_height = Inches(6.18)
    slide.shapes.add_picture(img_buffer, Inches(0.66), Inches(1.2), width=chart_width, height=chart_height)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_slide("dashboard_chart.pptx")
#
# # Example with different data and colors
# custom_data = {
#     'Q1 Sales': 250, 'Q2 Sales': 410, 
#     'Q3 Sales': 380, 'Q4 Sales': 520
# }
# create_slide(
#     "sales_dashboard.pptx",
#     title_text="Quarterly Sales Performance",
#     chart_data=custom_data,
#     highlight_category="Q4 Sales",
#     compare_categories=("Q1 Sales", "Q4 Sales"),
#     accent_color=(46, 172, 109) # Green
# )

