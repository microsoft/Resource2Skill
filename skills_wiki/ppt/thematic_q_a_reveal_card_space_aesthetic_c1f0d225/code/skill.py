def create_slide(
    output_pptx_path: str,
    title_text: str = "Who was the first person to walk on the moon?",
    answer_text: str = "Neil Armstrong",
    question_number: str = "1",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Thematic Q&A Reveal Card' visual effect.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import random
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Generation via PIL ===
    bg_path = "temp_starfield_bg.png"
    width_px, height_px = 1920, 1080
    
    # Deep space navy blue
    img = Image.new('RGB', (width_px, height_px), color=(12, 14, 22))
    draw = ImageDraw.Draw(img)
    
    # Generate random starfield
    for _ in range(600):
        x = random.randint(0, width_px)
        y = random.randint(0, height_px)
        radius = random.uniform(0.5, 2.5)
        intensity = random.randint(100, 255) # Varying brightness
        draw.ellipse([x - radius, y - radius, x + radius, y + radius], 
                     fill=(intensity, intensity, intensity))
    img.save(bg_path)
    
    # Insert background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Thematic Graphics (Moon & Craters) ===
    # Main Moon Body
    moon = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(8.25), Inches(1.5), Inches(4.5), Inches(4.5)
    )
    moon.fill.solid()
    moon.fill.fore_color.rgb = RGBColor(210, 210, 210)
    moon.line.fill.background() # No outline

    # Crater Geometries (relative to moon positioning)
    craters = [
        (9.0, 2.5, 0.8, 0.8),
        (10.5, 4.0, 1.2, 1.0),
        (11.5, 3.0, 0.6, 0.6),
        (9.5, 4.5, 1.0, 0.8),
        (11.0, 5.0, 0.7, 0.6)
    ]
    for cx, cy, cw, ch in craters:
        crater = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx), Inches(cy), Inches(cw), Inches(ch)
        )
        crater.fill.solid()
        crater.fill.fore_color.rgb = RGBColor(180, 180, 180) # Darker grey
        crater.line.fill.background()

    # === Layer 3: Text & Content Composition ===
    
    # 1. Sequence Indicator Box
    num_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.0), Inches(2.2), Inches(1.2), Inches(1.5)
    )
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = RGBColor(90, 60, 130) # Thematic purple
    num_box.line.fill.background()
    num_frame = num_box.text_frame
    num_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    n_p = num_frame.paragraphs[0]
    n_p.text = str(question_number)
    n_p.font.size = Pt(64)
    n_p.font.bold = True
    n_p.font.name = 'Arial'
    n_p.font.color.rgb = RGBColor(255, 255, 255)
    n_p.alignment = PP_ALIGN.CENTER

    # 2. Question Text Block
    q_box = slide.shapes.add_textbox(
        Inches(2.5), Inches(2.0), Inches(5.5), Inches(1.5)
    )
    q_frame = q_box.text_frame
    q_frame.word_wrap = True
    q_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    q_p = q_frame.paragraphs[0]
    q_p.text = title_text
    q_p.font.size = Pt(36)
    q_p.font.name = 'Arial'
    q_p.font.color.rgb = RGBColor(255, 255, 255)
    q_p.alignment = PP_ALIGN.LEFT

    # 3. Highlighted Answer Box
    a_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.5), Inches(4.2), Inches(5.5), Inches(0.8)
    )
    a_box.fill.solid()
    a_box.fill.fore_color.rgb = RGBColor(30, 30, 30) # Dark grey fill
    a_box.line.color.rgb = RGBColor(218, 165, 32)    # Gold highlight border
    a_box.line.width = Pt(2)
    
    a_frame = a_box.text_frame
    a_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    a_frame.margin_left = Inches(0.2)
    a_p = a_frame.paragraphs[0]
    a_p.alignment = PP_ALIGN.LEFT

    # Styling "Answer:" prefix slightly dimmer to create visual distinction
    run1 = a_p.add_run()
    run1.text = "Answer:  "
    run1.font.size = Pt(24)
    run1.font.name = 'Arial'
    run1.font.color.rgb = RGBColor(180, 180, 180)

    # Bright actual answer text
    run2 = a_p.add_run()
    run2.text = answer_text
    run2.font.size = Pt(24)
    run2.font.bold = True
    run2.font.name = 'Arial'
    run2.font.color.rgb = RGBColor(255, 255, 255)

    # Cleanup temp image
    if os.path.exists(bg_path):
        os.remove(bg_path)

    prs.save(output_pptx_path)
    return output_pptx_path
