def create_slide(
    output_pptx_path: str,
    title_text: str = "Testimonial Report",
    body_text: str = "",
    bg_palette: str = "business",  
    accent_color: tuple = (40, 40, 40),  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Picture-in-Picture Video Narration Layout.
    Generates a simulated report and a mock webcam overlay placeholder.
    
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Generate Simulated Document / Report (PIL) ===
    # We generate a graphic to represent the "Test Report" shown in the video
    report_img_path = "temp_simulated_report.png"
    
    # Create a white canvas
    img = Image.new('RGB', (1200, 800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    # Draw simulated UI / Report elements
    # Header bar
    draw.rectangle([(100, 50), (1100, 120)], fill=(240, 245, 250))
    # Simulated Text Lines
    for i in range(3):
        draw.rectangle([(120, 160 + i*40), (600, 180 + i*40)], fill=(220, 220, 220))
    
    # Simulated Data Tables / Colored Indicators (Red/Yellow/Green like the video)
    draw.rectangle([(700, 160), (1100, 300)], outline=(200, 200, 200), width=2)
    draw.rectangle([(1000, 180), (1080, 200)], fill=(255, 100, 100)) # Red indicator
    draw.rectangle([(1000, 220), (1080, 240)], fill=(255, 200, 50))  # Yellow indicator
    draw.rectangle([(1000, 260), (1080, 280)], fill=(100, 200, 100)) # Green indicator
    
    # Simulated Detailed Data Rows
    for i in range(8):
        y_pos = 350 + (i * 45)
        draw.rectangle([(100, y_pos), (1100, y_pos + 30)], fill=(245, 245, 245))
        # Simulated data bars within rows
        draw.rectangle([(800, y_pos + 5), (800 + (i*30 % 200) + 50, y_pos + 25)], fill=(150, 180, 220))
        
    img.save(report_img_path)

    # === Layer 2: Insert Document into Slide ===
    # Centered with slight margins
    slide.shapes.add_picture(
        report_img_path, 
        left=Inches(1.666), 
        top=Inches(0.4), 
        width=Inches(10)
    )

    # === Layer 3: Picture-in-Picture (PiP) Webcam Placeholder ===
    # Positioned in the bottom right corner
    cam_width = Inches(2.8)
    cam_height = Inches(1.8)
    cam_left = prs.slide_width - cam_width - Inches(0.4)
    cam_top = prs.slide_height - cam_height - Inches(0.4)
    
    # Create rounded rectangle to simulate webcam feed border
    cam_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, cam_left, cam_top, cam_width, cam_height
    )
    
    # Style the placeholder (Dark gray fill to represent video, white border)
    cam_shape.fill.solid()
    cam_shape.fill.fore_color.rgb = RGBColor(accent_color[0], accent_color[1], accent_color[2])
    cam_shape.line.color.rgb = RGBColor(255, 255, 255)
    cam_shape.line.width = Pt(3)
    
    # Add descriptive text to the placeholder
    tf = cam_shape.text_frame
    tf.text = "WEBCAM FEED\n(Added via PowerPoint Recording)"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    for p in tf.paragraphs:
        p.font.color.rgb = RGBColor(200, 200, 200)
        p.font.size = Pt(12)
        p.font.bold = True

    # Save the presentation
    prs.save(output_pptx_path)
    
    # Cleanup temporary image file
    if os.path.exists(report_img_path):
        os.remove(report_img_path)
        
    return output_pptx_path
