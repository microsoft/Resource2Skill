import os
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.xmlchemy import OxmlElement

def add_text_shadow(run):
    """Adds a subtle black shadow to a text run."""
    el = run._r.get_or_add_rPr()
    effect_lst = OxmlElement("a:effectLst")
    outer_shdw = OxmlElement("a:outerShdw")
    outer_shdw.set("blurRad", "40000")
    outer_shdw.set("dist", "20000")
    outer_shdw.set("dir", "2700000")
    outer_shdw.set("algn", "bl")
    
    srgb_clr = OxmlElement("a:srgbClr")
    srgb_clr.set("val", "000000")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "65000") # 65% opacity
    srgb_clr.append(alpha)
    
    outer_shdw.append(srgb_clr)
    effect_lst.append(outer_shdw)
    el.append(effect_lst)

def create_team_spotlight_slides(
    output_pptx_path: str,
    team_data: list,
    **kwargs
) -> str:
    """
    Creates a PPTX file with layouts for a Team Spotlight Morph effect.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        team_data: A list of dictionaries, where each dict contains 'name' and 'image_url'.

    Returns:
        Path to the saved PPTX file.
        
    **NOTE**: After generating, you must manually apply the 'Morph' transition
    to all slides in PowerPoint for the effect to work.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # --- Download images
    image_paths = []
    temp_dir = "temp_team_images"
    if not os.path.exists(temp_dir):
        os.makedirs(temp_dir)
        
    for i, member in enumerate(team_data):
        try:
            filename = os.path.join(temp_dir, f"member_{i}.jpg")
            urllib.request.urlretrieve(member['image_url'], filename)
            image_paths.append(filename)
        except Exception as e:
            print(f"Could not download image for {member['name']}: {e}")
            image_paths.append(None) # Handle download failure

    # --- Define Layout Constants ---
    SLIDE_WIDTH = prs.slide_width
    SLIDE_HEIGHT = prs.slide_height
    NUM_IMAGES = len(team_data)
    
    # Ratios for spotlight layout
    SPOTLIGHT_WIDTH_RATIO = 0.55
    
    # Calculate widths for the two states
    gallery_width = SLIDE_WIDTH / NUM_IMAGES
    spotlight_width = SLIDE_WIDTH * SPOTLIGHT_WIDTH_RATIO
    preview_width = (SLIDE_WIDTH - spotlight_width) / (NUM_IMAGES - 1)

    def add_pictures_to_slide(slide, widths, image_paths):
        current_left = 0
        for i, path in enumerate(image_paths):
            if path:
                slide.shapes.add_picture(path, current_left, 0, width=widths[i], height=SLIDE_HEIGHT)
            current_left += widths[i]

    # --- Slide 1: Gallery View (Initial State) ---
    slide_layout = prs.slide_layouts[6] # Blank layout
    gallery_slide = prs.slides.add_slide(slide_layout)
    gallery_widths = [gallery_width] * NUM_IMAGES
    add_pictures_to_slide(gallery_slide, gallery_widths, image_paths)

    # --- Slides 2 to N+1: Spotlight Views ---
    for i in range(NUM_IMAGES):
        spotlight_slide = prs.slides.add_slide(slide_layout)
        
        # Define widths for this specific spotlight slide
        spotlight_widths = []
        for j in range(NUM_IMAGES):
            spotlight_widths.append(spotlight_width if i == j else preview_width)
            
        add_pictures_to_slide(spotlight_slide, spotlight_widths, image_paths)
        
        # Add name text box over the spotlighted image
        spotlight_left_position = sum(spotlight_widths[:i])
        
        textbox = spotlight_slide.shapes.add_textbox(
            spotlight_left_position + Inches(0.2), 
            Inches(6.5), 
            spotlight_width - Inches(0.4), 
            Inches(0.75)
        )
        text_frame = textbox.text_frame
        text_frame.word_wrap = False
        text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        
        p = text_frame.paragraphs[0]
        run = p.add_run()
        run.text = team_data[i]['name'].upper()
        
        font = run.font
        font.name = 'Arial Black'
        font.size = Pt(36)
        font.color.rgb = RGBColor(255, 255, 255)
        add_text_shadow(run) # Apply subtle shadow for readability

    # --- Final Slide: Return to Gallery View ---
    final_gallery_slide = prs.slides.add_slide(slide_layout)
    add_pictures_to_slide(final_gallery_slide, gallery_widths, image_paths)

    prs.save(output_pptx_path)
    print(f"Presentation saved to {output_pptx_path}")
    print("\nIMPORTANT: Open the file in PowerPoint and apply the 'Morph' transition to all slides.")
    return output_pptx_path

# --- Example Usage ---
if __name__ == '__main__':
    # Sample data using actors from the video
    marvel_team = [
        {
            "name": "Chris Evans",
            "image_url": "https://i.pinimg.com/564x/f3/7b/3b/f37b3b3e21422a517035f0376aa13214.jpg"
        },
        {
            "name": "Elizabeth Olsen",
            "image_url": "https://i.pinimg.com/564x/0f/5d/28/0f5d28913d358172a39335a013c7bd1b.jpg"
        },
        {
            "name": "Robert Downey Jr",
            "image_url": "https://i.pinimg.com/564x/72/7a/12/727a124f33bf131345fb19a8618e974e.jpg"
        },
        {
            "name": "Zoe Saldana",
            "image_url": "https://i.pinimg.com/564x/d5/a2/33/d5a23348005b790a8862f1a300d23589.jpg"
        }
    ]

    output_file = "team_spotlight_presentation.pptx"
    create_team_spotlight_slides(output_file, marvel_team)

