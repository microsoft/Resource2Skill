import os
import random
import urllib.request
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "大力发展",
    subtitle_text: str = "特色乡村文化旅游",
    english_label: str = "Featured rural cultural tourism",
    body_text: str = "Pellentesque habitant morbi tristique senectus et netus et malesuada fames ac turpis egestas. Proin pharetra nonummy pede. Mauris et orci.",
    image_url: str = "https://picsum.photos/seed/tea_plantation/1200/1080", # Landscape/Nature placeholder
    accent_color: tuple = (46, 139, 87),  # Sea Green
    grid_rows: int = 6,
    grid_cols: int = 6,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Mosaic Grid Image Reveal' visual effect.
    """
    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide
    
    # Ensure background is white
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # 2. Dimensions for the image side (~60% of slide width)
    # PPTX resolution is conceptually 13.333 x 7.5 inches.
    # We will generate a high-res PIL image at 144 DPI for crispness.
    dpi = 144
    img_width_in = 8.0 # 13.333 - 5.333 (leaving ~5.3 inches for text)
    img_height_in = 7.5
    
    px_width = int(img_width_in * dpi)
    px_height = int(img_height_in * dpi)

    # 3. Fetch Base Image
    try:
        req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            base_img = Image.open(BytesIO(response.read())).convert("RGBA")
    except Exception as e:
        print(f"Image download failed ({e}), generating gradient fallback.")
        base_img = Image.new("RGBA", (px_width, px_height))
        draw = ImageDraw.Draw(base_img)
        for y in range(px_height):
            r = int(accent_color[0] * (1 - y/px_height))
            g = int(accent_color[1] * (1 - y/px_height))
            b = int(accent_color[2] * (1 - y/px_height))
            draw.line([(0, y), (px_width, y)], fill=(r, g, b, 255))

    # Center crop and resize image to exact dimensions
    img_ratio = base_img.width / base_img.height
    target_ratio = px_width / px_height
    
    if img_ratio > target_ratio:
        # Image is wider, crop width
        new_w = int(target_ratio * base_img.height)
        offset = (base_img.width - new_w) // 2
        base_img = base_img.crop((offset, 0, offset + new_w, base_img.height))
    else:
        # Image is taller, crop height
        new_h = int(base_img.width / target_ratio)
        offset = (base_img.height - new_h) // 2
        base_img = base_img.crop((0, offset, base_img.width, offset + new_h))
        
    base_img = base_img.resize((px_width, px_height), Image.Resampling.LANCZOS)

    # 4. Generate the Mosaic Grid Overlay using PIL
    overlay = Image.new("RGBA", (px_width, px_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    
    cell_w = px_width / grid_cols
    cell_h = px_height / grid_rows
    
    # Determine random cells for effects
    total_cells = grid_rows * grid_cols
    all_indices = list(range(total_cells))
    random.shuffle(all_indices)
    
    # ~15% completely missing (white background shows through)
    num_missing = int(total_cells * 0.15) 
    # ~25% semi-transparent white
    num_faded = int(total_cells * 0.25)   
    
    missing_cells = set(all_indices[:num_missing])
    faded_cells = set(all_indices[num_missing:num_missing + num_faded])
    
    for r in range(grid_rows):
        for c in range(grid_cols):
            idx = r * grid_cols + c
            x0, y0 = c * cell_w, r * cell_h
            x1, y1 = x0 + cell_w, y0 + cell_h
            
            if idx in missing_cells:
                # Solid white to simulate 'missing' cell showing slide background
                draw.rectangle([x0, y0, x1, y1], fill=(255, 255, 255, 255))
            elif idx in faded_cells:
                # Semi-transparent white
                # Randomize opacity slightly between 100 and 180 for texture
                alpha = random.randint(100, 180)
                draw.rectangle([x0, y0, x1, y1], fill=(255, 255, 255, alpha))

    # Draw grid lines (white, 2px)
    for r in range(1, grid_rows):
        y = r * cell_h
        draw.line([(0, y), (px_width, y)], fill=(255, 255, 255, 255), width=2)
    for c in range(1, grid_cols):
        x = c * cell_w
        draw.line([(x, 0), (x, px_height)], fill=(255, 255, 255, 255), width=2)

    # Composite base image with overlay
    final_img = Image.alpha_composite(base_img, overlay)
    
    # Save to buffer
    img_buffer = BytesIO()
    final_img.save(img_buffer, format="PNG")
    img_buffer.seek(0)

    # 5. Insert Image into Slide
    slide.shapes.add_picture(
        img_buffer,
        Inches(5.333), Inches(0), # Positioned on the right
        width=Inches(img_width_in), height=Inches(img_height_in)
    )

    # 6. Add Text Elements (Left Side)
    # Icon placeholder (green leaf concept)
    icon_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(1.0), Inches(0.5))
    tf_icon = icon_box.text_frame
    tf_icon.text = "🌿" # Using emoji as quick shape placeholder
    tf_icon.paragraphs[0].font.size = Pt(36)

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(4.0), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(50, 50, 50)
    p_title.alignment = PP_ALIGN.LEFT

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(4.0), Inches(0.6))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(28)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(*accent_color)
    p_sub.alignment = PP_ALIGN.LEFT

    # English Label
    eng_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(4.0), Inches(0.5))
    tf_eng = eng_box.text_frame
    tf_eng.word_wrap = True
    p_eng = tf_eng.paragraphs[0]
    p_eng.text = english_label.upper()
    p_eng.font.size = Pt(12)
    p_eng.font.bold = True
    p_eng.font.color.rgb = RGBColor(100, 100, 100)

    # Body Text
    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(3.5), Inches(1.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = RGBColor(120, 120, 120)
    p_body.line_spacing = 1.3

    # Save Presentation
    prs.save(output_pptx_path)
    return output_pptx_path
