import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN
from lxml import etree

def create_slide(
    output_pptx_path: str,
    title_text: str = "TABLE OF CONTENTS",
    list_items: list = None,
    bg_color: tuple = (48, 84, 150),
    accent_color: tuple = (255, 192, 0),
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with an 'Asymmetric Soft-UI List' design.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        title_text: The main title for the central circle.
        list_items: A list of strings for the list items. Defaults to a sample list.
        bg_color: RGB tuple for the slide background.
        accent_color: RGB tuple for the yellow accent color.

    Returns:
        Path to the saved PPTX file.
    """

    if list_items is None:
        list_items = ["Text Here", "Text Here", "Text Here", "Text Here"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Helper function for adding shadow via lxml ---
    def add_shadow_effect(shape, blur_radius=15, distance=3, direction=45, alpha=50):
        """Adds a soft outer shadow effect to a shape."""
        shape_element = shape.element
        spPr = shape_element.spPr
        
        # Create effect list if it doesn't exist
        effect_list = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        if effect_list is None:
            effect_list = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")

        # Define the outer shadow effect
        outer_shadow = etree.SubElement(effect_list, "{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw")
        outer_shadow.set("blurRad", str(Emu(Pt(blur_radius))))
        outer_shadow.set("dist", str(Emu(Pt(distance))))
        outer_shadow.set("dir", str(int(direction * 60000)))
        outer_shadow.set("algn", "bl") # Bottom-left alignment
        
        # Define shadow color
        shadow_color = etree.SubElement(outer_shadow, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        shadow_color.set("val", "000000") # Black shadow
        alpha_element = etree.SubElement(shadow_color, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
        alpha_element.set("val", str(alpha * 1000)) # Alpha is in 1000ths of a percent
        
    # === Layer 1: Background ===
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(*bg_color)
    background.line.fill.background()

    # === Layer 2: Core Visual Elements ===
    # --- Left Side: Title Crescent ---
    circ_diameter = Inches(3.5)
    y_center = (prs.slide_height - circ_diameter) / 2
    
    # Back yellow circle
    back_circle_shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1), y_center - Inches(0.1), circ_diameter, circ_diameter
    )
    back_circle_shape.fill.solid()
    back_circle_shape.fill.fore_color.rgb = RGBColor(*accent_color)
    back_circle_shape.line.fill.background()
    add_shadow_effect(back_circle_shape, blur_radius=20, distance=5, direction=45, alpha=40)

    # Front white circle for title
    front_circ_diameter = Inches(3.0)
    front_circle_shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(0.75), y_center + Inches(0.15), front_circ_diameter, front_circ_diameter
    )
    front_circle_shape.fill.solid()
    front_circle_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    front_circle_shape.line.fill.background()
    add_shadow_effect(front_circle_shape, blur_radius=15, distance=3, direction=45, alpha=30)
    
    # Add title text to the front circle
    tf = front_circle_shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = 'middle'
    font = run.font
    font.name = 'Arial Black'
    font.size = Pt(24)
    font.bold = True
    font.color.rgb = RGBColor(64, 64, 64)


    # --- Right Side: List Items ---
    item_height = Inches(0.9)
    item_width = Inches(6)
    v_spacing = Inches(0.3)
    total_list_height = (item_height * len(list_items)) + (v_spacing * (len(list_items) - 1))
    start_y = (prs.slide_height - total_list_height) / 2
    start_x = Inches(6)

    for i, item_text in enumerate(list_items):
        current_y = start_y + i * (item_height + v_spacing)
        
        # Base rounded rectangle (yellow)
        base_item = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, start_x, current_y, item_width, item_height
        )
        base_item.adjustments[0] = 0.5  # Max rounding
        base_item.fill.solid()
        base_item.fill.fore_color.rgb = RGBColor(*accent_color)
        base_item.line.fill.background()
        add_shadow_effect(base_item, blur_radius=12, distance=4, direction=45, alpha=35)

        # White content area
        content_width = item_width - Inches(1.1)
        content_height = item_height - Inches(0.2)
        content_y = current_y + Inches(0.1)
        content_x = start_x + Inches(1.0)
        content_area = slide.shapes.add_shape(
             MSO_SHAPE.ROUNDED_RECTANGLE, content_x, content_y, content_width, content_height
        )
        content_area.adjustments[0] = 0.5
        content_area.fill.solid()
        content_area.fill.fore_color.rgb = RGBColor(255, 255, 255)
        content_area.line.fill.background()
        
        # Add text to content area
        tf_content = content_area.text_frame
        tf_content.clear()
        p_content = tf_content.paragraphs[0]
        run_content = p_content.add_run()
        run_content.text = item_text
        tf_content.vertical_anchor = 'middle'
        font_content = run_content.font
        font_content.name = 'Arial'
        font_content.size = Pt(18)
        font_content.color.rgb = RGBColor(64, 64, 64)

        # Bullet circle (white)
        bullet_diameter = item_height - Inches(0.1)
        bullet_x = start_x + Inches(0.05)
        bullet_y = current_y + Inches(0.05)
        bullet_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, bullet_x, bullet_y, bullet_diameter, bullet_diameter
        )
        bullet_circle.fill.solid()
        bullet_circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
        bullet_circle.line.fill.background()
        add_shadow_effect(bullet_circle, blur_radius=10, distance=2, direction=45, alpha=30)
        
        # Add letter to bullet
        tf_bullet = bullet_circle.text_frame
        tf_bullet.clear()
        p_bullet = tf_bullet.paragraphs[0]
        run_bullet = p_bullet.add_run()
        run_bullet.text = chr(ord('A') + i)
        p_bullet.alignment = PP_ALIGN.CENTER
        tf_bullet.vertical_anchor = 'middle'
        font_bullet = run_bullet.font
        font_bullet.name = 'Arial Black'
        font_bullet.size = Pt(22)
        font_bullet.color.rgb = RGBColor(64, 64, 64)
        
        # Connector Line
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW,
            back_circle_shape.left + back_circle_shape.width,
            back_circle_shape.top + back_circle_shape.height / 2,
            base_item.left,
            base_item.top + base_item.height/2
        )
        line = connector.line
        line.color.rgb = RGBColor(255, 255, 255)
        line.width = Pt(1.5)
        
        # Move connector to back
        connector_xml = connector.element
        parent = connector_xml.getparent()
        parent.insert(0, connector_xml)


    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# if __name__ == '__main__':
#     file_path = "Asymmetric_Soft_UI_List.pptx"
#     create_slide(file_path, list_items=["Introduction", "Methodology", "Results", "Conclusion"])
#     if os.path.exists(file_path):
#         print(f"Presentation saved to {os.path.abspath(file_path)}")
#         # os.startfile(os.path.abspath(file_path)) # Uncomment to open the file on Windows
