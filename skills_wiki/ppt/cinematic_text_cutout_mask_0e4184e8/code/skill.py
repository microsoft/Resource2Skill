def create_slide(
    output_pptx_path: str,
    title_text: str = "THANK YOU",
    subtitle_text: str = "Do you have Any Question?",
    bg_theme_1: str = "landscape,nature",
    bg_theme_2: str = "japan,architecture",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Cinematic Text Cutout Mask' visual effect.
    """
    import os
    import urllib.request
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Helper function to download images
    def download_image(url, filename):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req) as response, open(filename, 'wb') as out_file:
                out_file.write(response.read())
            return filename
        except Exception as e:
            print(f"Failed to download {url}: {e}")
            # Create a fallback colored image
            img = Image.new('RGB', (800, 600), color=(100, 150, 200))
            img.save(filename)
            return filename

    img1_path = download_image(f"https://source.unsplash.com/featured/1280x720/?{bg_theme_1}", "bg1.jpg")
    img2_path = download_image(f"https://source.unsplash.com/featured/1280x720/?{bg_theme_2}", "bg2.jpg")

    # === Layer 1: Background Images ===
    # Place images side by side so they span across the back of the mask
    slide.shapes.add_picture(img1_path, Inches(0), Inches(0), width=Inches(8), height=Inches(7.5))
    slide.shapes.add_picture(img2_path, Inches(8), Inches(0), width=Inches(8), height=Inches(7.5))

    # === Layer 2: Text Cutout Mask using PIL ===
    mask_width, mask_height = int(13.333 * 96), int(7.5 * 96) # Standard 96 DPI
    mask_rgba = Image.new('RGBA', (mask_width, mask_height), (228, 228, 228, 255)) # Light Gray Overlay
    alpha_layer = Image.new('L', (mask_width, mask_height), 255) # 255 = solid, 0 = transparent hole
    draw = ImageDraw.Draw(alpha_layer)

    # Attempt to load a thick block font, fallback gracefully
    font = None
    for f_name in ["impact.ttf", "arialbd.ttf", "tahoma.ttf", "DejaVuSans-Bold.ttf"]:
        try:
            font = ImageFont.truetype(f_name, 280)
            break
        except OSError:
            continue
    if not font:
        font = ImageFont.load_default()

    # Calculate text bounding box to center it
    bbox = draw.textbbox((0, 0), title_text, font=font)
    t_w, t_h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    
    # Scale down if text is too wide for slide
    if t_w > mask_width * 0.95:
        try:
            scaled_size = int(280 * (mask_width * 0.9 / t_w))
            font = ImageFont.truetype(font.path, scaled_size)
            bbox = draw.textbbox((0, 0), title_text, font=font)
            t_w, t_h = bbox[2] - bbox[0], bbox[3] - bbox[1]
        except:
            pass

    # Position text slightly above true center
    x = (mask_width - t_w) / 2
    y = (mask_height - t_h) / 2 - 80

    # Draw text in black (0) on the alpha layer to punch the hole
    draw.text((x, y), title_text, fill=0, font=font)
    mask_rgba.putalpha(alpha_layer)
    mask_path = "cutout_mask.png"
    mask_rgba.save(mask_path)

    # Insert the mask exactly over the entire slide
    slide.shapes.add_picture(mask_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

    # === Layer 3: Accent Line and Subtitle ===
    
    # Add vertical divider line
    line_x = Inches(10)
    line_y = Inches(1.5)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_x, line_y, width=Inches(0.06), height=Inches(4.5))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(30, 30, 30) # Very dark gray/black
    line.line.fill.background() # Remove outline

    # Add subtitle with wide letter tracking (achieved via spacing interpolation)
    # E.g., "TEXT" -> "T  E  X  T"
    spaced_subtitle = "   ".join(list(subtitle_text)).replace("      ", "   ")
    
    tb = slide.shapes.add_textbox(Inches(0), Inches(5.2), Inches(13.333), Inches(1))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = spaced_subtitle
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(30, 30, 30)

    # Clean up local temporary files
    prs.save(output_pptx_path)
    for tmp_file in [img1_path, img2_path, mask_path]:
        if os.path.exists(tmp_file):
            try:
                os.remove(tmp_file)
            except:
                pass

    return output_pptx_path
