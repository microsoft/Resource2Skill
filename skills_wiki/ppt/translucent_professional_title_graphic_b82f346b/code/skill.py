import io
import random
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont

def create_moxie_title_slide(
    output_pptx_path: str,
    tip_number: int = 1,
    tip_text: str = "PREPARE",
    speaker_image_url: str = "https://images.pexels.com/photos/415829/pexels-photo-415829.jpeg?auto=compress&cs=tinysrgb&w=1260&h=750&dpr=1",
) -> str:
    """
    Creates a PPTX file with a title slide reproducing the "Moxie Talk" visual style.

    This style features a corporate blue background with a subtle bar chart motif,
    a translucent speaker image, and clean, high-contrast text.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        tip_number: The number of the tip to display.
        tip_text: The main text of the tip.
        speaker_image_url: URL for a portrait image (ideally with a simple background).

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- 1. Create Background with PIL ---
    SLIDE_W_PX, SLIDE_H_PX = 1280, 720
    BG_COLOR = (37, 105, 163)
    BAR_COLOR = (89, 148, 195, 80) # RGBA with alpha for transparency
    ACCENT_GREEN = (186, 218, 85)

    # Create the base image with solid blue
    bg_image = Image.new("RGBA", (SLIDE_W_PX, SLIDE_H_PX), BG_COLOR)
    draw = ImageDraw.Draw(bg_image)

    # Draw the subtle bar chart motif
    num_bars = 12
    bar_width = SLIDE_W_PX / (num_bars * 1.5)
    for i in range(num_bars):
        bar_height = random.randint(int(SLIDE_H_PX * 0.2), int(SLIDE_H_PX * 0.7))
        x0 = i * (bar_width * 1.5) + (bar_width * 0.25)
        y0 = SLIDE_H_PX - bar_height
        x1 = x0 + bar_width
        y1 = SLIDE_H_PX
        draw.rectangle([x0, y0, x1, y1], fill=BAR_COLOR)

    # --- 2. Add Translucent Speaker Image ---
    try:
        with urllib.request.urlopen(speaker_image_url) as url:
            f = io.BytesIO(url.read())
        speaker_img = Image.open(f).convert("RGBA")

        # Create a mask for opacity
        alpha = speaker_img.getchannel('A')
        new_alpha = alpha.point(lambda p: p * 0.3) # 30% opacity
        speaker_img.putalpha(new_alpha)

        # Resize and position the speaker image
        base_width = int(SLIDE_W_PX * 0.4)
        w_percent = (base_width / float(speaker_img.size[0]))
        h_size = int((float(speaker_img.size[1]) * float(w_percent)))
        speaker_img = speaker_img.resize((base_width, h_size), Image.LANCZOS)
        
        # Paste onto the background
        paste_x = SLIDE_W_PX - base_width
        paste_y = (SLIDE_H_PX - h_size) // 2
        bg_image.paste(speaker_img, (paste_x, paste_y), speaker_img)

    except Exception as e:
        print(f"Could not download or process speaker image: {e}. Skipping.")

    # Save the composite background to a memory buffer
    image_stream = io.BytesIO()
    bg_image.save(image_stream, format="PNG")
    image_stream.seek(0)

    # Add the background image to the slide
    slide.shapes.add_picture(image_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # --- 3. Add Text Elements ---
    # MOXIE
    left, top, width, height = Inches(1), Inches(1.5), Inches(5), Inches(2)
    tb = slide.shapes.add_textbox(left, top, width, height)
    p = tb.text_frame.paragraphs[0]
    p.text = "MOXIE"
    p.font.name = "Arial Black"
    p.font.size = Pt(100)
    p.font.color.rgb = RGBColor(255, 255, 255)

    # talk
    left, top, width, height = Inches(5.3), Inches(2.2), Inches(3), Inches(1)
    tb_talk = slide.shapes.add_textbox(left, top, width, height)
    p_talk = tb_talk.text_frame.paragraphs[0]
    p_talk.text = "talk"
    # A common cursive/script font; change if not available
    p_talk.font.name = "Brush Script MT" 
    p_talk.font.size = Pt(80)
    p_talk.font.color.rgb = RGBColor(*ACCENT_GREEN)
    
    # TIP Text
    left, top, width, height = Inches(0), Inches(5.5), prs.slide_width, Inches(1.5)
    tb_tip = slide.shapes.add_textbox(left, top, width, height)
    p_tip = tb_tip.text_frame.paragraphs[0]
    p_tip.text = f"TIP {tip_number}:\n{tip_text.upper()}"
    p_tip.font.name = "Arial"
    p_tip.font.bold = True
    p_tip.font.size = Pt(44)
    p_tip.font.color.rgb = RGBColor(255, 255, 255)
    
    # Moxie Institute Branding (simplified)
    left, top, width, height = Inches(10.5), Inches(0.5), Inches(2.5), Inches(0.5)
    tb_brand = slide.shapes.add_textbox(left, top, width, height)
    p_brand = tb_brand.text_frame.paragraphs[0]
    p_brand.text = "moxie INSTITUTE"
    p_brand.font.name = "Arial"
    p_brand.font.size = Pt(16)
    p_brand.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_moxie_title_slide("moxie_slide.pptx", tip_number=1, tip_text="PREPARE")
