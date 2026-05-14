def create_slide(
    output_pptx_path: str,
    title_text: str = "The Secret To\nExceptional\nCustomer Service",
    subtitle_text: str = "MrDogBrain Productions Inc.",
    credit_text: str = "Why Are You Reading This? LLC (2021)",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Cinematic Minimalist Title Card" visual effect.
    This script generates two slides: an opening title card and an end credit card,
    mimicking the flow of the provided video tutorial.

    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import os

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 Widescreen
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    blank_slide_layout = prs.slide_layouts[6]

    def set_black_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 0, 0)

    # ==========================================
    # SLIDE 1: Opening Cinematic Title
    # ==========================================
    slide_1 = prs.slides.add_slide(blank_slide_layout)
    set_black_background(slide_1)

    # Title Text Box
    # Positioned slightly above absolute center
    title_left = Inches(1.66)
    title_top = Inches(2.0)
    title_width = Inches(10.0)
    title_height = Inches(2.5)
    
    txBox = slide_1.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Calibri'
    p.font.size = Pt(44)
    # Using a muted grey, not stark white, to match the video's cinematic feel
    p.font.color.rgb = RGBColor(150, 150, 150) 

    # Subtitle Text Box
    # Positioned lower down, creating macro-whitespace
    if subtitle_text:
        sub_left = Inches(1.66)
        sub_top = Inches(5.5)
        sub_width = Inches(10.0)
        sub_height = Inches(1.0)
        
        sub_txBox = slide_1.shapes.add_textbox(sub_left, sub_top, sub_width, sub_height)
        sub_tf = sub_txBox.text_frame
        
        sub_p = sub_tf.paragraphs[0]
        sub_p.text = subtitle_text
        sub_p.alignment = PP_ALIGN.CENTER
        sub_p.font.name = 'Calibri'
        sub_p.font.size = Pt(24)
        sub_p.font.color.rgb = RGBColor(180, 180, 180)

    # ==========================================
    # SLIDE 2: End Credit Card
    # ==========================================
    slide_2 = prs.slides.add_slide(blank_slide_layout)
    set_black_background(slide_2)

    # Single small text in dead center
    credit_left = Inches(1.66)
    credit_top = Inches(3.25)
    credit_width = Inches(10.0)
    credit_height = Inches(1.0)
    
    c_txBox = slide_2.shapes.add_textbox(credit_left, credit_top, credit_width, credit_height)
    c_tf = c_txBox.text_frame
    
    c_p = c_tf.paragraphs[0]
    c_p.text = credit_text
    c_p.alignment = PP_ALIGN.CENTER
    c_p.font.name = 'Calibri'
    c_p.font.size = Pt(18)
    c_p.font.color.rgb = RGBColor(200, 200, 200)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_slide("cinematic_minimalist_cards.pptx")
