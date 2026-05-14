def create_slide(
    output_pptx_path: str,
    title_text: str = "Our\nMission",
    bg_image_url: str = "https://images.unsplash.com/photo-1486406146926-c627a92ad1ab?q=80&w=1200&auto=format&fit=crop",
    accent_color: tuple = (125, 181, 165),  # Mint Teal
    text_color: tuple = (38, 38, 38),       # Dark Charcoal
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Split-Pane B&W Hero with Oversized Accent Circle' style.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    import urllib.request
    from PIL import Image, ImageOps
    from io import BytesIO
    import tempfile
    import os

    # Initialize Widescreen Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Image (Processed via PIL for Grayscale) ===
    img_width_in = 7.0
    img_height_in = 7.5
    dpi = 150
    px_w = int(img_width_in * dpi)
    px_h = int(img_height_in * dpi)

    # Securely create a temporary file path
    fd, temp_img_path = tempfile.mkstemp(suffix=".jpg")
    os.close(fd)

    try:
        # Download, crop, and desaturate image
        req = urllib.request.Request(bg_image_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            img = Image.open(BytesIO(response.read()))
            img = ImageOps.fit(img, (px_w, px_h), Image.Resampling.LANCZOS)
            img = img.convert('L') # Convert to pure Grayscale
            img.save(temp_img_path, format="JPEG", quality=95)
    except Exception:
        # Fallback to a solid soft gray if network fails
        img = Image.new('RGB', (px_w, px_h), color=(230, 230, 230))
        img.save(temp_img_path, format="JPEG")

    # Insert processed image to cover left portion
    slide.shapes.add_picture(temp_img_path, 0, 0, Inches(img_width_in), Inches(img_height_in))
    os.remove(temp_img_path)

    # === Layer 2: The Oversized Hero Circle ===
    circle_dia = 4.5
    # Center the circle perfectly on the right edge of the image
    circle_x = img_width_in - (circle_dia / 2) 
    circle_y = (7.5 - circle_dia) / 2

    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(circle_x),
        Inches(circle_y),
        Inches(circle_dia),
        Inches(circle_dia)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(*accent_color)
    circle.line.fill.background() # Remove border

    # Configure centered title text inside the circle
    tf = circle.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = title_text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*text_color)
    p.font.name = "Arial"

    # === Layer 3: Right Side Content (Vertical List) ===
    # Start positioning elements comfortably to the right of the circle
    start_x = img_width_in + (circle_dia / 2) + 0.3 
    start_y = 1.6
    spacing = 1.6

    list_items = [
        ("Mission", "Adapt this section to highlight your core strategic mission and underlying principles."),
        ("Vision", "Outline your future goals and what the company aims to become in the coming years."),
        ("Goal", "Define the specific, measurable steps you will take to achieve your stated vision.")
    ]

    for i, (item_title, item_desc) in enumerate(list_items):
        y_offset = start_y + (i * spacing)

        # Small dark geometric marker to replace icons
        marker_size = 0.4
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(start_x),
            Inches(y_offset + 0.05),
            Inches(marker_size),
            Inches(marker_size)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(*text_color)
        marker.line.fill.background()

        # Item Title
        tx_title = slide.shapes.add_textbox(
            Inches(start_x + 0.6),
            Inches(y_offset),
            Inches(3.5),
            Inches(0.4)
        )
        pt = tx_title.text_frame.paragraphs[0]
        pt.text = item_title
        pt.font.size = Pt(22)
        pt.font.bold = True
        pt.font.color.rgb = RGBColor(*text_color)
        pt.font.name = "Arial"

        # Item Description
        tx_desc = slide.shapes.add_textbox(
            Inches(start_x + 0.6),
            Inches(y_offset + 0.4),
            Inches(3.5),
            Inches(1.0)
        )
        tx_desc.text_frame.word_wrap = True
        pd = tx_desc.text_frame.paragraphs[0]
        pd.text = item_desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = RGBColor(100, 100, 100)
        pd.font.name = "Arial"

    prs.save(output_pptx_path)
    return output_pptx_path
