from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.shapes.freeform import FreeformBuilder

def create_slide(
    output_pptx_path: str,
    title_text: str = "PDF 转换工具",
    subtitle_text: str = "WWW.ILOVEPDF.COM",
    bg_color: tuple = (211, 166, 106),
    bar_color_top: tuple = (255, 255, 255),
    bar_color_bottom: tuple = (15, 15, 15),
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a modern "Offset Bar Title Graphic".

    Args:
        output_pptx_path: Path to save the generated PPTX file.
        title_text: The main title for the top bar.
        subtitle_text: The subtitle or URL for the bottom bar.
        bg_color: RGB tuple for the slide background.
        bar_color_top: RGB tuple for the top bar (and bottom tag).
        bar_color_bottom: RGB tuple for the bottom bar (and top tag).

    Returns:
        Path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Visual Elements & Text ===
    # Define overall dimensions and position for the graphic
    graphic_width = Inches(8.5)
    bar_height = Inches(0.8)
    total_graphic_height = bar_height * 2
    offset = Inches(0.12)

    # Center the entire graphic on the slide
    center_x = prs.slide_width / 2
    center_y = prs.slide_height / 2
    
    # Calculate positions for the two bars
    top_bar_left = center_x - (graphic_width / 2)
    top_bar_top = center_y - (total_graphic_height / 2)
    
    bottom_bar_left = top_bar_left + offset
    bottom_bar_top = top_bar_top + offset

    # --- Bottom Bar (created first to be in the back) ---
    bottom_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        bottom_bar_left,
        bottom_bar_top + bar_height,
        graphic_width,
        bar_height
    )
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = RGBColor(*bar_color_bottom)
    bottom_bar.line.fill.background()

    # Add subtitle text to the bottom bar
    tf_bottom = bottom_bar.text_frame
    tf_bottom.text = subtitle_text
    tf_bottom.paragraphs[0].font.name = 'Helvetica'
    tf_bottom.paragraphs[0].font.size = Pt(24)
    tf_bottom.paragraphs[0].font.color.rgb = RGBColor(*bar_color_top)
    tf_bottom.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_bottom.vertical_anchor = 'middle'
    tf_bottom.margin_bottom = Inches(0)
    tf_bottom.margin_top = Inches(0)


    # --- Top Bar (created second to be in the front) ---
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        top_bar_left,
        top_bar_top,
        graphic_width,
        bar_height
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = RGBColor(*bar_color_top)
    top_bar.line.fill.background()
    
    # Add title text to the top bar
    tf_top = top_bar.text_frame
    tf_top.text = title_text
    tf_top.paragraphs[0].font.name = 'Helvetica Bold'
    tf_top.paragraphs[0].font.bold = True
    tf_top.paragraphs[0].font.size = Pt(28)
    tf_top.paragraphs[0].font.color.rgb = RGBColor(*bar_color_bottom)
    tf_top.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_top.vertical_anchor = 'middle'
    tf_top.margin_bottom = Inches(0)
    tf_top.margin_top = Inches(0)

    # --- Custom Tags using FreeformBuilder ---
    tag_width = Inches(0.2)
    
    # Tag for the top bar (dark color)
    with FreeformBuilder(
        slide.shapes,
        top_bar.left + graphic_width,  # Start X
        top_bar.top,                   # Start Y
    ) as builder:
        builder.add_line_segments([(tag_width, 0)])
        builder.add_line_segments([(0, bar_height)])
        builder.add_line_segments([(-tag_width, 0)])
        builder.close()
    
    tag_top_shape = builder.shape
    tag_top_shape.fill.solid()
    tag_top_shape.fill.fore_color.rgb = RGBColor(*bar_color_bottom)
    tag_top_shape.line.fill.background()

    # Tag for the bottom bar (light color)
    with FreeformBuilder(
        slide.shapes,
        bottom_bar.left + graphic_width,
        bottom_bar.top,
    ) as builder:
        builder.add_line_segments([(tag_width, 0)])
        builder.add_line_segments([(0, bar_height)])
        builder.add_line_segments([(-tag_width, 0)])
        builder.close()

    tag_bottom_shape = builder.shape
    tag_bottom_shape.fill.solid()
    tag_bottom_shape.fill.fore_color.rgb = RGBColor(*bar_color_top)
    tag_bottom_shape.line.fill.background()

    prs.save(output_pptx_path)
    return output_pptx_path
