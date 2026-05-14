def create_slide(
    output_pptx_path: str,
    title_text: str = "MODERN CORPORATE\nTEMPLATE STYLE",
    body_text: str = "This layout uses a geometric pattern overlay, bold angular blocks, and a high-contrast accent color to create a modern corporate aesthetic.\n\nThe off-center angled division adds forward momentum while preserving a clean, highly readable content area on the left.",
    bg_palette: str = "corporate",
    accent_color: tuple = (211, 75, 50),  # Burnt Orange
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Corporate Geometric Angular Template' visual effect.
    """
    import os
    import math
    import tempfile
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Geometric Mesh (via PIL) ===
    # Generate a subtle hexagonal tile pattern
    bg_width, bg_height = 1920, 1080
    img = Image.new('RGBA', (bg_width, bg_height), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)

    r = 45  # Hexagon radius
    dx = r * math.sqrt(3)
    dy = r * 1.5

    for row in range(int(bg_height / dy) + 2):
        for col in range(int(bg_width / dx) + 2):
            x = col * dx
            if row % 2 == 1:
                x += dx / 2
            y = row * dy
            
            pts = []
            for i in range(6):
                angle_deg = 60 * i + 30
                angle_rad = math.pi / 180 * angle_deg
                px = x + r * math.cos(angle_rad)
                py = y + r * math.sin(angle_rad)
                pts.append((px, py))
            draw.polygon(pts, outline=(242, 242, 242, 255), fill=None)

    bg_path = os.path.join(tempfile.gettempdir(), "hex_bg.png")
    img.save(bg_path)
    
    # Insert the generated mesh as the slide background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Angular Accent Shape (via FreeformBuilder) ===
    # Draw a massive trapezoid on the right edge
    start_x, start_y = Inches(9.5), Inches(0)
    ff_builder = slide.shapes.build_freeform(start_x, start_y)
    ff_builder.add_line_segments([
        (prs.slide_width, Inches(0)),
        (prs.slide_width, prs.slide_height),
        (Inches(7.5), prs.slide_height),
        (start_x, start_y)
    ])
    angled_shape = ff_builder.convert_to_shape()
    angled_shape.fill.solid()
    angled_shape.fill.fore_color.rgb = RGBColor(*accent_color)
    angled_shape.line.fill.background()  # Remove default outline

    # === Layer 3: Typography and Structure ===
    # Title Text
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(6.0), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p_title = title_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(60, 60, 60)
    p_title.alignment = PP_ALIGN.LEFT

    # Accent Divider Line
    divider = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(1.0), Inches(3.2), Inches(1.5), Pt(4)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(*accent_color)
    divider.line.fill.background()

    # Body Text
    body_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(5.5), Inches(3.0))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    p_body = body_frame.paragraphs[0]
    p_body.text = body_text
    p_body.font.name = "Arial"
    p_body.font.size = Pt(18)
    p_body.font.color.rgb = RGBColor(100, 100, 100)
    p_body.alignment = PP_ALIGN.LEFT
    
    # Placeholder Prompt Text (Simulating template behavior)
    prompt_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(5.0), Inches(0.5))
    prompt_frame = prompt_box.text_frame
    p_prompt = prompt_frame.paragraphs[0]
    p_prompt.text = "[Type custom bullet points here]"
    p_prompt.font.name = "Arial"
    p_prompt.font.size = Pt(14)
    p_prompt.font.italic = True
    p_prompt.font.color.rgb = RGBColor(*accent_color)

    # Save presentation
    prs.save(output_pptx_path)
    
    # Cleanup temporary image
    try:
        os.remove(bg_path)
    except OSError:
        pass
        
    return output_pptx_path
