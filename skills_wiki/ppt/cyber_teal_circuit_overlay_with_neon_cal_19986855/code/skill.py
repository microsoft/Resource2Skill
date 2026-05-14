import os
import random
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    title_text: str = "MY FIRST PRESENTATION",
    subtitle_text: str = "I WILL DO MY BEST HERE",
    callout_text: str = "Check HERE!!",
    bg_color_start: tuple = (11, 49, 66),  # Deep Navy
    bg_color_end: tuple = (15, 82, 87),    # Cyber Teal
    accent_color: tuple = (140, 210, 110), # Neon Green
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Cyber-Teal Circuit Overlay with Neon Callout' effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Generation via PIL ===
    bg_path = "temp_tech_bg.png"
    width, height = 1920, 1080
    base_img = Image.new('RGBA', (width, height), bg_color_start + (255,))
    draw = ImageDraw.Draw(base_img)

    # 1. Draw Vertical Gradient
    for y in range(height):
        r = int(bg_color_start[0] + (bg_color_end[0] - bg_color_start[0]) * (y / height))
        g = int(bg_color_start[1] + (bg_color_end[1] - bg_color_start[1]) * (y / height))
        b = int(bg_color_start[2] + (bg_color_end[2] - bg_color_start[2]) * (y / height))
        draw.line([(0, y), (width, y)], fill=(r, g, b, 255))

    # 2. Draw Procedural Circuit Overlay (Anchored to the Left)
    circuit_layer = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    c_draw = ImageDraw.Draw(circuit_layer)
    random.seed(42)  # Fixed seed for reproducibility

    for _ in range(35):
        # Start points heavily weighted to the left 30% of the screen
        x1 = random.randint(0, int(width * 0.3))
        y1 = random.randint(0, height)
        
        # Move horizontally to the right
        x2 = x1 + random.randint(50, 300)
        y2 = y1
        
        # Branch diagonally
        direction = random.choice([-1, 1])
        x3 = x2 + random.randint(50, 150)
        y3 = y2 + (random.randint(50, 150) * direction)
        
        # Draw lines
        c_draw.line([(x1, y1), (x2, y2), (x3, y3)], fill=(0, 255, 255, 25), width=3)
        
        # Draw node (terminal circle)
        radius = random.choice([4, 6, 8])
        c_draw.ellipse([x3 - radius, y3 - radius, x3 + radius, y3 + radius], 
                       outline=(0, 255, 255, 50), width=2)
        if random.random() > 0.5:
             c_draw.ellipse([x3 - 2, y3 - 2, x3 + 2, y3 + 2], fill=(0, 255, 255, 80))

    # Composite and save background
    final_bg = Image.alpha_composite(base_img, circuit_layer)
    final_bg.save(bg_path)

    # Insert as slide background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Typography ===
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(3.5), Inches(2.5), Inches(7.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text.upper()
    p.font.name = 'Tw Cen MT' # Matches the sleek style in the video
    p.font.size = Pt(54)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(3.5), Inches(3.8), Inches(7.0), Inches(0.8))
    stf = sub_box.text_frame
    p2 = stf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = 'Tw Cen MT'
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(220, 220, 220)
    p2.alignment = PP_ALIGN.CENTER

    # === Layer 3: Neon Callout Arrow with LXML Glow ===
    # Draw Arrow
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.LEFT_ARROW, 
        Inches(7.5), Inches(4.5), Inches(2.5), Inches(1.0)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(*accent_color)
    arrow.line.color.rgb = RGBColor(255, 255, 255) # White border
    arrow.line.width = Pt(1.5)

    # Add Text to Arrow
    arrow_tf = arrow.text_frame
    p3 = arrow_tf.paragraphs[0]
    p3.text = callout_text
    p3.font.name = 'Tw Cen MT'
    p3.font.size = Pt(20)
    p3.font.color.rgb = RGBColor(255, 255, 255)
    p3.font.bold = True
    p3.alignment = PP_ALIGN.CENTER

    # Inject LXML Glow Effect to Arrow
    # Converts accent_color tuple to hex string
    hex_color = '{:02x}{:02x}{:02x}'.format(*accent_color).upper()
    
    glow_xml = f"""
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:glow rad="150000">
            <a:srgbClr val="{hex_color}">
                <a:alpha val="60000"/>
            </a:srgbClr>
        </a:glow>
    </a:effectLst>
    """
    effect_element = parse_xml(glow_xml)
    arrow.element.spPr.append(effect_element)

    # Save presentation and clean up
    prs.save(output_pptx_path)
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
