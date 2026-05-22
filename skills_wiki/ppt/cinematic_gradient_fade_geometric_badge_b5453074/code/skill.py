def create_slide(
    output_pptx_path: str,
    title_text: str = "MOUNTAIN",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit.",
    badge_text: str = "Add Logo\nhere",
    bg_keyword: str = "mountain,landscape",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Cinematic Gradient Fade & Geometric Badge style.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # Initialize presentation (Widescreen 16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Temp file paths
    bg_img_path = "temp_bg.jpg"
    grad_img_path = "temp_gradient.png"

    # === Layer 1: Background Image ===
    try:
        url = f"https://source.unsplash.com/random/1920x1080/?{bg_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(bg_img_path, 'wb') as out_file:
                out_file.write(response.read())
    except Exception:
        # Fallback to a solid color block if download fails
        fallback = Image.new('RGB', (1920, 1080), color=(40, 50, 60))
        fallback.save(bg_img_path)

    # Insert background
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: PIL Gradient Overlay Mask ===
    # Generate a transparent-to-black gradient
    grad_width = 1920
    grad_height = int(1080 * 0.45) # Covers bottom 45%
    gradient = Image.new('RGBA', (grad_width, grad_height), color=0)
    draw = ImageDraw.Draw(gradient)

    for y in range(grad_height):
        # Calculate alpha: 0 at top to 255 at bottom
        alpha = int((y / grad_height) * 255)
        # Use dark base color (0,0,0)
        draw.line((0, y, grad_width, y), fill=(0, 0, 0, alpha))
    
    gradient.save(grad_img_path, format="PNG")

    # Insert gradient at the bottom of the slide
    grad_top = prs.slide_height - Inches(7.5 * 0.45)
    slide.shapes.add_picture(
        grad_img_path, 
        0, 
        grad_top, 
        prs.slide_width, 
        Inches(7.5 * 0.45)
    )

    # === Layer 3: Central Hexagon Badge ===
    hex_size = Inches(3.2)
    hex_left = (prs.slide_width - hex_size) / 2
    hex_top = (prs.slide_height - hex_size) / 2 - Inches(0.5)

    hexagon = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, hex_left, hex_top, hex_size, hex_size
    )
    
    # Format Hexagon
    hexagon.fill.solid()
    hexagon.fill.fore_color.rgb = RGBColor(0, 0, 0)
    hexagon.line.color.rgb = RGBColor(255, 255, 255)
    hexagon.line.width = Pt(4.5)

    # Add text to Hexagon
    tf_hex = hexagon.text_frame
    tf_hex.text = badge_text
    for paragraph in tf_hex.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(24)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 4: Typography (Bottom Text) ===
    # Main Title
    title_width = Inches(8)
    title_height = Inches(0.6)
    title_left = (prs.slide_width - title_width) / 2
    title_top = prs.slide_height - Inches(1.5)

    title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text.upper()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle / Body
    body_top = title_top + Inches(0.4)
    body_box = slide.shapes.add_textbox(title_left, body_top, title_width, title_height)
    tf_body = body_box.text_frame
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.alignment = PP_ALIGN.CENTER
    p_body.font.name = "Arial"
    p_body.font.size = Pt(14)
    p_body.font.color.rgb = RGBColor(230, 230, 230)

    # Save Presentation
    prs.save(output_pptx_path)

    # Cleanup temp files
    for file in [bg_img_path, grad_img_path]:
        if os.path.exists(file):
            os.remove(file)

    return output_pptx_path
