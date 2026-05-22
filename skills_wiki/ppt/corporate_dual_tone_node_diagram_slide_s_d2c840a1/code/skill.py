def create_slide(
    output_pptx_path: str,
    title_text: str = "Process Architecture",
    **kwargs
) -> str:
    """
    Creates a PPTX file featuring a dual-tone circular node diagram,
    extracting the aesthetic of the template deck shown in the tutorial.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import math

    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Style Palette ===
    color_bg = RGBColor(248, 248, 250)
    color_purple = RGBColor(43, 22, 114)
    color_magenta = RGBColor(216, 27, 96)
    color_white = RGBColor(255, 255, 255)

    # === Background ===
    # Set a very light gray background so the white shape borders stand out
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color_bg

    # === Title Section ===
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = color_purple

    # === Diagram Layout Parameters ===
    center_x = Inches(13.333 / 2)
    center_y = Inches(4.2)
    orbit_radius = Inches(2.3)
    
    hub_size = Inches(2.2)
    node_size = Inches(1.5)

    # === Connectors (Lines drawn first so they sit behind shapes) ===
    # Draw simple lines radiating from the center to the 4 compass points
    angles = [0, 90, 180, 270]
    for angle in angles:
        rad = math.radians(angle)
        end_x = center_x + orbit_radius * math.cos(rad)
        end_y = center_y + orbit_radius * math.sin(rad)
        
        # 1 = msoConnectorStraight
        connector = slide.shapes.add_connector(1, center_x, center_y, end_x, end_y)
        connector.line.color.rgb = color_purple
        connector.line.width = Pt(2.5)

    # === Central Hub (Purple) ===
    hub = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        center_x - hub_size/2, center_y - hub_size/2, 
        hub_size, hub_size
    )
    hub.fill.solid()
    hub.fill.fore_color.rgb = color_purple
    hub.line.color.rgb = color_white
    hub.line.width = Pt(4) # Thick white border is characteristic of this style
    
    tf_hub = hub.text_frame
    tf_hub.word_wrap = True
    p_hub = tf_hub.paragraphs[0]
    p_hub.text = "CORE\nSYSTEM"
    p_hub.alignment = PP_ALIGN.CENTER
    p_hub.font.size = Pt(18)
    p_hub.font.bold = True
    p_hub.font.color.rgb = color_white

    # === Satellite Nodes (Magenta) ===
    node_labels = ["Phase 1\nPlan", "Phase 2\nDesign", "Phase 3\nBuild", "Phase 4\nTest"]
    
    for i, angle in enumerate(angles):
        rad = math.radians(angle)
        nx = center_x + orbit_radius * math.cos(rad)
        ny = center_y + orbit_radius * math.sin(rad)
        
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            nx - node_size/2, ny - node_size/2, 
            node_size, node_size
        )
        node.fill.solid()
        node.fill.fore_color.rgb = color_magenta
        node.line.color.rgb = color_white
        node.line.width = Pt(3)
        
        tf_node = node.text_frame
        tf_node.word_wrap = True
        p_node = tf_node.paragraphs[0]
        p_node.text = node_labels[i]
        p_node.alignment = PP_ALIGN.CENTER
        p_node.font.size = Pt(14)
        p_node.font.bold = True
        p_node.font.color.rgb = color_white

    prs.save(output_pptx_path)
    return output_pptx_path
