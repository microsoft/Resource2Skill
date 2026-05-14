def create_slide(
    output_pptx_path: str,
    title_text: str = "Wishing You",
    subtitle_text: str = "PEACE, LOVE & JOY\nTHIS HOLIDAY SEASON",
    accent_color: tuple = (197, 160, 89),  # Gold
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Elegant Layered Collage' (Vellum Card) visual effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Colors
    c_bg_green = RGBColor(27, 77, 62)
    c_gold = RGBColor(*accent_color)
    c_ivory = RGBColor(253, 248, 231)
    c_crimson = RGBColor(160, 30, 30)
    c_white = RGBColor(255, 255, 255)

    # --- Helper: LXML Drop Shadow ---
    def add_drop_shadow(shape, blur_pt=5, dist_pt=4, alpha_pct=40):
        spPr = shape.element.spPr
        effectLst = OxmlElement('a:effectLst')
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', str(blur_pt * 12700)) # 1 pt = 12700 EMU
        outerShdw.set('dist', str(dist_pt * 12700))
        outerShdw.set('dir', '2700000') # 45 deg down/right
        outerShdw.set('algn', 'tl')
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', str(int(alpha_pct * 1000))) # e.g. 40000 for 40%
        srgbClr.append(alpha)
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)
        spPr.append(effectLst)

    # --- Helper: LXML Alpha Transparency ---
    def make_transparent(shape, alpha_pct):
        # Assumes shape has a solid fill applied already
        alpha_val = str(int(alpha_pct * 1000))
        for srgbClr in shape.element.iter('.//a:srgbClr'):
            # Remove existing alpha if present to avoid duplicates
            for existing_alpha in srgbClr.findall('.//a:alpha', namespaces=shape.element.nsmap):
                srgbClr.remove(existing_alpha)
            alpha = OxmlElement('a:alpha')
            alpha.set('val', alpha_val)
            srgbClr.append(alpha)

    # === Layer 0: Background Texture ===
    # Attempt to download a subtle paper/grunge texture, fallback to solid deep green
    bg_img_path = "temp_bg_texture.jpg"
    try:
        url = "https://images.unsplash.com/photo-1603513492128-ba7bfafcb3bf?q=80&w=1920&auto=format&fit=crop"
        urllib.request.urlretrieve(url, bg_img_path)
        slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)
    except:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg_green
        bg.line.fill.background()

    # === Layer 1: The Gold Mat ===
    mat_w, mat_h = Inches(8.5), Inches(6.5)
    mat_left = (prs.slide_width - mat_w) / 2
    mat_top = (prs.slide_height - mat_h) / 2
    mat = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, mat_left, mat_top, mat_w, mat_h)
    mat.fill.solid()
    mat.fill.fore_color.rgb = c_gold
    mat.line.fill.background()
    add_drop_shadow(mat, blur_pt=8, dist_pt=5, alpha_pct=50)

    # === Layer 2: The Ivory Card Body ===
    card_w, card_h = Inches(8.0), Inches(6.0)
    card_left = (prs.slide_width - card_w) / 2
    card_top = (prs.slide_height - card_h) / 2
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, card_top, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = c_ivory
    card.line.color.rgb = c_bg_green
    card.line.width = Pt(1)
    add_drop_shadow(card, blur_pt=4, dist_pt=2, alpha_pct=30)

    # === Layer 3: Crimson Ribbons (Embellishments) ===
    # Vertical ribbon
    v_ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, prs.slide_width/2 - Inches(0.5), card_top, Inches(1.0), card_h)
    v_ribbon.fill.solid()
    v_ribbon.fill.fore_color.rgb = c_crimson
    v_ribbon.line.fill.background()
    add_drop_shadow(v_ribbon, blur_pt=3, dist_pt=2, alpha_pct=40)
    
    # Horizontal ribbon
    h_ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, card_left, prs.slide_height/2 - Inches(0.5), card_w, Inches(1.0))
    h_ribbon.fill.solid()
    h_ribbon.fill.fore_color.rgb = c_crimson
    h_ribbon.line.fill.background()
    add_drop_shadow(h_ribbon, blur_pt=3, dist_pt=2, alpha_pct=40)

    # === Layer 4: The Translucent Vellum Tag ===
    # A bracket/ticket shape placed over the ribbons to show the translucency effect
    vellum_w, vellum_h = Inches(5.5), Inches(3.5)
    vellum_left = (prs.slide_width - vellum_w) / 2
    vellum_top = (prs.slide_height - vellum_h) / 2
    # Using ROUNDED_RECTANGLE to mimic die-cut sticker
    vellum = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, vellum_left, vellum_top, vellum_w, vellum_h)
    
    # Style the vellum
    vellum.fill.solid()
    vellum.fill.fore_color.rgb = c_white
    # Make it 80% opaque (20% transparent) so ribbons show through subtly
    make_transparent(vellum, 80) 
    
    vellum.line.color.rgb = c_gold
    vellum.line.width = Pt(2.5)
    # Strong shadow to separate the vellum from the ribbons
    add_drop_shadow(vellum, blur_pt=6, dist_pt=4, alpha_pct=50)

    # === Layer 5: Typography ===
    # Adding a text box exactly over the vellum
    tx_box = slide.shapes.add_textbox(vellum_left, vellum_top + Inches(0.4), vellum_w, vellum_h)
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    # Title Paragraph (Elegant Script/Serif feel)
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    run1 = p1.add_run()
    run1.text = title_text
    run1.font.name = "Georgia"
    run1.font.size = Pt(28)
    run1.font.italic = True
    run1.font.color.rgb = c_bg_green

    # Spacing paragraph
    p_space = tf.add_paragraph()
    p_space.font.size = Pt(12)

    # Subtitle Paragraph (Clean, tracked out Sans/Serif)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = subtitle_text
    run2.font.name = "Arial"
    run2.font.size = Pt(16)
    run2.font.bold = True
    run2.font.color.rgb = c_gold

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path
