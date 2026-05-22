def create_slide(
    output_pptx_path: str,
    title_text: str = "By combining a template, **custom colors**, a high-quality image, and a modern font, we create **real impact**.",
    bg_theme: str = "abstract colorful powder splash",
    accent_color: tuple = (255, 192, 0),  # RGB Gold
    brightness_factor: float = 0.35,      # Dim image to 35% of original brightness
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Cinematic Full-Bleed Impact Slide" visual effect.
    Use **double asterisks** around words in `title_text` to apply the accent color.
    
    Returns: path to the saved PPTX file.
    """
    import os
    import urllib.request
    import urllib.parse
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from PIL import Image, ImageDraw, ImageEnhance

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Image Fetch & Dimming ===
    img_path = "temp_background.jpg"
    
    # Attempt to fetch a thematic high-res image
    prompt = urllib.parse.quote(bg_theme)
    url = f"https://image.pollinations.ai/prompt/{prompt}?width=1920&height=1080&nologo=true"
    
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
        img = Image.open(img_path).convert("RGB")
    except Exception as e:
        # Fallback to a high-quality dark gradient if network fails
        img = Image.new("RGB", (1920, 1080))
        draw = ImageDraw.Draw(img)
        for y in range(1080):
            r = int(10 + (30 - 10) * y / 1080)
            g = int(15 + (40 - 15) * y / 1080)
            b = int(25 + (60 - 25) * y / 1080)
            draw.line([(0, y), (1920, y)], fill=(r, g, b))

    # Dim the image to create a legible canvas for text
    enhancer = ImageEnhance.Brightness(img)
    img_dimmed = enhancer.enhance(brightness_factor)
    img_dimmed.save(img_path)

    # Insert full-bleed background
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Visual Effect (Typography & Highlighting) ===
    # Create a centered text box with generous margins
    txBox = slide.shapes.add_textbox(
        left=Inches(1.66), 
        top=Inches(1.5), 
        width=Inches(10.0), 
        height=Inches(4.5)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Parse text for **markdown-style** highlights and apply runs
    paragraphs = title_text.split('\n')
    for p_idx, para_text in enumerate(paragraphs):
        if p_idx > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
            
        p.line_spacing = 1.2
        
        # Split by the double asterisk marker
        chunks = para_text.split("**")
        
        for i, chunk in enumerate(chunks):
            if not chunk:
                continue
            
            run = p.add_run()
            run.text = chunk
            run.font.name = "Arial" # Solid, universally available modern fallback
            run.font.size = Pt(48)
            
            # Odd indices are the highlighted text chunks
            if i % 2 == 1:
                run.font.color.rgb = RGBColor(*accent_color)
                run.font.bold = True
            else:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = False

    # Cleanup temporary image
    try:
        if os.path.exists(img_path):
            os.remove(img_path)
    except Exception:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path
