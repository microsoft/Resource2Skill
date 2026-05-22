def create_slide(
    output_pptx_path: str,
    title_text: str = "Slide Main title",
    subtitle_text: str = "Sub title text here",
    right_title: str = "Top quality design",
    right_body: str = "Lorem ipsum is simply dummy text of the printing and typesetting industry. Lorem Ipsum has been the industry's standard dummy text ever since the 1500s, when an unknown printer took a galley of type and scrambled it to make a type specimen book.",
    image_theme: str = "architecture,portrait",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Premium Industrial Dark Grid' visual effect.
    """
    import os
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageOps

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # --- Color Palette ---
    COLOR_BG = RGBColor(55, 55, 55)
    COLOR_GOLD = RGBColor(218, 176, 60)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_DARK_TEXT = RGBColor(40, 40, 40)
    COLOR_LIGHT_GREY = RGBColor(180, 180, 180)

    # --- Apply Dark Background ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

    # --- Helper: Create Text Box ---
    def add_text_box(slide, text, left, top, width, height, font_size, font_color, bold=False, align=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.font.name = "Arial"
        return txBox

    # --- Header Section ---
    add_text_box(slide, title_text, Inches(0.5), Inches(0.3), Inches(8), Inches(0.8), 36, COLOR_WHITE, bold=True)
    add_text_box(slide, subtitle_text, Inches(0.5), Inches(1.0), Inches(8), Inches(0.5), 18, COLOR_GOLD)

    # --- Horizontal Separator Line ---
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.6), Inches(12.333), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_LIGHT_GREY
    line.line.fill.background()

    # --- PIL Image Processor ---
    def fetch_and_process_image(idx, width_in, height_in):
        dpi = 150
        px_w = int(width_in * dpi)
        px_h = int(height_in * dpi)
        temp_path = f"temp_profile_{idx}.jpg"
        
        # Download image
        url = f"https://source.unsplash.com/featured/{px_w}x{px_h}/?{image_theme}&sig={idx}"
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=5) as response:
                img = Image.open(io.BytesIO(response.read()))
        except Exception:
            # Fallback if network fails
            img = Image.new('RGB', (px_w, px_h), color=(120, 120, 120))
            
        # Crop and force grayscale
        img = ImageOps.fit(img, (px_w, px_h), Image.Resampling.LANCZOS)
        img = img.convert('L') # Magic happens here: force grayscale
        img.save(temp_path, quality=90)
        return temp_path

    # --- Left/Center Grid Layout (3 Profiles) ---
    col_width = Inches(2.5)
    img_height = Inches(2.8)
    box_height = Inches(2.0)
    start_left = Inches(0.5)
    start_top = Inches(2.2)
    spacing = Inches(0.25)

    names = ["John Done", "Iren Rose", "John Done"]

    for i in range(3):
        cur_left = start_left + (i * (col_width + spacing))
        
        # 1. Process & Add Grayscale Image
        img_path = fetch_and_process_image(i, col_width.inches, img_height.inches)
        slide.shapes.add_picture(img_path, cur_left, start_top, col_width, img_height)
        if os.path.exists(img_path):
            os.remove(img_path) # Cleanup
            
        # 2. Add Gold Info Box underneath
        box_top = start_top + img_height
        gold_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cur_left, box_top, col_width, box_height)
        gold_box.fill.solid()
        gold_box.fill.fore_color.rgb = COLOR_GOLD
        gold_box.line.fill.background() # No border
        
        # 3. Add Text inside Gold Box
        tf = gold_box.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = "Manager\n"
        p1.font.size = Pt(12)
        p1.font.color.rgb = COLOR_DARK_TEXT
        
        p2 = tf.add_paragraph()
        p2.text = f"{names[i]}\n"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_DARK_TEXT
        
        p3 = tf.add_paragraph()
        p3.text = "Lorem ipsum is simply dummy text of the printing and typesetting industry."
        p3.font.size = Pt(10)
        p3.font.color.rgb = COLOR_DARK_TEXT

    # --- Right Side Typography Block ---
    right_left = Inches(9.2)
    right_top = Inches(2.2)
    right_width = Inches(3.6)

    # Large Side Title
    add_text_box(slide, right_title, right_left, right_top, right_width, Inches(1.0), 32, COLOR_WHITE, bold=True)
    
    # Yellow separator under side title
    sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_left, right_top + Inches(1.1), Inches(1.0), Inches(0.05))
    sep.fill.solid()
    sep.fill.fore_color.rgb = COLOR_GOLD
    sep.line.fill.background()

    # Body paragraph
    add_text_box(slide, right_body, right_left, right_top + Inches(1.3), right_width, Inches(4.0), 12, COLOR_LIGHT_GREY)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
