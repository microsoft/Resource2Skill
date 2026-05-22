import colorsys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_slide_with_custom_colors(
    output_pptx_path: str,
    rgb_color: tuple = (255, 87, 34),  # A sample vibrant orange
    hsl_color: tuple = (205, 0.85, 0.55), # A sample sky blue (Hue, Saturation, Lightness)
    **kwargs,
) -> str:
    """
    Creates a PPTX file demonstrating how to set shape colors using RGB and HSL models.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        rgb_color (tuple): A tuple of (R, G, B) values (0-255).
        hsl_color (tuple): A tuple of (Hue, Saturation, Lightness) values.
                           - Hue is in degrees (0-360).
                           - Saturation and Lightness are floats (0.0 - 1.0).

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout

    # --- Slide Title ---
    title = slide.shapes.title
    title.text = "Precision Color Control: RGB vs. HSL"
    title.text_frame.paragraphs[0].font.name = "Arial Black"
    title.text_frame.paragraphs[0].font.size = Pt(40)
    
    # --- Part 1: Setting Color with RGB ---
    # Label
    tx_box_rgb = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(5), Inches(0.8))
    p_rgb = tx_box_rgb.text_frame.paragraphs[0]
    p_rgb.text = f"Method 1: Direct RGB Input\nRGB: {rgb_color}"
    p_rgb.font.size = Pt(20)
    p_rgb.font.name = "Arial"

    # Shape
    shape_rgb = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.7), Inches(5), Inches(3)
    )
    fill_rgb = shape_rgb.fill
    fill_rgb.solid()
    fill_rgb.fore_color.rgb = RGBColor(rgb_color[0], rgb_color[1], rgb_color[2])
    shape_rgb.line.fill.background()  # Remove outline

    # --- Part 2: Setting Color with HSL (via conversion) ---
    # Label
    tx_box_hsl = slide.shapes.add_textbox(Inches(7.33), Inches(1.8), Inches(5), Inches(0.8))
    p_hsl = tx_box_hsl.text_frame.paragraphs[0]
    p_hsl.text = f"Method 2: Intuitive HSL Input\nHSL: ({int(hsl_color[0])}°, {int(hsl_color[1]*100)}%, {int(hsl_color[2]*100)}%)"
    p_hsl.font.size = Pt(20)
    p_hsl.font.name = "Arial"

    # Shape
    shape_hsl = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.33), Inches(2.7), Inches(5), Inches(3)
    )

    # HSL-to-RGB Conversion Logic
    h, s, l = hsl_color
    normalized_h = h / 360.0
    
    # colorsys.hls_to_rgb returns a tuple of floats (0.0-1.0)
    # Note the order for this specific function is H, L, S
    rgb_float = colorsys.hls_to_rgb(normalized_h, l, s)

    # Convert float tuple (0.0-1.0) to integer tuple (0-255)
    rgb_int_from_hsl = tuple(int(c * 255) for c in rgb_float)

    fill_hsl = shape_hsl.fill
    fill_hsl.solid()
    fill_hsl.fore_color.rgb = RGBColor(rgb_int_from_hsl[0], rgb_int_from_hsl[1], rgb_int_from_hsl[2])
    shape_hsl.line.fill.background()  # Remove outline

    prs.save(output_pptx_path)
    return output_pptx_path
