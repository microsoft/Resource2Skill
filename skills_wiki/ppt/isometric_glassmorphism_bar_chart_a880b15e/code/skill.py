def create_slide(
    output_pptx_path: str,
    title_text: str = "SALES REPORT",
    **kwargs,
) -> str:
    """
    Creates a slide featuring an Isometric Glassmorphism Bar Chart on a 3D stage.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.shapes.freeform import FreeformBuilder
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ==========================================
    # HELPER: XML Injection for Advanced Styling
    # ==========================================
    def apply_style(shape, stops, angle=5400000, no_line=True):
        """Applies solid/gradient fills with alpha transparency and removes borders."""
        spPr = shape._element.spPr
        
        # Strip existing fills
        for elem in spPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill"):
            spPr.remove(elem)
        for elem in spPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill"):
            spPr.remove(elem)
            
        if len(stops) == 1:
            # Solid Fill
            hex_color, _, alpha = stops[0]
            fill = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
            srgb = etree.SubElement(fill, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            srgb.set("val", hex_color)
            if alpha < 100000:
                a_tag = etree.SubElement(srgb, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
                a_tag.set("val", str(alpha))
        else:
            # Gradient Fill
            grad = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill")
            grad.set("rotWithShape", "1")
            gsLst = etree.SubElement(grad, "{http://schemas.openxmlformats.org/drawingml/2006/main}gsLst")
            for hex_color, pos, alpha in stops:
                gs = etree.SubElement(gsLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}gs")
                gs.set("pos", str(pos))
                srgb = etree.SubElement(gs, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
                srgb.set("val", hex_color)
                if alpha < 100000:
                    a_tag = etree.SubElement(srgb, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
                    a_tag.set("val", str(alpha))
            lin = etree.SubElement(grad, "{http://schemas.openxmlformats.org/drawingml/2006/main}lin")
            lin.set("ang", str(angle))
            lin.set("scaled", "1")
            
        # Strip line
        if no_line:
            for elem in spPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}ln"):
                spPr.remove(elem)
            ln = etree.SubElement(spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
            etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}noFill")


    # ==========================================
    # LAYER 1: Deep Atmospheric Background
    # ==========================================
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    apply_style(bg, [("1A0B2E", 0, 100000), ("0B132B", 100000, 100000)])

    # Wavy Background Mountains
    ff = slide.shapes.build_freeform(Inches(5.0), Inches(7.5))
    ff.add_nodes(FreeformBuilder.PT_BEZIERCURVE, [(Inches(6.0), Inches(4.0)), (Inches(8.0), Inches(2.5)), (Inches(9.0), Inches(4.5))])
    ff.add_nodes(FreeformBuilder.PT_BEZIERCURVE, [(Inches(10.0), Inches(6.5)), (Inches(11.5), Inches(3.5)), (Inches(13.5), Inches(4.0))])
    ff.add_line_segments([(Inches(13.5), Inches(7.5)), (Inches(5.0), Inches(7.5))])
    mountain = ff.convert_to_shape()
    # Semi-transparent fading gradient
    apply_style(mountain, [("06D6A0", 0, 70000), ("118AB2", 100000, 10000)])

    # ==========================================
    # LAYER 2: Fake 3D Stage
    # ==========================================
    stage_cx, stage_cy, stage_w, stage_h, depth = 8.0, 6.0, 8.0, 2.5, 0.5
    
    bot_oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(stage_cx - stage_w/2), Inches(stage_cy - stage_h/2 + depth), Inches(stage_w), Inches(stage_h))
    apply_style(bot_oval, [("050A1F", 0, 100000)])
    
    mid_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(stage_cx - stage_w/2), Inches(stage_cy), Inches(stage_w), Inches(depth))
    apply_style(mid_rect, [("0A1128", 0, 100000)])
    
    top_oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(stage_cx - stage_w/2), Inches(stage_cy - stage_h/2), Inches(stage_w), Inches(stage_h))
    apply_style(top_oval, [("1C2541", 0, 100000)])

    # ==========================================
    # LAYER 3: Volumetric Glass Pillars
    # ==========================================
    pillars = [
        {"cx": 5.5, "h": 2.0, "pct": "45%", "c1": "FFD166", "c2": "F4A261"},
        {"cx": 8.0, "h": 3.2, "pct": "65%", "c1": "06D6A0", "c2": "118AB2"},
        {"cx": 10.5, "h": 4.5, "pct": "80%", "c1": "EF476F", "c2": "8338EC"},
    ]
    
    w, d, by = 1.2, 0.6, 6.0  # Width, Isometric Depth, Base Y

    for p in pillars:
        cx, h, c1, c2 = p["cx"], p["h"], p["c1"], p["c2"]
        ty = by - h  # Top Y
        
        # 1. Main Base Pillar
        ff = slide.shapes.build_freeform(Inches(cx - w/2), Inches(ty))
        ff.add_line_segments([(Inches(cx), Inches(ty + d/2)), (Inches(cx + w/2), Inches(ty)), 
                              (Inches(cx + w/2), Inches(by)), (Inches(cx), Inches(by + d/2)), 
                              (Inches(cx - w/2), Inches(by)), (Inches(cx - w/2), Inches(ty))])
        main_shape = ff.convert_to_shape()
        apply_style(main_shape, [(c1, 0, 100000), (c2, 100000, 100000)])

        # 2. Right Highlight (Glass Overlay)
        ff = slide.shapes.build_freeform(Inches(cx), Inches(ty + d/2))
        ff.add_line_segments([(Inches(cx + w/2), Inches(ty)), (Inches(cx + w/2), Inches(by)),
                              (Inches(cx), Inches(by + d/2)), (Inches(cx), Inches(ty + d/2))])
        hl_shape = ff.convert_to_shape()
        apply_style(hl_shape, [("FFFFFF", 0, 30000), ("FFFFFF", 100000, 0)]) # 30% to 0% opaque

        # 3. Top Diamond Face
        ff = slide.shapes.build_freeform(Inches(cx), Inches(ty - d/2))
        ff.add_line_segments([(Inches(cx + w/2), Inches(ty)), (Inches(cx), Inches(ty + d/2)),
                              (Inches(cx - w/2), Inches(ty)), (Inches(cx), Inches(ty - d/2))])
        top_shape = ff.convert_to_shape()
        apply_style(top_shape, [(c1, 0, 100000)]) # Solid match to top gradient

        # 4. Data Line & Percentage
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - 0.02), Inches(ty - d/2 - 0.7), Inches(0.04), Inches(0.6))
        apply_style(line, [("FFFFFF", 0, 60000)]) # 60% opaque line

        tx = slide.shapes.add_textbox(Inches(cx - 0.6), Inches(ty - d/2 - 1.1), Inches(1.2), Inches(0.4))
        p_tf = tx.text_frame.paragraphs[0]
        p_tf.text = p["pct"]
        p_tf.alignment = PP_ALIGN.CENTER
        p_tf.font.size = Pt(22)
        p_tf.font.bold = True
        p_tf.font.color.rgb = RGBColor.from_string("FFFFFF")

    # ==========================================
    # LAYER 4: Floating Glass UI Panel
    # ==========================================
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.0), Inches(3.5), Inches(2.5))
    apply_style(panel, [("FFFFFF", 0, 20000), ("FFFFFF", 100000, 5000)], angle=3000000) # Diagonal glass fade
    
    tx = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(3.0), Inches(0.5))
    title_p = tx.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.font.color.rgb = RGBColor.from_string("FFFFFF")
    title_p.font.size = Pt(16)
    title_p.font.bold = True

    # Decorative Line Chart on Panel
    chart = slide.shapes.build_freeform(Inches(1.4), Inches(3.8))
    chart.add_line_segments([(Inches(1.8), Inches(3.2)), (Inches(2.4), Inches(3.6)), 
                             (Inches(3.0), Inches(2.6)), (Inches(3.6), Inches(2.9)), (Inches(4.2), Inches(2.2))])
    c_shape = chart.convert_to_shape()
    apply_style(c_shape, [("00FFFF", 0, 100000)], no_line=False)
    c_shape.line.color.rgb = RGBColor.from_string("00FFFF")
    c_shape.line.width = Pt(2.5)
    
    # Add neon glow to line chart
    effectLst = etree.SubElement(c_shape._element.spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    glow = etree.SubElement(effectLst, "{http://schemas.openxmlformats.org/drawingml/2006/main}glow")
    glow.set("rad", "63500") # 5pt radius
    srgbClr = etree.SubElement(glow, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    srgbClr.set("val", "00FFFF")
    a_tag = etree.SubElement(srgbClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
    a_tag.set("val", "50000")

    prs.save(output_pptx_path)
    return output_pptx_path
