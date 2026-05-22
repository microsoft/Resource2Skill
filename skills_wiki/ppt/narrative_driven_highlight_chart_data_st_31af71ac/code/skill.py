def create_slide(
    output_pptx_path: str,
    title_text: str = "Loyalty Members Spend 50% More During Peak Season",
    subtitle_text: str = "Average Order Value (AOV) comparison between loyal and non-loyal customer segments over 12 months.",
    accent_color: tuple = (0, 112, 192),  # Deep Azure Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Narrative-Driven Highlight Chart' data storytelling effect.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Storytelling Titles (The "So What?") ===
    # Headline
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 30, 30)

    # Subtitle / Metric definition
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(120, 120, 120)

    # === Layer 2: The Data Chart ===
    chart_data = CategoryChartData()
    chart_data.categories = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    
    # Dummy Data setup
    loyal_data = [210, 220, 240, 250, 270, 310, 340, 380, 350, 340, 360, 390]
    non_loyal_data = [200, 205, 215, 210, 220, 230, 240, 250, 235, 230, 240, 250]
    guest_data = [180, 185, 190, 185, 195, 200, 210, 220, 210, 205, 215, 220]

    # Add series (Signal first, then Noise)
    chart_data.add_series('Loyal Members', loyal_data)
    chart_data.add_series('Non-Loyal', non_loyal_data)
    chart_data.add_series('Guests', guest_data)

    x, y, cx, cy = Inches(0.5), Inches(1.8), Inches(11.5), Inches(5.2)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, x, y, cx, cy, chart_data
    ).chart

    # --- Data-to-Ink Formatting ---
    chart.has_legend = False  # We will use direct labeling

    # Y-Axis Formatting (Minimalist)
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor(235, 235, 235) # Very faint gray
    value_axis.major_tick_mark = XL_TICK_MARK.NONE
    value_axis.format.line.fill.background() # Remove axis line

    # X-Axis Formatting
    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    category_axis.major_tick_mark = XL_TICK_MARK.NONE
    category_axis.format.line.color.rgb = RGBColor(200, 200, 200)

    # --- Signal vs. Noise Series Styling ---
    # Highlight Series (Signal)
    series_loyal = chart.series[0]
    series_loyal.format.line.color.rgb = RGBColor(*accent_color)
    series_loyal.format.line.width = Pt(3.5)

    # Context Series 1 (Noise)
    series_non_loyal = chart.series[1]
    series_non_loyal.format.line.color.rgb = RGBColor(180, 180, 180)
    series_non_loyal.format.line.width = Pt(1.5)

    # Context Series 2 (Noise)
    series_guest = chart.series[2]
    series_guest.format.line.color.rgb = RGBColor(215, 215, 215)
    series_guest.format.line.width = Pt(1.5)

    # === Layer 3: Storytelling Enhancements (Direct Labels & Annotations) ===

    # Direct Labeling (Replaces Legend) - positions estimated near the end of the lines
    labels = [
        ("Loyal Members", RGBColor(*accent_color), Inches(12.1), Inches(2.2)),
        ("Non-Loyal", RGBColor(150, 150, 150), Inches(12.1), Inches(4.5)),
        ("Guests", RGBColor(180, 180, 180), Inches(12.1), Inches(5.1))
    ]

    for text, color, lx, ly in labels:
        lbl_box = slide.shapes.add_textbox(lx, ly, Inches(1.2), Inches(0.5))
        tf_lbl = lbl_box.text_frame
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.text = text
        p_lbl.font.size = Pt(12)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = color

    # Callout Annotation pointing to the specific insight (Peak in August)
    # Background Box
    callout = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2.0), Inches(2.2), Inches(0.8)
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = RGBColor(255, 248, 220) # Soft highlighting yellow
    callout.line.color.rgb = RGBColor(230, 220, 180)
    
    tf_callout = callout.text_frame
    tf_callout.word_wrap = True
    p_callout = tf_callout.paragraphs[0]
    p_callout.alignment = PP_ALIGN.CENTER
    p_callout.text = "Loyalty Promo launched, driving a 50% spike vs. non-loyals."
    p_callout.font.size = Pt(11)
    p_callout.font.bold = True
    p_callout.font.color.rgb = RGBColor(60, 60, 60)

    # Connector Line to the data point
    connector = slide.shapes.add_shape(
        MSO_SHAPE.LINE_INVERSE, Inches(7.6), Inches(2.8), Inches(0.4), Inches(0.8)
    )
    connector.line.color.rgb = RGBColor(150, 150, 150)
    connector.line.width = Pt(1.5)

    prs.save(output_pptx_path)
    return output_pptx_path
