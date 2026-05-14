AMBIENT_CAPABLE = True

def create_slide(
    output_pptx_path: str,
    title_text: str = "Orbital Motion Dynamics",
    body_text: str = "Custom paths and continuous animation loops",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Orbital Motion paths visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    import os
    import random
    from PIL import Image, ImageDraw

    # Import ambient animation helpers (with graceful fallback for standard environments)
    try:
        from _shell_helpers import add_orbital_motion, add_drift_motion
    except ImportError:
        def add_orbital_motion(*args, **kwargs): pass
        def add_drift_motion(*args, **kwargs): pass

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Generate Starry Background using PIL ===
    bg_img_path = "temp_starfield_bg.png"
    width_px, height_px = 1920, 1080
    bg_img = Image.new('RGBA', (width_px, height_px), (10, 14, 30, 255))
    draw = ImageDraw.Draw(bg_img)
    
    # Draw scattered stars
    for _ in range(300):
        x = random.randint(0, width_px)
        y = random.randint(0, height_px)
        radius = random.uniform(0.5, 2.5)
        alpha = random.randint(100, 255)
        draw.ellipse([x-radius, y-radius, x+radius, y+radius], fill=(255, 255, 255, alpha))
    
    bg_img.save(bg_img_path)
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Text Layout ===
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(8), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(170, 180, 200)

    # === Layer 3: System Canvas (Solar System concept) ===
    # Center coordinates
    cx_in = 13.333 / 2
    cy_in = 7.5 / 2 + 0.5  # Slightly offset downward to balance the top title

    # 1. The Sun (Center Hub)
    sun_radius = 1.0
    sun = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(cx_in - sun_radius),
        Inches(cy_in - sun_radius),
        Inches(sun_radius * 2),
        Inches(sun_radius * 2)
    )
    sun.fill.solid()
    sun.fill.fore_color.rgb = RGBColor(255, 204, 0)
    sun.line.fill.background()

    # Define Orbit Parameters
    orbits = [
        {"radius": 2.2, "color": RGBColor(100, 149, 237), "size": 0.3, "duration": 4000, "dir": "cw"},  # Earth
        {"radius": 3.6, "color": RGBColor(205, 92, 92), "size": 0.25, "duration": 7000, "dir": "cw"},   # Mars
        {"radius": 5.2, "color": RGBColor(200, 180, 150), "size": 0.5, "duration": 12000, "dir": "ccw"} # Jupiter
    ]

    for orbit in orbits:
        r = orbit["radius"]
        
        # Draw Orbit Path (Dashed Ring)
        ring = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx_in - r),
            Inches(cy_in - r),
            Inches(r * 2),
            Inches(r * 2)
        )
        ring.fill.background() # No fill
        ring.line.color.rgb = RGBColor(255, 255, 255)
        ring.line.width = Pt(1)
        ring.line.dash_style = 4 # Dashed line
        ring.shadow.inherit = False

        # Draw Planet Shape
        p_size = orbit["size"]
        planet = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx_in + r - p_size/2), # Initial position on the right edge of the ring
            Inches(cy_in - p_size/2),
            Inches(p_size),
            Inches(p_size)
        )
        planet.fill.solid()
        planet.fill.fore_color.rgb = orbit["color"]
        planet.line.fill.background()

        # Apply continuous orbital motion
        add_orbital_motion(
            slide=slide,
            shape=planet,
            center_xy=(cx_in, cy_in),
            radius_in=r,
            duration_ms=orbit["duration"],
            direction=orbit["dir"]
        )

    # === Layer 4: Drifting Element (Comet / Feather equivalent) ===
    comet_w, comet_h = 0.6, 0.2
    comet = slide.shapes.add_shape(
        MSO_SHAPE.TEARDROP,
        Inches(1.0),
        Inches(6.0),
        Inches(comet_w),
        Inches(comet_h)
    )
    comet.rotation = 45
    comet.fill.solid()
    comet.fill.fore_color.rgb = RGBColor(0, 255, 255)
    comet.line.fill.background()

    # Apply drifting motion (back and forth diagonally across the slide bottom)
    add_drift_motion(
        slide=slide,
        shape=comet,
        dx_in=10.0,
        dy_in=-2.0,
        duration_ms=8000,
        pingpong=True
    )

    # Save Presentation
    prs.save(output_pptx_path)

    # Cleanup temporary image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    return output_pptx_path
