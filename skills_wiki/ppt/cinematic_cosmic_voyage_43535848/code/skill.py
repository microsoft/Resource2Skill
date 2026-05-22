import os
import random
import urllib.request
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.util import Inches, Pt, Emu
from lxml import etree

# Helper for XML namespaces
def qn(tag):
    """
    Stands for 'qualified name', a utility function to turn a namespace prefixed
    tag name into a Clark-notation qualified tag name for lxml. For example,
    'p:cSld' becomes '{http://schemas.openxmlformats.org/presentationml/2006/main}cSld'.
    """
    nsmap = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    prefix, tagroot = tag.split(":")
    uri = nsmap[prefix]
    return f"{{{uri}}}{tagroot}"

# Animation Helper
class AnimationManager:
    """A helper class to manage the complexities of adding animations via LXML."""
    def __init__(self, slide):
        self.slide = slide
        self.tree = self.slide.element
        self.timing = self._get_or_create_timing()
        self.next_node_id = 1

    def _get_or_create_timing(self):
        timing = self.tree.find(qn("p:timing"))
        if timing is None:
            sld_element = self.tree.xpath('//p:cSld')[0]
            timing = etree.SubElement(sld_element, qn("p:timing"))
        
        tn_lst = timing.find(qn("p:tnLst"))
        if tn_lst is None:
            tn_lst = etree.SubElement(timing, qn("p:tnLst"))
        
        par = tn_lst.find(qn("p:par"))
        if par is None:
            par = etree.SubElement(tn_lst, qn("p:par"))
        
        return par

    def _add_common_time_node(self, shape_id, dur="indefinite", repeat_count="indefinite", fill="hold", delay=0):
        c_tn = etree.SubElement(self.timing, qn("p:cTn"), id=str(self.next_node_id), dur=dur, fill=fill)
        if repeat_count != "indefinite":
            c_tn.set("repeatCount", repeat_count)
        self.next_node_id += 1
        
        st_cond_lst = etree.SubElement(c_tn, qn("p:stCondLst"))
        etree.SubElement(st_cond_lst, qn("p:cond"), delay=str(delay * 1000)) # delay in ms
        
        child_tn_lst = etree.SubElement(c_tn, qn("p:childTnLst"))
        par_child = etree.SubElement(child_tn_lst, qn("p:par"))

        set_node = etree.SubElement(par_child, qn("p:set"))
        c_bhvr = etree.SubElement(set_node, qn("p:cBhvr"))
        c_tn_inner = etree.SubElement(c_bhvr, qn("p:cTn"), id=str(self.next_node_id), dur="1000", fill="hold")
        self.next_node_id += 1
        st_cond_lst_inner = etree.SubElement(c_tn_inner, qn("p:stCondLst"))
        etree.SubElement(st_cond_lst_inner, qn("p:cond"), delay="0")
        tgt_el = etree.SubElement(c_bhvr, qn("p:tgtEl"))
        etree.SubElement(tgt_el, qn("p:spTgt"), spid=str(shape_id))

        return par_child

    def add_motion_path(self, shape, path, dur, delay=0):
        shape_id = shape.shape_id
        c_tn_par = self._add_common_time_node(shape_id, dur=str(dur*1000), repeat_count="1", delay=delay)
        
        anim_motion = etree.SubElement(c_tn_par, qn("p:animMotion"), origin="layout")
        c_bhvr = etree.SubElement(anim_motion, qn("p:cBhvr"))
        
        c_tn_inner = etree.SubElement(c_bhvr, qn("p:cTn"), id=str(self.next_node_id), dur=str(dur * 1000))
        self.next_node_id += 1
        
        tgt_el = etree.SubElement(c_bhvr, qn("p:tgtEl"))
        etree.SubElement(tgt_el, qn("p:spTgt"), spid=str(shape_id))
        
        path_el = etree.SubElement(anim_motion, qn("p:path"))
        etree.SubElement(path_el, qn("p:path"),).set("path", path)
    
    def add_spin(self, shape, dur, delay=0):
        shape_id = shape.shape_id
        c_tn_par = self._add_common_time_node(shape_id, dur="indefinite", delay=delay)
        
        anim = etree.SubElement(c_tn_par, qn("p:anim"), by="360000", calcmode="lin")
        c_bhvr = etree.SubElement(anim, qn("p:cBhvr"), additive="base")
        
        c_tn_inner = etree.SubElement(c_bhvr, qn("p:cTn"), id=str(self.next_node_id), dur=str(dur*1000), repeatCount="indefinite")
        self.next_node_id += 1
        
        tgt_el = etree.SubElement(c_bhvr, qn("p:tgtEl"))
        etree.SubElement(tgt_el, qn("p:spTgt"), spid=str(shape_id))
        
        attr_name_lst = etree.SubElement(anim, qn("p:attrNameLst"))
        etree.SubElement(attr_name_lst, qn("p:attrName")).text = "r"


def create_slide(
    output_pptx_path: str,
    title_text: str = "穿梭蔚蓝",
    subtitle_text: str = "穿越浩瀚星空, 感受蓝色星球",
) -> str:
    """
    Creates a PPTX file with a cinematic cosmic voyage animation.

    Args:
        output_pptx_path: The path to save the generated PPTX file.
        title_text: The main title text.
        subtitle_text: The subtitle text.

    Returns:
        The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # --- Asset URLs ---
    assets = {
        "background": "https://images.unsplash.com/photo-1534796636912-3b95b3ab5986?ixlib=rb-4.0.3&q=85&fm=jpg&crop=entropy&cs=srgb&w=1920",
        "earth": "http://assets.stickpng.com/images/580b585b2edb36847703775d.png",
        "moon": "https://www.pngall.com/wp-content/uploads/2016/03/Moon-Vector-PNG.png",
        "spaceship": "https://www.pngmart.com/files/13/Spaceship-PNG-Pic.png"
    }
    
    # --- Layer 1: Background ---
    try:
        with urllib.request.urlopen(assets["background"]) as url:
            f = BytesIO(url.read())
        slide.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)
    except Exception as e:
        print(f"Could not download background, using solid fill: {e}")
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(10, 10, 30)
    
    # Animation Manager
    anim_manager = AnimationManager(slide)

    # --- Layer 2: Celestial Bodies & Vehicles ---
    def add_asset(asset_key, left, top, width):
        try:
            with urllib.request.urlopen(assets[asset_key]) as url:
                f = BytesIO(url.read())
            return slide.shapes.add_picture(f, left, top, width=width)
        except Exception as e:
            print(f"Could not download {asset_key}, skipping: {e}")
            return None

    earth = add_asset("earth", Inches(-4), Inches(0.5), Inches(8))
    moon = add_asset("moon", Inches(11), Inches(4), Inches(3))
    spaceship = add_asset("spaceship", Inches(-3), Inches(-1), Inches(3))
    
    # --- Layer 3: Typography ---
    title_box = slide.shapes.add_textbox(Inches(5.5), Inches(2.5), Inches(7), Inches(2))
    title_tf = title_box.text_frame
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text
    title_p.font.name = 'Arial Black'
    title_p.font.size = Pt(80)
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    subtitle_box = slide.shapes.add_textbox(Inches(5.8), Inches(4.5), Inches(6.5), Inches(1))
    subtitle_tf = subtitle_box.text_frame
    subtitle_p = subtitle_tf.paragraphs[0]
    subtitle_p.text = subtitle_text
    subtitle_p.font.name = 'Arial'
    subtitle_p.font.size = Pt(24)
    subtitle_p.font.color.rgb = RGBColor(220, 220, 220)

    # --- Layer 4: Meteor Effects ---
    for i in range(15):
        meteor = slide.shapes.add_shape(1, Inches(random.uniform(3, 14)), Inches(-1), Inches(1.5), Inches(0.02))
        meteor.rotation = 155
        fill = meteor.fill
        fill.gradient()
        fill.gradient_stops[0].color.rgb = RGBColor(255, 255, 255)
        fill.gradient_stops[0].position = 0.0
        fill.gradient_stops[1].color.rgb = RGBColor(255, 255, 255)
        fill.gradient_stops[1].alpha = 0
        fill.gradient_stops[1].position = 1.0
        fill.gradient_angle = 90
        meteor.line.fill.background()
        
        delay = random.uniform(0.5, 4.0)
        duration = random.uniform(2.5, 4.0)
        anim_manager.add_motion_path(meteor, "M 0 0 L -0.5 0.5", dur=duration, delay=delay)


    # --- Animation Setup ---
    if earth:
        anim_manager.add_motion_path(earth, "M 0 0 L 0.3 0", dur=8)
        anim_manager.add_spin(earth, dur=60)
    
    if moon:
        anim_manager.add_motion_path(moon, "M 0 0 L -0.4 -0.2", dur=10)
        anim_manager.add_spin(moon, dur=30, delay=1)
        
    if spaceship:
        anim_manager.add_motion_path(spaceship, "M 0 0 L 1.2 0.6", dur=6, delay=0.2)
    
    # Save the presentation
    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# if __name__ == '__main__':
#     output_file = "cosmic_voyage_animation.pptx"
#     create_slide(output_file)
#     print(f"Presentation saved to {output_file}")
#     # On Windows, you can open it directly
#     # import os
#     # os.startfile(output_file)

