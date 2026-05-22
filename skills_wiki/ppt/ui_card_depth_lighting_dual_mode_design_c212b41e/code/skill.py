def create_slide(
    output_pptx_path: str,
    title_text: str = "Sharp heading",
    body_text: str = "A muted shade for the rest of the text element, so it's not always in your face.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the dual-mode UI Card lighting and depth effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml import OxmlElement
    from pptx.oxml.ns import qn

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- XML Helpers ---
    def set_corner_radius(shape, radius_ratio=0.08):
        """Injects XML to set exact corner rounding."""
        prstGeom = shape.element.spPr.prstGeom
        avLst = prstGeom.find(qn('a:avLst'))
        if avLst is None:
            avLst = OxmlElement('a:avLst')
            prstGeom.insert(0, avLst)
        gd = OxmlElement('a:gd')
        gd.set('name', 'adj')
        gd.set('fmla', f'val {int(radius_ratio * 100000)}')
        avLst.append(gd)

    def add_soft_shadow(shape, blur_rad_pt=20, dist_pt=10, dir_deg=90, alpha_pct=15):
        """Injects XML to create modern, soft UI drop shadows."""
        spPr = shape.element.spPr
        effectLst = spPr.find(qn('a:effectLst'))
        if effectLst is None:
            effectLst = OxmlElement('a:effectLst')
            spPr.append(effectLst)
            
        outerShdw = OxmlElement('a:outerShdw')
        outerShdw.set('blurRad', str(int(blur_rad_pt * 12700)))
        outerShdw.set('dist', str(int(dist_pt * 12700)))
        outerShdw.set('dir', str(int(dir_deg * 60000)))
        outerShdw.set('algn', 'ctr')
        
        srgbClr = OxmlElement('a:srgbClr')
        srgbClr.set('val', '000000')  # base black shadow
        alpha = OxmlElement('a:alpha')
        alpha.set('val', str(int(alpha_pct * 1000)))
        srgbClr.append(alpha)
        
        outerShdw.append(srgbClr)
        effectLst.append(outerShdw)

    def style_text_frame(shape, top_pad=0.3, left_pad=0.3):
        """Adds CSS-like padding to the text container."""
        tf = shape.text_frame
        tf.margin_top = Inches(top_pad)
        tf.margin_bottom = Inches(top_pad)
        tf.margin_left = Inches(left_pad)
        tf.margin_right = Inches(left_pad)

    # --- 1. Dual Canvas Backgrounds ---
    # Left side (Dark Mode)
    left_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(6.666), Inches(7.5))
    left_bg.fill.solid()
    left_bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
    left_bg.line.fill.background()

    # Right side (Light Mode)
    right_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.666), 0, Inches(6.666), Inches(7.5))
    right_bg.fill.solid()
    right_bg.fill.fore_color.rgb = RGBColor(230, 230, 230)
    right_bg.line.fill.background()

    # --- 2. Build Card Helper ---
    def build_ui_card(x, y, is_dark=True):
        card_w, card_h = Inches(4.5), Inches(3.5)
        
        # Color Theme Definitions based on HSL lightness logic
        if is_dark:
            c_bg         = RGBColor(13, 13, 13)    # hsl(0,0,5%)
            c_border     = RGBColor(26, 26, 26)    # hsl(0,0,10%)
            c_highlight  = RGBColor(153, 153, 153) # hsl(0,0,60%)
            c_text       = RGBColor(242, 242, 242) # hsl(0,0,95%)
            c_text_muted = RGBColor(178, 178, 178) # hsl(0,0,70%)
            c_btn_bg     = RGBColor(26, 26, 26)    # hsl(0,0,10%)
            shadow_alpha = 40 # Stronger shadow for dark mode
        else:
            c_bg         = RGBColor(242, 242, 242) # hsl(0,0,95%)
            c_border     = RGBColor(217, 217, 217) # hsl(0,0,85%)
            c_highlight  = RGBColor(255, 255, 255) # hsl(0,0,100%) - catches max light
            c_text       = RGBColor(13, 13, 13)    # hsl(0,0,5%)
            c_text_muted = RGBColor(77, 77, 77)    # hsl(0,0,30%)
            c_btn_bg     = RGBColor(255, 255, 255) # hsl(0,0,100%)
            shadow_alpha = 15 # Softer shadow for light mode

        # A. Base Card Shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_bg
        card.line.color.rgb = c_border
        card.line.width = Pt(1)
        set_corner_radius(card, 0.08)
        add_soft_shadow(card, blur_rad_pt=25, dist_pt=12, alpha_pct=shadow_alpha)

        # B. CSS 'border-top' Highlight (Light catching effect)
        # We draw a 1px line perfectly over the top edge. 
        # (Inset slightly by corner radius so it doesn't poke out of the curves)
        inset = Inches(0.2)
        highlight = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x + inset, y, x + card_w - inset, y)
        highlight.line.color.rgb = c_highlight
        highlight.line.width = Pt(1)
        
        # C. Typography Container
        text_box = slide.shapes.add_textbox(x, y, card_w, Inches(2.0))
        style_text_frame(text_box, top_pad=0.4, left_pad=0.4)
        
        p_title = text_box.text_frame.add_paragraph()
        p_title.text = title_text
        p_title.font.bold = True
        p_title.font.size = Pt(22)
        p_title.font.color.rgb = c_text
        p_title.font.name = "Segoe UI"
        
        p_space = text_box.text_frame.add_paragraph()
        p_space.font.size = Pt(8) # spacing
        
        p_body = text_box.text_frame.add_paragraph()
        p_body.text = body_text
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = c_text_muted
        p_body.font.name = "Segoe UI"

        # D. Interactive Button Element
        btn_w, btn_h = Inches(1.8), Inches(0.45)
        btn_x = x + Inches(0.4)
        btn_y = y + card_h - Inches(0.85)
        
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, btn_x, btn_y, btn_w, btn_h)
        btn.fill.solid()
        btn.fill.fore_color.rgb = c_btn_bg
        btn.line.color.rgb = c_border
        btn.line.width = Pt(1)
        set_corner_radius(btn, 0.2)
        
        # Button Top Highlight
        btn_hl = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, btn_x + Inches(0.1), btn_y, btn_x + btn_w - Inches(0.1), btn_y)
        btn_hl.line.color.rgb = c_highlight
        btn_hl.line.width = Pt(0.75)

        btn_tf = btn.text_frame
        btn_tf.clear()
        btn_tf.word_wrap = False
        p_btn = btn_tf.paragraphs[0]
        p_btn.text = "Some Action"
        p_btn.alignment = PP_ALIGN.CENTER
        p_btn.font.bold = True
        p_btn.font.size = Pt(11)
        p_btn.font.color.rgb = c_text
        p_btn.font.name = "Segoe UI"
        
    # --- 3. Render Cards ---
    # Centered vertically, spaced evenly horizontally
    card_y = Inches(2.0)
    
    # Render Dark Mode Card (Left)
    build_ui_card(Inches(1.08), card_y, is_dark=True)
    
    # Render Light Mode Card (Right)
    build_ui_card(Inches(7.75), card_y, is_dark=False)

    prs.save(output_pptx_path)
    return output_pptx_path
