import math
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from lxml import etree
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFilter

def _create_radial_gradient_bg(filepath: str, width=1920, height=1080):
    """Generates a moody radial gradient background (dark brown to black)."""
    img = Image.new('RGB', (width, height))
    draw = ImageDraw.Draw(img)
    cx, cy = width / 2, height / 2
    max_radius = math.hypot(cx, cy)
    
    center_color = (89, 44, 18)  # Warm dark brown
    edge_color = (15, 8, 4)      # Near black
    
    for r in range(int(max_radius), 0, -5):
        ratio = r / max_radius
        rc = int(edge_color[0] * ratio + center_color[0] * (1 - ratio))
        gc = int(edge_color[1] * ratio + center_color[1] * (1 - ratio))
        bc = int(edge_color[2] * ratio + center_color[2] * (1 - ratio))
        draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(rc, gc, bc))
        
    img.save(filepath)
    return filepath

def _create_placeholder_assets():
    """Generates transparent PNG assets using PIL to ensure the script runs completely offline."""
    assets = {}
    
    # 1. Main Burger
    img_b = Image.new('RGBA', (500, 500), (0,0,0,0))
    draw_b = ImageDraw.Draw(img_b)
    draw_b.chord([100, 100, 400, 250], 180, 360, fill=(210, 140, 70)) # Top bun
    draw_b.rectangle([90, 250, 410, 280], fill=(50, 180, 50), radius=10) # Lettuce
    draw_b.rectangle([100, 280, 400, 310], fill=(220, 50, 50)) # Tomato
    draw_b.rectangle([110, 310, 390, 360], fill=(90, 50, 20), radius=15) # Patty
    draw_b.chord([110, 350, 390, 450], 0, 180, fill=(200, 130, 60)) # Bottom bun
    assets['burger'] = 'asset_burger.png'
    img_b.save(assets['burger'])
    
    # 2. Tomato Slice
    img_t = Image.new('RGBA', (150, 150), (0,0,0,0))
    draw_t = ImageDraw.Draw(img_t)
    draw_t.ellipse([10, 20, 140, 130], fill=(230, 40, 40))
    draw_t.ellipse([30, 40, 65, 110], fill=(180, 20, 20))
    draw_t.ellipse([85, 40, 120, 110], fill=(180, 20, 20))
    assets['tomato'] = 'asset_tomato.png'
    img_t.save(assets['tomato'])
    
    # 3. Onion Ring
    img_o = Image.new('RGBA', (120, 120), (0,0,0,0))
    draw_o = ImageDraw.Draw(img_o)
    draw_o.ellipse([10, 10, 110, 110], outline=(150, 80, 180), width=12)
    draw_o.ellipse([25, 25, 95, 95], outline=(200, 180, 220), width=4)
    assets['onion'] = 'asset_onion.png'
    img_o.save(assets['onion'])
    
    # 4. Lettuce Leaf
    img_l = Image.new('RGBA', (200, 100), (0,0,0,0))
    draw_l = ImageDraw.Draw(img_l)
    draw_l.polygon([(10, 50), (40, 10), (100, 30), (160, 10), (190, 60), (150, 90), (80, 80)], fill=(70, 200, 70))
    assets['leaf'] = 'asset_leaf.png'
    img_l.save(assets['leaf'])
    
    return assets

def apply_stroke_only_to_run(run, stroke_color="FFFFFF", stroke_width_emu=38100):
    """
    Uses lxml to inject <a:noFill/> and <a:ln> to make text transparent with a stroke outline.
    38100 EMU = 3 pt.
    """
    rPr = run._r.get_or_add_rPr()
    
    # Remove existing fill if present
    for child in list(rPr):
        if child.tag.endswith('Fill'):
            rPr.remove(child)
            
    # Add No Fill
    etree.SubElement(rPr, qn('a:noFill'))
    
    # Add Stroke (Line)
    ln = etree.SubElement(rPr, qn('a:ln'))
    ln.set('w', str(stroke_width_emu))
    
    solidFill = etree.SubElement(ln, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', stroke_color)

def create_slide(output_pptx_path: str, **kwargs) -> str:
    """
    Create a PPTX file reproducing the 3D Exploded Layering with Stroke Typography effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # --- 1. Background Layer ---
    bg_path = _create_radial_gradient_bg("temp_bg.png")
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    
    # Prepare assets
    assets = _create_placeholder_assets()
    
    # --- 2. Distant Elements Layer (Behind Text) ---
    # Placing a few ingredients in the far back
    pic = slide.shapes.add_picture(assets['onion'], Inches(2), Inches(1), width=Inches(1.2))
    pic.rotation = -20
    pic = slide.shapes.add_picture(assets['tomato'], Inches(9.5), Inches(0.5), width=Inches(1.8))
    pic.rotation = 45
    pic = slide.shapes.add_picture(assets['leaf'], Inches(10), Inches(5), width=Inches(2))
    pic.rotation = 110

    # --- 3. Typography Layer (Stroke Only) ---
    def add_hollow_text(text, top_inch):
        txBox = slide.shapes.add_textbox(Inches(0), Inches(top_inch), prs.slide_width, Inches(2.5))
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = 'Arial Black'
        run.font.size = Pt(130)
        
        # lxml hack: apply white outline, remove fill, add character spacing (spc)
        apply_stroke_only_to_run(run, stroke_color="FFFFFF", stroke_width_emu=25400) # 2pt
        run._r.get_or_add_rPr().set('spc', '20000') # Widen letter spacing
        
    add_hollow_text("FRESH", 0.5)
    add_hollow_text("BURGER", 4.2)
    
    # --- 4. Main Subject Layer ---
    # Center the burger over the text
    burger_width = Inches(4.5)
    burger_left = (prs.slide_width - burger_width) / 2
    burger_top = Inches(2.0)
    slide.shapes.add_picture(assets['burger'], burger_left, burger_top, width=burger_width)
    
    # --- 5. Foreground Elements Layer (In Front of Subject & Text) ---
    pic = slide.shapes.add_picture(assets['tomato'], Inches(2.5), Inches(4.5), width=Inches(2))
    pic.rotation = -35
    pic = slide.shapes.add_picture(assets['leaf'], Inches(1.5), Inches(3), width=Inches(2.5))
    pic.rotation = -15
    pic = slide.shapes.add_picture(assets['onion'], Inches(8.5), Inches(4.5), width=Inches(1.5))
    pic.rotation = 60
    pic = slide.shapes.add_picture(assets['leaf'], Inches(8.0), Inches(2.0), width=Inches(1.8))
    pic.rotation = -40

    # Save and clean up
    prs.save(output_pptx_path)
    
    # Clean up temp assets
    for f in [bg_path] + list(assets.values()):
        if os.path.exists(f):
            os.remove(f)
            
    return output_pptx_path

# To test:
# create_slide("exploded_burger_effect.pptx")
