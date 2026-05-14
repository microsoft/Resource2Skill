def create_slide(
    output_pptx_path: str,
    title_text: str = "Title Here",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit,\nsed do eiusmod tempor incididunt ut labore et dolore magna aliqua.",
    bg_color: tuple = (255, 255, 255),
    shadow_color: tuple = (105, 25, 45),  # Dark Burgundy
    front_color: tuple = (215, 150, 160), # Soft Rose/Pink
    text_color: tuple = (50, 20, 25),     # Dark font color
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Dual-Tone Layered Callout Offset" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml

    # Initialize presentation and blank slide
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # 6 is usually a blank slide layout
    slide = prs.slides.add_slide(slide_layout)

    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 1: Top Title ===
    title_box = slide.shapes.add_textbox(Inches(2), Inches(0.8), Inches(9.333), Inches(1))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.alignment = PP_ALIGN.CENTER
    run_title = p_title.runs[0]
    run_title.font.size = Pt(40)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(50, 50, 50)

    # Base coordinates & dimensions for the callout
    w = Inches(6.0)
    h = Inches(3.2)
    base_left = Inches(3.66)
    base_top = Inches(2.2)
    offset = Inches(0.25)

    # === Layer 2: Shadowed Back Callout ===
    # Placed with offset (bottom-right)
    back_shape = slide.shapes.add_shape(
        MSO_SHAPE.WEDGE_RECT_CALLOUT,
        base_left + offset, base_top + offset, w, h
    )
    back_shape.fill.solid()
    back_shape.fill.fore_color.rgb = RGBColor(*shadow_color)
    back_shape.line.fill.solid()
    back_shape.line.color.rgb = RGBColor(*shadow_color) # Hide border

    # Inject exact drop shadow using lxml
    shadow_xml = """
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:outerShdw blurRad="50000" dist="40000" dir="2700000" algn="tl" rotWithShape="0">
            <a:srgbClr val="000000">
                <a:alpha val="40000"/>
            </a:srgbClr>
        </a:outerShdw>
    </a:effectLst>
    """
    back_shape.element.spPr.append(parse_xml(shadow_xml))

    # === Layer 3: Front Content Callout ===
    # Placed at true base coordinates
    front_shape = slide.shapes.add_shape(
        MSO_SHAPE.WEDGE_RECT_CALLOUT,
        base_left, base_top, w, h
    )
    front_shape.fill.solid()
    front_shape.fill.fore_color.rgb = RGBColor(*front_color)
    front_shape.line.fill.solid()
    front_shape.line.color.rgb = RGBColor(*front_color) # Hide border

    # Add and format text inside the front callout
    tf_front = front_shape.text_frame
    tf_front.word_wrap = True
    tf_front.margin_left = Inches(0.4)
    tf_front.margin_right = Inches(0.4)
    tf_front.margin_top = Inches(0.6)
    tf_front.margin_bottom = Inches(0.6)

    p_front = tf_front.paragraphs[0]
    p_front.text = body_text
    p_front.alignment = PP_ALIGN.LEFT
    run_front = p_front.runs[0]
    run_front.font.size = Pt(20)
    run_front.font.color.rgb = RGBColor(*text_color)

    # === Layer 4: Lower Supplementary Text ===
    footer_box = slide.shapes.add_textbox(Inches(2.66), Inches(6.0), Inches(8), Inches(1))
    tf_foot = footer_box.text_frame
    tf_foot.word_wrap = True
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 2
    p_foot.alignment = PP_ALIGN.CENTER
    p_foot.font.size = Pt(14)
    p_foot.font.color.rgb = RGBColor(100, 100, 100)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
