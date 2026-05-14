def create_slide(
    output_pptx_path: str,
    title_text: str = "Four Step Process Flow",
    body_text: str = "",
    bg_palette: str = "light",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Modern Floating Orb Sequence' visual effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml
    
    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide
    
    # Optional Background Fill (Subtle off-white)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(248, 249, 250)

    # 2. Define Palette & Layout Parameters
    colors = [
        RGBColor(0, 173, 181),   # Teal
        RGBColor(57, 62, 70),    # Dark Gray/Blue
        RGBColor(255, 87, 34),   # Deep Orange
        RGBColor(144, 55, 73)    # Burgundy
    ]
    
    num_steps = 4
    canvas_width = 13.333
    margins = 1.5
    usable_width = canvas_width - (margins * 2)
    step_width = usable_width / num_steps
    
    orb_size = 1.6  # Diameter in inches
    line_y = 3.5    # Vertical center for the flow line
    
    # Add Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(33, 37, 41)
    p.alignment = PP_ALIGN.CENTER
    
    # 3. Draw Connecting Axis Line (Draw first so it sits behind orbs)
    # Start center of first orb, end center of last orb
    start_x = margins + (step_width / 2)
    end_x = margins + usable_width - (step_width / 2)
    
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(start_x), Inches(line_y - 0.05), 
        Inches(end_x - start_x), Inches(0.1)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(222, 226, 230)
    line.line.fill.background() # No border
    
    # Helper XML snippet for Drop Shadow
    shadow_xml = """
    <a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <a:outerShdw blurRad="254000" dist="127000" dir="5400000" algn="tl" rotWithShape="0">
            <a:srgbClr val="000000">
                <a:alpha val="20000"/>
            </a:srgbClr>
        </a:outerShdw>
    </a:effectLst>
    """
    
    # 4. Generate Sequence Steps
    for i in range(num_steps):
        center_x = margins + (i * step_width) + (step_width / 2)
        
        # --- Create Orb ---
        orb_x = center_x - (orb_size / 2)
        orb_y = line_y - (orb_size / 2)
        
        orb = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            Inches(orb_x), Inches(orb_y), 
            Inches(orb_size), Inches(orb_size)
        )
        orb.fill.solid()
        orb.fill.fore_color.rgb = colors[i]
        orb.line.fill.background() # Remove border
        
        # Apply LXML Drop Shadow
        shadow_element = parse_xml(shadow_xml)
        orb.spPr.append(shadow_element)
        
        # --- Add Number Text to Orb ---
        tf = orb.text_frame
        tf.clear() # clear default paragraph
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.name = "Arial"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        
        # --- Add Step Title ---
        title_width = 2.5
        title_x = center_x - (title_width / 2)
        title_y = line_y + (orb_size / 2) + 0.3
        
        step_title = slide.shapes.add_textbox(
            Inches(title_x), Inches(title_y), 
            Inches(title_width), Inches(0.5)
        )
        stf = step_title.text_frame
        stf.word_wrap = True
        p_title = stf.paragraphs[0]
        p_title.text = f"Phase {i+1} Setup"
        p_title.font.bold = True
        p_title.font.size = Pt(18)
        p_title.font.color.rgb = colors[i]
        p_title.alignment = PP_ALIGN.CENTER
        
        # --- Add Step Body Text ---
        body_y = title_y + 0.4
        step_body = slide.shapes.add_textbox(
            Inches(title_x), Inches(body_y), 
            Inches(title_width), Inches(1.5)
        )
        btf = step_body.text_frame
        btf.word_wrap = True
        p_body = btf.paragraphs[0]
        p_body.text = "Describe the key actions and deliverables required to successfully complete this phase of the process."
        p_body.font.size = Pt(12)
        p_body.font.color.rgb = RGBColor(108, 117, 125)
        p_body.alignment = PP_ALIGN.CENTER

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
