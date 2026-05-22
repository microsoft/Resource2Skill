import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR

def create_slide_declarative_title_card(
    output_pptx_path: str,
    title_text: str = "#1",
    body_text: str = "GOOD PRESENTATION\nSLIDES ARE CLEAR",
    bg_color: tuple = (27, 43, 75),  # Dark Navy Blue from video
    font_color: tuple = (255, 255, 255), # White
    font_family: str = "Montserrat ExtraBold", # A suitable modern, bold font
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with the "Declarative Title Card" style.

    This style features large, bold, centered text on a solid, high-contrast
    background to deliver a clear and impactful message.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        title_text (str): The smaller text, often a number or label, on top.
        body_text (str): The main message. Use '\n' for line breaks.
        bg_color (tuple): RGB tuple for the slide background.
        font_color (tuple): RGB tuple for the text color.
        font_family (str): The font to use for the text.

    Returns:
        str: The path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Text & Content ===
    # Use a single textbox spanning the entire slide for easy centering.
    left = Inches(0.5)
    top = Inches(0)
    width = prs.slide_width - Inches(1.0)
    height = prs.slide_height
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE # Center vertically
    tf.word_wrap = True

    # Title paragraph (e.g., "#1")
    p_title = tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = font_family
    p_title.font.size = Pt(66)
    p_title.font.bold = True # Explicitly set bold
    p_title.font.color.rgb = RGBColor(*font_color)
    p_title.alignment = PP_ALIGN.CENTER

    # Body paragraph (main message)
    p_body = tf.add_paragraph()
    p_body.text = body_text
    p_body.font.name = font_family
    p_body.font.size = Pt(88)
    p_body.font.bold = True # Explicitly set bold
    p_body.font.color.rgb = RGBColor(*font_color)
    p_body.alignment = PP_ALIGN.CENTER
    p_body.space_before = Pt(12) # Add space between title and body

    # Ensure the output directory exists
    output_dir = os.path.dirname(output_pptx_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
        
    prs.save(output_pptx_path)
    return output_pptx_path

