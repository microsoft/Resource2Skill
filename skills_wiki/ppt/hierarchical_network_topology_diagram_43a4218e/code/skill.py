import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def _create_icon_router(filename):
    """Draws a flat-design Router icon using PIL"""
    img = Image.new('RGBA', (120, 80), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    # Antennas
    draw.line([30, 20, 20, 0], fill=(52, 73, 94, 255), width=4)
    draw.line([90, 20, 100, 0], fill=(52, 73, 94, 255), width=4)
    # Box body
    draw.rounded_rectangle([10, 20, 110, 70], radius=10, fill=(41, 128, 185, 255), outline=(31, 97, 141, 255), width=3)
    # Status Lights
    draw.ellipse([45, 40, 50, 45], fill=(46, 204, 113, 255))
    draw.ellipse([55, 40, 60, 45], fill=(46, 204, 113, 255))
    draw.ellipse([65, 40, 70, 45], fill=(46, 204, 113, 255))
    img.save(filename)
    return filename

def _create_icon_switch(filename):
    """Draws a flat-design Network Switch icon using PIL"""
    img = Image.new('RGBA', (150, 60), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    # Switch body
    draw.rectangle([5, 10, 145, 50], fill=(52, 73, 94, 255), outline=(44, 62, 80, 255), width=3)
    # Ports
    for i in range(8):
        x = 18 + i * 15
        draw.rectangle([x, 25, x + 8, 35], fill=(46, 204, 113, 255))
    img.save(filename)
    return filename

def _create_icon_pc(filename):
    """Draws a flat-design Workstation/PC icon using PIL"""
    img = Image.new('RGBA', (100, 100), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    # Monitor frame
    draw.rounded_rectangle([10, 10, 90, 70], radius=5, fill=(149, 165, 166, 255), outline=(127, 140, 141, 255), width=2)
    # Screen display
    draw.rectangle([15, 15, 85, 65], fill=(236, 240, 241, 255))
    # Code/Text lines on screen
    draw.line([20, 25, 40, 25], fill=(189, 195, 199, 255), width=3)
    draw.line([20, 35, 60, 35], fill=(189, 195, 199, 255), width=3)
    draw.line([20, 45, 50, 45], fill=(189, 195, 199, 255), width=3)
    # Stand and Base
    draw.rectangle([45, 70, 55, 90], fill=(127, 140, 141, 255))
    draw.rounded_rectangle([30, 90, 70, 95], radius=2, fill=(127, 140, 141, 255))
    img.save(filename)
    return filename

def _draw_orthogonal_line(slide, x1, y1, x2, y2):
    """Draws an exact 90-degree elbow path between two points"""
    builder = slide.shapes.build_freeform(Inches(x1), Inches(y1))
    
    if abs(x1 - x2) < 0.01:
        # Perfectly vertical, no elbow needed
        builder.add_line_segments([(Inches(x2), Inches(y2))])
    else:
        # Calculate midpoint for the horizontal elbow segment
        mid_y = y1 + (y2 - y1) / 2.0
        builder.add_line_segments([
            (Inches(x1), Inches(mid_y)),
            (Inches(x2), Inches(mid_y)),
            (Inches(x2), Inches(y2))
        ])
        
    shape = builder.convert_to_shape()
    shape.line.color.rgb = RGBColor(80, 80, 80)
    shape.line.width = Pt(1.5)

def _add_label(slide, text, cx, cy_top, width_in=1.5):
    """Adds formatted text (Verdana 12pt) centered below an icon"""
    left = Inches(cx - width_in / 2)
    top = Inches(cy_top)
    width = Inches(width_in)
    height = Inches(0.5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    
    run = p.runs[0]
    run.font.name = 'Verdana'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(44, 62, 80)

def create_slide(
    output_pptx_path: str,
    title_text: str = "Corporate Network Architecture",
    **kwargs,
) -> str:
    """
    Creates a PPTX file containing a hierarchial network topology diagram
    with custom icons and orthogonal routing.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Generate Visual Assets (Icons) ===
    router_path = "temp_router_icon.png"
    switch_path = "temp_switch_icon.png"
    pc_path = "temp_pc_icon.png"
    
    _create_icon_router(router_path)
    _create_icon_switch(switch_path)
    _create_icon_pc(pc_path)
    
    icon_map = {
        'router': router_path,
        'switch': switch_path,
        'pc': pc_path
    }

    # === Add Title ===
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Verdana'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(44, 62, 80)

    # === Define Network Topology Layout ===
    # Hardcoded coordinates to ensure perfect symmetry and clustering
    nodes = [
        # Tier 1
        {'id': 'R1', 'type': 'router', 'label': 'Core Router', 'x': 6.666, 'y': 1.5, 'w': 1.2, 'h': 0.8},
        # Tier 2
        {'id': 'S1', 'type': 'switch', 'label': 'Switch A', 'x': 3.666, 'y': 3.5, 'w': 1.5, 'h': 0.6},
        {'id': 'S2', 'type': 'switch', 'label': 'Switch B', 'x': 9.666, 'y': 3.5, 'w': 1.5, 'h': 0.6},
        # Tier 3 (Clustered under Switch A)
        {'id': 'PC1', 'type': 'pc', 'label': 'Workstation 1', 'x': 1.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
        {'id': 'PC2', 'type': 'pc', 'label': 'Workstation 2', 'x': 3.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
        {'id': 'PC3', 'type': 'pc', 'label': 'Workstation 3', 'x': 5.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
        # Tier 3 (Clustered under Switch B)
        {'id': 'PC4', 'type': 'pc', 'label': 'Workstation 4', 'x': 7.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
        {'id': 'PC5', 'type': 'pc', 'label': 'Workstation 5', 'x': 9.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
        {'id': 'PC6', 'type': 'pc', 'label': 'Workstation 6', 'x': 11.666, 'y': 5.5, 'w': 1.0, 'h': 1.0},
    ]

    links = [
        ('R1', 'S1'), ('R1', 'S2'),
        ('S1', 'PC1'), ('S1', 'PC2'), ('S1', 'PC3'),
        ('S2', 'PC4'), ('S2', 'PC5'), ('S2', 'PC6')
    ]

    # Convert node list to dictionary for quick lookup during link drawing
    node_dict = {n['id']: n for n in nodes}

    # === Draw Orthogonal Connectors First (so they sit behind the icons) ===
    for source_id, target_id in links:
        source = node_dict[source_id]
        target = node_dict[target_id]
        
        # Calculate exact connector start (bottom of parent) and end (top of child)
        start_x = source['x']
        start_y = source['y'] + (source['h'] / 2)
        
        end_x = target['x']
        end_y = target['y'] - (target['h'] / 2)
        
        _draw_orthogonal_line(slide, start_x, start_y, end_x, end_y)

    # === Place Icons and Labels ===
    for node in nodes:
        img_path = icon_map[node['type']]
        left = Inches(node['x'] - node['w'] / 2)
        top = Inches(node['y'] - node['h'] / 2)
        
        # Insert custom PIL icon
        slide.shapes.add_picture(img_path, left, top, Inches(node['w']), Inches(node['h']))
        
        # Add Verdana label beneath the icon
        label_y = node['y'] + (node['h'] / 2) + 0.1  # Slight padding below icon
        _add_label(slide, node['label'], node['x'], label_y)

    # === Cleanup Temporary Images ===
    for path in icon_map.values():
        if os.path.exists(path):
            os.remove(path)

    prs.save(output_pptx_path)
    return output_pptx_path
