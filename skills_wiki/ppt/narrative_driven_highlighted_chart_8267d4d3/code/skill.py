def create_slide(
    output_pptx_path: str,
    title_text: str = "Monthly Sales Analysis",
    body_text: str = "",
    bg_palette: str = "white", 
    accent_color: tuple = (220, 53, 69),  # Crimson Red for highlight
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Narrative-Driven Highlighted Chart" visual effect.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank layout to build from scratch
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Title ===
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(40, 40, 40)

    # === Layer 2: Chart Data & Generation ===
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    # Simulated sales data showing a distinct drop in July
    sales = [560000, 580000, 620000, 610000, 690000, 750000, 380000, 650000, 680000, 710000, 780000, 820000]

    chart_data = CategoryChartData()
    chart_data.categories = months
    chart_data.add_series("Sales", sales)

    # Position chart in the middle
    x, y, cx, cy = Inches(0.8), Inches(1.3), Inches(11.7), Inches(4.5)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart

    # === Layer 3: Declutter & Format Chart ===
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = False
    
    # Format axes fonts
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.color.rgb = RGBColor(100, 100, 100)
    chart.value_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.color.rgb = RGBColor(100, 100, 100)

    # Thicker bars (reduce gap width) & add data labels
    plot = chart.plots[0]
    plot.gap_width = 80  
    plot.has_data_labels = True
    
    # Format data labels
    data_labels = plot.data_labels
    data_labels.font.size = Pt(10)
    data_labels.font.color.rgb = RGBColor(80, 80, 80)
    data_labels.number_format = '$#,##0'

    # === Layer 4: Color Storytelling (The Focal Highlight) ===
    highlight_index = 6  # Index for 'Jul'
    highlight_rgb = RGBColor(*accent_color)
    # Simulated 54% transparency of standard blue on a white background
    muted_rgb = RGBColor(176, 196, 222) 

    series = chart.series[0]
    for idx, point in enumerate(series.points):
        fill = point.format.fill
        fill.solid()
        if idx == highlight_index:
            fill.fore_color.rgb = highlight_rgb
            # Bold the data label for the highlight
            try:
                point.data_label.font.bold = True
                point.data_label.font.color.rgb = highlight_rgb
            except:
                pass # Safe fallback if individual data label formatting isn't exposed in current version
        else:
            fill.fore_color.rgb = muted_rgb

    # === Layer 5: Key Takeaway Box ===
    # Creates a grounded, bordered box tied visually to the highlight color
    takeaway_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.8)
    )
    takeaway_bg.fill.solid()
    takeaway_bg.fill.fore_color.rgb = RGBColor(250, 250, 250) # Very light off-white
    takeaway_bg.line.color.rgb = highlight_rgb # Border matches the anomaly bar
    takeaway_bg.line.width = Pt(1.5)

    tf_takeaway = takeaway_bg.text_frame
    tf_takeaway.word_wrap = True
    p_takeaway = tf_takeaway.paragraphs[0]
    p_takeaway.alignment = PP_ALIGN.CENTER
    
    # Bold identifier
    run1 = p_takeaway.add_run()
    run1.text = "Key Takeaway: "
    run1.font.bold = True
    run1.font.size = Pt(16)
    run1.font.color.rgb = RGBColor(0, 0, 0)

    # Contextual explanation
    run2 = p_takeaway.add_run()
    if body_text:
        run2.text = body_text
    else:
        run2.text = "Sales dropped significantly in July due to server outages, requiring immediate mitigation strategies."
    run2.font.size = Pt(16)
    run2.font.color.rgb = RGBColor(60, 60, 60)

    prs.save(output_pptx_path)
    return output_pptx_path
