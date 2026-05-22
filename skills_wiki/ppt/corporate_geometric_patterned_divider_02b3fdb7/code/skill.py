def create_slide(
    output_pptx_path: str,
    title_text: str = "Our Proposed Approach\n& Fees",
    body_text: str = "www.YourCompany.com",
    theme: str = "blue",  # Options: 'blue' or 'red'
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Corporate Geometric Patterned Divider' visual effect.
    Returns: path to the saved PPTX file.
    """
    import os
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # === Layer 1 & 2: Background & Geometric Visual Effect via PIL ===
    # Establish theme colors
    if theme.lower() == "red":
        bg_color = (220, 74, 93, 255)      # Crimson Red
        accent_color = (36, 68, 90, 255)   # Dark Blue
        dark_color = (180, 50, 70, 255)    # Deeper Red
    else:
        bg_color = (36, 68, 90, 255)       # Dark Blue
        accent_color = (220, 74, 93, 255)  # Crimson Red
        dark_color = (20, 45, 65, 255)     # Deeper Blue

    canvas_width, canvas_height = 1920, 1080
    base = Image.new('RGBA', (canvas_width, canvas_height), bg_color)
    
    # 1. Create horizontal hatch pattern (stripes) across a transparent canvas
    pattern = Image.new('RGBA', (canvas_width, canvas_height), (0, 0, 0, 0))
    p_draw = ImageDraw.Draw(pattern)
    for y in range(0, canvas_height, 12):
        p_draw.line([(0, y), (canvas_width, y)], fill=(255, 255, 255, 25), width=2)
        
    # 2. Apply pattern using triangular masks to isolate the stripes
    # Massive triangle emerging from bottom-right
    mask1 = Image.new('L', (canvas_width, canvas_height), 0)
    ImageDraw.Draw(mask1).polygon([(1000, canvas_height), (canvas_width, 150), (canvas_width, canvas_height)], fill=255)
    base.paste(pattern, mask=mask1)
    
    # Medium triangle emerging from top-right
    mask2 = Image.new('L', (canvas_width, canvas_height), 0)
    ImageDraw.Draw(mask2).polygon([(1300, 0), (canvas_width, 0), (canvas_width, 600)], fill=255)
    base.paste(pattern, mask=mask2)
    
    # 3. Add solid, semi-transparent overlapping triangles for depth
    solid_layer = Image.new('RGBA', (canvas_width, canvas_height), (0, 0, 0, 0))
    s_draw = ImageDraw.Draw(solid_layer)
    
    # Darker base shadow triangle overlapping the bottom edge
    s_draw.polygon([(700, canvas_height), (1300, canvas_height), (1000, 750)], fill=dark_color)
    
    # Bright semi-transparent white triangle on the right edge
    s_draw.polygon([(1600, canvas_height), (canvas_width, 700), (canvas_width, canvas_height)], fill=(255, 255, 255, 15))
    
    # Bold accent color triangle pointing left (mid-right area)
    s_draw.polygon([(1350, 450), (1500, 300), (1500, 600)], fill=accent_color)
    
    # Tiny floating decorative triangles (adds tech/consulting detail)
    s_draw.polygon([(900, 250), (940, 210), (940, 290)], fill=accent_color)
    s_draw.polygon([(1100, 850), (1150, 800), (1180, 880)], fill=(255, 255, 255, 60))
    s_draw.polygon([(1600, 150), (1630, 120), (1660, 160)], fill=(255, 255, 255, 40))

    # Composite shapes over background
    base = Image.alpha_composite(base, solid_layer)
    
    temp_bg = "temp_bg_geometric.png"
    base.save(temp_bg)

    # === Layer 3: Presentation Layout & Text via python-pptx ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Insert generated geometric background
    slide.shapes.add_picture(temp_bg, 0, 0, width=Inches(13.333), height=Inches(7.5))

    # Add Left Vertical Anchor Line
    anchor_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(1.2), Inches(2.6), Inches(0.06), Inches(2.2)
    )
    anchor_line.fill.solid()
    anchor_line.fill.fore_color.rgb = RGBColor(255, 255, 255)
    anchor_line.line.fill.background() # Remove border

    # Add Title Box
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(7.0), Inches(2.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = 'Arial'

    # Add Subtitle / Body text
    tx_box2 = slide.shapes.add_textbox(Inches(1.5), Inches(4.8), Inches(6.0), Inches(1.0))
    tf2 = tx_box2.text_frame
    p2 = tf2.add_paragraph()
    p2.text = body_text
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(200, 200, 200)
    p2.font.name = 'Arial'
    
    # Cleanup temp image
    if os.path.exists(temp_bg):
        os.remove(temp_bg)

    prs.save(output_pptx_path)
    return output_pptx_path
