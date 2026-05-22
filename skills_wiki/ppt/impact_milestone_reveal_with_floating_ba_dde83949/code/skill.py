def create_slide(
    output_pptx_path: str,
    title_text: str = "LAUNCHING DATE",
    badge_subtext: str = "05 DEC",
    badge_maintext: str = "2024",
    accent_color: tuple = (244, 176, 4),  # Vivid Yellow/Orange
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Impact Milestone Reveal with Floating Badge" effect.
    """
    import os
    import urllib.request
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shape import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # Initialize presentation (16:9 aspect ratio)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Image ===
    # Attempt to download a high-quality cityscape image, fallback to a PIL gradient
    bg_img_path = "temp_bg.jpg"
    try:
        req = urllib.request.Request(
            "https://images.unsplash.com/photo-1449824913935-59a10b8d2000?q=80&w=1920&auto=format&fit=crop",
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req) as response:
            img_data = response.read()
            with open(bg_img_path, "wb") as f:
                f.write(img_data)
    except Exception:
        # Fallback: Create a dark moody gradient if download fails
        fallback_img = Image.new('RGB', (1920, 1080), color=(40, 45, 50))
        draw = ImageDraw.Draw(fallback_img)
        for y in range(1080):
            r = int(40 - (y / 1080) * 20)
            g = int(45 - (y / 1080) * 20)
            b = int(50 - (y / 1080) * 20)
            draw.line([(0, y), (1920, y)], fill=(r, g, b))
        fallback_img.save(bg_img_path)

    # Insert background full-bleed
    slide.shapes.add_picture(bg_img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Translucent Overlay Band ===
    # Using PIL to create a transparent PNG ensures perfect cross-platform rendering in PPTX
    band_path = "temp_band.png"
    # Dark charcoal with ~65% opacity
    band_img = Image.new('RGBA', (100, 100), (25, 25, 28, 165)) 
    band_img.save(band_path)
    
    # Position: Middle of the screen
    band_y = Inches(2.25)
    band_height = Inches(2.75)
    slide.shapes.add_picture(band_path, 0, band_y, prs.slide_width, band_height)

    # === Layer 3: Typography & Accents ===
    # Thin accent line above the title
    line_width = Inches(1.5)
    line_x = (prs.slide_width - line_width) / 2
    line_y = band_y + Inches(0.4)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_x, line_y, line_width, Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background()

    # Main Title Text
    txBox = slide.shapes.add_textbox(0, line_y + Inches(0.2), prs.slide_width, Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.CENTER

    # === Layer 4: The Floating Badge ===
    # Position: Dead center horizontally, overlapping the bottom edge of the translucent band
    badge_size = Inches(2.5)
    badge_x = (prs.slide_width - badge_size) / 2
    badge_y = band_y + band_height - (badge_size / 2) # Exact middle overlap
    
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, badge_x, badge_y, badge_size, badge_size)
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(*accent_color)
    badge.line.fill.background() # Remove outline

    # Badge Text
    tf_badge = badge.text_frame
    tf_badge.word_wrap = True
    
    # Top text (e.g., Month/Day)
    bp1 = tf_badge.paragraphs[0]
    bp1.text = badge_subtext.upper()
    bp1.font.size = Pt(18)
    bp1.font.bold = True
    bp1.font.color.rgb = RGBColor(34, 34, 34)
    bp1.font.name = "Arial"
    bp1.alignment = PP_ALIGN.CENTER
    bp1.space_after = Pt(0)

    # Bottom text (e.g., Year/Metric)
    bp2 = tf_badge.add_paragraph()
    bp2.text = badge_maintext
    bp2.font.size = Pt(36)
    bp2.font.bold = True
    bp2.font.color.rgb = RGBColor(34, 34, 34)
    bp2.font.name = "Arial"
    bp2.alignment = PP_ALIGN.CENTER

    # Save presentation
    prs.save(output_pptx_path)

    # Cleanup temp files
    if os.path.exists(bg_img_path): os.remove(bg_img_path)
    if os.path.exists(band_path): os.remove(band_path)

    return output_pptx_path
