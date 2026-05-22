import os
import math
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "Core Strategy Pillars",
    body_text: str = "Deconstructing our approach into four distinct, actionable segments.",
    bg_color: tuple = (15, 23, 42),  # Deep Navy
    **kwargs,
) -> str:
    """
    Creates a PPTX file featuring a custom Fragmented Ring Infographic, 
    mimicking the 'Merge Shapes' / Custom SVG style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Create blank slide
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # === Layer 1: Solid Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # Palette for the fragmented ring
    segment_colors = [
        (45, 212, 191, 255),  # Teal
        (56, 189, 248, 255),  # Sky Blue
        (99, 102, 241, 255),  # Indigo
        (168, 85, 247, 255)   # Purple
    ]
    num_segments = len(segment_colors)
    
    # === Layer 2: Generate Fragmented Ring via PIL ===
    # We use a large canvas for anti-aliasing (downsampled in PPTX)
    img_size = 2000
    ring_img = Image.new('RGBA', (img_size, img_size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(ring_img)
    
    # Ring geometry parameters
    center = img_size / 2
    radius = 700
    thickness = 250
    gap_degrees = 8  # Negative space between fragments
    
    # Draw shadows first (on a separate layer to composite)
    shadow_img = Image.new('RGBA', (img_size, img_size), (0, 0, 0, 0))
    shadow_draw = ImageDraw.Draw(shadow_img)
    
    arc_bbox = [center - radius, center - radius, center + radius, center + radius]
    
    # Draw shadow arcs
    for i in range(num_segments):
        start_angle = i * (360 / num_segments) + (gap_degrees / 2)
        end_angle = (i + 1) * (360 / num_segments) - (gap_degrees / 2)
        shadow_draw.arc(arc_bbox, start=start_angle, end=end_angle, fill=(0, 0, 0, 100), width=thickness)
        
    # Blur shadow
    shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(radius=25))
    ring_img.alpha_composite(shadow_img)
    
    # Draw the actual colored segments
    for i, color in enumerate(segment_colors):
        start_angle = i * (360 / num_segments) + (gap_degrees / 2)
        end_angle = (i + 1) * (360 / num_segments) - (gap_degrees / 2)
        draw.arc(arc_bbox, start=start_angle, end=end_angle, fill=color, width=thickness)

    # Save PIL image to memory
    img_io = BytesIO()
    ring_img.save(img_io, format='PNG')
    img_io.seek(0)
    
    # Insert ring image into PPTX
    # Center it on the slide
    ring_display_size = Inches(5)
    ring_x = (prs.slide_width - ring_display_size) / 2
    ring_y = (prs.slide_height - ring_display_size) / 2 + Inches(0.5) # Shift down slightly to leave room for title
    slide.shapes.add_picture(img_io, ring_x, ring_y, width=ring_display_size, height=ring_display_size)

    # === Layer 3: Native PPTX Text Elements ===
    
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(2), Inches(1.2), Inches(9.333), Inches(0.6))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = body_text
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(148, 163, 184) # Light Slate
    p_sub.alignment = PP_ALIGN.CENTER
    
    # Central Anchor Text (inside the donut hole)
    center_box = slide.shapes.add_textbox(
        (prs.slide_width - Inches(2)) / 2, 
        ring_y + (ring_display_size - Inches(1)) / 2, 
        Inches(2), Inches(1)
    )
    tf_center = center_box.text_frame
    p_c = tf_center.paragraphs[0]
    p_c.text = "CORE\nMODEL"
    p_c.font.bold = True
    p_c.font.size = Pt(24)
    p_c.font.color.rgb = RGBColor(255, 255, 255)
    p_c.alignment = PP_ALIGN.CENTER

    # Satellite Text Boxes dynamically calculated via Trigonometry
    # Slide center coordinate for the ring
    cx = prs.slide_width / 2
    cy = ring_y + (ring_display_size / 2)
    
    # Distance from center to place the text boxes
    text_radius = Inches(3.3) 
    
    labels = ["DISCOVERY", "EXECUTION", "ANALYSIS", "ITERATION"]
    
    for i, color in enumerate(segment_colors):
        # Calculate mid-angle of the segment in radians
        # PIL angles: 0 is right, goes clockwise.
        mid_angle_deg = (i * (360 / num_segments)) + (360 / (num_segments * 2))
        mid_angle_rad = math.radians(mid_angle_deg)
        
        # Calculate X, Y. Y is inverted in screen space, but PIL angles match screen space.
        tx = cx + text_radius * math.cos(mid_angle_rad)
        ty = cy + text_radius * math.sin(mid_angle_rad)
        
        # Adjust placement so the text box center aligns with the point
        tb_width = Inches(2.2)
        tb_height = Inches(0.8)
        
        sat_box = slide.shapes.add_textbox(tx - (tb_width/2), ty - (tb_height/2), tb_width, tb_height)
        tf_sat = sat_box.text_frame
        tf_sat.word_wrap = True
        
        # Label Title
        p_sat = tf_sat.paragraphs[0]
        p_sat.text = f"0{i+1}. {labels[i]}"
        p_sat.font.bold = True
        p_sat.font.size = Pt(16)
        p_sat.font.color.rgb = RGBColor(color[0], color[1], color[2]) # Match slice color
        
        # Label Body
        p_desc = tf_sat.add_paragraph()
        p_desc.text = "Strategic phase overview and metrics."
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = RGBColor(200, 200, 200)
        
        # Align text based on position on the screen
        if math.cos(mid_angle_rad) > 0.1:
            p_sat.alignment = PP_ALIGN.LEFT
            p_desc.alignment = PP_ALIGN.LEFT
        elif math.cos(mid_angle_rad) < -0.1:
            p_sat.alignment = PP_ALIGN.RIGHT
            p_desc.alignment = PP_ALIGN.RIGHT
        else:
            p_sat.alignment = PP_ALIGN.CENTER
            p_desc.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path

# Example execution:
# create_slide("fragmented_infographic.pptx")
