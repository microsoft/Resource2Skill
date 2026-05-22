import os
import io
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageFilter

def create_slide(
    output_pptx_path: str,
    items: list = None,
    bg_color: tuple = (248, 248, 248),
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with the Sliced Number Infographic style.

    Args:
        output_pptx_path: The path to save the generated PPTX file.
        items: A list of dictionaries, each containing 'number', 'color', 'title', 'details', and 'icon_char'.
               If None, default items will be used.
        bg_color: The RGB background color for the slide and the visual effect.

    Returns:
        The path to the saved PPTX file.
    """

    # --- Default Content ---
    if items is None:
        items = [
            {'number': '1', 'color': (118, 187, 2), 'title': 'RESEARCH', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '🔍'},
            {'number': '2', 'color': (0, 198, 164), 'title': 'IDEA', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '💡'},
            {'number': '3', 'color': (0, 169, 133), 'title': 'STRATEGY', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '♟️'},
            {'number': '4', 'color': (0, 163, 171), 'title': 'ASPIRATION', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '🚀'},
            {'number': '5', 'color': (44, 130, 184), 'title': 'PROCESS', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '⚙️'},
            {'number': '6', 'color': (48, 87, 143), 'title': 'TIME', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '⏱️'},
            {'number': '7', 'color': (54, 59, 101), 'title': 'EXPERIENCE', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '👔'},
            {'number': '8', 'color': (52, 53, 58), 'title': 'GOAL', 'details': 'Lorem ipsum dolor sit amet, consectetuer adipiscing elit.', 'icon_char': '🎯'},
        ]

    # --- Font Handling ---
    def download_font(url, filename="Roboto-Black.ttf"):
        if not os.path.exists(filename):
            try:
                print(f"Downloading font: {filename}...")
                urllib.request.urlretrieve(url, filename)
                print("Download complete.")
            except Exception as e:
                print(f"Error downloading font: {e}")
                return None
        return filename

    font_url = "https://github.com/google/fonts/raw/main/apache/roboto/Roboto-Black.ttf"
    font_path = download_font(font_url)
    try:
        main_font = ImageFont.truetype(font_path, 250)
        icon_font = ImageFont.truetype(font_path, 40)
        detail_font = ImageFont.truetype(font_path.replace("Black", "Regular"), 14)
    except IOError:
        print("Fallback to default Arial font.")
        main_font = ImageFont.truetype("arialbd.ttf", 250)
        icon_font = ImageFont.truetype("arial.ttf", 40)
        detail_font = ImageFont.truetype("arial.ttf", 14)

    # --- PIL Helper Function to create one infographic item ---
    def create_sliced_number_image(number_str, color_rgb, bg_rgb):
        img_size = (400, 400)
        
        # 1. Main canvas with background color
        img = Image.new("RGB", img_size, bg_rgb)
        draw = ImageDraw.Draw(img)

        # 2. Draw the colored number
        text_bbox = draw.textbbox((0, 0), number_str, font=main_font)
        text_w = text_bbox[2] - text_bbox[0]
        text_h = text_bbox[3] - text_bbox[1]
        text_pos = ((img_size[0] - text_w) / 2, (img_size[1] - text_h) / 2 - 30)
        draw.text(text_pos, number_str, font=main_font, fill=color_rgb)

        # 3. Create the shadow
        shadow_canvas = Image.new("RGBA", img_size, (0, 0, 0, 0))
        shadow_draw = ImageDraw.Draw(shadow_canvas)
        oval_coords = [10, 240, 390, 310]
        shadow_draw.ellipse(oval_coords, fill=(0, 0, 0, 100))
        shadow_canvas = shadow_canvas.filter(ImageFilter.GaussianBlur(15))
        
        # Rotate shadow
        shadow_canvas = shadow_canvas.rotate(20, resample=Image.BICUBIC, center=(img_size[0]/2, img_size[1]/2))
        
        # Composite shadow onto the main image
        img.paste(shadow_canvas, (0, -40), shadow_canvas) # Y-offset for placement

        # 4. Draw the "cutting" shape (mask)
        cutter_poly = [(0, 170), (400, 280), (400, 400), (0, 400)]
        draw.polygon(cutter_poly, fill=bg_rgb)

        return img

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Set background color ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)
    
    # --- Generate and place items ---
    img_width_in = 2.5
    num_cols = 4
    col_spacing = (prs.slide_width.inches - (num_cols * img_width_in)) / (num_cols + 1)
    
    for i, item in enumerate(items):
        row = i // num_cols
        col = i % num_cols
        
        # Create image
        pil_img = create_sliced_number_image(item['number'], item['color'], bg_color)
        img_stream = io.BytesIO()
        pil_img.save(img_stream, format="PNG")
        img_stream.seek(0)
        
        # Place image on slide
        left = Inches(col_spacing * (col + 1) + img_width_in * col)
        top = Inches(0.5 if row == 0 else 4.0)
        pic = slide.shapes.add_picture(img_stream, left, top, width=Inches(img_width_in))
        
        # Add Icon
        icon_left = left + Inches(1.3)
        icon_top = top + Inches(1.5)
        tb = slide.shapes.add_textbox(icon_left, icon_top, Inches(0.5), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = item['icon_char']
        p.font.name = 'Segoe UI Emoji' # For emoji icons
        p.font.size = Pt(24)

        # Add Title
        title_left = icon_left + Inches(0.4)
        title_top = top + Inches(1.55)
        tb = slide.shapes.add_textbox(title_left, title_top, Inches(1.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = item['title']
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(80, 80, 80)
        
        # Add Details
        details_left = title_left
        details_top = title_top + Inches(0.2)
        tb = slide.shapes.add_textbox(details_left, details_top, Inches(1.5), Inches(0.8))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = item['details']
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(120, 120, 120)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_slide("sliced_number_infographic.pptx")
