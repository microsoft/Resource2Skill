def create_slide(
    output_pptx_path: str,
    title_text: str = "Credits",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Cinematic End Credits Scroll visual effect.
    The output represents the visual state of the text mid-scroll.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Cinematic Black Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)

    # === Layer 2: Credits Content Setup ===
    # A mix of roles and names to simulate a movie crew list
    credits_data = [
        "Writer", "Jane Doe", "",
        "Producer", "John Smith", "",
        "Executive Producer", "Alice Johnson", "",
        "Lead Cast 1", "Michael Chang", "",
        "Lead Cast 2", "Sarah Williams", "",
        "Lead Cast 3", "David Brown", "",
        "Secondary Cast 1", "Emily Davis", "",
        "Secondary Cast 2", "Chris Wilson", "",
        "Supporting Cast", "Alex Miller", "",
        "Director of Photography", "Robert Taylor", "",
        "Production Designer", "Jessica Moore", "",
        "Editor 1", "William Anderson", "",
        "Lead Editor", "Thomas Jackson", "",
        "Colour Grader", "Sophie White", ""
    ]

    # === Layer 3: Text Box Layout & Formatting ===
    # Set text box width to ~60% of screen to create cinematic margins
    tb_width = Inches(8)
    tb_height = Inches(10) # Arbitrary initial height; will expand downwards
    left = (prs.slide_width - tb_width) / 2
    
    # Position slightly offset from the top to simulate the "mid-scroll" snapshot
    top = Inches(0.5) 

    txBox = slide.shapes.add_textbox(left, top, tb_width, tb_height)
    tf = txBox.text_frame
    tf.word_wrap = True

    # Populate and style the text
    for i, line in enumerate(credits_data):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        
        # Base font settings
        p.font.name = "Calibri"
        p.font.size = Pt(22)
        
        # Apply subtle hierarchy (Roles vs Names) while maintaining the style
        if i % 3 == 0 and line != "": 
            # Roles: Light grey and slightly smaller to recede
            p.font.color.rgb = RGBColor(170, 170, 170) 
            p.font.size = Pt(18)
        elif line != "":
            # Names: Bright white and larger to pop
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.bold = True

    # Save presentation
    prs.save(output_pptx_path)
    
    # Instruction for the user to complete the animation (since python-pptx cannot add motion paths natively)
    print("Slide generated successfully.")
    print("To complete the cinematic effect in PowerPoint:")
    print("1. Drag the text box completely below the slide canvas.")
    print("2. Go to Animations > Add Animation > More Motion Paths > 'Up'.")
    print("3. Drag the red endpoint above the slide canvas.")
    print("4. Set Animation Duration to ~15.00 seconds.")
    print("5. In Effect Options, set 'Smooth start' and 'Smooth end' to 0 seconds.")

    return output_pptx_path
