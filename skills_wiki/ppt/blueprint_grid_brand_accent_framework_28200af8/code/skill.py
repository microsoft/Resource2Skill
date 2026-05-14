def create_slide(
    output_pptx_path: str,
    title_text: str = "Andrew Doe",
    body_text: str = "Lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Facilisis leo vel fringilla est ullamcorper. Habitant morbi tristique senectus et netus et malesuada fames ac turpis egestas.",
    bg_palette: str = "business,portrait", 
    accent_color: tuple = (240, 90, 0),  # Bright Orange
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Blueprint Grid & Brand Accent Framework' visual effect.
    """
    import os
    import math
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # === Initialize Presentation ===
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # === Colors ===
    ACCENT_RGB = RGBColor(*accent_color)
    DARK_TEXT_RGB = RGBColor(40, 40, 40)
    LIGHT_TEXT_RGB = RGBColor(100, 100, 100)
    GRID_COLOR = (235, 235, 235)

    # === Layer 1: Background Dashed Grid (PIL) ===
    grid_img_path = "temp_grid_bg.png"
    img_w, img_h = 1920, 1080
    bg_img = Image.new('RGB', (img_w, img_h), (255, 255, 255))
    draw = ImageDraw.Draw(bg_img)
    
    grid_spacing = 60
    dash_length = 6
    
    def draw_dashed_line(draw_obj, pt1, pt2, fill, width, dash_len):
        x1, y1 = pt1
        x2, y2 = pt2
        dist = math.hypot(x2 - x1, y2 - y1)
        if dist == 0: return
        dashes = int(dist / dash_len)
        for i in range(dashes):
            if i % 2 == 0:  # Draw on even intervals
                start = (x1 + (x2-x1)*i/dashes, y1 + (y2-y1)*i/dashes)
                end = (x1 + (x2-x1)*(i+1)/dashes, y1 + (y2-y1)*(i+1)/dashes)
                draw_obj.line([start, end], fill=fill, width=width)

    # Draw vertical dashed lines
    for x in range(0, img_w, grid_spacing):
        draw_dashed_line(draw, (x, 0), (x, img_h), GRID_COLOR, 2, dash_length)
    
    # Draw horizontal dashed lines
    for y in range(0, img_h, grid_spacing):
        draw_dashed_line(draw, (0, y), (img_w, y), GRID_COLOR, 2, dash_length)
        
    bg_img.save(grid_img_path)
    slide.shapes.add_picture(grid_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Anchor Borders ===
    # Top Bar (Thicker)
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.4))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_RGB
    top_bar.line.fill.background()

    # Bottom Bar (Thinner)
    bottom_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, prs.slide_height - Inches(0.2), prs.slide_width, Inches(0.2))
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = ACCENT_RGB
    bottom_bar.line.fill.background()

    # === Layer 3: Image & Corner Accents ===
    img_left, img_top = Inches(1.5), Inches(1.5)
    img_width, img_height = Inches(4.5), Inches(3.0)
    
    # Download Photo
    photo_path = "temp_profile.jpg"
    try:
        url = f"https://source.unsplash.com/random/800x600/?{bg_palette}"
        urllib.request.urlretrieve(url, photo_path)
        slide.shapes.add_picture(photo_path, img_left, img_top, width=img_width, height=img_height)
    except Exception:
        # Fallback placeholder if download fails
        placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, img_left, img_top, img_width, img_height)
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(200, 200, 200)
        placeholder.line.fill.background()

    # Add Offset L-Accent (Bottom Right)
    offset = Inches(0.15)
    line_thickness = Inches(0.04)
    accent_len = Inches(0.8)
    
    # Vertical leg of the L
    v_leg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        img_left + img_width + offset, 
        img_top + img_height - accent_len + offset + line_thickness, 
        line_thickness, 
        accent_len
    )
    v_leg.fill.solid()
    v_leg.fill.fore_color.rgb = ACCENT_RGB
    v_leg.line.fill.background()

    # Horizontal leg of the L
    h_leg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        img_left + img_width - accent_len + offset + line_thickness, 
        img_top + img_height + offset, 
        accent_len, 
        line_thickness
    )
    h_leg.fill.solid()
    h_leg.fill.fore_color.rgb = ACCENT_RGB
    h_leg.line.fill.background()

    # === Layer 4: Text Content ===
    
    # Name / Title
    name_box = slide.shapes.add_textbox(img_left, img_top + img_height + Inches(0.3), Inches(4.5), Inches(0.8))
    tf = name_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = DARK_TEXT_RGB

    # Attributes (Age / Nationality)
    attr_top = img_top + img_height + Inches(1.2)
    labels = ["Age:", "Nationality:"]
    values = ["46", "American"]
    
    for i in range(2):
        lbl_box = slide.shapes.add_textbox(img_left, attr_top + Inches(i*0.3), Inches(1.5), Inches(0.3))
        p_lbl = lbl_box.text_frame.paragraphs[0]
        p_lbl.text = labels[i]
        p_lbl.font.size = Pt(12)
        p_lbl.font.color.rgb = ACCENT_RGB
        
        val_box = slide.shapes.add_textbox(img_left + Inches(0.8), attr_top + Inches(i*0.3), Inches(2.0), Inches(0.3))
        p_val = val_box.text_frame.paragraphs[0]
        p_val.text = values[i]
        p_val.font.size = Pt(12)
        p_val.font.color.rgb = LIGHT_TEXT_RGB

    # Body Paragraph (Right side)
    body_box = slide.shapes.add_textbox(img_left + img_width + Inches(0.8), img_top, Inches(5.0), Inches(4.0))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = LIGHT_TEXT_RGB
    p_body.line_spacing = 1.5

    # Cleanup temp files
    if os.path.exists(grid_img_path):
        os.remove(grid_img_path)
    if os.path.exists(photo_path):
        os.remove(photo_path)

    prs.save(output_pptx_path)
    return output_pptx_path
