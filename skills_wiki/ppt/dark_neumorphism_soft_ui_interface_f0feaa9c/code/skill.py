import os
import urllib.request
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFilter

def create_slide(
    output_pptx_path: str,
    title_text: str = "Neumorphism",
    body_text: str = "",
    bg_palette: str = "portrait", 
    accent_color: tuple = (43, 214, 214),  # Cyan
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Dark Neumorphism (Soft UI) effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Color Palette ---
    BG_COLOR = (46, 50, 57)
    SHADOW_DARK = (20, 22, 25)
    SHADOW_LIGHT = (70, 75, 85)
    DEPRESSED_BG = (35, 38, 43)

    # 1. Set Slide Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*BG_COLOR)

    # --- Helper: Generate Neumorphic Panel (PIL) ---
    def make_neu_panel(w_in, h_in, pressed=False, radius=20, blur=12, offset=8):
        dpi = 150 # High res for crisp rendering
        w, h = int(w_in * dpi), int(h_in * dpi)
        pad = int(max(blur, offset) * 2.5) # Padding for shadow bleed
        
        img = Image.new("RGBA", (w + pad*2, h + pad*2), (0,0,0,0))
        
        if not pressed:
            # Highlight (Top-Left)
            light_layer = Image.new("RGBA", img.size, (0,0,0,0))
            ImageDraw.Draw(light_layer).rounded_rectangle(
                (pad - offset, pad - offset, pad + w - offset, pad + h - offset), 
                radius=radius, fill=SHADOW_LIGHT + (255,)
            )
            light_layer = light_layer.filter(ImageFilter.GaussianBlur(blur))
            
            # Dark Shadow (Bottom-Right)
            dark_layer = Image.new("RGBA", img.size, (0,0,0,0))
            ImageDraw.Draw(dark_layer).rounded_rectangle(
                (pad + offset, pad + offset, pad + w + offset, pad + h + offset), 
                radius=radius, fill=SHADOW_DARK + (255,)
            )
            dark_layer = dark_layer.filter(ImageFilter.GaussianBlur(blur))
            
            # Composite shadows
            img = Image.alpha_composite(img, light_layer)
            img = Image.alpha_composite(img, dark_layer)
            
            # Base Shape
            base_layer = Image.new("RGBA", img.size, (0,0,0,0))
            ImageDraw.Draw(base_layer).rounded_rectangle(
                (pad, pad, pad + w, pad + h), radius=radius, fill=BG_COLOR + (255,)
            )
            img = Image.alpha_composite(img, base_layer)
        else:
            # Simple depressed state (darker inner well)
            base_layer = Image.new("RGBA", img.size, (0,0,0,0))
            ImageDraw.Draw(base_layer).rounded_rectangle(
                (pad, pad, pad + w, pad + h), radius=radius, fill=DEPRESSED_BG + (255,)
            )
            img = Image.alpha_composite(img, base_layer)
            
        # Crop slightly to remove excess padding but keep shadow
        bbox = img.getbbox()
        if bbox:
            img = img.crop(bbox)
            
        temp_path = f"temp_neu_{w_in}_{h_in}_{pressed}.png"
        img.save(temp_path)
        return temp_path

    # --- Helper: Generate Circular Avatar (PIL) ---
    def make_circular_avatar(image_url):
        temp_path = "temp_avatar.png"
        try:
            req = urllib.request.Request(image_url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req) as response:
                img = Image.open(BytesIO(response.read())).convert("RGBA")
        except Exception:
            # Fallback if download fails
            img = Image.new("RGBA", (400, 400), (100, 100, 100, 255))
            
        # Make square
        min_dim = min(img.size)
        img = img.crop((0, 0, min_dim, min_dim))
        img = img.resize((300, 300), Image.Resampling.LANCZOS)
        
        # Apply circular mask
        mask = Image.new("L", img.size, 0)
        ImageDraw.Draw(mask).ellipse((0, 0) + img.size, fill=255)
        img.putalpha(mask)
        
        img.save(temp_path)
        return temp_path

    # --- Build UI: 1. Profile Card (Left) ---
    card_w, card_h = 3.5, 4.0
    card_x, card_y = 1.5, 1.0
    panel_img = make_neu_panel(card_w, card_h)
    # Insert panel (adjust position slightly to account for the PIL padding)
    slide.shapes.add_picture(panel_img, Inches(card_x - 0.2), Inches(card_y - 0.2), width=Inches(card_w + 0.4))

    # Add Avatar
    avatar_img = make_circular_avatar("https://images.unsplash.com/photo-1534528741775-53994a69daeb?auto=format&fit=crop&w=400&q=80")
    slide.shapes.add_picture(avatar_img, Inches(card_x + 0.75), Inches(card_y + 0.5), width=Inches(2.0), height=Inches(2.0))

    # Add Cyan Button Native Shape
    btn = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(card_x + 0.5), Inches(card_y + 2.8),
        Inches(2.5), Inches(0.6)
    )
    btn.fill.solid()
    btn.fill.fore_color.rgb = RGBColor(*accent_color)
    btn.line.fill.background() # No border
    
    tf = btn.text_frame
    tf.text = "Message"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.name = "Segoe UI"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # --- Build UI: 2. Media Controls (Bottom Left) ---
    ctrl_y = 5.5
    ctrl_size = 0.8
    spacing = 1.1
    symbols = ["\u25C0", "\u23F8", "\u25B6"] # Prev, Pause, Next
    
    ctrl_panel_img = make_neu_panel(ctrl_size, ctrl_size, radius=15)
    for i in range(3):
        cx = 1.5 + (i * spacing)
        slide.shapes.add_picture(ctrl_panel_img, Inches(cx - 0.15), Inches(ctrl_y - 0.15), width=Inches(ctrl_size + 0.3))
        
        # Add text symbol
        tx_box = slide.shapes.add_textbox(Inches(cx), Inches(ctrl_y), Inches(ctrl_size), Inches(ctrl_size))
        tx_frame = tx_box.text_frame
        tx_frame.text = symbols[i]
        tx_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        tx_frame.paragraphs[0].font.size = Pt(20)
        tx_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # --- Build UI: 3. Equalizer (Right) ---
    eq_w, eq_h = 4.5, 3.5
    eq_x, eq_y = 6.5, 1.2
    
    # Track metrics
    num_tracks = 5
    track_w = 0.4
    track_spacing = 0.8
    bar_heights = [1.5, 2.5, 1.2, 2.8, 2.0]
    
    track_img = make_neu_panel(track_w, eq_h, pressed=True, radius=10)
    
    for i in range(num_tracks):
        tx = eq_x + (i * track_spacing)
        # Insert depressed track
        slide.shapes.add_picture(track_img, Inches(tx - 0.1), Inches(eq_y - 0.1), width=Inches(track_w + 0.2))
        
        # Add Cyan active bar over it (bottom aligned)
        bh = bar_heights[i]
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(tx), Inches(eq_y + eq_h - bh),
            Inches(track_w), Inches(bh)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(*accent_color)
        bar.line.fill.background()

    # Clean up temp files
    try:
        os.remove(panel_img)
        os.remove(ctrl_panel_img)
        os.remove(track_img)
        os.remove(avatar_img)
    except OSError:
        pass

    prs.save(output_pptx_path)
    return output_pptx_path
