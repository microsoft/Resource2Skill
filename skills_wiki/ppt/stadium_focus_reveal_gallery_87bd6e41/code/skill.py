"""
Skill: Stadium Focus Reveal Gallery
skill_id: stadium_focus_reveal_gallery_87bd6e41

Source video: https://www.youtube.com/watch?v=gYkQUFhWqeI
  (Easy TEAM SLIDES in PowerPoint 😎 (step-by-step) — Luis Urrutia)

Wiki cross-refs:
  - text/overview.md       -- when & why to use
  - text/visual_design.md  -- visual breakdown of each layer
  - text/usage.md          -- parameter docs
  - visual/frames/         -- reference frames from source video
  - visual/reference_render.png -- output of THIS code
"""

def create_slide(
    output_pptx_path: str,
    title_text: str = "Meet our Team",
    active_index: int = 1,
    accent_color: tuple = (236, 64, 122),  # Vibrant Pink
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Stadium Focus Reveal Gallery.
    """
    import urllib.request
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageOps
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)

    # Team Data definition
    team_data = [
        {"name": "Bill Board", "role": "Content Marketing Lead", "url": "https://images.unsplash.com/photo-1560250097-0b93528c311a?w=400&q=80"},
        {"name": "Anita Break", "role": "Marketing Operations", "url": "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?w=400&q=80"},
        {"name": "Ella Vator", "role": "Social Media Manager", "url": "https://images.unsplash.com/photo-1573497019940-1c28c88b4f3e?w=400&q=80"},
        {"name": "Cliff Hanger", "role": "Events Manager", "url": "https://images.unsplash.com/photo-1519085360753-af0119f7cbe7?w=400&q=80"}
    ]

    # Helper: Fetch image with fallback
    def get_image(url):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=5) as resp:
                return Image.open(BytesIO(resp.read())).convert("RGBA")
        except Exception:
            # Fallback graphic if download fails
            img = Image.new('RGBA', (400, 400), (100, 100, 100, 255))
            d = ImageDraw.Draw(img)
            d.ellipse([100, 100, 300, 300], fill=(150, 150, 150, 255))
            return img

    # Helper: Build Stadium Card via PIL
    def build_stadium_card(url, is_active):
        width = 300 if is_active else 250
        height = 600 if is_active else 500
        bg_color = accent_color + (255,) if is_active else (40, 40, 40, 255)

        # 1. Base color fill
        combined = Image.new('RGBA', (width, height), bg_color)
        
        # 2. Fetch and format portrait
        portrait = get_image(url)
        min_dim = min(portrait.width, portrait.height)
        left = (portrait.width - min_dim) / 2
        top = (portrait.height - min_dim) / 2
        portrait = portrait.crop((left, top, left + min_dim, top + min_dim))
        portrait = portrait.resize((width, width))
        
        if not is_active:
            portrait = ImageOps.grayscale(portrait).convert('RGBA')
            
        # 3. Paste portrait at the top
        combined.paste(portrait, (0, 0))
        
        # 4. Create Stadium (Pill) Mask
        mask = Image.new('L', (width, height), 0)
        draw = ImageDraw.Draw(mask)
        draw.rounded_rectangle([0, 0, width, height], radius=width//2, fill=255)
        
        # 5. Apply Mask
        final = Image.new('RGBA', (width, height), (0, 0, 0, 0))
        final.paste(combined, (0, 0), mask)
        
        io = BytesIO()
        final.save(io, format='PNG')
        io.seek(0)
        return io

    # === Layer 2: Visual Effect (Inserting the Cards) ===
    centers_x = [1.66, 5.00, 8.33, 11.66]  # Evenly distributed horizontally
    
    for i, person in enumerate(team_data):
        is_active = (i == active_index)
        card_stream = build_stadium_card(person['url'], is_active)
        
        w_inch = 3.0 if is_active else 2.5
        h_inch = 6.0 if is_active else 5.0
        
        left_inch = centers_x[i] - (w_inch / 2)
        top_inch = 2.0 if is_active else 3.0  # Active card floats higher
        
        slide.shapes.add_picture(card_stream, Inches(left_inch), Inches(top_inch), 
                                 width=Inches(w_inch), height=Inches(h_inch))

        # === Add Highlight Text ===
        if is_active:
            # Place floating details panel to the right of the active card
            # (If it's the last card, place it to the left)
            text_left = centers_x[i] + (w_inch / 2) + 0.2
            if i == len(team_data) - 1:
                text_left = centers_x[i] - (w_inch / 2) - 3.2
                
            tb = slide.shapes.add_textbox(Inches(text_left), Inches(2.2), Inches(3.0), Inches(1.0))
            tf = tb.text_frame
            tf.clear()
            
            p1 = tf.paragraphs[0]
            p1.text = person['name']
            p1.font.bold = True
            p1.font.size = Pt(28)
            p1.font.color.rgb = RGBColor(*accent_color)
            
            p2 = tf.add_paragraph()
            p2.text = person['role']
            p2.font.size = Pt(16)
            p2.font.color.rgb = RGBColor(200, 200, 200)

    # === Layer 3: Bottom Fade Gradient Mask ===
    # Creates the illusion that the shapes fade smoothly into the black background
    grad_img = Image.new('RGBA', (2000, 400))
    draw_grad = ImageDraw.Draw(grad_img)
    for y in range(400):
        # Cubic curve for a smoother visual ease-in to total darkness
        alpha = int(255 * (y / 400)**3)
        draw_grad.line([(0, y), (2000, y)], fill=(0, 0, 0, alpha))
        
    grad_io = BytesIO()
    grad_img.save(grad_io, format='PNG')
    grad_io.seek(0)
    
    # Position gradient over the bottom 2.5 inches of the slide
    slide.shapes.add_picture(grad_io, 0, Inches(5.0), 
                             width=Inches(13.333), height=Inches(2.5))

    # === Layer 4: Global Typography ===
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.6), Inches(13.333), Inches(1.0))
    title_tf = title_box.text_frame
    title_p = title_tf.paragraphs[0]
    title_p.text = title_text
    title_p.alignment = PP_ALIGN.CENTER
    title_p.font.bold = True
    title_p.font.size = Pt(44)
    title_p.font.name = 'Arial'
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    prs.save(output_pptx_path)
    return output_pptx_path
