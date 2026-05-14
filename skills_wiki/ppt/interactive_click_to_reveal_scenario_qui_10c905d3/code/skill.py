def create_slide(
    output_pptx_path: str,
    title_text: str = "What does this function return?",
    subtitle_text: str = "=ROUND(258.26, -1)",
    options: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Interactive Click-to-Reveal Quiz effect.
    Generates the layout and uniquely names shapes so you can instantly add 
    'On Click' animation triggers in PowerPoint.

    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree

    # Default quiz options if none provided
    if options is None:
        options = [
            {"label": "A", "text": "An error", "correct": False, "color": (59, 130, 246)},
            {"label": "B", "text": "260", "correct": True, "color": (14, 116, 144)},
            {"label": "C", "text": "258.3", "correct": False, "color": (16, 185, 129)},
            {"label": "D", "text": "258", "correct": False, "color": (132, 204, 22)},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(31, 41, 55) # Dark Slate

    # --- Helper: Name Objects for PPT Selection Pane ---
    def assign_name(shape, name):
        try:
            shape.name = name
        except Exception:
            pass
        try:
            # Fallback lxml injection if native property fails
            for cNvPr in shape.element.xpath('.//p:cNvPr', namespaces={'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}):
                cNvPr.set('name', name)
        except Exception:
            pass

    # --- Helper: Add subtle shadow to shapes ---
    def add_shadow(shape):
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad="50800", dist="38100", dir="2700000", algn="tl", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="40000")

    # === Layer 2: Text & Content ===
    
    # Header
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11.3), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Subtitle / Code snippet
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(11.3), Inches(0.8))
    p = sub_box.text_frame.paragraphs[0]
    p.text = subtitle_text
    p.font.name = "Consolas"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Render Options
    start_y = 2.8
    spacing_y = 1.0

    for i, opt in enumerate(options):
        y = start_y + i * spacing_y
        
        # 1. Circle Button (The "Trigger")
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(y), Inches(0.6), Inches(0.6))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor(*opt["color"])
        circle.line.color.rgb = RGBColor(*opt["color"]) # hide border
        
        tf = circle.text_frame
        tf.text = opt["label"]
        tf.paragraphs[0].font.size = Pt(24)
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        add_shadow(circle)
        assign_name(circle, f"Trigger_Button_{opt['label']}")
        
        # 2. Answer Text
        text_box = slide.shapes.add_textbox(Inches(4.4), Inches(y), Inches(3.5), Inches(0.6))
        tf = text_box.text_frame
        tf.text = opt["text"]
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        # 3. Reveal Icon (The item that will appear on click)
        icon_box = slide.shapes.add_textbox(Inches(8.0), Inches(y - 0.1), Inches(0.8), Inches(0.8))
        tf = icon_box.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        
        if opt["correct"]:
            p.text = "\u2714" # Checkmark
            p.font.color.rgb = RGBColor(34, 197, 94) # Vibrant Green
        else:
            p.text = "\u2718" # Cross mark
            p.font.color.rgb = RGBColor(239, 68, 68) # Vibrant Red
            
        p.font.size = Pt(44)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        add_shadow(icon_box)
        assign_name(icon_box, f"Reveal_Icon_{opt['label']}")

    # Instructions for the user (added off-slide or as a note, handled locally)
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = (
        "INTERACTIVE SETUP INSTRUCTIONS:\n"
        "1. Open the 'Animations' pane.\n"
        "2. Select the green/red icon shapes on the slide.\n"
        "3. Add an 'Appear' animation to them.\n"
        "4. Click 'Trigger' -> 'On Click of' -> Choose the corresponding 'Trigger_Button_X' from the list!"
    )

    prs.save(output_pptx_path)
    return output_pptx_path

