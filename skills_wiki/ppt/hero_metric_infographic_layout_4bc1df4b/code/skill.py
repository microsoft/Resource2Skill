import os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide(
    output_pptx_path: str,
    title_text: str = "GLOBAL IMPACT & SCALE",
    metrics_data: list = None,
    bg_color: tuple = (15, 23, 42),      # Dark Slate/Navy
    grid_color: tuple = (30, 41, 59),    # Slightly lighter for grid texture
    accent_color: tuple = (0, 191, 255), # Cyan accent
    **kwargs,
) -> str:
    """
    Creates a PPTX slide featuring a high-impact 'Hero Metric Grid' style 
    extracted from professional presentation makeovers.
    """
    if metrics_data is None:
        metrics_data = [
            {"value": "180,000", "label": "Daily active vessels tracked globally"},
            {"value": "35B", "label": "Dollars generated in channel business"},
            {"value": "98%", "label": "Of Fortune 500 companies served"}
        ]

    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ==========================================
    # Layer 1: Generate & Apply Custom Background
    # ==========================================
    bg_img_path = "temp_hero_bg.png"
    # Create a 1920x1080 background
    img = Image.new('RGB', (1920, 1080), color=bg_color)
    draw = ImageDraw.Draw(img)
    
    # Draw a subtle "tech/data" grid pattern to give the background depth
    grid_spacing = 80
    for x in range(0, 1920, grid_spacing):
        draw.line([(x, 0), (x, 1080)], fill=grid_color, width=2)
    for y in range(0, 1080, grid_spacing):
        draw.line([(0, y), (1920, y)], fill=grid_color, width=2)
        
    img.save(bg_img_path)

    # Insert background
    slide.shapes.add_picture(bg_img_path, Inches(0), Inches(0), prs.slide_width, prs.slide_height)

    # ==========================================
    # Layer 2: Slide Title
    # ==========================================
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11), Inches(1))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Arial"
    
    # Add a title accent line
    title_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(1.5), Inches(0.06)
    )
    title_line.fill.solid()
    title_line.fill.fore_color.rgb = RGBColor(*accent_color)
    title_line.line.color.rgb = RGBColor(*accent_color)

    # ==========================================
    # Layer 3: Hero Metrics Grid Generation
    # ==========================================
    num_metrics = len(metrics_data)
    # Calculate horizontal spacing
    margin_x = 1.0
    usable_width = 13.333 - (margin_x * 2)
    col_width = usable_width / num_metrics
    start_y = 2.8

    for i, metric in enumerate(metrics_data):
        start_x = margin_x + (i * col_width)
        
        # 1. Accent dash for the metric
        dash = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(start_x), Inches(start_y), Inches(0.5), Inches(0.08)
        )
        dash.fill.solid()
        dash.fill.fore_color.rgb = RGBColor(*accent_color)
        dash.line.color.rgb = RGBColor(*accent_color)

        # 2. Big Hero Number
        num_box = slide.shapes.add_textbox(Inches(start_x), Inches(start_y + 0.1), Inches(col_width * 0.9), Inches(1.5))
        num_tf = num_box.text_frame
        num_tf.word_wrap = True
        num_p = num_tf.paragraphs[0]
        num_p.text = metric.get("value", "0")
        num_p.font.size = Pt(88) # Massive typography
        num_p.font.bold = True
        num_p.font.color.rgb = RGBColor(255, 255, 255)
        num_p.font.name = "Arial"
        
        # 3. Context/Description Text
        desc_box = slide.shapes.add_textbox(Inches(start_x), Inches(start_y + 1.8), Inches(col_width * 0.85), Inches(1.5))
        desc_tf = desc_box.text_frame
        desc_tf.word_wrap = True
        desc_p = desc_tf.paragraphs[0]
        desc_p.text = metric.get("label", "Description")
        desc_p.font.size = Pt(18)
        desc_p.font.color.rgb = RGBColor(200, 205, 215) # Light gray for contrast
        desc_p.font.name = "Arial"

    prs.save(output_pptx_path)
    
    # Cleanup temp file
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)
        
    return output_pptx_path

# Example execution:
# create_slide("hero_metrics_infographic.pptx")
