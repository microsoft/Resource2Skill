import io
import requests
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR
from lxml import etree
from PIL import Image, ImageDraw, ImageOps

def create_slide(
    output_pptx_path: str,
    expert_name: str = "ANGELA SMITH",
    expert_title: str = "WEB DESIGNER, DEVELOPER, CREATIVE",
    image_url: str = "https://images.unsplash.com/photo-1594744803329-e58b31de8bf5?q=80&w=1887",
    bg_image_url: str = "https://images.unsplash.com/photo-1605979854205-399564177716?q=80&w=1974",
    **kwargs,
) -> str:
    """
    Creates a single-slide PowerPoint presentation with a 'Geometric Profile Showcase' design.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        expert_name: The name of the expert.
        expert_title: The job title/role of the expert.
        image_url: URL to the expert's profile picture.
        bg_image_url: URL to the background cityscape image.

    Returns:
        The path to the saved .pptx file.
    """
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- 1. Create Composite Background with PIL ---
    width_px, height_px = 1920, 1080

    # Download background image with a fallback
    try:
        bg_response = requests.get(bg_image_url, stream=True)
        bg_response.raise_for_status()
        bg_img = Image.open(bg_response.raw).convert("RGBA")
    except requests.exceptions.RequestException:
        bg_img = Image.new("RGBA", (width_px, height_px), (50, 50, 60, 255)) # Fallback

    # Resize, desaturate, and darken the background image
    bg_img = bg_img.resize((width_px, height_px))
    bg_img = ImageOps.grayscale(bg_img)
    bg_img = ImageOps.colorize(bg_img, black="rgb(0,0,0)", white="rgb(150,150,150)")
    bg_img = bg_img.convert("RGBA")

    # Create a drawing canvas
    canvas = Image.new("RGBA", (width_px, height_px), (0, 0, 0, 0))
    draw = ImageDraw.Draw(canvas)

    # Define vertices for the polygons (as proportions of slide dimensions)
    poly_dark = [(0, 0.45), (1, 0.6), (1, 1), (0, 1)]
    poly_pink = [(0, 0), (0.7, 0), (0.35, 1), (0, 1)]
    poly_orange = [(0.4, 0), (1, 0), (1, 0.6), (0.75, 0.5)]

    # Convert proportional vertices to pixel coordinates
    poly_dark_px = [(x * width_px, y * height_px) for x, y in poly_dark]
    poly_pink_px = [(x * width_px, y * height_px) for x, y in poly_pink]
    poly_orange_px = [(x * width_px, y * height_px) for x, y in poly_orange]

    # Draw semi-transparent polygons
    draw.polygon(poly_dark_px, fill=(46, 33, 83, 140))
    draw.polygon(poly_pink_px, fill=(219, 39, 119, 140))
    draw.polygon(poly_orange_px, fill=(239, 113, 83, 160))

    # Composite the polygons over the background image
    final_bg = Image.alpha_composite(bg_img, canvas)

    # Save to a memory buffer
    img_stream = io.BytesIO()
    final_bg.save(img_stream, format="PNG")
    img_stream.seek(0)

    # Add the composite image as the new slide background
    slide.shapes.add_picture(img_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # --- 2. Add Content with python-pptx ---

    # Photo Frame and Picture
    frame_left, frame_top, frame_size = Inches(1), Inches(2), Inches(3.5)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, frame_left, frame_top, frame_size, frame_size)
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.fill.solid()
    frame.line.fill.fore_color.rgb = RGBColor(200, 200, 200)
    frame.shadow.inherit = False

    try:
        pic_response = requests.get(image_url, stream=True)
        pic_response.raise_for_status()
        pic_stream = io.BytesIO(pic_response.content)
        slide.shapes.add_picture(pic_stream, frame_left + Inches(0.1), frame_top + Inches(0.1), width=frame_size - Inches(0.2), height=frame_size - Inches(0.2))
    except requests.exceptions.RequestException:
        pass # If photo fails, the white frame remains as a placeholder

    # Main Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.7), Inches(8), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = "ABOUT OUR EXPERT"
    p.font.name = 'Agency FB'
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Sub-text for title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(8), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "Lorem ipsum dolor sit amet, consectetuer adipiscing elit. Maecenas porttitor congue massa."
    p.font.name = 'Calibri (Body)'
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(220, 220, 220)

    # Expert Name & Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(5.7), Inches(4), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = expert_name
    p.font.name = 'Agency FB'
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(6.1), Inches(4), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = expert_title
    p.font.name = 'Agency FB'
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(220, 220, 220)

    # Main Content Area
    txBox = slide.shapes.add_textbox(Inches(5.5), Inches(2.5), Inches(9.5), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "HEADING HERE"
    p.font.name = 'Agency FB'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    txBox = slide.shapes.add_textbox(Inches(5.5), Inches(3.2), Inches(9.5), Inches(2))
    p = txBox.text_frame.paragraphs[0]
    p.text = "Nunc viverra imperdiet enim. Fusce est. Vivamus a tellus. Lorem ipsum dolor sit amet, consectetuer adipiscing elit.\n\n" \
             "• Nunc viverra imperdiet enim. Fusce est. Vivamus a tellus.\n" \
             "• Nunc viverra imperdiet enim. Fusce est. Vivamus a tellus.\n" \
             "• Nunc viverra imperdiet enim. Fusce est. Vivamus a tellus."
    p.font.name = 'Calibri'
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(220, 220, 220)
    p.line_spacing = 1.5

    # Skill Bars
    def create_skill_bar(left, top, width, percentage, color1, color2):
        # Bar background
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.3))
        track.fill.solid()
        track.fill.fore_color.rgb = RGBColor(255, 255, 255)
        track.line.fill.solid()
        track.line.fill.fore_color.rgb = RGBColor(220, 220, 220)
        track.adjustments[0] = 0.5 # Fully rounded
        
        # Bar fill
        fill_width = width * (percentage / 100)
        fill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, fill_width, Inches(0.3))
        fill.line.fill.background()
        fill_grad = fill.fill
        fill_grad.gradient()
        fill_grad.gradient_stops[0].color.rgb = color1
        fill_grad.gradient_stops[1].color.rgb = color2
        fill_grad.gradient_angle = 0
        fill.adjustments[0] = 0.5

        # Percentage text
        txBox = slide.shapes.add_textbox(left + width - Inches(0.7), top, Inches(0.7), Inches(0.3))
        p = txBox.text_frame.paragraphs[0]
        p.text = f"{percentage}%"
        p.font.bold = True
        p.font.color.rgb = RGBColor(80, 80, 80)
        p.alignment = PP_ALIGN.RIGHT

    create_skill_bar(Inches(10.5), Inches(5.5), Inches(4.5), 85, RGBColor(227, 85, 134), RGBColor(244, 151, 107))
    create_skill_bar(Inches(10.5), Inches(6.2), Inches(4.5), 79, RGBColor(227, 85, 134), RGBColor(244, 151, 107))

    # Icons
    icons_y = Inches(1.3)
    icon_size = Inches(0.5)
    for i in range(6):
        icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12 + i * 0.7), icons_y, icon_size, icon_size)
        icon.fill.solid()
        icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
        icon.line.fill.solid()
        icon.line.fill.fore_color.rgb = RGBColor(255, 255, 255)
        icon.line.width = Pt(1.5)
        icon.fill.background() # Make transparent

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
# create_slide("geometric_profile_showcase.pptx")
