def create_slide(
    output_pptx_path: str,
    title_text: str = "Australia",
    body_text: str = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris.",
    bg_palette: str = "landscape,coast",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Alpha-Masked Typographic Reveal effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image, ImageDraw, ImageFont

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    width_px, height_px = 1920, 1080

    # --- Helper: Download Image ---
    bg_img_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/featured/1920x1080/?{bg_palette.replace(',', '%20')}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(bg_img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Fallback to creating a solid gray image if download fails
        fallback_img = Image.new("RGB", (width_px, height_px), (100, 110, 120))
        fallback_img.save(bg_img_path)

    # --- Helper: Download/Load Font for Mask ---
    font_path = "temp_font.ttf"
    try:
        # Arvo Bold - A heavy slab serif similar to ChunkFive
        font_url = "https://raw.githubusercontent.com/google/fonts/main/ofl/arvo/Arvo-Bold.ttf"
        req = urllib.request.Request(font_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(font_path, 'wb') as out_file:
            out_file.write(response.read())
        giant_font = ImageFont.truetype(font_path, 800)
    except Exception:
        # Fallback to system fonts
        try:
            giant_font = ImageFont.truetype("arialbd.ttf", 800) # Windows
        except:
            try:
                giant_font = ImageFont.truetype("HelveticaNeue-Bold.ttc", 800) # Mac
            except:
                giant_font = ImageFont.load_default()

    # --- Create Gradient Overlay with Text Cutout (PIL) ---
    overlay_path = "temp_overlay.png"
    
    # 1. Create the Alpha mask (L mode)
    alpha_img = Image.new("L", (width_px, height_px))
    draw_alpha = ImageDraw.Draw(alpha_img)
    
    # Draw horizontal gradient: 95% opaque left -> transparent right
    # Hex for 95% opacity is ~242
    for x in range(width_px):
        if x < width_px * 0.35:
            a = 242
        elif x < width_px * 0.8:
            progress = (x - width_px * 0.35) / (width_px * 0.45)
            a = int(242 * (1 - progress))
        else:
            a = 0
        draw_alpha.line([(x, 0), (x, height_px)], fill=a)
        
    # 2. Punch hole in the alpha mask using the first letter of the title
    first_letter = title_text[0].upper() if title_text else "A"
    
    # Calculate centering for the giant letter on the left side
    try:
        bbox = giant_font.getbbox(first_letter)
        text_w = bbox[2] - bbox[0]
        text_h = bbox[3] - bbox[1]
        x_pos = 150 # Margin from left
        y_pos = (height_px - text_h) // 2 - bbox[1]
    except AttributeError:
        x_pos, y_pos = 150, 100 # Fallback
        
    # Draw text with fill=0 (completely transparent hole)
    draw_alpha.text((x_pos, y_pos), first_letter, font=giant_font, fill=0)
    
    # 3. Apply alpha mask to solid color block
    # Dark navy/slate color for modern aesthetic
    overlay_img = Image.new("RGB", (width_px, height_px), (15, 23, 42))
    overlay_img.putalpha(alpha_img)
    overlay_img.save(overlay_path)

    # --- Assemble PowerPoint Slide ---
    
    # Layer 1: Background Image
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    
    # Layer 2: Gradient Overlay with Punch-out
    slide.shapes.add_picture(overlay_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # Layer 3: Typography (Title)
    left_margin = Inches(5.0)
    tx_box = slide.shapes.add_textbox(left_margin, Inches(3.0), Inches(6.0), Inches(1.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Arial Black"

    # Layer 4: Typography (Body)
    body_box = slide.shapes.add_textbox(left_margin, Inches(4.2), Inches(6.5), Inches(2.0))
    bf = body_box.text_frame
    bf.word_wrap = True
    
    p2 = bf.paragraphs[0]
    p2.text = body_text
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(226, 232, 240) # Light Slate Gray
    p2.font.name = "Calibri"

    # Save Presentation
    prs.save(output_pptx_path)

    # Cleanup temp files
    for f in [bg_img_path, font_path, overlay_path]:
        if os.path.exists(f):
            try:
                os.remove(f)
            except:
                pass

    return output_pptx_path
