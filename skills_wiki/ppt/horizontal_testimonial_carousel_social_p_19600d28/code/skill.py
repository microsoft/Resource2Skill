def create_slide(
    output_pptx_path: str,
    title_text: str = "Client Testimonials.",
    subtitle_text: str = "See what our partners have to say about working with us.",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing a modern Horizontal Testimonial Carousel.
    """
    import io
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw

    # --- Configuration ---
    bg_color = RGBColor(244, 240, 236)         # Soft warm beige
    card_color = RGBColor(255, 255, 255)       # Pure white
    text_dark = RGBColor(40, 40, 40)           # Dark gray for names
    text_light = RGBColor(100, 100, 100)       # Medium gray for reviews
    star_color = RGBColor(255, 180, 0)         # Golden yellow
    
    # Mock data for the slider
    testimonials = [
        {
            "name": "Hannah Morales",
            "text": "Presentations are communication tools that can be used as demonstrations, lectures, speeches, reports, and more.",
            "avatar_url": "https://images.unsplash.com/photo-1494790108377-be9c29b29330?auto=format&fit=crop&w=200&q=80"
        },
        {
            "name": "Olivia Wilson",
            "text": "The cleanest design templates I have ever used. They helped us close our Series A funding round with ease.",
            "avatar_url": "https://images.unsplash.com/photo-1438761681033-6461ffad8d80?auto=format&fit=crop&w=200&q=80"
        },
        {
            "name": "Morgan Maxwell",
            "text": "An absolute game changer for our marketing team. We create beautiful materials in half the time.",
            "avatar_url": "https://images.unsplash.com/photo-1507003211169-0a1dd7228f2d?auto=format&fit=crop&w=200&q=80"
        }
    ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background ===
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg_color
    bg_shape.line.fill.background() # No outline

    # === Layer 2: Main Title ===
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = text_dark
    p.alignment = PP_ALIGN.CENTER

    if subtitle_text:
        sub_p = tf.add_paragraph()
        sub_p.text = subtitle_text
        sub_p.font.size = Pt(14)
        sub_p.font.color.rgb = text_light
        sub_p.font.bold = False
        sub_p.alignment = PP_ALIGN.CENTER

    # === Helper Function: Generate Circular Avatar via PIL ===
    def get_circular_avatar(url, size=(200, 200)):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=5) as response:
                img = Image.open(io.BytesIO(response.read())).convert("RGBA")
        except Exception:
            # Fallback if download fails
            img = Image.new("RGBA", size, (200, 200, 200, 255))
        
        # Crop to square
        w, h = img.size
        min_dim = min(w, h)
        left = (w - min_dim)/2
        top = (h - min_dim)/2
        img = img.crop((left, top, left+min_dim, top+min_dim))
        img = img.resize(size, Image.Resampling.LANCZOS)
        
        # Create circular mask
        mask = Image.new("L", size, 0)
        draw = ImageDraw.Draw(mask)
        draw.ellipse((0, 0, size[0], size[1]), fill=255)
        
        # Apply mask
        output = Image.new("RGBA", size, (0, 0, 0, 0))
        output.paste(img, (0, 0), mask=mask)
        
        img_byte_arr = io.BytesIO()
        output.save(img_byte_arr, format='PNG')
        img_byte_arr.seek(0)
        return img_byte_arr

    # === Layer 3: Testimonial Cards ===
    num_cards = len(testimonials)
    card_width = Inches(3.4)
    card_height = Inches(4.2)
    gap = Inches(0.5)
    
    # Center the entire block of cards
    total_width = (num_cards * card_width) + ((num_cards - 1) * gap)
    start_x = (prs.slide_width - total_width) / 2
    start_y = Inches(2.2)

    for i, data in enumerate(testimonials):
        # Card X position
        curr_x = start_x + i * (card_width + gap)

        # 1. Card Container (Rounded Rectangle)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, curr_x, start_y, card_width, card_height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = card_color
        card.line.fill.background()
        card.adjustments[0] = 0.05  # Subtle rounding

        # 2. Circular Avatar Image
        avatar_size = Inches(0.8)
        avatar_x = curr_x + (card_width - avatar_size) / 2
        avatar_y = start_y + Inches(0.4)
        avatar_stream = get_circular_avatar(data["avatar_url"])
        slide.shapes.add_picture(avatar_stream, avatar_x, avatar_y, avatar_size, avatar_size)

        # 3. Name Label
        name_box = slide.shapes.add_textbox(curr_x, avatar_y + avatar_size + Inches(0.1), card_width, Inches(0.4))
        nf = name_box.text_frame
        nf.word_wrap = True
        np = nf.paragraphs[0]
        np.text = data["name"]
        np.font.size = Pt(14)
        np.font.bold = True
        np.font.color.rgb = text_dark
        np.alignment = PP_ALIGN.CENTER

        # 4. Star Ratings
        star_box = slide.shapes.add_textbox(curr_x, avatar_y + avatar_size + Inches(0.45), card_width, Inches(0.4))
        sf = star_box.text_frame
        sp = sf.paragraphs[0]
        sp.text = "★★★★★"
        sp.font.size = Pt(16)
        sp.font.color.rgb = star_color
        sp.alignment = PP_ALIGN.CENTER

        # 5. Review Text
        text_box = slide.shapes.add_textbox(
            curr_x + Inches(0.2), 
            avatar_y + avatar_size + Inches(0.9), 
            card_width - Inches(0.4), 
            Inches(2.0)
        )
        tframe = text_box.text_frame
        tframe.word_wrap = True
        tp = tframe.paragraphs[0]
        tp.text = data["text"]
        tp.font.size = Pt(11)
        tp.font.color.rgb = text_light
        tp.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
