def create_slide(
    output_pptx_path: str,
    title_text: str = "WHY US",
    points: list = None,
    bg_color: tuple = (255, 255, 255),
    accent_color: tuple = (237, 28, 36),
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a 'Corporate Process Timeline' design.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        title_text (str): The main title for the slide.
        points (list): A list of strings for the timeline points. Defaults to a sample list.
        bg_color (tuple): RGB tuple for the background.
        accent_color (tuple): RGB tuple for accent elements.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image, ImageDraw
    import io

    if points is None:
        points = [
            "Loads of experience in presentation designing",
            "Turnaround time is extra fast",
            "100% satisfaction or money back guarantee",
            "Transitions, animation & video in presentation slides"
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color)

    # === Layer 2: Stylized Graphic (Generated with PIL) ===
    def create_segmented_ring():
        img_size = 400
        image = Image.new("RGBA", (img_size, img_size), (0, 0, 0, 0))
        draw = ImageDraw.Draw(image)
        
        center_x, center_y = img_size / 2, img_size / 2
        radius = 160
        thickness = 60
        
        # Colors
        blue_color = (33, 98, 222)
        red_color = accent_color
        
        # Draw base blue ring segments
        for i in range(8):
            start_angle = i * 45
            end_angle = start_angle + 35
            bbox = [center_x - radius, center_y - radius, center_x + radius, center_y + radius]
            draw.arc(bbox, start=start_angle, end=end_angle, fill=blue_color, width=thickness)

        # Draw red accent segment
        start_angle_red = 90
        end_angle_red = start_angle_red + 35
        bbox_red = [center_x - radius, center_y - radius, center_x + radius, center_y + radius]
        draw.arc(bbox_red, start=start_angle_red, end=end_angle_red, fill=red_color, width=thickness)
        
        # Draw a smaller red block for a 3D illusion
        block_size = 80
        block_pos = (center_x - radius - thickness/2, center_y - block_size/2)
        draw.rectangle(
            [block_pos[0], block_pos[1], block_pos[0] + thickness, block_pos[1] + block_size],
            fill=red_color
        )

        image_stream = io.BytesIO()
        image.save(image_stream, format='PNG')
        image_stream.seek(0)
        return image_stream

    ring_image_stream = create_segmented_ring()
    slide.shapes.add_picture(ring_image_stream, Inches(0.5), Inches(4.5), height=Inches(2.5))

    # === Layer 3: Text & Content (Timeline) ===
    # Title
    title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(1))
    text_frame = title_shape.text_frame
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial Black'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(64, 64, 64)

    # Timeline Axis
    timeline_x = Inches(9.5)
    timeline_start_y = Inches(1.8)
    timeline_height = Inches(4.5)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        timeline_x - Pt(1),
        timeline_start_y,
        Pt(2),
        timeline_height
    ).fill.solid.fore_color.rgb = RGBColor(200, 200, 200)

    # Timeline Points
    num_points = len(points)
    spacing = timeline_height / (num_points - 1) if num_points > 1 else 0
    node_diameter = Inches(0.4)
    text_box_width = Inches(4)
    text_box_height = Inches(0.5)

    for i, point_text in enumerate(points):
        node_y = timeline_start_y + (i * spacing) - (node_diameter / 2)
        
        # Node Circle
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, timeline_x - node_diameter/2, node_y, node_diameter, node_diameter)
        node.fill.solid()
        node.fill.fore_color.rgb = RGBColor(64, 64, 64)
        node.line.fill.background()

        # Node Number
        tf = node.text_frame
        tf.text = f"0{i+1}"
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.margin_bottom = tf.margin_top = tf.margin_left = tf.margin_right = 0
        
        # Connecting line with accent
        line_start_x = timeline_x - node_diameter/2 - Inches(0.2)
        line_y_center = node_y + node_diameter/2
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            line_start_x,
            line_y_center - Pt(0.5),
            Inches(0.2),
            Pt(1)
        ).fill.solid.fore_color.rgb = RGBColor(200, 200, 200)
        
        dot_size = Inches(0.1)
        slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            line_start_x - dot_size/2,
            line_y_center - dot_size/2,
            dot_size,
            dot_size
        ).fill.solid.fore_color.rgb = RGBColor(*accent_color)


        # Text box
        text_shape = slide.shapes.add_textbox(
            line_start_x - text_box_width,
            line_y_center - text_box_height/2,
            text_box_width,
            text_box_height
        )
        p = text_shape.text_frame.paragraphs[0]
        p.text = point_text
        p.font.name = 'Calibri'
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(89, 89, 89)
        p.alignment = PP_ALIGN.RIGHT

    prs.save(output_pptx_path)
    return output_pptx_path
