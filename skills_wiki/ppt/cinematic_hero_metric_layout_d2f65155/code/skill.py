import os
import urllib.request
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_cinematic_metric_slide(
    output_pptx_path: str = "cinematic_metric.pptx",
    pre_text: str = "MORE THAN",
    main_metric: str = "35",
    post_text: str = "YEARS OF EXPERIENCE",
    image_keyword: str = "technology,server",
    accent_color: tuple = (244, 176, 4),  # Gold
    overlay_color: tuple = (15, 23, 42, 190)  # Dark navy, ~75% opacity
) -> str:
    """
    Creates a PPTX file reproducing the Cinematic Hero Metric visual effect.
    """
    # 1. Setup Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # 2. Fetch Background Image
    bg_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1920x1080/?{image_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(bg_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception as e:
        print(f"Failed to download image: {e}. Using fallback solid background.")
        # Create a solid dark gray image as fallback
        img = Image.new('RGB', (1920, 1080), color=(40, 40, 40))
        img.save(bg_path)

    # 3. Create Semi-Transparent Overlay using PIL
    overlay_path = "temp_overlay.png"
    # Create an RGBA image filled with the overlay color
    overlay_img = Image.new('RGBA', (1920, 1080), color=overlay_color)
    overlay_img.save(overlay_path, "PNG")

    # 4. Add Background and Overlay to Slide
    slide.shapes.add_picture(bg_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.shapes.add_picture(overlay_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # 5. Add Core Visual Elements (Typography & Accents)
    
    # Layout Coordinates
    left_margin = Inches(1.5)
    
    # Pre-text (Top)
    top_pre_text = Inches(2.2)
    tx_box = slide.shapes.add_textbox(left_margin, top_pre_text, Inches(5), Inches(0.8))
    tf = tx_box.text_frame
    p = tf.add_paragraph()
    p.text = pre_text.upper()
    p.font.name = 'Arial'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Accent Line (Under Pre-text)
    line_top = Inches(3.0)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left_margin, line_top, Inches(4.5), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*accent_color)
    line.line.fill.background() # No border

    # Main Metric (Huge Number)
    top_metric = Inches(2.9)
    tx_box_metric = slide.shapes.add_textbox(left_margin, top_metric, Inches(5), Inches(3.0))
    tf_metric = tx_box_metric.text_frame
    p_metric = tf_metric.add_paragraph()
    p_metric.text = main_metric
    p_metric.font.name = 'Arial'
    p_metric.font.size = Pt(180)
    p_metric.font.bold = True
    p_metric.font.color.rgb = RGBColor(255, 255, 255)
    
    # Post-text (Context, sitting next to or under the main metric)
    # Using a technique to position it nicely relative to the big number
    top_post = Inches(5.5)
    tx_box_post = slide.shapes.add_textbox(left_margin, top_post, Inches(8), Inches(1.0))
    tf_post = tx_box_post.text_frame
    p_post = tf_post.add_paragraph()
    p_post.text = post_text.upper()
    p_post.font.name = 'Arial'
    p_post.font.size = Pt(28)
    p_post.font.bold = False
    p_post.font.color.rgb = RGBColor(200, 200, 200) # Slightly dimmed white

    # 6. Cleanup & Save
    prs.save(output_pptx_path)
    
    # Remove temp files
    if os.path.exists(bg_path):
        os.remove(bg_path)
    if os.path.exists(overlay_path):
        os.remove(overlay_path)

    return output_pptx_path

if __name__ == "__main__":
    # Test the function
    create_cinematic_metric_slide()
    print("Slide generated successfully.")
