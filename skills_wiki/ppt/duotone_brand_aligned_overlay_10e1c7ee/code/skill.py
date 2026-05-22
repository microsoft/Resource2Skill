def create_slide(
    output_pptx_path: str,
    title_text: str = "9%",
    body_text: str = "Only\ngets recycled",
    bg_palette: str = "environment",
    accent_color: tuple = (44, 76, 104),  # Slate Navy Blue
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Duotone Brand Overlay visual effect.
    """
    import os
    import urllib.request
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from lxml import etree

    # Initialize presentation (16:9 format)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Image (Grayscale processing via PIL) ===
    bg_path = "temp_bg_gray.jpg"
    try:
        # Fetch placeholder image
        url = f"https://picsum.photos/seed/{bg_palette}/1920/1080"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open("temp_raw_bg.jpg", "wb") as f:
                f.write(response.read())
        
        # Process image with Pillow
        img = Image.open("temp_raw_bg.jpg")
        
        # Crop exactly to 16:9 to prevent distortion
        target_aspect = 16 / 9
        w, h = img.size
        aspect = w / h
        if aspect > target_aspect:
            new_w = int(h * target_aspect)
            left = (w - new_w) / 2
            img = img.crop((left, 0, left + new_w, h))
        elif aspect < target_aspect:
            new_h = int(w / target_aspect)
            top = (h - new_h) / 2
            img = img.crop((0, top, w, top + new_h))
            
        # Desaturate (Convert to Grayscale)
        img = img.convert('L')
        img.save(bg_path)
    except Exception as e:
        # Fallback if download fails
        img = Image.new('L', (1920, 1080), color=100)
        img.save(bg_path)

    # Insert background
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Transparent Brand Color Overlay (lxml XML injection) ===
    overlay = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(*accent_color)
    overlay.line.fill.background()  # Remove outline

    # Inject alpha transparency natively into OOXML (70% opaque = 30% transparent)
    spPr = overlay.element
    srgbClr = spPr.find('.//a:srgbClr', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if srgbClr is not None:
        alpha = etree.SubElement(srgbClr, qn('a:alpha'))
        alpha.set('val', '70000')  # PPTX expects 100000 base for opacity

    # === Layer 3: Typography ===
    body_parts = body_text.split('\n')
    top_text = body_parts[0] if len(body_parts) > 0 else ""
    bottom_text = body_parts[1] if len(body_parts) > 1 else ""

    left_margin = Inches(1.0)

    # Microcopy: Top
    if top_text:
        txBox1 = slide.shapes.add_textbox(left_margin, Inches(1.8), Inches(5), Inches(1))
        tf1 = txBox1.text_frame
        tf1.margin_left = 0
        p1 = tf1.paragraphs[0]
        p1.text = top_text
        p1.font.size = Pt(36)
        p1.font.name = 'Arial'
        p1.font.color.rgb = RGBColor(255, 255, 255)

    # Macro-metric: Center Huge
    txBox2 = slide.shapes.add_textbox(left_margin, Inches(2.2), Inches(8), Inches(2.5))
    tf2 = txBox2.text_frame
    tf2.margin_left = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(160)
    p2.font.bold = True
    p2.font.name = 'Arial'
    p2.font.color.rgb = RGBColor(255, 255, 255)

    # Microcopy: Bottom
    if bottom_text:
        txBox3 = slide.shapes.add_textbox(left_margin, Inches(4.7), Inches(8), Inches(1))
        tf3 = txBox3.text_frame
        tf3.margin_left = 0
        p3 = tf3.paragraphs[0]
        p3.text = bottom_text
        p3.font.size = Pt(36)
        p3.font.bold = True
        p3.font.name = 'Arial'
        p3.font.color.rgb = RGBColor(255, 255, 255)

    # Save and clean up
    prs.save(output_pptx_path)
    for tmp in ["temp_raw_bg.jpg", bg_path]:
        if os.path.exists(tmp):
            os.remove(tmp)
            
    return output_pptx_path
