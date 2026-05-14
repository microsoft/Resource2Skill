def create_slide(
    output_pptx_path: str,
    hero_prefix: str = "BE",
    hero_suffix: str = "COME",
    top_accent_text: str = "YOUR OWN BOSS",
    side_text: str = "INDEPENDENT",
    bg_color_rgb: tuple = (29, 23, 51),
    accent_color_rgb: tuple = (243, 54, 100),
    text_color_rgb: tuple = (255, 255, 255),
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the "Kinetic Brutalist Typography" visual effect.
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import OxmlElement

    # Helper function to guarantee shape background transparency via lxml
    def make_shape_transparent(shape):
        spPr = shape.element.spPr
        for fill_tag in ['a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:noFill']:
            for fill in spPr.xpath(fill_tag):
                spPr.remove(fill)
        spPr.append(OxmlElement('a:noFill'))

    # Helper function to remove internal margins from textboxes for a tight graphic feel
    def remove_text_margins(text_frame):
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*bg_color_rgb)

    # === Layer 2: Offset Outline Frame ===
    # Placed first so it renders behind the hero block
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.5), Inches(2.2), Inches(2.8), Inches(2.8)
    )
    make_shape_transparent(frame)
    frame.line.color.rgb = RGBColor(*text_color_rgb)
    frame.line.width = Pt(4.5)

    # === Layer 3: Solid Accent Hero Box ===
    hero_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.2), Inches(2.6), Inches(2.8), Inches(2.8)
    )
    hero_box.fill.solid()
    hero_box.fill.fore_color.rgb = RGBColor(*accent_color_rgb)
    hero_box.line.color.rgb = RGBColor(*accent_color_rgb) # match fill to hide line
    
    tf_hero = hero_box.text_frame
    tf_hero.text = hero_prefix
    p = tf_hero.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(120)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(*text_color_rgb)
    remove_text_margins(tf_hero)

    # === Layer 4: Top Subtext Highlight Bar ===
    top_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.1), Inches(2.6), Inches(3.8), Inches(0.6)
    )
    top_box.fill.solid()
    top_box.fill.fore_color.rgb = RGBColor(*accent_color_rgb)
    top_box.line.color.rgb = RGBColor(*accent_color_rgb)
    
    tf_top = top_box.text_frame
    tf_top.text = top_accent_text
    p = tf_top.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(24)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(*text_color_rgb)
    remove_text_margins(tf_top)

    # === Layer 5: Main Headline Text ===
    main_text = slide.shapes.add_textbox(
        Inches(4.9), Inches(3.2), Inches(5.0), Inches(1.5)
    )
    tf_main = main_text.text_frame
    tf_main.text = hero_suffix
    p = tf_main.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.font.size = Pt(110)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(*text_color_rgb)
    remove_text_margins(tf_main)

    # === Layer 6: Sideways Vertical Text ===
    side_text_shape = slide.shapes.add_textbox(
        Inches(8.0), Inches(4.3), Inches(3.0), Inches(0.5)
    )
    side_text_shape.rotation = 270.0 # Rotate -90 degrees
    tf_side = side_text_shape.text_frame
    tf_side.text = side_text
    p = tf_side.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.name = "Arial Black"
    p.font.color.rgb = RGBColor(*text_color_rgb)
    remove_text_margins(tf_side)

    prs.save(output_pptx_path)
    return output_pptx_path
