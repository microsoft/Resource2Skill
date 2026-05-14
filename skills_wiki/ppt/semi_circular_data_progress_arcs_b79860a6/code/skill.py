def create_slide(
    output_pptx_path: str,
    title_text: str = "PERCENTAGES",
    subtitle_text: str = "This is a demo text you may write a brief text here to explain the title or if you think you do\nnot need this you may consider deleting the text box.",
    segments: list = None,
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Semi-Circular Data Progress Arcs' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw
    import io

    # Default segment data mapping the tutorial's exact numbers and color palette
    if segments is None:
        segments = [
            {"pct": 60, "color": (233, 30, 99), "title": "GRAPHIC DESIGN", "desc": "Here You Should Add\nSome Brief Text to Explain\nMain Title"},
            {"pct": 70, "color": (0, 150, 136), "title": "WEB DESIGN", "desc": "Here You Should Add\nSome Brief Text to Explain\nMain Title"},
            {"pct": 50, "color": (139, 195, 74), "title": "VIDEO EDITING", "desc": "Here You Should Add\nSome Brief Text to Explain\nMain Title"},
            {"pct": 90, "color": (63, 81, 181), "title": "UX DESIGN", "desc": "Here You Should Add\nSome Brief Text to Explain\nMain Title"},
        ]

    # Initialize presentation (16:9 standard)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Helper: Text Box Generator
    def add_formatted_text(left, top, width, height, text, font_size, font_color, is_bold=False):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = is_bold
        run.font.name = "Century Gothic"
        run.font.color.rgb = font_color
        return txBox

    # Helper: PIL Arc Generator for smooth high-res progress rings
    def get_arc_stream(pct, color):
        scale = 4  # Render large for antialiasing
        base_size = 400
        size = base_size * scale
        thickness = 40 * scale
        
        img = Image.new('RGBA', (size, size), (255, 255, 255, 0))
        draw = ImageDraw.Draw(img)
        bbox = [thickness/2, thickness/2, size - thickness/2, size - thickness/2]
        
        # Draw background track (180 to 360 sweeps the top half)
        draw.arc(bbox, 180, 360, fill=(235, 235, 235, 255), width=thickness)
        
        # Draw percentage fill
        fill_end = 180 + (pct / 100.0) * 180
        draw.arc(bbox, 180, fill_end, fill=color + (255,), width=thickness)
        
        # Crop to the exact semi-circle bounding area
        crop_bottom = int(size/2 + thickness/2 + 2 * scale)
        img = img.crop((0, 0, size, crop_bottom))
        
        # Downscale for crisp anti-aliased edges
        img = img.resize((base_size, int(crop_bottom/scale)), Image.Resampling.LANCZOS)
        
        img_stream = io.BytesIO()
        img.save(img_stream, format='PNG')
        img_stream.seek(0)
        return img_stream

    # --- Layer 1: Global Titles ---
    # Convert "PERCENTAGES" to spaced out string "P  E  R  C  E  N  T  A  G  E  S"
    title_spaced = "  ".join(list(title_text.replace(" ", "")))
    add_formatted_text(Inches(2.0), Inches(0.6), Inches(9.33), Inches(0.8), 
                       title_spaced, 36, RGBColor(160, 160, 160), is_bold=True)
    
    add_formatted_text(Inches(2.0), Inches(1.4), Inches(9.33), Inches(0.8), 
                       subtitle_text, 12, RGBColor(158, 158, 158), is_bold=False)

    # --- Layer 2: Dashboard Grid Assembly ---
    num_items = len(segments)
    # Distribute centers evenly across the canvas
    centers = [2.166 + (i * 3.0) for i in range(num_items)]

    for i, seg in enumerate(segments):
        center_x = centers[i]
        seg_color = RGBColor(*seg['color'])
        
        # Insert PIL Progress Arc
        arc_stream = get_arc_stream(seg['pct'], seg['color'])
        arc_width = 2.4
        slide.shapes.add_picture(arc_stream, Inches(center_x - arc_width/2), Inches(3.0), width=Inches(arc_width))
        
        # Insert Percentage Text
        add_formatted_text(Inches(center_x - 1.0), Inches(4.4), Inches(2.0), Inches(0.6),
                           f"{seg['pct']}%", 32, seg_color, is_bold=True)
        
        # Insert Title Text
        add_formatted_text(Inches(center_x - 1.0), Inches(5.1), Inches(2.0), Inches(0.4),
                           seg['title'], 14, seg_color, is_bold=True)
        
        # Insert Description Text
        add_formatted_text(Inches(center_x - 1.2), Inches(5.4), Inches(2.4), Inches(1.0),
                           seg['desc'], 11, RGBColor(158, 158, 158), is_bold=False)

    prs.save(output_pptx_path)
    return output_pptx_path
