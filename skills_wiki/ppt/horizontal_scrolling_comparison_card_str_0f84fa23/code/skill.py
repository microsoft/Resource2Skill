def create_slide(
    output_pptx_path: str = "Horizontal_Comparison_Strip.pptx",
    title_text: str = "Comparison Video Template",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 'Horizontal Scrolling Comparison Card Strip' visual effect.
    Constructs a panoramic sequence of cards extending off-screen.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from PIL import Image, ImageDraw
    import io

    # 1. Initialize Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use a blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 2. Set dark background for the slide
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(25, 25, 25)

    # 3. Define Card Dimensions & Layout Data
    num_cards = 8  # Create 8 items to ensure it goes well off-screen
    col_width = Inches(3.333) # 4 items visible on a 13.333 inch slide
    
    # Y-positions and heights
    y_rank = Inches(0.5)
    h_rank = Inches(1.2)
    
    y_name = y_rank + h_rank
    h_name = Inches(1.0)
    
    y_img = y_name + h_name
    h_img = Inches(2.3)
    
    y_desc = y_img + h_img
    h_desc = Inches(2.0)

    # Sample data to simulate the comparison list
    data = [
        {"rank": "10", "name": "Brazil", "color": (0, 156, 59), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "9", "name": "Thailand", "color": (237, 28, 36), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "8", "name": "Japan", "color": (255, 255, 255), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "7", "name": "Netherlands", "color": (33, 70, 139), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "6", "name": "South Korea", "color": (255, 255, 255), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "5", "name": "Bolivia", "color": (244, 228, 0), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "4", "name": "South Africa", "color": (0, 119, 73), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."},
        {"rank": "3", "name": "Australia", "color": (1, 33, 105), "desc": "The quick brown fox jumps over the lazy dog.\nThe quick brown fox jumps over the lazy dog."}
    ]

    # Helper function to create a solid color dummy flag
    def create_dummy_flag(rgb_color):
        img = Image.new('RGB', (400, 300), color=rgb_color)
        draw = ImageDraw.Draw(img)
        # Add a subtle border to white flags
        if rgb_color == (255, 255, 255):
            draw.rectangle([0, 0, 399, 299], outline=(200, 200, 200), width=3)
        # Add a placeholder circle in the center to look more "flag-like"
        circle_color = (200, 0, 0) if rgb_color == (255, 255, 255) else (255, 255, 255)
        draw.ellipse([150, 100, 250, 200], fill=circle_color)
        
        img_io = io.BytesIO()
        img.save(img_io, format='PNG')
        img_io.seek(0)
        return img_io

    # 4. Generate the Panoramic Strip
    for i in range(num_cards):
        left = i * col_width
        item = data[i]

        # --- RANK BLOCK ---
        rank_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, y_rank, col_width, h_rank
        )
        rank_shape.fill.solid()
        rank_shape.fill.fore_color.rgb = RGBColor(192, 0, 0)  # Red
        rank_shape.line.fill.background() # No border
        
        rank_tf = rank_shape.text_frame
        rank_tf.text = item["rank"]
        rank_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        rank_tf.paragraphs[0].font.size = Pt(64)
        rank_tf.paragraphs[0].font.bold = True
        rank_tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

        # --- NAME BLOCK ---
        name_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, y_name, col_width, h_name
        )
        name_shape.fill.solid()
        name_shape.fill.fore_color.rgb = RGBColor(217, 217, 217)  # Light Grey
        name_shape.line.fill.background()
        
        name_tf = name_shape.text_frame
        name_tf.text = item["name"]
        name_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        name_tf.paragraphs[0].font.size = Pt(36)
        name_tf.paragraphs[0].font.bold = True
        name_tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)

        # --- IMAGE BLOCK ---
        # Create a visually pleasing placeholder flag
        flag_io = create_dummy_flag(item["color"])
        
        # We place the image centered in its block
        img_w = Inches(2.4)
        img_h = Inches(1.8)
        img_left = left + (col_width - img_w) / 2
        img_top = y_img + (h_img - img_h) / 2
        
        pic = slide.shapes.add_picture(flag_io, img_left, img_top, img_w, img_h)

        # --- DESCRIPTION TEXT ---
        desc_box = slide.shapes.add_textbox(left, y_desc, col_width, h_desc)
        desc_tf = desc_box.text_frame
        desc_tf.word_wrap = True
        p = desc_tf.paragraphs[0]
        p.text = item["desc"]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        # --- VERTICAL SEPARATOR LINE ---
        # Add a dark grey separator between columns (except the very first edge)
        if i > 0:
            sep_line = slide.shapes.add_connector(
                MSO_SHAPE.LINE_INVERSE, left, Inches(0.5), left, Inches(7.0)
            )
            sep_line.line.color.rgb = RGBColor(10, 10, 10)
            sep_line.line.width = Pt(4)

    # 5. Save the Presentation
    prs.save(output_pptx_path)
    print(f"Comparison panoramic template saved to {output_pptx_path}")
    return output_pptx_path

if __name__ == "__main__":
    create_slide()
