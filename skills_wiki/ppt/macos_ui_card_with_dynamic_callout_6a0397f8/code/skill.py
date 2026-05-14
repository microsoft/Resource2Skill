import os
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

def _apply_dropshadow(shape, blur_pt=15, distance_pt=5, angle_deg=90, alpha_pct=25):
    """
    Injects a high-quality outer shadow into a shape's XML properties to create Z-depth.
    Works on both standard shapes and textboxes (applying to the text).
    """
    spPr = shape.element.spPr
    effectLst = spPr.find(qn('a:effectLst'))
    if effectLst is None:
        effectLst = etree.SubElement(spPr, qn('a:effectLst'))
    
    # Clean up existing shadow if present
    for outerShdw in effectLst.findall(qn('a:outerShdw')):
        effectLst.remove(outerShdw)
        
    blur_emu = int(blur_pt * 12700)
    dist_emu = int(distance_pt * 12700)
    dir_val = int(angle_deg * 60000)
    alpha_val = int(alpha_pct * 1000)
    
    outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'), 
                                 blurRad=str(blur_emu), 
                                 dist=str(dist_emu), 
                                 dir=str(dir_val), 
                                 algn="ctr", 
                                 rotWithShape="0")
    srgbClr = etree.SubElement(outerShdw, qn('a:srgbClr'), val="000000")
    etree.SubElement(srgbClr, qn('a:alpha'), val=str(alpha_val))

def _generate_studio_bg(filepath: str, width_px=1920, height_px=1080):
    """Generates a soft, atmospheric vertical gradient image to simulate studio lighting."""
    # Create a 1D gradient and stretch it for maximum performance and smoothness
    base = Image.new('RGB', (1, 256))
    draw = ImageDraw.Draw(base)
    
    color_top = (245, 245, 248)    # Crisp cool-white
    color_bottom = (205, 205, 212) # Soft shadow gray
    
    for y in range(256):
        ratio = y / 255.0
        r = int(color_top[0] * (1 - ratio) + color_bottom[0] * ratio)
        g = int(color_top[1] * (1 - ratio) + color_bottom[1] * ratio)
        b = int(color_top[2] * (1 - ratio) + color_bottom[2] * ratio)
        draw.point((0, y), fill=(r, g, b))
        
    # Stretch linearly to fill the screen
    bg = base.resize((width_px, height_px), Image.BICUBIC)
    bg.save(filepath, format="PNG")
    return filepath

def create_slide(
    output_pptx_path: str,
    callout_text: str = "NOTABLE",
    window_title: str = "London Trip ☕️",
    window_body: str = "Dates: 25-31 October\n\nDay 1: Arrive in the evening, walk around Covent Garden, and have dinner at Dishoom.\n\nDay 2: Visit the British Museum in the morning, grab lunch nearby, and explore Soho in the afternoon.",
    callout_color: tuple = (212, 255, 0), # Neon Yellow-Green
    **kwargs,
) -> str:
    """
    Creates a slide featuring a faux macOS UI window overlapping a soft background,
    accented by a dynamic, rotated callout text.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # === Layer 1: Atmospheric Background ===
    bg_path = "temp_studio_bg.png"
    _generate_studio_bg(bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    if os.path.exists(bg_path):
        os.remove(bg_path)

    # === Layer 2: macOS UI Card ===
    card_width = Inches(8.5)
    card_height = Inches(5.5)
    card_left = Inches(3.5)
    card_top = Inches(1.0)
    
    # Draw main window container
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_width, card_height)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(230, 230, 230) # Subtle border
    card.line.width = Pt(1)
    
    # Refine corner radius to mimic macOS (approx 3%)
    if len(card.adjustments) > 0:
        card.adjustments[0] = 0.03
        
    # Apply soft ambient drop shadow to elevate the window
    _apply_dropshadow(card, blur_pt=25, distance_pt=8, angle_deg=90, alpha_pct=15)

    # Add Traffic Light Window Controls
    dot_radius = Inches(0.12)
    dot_spacing = Inches(0.08)
    dot_y = card_top + Inches(0.15)
    start_x = card_left + Inches(0.2)
    
    colors = [
        RGBColor(255, 95, 86),  # Close Red
        RGBColor(255, 189, 46), # Minimize Yellow
        RGBColor(39, 201, 63)   # Expand Green
    ]
    
    for i, color in enumerate(colors):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            start_x + (i * (dot_radius + dot_spacing)), 
            dot_y, 
            dot_radius, 
            dot_radius
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background() # No border

    # Add Window Content (Text)
    content_box = slide.shapes.add_textbox(card_left + Inches(0.5), card_top + Inches(0.6), card_width - Inches(1), card_height - Inches(1))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    # Title Paragraph
    p_title = tf.paragraphs[0]
    p_title.text = window_title
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(30, 30, 30)
    
    # Body Paragraph
    p_body = tf.add_paragraph()
    p_body.text = "\n" + window_body
    p_body.font.name = "Calibri"
    p_body.font.size = Pt(16)
    p_body.font.color.rgb = RGBColor(80, 80, 80)
    p_body.line_spacing = 1.3

    # === Layer 3: Dynamic Callout Text ===
    # Positioned to overlap the bottom-left edge of the UI card
    callout_left = Inches(0.8)
    callout_top = Inches(4.5)
    callout_width = Inches(6.0)
    callout_height = Inches(2.0)
    
    callout = slide.shapes.add_textbox(callout_left, callout_top, callout_width, callout_height)
    # Tilt the text for dynamic energy
    callout.rotation = -8 
    
    c_tf = callout.text_frame
    c_p = c_tf.paragraphs[0]
    c_p.text = callout_text.upper()
    c_p.font.name = "Arial Black" # Standard ultra-bold font
    c_p.font.size = Pt(72)
    c_p.font.bold = True
    c_p.font.color.rgb = RGBColor(*callout_color)
    
    # Add a tighter, darker drop shadow directly to the text to make the neon color pop
    _apply_dropshadow(callout, blur_pt=10, distance_pt=6, angle_deg=90, alpha_pct=30)

    prs.save(output_pptx_path)
    return output_pptx_path
