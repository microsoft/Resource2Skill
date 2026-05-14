import os
import math
from typing import Tuple
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFilter

def draw_hex_grid(draw: ImageDraw.Draw, width: int, height: int, hex_size: int, color: Tuple[int, int, int, int]):
    """Draws a faint hexagonal grid over the entire image."""
    hex_width = math.sqrt(3) * hex_size
    hex_height = 2 * hex_size
    
    col_spacing = hex_width
    row_spacing = hex_height * 0.75
    
    cols = int(width / col_spacing) + 2
    rows = int(height / row_spacing) + 2
    
    for row in range(rows):
        for col in range(cols):
            # Offset every other row
            x_offset = (hex_width / 2) if row % 2 == 1 else 0
            cx = col * col_spacing + x_offset
            cy = row * row_spacing
            
            # Calculate points for a pointy-topped hexagon
            points = []
            for i in range(6):
                angle_deg = 60 * i - 30
                angle_rad = math.radians(angle_deg)
                px = cx + hex_size * math.cos(angle_rad)
                py = cy + hex_size * math.sin(angle_rad)
                points.append((px, py))
                
            draw.polygon(points, outline=color, width=1)

def create_slide(
    output_pptx_path: str,
    title_text: str = "ARTIFICIAL\nINTELLIGENCE (AI)",
    subtitle_text: str = "for research and academia\n\nDr. AI Assistant",
    bg_color: tuple = (18, 14, 31),
    blob_color_1: tuple = (150, 0, 255), # Deep Purple
    blob_color_2: tuple = (0, 255, 255), # Neon Cyan
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Neon Cyber-Academic Aesthetic.
    Generates a custom background with a hex grid and blurred organic shapes.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Layer 1: Background Generation (PIL) ===
    # Dimensions based on slide size at 150 DPI for good resolution
    img_width = int(13.333 * 150)
    img_height = int(7.5 * 150)
    
    # Base dark background
    base_img = Image.new('RGBA', (img_width, img_height), bg_color + (255,))
    
    # Draw Hex Grid (very faint)
    grid_overlay = Image.new('RGBA', (img_width, img_height), (0, 0, 0, 0))
    grid_draw = ImageDraw.Draw(grid_overlay)
    draw_hex_grid(grid_draw, img_width, img_height, hex_size=40, color=(255, 255, 255, 8)) # 8/255 alpha is very faint
    
    base_img = Image.alpha_composite(base_img, grid_overlay)

    # === Layer 2: Organic Blurred Blobs (PIL) ===
    blob_layer = Image.new('RGBA', (img_width, img_height), (0, 0, 0, 0))
    blob_draw = ImageDraw.Draw(blob_layer)
    
    # Draw large ellipses
    # Top Right Blob (Purple/Magenta)
    blob_draw.ellipse(
        [img_width * 0.6, -img_height * 0.2, img_width * 1.1, img_height * 0.6], 
        fill=blob_color_1 + (180,)
    )
    
    # Bottom Left Blob (Cyan)
    blob_draw.ellipse(
        [-img_width * 0.1, img_height * 0.5, img_width * 0.4, img_height * 1.2], 
        fill=blob_color_2 + (150,)
    )
    
    # Top Left small accent blob
    blob_draw.ellipse(
        [img_width * 0.1, img_height * 0.1, img_width * 0.3, img_height * 0.4], 
        fill=(255, 0, 255, 100) # Magenta
    )

    # Apply massive blur to create the light-leak/organic effect
    blob_layer = blob_layer.filter(ImageFilter.GaussianBlur(radius=120))
    
    # Composite blobs over background
    final_bg = Image.alpha_composite(base_img, blob_layer)
    
    # Save temp background
    bg_path = "temp_cyber_bg.png"
    final_bg.save(bg_path)
    
    # Insert Background into PPTX
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 3: Text & Content (python-pptx) ===
    
    # Title Text Box
    title_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2.5))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    
    p_title = title_tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Arial Black' # Use a heavy, blocky font
    p_title.font.size = Pt(64)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)
    
    # Subtitle Text Box
    sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.2), Inches(8), Inches(1.5))
    sub_tf = sub_box.text_frame
    sub_tf.word_wrap = True
    
    p_sub = sub_tf.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(28)
    p_sub.font.color.rgb = RGBColor(220, 220, 230) # Slightly off-white

    # Decorative Neon Accent Line
    line = slide.shapes.add_shape(
        1, # msoShapeRectangle
        Inches(1), Inches(4.0), Inches(4), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(blob_color_2[0], blob_color_2[1], blob_color_2[2]) # Match Cyan
    line.line.fill.background() # No border

    prs.save(output_pptx_path)
    
    # Clean up temp image
    if os.path.exists(bg_path):
        os.remove(bg_path)
        
    return output_pptx_path
