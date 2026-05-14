def create_slide(
    output_pptx_path: str,
    title_text: str = "2024",
    body_text: str = "TITLE SLIDE\nINTRO",
    bg_palette: str = "city,architecture,night",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Dynamic Layered Chevron Reveal' visual effect.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageDraw

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # --- Color Palette ---
    COLOR_BG_NAVY = (43, 50, 60)
    COLOR_YELLOW = (242, 194, 0)
    COLOR_SLATE = (89, 102, 117)
    COLOR_WHITE = (255, 255, 255)

    # --- Set Slide Background ---
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*COLOR_BG_NAVY)

    # --- Helper: Image Downloader with Fallback ---
    img_path = "temp_chevron_bg.jpg"
    try:
        # Try fetching a high-quality relevant image
        url = f"https://images.unsplash.com/photo-1449844908441-8829872d2607?q=80&w=1600&auto=format&fit=crop"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Fallback: Create a gradient image using PIL if download fails
        img = Image.new('RGB', (1600, 900))
        draw = ImageDraw.Draw(img)
        for y in range(900):
            r = int(20 + (y / 900) * 40)
            g = int(40 + (y / 900) * 60)
            b = int(80 + (y / 900) * 120)
            draw.line([(0, y), (1600, y)], fill=(r, g, b))
        img.save(img_path)

    # --- Helper: Build Chevron Polygon ---
    def add_chevron(x_base_inches, tip_depth_inches, color_rgb=None, fill_image=None):
        """Draws a right-pointing chevron anchored to the left side."""
        ff_builder = slide.shapes.build_freeform()
        # Define vertices
        p1 = (Inches(0), Inches(0))
        p2 = (Inches(x_base_inches), Inches(0))
        p3 = (Inches(x_base_inches + tip_depth_inches), Inches(7.5 / 2)) # Middle point
        p4 = (Inches(x_base_inches), Inches(7.5))
        p5 = (Inches(0), Inches(7.5))
        
        ff_builder.add_line_segments([p1, p2, p3, p4, p5], close=True)
        shape = ff_builder.convert_to_shape()
        
        # Remove outline
        shape.line.fill.background()
        
        # Apply fill
        if fill_image:
            shape.fill.user_picture(fill_image)
        elif color_rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*color_rgb)
            shape.line.color.rgb = RGBColor(*color_rgb) # Clean edge
            
        return shape

    # --- Create Layered Chevrons ---
    # Tip depth must be constant to maintain parallel lines
    TIP_DEPTH = 2.5 
    
    # Layer 1: Yellow (Base)
    add_chevron(x_base_inches=8.5, tip_depth_inches=TIP_DEPTH, color_rgb=COLOR_YELLOW)
    
    # Layer 2: Slate Grey
    add_chevron(x_base_inches=7.7, tip_depth_inches=TIP_DEPTH, color_rgb=COLOR_SLATE)
    
    # Layer 3: Image Container
    img_chevron = add_chevron(x_base_inches=6.9, tip_depth_inches=TIP_DEPTH, fill_image=img_path)

    # --- Add Text Overlay ---
    # Title Box (e.g., "2019")
    tx_box_title = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.0), Inches(1.5))
    tf_title = tx_box_title.text_frame
    tf_title.clear()
    p = tf_title.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial Black" # Standard fallback for heavy bold
    p.font.size = Pt(88)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*COLOR_WHITE)
    
    # Subtitle Box (e.g., "Intro Slide")
    tx_box_sub = slide.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(4.0), Inches(1.0))
    tf_sub = tx_box_sub.text_frame
    tf_sub.clear()
    p2 = tf_sub.paragraphs[0]
    p2.text = body_text
    p2.font.name = "Arial"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(*COLOR_WHITE)

    # --- Cleanup & Save ---
    prs.save(output_pptx_path)
    
    # Clean up temp image
    if os.path.exists(img_path):
        os.remove(img_path)
        
    return output_pptx_path
