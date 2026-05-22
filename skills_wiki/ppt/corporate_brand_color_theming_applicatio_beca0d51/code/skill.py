def create_slide(
    output_pptx_path: str,
    title_text: str = "Sales Report Q1",
    palette: list = [(2, 9, 77), (62, 137, 255), (173, 205, 255), (214, 213, 215)],
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Corporate Brand Color Theming layout.
    
    :param output_pptx_path: Path to save the PPTX file.
    :param title_text: Main title of the slide.
    :param palette: A list of 4 RGB tuples (Dark, Medium, Light, Grey) representing the brand.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    # Map palette hierarchically
    color_dark = RGBColor(*palette[0])
    color_medium = RGBColor(*palette[1])
    color_light = RGBColor(*palette[2])
    color_grey = RGBColor(*palette[3])
    color_white = RGBColor(255, 255, 255)
    
    # === 1. Header & Style Guide Strip ===
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = color_dark
    
    # Mini style guide swatches at top right
    strip_w, strip_h = Inches(0.5), Inches(0.15)
    start_x = Inches(13.333) - 4 * strip_w - Inches(0.8)
    for i, c in enumerate(palette):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, start_x + i*strip_w, Inches(0.7), strip_w, strip_h
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*c)
        shape.line.fill.background()
        
    # === 2. Left Column: KPI List ===
    data_points = [
        ("+10K", "Sales Data 1", "Presentations are communication tools that can be used as demonstrations.", color_light, color_dark),
        ("+50K", "Sales Data 2", "Presentations are communication tools that can be used as demonstrations.", color_medium, color_white),
        ("+19K", "Sales Data 3", "Presentations are communication tools that can be used as demonstrations.", color_dark, color_white)
    ]
    
    y_offset = Inches(2.0)
    for value, sub_title, desc, bg_col, text_col in data_points:
        # KPI Circle Badge
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y_offset, Inches(1.0), Inches(1.0))
        circ.fill.solid()
        circ.fill.fore_color.rgb = bg_col
        circ.line.fill.background()
        
        tf_circ = circ.text_frame
        tf_circ.text = value
        tf_circ.paragraphs[0].font.size = Pt(16)
        tf_circ.paragraphs[0].font.bold = True
        tf_circ.paragraphs[0].font.color.rgb = text_col
        tf_circ.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # KPI Text
        tb = slide.shapes.add_textbox(Inches(2.0), y_offset - Inches(0.1), Inches(3.2), Inches(1.2))
        tb.text_frame.word_wrap = True
        
        p1 = tb.text_frame.paragraphs[0]
        p1.text = sub_title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = color_dark
        
        p2 = tb.text_frame.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(100, 100, 100)
        
        y_offset += Inches(1.6)
        
    # === 3. Right Column: Brand Themed Chart ===
    chart_data = CategoryChartData()
    chart_data.categories = ['Item 1', 'Item 2', 'Item 3', 'Item 4', 'Item 5']
    chart_data.add_series('Series 1', (16.2, 18.5, 15.0, 17.3, 18.1))
    chart_data.add_series('Series 2', (20.5, 22.0, 19.4, 22.1, 23.0))
    
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.5), Inches(1.5), Inches(7.2), Inches(5.0), chart_data
    )
    chart = chart_shape.chart
    chart.has_legend = False
    
    # Apply brand colors to chart series
    series_colors = [color_light, color_medium]
    for i, series in enumerate(chart.plots[0].series):
        fill = series.format.fill
        fill.solid()
        fill.fore_color.rgb = series_colors[i]
        
    # Style chart axes and gridlines
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    try:
        value_axis.major_gridlines.format.line.color.rgb = color_grey
    except Exception:
        pass  # Failsafe for older pptx versions
        
    # === 4. Chart Annotation Overlay ===
    # Emphasis Badge
    ann_circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.8), Inches(1.8), Inches(0.8), Inches(0.8))
    ann_circ.fill.solid()
    ann_circ.fill.fore_color.rgb = color_dark
    ann_circ.line.fill.background()
    
    tf_ann = ann_circ.text_frame
    tf_ann.text = "100%"
    tf_ann.paragraphs[0].font.size = Pt(16)
    tf_ann.paragraphs[0].font.bold = True
    tf_ann.paragraphs[0].font.color.rgb = color_white
    tf_ann.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Emphasis Connector Arrow
    connector = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.7), Inches(2.1), Inches(2.0), Inches(0.2))
    connector.fill.solid()
    connector.fill.fore_color.rgb = color_dark
    connector.line.color.rgb = color_white
    connector.line.width = Pt(1)
    
    prs.save(output_pptx_path)
    return output_pptx_path
