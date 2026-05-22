from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_slide(
    output_pptx_path: str,
    company_name: str = "Company Name",
    tagline: str = "Your Local experts in\nCity State",
    phone_number: str = "(xxx) xxx-xxxx",
    website_url: str = "www.yourwebsite.com",
    **kwargs
) -> str:
    """
    Creates a PPTX file with a single, minimalist call-to-action slide.

    This reproduces the "Minimalist Contact Outro" style, featuring high-contrast,
    centered white text on a black background.

    Args:
        output_pptx_path: The path to save the generated .pptx file.
        company_name: The main company name (top line).
        tagline: A descriptor or location, can include a newline character.
        phone_number: The contact phone number.
        website_url: The company website URL.
        kwargs: Not used in this implementation, but included for signature consistency.

    Returns:
        The path to the saved PPTX file.
    """
    # --- Presentation and Slide Setup (16:9 Aspect Ratio) ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # === Layer 1: Background ===
    # Set the background to solid black.
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)

    # === Layer 2: Text Content ===
    # A single textbox is used for easy centering of the entire block.
    # Dimensions are approximate; the key is centering it on the slide.
    width = Inches(10)
    height = Inches(5)
    left = (prs.slide_width - width) / 2
    top = (prs.slide_height - height) / 2
    
    textbox = slide.shapes.add_textbox(left, top, width, height)
    text_frame = textbox.text_frame
    text_frame.word_wrap = True # Ensure text wraps if it's too long

    # --- Text Hierarchy and Styling ---
    # Paragraph 1: Company Name
    p1 = text_frame.paragraphs[0]
    p1.text = company_name
    p1.alignment = PP_ALIGN.CENTER
    font1 = p1.font
    font1.name = 'Arial'
    font1.size = Pt(60)
    font1.bold = True
    font1.color.rgb = RGBColor(255, 255, 255)

    # Paragraph 2: Tagline
    p2 = text_frame.add_paragraph()
    p2.text = tagline
    p2.alignment = PP_ALIGN.CENTER
    font2 = p2.font
    font2.name = 'Arial'
    font2.size = Pt(40)
    font2.bold = True
    font2.color.rgb = RGBColor(255, 255, 255)

    # Add extra spacing before contact info
    p_spacer = text_frame.add_paragraph()
    p_spacer.text = ""
    p_spacer.font.size = Pt(24)

    # Paragraph 3: Phone Number
    p3 = text_frame.add_paragraph()
    p3.text = phone_number
    p3.alignment = PP_ALIGN.CENTER
    font3 = p3.font
    font3.name = 'Arial'
    font3.size = Pt(36)
    font3.bold = True
    font3.color.rgb = RGBColor(255, 255, 255)

    # Paragraph 4: Website URL
    p4 = text_frame.add_paragraph()
    p4.text = website_url
    p4.alignment = PP_ALIGN.CENTER
    font4 = p4.font
    font4.name = 'Arial'
    font4.size = Pt(36)
    font4.bold = True
    font4.color.rgb = RGBColor(255, 255, 255)

    # --- Save the Presentation ---
    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# if __name__ == '__main__':
#     create_slide(
#         output_pptx_path="Minimalist_Contact_Outro.pptx",
#         company_name="123 SEO NOW",
#         tagline="Internet Marketing Training Videos\nfor Small Businesses",
#         phone_number="(888) 555-1234",
#         website_url="www.123seonow.com"
#     )
