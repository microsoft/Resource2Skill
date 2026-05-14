def create_slide(
    output_pptx_path: str,
    title_text: str = "Cross-Functional E-Commerce Process Flow",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing a Cross-Functional Swimlane Flowchart.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.dml import MSO_ARROWHEAD_STYLE
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Set background to pure white
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Add Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11.5), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(44, 62, 80)

    # === Layer 1: Swimlane Backgrounds & Headers ===
    lanes = [
        {"name": "Customer", "bg": (234, 242, 248), "header_bg": (41, 128, 185)},
        {"name": "Sales Dept", "bg": (233, 247, 239), "header_bg": (39, 174, 96)},
        {"name": "Warehouse", "bg": (235, 237, 239), "header_bg": (52, 73, 94)},
        {"name": "Accounting", "bg": (253, 242, 233), "header_bg": (211, 84, 0)},
    ]

    lane_y_start = 1.2
    lane_height = 1.4
    slide_width = 13.333
    header_width = 0.8

    for i, lane in enumerate(lanes):
        y = lane_y_start + i * lane_height
        
        # Lane Body
        body = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(header_width), Inches(y), 
            Inches(slide_width - header_width), Inches(lane_height)
        )
        body.fill.solid()
        body.fill.fore_color.rgb = RGBColor(*lane["bg"])
        body.line.color.rgb = RGBColor(255, 255, 255)
        body.line.width = Pt(1.5)
        
        # Lane Header Background
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), 
            Inches(header_width), Inches(lane_height)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor(*lane["header_bg"])
        header.line.color.rgb = RGBColor(255, 255, 255)
        header.line.width = Pt(1.5)
        
        # Rotated Header Textbox
        tb_w, tb_h = lane_height, header_width
        tb_x = (header_width / 2) - (tb_w / 2)
        tb_y = y + (lane_height / 2) - (tb_h / 2)
        
        tb = slide.shapes.add_textbox(Inches(tb_x), Inches(tb_y), Inches(tb_w), Inches(tb_h))
        tb.text = lane["name"]
        tb.rotation = -90
        tb.text_frame.word_wrap = False
        
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # === Layer 2: Flowchart Nodes ===
    nodes_config = {
        "Start":   {"type": MSO_SHAPE.ROUNDED_RECTANGLE, "cx": 1.8,  "cy": 1.9, "w": 1.0, "h": 0.5, "text": "Start"},
        "Order":   {"type": MSO_SHAPE.RECTANGLE,         "cx": 3.8,  "cy": 1.9, "w": 1.4, "h": 0.6, "text": "Place Order"},
        "Review":  {"type": MSO_SHAPE.RECTANGLE,         "cx": 3.8,  "cy": 3.3, "w": 1.4, "h": 0.6, "text": "Review Order"},
        "Check":   {"type": MSO_SHAPE.DIAMOND,           "cx": 6.0,  "cy": 4.7, "w": 1.6, "h": 0.9, "text": "In Stock?"},
        "Procure": {"type": MSO_SHAPE.RECTANGLE,         "cx": 8.5,  "cy": 4.7, "w": 1.4, "h": 0.6, "text": "Restock"},
        "Invoice": {"type": MSO_SHAPE.RECTANGLE,         "cx": 6.0,  "cy": 6.1, "w": 1.4, "h": 0.6, "text": "Process Payment"},
        "Ship":    {"type": MSO_SHAPE.RECTANGLE,         "cx": 11.0, "cy": 4.7, "w": 1.4, "h": 0.6, "text": "Ship Goods"},
        "Receive": {"type": MSO_SHAPE.ROUNDED_RECTANGLE, "cx": 11.0, "cy": 1.9, "w": 1.2, "h": 0.5, "text": "Receive Order"}
    }

    def apply_shadow(shape):
        spPr = shape.element.spPr
        effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw', 
                                     blurRad="40000", dist="38100", dir="2700000", algn="tl", rotWithShape="0")
        srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val="000000")
        etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', val="15000")

    shapes_dict = {}
    for name, data in nodes_config.items():
        left = Inches(data["cx"] - data["w"]/2)
        top = Inches(data["cy"] - data["h"]/2)
        
        shape = slide.shapes.add_shape(data["type"], left, top, Inches(data["w"]), Inches(data["h"]))
        shape.text = data["text"]
        
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(89, 89, 89)
        shape.line.width = Pt(1.5)
        
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.color.rgb = RGBColor(50, 50, 50)
                run.font.size = Pt(11)
                run.font.bold = True
                
        apply_shadow(shape)
        shapes_dict[name] = shape

    # === Layer 3: Connectors & Overlaid Labels ===
    # Edge format: (from_node, to_node, from_site, to_site, text_label)
    # Sites: 0=Top, 1=Left, 2=Bottom, 3=Right
    edges = [
        ("Start",   "Order",   3, 1, ""),      
        ("Order",   "Review",  2, 0, ""),      
        ("Review",  "Check",   2, 0, ""),      
        ("Check",   "Procure", 3, 1, "No"),    
        ("Check",   "Invoice", 2, 0, "Yes"),   
        ("Procure", "Ship",    3, 1, ""),      
        ("Invoice", "Ship",    3, 2, ""),      
        ("Ship",    "Receive", 0, 2, ""),      
    ]

    for start_node, end_node, start_site, end_site, label in edges:
        s1 = shapes_dict[start_node]
        s2 = shapes_dict[end_node]
        
        connector = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(0), Inches(0), Inches(1), Inches(1))
        connector.begin_connect(s1, start_site)
        connector.end_connect(s2, end_site)
        
        connector.line.color.rgb = RGBColor(89, 89, 89)
        connector.line.width = Pt(1.5)
        connector.line.end_arrowhead = MSO_ARROWHEAD_STYLE.TRIANGLE
        
        # Generate logical positions for overlay labels
        if label:
            start_data = nodes_config[start_node]
            lx, ly = 0, 0
            
            if start_site == 3: # Exiting Right
                lx = start_data["cx"] + start_data["w"]/2 + 0.1
                ly = start_data["cy"] - 0.25
            elif start_site == 2: # Exiting Bottom
                lx = start_data["cx"] + 0.1
                ly = start_data["cy"] + start_data["h"]/2 + 0.05
                
            tb = slide.shapes.add_textbox(Inches(lx), Inches(ly), Inches(0.4), Inches(0.25))
            
            # Solid white fill to mask the line beneath the text
            tb.fill.solid()
            tb.fill.fore_color.rgb = RGBColor(255, 255, 255)
            
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            
            p = tf.paragraphs[0]
            p.text = label
            p.alignment = PP_ALIGN.CENTER
            
            run = p.runs[0]
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(192, 57, 43) # Highlight red for decisions
            run.font.bold = True

    prs.save(output_pptx_path)
    return output_pptx_path
