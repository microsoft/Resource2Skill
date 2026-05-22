import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw

def create_radial_agenda_graphic(img_path: str, arc_data: list):
    """
    Generates the central clock and colored arcs as a high-res PNG.
    """
    # Create a transparent high-res canvas
    size = 1000
    center = size // 2
    img = Image.new('RGBA', (size, size), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)

    # Base clock geometry
    clock_radius = 250
    clock_bbox = [center - clock_radius, center - clock_radius, 
                  center + clock_radius, center + clock_radius]
    
    # Draw clock outer ring
    draw.ellipse(clock_bbox, outline=(30, 30, 30, 255), width=35)
    
    # Draw clock center dot
    dot_radius = 25
    draw.ellipse([center - dot_radius, center - dot_radius, 
                  center + dot_radius, center + dot_radius], fill=(30, 30, 30, 255))
    
    # Draw clock hands (pointing approx to 10 and 4)
    # Minute hand
    draw.line([(center, center), (center - 120, center - 140)], fill=(30, 30, 30, 255), width=25)
    # Hour hand
    draw.line([(center, center), (center + 80, center + 90)], fill=(30, 30, 30, 255), width=35)

    # Draw the colored segmented arcs around the clock
    arc_offset = 60
    arc_width = 45
    arc_bbox = [
        clock_bbox[0] - arc_offset, clock_bbox[1] - arc_offset,
        clock_bbox[2] + arc_offset, clock_bbox[3] + arc_offset
    ]

    for item in arc_data:
        # Pillow arc angles: 0 is at 3 o'clock, increasing clockwise.
        start_angle = item['start_angle']
        end_angle = item['end_angle']
        color = item['color'] # RGB tuple
        rgba_color = (color[0], color[1], color[2], 255)
        
        # Draw arc
        draw.arc(arc_bbox, start=start_angle, end=end_angle, fill=rgba_color, width=arc_width)

    # Save to file
    img.save(img_path, format="PNG")
    return img_path

def create_slide(
    output_pptx_path: str,
    title_text: str = "Agenda Slide",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Radial Chrono-Segment Agenda style.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Define Agenda Data (Maps to quadrants and angles)
    # Pillow Angles: 0=Right, 90=Bottom, 180=Left, 270=Top
    agenda_data = [
        {
            "title": "Sub Heading Here",
            "desc": "Introduction and best practices to learn Project Management concepts.",
            "color": (255, 192, 0),    # Yellow
            "start_angle": 280,        # Top-Right quadrant
            "end_angle": 350,
            "position": "top_right"
        },
        {
            "title": "Sub Heading Here",
            "desc": "Deep dive into agile methodologies and slide templates.",
            "color": (0, 176, 240),    # Blue
            "start_angle": 10,         # Bottom-Right quadrant
            "end_angle": 80,
            "position": "bottom_right"
        },
        {
            "title": "Sub Heading Here",
            "desc": "Workshop exercises, resource sharing, and Q&A sessions.",
            "color": (0, 176, 80),     # Green
            "start_angle": 100,        # Bottom-Left quadrant
            "end_angle": 170,
            "position": "bottom_left"
        },
        {
            "title": "Sub Heading Here",
            "desc": "Conclusion, next steps, and project sign-off procedures.",
            "color": (112, 48, 160),   # Purple
            "start_angle": 190,        # Top-Left quadrant
            "end_angle": 260,
            "position": "top_left"
        }
    ]

    # 1. Generate and Insert Central Graphic
    temp_img_path = "temp_agenda_clock.png"
    create_radial_agenda_graphic(temp_img_path, agenda_data)
    
    img_size = Inches(4.5)
    img_left = (prs.slide_width - img_size) / 2
    img_top = Inches(1.8)
    slide.shapes.add_picture(temp_img_path, img_left, img_top, width=img_size, height=img_size)

    # 2. Add Main Title
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.5), prs.slide_width, Inches(1))
    tf = title_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(44)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.name = "Arial"
    tf.paragraphs[0].font.color.rgb = RGBColor(30, 30, 30)

    # 3. Add Text Blocks and Color Bullets
    # Layout configuration
    text_width = Inches(3.0)
    text_height = Inches(1.2)
    bullet_size = Inches(0.25)
    
    # Coordinates mapped to positions
    coords = {
        "top_right":    {"left": Inches(9.2), "top": Inches(2.2), "align": PP_ALIGN.LEFT, "bullet_offset": Inches(-0.4)},
        "bottom_right": {"left": Inches(9.2), "top": Inches(4.8), "align": PP_ALIGN.LEFT, "bullet_offset": Inches(-0.4)},
        "top_left":     {"left": Inches(1.1), "top": Inches(2.2), "align": PP_ALIGN.RIGHT, "bullet_offset": Inches(3.1)},
        "bottom_left":  {"left": Inches(1.1), "top": Inches(4.8), "align": PP_ALIGN.RIGHT, "bullet_offset": Inches(3.1)},
    }

    for item in agenda_data:
        pos = coords[item["position"]]
        
        # Add Bullet Shape
        bullet = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            pos["left"] + pos["bullet_offset"], 
            pos["top"] + Inches(0.05), 
            bullet_size, 
            bullet_size
        )
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = RGBColor(*item["color"])
        bullet.line.color.rgb = RGBColor(255, 255, 255) # White border
        bullet.line.width = Pt(2)

        # Add Text Box
        txBox = slide.shapes.add_textbox(pos["left"], pos["top"], text_width, text_height)
        tf = txBox.text_frame
        tf.word_wrap = True
        
        # Add Title
        p1 = tf.paragraphs[0]
        p1.text = item["title"]
        p1.alignment = pos["align"]
        p1.font.bold = True
        p1.font.size = Pt(18)
        p1.font.color.rgb = MSO_THEME_COLOR_INDEX = RGBColor(30, 30, 30)
        
        # Add Description
        p2 = tf.add_paragraph()
        p2.text = item["desc"]
        p2.alignment = pos["align"]
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(100, 100, 100)

    prs.save(output_pptx_path)
    
    # Cleanup temp image
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
        
    return output_pptx_path

# Example execution:
# create_slide("agenda_slide.pptx", title_text="Meeting Agenda")
