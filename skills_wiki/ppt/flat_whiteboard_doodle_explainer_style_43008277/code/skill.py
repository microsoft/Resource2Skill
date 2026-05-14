def create_slide(
    output_pptx_path: str,
    title_text: str = "用PPT也能做出\n好看的视频！",
    pain_points: list = ["不会拍摄", "不想出镜"],
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the "Flat Whiteboard Doodle Explainer" visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # Colors
    COLOR_BG = RGBColor(248, 249, 250)
    COLOR_BLACK = RGBColor(30, 30, 30)
    COLOR_GREEN = RGBColor(85, 175, 75)
    COLOR_RED = RGBColor(235, 50, 50)
    COLOR_WHITE = RGBColor(255, 255, 255)

    # --- Layer 1: Background ---
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()

    # --- Cinematic Film Strip Borders ---
    def draw_film_strip(top_y):
        # Black bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top_y, prs.slide_width, Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_BLACK
        bar.line.fill.background()
        
        # White sprocket holes
        num_holes = 30
        spacing = prs.slide_width / num_holes
        for i in range(num_holes):
            hole = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 
                i * spacing + Inches(0.15), 
                top_y + Inches(0.1), 
                Inches(0.15), 
                Inches(0.2)
            )
            hole.fill.solid()
            hole.fill.fore_color.rgb = COLOR_WHITE
            hole.line.fill.background()

    draw_film_strip(0) # Top border
    draw_film_strip(prs.slide_height - Inches(0.4)) # Bottom border

    # --- Layer 2: Pain Points (Top Left) ---
    start_x = Inches(1.5)
    start_y = Inches(1.2)
    
    for i, point in enumerate(pain_points):
        # Bubble
        bubble_width = Inches(2.2)
        bubble = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            start_x + (i * (bubble_width + Inches(0.5))), 
            start_y, 
            bubble_width, 
            Inches(0.8)
        )
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = COLOR_WHITE
        bubble.line.color.rgb = COLOR_GREEN
        bubble.line.width = Pt(3)
        
        # Text
        tf = bubble.text_frame
        tf.text = point
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_BLACK
        
        # Red X (Marker strikethrough effect)
        x_center_x = start_x + (i * (bubble_width + Inches(0.5))) + bubble_width - Inches(0.2)
        x_center_y = start_y + Inches(0.1)
        
        line1 = slide.shapes.add_connector(
            1, x_center_x - Inches(0.2), x_center_y - Inches(0.2), 
            x_center_x + Inches(0.2), x_center_y + Inches(0.2)
        )
        line1.line.color.rgb = COLOR_RED
        line1.line.width = Pt(6)
        
        line2 = slide.shapes.add_connector(
            1, x_center_x + Inches(0.2), x_center_y - Inches(0.2), 
            x_center_x - Inches(0.2), x_center_y + Inches(0.2)
        )
        line2.line.color.rgb = COLOR_RED
        line2.line.width = Pt(6)

    # --- Layer 3: Main Hook Text (Center Left) ---
    txt_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(8), Inches(3))
    tf = txt_box.text_frame
    tf.word_wrap = True

    # First line: "用PPT也能做出"
    p1 = tf.paragraphs[0]
    run1 = p1.add_run()
    run1.text = "用"
    run1.font.size = Pt(65)
    run1.font.bold = True
    run1.font.color.rgb = COLOR_BLACK

    run2 = p1.add_run()
    run2.text = "PPT"
    run2.font.size = Pt(85)
    run2.font.bold = True
    run2.font.color.rgb = COLOR_GREEN

    run3 = p1.add_run()
    run3.text = "也能做出"
    run3.font.size = Pt(65)
    run3.font.bold = True
    run3.font.color.rgb = COLOR_BLACK

    # Second line: "好看的视频！"
    p2 = tf.add_paragraph()
    p2.space_before = Pt(20)
    run4 = p2.add_run()
    run4.text = "好看的视频！"
    run4.font.size = Pt(90)
    run4.font.bold = True
    run4.font.color.rgb = COLOR_BLACK

    # --- Layer 4: Graphic Icon - Film Reel (Right Side) ---
    icon_center_x = Inches(10)
    icon_center_y = Inches(3.8)
    
    # Base Thick Outline Circle (Black)
    outer_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        icon_center_x - Inches(2), 
        icon_center_y - Inches(2), 
        Inches(4), Inches(4)
    )
    outer_circle.fill.solid()
    outer_circle.fill.fore_color.rgb = COLOR_WHITE
    outer_circle.line.color.rgb = COLOR_BLACK
    outer_circle.line.width = Pt(12)

    # Inner Filled Circle (Green)
    inner_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        icon_center_x - Inches(1.5), 
        icon_center_y - Inches(1.5), 
        Inches(3), Inches(3)
    )
    inner_circle.fill.solid()
    inner_circle.fill.fore_color.rgb = COLOR_GREEN
    inner_circle.line.fill.background()

    # Small Reel Holes (White)
    hole_positions = [
        (icon_center_x - Inches(0.6), icon_center_y - Inches(0.6)), # Top left
        (icon_center_x + Inches(0.6), icon_center_y - Inches(0.6)), # Top right
        (icon_center_x, icon_center_y + Inches(0.6))                # Bottom center
    ]
    
    for hx, hy in hole_positions:
        hole = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, 
            hx - Inches(0.4), hy - Inches(0.4), 
            Inches(0.8), Inches(0.8)
        )
        hole.fill.solid()
        hole.fill.fore_color.rgb = COLOR_BG # Match background to simulate cutout
        hole.line.color.rgb = COLOR_BLACK
        hole.line.width = Pt(4)

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
