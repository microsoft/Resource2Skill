def create_slide(
    output_pptx_path: str,
    chart_data: dict = None,
    title_text: str = "Technology enabled up to 71 percent of the value derived in business transformations across different sectors.",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a McKinsey-style Radial Bar Chart.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        chart_data (dict): A dictionary of data to plot, e.g.,
                           {'Financial Services': 71, 'Telecommunications': 70, ...}.
        title_text (str): The main title for the slide.

    Returns:
        str: The path to the saved PPTX file.
    """
    import matplotlib.pyplot as plt
    import numpy as np
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    import io

    # --- Use default data if none is provided ---
    if chart_data is None:
        chart_data = {
            "Financial Services": 71,
            "Telecommunications, media, and technology": 70,
            "Consumer": 39,
 богаты "Life sciences": 28,
            "Travel, logistics, and infrastructure": 22,
            "Global energy and materials": 15,
            "Advanced industries": 9,
        }

    # --- Chart Generation with Matplotlib ---
    categories = list(chart_data.keys())
    values = list(chart_data.values())
    num_categories = len(categories)

    # Colors and styles
    dark_blue = "#004685"
    light_blue = "#00AFEF"
    font_family = "Calibri"

    # Create a polar plot
    fig, ax = plt.subplots(figsize=(10, 10), subplot_kw=dict(polar=True))
    fig.patch.set_alpha(0)  # Transparent background for the figure

    # Define the angle for the bars (e.g., 270 degrees)
    # We convert percentages to radians for plotting
    max_angle = 270
    theta = np.deg2rad(np.linspace(0, max_angle, 100))

    # Define radii for the concentric circles
    radii = np.linspace(2, 2 + num_categories * 0.7, num_categories)
    ring_thickness = 0.6

    # Plot the bars/arcs
    for i, (category, value) in enumerate(chart_data.items()):
        radius = radii[i]
        angle_value = np.deg2rad(value / 100 * max_angle)
        
        # Plot the main data arc
        ax.bar(x=0, height=ring_thickness, width=angle_value, bottom=radius,
               color=dark_blue if i < 2 else light_blue,  # Style first two differently
               align='edge', edgecolor='white', linewidth=0)

        # Add percentage label at the end of the arc
        label_angle = angle_value
        label_radius = radius + ring_thickness / 2
        ax.text(label_angle, label_radius, f'{value}',
                ha='center', va='center',
                fontsize=14, fontweight='bold', color='black', family=font_family,
                rotation=-(value / 100 * max_angle))
                
        # Add category label to the left
        ax.text(np.deg2rad(-5), label_radius, category,
                ha='right', va='center',
                fontsize=14, color='black', family=font_family)


    # --- Formatting the plot ---
    ax.set_yticklabels([])  # Hide radial ticks
    ax.set_xticklabels([])  # Hide angular ticks
    ax.spines['polar'].set_visible(False) # Hide the outer circle
    ax.grid(False) # Hide grid lines
    ax.set_theta_zero_location('N') # Set 0 degrees to the top
    ax.set_theta_direction(-1) # Clockwise
    
    # Set the angular view limits to only show the 270-degree arc
    ax.set_thetamin(0)
    ax.set_thetamax(max_angle)
    
    # Adjust radial limits to fit the chart
    ax.set_rmax(radii[-1] + ring_thickness + 0.5)

    # --- Save plot to a memory buffer ---
    img_stream = io.BytesIO()
    plt.savefig(img_stream, format='png', bbox_inches='tight', pad_inches=0.1, dpi=300, transparent=True)
    img_stream.seek(0)
    plt.close(fig)

    # --- Create PowerPoint Slide ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Add title
    title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    title_tf = title_shape.text_frame
    title_tf.text = title_text
    p = title_tf.paragraphs[0]
    p.font.name = font_family
    p.font.size = Pt(24)
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

    # Insert the chart image
    chart_pic = slide.shapes.add_picture(img_stream, Inches(2), Inches(1.5), height=Inches(5.5))

    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function:
# create_slide("McKinsey_Radial_Chart.pptx")
