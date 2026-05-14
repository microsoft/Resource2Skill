def create_slide(
    output_pptx_path: str,
    title_text: str = "TITLE LOREM\nIPSUM",
    body_text: str = "Sit Dolor Amet",
    bg_palette: str = "colorful",  
    accent_color: tuple = (40, 40, 40),  
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Framed Minimal Overlay visual effect.
    """
    import os
    import random
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # === Layer 1: Background Generation via PIL ===
    # Generating a complex, colorful geometric background to mimic the tutorial's 3D vibe
    bg_img_path = "temp_complex_bg.png"
    img_width, img_height = int(13.333 * 150), int(7.5 * 150) # 150 DPI
    bg_img = Image.new('RGB', (img_width, img_height), (245, 245, 245))
    draw = ImageDraw.Draw(bg_img, 'RGBA')
    
    # Palette mimicking the colorful blocks in the video
    colors = [
        (220, 50, 50, 200),   # Red
        (50, 180, 80, 200),   # Green
        (50, 100, 220, 200),  # Blue
        (240, 180, 30, 200),  # Yellow
        (100, 200, 220, 200)  # Cyan
    ]
    
    # Draw overlapping abstract geometric shards
    for _ in range(60):
        x1 = random.randint(-300, img_width)
        y1 = random.randint(-300, img_height)
        size_w = random.randint(200, 800)
        size_h = random.randint(200, 800)
        offset = random.randint(-200, 200)
        color = random.choice(colors)
        poly = [
            (x1, y1), 
            (x1 + size_w, y1 + offset), 
            (x1 + size_w - offset, y1 + size_h), 
            (x1 - offset, y1 + size_h - offset)
        ]
        draw.polygon(poly, fill=color)
        
    bg_img.save(bg_img_path)
    
    # Insert background
    slide.shapes.add_picture(bg_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: The Framed Minimal Overlay ===
    
    # Dimensions for the dark plate
    box_width = Inches(5.8)
    box_height = Inches(3.2)
    box_left = Inches(6.8)  # Positioned on the right
    box_top = Inches(3.5)   # Positioned slightly below center
    
    # 1. Main Dark Plate
    plate = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box_left, box_top, box_width, box_height)
    plate.fill.solid()
    plate.fill.fore_color.rgb = RGBColor(*accent_color)
    plate.line.fill.background() # No outline on the main plate
    
    # 2. Inner White Frame
    margin = Inches(0.25)
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        box_left + margin, 
        box_top + margin, 
        box_width - (margin * 2), 
        box_height - (margin * 2)
    )
    frame.fill.background() # Transparent fill
    frame.line.color.rgb = RGBColor(255, 255, 255)
    frame.line.width = Pt(1.5)

    # === Layer 3: Text & Separator ===
    
    # Main Title
    txBox = slide.shapes.add_textbox(
        box_left + margin, 
        box_top + margin + Inches(0.2), 
        box_width - (margin * 2), 
        Inches(1.5)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Segoe UI Light"

    # Horizontal Separator Line
    line_y = box_top + box_height - Inches(1.0)
    line_width = Inches(1.5)
    line_left = box_left + (box_width / 2) - (line_width / 2)
    
    sep_line = slide.shapes.add_shape(MSO_SHAPE.LINE, line_left, line_y, line_width, 0)
    sep_line.line.color.rgb = RGBColor(255, 255, 255)
    sep_line.line.width = Pt(1.0)

    # Subtitle
    subBox = slide.shapes.add_textbox(
        box_left + margin, 
        line_y + Inches(0.1), 
        box_width - (margin * 2), 
        Inches(0.6)
    )
    sub_tf = subBox.text_frame
    sub_tf.word_wrap = True
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = body_text
    sub_p.alignment = PP_ALIGN.CENTER
    sub_p.font.size = Pt(16)
    sub_p.font.color.rgb = RGBColor(255, 255, 255)
    sub_p.font.name = "Segoe UI"
    # Make subtitle slightly transparent looking by using light grey
    sub_p.font.color.rgb = RGBColor(200, 200, 200) 

    # Clean up temporary PIL image
    if os.path.exists(bg_img_path):
        os.remove(bg_img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
