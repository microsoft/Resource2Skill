def create_slide(
    output_pptx_path: str,
    title_text: str = "Chart created in PowerPoint",
    chart_data: dict = None,
    chart_title_text: str = "Units",
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with a native column chart populated from a
    Python data structure, mimicking the result of pasting Excel data.

    Args:
        output_pptx_path (str): The path to save the generated .pptx file.
        title_text (str): The main title for the PowerPoint slide.
        chart_data (dict): A dictionary with 'categories' and 'values' keys.
                           If None, default data is used.
        chart_title_text (str): The title to display on the chart itself.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only layout

    # --- Slide Title ---
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.size = Pt(36)
    title_shape.top = Inches(0.2)
    title_shape.left = Inches(0.5)

    # --- Chart Data (using data from the video tutorial) ---
    if chart_data is None:
        chart_data = {
            "categories": ["Region A", "Region B", "Region C", "Region D", "Region E"],
            "values": [24, 65, 36, 48, 51]
        }

    # --- Create and Populate Chart Data Object ---
    data_for_chart = CategoryChartData()
    data_for_chart.categories = chart_data['categories']
    # The series name here is what would be in the legend if it were visible.
    data_for_chart.add_series(chart_title_text, chart_data['values'])

    # --- Add Chart to Slide ---
    x, y, cx, cy = Inches(1.5), Inches(1.5), Inches(10), Inches(5.5)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, data_for_chart
    )
    chart = graphic_frame.chart

    # --- Style the Chart ---
    chart.has_legend = False # As seen in the video, single-series charts don't need a legend

    # Set chart title
    chart.chart_title.text_frame.text = chart_title_text
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(20)

    # Style category axis (X-axis)
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(12)

    # Style value axis (Y-axis)
    value_axis = chart.value_axis
    value_axis.has_title = False # No Y-axis title in the video
    value_axis.tick_labels.font.size = Pt(12)
    value_axis.maximum_scale = 70.0 # Set max value to give some headroom, as in video
    
    # Style the data series plot
    plot = chart.plots[0]
    plot.vary_colors_by_category = False # Ensure all bars are the same color
    
    # Set the color of the bars (orange from the video)
    series = plot.series[0]
    fill = series.format.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(244, 122, 32) # An orange similar to the video's theme

    # --- Save the Presentation ---
    prs.save(output_pptx_path)
    return output_pptx_path

# Example of how to run the function
# create_slide("powerpoint_chart_from_data.pptx")
