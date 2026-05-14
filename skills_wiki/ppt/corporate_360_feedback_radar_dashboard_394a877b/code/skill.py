def create_slide(
    output_pptx_path: str,
    title_text: str = "360 Degree Feedback Dashboard",
    employee_name: str = "Jane Doe",
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the 360 Degree Feedback Radar Dashboard.
    Uses matplotlib to render the radar chart and python-pptx/lxml for the dashboard layout.
    """
    import io
    import math
    import numpy as np
    import matplotlib.pyplot as plt
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml import parse_xml

    # --- Setup Presentation ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # --- Color Palette ---
    COLOR_TEAL_DARK = RGBColor(0, 105, 115)
    COLOR_TEAL_LIGHT = RGBColor(95, 180, 175)
    COLOR_GREEN = RGBColor(150, 200, 80)
    COLOR_GRAY = RGBColor(100, 100, 100)

    # --- Data Definition ---
    categories = ['Leadership', 'Adaptability', 'Relationships', 
                  'Analytical Thinking', 'Integrity', 'Teamwork', 
                  'Decision Making', 'Communication']
    N = len(categories)
    
    # Mock scores out of 5
    self_scores = [4.0, 3.5, 4.0, 3.0, 5.0, 4.5, 3.5, 4.0]
    other_scores = [4.5, 4.0, 3.5, 4.0, 4.8, 4.0, 4.5, 3.8]

    # Calculate averages
    avg_self = sum(self_scores) / len(self_scores)
    avg_other = sum(other_scores) / len(other_scores)

    # --- MATPLOTLIB: Generate Radar Chart ---
    # Repeat first value to close the circular polygon
    self_scores_plot = self_scores + [self_scores[0]]
    other_scores_plot = other_scores + [other_scores[0]]
    angles = [n / float(N) * 2 * math.pi for n in range(N)]
    angles += angles[:1]

    # Matplotlib styling
    plt.style.use('default')
    fig, ax = plt.subplots(figsize=(6, 6), subplot_kw=dict(polar=True))
    fig.patch.set_alpha(0.0) # Transparent background
    ax.patch.set_alpha(0.0)

    # Offset to start at top
    ax.set_theta_offset(math.pi / 2)
    ax.set_theta_direction(-1)

    # Draw grid/labels
    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(categories, color='#006973', size=12, weight='bold')
    ax.set_rlabel_position(0)
    plt.yticks([1, 2, 3, 4, 5], ["1", "2", "3", "4", "5"], color="grey", size=10)
    plt.ylim(0, 5)

    # Plot Series 1 (Self) - Teal Line
    ax.plot(angles, self_scores_plot, linewidth=2, linestyle='solid', color='#5fb4af', label='Self Assessment')
    # Plot Series 2 (Others) - Green Fill
    ax.plot(angles, other_scores_plot, linewidth=2, linestyle='solid', color='#96c850', label='Others Assessment')
    ax.fill(angles, other_scores_plot, color='#96c850', alpha=0.25)

    # Hide outer spine
    ax.spines['polar'].set_visible(False)
    
    # Save chart to memory buffer
    chart_img_buf = io.BytesIO()
    plt.tight_layout()
    plt.savefig(chart_img_buf, format='png', dpi=300, transparent=True)
    chart_img_buf.seek(0)
    plt.close(fig)

    # --- PPTX: Insert Elements ---

    # 1. Slide Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(28)
    p.font.color.rgb = COLOR_TEAL_DARK
    p.font.bold = True

    # 2. Insert Radar Chart
    slide.shapes.add_picture(chart_img_buf, Inches(0.5), Inches(1.2), width=Inches(6.0), height=Inches(6.0))

    # --- Helper Function for KPI Cards ---
    def add_kpi_card(x, y, w, h, title, score, color, description, icon_char):
        # Base Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(250, 250, 250)
        shape.line.color.rgb = RGBColor(230, 230, 230)
        
        # lxml: Add drop shadow to the shape
        spPr = shape.element.spPr
        shadow_xml = """
            <a:outerShdw xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" 
                         blurRad="100000" dist="30000" dir="5400000" algn="b" rotWithShape="0">
                <a:srgbClr val="000000">
                    <a:alpha val="15000"/>
                </a:srgbClr>
            </a:outerShdw>
        """
        effectLst = OxmlElement('a:effectLst')
        effectLst.append(parse_xml(shadow_xml))
        spPr.append(effectLst)

        # Title TextBox
        tb_title = slide.shapes.add_textbox(Inches(x+0.2), Inches(y+0.1), Inches(w-0.4), Inches(0.5))
        p_title = tb_title.text_frame.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_TEAL_DARK
        p_title.font.bold = True

        # Score TextBox
        tb_score = slide.shapes.add_textbox(Inches(x+0.2), Inches(y+0.6), Inches(2), Inches(0.8))
        p_score = tb_score.text_frame.add_paragraph()
        p_score.text = f"{score:.2f}"
        p_score.font.size = Pt(36)
        p_score.font.color.rgb = color
        p_score.font.bold = True

        # Description TextBox
        tb_desc = slide.shapes.add_textbox(Inches(x+0.2), Inches(y+1.5), Inches(w-0.4), Inches(1.5))
        tb_desc.text_frame.word_wrap = True
        p_desc = tb_desc.text_frame.add_paragraph()
        p_desc.text = description
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_GRAY

    # 3. Create Self Assessment Card
    desc_self = ("My goal is to enhance my collaboration with cross-functional teams "
                 "to contribute more effectively to company-wide projects.")
    add_kpi_card(
        x=7.5, y=1.2, w=5.0, h=2.5,
        title="Self Assessment", 
        score=avg_self, 
        color=COLOR_TEAL_LIGHT, 
        description=desc_self,
        icon_char="👤"
    )

    # 4. Create Others Assessment Card
    desc_others = ("Consider time management training to help balance tasks "
                   "and responsibilities more effectively. Great leadership potential demonstrated.")
    add_kpi_card(
        x=7.5, y=4.2, w=5.0, h=2.5,
        title="Others Assessment", 
        score=avg_other, 
        color=COLOR_GREEN, 
        description=desc_others,
        icon_char="👥"
    )

    # Save and return
    prs.save(output_pptx_path)
    return output_pptx_path
