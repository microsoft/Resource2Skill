import os
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw
from lxml import etree

def create_dot_grid(output_path, width_px=1920, height_px=1080, spacing=40, dot_radius=2):
    """Generates a subtle dot grid image to reinforce the 'Grid Systems' law visually."""
    img = Image.new('RGBA', (width_px, height_px), (240, 244, 248, 255)) # Light Ice Blue bg
    draw = ImageDraw.Draw(img)
    color = (220, 225, 230, 255) # Subtle gray dots
    for x in range(0, width_px, spacing):
        for y in range(0, height_px, spacing):
            draw.ellipse((x-dot_radius, y-dot_radius, x+dot_radius, y+dot_radius), fill=color)
    img.save(output_path)

def fetch_and_crop_image(url, output_path, target_width_in, target_height_in, dpi=300):
    """Fetches an image and center-crops it to strictly match the layout grid ratio."""
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(output_path, 'wb') as f:
                f.write(response.read())
        
        img = Image.open(output_path).convert('RGB')
        img_w, img_h = img.size
        
        target_aspect = target_width_in / target_height_in
        img_aspect = img_w / img_h
        
        # Center crop
        if img_aspect > target_aspect:
            new_w = int(img_h * target_aspect)
            left = (img_w - new_w) // 2
            img = img.crop((left, 0, left + new_w, img_h))
        else:
            new_h = int(img_w / target_aspect)
            top = (img_h - new_h) // 2
            img = img.crop((0, top, img_w, top + new_h))
            
        img = img.resize((int(target_width_in * dpi), int(target_height_in * dpi)), Image.Resampling.LANCZOS)
        img.save(output_path)
    except Exception as e:
        print(f"Image fetch failed: {e}. Using fallback geometry.")
        img = Image.new('RGB', (int(target_width_in * dpi), int(target_height_in * dpi)), color=(200, 210, 220))
        draw = ImageDraw.Draw(img)
        draw.line((0, 0, img.size[0], img.size[1]), fill=(180, 190, 200), width=5)
        draw.line((0, img.size[1], img.size[0], 0), fill=(180, 190, 200), width=5)
        img.save(output_path)

def make_pill_shape(shape):
    """XML Injection: Turns a rounded rectangle into a perfect pill shape (50% radius)."""
    geom = shape._element.spPr.prstGeom
    avLst = geom.get_or_add_avLst()
    gd = etree.SubElement(avLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd')
    gd.set('name', 'adj')
    gd.set('fmla', 'val 50000')

def add_shadow(shape):
    """XML Injection: Adds a soft outer drop-shadow for depth and layering."""
    spPr = shape._element.spPr
    effectLst = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
    outerShdw = etree.SubElement(effectLst, '{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw')
    outerShdw.set('blurRad', '190500') # 15 pt
    outerShdw.set('dist', '63500') # 5 pt
    outerShdw.set('dir', '5400000') # 90 degrees (down)
    outerShdw.set('algn', 'tl')
    srgbClr = etree.SubElement(outerShdw, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
    srgbClr.set('val', '000000')
    alpha = etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
    alpha.set('val', '30000') # 30% opacity

def create_slide(
    output_pptx_path: str,
    title_text: str = "It always seems impossible until it's done.",
    body_text: str = "Career coaching can help you with your current job, helping you to establish professional goals and feel more fulfilled, or support you in making a career change.",
    bg_palette: str = "portrait", 
    accent_color: tuple = (255, 193, 7),
    **kwargs,
) -> str:
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Palette
    cyan = RGBColor(0, 191, 255)
    navy = RGBColor(13, 71, 161)
    slate = RGBColor(69, 90, 100)
    yellow = RGBColor(accent_color[0], accent_color[1], accent_color[2])

    # === LAYER 1: Background & Grid System ===
    grid_path = "temp_grid.png"
    create_dot_grid(grid_path)
    slide.shapes.add_picture(grid_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === LAYER 2: Giant Decorative Typography (Layering) ===
    tx_a = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(5))
    p_a = tx_a.text_frame.paragraphs[0]
    p_a.text = "A"
    p_a.font.size = Pt(450)
    p_a.font.name = "Arial Black"
    p_a.font.bold = True
    p_a.font.color.rgb = yellow

    tx_w = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(4), Inches(5))
    p_w = tx_w.text_frame.paragraphs[0]
    p_w.text = "W"
    p_w.font.size = Pt(450)
    p_w.font.name = "Arial Black"
    p_w.font.bold = True
    p_w.font.color.rgb = cyan

    # === LAYER 3: Core Image ===
    # 12-column grid math: 1 col = 1.111 inches
    img_path = "temp_portrait.jpg"
    img_url = "https://images.unsplash.com/photo-1560250097-0b93528c311a?q=80&w=800&auto=format&fit=crop"
    fetch_and_crop_image(img_url, img_path, target_width_in=4.444, target_height_in=5.5)
    
    pic = slide.shapes.add_picture(img_path, Inches(1.111), Inches(1.0), width=Inches(4.444), height=Inches(5.5))
    add_shadow(pic)

    # === LAYER 4: Text Hierarchy & Proximity ===
    # Starts at Col 6 (X = 6.666)
    
    # Large Quote
    title_box = slide.shapes.add_textbox(Inches(6.666), Inches(1.0), Inches(5.5), Inches(2.5))
    title_box.text_frame.word_wrap = True
    title_box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    p = title_box.text_frame.paragraphs[0]
    p.text = f'"{title_text}"'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = "Georgia"
    p.font.color.rgb = navy

    # Subtitle / Author
    sub_box = slide.shapes.add_textbox(Inches(6.666), Inches(3.8), Inches(5.5), Inches(0.5))
    p = sub_box.text_frame.paragraphs[0]
    p.text = "Nelson Mandela"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = cyan

    # Body Text (Proximity grouping)
    body_box = slide.shapes.add_textbox(Inches(6.666), Inches(4.3), Inches(5.0), Inches(1.5))
    body_box.text_frame.word_wrap = True
    p = body_box.text_frame.paragraphs[0]
    p.text = body_text
    p.font.size = Pt(12)
    p.font.color.rgb = slate

    # === LAYER 5: Focal Point CTA Button ===
    cta = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.666), Inches(6.0), Inches(3.0), Inches(0.5))
    cta.fill.solid()
    cta.fill.fore_color.rgb = yellow
    cta.line.fill.background()
    make_pill_shape(cta)

    p = cta.text_frame.paragraphs[0]
    p.text = "MAKE AN ENQUIRY TODAY"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = navy

    prs.save(output_pptx_path)
    
    # Cleanup temp files
    if os.path.exists(grid_path): os.remove(grid_path)
    if os.path.exists(img_path): os.remove(img_path)
    
    return output_pptx_path
