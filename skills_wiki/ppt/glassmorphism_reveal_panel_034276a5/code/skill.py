import os
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml import qn
from PIL import Image, ImageDraw, ImageFilter, ImageFont

def create_slide(
    output_pptx_path: str,
    title_text: str = "Future of design",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Glassmorphism Reveal Panel effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # === Step 1: Generate Base Gradient Background via PIL ===
    img_w, img_h = 1920, 1080
    grad = Image.new('RGB', (512, 512), (15, 15, 25))
    draw = ImageDraw.Draw(grad)
    
    # Draw soft glowing orbs to simulate a mesh gradient
    draw.ellipse((-100, 200, 300, 600), fill=(40, 60, 200)) # Blue glow bottom left
    draw.ellipse((200, -50, 600, 300), fill=(100, 30, 150)) # Purple glow top right
    grad = grad.filter(ImageFilter.GaussianBlur(60))
    grad = grad.resize((img_w, img_h), Image.Resampling.BICUBIC)

    # === Step 2: Bake Title Text and Blur for the "Refracted" Layer ===
    bg_blurred = grad.copy()
    draw_blur = ImageDraw.Draw(bg_blurred)
    
    # Attempt to load a serif font, fallback to default
    font_to_use = None
    for f in ["georgia.ttf", "times.ttf", "arial.ttf", "DejaVuSans.ttf"]:
        try:
            font_to_use = ImageFont.truetype(f, 160)
            break
        except IOError:
            continue
            
    if font_to_use is None:
        font_to_use = ImageFont.load_default()

    # Center the text horizontally in PIL
    try:
        bbox = draw_blur.textbbox((0, 0), title_text, font=font_to_use)
        tw = bbox[2] - bbox[0]
    except AttributeError:
        tw = font_to_use.getsize(title_text)[0] # Legacy PIL fallback
        
    x_pos = (img_w - tw) / 2
    y_pos = 150
    draw_blur.text((x_pos, y_pos), title_text, font=font_to_use, fill=(255, 255, 255))
    
    # Apply heavy blur for the frosted glass effect
    bg_blurred = bg_blurred.filter(ImageFilter.GaussianBlur(40))

    # Save to memory streams
    clear_stream = BytesIO()
    grad.save(clear_stream, format='PNG')
    clear_stream.seek(0)
    
    blur_stream = BytesIO()
    bg_blurred.save(blur_stream, format='PNG')
    blur_stream.seek(0)

    # === Step 3: Inject Blurred Image as the Core Slide Background ===
    # Trick: Add picture to get rId, construct background XML, then delete picture
    pic_blur = slide.shapes.add_picture(blur_stream, 0, 0, prs.slide_width, prs.slide_height)
    blip = pic_blur.element.xpath('.//a:blip')[0]
    rId = blip.get(qn('r:embed'))
    
    bg = OxmlElement('p:bg')
    bgPr = OxmlElement('p:bgPr')
    blipFill = OxmlElement('a:blipFill')
    blip_new = OxmlElement('a:blip')
    blip_new.set(qn('r:embed'), rId)
    stretch = OxmlElement('a:stretch')
    fillRect = OxmlElement('a:fillRect')
    stretch.append(fillRect)
    blipFill.append(blip_new)
    blipFill.append(stretch)
    bgPr.append(blipFill)
    bg.append(bgPr)
    
    slide.element.insert(0, bg)
    slide.element.shapes.remove(pic_blur.element)

    # === Step 4: Add Clear Background as Base Layer ===
    slide.shapes.add_picture(clear_stream, 0, 0, prs.slide_width, prs.slide_height)

    # === Step 5: Add Clear Ambient Title Text ===
    tb_top = Inches(y_pos / 1080 * 7.5) # Match vertical position from PIL
    tb_title = slide.shapes.add_textbox(0, tb_top, prs.slide_width, Inches(2))
    p = tb_title.text_frame.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.name = "Georgia"
    run.font.size = Pt(110)
    run.font.italic = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    # === Step 6: Create the Glassmorphism Panel ===
    glass_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(2), Inches(2.2), Inches(9.333), Inches(4.5)
    )
    
    spPr = glass_shape.element.spPr
    # Clear solid fill to apply the magical bgFill (Slide Background Fill)
    for el in spPr.xpath('.//a:solidFill | .//a:gradFill | .//a:noFill | .//a:blipFill'):
        el.getparent().remove(el)
        
    bgFill = OxmlElement('a:bgFill')
    prstGeom = spPr.find(qn('a:prstGeom'))
    if prstGeom is not None:
        prstGeom.addnext(bgFill)
    else:
        spPr.insert(0, bgFill)
        
    # Add an ultra-thin gradient outline to simulate glass rim reflection
    ln = OxmlElement('a:ln')
    ln.set('w', '19050') # 1.5 pt
    gradFill = OxmlElement('a:gradFill')
    gradFill.set('rotWithShape', '1')
    gsLst = OxmlElement('a:gsLst')
    
    # White 80% opacity
    gs1 = OxmlElement('a:gs'); gs1.set('pos', '0')
    c1 = OxmlElement('a:srgbClr'); c1.set('val', 'FFFFFF')
    a1 = OxmlElement('a:alpha'); a1.set('val', '80000')
    c1.append(a1); gs1.append(c1); gsLst.append(gs1)
    
    # White 20% opacity
    gs2 = OxmlElement('a:gs'); gs2.set('pos', '100000')
    c2 = OxmlElement('a:srgbClr'); c2.set('val', 'FFFFFF')
    a2 = OxmlElement('a:alpha'); a2.set('val', '20000')
    c2.append(a2); gs2.append(c2); gsLst.append(gs2)
    
    lin = OxmlElement('a:lin'); lin.set('ang', '3150000'); lin.set('scaled', '1')
    gradFill.append(gsLst); gradFill.append(lin); ln.append(gradFill)
    
    for el in spPr.xpath('.//a:ln'): el.getparent().remove(el)
    spPr.append(ln)

    # === Step 7: Add Foreground Content (Metrics) ===
    stats = [
        ("44M", "Lorem ipsum", "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor."),
        ("72%", "Lorem ipsum", "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor."),
        ("84K", "Lorem ipsum", "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Maecenas porttitor.")
    ]
    
    left_offsets = [Inches(2.5), Inches(5.5), Inches(8.5)]
    
    for idx, (val, title, body) in enumerate(stats):
        tb = slide.shapes.add_textbox(left_offsets[idx], Inches(2.7), Inches(2.8), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = val
        p1.font.size = Pt(54)
        p1.font.name = "Georgia"
        p1.font.color.rgb = RGBColor(255, 255, 255)
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.name = "Arial"
        p2.font.color.rgb = RGBColor(220, 220, 230)
        p2.space_before = Pt(10)
        
        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(12)
        p3.font.name = "Arial"
        p3.font.color.rgb = RGBColor(200, 200, 210)
        p3.space_before = Pt(5)

    prs.save(output_pptx_path)
    return output_pptx_path
