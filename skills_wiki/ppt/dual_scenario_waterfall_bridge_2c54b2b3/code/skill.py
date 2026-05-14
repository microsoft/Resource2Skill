def create_slide(
    output_pptx_path: str,
    title_text: str = "Aggressive Scenario 2 adds $262M Sales to Budget 20X2 compared to Scenario 1",
    subtitle_text: str = "Scenarios for Sales Budget 20X2",
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a Dual-Scenario Waterfall Bridge chart.

    This function manually constructs the waterfall chart using individual shapes
    to precisely replicate the visual comparison of two scenarios.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Data Definition ---
    # Data for the two scenarios
    start_value = 1500
    scenario1_changes = [
        {"label": "New Products", "value": 35},
        {"label": "M&A", "value": 88},
        {"label": "Online", "value": 100},
    ]
    scenario2_changes = [
        {"label": "New Products", "value": 150},
        {"label": "M&A", "value": 150},
        {"label": "Online", "value": 185},
    ]

    # --- Charting Parameters ---
    chart_area = {
        "x": Inches(1.0), "y": Inches(2.5),
        "width": Inches(11.333), "height": Inches(4.5)
    }
    max_y_value = 2500 # This sets the scale of the Y-axis
    y_scale = chart_area["height"] / max_y_value
    bar_width = Inches(1.0)
    category_gap = Inches(0.5)
    x_offset_s2 = Inches(0.1) # Horizontal offset for scenario 2 bars

    # --- Color Palette ---
    colors = {
        "s1_total": RGBColor(128, 128, 128),
        "s1_increment": RGBColor(47, 84, 150),
        "s2_total": RGBColor(191, 191, 191),
        "s2_increment": RGBColor(155, 194, 230),
        "text": RGBColor(0, 0, 0),
        "connector": RGBColor(89, 89, 89),
    }

    # --- Helper function to add labels ---
    def add_label(shape, text, font_size=12, bold=False, color=colors["text"]):
        text_frame = shape.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.text = str(text)
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = PP_ALIGN.CENTER
        text_frame.margin_bottom = text_frame.margin_top = Pt(0)
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # --- Draw Chart Elements ---
    num_categories = len(scenario1_changes) + 2  # Start, changes, End
    current_x = chart_area["x"]

    # 1. Draw Start Bar
    start_height = start_value * y_scale
    start_top = chart_area["y"] + chart_area["height"] - start_height
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, current_x, start_top, bar_width, start_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = colors["s1_total"]
    shape.line.color.rgb = RGBColor(0,0,0)
    add_label(shape, f"{start_value:,.0f}")
    
    # Category Label for start bar
    tb = slide.shapes.add_textbox(current_x, chart_area["y"] + chart_area["height"] + Inches(0.1), bar_width, Inches(0.3))
    add_label(tb, "FC 20X1", font_size=10)


    # 2. Process and Draw Waterfall segments
    running_total_s1 = start_value
    running_total_s2 = start_value
    
    for i, (s1_change, s2_change) in enumerate(zip(scenario1_changes, scenario2_changes)):
        current_x += bar_width + category_gap

        # Scenario 1 Bar
        s1_height = s1_change["value"] * y_scale
        s1_top = chart_area["y"] + chart_area["height"] - (running_total_s1 * y_scale) - s1_height
        shape_s1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, current_x, s1_top, bar_width, s1_height)
        shape_s1.fill.solid()
        shape_s1.fill.fore_color.rgb = colors["s1_increment"]
        shape_s1.line.color.rgb = RGBColor(0,0,0)
        add_label(shape_s1, f'{s1_change["value"]:.0f}')

        # Scenario 2 Bar (Delta)
        s2_delta_val = s2_change["value"] - s1_change["value"]
        s2_height = s2_delta_val * y_scale
        s2_top = s1_top - s2_height
        shape_s2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, current_x + x_offset_s2, s2_top, bar_width, s2_height)
        shape_s2.fill.solid()
        shape_s2.fill.fore_color.rgb = colors["s2_increment"]
        shape_s2.line.color.rgb = RGBColor(0,0,0)
        add_label(shape_s2, f'{s2_change["value"]:.0f}') # The video shows the full value, not the delta

        # Category Label
        tb = slide.shapes.add_textbox(current_x, chart_area["y"] + chart_area["height"] + Inches(0.1), bar_width, Inches(0.3))
        add_label(tb, s1_change["label"], font_size=10)

        running_total_s1 += s1_change["value"]
        running_total_s2 += s2_change["value"]

    # 3. Draw End Bars
    current_x += bar_width + category_gap

    # Scenario 1 End Bar
    s1_end_height = running_total_s1 * y_scale
    s1_end_top = chart_area["y"] + chart_area["height"] - s1_end_height
    shape_s1_end = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, current_x, s1_end_top, bar_width, s1_end_height)
    shape_s1_end.fill.solid()
    shape_s1_end.fill.fore_color.rgb = colors["s1_total"]
    shape_s1_end.line.color.rgb = RGBColor(0,0,0)
    add_label(shape_s1_end, f"{running_total_s1:,.0f}")
    
    # Scenario 2 End Bar
    s2_end_height = running_total_s2 * y_scale
    s2_end_top = chart_area["y"] + chart_area["height"] - s2_end_height
    shape_s2_end = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, current_x + bar_width, s2_end_top, bar_width, s2_end_height)
    shape_s2_end.fill.solid()
    shape_s2_end.fill.fore_color.rgb = colors["s2_total"]
    shape_s2_end.line.color.rgb = RGBColor(0,0,0)
    add_label(shape_s2_end, f"{running_total_s2:,.0f}")
    
    # Category Labels for end bars
    tb1 = slide.shapes.add_textbox(current_x, chart_area["y"] + chart_area["height"] + Inches(0.1), bar_width, Inches(0.3))
    add_label(tb1, "Bud 20X2", font_size=10)
    tb2 = slide.shapes.add_textbox(current_x + bar_width, chart_area["y"] + chart_area["height"] + Inches(0.1), bar_width, Inches(0.3))
    add_label(tb2, "Bud 20X2", font_size=10)

    # 4. Draw Connector Arrow
    arrow_start_x = chart_area["x"]
    arrow_start_y = start_top
    arrow_end_x = shape_s2_end.left + shape_s2_end.width
    arrow_end_y = s2_end_top

    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, arrow_start_x, arrow_start_y - Inches(0.2), arrow_end_x, arrow_start_y - Inches(0.2))
    connector.line.color.rgb = colors["connector"]
    connector.line.width = Pt(1)

    # Vertical lines for arrow
    line1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, arrow_start_x, arrow_start_y - Inches(0.2), arrow_start_x, arrow_start_y)
    line1.line.color.rgb = colors["connector"]
    line2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, arrow_end_x, arrow_start_y - Inches(0.2), arrow_end_x, arrow_end_y)
    line2.line.color.rgb = colors["connector"]
    line2.line.end_arrow_type = MSO_ARROWHEAD.TRIANGLE
    
    # Arrow label
    growth_percent = (running_total_s2 / start_value - 1) * 100
    arrow_label_tb = slide.shapes.add_shape(MSO_SHAPE.OVAL, (arrow_start_x + arrow_end_x)/2 - Inches(0.4), arrow_start_y - Inches(0.4), Inches(0.8), Inches(0.4))
    arrow_label_tb.fill.solid()
    arrow_label_tb.fill.fore_color.rgb = RGBColor(255, 255, 255)
    arrow_label_tb.line.color.rgb = colors["connector"]
    add_label(arrow_label_tb, f"+{growth_percent:.0f}%", font_size=10)

    # --- Titles and Legend ---
    title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    p = title_shape.text_frame.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(24)

    subtitle_shape = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(0.5))
    p = subtitle_shape.text_frame.paragraphs[0]
    p.text = subtitle_text
    p.font.size = Pt(18)

    # --- Save Presentation ---
    prs.save(output_pptx_path)
    return output_pptx_path
