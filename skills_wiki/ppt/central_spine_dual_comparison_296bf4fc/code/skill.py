def create_slide(
    output_pptx_path: str,
    title_text: str = "Option A vs Option B Comparison",
    body_text: str = "",
    bg_palette: str = "gray",
    accent_color: tuple = (41, 153, 175),
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Central Spine Dual Comparison' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml import parse_xml

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Color Palette ---
    A_DARK = RGBColor(205, 83, 76)
    A_LIGHT = RGBColor(234, 130, 104)
    B_DARK = RGBColor(*accent_color)
    B_LIGHT = RGBColor(min(accent_color[0]+77, 255), min(accent_color[1]+42, 255), min(accent_color[2]+37, 255))
    GRAY_BG = RGBColor(240, 240, 240)

    # --- Helper Functions ---
    def add_shadow(shape):
        """Injects a native PowerPoint drop shadow via lxml."""
        spPr = shape.element.spPr
        shadow_xml = """
        <a:outerShdw xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" 
                     blurRad="40000" dist="35000" dir="2700000" algn="tl" rotWithShape="0">
            <a:srgbClr val="000000">
                <a:alpha val="20000"/>
            </a:srgbClr>
        </a:outerShdw>
        """
        effectLst = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        if effectLst is None:
            effectLst = OxmlElement('a:effectLst')
            spPr.append(effectLst)
        effectLst.append(parse_xml(shadow_xml))

    def draw_polygon(slide, points, color):
        """Draws a custom polygon shape using FreeformBuilder."""
        builder = slide.shapes.build_freeform()
        start_pt = points[0]
        builder.add_line_segments(
            [(Inches(x), Inches(y)) for x, y in points[1:]],
            start=(Inches(start_pt[0]), Inches(start_pt[1]))
        )
        shape = builder.convert_to_shape()
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color  # Hide border
        return shape

    # --- Global Title ---
    title_box = slide.shapes.add_textbox(Inches(0), Inches(0.2), Inches(13.333), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(80, 80, 80)
    p.alignment = PP_ALIGN.CENTER

    # --- Option Headers ---
    for x_pos, text, color, align in [(1.0, "OPTION A", A_DARK, PP_ALIGN.RIGHT), 
                                      (7.333, "OPTION B", B_DARK, PP_ALIGN.LEFT)]:
        hdr = slide.shapes.add_textbox(Inches(x_pos), Inches(0.8), Inches(5.0), Inches(0.5))
        hp = hdr.text_frame.paragraphs[0]
        hp.text = text
        hp.font.size = Pt(22)
        hp.font.bold = True
        hp.font.color.rgb = color
        hp.alignment = align

    # --- Background Horizontal Divider Bars ---
    y_bg_positions = [1.2, 2.6, 4.0, 5.4]
    for y in y_bg_positions:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), Inches(13.333), Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GRAY_BG
        bar.line.color.rgb = GRAY_BG

    # --- Geometry Layout Math ---
    x_center = 6.666
    spine_w = 0.5
    h_spine = 1.0
    h_outer = 0.5
    outer_x_list = [3.0, 4.0, 5.0] # Creates the funnel effect

    # Arrays to hold spine blocks for shadow application later (Z-order management)
    spine_shapes = []

    # --- Build Comparison Levels ---
    for i in range(3):
        outer_x = outer_x_list[i]
        y_top = 1.6 + i * 1.4
        y_bottom_inner = y_top + h_spine
        y_bottom_outer = y_top + h_outer

        # 1. Left Rib (Polygon)
        pts_l = [(outer_x, y_top), (x_center - spine_w, y_top), 
                 (x_center - spine_w, y_bottom_inner), (outer_x, y_bottom_outer), (outer_x, y_top)]
        draw_polygon(slide, pts_l, A_LIGHT)

        # 2. Right Rib (Polygon)
        right_outer_x = 13.333 - outer_x
        pts_r = [(right_outer_x, y_top), (x_center + spine_w, y_top), 
                 (x_center + spine_w, y_bottom_inner), (right_outer_x, y_bottom_outer), (right_outer_x, y_top)]
        draw_polygon(slide, pts_r, B_LIGHT)

        # 3. Text Boxes
        # Left
        tb_l = slide.shapes.add_textbox(Inches(0.5), Inches(y_top), Inches(outer_x - 1.0), Inches(1.0))
        p_l1 = tb_l.text_frame.paragraphs[0]
        p_l1.text = f"Feature Point {i+1}"
        p_l1.font.bold = True
        p_l1.font.size = Pt(14)
        p_l1.alignment = PP_ALIGN.RIGHT
        p_l2 = tb_l.text_frame.add_paragraph()
        p_l2.text = "Strategic advantage highlighted here with supporting detail."
        p_l2.font.size = Pt(11)
        p_l2.font.color.rgb = RGBColor(120, 120, 120)
        p_l2.alignment = PP_ALIGN.RIGHT

        # Right
        tb_r = slide.shapes.add_textbox(Inches(13.333 - outer_x + 0.5), Inches(y_top), Inches(outer_x - 1.0), Inches(1.0))
        p_r1 = tb_r.text_frame.paragraphs[0]
        p_r1.text = f"Alternative Point {i+1}"
        p_r1.font.bold = True
        p_r1.font.size = Pt(14)
        p_r1.alignment = PP_ALIGN.LEFT
        p_r2 = tb_r.text_frame.add_paragraph()
        p_r2.text = "Strategic advantage highlighted here with supporting detail."
        p_r2.font.size = Pt(11)
        p_r2.font.color.rgb = RGBColor(120, 120, 120)
        p_r2.alignment = PP_ALIGN.LEFT

        # 4. Center Spine Blocks (Rendered after ribs to sit on top)
        spine_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_center - spine_w), Inches(y_top), Inches(spine_w), Inches(h_spine))
        spine_l.fill.solid()
        spine_l.fill.fore_color.rgb = A_DARK
        spine_l.line.color.rgb = A_DARK
        spine_shapes.append(spine_l)

        spine_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_center), Inches(y_top), Inches(spine_w), Inches(h_spine))
        spine_r.fill.solid()
        spine_r.fill.fore_color.rgb = B_DARK
        spine_r.line.color.rgb = B_DARK
        spine_shapes.append(spine_r)

        # 5. Icon Placeholders (Small white circles)
        for cx in [x_center - spine_w/2, x_center + spine_w/2]:
            icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.15), Inches(y_top + 0.35), Inches(0.3), Inches(0.3))
            icon.fill.solid()
            icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
            icon.line.color.rgb = RGBColor(255, 255, 255)
            add_shadow(icon)

    # --- Bottom Trunks & Option Circles ---
    y_base = 1.6 + 2 * 1.4 + h_spine # bottom of last level
    
    trunk_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_center - spine_w), Inches(y_base), Inches(spine_w), Inches(0.8))
    trunk_l.fill.solid()
    trunk_l.fill.fore_color.rgb = A_LIGHT
    trunk_l.line.color.rgb = A_LIGHT

    trunk_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_center), Inches(y_base), Inches(spine_w), Inches(0.8))
    trunk_r.fill.solid()
    trunk_r.fill.fore_color.rgb = B_LIGHT
    trunk_r.line.color.rgb = B_LIGHT

    for label, color, offset in [("A", A_DARK, -spine_w/2), ("B", B_DARK, spine_w/2)]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x_center + offset - 0.4), Inches(y_base + 0.4), Inches(0.8), Inches(0.8))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.color.rgb = color
        tf = circle.text_frame
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        spine_shapes.append(circle) # Add to shadow list

    # Apply Z-depth shadows to all spine and circle elements
    for shape in spine_shapes:
        add_shadow(shape)

    prs.save(output_pptx_path)
    return output_pptx_path
