def create_slide(
    output_pptx_path: str,
    title_text: str = "System Authentication Flow",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing a professional, draw.io-style flowchart 
    using native shapes, precise elbow routing, and injected arrowheads.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.shapes.freeform import FreeformBuilder
    from pptx.oxml import OxmlElement

    # Initialize presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title_text
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(51, 51, 51)
    tf.paragraphs[0].font.name = "Arial"

    # --- HELPER 1: Draw Nodes ---
    def add_node(text, shape_type, cx, cy, w, h, bg_rgb):
        """Creates a perfectly centered geometric shape representing a flowchart node."""
        left = cx - w / 2
        top = cy - h / 2
        shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(w), Inches(h))
        
        # Apply standard draw.io style (pastel fill, dark gray border)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*bg_rgb)
        shape.line.color.rgb = RGBColor(102, 102, 102)
        shape.line.width = Pt(1.5)
        
        # Apply text formatting
        shape.text_frame.text = text
        shape.text_frame.word_wrap = True
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = RGBColor(51, 51, 51)
            p.font.size = Pt(13)
            p.font.name = "Arial"
            p.font.bold = True
            
        # Return exact boundaries for perfect line routing
        return {"top": top, "bottom": top + h, "left": left, "right": left + w, "cx": cx, "cy": cy}

    # --- HELPER 2: Draw Connectors with Arrowheads ---
    def add_arrow(p1, p2, elbow=False, orientation="v"):
        """Draws an orthogonal line connecting two points, adding an XML arrowhead."""
        fb = FreeformBuilder(slide.shapes)
        fb.move_to(Inches(p1[0]), Inches(p1[1]))
        
        # Orthogonal elbow routing
        if elbow:
            if orientation == "v":
                mid_y = (p1[1] + p2[1]) / 2
                fb.line_to(Inches(p1[0]), Inches(mid_y))
                fb.line_to(Inches(p2[0]), Inches(mid_y))
                fb.line_to(Inches(p2[0]), Inches(p2[1]))
            elif orientation == "h":
                mid_x = (p1[0] + p2[0]) / 2
                fb.line_to(Inches(mid_x), Inches(p1[1]))
                fb.line_to(Inches(mid_x), Inches(p2[1]))
                fb.line_to(Inches(p2[0]), Inches(p2[1]))
        else:
            fb.line_to(Inches(p2[0]), Inches(p2[1]))

        # Convert to line shape
        shape = fb.convert_to_shape()
        shape.line.color.rgb = RGBColor(102, 102, 102)
        shape.line.width = Pt(1.5)

        # LXML Injection: Add triangular arrowhead to the path end
        ln = shape.element.spPr.ln
        if ln is not None:
            tailEnd = OxmlElement('a:tailEnd')
            tailEnd.set('type', 'triangle')
            tailEnd.set('w', 'med')
            tailEnd.set('len', 'med')
            ln.append(tailEnd)
            
    # --- HELPER 3: Add Floating Labels ---
    def add_label(text, cx, cy):
        """Adds condition labels (e.g., 'Yes', 'No') on branches."""
        txBox = slide.shapes.add_textbox(Inches(cx - 0.5), Inches(cy - 0.25), Inches(1), Inches(0.5))
        txBox.text_frame.text = text
        for p in txBox.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(102, 102, 102)
            p.font.italic = True
            p.font.name = "Arial"

    # === CORE EXECUTION: Render Nodes ===
    # Using standardized draw.io pastel palette mapping
    C_GREEN = (213, 232, 212)
    C_YELLOW = (255, 242, 204)
    C_PURPLE = (225, 213, 231)
    C_BLUE = (218, 232, 252)
    C_RED = (248, 206, 204)

    # Coordinates structure the flow hierarchically
    n_start = add_node("Receive Request", MSO_SHAPE.ROUNDED_RECTANGLE, 6.66, 1.5, 2.2, 0.8, C_GREEN)
    n_check = add_node("Auth Token\nValid?", MSO_SHAPE.DIAMOND, 6.66, 3.5, 2.0, 1.5, C_YELLOW)
    n_db = add_node("User DB", MSO_SHAPE.CAN, 10.5, 3.5, 1.2, 1.5, C_PURPLE)
    n_app = add_node("Process Payload", MSO_SHAPE.RECTANGLE, 3.5, 6.0, 2.2, 1.0, C_BLUE)
    n_rej = add_node("Reject Request", MSO_SHAPE.RECTANGLE, 9.8, 6.0, 2.2, 1.0, C_RED)

    # === CORE EXECUTION: Route Lines ===
    # Downward straight flow
    add_arrow((n_start["cx"], n_start["bottom"]), (n_check["cx"], n_check["top"]), elbow=False)
    
    # Horizontal straight flow to DB
    add_arrow((n_check["right"], n_check["cy"]), (n_db["left"], n_db["cy"]), elbow=False)
    
    # Orthogonal branches out of the decision diamond
    # Branch 1 (Valid)
    add_arrow((n_check["cx"], n_check["bottom"]), (n_app["cx"], n_app["top"]), elbow=True, orientation="v")
    # Branch 2 (Invalid)
    add_arrow((n_check["cx"], n_check["bottom"]), (n_rej["cx"], n_rej["top"]), elbow=True, orientation="v")

    # === CORE EXECUTION: Add Labels ===
    add_label("Verify", 8.5, 3.25)
    add_label("Yes", 5.0, 4.5)
    add_label("No", 8.2, 4.5)

    # Save output
    prs.save(output_pptx_path)
    return output_pptx_path
