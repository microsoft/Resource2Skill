import os
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_cyberpunk_hud_slide(
    output_pptx_path: str,
    title_text: str = "WELCOME TO CYBER SECURITY CLASS",
    subtitle_text: str = "SYSTEM INITIALIZATION COMPLETE",
    panel_titles: list = ["01: ENCRYPTION", "02: FIREWALLS", "03: PROTOCOLS"],
    bg_color: tuple = (10, 5, 25),      # Deep void purple
    neon_cyan: tuple = (0, 255, 255),   # Glow color 1
    neon_pink: tuple = (255, 0, 128)    # Glow color 2
) -> str:
    """
    Creates a PPTX file reproducing the Cyberpunk Neon HUD Dashboard effect.
    Returns: path to the saved PPTX file.
    """
    # Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # ==========================================
    # Layer 1: Generate Cyber Grid Background (PIL)
    # ==========================================
    bg_path = "temp_cyber_bg.png"
    bg_img = Image.new("RGBA", (1920, 1080), bg_color + (255,))
    draw_bg = ImageDraw.Draw(bg_img, "RGBA")
    
    # Draw faint perspective/tech grid
    grid_color = neon_cyan + (30,) # Low opacity
    grid_spacing = 80
    for x in range(0, 1920, grid_spacing):
        draw_bg.line([(x, 0), (x, 1080)], fill=grid_color, width=2)
    for y in range(0, 1080, grid_spacing):
        draw_bg.line([(0, y), (1920, y)], fill=grid_color, width=2)
        
    # Add a subtle vignette/dark edge
    vignette = Image.new("RGBA", (1920, 1080), (0, 0, 0, 0))
    v_draw = ImageDraw.Draw(vignette)
    v_draw.rectangle([0, 0, 1920, 1080], outline=(0, 0, 0, 150), width=150)
    vignette = vignette.filter(ImageFilter.GaussianBlur(100))
    bg_img = Image.alpha_composite(bg_img, vignette)
    
    bg_img.save(bg_path)
    slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)

    # ==========================================
    # Helper: Generate Glowing HUD Panel (PIL)
    # ==========================================
    def create_hud_panel(filename, w_px, h_px, glow_color, fill_color=(10, 15, 30, 200)):
        panel = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
        
        # Base glass fill
        draw = ImageDraw.Draw(panel, "RGBA")
        margin = 30
        panel_rect = [margin, margin, w_px-margin, h_px-margin]
        draw.rectangle(panel_rect, fill=fill_color)
        
        # Glow Layer (thick blurred line)
        glow_layer = Image.new("RGBA", (w_px, h_px), (0, 0, 0, 0))
        g_draw = ImageDraw.Draw(glow_layer, "RGBA")
        g_draw.rectangle(panel_rect, outline=glow_color + (255,), width=15)
        glow_layer = glow_layer.filter(ImageFilter.GaussianBlur(15))
        
        # Composite glow onto base
        panel = Image.alpha_composite(panel, glow_layer)
        
        # Core sharp line
        draw = ImageDraw.Draw(panel, "RGBA")
        draw.rectangle(panel_rect, outline=(255, 255, 255, 255), width=2)
        
        # HUD Accents (Corners)
        accent_len = 30
        accent_w = 6
        # Top Left
        draw.line([(margin, margin), (margin+accent_len, margin)], fill=glow_color+(255,), width=accent_w)
        draw.line([(margin, margin), (margin, margin+accent_len)], fill=glow_color+(255,), width=accent_w)
        # Bottom Right
        draw.line([(w_px-margin, h_px-margin), (w_px-margin-accent_len, h_px-margin)], fill=glow_color+(255,), width=accent_w)
        draw.line([(w_px-margin, h_px-margin), (w_px-margin, h_px-margin-accent_len)], fill=glow_color+(255,), width=accent_w)
        
        panel.save(filename)
        return filename

    # ==========================================
    # Layer 2 & 3: Insert Panels and Native Text
    # ==========================================
    
    # --- Main Title Panel ---
    main_panel_path = "temp_main_panel.png"
    create_hud_panel(main_panel_path, 1200, 400, neon_cyan)
    # Center horizontally, near top
    pic_w, pic_h = Inches(8.33), Inches(2.77)
    pic_left = (prs.slide_width - pic_w) / 2
    pic_top = Inches(0.8)
    slide.shapes.add_picture(main_panel_path, pic_left, pic_top, pic_w, pic_h)
    
    # Main Title Text
    tx_box = slide.shapes.add_textbox(pic_left, pic_top + Inches(0.6), pic_w, Inches(1))
    tf = tx_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial Black"
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(*neon_cyan)
    
    # Subtitle Text
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Consolas" # Tech font
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(255, 255, 255)

    # --- Feature Panels (Bottom Row) ---
    panel_w_in, panel_h_in = Inches(3.5), Inches(2.5)
    spacing_in = Inches(0.6)
    total_w = (3 * panel_w_in) + (2 * spacing_in)
    start_left = (prs.slide_width - total_w) / 2
    top_pos = Inches(4.2)
    
    sub_panel_path = "temp_sub_panel.png"
    create_hud_panel(sub_panel_path, 600, 450, neon_pink) # Pink glow for contrast

    for i in range(3):
        curr_left = start_left + i * (panel_w_in + spacing_in)
        # Add PIL Graphic
        slide.shapes.add_picture(sub_panel_path, curr_left, top_pos, panel_w_in, panel_h_in)
        
        # Add native text over the graphic
        t_box = slide.shapes.add_textbox(curr_left, top_pos + Inches(0.4), panel_w_in, panel_h_in)
        t_f = t_box.text_frame
        t_f.word_wrap = True
        
        sp = t_f.paragraphs[0]
        sp.text = panel_titles[i] if i < len(panel_titles) else f"NODE {i+1}"
        sp.alignment = PP_ALIGN.CENTER
        sp.font.name = "Arial Black"
        sp.font.size = Pt(20)
        sp.font.color.rgb = RGBColor(*neon_pink)
        
        bp = t_f.add_paragraph()
        bp.text = "\nSystem parameters initialized. Data stream encrypted and secure. Awaiting command."
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = "Arial"
        bp.font.size = Pt(12)
        bp.font.color.rgb = RGBColor(200, 200, 220)

    # Cleanup temp files
    prs.save(output_pptx_path)
    
    for tmp in [bg_path, main_panel_path, sub_panel_path]:
        if os.path.exists(tmp):
            os.remove(tmp)

    return output_pptx_path
