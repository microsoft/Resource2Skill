def create_slide(
    output_pptx_path: str,
    title_text: str = "Product Overview",
    subtitle_text: str = "Enter your subhead line here",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Vibrant Corporate Mosaic Info Grid'.
    
    Returns: path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    # Use standard widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank slide

    # === Slide Background ===
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(250, 250, 250) # Very light gray for contrast

    # === Header Section ===
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(1))
    tf = tx_box.text_frame
    
    # Main Title
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = 'Arial'
    p.font.color.rgb = RGBColor(60, 60, 60)
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(14)
    p2.font.name = 'Arial'
    p2.font.color.rgb = RGBColor(120, 120, 120)

    # === Mosaic Grid Configuration ===
    # Colors matching the vibrant corporate style
    palette = [
        RGBColor(45, 168, 216),  # Light Blue
        RGBColor(36, 123, 160),  # Dark Blue
        RGBColor(117, 192, 67),  # Green
        RGBColor(246, 174, 45),  # Yellow
        RGBColor(242, 100, 25),  # Orange
        RGBColor(217, 63, 76),   # Red/Crimson
        RGBColor(142, 68, 173),  # Purple
        RGBColor(43, 172, 184)   # Teal
    ]

    # Data for the cards (matching the icons and concepts)
    card_data = [
        {"title": "Problem Statement", "desc": "Describe the business reason(s) for initiating the project specifically stating the business problem.", "icon": MSO_SHAPE.LIGHTNING_BOLT},
        {"title": "Project Description", "desc": "Describe the approach that project will use to address the business problem.", "icon": MSO_SHAPE.FLOWCHART_DOCUMENT},
        {"title": "Goals & Objectives", "desc": "Describe the business goals and objectives of the project. Refine the goals stated in the business case.", "icon": MSO_SHAPE.TARGET},
        {"title": "Assumptions", "desc": "State the critical assumptions that have been considered for this project.", "icon": MSO_SHAPE.STAR_5_POINT},
        {"title": "Project Scope", "desc": "The scope defines project limits & identifies the product/service delivered by the project.", "icon": MSO_SHAPE.BULLSEYE},
        {"title": "Project Inclusions", "desc": "This is a sample text. You simply add your own text and description here. This text is fully editable.", "icon": MSO_SHAPE.MATH_PLUS},
        {"title": "Project Exclusions", "desc": "This is a sample text. You simply add your own text and description here. This text is fully editable.", "icon": MSO_SHAPE.MATH_MINUS},
        {"title": "Critical Success", "desc": "Describe the certain factors which are so critical that in their absence, the project might fail.", "icon": MSO_SHAPE.CHEVRON},
    ]

    # Grid Math
    num_cols = 4
    num_rows = 2
    
    # Slide dimensions available for grid
    start_x = Inches(0.5)
    start_y = Inches(1.8)
    slide_w_avail = Inches(13.333) - Inches(1.0) # 0.5 inch margins
    slide_h_avail = Inches(7.5) - Inches(2.2)    # account for header and bottom margin
    
    gutter = Inches(0.1) # Thin gutter for modern mosaic look
    
    card_width = (slide_w_avail - (gutter * (num_cols - 1))) / num_cols
    card_height = (slide_h_avail - (gutter * (num_rows - 1))) / num_rows

    # Generate Grid
    for i, data in enumerate(card_data):
        row = i // num_cols
        col = i % num_cols
        
        x = start_x + (col * (card_width + gutter))
        y = start_y + (row * (card_height + gutter))
        color = palette[i % len(palette)]
        
        # 1. Base Card Shape (Rounded Rectangle)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_width, card_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background() # No border
        
        # Modify roundness (adjusting adjust_value)
        if shape.adjustments:
            shape.adjustments[0] = 0.05 # slight rounding

        # 2. Add Icon
        icon_size = Inches(0.6)
        icon_x = x + (card_width - icon_size) / 2
        icon_y = y + Inches(0.3)
        
        icon = slide.shapes.add_shape(
            data["icon"], icon_x, icon_y, icon_size, icon_size
        )
        icon.fill.solid()
        icon.fill.fore_color.rgb = RGBColor(255, 255, 255)
        icon.line.fill.background()

        # 3. Add Title
        tx_width = card_width - Inches(0.4)
        tx_x = x + Inches(0.2)
        tx_y = icon_y + icon_size + Inches(0.1)
        
        title_box = slide.shapes.add_textbox(tx_x, tx_y, tx_width, Inches(0.4))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.name = 'Arial'
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)
        p_title.alignment = PP_ALIGN.CENTER

        # 4. Add Body Text
        desc_y = tx_y + Inches(0.4)
        desc_box = slide.shapes.add_textbox(tx_x, desc_y, tx_width, Inches(1.0))
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = data["desc"]
        p_desc.font.name = 'Arial'
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = RGBColor(255, 255, 255)
        p_desc.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    return output_pptx_path
