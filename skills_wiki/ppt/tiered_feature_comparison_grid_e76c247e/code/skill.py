def create_slide(
    output_pptx_path: str,
    title_text: str = "Product capability comparison for photo and videos editing website",
    company_name: str = "ABC Pvt. Ltd.",
    service_name: str = "Photo & Videos editing website",
    tiers: list = None,
    features: list = None,
    **kwargs,
) -> str:
    """
    Creates a PowerPoint slide with a tiered feature comparison grid.

    Args:
        output_pptx_path (str): The path to save the generated PPTX file.
        title_text (str): The main title of the slide.
        company_name (str): The name of the company.
        service_name (str): The name of the service being compared.
        tiers (list): A list of tier names (e.g., ["Basic", "Portfolio", "Business"]).
        features (list): A list of dictionaries, where each dict contains a "name"
                         and a "values" list of booleans corresponding to the tiers.

    Returns:
        str: The path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.shapes import MSO_SHAPE

    # --- Default Data if not provided ---
    if tiers is None:
        tiers = ["Basic", "Portfolio", "Business"]
    if features is None:
        features = [
            {"name": "Customizable website", "values": [True, True, True]},
            {"name": "Unlimited photos and videos upload", "values": [True, True, True]},
            {"name": "Responsive design", "values": [True, True, True]},
            {"name": "Free mobile app to edit", "values": [True, True, True]},
            {"name": "Fully hosted, unlimited traffic", "values": [False, True, True]},
            {"name": "Ads and spam", "values": [True, True, True]},
            {"name": "Share on the go", "values": [False, False, True]},
            {"name": "Add text here", "values": [False, False, True]},
            {"name": "Add text here", "values": [False, False, True]},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Color Palette ---
    ACCENT_COLOR = RGBColor(79, 129, 133)
    WHITE_COLOR = RGBColor(255, 255, 255)
    DARK_GRAY_COLOR = RGBColor(64, 64, 64)
    GREEN_CHECK_COLOR = RGBColor(0, 176, 80)
    RED_CROSS_COLOR = RGBColor(255, 0, 0)
    LIGHT_GRAY_BORDER = RGBColor(217, 217, 217)

    # === Slide Title ===
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.5))
    title_shape.text_frame.text = title_text
    p = title_shape.text_frame.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = DARK_GRAY_COLOR

    # === Top Info Bar ===
    company_box = slide.shapes.add_textbox(Inches(2.8), Inches(1.2), Inches(2), Inches(0.3))
    company_box.text_frame.text = f"Company Name: {company_name}"
    company_box.text_frame.paragraphs[0].font.size = Pt(10)

    service_box = slide.shapes.add_textbox(Inches(6.5), Inches(1.2), Inches(3), Inches(0.3))
    service_box.text_frame.text = f"Services: {service_name}"
    service_box.text_frame.paragraphs[0].font.size = Pt(10)

    # === Vertical "Features" Bar ===
    features_bar = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(4), Inches(0.5))
    features_bar.rotation = 270
    features_bar.text_frame.text = "Features"
    features_bar.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    features_bar.text_frame.paragraphs[0].font.color.rgb = WHITE_COLOR
    features_bar.text_frame.paragraphs[0].font.bold = True
    features_bar.text_frame.paragraphs[0].font.size = Pt(18)
    features_bar.fill.solid()
    features_bar.fill.fore_color.rgb = ACCENT_COLOR
    features_bar.line.fill.background()

    # === Feature Name Text Boxes ===
    feature_y_start = Inches(2.5)
    feature_row_height = Inches(0.45)
    for i, feature in enumerate(features):
        y_pos = feature_y_start + (i * feature_row_height)
        tx_box = slide.shapes.add_textbox(Inches(1.2), y_pos, Inches(3.5), feature_row_height)
        tx_box.text_frame.text = feature["name"]
        tx_box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        p = tx_box.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_GRAY_COLOR
        p.alignment = PP_ALIGN.LEFT
        tx_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # === Comparison Table ===
    table_cols = len(tiers)
    table_rows = len(features) + 1
    table_left = Inches(4.8)
    table_top = Inches(2.0)
    table_width = Inches(8)
    table_height = Inches(0.5) + (len(features) * feature_row_height)

    table_shape = slide.shapes.add_table(table_rows, table_cols, table_left, table_top, table_width, table_height)
    table = table_shape.table
    
    # --- Set Column Widths ---
    for i in range(table_cols):
        table.columns[i].width = int(table_width / table_cols)

    # --- Format Table Header ---
    for i, tier_name in enumerate(tiers):
        cell = table.cell(0, i)
        cell.text = tier_name
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(16)
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE_COLOR
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_COLOR
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # --- Populate Table Body ---
    for row_idx, feature in enumerate(features, start=1):
        for col_idx, value in enumerate(feature["values"]):
            cell = table.cell(row_idx, col_idx)
            paragraph = cell.text_frame.paragraphs[0]
            if value:
                paragraph.text = "✓"
                paragraph.font.color.rgb = GREEN_CHECK_COLOR
            else:
                paragraph.text = "✗"
                paragraph.font.color.rgb = RED_CROSS_COLOR
            paragraph.font.name = "Segoe UI Symbol" # A font that reliably renders these symbols
            paragraph.font.size = Pt(24)
            paragraph.font.bold = True
            paragraph.alignment = PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    # --- Apply Borders to all cells ---
    for r in range(table_rows):
        for c in range(table_cols):
            cell = table.cell(r, c)
            for border_part in ['top', 'bottom', 'left', 'right']:
                border = getattr(cell, f'border_{border_part}')
                border.fill.solid()
                border.fill.fore_color.rgb = LIGHT_GRAY_BORDER
                border.width = Pt(1)

    prs.save(output_pptx_path)
    return output_pptx_path

