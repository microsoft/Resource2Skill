def create_slide(
    output_pptx_path: str,
    title_text: str = "INFOGRAPHIC\nSIX STEPS",
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Radial Exploded Donut layout.
    """
    import math
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageFilter
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Set Slide Background ---
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(52, 66, 75)  # Dark Slate

    # --- Generate Core Graphic via PIL ---
    # We draw at 2x scale to ensure perfectly smooth, anti-aliased arcs, then downsample.
    scale = 2
    canvas_w, canvas_h = 1920 * scale, 1080 * scale
    cx, cy = 960 * scale, 540 * scale
    outer_R, inner_R = 300 * scale, 120 * scale
    explode = 15 * scale

    # Angles and Colors (RGBA) - Clockwise from 135 deg (Bottom Left)
    segments = [
        (135, 180, (242, 107, 56, 255)),   # Orange (Bottom Left)
        (180, 225, (242, 166, 56, 255)),   # Yellow-Orange
        (225, 270, (242, 203, 56, 255)),   # Yellow
        (270, 315, (155, 191, 64, 255)),   # Light Green
        (315, 360, (38, 140, 96, 255)),    # Dark Green
        (0, 45, (25, 126, 140, 255))       # Teal (Bottom Right)
    ]

    shadow_base = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 0))
    color_base = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 0))
    
    node_coords = [] # To store positions for PPTX text

    for i, (start_a, end_a, color) in enumerate(segments):
        mid_a = math.radians((start_a + end_a) / 2)
        
        # Calculate exploded center
        dx = explode * math.cos(mid_a)
        dy = explode * math.sin(mid_a)
        sx, sy = cx + dx, cy + dy

        # 1. Base Mask for the Donut Segment
        mask = Image.new('L', (canvas_w, canvas_h), 0)
        mdraw = ImageDraw.Draw(mask)
        mdraw.pieslice([sx - outer_R, sy - outer_R, sx + outer_R, sy + outer_R], start_a, end_a, fill=255)
        mdraw.pieslice([sx - inner_R, sy - inner_R, sx + inner_R, sy + inner_R], start_a, end_a, fill=0)

        # 2. Outer Drop Shadow (Offset slightly to the bottom right)
        shadow_mask = mask.filter(ImageFilter.GaussianBlur(12 * scale))
        shadow_layer = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 140))
        shadow_comp = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 0))
        shadow_comp.paste(shadow_layer, (6 * scale, 6 * scale), shadow_mask)
        shadow_base.alpha_composite(shadow_comp)

        # 3. Main Color Fill
        color_layer = Image.new('RGBA', (canvas_w, canvas_h), color)
        color_base.paste(color_layer, (0, 0), mask)

        # 4. Inner Depth Overlay (The 3D lip effect)
        band_width = (outer_R - inner_R) * 0.35
        overlay_mask = Image.new('L', (canvas_w, canvas_h), 0)
        odraw = ImageDraw.Draw(overlay_mask)
        # Draw a faded black band spanning outward from inner radius
        odraw.pieslice([sx - (inner_R + band_width), sy - (inner_R + band_width), 
                        sx + (inner_R + band_width), sy + (inner_R + band_width)], start_a, end_a, fill=110)
        # Clear out the actual hole
        odraw.pieslice([sx - inner_R, sy - inner_R, sx + inner_R, sy + inner_R], start_a, end_a, fill=0)
        
        overlay_layer = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 255))
        color_base.paste(overlay_layer, (0, 0), overlay_mask)

        # 5. Outer Node Pins
        node_dist = outer_R
        nx = sx + node_dist * math.cos(mid_a)
        ny = sy + node_dist * math.sin(mid_a)

        ndraw = ImageDraw.Draw(color_base)
        # Outer dark ring shadow
        node_bg_r = 45 * scale
        ndraw.ellipse([nx - node_bg_r, ny - node_bg_r, nx + node_bg_r, ny + node_bg_r], fill=(30, 40, 50, 255))
        # White ring
        node_w_r = 38 * scale
        ndraw.ellipse([nx - node_w_r, ny - node_w_r, nx + node_w_r, ny + node_w_r], fill=(255, 255, 255, 255))
        # Inner color fill
        node_c_r = 28 * scale
        ndraw.ellipse([nx - node_c_r, ny - node_c_r, nx + node_c_r, ny + node_c_r], fill=color)

        # Save coordinates (downscaled back to 1x) for PowerPoint
        node_coords.append((nx / scale, ny / scale, mid_a, color))

    # Composite layers and resize
    main_img = Image.new('RGBA', (canvas_w, canvas_h), (0, 0, 0, 0))
    main_img.alpha_composite(shadow_base)
    main_img.alpha_composite(color_base)
    main_img = main_img.resize((1920, 1080), Image.Resampling.LANCZOS)

    # Insert Image into Slide
    img_io = BytesIO()
    main_img.save(img_io, format='PNG')
    img_io.seek(0)
    slide.shapes.add_picture(img_io, 0, 0, prs.slide_width, prs.slide_height)

    # --- Add Text & Icons via python-pptx ---
    
    # Center Hub Icon and Title
    center_w, center_h = Inches(3), Inches(2)
    center_x = (prs.slide_width - center_w) / 2
    center_y = (prs.slide_height - center_h) / 2
    
    hub_tb = slide.shapes.add_textbox(center_x, center_y, center_w, center_h)
    hub_tf = hub_tb.text_frame
    hub_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p_gear = hub_tf.paragraphs[0]
    p_gear.text = "⚙"  # Unicode Gear
    p_gear.font.size = Pt(54)
    p_gear.font.color.rgb = RGBColor(180, 180, 180)
    p_gear.alignment = PP_ALIGN.CENTER
    
    p_title = hub_tf.add_paragraph()
    p_title.text = title_text
    p_title.font.bold = True
    p_title.font.size = Pt(14)
    p_title.font.color.rgb = RGBColor(255, 255, 255)
    p_title.alignment = PP_ALIGN.CENTER

    # Segment Text Boxes
    for i, (nx, ny, mid_a, color) in enumerate(node_coords):
        # Determine Left/Right Side
        is_left = math.cos(mid_a) < 0
        
        # 1. Insert Node Number
        num_size = Inches(0.5)
        num_box = slide.shapes.add_textbox(Inches(nx / 144) - (num_size / 2), 
                                           Inches(ny / 144) - (num_size / 2) + Inches(0.03), # optical alignment
                                           num_size, num_size)
        num_tf = num_box.text_frame
        num_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        num_tf.margin_left = num_tf.margin_right = num_tf.margin_top = num_tf.margin_bottom = 0
        
        p = num_tf.paragraphs[0]
        p.text = str(i + 1)
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        # 2. Insert Titles and Descriptions
        text_w = Inches(2.2)
        text_h = Inches(1.0)
        margin = Inches(0.7)  # Spacing from the node circle to the text box

        if is_left:
            tx = Inches(nx / 144) - margin - text_w
            align = PP_ALIGN.RIGHT
        else:
            tx = Inches(nx / 144) + margin
            align = PP_ALIGN.LEFT

        ty = Inches(ny / 144) - text_h / 2
        tb = slide.shapes.add_textbox(tx, ty, text_w, text_h)
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = 0
        
        p_title = tf.paragraphs[0]
        p_title.text = f"Neque porro quisquam"
        p_title.font.bold = True
        p_title.font.size = Pt(13)
        p_title.font.color.rgb = RGBColor(255, 255, 255)
        p_title.alignment = align

        p_desc = tf.add_paragraph()
        p_desc.text = "Lorem ipsum is simply dummy text of the printing and typesetting industry."
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = RGBColor(160, 175, 185)
        p_desc.alignment = align

    prs.save(output_pptx_path)
    return output_pptx_path
