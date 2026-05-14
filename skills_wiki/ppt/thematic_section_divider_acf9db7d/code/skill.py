def create_slide(
    output_pptx_path: str,
    section_title: str = "Product Idea Screening",
    agenda_items: list = None,
    highlight_index: int = 0,
    **kwargs,
) -> str:
    """
    Creates a PPTX slide with a thematic section divider layout.

    Args:
        output_pptx_path: Path to save the generated PPTX file.
        section_title: The main title for the presentation section.
        agenda_items: A list of strings for the sub-topic agenda.
        highlight_index: The 0-based index of the agenda item to highlight.

    Returns:
        Path to the saved PPTX file.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    if agenda_items is None:
        agenda_items = [
            "New Product Introduction",
            "New Product Detailed Overview",
            "Understanding Customer Needs",
            "External Sources of Ideas",
            "Internal Sources of Ideas",
            "Product Roadmap"
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Define Colors
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_DARK_BLUE = RGBColor(30, 50, 80)
    COLOR_ACCENT_PINK = RGBColor(231, 108, 114)
    COLOR_WHITE = RGBColor(255, 255, 255)

    # Set slide background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

    # === Left Section: Title and Icon Placeholder ===
    left_margin = Inches(0.8)
    
    # Title Banner
    banner_height = Inches(0.6)
    banner_width = Inches(4.5)
    banner_top = Inches(1.5)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_margin, banner_top, banner_width, banner_height)
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_ACCENT_PINK
    banner.line.fill.background()

    # Title Text
    title_box = slide.shapes.add_textbox(left_margin, banner_top, banner_width, banner_height)
    p = title_box.text_frame.paragraphs[0]
    p.text = section_title
    p.font.name = 'Arial'
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    title_box.text_frame.margin_bottom = Inches(0)
    title_box.text_frame.margin_top = Inches(0.1)

    # Icon Container
    container_top = banner_top + banner_height
    container_height = Inches(3.5)
    container = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_margin, container_top, banner_width, container_height)
    container.fill.solid()
    container.fill.fore_color.rgb = COLOR_BG
    container.line.color.rgb = COLOR_DARK_BLUE
    container.line.width = Pt(1.5)

    # Icon Placeholder (Using a simple shape as an example)
    # INSTRUCTION: To use a real icon, replace this section with:
    # slide.shapes.add_picture('your_icon.png', icon_left, icon_top, width=icon_size)
    icon_size = Inches(1.8)
    icon_left = left_margin + (banner_width - icon_size) / 2
    icon_top = container_top + (container_height - icon_size) / 2
    icon_placeholder = slide.shapes.add_shape(MSO_SHAPE.ACTION_BUTTON_HOME, icon_left, icon_top, icon_size, icon_size)
    icon_placeholder.fill.solid()
    icon_placeholder.fill.fore_color.rgb = COLOR_DARK_BLUE
    icon_placeholder.line.fill.background()
    
    # === Right Section: Agenda List ===
    list_start_left = Inches(6.5)
    list_start_top = Inches(1.5)
    item_height = Inches(0.8)
    circle_diameter = Inches(0.4)
    text_left_margin = Inches(0.6)

    # Connector Line
    total_list_height = len(agenda_items) * item_height
    line_left = list_start_left + circle_diameter / 2
    line = slide.shapes.add_shape(MSO_SHAPE.LINE_UP, line_left, list_start_top, Pt(2), total_list_height)
    line.line.color.rgb = COLOR_DARK_BLUE
    line.line.width = Pt(1)

    # Add list items
    for i, item_text in enumerate(agenda_items):
        current_top = list_start_top + (i * item_height)
        
        # Circle for number
        is_highlighted = (i == highlight_index)
        circle_color = COLOR_ACCENT_PINK if is_highlighted else COLOR_DARK_BLUE
        
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, list_start_left, current_top, circle_diameter, circle_diameter)
        circle.fill.solid()
        circle.fill.fore_color.rgb = circle_color
        circle.line.fill.background()

        # Number in circle
        num_box = slide.shapes.add_textbox(list_start_left, current_top, circle_diameter, circle_diameter)
        p_num = num_box.text_frame.paragraphs[0]
        p_num.text = f"{(i+1):02}"
        p_num.font.name = 'Arial'
        p_num.font.bold = True
        p_num.font.size = Pt(12)
        p_num.font.color.rgb = COLOR_WHITE
        p_num.alignment = PP_ALIGN.CENTER
        num_box.text_frame.margin_bottom = Inches(0)
        num_box.text_frame.margin_top = Inches(0.08)

        # Item text
        text_box_left = list_start_left + text_left_margin
        text_box_width = Inches(5.5)
        item_box = slide.shapes.add_textbox(text_box_left, current_top - Inches(0.05), text_box_width, circle_diameter)
        p_item = item_box.text_frame.paragraphs[0]
        p_item.text = item_text
        p_item.font.name = 'Arial'
        p_item.font.size = Pt(18)
        p_item.font.color.rgb = COLOR_DARK_BLUE
        item_box.text_frame.margin_bottom = Inches(0)
        item_box.text_frame.margin_top = Inches(0)
        
        # Connector Spoke
        spoke_left = line_left + Pt(1)
        spoke_top = current_top + circle_diameter / 2
        spoke_width = text_left_margin - circle_diameter / 2 - Pt(1)
        spoke = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, spoke_left, spoke_top, spoke_width, Pt(2))
        spoke.line.color.rgb = COLOR_DARK_BLUE
        spoke.line.width = Pt(1)

    prs.save(output_pptx_path)
    return output_pptx_path

# Example usage:
# create_slide(
#     "thematic_divider_slide.pptx",
#     section_title="Market Analysis",
#     agenda_items=["Market Segmentation", "Product Market Mapping", "Competitive Strategies", "Market Attractiveness"],
#     highlight_index=1
# )

