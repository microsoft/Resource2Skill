def create_slide(
    output_pptx_path: str,
    title_text: str = "DUOTONE",
    subtitle_text: str = "EDITORIAL STYLE",
    bg_keyword: str = "portrait,fashion",
    shadow_color: tuple = (32, 18, 77),     # Deep Violet
    highlight_color: tuple = (255, 170, 85),  # Bright Peach
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the Editorial Duotone visual effect.
    Uses PIL to download and pixel-process an image, mapping shadows and 
    highlights to a custom 2-color palette before overlaying text.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from PIL import Image, ImageOps, ImageDraw

    # 1. Initialize Presentation (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    temp_img_path = "temp_source.jpg"
    duotone_img_path = "temp_duotone.jpg"
    
    # 2. Download Base Image (with fallback)
    try:
        url = f"https://source.unsplash.com/featured/1920x1080/?{bg_keyword}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response, open(temp_img_path, 'wb') as out_file:
            out_file.write(response.read())
        img = Image.open(temp_img_path)
    except Exception as e:
        print(f"Download failed, generating fallback image: {e}")
        # Fallback: Create a procedural gradient/pattern image to colorize
        img = Image.new('RGB', (1920, 1080), color='black')
        draw = ImageDraw.Draw(img)
        for i in range(1080):
            val = int(255 * (i / 1080))
            draw.line([(0, i), (1920, i)], fill=(val, val, val))
        img.save(temp_img_path)

    # 3. Apply Duotone Effect using PIL
    # Crop to exact 16:9 aspect ratio to avoid PPTX distortion
    img = ImageOps.fit(img, (1920, 1080), method=Image.Resampling.LANCZOS)
    
    # Convert to grayscale
    gray_img = img.convert('L')
    
    # Apply Duotone: remap black to shadow_color, white to highlight_color
    duotone_img = ImageOps.colorize(gray_img, black=shadow_color, white=highlight_color)
    
    # Add a subtle dark gradient vignette on the left to ensure text readability
    vignette = Image.new('RGBA', (1920, 1080), (0, 0, 0, 0))
    v_draw = ImageDraw.Draw(vignette)
    for x in range(800):
        alpha = int(200 * (1 - (x / 800))) # Fade from 200 opacity to 0
        v_draw.line([(x, 0), (x, 1080)], fill=(shadow_color[0], shadow_color[1], shadow_color[2], alpha))
    
    # Composite vignette over duotone
    duotone_img = duotone_img.convert('RGBA')
    final_img = Image.alpha_composite(duotone_img, vignette).convert('RGB')
    final_img.save(duotone_img_path, quality=95)

    # 4. Add Duotone Image to Slide Canvas
    slide.shapes.add_picture(
        duotone_img_path, 
        0, 0, 
        width=prs.slide_width, 
        height=prs.slide_height
    )

    # 5. Add Typography Overlay
    # Title Text
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(8), Inches(2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial Black'  # Thick, bold font
    p.font.size = Pt(88)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255) # Pure white
    
    # Subtitle Text (Colored with the Highlight tone)
    subtitle_box = slide.shapes.add_textbox(Inches(1.05), Inches(4.0), Inches(8), Inches(1))
    tf_sub = subtitle_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(28)
    p_sub.font.bold = True
    p_sub.font.letter_spacing = Pt(3)
    p_sub.font.color.rgb = RGBColor(*highlight_color) 

    # Clean up temp files
    prs.save(output_pptx_path)
    if os.path.exists(temp_img_path):
        os.remove(temp_img_path)
    if os.path.exists(duotone_img_path):
        os.remove(duotone_img_path)

    return output_pptx_path
