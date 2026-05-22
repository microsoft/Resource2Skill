def create_slide(
    output_pptx_path: str,
    title_text: str = "Company Profile",
    body_text: str = "",
    bg_palette: str = "corporate architecture",  
    accent_color: tuple = (45, 160, 150),  # Teal
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Diagonal Corporate Hero Banner' visual effect.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from PIL import Image, ImageDraw
    import urllib.request
    import os

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    
    W = int(prs.slide_width)
    H = int(prs.slide_height)

    # === Layer 1: Background Image Processing via PIL ===
    bg_path = "temp_bg_123.png"
    tinted_bg_path = "tinted_bg_123.png"
    try:
        # Fetch placeholder image
        req = urllib.request.Request("https://picsum.photos/1920/1080", headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(bg_path, 'wb') as f:
                f.write(response.read())
        
        img = Image.open(bg_path).convert("RGBA")
        # For compatibility with older PIL versions use Image.LANCZOS instead of Image.Resampling.LANCZOS if needed
        img = img.resize((1920, 1080), getattr(Image, 'Resampling', Image).LANCZOS)
        
        # Apply a light greenish-white wash to fade the background
        overlay = Image.new("RGBA", img.size, (230, 245, 240, 210))
        tinted_img = Image.alpha_composite(img, overlay)
        tinted_img.save(tinted_bg_path)
    except Exception as e:
        print(f"Image fetch failed: {e}. Using fallback gradient.")
        # Fallback background if network fails
        img = Image.new("RGBA", (1920, 1080), (230, 245, 240, 255))
        draw = ImageDraw.Draw(img)
        for i in range(1080):
            alpha = int(255 * (i / 1080))
            draw.line([(0, i), (1920, i)], fill=(200, 220, 215, alpha))
        img.save(tinted_bg_path)

    # Insert background
    slide.shapes.add_picture(tinted_bg_path, 0, 0, width=W, height=H)

    # === Layer 2: Geometric Polygon Overlays ===
    # Left dark grey anchoring block
    grey_pts = [
        (0, 0),
        (int(W * 0.40), 0),
        (int(W * 0.15), H),
        (0, H),
        (0, 0)
    ]
    ff_grey = slide.shapes.build_freeform()
    ff_grey.add_line_segments(grey_pts)
    shape_grey = ff_grey.convert_to_shape()
    shape_grey.fill.solid()
    shape_grey.fill.fore_color.rgb = RGBColor(50, 55, 60)
    shape_grey.line.fill.background()

    # Teal diagonal separator stripe
    teal_pts = [
        (int(W * 0.40), 0),
        (int(W * 0.48), 0),
        (int(W * 0.23), H),
        (int(W * 0.15), H),
        (int(W * 0.40), 0)
    ]
    ff_teal = slide.shapes.build_freeform()
    ff_teal.add_line_segments(teal_pts)
    shape_teal = ff_teal.convert_to_shape()
    shape_teal.fill.solid()
    shape_teal.fill.fore_color.rgb = RGBColor(*accent_color)
    shape_teal.line.fill.background()

    # === Layer 3: Center Parallelogram Banner ===
    y_top = int(H * 0.55)
    y_bot = int(H * 0.72)
    
    # Shadow Banner (Accent Color)
    shadow_pts = [
        (int(Inches(1.5)), y_top),
        (int(W - Inches(1.5)), y_top),
        (int(W - Inches(3.5)), y_bot),
        (int(Inches(-0.5)), y_bot),
        (int(Inches(1.5)), y_top)
    ]
    ff_shadow = slide.shapes.build_freeform()
    ff_shadow.add_line_segments(shadow_pts)
    shape_shadow = ff_shadow.convert_to_shape()
    shape_shadow.fill.solid()
    shape_shadow.fill.fore_color.rgb = RGBColor(*accent_color)
    shape_shadow.line.fill.background()

    # Main Banner (White, offset up & left)
    offset = int(Inches(0.12))
    main_pts = [(x - offset, y - offset) for x, y in shadow_pts]
    ff_main = slide.shapes.build_freeform()
    ff_main.add_line_segments(main_pts)
    shape_main = ff_main.convert_to_shape()
    shape_main.fill.solid()
    shape_main.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape_main.line.fill.background()

    # === Layer 4: Title Text ===
    # Overlay standard textbox seamlessly inside the white banner area
    tx_box = slide.shapes.add_textbox(
        int(Inches(1.8)), y_top - offset, int(W * 0.6), y_bot - y_top
    )
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial'
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*accent_color)

    # Cleanup temp files
    if os.path.exists(bg_path): os.remove(bg_path)
    if os.path.exists(tinted_bg_path): os.remove(tinted_bg_path)

    prs.save(output_pptx_path)
    return output_pptx_path
