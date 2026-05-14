import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def _draw_trend_icon(slide, left, top, width, height, color_rgb, trend_up=True):
    """Draws a simple bar chart trend icon using rectangular shapes."""
    bar_count = 4
    bar_width = width / (bar_count * 2 - 1)  # Includes spacing
    max_height = height
    
    start_left = left
    for i in range(bar_count):
        if trend_up:
            bar_height = max_height * ((i + 1) / bar_count) * 0.8 + max_height * 0.2
        else:
            bar_height = max_height * ((bar_count - i) / bar_count) * 0.8 + max_height * 0.2
            
        bar_top = top + (max_height - bar_height)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, start_left, bar_top, bar_width, bar_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color_rgb)
        shape.line.fill.background()
        
        start_left += bar_width * 1.5

def _add_list_item(slide, left, top, number_str, title_text, body_text, color_rgb):
    """Adds a numbered list item with a circular icon."""
    icon_diameter = Inches(0.35)
    
    # Circle icon
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, icon_diameter, icon_diameter)
    icon.fill.solid()
    icon.fill.fore_color.rgb = RGBColor(*color_rgb)
    icon.line.fill.background()
    
    # Number inside circle
    text_box = slide.shapes.add_textbox(left, top, icon_diameter, icon_diameter)
    text_frame = text_box.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"{number_str}"
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0

    # Title text
    title_left = left + icon_diameter + Inches(0.2)
    title_width = Inches(5.0)
    title_box = slide.shapes.add_textbox(title_left, top, title_width, Inches(0.35))
    p_title = title_box.text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.bold = True
    p_title.font.size = Pt(16)
    p_title.font.color.rgb = RGBColor(51, 51, 51)
    title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Body text
    body_top = top + Inches(0.35)
    body_box = slide.shapes.add_textbox(title_left, body_top, title_width, Inches(0.6))
    p_body = body_box.text_frame.paragraphs[0]
    p_body.text = body_text
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = RGBColor(128, 128, 128)
    body_box.text_frame.word_wrap = True

def create_slide(
    output_pptx_path: str,
    slide_title: str = "Pros & Cons",
    slide_subtitle: str = "A balanced analysis of the key factors involved.",
    pros_title: str = "Pros",
    cons_title: str = "Cons",
    pros_data: list = None,
    cons_data: list = None,
    pros_color_rgb: tuple = (74, 105, 189),
    cons_color_rgb: tuple = (51, 51, 51),
    **kwargs
) -> str:
    """
    Create a PPTX file reproducing the Bifurcated Comparison Layout.

    Returns: path to the saved PPTX file.
    """
    
    # Default data if not provided
    if pros_data is None:
        pros_data = [
            {"title": "Benefit One", "body": "Studio is a fast way to start your responsive web design projects."},
            {"title": "Benefit Two", "body": "Harnesses the power of Sass and Compass for streamlined development."},
            {"title": "Benefit Three", "body": "Offers a wide range of pre-built components and templates."},
            {"title": "Benefit Four", "body": "Improves collaboration between designers and developers."},
        ]
    if cons_data is None:
        cons_data = [
            {"title": "Risk One", "body": "Studio is a fast way to start your responsive web design projects."},
            {"title": "Risk Two", "body": "May have a learning curve for beginners new to the ecosystem."},
            {"title": "Risk Three", "body": "Certain advanced customizations might require extra work."},
            {"title": "Risk Four", "body": "Dependency on the platform for updates and support."},
        ]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.75))
    p_title = title_box.text_frame.paragraphs[0]
    p_title.text = slide_title
    p_title.font.name = 'Calibri Light'
    p_title.font.size = Pt(36)

    # Slide Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12.333), Inches(0.5))
    p_subtitle = subtitle_box.text_frame.paragraphs[0]
    p_subtitle.text = slide_subtitle
    p_subtitle.font.size = Pt(14)
    p_subtitle.font.color.rgb = RGBColor(128, 128, 128)

    # --- Column Layout ---
    column_width = Inches(6.0)
    gutter = Inches(0.5)
    pros_left = (prs.slide_width - (2 * column_width + gutter)) / 2
    cons_left = pros_left + column_width + gutter
    
    column_top = Inches(1.5)
    column_height = Inches(5.5)
    header_height = Inches(0.8)
    content_top = column_top + header_height

    # --- Pros Column ---
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pros_left, content_top, column_width, column_height - header_height).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RGBColor(245, 245, 245)
    slide.shapes[-1].line.fill.background()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pros_left, column_top, column_width, header_height).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RGBColor(*pros_color_rgb)
    slide.shapes[-1].line.fill.background()
    
    pros_title_box = slide.shapes.add_textbox(pros_left + Inches(0.8), column_top, column_width - Inches(0.9), header_height)
    pros_title_box.text_frame.paragraphs[0].text = pros_title
    pros_title_box.text_frame.paragraphs[0].font.bold = True
    pros_title_box.text_frame.paragraphs[0].font.size = Pt(24)
    pros_title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    pros_title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _draw_trend_icon(slide, pros_left + Inches(0.2), column_top + Inches(0.2), Inches(0.5), Inches(0.4), (255, 255, 255), trend_up=True)

    item_top = content_top + Inches(0.4)
    item_left = pros_left + Inches(0.4)
    for i, item in enumerate(pros_data):
        _add_list_item(slide, item_left, item_top, f"{i+1:02d}", item['title'], item['body'], pros_color_rgb)
        item_top += Inches(1.2)

    # --- Cons Column ---
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cons_left, content_top, column_width, column_height - header_height).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RGBColor(245, 245, 245)
    slide.shapes[-1].line.fill.background()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cons_left, column_top, column_width, header_height).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RGBColor(*cons_color_rgb)
    slide.shapes[-1].line.fill.background()
    
    cons_title_box = slide.shapes.add_textbox(cons_left + Inches(0.8), column_top, column_width - Inches(0.9), header_height)
    cons_title_box.text_frame.paragraphs[0].text = cons_title
    cons_title_box.text_frame.paragraphs[0].font.bold = True
    cons_title_box.text_frame.paragraphs[0].font.size = Pt(24)
    cons_title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    cons_title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _draw_trend_icon(slide, cons_left + Inches(0.2), column_top + Inches(0.2), Inches(0.5), Inches(0.4), (255, 255, 255), trend_up=False)

    item_top = content_top + Inches(0.4)
    item_left = cons_left + Inches(0.4)
    for i, item in enumerate(cons_data):
        _add_list_item(slide, item_left, item_top, f"{i+1:02d}", item['title'], item['body'], cons_color_rgb)
        item_top += Inches(1.2)

    prs.save(output_pptx_path)
    return output_pptx_path

