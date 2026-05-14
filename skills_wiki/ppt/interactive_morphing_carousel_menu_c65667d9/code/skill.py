import os
import io
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageOps

def create_slide(
    output_pptx_path: str,
    **kwargs,
) -> str:
    """
    Creates a PPTX file reproducing the Interactive Morphing Carousel effect.

    **Post-generation steps are required in PowerPoint:**
    1.  Select all slides (Ctrl+A in the slide sorter view).
    2.  Go to the "Transitions" tab and select "Morph". Set a duration (e.g., 0.5s).
    3.  Follow the hyperlink instructions printed to the console after generation.

    Returns: path to the saved PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- Configuration ---
    image_keywords = ['food burger', 'coca-cola can', 'french fries', 'colorful macaron', 'pepperoni pizza', 'sprinkle donut', 'strawberry smoothie']
    item_names = ['BURGER', 'COKE', 'FRIES', 'MACARON', 'PIZZA', 'DONUT', 'SMOOTHIES']
    item_descriptions = {
        'BURGER': "Burgers can be served with various toppings and condiments, such as lettuce, tomatoes, onions, cheese, pickles, ketchup, mustard, mayonnaise, and more.",
        'COKE': "Coca-Cola is known for its distinctive taste, which combines sweet and tangy flavors, and it is one of the most recognized and consumed beverages in the world.",
        'FRIES': "They are typically cut into long, thin strips and then fried until they become crispy and golden brown. French fries are a common side dish or snack in many parts of the world.",
        'MACARON': "Macarons have a crisp exterior and a soft, chewy interior. They are often characterized by their vibrant colors and are typically filled with various flavored fillings.",
        'PIZZA': "A dish of Italian origin consisting of a usually round, flat base of leavened wheat-based dough topped with tomatoes, cheese, and often various other ingredients.",
        'DONUT': "It is made from dough that is often sweetened and enriched with ingredients like sugar, eggs, and butter. Donuts can come in a variety of flavors, sizes, and shapes.",
        'SMOOTHIES': "The quick brown fox jumps over the lazy dog. The quick brown fox jumps over the lazy dog. The quick brown fox jumps over the lazy dog."
    }

    # --- Image Fetching and Processing ---
    image_size = Inches(2.5)
    images_color = {}
    images_bw = {}

    print("Downloading and processing images...")
    for i, keyword in enumerate(image_keywords):
        try:
            url = f"https://source.unsplash.com/400x500/?{keyword}"
            with urllib.request.urlopen(url) as response:
                img_data = response.read()
                img = Image.open(io.BytesIO(img_data)).convert("RGB")
                
                # Store color version
                color_stream = io.BytesIO()
                img.save(color_stream, format='PNG')
                images_color[i] = color_stream
                
                # Store grayscale version
                bw_img = ImageOps.grayscale(img)
                bw_stream = io.BytesIO()
                bw_img.save(bw_stream, format='PNG')
                images_bw[i] = bw_stream
        except Exception as e:
            # Handle failed download if necessary
            print(f"Warning: Could not download image for '{keyword}'. Skipping. {e}")
            continue
    
    num_images = len(images_color)
    center_x = (prs.slide_width - image_size) / 2
    y_pos = (prs.slide_height - image_size) / 2

    # --- Function to add shared UI elements ---
    def add_ui_elements(slide, item_name):
        # Top mask
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(1.5), Inches(17.33), Inches(2.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.fill.background()
        # Bottom mask
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(3.8), Inches(17.33), Inches(2.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.fill.background()

        # Top Title
        tb = slide.shapes.add_textbox(Inches(5.16), Inches(0.5), Inches(3), Inches(0.5))
        tb.text_frame.paragraphs[0].text = "OUR MENU"
        tb.text_frame.paragraphs[0].font.name = 'Calibri Light'; tb.text_frame.paragraphs[0].font.size = Pt(16)
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        tb.line.color.rgb = RGBColor(0,0,0); tb.line.width = Pt(0.75)
        tb.fill.background()

        # Item Title
        tb = slide.shapes.add_textbox(Inches(4.66), Inches(5.5), Inches(4), Inches(0.5))
        tb.text_frame.paragraphs[0].text = item_name
        tb.text_frame.paragraphs[0].font.name = 'Calibri Light'; tb.text_frame.paragraphs[0].font.size = Pt(18)
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        tb.line.color.rgb = RGBColor(0,0,0); tb.line.width = Pt(0.75)
        tb.fill.background()

        # Navigation Arrows
        slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(3.9), Inches(5.6), Inches(0.5), Inches(0.3)).fill.background()
        slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.9), Inches(5.6), Inches(0.5), Inches(0.3)).fill.background()

    # --- Create Main Carousel and Detail Slides ---
    carousel_slide_indices = {}
    detail_slide_indices = {}

    for i in range(num_images):
        # Create Carousel Slide
        slide = prs.slides.add_slide(blank_layout)
        carousel_slide_indices[i] = len(prs.slides)
        
        start_pos_x = center_x - (i * image_size)
        for j in range(num_images):
            img_stream = images_color[j] if i == j else images_bw[j]
            img_stream.seek(0)
            slide.shapes.add_picture(img_stream, start_pos_x + (j * image_size), y_pos, height=image_size)
        
        add_ui_elements(slide, item_names[i])
    
    for i in range(num_images):
        # Create Detail Slide
        slide = prs.slides.add_slide(blank_layout)
        detail_slide_indices[i] = len(prs.slides)
        
        # Move other UI elements off-screen
        add_ui_elements(slide, item_names[i])
        for shape in slide.shapes:
             shape.left -= Inches(20)

        # Enlarged image
        img_width_detail = Inches(5.5)
        left_pos = Inches(0.5)
        top_pos = (prs.slide_height - img_width_detail) / 2
        images_color[i].seek(0)
        pic = slide.shapes.add_picture(images_color[i], left_pos, top_pos, height=img_width_detail)
        
        # Text
        title_box = slide.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6), Inches(1.2))
        p = title_box.text_frame.paragraphs[0]
        p.text = item_names[i].title()
        p.font.name = 'Garamond'; p.font.size = Pt(60)

        desc_box = slide.shapes.add_textbox(Inches(6.5), Inches(2.8), Inches(6.5), Inches(3))
        p = desc_box.text_frame.paragraphs[0]
        p.text = item_descriptions[item_names[i]]
        p.font.name = 'Calibri Light'; p.font.size = Pt(16)
        
        # Back Arrow for hyperlink
        slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(6.5), Inches(0.8), Inches(0.5), Inches(0.3)).fill.background()


    # --- Final Save and Instructions ---
    prs.save(output_pptx_path)
    
    print("\n" + "="*80)
    print(f"SUCCESS: Presentation saved to '{output_pptx_path}'")
    print("Please follow these MANUAL steps in PowerPoint to complete the effect:")
    print("1. Open the file. Select all slides (Ctrl+A).")
    print("2. Go to 'Transitions' tab -> Click 'Morph'.")
    print("\n3. Add the following HYPERLINKS:")
    for i in range(num_images):
        print(f"\n--- On Slide {carousel_slide_indices[i]} ({item_names[i]} Carousel) ---")
        if i > 0:
            print(f"  - Link the LEFT arrow to Slide {carousel_slide_indices[i-1]} (Previous Slide).")
        if i < num_images - 1:
            print(f"  - Link the RIGHT arrow to Slide {carousel_slide_indices[i+1]} (Next Slide).")
        print(f"  - Link the central IMAGE to Slide {detail_slide_indices[i]} ({item_names[i]} Detail).")

    for i in range(num_images):
        print(f"\n--- On Slide {detail_slide_indices[i]} ({item_names[i]} Detail) ---")
        print(f"  - Link the BACK arrow to Slide {carousel_slide_indices[i]} ({item_names[i]} Carousel).")
    print("="*80)

    return output_pptx_path

