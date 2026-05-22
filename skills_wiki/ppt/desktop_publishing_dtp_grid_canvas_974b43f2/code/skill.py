import os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def _create_placeholder_image(filepath, text, width_in, height_in, bg_color=(230, 230, 240)):
    """Helper to generate a clean mockup graphic using PIL"""
    dpi = 300
    w, h = int(width_in * dpi), int(height_in * dpi)
    img = Image.new('RGB', (w, h), color=bg_color)
    draw = ImageDraw.Draw(img)
    
    # Draw border
    draw.rectangle([(0, 0), (w-1, h-1)], outline=(150, 150, 160), width=4)
    
    # Draw diagonal lines for "placeholder" look
    draw.line([(0, 0), (w, h)], fill=(200, 200, 210), width=3)
    draw.line([(0, h), (w, 0)], fill=(200, 200, 210), width=3)
    
    # We won't load a custom font to avoid cross-platform font missing errors,
    # we'll just rely on the visual geometry.
    img.save(filepath)
    return filepath

def create_slide(
    output_pptx_path: str,
    title_text: str = "Chapter 4: Advanced Multimedia",
    body_text: str = "",
    bg_palette: str = "corporate",
    accent_color: tuple = (204, 34, 41),  # Red accent matching the book cover in video
    **kwargs,
) -> str:
    """
    Creates a two-page Desktop Publishing (DTP) spread in PowerPoint.
    Simulates a book layout with exact grid constraints, headers, and columns.
    """
    prs = Presentation()
    
    # Set to 11x8.5 (US Letter Landscape) to simulate a two-page spread (5.5 x 8.5 per page)
    prs.slide_width = Inches(11.0)
    prs.slide_height = Inches(8.5)
    
    # Use blank layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Standard dummy text for dense layout testing
    lorem_ipsum = (
        "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod "
        "tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, "
        "quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. "
        "Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore "
        "eu fugiat nulla pariatur. Excepteur sint occaecat cupidatat non proident."
    )
    if not body_text:
        body_text = lorem_ipsum
        
    # ==========================================
    # GRID MATH & CONSTANTS
    # ==========================================
    margin = Inches(0.5)
    page_width = Inches(5.5)
    content_width = Inches(4.5)  # 5.5 - (0.5 * 2)
    col_width = Inches(2.15)     # (4.5 - 0.2) / 2
    gutter = Inches(0.2)
    
    # Colors
    accent = RGBColor(*accent_color)
    gray_header = RGBColor(235, 235, 235)
    text_dark = RGBColor(40, 40, 40)
    
    # ==========================================
    # BACKGROUND / SPINE
    # ==========================================
    # Draw spine divider (subtle gray line down the middle)
    spine = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        page_width, Inches(0.2), Inches(0.02), Inches(8.1)
    )
    spine.fill.solid()
    spine.fill.fore_color.rgb = RGBColor(200, 200, 200)
    spine.line.fill.background()

    # ==========================================
    # LEFT PAGE COMPOSITION
    # ==========================================
    # 1. Header Band
    left_header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, margin, Y=margin, width=content_width, height=Inches(0.4))
    left_header.fill.solid()
    left_header.fill.fore_color.rgb = gray_header
    left_header.line.fill.background()
    
    tf = left_header.text_frame
    tf.text = "CLICK TO EDIT MASTER TEXT STYLES"
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = text_dark
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    # 2. Title
    title_box = slide.shapes.add_textbox(margin, Inches(1.1), content_width, Inches(0.8))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = accent
    
    # 3. Two Columns of Text
    col1 = slide.shapes.add_textbox(margin, Inches(2.0), col_width, Inches(3.0))
    col1.text_frame.word_wrap = True
    p = col1.text_frame.paragraphs[0]
    p.text = body_text
    p.font.size = Pt(11)
    p.font.color.rgb = text_dark
    
    col2 = slide.shapes.add_textbox(margin + col_width + gutter, Inches(2.0), col_width, Inches(3.0))
    col2.text_frame.word_wrap = True
    p = col2.text_frame.paragraphs[0]
    p.text = body_text
    p.font.size = Pt(11)
    p.font.color.rgb = text_dark
    
    # 4. Large Image Figure at the bottom (Spanning both columns)
    img_path = "temp_placeholder.png"
    _create_placeholder_image(img_path, "Figure 1", 4.5, 2.5)
    slide.shapes.add_picture(img_path, margin, Inches(5.2), width=content_width, height=Inches(2.5))
    
    # Figure Caption
    cap_box = slide.shapes.add_textbox(margin, Inches(7.7), content_width, Inches(0.3))
    p = cap_box.text_frame.paragraphs[0]
    p.text = "Figure 1.1: Example of an inserted graphic bridging the grid."
    p.font.size = Pt(9)
    p.font.italic = True
    p.font.color.rgb = RGBColor(100, 100, 100)

    # ==========================================
    # RIGHT PAGE COMPOSITION
    # ==========================================
    right_x_offset = page_width + margin
    
    # 1. Header Band
    right_header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x_offset, Y=margin, width=content_width, height=Inches(0.4))
    right_header.fill.solid()
    right_header.fill.fore_color.rgb = gray_header
    right_header.line.fill.background()
    tf = right_header.text_frame
    tf.text = "SECTION 2: WORKFLOW LOGIC"
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = text_dark
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    # 2. Right Page Top Image (Small, aligned left column)
    _create_placeholder_image(img_path, "UI Screen", 2.15, 2.0)
    slide.shapes.add_picture(img_path, right_x_offset, Inches(1.1), width=col_width, height=Inches(2.0))
    
    # 3. Right Page Text (Wrapping around the image conceptually)
    # Text next to image
    col3 = slide.shapes.add_textbox(right_x_offset + col_width + gutter, Inches(1.0), col_width, Inches(2.2))
    col3.text_frame.word_wrap = True
    p = col3.text_frame.paragraphs[0]
    p.text = body_text[:200] + "..."
    p.font.size = Pt(11)
    p.font.color.rgb = text_dark
    
    # Text below image (Full width)
    wide_text = slide.shapes.add_textbox(right_x_offset, Inches(3.3), content_width, Inches(1.5))
    wide_text.text_frame.word_wrap = True
    p = wide_text.text_frame.paragraphs[0]
    p.text = body_text + " " + body_text[:100]
    p.font.size = Pt(11)
    p.font.color.rgb = text_dark
    
    # 4. Callout/Tip Box
    tip_box_y = Inches(5.2)
    tip_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x_offset, tip_box_y, content_width, Inches(1.8))
    tip_box.fill.solid()
    tip_box.fill.fore_color.rgb = RGBColor(245, 248, 255) # Light blue
    tip_box.line.color.rgb = accent
    tip_box.line.width = Pt(1.5)
    
    tf = tip_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)
    
    p1 = tf.paragraphs[0]
    p1.text = "PRO TIP: PPT AS A DTP TOOL"
    p1.font.bold = True
    p1.font.size = Pt(12)
    p1.font.color.rgb = accent
    
    p2 = tf.add_paragraph()
    p2.text = (
        "PowerPoint's absolute positioning engine allows you to place text boxes and graphics exactly "
        "where you need them without dealing with Microsoft Word's flow disruptions. It's excellent "
        "for visually dense instruction manuals."
    )
    p2.font.size = Pt(11)
    p2.font.color.rgb = text_dark
    
    # Cleanup temporary image
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
