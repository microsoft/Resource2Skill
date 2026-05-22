def create_slide(
    output_pptx_path: str,
    title_text: str = "Cash Flow Strategy Sessions",
    subtitle_text: str = "By Decisions Plus Strategic",
    tagline_text: str = "Connect with us",
    cta_email: str = "hello@decisionsplusstrategic.com",
    badge_text: str = "B1G1",
    body_text: str = "Each time we hold a Cash Flow Strategy Session, we give our customers an opportunity to choose our giving project.",
    accent_color_rgb: tuple = (65, 168, 95),  # Corporate Green
) -> str:
    """
    Creates a PPTX file reproducing the 'Corporate Impact Split-Layout'.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank layout
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Convert tuple to RGBColor object
    accent_color = RGBColor(*accent_color_rgb)
    dark_gray = RGBColor(80, 80, 80)
    white = RGBColor(255, 255, 255)

    # === 1. Header Band ===
    header_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        0, 0, Inches(13.333), Inches(1.4)
    )
    header_rect.fill.solid()
    header_rect.fill.fore_color.rgb = RGBColor(245, 245, 245)
    header_rect.line.fill.background() # No border

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(6), Inches(0.6))
    tf_title = title_box.text_frame
    p_title = tf_title.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(36)
    p_title.font.color.rgb = accent_color
    p_title.font.name = "Arial"

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.8), Inches(6), Inches(0.4))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.add_paragraph()
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = dark_gray
    p_sub.font.name = "Arial"

    # Tagline (Right aligned)
    tagline_box = slide.shapes.add_textbox(Inches(8), Inches(0.4), Inches(4.8), Inches(0.6))
    tf_tag = tagline_box.text_frame
    p_tag = tf_tag.add_paragraph()
    p_tag.text = tagline_text
    p_tag.alignment = PP_ALIGN.RIGHT
    p_tag.font.size = Pt(28)
    p_tag.font.color.rgb = dark_gray
    p_tag.font.name = "Arial"
    p_tag.font.italic = True

    # === 2. Call to Action (Center) ===
    # Small @ badge
    at_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        Inches(3.8), Inches(2.3), Inches(0.6), Inches(0.6)
    )
    at_circle.fill.background() # Hollow
    at_circle.line.color.rgb = dark_gray
    at_circle.line.width = Pt(2)
    tf_at = at_circle.text_frame
    tf_at.text = "@"
    tf_at.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf_at.paragraphs[0].font.size = Pt(24)
    tf_at.paragraphs[0].font.color.rgb = dark_gray

    # Email Text
    email_box = slide.shapes.add_textbox(Inches(4.5), Inches(2.2), Inches(6), Inches(0.8))
    tf_email = email_box.text_frame
    p_email = tf_email.add_paragraph()
    p_email.text = cta_email
    p_email.font.size = Pt(28)
    p_email.font.color.rgb = accent_color
    p_email.font.name = "Arial"
    p_email.font.underline = True

    # Subtle horizontal divider line
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(2), Inches(3.2), Inches(9.333), Inches(0.02)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(220, 220, 220)
    divider.line.fill.background()

    # === 3. Lower Split Section ===
    # Left: Big Impact Circle
    big_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, 
        Inches(1.5), Inches(3.8), Inches(2.8), Inches(2.8)
    )
    big_circle.fill.solid()
    big_circle.fill.fore_color.rgb = accent_color
    big_circle.line.fill.background()
    
    tf_badge = big_circle.text_frame
    tf_badge.vertical_anchor = 3 # Middle alignment
    p_badge = tf_badge.paragraphs[0]
    p_badge.text = badge_text
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.font.size = Pt(54)
    p_badge.font.color.rgb = white
    p_badge.font.bold = True
    p_badge.font.name = "Arial"

    # Right: Body Text Paragraph
    body_box = slide.shapes.add_textbox(Inches(5.0), Inches(4.0), Inches(7.5), Inches(2.5))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    p_body = tf_body.add_paragraph()
    p_body.text = body_text
    p_body.font.size = Pt(22)
    p_body.font.color.rgb = dark_gray
    p_body.font.name = "Arial"
    
    # Adjust line spacing for readability
    p_body.line_spacing = 1.2 

    # Save presentation
    prs.save(output_pptx_path)
    return output_pptx_path
