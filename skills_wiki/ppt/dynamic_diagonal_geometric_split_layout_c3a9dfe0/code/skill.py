def create_slide(
    output_pptx_path: str,
    title_text: str = "TITLE OF YOUR SECTION",
    subtitle_text: str = "GEOMETRIC GRAPHIC DESIGN",
    body_text: str = "Welcome to My YouTube Channel dedicated to sharing PowerPoint skills and Tutorials! Whether you're new to PowerPoint or a Seasoned Pro, our Channel has something for you.",
    bg_theme: str = "mountains,nature",
    main_color: tuple = (218, 73, 45),   # Red-Orange
    accent_color: tuple = (242, 166, 138) # Light Peach
) -> str:
    """
    Creates a PPTX file reproducing the Dynamic Diagonal Geometric Split Layout.
    """
    import os
    import urllib.request
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    # Initialize presentation (16:9 aspect ratio)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout

    # Colors
    c_main = RGBColor(*main_color)
    c_accent = RGBColor(*accent_color)
    c_white = RGBColor(255, 255, 255)
    c_light_gray = RGBColor(245, 245, 245)

    # Helper function to convert inches to EMUs for the FreeformBuilder
    def emu(inches_val):
        return int(inches_val * 914400)

    # Helper function to draw a custom polygon
    def add_polygon(slide, vertices_in_inches, fill_color):
        start_x = emu(vertices_in_inches[0][0])
        start_y = emu(vertices_in_inches[0][1])
        
        ff_builder = slide.shapes.build_freeform(start_x, start_y)
        
        # Add remaining points
        for x, y in vertices_in_inches[1:]:
            ff_builder.add_line_segments([(emu(x), emu(y))], close=False)
        
        # Close the shape
        ff_builder.add_line_segments([(start_x, start_y)], close=True)
        
        shape = ff_builder.convert_to_shape()
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        # Remove outline
        shape.line.fill.background()
        return shape

    # === Layer 1: Background Image (Right Side) ===
    # We place a standard rectangle image, which will be diagonally masked by our polygons
    img_path = "temp_bg.jpg"
    try:
        url = f"https://source.unsplash.com/1280x720/?{bg_theme}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(img_path, 'wb') as f:
                f.write(response.read())
        # Place image on the right portion of the slide
        slide.shapes.add_picture(img_path, Inches(4), Inches(0), Inches(9.333), Inches(7.5))
    except Exception as e:
        print(f"Image download failed, using solid fallback: {e}")
        fallback = slide.shapes.add_shape(1, Inches(4), Inches(0), Inches(9.333), Inches(7.5))
        fallback.fill.solid()
        fallback.fill.fore_color.rgb = RGBColor(30, 144, 255) # Blue fallback
        fallback.line.fill.background()

    # === Layer 2: Main Diagonal Polygons ===
    
    # 1. Main Left Block (Red-Orange)
    # Creates the large area on the left with a forward-leaning diagonal right edge
    left_block_vertices = [
        (0.0, 0.0),    # Top Left
        (6.5, 0.0),    # Top Middle-Right
        (8.5, 7.5),    # Bottom Middle-Right
        (0.0, 7.5)     # Bottom Left
    ]
    add_polygon(slide, left_block_vertices, c_main)

    # 2. Diagonal Accent Strip (Light Peach)
    # A parallel band laying exactly over the edge of the red block
    strip_vertices = [
        (6.5, 0.0),    # Touches top right of red block
        (7.5, 0.0),    # 1 inch wide at the top
        (9.5, 7.5),    # 1 inch wide at the bottom
        (8.5, 7.5)     # Touches bottom right of red block
    ]
    add_polygon(slide, strip_vertices, c_accent)

    # 3. Bottom Title/Accent Block (Light Gray)
    # An angled box at the bottom right/center for secondary text positioning
    bottom_block_vertices = [
        (3.5, 7.5),    # Bottom left (anchored in red area)
        (5.5, 5.0),    # Slants up and right
        (13.333, 5.0), # Straight to right edge
        (13.333, 7.5)  # Down to bottom right corner
    ]
    add_polygon(slide, bottom_block_vertices, c_light_gray)

    # === Layer 3: Typography & Content ===

    # Main Title on Red Block
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(5.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.name = "Arial Black"
    p.font.color.rgb = c_white

    # Body Text underneath Title
    body_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(4.5), Inches(1.5))
    btf = body_box.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.text = body_text
    bp.font.size = Pt(12)
    bp.font.name = "Arial"
    bp.font.color.rgb = c_white

    # Bottom Banner Text (Title)
    banner_title_box = slide.shapes.add_textbox(Inches(5.6), Inches(5.5), Inches(7.0), Inches(1.0))
    bt_tf = banner_title_box.text_frame
    bt_p = bt_tf.paragraphs[0]
    bt_p.text = "Engineer Haroon Mentor"
    bt_p.font.bold = True
    bt_p.font.size = Pt(36)
    bt_p.font.name = "Arial Black"
    bt_p.font.color.rgb = c_main
    bt_p.alignment = PP_ALIGN.CENTER

    # Bottom Banner Subtitle (Spaced out)
    banner_sub_box = slide.shapes.add_textbox(Inches(5.6), Inches(6.3), Inches(7.0), Inches(0.5))
    bs_tf = banner_sub_box.text_frame
    bs_p = bs_tf.paragraphs[0]
    # Creating a faux tracked-out text effect by adding spaces
    tracked_subtitle = "  ".join(subtitle_text.upper()) 
    bs_p.text = tracked_subtitle
    bs_p.font.bold = True
    bs_p.font.size = Pt(12)
    bs_p.font.name = "Arial"
    bs_p.font.color.rgb = RGBColor(150, 150, 150) # Gray
    bs_p.alignment = PP_ALIGN.CENTER

    # Cleanup temp image
    if os.path.exists(img_path):
        os.remove(img_path)

    prs.save(output_pptx_path)
    return output_pptx_path
