import os
import urllib.request
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml

def create_slide(
    output_pptx_path: str,
    title_text: str = "THANK YOU",
    body_text: str = "We recognize the importance of a business roadmap. This is why our designers have made this slide as visually attractive as possible. You can use this to outline your plans, predict market growth, marketing strategy, and more.",
    bg_palette: str = "business,workspace",
    accent_color: tuple = (255, 192, 0),  # Gold/Yellow
    **kwargs,
) -> str:
    """
    Create a PPTX file reproducing the 'Interlocking Geometric Motif Overlay' visual effect.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # === Layer 1: Background Image ===
    img_path = "bg_image_temp.jpg"
    try:
        # Attempt to fetch a contextual background image
        url = f"https://source.unsplash.com/1600x900/?{bg_palette.replace(' ', ',')}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response, open(img_path, 'wb') as out_file:
            out_file.write(response.read())
    except Exception:
        # Fallback: Generate a sleek dark gradient background using PIL if network fails
        bg = Image.new('RGB', (1600, 900))
        draw = ImageDraw.Draw(bg)
        for y in range(900):
            r, g, b = int(15 + (20 * y/900)), int(20 + (25 * y/900)), int(30 + (35 * y/900))
            draw.line([(0, y), (1600, y)], fill=(r, g, b))
        bg.save(img_path)
        
    slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)

    # === Layer 2: Semi-Transparent Content Overlay (PIL) ===
    overlay_height = prs.slide_height * 0.35
    overlay_top = prs.slide_height - overlay_height
    
    overlay_img_path = "overlay_temp.png"
    # Create an 82% opacity white rectangle (alpha=210)
    overlay = Image.new('RGBA', (100, 100), (255, 255, 255, 210))
    overlay.save(overlay_img_path)
    slide.shapes.add_picture(overlay_img_path, 0, overlay_top, prs.slide_width, overlay_height)

    # === Layer 3: Body Text ===
    tb = slide.shapes.add_textbox(Inches(2), overlay_top + Inches(0.4), prs.slide_width - Inches(4), overlay_height - Inches(0.8))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.add_paragraph()
    p.text = body_text
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(60, 60, 60)
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.3

    # === Layer 4: Interlocking Geometric Motif ===
    def apply_no_fill(shape):
        """Helper to inject XML <a:noFill/> into a shape for true transparency."""
        spPr = shape.element.spPr
        for child in list(spPr):
            if child.tag.endswith('Fill'):
                spPr.remove(child)
        noFill = parse_xml('<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
        spPr.insert(1, noFill)

    # Configuration for geometry
    cx = prs.slide_width / 2
    cy = overlay_top / 2  # Vertically centered in the space *above* the text panel
    center_d_size = Inches(2.2)
    small_d_size = Inches(1.3)
    offset = Inches(1.2)  # X-axis offset for flanking diamonds
    
    accent_rgb = RGBColor(*accent_color)
    line_weight = Pt(4.5)

    # 4a. Connectors (Lines)
    left_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, cy, cx - offset - small_d_size/2, cy)
    left_line.line.color.rgb = accent_rgb
    left_line.line.width = line_weight

    right_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx + offset + small_d_size/2, cy, prs.slide_width, cy)
    right_line.line.color.rgb = accent_rgb
    right_line.line.width = line_weight

    # 4b. Small Flanking Diamonds
    left_dia = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, cx - offset - small_d_size/2, cy - small_d_size/2, small_d_size, small_d_size)
    apply_no_fill(left_dia)
    left_dia.line.color.rgb = accent_rgb
    left_dia.line.width = line_weight

    right_dia = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, cx + offset - small_d_size/2, cy - small_d_size/2, small_d_size, small_d_size)
    apply_no_fill(right_dia)
    right_dia.line.color.rgb = accent_rgb
    right_dia.line.width = line_weight

    # 4c. Main Center Diamond (Drawn last so its stroke overlays the smaller ones)
    center_dia = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, cx - center_d_size/2, cy - center_d_size/2, center_d_size, center_d_size)
    apply_no_fill(center_dia)
    center_dia.line.color.rgb = accent_rgb
    center_dia.line.width = line_weight

    # === Layer 5: Hero Title Typography ===
    title_width = Inches(4)
    title_height = Inches(1)
    tb_title = slide.shapes.add_textbox(cx - title_width/2, cy - title_height/2, title_width, title_height)
    tb_title.text_frame.word_wrap = False
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.alignment = PP_ALIGN.CENTER
    
    font = p_title.font
    font.bold = True
    font.size = Pt(28)
    font.color.rgb = RGBColor(255, 255, 255)
    
    # Inject soft drop shadow to ensure title readability over varied images
    rPr = p_title.runs[0]._r.get_or_add_rPr()
    effectLst = parse_xml(
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:outerShdw blurRad="40000" dist="20000" dir="2700000" algn="tl">'
        '<a:srgbClr val="000000"><a:alpha val="60000"/></a:srgbClr>'
        '</a:outerShdw>'
        '</a:effectLst>'
    )
    rPr.append(effectLst)

    # Save and clean up
    prs.save(output_pptx_path)
    
    if os.path.exists(img_path): os.remove(img_path)
    if os.path.exists(overlay_img_path): os.remove(overlay_img_path)
    
    return output_pptx_path
