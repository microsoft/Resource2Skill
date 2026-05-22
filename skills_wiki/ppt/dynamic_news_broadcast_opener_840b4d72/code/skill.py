import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw

def create_slide(
    output_pptx_path: str,
    logo_text: str = "SET iNEWS",
    labels: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX slide reproducing the 'Dynamic News Broadcast Opener' style.

    Args:
        output_pptx_path: Path to save the generated .pptx file.
        logo_text: The main text for the central logo.
        labels: A list of strings for the info labels (e.g., ["Travel", "International", "Finance"]).

    Returns:
        Path to the saved PPTX file.
    """
    if labels is None:
        labels = ["Travel", "International", "Finance"]

    # --- Slide Setup ---
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # --- Color Palette ---
    BG_BLUE_DARK = (1, 82, 204)
    BG_BLUE_LIGHT = (0, 140, 255)
    ACCENT_YELLOW = (255, 220, 0)
    ACCENT_CYAN = (0, 204, 238)
    TEXT_WHITE = (255, 255, 255)
    TEXT_GREY = (200, 200, 200)

    # === Layer 1: Generate Background with PIL ===
    img_width, img_height = 1920, 1080
    im = Image.new('RGB', (img_width, img_height), BG_BLUE_DARK)
    draw = ImageDraw.Draw(im)

    # Draw light blue geometric panels
    # Top-left triangle
    draw.polygon([(0, 0), (img_width * 0.4, 0), (0, img_height * 0.7)], fill=BG_BLUE_LIGHT)
    # Bottom-right shape
    draw.polygon([(img_width, img_height), (img_width, img_height * 0.3), (img_width * 0.7, img_height)], fill=BG_BLUE_LIGHT)

    # Draw diagonal accent lines
    draw.line([(0, img_height * 0.1), (img_width * 0.9, img_height)], fill=ACCENT_YELLOW, width=25)
    draw.line([(img_width, img_height * 0.2), (img_width * 0.1, img_height)], fill=ACCENT_CYAN, width=25)
    
    # Draw dot patterns
    def draw_dot_pattern(x_start, y_start, rows, cols, dot_size=4, spacing=20):
        for r in range(rows):
            for c in range(cols):
                x = x_start + c * spacing
                y = y_start + r * spacing
                draw.ellipse([(x, y), (x + dot_size, y + dot_size)], fill=TEXT_WHITE)

    draw_dot_pattern(img_width * 0.85, img_height * 0.05, 5, 10) # Top-right
    draw_dot_pattern(img_width * 0.05, img_height * 0.4, 10, 5)  # Left-middle
    
    background_path = "news_opener_background.png"
    im.save(background_path)

    # Add background image to slide
    slide.shapes.add_picture(background_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # === Layer 2: Text and UI Elements ===
    
    # Vertical side text
    def add_vertical_text(text, left_inch, top_inch, color):
        tx_box = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(1), Inches(4))
        tf = tx_box.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = '\n'.join(list(text))
        p.font.name = 'Arial Black'
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(*color)
        p.alignment = PP_ALIGN.CENTER
        
    add_vertical_text("LIVE", 0.3, 1, TEXT_GREY)
    add_vertical_text("STREAM", 15, 2.5, TEXT_GREY)

    # Main Logo
    tx_box = slide.shapes.add_textbox(Inches(4), Inches(3.5), Inches(8), Inches(2))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = logo_text
    p.font.name = 'Arial Black'
    p.font.size = Pt(80)
    p.font.color.rgb = RGBColor(*TEXT_WHITE)
    p.alignment = PP_ALIGN.CENTER

    # Info Labels and UI Elements
    start_top = 2.0
    for i, label_text in enumerate(labels):
        top = Inches(start_top + i * 1.5)
        # Label text
        tx_box = slide.shapes.add_textbox(Inches(12), top, Inches(3), Inches(0.5))
        tf = tx_box.text_frame
        tf.paragraphs[0].text = f"× {label_text}"
        tf.paragraphs[0].font.name = 'Arial'
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.color.rgb = RGBColor(*TEXT_WHITE)
        
        # Associated UI element
        if i == 0: # Chat icon
            slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(14.5), top, Inches(0.8), Inches(0.5))
            slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(14.7), top + Inches(0.4), Inches(0.4), Inches(0.25))
        if i == 1: # Progress bar
            bar_back = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.1), top + Inches(0.5), Inches(2), Inches(0.1))
            bar_back.fill.solid()
            bar_back.fill.fore_color.rgb = RGBColor(100, 100, 100)
            bar_back.line.fill.background()
            
            bar_front = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.1), top + Inches(0.5), Inches(1.5), Inches(0.1))
            bar_front.fill.solid()
            bar_front.fill.fore_color.rgb = RGBColor(*ACCENT_YELLOW)
            bar_front.line.fill.background()

    # --- Save and Cleanup ---
    prs.save(output_pptx_path)
    if os.path.exists(background_path):
        os.remove(background_path)
    
    return output_pptx_path

# # Example Usage:
# if __name__ == '__main__':
#     file_path = "Dynamic_News_Opener.pptx"
#     create_slide(
#         output_pptx_path=file_path,
#         logo_text="SET iNEWS",
#         labels=["Travel", "International", "Finance"]
#     )
#     print(f"Presentation saved to {file_path}")

