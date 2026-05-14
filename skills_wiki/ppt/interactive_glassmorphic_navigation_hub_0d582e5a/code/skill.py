import io
import math
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.action import PP_ACTION


def create_interactive_glassmorphism_hub(
    output_pptx_path: str,
    title_text: str = "How to Add Hyperlink to a Slide",
    topics: list = None,
    **kwargs,
) -> str:
    """
    Creates a PPTX file with an interactive navigation hub using a glassmorphism style.

    Args:
        output_pptx_path: The path to save the generated .pptx file.
        title_text: The main title displayed on the hub slide.
        topics: A list of topic names for the navigation buttons and slide titles.

    Returns:
        The path to the saved PPTX file.
    """
    if topics is None:
        topics = ["Introduction", "Goals", "Topic", "Examples", "Analysis"]

    # --- Presentation Setup ---
    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 16:9 aspect ratio
    prs.slide_height = Emu(6858000)
    
    # --- Helper Functions ---
    def create_gradient_background(width, height):
        """Generates a vibrant mesh gradient image using PIL."""
        img = Image.new("RGB", (width, height), "#FFFFFF")
        draw = ImageDraw.Draw(img)
        
        # Define gradient colors (R, G, B)
        colors = {
            (0, 0): (43, 88, 118),
            (width, 0): (78, 29, 68),
            (0, height): (3, 111, 113),
            (width, height): (170, 75, 53)
        }
        
        for y in range(height):
            for x in range(width):
                # Bilinear interpolation for a smooth gradient
                dx = x / width
                dy = y / height
                
                c00 = colors[(0, 0)]
                c10 = colors[(width, 0)]
                c01 = colors[(0, height)]
                c11 = colors[(width, height)]
                
                r = (c00[0] * (1 - dx) * (1 - dy) + c10[0] * dx * (1 - dy) +
                     c01[0] * (1 - dx) * dy + c11[0] * dx * dy)
                g = (c00[1] * (1 - dx) * (1 - dy) + c10[1] * dx * (1 - dy) +
                     c01[1] * (1 - dx) * dy + c11[1] * dx * dy)
                b = (c00[2] * (1 - dx) * (1 - dy) + c10[2] * dx * (1 - dy) +
                     c01[2] * (1 - dx) * dy + c11[2] * dx * dy)
                
                draw.point((x, y), fill=(int(r), int(g), int(b)))
        
        return img

    def add_zoom_transition(slide, direction="in"):
        """Injects XML for a zoom transition."""
        slide_xml = slide.element
        transition_tag = "{http://schemas.openxmlformats.org/presentationml/2006/main}transition"
        
        # Remove existing transition if any
        for el in slide_xml.findall(transition_tag):
            slide_xml.remove(el)

        # Create new transition element
        transition_node = etree.Element(transition_tag)
        zoom_node = etree.SubElement(transition_node, "{http://schemas.openxmlformats.org/presentationml/2006/main}zoom")
        
        if direction == "out":
            zoom_node.set("transition", "out")
            
        slide_xml.insert(0, transition_node)

    # Generate background image
    bg_image = create_gradient_background(int(prs.slide_width / Emu(9600)), int(prs.slide_height / Emu(9600)))
    bg_image_bytes = io.BytesIO()
    bg_image.save(bg_image_bytes, format='PNG')
    bg_image_bytes.seek(0)
    
    # --- Slide 1: Navigation Hub ---
    hub_slide = prs.slides.add_slide(prs.slide_layouts[6])
    hub_slide.shapes.add_picture(io.BytesIO(bg_image_bytes.getvalue()), 0, 0, width=prs.slide_width, height=prs.slide_height)
    add_zoom_transition(hub_slide, direction="out")

    title_shape = hub_slide.shapes.add_textbox(Inches(9.5), Inches(0.5), Inches(3.5), Inches(0.5))
    title_p = title_shape.text_frame.paragraphs[0]
    title_p.text = title_text
    title_p.font.name = 'Segoe UI'
    title_p.font.size = Pt(14)
    title_p.font.color.rgb = RGBColor(255, 255, 255)
    title_shape.line.fill.background()
    
    # Navigation Buttons
    num_topics = len(topics)
    total_width = num_topics * 1.8 + (num_topics - 1) * 0.2
    start_left = (13.333 - total_width) / 2
    
    content_slides = []
    
    # First, create all the content slides so we can link to them
    for i, topic in enumerate(topics):
        content_slide = prs.slides.add_slide(prs.slide_layouts[6])
        content_slide.shapes.add_picture(io.BytesIO(bg_image_bytes.getvalue()), 0, 0, width=prs.slide_width, height=prs.slide_height)
        add_zoom_transition(content_slide, direction="in")
        
        # Glassmorphism panel
        panel_left, panel_top, panel_width, panel_height = Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1)
        
        # 1. Blurred background crop
        crop_box = (
            int(panel_left / Inches(1) * 96), int(panel_top / Inches(1) * 96),
            int((panel_left + panel_width) / Inches(1) * 96), int((panel_top + panel_height) / Inches(1) * 96)
        )
        bg_crop = bg_image.crop(crop_box).filter(ImageFilter.GaussianBlur(15))
        crop_bytes = io.BytesIO()
        bg_crop.save(crop_bytes, format='PNG')
        content_slide.shapes.add_picture(crop_bytes, panel_left, panel_top, width=panel_width, height=panel_height)
        
        # 2. Semi-transparent overlay
        panel = content_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, panel_left, panel_top, panel_width, panel_height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        panel.fill.transparency = 0.85
        panel.line.fill.background()
        
        # Content text
        title_box = content_slide.shapes.add_textbox(Inches(2), Inches(1.8), Inches(8), Inches(1))
        title_box.text_frame.text = topic
        title_box.text_frame.paragraphs[0].font.name = 'Segoe UI Bold'
        title_box.text_frame.paragraphs[0].font.size = Pt(44)
        title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

        body_box = content_slide.shapes.add_textbox(Inches(2), Inches(2.8), Inches(8), Inches(3))
        body_box.text_frame.text = "The quick brown fox jumps over the lazy dog. " * 3
        body_box.text_frame.paragraphs[0].font.name = 'Segoe UI'
        body_box.text_frame.paragraphs[0].font.size = Pt(18)
        body_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        content_slides.append(content_slide)

    # Now, add hub buttons and link them
    for i, topic in enumerate(topics):
        left = Inches(start_left + i * 2.0)
        top = Inches(2.5)
        width = height = Inches(1.8)
        
        button_bg = hub_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        button_bg.fill.solid()
        button_bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
        button_bg.line.fill.background()
        
        # Link the button shape to the content slide
        hlink = button_bg.click_action.hyperlink
        hlink.action = PP_ACTION.HYPERLINK_TO_SLIDE
        hlink.target_slide = content_slides[i]

        label_box = hub_slide.shapes.add_textbox(left, top + height - Inches(0.2), width, Inches(0.5))
        label_p = label_box.text_frame.paragraphs[0]
        label_p.text = topic
        label_p.font.name = 'Segoe UI Bold'
        label_p.font.size = Pt(16)
        label_p.font.color.rgb = RGBColor(255, 255, 255)
        
    # Add home buttons to all content slides
    for slide in content_slides:
        home_button = slide.shapes.add_shape(MSO_SHAPE.HOME, Inches(0.3), Inches(0.3), Inches(0.5), Inches(0.5))
        home_button.fill.solid()
        home_button.fill.fore_color.rgb = RGBColor(255, 255, 255)
        home_button.line.fill.background()
        
        # Link home button back to the hub slide
        hlink = home_button.click_action.hyperlink
        hlink.action = PP_ACTION.HYPERLINK_TO_SLIDE
        hlink.target_slide = hub_slide

    prs.save(output_pptx_path)
    return output_pptx_path

# Example Usage:
# create_interactive_glassmorphism_hub("interactive_presentation.pptx")

