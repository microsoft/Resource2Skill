import openpyxl
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.label import DataLabelList

def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    # Default palette fallback
    theme_colors = {
        "female_accent": "70AD47", # Green
        "male_accent": "5B9BD5",   # Blue
        "bg_track": "E7E6E6"       # Light gray
    }
    
    # Set up data
    start_row = ws.max_row + 2 if ws.max_row > 1 else 1
    
    data = [
        ["Category", "Full", "Filled"],
        ["Female", 1.0, 0.75],
        ["Male", 1.0, 0.25]
    ]
    
    for r_idx, row in enumerate(data, start=start_row):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if r_idx > start_row and c_idx > 1:
                cell.number_format = '0%'

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Demographic Split"
    
    # Lock axes and clean up gridlines for an infographic look
    chart.y_axis.scaling.max = 1.0
    chart.y_axis.scaling.min = 0.0
    chart.y_axis.majorGridlines = None
    chart.y_axis.delete = True 
    
    # Overlap 100% to stack the Filled series directly on top of the Full series
    chart.overlap = 100
    chart.gapWidth = 50

    cats = Reference(ws, min_col=1, min_row=start_row+1, max_row=start_row+2)
    
    # Series 1: Background track (added first to render in back)
    data_full = Reference(ws, min_col=2, min_row=start_row+1, max_row=start_row+2)
    s1 = Series(data_full, title_from_data=False)
    s1.title = "Full"
    s1.graphicalProperties.solidFill = theme_colors["bg_track"]
    chart.series.append(s1)
    
    # Series 2: Foreground value
    data_filled = Reference(ws, min_col=3, min_row=start_row+1, max_row=start_row+2)
    s2 = Series(data_filled, title_from_data=False)
    s2.title = "Filled"
    
    # Target specific colors to individual data points in the series
    dp_female = DataPoint(idx=0)
    dp_female.graphicalProperties.solidFill = theme_colors["female_accent"]
    dp_male = DataPoint(idx=1)
    dp_male.graphicalProperties.solidFill = theme_colors["male_accent"]
    s2.dPt = [dp_female, dp_male]
    
    # Add data labels cleanly on top of the columns
    s2.dLbls = DataLabelList()
    s2.dLbls.showVal = True
    s2.dLbls.position = "outEnd"
    
    chart.series.append(s2)
    chart.set_categories(cats)
    chart.legend = None 

    ws.add_chart(chart, anchor)
