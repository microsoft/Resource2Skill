import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

def render_workbook(wb, *, title: str = "KPI Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders an interactive KPI Dashboard driven by a central dropdown.
    Generates a 'Data' sheet for storage and a 'Dashboard' sheet for visualization.
    """
    
    # Theme resolution (fallback provided for self-containment)
    try:
        from _helpers import get_theme_palette
        palette = get_theme_palette(theme)
    except ImportError:
        palette = {
            "primary": "2F5597",
            "secondary": "ED7D31",
            "bg": "F2F2F2",
            "text": "000000",
            "success_bg": "C6EFCE",
            "success_fg": "006100",
            "danger_bg": "FFC7CE",
            "danger_fg": "9C0006"
        }

    # Clean up default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    ws_data = wb.create_sheet("Data")
    ws_dash = wb.create_sheet("Dashboard")
    
    # --- 1. Populate Staging Data ---
    months = ["Jan-20", "Feb-20", "Mar-20", "Apr-20"]
    data_rows = [
        ["Metric"] + months,
        ["DSO Actual", 41, 54, 58, 56],
        ["DSO Target", 45, 45, 45, 45],
        ["DSO Prior Month", 35, 41, 54, 58],
        ["DPO Actual", 90, 100, 95, 89],
        ["DPO Target", 90, 90, 90, 90],
        ["DPO Prior Month", 85, 90, 100, 95],
        ["CAC Actual", 12000, 13500, 14444, 11910],
        ["CAC Target", 15000, 15000, 15000, 15000],
        ["CAC Prior Month", 11000, 12000, 13500, 14444],
        ["Gross Margin Actual", 0.45, 0.40, 0.38, 0.42],
        ["Gross Margin Target", 0.40, 0.40, 0.40, 0.40],
        ["Gross Margin Prior Month", 0.46, 0.45, 0.40, 0.38],
    ]
    
    for r in data_rows:
        ws_data.append(r)
        
    # Apply basic formatting to Data sheet
    ws_data.column_dimensions["A"].width = 25
    for cell in ws_data[1]:
        cell.font = Font(bold=True)
        
    # --- 2. Build Dashboard Layout ---
    ws_dash.sheet_view.showGridLines = False
    ws_dash.column_dimensions["A"].width = 2
    
    # Dashboard Title
    ws_dash["B1"] = title
    ws_dash["B1"].font = Font(size=20, bold=True, color=palette.get("primary", "2F5597"))
    
    # Month Dropdown Selector
    ws_dash["B3"] = "For the month of:"
    ws_dash["B3"].font = Font(bold=True)
    ws_dash["B3"].alignment = Alignment(horizontal="right")
    
    ws_dash["D3"] = "Feb-20"  # Default selected value
    ws_dash["D3"].fill = PatternFill("solid", fgColor="FFFFFF")
    ws_dash["D3"].border = Border(bottom=Side(style="medium", color=palette.get("secondary", "ED7D31")))
    ws_dash["D3"].font = Font(bold=True)
    
    # Apply Data Validation
    dv = DataValidation(type="list", formula1="Data!$B$1:$E$1", allow_blank=False)
    ws_dash.add_data_validation(dv)
    dv.add(ws_dash["D3"])
    
    # --- 3. KPI Card Generator Logic ---
    def make_kpi_card(row: int, col: int, card_title: str, data_row_idx: int, is_lower_better: bool = True, is_percent: bool = False):
        """
        Builds a 4-column wide KPI card pulling data dynamically using INDEX/MATCH.
        """
        # Header (Merged)
        ws_dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+3)
        h_cell = ws_dash.cell(row=row, column=col)
        h_cell.value = card_title
        h_cell.font = Font(bold=True, color="FFFFFF")
        h_cell.fill = PatternFill("solid", fgColor=palette.get("primary", "2F5597"))
        h_cell.alignment = Alignment(horizontal="center")
        
        # Big Value (Merged)
        ws_dash.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+3)
        v_cell = ws_dash.cell(row=row+1, column=col)
        # Dynamic formula referencing the dropdown at D3
        v_cell.value = f'=INDEX(Data!$B${data_row_idx}:$E${data_row_idx}, 1, MATCH($D$3, Data!$B$1:$E$1, 0))'
        v_cell.font = Font(size=24, bold=True)
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell.fill = PatternFill("solid", fgColor="FFFFFF")
        ws_dash.row_dimensions[row+1].height = 40
        
        # Formatting
        if is_percent:
            v_cell.number_format = "0%"
        elif "CAC" in card_title:
            v_cell.number_format = "$#,##0"
            
        # Target / Prior Info Row (Unmerged to fit 4 distinct fields)
        labels = ["Vs. Target", "Vs. Prior Month"]
        data_indices = [data_row_idx + 1, data_row_idx + 2]
        
        for i in range(2):
            lbl_col = col + (i * 2)
            val_col = lbl_col + 1
            
            lbl_cell = ws_dash.cell(row=row+2, column=lbl_col)
            lbl_cell.value = labels[i]
            lbl_cell.font = Font(size=9, color="555555", italic=True)
            lbl_cell.fill = PatternFill("solid", fgColor="F5F5F5")
            lbl_cell.alignment = Alignment(horizontal="right")
            
            val_cell = ws_dash.cell(row=row+2, column=val_col)
            val_cell.value = f'=INDEX(Data!$B${data_indices[i]}:$E${data_indices[i]}, 1, MATCH($D$3, Data!$B$1:$E$1, 0))'
            val_cell.font = Font(size=10, bold=True)
            val_cell.fill = PatternFill("solid", fgColor="F5F5F5")
            val_cell.alignment = Alignment(horizontal="left")
            
            if is_percent:
                val_cell.number_format = "0%"
            elif "CAC" in card_title:
                val_cell.number_format = "$#,##0"
                
            if i == 0:
                t_cell = val_cell  # Save the target cell to use in conditional formatting

        # Conditional Formatting for Big Value
        t_cell_abs = f"${t_cell.column_letter}${t_cell.row}"
        
        success_fill = PatternFill(start_color=palette.get("success_bg", "C6EFCE"), end_color=palette.get("success_bg", "C6EFCE"), fill_type="solid")
        success_font = Font(color=palette.get("success_fg", "006100"), size=24, bold=True)
        danger_fill = PatternFill(start_color=palette.get("danger_bg", "FFC7CE"), end_color=palette.get("danger_bg", "FFC7CE"), fill_type="solid")
        danger_font = Font(color=palette.get("danger_fg", "9C0006"), size=24, bold=True)

        if is_lower_better:
            rule_good = CellIsRule(operator='lessThanOrEqual', formula=[t_cell_abs], stopIfTrue=True, fill=success_fill, font=success_font)
            rule_bad = CellIsRule(operator='greaterThan', formula=[t_cell_abs], stopIfTrue=True, fill=danger_fill, font=danger_font)
        else:
            rule_good = CellIsRule(operator='greaterThanOrEqual', formula=[t_cell_abs], stopIfTrue=True, fill=success_fill, font=success_font)
            rule_bad = CellIsRule(operator='lessThan', formula=[t_cell_abs], stopIfTrue=True, fill=danger_fill, font=danger_font)
            
        ws_dash.conditional_formatting.add(f"{v_cell.coordinate}:{v_cell.coordinate}", rule_good)
        ws_dash.conditional_formatting.add(f"{v_cell.coordinate}:{v_cell.coordinate}", rule_bad)
        
        # Border boxing
        thin = Side(style="thin", color="CCCCCC")
        for r in range(row, row+3):
            for c in range(col, col+4):
                c_obj = ws_dash.cell(row=r, column=c)
                c_obj.border = Border(
                    top=thin if r == row else None,
                    bottom=thin if r == row+2 else None,
                    left=thin if c == col else None,
                    right=thin if c == col+3 else None
                )

    # --- 4. Render Dashboard Components ---
    
    # Section: Working Capital Efficiency
    ws_dash.merge_cells("B5:J5")
    s1 = ws_dash["B5"]
    s1.value = "Working Capital Efficiency"
    s1.font = Font(size=14, bold=True, color="FFFFFF")
    s1.fill = PatternFill("solid", fgColor=palette.get("secondary", "ED7D31"))
    s1.alignment = Alignment(horizontal="center")
    
    # KPIs -> row 7, start col 2 (B) and 7 (G)
    make_kpi_card(7, 2, "DSO (Days Sales Outstanding)", data_row_idx=2, is_lower_better=True)
    make_kpi_card(7, 7, "DPO (Days Payables)", data_row_idx=5, is_lower_better=False)
    
    # Section: Sales KPIs
    ws_dash.merge_cells("B12:J12")
    s2 = ws_dash["B12"]
    s2.value = "Sales KPIs"
    s2.font = Font(size=14, bold=True, color="FFFFFF")
    s2.fill = PatternFill("solid", fgColor=palette.get("secondary", "ED7D31"))
    s2.alignment = Alignment(horizontal="center")
    
    make_kpi_card(14, 2, "CAC (Customer Acq Cost)", data_row_idx=8, is_lower_better=True)
    make_kpi_card(14, 7, "Gross Margin", data_row_idx=11, is_lower_better=False, is_percent=True)
    
    # Format dashboard grid columns appropriately
    for col_letter in ["B", "C", "D", "E", "G", "H", "I", "J"]:
        ws_dash.column_dimensions[col_letter].width = 15
    ws_dash.column_dimensions["F"].width = 4  # Spacer column

    wb.active = ws_dash
