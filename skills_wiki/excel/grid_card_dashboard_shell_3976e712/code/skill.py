from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a structural shell for a modern dashboard.
    Uses cell background colors to create 'cards' on a canvas.
    """
    ws = wb.create_sheet(sheet_name)
    
    # 1. Hide gridlines for a clean UI
    ws.sheet_view.showGridLines = False
    
    # 2. Define Theme Colors (using modern hex values as fallback)
    # In a full framework, these would be fetched via `theme` palette hooks.
    bg_color = "F3F4F6"       # Light gray canvas (creates the "card" effect)
    card_color = "FFFFFF"     # Bright white cards
    sidebar_color = "1E3A8A"  # Dark blue sidebar
    text_color = "111827"     # Dark text
    
    canvas_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    card_fill = PatternFill(start_color=card_color, end_color=card_color, fill_type="solid")
    sidebar_fill = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    
    # 3. Paint entire canvas region
    for row in ws.iter_rows(min_row=1, max_row=25, min_col=1, max_col=25):
        for cell in row:
            cell.fill = canvas_fill
            
    # 4. Paint Sidebar (Column A)
    ws.column_dimensions['A'].width = 8
    for row in range(1, 26):
        ws.cell(row=row, column=1).fill = sidebar_fill
        
    # 5. Add Main Dashboard Title
    main_title = ws.cell(row=2, column=3, value=title)
    main_title.font = Font(name="Calibri", size=24, bold=True, color=sidebar_color)
    main_title.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 40
        
    # 6. Define Card Regions
    # Format: (min_col, min_row, max_col, max_row, card_title)
    cards = [
        # Top KPI Row (3 small cards)
        (3, 4, 6, 8, "Sales"),
        (8, 4, 11, 8, "Profit"),
        (13, 4, 16, 8, "# of Customers"),
        
        # Right Tall Card (Map/Breakdown)
        (18, 4, 23, 21, "Sales by Country"),
        
        # Bottom Row (Wide cards)
        (3, 10, 11, 21, "2021-2022 Sales Trend"),
        (13, 10, 16, 21, "Customer Satisfaction")
    ]
    
    title_font = Font(name="Calibri", size=14, bold=True, color=text_color)
    
    for c_min, r_min, c_max, r_max, c_title in cards:
        # Paint card background
        for r in range(r_min, r_max + 1):
            for c in range(c_min, c_max + 1):
                ws.cell(row=r, column=c).fill = card_fill
        
        # Add card header
        title_cell = ws.cell(row=r_min, column=c_min, value=c_title)
        title_cell.font = title_font
        
        # Apply a subtle indent for professional padding
        title_cell.alignment = Alignment(vertical="center", indent=1)
        
        # Give the header row a bit of breathing room
        ws.row_dimensions[r_min].height = 25
        
    # 7. Size columns to create the proportional grid structure
    # Narrow gaps between cards (Cols B, G, L, Q)
    for col in [2, 7, 12, 17]:
        ws.column_dimensions[get_column_letter(col)].width = 3
        
    # Standardized card block columns
    card_columns = [
        3, 4, 5, 6,             # Block 1
        8, 9, 10, 11,           # Block 2
        13, 14, 15, 16,         # Block 3
        18, 19, 20, 21, 22, 23  # Block 4
    ]
    for col in card_columns:
        ws.column_dimensions[get_column_letter(col)].width = 8.5
