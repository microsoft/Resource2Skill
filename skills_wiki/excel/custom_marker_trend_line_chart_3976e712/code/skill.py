import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str = "2021-2022 Sales Trend (in millions)", data: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a cleanly styled, transparent line chart with custom web-like markers.
    """
    # Example dataset matching the tutorial if none is provided
    if not data:
        data = [
            ["Month", "2021", "2022"],
            ["Jan", 201.9, 215.3],
            ["Feb", 204.2, 217.6],
            ["Mar", 198.6, 220.1],
            ["Apr", 199.2, 206.4],
            ["May", 206.4, 204.3],
            ["Jun", 195.1, 201.0],
        ]

    coords = coordinate_from_string(anchor)
    col_idx = column_index_from_string(coords[0])
    row_idx = coords[1]
    
    # 1. Write data to worksheet
    for i, row_data in enumerate(data):
        for j, val in enumerate(row_data):
            ws.cell(row=row_idx + i, column=col_idx + j, value=val)
            
    # 2. Initialize Line Chart
    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.width = 16
    chart.height = 8
    
    # 3. Apply transparent background and border for dashboard integration
    chart.graphical_properties = GraphicalProperties(noFill=True)
    chart.graphical_properties.line = LineProperties(noFill=True)
    
    # 4. Define data references
    max_row = row_idx + len(data) - 1
    max_col = col_idx + len(data[0]) - 1
    
    chart_data = Reference(ws, min_col=col_idx + 1, min_row=row_idx, max_col=max_col, max_row=max_row)
    categories = Reference(ws, min_col=col_idx, min_row=row_idx + 1, max_row=max_row)
    
    chart.add_data(chart_data, titles_from_data=True)
    chart.set_categories(categories)
    
    # Define series colors (Red and Dark Blue as seen in tutorial)
    series_colors = ["C00000", "1F4E78", "E26B0A", "548235"] 
    
    # 5. Customize each series with modern, smooth lines and polished markers
    for idx, series in enumerate(chart.series):
        color = series_colors[idx % len(series_colors)]
        
        # Line styling
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 20000  # ~1.5 pt
        series.smooth = True
        
        # Marker styling: Circle with white interior, colored border matching the line
        series.marker = Marker(symbol="circle", size=5)
        series.marker.graphicalProperties.solidFill = "FFFFFF"
        series.marker.graphicalProperties.line.solidFill = color
        series.marker.graphicalProperties.line.width = 15000  # ~1.2 pt

    # 6. Axis and layout formatting
    if "y_min" in kwargs:
        chart.y_axis.scaling.min = kwargs["y_min"]
    if "y_max" in kwargs:
        chart.y_axis.scaling.max = kwargs["y_max"]

    chart.legend.position = "b"
    
    # Position the chart immediately to the right of the data table
    chart_anchor_col = get_column_letter(max_col + 2)
    ws.add_chart(chart, f"{chart_anchor_col}{row_idx}")
