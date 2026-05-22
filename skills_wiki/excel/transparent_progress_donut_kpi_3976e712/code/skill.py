def render(ws, anchor: str, *, title: str = "Completion Target", actual: float = 0.85, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import DoughnutChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

    col_str, row = coordinate_from_string(anchor)
    col = column_index_from_string(col_str)

    # 1. Define Theme Colors
    theme_colors = {
        "corporate_blue": ("003366", "E0E0E0"),
        "modern_green": ("2CA02C", "E8F5E9"),
        "warm_red": ("D62728", "FFEBEE")
    }
    accent_color, bg_color = theme_colors.get(theme, ("003366", "E0E0E0"))

    # 2. Setup the "Dashboard Tile" Text (Title and Center KPI Value)
    title_cell = ws.cell(row=row, column=col)
    title_cell.value = title
    title_cell.font = Font(size=12, bold=True, color="333333")

    # Chart will anchor at row+1, col.
    # A 4.5cm x 4.5cm chart covers roughly 2.5 columns and 8.5 rows at default cell sizing.
    # The center of the donut hole will naturally fall around col+1, row+5.
    center_col = col + 1
    center_row = row + 5
    center_cell = ws.cell(row=center_row, column=center_col)
    
    center_cell.value = actual
    center_cell.number_format = '0%'
    center_cell.font = Font(size=20, bold=True, color=accent_color)
    center_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Write Chart Data in an Offset Range (so it doesn't appear under the chart)
    data_col = col + 10
    ws.cell(row=row, column=data_col, value="Metric")
    ws.cell(row=row, column=data_col+1, value="Value")
    
    ws.cell(row=row+1, column=data_col, value="Actual")
    ws.cell(row=row+1, column=data_col+1, value=actual)
    
    ws.cell(row=row+2, column=data_col, value="Remainder")
    ws.cell(row=row+2, column=data_col+1, value=1 - actual)

    # 4. Create and Configure Doughnut Chart
    chart = DoughnutChart()
    chart.holeSize = 65
    chart.width = 4.5  # approx 127 points
    chart.height = 4.5 
    
    labels = Reference(ws, min_col=data_col, min_row=row+1, max_row=row+2)
    data = Reference(ws, min_col=data_col+1, min_row=row+1, max_row=row+2)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(labels)
    
    chart.title = None
    chart.legend = None
    
    # 5. Make Chart Background Transparent (The Secret Sauce)
    # This allows the large text cell to show through the center of the donut
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line = LineProperties(noFill=True)
    
    # 6. Style the Donut Slices
    series = chart.series[0]
    
    pt_actual = DataPoint(idx=0)
    pt_actual.graphicalProperties = GraphicalProperties(solidFill=accent_color)
    
    pt_remainder = DataPoint(idx=1)
    pt_remainder.graphicalProperties = GraphicalProperties(solidFill=bg_color)
    
    series.dPt.append(pt_actual)
    series.dPt.append(pt_remainder)
    
    # 7. Add Chart to Worksheet
    chart_anchor = f"{get_column_letter(col)}{row + 1}"
    ws.add_chart(chart, chart_anchor)
