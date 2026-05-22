from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker

def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a complete, app-like interactive dashboard shell with navigation sidebars
    and floating 'card' style chart containers.
    """
    # 1. Initialize Workbook Sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_inputs = wb.create_sheet("Inputs")
    ws_contacts = wb.create_sheet("Contacts")
    
    # 2. Extract Palette (mocked here for standalone execution)
    # In a full framework, these would map to theme.sidebar_bg, theme.canvas_bg, etc.
    nav_bg = "002060"      # Dark blue sidebar
    text_fg = "FFFFFF"     # White text
    content_bg = "F2F2F2"  # Light gray dashboard canvas
    card_bg = "FFFFFF"     # White chart container panel
    
    # 3. Populate Input Data
    data = [
        ["Month", "2021", "2022"],
        ["Jan", 201.9, 215.3],
        ["Feb", 204.2, 217.6],
        ["Mar", 198.6, 220.1],
        ["Apr", 199.2, 206.4],
        ["May", 206.4, 204.3],
        ["Jun", 195.1, 203.0]
    ]
    for row in data:
        ws_inputs.append(row)
        
    ws_contacts.append(["Name", "Role", "Email"])
    ws_contacts.append(["John Doe", "Regional Manager", "john@mcdonalds.example.com"])
    
    # 4. Build Dashboard Shell Layout
    ws_dash.sheet_view.showGridLines = False
    ws_dash.column_dimensions['A'].width = 14
    
    fill_sidebar = PatternFill(start_color=nav_bg, end_color=nav_bg, fill_type="solid")
    fill_canvas = PatternFill(start_color=content_bg, end_color=content_bg, fill_type="solid")
    fill_card = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
    
    # Apply Canvas & Sidebar Colors
    for row in range(1, 25):
        ws_dash.cell(row=row, column=1).fill = fill_sidebar
        for col in range(2, 14):
            ws_dash.cell(row=row, column=col).fill = fill_canvas
            
    # Create a "White Card" Panel to hold the chart (simulates UI shapes)
    for row in range(5, 20):
        for col in range(3, 11):
            ws_dash.cell(row=row, column=col).fill = fill_card
            
    # 5. Add Dashboard Title
    ws_dash.cell(row=2, column=3, value=title).font = Font(size=22, bold=True, color=nav_bg)
    ws_dash.cell(row=3, column=3, value="Figures in millions of USD").font = Font(size=10, italic=True, color="595959")
    
    # 6. Build Interactive Navigation Menu
    nav_font = Font(color=text_fg, bold=True, underline="single")
    nav_align = Alignment(horizontal="center", vertical="center")
    
    # [(Display Name, Row Index, Hyperlink target)]
    menus = [
        ("Dashboard", 6, "#'Dashboard'!A1"),
        ("Inputs", 9, "#'Inputs'!A1"),
        ("Contacts", 12, "#'Contacts'!A1")
    ]
    
    for name, r_idx, link in menus:
        cell = ws_dash.cell(row=r_idx, column=1, value=name)
        cell.hyperlink = link
        cell.font = nav_font
        cell.alignment = nav_align
        
    # 7. Add Transparent Modern Chart to the Card Panel
    chart = LineChart()
    chart.title = "2021-2022 Sales Trend"
    chart.style = 13  # Clean built-in style
    chart.width = 14.5
    chart.height = 7.5
    
    # Make chart background and border invisible to blend seamlessly with the card panel
    chart.graphical_properties = GraphicalProperties(noFill=True)
    chart.graphical_properties.line = LineProperties(noFill=True)
    
    # Reference the data dynamically
    data_ref = Reference(ws_inputs, min_col=2, min_row=1, max_col=3, max_row=7)
    cats_ref = Reference(ws_inputs, min_col=1, min_row=2, max_row=7)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    
    # Add modern circle markers for readability
    for series in chart.series:
        series.marker = Marker(symbol="circle", size=5)
        
    # Place chart exactly on top of the white panel
    ws_dash.add_chart(chart, "C5")
