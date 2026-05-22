from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, kpis: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a two-pane dashboard shell with a dark left sidebar for KPIs 
    and a light main content area for charts.
    """
    # 1. Theme Configuration (Fallback palettes)
    palettes = {
        "corporate_blue": {
            "sidebar_bg": "102A43", "main_bg": "F0F4F8", 
            "text_sidebar": "FFFFFF", "text_main": "102A43", 
            "accent": "829AB1"
        },
        "forest_green": {
            "sidebar_bg": "1A3636", "main_bg": "E8F0E8", 
            "text_sidebar": "FFFFFF", "text_main": "1A3636", 
            "accent": "688B8B"
        }
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # 2. Worksheet Initialization
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # 3. Column Width Geometry
    ws.column_dimensions['A'].width = 2   # Left margin
    ws.column_dimensions['B'].width = 28  # Sidebar content
    ws.column_dimensions['C'].width = 3   # Middle margin
    
    # Expand main area columns
    for col in range(4, 16):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # 4. Apply Themed Background Blocks
    # Best practice: Apply fill to a finite range (e.g., 50 rows) to prevent file bloat
    sidebar_fill = PatternFill(start_color=palette["sidebar_bg"], end_color=palette["sidebar_bg"], fill_type="solid")
    main_fill = PatternFill(start_color=palette["main_bg"], end_color=palette["main_bg"], fill_type="solid")

    for row in range(1, 51):
        ws.cell(row=row, column=1).fill = sidebar_fill
        ws.cell(row=row, column=2).fill = sidebar_fill
        for col in range(3, 16):
            ws.cell(row=row, column=col).fill = main_fill

    # 5. Render Main Dashboard Title
    title_cell = ws['D2']
    title_cell.value = title.upper()
    title_cell.font = Font(name="Arial", size=24, bold=True, color=palette["text_main"])
    ws.merge_cells('D2:K3')
    title_cell.alignment = Alignment(vertical="center")

    # 6. Render Sidebar KPIs
    if not kpis:
        # Default sample data if none provided
        kpis = [
            {"label": "TOTAL ORDERS", "value": "2,400"},
            {"label": "REVENUE", "value": "$649.0K"},
            {"label": "AVG RATING", "value": "4.0"},
            {"label": "DAYS TO DELIVER", "value": "2.3"}
        ]

    start_row = 6
    lbl_font = Font(name="Arial", size=10, bold=True, color=palette["accent"])
    val_font = Font(name="Arial", size=22, bold=True, color=palette["text_sidebar"])

    for kpi in kpis:
        # Label cell
        lbl_cell = ws.cell(row=start_row, column=2)
        lbl_cell.value = kpi["label"]
        lbl_cell.font = lbl_font
        lbl_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Value cell
        val_cell = ws.cell(row=start_row+1, column=2)
        val_cell.value = kpi["value"]
        val_cell.font = val_font
        val_cell.alignment = Alignment(horizontal="left", vertical="center")

        start_row += 5  # Spacing between KPIs
