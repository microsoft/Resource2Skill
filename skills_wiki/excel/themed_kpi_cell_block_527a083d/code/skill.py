from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", primary_val: float = 369989.0, secondary_val: float = 0.05, theme: str = "corporate_blue", **kwargs) -> None:
    # Standard theme fallback (simulating theme hook extraction)
    themes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF", "accent": "D9E1F2"},
        "dark_mode": {"bg": "262626", "fg": "FFFFFF", "accent": "44546A"},
        "emerald": {"bg": "0F5132", "fg": "FFFFFF", "accent": "A3CFBB"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    row, col = coordinate_to_tuple(anchor)

    # Merge top row for the KPI title
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)

    # Inject values into the anchor grid
    title_cell = ws.cell(row=row, column=col, value=title)
    primary_cell = ws.cell(row=row+1, column=col, value=primary_val)
    secondary_cell = ws.cell(row=row+1, column=col+1, value=secondary_val)

    # Apply data formatting matching the tutorial's metrics
    primary_cell.number_format = '"$"#,##0'
    secondary_cell.number_format = '0%'

    # Define visual styles
    title_font = Font(color=palette["fg"], size=12, bold=False)
    primary_font = Font(color=palette["fg"], size=18, bold=True)
    secondary_font = Font(color=palette["accent"], size=14, bold=True)

    fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # Paint the 2x2 block and align distinct quadrants
    for r in range(row, row + 2):
        for c in range(col, col + 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            
            if r == row:
                # Title styling (top merged row)
                cell.font = title_font
                cell.alignment = center_align
            elif r == row + 1 and c == col:
                # Primary metric styling (bottom left)
                cell.font = primary_font
                cell.alignment = center_align
            elif r == row + 1 and c == col + 1:
                # Secondary metric styling (bottom right)
                cell.font = secondary_font
                cell.alignment = right_align

    # Artificially expand cell dimensions to resemble a UI card element
    ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions[get_column_letter(col+1)].width = 12
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row+1].height = 35
