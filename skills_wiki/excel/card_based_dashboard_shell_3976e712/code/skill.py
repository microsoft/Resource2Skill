import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a card-based dashboard shell layout using cell background colors to emulate 
    floating shapes. Provides designated drop-zones for KPIs and Charts.
    """
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        
    # Disable gridlines for a clean canvas look
    ws.sheet_view.showGridLines = False
    
    # Theme fallbacks
    primary_color = "1F4E78" # Dark Blue
    canvas_color = "F3F4F6"  # Light Gray
    card_color = "FFFFFF"    # White
    border_color = "D1D5DB"  # Soft Gray
    text_color = "333333"    # Dark Charcoal
    
    fill_sidebar = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    fill_canvas = PatternFill(start_color=canvas_color, end_color=canvas_color, fill_type="solid")
    fill_card = PatternFill(start_color=card_color, end_color=card_color, fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color=border_color)
    
    # Paint entire canvas area
    for row in range(1, 26):
        for col in range(2, 15): # Cols B to N
            ws.cell(row=row, column=col).fill = fill_canvas
            
    # Paint sidebar
    for row in range(1, 26):
        ws.cell(row=row, column=1).fill = fill_sidebar
        
    # Configure column widths (Gutters and Cards)
    ws.column_dimensions['A'].width = 8   # Sidebar
    ws.column_dimensions['B'].width = 3   # Gutter
    ws.column_dimensions['F'].width = 3   # Gutter
    ws.column_dimensions['J'].width = 3   # Gutter
    ws.column_dimensions['N'].width = 3   # Gutter
    
    for col in ['C', 'D', 'E', 'G', 'H', 'I', 'K', 'L', 'M']:
        ws.column_dimensions[col].width = 12
        
    # Main Dashboard Title
    ws.merge_cells('C2:M2')
    title_cell = ws['C2']
    title_cell.value = title
    title_cell.font = Font(size=20, bold=True, color=primary_color)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 30
    
    # Define Cards: (start_col, start_row, end_col, end_row, title)
    cards = [
        (3, 4, 5, 8, "Sales"),                   # KPI 1 (C4:E8)
        (7, 4, 9, 8, "Profit"),                  # KPI 2 (G4:I8)
        (11, 4, 13, 8, "# of Customers"),        # KPI 3 (K4:M8)
        (3, 10, 9, 23, "2021-2022 Sales Trend"), # Chart 1 (C10:I23)
        (11, 10, 13, 23, "Sales by Country"),    # Chart 2 (K10:M23)
    ]
    
    # Draw Cards
    for start_col, start_row, end_col, end_row, card_title in cards:
        # Fill card area and apply borders to the perimeter
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_card
                
                # Perimeter borders
                top = thin_border_side if r == start_row else None
                bottom = thin_border_side if r == end_row else None
                left = thin_border_side if c == start_col else None
                right = thin_border_side if c == end_col else None
                
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)
                
        # Card Header Styling
        card_header = ws.cell(row=start_row, column=start_col)
        card_header.value = f"  {card_title}" # Left padding
        card_header.font = Font(bold=True, size=12, color=text_color)
        card_header.alignment = Alignment(vertical="center")
        
        # Merge top row of card for title header
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
        ws.row_dimensions[start_row].height = 25
