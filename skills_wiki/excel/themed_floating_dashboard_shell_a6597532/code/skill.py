def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    subtitle = kwargs.get("subtitle", "Performance Dashboard")
    kpis = kwargs.get("kpis", [
        {"label": "Total Revenue", "value": 1250000, "format": "$#,##0"},
        {"label": "Active Users", "value": 45200, "format": "#,##0"},
        {"label": "Conversion", "value": 0.042, "format": "0.0%"},
        {"label": "Avg Order", "value": 275, "format": "$#,##0"}
    ])

    # Try to load theme, fallback to a robust defined palette
    try:
        from _helpers import get_theme
        palette = get_theme(theme)
    except ImportError:
        palette = {
            "primary": "4B286D",      # Deep Purple
            "secondary": "F2EFF5",    # Light Lilac
            "accent": "F2C811",       # Gold
            "background": "FFFFFF",
            "text": "4B286D",
            "text_light": "FFFFFF"
        }

    header_bg = palette.get("primary", "4B286D")
    header_fg = palette.get("text_light", "FFFFFF")
    body_bg = palette.get("secondary", "F2EFF5")
    card_bg = palette.get("background", "FFFFFF")
    text_main = palette.get("text", "4B286D")
    text_muted = "808080"
    accent = palette.get("accent", "F2C811")

    header_fill = PatternFill("solid", fgColor=header_bg)
    body_fill = PatternFill("solid", fgColor=body_bg)
    card_fill = PatternFill("solid", fgColor=card_bg)

    # 1. Base Grid & Backgrounds
    for c in range(1, 20):
        ws.column_dimensions[get_column_letter(c)].width = 8
    ws.column_dimensions['A'].width = 2 # Left margin
    ws.column_dimensions['B'].width = 2 # Left margin

    # Draw two-tone background
    for r in range(1, 8):
        for c in range(1, 20):
            ws.cell(row=r, column=c).fill = header_fill

    for r in range(8, 40):
        for c in range(1, 20):
            ws.cell(row=r, column=c).fill = body_fill

    # 2. Header text (Aligns with left edge of first KPI/Panel -> Col 3)
    t_cell = ws.cell(row=2, column=3, value=title)
    t_cell.font = Font(name="Calibri", size=22, bold=True, color=header_fg)
    
    s_cell = ws.cell(row=4, column=3, value=subtitle)
    s_cell.font = Font(name="Calibri", size=12, color=accent)

    # 3. Floating KPI Strip
    kpi_width = 3
    spacer = 1
    thick_top = Border(top=Side(style="thick", color=accent))
    
    for i, kpi in enumerate(kpis):
        c_start = 3 + i * (kpi_width + spacer)
        c_end = c_start + kpi_width - 1
        
        # Draw card background & top accent border overlapping the background split
        for r in range(6, 10):
            for c in range(c_start, c_end + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                if r == 6:
                    cell.border = thick_top

        # Value section
        ws.merge_cells(start_row=6, start_column=c_start, end_row=7, end_column=c_end)
        val_cell = ws.cell(row=6, column=c_start)
        val_cell.value = kpi.get("value", "")
        val_cell.font = Font(name="Calibri", size=18, bold=True, color=text_main)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        if "format" in kpi:
            val_cell.number_format = kpi["format"]

        # Label section
        ws.merge_cells(start_row=8, start_column=c_start, end_row=9, end_column=c_end)
        lbl_cell = ws.cell(row=8, column=c_start)
        lbl_cell.value = kpi.get("label", "").upper()
        lbl_cell.font = Font(name="Calibri", size=10, bold=True, color=text_muted)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")

    # 4. Chart Placeholder Panels (Aligned strictly to KPI widths)
    panels = kwargs.get("panels", [
        {"start_row": 12, "end_row": 23, "start_col": 3, "end_col": 9, "title": "KPI Breakdown"},
        {"start_row": 12, "end_row": 23, "start_col": 11, "end_col": 17, "title": "Trend Analysis"},
        {"start_row": 25, "end_row": 36, "start_col": 3, "end_col": 9, "title": "Duration Metrics"},
        {"start_row": 25, "end_row": 36, "start_col": 11, "end_col": 17, "title": "Drop Rates"},
    ])

    thin_bottom = Border(bottom=Side(style="thin", color="E0E0E0"))

    for panel in panels:
        pr = panel["start_row"]
        pc = panel["start_col"]
        
        for r in range(pr, panel["end_row"] + 1):
            for c in range(pc, panel["end_col"] + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                # Add subtle separator line under panel title
                if r == pr:
                    cell.border = thin_bottom

        # Panel Title
        p_title = ws.cell(row=pr, column=pc, value=panel["title"])
        p_title.font = Font(name="Calibri", size=11, bold=True, color=text_main)
        # Indent slightly for title padding
        p_title.alignment = Alignment(indent=1, vertical="center")
