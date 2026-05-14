def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list, panels: list, theme: str = "purple_gold", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple

    # Ensure a clean slate
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Built-in theme presets 
    themes = {
        "corporate_blue": {"primary": "2C3E50", "accent": "3498DB", "bg": "ECF0F1", "card_bg": "FFFFFF", "text": "2C3E50", "text_light": "7F8C8D"},
        "purple_gold": {"primary": "4A235A", "accent": "F1C40F", "bg": "F5EEF8", "card_bg": "FFFFFF", "text": "4A235A", "text_light": "888888"}
    }
    t = themes.get(theme, themes["corporate_blue"])

    # 1. Canvas Background
    fill_bg = PatternFill("solid", fgColor=t["bg"])
    for row in range(1, 40):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = fill_bg

    # 2. Hero Header Banner
    fill_banner = PatternFill("solid", fgColor=t["primary"])
    for row in range(1, 4):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = fill_banner
            
    ws.merge_cells("B2:H2")
    title_cell = ws.cell(row=2, column=2, value=title)
    title_cell.font = Font(name="Arial", size=24, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(vertical="center")
    
    ws.merge_cells("B3:H3")
    subtitle_cell = ws.cell(row=3, column=2, value=subtitle)
    subtitle_cell.font = Font(name="Arial", size=12, color=t["accent"])
    
    # 3. KPI Cards Strip
    start_col = 2
    for kpi in kpis:
        end_col = start_col + 2
        
        # Paint card background and accent stripe
        for r in range(5, 8):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor=t["accent"] if r == 5 else t["card_bg"])
                
                # Subtle outer border
                border_side = Side(border_style="thin", color="DDDDDD")
                cell.border = Border(
                    top=border_side if r == 5 else None,
                    bottom=border_side if r == 7 else None,
                    left=border_side if c == start_col else None,
                    right=border_side if c == end_col else None
                )

        # KPI Value
        ws.merge_cells(start_row=6, start_column=start_col, end_row=6, end_column=end_col)
        val_cell = ws.cell(row=6, column=start_col, value=kpi.get('value', ''))
        val_cell.font = Font(name="Arial", size=18, bold=True, color=t["primary"])
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # KPI Label
        ws.merge_cells(start_row=7, start_column=start_col, end_row=7, end_column=end_col)
        lbl_cell = ws.cell(row=7, column=start_col, value=kpi.get('label', '').upper())
        lbl_cell.font = Font(name="Arial", size=9, bold=True, color=t["text_light"])
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        start_col += 4  # Move right for the next card (3 cols wide + 1 col gap)
        
    # 4. Chart / Data Panel Placeholders
    for panel in panels:
        start_cell, end_cell = panel["range"].split(":")
        start_r, start_c = coordinate_to_tuple(start_cell)
        end_r, end_c = coordinate_to_tuple(end_cell)
        
        for r in range(start_r, end_r + 1):
            for c in range(start_c, end_c + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor=t["card_bg"])
                
                # Box the panel in
                border_side = Side(border_style="thin", color="E0E0E0")
                cell.border = Border(
                    top=border_side if r == start_r else None,
                    bottom=border_side if r == end_r else None,
                    left=border_side if c == start_c else None,
                    right=border_side if c == end_c else None
                )
                
        # Panel Title Bar
        ws.merge_cells(start_row=start_r, start_column=start_c, end_row=start_r, end_column=end_c)
        pt_cell = ws.cell(row=start_r, column=start_c, value=panel.get("title", ""))
        pt_cell.font = Font(name="Arial", size=11, bold=True, color=t["primary"])
        pt_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # 5. Grid Structural Formatting
    ws.column_dimensions['A'].width = 3
    for c in range(2, 22):
        ws.column_dimensions[get_column_letter(c)].width = 8.5
