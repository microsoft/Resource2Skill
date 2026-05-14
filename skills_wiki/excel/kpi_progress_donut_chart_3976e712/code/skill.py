from openpyxl.chart import DoughnutChart, Reference, Series
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Font, Alignment
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str, percentage: float, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a KPI Progress Donut Chart and its textual callouts.
    """
    col_str, row_idx = coordinate_from_string(anchor)
    col_idx = column_index_from_string(col_str)
    
    # 1. Render textual callouts (Title and Value)
    title_cell = ws[anchor]
    title_cell.value = title
    title_cell.font = Font(bold=True, size=12, color="333333")
    title_cell.alignment = Alignment(horizontal="center")
    
    pct_cell = ws.cell(row=row_idx + 1, column=col_idx)
    pct_cell.value = percentage
    pct_cell.number_format = "0%"
    pct_cell.font = Font(bold=True, size=20, color="003366")
    pct_cell.alignment = Alignment(horizontal="center")
    
    # 2. Write chart backing data to an offset area (hidden or out of print area)
    data_col = col_idx + 15
    ws.cell(row=row_idx, column=data_col, value="Actual")
    ws.cell(row=row_idx + 1, column=data_col, value="Remainder")
    ws.cell(row=row_idx, column=data_col + 1, value=percentage)
    ws.cell(row=row_idx + 1, column=data_col + 1, value=1.0 - percentage)
    
    # 3. Create the Doughnut Chart
    chart = DoughnutChart()
    chart.title = None
    chart.legend = None
    chart.holeSize = 65  # Key visual tweak for modern dashboards
    chart.width = 6.0    # Compact size for a KPI card
    chart.height = 6.0
    
    # 4. Bind data and configure series
    data_ref = Reference(ws, min_col=data_col + 1, min_row=row_idx, max_row=row_idx + 1)
    series = Series(data_ref)
    
    # Color the slices: Primary for actual, muted for remainder
    try:
        dp_actual = DataPoint(idx=0)
        dp_actual.graphicalProperties.solidFill = "003366"  # Dark Blue
        
        dp_remainder = DataPoint(idx=1)
        dp_remainder.graphicalProperties.solidFill = "D9E1E8"  # Light Blue
        
        series.dPt = [dp_actual, dp_remainder]
    except AttributeError:
        pass  # Graceful fallback if the specific openpyxl version lacks deep graphicalProperties support
        
    chart.append(series)
    
    # 5. Position the chart below the text callouts
    chart_anchor = f"{get_column_letter(col_idx)}{row_idx + 2}"
    ws.add_chart(chart, chart_anchor)
