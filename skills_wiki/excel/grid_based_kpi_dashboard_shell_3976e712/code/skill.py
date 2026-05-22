import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import DoughnutChart, LineChart, RadarChart, Reference, Series
from openpyxl.drawing.line import LineProperties

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", data: dict = None, **kwargs) -> None:
    ws = wb.create_sheet(sheet_name) if sheet_name not in wb.sheetnames else wb[sheet_name]
    ws.sheet_view.showGridLines = False
    
    # --- Theme Loading ---
    palette = {
        "bg_color": "F3F3F3",
        "card_bg": "FFFFFF",
        "nav_bg": "1F3864",
        "text_main": "000000",
        "text_muted": "595959",
        "accent": "4472C4",
        "border": "CCCCCC"
    }
    try:
        from _helpers import get_theme_palette
        theme_pal = get_theme_palette(theme)
        if theme_pal:
            palette["bg_color"] = theme_pal.get("background", palette["bg_color"])
            palette["card_bg"] = theme_pal.get("surface", palette["card_bg"])
            palette["nav_bg"] = theme_pal.get("primary", palette["nav_bg"])
            palette["text_main"] = theme_pal.get("text_primary", palette["text_main"])
            palette["text_muted"] = theme_pal.get("text_secondary", palette["text_muted"])
            palette["accent"] = theme_pal.get("primary", palette["accent"])
            palette["border"] = theme_pal.get("border", palette["border"])
    except ImportError:
        pass
        
    # --- Default Data ---
    if not data:
        data = {
            "kpis": [
                {"title": "Sales", "actual": 2544, "target": 3000, "format": "$#,##0"},
                {"title": "Profit", "actual": 890, "target": 1000, "format": "$#,##0"},
                {"title": "# of Customers", "actual": 87, "target": 100, "format": "#,##0"}
            ],
            "trend": [
                {"month": "Jan", "actual": 201, "target": 215},
                {"month": "Feb", "actual": 204, "target": 217},
                {"month": "Mar", "actual": 198, "target": 220},
                {"month": "Apr", "actual": 199, "target": 206},
                {"month": "May", "actual": 206, "target": 204},
                {"month": "Jun", "actual": 195, "target": 201}
            ],
            "satisfaction": [
                {"category": "Speed", "score": 0.54},
                {"category": "Quality", "score": 0.86},
                {"category": "Hygiene", "score": 0.91},
                {"category": "Service", "score": 0.53},
                {"category": "Availability", "score": 0.95}
            ]
        }
        
    # --- Setup Hidden Data Sheet ---
    data_ws = wb.create_sheet("_DashboardData")
    data_ws.sheet_state = 'hidden'
    
    # KPI Data (Calculates Remaining for Doughnut Chart)
    data_ws.append(["KPI", "Actual", "Remaining", "Target", "Pct"])
    for kpi in data["kpis"]:
        rem = max(0, kpi["target"] - kpi["actual"])
        pct = kpi["actual"] / kpi["target"] if kpi["target"] else 0
        data_ws.append([kpi["title"], kpi["actual"], rem, kpi["target"], pct])
        
    # Trend Data
    trend_start_row = len(data["kpis"]) + 3
    data_ws.cell(row=trend_start_row, column=1, value="Month")
    data_ws.cell(row=trend_start_row, column=2, value="Actual")
    data_ws.cell(row=trend_start_row, column=3, value="Target")
    for i, row in enumerate(data["trend"]):
        data_ws.cell(row=trend_start_row + 1 + i, column=1, value=row["month"])
        data_ws.cell(row=trend_start_row + 1 + i, column=2, value=row["actual"])
        data_ws.cell(row=trend_start_row + 1 + i, column=3, value=row["target"])
        
    # Satisfaction Data
    sat_start_row = trend_start_row + len(data["trend"]) + 2
    data_ws.cell(row=sat_start_row, column=1, value="Category")
    data_ws.cell(row=sat_start_row, column=2, value="Score")
    for i, row in enumerate(data["satisfaction"]):
        data_ws.cell(row=sat_start_row + 1 + i, column=1, value=row["category"])
        data_ws.cell(row=sat_start_row + 1 + i, column=2, value=row["score"])

    # --- Grid Layout & Base Styling ---
    # Global Background
    gray_fill = PatternFill(start_color=palette["bg_color"], end_color=palette["bg_color"], fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=28, min_col=1, max_col=15):
        for cell in row:
            cell.fill = gray_fill
            
    # Simulated Sidebar
    nav_fill = PatternFill(start_color=palette["nav_bg"], end_color=palette["nav_bg"], fill_type="solid")
    for row in range(1, 29):
        ws.cell(row=row, column=1).fill = nav_fill
        
    # Column Dimensions (incorporating gap columns)
    col_widths = {
        'A': 8, 'B': 3,
        'C': 12, 'D': 6, 'E': 6,
        'F': 3,
        'G': 12, 'H': 6, 'I': 6,
        'J': 3,
        'K': 12, 'L': 6, 'M': 6,
        'N': 3
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Row Spacing
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 15
    for r in range(4, 9):
        ws.row_dimensions[r].height = 15
    ws.row_dimensions[6].height = 25  # Enlarged for big numbers
    ws.row_dimensions[9].height = 15

    # Title
    title_cell = ws.cell(row=2, column=3, value=title)
    title_cell.font = Font(size=24, bold=True, color=palette["nav_bg"])
    title_cell.alignment = Alignment(vertical="center")

    # --- Card Renderer Helper ---
    def style_card(min_col, max_col, min_row, max_row):
        card_fill = PatternFill(start_color=palette["card_bg"], end_color=palette["card_bg"], fill_type="solid")
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Perimeter borders
                top = Side(style='thin', color=palette["border"]) if r == min_row else None
                bottom = Side(style='thin', color=palette["border"]) if r == max_row else None
                left = Side(style='thin', color=palette["border"]) if c == min_col else None
                right = Side(style='thin', color=palette["border"]) if c == max_col else None
                
                if any([top, bottom, left, right]):
                    cell.border = Border(top=top, bottom=bottom, left=left, right=right)

    # --- KPI Cards (Rows 4-8) ---
    col_starts = [3, 7, 11] # C, G, K
    
    for i, col_idx in enumerate(col_starts):
        style_card(col_idx, col_idx+2, 4, 8)
        kpi = data["kpis"][i]
        
        # Labels
        title_cell = ws.cell(row=4, column=col_idx, value=kpi["title"])
        title_cell.font = Font(size=11, color=palette["text_muted"], bold=True)
        
        act_cell = ws.cell(row=6, column=col_idx, value=kpi["actual"])
        act_cell.number_format = kpi["format"]
        act_cell.font = Font(size=18, bold=True, color=palette["text_main"])
        
        pct_cell = ws.cell(row=7, column=col_idx, value=f"{int((kpi['actual']/kpi['target'])*100)}% to target")
        pct_cell.font = Font(size=9, color=palette["accent"], bold=True)
        
        for c_row in [4, 6, 7]:
            ws.cell(row=c_row, column=col_idx).alignment = Alignment(vertical="center")
        
        # Donut Chart
        d_chart = DoughnutChart()
        d_chart.width = 4.0
        d_chart.height = 3.0
        d_chart.legend = None
        d_chart.title = None
        d_chart.holeSize = 65
        
        # Remove borders so it blends into the card
        d_chart.graphical_properties.line = LineProperties(noFill=True)
        
        # Data bindings
        data_row = i + 2
        values = Reference(data_ws, min_col=2, max_col=3, min_row=data_row, max_row=data_row)
        d_chart.add_data(values, from_rows=True)
        
        ws.add_chart(d_chart, f"{openpyxl.utils.get_column_letter(col_idx+1)}4")

    # --- Primary Trend Chart (Rows 10-25) ---
    style_card(3, 9, 10, 25)
    t_chart = LineChart()
    t_chart.title = "Sales Trend vs Target"
    t_chart.width = 14.5
    t_chart.height = 7.5
    t_chart.legend.position = "b"
    t_chart.style = 13 
    t_chart.graphical_properties.line = LineProperties(noFill=True)
    
    t_data = Reference(data_ws, min_col=2, max_col=3, min_row=trend_start_row, max_row=trend_start_row + len(data["trend"]))
    t_cats = Reference(data_ws, min_col=1, max_row=trend_start_row + len(data["trend"]), min_row=trend_start_row + 1)
    
    t_chart.add_data(t_data, titles_from_data=True)
    t_chart.set_categories(t_cats)
    
    # Optional smooth line formatting
    for s in t_chart.series:
        s.smooth = True
        
    ws.add_chart(t_chart, "C11") 

    # --- Secondary Radar Chart (Rows 10-25) ---
    style_card(11, 13, 10, 25)
    r_chart = RadarChart()
    r_chart.type = "filled"
    r_chart.title = "Customer Satisfaction"
    r_chart.width = 7.0
    r_chart.height = 7.5
    r_chart.legend = None
    r_chart.graphical_properties.line = LineProperties(noFill=True)
    
    r_data = Reference(data_ws, min_col=2, max_row=sat_start_row + len(data["satisfaction"]), min_row=sat_start_row)
    r_cats = Reference(data_ws, min_col=1, max_row=sat_start_row + len(data["satisfaction"]), min_row=sat_start_row + 1)
    
    r_chart.add_data(r_data, titles_from_data=True)
    r_chart.set_categories(r_cats)
    
    ws.add_chart(r_chart, "K11")
