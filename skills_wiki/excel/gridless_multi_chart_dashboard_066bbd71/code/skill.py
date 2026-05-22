def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import PatternFill, Font, Alignment

    # 1. Setup Theme colors (fallback to corporate blue if no dict provided)
    primary_color = "1F4E78" 
    text_color = "FFFFFF"
    
    if isinstance(theme, dict):
        primary_color = theme.get("primary", primary_color).replace("#", "")
        text_color = theme.get("text_on_primary", text_color).replace("#", "")
        
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_font = Font(color=text_color, size=24, bold=True)
    
    # 2. Configure Dashboard Presentation Sheet
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    # Header banner
    ws_dash.merge_cells("A1:P3")
    header_cell = ws_dash["A1"]
    header_cell.value = title
    header_cell.fill = header_fill
    header_cell.font = header_font
    header_cell.alignment = Alignment(vertical="center", indent=1)
    
    # 3. Create Hidden Data Sheet
    ws_data = wb.create_sheet("ChartData")
    ws_data.sheet_state = "hidden"
    
    # Table 1: Profit by Country & Product (Rows 1-5)
    ws_data.append(["Country", "Chocolate Chip", "Fortune Cookie", "Sugar"])
    ws_data.append(["India", 62349, 4872, 18560])
    ws_data.append(["United States", 36657, 6368, 9937])
    ws_data.append(["United Kingdom", 46530, 5220, 14620])
    ws_data.append(["Philippines", 54618, 7026, 8313])
    
    # Table 2: Units by Month (Rows 6-11)
    ws_data.append([]) 
    ws_data.append(["Month", "Units Sold"]) 
    ws_data.append(["Sep", 50601])
    ws_data.append(["Oct", 95622])
    ws_data.append(["Nov", 65481])
    ws_data.append(["Dec", 52970])
    
    # Table 3: Profit by Month (Rows 12-17)
    ws_data.append([]) 
    ws_data.append(["Month", "Profit"]) 
    ws_data.append(["Sep", 124812])
    ws_data.append(["Oct", 228275])
    ws_data.append(["Nov", 160228])
    ws_data.append(["Dec", 136337])
    
    # 4. Create and Position Charts
    
    # Chart 1: Stacked Column (Market & Product)
    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Cookie Type"
    c1.y_axis.number_format = '$#,##0'
    c1.legend.position = "b"
    
    data1 = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=5)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    c1.add_data(data1, titles_from_data=True)
    c1.set_categories(cats1)
    
    c1.width = 16
    c1.height = 10
    ws_dash.add_chart(c1, "B5")
    
    # Chart 2: Line (Units over Time)
    c2 = LineChart()
    c2.title = "Units sold each month"
    c2.y_axis.number_format = '#,##0'
    c2.legend = None  # Hide legend to save space
    
    data2 = Reference(ws_data, min_col=2, min_row=7, max_row=11)
    cats2 = Reference(ws_data, min_col=1, min_row=8, max_row=11)
    c2.add_data(data2, titles_from_data=True)
    c2.set_categories(cats2)
    
    c2.width = 12
    c2.height = 7
    ws_dash.add_chart(c2, "J5")
    
    # Chart 3: Line (Profit over Time)
    c3 = LineChart()
    c3.title = "Profit by month"
    c3.y_axis.number_format = '$#,##0'
    c3.legend = None
    
    data3 = Reference(ws_data, min_col=2, min_row=13, max_row=17)
    cats3 = Reference(ws_data, min_col=1, min_row=14, max_row=17)
    c3.add_data(data3, titles_from_data=True)
    c3.set_categories(cats3)
    
    c3.width = 12
    c3.height = 7
    ws_dash.add_chart(c3, "J15")
