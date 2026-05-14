from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_workbook(wb: Workbook, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a clean-canvas dashboard workbook with separated data and presentation layers.
    """
    # 1. Setup Data Sheet
    ws_data = wb.active
    ws_data.title = "Data"
    
    # Sample Aggregated Data: Profit by Market & Product (For Stacked Bar)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Sugar", "Snickerdoodle"],
        ["India", 62349, 4872, 18561, 25085],
        ["Philippines", 54618, 7026, 14947, 8313],
        ["United Kingdom", 46530, 5220, 19446, 14620],
        ["United States", 36657, 6368, 9186, 9937],
    ]
    
    for row in market_data:
        ws_data.append(row)
        
    # Sample Aggregated Data: Monthly Trends (For Line Charts)
    # Adding an empty row for separation, then the next table
    ws_data.append([])
    
    trend_start_row = ws_data.max_row + 1
    trend_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    
    for row in trend_data:
        ws_data.append(row)

    # 2. Setup Dashboard Sheet
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False
    
    # Dashboard Header
    ws_dash.merge_cells("B2:M3")
    header_cell = ws_dash["B2"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color="1F4E78")
    header_cell.alignment = Alignment(vertical="center")
    
    # 3. Create Main Stacked Bar Chart (Profit by Market & Cookie)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 2  # Hooks into standard Excel theme palettes
    bar_chart.height = 12
    bar_chart.width = 18
    
    data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    
    # 4. Create Line Chart 1 (Units Sold each month)
    line_units = LineChart()
    line_units.title = "Units sold each month"
    line_units.style = 2
    line_units.height = 5.5
    line_units.width = 12
    
    units_data_ref = Reference(ws_data, min_col=2, min_row=trend_start_row, max_row=trend_start_row+4)
    units_cats_ref = Reference(ws_data, min_col=1, min_row=trend_start_row+1, max_row=trend_start_row+4)
    line_units.add_data(units_data_ref, titles_from_data=True)
    line_units.set_categories(units_cats_ref)
    line_units.legend = None # Clean look, title explains it
    
    # 5. Create Line Chart 2 (Profit by month)
    line_profit = LineChart()
    line_profit.title = "Profit by month"
    line_profit.style = 2
    line_profit.height = 5.5
    line_profit.width = 12
    
    profit_data_ref = Reference(ws_data, min_col=3, min_row=trend_start_row, max_row=trend_start_row+4)
    line_profit.add_data(profit_data_ref, titles_from_data=True)
    line_profit.set_categories(units_cats_ref)
    line_profit.legend = None

    # 6. Arrange Dashboard Grid
    # Align stacked chart on the left, line charts stacked vertically on the right
    ws_dash.add_chart(bar_chart, "B5")
    ws_dash.add_chart(line_units, "K5")
    ws_dash.add_chart(line_profit, "K16")
