from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str = "Revenue", metric_val: str = "$369,989", sub_metric_val: str = "5% Market Share", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a KPI card component using a merged cell grid to emulate the
    visual effect of the floating shapes demonstrated in the tutorial.
    """
    # Simple theme palette resolution (fallback to defaults)
    palettes = {
        "corporate_blue": {"bg": "17365D", "text": "FFFFFF"}, # Navy Blue
        "executive_dark": {"bg": "262626", "text": "FFFFFF"}  # Dark Charcoal
    }
    theme_colors = palettes.get(theme, palettes["corporate_blue"])
    bg_color = theme_colors["bg"]
    text_color = theme_colors["text"]

    # Parse anchor coordinate (e.g., "B2")
    col_str, row = coordinate_from_string(anchor)
    col_idx = column_index_from_string(col_str)

    # Define the 2x3 block size for the KPI card
    col_end_idx = col_idx + 1
    col_end_str = get_column_letter(col_end_idx)
    row_end = row + 2

    # Initialize the specific cells
    c_title = ws.cell(row=row, column=col_idx)
    c_metric = ws.cell(row=row+1, column=col_idx)
    c_sub = ws.cell(row=row+2, column=col_idx)

    # Assign values (these can be strings, ints, or formulas like "=Data!B2")
    c_title.value = title
    c_metric.value = metric_val
    if sub_metric_val:
        c_sub.value = sub_metric_val

    # Apply merges row by row to create full-width card lines
    ws.merge_cells(f"{col_str}{row}:{col_end_str}{row}")
    ws.merge_cells(f"{col_str}{row+1}:{col_end_str}{row+1}")
    ws.merge_cells(f"{col_str}{row+2}:{col_end_str}{row+2}")

    # Set cell dimensions for a "card" feel
    ws.column_dimensions[col_str].width = 12
    ws.column_dimensions[col_end_str].width = 12
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row+1].height = 30
    ws.row_dimensions[row+2].height = 20

    # Styling objects
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    thick_border = Side(style="medium", color=bg_color)
    border = Border(left=thick_border, right=thick_border, top=thick_border, bottom=thick_border)
    center_align = Alignment(horizontal="center", vertical="center")

    # Apply background, border, and alignment to the entire merged block
    for r in range(row, row_end + 1):
        for c in range(col_idx, col_end_idx + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.border = border
            cell.alignment = center_align

    # Apply typography hierarchy
    c_title.font = Font(color=text_color, size=12, bold=False)
    c_metric.font = Font(color=text_color, size=20, bold=True)
    c_sub.font = Font(color=text_color, size=11, bold=True, italic=True)
