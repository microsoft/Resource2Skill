from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders an app-like dark mode dashboard shell, hiding Excel gridlines and headers,
    and setting up a top navigation bar with placeholder macro buttons.
    """
    ws = wb.create_sheet(sheet_name)
    
    # 1. App-like View Configuration
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    
    # 2. Palette Setup (Dark Mode aesthetic inspired by the tutorial)
    bg_main = "1A1A1A"
    bg_nav = "262626"
    text_main = "FFFFFF"
    accent = "0078D4" # Blue accent for interactive buttons
    
    fill_main = PatternFill("solid", fgColor=bg_main)
    fill_nav = PatternFill("solid", fgColor=bg_nav)
    fill_btn = PatternFill("solid", fgColor=accent)
    
    font_title = Font(name="Segoe UI", size=18, bold=True, color=text_main)
    font_btn = Font(name="Segoe UI", size=10, bold=True, color=text_main)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # 3. Apply Background Canvas
    # Coloring a predefined area to act as the dashboard screen
    for row in ws.iter_rows(min_row=1, max_row=40, min_col=1, max_col=20):
        for cell in row:
            cell.fill = fill_main
            
    # 4. Construct Top Navigation Bar
    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=20):
        for cell in row:
            cell.fill = fill_nav
            
    # Add Dashboard Title
    ws.merge_cells("B2:H2")
    ws["B2"].value = title.upper()
    ws["B2"].font = font_title
    ws["B2"].alignment = align_left
    
    # Add UI Buttons (Placeholder styling for VBA macro assignment)
    btn_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    # Button 1: Refresh
    ws.merge_cells("P2:Q2")
    btn1 = ws["P2"]
    btn1.value = "↻ REFRESH"
    btn1.fill = fill_btn
    btn1.font = font_btn
    btn1.alignment = align_center
    for col in ["P", "Q"]:
        ws[f"{col}2"].border = btn_border

    # Button 2: Full Screen
    ws.merge_cells("R2:S2")
    btn2 = ws["R2"]
    btn2.value = "⛶ FULL SCREEN"
    btn2.fill = fill_btn
    btn2.font = font_btn
    btn2.alignment = align_center
    for col in ["R", "S"]:
        ws[f"{col}2"].border = btn_border
        
    # 5. Spacing and Structure Optimization
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 10
    
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['P'].width = 12
    ws.column_dimensions['Q'].width = 12
    ws.column_dimensions['R'].width = 12
    ws.column_dimensions['S'].width = 12
    
    # Remove default sheet if it exists
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
