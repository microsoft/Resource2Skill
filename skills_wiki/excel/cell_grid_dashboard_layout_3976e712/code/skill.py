from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
    
    ws.sheet_view.showGridLines = False
    
    # 2. Colors (Fallback mapping for standard themes)
    bg_color = "F3F3F3"       # Light gray canvas
    panel_color = "FFFFFF"    # White cards
    sidebar_color = "1F4E78"  # Corporate blue accent
    text_muted = "595959"     # Gray text for labels
    border_color = "D9D9D9"   # Soft border
    
    # 3. Apply Canvas Background
    fill_bg = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=15):
        for cell in row:
            cell.fill = fill_bg
            
    # 4. Create Navigation Sidebar (Column A)
    fill_sidebar = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    ws.column_dimensions['A'].width = 8
    for row in range(1, 31):
        ws.cell(row=row, column=1).fill = fill_sidebar
        
    # Set default column widths for main content area
    for col in range(2, 15):
        ws.column_dimensions[get_column_letter(col)].width = 11

    # 5. Panel Builder Helper
    def create_panel(min_col, min_row, max_col, max_row, panel_title=""):
        fill_panel = PatternFill(start_color=panel_color, end_color=panel_color, fill_type="solid")
        
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_panel
                
                # Calculate outer borders for the panel block
                top = Side(style='thin', color=border_color) if r == min_row else None
                bottom = Side(style='thin', color=border_color) if r == max_row else None
                left = Side(style='thin', color=border_color) if c == min_col else None
                right = Side(style='thin', color=border_color) if c == max_col else None
                
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)
                
        if panel_title:
            title_cell = ws.cell(row=min_row, column=min_col)
            title_cell.value = panel_title
            title_cell.font = Font(bold=True, size=11, color=text_muted)
            # Use native indenting instead of space-padding for cleaner UI
            title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # 6. Construct Dashboard Layout Architecture
    
    # Header Panel
    create_panel(2, 2, 10, 4)
    header_cell = ws['B2']
    header_cell.value = title
    header_cell.font = Font(bold=True, size=18, color=sidebar_color)
    header_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    # Tall Right Panel (Reserved for Map/List)
    create_panel(11, 2, 13, 22, "Sales by Country 2022")
    
    # 3-KPI Strip
    create_panel(2, 6, 4, 9, "Sales")
    ws['B7'] = "$2,544"
    ws['B7'].font = Font(bold=True, size=16, color="000000")
    ws['B7'].alignment = Alignment(horizontal="left", indent=1)
    
    create_panel(5, 6, 7, 9, "Profit")
    ws['E7'] = "$890"
    ws['E7'].font = Font(bold=True, size=16, color="000000")
    ws['E7'].alignment = Alignment(horizontal="left", indent=1)
    
    create_panel(8, 6, 10, 9, "# of Customers")
    ws['H7'] = "87.0"
    ws['H7'].font = Font(bold=True, size=16, color="000000")
    ws['H7'].alignment = Alignment(horizontal="left", indent=1)
    
    # Main Chart Areas
    create_panel(2, 11, 6, 22, "2021-2022 Sales Trend (in millions)")
    create_panel(7, 11, 10, 22, "Customer Satisfaction")
