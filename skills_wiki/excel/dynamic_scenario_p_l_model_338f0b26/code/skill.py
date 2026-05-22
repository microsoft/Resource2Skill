def render_sheet(wb, sheet_name: str, *, title: str = "Scenario-Driven P&L", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(sheet_name)
    
    # Theme configuration
    themes = {
        "corporate_blue": {"primary": "2F5597", "text": "FFFFFF", "highlight": "FFF2CC"},
        "slate_gray": {"primary": "374151", "text": "FFFFFF", "highlight": "FDE68A"}
    }
    palette = themes.get(theme, themes["corporate_blue"])
    
    primary_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    primary_font = Font(bold=True, color=palette["text"])
    highlight_fill = PatternFill(start_color=palette["highlight"], end_color=palette["highlight"], fill_type="solid")
    
    # 1. Setup Headers & Layout
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color=palette["primary"])
    
    ws["I2"] = "Scenario:"
    ws["I2"].font = Font(bold=True)
    ws["I2"].alignment = Alignment(horizontal="right")
    
    ws["J2"] = 1
    ws["J2"].fill = highlight_fill
    ws["J2"].border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    ws["J2"].alignment = Alignment(horizontal="center")
    
    # Data Validation Dropdown for Scenario Selector
    dv = DataValidation(type="list", formula1='"1,2,3"', allowBlank=False)
    ws.add_data_validation(dv)
    dv.add("J2")
    
    # Year Headers
    years = ["Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for i, year in enumerate(years, start=3):
        cell = ws.cell(row=3, column=i)
        cell.value = year
        cell.font = primary_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center")
        
    # P&L Labels
    labels = [
        ("Revenue", 4), ("Cost of Goods Sold", 5), ("Gross Profit", 6),
        ("Operating Expenses", 7), ("Operating Profit", 8),
        ("Corporate Tax", 9), ("Net Income", 10)
    ]
    for label, r in labels:
        ws.cell(row=r, column=1, value=label)
        
    # Format P&L Subtotals
    for r in [6, 8, 10]:
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(3, 8):
            ws.cell(row=r, column=c).font = Font(bold=True)
            bottom_style = "double" if r == 10 else "thin"
            ws.cell(row=r, column=c).border = Border(top=Side(style="thin"), bottom=Side(style=bottom_style))
            
    # Assumptions Labels
    ws["A13"] = "Live Case Assumptions"
    ws["A13"].font = Font(bold=True, color=palette["primary"])
    
    assumptions_labels = [
        "Order Growth Rate", "Number of Orders", "Average Order Value",
        "Cost per Order", "Operating Expenses", "Tax Rate"
    ]
    for i, label in enumerate(assumptions_labels, start=14):
        ws.cell(row=i, column=1, value=label)
        
    # Insert Hardcoded Scenarios
    scenarios = [
        ("Scenario 1: Base Case", 21),
        ("Scenario 2: Best Case", 29),
        ("Scenario 3: Worst Case", 37)
    ]
    
    for title_text, r in scenarios:
        ws.cell(row=r, column=1, value=title_text).font = Font(bold=True, italic=True)
        for i, label in enumerate(assumptions_labels, start=r+1):
            ws.cell(row=i, column=1, value=label)
            
    # Populate Scenario 1 (Base Case)
    base_growth = [0.0, 0.5, 0.4, 0.3, 0.2]
    for i, val in enumerate(base_growth, start=3):
        ws.cell(row=22, column=i, value=val)
    ws["C23"] = 3000
    for c in range(4, 8): 
        ws.cell(row=23, column=c, value=f"={get_column_letter(c-1)}23*(1+{get_column_letter(c)}22)")
    for c in range(3, 8):
        ws.cell(row=24, column=c, value=40)
        ws.cell(row=25, column=c, value=10)
        ws.cell(row=26, column=c, value=50000)
        ws.cell(row=27, column=c, value=0.2)
        
    # Populate Scenario 2 (Best Case)
    best_growth = [0.0, 0.8, 0.6, 0.4, 0.3]
    for i, val in enumerate(best_growth, start=3):
        ws.cell(row=30, column=i, value=val)
    ws["C31"] = 4000
    for c in range(4, 8): 
        ws.cell(row=31, column=c, value=f"={get_column_letter(c-1)}31*(1+{get_column_letter(c)}30)")
    for c in range(3, 8):
        ws.cell(row=32, column=c, value=45)
        ws.cell(row=33, column=c, value=9)
        ws.cell(row=34, column=c, value=60000)
        ws.cell(row=35, column=c, value=0.2)
        
    # Populate Scenario 3 (Worst Case)
    worst_growth = [0.0, 0.2, 0.1, 0.05, 0.05]
    for i, val in enumerate(worst_growth, start=3):
        ws.cell(row=38, column=i, value=val)
    ws["C39"] = 2000
    for c in range(4, 8): 
        ws.cell(row=39, column=c, value=f"={get_column_letter(c-1)}39*(1+{get_column_letter(c)}38)")
    for c in range(3, 8):
        ws.cell(row=40, column=c, value=35)
        ws.cell(row=41, column=c, value=12)
        ws.cell(row=42, column=c, value=40000)
        ws.cell(row=43, column=c, value=0.2)

    # Link Live Case using CHOOSE
    for r_live, r_s1, r_s2, r_s3 in zip(range(14, 20), range(22, 28), range(30, 36), range(38, 44)):
        for c in range(3, 8):
            col = get_column_letter(c)
            ws[f"{col}{r_live}"] = f"=CHOOSE($J$2, {col}{r_s1}, {col}{r_s2}, {col}{r_s3})"
            
    # Build P&L Formulas
    for c in range(3, 8):
        col = get_column_letter(c)
        ws[f"{col}4"] = f"={col}15*{col}16" # Revenue = Orders * AOV
        ws[f"{col}5"] = f"={col}15*{col}17" # COGS = Orders * CPO
        ws[f"{col}6"] = f"={col}4-{col}5"   # Gross Profit = Revenue - COGS
        ws[f"{col}7"] = f"={col}18"         # Opex = Live Opex
        ws[f"{col}8"] = f"={col}6-{col}7"   # Operating Profit = GP - Opex
        ws[f"{col}9"] = f"=IF({col}8>0, {col}8*{col}19, 0)" # Corporate Tax (no tax on loss)
        ws[f"{col}10"] = f"={col}8-{col}9"  # Net Income = OP - Tax
        
    # Apply Standard Financial Number Formats
    num_fmt = "_($* #,##0_);_($* (#,##0);_($* \"-\"_);_(@_)"
    for r in [4, 5, 6, 7, 8, 9, 10, 16, 17, 18, 24, 25, 26, 32, 33, 34, 40, 41, 42]:
        for c in range(3, 8):
            ws.cell(row=r, column=c).number_format = num_fmt
            
    pct_fmt = "0.0%"
    for r in [14, 19, 22, 27, 30, 35, 38, 43]:
        for c in range(3, 8):
            ws.cell(row=r, column=c).number_format = pct_fmt
            
    int_fmt = "#,##0"
    for r in [15, 23, 31, 39]:
        for c in range(3, 8):
            ws.cell(row=r, column=c).number_format = int_fmt
            
    # Set Column Widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["I"].width = 12
    for c in range(3, 8):
        ws.column_dimensions[get_column_letter(c)].width = 16
