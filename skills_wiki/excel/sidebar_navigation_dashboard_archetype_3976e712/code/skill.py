from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # Minimal palette fallback
    palettes = {
        "corporate_blue": {"primary": "1F4E78", "accent": "F39C12", "bg": "F2F2F2", "card": "FFFFFF", "text": "333333", "nav_text": "D9E1F2"},
        "modern_dark": {"primary": "2C3E50", "accent": "18BC9C", "bg": "1E1E1E", "card": "2D2D30", "text": "E0E0E0", "nav_text": "AAAAAA"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    sheet_names = ["Dashboard", "Data Inputs", "Contacts"]
    
    # Initialize sheets
    wb.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(name)
        
    for current_sheet in sheet_names:
        ws = wb[current_sheet]
        ws.sheet_view.showGridLines = False
        
        # 1. Apply canvas background
        bg_fill = PatternFill("solid", fgColor=palette["bg"])
        for row in range(1, 41):
            for col in range(2, 21):
                ws.cell(row=row, column=col).fill = bg_fill
                
        # 2. Build Sidebar (Column A)
        ws.column_dimensions['A'].width = 16
        nav_fill = PatternFill("solid", fgColor=palette["primary"])
        for row in range(1, 41):
            ws.cell(row=row, column=1).fill = nav_fill
            
        nav_row = 6
        for sheet in sheet_names:
            cell = ws.cell(row=nav_row, column=1)
            # Active vs Inactive state styling
            if sheet == current_sheet:
                cell.value = f"  ▶ {sheet}"
                cell.font = Font(color=palette["accent"], bold=True)
            else:
                cell.value = f'=HYPERLINK("#\'{sheet}\'!A1", "    {sheet}")'
                cell.font = Font(color=palette["nav_text"])
            cell.alignment = Alignment(vertical="center")
            nav_row += 3
            
    # 3. Build Dashboard Cards
    ws = wb["Dashboard"]
    
    # Title
    title_cell = ws.cell(row=2, column=3)
    title_cell.value = title
    title_cell.font = Font(size=20, bold=True, color=palette["primary"])
    
    def draw_card(ws, start_row, start_col, end_row, end_col, title_text):
        card_fill = PatternFill("solid", fgColor=palette["card"])
        thin_edge = Side(style='thin', color="CCCCCC")
        
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Apply boundary borders
                cell.border = Border(
                    top=thin_edge if r == start_row else None,
                    bottom=thin_edge if r == end_row else None,
                    left=thin_edge if c == start_col else None,
                    right=thin_edge if c == end_col else None
                )
                
        # Card Header
        header_cell = ws.cell(row=start_row, column=start_col)
        header_cell.value = f"  {title_text}"
        header_cell.font = Font(bold=True, color=palette["text"])
        header_cell.alignment = Alignment(vertical="center")

    # Layout UI Cards (KPIs and Charts)
    draw_card(ws, start_row=5, start_col=3,  end_row=10, end_col=6,  title_text="Total Sales")
    draw_card(ws, start_row=5, start_col=8,  end_row=10, end_col=11, title_text="Profit Margin")
    draw_card(ws, start_row=5, start_col=13, end_row=10, end_col=16, title_text="Active Customers")
    
    draw_card(ws, start_row=12, start_col=3,  end_row=25, end_col=11, title_text="Sales Trend (YTD)")
    draw_card(ws, start_row=12, start_col=13, end_row=25, end_col=16, title_text="Satisfaction Score")
