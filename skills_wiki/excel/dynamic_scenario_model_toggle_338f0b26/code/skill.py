from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Theme setup
    theme_colors = {
        "corporate_blue": {"primary": "4F81BD", "accent": "DCE6F1"},
        "emerald_green": {"primary": "00B050", "accent": "E2EFDA"},
    }.get(theme, {"primary": "4F81BD", "accent": "DCE6F1"})
    
    header_fill = PatternFill("solid", fgColor=theme_colors["primary"])
    header_font = Font(color="FFFFFF", bold=True)
    toggle_fill = PatternFill("solid", fgColor=theme_colors["accent"])
    bold_font = Font(bold=True)
    input_font = Font(color="0000FF") # Financial modeling standard for hard-coded inputs
    
    top_border = Border(top=Side(style="thin"))
    
    # 1. Set up Scenario Toggle Dropdown
    ws["H2"] = "Scenario:"
    ws["H2"].font = bold_font
    ws["H2"].alignment = Alignment(horizontal="right")
    
    toggle_cell = ws["I2"]
    toggle_cell.value = 1
    toggle_cell.fill = toggle_fill
    toggle_cell.alignment = Alignment(horizontal="center")
    toggle_cell.border = Border(outline=Side(style="thin", color="000000"))
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(toggle_cell)
    
    ws["J2"] = "1 = Base Case, 2 = Downside"
    ws["J2"].font = Font(italic=True, color="808080")
    
    # 2. Section Headers
    sections = {
        4: "Income Statement",
        14: "Live Case Assumptions",
        23: "Scenario 1: Base Case",
        32: "Scenario 2: Downside Case"
    }
    
    years = ["Year 1", "Year 2", "Year 3"]
    
    for row_idx, sec_title in sections.items():
        ws.cell(row=row_idx, column=2, value=sec_title).font = header_font
        ws.cell(row=row_idx, column=2).fill = header_fill
        for c_idx, year in enumerate(years):
            col = 3 + c_idx
            c = ws.cell(row=row_idx, column=col, value=year)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
            
    # 3. Model Layout & Labels
    labels = ["Orders", "Order Growth", "AOV", "COGS per order", "Fixed Opex", "Tax Rate"]
    is_labels = ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "Operating Profit", "Tax", "Net Profit"]
    
    for i, label in enumerate(is_labels):
        ws.cell(row=5+i, column=2, value=label)
        
    # 4. Fill Scenarios (Hardcoded inputs)
    def set_scenario_vals(start_row, vals_matrix):
        for r_idx, (label, vals, is_pct) in enumerate(vals_matrix):
            row = start_row + r_idx
            ws.cell(row=row, column=2, value=label)
            for c_idx, val in enumerate(vals):
                if val is not None:
                    cell = ws.cell(row=row, column=3+c_idx, value=val)
                    cell.font = input_font
                    if is_pct:
                        cell.number_format = "0%"
                    else:
                        cell.number_format = "#,##0.00" if isinstance(val, float) else "#,##0"

    base_case = [
        ("Orders", [3000, 6000, 10500], False),
        ("Order Growth", [None, 1.0, 0.75], True),
        ("AOV", [39.95, 39.95, 39.95], False),
        ("COGS per order", [8.75, 8.75, 8.75], False),
        ("Fixed Opex", [100000, 100000, 100000], False),
        ("Tax Rate", [0.2, 0.2, 0.2], True)
    ]
    
    downside_case = [
        ("Orders", [2000, 3000, 4500], False),
        ("Order Growth", [None, 0.5, 0.5], True),
        ("AOV", [34.95, 34.95, 34.95], False),
        ("COGS per order", [9.50, 9.50, 9.50], False),
        ("Fixed Opex", [100000, 100000, 100000], False),
        ("Tax Rate", [0.25, 0.25, 0.25], True)
    ]
    
    set_scenario_vals(24, base_case)
    set_scenario_vals(33, downside_case)
    
    # 5. Build Live Case with CHOOSE formulas
    for r_idx, label in enumerate(labels):
        row = 15 + r_idx
        ws.cell(row=row, column=2, value=label)
        
        for c_idx in range(3):
            col = 3 + c_idx
            # Map correctly to the rows configured in the data dumps above
            scen1_cell = f"{get_column_letter(col)}{row+9}"
            scen2_cell = f"{get_column_letter(col)}{row+18}"
            
            cell = ws.cell(row=row, column=col, value=f"=CHOOSE($I$2, {scen1_cell}, {scen2_cell})")
            
            # Formatting
            is_pct = (label in ["Order Growth", "Tax Rate"])
            if is_pct:
                cell.number_format = "0%"
            else:
                cell.number_format = "#,##0.00" if "AOV" in label or "COGS" in label else "#,##0"

    # 6. Build Income Statement Formulas
    for c_idx in range(3):
        col = 3 + c_idx
        col_ltr = get_column_letter(col)
        
        # Revenue = Orders * AOV
        ws.cell(row=5, column=col, value=f"={col_ltr}15*{col_ltr}17")
        # COGS = Orders * COGS per order
        ws.cell(row=6, column=col, value=f"={col_ltr}15*{col_ltr}18")
        
        # Gross Profit
        gp = ws.cell(row=7, column=col, value=f"={col_ltr}5-{col_ltr}6")
        gp.border = top_border
        gp.font = bold_font
        
        # Opex
        ws.cell(row=8, column=col, value=f"={col_ltr}19")
        
        # Op Profit
        op = ws.cell(row=9, column=col, value=f"={col_ltr}7-{col_ltr}8")
        op.border = top_border
        op.font = bold_font
        
        # Tax = MAX(Op_Profit * Tax_Rate, 0)
        ws.cell(row=10, column=col, value=f"=MAX({col_ltr}9*{col_ltr}20, 0)")
        
        # Net Profit
        np = ws.cell(row=11, column=col, value=f"={col_ltr}9-{col_ltr}10")
        np.border = Border(top=Side(style="thin"), bottom=Side(style="double"))
        np.font = bold_font
        
        for r in range(5, 12):
            ws.cell(row=r, column=col).number_format = "$#,##0"

    # Polish Layout Widths
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 30
