import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import DoughnutChart, LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Base Theme Configuration
    colors = {
        "bg": "F3F4F6",         # Light grey canvas
        "card": "FFFFFF",       # White cards
        "sidebar": "1E3A8A",    # Dark blue sidebar
        "text": "1F2937",       # Dark slate text
        "border": "E5E7EB",     # Light grey borders
        "accents": ["2563EB", "10B981", "F59E0B"]  # Blue, Green, Orange
    }

    # 2. Data Setup (Inputs Sheet)
    ws_data = wb.active
    ws_data.title = "Inputs"

    ws_data.append(["KPI Data"])
    ws_data.append(["Metric", "Complete", "Remaining"])
    ws_data.append(["Sales", 0.85, 0.15])
    ws_data.append(["Profit", 0.89, 0.11])
    ws_data.append(["Customers", 0.87, 0.13])

    ws_data.append([])
    ws_data.append(["Trend Data"])
    ws_data.append(["Month", "2021", "2022"])
    for row in [
        ["Jan", 201, 215], ["Feb", 204, 217], ["Mar", 198, 220],
        ["Apr", 199, 206], ["May", 206, 204], ["Jun", 195, 203]
    ]:
        ws_data.append(row)

    ws_data.append([])
    ws_data.append(["Category Data"])
    ws_data.append(["Country", "Sales"])
    for row in [
        ["Brazil", 553], ["Colombia", 432], ["Argentina", 953],
        ["Peru", 425], ["Chile", 253]
    ]:
        ws_data.append(row)

    # 3. Dashboard Layout Setup
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False

    # Structural Grid
    ws_dash.column_dimensions['A'].width = 8
    for col in ['B', 'F', 'J']: # Spacers
        ws_dash.column_dimensions[col].width = 3
    for col in ['C', 'D', 'E', 'G', 'H', 'I', 'K', 'L', 'M']: # Content blocks
        ws_dash.column_dimensions[col].width = 12

    # Canvas Fill
    bg_fill = PatternFill(start_color=colors["bg"], end_color=colors["bg"], fill_type="solid")
    for row in range(1, 40):
        for col in range(1, 16):
            ws_dash.cell(row=row, column=col).fill = bg_fill

    # Sidebar Fill
    sidebar_fill = PatternFill(start_color=colors["sidebar"], end_color=colors["sidebar"], fill_type="solid")
    for row in range(1, 40):
        ws_dash.cell(row=row, column=1).fill = sidebar_fill

    # Header
    ws_dash.cell(row=2, column=3, value=title).font = Font(size=24, bold=True, color=colors["sidebar"])
    ws_dash.cell(row=3, column=3, value="Figures in millions of USD").font = Font(size=12, italic=True, color="6B7280")

    # 4. Card Drawing Engine
    card_fill = PatternFill(start_color=colors["card"], end_color=colors["card"], fill_type="solid")
    
    def draw_card(start_col, start_row, end_col, end_row, title_text):
        thin_border = Side(style='thin', color=colors["border"])
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws_dash.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Intelligent perimeter borders
                b_top = thin_border if r == start_row else None
                b_bottom = thin_border if r == end_row else None
                b_left = thin_border if c == start_col else None
                b_right = thin_border if c == end_col else None
                
                if b_top or b_bottom or b_left or b_right:
                    cell.border = Border(top=b_top, bottom=b_bottom, left=b_left, right=b_right)

        t_cell = ws_dash.cell(row=start_row + 1, column=start_col + 1)
        t_cell.value = title_text
        t_cell.font = Font(name="Calibri", size=14, bold=True, color=colors["text"])

    # 5. KPI Cards & Doughnut Charts
    def add_kpi(start_col, title_text, pct_val, data_row, color):
        end_col = start_col + 2
        draw_card(start_col, 5, end_col, 12, title_text)

        # Highlight Value
        val_cell = ws_dash.cell(row=6, column=end_col)
        val_cell.value = pct_val
        val_cell.number_format = "0%"
        val_cell.font = Font(size=20, bold=True, color=color)
        val_cell.alignment = Alignment(horizontal="right")

        # Doughnut ring
        chart = DoughnutChart()
        chart.width = 4.5
        chart.height = 3.2
        chart.legend = None
        # Make chart transparent so the white card shows through
        chart.graphical_properties.noFill = True
        chart.graphical_properties.line.noFill = True

        data = Reference(ws_data, min_col=2, min_row=data_row, max_col=3)
        chart.add_data(data, titles_from_data=False)
        ws_dash.add_chart(chart, f"{get_column_letter(start_col)}7")

    add_kpi(3, "Sales", 0.85, 3, colors["accents"][0])
    add_kpi(7, "Profit", 0.89, 4, colors["accents"][1])
    add_kpi(11, "Customers", 0.87, 5, colors["accents"][2])

    # 6. Trend Line Chart
    draw_card(3, 14, 9, 28, "2021-2022 Sales Trend")
    line_chart = LineChart()
    line_chart.width = 13
    line_chart.height = 6.5
    line_chart.legend.position = "b"
    line_chart.graphical_properties.noFill = True
    line_chart.graphical_properties.line.noFill = True
    
    line_data = Reference(ws_data, min_col=2, min_row=8, max_col=3, max_row=14)
    line_cats = Reference(ws_data, min_col=1, min_row=9, max_row=14)
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)
    ws_dash.add_chart(line_chart, "C16")

    # 7. Category Bar Chart
    draw_card(11, 14, 13, 28, "Sales by Country")
    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.barDir = "bar"
    bar_chart.width = 5.5
    bar_chart.height = 6.5
    bar_chart.legend = None
    bar_chart.graphical_properties.noFill = True
    bar_chart.graphical_properties.line.noFill = True
    
    bar_data = Reference(ws_data, min_col=2, min_row=17, max_row=22)
    bar_cats = Reference(ws_data, min_col=1, min_row=18, max_row=22)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    ws_dash.add_chart(bar_chart, "K16")
