def render_workbook(wb, *, title: str = "Interactive Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    
    # Define pages and their corresponding icons
    pages = [
        {"name": "Dashboard", "icon": "🏠 "},
        {"name": "Details", "icon": "📊 "},
        {"name": "Settings", "icon": "⚙️ "}
    ]
    sheet_names = [p["name"] for p in pages]
    
    # Safely initialize the workbook sheets
    if len(wb.sheetnames) == 1 and wb.sheetnames[0] == "Sheet":
        wb.active.title = sheet_names[0]
        
    for sn in sheet_names:
        if sn not in wb.sheetnames:
            wb.create_sheet(sn)
            
    # Theme colors (Dark mode sidebar layout)
    sidebar_bg = "1A233A"     # Dark navy
    active_bg = "27344F"      # Lighter navy for active item background
    text_color = "E2E8F0"     # Light gray text
    active_text = "38BDF8"    # Bright blue for active text
    
    sidebar_fill = PatternFill("solid", fgColor=sidebar_bg)
    active_fill = PatternFill("solid", fgColor=active_bg)
    
    font_normal = Font(name="Calibri", size=12, bold=True, color=text_color)
    font_active = Font(name="Calibri", size=12, bold=True, color=active_text)
    
    for current_sheet in sheet_names:
        ws = wb[current_sheet]
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 25
        
        # Color sidebar (fill rows 1 through 40)
        for r in range(1, 41):
            ws.cell(row=r, column=1).fill = sidebar_fill
            
        # Logo / Title Area
        logo_cell = ws.cell(row=3, column=1, value="APP MENU")
        logo_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        logo_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Build Navigation Links
        start_row = 7
        for idx, page in enumerate(pages):
            row_idx = start_row + (idx * 3)
            cell = ws.cell(row=row_idx, column=1, value=page["icon"] + page["name"])
            
            # Clean layout for the navigation item
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=2)
            
            if page["name"] == current_sheet:
                # Active state styling
                cell.fill = active_fill
                cell.font = font_active
            else:
                # Inactive state styling + Internal Hyperlink
                cell.font = font_normal
                cell.hyperlink = f"#'{page['name']}'!A1"
                
        # Main content area placeholder styling
        ws.column_dimensions['B'].width = 5
        content_header = ws.cell(row=3, column=3, value=current_sheet.upper())
        content_header.font = Font(name="Calibri", size=24, bold=True, color="333333")
