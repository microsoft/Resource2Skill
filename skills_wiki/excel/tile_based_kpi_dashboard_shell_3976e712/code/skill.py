from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import LineChart, RadarChart, Reference
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a tile-based executive dashboard with a faux navigation sidebar, 
    top-level KPI cards, and designated areas for Line and Radar charts.
    """
    # 1. Theme Configuration
    # In a real system, these would load from a theme registry based on the `theme` arg.
    palette = {
        "sidebar_bg": "1F3864", # Dark Navy
        "canvas_bg": "F2F2F2",  # Light Gray
        "card_bg": "FFFFFF",    # White
        "border": "D9D9D9",     # Subtle Gray
        "title_fg": "203764",   # Deep Blue
        "text_main": "000000",
        "text_muted": "7F7F7F"
    }

    ws = wb.create_sheet(sheet_name)
    ws_data = wb.create_sheet(f"{sheet_name}_Data")
    
    # Disable gridlines for a clean dashboard look
    ws.sheet_view.showGridLines = False

    # 2. Populate Backend Data
    sales_data = [
        ["Month", "2021", "2022"],
        ["Jan", 201.9, 215.3],
        ["Feb", 204.2, 217.6],
        ["Mar", 198.6, 220.1],
        ["Apr", 199.2, 206.4],
        ["May", 206.4, 204.3],
        ["Jun", 195.1, 201.0],
        ["Jul", 192.4, 205.5],
        ["Aug", 189.3, 200.6],
        ["Sep", 194.2, 210.6]
    ]
    for row in sales_data:
        ws_data.append(row)

    radar_data = [
        ["Metric", "Score"],
        ["Speed", 54],
        ["Quality", 86],
        ["Service", 53],
        ["Hygiene", 93],
        ["Availability", 95]
    ]
    # Place radar data starting at column E
    for i, row in enumerate(radar_data, start=1):
        ws_data.cell(row=i, column=5, value=row[0])
        ws_data.cell(row=i, column=6, value=row[1])

    # 3. Layout Dimensions & Canvas Painting
    ws.column_dimensions['A'].width = 8   # Sidebar
    ws.column_dimensions['B'].width = 3   # Margin
    for col in range(3, 11): # C through J
        ws.column_dimensions[get_column_letter(col)].width = 12

    sidebar_fill = PatternFill(start_color=palette["sidebar_bg"], fill_type="solid")
    canvas_fill = PatternFill(start_color=palette["canvas_bg"], fill_type="solid")
    card_fill = PatternFill(start_color=palette["card_bg"], fill_type="solid")
    thin_border = Side(border_style="thin", color=palette["border"])
    card_border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)

    # Paint the canvas and sidebar for the first 30 rows
    for row in ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=11):
        for cell in row:
            if cell.column == 1:
                cell.fill = sidebar_fill
            else:
                cell.fill = canvas_fill

    # Dashboard Title
    title_cell = ws['C1']
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=20, bold=True, color=palette["title_fg"])
    ws.row_dimensions[1].height = 30

    # Helper to draw KPI Cards
    def draw_kpi_card(min_col, min_row, max_col, max_row, header, value, subtext):
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = card_fill
                cell.border = card_border
        
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=min_row, end_column=max_col)
        ws.merge_cells(start_row=min_row+1, start_column=min_col, end_row=max_row-1, end_column=max_col)
        
        h_cell = ws.cell(row=min_row, column=min_col)
        h_cell.value = header
        h_cell.font = Font(bold=True, color=palette["text_muted"])
        h_cell.alignment = Alignment(vertical="center", indent=1)

        v_cell = ws.cell(row=min_row+1, column=min_col)
        v_cell.value = value
        v_cell.font = Font(size=24, bold=True, color=palette["text_main"])
        v_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 4. Construct KPI Cards
    draw_kpi_card(3, 3, 4, 6, "Sales", "$2,544", "M")         # C3:D6
    draw_kpi_card(5, 3, 6, 6, "Profit", "$890", "M")          # E3:F6
    draw_kpi_card(7, 3, 8, 6, "# of Customers", "87.0", "M")  # G3:H6

    # 5. Construct Chart Background Cards
    def draw_chart_bg(min_col, min_row, max_col, max_row):
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = card_fill
                cell.border = card_border

    draw_chart_bg(3, 8, 6, 22) # C8:F22 for Trend Line
    draw_chart_bg(7, 8, 10, 22) # G8:J22 for Radar Chart

    # 6. Trend Line Chart
    line_chart = LineChart()
    line_chart.title = "2021-2022 Sales Trend (in millions)"
    line_chart.style = 13
    line_chart.width = 13.5
    line_chart.height = 7.5
    
    data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=3, max_row=10)
    cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=10)
    line_chart.add_data(data_ref, titles_from_data=True)
    line_chart.set_categories(cats_ref)
    
    # Format markers as seen in the tutorial
    for series in line_chart.series:
        series.marker.symbol = "circle"
        series.marker.size = 5
        series.graphicalProperties.line.width = 25000

    ws.add_chart(line_chart, "C8")

    # 7. Customer Satisfaction Radar Chart
    radar_chart = RadarChart()
    radar_chart.type = "filled"
    radar_chart.title = "Customer Satisfaction"
    radar_chart.style = 26
    radar_chart.width = 11.5
    radar_chart.height = 7.5

    r_labels = Reference(ws_data, min_col=5, min_row=2, max_row=6)
    r_data = Reference(ws_data, min_col=6, min_row=1, max_row=6)
    
    radar_chart.add_data(r_data, titles_from_data=True)
    radar_chart.set_categories(r_labels)
    # Remove legend for a cleaner KPI look
    radar_chart.legend = None 

    ws.add_chart(radar_chart, "G8")
