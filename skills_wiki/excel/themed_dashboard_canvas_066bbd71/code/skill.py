def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", charts: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Theme Configuration
    palettes = {
        "corporate_blue": {"primary": "003366", "accent": "4F81BD", "text": "FFFFFF"},
        "modern_dark": {"primary": "262626", "accent": "00B050", "text": "FFFFFF"},
        "vibrant": {"primary": "5E17EB", "accent": "FF007F", "text": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        
    # 2. Clean Canvas Setup (The Dashboard Aesthetic)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    
    # 3. Themed Title Banner
    # Merging across the common viewable area (A1:R3 provides a wide top bar)
    ws.merge_cells("A1:R3")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["text"])
    title_cell.fill = PatternFill(fill_type="solid", fgColor=palette["primary"])
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Generate Fallback Charts if None Provided
    # This allows the skill to be tested independently and shows standard chart prep
    if not charts:
        charts = []
        # Store dummy calculation data out of sight (Column AA)
        data = [
            ["Month", "Revenue", "Profit"],
            ["Jan", 120000, 45000],
            ["Feb", 155000, 52000],
            ["Mar", 180000, 71000],
            ["Apr", 135000, 38000],
            ["May", 142000, 41000]
        ]
        for r_idx, row in enumerate(data, start=1):
            for c_idx, val in enumerate(row, start=27):
                ws.cell(row=r_idx, column=c_idx, value=val)
                
        # Bar Chart
        bc = BarChart()
        bc.title = "Revenue by Month"
        bc.style = 11  # Built-in Excel style
        data_ref = Reference(ws, min_col=28, min_row=1, max_row=6)
        cats_ref = Reference(ws, min_col=27, min_row=2, max_row=6)
        bc.add_data(data_ref, titles_from_data=True)
        bc.set_categories(cats_ref)
        charts.append(bc)
        
        # Line Chart
        lc = LineChart()
        lc.title = "Profit Trend"
        lc.style = 12
        data_ref2 = Reference(ws, min_col=29, min_row=1, max_row=6)
        lc.add_data(data_ref2, titles_from_data=True)
        lc.set_categories(cats_ref)
        charts.append(lc)
        
    # 5. Auto-Arrange Charts in a Uniform Grid
    col_start = 2       # Start at column B
    row_start = 5       # Start below the title banner
    charts_per_row = 2
    col_spacing = 10    # Columns to jump per chart horizontally
    row_spacing = 16    # Rows to jump per chart vertically
    
    for i, chart in enumerate(charts):
        # Enforce exact uniform sizing for a clean, aligned dashboard layout (sizes in cm)
        chart.width = 15
        chart.height = 7.5
        
        # Calculate grid position dynamically
        row_offset = i // charts_per_row
        col_offset = i % charts_per_row
        
        anchor_col = col_start + col_offset * col_spacing
        anchor_row = row_start + row_offset * row_spacing
        anchor = f"{get_column_letter(anchor_col)}{anchor_row}"
        
        # Inject chart into the presentation sheet
        ws.add_chart(chart, anchor)
