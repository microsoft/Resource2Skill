from openpyxl.styles import PatternFill, Font, Alignment

def render_workbook(wb, *, title: str = "Sales Dashboard 2024", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Transforms a workbook into an interactive, multi-sheet dashboard app
    with a persistent sidebar navigation menu.
    """
    # 1. Define the necessary application "views" (sheets)
    tabs = ["Dashboard", "Inputs", "Contacts"]
    
    # Clean up the default empty sheet if it's a fresh workbook
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb["Sheet"].title = tabs[0]
    
    # Ensure all target tabs exist
    for tab in tabs:
        if tab not in wb.sheetnames:
            wb.create_sheet(tab)
            
    # 2. Setup theme-driven styling parameters
    # (In a full implementation, these are extracted from the theme dictionary)
    sidebar_bg = "2C3E50"      # Dark slate primary
    sidebar_active = "34495E"  # Slightly lighter slate for the active state
    text_light = "ECF0F1"      # Off-white for high contrast
    title_color = "2C3E50"
    
    icon_font = Font(name="Segoe UI Emoji", size=18, color=text_light)
    active_font = Font(name="Segoe UI Emoji", size=20, color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Navigation items configuration: (Icon, Target Sheet Name)
    nav_menu = [
        ("🏠", "Dashboard"),
        ("📊", "Inputs"),
        ("👤", "Contacts")
    ]
    
    # 3. Apply the universal dashboard shell to every tab
    for current_tab in tabs:
        ws = wb[current_tab]
        
        # Disable gridlines for a clean canvas
        ws.sheet_view.showGridLines = False
        
        # Construct the physical sidebar (Column A)
        ws.column_dimensions['A'].width = 8
        base_fill = PatternFill(start_color=sidebar_bg, end_color=sidebar_bg, fill_type="solid")
        
        # Extend the sidebar fill down to cover a standard screen height
        for row in range(1, 45):
            ws.cell(row=row, column=1).fill = base_fill
            
        # Add dynamic Page Headers
        title_cell = ws.cell(row=2, column=3)
        title_cell.value = current_tab
        title_cell.font = Font(size=22, bold=True, color=title_color)
        
        subtitle_cell = ws.cell(row=3, column=3)
        subtitle_cell.value = title
        subtitle_cell.font = Font(size=12, italic=True, color="7F8C8D")
        
        # 4. Inject Interactive Navigation Icons
        start_row = 6
        spacing = 4
        
        for i, (icon, target) in enumerate(nav_menu):
            row_idx = start_row + (i * spacing)
            cell = ws.cell(row=row_idx, column=1)
            
            # Setup content and internal linking
            cell.value = icon
            cell.alignment = center_align
            cell.hyperlink = f"#'{target}'!A1"
            
            # Style the active vs inactive menu states
            if target == current_tab:
                cell.fill = PatternFill(start_color=sidebar_active, end_color=sidebar_active, fill_type="solid")
                cell.font = active_font
            else:
                cell.font = icon_font
