def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard 2022", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Border, Side, Font, Alignment

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    # 1. Canvas Setup: Turn off gridlines for a clean UI look
    ws.sheet_view.showGridLines = False

    # Theme Palette (fallback matching the video's dark/light aesthetic)
    colors = {
        "sidebar": "1C2A39",     # Dark navy for the navigation bar
        "canvas": "F3F4F6",      # Light gray for the main dashboard area
        "card_bg": "FFFFFF",     # White for the data cards
        "card_border": "E5E7EB", # Subtle border for cards
        "text_main": "111827",
        "text_muted": "6B7280"
    }

    # Apply global background colors
    fill_sidebar = PatternFill("solid", fgColor=colors["sidebar"])
    fill_canvas = PatternFill("solid", fgColor=colors["canvas"])

    # Paint 40 rows deep to cover the typical viewable area
    for row in range(1, 41):
        ws.cell(row=row, column=1).fill = fill_sidebar
        for col in range(2, 16):
            ws.cell(row=row, column=col).fill = fill_canvas

    # Set column widths for a grid system (A=sidebar; B,F,J=spacers)
    ws.column_dimensions['A'].width = 8
    for col_letter in ["B", "F", "J", "N"]:
        ws.column_dimensions[col_letter].width = 3
    for col_letter in ["C", "D", "E", "G", "H", "I", "K", "L", "M"]:
        ws.column_dimensions[col_letter].width = 11

    # 2. Card Generator Helper
    def draw_card(start_col, start_row, end_col, end_row, title_text):
        """Paints a white card with an outer border over the gray canvas."""
        fill_card = PatternFill("solid", fgColor=colors["card_bg"])
        border_side = Side(style="thin", color=colors["card_border"])

        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_card

                # Apply outer borders to create a bounding box effect
                cell.border = Border(
                    top=border_side if r == start_row else None,
                    bottom=border_side if r == end_row else None,
                    left=border_side if c == start_col else None,
                    right=border_side if c == end_col else None
                )

        # Insert Card Title in the top-left corner
        title_cell = ws.cell(row=start_row, column=start_col)
        title_cell.value = title_text
        title_cell.font = Font(bold=True, size=11, color=colors["text_main"])
        title_cell.alignment = Alignment(vertical="top", horizontal="left")

    # 3. Render Dashboard Structure (matching video quadrants)
    
    # Header Card
    draw_card(3, 2, 13, 3, title)
    # Override font size for the main dashboard title
    ws.cell(row=2, column=3).font = Font(bold=True, size=18, color=colors["text_main"])
    # Add subtitle
    ws.cell(row=3, column=3).value = "Figures in millions of USD"
    ws.cell(row=3, column=3).font = Font(italic=True, size=10, color=colors["text_muted"])

    # Top KPI Cards
    draw_card(3, 5, 5, 8, "Sales")
    draw_card(7, 5, 9, 8, "Profit")
    draw_card(11, 5, 13, 8, "# of Customers")

    # Main Visualization Cards
    draw_card(3, 10, 9, 21, "2021-2022 Sales Trend (in millions)")
    draw_card(11, 10, 13, 21, "Customer Satisfaction")
    draw_card(3, 23, 13, 34, "Sales by Country 2022")
