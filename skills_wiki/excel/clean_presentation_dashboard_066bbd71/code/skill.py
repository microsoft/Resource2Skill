from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Theme Colors (Fallback to corporate_blue)
    themes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF"},
        "dark_mode": {"bg": "262626", "fg": "E0E0E0"},
        "emerald": {"bg": "0F7A4D", "fg": "FFFFFF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])
    
    # 2. Construct Data Sheet
    data_ws = wb.active
    data_ws.title = "Chart Data"
    
    # Populate Stacked Column Data (Mock PivotTable output)
    data_ws.append(["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"])
    data_ws.append(["India", 62349, 23621, 21028])
    data_ws.append(["United States", 36657, 19446, 22260])
    data_ws.append(["United Kingdom", 46530, 26731, 11497])
    data_ws.append(["Philippines", 54618, 24567, 22005])
    
    # Populate Line Chart Data (Mock PivotTable output)
    data_ws.append([])
    data_ws.append(["Month", "Units Sold"])
    for m, val in [("Sep", 50601), ("Oct", 95622), ("Nov", 65481), ("Dec", 52970)]:
        data_ws.append([m, val])
        
    # 3. Construct Dashboard Sheet
    dash_ws = wb.create_sheet("Dashboard", 0)
    dash_ws.sheet_view.showGridLines = False
    
    # Create Title Banner
    dash_ws.merge_cells("B2:P4")
    title_cell = dash_ws["B2"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["fg"])
    title_cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Create and Place Charts
    # Stacked Column Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.height = 10
    bar_chart.width = 15
    
    bar_data = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=5)
    bar_cats = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    dash_ws.add_chart(bar_chart, "B6")
    
    # Line Chart
    line_chart = LineChart()
    line_chart.title = "Units sold each month"
    line_chart.height = 10
    line_chart.width = 15
    
    line_data = Reference(data_ws, min_col=2, min_row=7, max_row=11)
    line_cats = Reference(data_ws, min_col=1, min_row=8, max_row=11)
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)
    dash_ws.add_chart(line_chart, "J6")
    
    # 5. Hide the Data Sheet to create a clean user experience
    data_ws.sheet_state = 'hidden'
