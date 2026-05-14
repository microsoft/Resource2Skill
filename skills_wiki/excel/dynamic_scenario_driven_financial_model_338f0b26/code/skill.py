from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.worksheet.datavalidation import DataValidation
from skills_library.excel.components._helpers import get_theme_palette

def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Financial Model", base_year: int = 2024, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    palette = get_theme_palette(theme)
    
    header_fill = PatternFill(start_color=palette["primary"], fill_type="solid")
    header_font = Font(color=palette["bg"], bold=True)
    bold_font = Font(bold=True)
    
    # 1. Column Sizing
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 8
    for col in ['D', 'E', 'F', 'G', 'H']:
        ws.column_dimensions[col].width = 13
        
    # 2. Title & Control Panel
    ws["B2"] = title
    ws["B2"].font = Font(size=14, bold=True, color=palette["primary"])
    
    ws["J3"] = "Active Scenario:"
    ws["J3"].font = bold_font
    ws["J3"].alignment = Alignment(horizontal="right")
    ws["K3"] = 1
    
    # Dropdown validation
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["K3"])
    
    ws["K3"].fill = PatternFill(start_color="FFF2CC", fill_type="solid") # Yellow highlight
    ws["K3"].border = Border(
        outline=True, 
        top=Side(style='thin'), bottom=Side(style='thin'), 
        left=Side(style='thin'), right=Side(style='thin')
    )
    ws["K3"].alignment = Alignment(horizontal="center")
    
    # 3. Headers
    headers = ["Line Item", "Unit", f"{base_year}", f"{base_year+1}", f"{base_year+2}", f"{base_year+3}", f"{base_year+4}"]
    for col_idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=4, column=col_idx)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    # 4. Income Statement Section
    ws["B5"] = "Income Statement"
    ws["B5"].font = bold_font
    
    row_labels = [
        ("Revenue", "$"),
        ("COGS", "$"),
        ("Gross Profit", "$"),
        ("Gross Margin", "%"),
        ("Operating Expenses", "$"),
        ("Operating Profit", "$"),
        ("Tax", "$"),
        ("Net Profit", "$")
    ]
    
    r_is = 6
    for i, (lbl, unit) in enumerate(row_labels):
        ws.cell(row=r_is+i, column=2, value=lbl)
        ws.cell(row=r_is+i, column=3, value=unit).alignment = Alignment(horizontal="center")
        
    # Apply formulas dependent strictly on the "Live" rows (16-21)
    for c in range(4, 9): 
        col_let = ws.cell(row=1, column=c).column_letter
        
        ws[f"{col_let}6"] = f"={col_let}16*{col_let}18"  # Revenue = Vol * Price
        ws[f"{col_let}7"] = f"={col_let}16*{col_let}19"  # COGS = Vol * Unit COGS
        ws[f"{col_let}8"] = f"={col_let}6-{col_let}7"    # Gross Profit
        ws[f"{col_let}8"].font = bold_font
        ws[f"{col_let}9"] = f"=IFERROR({col_let}8/{col_let}6, 0)" # Margin
        ws[f"{col_let}9"].number_format = numbers.FORMAT_PERCENTAGE_00
        ws[f"{col_let}10"] = f"={col_let}20"             # Opex
        ws[f"{col_let}11"] = f"={col_let}8-{col_let}10"  # Operating Profit
        ws[f"{col_let}11"].font = bold_font
        ws[f"{col_let}12"] = f"=MAX({col_let}11*{col_let}21, 0)" # Tax
        ws[f"{col_let}13"] = f"={col_let}11-{col_let}12" # Net Profit
        ws[f"{col_let}13"].font = bold_font
        
        # Accounting borders
        ws[f"{col_let}8"].border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
        ws[f"{col_let}11"].border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
        ws[f"{col_let}13"].border = Border(top=Side(style='thin'), bottom=Side(style='double'))
        
        for r in [6, 7, 8, 10, 11, 12, 13]:
            ws.cell(row=r, column=c).number_format = '#,##0'

    # 5. Live Assumptions Engine
    r_live = 15
    ws.cell(row=r_live, column=2, value="Live Assumptions").font = Font(bold=True, color=palette["bg"])
    ws.cell(row=r_live, column=2).fill = PatternFill(start_color=palette["primary"], fill_type="solid")
    
    assmp_labels = [
        ("Volume", "#"),
        ("Volume Growth", "%"),
        ("Price", "$"),
        ("Unit COGS", "$"),
        ("Fixed Opex", "$"),
        ("Tax Rate", "%")
    ]
    
    for i, (lbl, unit) in enumerate(assmp_labels):
        ws.cell(row=r_live+1+i, column=2, value=lbl)
        ws.cell(row=r_live+1+i, column=3, value=unit).alignment = Alignment(horizontal="center")
        
    r_scen1 = 24
    r_scen2 = 33
    
    # Inject the core CHOOSE mechanism
    for c in range(4, 9):
        col_let = ws.cell(row=1, column=c).column_letter
        for i, (lbl, unit) in enumerate(assmp_labels):
            row = r_live + 1 + i
            ws[f"{col_let}{row}"] = f"=CHOOSE($K$3, {col_let}{r_scen1+1+i}, {col_let}{r_scen2+1+i})"
            
            if unit == "%":
                ws[f"{col_let}{row}"].number_format = numbers.FORMAT_PERCENTAGE_00
            elif unit == "$":
                ws[f"{col_let}{row}"].number_format = '#,##0.00' if lbl in ["Price", "Unit COGS"] else '#,##0'
            else:
                ws[f"{col_let}{row}"].number_format = '#,##0'
                
    # 6. Scenario Injector
    def fill_scenario(start_r, title_str, base_vol, growth, price, cogs, opex, tax):
        ws.cell(row=start_r, column=2, value=title_str).font = bold_font
        for i, (lbl, unit) in enumerate(assmp_labels):
            ws.cell(row=start_r+1+i, column=2, value=lbl)
            ws.cell(row=start_r+1+i, column=3, value=unit).alignment = Alignment(horizontal="center")
            
        for idx, col in enumerate(['D', 'E', 'F', 'G', 'H']):
            # Growth (Hardcoded = Blue)
            ws[f"{col}{start_r+2}"] = growth[idx]
            ws[f"{col}{start_r+2}"].number_format = '0%'
            ws[f"{col}{start_r+2}"].font = Font(color="0000FF")
            
            # Volume (Dynamic via previous year)
            if idx == 0:
                ws[f"{col}{start_r+1}"] = base_vol
                ws[f"{col}{start_r+1}"].font = Font(color="0000FF")
            else:
                prev_col = chr(ord(col)-1)
                ws[f"{col}{start_r+1}"] = f"={prev_col}{start_r+1}*(1+{col}{start_r+2})"
                ws[f"{col}{start_r+1}"].font = Font(color="000000") # Calculated
            ws[f"{col}{start_r+1}"].number_format = '#,##0'
            
            # Price, COGS, Opex, Tax (Hardcoded = Blue)
            ws[f"{col}{start_r+3}"] = price[idx]
            ws[f"{col}{start_r+3}"].font = Font(color="0000FF")
            ws[f"{col}{start_r+3}"].number_format = '#,##0.00'
            
            ws[f"{col}{start_r+4}"] = cogs[idx]
            ws[f"{col}{start_r+4}"].font = Font(color="0000FF")
            ws[f"{col}{start_r+4}"].number_format = '#,##0.00'
            
            ws[f"{col}{start_r+5}"] = opex[idx]
            ws[f"{col}{start_r+5}"].font = Font(color="0000FF")
            ws[f"{col}{start_r+5}"].number_format = '#,##0'
            
            ws[f"{col}{start_r+6}"] = tax[idx]
            ws[f"{col}{start_r+6}"].font = Font(color="0000FF")
            ws[f"{col}{start_r+6}"].number_format = '0%'

    # Populate Scenario 1 (Aggressive)
    fill_scenario(r_scen1, "Scenario 1: Best Case", 3000, 
                  [0, 1.0, 0.75, 0.50, 0.35], [39.95]*5, [8.75]*5, 
                  [100000, 100000, 185000, 235000, 235000], [0.2]*5)
                  
    # Populate Scenario 2 (Conservative)
    fill_scenario(r_scen2, "Scenario 2: Worst Case", 2000, 
                  [0, 0.20, 0.15, 0.10, 0.05], [34.95]*5, [10.25]*5, 
                  [100000]*5, [0.25]*5)
