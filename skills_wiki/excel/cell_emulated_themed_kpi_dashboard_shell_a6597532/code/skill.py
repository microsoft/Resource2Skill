from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list[dict], theme: str = "purple_gold", **kwargs) -> None:
    """
    Renders a two-tone dashboard shell with a top strip of cell-emulated KPI cards.
    
    :param kpis: List of dicts, e.g., [{"label": "CALLS", "value": "16,749", "icon": "📞"}, ...]
    """
    ws = wb.create_sheet(sheet_name)
    
    # Theme palette resolution (fallback matches the video's aesthetic)
    palettes = {
        "purple_gold": {
            "header_bg": "4B2E5D", # Dark purple
            "body_bg": "F2EDF4",   # Light pastel purple
            "card_bg": "FFFFFF",   # White
            "accent": "F4B41A",    # Gold/Yellow
            "text_light": "FFFFFF",
            "text_dark": "333333",
            "border": "D0C9D6"
        },
        "corporate_blue": {
            "header_bg": "2C3E50",
            "body_bg": "ECF0F1",
            "card_bg": "FFFFFF",
            "accent": "F39C12",
            "text_light": "FFFFFF",
            "text_dark": "2C3E50",
            "border": "BDC3C7"
        }
    }
    t = palettes.get(theme, palettes["purple_gold"])
    
    fill_header = PatternFill("solid", fgColor=t["header_bg"])
    fill_body = PatternFill("solid", fgColor=t["body_bg"])
    fill_card = PatternFill("solid", fgColor=t["card_bg"])
    fill_accent = PatternFill("solid", fgColor=t["accent"])
    
    thin_border = Border(
        left=Side(style='thin', color=t["border"]),
        right=Side(style='thin', color=t["border"]),
        top=Side(style='thin', color=t["border"]),
        bottom=Side(style='thin', color=t["border"])
    )

    # 1. Paint Background
    for row in range(1, 9):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = fill_header
    for row in range(9, 45):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = fill_body

    # 2. Add Title & Subtitle
    ws["B2"] = title
    ws["B2"].font = Font(size=26, color=t["text_light"], bold=True)
    ws["B3"] = subtitle
    ws["B3"].font = Font(size=14, color=t["accent"])

    # 3. Construct Cell-Emulated KPI Cards
    # Adjust heights to shape the cards
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20
    ws.row_dimensions[6].height = 18
    
    start_col = 5 # Start cards at column E
    for i, kpi in enumerate(kpis):
        # Calculate matrix coordinates for this card
        col_idx = start_col + (i * 4) # 4 columns per card
        
        # Adjust column widths for card proportions
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 6   # Icon col
        ws.column_dimensions[ws.cell(row=1, column=col_idx+1).column_letter].width = 9 # Val col 1
        ws.column_dimensions[ws.cell(row=1, column=col_idx+2).column_letter].width = 9 # Val col 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx+3).column_letter].width = 3 # Spacer
        
        # --- LEFT PANEL (Icon / Accent Bar) ---
        ws.merge_cells(start_row=4, start_column=col_idx, end_row=6, end_column=col_idx)
        icon_cell = ws.cell(row=4, column=col_idx)
        icon_cell.value = kpi.get("icon", "")
        icon_cell.fill = fill_accent
        icon_cell.font = Font(size=20, color=t["text_dark"])
        icon_cell.alignment = Alignment(horizontal="center", vertical="center")
        icon_cell.border = thin_border
        
        # --- RIGHT TOP (Value) ---
        ws.merge_cells(start_row=4, start_column=col_idx+1, end_row=5, end_column=col_idx+2)
        val_cell = ws.cell(row=4, column=col_idx+1)
        val_cell.value = kpi.get("value", 0)
        val_cell.fill = fill_card
        val_cell.font = Font(size=22, color=t["text_dark"], bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # --- RIGHT BOTTOM (Label) ---
        ws.merge_cells(start_row=6, start_column=col_idx+1, end_row=6, end_column=col_idx+2)
        lbl_cell = ws.cell(row=6, column=col_idx+1)
        lbl_cell.value = str(kpi.get("label", "")).upper()
        lbl_cell.fill = fill_card
        lbl_cell.font = Font(size=10, color=t["text_dark"])
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")
        
        # Apply borders around the right panel pieces
        for r in range(4, 7):
            for c in range(col_idx+1, col_idx+3):
                ws.cell(row=r, column=c).border = thin_border
                
    # Hide gridlines for a clean dashboard look
    ws.sheet_view.showGridLines = False
