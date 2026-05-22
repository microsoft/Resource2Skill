def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    
    # Clear default sheets
    for sheet in wb.sheetnames:
        wb.remove(wb[sheet])
        
    pages = ["Dashboard", "Inputs", "Contacts"]
    
    # Theme palette (simulating a clean, modern UI)
    sidebar_bg = "1E293B" # Dark slate
    sidebar_fg = "FFFFFF"
    body_bg = "F3F4F6"    # Light gray
    card_bg = "FFFFFF"
    card_border = "E2E8F0"
    text_main = "0F172A"
    
    for page in pages:
        ws = wb.create_sheet(title=page)
        ws.sheet_view.showGridLines = False
        
        # 1. Set Body Background
        for col_letter in "BCDEFGHIJKLMNO":
            ws.column_dimensions[col_letter].width = 12
            
        for row in ws.iter_rows(min_row=1, max_row=40, min_col=2, max_col=15):
            for cell in row:
                cell.fill = PatternFill(start_color=body_bg, end_color=body_bg, fill_type="solid")
                
        # 2. Setup Sidebar
        ws.column_dimensions['A'].width = 18
        for row in range(1, 41):
            cell = ws.cell(row=row, column=1)
            cell.fill = PatternFill(start_color=sidebar_bg, end_color=sidebar_bg, fill_type="solid")
            
        # 3. Add Navigation Links to Sidebar
        menu_header = ws.cell(row=3, column=1, value="MAIN MENU")
        menu_header.font = Font(color="94A3B8", bold=True, size=10)
        menu_header.alignment = Alignment(horizontal="center")
        
        for i, target_page in enumerate(pages):
            nav_cell = ws.cell(row=5 + (i * 2), column=1, value=target_page)
            # Highlight current page
            if target_page == page:
                nav_cell.font = Font(color=sidebar_bg, bold=True)
                nav_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            else:
                nav_cell.font = Font(color=sidebar_fg, underline="single")
                # Internal hyperlink syntax: #'SheetName'!A1
                nav_cell.hyperlink = f"#'{target_page}'!A1"
                
            nav_cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # 4. Add Page Header
        title_cell = ws.cell(row=2, column=3, value=f"{title} - {page}")
        title_cell.font = Font(size=20, bold=True, color=text_main)
        
        # 5. Dashboard Specific Layout (Cards)
        if page == "Dashboard":
            # [min_col, max_col, min_row, max_row, Title, Value]
            cards = [
                (3, 5, 5, 9, "Total Sales", "$2,544M"),
                (6, 8, 5, 9, "Total Profit", "$890M"),
                (9, 11, 5, 9, "Active Customers", "87.0K"),
                (3, 7, 11, 22, "2021-2022 Sales Trend", "View Chart Area"),
                (8, 11, 11, 22, "Customer Satisfaction", "View Radar Area")
            ]
            
            for min_c, max_c, min_r, max_r, c_title, c_val in cards:
                # Draw Card Background and Border
                for r in range(min_r, max_r + 1):
                    for c in range(min_c, max_c + 1):
                        cell = ws.cell(row=r, column=c)
                        cell.fill = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
                        
                        # Add border exclusively to the outer edges of the card range
                        b_left = Side(style='thin', color=card_border) if c == min_c else None
                        b_right = Side(style='thin', color=card_border) if c == max_c else None
                        b_top = Side(style='thin', color=card_border) if r == min_r else None
                        b_bottom = Side(style='thin', color=card_border) if r == max_r else None
                        cell.border = Border(left=b_left, right=b_right, top=b_top, bottom=b_bottom)
                
                # Insert Card Content
                header_cell = ws.cell(row=min_r + 1, column=min_c + 1, value=c_title)
                header_cell.font = Font(bold=True, color="64748B")
                
                val_cell = ws.cell(row=min_r + 3, column=min_c + 1, value=c_val)
                val_cell.font = Font(size=18, bold=True, color=text_main)
