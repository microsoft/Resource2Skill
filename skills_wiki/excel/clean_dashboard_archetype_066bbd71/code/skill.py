from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def get_theme_palette(theme_name: str) -> dict:
    """Mock theme loader to satisfy self-contained requirement."""
    palettes = {
        "corporate_blue": {
            "primary": "003366", 
            "accent": "4F81BD", 
            "bg": "FFFFFF", 
            "text": "000000", 
            "header_bg": "003366", 
            "header_fg": "FFFFFF"
        }
    }
    return palettes.get(theme_name, palettes["corporate_blue"])

def render_workbook(wb, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    palette = get_theme_palette(theme)
    
    # 1. Create Isolated Data Sheet
    ws_data = wb.create_sheet("Dashboard Data")
    
    # Sample Data for Stacked Column Chart (Profit by Market & Cookie)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["United States", 36657, 6368, 22260],
        ["United Kingdom", 46530, 5220, 11497],
    ]
    for r in market_data:
        ws_data.append(r)
        
    # Leave blank rows between datasets to keep references clean
    ws_data.append([])
    ws_data.append([])
    
    # Sample Data for Line Chart (Units sold each month)
    start_row_line = len(market_data) + 3  # Calculated start row for the second block
    line_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970],
    ]
    for r in line_data:
        ws_data.append(r)
        
    # 2. Setup Dashboard Presentation Sheet
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    
    # Disable gridlines to make it look like a standalone dashboard app
    ws_dash.sheet_view.showGridLines = False
    
    # Create Title Banner
    ws_dash.merge_cells("B2:M4")
    title_cell = ws_dash["B2"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["header_fg"])
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    for row in ws_dash["B2:M4"]:
        for cell in row:
            cell.fill = fill
            
    # 3. Create Stacked Column Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.width = 15
    bar_chart.height = 10
    
    data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=len(market_data))
    cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=len(market_data))
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    
    ws_dash.add_chart(bar_chart, "B6")
    
    # 4. Create Line Chart
    line_chart = LineChart()
    line_chart.title = "Units sold each month"
    line_chart.style = 13
    line_chart.width = 15
    line_chart.height = 10
    
    l_data_ref = Reference(ws_data, min_col=2, min_row=start_row_line, max_col=2, max_row=start_row_line + len(line_data) - 1)
    l_cats_ref = Reference(ws_data, min_col=1, min_row=start_row_line + 1, max_row=start_row_line + len(line_data) - 1)
    line_chart.add_data(l_data_ref, titles_from_data=True)
    line_chart.set_categories(l_cats_ref)
    
    ws_dash.add_chart(line_chart, "H6")
