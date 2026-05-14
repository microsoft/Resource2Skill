from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Financial Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # 1. Styles & Themes
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    highlight_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), 
                         top=Side(style="thin"), bottom=Side(style="thin"))
    
    # Number Formats
    fmt_currency = '"$"#,##0.00'
    fmt_whole = '#,##0'
    fmt_pct = '0.0%'

    # 2. Scenario Toggle
    ws["F2"] = "Scenario:"
    ws["F2"].font = bold_font
    ws["F2"].alignment = Alignment(horizontal="right")
    ws["G2"] = 1
    ws["G2"].fill = highlight_fill
    ws["G2"].border = thin_border
    ws["G2"].alignment = Alignment(horizontal="center")
    
    # Data Validation for Toggle (1 = Upper Case, 2 = Lower Case)
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["G2"])
    
    # 3. Headers
    headers = ["Items", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for col_idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    ws.column_dimensions["B"].width = 25
    for c in range(3, 8):
        ws.column_dimensions[get_column_letter(c)].width = 15

    # 4. Income Statement Block
    ws["B5"] = "Income Statement"
    ws["B5"].font = bold_font
    
    is_labels = ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "Operating Profit", "Corporate Tax", "Net Profit"]
    for i, label in enumerate(is_labels, start=6):
        ws.cell(row=i, column=2, value=label)
        
    ws["B8"].font = bold_font
    ws["B10"].font = bold_font
    ws["B12"].font = bold_font
    
    for col in range(3, 8):
        col_ltr = get_column_letter(col)
        # Revenue = Orders * AOV
        ws[f"{col_ltr}6"] = f"={col_ltr}16*{col_ltr}17"
        # COGS = Orders * COGS/order
        ws[f"{col_ltr}7"] = f"={col_ltr}16*{col_ltr}18"
        # Gross Profit
        ws[f"{col_ltr}8"] = f"={col_ltr}6-{col_ltr}7"
        # Operating Expenses
        ws[f"{col_ltr}9"] = f"={col_ltr}19"
        # Operating Profit
        ws[f"{col_ltr}10"] = f"={col_ltr}8-{col_ltr}9"
        # Tax (Don't tax losses)
        ws[f"{col_ltr}11"] = f"=IF({col_ltr}10>0, {col_ltr}10*{col_ltr}20, 0)"
        # Net Profit
        ws[f"{col_ltr}12"] = f"={col_ltr}10-{col_ltr}11"
        
        # Formatting
        for r in [6, 7, 8, 9, 10, 11, 12]:
            ws[f"{col_ltr}{r}"].number_format = fmt_currency

    # 5. Live Case Assumptions Block (Driven by CHOOSE)
    ws["B15"] = "Assumptions (Live Case)"
    ws["B15"].font = header_font
    ws["B15"].fill = header_fill
    
    assumptions_labels = ["Orders", "AOV", "COGS per order", "Fixed Opex", "Tax Rate"]
    for i, label in enumerate(assumptions_labels, start=16):
        ws.cell(row=i, column=2, value=label)
        for col in range(3, 8):
            col_ltr = get_column_letter(col)
            # CHOOSE dynamically pulls from Scenario 1 (Upper) or Scenario 2 (Lower)
            ws[f"{col_ltr}{i}"] = f"=CHOOSE($G$2, {col_ltr}{i+8}, {col_ltr}{i+16})"
            
            # Format row
            if i == 16: ws[f"{col_ltr}{i}"].number_format = fmt_whole
            elif i in [17, 18, 19]: ws[f"{col_ltr}{i}"].number_format = fmt_currency
            elif i == 20: ws[f"{col_ltr}{i}"].number_format = fmt_pct

    # 6. Setup Hardcoded Scenarios
    scenarios = [
        {
            "name": "Scenario 1 (Upper Case)",
            "start_row": 23,
            "data": [
                [3000, 6000, 10500, 15750, 21263],  # Orders
                [39.95, 39.95, 39.95, 39.95, 39.95], # AOV
                [8.75, 8.75, 8.75, 8.75, 8.75],      # COGS
                [100000, 100000, 185000, 235000, 235000], # Opex
                [0.2, 0.2, 0.2, 0.2, 0.2]            # Tax
            ]
        },
        {
            "name": "Scenario 2 (Lower Case)",
            "start_row": 31,
            "data": [
                [2000, 4000, 7000, 10500, 14175],  # Orders
                [34.95, 34.95, 34.95, 34.95, 34.95], # AOV
                [8.75, 8.75, 8.75, 8.75, 8.75],      # COGS
                [100000, 100000, 185000, 235000, 235000], # Opex
                [0.25, 0.25, 0.25, 0.25, 0.25]       # Tax
            ]
        }
    ]

    for scn in scenarios:
        r_start = scn["start_row"]
        ws[f"B{r_start}"] = scn["name"]
        ws[f"B{r_start}"].font = bold_font
        
        for i, label in enumerate(assumptions_labels, start=r_start+1):
            ws.cell(row=i, column=2, value=label)
            row_data = scn["data"][i - (r_start + 1)]
            for col_idx, val in enumerate(row_data, start=3):
                cell = ws.cell(row=i, column=col_idx, value=val)
                # Apply matching formats
                if label == "Orders": cell.number_format = fmt_whole
                elif label == "Tax Rate": cell.number_format = fmt_pct
                else: cell.number_format = fmt_currency
