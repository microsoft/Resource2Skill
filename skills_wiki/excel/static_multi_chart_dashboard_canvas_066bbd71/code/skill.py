from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, Alignment

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Dashboard Sheet (Canvas)
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    # Theme palette fallback
    colors = {
        "corporate_blue": {"primary": "002060", "text": "000000"}
    }.get(theme, {"primary": "002060", "text": "000000"})

    # Dashboard Title Setup
    ws_dash["B1"] = title
    ws_dash["B1"].font = Font(size=24, bold=True, color=colors["primary"])
    ws_dash.merge_cells("B1:N2")
    ws_dash["B1"].alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Setup Hidden Data Sheet
    ws_data = wb.create_sheet("Data")
    ws_data.sheet_state = "hidden"
    
    # --- Chart 1 Data: Profit by Market & Cookie (Stacked Column) ---
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62000, 23000, 21000],
        ["United States", 36000, 32000, 11000],
        ["United Kingdom", 46000, 26000, 14000]
    ]
    for r in market_data:
        ws_data.append(r)
        
    # Format as currency so the chart axis inherits it
    for row in ws_data.iter_rows(min_row=2, max_row=4, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = '$#,##0'
            
    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Cookie Type"
    data1 = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=4)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=4)
    c1.add_data(data1, titles_from_data=True)
    c1.set_categories(cats1)
    c1.height = 14
    c1.width = 16
    ws_dash.add_chart(c1, "B4")
    
    # --- Chart 2 Data: Units sold each month ---
    units_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for r in units_data:
        ws_data.append(r)
        
    for row in ws_data.iter_rows(min_row=6, max_row=9, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = '#,##0'
            
    c2 = LineChart()
    c2.title = "Units sold each month"
    c2.legend = None  # Replicates removing the default "Total" legend in the video
    data2 = Reference(ws_data, min_col=2, min_row=5, max_col=2, max_row=9)
    cats2 = Reference(ws_data, min_col=1, min_row=6, max_row=9)
    c2.add_data(data2, titles_from_data=True)
    c2.set_categories(cats2)
    c2.height = 7
    c2.width = 13
    ws_dash.add_chart(c2, "J4")
    
    # --- Chart 3 Data: Profit by month ---
    profit_data = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for r in profit_data:
        ws_data.append(r)
        
    for row in ws_data.iter_rows(min_row=11, max_row=14, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = '$#,##0'
            
    c3 = LineChart()
    c3.title = "Profit by month"
    c3.legend = None
    data3 = Reference(ws_data, min_col=2, min_row=10, max_col=2, max_row=14)
    cats3 = Reference(ws_data, min_col=1, min_row=11, max_row=14)
    c3.add_data(data3, titles_from_data=True)
    c3.set_categories(cats3)
    c3.height = 7
    c3.width = 13
    ws_dash.add_chart(c3, "J18")
