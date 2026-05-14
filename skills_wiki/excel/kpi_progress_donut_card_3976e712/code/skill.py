from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils import get_column_letter

def render(ws, anchor: str, *, title: str = "Sales Target", actual: float = 2544, target: float = 3000, theme: str = "corporate_blue", **kwargs) -> None:
    row, col = coordinate_to_tuple(anchor)
    
    # Theme palette fallback hooks
    accent_color = kwargs.get("accent_color", "1F4E78")
    text_color = kwargs.get("text_color", "000000")
    muted_color = kwargs.get("muted_color", "595959")
    card_bg = kwargs.get("card_bg", "FFFFFF")
    chart_gray = "E7E6E6"
    border_color = "D9D9D9"
    
    # 1. Write Data (off-screen, 20 cols to the right to keep dashboard clean)
    data_col = col + 20
    pct = actual / target if target else 0
    rem = max(0, 1 - pct)
    
    ws.cell(row=row, column=data_col, value="Category")
    ws.cell(row=row+1, column=data_col, value="Actual")
    ws.cell(row=row+2, column=data_col, value="Remainder")
    
    ws.cell(row=row, column=data_col+1, value="Value")
    ws.cell(row=row+1, column=data_col+1, value=pct)
    ws.cell(row=row+2, column=data_col+1, value=rem)
    
    # 2. Format Card Area (5 rows tall, 4 columns wide)
    thin_side = Side(style='thin', color=border_color)
    card_fill = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
    
    for r in range(row, row+5):
        ws.row_dimensions[r].height = 18
        for c in range(col, col+4):
            cell = ws.cell(row=r, column=c)
            cell.fill = card_fill
            
            # Apply simple outer box border
            top_side = thin_side if r == row else None
            bottom_side = thin_side if r == row+4 else None
            left_side = thin_side if c == col else None
            right_side = thin_side if c == col+3 else None
            
            cell.border = Border(top=top_side, bottom=bottom_side, left=left_side, right=right_side)

    # 3. Write Text Content Hierarchy
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
    t_cell = ws.cell(row=row, column=col, value=title)
    t_cell.font = Font(size=11, color=muted_color, bold=True)
    t_cell.alignment = Alignment(vertical="center")
    
    ws.merge_cells(start_row=row+1, start_column=col, end_row=row+2, end_column=col+1)
    val_cell = ws.cell(row=row+1, column=col, value=actual)
    val_cell.font = Font(size=20, color=text_color, bold=True)
    val_cell.number_format = '$#,##0'
    val_cell.alignment = Alignment(vertical="center", horizontal="left")
    
    ws.merge_cells(start_row=row+3, start_column=col, end_row=row+3, end_column=col+1)
    sub_cell = ws.cell(row=row+3, column=col, value=f"Target: ${target:,.0f}")
    sub_cell.font = Font(size=9, color=muted_color)
    sub_cell.alignment = Alignment(vertical="center")
    
    ws.merge_cells(start_row=row+4, start_column=col, end_row=row+4, end_column=col+1)
    pct_cell = ws.cell(row=row+4, column=col, value=f"{pct*100:.1f}% Complete")
    pct_cell.font = Font(size=9, color=accent_color, bold=True)
    pct_cell.alignment = Alignment(vertical="center")

    # 4. Insert Doughnut Chart
    chart = DoughnutChart()
    data = Reference(ws, min_col=data_col+1, min_row=row, max_row=row+2)
    cats = Reference(ws, min_col=data_col, min_row=row+1, max_row=row+2)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    chart.title = None
    chart.legend = None
    chart.holeSize = 70
    
    # Resize and position chart to fit in right side of the card
    chart.width = 3.5  
    chart.height = 2.5 
    
    # 5. Format Slices (Force colors using DataPoints)
    series = chart.series[0]
    
    dp_actual = DataPoint(idx=0)
    dp_actual.graphicalProperties.solidFill = accent_color 
    
    dp_rem = DataPoint(idx=1)
    dp_rem.graphicalProperties.solidFill = chart_gray
    
    series.dPt.append(dp_actual)
    series.dPt.append(dp_rem)
    
    # Anchor the chart into the card
    chart_anchor = f"{get_column_letter(col+2)}{row}"
    ws.add_chart(chart, chart_anchor)
