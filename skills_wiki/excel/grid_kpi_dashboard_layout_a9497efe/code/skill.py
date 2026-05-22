def render_sheet(wb, sheet_name: str, *, title: str = "KPI Dashboard", theme: str = "corporate_blue", data: list = None, **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter

    if not data:
        data = [
            {
                "section": "Working Capital Efficiency",
                "kpis": [
                    {"name": "DSO (Days Sales Outstanding)", "value": 41, "target": 45, "prior": 53, "format": "0", "good_dir": "down"},
                    {"name": "DPO (Days Payables Outstanding)", "value": 90, "target": 90, "prior": 89, "format": "0", "good_dir": "up"},
                    {"name": "Non-Current AR %", "value": 0.12, "target": 0.03, "prior": 0.11, "format": "0%", "good_dir": "down"}
                ]
            },
            {
                "section": "Sales KPIs",
                "kpis": [
                    {"name": "CAC (Customer Acquisition)", "value": 14444, "target": 15000, "prior": 13750, "format": "$#,##0", "good_dir": "down"},
                    {"name": "Sales vs. Budget %", "value": 0.92, "target": 1.00, "prior": 0.91, "format": "0%", "good_dir": "up"},
                    {"name": "Gross Margin", "value": 0.40, "target": 0.38, "prior": 0.40, "format": "0%", "good_dir": "up"}
                ]
            }
        ]

    # Theme setup with fallback logic
    themes = {
        "corporate_blue": {
            "header_bg": "2F5597", "header_fg": "FFFFFF",
            "section_bg": "D9E1F2", "section_fg": "000000",
            "kpi_bg": "F2F2F2",
            "good_bg": "C6EFCE", "good_fg": "006100",
            "bad_bg": "FFC7CE", "bad_fg": "9C0006"
        }
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # Fills, Fonts, and Borders
    header_fill = PatternFill("solid", fgColor=palette["header_bg"])
    header_font = Font(color=palette["header_fg"], bold=True, size=18)
    section_fill = PatternFill("solid", fgColor=palette["section_bg"])
    section_font = Font(color=palette["section_fg"], bold=True, size=14)
    kpi_title_fill = PatternFill("solid", fgColor=palette["kpi_bg"])
    kpi_title_font = Font(bold=True, size=11)
    val_font = Font(size=24, bold=True)
    sub_font = Font(size=9, color="595959")

    good_fill = PatternFill("solid", fgColor=palette["good_bg"])
    good_font = Font(color=palette["good_fg"], size=24, bold=True)
    bad_fill = PatternFill("solid", fgColor=palette["bad_bg"])
    bad_font = Font(color=palette["bad_fg"], size=24, bold=True)

    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Main Dashboard Title
    ws["B2"] = title
    ws["B2"].font = header_font
    ws["B2"].fill = header_fill
    ws.merge_cells("B2:M3")
    ws["B2"].alignment = center_align

    # Setup 3-wide flex column widths (A: Spacer, [B, C, D]: KPI1, E: Spacer...)
    ws.column_dimensions["A"].width = 3
    for col in ["B", "C", "D", "F", "G", "H", "J", "K", "L"]:
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["E"].width = 3  # Spacer
    ws.column_dimensions["I"].width = 3  # Spacer
    ws.column_dimensions["M"].width = 3  # Right margin

    current_row = 5

    for section in data:
        # 1. Section Header (Spanning full grid width B:L)
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=12)
        cell = ws.cell(row=current_row, column=2, value=section.get("section", "Metrics"))
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = center_align
        current_row += 2

        kpis = section.get("kpis", [])
        
        # 2. Render up to 3 KPIs per grid row
        for idx, kpi in enumerate(kpis[:3]):
            start_col = 2 + (idx * 4) # Iterates to columns 2(B), 6(F), 10(J)
            end_col = start_col + 2

            # KPI Header Banner
            ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
            name_cell = ws.cell(row=current_row, column=start_col, value=kpi["name"])
            name_cell.fill = kpi_title_fill
            name_cell.font = kpi_title_font
            name_cell.alignment = center_align
            
            for c in range(start_col, end_col + 1):
                ws.cell(row=current_row, column=c).border = Border(top=thin_border.top, left=thin_border.left if c == start_col else None, right=thin_border.right if c == end_col else None)

            # KPI Main Value Field
            val_row = current_row + 1
            ws.merge_cells(start_row=val_row, start_column=start_col, end_row=val_row, end_column=end_col)
            val_cell = ws.cell(row=val_row, column=start_col, value=kpi["value"])
            val_cell.font = val_font
            val_cell.alignment = center_align
            val_cell.number_format = kpi.get("format", "0")
            
            for c in range(start_col, end_col + 1):
                ws.cell(row=val_row, column=c).border = Border(left=thin_border.left if c == start_col else None, right=thin_border.right if c == end_col else None)

            # KPI Sub-metrics (Vs Target / Vs Prior)
            sub_row = current_row + 2
            t_lbl_cell = ws.cell(row=sub_row, column=start_col, value="Target:")
            t_lbl_cell.font = sub_font
            t_lbl_cell.alignment = Alignment(horizontal="left")

            t_val_cell = ws.cell(row=sub_row, column=start_col+1, value=kpi["target"])
            t_val_cell.font = Font(size=9, bold=True)
            t_val_cell.number_format = kpi.get("format", "0")
            t_val_cell.alignment = Alignment(horizontal="left")

            p_lbl_cell = ws.cell(row=sub_row, column=start_col+2, value=f"Prior: {kpi.get('prior', '-')}")
            p_lbl_cell.font = sub_font
            p_lbl_cell.alignment = Alignment(horizontal="right")
            
            for c in range(start_col, end_col + 1):
                ws.cell(row=sub_row, column=c).border = Border(bottom=thin_border.bottom, left=thin_border.left if c == start_col else None, right=thin_border.right if c == end_col else None)

            # Link Conditional Formatting to Target Cell
            val_coord = val_cell.coordinate
            target_coord = f"${get_column_letter(t_val_cell.column)}${t_val_cell.row}"

            if kpi.get("good_dir") == "up":
                ws.conditional_formatting.add(val_coord, CellIsRule(operator="greaterThanOrEqual", formula=[target_coord], fill=good_fill, font=good_font))
                ws.conditional_formatting.add(val_coord, CellIsRule(operator="lessThan", formula=[target_coord], fill=bad_fill, font=bad_font))
            else:
                ws.conditional_formatting.add(val_coord, CellIsRule(operator="lessThanOrEqual", formula=[target_coord], fill=good_fill, font=good_font))
                ws.conditional_formatting.add(val_coord, CellIsRule(operator="greaterThan", formula=[target_coord], fill=bad_fill, font=bad_font))

        current_row += 4
