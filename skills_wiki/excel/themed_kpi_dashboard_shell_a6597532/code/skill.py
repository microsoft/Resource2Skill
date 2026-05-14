from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list[dict], theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a modern KPI dashboard shell.
    
    :param kpis: List of dicts with 'label' and 'value' keys, e.g., 
                 [{"label": "Total Calls", "value": "16,749"}, ...]
    """
    ws = wb.create_sheet(sheet_name)
    
    # Hide gridlines for a clean dashboard look
    ws.sheet_view.showGridLines = False
    
    # Mock theme palette extraction (in production, load via theme utility)
    # Using the purple/gold palette from the tutorial
    primary_color = "4B286D"  # Dark purple header
    bg_light = "F2EFF5"       # Pale canvas background
    accent1 = "E5B82A"        # Gold accent
    text_light = "FFFFFF"
    text_dark = "000000"
    
    fill_bg = PatternFill(start_color=bg_light, end_color=bg_light, fill_type="solid")
    fill_header = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    fill_card = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    accent_top_border = Border(top=Side(style="thick", color=accent1))
    
    # 1. Fill entire visible canvas with light background
    for row in ws.iter_rows(min_row=1, max_row=40, min_col=1, max_col=20):
        for cell in row:
            cell.fill = fill_bg
            
    # 2. Header Banner (Rows 1 to 8)
    for row in ws.iter_rows(min_row=1, max_row=8, min_col=1, max_col=20):
        for cell in row:
            cell.fill = fill_header

    # 3. Titles
    ws["B2"] = title
    ws["B2"].font = Font(name="Arial", size=24, color=text_light, bold=True)
    ws["B3"] = subtitle
    ws["B3"].font = Font(name="Arial", size=14, color=text_light, italic=True)
    
    # Set default column width for proportional KPI cards
    for c in range(1, 20):
        ws.column_dimensions[get_column_letter(c)].width = 12
        
    # 4. KPI Cards Strip
    start_col = 2
    for kpi in kpis:
        # Card Background and Accent Border
        for r in range(5, 8):
            for c in range(start_col, start_col + 2):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_card
                if r == 5:
                    cell.border = accent_top_border

        # Value Region (Merged top half of card)
        ws.merge_cells(start_row=5, start_column=start_col, end_row=6, end_column=start_col+1)
        val_cell = ws.cell(row=5, column=start_col)
        val_cell.value = kpi["value"]
        val_cell.font = Font(name="Arial", size=20, bold=True, color=text_dark)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Label Region (Merged bottom half of card)
        ws.merge_cells(start_row=7, start_column=start_col, end_row=7, end_column=start_col+1)
        lbl_cell = ws.cell(row=7, column=start_col)
        lbl_cell.value = str(kpi["label"]).upper()
        lbl_cell.font = Font(name="Arial", size=10, bold=True, color="555555")
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Advance by 2 columns for card width + 1 column for gutter space
        start_col += 3  

    # Adjust row heights to give the cards vertical breathing room
    ws.row_dimensions[5].height = 25
    ws.row_dimensions[6].height = 25
    ws.row_dimensions[7].height = 20
