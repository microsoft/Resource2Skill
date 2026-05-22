def render_workbook(wb, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule

    # Theme palette definition
    themes = {
        "corporate_blue": {
            "header_bg": "1F4E78", "header_fg": "FFFFFF", 
            "input_bg": "F2F2F2", "accent": "4472C4",
            "success": "00B050", "error": "FF0000"
        }
    }
    palette = themes.get(theme, themes["corporate_blue"])

    header_fill = PatternFill(start_color=palette["header_bg"], fill_type="solid")
    header_font = Font(color=palette["header_fg"], bold=True)
    input_fill = PatternFill(start_color=palette["input_bg"], fill_type="solid")
    bold_font = Font(bold=True)
    
    thin_border = Border(bottom=Side(style='thin', color='000000'))

    # ---------------------------------------------------------
    # SHEET 1: ASSUMPTIONS
    # ---------------------------------------------------------
    ws_assumptions = wb.active
    ws_assumptions.title = "Assumptions"
    
    ws_assumptions.column_dimensions['A'].width = 30
    ws_assumptions.column_dimensions['B'].width = 15

    # Define assumptions layout
    assumptions_layout = [
        ("E-Commerce Scenario Assumptions", ""),
        ("Accessories", ""),
        ("Average Order Value ($)", 45),
        ("Starting Monthly Orders", 2000),
        ("Monthly Growth (%)", 0.06),
        ("Gross Margin (%)", 0.35),
        ("", ""),
        ("Devices", ""),
        ("Average Order Value ($)", 280),
        ("Starting Monthly Orders", 500),
        ("Monthly Growth (%)", 0.04),
        ("Gross Margin (%)", 0.22),
        ("", ""),
        ("General", ""),
        ("Marketing Spend (% of Rev)", 0.12)
    ]

    for row_idx, (label, val) in enumerate(assumptions_layout, start=1):
        cell_label = ws_assumptions.cell(row=row_idx, column=1, value=label)
        cell_val = ws_assumptions.cell(row=row_idx, column=2, value=val)
        
        if val == "":
            cell_label.font = bold_font
            if label.startswith("E-Commerce"):
                cell_label.font = Font(size=14, bold=True)
        else:
            # Format inputs
            cell_val.fill = input_fill
            if "%" in label:
                cell_val.number_format = '0.0%'
            elif "$" in label:
                cell_val.number_format = '$#,##0.00'
            else:
                cell_val.number_format = '#,##0'

    # ---------------------------------------------------------
    # SHEET 2: MODEL FORECAST
    # ---------------------------------------------------------
    ws_model = wb.create_sheet("Model")
    ws_model.column_dimensions['A'].width = 35
    
    model_labels = [
        (2, "Accessories", True),
        (3, "Revenue", False),
        (4, "Gross Profit", False),
        (5, "Devices", True),
        (6, "Revenue", False),
        (7, "Gross Profit", False),
        (8, "COMBINED TOTALS", True),
        (9, "Total Revenue", False),
        (10, "Total Gross Profit", False),
        (11, "Marketing Spend", False),
        (12, "Contribution Margin", True),
        (13, "SANITY CHECK", True),
        (14, "Acc Rev + Dev Rev = Total Rev?", False)
    ]
    
    for row, label, is_bold in model_labels:
        c = ws_model.cell(row=row, column=1, value=label)
        if is_bold:
            c.font = bold_font

    months = 24
    for m in range(1, months + 1):
        col_idx = m + 1
        col_letter = get_column_letter(col_idx)
        prev_col = get_column_letter(col_idx - 1) if col_idx > 2 else None

        # Header
        header_cell = ws_model.cell(row=1, column=col_idx, value=f"Month {m}")
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal="center")
        ws_model.column_dimensions[col_letter].width = 14

        # Formulas matching Assumptions coordinates:
        # Acc: AOV=$B$3, Orders=$B$4, Growth=$B$5, Margin=$B$6
        # Dev: AOV=$B$9, Orders=$B$10, Growth=$B$11, Margin=$B$12
        # Gen: Mktg=$B$15
        
        if m == 1:
            ws_model[f"{col_letter}3"] = "=Assumptions!$B$3*Assumptions!$B$4"
            ws_model[f"{col_letter}6"] = "=Assumptions!$B$9*Assumptions!$B$10"
        else:
            ws_model[f"{col_letter}3"] = f"={prev_col}3*(1+Assumptions!$B$5)"
            ws_model[f"{col_letter}6"] = f"={prev_col}6*(1+Assumptions!$B$11)"

        ws_model[f"{col_letter}4"] = f"={col_letter}3*Assumptions!$B$6"
        ws_model[f"{col_letter}7"] = f"={col_letter}6*Assumptions!$B$12"

        # Totals
        ws_model[f"{col_letter}9"] = f"={col_letter}3+{col_letter}6"
        ws_model[f"{col_letter}10"] = f"={col_letter}4+{col_letter}7"
        ws_model[f"{col_letter}11"] = f"={col_letter}9*Assumptions!$B$15"
        ws_model[f"{col_letter}12"] = f"={col_letter}10-{col_letter}11"
        ws_model[f"{col_letter}12"].border = thin_border
        ws_model[f"{col_letter}12"].font = bold_font

        # Sanity Check
        ws_model[f"{col_letter}14"] = f'=IF(ROUND({col_letter}3+{col_letter}6,2)=ROUND({col_letter}9,2), "✓ Match", "✗ Error")'
        ws_model[f"{col_letter}14"].alignment = Alignment(horizontal="center")

        # Number Formatting
        for r in [3, 4, 6, 7, 9, 10, 11, 12]:
            ws_model[f"{col_letter}{r}"].number_format = '$#,##0'

    # Apply Conditional Formatting to Sanity Check Row
    green_font = Font(color=palette["success"], bold=True)
    red_font = Font(color=palette["error"], bold=True)
    ws_model.conditional_formatting.add(
        f"B14:{get_column_letter(months+1)}14",
        CellIsRule(operator='equal', formula=['"✓ Match"'], font=green_font)
    )
    ws_model.conditional_formatting.add(
        f"B14:{get_column_letter(months+1)}14",
        CellIsRule(operator='equal', formula=['"✗ Error"'], font=red_font)
    )

    # ---------------------------------------------------------
    # CHART: REVENUE & MARGIN VISUALIZATION
    # ---------------------------------------------------------
    chart = LineChart()
    chart.title = "Monthly Revenue & Contribution Margin by Product Line"
    chart.style = 13
    chart.y_axis.title = "Amount ($)"
    chart.x_axis.title = "Month"
    chart.height = 10
    chart.width = 25

    data_acc = Reference(ws_model, min_col=1, min_row=3, max_col=months+1, max_row=3)
    data_dev = Reference(ws_model, min_col=1, min_row=6, max_col=months+1, max_row=6)
    data_cm = Reference(ws_model, min_col=1, min_row=12, max_col=months+1, max_row=12)

    chart.add_data(data_acc, titles_from_data=True)
    chart.add_data(data_dev, titles_from_data=True)
    chart.add_data(data_cm, titles_from_data=True)

    cats = Reference(ws_model, min_col=2, min_row=1, max_col=months+1, max_row=1)
    chart.set_categories(cats)

    ws_model.add_chart(chart, "B16")
