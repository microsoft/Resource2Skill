def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    # Standard fallback palette mapping
    palettes = {
        "corporate_blue": {"primary": "1F497D", "text": "FFFFFF"},
        "tech_green": {"primary": "0F9D58", "text": "FFFFFF"},
        "neutral_dark": {"primary": "333333", "text": "FFFFFF"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])

    ws = wb.create_sheet(title=sheet_name)
    
    # 1. Styles
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    header_font = Font(color=colors["text"], bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Financial modeling standard: pure inputs are strictly blue, formulas are black
    input_font = Font(color="0000FF") 
    
    top_border = Border(top=Side(style='thin', color='000000'))
    top_bottom_border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    
    number_fmt = '#,##0'
    percent_fmt = '0.0%'
    
    # 2. Write Title & Headers
    ws['A1'] = title
    ws['A1'].font = Font(size=14, bold=True)
    
    years = [2023, 2024, 2025, 2026]
    actual_years = 1 # 2023 is actual, rest are forecast
    
    ws['A3'] = "Income Statement"
    ws['A3'].font = bold_font
    
    for i, year in enumerate(years):
        col_letter = chr(ord('B') + i)
        cell = ws[f'{col_letter}3']
        cell.value = year
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        # Custom format for Actual (A) vs Estimate (E)
        if i < actual_years:
            cell.number_format = '0"A"'
        else:
            cell.number_format = '0"E"'

    # 3. Build Income Statement Structure
    line_items = [
        ("Revenue", 4),
        ("Cost of Goods Sold", 5),
        ("Gross Profit", 6),
        ("SG&A Expense", 7),
        ("Operating Income", 8),
        ("Taxes", 9),
        ("Net Income", 10)
    ]
    
    for name, row in line_items:
        ws[f'A{row}'] = name
        if "Profit" in name or "Income" in name:
            ws[f'A{row}'].font = bold_font

    # 4. Insert Hardcoded Historicals (2023) - Blue Font
    historical_inputs = {
        'B4': 5210, # Revenue
        'B5': 3345, # COGS
        'B7': 850,  # SG&A
        'B9': 228   # Taxes
    }
    for cell_ref, val in historical_inputs.items():
        ws[cell_ref] = val
        ws[cell_ref].font = input_font
        ws[cell_ref].number_format = number_fmt

    # 5. Insert Subtotal/Total Formulas for all years
    for i in range(len(years)):
        col = chr(ord('B') + i)
        # Gross Profit = Revenue - COGS
        ws[f'{col}6'] = f'={col}4-{col}5'
        ws[f'{col}6'].border = top_border
        
        # Operating Income = Gross Profit - SG&A
        ws[f'{col}8'] = f'={col}6-{col}7'
        ws[f'{col}8'].border = top_border
        
        # Net Income = Operating Income - Taxes
        ws[f'{col}10'] = f'={col}8-{col}9'
        ws[f'{col}10'].border = top_bottom_border
        
        # Apply standard number formatting to all rows
        for row in range(4, 11):
            ws[f'{col}{row}'].number_format = number_fmt

    # 6. Build Assumptions / Drivers Block
    ws['A13'] = "Income Statement Assumptions"
    ws['A13'].font = bold_font
    
    assumptions = [
        ("Revenue Growth Rate", 14),
        ("COGS % of Revenue", 15),
        ("SG&A % of Revenue", 16),
        ("Tax Rate", 17)
    ]
    
    for name, row in assumptions:
        ws[f'A{row}'] = name

    # 7. Fill Assumptions: Calculated for Historical, Hardcoded (Blue) for Forecast
    # Historical logic (2023)
    ws['B14'] = "" # No prior year to calculate growth
    ws['B15'] = "=B5/B4"
    ws['B16'] = "=B7/B4"
    ws['B17'] = "=B9/B8"
    
    # Forecast Hardcodes (2024, 2025, 2026)
    forecast_drivers = [
        (14, [0.043, 0.051, 0.047]), # Revenue Growth
        (15, [0.642, 0.616, 0.622]), # COGS Margin
        (16, [0.163, 0.160, 0.158]), # SG&A Margin
        (17, [0.250, 0.250, 0.250])  # Tax Rate
    ]
    
    for row, values in forecast_drivers:
        for i, val in enumerate(values):
            col = chr(ord('C') + i)
            ws[f'{col}{row}'] = val
            ws[f'{col}{row}'].font = input_font

    # Format all assumptions as percentages
    for row in range(14, 18):
        for i in range(len(years)):
            col = chr(ord('B') + i)
            ws[f'{col}{row}'].number_format = percent_fmt

    # 8. Link Forecast Income Statement to Drivers
    for i in range(1, len(years)):
        col = chr(ord('B') + i)
        prev_col = chr(ord('B') + i - 1)
        
        # Revenue = Prev Revenue * (1 + Growth)
        ws[f'{col}4'] = f'={prev_col}4*(1+{col}14)'
        # COGS = Revenue * COGS%
        ws[f'{col}5'] = f'={col}4*{col}15'
        # SG&A = Revenue * SG&A%
        ws[f'{col}7'] = f'={col}4*{col}16'
        # Taxes = Operating Income * Tax Rate
        ws[f'{col}9'] = f'={col}8*{col}17'

    # Column sizing
    ws.column_dimensions['A'].width = 35
    for i in range(len(years)):
        ws.column_dimensions[chr(ord('B') + i)].width = 12
        
    # Set view to freeze panes for headers and row labels
    ws.freeze_panes = 'B4'
