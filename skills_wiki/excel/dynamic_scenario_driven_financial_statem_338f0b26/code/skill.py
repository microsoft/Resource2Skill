def render_sheet(wb, sheet_name: str, *, title: str = "Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(sheet_name)
    
    # Theme palette fallback
    THEMES = {
        "corporate_blue": {"header_bg": "D9E1F2", "accent": "4F81BD", "input_fg": "0000FF"}
    }
    palette = THEMES.get(theme, THEMES["corporate_blue"])
    
    # Styles
    font_bold = Font(bold=True)
    font_blue = Font(color=palette["input_fg"])
    fill_header = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    fill_input = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # yellow for dropdown
    border_tb = Border(top=Side(style='thin'), bottom=Side(style='thin'))
    border_all = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Number formats
    fmt_usd = '_($* #,##0_);_($* (#,##0);_($* "-"_);_(@_)'
    fmt_pct = '0%'
    fmt_num = '#,##0'

    # Column widths
    ws.column_dimensions["B"].width = 25
    for col in ["C", "D", "E", "F", "G"]:
        ws.column_dimensions[col].width = 13
        
    # Title & Headers
    ws["B1"] = title
    ws["B1"].font = Font(size=14, bold=True)
    ws["B2"] = "Figures in USD"
    ws["B2"].font = font_bold
    
    for i, col in enumerate(["C", "D", "E", "F", "G"]):
        ws[f"{col}2"] = f"Year {i+1}"
        ws[f"{col}2"].font = font_bold
        ws[f"{col}2"].fill = fill_header
        ws[f"{col}2"].alignment = Alignment(horizontal="center")

    # Interactive Scenario Toggle
    ws["I2"] = "Scenario:"
    ws["I2"].font = font_bold
    ws["I2"].alignment = Alignment(horizontal="right")
    
    ws["J2"] = 1
    ws["J2"].fill = fill_input
    ws["J2"].border = border_all
    ws["J2"].alignment = Alignment(horizontal="center")
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["J2"])

    # Line Item Labels
    labels = {
        4: "Revenue",
        5: "Cost of Goods Sold",
        6: "Gross Profit",
        7: "Gross Profit Margin",
        8: "Operating Expenses",
        9: "Operating Profit",
        10: "Corporate Tax",
        11: "Profit / (Loss)",
        14: "Orders",
        15: "Average Order Value",
        16: "COGS per order",
        17: "Fixed Operating Expenses",
        18: "Corporate Tax Rate"
    }
    
    for r, lbl in labels.items():
        ws[f"B{r}"] = lbl
    for r in range(14, 19):
        ws[f"B{r+7}"] = labels[r]   # Copy labels to Scenario 1
        ws[f"B{r+14}"] = labels[r]  # Copy labels to Scenario 2

    ws["B3"] = "Income Statement"
    ws["B13"] = "Live Case Assumptions"
    ws["B20"] = "Scenario 1 (Base Case)"
    ws["B27"] = "Scenario 2 (Upside Case)"
    
    for r in [3, 13, 20, 27]:
        ws[f"B{r}"].font = font_bold
    ws["B20"].fill = fill_header
    ws["B27"].fill = fill_header

    # Hardcoded Scenario Data
    scen1_data = [
        [3000, 6000, 10500, 15750, 21263],      # Orders
        [40, 40, 40, 40, 40],                   # AOV
        [10, 10, 10, 10, 10],                   # COGS
        [50000, 50000, 100000, 100000, 100000], # Fixed Opex
        [0.2, 0.2, 0.2, 0.2, 0.2]               # Tax Rate
    ]
    scen2_data = [
        [3000, 8000, 15000, 25000, 35000],      # Orders (Optimistic)
        [45, 45, 45, 45, 45],                   # AOV (Higher pricing)
        [9, 9, 9, 9, 9],                        # COGS (Economies of scale)
        [60000, 60000, 120000, 120000, 120000], # Fixed Opex (Higher baseline)
        [0.25, 0.25, 0.25, 0.25, 0.25]          # Tax Rate (Higher bracket)
    ]

    for row_offset in range(5):
        s1_row = 21 + row_offset
        s2_row = 28 + row_offset
        for c_idx, col in enumerate(["C", "D", "E", "F", "G"]):
            ws[f"{col}{s1_row}"].value = scen1_data[row_offset][c_idx]
            ws[f"{col}{s1_row}"].font = font_blue
            
            ws[f"{col}{s2_row}"].value = scen2_data[row_offset][c_idx]
            ws[f"{col}{s2_row}"].font = font_blue

    # Formula Logic Layer
    for col in ["C", "D", "E", "F", "G"]:
        # Live Case evaluates the CHOOSE function based on J2 Dropdown
        ws[f"{col}14"] = f"=CHOOSE($J$2, {col}21, {col}28)"
        ws[f"{col}15"] = f"=CHOOSE($J$2, {col}22, {col}29)"
        ws[f"{col}16"] = f"=CHOOSE($J$2, {col}23, {col}30)"
        ws[f"{col}17"] = f"=CHOOSE($J$2, {col}24, {col}31)"
        ws[f"{col}18"] = f"=CHOOSE($J$2, {col}25, {col}32)"
        
        # P&L evaluates off the Live Case
        ws[f"{col}4"] = f"={col}14*{col}15"
        ws[f"{col}5"] = f"={col}14*{col}16"
        ws[f"{col}6"] = f"={col}4-{col}5"
        ws[f"{col}7"] = f"=IFERROR({col}6/{col}4, 0)"
        ws[f"{col}8"] = f"={col}17"
        ws[f"{col}9"] = f"={col}6-{col}8"
        ws[f"{col}10"] = f"=IF({col}9>0, {col}9*{col}18, 0)"
        ws[f"{col}11"] = f"={col}9-{col}10"

    # Formatting Helper
    def format_row(row_idx, fmt, is_bold=False, border=None):
        for col in ["C", "D", "E", "F", "G"]:
            cell = ws[f"{col}{row_idx}"]
            cell.number_format = fmt
            if is_bold:
                # Retain existing color (like blue inputs) if making bold
                current_color = cell.font.color if cell.font else None
                cell.font = Font(bold=True, color=current_color)
            if border:
                cell.border = border

    # Apply Formats to P&L
    format_row(4, fmt_usd)
    format_row(5, fmt_usd)
    format_row(6, fmt_usd, is_bold=True, border=border_tb)
    format_row(7, fmt_pct)
    format_row(8, fmt_usd)
    format_row(9, fmt_usd, is_bold=True)
    format_row(10, fmt_usd)
    format_row(11, fmt_usd, is_bold=True, border=border_tb)

    # Apply Formats to Assumption Blocks
    for block_start in [14, 21, 28]:
        format_row(block_start, fmt_num)         # Orders
        format_row(block_start + 1, fmt_usd)     # AOV
        format_row(block_start + 2, fmt_usd)     # COGS
        format_row(block_start + 3, fmt_usd)     # Fixed Opex
        format_row(block_start + 4, fmt_pct)     # Tax Rate
