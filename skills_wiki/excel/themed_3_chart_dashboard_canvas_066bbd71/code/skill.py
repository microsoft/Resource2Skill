def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Theme Configuration
    theme_colors = {
        "corporate_blue": {"primary": "002060", "accent": "4F81BD"},
        "modern_dark": {"primary": "333333", "accent": "00B050"},
        "cookie_brand": {"primary": "0033A0", "accent": "F2A900"} 
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # 2. Setup Dashboard Canvas
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Dashboard Title
    ws["B2"] = title
    ws["B2"].font = Font(size=24, bold=True, color=palette["primary"])

    # 3. Inject Hidden Backing Data (simulating pre-aggregated Pivot data)
    # Bar Chart Data: Market & Cookie Type
    bar_data = [
        ["Country", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62000, 4800, 21000],
        ["Philippines", 54000, 7000, 22000],
        ["United Kingdom", 46000, 5200, 11000],
        ["United States", 36000, 6300, 22000]
    ]
    
    # Line Chart 1 Data: Units sold each month
    line1_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601], ["Oct", 95622], ["Nov", 65481], ["Dec", 52970]
    ]

    # Line Chart 2 Data: Profit by month
    line2_data = [
        ["Month", "Profit"],
        ["Sep", 124812], ["Oct", 228275], ["Nov", 160228], ["Dec", 136337]
    ]

    # Write data to "off-screen" columns starting at AA (Col 27)
    for r_idx, row in enumerate(bar_data, start=1):
        for c_idx, val in enumerate(row, start=27):
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    for r_idx, row in enumerate(line1_data, start=1):
        for c_idx, val in enumerate(row, start=32):
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    for r_idx, row in enumerate(line2_data, start=1):
        for c_idx, val in enumerate(row, start=35):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Hide the backing data columns
    ws.column_dimensions.group('AA', 'AJ', hidden=True)

    # 4. Create Main Stacked Column Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 10  # Standard clean Excel style
    bar_chart.height = 12
    bar_chart.width = 16

    bar_data_ref = Reference(ws, min_col=28, min_row=1, max_col=30, max_row=5)
    bar_cats_ref = Reference(ws, min_col=27, min_row=2, max_row=5)
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)
    ws.add_chart(bar_chart, "B4")

    # 5. Create Top Right Line Chart (Units Sold)
    line1_chart = LineChart()
    line1_chart.title = "Units sold each month"
    line1_chart.style = 13
    line1_chart.height = 6
    line1_chart.width = 14
    line1_chart.legend = None # Remove legend for cleanliness
    
    l1_data_ref = Reference(ws, min_col=33, min_row=1, max_row=5)
    l1_cats_ref = Reference(ws, min_col=32, min_row=2, max_row=5)
    line1_chart.add_data(l1_data_ref, titles_from_data=True)
    line1_chart.set_categories(l1_cats_ref)
    ws.add_chart(line1_chart, "K4")

    # 6. Create Bottom Right Line Chart (Profit)
    line2_chart = LineChart()
    line2_chart.title = "Profit by month"
    line2_chart.style = 13
    line2_chart.height = 6
    line2_chart.width = 14
    line2_chart.legend = None
    
    l2_data_ref = Reference(ws, min_col=36, min_row=1, max_row=5)
    l2_cats_ref = Reference(ws, min_col=35, min_row=2, max_row=5)
    line2_chart.add_data(l2_data_ref, titles_from_data=True)
    line2_chart.set_categories(l2_cats_ref)
    ws.add_chart(line2_chart, "K14")
