from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import range_boundaries

def render_sheet(wb, sheet_name: str, *, title: str = "Executive Dashboard", nav_links: list[str] = None, theme: str = "corporate_blue", **kwargs) -> None:
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    # Disable standard Excel gridlines for a clean UI
    ws.sheet_view.showGridLines = False

    # Standardized modern UI colors
    bg_color = "F3F4F6"
    sidebar_color = "1F2937"
    card_bg = "FFFFFF"
    text_color = "111827"
    border_color = "E5E7EB"

    # 1. Set column widths and row heights for the grid layout
    ws.column_dimensions['A'].width = 8
    for col in "BFJ":
        ws.column_dimensions[col].width = 3  # Spacers
    for col in "CDEGHIKLM":
        ws.column_dimensions[col].width = 12 # Content columns

    ws.row_dimensions[3].height = 15 # Spacer row below title
    ws.row_dimensions[9].height = 15 # Spacer row between top and bottom cards

    # 2. Paint canvas background
    light_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=25, min_col=2, max_col=14):
        for cell in row:
            cell.fill = light_fill

    # 3. Paint sidebar
    sidebar_fill = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    for row in range(1, 26):
        ws.cell(row=row, column=1).fill = sidebar_fill

    # 4. Add Sidebar Navigation Buttons
    nav_links = nav_links or ["Dashboard", "Inputs", "Settings"]
    nav_font = Font(color="FFFFFF", bold=True, size=10)
    nav_align = Alignment(horizontal="center", vertical="center")
    
    for i, link in enumerate(nav_links):
        # Place navigation links starting at row 5, spaced by 3 rows
        cell = ws.cell(row=5 + (i * 3), column=1, value=link[:3].upper())
        cell.font = nav_font
        cell.alignment = nav_align
        cell.hyperlink = f"#'{link}'!A1"

    # 5. Dashboard Title
    t_cell = ws.cell(row=2, column=3, value=title)
    t_cell.font = Font(size=20, bold=True, color=text_color)

    # 6. Define & Render Cards (3 Top KPIs, 2 Bottom Chart areas)
    cards = [
        {"range": "C4:E8", "title": "Sales"},
        {"range": "G4:I8", "title": "Profit"},
        {"range": "K4:M8", "title": "Customers"},
        {"range": "C10:I20", "title": "Sales Trend"},
        {"range": "K10:M20", "title": "Customer Satisfaction"}
    ]

    card_fill = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
    
    for card in cards:
        min_col, min_row, max_col, max_row = range_boundaries(card["range"])
        
        # Apply fills and outer borders to the card range
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                border_kwargs = {}
                if r == min_row: border_kwargs['top'] = Side(style='thin', color=border_color)
                if r == max_row: border_kwargs['bottom'] = Side(style='thin', color=border_color)
                if c == min_col: border_kwargs['left'] = Side(style='thin', color=border_color)
                if c == max_col: border_kwargs['right'] = Side(style='thin', color=border_color)
                
                if border_kwargs:
                    cell.border = Border(**border_kwargs)

        # Format Card Title
        c_title = ws.cell(row=min_row, column=min_col, value=card["title"])
        c_title.font = Font(bold=True, color=text_color, size=12)
        c_title.alignment = Alignment(vertical="center")
        
        # Merge top row for the card title area
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=min_row, end_column=max_col)
