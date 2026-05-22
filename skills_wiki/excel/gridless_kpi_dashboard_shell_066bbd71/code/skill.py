def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    
    # Simple theme router fallback
    theme_colors = {
        "corporate_blue": {"bg": "203764", "fg": "FFFFFF"},
        "midnight": {"bg": "1A1A1A", "fg": "F2F2F2"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # 1. Setup the Presentation Dashboard
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    # Hide gridlines for the clean application-like feel shown in the tutorial
    ws_dash.sheet_view.showGridLines = False
    
    # 2. Setup the Hidden Data Engine
    ws_data = wb.create_sheet("ChartData")
    ws_data.sheet_state = 'hidden'
    
    # --- Inject Mock Summary Data ---
    
    # Table 1: Market by Cookie Type (Rows 1 to 5)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle"],
        ["India", 62349, 4872, 21028, 25085],
        ["Philippines", 54618, 7026, 22005, 8313],
        ["United Kingdom", 46530, 5220, 11497, 14620],
        ["United States", 36657, 6368, 22260, 9937],
    ]
    for row in market_data:
        ws_data.append(row)
        
    ws_data.append([]) # Row 6: Spacer
    ws_data.append([]) # Row 7: Spacer
    
    # Table 2: Units Sold Trend (Rows 8 to 12)
    units_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970],
    ]
    for row in units_data:
        ws_data.append(row)
        
    ws_data.append([]) # Row 13: Spacer
    ws_data.append([]) # Row 14: Spacer
    
    # Table 3: Profit Trend (Rows 15 to 19)
    profit_data = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337],
    ]
    for row in profit_data:
        ws_data.append(row)
        
    # 3. Build the Dashboard UI
    
    # Dashboard Header Banner
    ws_dash.merge_cells("B2:P3")
    header_cell = ws_dash["B2"]
    header_cell.value = title
    header_cell.font = Font(size=22, bold=True, color=palette["fg"])
    header_cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Create Chart 1: Stacked Column (Market Breakdown)
    bc = BarChart()
    bc.type = "col"
    bc.grouping = "stacked"
    bc.overlap = 100
    bc.title = "Profit by Market & Cookie Type"
    bc.style = 10
    data_ref_1 = Reference(ws_data, min_col=2, max_col=5, min_row=1, max_row=5)
    cats_ref_1 = Reference(ws_data, min_col=1, max_col=1, min_row=2, max_row=5)
    bc.add_data(data_ref_1, titles_from_data=True)
    bc.set_categories(cats_ref_1)
    bc.height = 12
    bc.width = 16
    ws_dash.add_chart(bc, "B5")
    
    # Create Chart 2: Line Chart (Units Sold)
    lc1 = LineChart()
    lc1.title = "Units sold each month"
    lc1.style = 13
    data_ref_2 = Reference(ws_data, min_col=2, max_col=2, min_row=8, max_row=12)
    cats_ref_2 = Reference(ws_data, min_col=1, max_col=1, min_row=9, max_row=12)
    lc1.add_data(data_ref_2, titles_from_data=True)
    lc1.set_categories(cats_ref_2)
    lc1.height = 7.5
    lc1.width = 13
    lc1.legend = None # Clean up chart clutter as shown in tutorial
    ws_dash.add_chart(lc1, "K5")
    
    # Create Chart 3: Line Chart (Profit)
    lc2 = LineChart()
    lc2.title = "Profit by month"
    lc2.style = 13
    data_ref_3 = Reference(ws_data, min_col=2, max_col=2, min_row=15, max_row=19)
    cats_ref_3 = Reference(ws_data, min_col=1, max_col=1, min_row=16, max_row=19)
    lc2.add_data(data_ref_3, titles_from_data=True)
    lc2.set_categories(cats_ref_3)
    lc2.height = 7.5
    lc2.width = 13
    lc2.legend = None
    ws_dash.add_chart(lc2, "K14")
    
    # Polish Column Widths for Margins
    ws_dash.column_dimensions["A"].width = 3
    ws_dash.column_dimensions["J"].width = 2
