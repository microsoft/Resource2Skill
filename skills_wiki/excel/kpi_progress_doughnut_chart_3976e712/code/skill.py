from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.series import DataPoint

def render(ws, anchor: str, *, kpi_name: str, actual: float, target: float, theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Map theme palette to chart colors
    colors = {
        "corporate_blue": {"primary": "003366", "muted": "D9D9D9"},
        "modern_dark": {"primary": "4ECDC4", "muted": "444444"}
    }
    palette = colors.get(theme, colors["corporate_blue"])

    # 2. Write supporting data under the chart anchor footprint
    row = ws[anchor].row
    col = ws[anchor].column

    ws.cell(row=row, column=col, value="Actual")
    ws.cell(row=row, column=col+1, value=actual)

    ws.cell(row=row+1, column=col, value="Remaining")
    remaining = max(0, target - actual)
    ws.cell(row=row+1, column=col+1, value=remaining)

    # 3. Initialize Doughnut Chart
    chart = DoughnutChart()
    pct_complete = int((actual / target) * 100) if target else 0
    chart.title = f"{kpi_name} ({pct_complete}%)"
    
    # 4. Apply modern dashboard styling (wide hole, no legend)
    chart.holeSize = 65
    chart.legend = None  
    chart.width = 6
    chart.height = 4

    # 5. Bind Data & Categories
    data = Reference(ws, min_col=col+1, min_row=row, max_row=row+1)
    cats = Reference(ws, min_col=col, min_row=row, max_row=row+1)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    # 6. Style Individual Slices (Actual vs Remaining)
    series = chart.series[0]

    dp0 = DataPoint(idx=0)
    dp0.graphicalProperties.solidFill = palette["primary"]

    dp1 = DataPoint(idx=1)
    dp1.graphicalProperties.solidFill = palette["muted"]

    # Assign customized points back to the series
    series.dPt = [dp0, dp1]

    # 7. Mount chart
    ws.add_chart(chart, anchor)
