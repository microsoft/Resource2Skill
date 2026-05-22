def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    import random

    # 1. Setup mock time-series data
    col_str, row_idx = coordinate_from_string(anchor)
    col_idx = column_index_from_string(col_str)
    
    data_start_row = row_idx + 15
    data_start_col = col_idx
    
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    
    ws.cell(row=data_start_row, column=data_start_col, value="Month")
    ws.cell(row=data_start_row, column=data_start_col+1, value="Visits")
    
    for i, month in enumerate(months):
        ws.cell(row=data_start_row+i+1, column=data_start_col, value=month)
        ws.cell(row=data_start_row+i+1, column=data_start_col+1, value=random.randint(15000, 80000))
        
    cats = Reference(ws, min_col=data_start_col, min_row=data_start_row+1, max_row=data_start_row+len(months))
    data = Reference(ws, min_col=data_start_col+1, min_row=data_start_row, max_row=data_start_row+len(months))
    
    # 2. Configure Base Line Chart
    chart = LineChart()
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    # 3. Apply Transparency (Core Technique)
    # Remove chart area background and bounding border
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line.noFill = True
    
    # Remove inner plot area background and border
    chart.plot_area.graphicalProperties.noFill = True
    chart.plot_area.graphicalProperties.line.noFill = True
    
    # 4. Minimalism Settings
    chart.y_axis.majorGridlines = None
    chart.x_axis.majorGridlines = None
    chart.legend = None
    
    # Format Y-axis to thousands with "K" suffix (e.g., 50,000 -> 50K)
    chart.y_axis.number_format = '#,##0,"K"'
    
    # 5. Series Styling
    series = chart.series[0]
    theme_color = "4A90E2"  # Usually fetched from a theme dict: theme.get("primary", "4A90E2")
    
    series.graphicalProperties.line.solidFill = theme_color
    series.graphicalProperties.line.width = 25000  # ~2pt line thickness
    
    # Add colored circle markers for data points
    series.marker = Marker(symbol="circle", size=5)
    series.marker.graphicalProperties.solidFill = theme_color
    series.marker.graphicalProperties.line.noFill = True
    
    # Shape properties
    chart.width = 14
    chart.height = 6
    
    ws.add_chart(chart, anchor)
