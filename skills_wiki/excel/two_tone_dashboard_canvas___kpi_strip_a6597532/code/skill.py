from openpyxl.utils.cell import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment

def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str = "Evaluating Agent Performance", kpis: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Generates a two-tone dashboard canvas with a header, subtitle, and a row of KPI cards.
    """
    # 1. Resolve Theme Palette
    try:
        from skills_library.excel.components._helpers import get_palette
        palette = get_palette(theme)
    except ImportError:
        # Fallback to the specific Deep Purple / Gold theme seen in the video
        palette = {
            "primary": "4A235A",      # Deep purple header
            "secondary": "F5EEF8",    # Pale purple body
            "accent": "F4D03F",       # Gold KPI accent
            "bg": "FFFFFF",           # White cards
            "text": "000000",         # Black
            "text_light": "FFFFFF"    # White text
        }

    # 2. Setup Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    ws.sheet_view.showGridLines = False

    header_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    body_fill = PatternFill(start_color=palette["secondary"], end_color=palette["secondary"], fill_type="solid")

    # 3. Apply Two-Tone Background
    # Header Band: Rows 1 to 8
    for row in range(1, 9):
        ws.row_dimensions[row].height = 20
        for col in range(1, 25):  # Cover A to X
            ws.cell(row=row, column=col).fill = header_fill

    # Body Canvas: Rows 9 to 40
    for row in range(9, 41):
        ws.row_dimensions[row].height = 18
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = body_fill

    # 4. Insert Title & Subtitle
    title_cell = ws.cell(row=2, column=2)
    title_cell.value = title
    title_cell.font = Font(name="Arial", size=24, bold=True, color=palette["text_light"])

    sub_cell = ws.cell(row=3, column=2)
    sub_cell.value = subtitle
    sub_cell.font = Font(name="Arial", size=12, italic=True, color=palette["accent"])

    # 5. Default KPI Data (if none provided)
    if kpis is None:
        kpis = [
            {"label": "Total Calls", "value": 16749, "format": "#,##0"},
            {"label": "Calls Reached", "value": 3328, "format": "#,##0"},
            {"label": "Deals Closed", "value": 1203, "format": "#,##0"},
            {"label": "Deal Value", "value": 646979, "format": "$#,##0"}
        ]

    # 6. Render KPI Strip (Anchored at B5)
    start_row = 5
    start_col = 2

    # Card layout row heights
    ws.row_dimensions[start_row].height = 32
    ws.row_dimensions[start_row + 1].height = 18

    accent_fill = PatternFill(start_color=palette["accent"], end_color=palette["accent"], fill_type="solid")
    card_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    
    val_font = Font(name="Arial", size=18, bold=True, color=palette["primary"])
    lbl_font = Font(name="Arial", size=10, bold=True, color=palette["primary"])
    
    align_val = Alignment(horizontal="center", vertical="bottom")
    align_lbl = Alignment(horizontal="center", vertical="top")

    curr_col = start_col
    for kpi in kpis:
        col_accent = curr_col
        col_data = curr_col + 1
        col_spacer = curr_col + 2

        # Configure card column widths
        ws.column_dimensions[get_column_letter(col_accent)].width = 3.0
        ws.column_dimensions[get_column_letter(col_data)].width = 18.0
        ws.column_dimensions[get_column_letter(col_spacer)].width = 2.0

        # Build Accent Color Block (Left edge of the card)
        ws.merge_cells(start_row=start_row, start_column=col_accent, end_row=start_row+1, end_column=col_accent)
        for r in range(start_row, start_row + 2):
            ws.cell(row=r, column=col_accent).fill = accent_fill

        # Main Value Cell (Top right of the card)
        val_cell = ws.cell(row=start_row, column=col_data)
        val_cell.value = kpi.get("value", "")
        if "format" in kpi:
            val_cell.number_format = kpi["format"]
        val_cell.font = val_font
        val_cell.fill = card_fill
        val_cell.alignment = align_val

        # Subtitle Label Cell (Bottom right of the card)
        lbl_cell = ws.cell(row=start_row + 1, column=col_data)
        lbl_cell.value = kpi.get("label", "").upper()
        lbl_cell.font = lbl_font
        lbl_cell.fill = card_fill
        lbl_cell.alignment = align_lbl

        # Shift anchor for the next KPI card
        curr_col += 3
