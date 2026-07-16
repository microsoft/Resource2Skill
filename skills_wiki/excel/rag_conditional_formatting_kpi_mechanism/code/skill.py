from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule


def build_rag_kpi_strip(path="rag_kpi_strip.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append(["KPI", "Actual", "Target", "Variance", "Status"])
    rows = [
        ["Revenue", 31_000_000, 30_000_000, None, None],
        ["Churn", 0.07, 0.08, None, None],
        ["NPS", 52, 50, None, None],
    ]
    for row in rows:
        ws.append(row)

    ws["G1"] = "Green threshold"
    ws["H1"] = 0
    for r in range(2, 5):
        ws[f"D{r}"] = f"=B{r}/C{r}-1"
    ws["E2"] = '=IF(D2>=0,"Green",IF(D2>=-0.05,"Amber","Red"))'
    ws["E3"] = '=IF(D3<=0,"Green",IF(D3<=0.10,"Amber","Red"))'
    ws["E4"] = '=IF(D4>=0,"Green",IF(D4>=-0.10,"Amber","Red"))'

    fills = {"Green": "C6EFCE", "Amber": "FFEB9C", "Red": "FFC7CE"}
    for status, color in fills.items():
        ws.conditional_formatting.add(
            "E2:E4",
            FormulaRule(formula=[f'$E2="{status}"'], fill=PatternFill("solid", fgColor=color)),
        )
    ws.conditional_formatting.add(
        "D2:D4",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(path)
    return path
