from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import PatternFill, Font, Alignment

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Prepare Hidden Data Sheet
    ws_data = wb.create_sheet(f"{sheet_name}_Data")
    ws_data.sheet_state = 'hidden'
    
    # Sample time-series tank capacity data
    headers = ["Date", "Tank A (Raw)", "Tank B (Mix)", "Tank C (Filter)", "Tank D (Pack)"]
    data = [
        ["2023-10-01", 0.85, 0.40, 0.90, 0.20],
        ["2023-10-02", 0.70, 0.55, 0.85, 0.35],
        ["2023-10-03", 0.45, 0.80, 0.60, 0.70],
        ["2023-10-04", 0.20, 0.95, 0.30, 0.90],
    ]
    ws_data.append(headers)
    for row in data:
        ws_data.append(row)

    # 2. Setup Dashboard Sheet
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.showGridLines = False
    
    # Theme configuration
    theme_colors = {
        "corporate_blue": "1F4E78",
        "emerald_green": "228B22",
        "industrial_grey": "4F4F4F"
    }
    primary_color = theme_colors.get(theme, "1F4E78")
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    
    # Dashboard Header
    ws["B2"] = title
    ws["B2"].font = Font(size=18, bold=True, color=primary_color)
    
    # Date Selection Dropdown
    ws["B4"] = "Select Production Date:"
    ws["B4"].font = Font(bold=True)
    ws["B4"].alignment = Alignment(horizontal="right")
    ws["C4"] = "2023-10-03" # Default
    ws["C4"].font = Font(bold=True, color="0000FF")
    
    dv = DataValidation(type="list", formula1=f"'{ws_data.title}'!$A$2:$A$5", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["C4"])

    # 3. Build Process Flow Tanks
    tanks = headers[1:]
    start_col = 2 # Column B
    
    for i, tank_name in enumerate(tanks):
        col_letter = chr(64 + start_col + (i * 3)) # B, E, H, K
        val_col_idx = i + 2
        
        # Tank Header
        ws[f"{col_letter}6"] = tank_name
        ws[f"{col_letter}6"].fill = header_fill
        ws[f"{col_letter}6"].font = white_font
        ws[f"{col_letter}6"].alignment = Alignment(horizontal="center")
        ws.column_dimensions[col_letter].width = 15
        
        # Dynamic VLOOKUP Value
        val_cell = f"{col_letter}7"
        ws[val_cell] = f'=VLOOKUP($C$4, \'{ws_data.title}\'!$A$1:$E$5, {val_col_idx}, FALSE)'
        ws[val_cell].number_format = '0%'
        ws[val_cell].font = Font(size=12, bold=True)
        ws[val_cell].alignment = Alignment(horizontal="center")
        
        # "Tank" Chart (Column chart behaving like a fill gauge)
        chart = BarChart()
        chart.type = "col"
        chart.height = 5.0
        chart.width = 2.8
        chart.legend = None
        chart.gapWidth = 0  # Make bar fill the whole container width
        
        # Fix Y-axis to 0-100% capacity and hide axes to look like a raw container
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.y_axis.delete = True
        chart.x_axis.delete = True
        
        # Embed data label inside the bar
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        
        # Add Data
        data_ref = Reference(ws, min_col=ord(col_letter)-64, min_row=7, max_row=7)
        chart.add_data(data_ref)
        
        # Position chart directly below the value
        ws.add_chart(chart, f"{col_letter}8")
        
        # Add Flow Arrow (between tanks)
        if i < len(tanks) - 1:
            arrow_col = chr(64 + start_col + (i * 3) + 1) # C, F, I
            arrow_cell = f"{arrow_col}12"
            ws[arrow_cell] = "➔"
            ws[arrow_cell].font = Font(size=28, color="888888")
            ws[arrow_cell].alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[arrow_col].width = 8
            
    # Clean up default sheets if creating a fresh workbook payload
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 2:
        del wb["Sheet"]
