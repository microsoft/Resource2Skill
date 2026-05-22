from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean, executive dashboard layout with a title banner, 
    a left-hand filter panel placeholder, and a 3-chart grid layout.
    """
    ws = wb.create_sheet(sheet_name)
    
    # Clean up worksheet canvas
    ws.sheet_view.showGridLines = False
    
    # Theme palette fallback
    theme_colors = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF", "panel": "F2F2F2"},
        "midnight": {"bg": "203764", "fg": "FFFFFF", "panel": "E7E6E6"},
        "emerald": {"bg": "375623", "fg": "FFFFFF", "panel": "E2EFDA"},
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    # 1. Title Banner
    ws.merge_cells("B2:Q3")
    title_cell = ws["B2"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["fg"])
    title_cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Slicer/Filter Panel Placeholder
    ws.column_dimensions["B"].width = 22
    ws.merge_cells("B5:B25")
    panel = ws["B5"]
    panel.value = "[ Slicer / Filter Panel ]\n\nReserve this space for interactive controls."
    panel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    panel.fill = PatternFill(start_color=palette["panel"], end_color=palette["panel"], fill_type="solid")
    panel.font = Font(color="7F7F7F", italic=True)
    
    # Spacer columns to frame the dashboard
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["C"].width = 2
    ws.column_dimensions["K"].width = 2
    
    # 3. Generate Chart Data (Stored in hidden columns Z onwards)
    bar_data = [
        ["Market", "Chocolate Chip", "Oatmeal Raisin", "Sugar Cookie"],
        ["United States", 150000, 120000, 80000],
        ["United Kingdom", 100000, 110000, 90000],
        ["India", 80000, 70000, 60000],
        ["Philippines", 60000, 50000, 40000]
    ]
    for r_idx, row in enumerate(bar_data, start=1):
        for c_idx, val in enumerate(row, start=26): # Column Z
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    line_data = [
        ["Month", "Units Sold", "Profit"],
        ["Jan", 50000, 120000],
        ["Feb", 55000, 130000],
        ["Mar", 48000, 110000],
        ["Apr", 60000, 150000],
        ["May", 65000, 160000],
        ["Jun", 70000, 180000]
    ]
    for r_idx, row in enumerate(line_data, start=10):
        for c_idx, val in enumerate(row, start=26):
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    # Hide data staging columns
    for col in ["Z", "AA", "AB", "AC", "AD"]:
        ws.column_dimensions[col].hidden = True

    # 4. Main Chart (Stacked Column)
    bc = BarChart()
    bc.type = "col"
    bc.style = 11
    bc.grouping = "stacked"
    bc.overlap = 100
    bc.title = "Profit by Market & Cookie Type"
    bc_data = Reference(ws, min_col=27, min_row=1, max_col=29, max_row=5)
    bc_cats = Reference(ws, min_col=26, min_row=2, max_row=5)
    bc.add_data(bc_data, titles_from_data=True)
    bc.set_categories(bc_cats)
    bc.width = 16
    bc.height = 12.5
    ws.add_chart(bc, "D5")
    
    # 5. Top Right Chart (Line)
    lc1 = LineChart()
    lc1.style = 12
    lc1.title = "Units Sold Each Month"
    lc1_data = Reference(ws, min_col=27, min_row=10, max_row=16)
    lc1_cats = Reference(ws, min_col=26, min_row=11, max_row=16)
    lc1.add_data(lc1_data, titles_from_data=True)
    lc1.set_categories(lc1_cats)
    lc1.width = 14
    lc1.height = 6
    ws.add_chart(lc1, "L5")
    
    # 6. Bottom Right Chart (Line)
    lc2 = LineChart()
    lc2.style = 13
    lc2.title = "Profit by Month"
    lc2_data = Reference(ws, min_col=28, min_row=10, max_row=16)
    lc2_cats = Reference(ws, min_col=26, min_row=11, max_row=16)
    lc2.add_data(lc2_data, titles_from_data=True)
    lc2.set_categories(lc2_cats)
    lc2.width = 14
    lc2.height = 6
    ws.add_chart(lc2, "L16")
