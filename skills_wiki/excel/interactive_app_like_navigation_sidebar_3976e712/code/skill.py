from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

def render_workbook(wb: Workbook, *, title: str = "Sales Dashboard", sheets: list[str] = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Creates a multi-tab interactive workbook with a persistent left-hand navigation sidebar.
    """
    if sheets is None:
        # Based on the video's three main sections
        sheets = ["Dashboard", "Inputs", "Contacts"]
        
    # Simulate theme colors
    sidebar_bg = "1F4E78" # Dark blue
    sidebar_fg = "FFFFFF" # White
    active_bg = "4472C4"  # Lighter blue 
    body_bg = "F2F2F2"    # Light gray wash for dashboard canvas
    
    sidebar_fill = PatternFill("solid", fgColor=sidebar_bg)
    active_fill = PatternFill("solid", fgColor=active_bg)
    body_fill = PatternFill("solid", fgColor=body_bg)
    
    sidebar_font = Font(color=sidebar_fg, bold=True)
    active_font = Font(color=sidebar_fg, bold=True, underline="single")
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    for sheet_name in sheets:
        ws = wb.create_sheet(sheet_name)
        
        # 1. Turn off native gridlines for an "app" look
        ws.sheet_view.showGridLines = False
        
        # 2. Apply body background color to the main canvas (Columns B to P, Rows 1 to 40)
        for row in ws.iter_rows(min_row=1, max_row=40, min_col=2, max_col=16):
            for cell in row:
                cell.fill = body_fill
                
        # 3. Setup the persistent sidebar in Column A
        ws.column_dimensions['A'].width = 18
        for r in range(1, 41):
            ws.cell(row=r, column=1).fill = sidebar_fill
            
        # 4. Add sidebar header/logo area
        logo_cell = ws.cell(row=2, column=1)
        logo_cell.value = "MENU"
        logo_cell.font = Font(color=sidebar_fg, bold=True, size=14)
        logo_cell.alignment = center_align
            
        # 5. Add Navigation Links
        start_row = 6
        for i, target_sheet in enumerate(sheets):
            link_cell = ws.cell(row=start_row + (i * 3), column=1)
            
            # The HYPERLINK formula targeting cell A1 of the specific sheet
            link_cell.value = f'=HYPERLINK("#\'{target_sheet}\'!A1", "{target_sheet}")'
            link_cell.alignment = center_align
            
            # Highlight the link if it corresponds to the current sheet
            if target_sheet == sheet_name:
                link_cell.fill = active_fill
                link_cell.font = active_font
            else:
                link_cell.fill = sidebar_fill
                link_cell.font = sidebar_font
                
        # 6. Add Page Title to the main canvas
        title_cell = ws.cell(row=3, column=3)
        title_cell.value = f"{title} - {sheet_name}"
        title_cell.font = Font(size=24, bold=True, color=sidebar_bg)
