from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

def render_sheet(wb, sheet_name: str, *, title: str = "Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Theme resolution fallback
    themes = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF", "highlight": "FFF2CC", "input_fg": "0000FF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])
    
    # Styles
    input_font = Font(color=palette["input_fg"])
    calc_font = Font(color="000000")
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_font = Font(color=palette["header_fg"], bold=True)
    scen_fill = PatternFill(start_color=palette["highlight"], end_color=palette["highlight"], fill_type="solid")
    
    thin_border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
    double_bottom = Border(top=Side(style='thin'), bottom=Side(style='double'))
    
    # Formats
    fmt_usd = '"$"#,##0'
    fmt_usd_cents = '"$"#,##0.00'
    fmt_pct = '0%'
    fmt_num = '#,##0'

    # Column Widths
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 8
    for col in ['D', 'E', 'F', 'G', 'H']:
        ws.column_dimensions[col].width = 13

    # --- Header ---
    headers = ["Figures in USD", "Unit", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for c_idx, val in enumerate(headers, start=2):
        cell = ws.cell(row=2, column=c_idx, value=val)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # --- Scenario Toggle ---
    ws['J2'] = "Scenario"
    ws['J2'].font = bold_font
    ws['J2'].alignment = Alignment(horizontal="right")
    
    ws['K2'] = 1
    ws['K2'].fill = scen_fill
    ws['K2'].font = input_font
    ws['K2'].alignment = Alignment(horizontal="center")
    ws['K2'].border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws['K2'])

    # --- Income Statement ---
    ws['B3'] = "Income Statement"
    ws['B3'].font = bold_font

    line_items = [
        (4, "Revenue", "$", fmt_usd, "={col}22 * {col}24", None),
        (5, "Cost of Goods Sold", "$", fmt_usd, "={col}22 * {col}25", None),
        (6, "Gross Profit", "$", fmt_usd, "={col}4 - {col}5", thin_border),
        (7, "Gross Margin", "%", fmt_pct, "=IF({col}4>0, {col}6/{col}4, 0)", None),
        (9, "Operating Expenses", "$", fmt_usd, "={col}26", None),
        (11, "Operating Profit", "$", fmt_usd, "={col}6 - {col}9", thin_border),
        (12, "Taxes", "$", fmt_usd, "=MAX(0, {col}11 * {col}27)", None),
        (14, "Net Profit", "$", fmt_usd, "={col}11 - {col}12", double_bottom)
    ]

    for row, lbl, unit, fmt, formula, border in line_items:
        ws[f'B{row}'] = lbl
        ws[f'C{row}'] = unit
        if border: 
            ws[f'B{row}'].font = bold_font
            
        for col in ['D', 'E', 'F', 'G', 'H']:
            cell = ws[f'{col}{row}']
            cell.value = formula.format(col=col)
            cell.number_format = fmt
            if border:
                cell.border = border
                cell.font = bold_font

    # --- Assumptions: Live Case ---
    ws['B18'] = "Assumptions: Live Case"
    ws['B18'].font = bold_font
    
    assumptions_map = [
        ("Number of Orders", "#", fmt_num),
        ("Order Growth Rate", "%", fmt_pct),
        ("Average Order Value", "$", fmt_usd_cents),
        ("COGS (per order)", "$", fmt_usd_cents),
        ("Fixed Operating Expenses", "$", fmt_usd),
        ("Corporate Tax Rate", "%", fmt_pct)
    ]
    
    for idx, (lbl, unit, fmt) in enumerate(assumptions_map):
        r = 22 + idx
        ws[f'B{r}'] = lbl
        ws[f'C{r}'] = unit
        for col in ['D', 'E', 'F', 'G', 'H']:
            # Engine: CHOOSE drives the Live Case by pulling from offset scenario blocks below
            ws[f'{col}{r}'] = f'=CHOOSE($K$2, {col}{r+10}, {col}{r+20})'
            ws[f'{col}{r}'].number_format = fmt
            ws[f'{col}{r}'].font = calc_font

    # --- Scenario 1: Upper Case ---
    ws['B28'] = "Scenario 1: Upper Case"
    ws['B28'].font = bold_font
    
    scen1_data = [
        [3000, "=D32*(1+E33)", "=E32*(1+F33)", "=F32*(1+G33)", "=G32*(1+H33)"],
        ["", 1.0, 0.75, 0.50, 0.35], 
        [39.95, 39.95, 39.95, 39.95, 39.95],
        [8.75, 8.75, 8.75, 8.75, 8.75], 
        [50000, 50000, 60000, 70000, 80000],
        [0.20, 0.20, 0.20, 0.20, 0.20] 
    ]
    
    for idx, (lbl, unit, fmt) in enumerate(assumptions_map):
        r = 32 + idx
        ws[f'B{r}'] = lbl
        ws[f'C{r}'] = unit
        for i, col in enumerate(['D', 'E', 'F', 'G', 'H']):
            val = scen1_data[idx][i]
            if val == "": continue
            cell = ws[f'{col}{r}']
            cell.value = val
            cell.number_format = fmt
            if not str(val).startswith("="):
                cell.font = input_font # Color hardcoded inputs blue

    # --- Scenario 2: Lower Case ---
    ws['B38'] = "Scenario 2: Lower Case"
    ws['B38'].font = bold_font
    
    scen2_data = [
        [2000, "=D42*(1+E43)", "=E42*(1+F43)", "=F42*(1+G43)", "=G42*(1+H43)"], 
        ["", 0.50, 0.35, 0.25, 0.15],
        [34.95, 34.95, 34.95, 34.95, 34.95], 
        [8.00, 8.00, 8.00, 8.00, 8.00], 
        [50000, 50000, 50000, 50000, 50000], 
        [0.25, 0.25, 0.25, 0.25, 0.25] 
    ]
    
    for idx, (lbl, unit, fmt) in enumerate(assumptions_map):
        r = 42 + idx
        ws[f'B{r}'] = lbl
        ws[f'C{r}'] = unit
        for i, col in enumerate(['D', 'E', 'F', 'G', 'H']):
            val = scen2_data[idx][i]
            if val == "": continue
            cell = ws[f'{col}{r}']
            cell.value = val
            cell.number_format = fmt
            if not str(val).startswith("="):
                cell.font = input_font
