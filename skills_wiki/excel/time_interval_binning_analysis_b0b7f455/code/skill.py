from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Time Interval Analysis", theme: str = "corporate_blue", data: list = None, **kwargs) -> None:
    """
    Renders a complete time binning analysis sheet.
    Takes raw timestamped data, bins it into 30-minute intervals, 
    and generates a summary table and line chart.
    """
    ws = wb.create_sheet(sheet_name)

    # Basic styling elements
    header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border_side = Side(style="thin", color="CCCCCC")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    time_format = "h:mm"
    datetime_format = "yyyy-mm-dd h:mm"

    # 1. Setup Reference Lookup Table (Column G)
    ws['G1'] = "Interval Start"
    ws['G1'].fill = header_fill
    ws['G1'].font = header_font

    # Generate 48 half-hour intervals as Excel time fractions
    intervals = []
    for h in range(24):
        for m in (0, 30):
            intervals.append((h * 60 + m) / 1440.0)

    for i, time_val in enumerate(intervals, start=2):
        cell = ws.cell(row=i, column=7)
        cell.value = time_val
        cell.number_format = time_format
        cell.border = border

    # 2. Setup Data Table (Columns A-D)
    headers = ["Timestamp", "Sales Value", "Extracted Time", "Half-Hour Bin"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    # Sample Data if none provided
    if not data:
        data = [
            ("2023-07-20 08:15:00", 150),
            ("2023-07-20 08:45:00", 200),
            ("2023-07-20 09:05:00", 350),
            ("2023-07-20 12:20:00", 500),
            ("2023-07-20 12:55:00", 450),
            ("2023-07-20 18:10:00", 800),
            ("2023-07-20 18:35:00", 950),
            ("2023-07-20 22:15:00", 120),
        ]

    for i, row_data in enumerate(data, start=2):
        ts_cell = ws.cell(row=i, column=1, value=row_data[0])
        ws.cell(row=i, column=2, value=row_data[1]) 

        # Extracted Time: MOD removes the date portion, leaving only the time fraction
        time_cell = ws.cell(row=i, column=3, value=f"=MOD(A{i}, 1)")
        time_cell.number_format = time_format

        # Half-Hour Bin: VLOOKUP with TRUE finds the nearest interval below the extracted time
        bin_cell = ws.cell(row=i, column=4, value=f"=VLOOKUP(C{i}, $G$2:$G$49, 1, TRUE)")
        bin_cell.number_format = time_format

    # 3. Setup Summary Table (Columns I-J)
    ws['I1'] = "Time Bin"
    ws['J1'] = "Total Sales"
    ws['I1'].fill = header_fill
    ws['I1'].font = header_font
    ws['J1'].fill = header_fill
    ws['J1'].font = header_font

    for i in range(48):
        row = i + 2
        # Link to the reference table to ensure exact floating-point matching
        bin_cell = ws.cell(row=row, column=9, value=f"=$G${row}")
        bin_cell.number_format = time_format
        bin_cell.border = border

        # SUMIFS to aggregate values matching the calculated bins
        sum_cell = ws.cell(row=row, column=10, value=f"=SUMIFS(B:B, D:D, I{row})")
        sum_cell.border = border

    # 4. Add Intraday Trend Chart
    chart = LineChart()
    chart.title = "Sales by Half-Hour Interval"
    chart.style = 13
    chart.y_axis.title = "Total Sales"
    chart.x_axis.title = "Time of Day"

    # Data from J1:J49, Categories from I2:I49
    dat = Reference(ws, min_col=10, min_row=1, max_row=49)
    cats = Reference(ws, min_col=9, min_row=2, max_row=49)
    chart.add_data(dat, titles_from_data=True)
    chart.set_categories(cats)
    
    chart.width = 22
    chart.height = 12

    ws.add_chart(chart, "L2")

    # Layout sizing
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 15
