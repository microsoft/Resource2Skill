from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", tabs: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Creates an interactive dashboard shell with a left navigation sidebar and a 
    structured layout of widget containers (simulating shape cards).
    """
    if tabs is None:
        tabs = ["Dashboard", "Inputs", "Contacts", "Help"]
        
    # Ensure destination sheets exist so hyperlinks resolve correctly
    for tab in tabs:
        if tab not in wb.sheetnames:
            wb.create_sheet(tab)
            
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    
    # Standard theme fallback variables
    colors = {
        "sidebar": "1F4E78",     # Dark Blue
        "background": "F2F2F2",  # Light Gray
        "widget": "FFFFFF",      # White
        "text_main": "262626",   # Dark Gray
        "text_sub": "595959",    # Medium Gray
        "border": "D9D9D9"       # Light Gray Border
    }
    
    sidebar_fill = PatternFill(start_color=colors["sidebar"], end_color=colors["sidebar"], fill_type="solid")
    bg_fill = PatternFill(start_color=colors["background"], end_color=colors["background"], fill_type="solid")
    widget_fill = PatternFill(start_color=colors["widget"], end_color=colors["widget"], fill_type="solid")
    
    border_side = Side(border_style="thin", color=colors["border"])
    
    # 1. Paint global background
    for row in ws.iter_rows(min_row=1, max_row=35, min_col=1, max_col=30):
        for cell in row:
            cell.fill = bg_fill

    # 2. Build Navigation Sidebar (Column A)
    ws.column_dimensions['A'].width = 8
    for r in range(1, 36):
        ws.cell(row=r, column=1).fill = sidebar_fill
        
    icons = ["🏠", "📊", "✉️", "❓"]
    for i, tab in enumerate(tabs):
        icon = icons[i] if i < len(icons) else "📄"
        cell = ws.cell(row=4 + i*4, column=1)
        cell.value = icon
        cell.font = Font(size=20, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.hyperlink = f"#'{tab}'!A1"
        
    # 3. Define and paint widget containers (acting as chart drop-zones)
    widgets = [
        {"min_col": 3, "max_col": 20, "min_row": 2, "max_row": 4, "title": ""}, # Header box
        {"min_col": 3, "max_col": 8, "min_row": 6, "max_row": 11, "title": "Sales"},
        {"min_col": 9, "max_col": 14, "min_row": 6, "max_row": 11, "title": "Profit"},
        {"min_col": 15, "max_col": 20, "min_row": 6, "max_row": 11, "title": "# of Customers"},
        {"min_col": 3, "max_col": 14, "min_row": 13, "max_row": 26, "title": "2021-2022 Sales Trend"},
        {"min_col": 15, "max_col": 20, "min_row": 13, "max_row": 26, "title": "Customer Satisfaction"},
        {"min_col": 22, "max_col": 27, "min_row": 2, "max_row": 26, "title": "Sales by Country 2022"}
    ]
    
    for w in widgets:
        for row in range(w["min_row"], w["max_row"] + 1):
            for col in range(w["min_col"], w["max_col"] + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = widget_fill
                
                # Perimeter borders only to create the "card" effect
                top = border_side if row == w["min_row"] else None
                bottom = border_side if row == w["max_row"] else None
                left = border_side if col == w["min_col"] else None
                right = border_side if col == w["max_col"] else None
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)
        
        # Inject standard widget titles
        if w["title"]:
            title_cell = ws.cell(row=w["min_row"] + 1, column=w["min_col"] + 1)
            title_cell.value = w["title"]
            title_cell.font = Font(size=12, bold=True, color=colors["text_main"])
            
    # 4. Inject Dashboard Header text into the top container
    header_title = ws.cell(row=2, column=4)
    header_title.value = title
    header_title.font = Font(size=18, bold=True, color=colors["sidebar"])
    
    header_sub = ws.cell(row=3, column=4)
    header_sub.value = "Figures in millions of USD"
    header_sub.font = Font(size=10, italic=True, color=colors["text_sub"])
        
    # 5. Format spacing columns and disable gridlines
    ws.column_dimensions['B'].width = 2
    ws.column_dimensions['U'].width = 2
    ws.sheet_view.showGridLines = False
