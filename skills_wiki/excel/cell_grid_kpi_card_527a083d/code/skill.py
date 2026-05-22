def render(ws, anchor: str, *, title: str, subtitle: str, primary_value, secondary_value, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

    # Fallback palette lookup
    palettes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF"},
        "emerald_green": {"bg": "27AE60", "fg": "FFFFFF"},
        "slate_gray": {"bg": "2C3E50", "fg": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    
    # Typography hierarchy
    title_font = Font(color=palette["fg"], size=14, bold=True)
    sub_font = Font(color=palette["fg"], size=11)
    val_font = Font(color=palette["fg"], size=18, bold=True)
    sec_val_font = Font(color=palette["fg"], size=16, bold=True)
    
    center_align = Alignment(horizontal="center", vertical="center")
    
    row, col = coordinate_to_tuple(anchor)
    
    # Apply baseline fill to the entire 3x2 grid block
    for r in range(row, row + 3):
        for c in range(col, col + 2):
            ws.cell(row=r, column=c).fill = fill
    
    # Title (Row 1, merged Col 1 & 2)
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
    title_cell = ws.cell(row=row, column=col)
    title_cell.value = title
    title_cell.font = title_font
    title_cell.alignment = center_align
    
    # Subtitle (Row 2, Col 1)
    sub_cell = ws.cell(row=row+1, column=col)
    sub_cell.value = subtitle
    sub_cell.font = sub_font
    sub_cell.alignment = center_align
    
    # Primary Value (Row 3, Col 1)
    val_cell = ws.cell(row=row+2, column=col)
    val_cell.value = primary_value
    val_cell.font = val_font
    val_cell.alignment = center_align
    if isinstance(primary_value, (int, float)):
        val_cell.number_format = '"$"#,##0'
        
    # Secondary Value (Row 2-3, Col 2 merged)
    ws.merge_cells(start_row=row+1, start_column=col+1, end_row=row+2, end_column=col+1)
    sec_cell = ws.cell(row=row+1, column=col+1)
    sec_cell.value = secondary_value
    sec_cell.font = sec_val_font
    sec_cell.alignment = center_align
    if isinstance(secondary_value, (int, float)):
        sec_cell.number_format = '0%'
        
    # Format grid dimensions to simulate a card shape
    ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions[get_column_letter(col+1)].width = 12
    ws.row_dimensions[row].height = 25
    ws.row_dimensions[row+1].height = 18
    ws.row_dimensions[row+2].height = 25
    
    # Apply a solid border wrapper to emulate a unified shape outline
    card_border_side = Side(style='medium', color=palette["bg"])
    
    for r in range(row, row + 3):
        for c in range(col, col + 2):
            cell = ws.cell(row=r, column=c)
            b_left = card_border_side if c == col else None
            b_right = card_border_side if c == col + 1 else None
            b_top = card_border_side if r == row else None
            b_bottom = card_border_side if r == row + 2 else None
            cell.border = Border(left=b_left, right=b_right, top=b_top, bottom=b_bottom)
