from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import range_boundaries

def render_sheet(wb, sheet_name: str, *, title: str = "Heroic Insights Dashboard", theme: str = "dark_teal", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Conceal gridlines for an app-like feel
    ws.sheet_view.showGridLines = False
    
    # Theme palette (emulating the tutorial's aesthetic)
    themes = {
        "dark_teal": {
            "bg": "1A3B47", 
            "panel_bg": "FFFFFF", 
            "header_bg": "0D232A", 
            "header_fg": "FFFFFF", 
            "border": "2A5A6B", 
            "muted": "808080"
        },
        "corporate_blue": {
            "bg": "E6F0FA",
            "panel_bg": "FFFFFF",
            "header_bg": "1F4E78",
            "header_fg": "FFFFFF",
            "border": "B0C4DE",
            "muted": "808080"
        }
    }
    t = themes.get(theme, themes["corporate_blue"])
    
    bg_fill = PatternFill("solid", fgColor=t["bg"])
    panel_fill = PatternFill("solid", fgColor=t["panel_bg"])
    header_fill = PatternFill("solid", fgColor=t["header_bg"])
    
    # 1. Paint entire dashboard backdrop (A1:N35)
    for row in ws.iter_rows(min_row=1, max_row=35, min_col=1, max_col=14):
        for cell in row:
            cell.fill = bg_fill
            
    # 2. Header Banner
    ws.merge_cells("B2:M3")
    header_cell = ws["B2"]
    header_cell.value = title
    header_cell.font = Font(name="Calibri", size=24, bold=True, color=t["header_fg"])
    header_cell.fill = header_fill
    header_cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # 3. Render Layout Panels
    panels = [
        ("B5:C32", "Sidebar / Filters"),
        ("D5:M16", "Main Trend Chart"),
        ("D18:H32", "Secondary Chart 1"),
        ("I18:M32", "Secondary Chart 2")
    ]
    
    for panel_range, label in panels:
        min_col, min_row, max_col, max_row = range_boundaries(panel_range)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = panel_fill
                
                # Apply outer borders to the edge of the panel
                top = Side(style="thin", color=t["border"]) if r == min_row else None
                bottom = Side(style="thin", color=t["border"]) if r == max_row else None
                left = Side(style="thin", color=t["border"]) if c == min_col else None
                right = Side(style="thin", color=t["border"]) if c == max_col else None
                
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)
        
        # Insert a placeholder label for the panel
        label_cell = ws.cell(row=min_row, column=min_col)
        label_cell.value = label
        label_cell.font = Font(color=t["muted"], italic=True)
        label_cell.alignment = Alignment(horizontal="left", vertical="top")

    # 4. Column sizing for standard dashboard proportions
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    for col in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
        ws.column_dimensions[col].width = 12
    ws.column_dimensions['N'].width = 2
