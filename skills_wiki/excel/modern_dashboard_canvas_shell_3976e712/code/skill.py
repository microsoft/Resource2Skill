from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a modern dashboard layout with a sidebar and container 'cards'.
    """
    ws = wb.create_sheet(sheet_name)

    # Standard clean UI colors (fallback palette mimicking the video's design)
    bg_color = "F3F4F6"      # Soft light gray dashboard background
    sidebar_color = "1F2937" # Dark slate navigation sidebar
    card_color = "FFFFFF"    # White card container background
    border_color = "E5E7EB"  # Subtle gray card border
    text_dark = "111827"
    text_light = "FFFFFF"

    # 1. Clean Canvas Setup
    ws.sheet_view.showGridLines = False

    # Apply background to typical viewable area (e.g., rows 1-50, cols A-T)
    bg_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=20):
        for cell in row:
            cell.fill = bg_fill

    # Adjust standard column widths to make the grid more flexible for cards
    for col in range(2, 21):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # 2. Sidebar Navigation (Column A)
    ws.column_dimensions['A'].width = 8
    sidebar_fill = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    for row in range(1, 51):
        ws.cell(row=row, column=1).fill = sidebar_fill

    # Add generic navigation icons/text
    nav_items = ["🏠", "📊", "⚙️", "❓"]
    for i, item in enumerate(nav_items):
        cell = ws.cell(row=4 + (i * 4), column=1, value=item)
        cell.font = Font(color=text_light, size=16)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Dashboard Title
    title_cell = ws.cell(row=2, column=3, value=title)
    title_cell.font = Font(size=24, bold=True, color=text_dark)

    # 4. Card Generator Helper
    def create_card(start_row: int, start_col: int, end_row: int, end_col: int, card_title: str):
        card_fill = PatternFill(start_color=card_color, end_color=card_color, fill_type="solid")
        border_side = Side(style='thin', color=border_color)

        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill

                # Apply perimeter borders only to the outside edges of the card
                left = border_side if c == start_col else None
                right = border_side if c == end_col else None
                top = border_side if r == start_row else None
                bottom = border_side if r == end_row else None
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)

        # Card Header
        header_cell = ws.cell(row=start_row, column=start_col + 1, value=card_title)
        header_cell.font = Font(size=12, bold=True, color=text_dark)
        ws.row_dimensions[start_row].height = 24
        header_cell.alignment = Alignment(vertical="center")

    # 5. Lay out the Dashboard Grid
    # KPI Row
    create_card(5, 3, 9, 7, "Total Sales")
    create_card(5, 9, 9, 13, "Net Profit")
    create_card(5, 15, 9, 19, "Active Customers")

    # Main Content Row
    create_card(11, 3, 25, 13, "Sales Trend (2022-2023)")
    create_card(11, 15, 25, 19, "Regional Performance")
