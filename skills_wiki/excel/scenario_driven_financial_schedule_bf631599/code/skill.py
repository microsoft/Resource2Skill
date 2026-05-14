from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a Scenario-Driven Income Statement with dynamic forecasting.
    
    :param ws: openpyxl Worksheet object.
    :param anchor: Top-left cell coordinate (e.g., "B2") for the component.
    :param theme: Standard palette name.
    """
    # 1. Setup Theme & Palette
    palettes = {
        "corporate_blue": {"header_bg": "003366", "header_fg": "FFFFFF", "subtotal_bg": "E6F0FA", "accent": "4F81BD"},
        "modern_emerald": {"header_bg": "004D40", "header_fg": "FFFFFF", "subtotal_bg": "E0F2F1", "accent": "00897B"},
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    xy = coordinate_from_string(anchor)
    start_col = column_index_from_string(xy[0])
    start_row = xy[1]

    # 2. Define Layout Map
    c_hist1, c_hist2 = start_col + 1, start_col + 2
    c_fcst1, c_fcst2, c_fcst3 = start_col + 3, start_col + 4, start_col + 5
    
    r_hdr = start_row
    r_rev = r_hdr + 1
    r_cogs = r_hdr + 2
    r_gp = r_hdr + 3
    r_sga = r_hdr + 4
    r_opinc = r_hdr + 5
    r_tax = r_hdr + 6
    r_ni = r_hdr + 7
    r_assump_hdr = r_hdr + 9
    r_toggle = r_hdr + 10
    r_gr_best = r_hdr + 11
    r_gr_base = r_hdr + 12
    r_gr_worst = r_hdr + 13
    r_gr_active = r_hdr + 14
    r_cogs_pct = r_hdr + 15
    r_sga_pct = r_hdr + 16
    r_tax_pct = r_hdr + 17

    labels = {
        r_hdr: "Income Statement",
        r_rev: "Revenue",
        r_cogs: "Cost of Goods Sold",
        r_gp: "Gross Profit",
        r_sga: "Selling, General & Admin",
        r_opinc: "Operating Income",
        r_tax: "Taxes",
        r_ni: "Net Income",
        r_assump_hdr: "Assumptions & Drivers",
        r_toggle: "Live Scenario (1=Best, 2=Base, 3=Worst)",
        r_gr_best: "Revenue Growth - Best Case",
        r_gr_base: "Revenue Growth - Base Case",
        r_gr_worst: "Revenue Growth - Worst Case",
        r_gr_active: "Active Revenue Growth (Calculated)",
        r_cogs_pct: "COGS % of Revenue",
        r_sga_pct: "SG&A % of Revenue",
        r_tax_pct: "Tax Rate"
    }

    # Write Labels
    for r_idx, text in labels.items():
        ws.cell(row=r_idx, column=start_col, value=text)

    # Write Timeline (Years with custom formats)
    years = [2023, 2024, 2025, 2026, 2027]
    for i, year in enumerate(years):
        col = start_col + 1 + i
        cell = ws.cell(row=r_hdr, column=col, value=year)
        if i < 2:
            cell.number_format = '0"A"' # Formats exactly as e.g., 2023A
        else:
            cell.number_format = '0"E"' # Formats exactly as e.g., 2025E

    # 3. Populate Historical Hardcodes
    hist_data = {
        r_rev: [5210, 5435],
        r_cogs: [3345, 3350],
        r_sga: [850, 870],
        r_tax: [228, 186]
    }
    for r_idx, vals in hist_data.items():
        ws.cell(row=r_idx, column=c_hist1, value=vals[0])
        ws.cell(row=r_idx, column=c_hist2, value=vals[1])

    # Populate Assumptions Hardcodes (Anchored in c_hist1 for simplicity)
    assumptions = {
        r_toggle: 2,           # Base Case Default
        r_gr_best: 0.067,
        r_gr_base: 0.047,
        r_gr_worst: 0.027,
        r_cogs_pct: 0.616,
        r_sga_pct: 0.160,
        r_tax_pct: 0.250
    }
    for r_idx, val in assumptions.items():
        ws.cell(row=r_idx, column=c_hist1, value=val)

    # 4. Generate Core Formulas
    assump_col = get_column_letter(c_hist1)
    
    # Universal Structure Formulas (Applies to History + Forecasts)
    for c in range(c_hist1, c_fcst3 + 1):
        col_let = get_column_letter(c)
        ws.cell(row=r_gp, column=c, value=f"={col_let}{r_rev}-{col_let}{r_cogs}")
        ws.cell(row=r_opinc, column=c, value=f"={col_let}{r_gp}-{col_let}{r_sga}")
        ws.cell(row=r_ni, column=c, value=f"={col_let}{r_opinc}-{col_let}{r_tax}")

    # Forecast-only Projection Formulas
    for c in range(c_fcst1, c_fcst3 + 1):
        col_let = get_column_letter(c)
        prev_let = get_column_letter(c - 1)

        # Scenario Engine: Maps the integer toggle to the corresponding assumption cell
        choose_fmla = f"=CHOOSE(${assump_col}${r_toggle}, ${assump_col}${r_gr_best}, ${assump_col}${r_gr_base}, ${assump_col}${r_gr_worst})"
        ws.cell(row=r_gr_active, column=c, value=choose_fmla)

        # Projections
        ws.cell(row=r_rev, column=c, value=f"={prev_let}{r_rev}*(1+{col_let}{r_gr_active})")
        ws.cell(row=r_cogs, column=c, value=f"={col_let}{r_rev}*${assump_col}${r_cogs_pct}")
        ws.cell(row=r_sga, column=c, value=f"={col_let}{r_rev}*${assump_col}${r_sga_pct}")
        ws.cell(row=r_tax, column=c, value=f"={col_let}{r_opinc}*${assump_col}${r_tax_pct}")

    # 5. Formatting & Styles
    fmt_acct = '#,##0'
    fmt_pct = '0.0%'

    # Apply standard accounting numbers
    for r in [r_rev, r_cogs, r_gp, r_sga, r_opinc, r_tax, r_ni]:
        for c in range(c_hist1, c_fcst3 + 1):
            ws.cell(row=r, column=c).number_format = fmt_acct

    # Apply percentages to drivers
    for r in [r_gr_best, r_gr_base, r_gr_worst, r_cogs_pct, r_sga_pct, r_tax_pct]:
        ws.cell(row=r, column=c_hist1).number_format = fmt_pct
    for c in range(c_fcst1, c_fcst3 + 1):
        ws.cell(row=r_gr_active, column=c).number_format = fmt_pct

    # Data Validation on Toggle Cell
    toggle_cell = ws.cell(row=r_toggle, column=c_hist1)
    toggle_cell.number_format = '0'
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False, showErrorMessage=True)
    dv.error = "Please select scenario 1 (Best), 2 (Base), or 3 (Worst)"
    ws.add_data_validation(dv)
    dv.add(toggle_cell)

    # Visual Layout
    header_fill = PatternFill("solid", fgColor=palette["header_bg"])
    header_font = Font(bold=True, color=palette["header_fg"])
    tb_border = Border(top=Side(style='thin'), bottom=Side(style='thin'))

    # Format Headers
    for r_h in [r_hdr, r_assump_hdr]:
        ws.cell(row=r_h, column=start_col).font = header_font
        ws.cell(row=r_h, column=start_col).fill = header_fill
        for c in range(c_hist1, c_fcst3 + 1):
            cell = ws.cell(row=r_h, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    # Format Subtotals
    for r in [r_gp, r_opinc, r_ni]:
        ws.cell(row=r, column=start_col).font = Font(bold=True)
        for c in range(c_hist1, c_fcst3 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True)
            cell.border = tb_border

    # Format Toggle Switch Cell
    toggle_cell.fill = PatternFill("solid", fgColor="FFF2CC") # Light yellow attention grabber
    toggle_cell.border = Border(outline=True, top=Side(style='thin'), bottom=Side(style='thin'), left=Side(style='thin'), right=Side(style='thin'))
    toggle_cell.font = Font(bold=True, color="C00000")
    toggle_cell.alignment = Alignment(horizontal="center")

    # Column Widths & View Setup
    ws.column_dimensions[get_column_letter(start_col)].width = 36
    for c in range(c_hist1, c_fcst3 + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    # Group the assumption blocks for clean presentation
    ws.row_dimensions.group(r_gr_best, r_tax_pct, hidden=False)
    
    # Optional: Freeze panes (keep labels & header visible)
    ws.freeze_panes = ws.cell(row=r_rev, column=c_hist1)
