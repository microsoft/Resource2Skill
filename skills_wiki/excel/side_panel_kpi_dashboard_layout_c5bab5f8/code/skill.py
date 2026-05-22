from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, kpis: list[dict], theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a structural dashboard layout with a dark side-panel for KPIs.
    
    Expected kpis format:
    [
        {"label": "Total Revenue", "value": 649019, "format": "$#,##0.0,\"k\""},
        {"label": "Total Orders", "value": 2400, "format": "#,##0"},
        {"label": "Avg. Days to Deliver", "value": 2.3, "format": "0.0"}
    ]
    """
    # 1. Theme setup (simulated palette load)
    colors = {
        "panel_bg": "1E4B3D",    # Deep Forest Green (matching tutorial)
        "panel_fg": "FFFFFF",    # White
        "canvas_bg": "F8FAFC",   # Very Light Gray
        "title_fg": "0F172A",    # Slate 900
        "kpi_accent": "A7F3D0"   # Light Green Accent
    }
    
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # 2. Layout Grid & Margins
    ws.column_dimensions['A'].width = 2   # Left outer gutter
    ws.column_dimensions['B'].width = 22  # KPI Panel
    ws.column_dimensions['C'].width = 3   # Inner gutter separating panel and canvas
    
    # Expand main canvas columns to give breathing room for future charts
    for col_idx in range(4, 15):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
        
    # 3. Apply Background Fills
    panel_fill = PatternFill(start_color=colors["panel_bg"], end_color=colors["panel_bg"], fill_type="solid")
    canvas_fill = PatternFill(start_color=colors["canvas_bg"], end_color=colors["canvas_bg"], fill_type="solid")
    
    for row in range(1, 45): # Apply to standard visible screen area
        for col in range(1, 15):
            cell = ws.cell(row=row, column=col)
            if col == 2:
                cell.fill = panel_fill
            else:
                cell.fill = canvas_fill

    # 4. Brand / Title in the Panel
    brand_cell = ws.cell(row=2, column=2, value=title.upper())
    brand_cell.font = Font(name="Arial", bold=True, size=15, color=colors["panel_fg"])
    brand_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 5. Inject KPIs vertically into the side panel
    start_row = 6
    for kpi in kpis:
        # KPI Label
        lbl_cell = ws.cell(row=start_row, column=2, value=kpi.get("label", "Metric"))
        lbl_cell.font = Font(bold=True, size=10, color=colors["kpi_accent"])
        lbl_cell.alignment = Alignment(indent=1)
        
        # KPI Value
        val_cell = ws.cell(row=start_row + 1, column=2, value=kpi.get("value", 0))
        val_cell.font = Font(bold=True, size=22, color=colors["panel_fg"])
        val_cell.alignment = Alignment(indent=1)
        
        # Apply custom number formatting (e.g. thousand separators/abbreviations)
        if "format" in kpi:
            val_cell.number_format = kpi["format"]
            
        start_row += 4
        
    # 6. Main Canvas Title
    main_title = ws.cell(row=2, column=4, value="Dashboard Overview")
    main_title.font = Font(size=18, bold=True, color=colors["title_fg"])
    
    # Subtitle or instructions
    sub_title = ws.cell(row=3, column=4, value="Key trends and breakdowns for the selected period.")
    sub_title.font = Font(size=11, color="64748B", italic=True)
