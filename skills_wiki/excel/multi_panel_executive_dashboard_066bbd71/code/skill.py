def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Clean existing default sheets
    for sheet in wb.sheetnames:
        del wb[sheet]
        
    # Theme palette fallback map
    theme_colors = {
        "corporate_blue": {"primary": "1F4E78", "text": "FFFFFF"},
        "executive_dark": {"primary": "262626", "text": "FFFFFF"},
        "emerald_green": {"primary": "27AE60", "text": "FFFFFF"}
    }
    colors = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    # 2. Create a hidden Calc sheet for driving chart data
    calc_ws = wb.create_sheet("CalcData")
    calc_ws.sheet_state = 'hidden'
    
    # Chart 1 Data: Stacked Column (Row 1-5)
    calc_ws.append(["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"])
    calc_ws.append(["India", 62349, 23621, 21028])
    calc_ws.append(["United Kingdom", 46530, 26731, 11497])
    calc_ws.append(["United States", 36657, 32910, 22260])
    calc_ws.append(["Malaysia", 46587, 20452, 17536])
    
    calc_ws.append([]) # Spacer row
    
    # Chart 2 Data: Line 1 (Row 7-11)
    calc_ws.append(["Month", "Units Sold"])
    calc_ws.append(["Sep", 50601])
    calc_ws.append(["Oct", 95622])
    calc_ws.append(["Nov", 65481])
    calc_ws.append(["Dec", 52970])
    
    calc_ws.append([]) # Spacer row
    
    # Chart 3 Data: Line 2 (Row 13-17)
    calc_ws.append(["Month", "Profit"])
    calc_ws.append(["Sep", 124812])
    calc_ws.append(["Oct", 228275])
    calc_ws.append(["Nov", 160228])
    calc_ws.append(["Dec", 136337])
    
    # 3. Create clean Dashboard presentation sheet
    dash_ws = wb.create_sheet("Dashboard", 0)
    dash_ws.sheet_view.showGridLines = False
    dash_ws.column_dimensions['A'].width = 3  # Margin padding
    
    # Build Top Title Banner
    dash_ws.merge_cells("A1:P3")
    banner_cell = dash_ws["A1"]
    banner_cell.value = title
    banner_cell.font = Font(size=24, bold=True, color=colors["text"])
    banner_cell.fill = PatternFill(fill_type="solid", start_color=colors["primary"])
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add Chart 1: Stacked Bar
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    chart1.style = 10
    
    data1 = Reference(calc_ws, min_col=2, min_row=1, max_col=4, max_row=5)
    cats1 = Reference(calc_ws, min_col=1, min_row=2, max_row=5)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    
    chart1.width = 16
    chart1.height = 11
    dash_ws.add_chart(chart1, "B5")
    
    # Add Chart 2: Line (Units Sold)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.style = 13
    
    data2 = Reference(calc_ws, min_col=2, min_row=7, max_col=2, max_row=11)
    cats2 = Reference(calc_ws, min_col=1, min_row=8, max_row=11)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.legend = None  # Clean up view since there's only one metric
    
    chart2.width = 13
    chart2.height = 5.3
    dash_ws.add_chart(chart2, "J5")
    
    # Add Chart 3: Line (Profit by month)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.style = 13
    
    data3 = Reference(calc_ws, min_col=2, min_row=13, max_col=2, max_row=17)
    cats3 = Reference(calc_ws, min_col=1, min_row=14, max_row=17)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.legend = None
    
    chart3.width = 13
    chart3.height = 5.3
    dash_ws.add_chart(chart3, "J16")
    
    # Finalize active view
    wb.active = dash_ws
