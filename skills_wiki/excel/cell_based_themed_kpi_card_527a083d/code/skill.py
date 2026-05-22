from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", main_label: str = "Revenue", main_val: float = 369989.0, sub_val: float = 0.05, theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Define Fallback Theme Palette
    themes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF", "badge_bg": "2F75B5"},
        "executive_dark": {"bg": "262626", "fg": "FFFFFF", "badge_bg": "595959"},
        "forest_green": {"bg": "375623", "fg": "FFFFFF", "badge_bg": "548235"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 2. Calculate layout coordinates (2 columns x 3 rows block)
    start_row, start_col = coordinate_to_tuple(anchor)
    end_row = start_row + 2
    end_col = start_col + 1

    # 3. Inject Values
    ws.cell(row=start_row, column=start_col, value=title)
    ws.cell(row=start_row+1, column=start_col, value=main_label)
    ws.cell(row=start_row+2, column=start_col, value=main_val)
    ws.cell(row=start_row+1, column=start_col+1, value=sub_val)

    # 4. Merge Cells for Layout
    # Title spans the top row
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col) 
    # Badge spans the bottom two rows on the right
    ws.merge_cells(start_row=start_row+1, start_column=start_col+1, end_row=end_row, end_column=end_col)

    # 5. Setup Styles
    card_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    badge_fill = PatternFill(start_color=palette["badge_bg"], end_color=palette["badge_bg"], fill_type="solid")
    
    # Simulates the distinct shape edge for the badge
    badge_border = Border(
        left=Side(style='thick', color=palette["bg"]),
        top=Side(style='thick', color=palette["bg"])
    )

    # Apply base fills to the grid block
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            if c == start_col + 1 and r > start_row:
                cell.fill = badge_fill
                cell.border = badge_border
            else:
                cell.fill = card_fill

    # 6. Apply Typography & Number Formats
    # Title
    title_cell = ws.cell(row=start_row, column=start_col)
    title_cell.font = Font(name="Calibri", size=14, bold=True, color=palette["fg"])
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Main Label
    label_cell = ws.cell(row=start_row+1, column=start_col)
    label_cell.font = Font(name="Calibri", size=10, color=palette["fg"])
    label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Main Value
    val_cell = ws.cell(row=start_row+2, column=start_col)
    val_cell.font = Font(name="Calibri", size=16, bold=True, color=palette["fg"])
    val_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    val_cell.number_format = '"$"#,##0'

    # Secondary Badge Value
    badge_cell = ws.cell(row=start_row+1, column=start_col+1)
    badge_cell.font = Font(name="Calibri", size=18, bold=True, color=palette["fg"])
    badge_cell.alignment = Alignment(horizontal="center", vertical="center")
    badge_cell.number_format = '0%'

    # 7. Size Columns and Rows to create "Card" proportions
    ws.column_dimensions[get_column_letter(start_col)].width = 18
    ws.column_dimensions[get_column_letter(start_col+1)].width = 12
    
    ws.row_dimensions[start_row].height = 25
    ws.row_dimensions[start_row+1].height = 18
    ws.row_dimensions[start_row+2].height = 30
