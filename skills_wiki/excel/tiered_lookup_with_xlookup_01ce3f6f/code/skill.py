from openpyxl.styles import Font, PatternFill
from openpyxl.utils import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, theme: str = "corporate_blue", data: list[dict] = None, lookup_tiers: list[dict] = None, **kwargs) -> None:
    """
    Renders an employee bonus calculation table using XLOOKUP's 'next smaller item' feature.
    """
    # 1. Setup Defaults and Theme
    if data is None:
        data = [
            {"name": "Kim West", "salary": 60200},
            {"name": "James Willard", "salary": 39627},
            {"name": "Stevie Bridge", "salary": 93668},
            {"name": "Roger Mun", "salary": 134000},
            {"name": "Natalie Porter", "salary": 45000}
        ]
        
    if lookup_tiers is None:
        # Note: XLOOKUP does NOT require this list to be sorted!
        lookup_tiers = [
            {"threshold": 100000, "bonus": 0.15},
            {"threshold": 10000, "bonus": 0.0},
            {"threshold": 60000, "bonus": 0.10},
            {"threshold": 50000, "bonus": 0.08},
            {"threshold": 30000, "bonus": 0.05}
        ]

    theme_palettes = {
        "corporate_blue": {"header_bg": "003366"},
        "emerald_green": {"header_bg": "006633"},
        "warm_sunset": {"header_bg": "CC5500"}
    }
    palette = theme_palettes.get(theme, theme_palettes["corporate_blue"])
    
    header_fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    col_str, row_str = coordinate_from_string(anchor)
    start_col = column_index_from_string(col_str)
    start_row = int(row_str)

    # 2. Render the Reference/Lookup Table (Offset to the right)
    tier_col_start = start_col + 4
    
    ws.cell(row=start_row, column=tier_col_start, value="Salary Threshold").fill = header_fill
    ws.cell(row=start_row, column=tier_col_start).font = header_font
    ws.cell(row=start_row, column=tier_col_start+1, value="Bonus %").fill = header_fill
    ws.cell(row=start_row, column=tier_col_start+1).font = header_font

    for i, tier in enumerate(lookup_tiers):
        r = start_row + 1 + i
        thresh_cell = ws.cell(row=r, column=tier_col_start, value=tier["threshold"])
        thresh_cell.number_format = '"$"#,##0'
        
        bonus_val_cell = ws.cell(row=r, column=tier_col_start+1, value=tier["bonus"])
        bonus_val_cell.number_format = '0%'

    # Calculate absolute bounds for the XLOOKUP formula
    threshold_col_ltr = get_column_letter(tier_col_start)
    bonus_col_ltr = get_column_letter(tier_col_start + 1)
    range_start = start_row + 1
    range_end = start_row + len(lookup_tiers)
    
    threshold_range = f"${threshold_col_ltr}${range_start}:${threshold_col_ltr}${range_end}"
    bonus_range = f"${bonus_col_ltr}${range_start}:${bonus_col_ltr}${range_end}"

    # 3. Render the Main Data Table with XLOOKUP
    headers = ["Employee Name", "Salary", "Bonus Rate"]
    for i, h in enumerate(headers):
        c = ws.cell(row=start_row, column=start_col + i, value=h)
        c.fill = header_fill
        c.font = header_font

    for i, emp in enumerate(data):
        r = start_row + 1 + i
        ws.cell(row=r, column=start_col, value=emp["name"])
        
        sal_cell = ws.cell(row=r, column=start_col+1, value=emp["salary"])
        sal_cell.number_format = '"$"#,##0'

        # XLOOKUP syntax: =XLOOKUP(lookup_value, lookup_array, return_array, [if_not_found], [match_mode])
        # match_mode=-1 -> "Exact match or next smaller item"
        lookup_target = f"{get_column_letter(start_col+1)}{r}"
        formula = f"=XLOOKUP({lookup_target}, {threshold_range}, {bonus_range}, 0, -1)"
        
        bonus_cell = ws.cell(row=r, column=start_col+2, value=formula)
        bonus_cell.number_format = '0%'

    # 4. Cleanup & Formatting
    ws.column_dimensions[get_column_letter(start_col)].width = 18
    ws.column_dimensions[get_column_letter(start_col+1)].width = 14
    ws.column_dimensions[get_column_letter(start_col+2)].width = 14
    ws.column_dimensions[threshold_col_ltr].width = 18
    ws.column_dimensions[bonus_col_ltr].width = 12
