from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str = "Executive Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean, interactive-looking dashboard shell containing multiple charts 
    arranged in a grid, drawing data from a hidden backing sheet.
    """
    # 1. Setup Dashboard Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
    
    # Hide gridlines for a clean dashboard look
    ws.sheet_view.showGridLines = False
    
    # Simple theme palette fallback
    palettes = {
        "corporate_blue": {"primary": "003366", "accent": "4F81BD", "bg": "F2F2F2", "text": "FFFFFF"},
        "vibrant_red": {"primary": "C00000", "accent": "FF0000", "bg": "F2F2F2", "text": "FFFFFF"},
        "forest_green": {"primary": "2E7D32", "accent": "4CAF50", "bg": "F2F2F2", "text": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    # 2. Render Header
    ws.merge_cells("A1:N3")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(name="Arial", size=24, bold=True, color=palette["text"])
    header_cell.fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Setup Data Sheet (Hidden)
    data_sheet_name = f"{sheet_name}_Data"
    if data_sheet_name in wb.sheetnames:
        data_ws = wb[data_sheet_name]
    else:
        data_ws = wb.create_sheet(data_sheet_name)
        data_ws.sheet_state = 'hidden' # Hide raw data from the end user
        
        # Populate Mock Data for Charts
        
        # --- Dataset 1: Monthly Revenue Trend ---
        trend_data = [
            ("Month", "Revenue"),
            ("Jan", 15000), ("Feb", 18000), ("Mar", 21000), 
            ("Apr", 20000), ("May", 25000), ("Jun", 27000),
            ("Jul", 30000), ("Aug", 29000), ("Sep", 32000),
            ("Oct", 35000), ("Nov", 38000), ("Dec", 42000)
        ]
        for r, row in enumerate(trend_data, 1):
            data_ws.cell(row=r, column=1, value=row[0])
            data_ws.cell(row=r, column=2, value=row[1])
            
        # --- Dataset 2: Category Comparison ---
        cat_data = [
            ("Year", "Hoodies", "T-shirts"),
            ("2023", 120000, 95000),
            ("2024", 150000, 110000)
        ]
        for r, row in enumerate(cat_data, 1):
            data_ws.cell(row=r, column=4, value=row[0])
            data_ws.cell(row=r, column=5, value=row[1])
            data_ws.cell(row=r, column=6, value=row[2])
            
        # --- Dataset 3: Top 5 Performers ---
        state_data = [
            ("State", "Profit"),
            ("California", 38000),
            ("Texas", 34000),
            ("New York", 31000),
            ("Florida", 29000),
            ("Illinois", 25000)
        ]
        for r, row in enumerate(state_data, 1):
            data_ws.cell(row=r, column=8, value=row[0])
            data_ws.cell(row=r, column=9, value=row[1])
            
    # 4. Create and Position Charts
    
    # Chart 1: Line Chart (Trend) anchored at B5
    line_chart = LineChart()
    line_chart.title = "Monthly Revenue Trend"
    line_chart.style = 13
    line_chart.y_axis.title = "Revenue ($)"
    data_ref1 = Reference(data_ws, min_col=2, min_row=1, max_row=13)
    cats_ref1 = Reference(data_ws, min_col=1, min_row=2, max_row=13)
    line_chart.add_data(data_ref1, titles_from_data=True)
    line_chart.set_categories(cats_ref1)
    line_chart.width = 16
    line_chart.height = 8
    ws.add_chart(line_chart, "B5")
    
    # Chart 2: Clustered Column Chart (Category) anchored at J5
    bar_chart1 = BarChart()
    bar_chart1.type = "col"
    bar_chart1.style = 10
    bar_chart1.title = "Units Sold: Hoodies vs T-shirts"
    bar_chart1.grouping = "clustered"
    data_ref2 = Reference(data_ws, min_col=5, max_col=6, min_row=1, max_row=3)
    cats_ref2 = Reference(data_ws, min_col=4, min_row=2, max_row=3)
    bar_chart1.add_data(data_ref2, titles_from_data=True)
    bar_chart1.set_categories(cats_ref2)
    bar_chart1.width = 12
    bar_chart1.height = 8
    ws.add_chart(bar_chart1, "J5")
    
    # Chart 3: Clustered Column Chart (Top 5) anchored at B16
    bar_chart2 = BarChart()
    bar_chart2.type = "col"
    bar_chart2.style = 10
    bar_chart2.title = "Top 5 States by Profit"
    bar_chart2.legend = None # Hide legend for single-series bar charts
    data_ref3 = Reference(data_ws, min_col=9, min_row=1, max_row=6)
    cats_ref3 = Reference(data_ws, min_col=8, min_row=2, max_row=6)
    bar_chart2.add_data(data_ref3, titles_from_data=True)
    bar_chart2.set_categories(cats_ref3)
    bar_chart2.width = 16
    bar_chart2.height = 8
    ws.add_chart(bar_chart2, "B16")
