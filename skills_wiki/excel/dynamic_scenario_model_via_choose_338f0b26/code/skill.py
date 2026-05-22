from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Formatting Constants (Modeling Conventions)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    input_font = Font(color="0000FF") # Convention: Blue = Hard-coded input
    input_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Interactivity cue
    border_thin = Border(
        left=Side(style='thin'), right=Side(style='thin'), 
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # 1. Title and Scenario Toggle Setup
    ws["B2"] = title
    ws["B2"].font = Font(size=14, bold=True)
    
    ws["D2"] = "Active Scenario:"
    ws["D2"].font = bold_font
    ws["D2"].alignment = Alignment(horizontal="right")
    
    ws["E2"] = 1
    ws["E2"].fill = input_fill
    ws["E2"].border = border_thin
    ws["E2"].font = input_font
    ws["E2"].alignment = Alignment(horizontal="center")
    
    # Apply Data Validation dropdown to toggle cell
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    dv.error = 'Your entry is not in the list. Choose 1 or 2.'
    dv.errorTitle = 'Invalid Scenario'
    ws.add_data_validation(dv)
    dv.add(ws["E2"])
    
    # 2. Main Model Structure (Income Statement)
    headers = ["Line Item", "Year 1", "Year 2", "Year 3"]
    for col_num, header in enumerate(headers, start=2):
        cell = ws.cell(row=4, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    ws["B5"] = "Revenue"
    ws["B6"] = "COGS"
    ws["B7"] = "Gross Profit"
    ws["B7"].font = bold_font
    
    # Financial Formulas
    for col in range(3, 6):
        col_let = get_column_letter(col)
        ws[f"{col_let}5"] = f"={col_let}11*{col_let}12" # Revenue = Orders * AOV
        ws[f"{col_let}6"] = f"={col_let}11*15"         # COGS = Orders * $15 per unit
        ws[f"{col_let}7"] = f"={col_let}5-{col_let}6"  # Gross Profit = Revenue - COGS
        
    # 3. Dynamic Live Assumptions Block
    ws["B9"] = "Live Assumptions (Driven by Scenario)"
    ws["B9"].font = bold_font
    ws["B10"] = "Scenario Chosen:"
    ws["B11"] = "Orders"
    ws["B12"] = "Average Order Value (AOV)"
    
    for col in range(3, 6):
        col_let = get_column_letter(col)
        # CHOOSE pulls from Scenario 1 (Row 16/17) or Scenario 2 (Row 21/22) based on E2 toggle
        ws[f"{col_let}10"] = "=$E$2"
        ws[f"{col_let}11"] = f"=CHOOSE($E$2, {col_let}16, {col_let}21)"
        ws[f"{col_let}12"] = f"=CHOOSE($E$2, {col_let}17, {col_let}22)"
        
    # 4. Hard-coded Scenario Cases
    # Scenario 1 (Base Case)
    ws["B14"] = "Scenario 1: Base Case"
    ws["B14"].font = bold_font
    ws["B15"] = "Growth YoY"
    ws["B16"] = "Orders"
    ws["B17"] = "Average Order Value (AOV)"
    
    s1_orders = [3000, 3300, 3630]
    s1_aov = [39.95, 39.95, 40.50]
    for i, col in enumerate(range(3, 6)):
        if i == 0:
            ws.cell(row=15, column=col, value="-").alignment = Alignment(horizontal="center")
        else:
            ws.cell(row=15, column=col, value=0.10).font = input_font
            ws.cell(row=15, column=col).number_format = '0%'
        
        ws.cell(row=16, column=col, value=s1_orders[i]).font = input_font
        ws.cell(row=17, column=col, value=s1_aov[i]).font = input_font
        
    # Scenario 2 (Upside Case)
    ws["B19"] = "Scenario 2: Upside Case"
    ws["B19"].font = bold_font
    ws["B20"] = "Growth YoY"
    ws["B21"] = "Orders"
    ws["B22"] = "Average Order Value (AOV)"
    
    s2_orders = [3500, 4200, 5040]
    s2_aov = [45.00, 45.00, 48.00]
    for i, col in enumerate(range(3, 6)):
        if i == 0:
            ws.cell(row=20, column=col, value="-").alignment = Alignment(horizontal="center")
        else:
            ws.cell(row=20, column=col, value=0.20).font = input_font
            ws.cell(row=20, column=col).number_format = '0%'
            
        ws.cell(row=21, column=col, value=s2_orders[i]).font = input_font
        ws.cell(row=22, column=col, value=s2_aov[i]).font = input_font
        
    # 5. Number Formatting & Column Dimensions
    currency_rows = [5, 6, 7, 12, 17, 22]
    integer_rows = [11, 16, 21]
    
    for row in currency_rows:
        for col in range(3, 6):
            ws.cell(row=row, column=col).number_format = '"$"#,##0.00'
            
    for row in integer_rows:
        for col in range(3, 6):
            ws.cell(row=row, column=col).number_format = '#,##0'
            
    ws.column_dimensions['B'].width = 35
    for col in ['C', 'D', 'E']:
        ws.column_dimensions[col].width = 16
