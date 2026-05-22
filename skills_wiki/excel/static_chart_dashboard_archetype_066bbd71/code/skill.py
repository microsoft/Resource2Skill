import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import PatternFill, Font, Alignment

def render_workbook(wb, *, title: str = "Performance Dashboard", dashboard_data: dict = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a presentation dashboard with a clean layout, banner, and multiple charts.
    Data is stored on a hidden sheet to keep the dashboard pristine.
    """
    # 1. Establish Theme Palette Fallbacks
    primary_bg = "1F4E78" # Corporate Blue
    text_fg = "FFFFFF"
    
    if theme == "light":
        primary_bg = "F2F2F2"
        text_fg = "000000"

    # Default data mirroring the tutorial's domain
    if dashboard_data is None:
        dashboard_data = {
            "bar_chart": {
                "title": "Profit by Market & Cookie Type",
                "headers": ["Market", "Fortune Cookie", "Oatmeal Raisin", "Sugar"],
                "rows": [
                    ["India", 100000, 45000, 30000],
                    ["Philippines", 80000, 35000, 25000],
                    ["United Kingdom", 90000, 50000, 40000],
                    ["United States", 150000, 80000, 60000]
                ]
            },
            "line_chart_1": {
                "title": "Units Sold Each Month",
                "headers": ["Month", "Units"],
                "rows": [
                    ["Sep", 50601],
                    ["Oct", 95622],
                    ["Nov", 65481],
                    ["Dec", 52970]
                ]
            },
            "line_chart_2": {
                "title": "Profit By Month",
                "headers": ["Month", "Profit"],
                "rows": [
                    ["Sep", 124812],
                    ["Oct", 228275],
                    ["Nov", 160228],
                    ["Dec", 136337]
                ]
            }
        }

    # 2. Setup Dashboard Sheet
    if "Sheet" in wb.sheetnames:
        ws_dash = wb["Sheet"]
        ws_dash.title = "Dashboard"
    else:
        ws_dash = wb.create_sheet("Dashboard")
        
    # Create the "blank canvas" effect
    ws_dash.sheet_view.showGridLines = False
    
    # Render the Banner
    ws_dash.merge_cells("A1:R3")
    banner = ws_dash["A1"]
    banner.value = title
    banner.fill = PatternFill(start_color=primary_bg, end_color=primary_bg, fill_type="solid")
    banner.font = Font(color=text_fg, size=24, bold=True)
    banner.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Setup Hidden Data Sheet
    ws_data = wb.create_sheet("ChartData")
    ws_data.sheet_state = 'hidden'
    current_row = 1
    
    # Helper to sequentially write data and generate precise Chart References
    def write_chart_data(data_dict):
        nonlocal current_row
        start_row = current_row
        headers = data_dict.get("headers", [])
        rows = data_dict.get("rows", [])
        
        # Write headers
        for c_idx, h in enumerate(headers, start=1):
            ws_data.cell(row=start_row, column=c_idx, value=h)
            
        # Write rows
        for r_idx, row_data in enumerate(rows, start=start_row + 1):
            for c_idx, val in enumerate(row_data, start=1):
                ws_data.cell(row=r_idx, column=c_idx, value=val)
                
        end_row = start_row + len(rows)
        max_col = len(headers)
        
        cats = Reference(ws_data, min_col=1, min_row=start_row+1, max_row=end_row)
        data_ref = Reference(ws_data, min_col=2, min_row=start_row, max_col=max_col, max_row=end_row)
        
        current_row = end_row + 2
        return cats, data_ref

    # 4. Render Bar Chart
    if "bar_chart" in dashboard_data:
        b_cats, b_data = write_chart_data(dashboard_data["bar_chart"])
        bar_chart = BarChart()
        bar_chart.type = "col"
        bar_chart.grouping = "stacked"
        bar_chart.overlap = 100
        bar_chart.title = dashboard_data["bar_chart"].get("title", "")
        bar_chart.add_data(b_data, titles_from_data=True)
        bar_chart.set_categories(b_cats)
        bar_chart.width = 18
        bar_chart.height = 12
        bar_chart.graphical_properties.line.noFill = True # Hide chart border
        ws_dash.add_chart(bar_chart, "B6")

    # 5. Render Line Chart 1
    if "line_chart_1" in dashboard_data:
        l1_cats, l1_data = write_chart_data(dashboard_data["line_chart_1"])
        line1 = LineChart()
        line1.title = dashboard_data["line_chart_1"].get("title", "")
        line1.add_data(l1_data, titles_from_data=True)
        line1.set_categories(l1_cats)
        line1.legend = None # Remove legend for simple time series
        line1.width = 14
        line1.height = 8
        line1.graphical_properties.line.noFill = True
        ws_dash.add_chart(line1, "K6")
        
    # 6. Render Line Chart 2
    if "line_chart_2" in dashboard_data:
        l2_cats, l2_data = write_chart_data(dashboard_data["line_chart_2"])
        line2 = LineChart()
        line2.title = dashboard_data["line_chart_2"].get("title", "")
        line2.add_data(l2_data, titles_from_data=True)
        line2.set_categories(l2_cats)
        line2.legend = None
        line2.width = 14
        line2.height = 8
        line2.graphical_properties.line.noFill = True
        ws_dash.add_chart(line2, "K19")
