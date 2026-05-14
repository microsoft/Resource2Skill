def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Fallback Theme Palette
    theme_colors = {
        "corporate_blue": {"header_bg": "203764", "header_fg": "FFFFFF"},
        "emerald_green": {"header_bg": "004D40", "header_fg": "FFFFFF"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    # 2. Create the Data Layer
    ws_data = wb.active
    ws_data.title = "Dashboard Data"
    
    # Sample Data for Stacked Bar (Profit by Market & Product)
    bar_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Sugar Cookie"],
        ["India", 62349, 4872, 21028, 25085],
        ["Philippines", 54618, 7026, 22005, 8313],
        ["United Kingdom", 46530, 5220, 11497, 14620],
        ["United States", 36657, 6369, 22260, 9938]
    ]
    
    for row in bar_data:
        ws_data.append(row)
        
    ws_data.append([]) # Empty row gap
    
    # Sample Data for Line Chart (Units sold each month)
    line_start_row = ws_data.max_row + 1
    line_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    
    for row in line_data:
        ws_data.append(row)
        
    # 3. Create the Presentation Layer (Dashboard)
    ws_dash = wb.create_sheet(title="Dashboard", index=0)
    ws_dash.sheet_view.showGridLines = False
    
    # Header Banner
    ws_dash.merge_cells("A1:Q4")
    header_cell = ws_dash["A1"]
    header_cell.value = title
    header_cell.font = Font(size=28, bold=True, color=palette["header_fg"])
    header_cell.fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Layout Columns (Left sidebar reserved for interactive slicers)
    ws_dash.column_dimensions['A'].width = 2
    ws_dash.column_dimensions['B'].width = 22
    ws_dash.column_dimensions['C'].width = 2
    
    sidebar_title = ws_dash["B6"]
    sidebar_title.value = "Controls / Filters"
    sidebar_title.font = Font(bold=True, color="595959")
    sidebar_title.border = Border(bottom=Side(style="thick", color="595959"))
    
    # 5. Build and Anchor Charts
    # Stacked Bar Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.y_axis.title = "Profit ($)"
    bar_chart.height = 10
    bar_chart.width = 16
    bar_chart.legend.position = "b"
    
    bar_data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    bar_cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)
    
    ws_dash.add_chart(bar_chart, "D6")
    
    # Line Chart
    line_chart = LineChart()
    line_chart.title = "Units Sold Each Month"
    line_chart.style = 13 # Standard styling preset
    line_chart.height = 10
    line_chart.width = 16
    line_chart.legend.position = "b"
    
    line_data_ref = Reference(ws_data, min_col=2, min_row=line_start_row, max_row=line_start_row+4)
    line_cats_ref = Reference(ws_data, min_col=1, min_row=line_start_row+1, max_row=line_start_row+4)
    line_chart.add_data(line_data_ref, titles_from_data=True)
    line_chart.set_categories(line_cats_ref)
    
    ws_dash.add_chart(line_chart, "L6")
