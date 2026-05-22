import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list, table_headers: list, table_data: list, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # 1. Define Theme
    theme_colors = {
        "corporate_blue": {
            "primary": "2C3E50",
            "primary_light": "ECF0F1",
            "accent": "F39C12",
            "text_light": "FFFFFF",
            "text_dark": "34495E"
        },
        "purple_aspect": {
            "primary": "5C4084",      # Dark purple header
            "primary_light": "F2EFF5",# Light purple body
            "accent": "F4B41A",       # Gold
            "text_light": "FFFFFF",
            "text_dark": "5C4084"
        }
    }
    colors = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    fill_primary = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    fill_light = PatternFill(start_color=colors["primary_light"], end_color=colors["primary_light"], fill_type="solid")
    fill_accent = PatternFill(start_color=colors["accent"], end_color=colors["accent"], fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Fonts
    font_title = Font(name="Arial", size=28, color=colors["text_light"], bold=True)
    font_subtitle = Font(name="Arial", size=14, color=colors["accent"])
    font_kpi_icon = Font(name="Segoe UI Emoji", size=24, color=colors["text_light"])
    font_kpi_value = Font(name="Arial", size=18, color=colors["text_dark"], bold=True)
    font_kpi_label = Font(name="Arial", size=10, color=colors["text_dark"])
    font_th = Font(color=colors["text_light"], bold=True)
    
    # 2. Canvas Backgrounds
    # Header area: Rows 1 to 8
    for row in range(1, 9):
        ws.row_dimensions[row].height = 20
        for col in range(1, 22):
            ws.cell(row=row, column=col).fill = fill_primary

    # Body area: Rows 9 to 40
    for row in range(9, 41):
        ws.row_dimensions[row].height = 20
        for col in range(1, 22):
            ws.cell(row=row, column=col).fill = fill_light
            
    ws.sheet_view.showGridLines = False
    
    # 3. Titles
    ws["B2"] = title
    ws["B2"].font = font_title
    ws.row_dimensions[2].height = 36
    
    ws["B3"] = subtitle
    ws["B3"].font = font_subtitle
    
    # 4. KPI Cards
    start_col = 3
    kpi_row_start = 5
    kpi_row_end = 7
    thin = Side(border_style="thin", color="CCCCCC")
    border_box = Border(top=thin, left=thin, right=thin, bottom=thin)
    
    for i, kpi in enumerate(kpis):
        col_idx = start_col + (i * 4)
        c_icon = get_column_letter(col_idx)
        c_val1 = get_column_letter(col_idx + 1)
        c_val2 = get_column_letter(col_idx + 2)
        
        ws.column_dimensions[c_icon].width = 8
        ws.column_dimensions[c_val1].width = 10
        ws.column_dimensions[c_val2].width = 10
        
        # Format the whole 3x3 block background & border first
        for r in range(kpi_row_start, kpi_row_end + 1):
            ws.cell(row=r, column=col_idx).fill = fill_accent
            ws.cell(row=r, column=col_idx+1).fill = fill_white
            ws.cell(row=r, column=col_idx+2).fill = fill_white
            for c in range(col_idx, col_idx + 3):
                ws.cell(row=r, column=c).border = border_box
                
        # Icon
        ws.merge_cells(f"{c_icon}{kpi_row_start}:{c_icon}{kpi_row_end}")
        cell_icon = ws[f"{c_icon}{kpi_row_start}"]
        cell_icon.value = kpi.get("icon", "•")
        cell_icon.font = font_kpi_icon
        cell_icon.alignment = Alignment(horizontal="center", vertical="center")
        
        # Value
        ws.merge_cells(f"{c_val1}{kpi_row_start}:{c_val2}{kpi_row_end-1}")
        cell_val = ws[f"{c_val1}{kpi_row_start}"]
        cell_val.value = kpi.get("value", "")
        cell_val.font = font_kpi_value
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        
        # Label
        ws.merge_cells(f"{c_val1}{kpi_row_end}:{c_val2}{kpi_row_end}")
        cell_lbl = ws[f"{c_val1}{kpi_row_end}"]
        cell_lbl.value = kpi.get("label", "")
        cell_lbl.font = font_kpi_label
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        
    # 5. Data Table with Data Bars
    table_start_row = 10
    table_start_col = 3
    col_widths = [15, 12, 12, 12, 15]
    
    # Write Headers
    for c_idx, header in enumerate(table_headers):
        cell = ws.cell(row=table_start_row, column=table_start_col + c_idx)
        cell.value = header
        cell.fill = fill_primary
        cell.font = font_th
        cell.alignment = Alignment(horizontal="center")
        width = col_widths[c_idx] if c_idx < len(col_widths) else 12
        ws.column_dimensions[get_column_letter(table_start_col + c_idx)].width = width
        
    # Write Data
    for r_idx, row_data in enumerate(table_data):
        current_row = table_start_row + 1 + r_idx
        for c_idx, val in enumerate(row_data):
            cell = ws.cell(row=current_row, column=table_start_col + c_idx)
            cell.value = val
            cell.fill = fill_white
            cell.border = border_box
            
            # Format numeric columns
            if isinstance(val, (int, float)):
                if "Value" in table_headers[c_idx]:
                    cell.number_format = '"$"#,##0'
                else:
                    cell.number_format = '#,##0'
                
    # Conditional Formatting (Data Bars)
    last_row = table_start_row + len(table_data)
    
    # Deals Closed Data Bars (assumes Deals Closed is at index 3)
    col_deals = get_column_letter(table_start_col + 3)
    rule_deals = DataBarRule(start_type='min', end_type='max', color=f"FF{colors['primary']}")
    ws.conditional_formatting.add(f"{col_deals}{table_start_row+1}:{col_deals}{last_row}", rule_deals)
    
    # Deal Value Data Bars (assumes Deal Value is at index 4)
    col_value = get_column_letter(table_start_col + 4)
    rule_value = DataBarRule(start_type='min', end_type='max', color=f"FF{colors['accent']}")
    ws.conditional_formatting.add(f"{col_value}{table_start_row+1}:{col_value}{last_row}", rule_value)
