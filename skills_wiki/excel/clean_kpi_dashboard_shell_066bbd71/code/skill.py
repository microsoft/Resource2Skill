from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Generates a multi-sheet workbook containing a clean presentation dashboard 
    and a backend data sheet, populating stacked bar and line charts.
    """
    # 1. Setup Sheets
    dashboard_ws = wb.active
    dashboard_ws.title = "Dashboard"
    data_ws = wb.create_sheet("Data")

    # 2. Populate Data: Market / Cookie Profit
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Sugar"],
        ["India", 62349, 4872, 18560],
        ["United Kingdom", 46530, 5220, 14620],
        ["United States", 36657, 6368, 9937],
        ["Philippines", 54618, 7026, 8313]
    ]
    for row in market_data:
        data_ws.append(row)

    # 3. Populate Data: Monthly Units
    data_ws.append([])  # Spacer row
    monthly_units_start = data_ws.max_row + 1
    monthly_units_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for row in monthly_units_data:
        data_ws.append(row)

    # 4. Populate Data: Monthly Profit
    data_ws.append([])  # Spacer row
    monthly_profit_start = data_ws.max_row + 1
    monthly_profit_data = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for row in monthly_profit_data:
        data_ws.append(row)

    # 5. Build Dashboard Header & Layout
    dashboard_ws.sheet_view.showGridLines = False
    dashboard_ws.merge_cells("A1:O2")
    header_cell = dashboard_ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color="FFFFFF")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Simple theme hook fallback
    theme_colors = {
        "corporate_blue": "1F4E78",
        "emerald_green": "107C41",
        "crimson_red": "A4262C"
    }
    bg_color = theme_colors.get(theme, "1F4E78")
    header_cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

    # 6. Chart 1: Stacked Bar (Categorical Data)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 10  # Built-in Excel chart style
    
    data_ref = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=5)
    cats_ref = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    bar_chart.height = 11
    bar_chart.width = 16
    dashboard_ws.add_chart(bar_chart, "B4")

    # 7. Chart 2: Line Chart (Time Series 1)
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.style = 13
    line1.legend = None  # Remove legend for clean look
    
    l1_data = Reference(data_ws, min_col=2, min_row=monthly_units_start, max_row=monthly_units_start+4)
    l1_cats = Reference(data_ws, min_col=1, min_row=monthly_units_start+1, max_row=monthly_units_start+4)
    line1.add_data(l1_data, titles_from_data=True)
    line1.set_categories(l1_cats)
    line1.height = 7
    line1.width = 12
    dashboard_ws.add_chart(line1, "J4")

    # 8. Chart 3: Line Chart (Time Series 2)
    line2 = LineChart()
    line2.title = "Profit by month"
    line2.style = 13
    line2.legend = None
    
    l2_data = Reference(data_ws, min_col=2, min_row=monthly_profit_start, max_row=monthly_profit_start+4)
    l2_cats = Reference(data_ws, min_col=1, min_row=monthly_profit_start+1, max_row=monthly_profit_start+4)
    line2.add_data(l2_data, titles_from_data=True)
    line2.set_categories(l2_cats)
    line2.height = 7
    line2.width = 12
    dashboard_ws.add_chart(line2, "J15")
