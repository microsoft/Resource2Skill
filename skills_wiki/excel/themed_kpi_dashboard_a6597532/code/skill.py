from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", subtitle: str = "Evaluating Sales Agent Performance", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Theme color definitions
    themes = {
        "corporate_blue": {"header": "1F4E78", "bg": "F2F2F2", "accent": "00B050", "text": "FFFFFF"},
        "purple_gold": {"header": "4B286D", "bg": "F2EFF5", "accent": "FFC000", "text": "FFFFFF"}
    }
    t = themes.get(theme, themes["corporate_blue"]) 
    
    header_fill = PatternFill("solid", fgColor=t["header"])
    bg_fill = PatternFill("solid", fgColor=t["bg"])
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    
    # Apply background color to entire viewable area and hide gridlines
    ws.sheet_view.showGridLines = False
    for row in range(1, 40):
        for col in range(1, 20):
            ws.cell(row=row, column=col).fill = bg_fill

    # Render dark full-bleed header banner
    for row in range(1, 5):
        for col in range(1, 20):
            ws.cell(row=row, column=col).fill = header_fill
            
    ws.cell(row=2, column=2, value=title).font = Font(size=24, bold=True, color=t["text"])
    ws.cell(row=3, column=2, value=subtitle).font = Font(size=12, italic=True, color=t["accent"])
    
    # Define KPI data
    kpis = [
        {"label": "CALLS", "value": 16749, "fmt": "#,##0"},
        {"label": "REACHED", "value": 3328, "fmt": "#,##0"},
        {"label": "CLOSED", "value": 1203, "fmt": "#,##0"},
        {"label": "VALUE", "value": 646979, "fmt": "$#,##0"}
    ]
    
    # Render KPI Cards (Simulated shapes using formatted cells)
    col_idx = 2
    accent_border = Border(top=Side(style='thick', color=t["accent"]))
    box_border = Border(
        bottom=Side(style='thin', color="CCCCCC"),
        left=Side(style='thin', color="CCCCCC"),
        right=Side(style='thin', color="CCCCCC")
    )
    
    for kpi in kpis:
        ws.merge_cells(start_row=6, start_column=col_idx, end_row=7, end_column=col_idx+2)
        ws.merge_cells(start_row=8, start_column=col_idx, end_row=8, end_column=col_idx+2)
        
        # Large Value
        val_cell = ws.cell(row=6, column=col_idx, value=kpi["value"])
        val_cell.font = Font(size=20, bold=True, color=t["header"])
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.number_format = kpi["fmt"]
        
        # Subtitle Label
        lbl_cell = ws.cell(row=8, column=col_idx, value=kpi["label"])
        lbl_cell.font = Font(size=10, bold=True, color="595959")
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")
        
        # Apply KPI box fills and borders
        for r in range(6, 9):
            for c in range(col_idx, col_idx+3):
                cell = ws.cell(row=r, column=c)
                cell.fill = white_fill
                if r == 6:
                    cell.border = accent_border
                elif r == 8:
                    cell.border = box_border
                    
        col_idx += 4

    # Render Data Table
    ws.cell(row=11, column=2, value="Agent Performance").font = Font(size=14, bold=True, color=t["header"])
    headers = ["Agent", "Calls", "Reached", "Closed"]
    data = [
        ["Alice", 1031, 128, 49],
        ["Bob", 661, 73, 28],
        ["Charlie", 610, 86, 67],
        ["Diana", 566, 163, 26]
    ]
    
    for i, h in enumerate(headers):
        cell = ws.cell(row=13, column=2+i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(2+i)].width = 15
        
    for r_idx, row_data in enumerate(data, start=14):
        for c_idx, val in enumerate(row_data):
            cell = ws.cell(row=r_idx, column=2+c_idx, value=val)
            cell.fill = white_fill
            if c_idx > 0:
                cell.number_format = "#,##0"

    # Add Data Bars conditional formatting to the "Closed" column
    rule = DataBarRule(start_type="min", end_type="max", color=t["accent"])
    ws.conditional_formatting.add("E14:E17", rule)

    # Render simple chart on dashboard
    chart = BarChart()
    chart.type = "col"
    chart.style = 11
    chart.title = "Calls vs Reached"
    chart.height = 8.5
    chart.width = 14
    
    cats = Reference(ws, min_col=2, min_row=14, max_row=17)
    data_ref = Reference(ws, min_col=3, min_row=13, max_col=4, max_row=17)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    
    ws.add_chart(chart, "G11")
