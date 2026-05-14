from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Income Statement", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Financial modeling standard fonts
    blue_font = Font(color="0000FF")    # Hardcoded inputs
    black_font = Font(color="000000")   # Formulas / Calculations
    bold_font = Font(bold=True)
    
    thin = Side(style="thin")
    double = Side(style="double")
    
    ws.column_dimensions['B'].width = 25
    for col in ['C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws.column_dimensions[col].width = 14

    # --- 1. Scenario Toggle Setup ---
    ws['H2'] = "Scenario Selector:"
    ws['H2'].font = bold_font
    ws['H2'].alignment = Alignment(horizontal="right")
    
    toggle_cell = ws['I2']
    toggle_cell.value = 1
    toggle_cell.fill = PatternFill(start_color="FFF2CC", fill_type="solid") # Warning yellow to indicate input
    toggle_cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    toggle_cell.alignment = Alignment(horizontal="center")
    
    # Restrict input to 1 (Base Case) or 2 (Upside Case)
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(toggle_cell)
    
    # --- 2. Income Statement Header ---
    ws['B3'] = title
    ws['B3'].font = bold_font
    for i in range(5):
        c = ws.cell(row=3, column=3+i, value=f"Year {i+1}")
        c.font = bold_font
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=thin)
        
    # --- 3. Income Statement Layout ---
    # Maps labels to formulas referencing the Live Case Assumptions block
    pl_layout = [
        ("Revenue", "=C14*C16"),                   # Orders * AOV
        ("COGS", "=C14*C17"),                      # Orders * CPO
        ("Gross Profit", "=C4-C5"),                # Revenue - COGS
        ("Operating Expenses", "=C18"),            # Fixed Opex
        ("Operating Profit", "=C6-C7"),            # Gross Profit - Opex
        ("Corporate Tax", "=IFERROR(IF(C8>0, C8*C19, 0), 0)"), # Tax on positive profit only
        ("Net Income", "=C8-C9")                   # Operating Profit - Tax
    ]
    
    for row_idx, (label, frmla) in enumerate(pl_layout, start=4):
        ws.cell(row=row_idx, column=2, value=label)
        for col_idx in range(3, 8):
            col_letter = get_column_letter(col_idx)
            # Adapt the formula reference to the current column dynamically
            adjusted_frmla = frmla.replace('C', col_letter)
            c = ws.cell(row=row_idx, column=col_idx, value=adjusted_frmla)
            c.number_format = '#,##0'
            c.font = black_font
            
        # Add formatting for summary lines
        if "Profit" in label or "Income" in label:
            ws.cell(row=row_idx, column=2).font = bold_font
            for col_idx in range(3, 8):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = bold_font
                if "Income" in label:
                    c.border = Border(top=thin, bottom=double)
                else:
                    c.border = Border(top=thin)

    # --- 4. Assumptions: Live Case (Driven by CHOOSE) ---
    ws.cell(row=13, column=2, value="Live Case Assumptions").font = bold_font
    assumptions = [
        "Number of Orders",
        "Order Growth Rate",
        "Average Order Value (AOV)",
        "Cost per Order (CPO)",
        "Operating Expenses",
        "Tax Rate"
    ]
    
    s1_start, s2_start = 22, 30
    
    for i, label in enumerate(assumptions):
        r = 14 + i
        ws.cell(row=r, column=2, value=label)
        for col_idx in range(3, 8):
            col_letter = get_column_letter(col_idx)
            # Build the dynamic link: =CHOOSE($I$2, BaseCaseCell, UpsideCaseCell)
            s1_ref = f"{col_letter}{s1_start + 1 + i}"
            s2_ref = f"{col_letter}{s2_start + 1 + i}"
            frmla = f"=CHOOSE($I$2, {s1_ref}, {s2_ref})"
            
            c = ws.cell(row=r, column=col_idx, value=frmla)
            c.font = black_font
            c.number_format = '0.0%' if "Rate" in label else '#,##0'

    # --- 5. Helper Function for Scenario Blocks ---
    def build_scenario_block(start_row, block_title, y1_orders, growths, aov, cpo, opex, tax):
        ws.cell(row=start_row, column=2, value=block_title).font = bold_font
        for idx, lbl in enumerate(assumptions):
            ws.cell(row=start_row + 1 + idx, column=2, value=lbl)
            
        # Y1 Orders (Hardcoded)
        ws.cell(row=start_row+1, column=3, value=y1_orders).font = blue_font
        ws.cell(row=start_row+1, column=3).number_format = '#,##0'
        
        # Y2-Y5 Growth Rates (Hardcoded)
        for j, g in enumerate(growths):
            c = ws.cell(row=start_row+2, column=4+j, value=g)
            c.font = blue_font
            c.number_format = '0.0%'
            
        # Y2-Y5 Orders (Formula: Prev Year * (1 + Growth))
        for col_idx in range(4, 8):
            prev_c = get_column_letter(col_idx - 1)
            curr_c = get_column_letter(col_idx)
            c = ws.cell(row=start_row+1, column=col_idx, value=f"={prev_c}{start_row+1}*(1+{curr_c}{start_row+2})")
            c.font = black_font
            c.number_format = '#,##0'
            
        # Flat Assumptions (Hardcoded)
        for col_idx in range(3, 8):
            ws.cell(row=start_row+3, column=col_idx, value=aov).font = blue_font
            ws.cell(row=start_row+3, column=col_idx).number_format = '#,##0'
            
            ws.cell(row=start_row+4, column=col_idx, value=cpo).font = blue_font
            ws.cell(row=start_row+4, column=col_idx).number_format = '#,##0'
            
            ws.cell(row=start_row+5, column=col_idx, value=opex).font = blue_font
            ws.cell(row=start_row+5, column=col_idx).number_format = '#,##0'
            
            t = ws.cell(row=start_row+6, column=col_idx, value=tax)
            t.font = blue_font
            t.number_format = '0.0%'

    # Build Base Case
    build_scenario_block(
        start_row=s1_start, 
        block_title="Base Case (Scenario 1)",
        y1_orders=3000, 
        growths=[1.0, 0.75, 0.50, 0.35], 
        aov=40, cpo=10, opex=50000, tax=0.20
    )
    
    # Build Upside Case
    build_scenario_block(
        start_row=s2_start, 
        block_title="Upside Case (Scenario 2)",
        y1_orders=4000, 
        growths=[1.2, 0.90, 0.60, 0.40], 
        aov=45, cpo=9, opex=60000, tax=0.25
    )
