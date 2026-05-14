from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # 1. Theme Setup
    theme_colors = {
        "corporate_blue": {"header": "1F4E78", "accent": "DDEBF7", "text": "000000", "highlight": "FFF2CC"},
        "emerald": {"header": "276A44", "accent": "E2EFDA", "text": "000000", "highlight": "FFF2CC"},
        "slate": {"header": "3B3838", "accent": "D9D9D9", "text": "000000", "highlight": "FFF2CC"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    header_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill("solid", fgColor=palette["header"])
    bold_font = Font(bold=True)
    highlight_fill = PatternFill("solid", fgColor=palette["highlight"])
    accent_fill = PatternFill("solid", fgColor=palette["accent"])
    
    border_bottom = Border(bottom=Side(style="thin", color="000000"))
    
    # 2. Setup Header & Scenario Dropdown
    ws["B2"] = title
    ws["B2"].font = bold_font
    
    cols = ["C", "D", "E", "F", "G"]
    for i, col in enumerate(cols, start=1):
        ws[f"{col}2"] = f"Year {i}"
        ws[f"{col}2"].font = header_font
        ws[f"{col}2"].fill = header_fill
        ws[f"{col}2"].alignment = Alignment(horizontal="center")
        
    ws["I2"] = "Scenario:"
    ws["I2"].font = bold_font
    ws["I2"].alignment = Alignment(horizontal="right")
    
    ws["J2"] = 1
    ws["J2"].fill = highlight_fill
    ws["J2"].font = bold_font
    ws["J2"].alignment = Alignment(horizontal="center")
    ws["J2"].border = Border(outline=True, top=Side(style="thin"), bottom=Side(style="thin"), left=Side(style="thin"), right=Side(style="thin"))
    
    # Data Validation for the dropdown toggle
    dv = DataValidation(type="list", formula1='"1,2,3"', allowBlank=False)
    ws.add_data_validation(dv)
    dv.add(ws["J2"])
    
    # 3. Financial Output Matrix
    output_labels = ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "Net Profit"]
    for i, label in enumerate(output_labels, start=3):
        ws[f"B{i}"] = label
        
    for col in cols:
        ws[f"{col}3"] = f"={col}10*{col}11"
        ws[f"{col}4"] = f"={col}10*{col}12"
        ws[f"{col}5"] = f"={col}3-{col}4"
        ws[f"{col}6"] = f"={col}13"
        ws[f"{col}7"] = f"={col}5-{col}6"
        
        for r in range(3, 8):
            ws[f"{col}{r}"].number_format = "#,##0"
        ws[f"{col}5"].border = border_bottom
        ws[f"{col}7"].border = Border(top=Side(style="thin"), bottom=Side(style="double"))
        
    ws["B5"].font = bold_font
    ws["B7"].font = bold_font
    
    # 4. Live Case Drivers (Dynamic CHOOSE formulas)
    driver_labels = ["Volume (Units)", "Price per Unit", "Cost per Unit", "Fixed Opex"]
    
    ws["B9"] = "Live Case Drivers"
    ws["B9"].font = header_font
    ws["B9"].fill = header_fill
    
    for i, label in enumerate(driver_labels, start=10):
        ws[f"B{i}"] = label
        ws[f"B{i}"].fill = accent_fill
        for col in cols:
            # CHOOSE swaps to Scenario 1 (+6 rows), Scenario 2 (+12 rows), or Scenario 3 (+18 rows)
            ws[f"{col}{i}"] = f"=CHOOSE($J$2, {col}{i+6}, {col}{i+12}, {col}{i+18})"
            ws[f"{col}{i}"].fill = accent_fill
            if "Price" in label or "Cost" in label:
                ws[f"{col}{i}"].number_format = "$#,##0.00"
            else:
                ws[f"{col}{i}"].number_format = "#,##0"
                
    # 5. Static Scenario Data Blocks
    scenarios = [
        ("Scenario 1: Base Case", 15, [
            [1000, 1100, 1210, 1331, 1464],      # Volume
            [50, 50, 50, 50, 50],                # Price
            [20, 20, 20, 20, 20],                # Cost
            [10000, 10500, 11000, 11500, 12000]  # Opex
        ]),
        ("Scenario 2: Upside", 21, [
            [1000, 1250, 1560, 1950, 2440],
            [55, 55, 55, 55, 55],
            [19, 18, 18, 17, 17],
            [10000, 10500, 11000, 11500, 12000]
        ]),
        ("Scenario 3: Downside", 27, [
            [1000, 900, 810, 730, 650],
            [45, 45, 45, 45, 45],
            [22, 23, 24, 25, 26],
            [10000, 10000, 10000, 10000, 10000]
        ])
    ]
    
    for block_title, start_row, data in scenarios:
        ws[f"B{start_row}"] = block_title
        ws[f"B{start_row}"].font = bold_font
        ws[f"B{start_row}"].border = border_bottom
        
        for i, label in enumerate(driver_labels, start=start_row+1):
            ws[f"B{i}"] = label
            row_data = data[i - start_row - 1]
            for col_idx, col in enumerate(cols):
                ws[f"{col}{i}"] = row_data[col_idx]
                if "Price" in label or "Cost" in label:
                    ws[f"{col}{i}"].number_format = "$#,##0.00"
                else:
                    ws[f"{col}{i}"].number_format = "#,##0"
                    
    # Resize columns for readability
    ws.column_dimensions["B"].width = 25
    for col in cols:
        ws.column_dimensions[col].width = 15
