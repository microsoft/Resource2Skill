def render_sheet(wb, sheet_name: str, *, dashboard_data: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Default data structure if none provided
    if dashboard_data is None:
        dashboard_data = [
            {
                "category": "Working Capital Efficiency",
                "kpis": [
                    {
                        "title": "DSO (Days Sales Outstanding)",
                        "value": 31,
                        "target": 45,
                        "prior": 41,
                        "format": "number",
                        "good_direction": "down"
                    },
                    {
                        "title": "DPO (Days Payables Outstanding)",
                        "value": 89,
                        "target": 90,
                        "prior": 90,
                        "format": "number",
                        "good_direction": "up"
                    },
                    {
                        "title": "Non-Current AR %",
                        "value": 0.12,
                        "target": 0.03,
                        "prior": 0.12,
                        "format": "percent",
                        "good_direction": "down"
                    }
                ]
            },
            {
                "category": "Sales KPIs",
                "kpis": [
                    {
                        "title": "CAC (Customer Acq. Cost)",
                        "value": 26319,
                        "target": 15000,
                        "prior": 17725,
                        "format": "currency",
                        "good_direction": "down"
                    },
                    {
                        "title": "Sales vs. Budget %",
                        "value": 1.62,
                        "target": 1.00,
                        "prior": 1.27,
                        "format": "percent",
                        "good_direction": "up"
                    },
                    {
                        "title": "Gross Margin",
                        "value": 0.20,
                        "target": 0.38,
                        "prior": 0.26,
                        "format": "percent",
                        "good_direction": "up"
                    }
                ]
            }
        ]

    # Theme definitions
    header_color = "4F81BD"
    text_light = "FFFFFF"
    text_dark = "000000"
    bg_light = "F2F2F2"
    success_color = "C6EFCE"  # Light green
    danger_color = "FFC7CE"   # Light red
    
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    bg_light_fill = PatternFill(start_color=bg_light, end_color=bg_light, fill_type="solid")
    success_fill = PatternFill(start_color=success_color, end_color=success_color, fill_type="solid")
    danger_fill = PatternFill(start_color=danger_color, end_color=danger_color, fill_type="solid")
    
    thin_border = Border(left=Side(style='thin', color="BFBFBF"), 
                         right=Side(style='thin', color="BFBFBF"), 
                         top=Side(style='thin', color="BFBFBF"), 
                         bottom=Side(style='thin', color="BFBFBF"))

    # Set column widths for a clean card layout
    for i in range(1, 20):
        ws.column_dimensions[get_column_letter(i)].width = 13
    
    ws.column_dimensions['A'].width = 3
    # Gaps between standard 3-KPI layout (cols F, K, P)
    ws.column_dimensions['F'].width = 3
    ws.column_dimensions['K'].width = 3
    ws.column_dimensions['P'].width = 3

    # Main Dashboard Title
    ws.merge_cells("B2:O2")
    main_title = ws.cell(row=2, column=2, value="Executive KPI Dashboard")
    main_title.font = Font(size=20, bold=True, color=text_dark)
    main_title.alignment = Alignment(horizontal="left", vertical="center")

    start_row = 4
    
    for section in dashboard_data:
        cat_title = section.get("category", "Category")
        kpis = section.get("kpis", [])
        num_kpis = len(kpis)
        
        # Draw Category Header
        end_col = 2 + (num_kpis * 5) - 2  # 4 columns per KPI card + 1 column gap
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=end_col)
        cat_cell = ws.cell(row=start_row, column=2, value=cat_title)
        cat_cell.fill = header_fill
        cat_cell.font = Font(color=text_light, bold=True, size=14)
        cat_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        start_row += 2 # Leave a blank row before rendering KPIs
        current_col = 2
        
        for kpi in kpis:
            val_row = start_row + 1
            sub_row = val_row + 1
            
            # --- 1. KPI Title Row ---
            ws.merge_cells(start_row=start_row, start_column=current_col, end_row=start_row, end_column=current_col+3)
            title_cell = ws.cell(row=start_row, column=current_col, value=kpi.get("title"))
            title_cell.font = Font(bold=True, size=11, color=text_dark)
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # --- 2. KPI Primary Value Row ---
            ws.merge_cells(start_row=val_row, start_column=current_col, end_row=val_row, end_column=current_col+3)
            val_cell = ws.cell(row=val_row, column=current_col, value=kpi.get("value"))
            val_cell.font = Font(bold=True, size=24, color=text_dark)
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            fmt = kpi.get("format", "number")
            if fmt == "percent":
                val_cell.number_format = "0%"
            elif fmt == "currency":
                val_cell.number_format = "$#,##0"
            else:
                val_cell.number_format = "#,##0"
                
            # --- 3. Sub-metrics Row ---
            ws.cell(row=sub_row, column=current_col, value="Vs. Target").font = Font(size=9, color=text_dark)
            t_cell = ws.cell(row=sub_row, column=current_col+1, value=kpi.get("target"))
            t_cell.font = Font(size=9, bold=True, color=text_dark)
            
            ws.cell(row=sub_row, column=current_col+2, value="Vs. Prior").font = Font(size=9, color=text_dark)
            p_cell = ws.cell(row=sub_row, column=current_col+3, value=kpi.get("prior"))
            p_cell.font = Font(size=9, bold=True, color=text_dark)
            
            if fmt == "percent":
                t_cell.number_format = "0%"
                p_cell.number_format = "0%"
            elif fmt == "currency":
                t_cell.number_format = "$#,##0"
                p_cell.number_format = "$#,##0"
            
            # Formatting standard borders and fills for the card block
            for row_idx in range(start_row, sub_row + 1):
                for col_idx in range(current_col, current_col + 4):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    # Apply background color to Title and Sub-metric rows to frame the card
                    if row_idx == start_row or row_idx == sub_row:
                        cell.fill = bg_light_fill
                    if row_idx == sub_row:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            # --- 4. Dynamic Conditional Formatting ---
            target_val = kpi.get("target", 0)
            good_dir = kpi.get("good_direction", "up")
            
            if good_dir == "down":
                rule_green = CellIsRule(operator='lessThanOrEqual', formula=[str(target_val)], stopIfTrue=True, fill=success_fill)
                rule_red = CellIsRule(operator='greaterThan', formula=[str(target_val)], stopIfTrue=True, fill=danger_fill)
            else:
                rule_green = CellIsRule(operator='greaterThanOrEqual', formula=[str(target_val)], stopIfTrue=True, fill=success_fill)
                rule_red = CellIsRule(operator='lessThan', formula=[str(target_val)], stopIfTrue=True, fill=danger_fill)
            
            # Apply condition to the big number
            ws.conditional_formatting.add(val_cell.coordinate, rule_green)
            ws.conditional_formatting.add(val_cell.coordinate, rule_red)
            
            current_col += 5 # Advance by 4 card columns + 1 gap column
            
        start_row += 4 # Move down for the next category block
