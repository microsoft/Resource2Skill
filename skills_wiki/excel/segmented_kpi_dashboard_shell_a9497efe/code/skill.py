from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

def render_sheet(wb, sheet_name: str, *, title: str = "KPI Dashboard", theme: str = "corporate_blue", dashboard_data: list = None, **kwargs) -> None:
    """
    Renders a structured, segmented KPI dashboard with conditional formatting.
    """
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    # Theme palette fallback
    palettes = {
        "corporate_blue": {"primary": "4472C4", "secondary": "D9E1F2", "text": "000000", "header_text": "FFFFFF"},
        "executive_grey": {"primary": "404040", "secondary": "D9D9D9", "text": "000000", "header_text": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # Colors for conditional formatting
    green_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    red_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")

    # Default realistic dataset if none provided
    if not dashboard_data:
        dashboard_data = [
            {
                "section": "Working Capital Efficiency",
                "kpis": [
                    {"name": "DSO (Days Sales Outstanding)", "value": 31, "target": 45, "prior": 41, "format": "0", "lower_is_better": True},
                    {"name": "DPO (Days Payables Outstanding)", "value": 89, "target": 90, "prior": 90, "format": "0", "lower_is_better": False},
                    {"name": "Non-Current AR %", "value": 0.12, "target": 0.03, "prior": 0.12, "format": "0%", "lower_is_better": True}
                ]
            },
            {
                "section": "Sales KPIs",
                "kpis": [
                    {"name": "CAC (Customer Acquisition Cost)", "value": 26319, "target": 15000, "prior": 17725, "format": "$#,##0", "lower_is_better": True},
                    {"name": "Sales vs. Budget %", "value": 1.62, "target": 1.00, "prior": 0.94, "format": "0%", "lower_is_better": False},
                    {"name": "Gross Margin", "value": 0.20, "target": 0.38, "prior": 0.26, "format": "0%", "lower_is_better": False}
                ]
            }
        ]

    ws.sheet_view.showGridLines = False

    # Standard styling
    title_font = Font(size=18, bold=True, color=palette["header_text"])
    section_font = Font(size=14, bold=True, color=palette["header_text"])
    card_header_font = Font(size=11, bold=True, color=palette["text"])
    value_font = Font(size=24, bold=True, color=palette["text"])
    footer_font = Font(size=9, color=palette["text"])
    
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color="BFBFBF"), 
                         right=Side(style='thin', color="BFBFBF"), 
                         top=Side(style='thin', color="BFBFBF"), 
                         bottom=Side(style='thin', color="BFBFBF"))

    primary_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    secondary_fill = PatternFill(start_color=palette["secondary"], end_color=palette["secondary"], fill_type="solid")

    # Render Dashboard Header & Controls
    ws['B2'] = "For the month of:"
    ws['B2'].font = Font(bold=True)
    ws['C2'] = "Jul-20"
    ws['C2'].fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    ws['C2'].border = thin_border
    
    # Title
    ws.merge_cells("G2:J2")
    ws['G2'] = title
    ws['G2'].font = title_font
    ws['G2'].fill = primary_fill
    ws['G2'].alignment = center_align

    current_row = 4

    for section in dashboard_data:
        # Render Section Banner
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=15)
        banner_cell = ws.cell(row=current_row, column=2)
        banner_cell.value = section["section"]
        banner_cell.font = section_font
        banner_cell.fill = primary_fill
        banner_cell.alignment = center_align
        
        current_row += 2
        col_start = 2

        for kpi in section["kpis"]:
            # Setup Card Columns (4 cols per card)
            ws.column_dimensions[ws.cell(row=1, column=col_start).column_letter].width = 12
            ws.column_dimensions[ws.cell(row=1, column=col_start+1).column_letter].width = 10
            ws.column_dimensions[ws.cell(row=1, column=col_start+2).column_letter].width = 14
            ws.column_dimensions[ws.cell(row=1, column=col_start+3).column_letter].width = 10
            
            # --- KPI Header ---
            ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_start+3)
            header_cell = ws.cell(row=current_row, column=col_start)
            header_cell.value = kpi["name"]
            header_cell.font = card_header_font
            header_cell.fill = secondary_fill
            header_cell.alignment = center_align

            # --- KPI Value ---
            val_row = current_row + 1
            ws.merge_cells(start_row=val_row, start_column=col_start, end_row=val_row, end_column=col_start+3)
            val_cell = ws.cell(row=val_row, column=col_start)
            val_cell.value = kpi["value"]
            val_cell.font = value_font
            val_cell.alignment = center_align
            val_cell.number_format = kpi["format"]

            # --- KPI Footers ---
            footer_row = current_row + 2
            lbl_tgt = ws.cell(row=footer_row, column=col_start, value="Vs. Target")
            val_tgt = ws.cell(row=footer_row, column=col_start+1, value=kpi["target"])
            lbl_pri = ws.cell(row=footer_row, column=col_start+2, value="Vs. Prior Month")
            val_pri = ws.cell(row=footer_row, column=col_start+3, value=kpi["prior"])

            for cell in [lbl_tgt, val_tgt, lbl_pri, val_pri]:
                cell.font = footer_font
                cell.alignment = center_align
                if cell in [val_tgt, val_pri]:
                    cell.number_format = kpi["format"]

            # Apply Border to whole card
            for r in range(current_row, footer_row + 1):
                for c in range(col_start, col_start + 4):
                    ws.cell(row=r, column=c).border = thin_border

            # --- Conditional Formatting ---
            # Based on Target cell within the same card footer
            target_coord = val_tgt.coordinate
            val_coord = val_cell.coordinate

            if kpi["lower_is_better"]:
                ws.conditional_formatting.add(val_coord, CellIsRule(operator='lessThanOrEqual', formula=[target_coord], fill=green_fill))
                ws.conditional_formatting.add(val_coord, CellIsRule(operator='greaterThan', formula=[target_coord], fill=red_fill))
            else:
                ws.conditional_formatting.add(val_coord, CellIsRule(operator='greaterThanOrEqual', formula=[target_coord], fill=green_fill))
                ws.conditional_formatting.add(val_coord, CellIsRule(operator='lessThan', formula=[target_coord], fill=red_fill))

            col_start += 5 # Move to next card position (4 cols + 1 spacing)

        current_row += 4 # Move down for the next section banner
