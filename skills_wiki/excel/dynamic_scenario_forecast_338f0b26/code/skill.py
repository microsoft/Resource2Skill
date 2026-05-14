import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Stylistic elements (would typically map to theme tokens)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    highlight_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    # Title
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    
    # Scenario Toggle Mechanism
    ws["G2"] = "Active Scenario:"
    ws["G2"].font = bold_font
    ws["G2"].alignment = Alignment(horizontal="right")
    
    ws["H2"] = 1
    ws["H2"].fill = highlight_fill
    ws["H2"].font = bold_font
    ws["H2"].alignment = Alignment(horizontal="center")
    ws["H2"].border = Border(
        left=Side(style='thin'), right=Side(style='thin'), 
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Add Data Validation (Drop-down list)
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    dv.error = "Please select a valid scenario number (1 or 2)."
    dv.errorTitle = "Invalid Scenario"
    ws.add_data_validation(dv)
    dv.add(ws["H2"])
    
    years = ["Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    
    # Reusable block renderer for assumptions
    def render_block(start_row, block_title, is_input=False, data_rows=None):
        ws.cell(row=start_row, column=1, value=block_title).font = bold_font
        for i, year in enumerate(years, start=3):
            c = ws.cell(row=start_row, column=i, value=year)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")
            
        labels = ["Number of Orders", "Average Order Value", "Cost per Order"]
        for idx, label in enumerate(labels):
            r = start_row + 1 + idx
            ws.cell(row=r, column=1, value=label)
            
            for col_idx in range(3, 8):
                cell = ws.cell(row=r, column=col_idx)
                if data_rows:
                    # Hardcoded assumption block
                    cell.value = data_rows[idx][col_idx-3]
                    if is_input:
                        cell.font = Font(color="0000FF") # Financial modeling convention for inputs
                else:
                    # Live block utilizing the CHOOSE formula
                    col_let = get_column_letter(col_idx)
                    scen1_ref = f"{col_let}{r + 6}"
                    scen2_ref = f"{col_let}{r + 12}"
                    cell.value = f"=CHOOSE($H$2, {scen1_ref}, {scen2_ref})"
                    cell.font = Font(color="000000") 
                
                # Number formatting
                if idx == 0:
                    cell.number_format = '#,##0'
                else:
                    cell.number_format = '"$"#,##0.00'

    # 1. Forecast Output Section
    ws["A4"] = "Income Statement"
    ws["A4"].font = bold_font
    for i, year in enumerate(years, start=3):
        c = ws.cell(row=4, column=i, value=year)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    line_items = ["Revenue", "COGS", "Gross Profit"]
    for idx, item in enumerate(line_items, start=5):
        ws.cell(row=idx, column=1, value=item)
        for col_idx in range(3, 8):
            col_let = get_column_letter(col_idx)
            cell = ws.cell(row=idx, column=col_idx)
            # Link to Live Case assumptions
            if item == "Revenue":
                cell.value = f"={col_let}11*{col_let}12"
            elif item == "COGS":
                cell.value = f"={col_let}11*{col_let}13"
            elif item == "Gross Profit":
                cell.value = f"={col_let}5-{col_let}6"
                cell.font = bold_font
            cell.number_format = '"$"#,##0'
            
    # 2. Render Live Case (driven by dropdown)
    render_block(10, "Live Case Assumptions")
    
    # 3. Render Scenario 1
    scen1_data = [
        [3000, 6000, 10500, 15750, 21263],
        [39.95, 39.95, 39.95, 39.95, 39.95],
        [6.50, 6.50, 6.50, 6.50, 6.50]
    ]
    render_block(16, "Upper Case (Scenario 1)", is_input=True, data_rows=scen1_data)
    
    # 4. Render Scenario 2
    scen2_data = [
        [2000, 4000, 7000, 10500, 14175],
        [34.95, 34.95, 34.95, 34.95, 34.95],
        [8.00, 8.00, 8.00, 8.00, 8.00]
    ]
    render_block(22, "Lower Case (Scenario 2)", is_input=True, data_rows=scen2_data)
    
    # Final cleanup & sizing
    ws.column_dimensions['A'].width = 25
    for col in ['C','D','E','F','G']:
        ws.column_dimensions[col].width = 15
