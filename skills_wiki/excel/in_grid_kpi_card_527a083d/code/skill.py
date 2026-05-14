def render(ws, anchor: str, *, title: str = "Asia Revenue", value_expr: str = "=369989", percent_expr: str = "=0.05", val_fmt: str = '"$"#,##0', pct_fmt: str = '0%', theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    from openpyxl.utils import get_column_letter

    # 1. Theme configuration fallback
    themes = {
        "corporate_blue": {"bg": "002060", "fg": "FFFFFF", "accent": "0070C0"},
        "modern_dark": {"bg": "262626", "fg": "FFFFFF", "accent": "00B050"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 2. Anchor resolution
    xy = coordinate_from_string(anchor)
    c_idx = column_index_from_string(xy[0])
    r_idx = xy[1]

    # 3. Layout allocation
    title_cell = ws.cell(row=r_idx, column=c_idx)
    val_cell = ws.cell(row=r_idx+1, column=c_idx)
    pct_cell = ws.cell(row=r_idx, column=c_idx+1)

    # 4. Values application
    title_cell.value = title
    val_cell.value = value_expr
    pct_cell.value = percent_expr

    # Merge right column for the secondary metric block
    ws.merge_cells(start_row=r_idx, start_column=c_idx+1, end_row=r_idx+1, end_column=c_idx+1)

    # 5. Styling Definitions
    bg_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    pct_fill = PatternFill(start_color=palette["accent"], end_color=palette["accent"], fill_type="solid")
    
    title_font = Font(color=palette["fg"], bold=False, size=11)
    val_font = Font(color=palette["fg"], bold=True, size=14)
    pct_font = Font(color=palette["fg"], bold=True, size=14)

    # Title Styling
    title_cell.fill = bg_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Main Value Styling
    val_cell.fill = bg_fill
    val_cell.font = val_font
    val_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    val_cell.number_format = val_fmt

    # Percentage / Secondary Metric Styling
    pct_cell.fill = pct_fill
    pct_cell.font = pct_font
    pct_cell.alignment = Alignment(horizontal="center", vertical="center")
    pct_cell.number_format = pct_fmt
    
    # Ensure merged child cell has underlying formatting applied
    ws.cell(row=r_idx+1, column=c_idx+1).fill = pct_fill

    # 6. Card Enclosure Border
    thick = Side(border_style="medium", color=palette["bg"])
    for r in range(r_idx, r_idx+2):
        for c in range(c_idx, c_idx+2):
            border_kwargs = {}
            if r == r_idx: border_kwargs["top"] = thick
            if r == r_idx+1: border_kwargs["bottom"] = thick
            if c == c_idx: border_kwargs["left"] = thick
            if c == c_idx+1: border_kwargs["right"] = thick
            ws.cell(row=r, column=c).border = Border(**border_kwargs)

    # 7. Sizing Adjustments
    ws.column_dimensions[get_column_letter(c_idx)].width = 18
    ws.column_dimensions[get_column_letter(c_idx+1)].width = 12
    ws.row_dimensions[r_idx].height = 20
    ws.row_dimensions[r_idx+1].height = 25
