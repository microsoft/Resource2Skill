from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule


def build_two_variable_sensitivity(path="two_variable_sensitivity.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sensitivity"

    ws["A1"] = "Base assumptions"
    ws["A2"] = "Unit cost"
    ws["B2"] = 14
    ws["A3"] = "Base price"
    ws["B3"] = 25
    ws["A4"] = "Base volume"
    ws["B4"] = 10000
    ws["A6"] = "Base profit"
    ws["B6"] = "=(B3-B2)*B4"

    prices = [20, 22, 24, 26, 28]
    volumes = [8000, 9000, 10000, 11000, 12000]
    ws["D2"] = "Profit sensitivity"
    for c, price in enumerate(prices, 5):
        ws.cell(2, c, price)
    for r, volume in enumerate(volumes, 3):
        ws.cell(r, 4, volume)
        for c in range(5, 10):
            price_cell = ws.cell(2, c).coordinate
            volume_cell = ws.cell(r, 4).coordinate
            ws.cell(r, c, f"=({price_cell}-$B$2)*{volume_cell}")

    ws.conditional_formatting.add(
        "E3:I7",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    wb.save(path)
    return path
