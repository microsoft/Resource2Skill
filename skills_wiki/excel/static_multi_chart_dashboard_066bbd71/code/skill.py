from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a clean, multi-chart presentation dashboard reading from a dedicated data sheet.
    """
    
    # Mocking standard theme extraction pattern
    theme_colors = {
        "corporate_blue": {"primary": "1F4E78", "text": "FFFFFF"},
        "midnight": {"primary": "203764", "text": "FFFFFF"},
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    # 1. Setup Data Sheet (Pre-aggregated metrics)
    ws_data = wb.active
    ws_data.title = "Data"
    
    # Segment Data (For Stacked Bar)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["Philippines", 54618, 7026, 22005],
        ["United Kingdom", 46530, 5220, 11497],
        ["Malaysia", 46587, 5538, 17536],
        ["United States", 36657, 6369, 22260],
    ]
    for row in market_data:
        ws_data.append(row)
        
    ws_data.append([]) # Empty row separator
    
    # Trend Data (For Line Charts)
    time_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    time_start_row = ws_data.max_row + 1
    for row in time_data:
        ws_data.append(row)
        
    # 2. Setup Dashboard Sheet
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False
    
    # 3. Apply Theme to Dashboard Header
    ws_dash.merge_cells("A1:N3")
    header_cell = ws_dash["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["text"])
    header_cell.fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Create Chart 1: Stacked Column (Market vs Product)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    
    # Data range: cols 2-4 (Products), rows 1-6 (Header + Markets)
    data1 = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=6)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=6)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    chart1.height = 10
    chart1.width = 15
    ws_dash.add_chart(chart1, "B5")
    
    # 5. Create Chart 2: Line Chart (Units Sold)
    chart2 = LineChart()
    chart2.title = "Units Sold Each Month"
    
    data2 = Reference(ws_data, min_col=2, min_row=time_start_row, max_row=time_start_row+4)
    cats2 = Reference(ws_data, min_col=1, min_row=time_start_row+1, max_row=time_start_row+4)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.height = 7
    chart2.width = 12
    chart2.legend = None # Remove legend for single series to maximize plot area
    ws_dash.add_chart(chart2, "H5")
    
    # 6. Create Chart 3: Line Chart (Profit)
    chart3 = LineChart()
    chart3.title = "Profit by Month"
    
    data3 = Reference(ws_data, min_col=3, min_row=time_start_row, max_row=time_start_row+4)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats2)
    chart3.height = 7
    chart3.width = 12
    chart3.legend = None
    ws_dash.add_chart(chart3, "H19")

    # Lock focus on Dashboard sheet and optionally hide data processing
    wb.active = ws_dash
    # ws_data.sheet_state = "hidden"
