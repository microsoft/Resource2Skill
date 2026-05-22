from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a complete dashboard workbook structure with raw data, a hidden 
    calculation sheet, and a gridless, stylized dashboard canvas.
    """
    
    # 1. Setup Data Sheet
    ws_data = wb.active
    ws_data.title = "Data"
    
    headers = ["Date", "Country", "Product", "Units Sold", "Profit"]
    rows = [
        ["2023-01-01", "India", "Chocolate Chip", 1500, 5000],
        ["2023-01-01", "USA", "Sugar", 2000, 7000],
        ["2023-02-01", "India", "Chocolate Chip", 1600, 5500],
        ["2023-02-01", "USA", "Sugar", 1800, 6500],
        ["2023-03-01", "India", "Chocolate Chip", 1700, 6000],
        ["2023-03-01", "USA", "Sugar", 2100, 7500],
    ]
    
    ws_data.append(headers)
    for row in rows:
        ws_data.append(row)
        
    tab = Table(displayName="SalesData", ref=f"A1:E{len(rows)+1}")
    style = TableStyleInfo(
        name="TableStyleMedium9", 
        showFirstColumn=False, 
        showLastColumn=False, 
        showRowStripes=True, 
        showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws_data.add_table(tab)
    
    # 2. Setup Calculation Sheet (for charts)
    ws_calc = wb.create_sheet("Aggregations")
    ws_calc.sheet_state = 'hidden'  # Hide calculation sheet to keep UX clean
    
    # Simulated Pivot 1: Profit by Country & Product
    calc_data_1 = [
        ["Country", "Chocolate Chip", "Sugar"],
        ["India", 16500, 0],
        ["USA", 0, 21000]
    ]
    for r, row in enumerate(calc_data_1, 1):
        for c, val in enumerate(row, 1):
            ws_calc.cell(row=r, column=c, value=val)
            
    # Simulated Pivot 2: Monthly Units
    calc_data_2 = [
        ["Month", "Units Sold"],
        ["Jan", 3500],
        ["Feb", 3400],
        ["Mar", 3800]
    ]
    for r, row in enumerate(calc_data_2, 6):
        for c, val in enumerate(row, 1):
            ws_calc.cell(row=r, column=c, value=val)
            
    # Simulated Pivot 3: Monthly Profit
    calc_data_3 = [
        ["Month", "Profit"],
        ["Jan", 12000],
        ["Feb", 12000],
        ["Mar", 13500]
    ]
    for r, row in enumerate(calc_data_3, 12):
        for c, val in enumerate(row, 1):
            ws_calc.cell(row=r, column=c, value=val)
            
    # 3. Setup Dashboard Sheet
    ws_dash = wb.create_sheet("Dashboard")
    
    # The crucial "Canvas" settings: hide gridlines and row/col headers
    ws_dash.sheet_view.showGridLines = False
    ws_dash.sheet_view.showRowColHeaders = False
    
    # Title Banner
    ws_dash.merge_cells("A1:R3")
    title_cell = ws_dash["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Chart 1: Stacked Bar (Profit by Country & Product)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    chart1.style = 11  # Built-in Excel style
    chart1.height = 12
    chart1.width = 16
    
    data1 = Reference(ws_calc, min_col=2, min_row=1, max_col=3, max_row=3)
    cats1 = Reference(ws_calc, min_col=1, min_row=2, max_row=3)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    
    ws_dash.add_chart(chart1, "B5")
    
    # Chart 2: Line (Units Sold)
    chart2 = LineChart()
    chart2.title = "Units Sold Each Month"
    chart2.style = 12
    chart2.height = 8
    chart2.width = 12
    
    data2 = Reference(ws_calc, min_col=2, min_row=6, max_row=9)
    cats2 = Reference(ws_calc, min_col=1, min_row=7, max_row=9)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.legend = None  # Cleaner UI for single-series line charts
    
    ws_dash.add_chart(chart2, "K5")
    
    # Chart 3: Line (Profit)
    chart3 = LineChart()
    chart3.title = "Profit by Month"
    chart3.style = 12
    chart3.height = 8
    chart3.width = 12
    
    data3 = Reference(ws_calc, min_col=2, min_row=12, max_row=15)
    cats3 = Reference(ws_calc, min_col=1, min_row=13, max_row=15)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.legend = None
    
    ws_dash.add_chart(chart3, "K17")
    
    # Make Dashboard the active sheet upon opening
    wb.active = ws_dash
