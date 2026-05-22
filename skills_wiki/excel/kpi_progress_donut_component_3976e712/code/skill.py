from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.series import DataPoint

def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    # 1. Setup anchor coordinates
    col_letter, row = coordinate_from_string(anchor)
    col = column_index_from_string(col_letter)

    # 2. Write KPI data directly into the sheet 
    # In a dashboard, this data usually sits on an "Inputs" tab or behind the chart
    ws.cell(row=row, column=col, value="Status")
    ws.cell(row=row, column=col+1, value="Percentage")

    ws.cell(row=row+1, column=col, value="Achieved")
    ws.cell(row=row+1, column=col+1, value=0.85).number_format = "0%"

    ws.cell(row=row+2, column=col, value="Remaining")
    ws.cell(row=row+2, column=col+1, value=0.15).number_format = "0%"

    # 3. Initialize & Configure the Doughnut Chart
    chart = DoughnutChart()
    chart.title = "Customer Satisfaction"
    chart.legend = None  # Hide legend to save space
    chart.holeSize = 65  # Specific setting from the tutorial to make it a thick ring

    # 4. Map the Data
    data = Reference(ws, min_col=col+1, min_row=row+1, max_row=row+2)
    cats = Reference(ws, min_col=col, min_row=row+1, max_row=row+2)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    # 5. Style the Slices (Achieved vs Remaining)
    # Manually addressing DataPoints overrides Excel's default multi-color behavior
    if chart.series:
        series = chart.series[0]

        # Achieved slice (Primary Theme Color)
        pt_achieved = DataPoint(idx=0)
        pt_achieved.graphicalProperties.solidFill = "0055A4" # Fallback corporate blue
        series.dPt.append(pt_achieved)

        # Remaining slice (Neutral Track Color)
        pt_remaining = DataPoint(idx=1)
        pt_remaining.graphicalProperties.solidFill = "D9D9D9" # Light Grey
        series.dPt.append(pt_remaining)

    # 6. Clean up the chart area for seamless dashboard integration
    # Removes the background fill and outline of the chart container
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line.noFill = True

    # 7. Position chart
    ws.add_chart(chart, anchor)
