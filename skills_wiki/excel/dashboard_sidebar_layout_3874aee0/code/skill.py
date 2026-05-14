from openpyxl.styles import PatternFill, Font, Border, Side, Alignment

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Hide gridlines to create a clean dashboard canvas
    ws.sheet_view.showGridLines = False
    
    # Theme fallbacks (standard grays for UI structure)
    bg_sidebar = "F3F4F6"
    border_color = "D1D5DB"
    text_main = "111827"
    text_muted = "6B7280"
    
    sidebar_fill = PatternFill("solid", fgColor=bg_sidebar)
    divider_border = Border(right=Side(style="medium", color=border_color))
    
    # 1. Paint the Sidebar (Cols A-C) and Divider (Col D)
    for row in range(1, 100):
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = sidebar_fill
        
        div_cell = ws.cell(row=row, column=4)
        div_cell.fill = sidebar_fill
        div_cell.border = divider_border
        
    # 2. Configure Column Widths
    for col_letter in ['A', 'B', 'C']:
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions['D'].width = 2
    
    # 3. Add Sidebar Header
    ws.merge_cells("A2:C2")
    sidebar_header = ws["A2"]
    sidebar_header.value = "Filters & Controls"
    sidebar_header.font = Font(name="Calibri", size=14, bold=True, color=text_main)
    sidebar_header.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Add Placeholder Slots for Slicers
    ws.merge_cells("A4:C10")
    hint1 = ws["A4"]
    hint1.value = "[ Insert Date Slicer ]"
    hint1.font = Font(name="Calibri", size=11, color=text_muted, italic=True)
    hint1.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A12:C18")
    hint2 = ws["A12"]
    hint2.value = "[ Insert Category Slicer ]"
    hint2.font = Font(name="Calibri", size=11, color=text_muted, italic=True)
    hint2.alignment = Alignment(horizontal="center", vertical="center")
    
    # 5. Main Content Area Setup
    ws["F2"] = title
    ws["F2"].font = Font(name="Calibri", size=20, bold=True, color=text_main)
    ws["F3"] = "Main Reporting Canvas"
    ws["F3"].font = Font(name="Calibri", size=12, color=text_muted)
    
    # 6. Freeze Panes
    # Freezing at E4 locks Rows 1-3 and Columns A-D in place.
    # This ensures the sidebar and main header remain visible when scrolling right or down.
    ws.freeze_panes = "E4"
