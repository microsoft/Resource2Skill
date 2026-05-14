from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # Set up theme palette
    themes = {
        "corporate_blue": {"bg": "2F5597", "fg": "FFFFFF"},
        "emerald": {"bg": "005A36", "fg": "FFFFFF"},
        "slate": {"bg": "475569", "fg": "FFFFFF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 1. Setup sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_data = wb.create_sheet("Data")
    
    # Clean up dashboard canvas
    ws_dash.sheet_view.showGridLines = False
    
    # 2. Add Title Header
    ws_dash.merge_cells("A1:P3")
    title_cell = ws_dash["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["fg"])
    title_cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Insert backing data into the Data sheet
    # Dataset 1: Stacked Bar Data
    data_bar = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Sugar"],
        ["India", 62000, 23000, 21000, 25000],
        ["Philippines", 54000, 24000, 22000, 8000],
        ["United Kingdom", 46000, 26000, 11000, 14000],
        ["United States", 36000, 32000, 9000, 19000]
    ]
    for row in data_bar:
        ws_data.append(row)
        
    # Dataset 2: Line Chart 1 Data
    ws_data.append([]) # Blank row spacer
    data_line1 = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for row in data_line1:
        ws_data.append(row)
        
    # Dataset 3: Line Chart 2 Data
    ws_data.append([]) # Blank row spacer
    data_line2 = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for row in data_line2:
        ws_data.append(row)
        
    # 4. Render Charts on Dashboard
    
    # Main Stacked Bar Chart
    bc = BarChart()
    bc.type = "col"
    bc.grouping = "stacked"
    bc.overlap = 100
    bc.title = "Profit by Market & Cookie Type"
    
    bc_data = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    bc_cats = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    bc.add_data(bc_data, titles_from_data=True)
    bc.set_categories(bc_cats)
    
    bc.width = 16
    bc.height = 10
    ws_dash.add_chart(bc, "B5")
    
    # First Line Chart
    lc1 = LineChart()
    lc1.title = "Units sold each month"
    
    lc1_data = Reference(ws_data, min_col=2, min_row=7, max_col=2, max_row=11)
    lc1_cats = Reference(ws_data, min_col=1, min_row=8, max_row=11)
    lc1.add_data(lc1_data, titles_from_data=True)
    lc1.set_categories(lc1_cats)
    
    lc1.width = 14
    lc1.height = 7
    lc1.legend = None # Remove legend to maximize chart area
    ws_dash.add_chart(lc1, "K5")
    
    # Second Line Chart
    lc2 = LineChart()
    lc2.title = "Profit by month"
    
    lc2_data = Reference(ws_data, min_col=2, min_row=13, max_col=2, max_row=17)
    lc2_cats = Reference(ws_data, min_col=1, min_row=14, max_row=17)
    lc2.add_data(lc2_data, titles_from_data=True)
    lc2.set_categories(lc2_cats)
    
    lc2.width = 14
    lc2.height = 7
    lc2.legend = None
    ws_dash.add_chart(lc2, "K13")

    # Hide Data sheet so the end user only sees the dashboard
    ws_data.sheet_state = "hidden"
