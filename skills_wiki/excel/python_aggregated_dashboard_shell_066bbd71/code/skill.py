import datetime
from collections import defaultdict
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_workbook(wb, *, title: str = "Performance Dashboard", raw_data: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a complete dashboard workbook by taking raw transaction data,
    aggregating it into a hidden 'ChartData' sheet using Python, and 
    generating a clean 'Dashboard' sheet with layout-managed charts.
    """
    if not raw_data:
        raw_data = [
            {'Date': datetime.date(2019, 9, 1), 'Country': 'India', 'Product': 'Chocolate Chip', 'Units Sold': 1725, 'Profit': 5175},
            {'Date': datetime.date(2019, 9, 1), 'Country': 'India', 'Product': 'Fortune Cookie', 'Units Sold': 345, 'Profit': 1220},
            {'Date': datetime.date(2019, 10, 1), 'Country': 'United States', 'Product': 'Chocolate Chip', 'Units Sold': 2152, 'Profit': 6456},
            {'Date': datetime.date(2019, 11, 1), 'Country': 'United Kingdom', 'Product': 'Sugar', 'Units Sold': 1800, 'Profit': 5400},
            {'Date': datetime.date(2019, 12, 1), 'Country': 'Malaysia', 'Product': 'Oatmeal Raisin', 'Units Sold': 2200, 'Profit': 6600},
            {'Date': datetime.date(2019, 12, 1), 'Country': 'Philippines', 'Product': 'Snickerdoodle', 'Units Sold': 1500, 'Profit': 4500},
        ]

    # Theme palette resolver
    theme_colors = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF"},
        "emerald_green": {"header_bg": "2E7D32", "header_fg": "FFFFFF"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # 1. Aggregate Data in Python (Replacing PivotTables)
    profit_by_cp = defaultdict(lambda: defaultdict(float))
    units_by_m = defaultdict(float)
    profit_by_m = defaultdict(float)
    products = set()
    
    for row in raw_data:
        d = row['Date']
        month_str = d.strftime('%b %Y') 
        c = row['Country']
        p = row['Product']
        
        products.add(p)
        profit_by_cp[c][p] += row['Profit']
        units_by_m[month_str] += row['Units Sold']
        profit_by_m[month_str] += row['Profit']
        
    products = sorted(list(products))
    countries = sorted(list(profit_by_cp.keys()))
    unique_dates = sorted(list(set(r['Date'].replace(day=1) for r in raw_data)))
    months = [d.strftime('%b %Y') for d in unique_dates]
    
    # 2. Setup Sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_data = wb.create_sheet(title="ChartData")
    ws_data.sheet_state = 'hidden'
    
    # 3. Write Data to ChartData Sheet
    # a. Profit by Country & Product
    ws_data.cell(row=1, column=1, value="Country")
    for c_idx, p in enumerate(products, start=2):
        ws_data.cell(row=1, column=c_idx, value=p)
        
    for r_idx, c in enumerate(countries, start=2):
        ws_data.cell(row=r_idx, column=1, value=c)
        for c_idx, p in enumerate(products, start=2):
            ws_data.cell(row=r_idx, column=c_idx, value=profit_by_cp[c][p])
            
    # b. Units by Month
    row_offset_units = len(countries) + 5
    ws_data.cell(row=row_offset_units, column=1, value="Month")
    ws_data.cell(row=row_offset_units, column=2, value="Units Sold")
    for r_idx, m in enumerate(months, start=row_offset_units + 1):
        ws_data.cell(row=r_idx, column=1, value=m)
        ws_data.cell(row=r_idx, column=2, value=units_by_m[m])
        
    # c. Profit by Month
    row_offset_profit = row_offset_units
    col_offset_profit = 4
    ws_data.cell(row=row_offset_profit, column=col_offset_profit, value="Month")
    ws_data.cell(row=row_offset_profit, column=col_offset_profit+1, value="Profit")
    for r_idx, m in enumerate(months, start=row_offset_profit + 1):
        ws_data.cell(row=r_idx, column=col_offset_profit, value=m)
        ws_data.cell(row=r_idx, column=col_offset_profit+1, value=profit_by_m[m])

    # 4. Style Dashboard Sheet
    ws_dash.sheet_view.showGridLines = False
    ws_dash.merge_cells("B2:O3")
    title_cell = ws_dash["B2"]
    title_cell.value = title
    title_cell.font = Font(size=20, bold=True, color=palette["header_fg"])
    title_cell.fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # 5. Create Charts
    # Chart 1: Stacked Bar (Profit by Country & Product)
    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Product Type"
    c1.height = 12
    c1.width = 16
    
    data_ref1 = Reference(ws_data, min_col=2, min_row=1, max_col=1+len(products), max_row=1+len(countries))
    cats_ref1 = Reference(ws_data, min_col=1, min_row=2, max_row=1+len(countries))
    c1.add_data(data_ref1, titles_from_data=True)
    c1.set_categories(cats_ref1)
    ws_dash.add_chart(c1, "B5")
    
    # Chart 2: Line (Units by Month)
    c2 = LineChart()
    c2.title = "Units Sold Each Month"
    c2.height = 5.8
    c2.width = 11
    c2.legend = None
    data_ref2 = Reference(ws_data, min_col=2, min_row=row_offset_units, max_row=row_offset_units+len(months))
    cats_ref2 = Reference(ws_data, min_col=1, min_row=row_offset_units+1, max_row=row_offset_units+len(months))
    c2.add_data(data_ref2, titles_from_data=True)
    c2.set_categories(cats_ref2)
    ws_dash.add_chart(c2, "K5")
    
    # Chart 3: Line (Profit by Month)
    c3 = LineChart()
    c3.title = "Profit By Month"
    c3.height = 5.8
    c3.width = 11
    c3.legend = None
    data_ref3 = Reference(ws_data, min_col=col_offset_profit+1, min_row=row_offset_profit, max_row=row_offset_profit+len(months))
    cats_ref3 = Reference(ws_data, min_col=col_offset_profit, min_row=row_offset_profit+1, max_row=row_offset_profit+len(months))
    c3.add_data(data_ref3, titles_from_data=True)
    c3.set_categories(cats_ref3)
    ws_dash.add_chart(c3, "K16")
