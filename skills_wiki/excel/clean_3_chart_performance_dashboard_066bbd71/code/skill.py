from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Creates a clean, 3-chart dashboard layout using a hidden data backend.
    Mimics the layout of an interactive Pivot/Slicer dashboard.
    """
    # 1. Create Dashboard Presentation Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
    
    ws.sheet_view.showGridLines = False
    
    # 2. Create Hidden Data Sheet
    data_ws_name = f"{sheet_name}_Data"
    if data_ws_name in wb.sheetnames:
        data_ws = wb[data_ws_name]
    else:
        data_ws = wb.create_sheet(data_ws_name)
    data_ws.sheet_state = 'hidden'

    # 3. Populate Mock Aggregated Data for Charts
    # Chart 1 Data: Stacked Bar (Markets x Products)
    markets = ["India", "Philippines", "United Kingdom", "Malaysia", "United States"]
    products = ["Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle", "Sugar"]
    
    data_ws.append(["Market"] + products)
    
    import random
    random.seed(42) # For reproducible visual
    for m in markets:
        data_ws.append([m] + [random.randint(10000, 60000) for _ in products])
        
    # Chart 2 & 3 Data: Monthly Trends
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    data_ws.append([]) # Spacer row
    start_row_trends = len(markets) + 3
    data_ws.append(["Month", "Units Sold", "Profit"])
    
    for i, m in enumerate(months):
        units = random.randint(20000, 100000)
        profit = units * random.uniform(1.5, 3.0)
        data_ws.append([m, units, profit])

    # 4. Dashboard Title and UI Elements (Slicer Placeholders)
    ws["B2"] = title
    ws["B2"].font = Font(size=24, bold=True, color="333333")
    
    ws["B6"] = "Filters"
    ws["B6"].font = Font(bold=True, size=12)
    
    slicer_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    slicer_border = Border(
        top=Side(style="thin", color="CCCCCC"),
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )
    
    def draw_fake_slicer(start_row, height, label):
        ws.cell(row=start_row, column=2, value=label).font = Font(bold=True, size=10)
        for r in range(start_row + 1, start_row + height):
            for c in range(2, 4):
                cell = ws.cell(row=r, column=c)
                cell.fill = slicer_fill
                cell.border = slicer_border

    draw_fake_slicer(7, 4, "Date Timeline")
    draw_fake_slicer(12, 6, "Market")
    draw_fake_slicer(19, 6, "Product")

    # 5. Create Chart 1 (Stacked Column)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Product"
    chart1.height = 12.0
    chart1.width = 16.0
    chart1.style = 2  # Standard clean Excel style
    
    data1 = Reference(data_ws, min_col=2, min_row=1, max_col=len(products)+1, max_row=len(markets)+1)
    cats1 = Reference(data_ws, min_col=1, min_row=2, max_row=len(markets)+1)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    ws.add_chart(chart1, "E6")

    # 6. Create Chart 2 (Line - Units Sold)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.height = 5.8
    chart2.width = 13.0
    chart2.style = 2
    
    data2 = Reference(data_ws, min_col=2, min_row=start_row_trends, max_col=2, max_row=start_row_trends+12)
    cats2 = Reference(data_ws, min_col=1, min_row=start_row_trends+1, max_row=start_row_trends+12)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.legend = None  # Suppress legend for single-series line charts
    ws.add_chart(chart2, "O6")

    # 7. Create Chart 3 (Line - Profit)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.height = 5.8
    chart3.width = 13.0
    chart3.style = 2
    
    data3 = Reference(data_ws, min_col=3, min_row=start_row_trends, max_col=3, max_row=start_row_trends+12)
    cats3 = Reference(data_ws, min_col=1, min_row=start_row_trends+1, max_row=start_row_trends+12)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.legend = None
    ws.add_chart(chart3, "O19")
