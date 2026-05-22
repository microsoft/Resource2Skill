from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", main_value: str = "$369,989", sub_value: str = "5% Market Share", theme: str = "corporate_blue", **kwargs) -> None:
    # Theme palette fallback
    palettes = {
        "corporate_blue": {"primary": "002060", "text_light": "FFFFFF", "accent": "0070C0"},
        "modern_dark": {"primary": "262626", "text_light": "FFFFFF", "accent": "00B050"},
        "vibrant_navy": {"primary": "0A2540", "text_light": "FFFFFF", "accent": "635BFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    col = ws[anchor].column
    row = ws[anchor].row

    # Define card dimensions
    card_width = 3  # Spans 3 columns
    card_height = 3 # Spans 3 rows

    # Setup hierarchical typography and content
    content = [
        (title, Font(size=12, color=palette["text_light"], bold=False), Alignment(horizontal="center", vertical="center")),
        (main_value, Font(size=22, color=palette["text_light"], bold=True), Alignment(horizontal="center", vertical="center")),
        (sub_value, Font(size=12, color=palette["text_light"], italic=True), Alignment(horizontal="center", vertical="center"))
    ]

    # Style definitions
    fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    border_color = Side(style="medium", color=palette["primary"])

    # Apply styles, merges, and borders
    for r_offset in range(card_height):
        current_row = row + r_offset
        start_col_letter = get_column_letter(col)
        end_col_letter = get_column_letter(col + card_width - 1)

        # Merge the horizontal slice for the current text element
        ws.merge_cells(f"{start_col_letter}{current_row}:{end_col_letter}{current_row}")
        target_cell = ws.cell(row=current_row, column=col)

        # Inject text and typography
        text, font, alignment = content[r_offset]
        target_cell.value = text
        target_cell.font = font
        target_cell.alignment = alignment

        # Apply block formatting to ensure the "card" looks like a single unified shape
        for c_offset in range(card_width):
            c = ws.cell(row=current_row, column=col + c_offset)
            c.fill = fill
            
            # Determine outer borders to create a crisp bounding box
            top = border_color if r_offset == 0 else None
            bottom = border_color if r_offset == card_height - 1 else None
            left = border_color if c_offset == 0 else None
            right = border_color if c_offset == card_width - 1 else None
            c.border = Border(top=top, bottom=bottom, left=left, right=right)

    # Adjust row heights to simulate the padding of a shape
    ws.row_dimensions[row].height = 20     # Title spacing
    ws.row_dimensions[row + 1].height = 35 # Main value spacing
    ws.row_dimensions[row + 2].height = 20 # Sub value spacing
