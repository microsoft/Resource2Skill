from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Creates a split-panel dashboard layout featuring a dark KPI sidebar and a 
    light main canvas organized into bordered 'Cards' for charts.
    """
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Standardized theme fallback (can be dynamically overridden by kwargs)
    colors = kwargs.get("theme_palette", {
        "sidebar_bg": "1E293B",      # Slate 800
        "sidebar_fg": "F8FAFC",      # Slate 50
        "sidebar_accent": "38BDF8",  # Sky 400
        "canvas_bg": "F1F5F9",       # Slate 100
        "card_bg": "FFFFFF",         # White
        "card_border": "CBD5E1",     # Slate 300
        "text_main": "334155"        # Slate 700
    })

    # 1. Base Grid Setup: Configure widths for sidebar and main canvas
    ws.column_dimensions['A'].width = 2   # Left margin
    ws.column_dimensions['B'].width = 28  # Sidebar width
    ws.column_dimensions['C'].width = 2   # Gutter separating sidebar and canvas
    
    # Configure canvas grid (Columns D through N)
    for col in range(4, 15): 
        ws.column_dimensions[get_column_letter(col)].width = 12

    # 2. Paint the Backdrop
    sidebar_fill = PatternFill(start_color=colors["sidebar_bg"], end_color=colors["sidebar_bg"], fill_type="solid")
    canvas_fill = PatternFill(start_color=colors["canvas_bg"], end_color=colors["canvas_bg"], fill_type="solid")

    for row in range(1, 41):
        ws.row_dimensions[row].height = 18
        ws.cell(row=row, column=2).fill = sidebar_fill
        for col in range(3, 15):
            ws.cell(row=row, column=col).fill = canvas_fill

    # 3. Sidebar Header
    title_cell = ws['B2']
    title_cell.value = title.upper()
    title_cell.font = Font(name="Segoe UI", size=16, bold=True, color=colors["sidebar_fg"])
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # 4. Populate Sidebar KPIs
    kpis = kwargs.get("kpis", [
        {"label": "TOTAL ORDERS", "value": "2,400"},
        {"label": "UNITS SOLD", "value": "11,997"},
        {"label": "REVENUE", "value": "$649.0K"},
        {"label": "AVG RATING", "value": "4.0 / 5.0"}
    ])

    current_row = 5
    for kpi in kpis:
        lbl_cell = ws.cell(row=current_row, column=2)
        lbl_cell.value = kpi["label"]
        lbl_cell.font = Font(name="Segoe UI", size=9, bold=True, color=colors["sidebar_accent"])
        lbl_cell.alignment = Alignment(indent=1)

        val_cell = ws.cell(row=current_row+1, column=2)
        val_cell.value = kpi["value"]
        val_cell.font = Font(name="Segoe UI", size=22, bold=True, color=colors["sidebar_fg"])
        val_cell.alignment = Alignment(indent=1)

        current_row += 4

    # 5. Helper function to generate 'Cards' on the canvas
    def create_card(start_col: int, start_row: int, end_col: int, end_row: int, card_title: str):
        card_fill = PatternFill(start_color=colors["card_bg"], end_color=colors["card_bg"], fill_type="solid")
        
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Resolve precise outer borders
                top = Side(style='thin', color=colors["card_border"]) if r == start_row else None
                bottom = Side(style='thin', color=colors["card_border"]) if r == end_row else None
                left = Side(style='thin', color=colors["card_border"]) if c == start_col else None
                right = Side(style='thin', color=colors["card_border"]) if c == end_col else None
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)

        # Set title in the merged top row of the card
        t_cell = ws.cell(row=start_row, column=start_col)
        t_cell.value = card_title
        t_cell.font = Font(name="Segoe UI", size=11, bold=True, color=colors["text_main"])
        t_cell.alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)

    # 6. Instantiate Dashboard Layout Cards (Staging areas for future content)
    create_card(4, 3, 8, 14, "Last 13 Weeks Trend - Qty & Amount")
    create_card(10, 3, 14, 14, "Customer Acquisition Modes")
    create_card(4, 16, 9, 29, "Popular Products Breakdown")
    create_card(11, 16, 14, 29, "Geographical Analysis")
