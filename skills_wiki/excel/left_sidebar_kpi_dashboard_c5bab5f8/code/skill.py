from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Brand Performance", theme: str = "corporate_blue", kpi_data: list = None, **kwargs) -> None:
    """
    Renders the Left Sidebar Dashboard layout scaffolding.
    """
    # 1. Setup Theme Palette Fallbacks
    palettes = {
        "corporate_blue": {
            "sidebar_bg": "1F4E78", "sidebar_text": "FFFFFF", "sidebar_muted": "9BC2E6", 
            "main_bg": "F8F9FA", "main_text": "2C3E50"
        },
        "botanical_green": {
            "sidebar_bg": "2A4034", "sidebar_text": "FFFFFF", "sidebar_muted": "A3B8AD", 
            "main_bg": "E9EFEA", "main_text": "1A261F"
        }
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    # 2. Default realistic KPI data
    kpis = kpi_data or [
        ("Total Orders", "2,400"),
        ("Quantity Sold", "11,997"),
        ("Gross Revenue", "$649.0K"),
        ("Avg. Rating", "4.0"),
        ("Days to Deliver", "2.3")
    ]

    # 3. Create Sheet and clear gridlines
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # 4. Set Layout Dimensions
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 2
    
    for i in range(4, 16):
        ws.column_dimensions[get_column_letter(i)].width = 12

    # 5. Apply Block Fills (Sidebar vs Main Area)
    sidebar_fill = PatternFill(fgColor=palette["sidebar_bg"], fill_type="solid")
    main_fill = PatternFill(fgColor=palette["main_bg"], fill_type="solid")

    for row in range(1, 51):
        # Sidebar fill
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = sidebar_fill
        # Main area fill
        for col in range(4, 16):
            ws.cell(row=row, column=col).fill = main_fill

    # 6. Build Sidebar Header
    title_cell = ws['B2']
    title_cell.value = title.upper()
    title_cell.font = Font(color=palette["sidebar_text"], size=20, bold=True)
    title_cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 40

    # 7. Build Vertical KPI Strip
    current_row = 5
    for label, value in kpis:
        # Label (small, muted)
        lbl_cell = ws.cell(row=current_row, column=2, value=label)
        lbl_cell.font = Font(color=palette["sidebar_muted"], size=10, bold=True)
        
        # Value (large, bright)
        val_cell = ws.cell(row=current_row+1, column=2, value=value)
        val_cell.font = Font(color=palette["sidebar_text"], size=18, bold=True)
        
        current_row += 3

    # 8. Slicer / Filter Placemarker
    current_row += 1
    filter_hdr = ws.cell(row=current_row, column=2, value="FILTERS / SLICERS")
    filter_hdr.font = Font(color=palette["sidebar_muted"], size=10, bold=True)
    
    filter_box = ws.cell(row=current_row+2, column=2, value="[ Place Slicers Here ]")
    filter_box.font = Font(color=palette["sidebar_bg"], size=10, italic=True)
    filter_box.fill = PatternFill(fgColor=palette["sidebar_muted"], fill_type="solid")
    filter_box.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[current_row+2].height = 60

    # 9. Main Area Placeholder
    main_hdr = ws['E2']
    main_hdr.value = "Dashboard Visuals Canvas"
    main_hdr.font = Font(color=palette["main_text"], size=16, bold=True)
