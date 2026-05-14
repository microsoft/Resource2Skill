from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.worksheet import Worksheet

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a polished dashboard shell with a title banner, control panel placeholder, 
    and a 3-chart layout (1 main stacked column, 2 secondary line charts).
    """
    # 1. Setup Theme Colors
    themes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF", "accent": "D9E1F2", "text": "000000"},
        "dark_mode": {"bg": "202020", "fg": "FFFFFF", "accent": "333333", "text": "E0E0E0"},
        "emerald": {"bg": "005A36", "fg": "FFFFFF", "accent": "C6E0B4", "text": "000000"}
    }
    colors = themes.get(theme, themes["corporate_blue"])

    # 2. Seed Hidden Backend Data for Charts
    data_ws = wb.create_sheet(f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'

    # Main Chart Data (Stacked Column)
    main_data = [
        ["Market", "Product A", "Product B", "Product C"],
        ["North America", 62000, 24000, 12000],
        ["Europe", 54000, 27000, 14000],
        ["Asia", 46000, 31000, 21000],
        ["Latin America", 36000, 35000, 26000]
    ]
    for row in main_data:
        data_ws.append(row)
    
    # Line Chart 1 Data
    data_ws.append([]) 
    l1_start_row = data_ws.max_row + 1
    l1_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for row in l1_data:
        data_ws.append(row)
        
    # Line Chart 2 Data
    data_ws.append([])
    l2_start_row = data_ws.max_row + 1
    l2_data = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for row in l2_data:
        data_ws.append(row)

    # 3. Build Presentation Layer (Dashboard)
    ws = wb.create_sheet(sheet_name)
    
    # Hide gridlines to give an application-like dashboard feel
    ws.sheet_view.showGridLines = False
    if hasattr(ws.sheet_view, 'showRowColHeaders'):
        ws.sheet_view.showRowColHeaders = False

    header_fill = PatternFill(start_color=colors["bg"], end_color=colors["bg"], fill_type="solid")
    header_font = Font(color=colors["fg"], size=24, bold=True)
    
    # Title Banner (Rows 1 to 3)
    ws.merge_cells("A1:S3")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.fill = header_fill
    title_cell.font = header_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=19):
        for cell in row:
            cell.fill = header_fill
            
    # Sidebar Slicer/Control Panel Placeholder
    sidebar_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    sidebar_font = Font(color="808080", size=12, bold=True)
    ws.merge_cells("A5:C25")
    sidebar_cell = ws["A5"]
    sidebar_cell.value = "Filter Controls\n(Slicers & Timelines)"
    sidebar_cell.fill = sidebar_fill
    sidebar_cell.font = sidebar_font
    sidebar_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for row in ws.iter_rows(min_row=5, max_row=25, min_col=1, max_col=3):
        for cell in row:
            cell.fill = sidebar_fill

    # 4. Inject & Layout Charts
    
    # Main Chart (Stacked Column)
    main_chart = BarChart()
    main_chart.type = "col"
    main_chart.style = 10
    main_chart.grouping = "stacked"
    main_chart.overlap = 100
    main_chart.title = "Profit by Market & Product"
    main_chart.height = 10.5
    main_chart.width = 14
    
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    data = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=5)
    main_chart.add_data(data, titles_from_data=True)
    main_chart.set_categories(cats)
    ws.add_chart(main_chart, "D5")
    
    # Secondary Line Chart 1 (Top Right)
    lc1 = LineChart()
    lc1.title = "Units Sold Each Month"
    lc1.style = 13
    lc1.height = 5
    lc1.width = 14
    lc1.legend = None  # Cleaner without legend for simple trends
    lc1_data = Reference(data_ws, min_col=2, min_row=l1_start_row, max_row=l1_start_row+4)
    lc1_cats = Reference(data_ws, min_col=1, min_row=l1_start_row+1, max_row=l1_start_row+4)
    lc1.add_data(lc1_data, titles_from_data=True)
    lc1.set_categories(lc1_cats)
    ws.add_chart(lc1, "L5")

    # Secondary Line Chart 2 (Bottom Right)
    lc2 = LineChart()
    lc2.title = "Profit By Month"
    lc2.style = 13
    lc2.height = 5
    lc2.width = 14
    lc2.legend = None
    lc2_data = Reference(data_ws, min_col=2, min_row=l2_start_row, max_row=l2_start_row+4)
    lc2_cats = Reference(data_ws, min_col=1, min_row=l2_start_row+1, max_row=l2_start_row+4)
    lc2.add_data(lc2_data, titles_from_data=True)
    lc2.set_categories(lc2_cats)
    ws.add_chart(lc2, "L16")
