from openpyxl.chart import BarChart3D, Reference
from openpyxl.chart.data_source import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.colors import ColorChoice
from openpyxl.chart.view3D import View3D

def render(ws, anchor: str, *, theme: str = "corporate_blue", data: list = None, **kwargs) -> None:
    """
    Renders a 3D Container Chart (Cylinder) showing fill levels.
    """
    if not data:
        data = [
            ("Category", "Service Level"),
            ("Qtr1", 0.86),
            ("Qtr2", 0.82),
            ("Qtr3", 0.68),
            ("Qtr4", 0.64),
        ]
        
    try:
        from _helpers import get_theme_palette
        palette = get_theme_palette(theme)
        colors = [getattr(palette, f"accent{i}").replace("#", "") for i in range(1, 7)]
    except (ImportError, AttributeError):
        colors = ["E74C3C", "3498DB", "F1C40F", "2ECC71", "9B59B6", "34495E"]
        
    data_start_col = 20
    data_start_row = 1
    
    headers = ["Category", "Value", "Blank", "Lower Cap", "Mid 1a", "Mid 2a", "Mid 2b", "Mid 1b", "Upper Cap"]
    
    for i, h in enumerate(headers):
        ws.cell(row=data_start_row, column=data_start_col+i, value=h)
        
    for r_idx, row in enumerate(data[1:], start=data_start_row+1):
        cat, val = row
        blank = 1.0 - val
        # Proportions: Lower(0.20), Mid1(0.05), Mid2(0.02), Value(val), Blank(blank), Mid2(0.02), Mid1(0.05), Upper(0.15)
        row_vals = [cat, val, blank, 0.20, 0.05, 0.02, 0.02, 0.05, 0.15]
        for c_idx, cell_val in enumerate(row_vals):
            cell = ws.cell(row=r_idx, column=data_start_col+c_idx, value=cell_val)
            if c_idx > 0:
                cell.number_format = '0%'

    chart = BarChart3D()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.shape = "cylinder"
    chart.gapWidth = 80
    chart.title = "Quarterly Service Level"
    chart.legend = None
    chart.view3D = View3D(rotX=15, rotY=15, rightAngleAxes=False)
    
    # Hide axes to complete the clean infographic look
    chart.y_axis.delete = True
    chart.x_axis.delete = True
    
    # Add data
    data_ref = Reference(ws, min_col=data_start_col+1, min_row=data_start_row, max_col=data_start_col+8, max_row=data_start_row+len(data)-1)
    chart.add_data(data_ref, titles_from_data=True)
    
    # Reorder series to build the physical container from bottom to top
    # Initial openpyxl order: 0:Value, 1:Blank, 2:Lower Cap, 3:Mid 1a, 4:Mid 2a, 5:Mid 2b, 6:Mid 1b, 7:Upper Cap
    desired_order = [2, 3, 4, 0, 1, 5, 6, 7]
    chart.series = [chart.series[i] for i in desired_order]
    
    cats = Reference(ws, min_col=data_start_col, min_row=data_start_row+1, max_row=data_start_row+len(data)-1)
    chart.set_categories(cats)
    
    # Apply individual data point formatting to give each category a distinct color
    for s_idx, series in enumerate(chart.series):
        for pt_idx in range(len(data) - 1):
            base_color = colors[pt_idx % len(colors)]
            
            # Map visual roles based on the reordered series index
            if s_idx in (0, 3):    # Lower Cap, Value
                fill_hex = base_color
            elif s_idx == 4:       # Blank
                fill_hex = "E0E0E0"
            elif s_idx == 7:       # Upper Cap
                fill_hex = "D9D9D9"
            else:                  # Mid lines (acts as 3D highlights/rims)
                fill_hex = "FFFFFF"
                
            dp = DataPoint(idx=pt_idx)
            dp.spPr = GraphicalProperties(solidFill=ColorChoice(srgbClr=fill_hex))
            series.dPt.append(dp)
            
        # Configure data labels: values on the fluid layer, categories on the lid
        if s_idx == 3: # Value
            series.dLbls = DataLabelList(showVal=True, showCatName=False)
        elif s_idx == 7: # Upper Cap
            series.dLbls = DataLabelList(showCatName=True, showVal=False)
            
    chart.width = 16
    chart.height = 9
    ws.add_chart(chart, anchor)
    ws.sheet_view.showGridLines = False
