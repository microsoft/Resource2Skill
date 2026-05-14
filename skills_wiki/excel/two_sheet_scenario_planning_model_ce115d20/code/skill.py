from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import Series
from openpyxl.utils import get_column_letter

def render_workbook(wb, *, title: str = "Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Theme Configuration
    theme_colors = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF", "accent": "D9E1F2"},
        "emerald": {"header_bg": "005A36", "header_fg": "FFFFFF", "accent": "A8E6CF"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    header_fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_font = Font(color=palette["header_fg"], bold=True)
    bold_font = Font(bold=True)
    accent_fill = PatternFill(start_color=palette["accent"], end_color=palette["accent"], fill_type="solid")

    # 2. Setup Sheets
    if "Sheet" in wb.sheetnames:
        ws_assumptions = wb["Sheet"]
        ws_assumptions.title = "Assumptions"
    else:
        ws_assumptions = wb.create_sheet("Assumptions")
        
    ws_model = wb.create_sheet("Model")

    # 3. Build Assumptions Sheet
    ws_assumptions.column_dimensions['A'].width = 30
    ws_assumptions.column_dimensions['B'].width = 15

    ws_assumptions["A1"] = "Assumption"
    ws_assumptions["B1"] = "Value"
    ws_assumptions["A1"].fill = header_fill
    ws_assumptions["A1"].font = header_font
    ws_assumptions["B1"].fill = header_fill
    ws_assumptions["B1"].font = header_font

    assumptions_data = [
        ("ACCESSORIES", None),
        ("Starting Monthly Orders", 2000),
        ("Monthly Growth (%)", 0.06),
        ("Average Order Value ($)", 45),
        ("Gross Margin (%)", 0.35),
        ("DEVICES", None),
        ("Starting Monthly Orders", 500),
        ("Monthly Growth (%)", 0.04),
        ("Average Order Value ($)", 280),
        ("Gross Margin (%)", 0.22)
    ]

    for i, (label, val) in enumerate(assumptions_data, start=2):
        cell_A = ws_assumptions[f"A{i}"]
        cell_B = ws_assumptions[f"B{i}"]
        cell_A.value = label

        if val is None:
            # Subheader formatting
            cell_A.font = bold_font
            cell_A.fill = accent_fill
            cell_B.fill = accent_fill
        else:
            cell_B.value = val
            if "%" in label:
                cell_B.number_format = numbers.FORMAT_PERCENTAGE_00
            elif "$" in label:
                cell_B.number_format = '"$"#,##0.00'
            else:
                cell_B.number_format = '#,##0'

    # 4. Build Forecast Model Sheet
    months = 12
    ws_model.column_dimensions['A'].width = 25

    # Header Row
    for m in range(1, months + 1):
        col_letter = get_column_letter(m + 1)
        cell = ws_model[f"{col_letter}1"]
        cell.value = f"Month {m}"
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws_model.column_dimensions[col_letter].width = 15

    # Accessories Block
    ws_model["A2"] = "ACCESSORIES"
    ws_model["A2"].font = bold_font
    ws_model["A2"].fill = accent_fill
    for m in range(1, months + 1):
        ws_model[f"{get_column_letter(m+1)}2"].fill = accent_fill

    ws_model["A3"] = "Orders"
    ws_model["A4"] = "Revenue"
    ws_model["A5"] = "Gross Profit"

    # Devices Block
    ws_model["A7"] = "DEVICES"
    ws_model["A7"].font = bold_font
    ws_model["A7"].fill = accent_fill
    for m in range(1, months + 1):
        ws_model[f"{get_column_letter(m+1)}7"].fill = accent_fill

    ws_model["A8"] = "Orders"
    ws_model["A9"] = "Revenue"
    ws_model["A10"] = "Gross Profit"

    # Combined Totals Block
    ws_model["A12"] = "COMBINED TOTALS"
    ws_model["A12"].font = bold_font
    ws_model["A12"].fill = header_fill
    ws_model["A12"].font = header_font
    for m in range(1, months + 1):
        ws_model[f"{get_column_letter(m+1)}12"].fill = header_fill

    ws_model["A13"] = "Total Revenue"
    ws_model["A14"] = "Total Gross Profit"

    # Insert Cross-Sheet Dynamic Formulas
    for m in range(1, months + 1):
        col_letter = get_column_letter(m + 1)
        prev_col = get_column_letter(m)

        # Accessories Logic
        if m == 1:
            ws_model[f"{col_letter}3"] = "=Assumptions!$B$3"
        else:
            ws_model[f"{col_letter}3"] = f"={prev_col}3*(1+Assumptions!$B$4)"
        ws_model[f"{col_letter}4"] = f"={col_letter}3*Assumptions!$B$5"
        ws_model[f"{col_letter}5"] = f"={col_letter}4*Assumptions!$B$6"

        # Devices Logic
        if m == 1:
            ws_model[f"{col_letter}8"] = "=Assumptions!$B$8"
        else:
            ws_model[f"{col_letter}8"] = f"={prev_col}8*(1+Assumptions!$B$9)"
        ws_model[f"{col_letter}9"] = f"={col_letter}8*Assumptions!$B$10"
        ws_model[f"{col_letter}10"] = f"={col_letter}9*Assumptions!$B$11"

        # Combined Logic
        ws_model[f"{col_letter}13"] = f"={col_letter}4+{col_letter}9"
        ws_model[f"{col_letter}14"] = f"={col_letter}5+{col_letter}10"

        # Output Formatting
        for row in [3, 8]:
            ws_model[f"{col_letter}{row}"].number_format = '#,##0'
        for row in [4, 5, 9, 10, 13, 14]:
            ws_model[f"{col_letter}{row}"].number_format = '"$"#,##0'

    # 5. Build Summary Chart
    chart = LineChart()
    chart.title = "Monthly Revenue by Product Line"
    chart.style = 13
    chart.y_axis.title = "Revenue ($)"
    chart.x_axis.title = "Month"
    chart.height = 10
    chart.width = 20

    # Explicitly map series to avoid header row interpretation issues
    acc_rev_data = Reference(ws_model, min_col=2, min_row=4, max_col=months+1, max_row=4)
    acc_series = Series(acc_rev_data, title="Accessories Revenue")
    chart.series.append(acc_series)

    dev_rev_data = Reference(ws_model, min_col=2, min_row=9, max_col=months+1, max_row=9)
    dev_series = Series(dev_rev_data, title="Devices Revenue")
    chart.series.append(dev_series)

    cats = Reference(ws_model, min_col=2, min_row=1, max_col=months+1, max_row=1)
    chart.set_categories(cats)

    ws_model.add_chart(chart, "B16")
