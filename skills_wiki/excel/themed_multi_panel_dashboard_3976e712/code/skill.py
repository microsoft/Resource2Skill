def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Border, Side, Font
    from openpyxl.chart import DoughnutChart, LineChart, RadarChart, Reference
    from openpyxl.utils import get_column_letter

    # 1. Colors (Mocking a theme palette)
    bg_color = "F2F2F2"       # Canvas background
    card_color = "FFFFFF"     # Card background
    nav_color = "1F4E78"      # Left nav bar (Dark Blue)
    text_color = "1F4E78"     # Title text
    border_color = "CCCCCC"   # Card outlines
    
    # Setup sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    ws_inputs = wb.create_sheet("Inputs")
    
    # 2. Populate Inputs
    inputs_data = [
        ["KPI", "Actual", "Target", "% Complete", "Remainder"],
        ["Sales", 2544, 3000, 0.85, 0.15],
        ["Profit", 890, 1000, 0.89, 0.11],
        ["Customers", 87, 100, 0.87, 0.13],
        [],
        ["Month", "2021", "2022"],
        ["Jan", 201.9, 215.3],
        ["Feb", 204.2, 217.6],
        ["Mar", 198.6, 220.1],
        ["Apr", 199.2, 206.4],
        ["May", 206.4, 204.3],
        ["Jun", 195.1, 203.0],
        ["Jul", 192.4, 201.5],
        ["Aug", 186.3, 200.6],
        ["Sep", 194.2, 210.6],
        ["Oct", 225.0, 250.0],
        ["Nov", 205.2, 222.3],
        ["Dec", 204.3, 225.8],
        [],
        ["Metric", "Score"],
        ["Speed", 0.54],
        ["Quality", 0.86],
        ["Hygiene", 0.93],
        ["Service", 0.53],
        ["Availability", 0.95],
    ]
    for row in inputs_data:
        ws_inputs.append(row)
        
    # Format Inputs %
    for r in range(2, 5):
        ws_inputs.cell(row=r, column=4).number_format = '0%'
        ws_inputs.cell(row=r, column=5).number_format = '0%'
    for r in range(21, 26):
        ws_inputs.cell(row=r, column=2).number_format = '0%'

    # 3. Build Dashboard Layout Canvas
    bg_fill = PatternFill("solid", fgColor=bg_color)
    nav_fill = PatternFill("solid", fgColor=nav_color)
    card_fill = PatternFill("solid", fgColor=card_color)
    
    # Apply global gray background
    for row in ws_dash.iter_rows(min_row=1, max_row=25, min_col=2, max_col=18):
        for cell in row:
            cell.fill = bg_fill
            
    # Draw Left Nav
    for row in ws_dash.iter_rows(min_row=1, max_row=25, min_col=1, max_col=1):
        for cell in row:
            cell.fill = nav_fill
    ws_dash.column_dimensions['A'].width = 8
    
    # Helper to draw bordered white cards
    def draw_card(min_col, min_row, max_col, max_row):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws_dash.cell(row=r, column=c)
                cell.fill = card_fill
                
                top = Side(style="thin", color=border_color) if r == min_row else None
                bottom = Side(style="thin", color=border_color) if r == max_row else None
                left = Side(style="thin", color=border_color) if c == min_col else None
                right = Side(style="thin", color=border_color) if c == max_col else None
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)

    # 4. Draw Cards & Spacing
    ws_dash.column_dimensions['B'].width = 3
    ws_dash.column_dimensions['G'].width = 3
    ws_dash.column_dimensions['L'].width = 3

    draw_card(3, 3, 6, 8)    # Sales KPI Card
    draw_card(8, 3, 11, 8)   # Profit KPI Card
    draw_card(13, 3, 16, 8)  # Customers KPI Card
    draw_card(3, 10, 11, 22) # Trend Chart Card
    draw_card(13, 10, 16, 22)# Radar Chart Card
    
    # Title
    ws_dash["C1"] = title
    ws_dash["C1"].font = Font(size=20, bold=True, color=text_color)
    
    # 5. Populate KPI Cards & Doughnut Charts
    kpis = [
        {"name": "Sales", "val_row": 2, "col_offset": 3, "format": '"$"#,##0'},
        {"name": "Profit", "val_row": 3, "col_offset": 8, "format": '"$"#,##0'},
        {"name": "Customers", "val_row": 4, "col_offset": 13, "format": '#,##0'}
    ]
    
    for kpi in kpis:
        c = kpi["col_offset"]
        
        # KPI Title
        title_cell = ws_dash.cell(row=3, column=c)
        title_cell.value = kpi["name"]
        title_cell.font = Font(size=14, bold=True, color=text_color)
        
        # KPI Value
        val_cell = ws_dash.cell(row=5, column=c)
        val_cell.value = f"=Inputs!B{kpi['val_row']}"
        val_cell.font = Font(size=22, bold=True)
        val_cell.number_format = kpi["format"]

        # KPI % Complete Sub-label
        pct_cell = ws_dash.cell(row=6, column=c)
        pct_cell.value = f"=Inputs!D{kpi['val_row']}"
        pct_cell.font = Font(size=12, color="7F7F7F")
        pct_cell.number_format = '0% "Completed"'
        
        # Associated Doughnut Chart
        donut = DoughnutChart()
        donut.holeSize = 65
        d_data = Reference(ws_inputs, min_col=4, max_col=5, min_row=kpi["val_row"], max_row=kpi["val_row"])
        d_cats = Reference(ws_inputs, min_col=4, max_col=5, min_row=1, max_row=1)
        donut.add_data(d_data, from_rows=True)
        donut.set_categories(d_cats)
        
        donut.legend = None
        donut.title = None
        donut.width = 4.5
        donut.height = 4.5
        
        # Blend chart background into the white card
        donut.graphical_properties.noFill = True
        donut.graphical_properties.line.noFill = True
        
        ws_dash.add_chart(donut, f"{get_column_letter(c+2)}3")

    # 6. Add Trend Line Chart
    line_chart = LineChart()
    line_data = Reference(ws_inputs, min_col=2, max_col=3, min_row=6, max_row=18)
    line_cats = Reference(ws_inputs, min_col=1, min_row=7, max_row=18)
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)
    line_chart.title = "2021-2022 Sales Trend (in millions)"
    line_chart.width = 15
    line_chart.height = 6.5
    line_chart.legend.position = 't'
    
    line_chart.graphical_properties.noFill = True
    line_chart.graphical_properties.line.noFill = True
    
    if len(line_chart.series) > 0:
        line_chart.series[0].marker.symbol = "circle"
        line_chart.series[1].marker.symbol = "circle"
        
    ws_dash.add_chart(line_chart, "C10")
    
    # 7. Add Radar Chart
    radar = RadarChart()
    radar.type = "filled"
    r_data = Reference(ws_inputs, min_col=2, min_row=20, max_row=25)
    r_cats = Reference(ws_inputs, min_col=1, min_row=21, max_row=25)
    radar.add_data(r_data, titles_from_data=True)
    radar.set_categories(r_cats)
    radar.title = "Customer Satisfaction"
    radar.legend = None
    radar.width = 6.5
    radar.height = 6.5
    
    radar.graphical_properties.noFill = True
    radar.graphical_properties.line.noFill = True
    
    ws_dash.add_chart(radar, "M10")
