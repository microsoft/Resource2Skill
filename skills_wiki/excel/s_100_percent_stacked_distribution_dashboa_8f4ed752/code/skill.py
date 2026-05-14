from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Distribution by Region", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a worksheet simulating a pivot table aggregation paired with a 100% Stacked Column Chart.
    """
    ws = wb.create_sheet(sheet_name)

    # Sample aggregated data simulating a Pivot Table output (Years down, Regions across)
    data = [
        ["Year", "Canada", "China", "Germany", "United Kingdom", "United States"],
        [2020, 3135000, 11160000, 5245000, 4090000, 134250000],
        [2021, 3660000, 12520000, 5920000, 4615000, 147000000],
        [2022, 4185000, 13800000, 6645000, 5140000, 160950000],
        [2023, 4710000, 15325000, 7370000, 5665000, 177400000],
        [2024, 5685000, 16130000, 10775000, 9840000, 188150000]
    ]

    # Theme definitions (Fallback to corporate blue palette)
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    # Inject data and apply basic pivot-like formatting
    for r_idx, row in enumerate(data, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            
            if r_idx == 1:
                # Format headers
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                # Format numbers as currency with zero decimals
                if c_idx > 1:
                    cell.number_format = '"$"#,##0'
                else:
                    cell.alignment = Alignment(horizontal="center")

    # Adjust column widths for clean presentation
    for c in range(1, len(data[0]) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18

    # Build the 100% Stacked Column Chart
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    chart.title = title
    chart.height = 12
    chart.width = 20

    # Define Data and Categories
    # Series (Countries) span columns 2 through 6, including the header row for titles
    chart_data = Reference(ws, min_col=2, max_col=6, min_row=1, max_row=len(data))
    # Categories (Years) span rows 2 through 6 in column 1
    cats = Reference(ws, min_col=1, max_col=1, min_row=2, max_row=len(data))

    # Assemble chart
    chart.add_data(chart_data, titles_from_data=True)
    chart.set_categories(cats)

    # Position chart below the data grid
    anchor_cell = f"A{len(data) + 3}"
    ws.add_chart(chart, anchor_cell)
