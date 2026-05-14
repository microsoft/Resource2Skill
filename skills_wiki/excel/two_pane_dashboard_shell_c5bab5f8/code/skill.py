from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, kpis: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a two-pane dashboard shell with a dark left KPI sidebar and a light main canvas.
    
    :param wb: The openpyxl Workbook object.
    :param sheet_name: Name of the new dashboard sheet.
    :param title: The title of the dashboard (placed at top of sidebar).
    :param kpis: List of tuples [(label, value), ...] to display in the sidebar.
    :param theme: The color theme to apply.
    """
    # 1. Theme Setup (Fallback mapping for self-containment)
    theme_palettes = {
        "corporate_blue": {"primary": "1F4E78", "text_light": "FFFFFF", "bg_light": "F2F2F2"},
        "botanical_green": {"primary": "2D4A22", "text_light": "FFFFFF", "bg_light": "EAF0E6"}, # Matches the video's 'Viva Calif' brand
        "dark_mode": {"primary": "1A1A1A", "text_light": "E0E0E0", "bg_light": "2D2D2D"}
    }
    palette = theme_palettes.get(theme, theme_palettes["corporate_blue"])

    # 2. Sheet Initialization
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # 3. Layout Dimensions
    ws.column_dimensions['A'].width = 3   # Left padding
    ws.column_dimensions['B'].width = 25  # KPI Sidebar
    ws.column_dimensions['C'].width = 3   # Right padding for sidebar
    ws.column_dimensions['D'].width = 3   # Left padding for main area
    
    for col_idx in range(5, 21): # E through T for main canvas
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 12

    # 4. Apply Background Fills (Simulating the large rectangle shapes from the tutorial)
    sidebar_fill = PatternFill(start_color=palette["primary"], fill_type="solid")
    main_fill = PatternFill(start_color=palette["bg_light"], fill_type="solid")

    for row in range(1, 50):
        # Sidebar Area (Cols A-C)
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = sidebar_fill
        # Main Canvas Area (Cols D-T)
        for col in range(4, 21):
            ws.cell(row=row, column=col).fill = main_fill

    # 5. Add Dashboard Title in Sidebar
    title_cell = ws['B2']
    title_cell.value = title.upper()
    title_cell.font = Font(color=palette["text_light"], size=22, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 6. Inject KPI Strip
    if not kpis:
        # Default realistic data if none provided
        kpis = [
            ("Orders", "2,400"), 
            ("Quantity", "11,997"), 
            ("Amount", "$649.0K"), 
            ("Avg. Rating", "4.0")
        ]

    start_row = 6
    for label, val in kpis:
        # KPI Label
        lbl_cell = ws.cell(row=start_row, column=2)
        lbl_cell.value = label
        lbl_cell.font = Font(color=palette["text_light"], size=12, italic=True)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")

        # KPI Value
        val_cell = ws.cell(row=start_row + 1, column=2)
        val_cell.value = val
        val_cell.font = Font(color=palette["text_light"], size=20, bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")

        start_row += 4 # Spacing between KPIs
