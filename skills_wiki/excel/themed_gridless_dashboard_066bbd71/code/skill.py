def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Set up data sheet (separated from presentation view)
    ws_data = wb.active
    ws_data.title = "Data"
    
    # Sample realistic performance data
    dataset = [
        ["Month", "Profit", "Units Sold"],
        ["Jan", 124812, 50601],
        ["Feb", 228275, 95622],
        ["Mar", 160228, 65481],
        ["Apr", 136337, 52970],
        ["May", 210500, 80200],
        ["Jun", 195000, 75000],
    ]
    for row in dataset:
        ws_data.append(row)
        
    # 2. Set up Dashboard presentation sheet
    ws_dash = wb.create_sheet("Dashboard")
    
    # Clean up the canvas: hide gridlines and row/col headers to create a "software" feel
    ws_dash.sheet_view.showGridLines = False
    ws_dash.sheet_view.showRowColHeaders = False
    
    # 3. Dashboard Banner
    # In a full framework implementation, fetch these from the active theme palette
    banner_fill = PatternFill("solid", fgColor="1F4E78")
    banner_font = Font(size=24, bold=True, color="FFFFFF")
    
    ws_dash.merge_cells("A1:O3")
    banner = ws_dash["A1"]
    banner.value = title
    banner.font = banner_font
    banner.fill = banner_fill
    banner.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Chart 1: Profit by Month (Bar Chart)
    bar_chart = BarChart()
    bar_chart.title = "Profit by Month"
    bar_chart.style = 11 # Clean modern preset style
    bar_chart.height = 9.5
    bar_chart.width = 16
    
    cats = Reference(ws_data, min_col=1, min_row=2, max_row=len(dataset))
    profit_data = Reference(ws_data, min_col=2, min_row=1, max_row=len(dataset))
    
    bar_chart.add_data(profit_data, titles_from_data=True)
    bar_chart.set_categories(cats)
    bar_chart.legend = None  # Remove legend to match video's clean look
    
    ws_dash.add_chart(bar_chart, "B5")
    
    # 5. Chart 2: Units Sold (Line Chart)
    line_chart = LineChart()
    line_chart.title = "Units Sold each month"
    line_chart.style = 12
    line_chart.height = 9.5
    line_chart.width = 16
    
    units_data = Reference(ws_data, min_col=3, min_row=1, max_row=len(dataset))
    
    line_chart.add_data(units_data, titles_from_data=True)
    line_chart.set_categories(cats)
    line_chart.legend = None # Remove legend
    
    ws_dash.add_chart(line_chart, "J5")
    
    # Ensure the dashboard is the active sheet when the workbook is opened
    wb.active = ws_dash
