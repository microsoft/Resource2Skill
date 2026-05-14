def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # Fallback palette if a standard theme dict isn't provided via kwargs
    palette = kwargs.get("palette", {
        "primary": "4F81BD",
        "text_on_primary": "FFFFFF"
    })

    # 1. Setup Dashboard Sheet
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    # Create the full-width Dashboard Banner
    ws_dash.merge_cells("A1:R2")
    title_cell = ws_dash["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette.get("text_on_primary", "FFFFFF"))
    title_cell.fill = PatternFill(start_color=palette.get("primary", "4F81BD"), end_color=palette.get("primary", "4F81BD"), fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Setup Hidden Data Sheet
    ws_data = wb.create_sheet("ChartData")
    ws_data.sheet_state = "hidden"
    
    # Sample Aggregation: Profit by Market & Cookie Type
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Sugar"],
        ["India", 62349, 4872, 21028, 25085],
        ["United Kingdom", 46530, 5220, 11497, 14620],
        ["United States", 36657, 6368, 22260, 9937],
        ["Philippines", 54618, 7026, 22005, 8313],
    ]
    for r_idx, row in enumerate(market_data, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
            if r_idx > 1 and c_idx > 1:
                cell.number_format = "$#,##0"
                
    # Sample Aggregation: Units & Profit by Month
    month_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    for r_idx, row in enumerate(month_data, 10):
        for c_idx, val in enumerate(row, 1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
            if r_idx > 10 and c_idx == 2:
                cell.number_format = "#,##0"
            elif r_idx > 10 and c_idx == 3:
                cell.number_format = "$#,##0"
                
    # 3. Create & Position Charts
    # Chart 1: Stacked Column (Profit by Market)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    
    bar_data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    bar_cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)
    bar_chart.height = 11
    bar_chart.width = 16
    ws_dash.add_chart(bar_chart, "B4")
    
    # Chart 2: Line Chart (Units by Month)
    line_units = LineChart()
    line_units.title = "Units sold each month"
    
    units_data_ref = Reference(ws_data, min_col=2, min_row=10, max_col=2, max_row=14)
    month_cats_ref = Reference(ws_data, min_col=1, min_row=11, max_row=14)
    line_units.add_data(units_data_ref, titles_from_data=True)
    line_units.set_categories(month_cats_ref)
    line_units.legend = None  # Clean interface, hide legend if only 1 series
    line_units.height = 7
    line_units.width = 12
    ws_dash.add_chart(line_units, "J4")
    
    # Chart 3: Line Chart (Profit by Month)
    line_profit = LineChart()
    line_profit.title = "Profit by month"
    
    profit_data_ref = Reference(ws_data, min_col=3, min_row=10, max_col=3, max_row=14)
    line_profit.add_data(profit_data_ref, titles_from_data=True)
    line_profit.set_categories(month_cats_ref)
    line_profit.legend = None
    line_profit.height = 7
    line_profit.width = 12
    ws_dash.add_chart(line_profit, "J16")
