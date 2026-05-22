from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", main_label: str = "Revenue", main_value_ref: str = "B2", sec_value_ref: str = "C2", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a dynamic KPI card using cell formatting to emulate a floating shape widget.
    Links directly to source data cells for live updates.
    """
    # Self-contained theme fallback
    theme_colors = {
        "corporate_blue": {"bg": "17365D", "fg": "FFFFFF", "sec_bg": "366092"},
        "executive_gray": {"bg": "404040", "fg": "FFFFFF", "sec_bg": "7F7F7F"},
        "emerald_green": {"bg": "215967", "fg": "FFFFFF", "sec_bg": "31859C"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # Parse anchor coordinate
    col_str, row_str = coordinate_from_string(anchor)
    start_col = column_index_from_string(col_str)
    start_row = int(row_str)

    # Apply structural merges
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+1)
    ws.merge_cells(start_row=start_row+1, start_column=start_col+1, end_row=start_row+2, end_column=start_col+1)

    # Define Styles
    main_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    sec_fill = PatternFill(start_color=palette["sec_bg"], end_color=palette["sec_bg"], fill_type="solid")
    
    title_font = Font(color=palette["fg"], size=14, bold=True)
    label_font = Font(color=palette["fg"], size=10)
    value_font = Font(color=palette["fg"], size=16, bold=True)
    sec_value_font = Font(color=palette["fg"], size=18, bold=True)
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="center", vertical="bottom") # Centered relative to its column
    val_align = Alignment(horizontal="center", vertical="top")

    # 1. Title (Top Merged Row)
    c_title = ws.cell(row=start_row, column=start_col, value=title)
    c_title.font = title_font
    c_title.alignment = center_align
    # Fill the entire merged area
    for c in range(start_col, start_col + 2):
        ws.cell(row=start_row, column=c).fill = main_fill

    # 2. Main Label (Bottom Left, Top Half)
    c_label = ws.cell(row=start_row+1, column=start_col, value=main_label)
    c_label.font = label_font
    c_label.fill = main_fill
    c_label.alignment = left_align

    # 3. Main Value (Bottom Left, Bottom Half)
    c_value = ws.cell(row=start_row+2, column=start_col, value=f"={main_value_ref}")
    c_value.font = value_font
    c_value.fill = main_fill
    c_value.alignment = val_align
    c_value.number_format = "$#,##0"

    # 4. Secondary Value Badge (Bottom Right, Merged Vertically)
    c_sec = ws.cell(row=start_row+1, column=start_col+1, value=f"={sec_value_ref}")
    c_sec.font = sec_value_font
    c_sec.alignment = center_align
    c_sec.number_format = "0%"
    # Fill the entire merged badge area
    ws.cell(row=start_row+1, column=start_col+1).fill = sec_fill
    ws.cell(row=start_row+2, column=start_col+1).fill = sec_fill

    # Adjust dimensions for widget proportions
    ws.column_dimensions[get_column_letter(start_col)].width = 18
    ws.column_dimensions[get_column_letter(start_col+1)].width = 12
    ws.row_dimensions[start_row].height = 25
    ws.row_dimensions[start_row+1].height = 18
    ws.row_dimensions[start_row+2].height = 25
