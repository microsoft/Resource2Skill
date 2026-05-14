def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list[dict], theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    ws = wb.create_sheet(sheet_name)
    
    # Safely load theme colors
    try:
        from skills_library.excel.components._helpers import get_theme
        t = get_theme(theme)
        primary = getattr(t, "primary", "4F2D7F").replace("#", "")
        primary_light = getattr(t, "background", "F2EFF5").replace("#", "")
        accent = getattr(t, "accent", "E4A925").replace("#", "")
    except ImportError:
        # Fallback to the tutorial's Purple/Gold palette
        primary = "4F2D7F"
        primary_light = "F2EFF5"
        accent = "E4A925"

    text_light = "FFFFFF"
    text_dark = "666666"

    # 1. Paint the split background canvas
    header_fill = PatternFill(start_color=primary, fill_type="solid")
    body_fill = PatternFill(start_color=primary_light, fill_type="solid")
    
    for row in range(1, 10):
        ws.row_dimensions[row].height = 25
        for col in range(1, 30):
            ws.cell(row=row, column=col).fill = header_fill
            
    for row in range(10, 40):
        for col in range(1, 30):
            ws.cell(row=row, column=col).fill = body_fill
            
    # 2. Add Title & Subtitle
    title_cell = ws.cell(row=2, column=2, value=title)
    title_cell.font = Font(size=32, color=text_light, bold=True)
    
    sub_cell = ws.cell(row=4, column=2, value=subtitle)
    sub_cell.font = Font(size=14, color=accent)
    
    # 3. Render In-Cell KPI Cards
    # Example input: [{"label": "CALLS", "value": "16,749", "icon": "📞"}, ...]
    # KPI cards start at column H (8) and span rows 3 to 6
    start_col = 8
    
    for kpi in kpis:
        icon_fill = PatternFill(start_color=accent, fill_type="solid")
        data_fill = PatternFill(start_color="FFFFFF", fill_type="solid")
        border_edge = Side(style="thin", color="D9D9D9")
        
        # Set column widths for the card geometry
        ws.column_dimensions[get_column_letter(start_col)].width = 8
        ws.column_dimensions[get_column_letter(start_col+1)].width = 12
        ws.column_dimensions[get_column_letter(start_col+2)].width = 12
        
        # Apply fills and outer borders to the 3x4 block
        for r in range(3, 7):
            for c in range(start_col, start_col+3):
                cell = ws.cell(row=r, column=c)
                cell.fill = icon_fill if c == start_col else data_fill
                
                b_top = border_edge if r == 3 else None
                b_bottom = border_edge if r == 6 else None
                b_left = border_edge if c == start_col else None
                b_right = border_edge if c == start_col + 2 else None
                cell.border = Border(top=b_top, bottom=b_bottom, left=b_left, right=b_right)
        
        # Merge and populate Icon (Left column)
        ws.merge_cells(start_row=3, start_column=start_col, end_row=6, end_column=start_col)
        icon_cell = ws.cell(row=3, column=start_col, value=kpi.get("icon", ""))
        icon_cell.font = Font(size=24, color=text_light)
        icon_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Merge and populate Value (Top right)
        ws.merge_cells(start_row=3, start_column=start_col+1, end_row=4, end_column=start_col+2)
        val_cell = ws.cell(row=3, column=start_col+1, value=kpi.get("value", ""))
        val_cell.font = Font(size=22, color=primary, bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="bottom")
        
        # Merge and populate Label (Bottom right)
        ws.merge_cells(start_row=5, start_column=start_col+1, end_row=6, end_column=start_col+2)
        lbl_cell = ws.cell(row=5, column=start_col+1, value=kpi.get("label", ""))
        lbl_cell.font = Font(size=12, color=text_dark, bold=True)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")
        
        # Step forward for the next card (3 columns for card + 1 column spacing)
        start_col += 4

    # Remove gridlines for a clean dashboard look
    ws.sheet_view.showGridLines = False
