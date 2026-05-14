from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter, coordinate_from_string, column_index_from_string

def render(
    ws, 
    anchor: str, 
    *, 
    title: str = "DSO (Days Sales Outstanding)", 
    current_val: float = 31, 
    target_val: float = 45, 
    prior_val: float = 41, 
    higher_is_better: bool = False, 
    number_format: str = "#,##0",
    theme: str = "corporate_blue", 
    **kwargs
) -> None:
    """
    Renders a 4x3 cell KPI card with large font and conditional formatting.
    """
    # Fallback theme palette
    theme_colors = {
        "corporate_blue": {"accent": "4F81BD", "text": "FFFFFF", "sub_bg": "F2F2F2", "sub_text": "595959"},
        "executive_gray": {"accent": "404040", "text": "FFFFFF", "sub_bg": "D9D9D9", "sub_text": "262626"}
    }
    colors = theme_colors.get(theme, theme_colors["corporate_blue"])

    # Base Styles
    header_fill = PatternFill(start_color=colors["accent"], end_color=colors["accent"], fill_type="solid")
    header_font = Font(color=colors["text"], bold=True)
    sub_fill = PatternFill(start_color=colors["sub_bg"], end_color=colors["sub_bg"], fill_type="solid")
    sub_font = Font(color=colors["sub_text"], size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color="BFBFBF"),
        right=Side(style='thin', color="BFBFBF"),
        top=Side(style='thin', color="BFBFBF"),
        bottom=Side(style='thin', color="BFBFBF")
    )

    # Parse anchor coordinate
    col_str, row_str = coordinate_from_string(anchor)
    anchor_col = column_index_from_string(col_str)
    anchor_row = int(row_str)

    # Set column widths for proper display
    for i in range(4):
        ws.column_dimensions[get_column_letter(anchor_col + i)].width = 14

    # ---------------------------------------------------------
    # Row 0: Title Header
    # ---------------------------------------------------------
    ws.cell(row=anchor_row, column=anchor_col, value=title)
    ws.merge_cells(start_row=anchor_row, start_column=anchor_col, end_row=anchor_row, end_column=anchor_col+3)
    
    for c in range(anchor_col, anchor_col+4):
        cell = ws.cell(row=anchor_row, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # ---------------------------------------------------------
    # Row 1: Main Metric Value
    # ---------------------------------------------------------
    val_cell = ws.cell(row=anchor_row+1, column=anchor_col)
    val_cell.value = current_val
    ws.merge_cells(start_row=anchor_row+1, start_column=anchor_col, end_row=anchor_row+1, end_column=anchor_col+3)
    
    for c in range(anchor_col, anchor_col+4):
        cell = ws.cell(row=anchor_row+1, column=c)
        cell.font = Font(size=24, bold=True)
        cell.alignment = center_align
        cell.border = thin_border
        cell.number_format = number_format

    # ---------------------------------------------------------
    # Row 2: Sub-Metrics (Target & Prior)
    # ---------------------------------------------------------
    subs = [
        ("Vs. Target", anchor_col),
        (target_val, anchor_col+1),
        ("Vs. Prior", anchor_col+2),
        (prior_val, anchor_col+3)
    ]
    
    for val, col in subs:
        cell = ws.cell(row=anchor_row+2, column=col)
        cell.value = val
        cell.fill = sub_fill
        cell.font = sub_font
        cell.alignment = center_align
        cell.border = thin_border
        if isinstance(val, (int, float)):
            cell.number_format = number_format

    # ---------------------------------------------------------
    # Conditional Formatting for Main Metric
    # ---------------------------------------------------------
    target_col_letter = get_column_letter(anchor_col + 1)
    abs_target_coord = f"${target_col_letter}${anchor_row + 2}"
    
    # Range covering the merged Main Value cells
    cf_range = f"{get_column_letter(anchor_col)}{anchor_row+1}:{get_column_letter(anchor_col+3)}{anchor_row+1}"

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(color="006100", size=24, bold=True)
    
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006", size=24, bold=True)

    if higher_is_better:
        good_op, bad_op = 'greaterThanOrEqual', 'lessThan'
    else:
        # e.g., DSO (Days Sales Outstanding) is better when lower
        good_op, bad_op = 'lessThanOrEqual', 'greaterThan'

    # Add Good Condition
    ws.conditional_formatting.add(
        cf_range,
        CellIsRule(operator=good_op, formula=[abs_target_coord], fill=green_fill, font=green_font)
    )
    # Add Bad Condition
    ws.conditional_formatting.add(
        cf_range,
        CellIsRule(operator=bad_op, formula=[abs_target_coord], fill=red_fill, font=red_font)
    )
