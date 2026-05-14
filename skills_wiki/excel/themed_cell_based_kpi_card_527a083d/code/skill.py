from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, 
           title: str = "Asia", 
           label: str = "Revenue", 
           value_ref: str = None, 
           value_val: float = 369989, 
           badge_ref: str = None, 
           badge_val: float = 0.05, 
           theme: str = "corporate_blue", 
           **kwargs) -> None:
    """
    Renders a 3x2 cell-based KPI card with a main value and a side badge.
    """
    # Theme palette definition
    palettes = {
        "corporate_blue": {"card_bg": "172B4D", "card_fg": "FFFFFF", "badge_bg": "091E42"},
        "emerald": {"card_bg": "065F46", "card_fg": "FFFFFF", "badge_bg": "047857"},
        "slate": {"card_bg": "334155", "card_fg": "FFFFFF", "badge_bg": "0F172A"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])
    
    col_str, row = coordinate_from_string(anchor)
    col = column_index_from_string(col_str)
    
    card_fill = PatternFill(start_color=colors["card_bg"], end_color=colors["card_bg"], fill_type="solid")
    badge_fill = PatternFill(start_color=colors["badge_bg"], end_color=colors["badge_bg"], fill_type="solid")
    
    font_title = Font(color=colors["card_fg"], size=14, bold=True)
    font_label = Font(color=colors["card_fg"], size=10)
    font_value = Font(color=colors["card_fg"], size=16, bold=True)
    font_badge = Font(color=colors["card_fg"], size=14, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Apply baseline fills to all cells in the 3x2 grid to prevent unstyled gaps
    for r in range(row, row + 3):
        ws.cell(row=r, column=col).fill = card_fill
        ws.cell(row=r, column=col+1).fill = badge_fill

    # Row 1: Title
    cell_title = ws.cell(row=row, column=col, value=title)
    cell_title.font = font_title
    cell_title.alignment = align_center
    
    # Row 2: Label
    cell_label = ws.cell(row=row+1, column=col, value=label)
    cell_label.font = font_label
    cell_label.alignment = align_center
    
    # Row 3: Value
    cell_val = ws.cell(row=row+2, column=col)
    cell_val.value = f"={value_ref}" if value_ref else value_val
    cell_val.font = font_value
    cell_val.alignment = align_center
    cell_val.number_format = '"$"#,##0'
    
    # Badge (Column 2, spanning Row 1 to Row 3)
    ws.merge_cells(start_row=row, start_column=col+1, end_row=row+2, end_column=col+1)
    cell_badge = ws.cell(row=row, column=col+1)
    cell_badge.value = f"={badge_ref}" if badge_ref else badge_val
    cell_badge.font = font_badge
    cell_badge.alignment = align_center
    cell_badge.number_format = '0%'
    
    # Adjust sizing to maintain card proportions
    ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions[get_column_letter(col+1)].width = 10
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row+1].height = 14
    ws.row_dimensions[row+2].height = 24
