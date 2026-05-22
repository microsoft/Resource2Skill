from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.styles import Font

def render(ws, anchor: str, *, kpi_name: str = "Sales", actual_val: float = 2544, target_val: float = 3000, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a KPI data block and an accompanying 'progress ring' style Donut Chart.
    """
    row, col = coordinate_to_tuple(anchor)

    # 1. Write KPI Title Header
    ws.cell(row=row, column=col, value=kpi_name).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)

    # 2. Write Labels
    labels = ["Actual", "Target", "% Complete", "% Remainder"]
    for i, label in enumerate(labels):
        ws.cell(row=row+1+i, column=col, value=label)

    # 3. Write Values & Format
    c_actual = ws.cell(row=row+1, column=col+1, value=actual_val)
    c_actual.number_format = "$#,##0"
    
    c_target = ws.cell(row=row+2, column=col+1, value=target_val)
    c_target.number_format = "$#,##0"

    # 4. Write Dynamic Formulas
    actual_ref = f"{get_column_letter(col+1)}{row+1}"
    target_ref = f"{get_column_letter(col+1)}{row+2}"
    pct_complete_ref = f"{get_column_letter(col+1)}{row+3}"

    c_pct = ws.cell(row=row+3, column=col+1, value=f"={actual_ref}/{target_ref}")
    c_pct.number_format = "0%"

    c_rem = ws.cell(row=row+4, column=col+1, value=f"=1-{pct_complete_ref}")
    c_rem.number_format = "0%"

    # 5. Generate and Customize the Doughnut Chart
    chart = DoughnutChart()
    chart.title = None
    chart.legend = None
    chart.holeSize = 65  # Expands the inner hole for a modern "progress ring" look

    # Point chart to the calculated percentages
    data = Reference(ws, min_col=col+1, min_row=row+3, max_row=row+4)
    cats = Reference(ws, min_col=col, min_row=row+3, max_row=row+4)

    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    # Resize chart to act as a compact card component
    chart.width = 6
    chart.height = 4.5

    # Anchor the chart directly to the right of the data block
    chart_anchor = f"{get_column_letter(col+2)}{row}"
    ws.add_chart(chart, chart_anchor)
