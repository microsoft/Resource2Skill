from datetime import date, timedelta
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import PatternFill, Font

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.title = title if title else "Task Tracker"
    
    # 1. Define Headers
    headers = ["Task", "Subject", "Type", "Due", "Time Req'd", "Status", "Days Available", "Notes"]
    ws.append(headers)
    
    # 2. Populate Realistic Dummy Data
    today = date.today()
    tasks = [
        ["Numerical Methods Report", "Applied Math", "Assignment", today + timedelta(days=10), 5, "Not Started", "=[@Due]-TODAY()", ""],
        ["Logic & Proofs Homework", "Pure Math", "Course Work", today - timedelta(days=2), 3, "Not Started", "=[@Due]-TODAY()", "Overdue example"],
        ["Data Structures Lab", "Algorithms", "Assignment", today + timedelta(days=5), 4, "Started", "=[@Due]-TODAY()", "Halfway done"],
        ["Encryption Algorithms", "Cyber Security", "Exam", today + timedelta(days=20), 2, "Completed", "=[@Due]-TODAY()", "Review notes"],
        ["Sorting Algorithm Comparison", "Algorithms", "Assignment", today + timedelta(days=1), 4, "Started", "=[@Due]-TODAY()", "Focus on quicksort"],
    ]
    
    for row in tasks:
        ws.append(row)
        
    # Format Date column for readability
    for cell in ws["D"][1:]:
        cell.number_format = 'DD/MM/YYYY'
        
    # 3. Apply Excel Table Structuring
    table_ref = f"A1:H{len(tasks) + 1}"
    tab = Table(displayName="TaskTracker", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleLight13", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)
    
    # Adjust Column Widths
    column_widths = {"A": 32, "B": 20, "C": 15, "D": 14, "E": 12, "F": 15, "G": 16, "H": 30}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # 4. Insert Data Validation Dropdowns
    dv_subject = DataValidation(type="list", formula1='"Applied Math,Pure Math,Algorithms,Cyber Security"', allow_blank=True)
    dv_type = DataValidation(type="list", formula1='"Assignment,Course Work,Exam"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"Not Started,Started,Completed"', allow_blank=True)
    
    ws.add_data_validation(dv_subject)
    ws.add_data_validation(dv_type)
    ws.add_data_validation(dv_status)
    
    dv_subject.add("B2:B100")
    dv_type.add("C2:C100")
    dv_status.add("F2:F100")

    # 5. Apply Priority-layered Conditional Formatting
    # Rule 1: Completed Row - Strikethrough & Gray (Top priority, stopIfTrue overrides subsequent rules)
    gray_strike = Font(color="A6A6A6", strikethrough=True)
    ws.conditional_formatting.add("A2:H100", FormulaRule(formula=['=$F2="Completed"'], font=gray_strike, stopIfTrue=True))
    
    # Rule 2: Overdue - Negative Days Available
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_text = Font(color="9C0006")
    ws.conditional_formatting.add("G2:G100", CellIsRule(operator="lessThan", formula=["0"], fill=red_fill, font=red_text))
    
    # Rule 3: Not Started - Dark Blue Fill, White Text
    blue_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    white_text = Font(color="FFFFFF")
    ws.conditional_formatting.add("F2:F100", CellIsRule(operator="equal", formula=['"Not Started"'], fill=blue_fill, font=white_text))
    
    # Rule 4: Started - Purple Fill, White Text
    purple_fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
    ws.conditional_formatting.add("F2:F100", CellIsRule(operator="equal", formula=['"Started"'], fill=purple_fill, font=white_text))
