from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
import re

def render(ws, anchor: str, *, title: str = "Sales Dashboard", subtitle: str = "Evaluating Sales Agent Performance", kpis: list = None, theme: dict = None, **kwargs) -> None:
    """
    Renders a deep-colored top banner containing a title, subtitle, and a strip of
    white KPI cards with an accent-colored left edge.
    """
    if theme is None:
        theme = {}
        
    # Default data matching the video's context
    if kpis is None:
        kpis = [
            {"label": "CALLS", "value": 16749, "format": "#,##0"},
            {"label": "REACHED", "value": 3328, "format": "#,##0"},
            {"label": "CLOSED", "value": 1203, "format": "#,##0"},
            {"label": "VALUE", "value": 646979, "format": "$#,##0"}
        ]

    # Theme hook bindings
    primary_color = theme.get("primary", "4B286D")  # Deep Purple
    accent_color = theme.get("accent", "FFC000")    # Gold
    card_bg = "FFFFFF"
    
    primary_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    accent_fill = PatternFill(start_color=accent_color, end_color=accent_color, fill_type="solid")
    white_fill = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")

    title_font = Font(name="Calibri", size=28, color="FFFFFF", bold=True)
    subtitle_font = Font(name="Calibri", size=14, color=accent_color)
    value_font = Font(name="Calibri", size=20, color=primary_color, bold=True)
    label_font = Font(name="Calibri", size=10, color="595959", bold=True)

    align_val = Alignment(horizontal="center", vertical="bottom")
    align_lbl = Alignment(horizontal="center", vertical="top")
    align_left = Alignment(horizontal="left", vertical="center")

    match = re.match(r"([A-Z]+)([0-9]+)", anchor)
    start_col = column_index_from_string(match.group(1))
    start_row = int(match.group(2))

    # Calculate required banner width (left pad + title + gap + [cards * 3 cols] + right pad)
    total_cols = 4 + (len(kpis) * 3)

    # 1. Paint the background banner
    for r in range(start_row, start_row + 4):
        for c in range(start_col, start_col + total_cols):
            ws.cell(row=r, column=c).fill = primary_fill

    # 2. Structure Row Heights for the banner
    ws.row_dimensions[start_row].height = 15     # Top padding
    ws.row_dimensions[start_row + 1].height = 35 # Value / Title row
    ws.row_dimensions[start_row + 2].height = 20 # Label / Subtitle row
    ws.row_dimensions[start_row + 3].height = 15 # Bottom padding

    # 3. Configure Title Area
    col_idx = start_col
    ws.column_dimensions[get_column_letter(col_idx)].width = 3 # Left padding
    col_idx += 1
    ws.column_dimensions[get_column_letter(col_idx)].width = 35 # Title block width

    t_cell = ws.cell(row=start_row + 1, column=col_idx, value=title)
    t_cell.font, t_cell.fill, t_cell.alignment = title_font, primary_fill, align_left

    s_cell = ws.cell(row=start_row + 2, column=col_idx, value=subtitle)
    s_cell.font, s_cell.fill, s_cell.alignment = subtitle_font, primary_fill, align_left

    col_idx += 1
    ws.column_dimensions[get_column_letter(col_idx)].width = 3 # Gap before KPIs
    col_idx += 1

    # 4. Render KPI Cards horizontally
    for kpi in kpis:
        # Card Segment A: The Accent Bar
        ws.column_dimensions[get_column_letter(col_idx)].width = 2
        ws.merge_cells(start_row=start_row+1, start_column=col_idx, end_row=start_row+2, end_column=col_idx)
        ws.cell(row=start_row+1, column=col_idx).fill = accent_fill

        col_idx += 1
        
        # Card Segment B: The White Content Block
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

        v_cell = ws.cell(row=start_row+1, column=col_idx, value=kpi.get("value", 0))
        v_cell.font, v_cell.fill, v_cell.alignment = value_font, white_fill, align_val
        if "format" in kpi:
            v_cell.number_format = kpi["format"]

        l_cell = ws.cell(row=start_row+2, column=col_idx, value=kpi.get("label", "").upper())
        l_cell.font, l_cell.fill, l_cell.alignment = label_font, white_fill, align_lbl

        col_idx += 1
        
        # Card Segment C: The Inter-Card Gap
        ws.column_dimensions[get_column_letter(col_idx)].width = 2
        col_idx += 1
