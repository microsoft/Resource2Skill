from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference
import random

def render_workbook(wb, *, title: str = "Budget vs Actual", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Theme Configuration Fallback
    theme_colors = {
        "corporate_blue": {"primary": "003366", "secondary": "4F81BD", "accent": "FFC000", "text": "FFFFFF", "success": "00B050", "danger": "FF0000"}
    }.get(theme, {"primary": "003366", "secondary": "4F81BD", "accent": "FFC000", "text": "FFFFFF", "success": "00B050", "danger": "FF0000"})

    header_font = Font(color=theme_colors["text"], bold=True)
    header_fill = PatternFill("solid", fgColor=theme_colors["primary"])
    input_fill = PatternFill("solid", fgColor=theme_colors["accent"])
    bold_font = Font(bold=True)
    thin_border = Border(bottom=Side(style='thin', color='CCCCCC'))

    months = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
    categories = ["Base Salary", "Bonus", "Side Hustle", "Rent", "Utilities", "Groceries", "Leisure", "Transport"]

    # --- SETUP BUDGET SHEET (Matrix) ---
    ws_budget = wb.active
    ws_budget.title = "Budget"
    ws_budget.append(["Category"] + months)
    
    for cat in categories:
        # Generate some realistic-looking budget numbers
        base_val = random.randint(500, 3000)
        row = [cat] + [base_val + random.randint(-100, 100) for _ in months]
        ws_budget.append(row)
        
    for cell in ws_budget[1]:
        cell.font = header_font
        cell.fill = header_fill

    # --- SETUP ACTUALS SHEET (Transactional) ---
    ws_actuals = wb.create_sheet("Actuals")
    ws_actuals.append(["Date", "Month", "Category", "Description", "Amount"])
    
    # Populate dummy transactions for Jan-May
    for m_idx, month in enumerate(months[:5], start=1):
        for cat in categories:
            # Create 1-3 transactions per category per month
            for _ in range(random.randint(1, 3)):
                amount = random.randint(100, 2000)
                ws_actuals.append([f"2023-{m_idx:02d}-15", month, cat, "Logged Transaction", amount])

    for cell in ws_actuals[1]:
        cell.font = header_font
        cell.fill = header_fill

    # --- SETUP DASHBOARD SHEET (Presentation & Logic) ---
    ws_dash = wb.create_sheet("Dashboard")
    wb.move_sheet(ws_dash, offset=-2) # Move to front
    
    # Title & Control Cell
    ws_dash["B2"] = "Current Month ->"
    ws_dash["B2"].font = bold_font
    ws_dash["B2"].alignment = Alignment(horizontal="right")
    
    ws_dash["C2"] = "April"
    ws_dash["C2"].fill = input_fill
    ws_dash["C2"].font = bold_font
    ws_dash["C2"].border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Data Validation for Month Dropdown
    dv = DataValidation(type="list", formula1=f'"{",".join(months)}"', allow_blank=False)
    ws_dash.add_data_validation(dv)
    dv.add(ws_dash["C2"])

    # Table Headers
    headers = ["Figures in USD", "Budget", "Actual", "Var. Abs.", "Var. %"]
    for col, h in enumerate(headers, start=2):
        cell = ws_dash.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Table Rows & Formulas
    start_row = 5
    for i, cat in enumerate(categories, start=start_row):
        # Category Label
        ws_dash.cell(row=i, column=2, value=cat)
        ws_dash.cell(row=i, column=2).border = thin_border
        
        # Budget Formula: INDEX(Budget Data, MATCH(Category), MATCH(Month))
        budget_formula = f'=INDEX(Budget!$B$2:$M$20, MATCH($B{i}, Budget!$A$2:$A$20, 0), MATCH($C$2, Budget!$B$1:$M$1, 0))'
        b_cell = ws_dash.cell(row=i, column=3, value=budget_formula)
        b_cell.number_format = '#,##0'
        
        # Actuals Formula: SUMIFS(Amount, CategoryRange, CategoryLabel, MonthRange, SelectedMonth)
        actual_formula = f'=SUMIFS(Actuals!$E:$E, Actuals!$C:$C, $B{i}, Actuals!$B:$B, $C$2)'
        a_cell = ws_dash.cell(row=i, column=4, value=actual_formula)
        a_cell.number_format = '#,##0'
        
        # Variances
        v_abs = ws_dash.cell(row=i, column=5, value=f'=D{i}-C{i}')
        v_abs.number_format = '#,##0'
        
        v_pct = ws_dash.cell(row=i, column=6, value=f'=IF(C{i}<>0, (D{i}/C{i})-1, 0)')
        v_pct.number_format = '0.0%'
        
        for col in range(3, 7):
            ws_dash.cell(row=i, column=col).border = thin_border

    # Conditional Formatting for Variance % (Column F)
    cf_range = f"F{start_row}:F{start_row + len(categories) - 1}"
    green_font = Font(color=theme_colors["success"], bold=True)
    red_font = Font(color=theme_colors["danger"], bold=True)
    
    ws_dash.conditional_formatting.add(cf_range, CellIsRule(operator='greaterThan', formula=['0'], font=green_font))
    ws_dash.conditional_formatting.add(cf_range, CellIsRule(operator='lessThan', formula=['0'], font=red_font))

    # Column Widths
    ws_dash.column_dimensions['B'].width = 18
    ws_dash.column_dimensions['C'].width = 12
    ws_dash.column_dimensions['D'].width = 12
    ws_dash.column_dimensions['E'].width = 12
    ws_dash.column_dimensions['F'].width = 12

    # --- ADD CHART ---
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Budget vs Actuals"
    chart.y_axis.title = "USD"
    chart.grouping = "clustered"
    chart.overlap = -20
    
    # Data ranges
    cats_ref = Reference(ws_dash, min_col=2, min_row=start_row, max_row=start_row + len(categories) - 1)
    data_ref = Reference(ws_dash, min_col=3, max_col=4, min_row=4, max_row=start_row + len(categories) - 1)
    
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    
    ws_dash.add_chart(chart, "H4")
