from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a standard Side-Panel Dashboard Canvas shell.
    
    kwargs:
        kpis (list[dict]): List of KPI dictionaries with 'label' and 'value' keys.
    """
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False

    # 1. Resolve Theme Palette (fallback to defaults if standard tokens missing)
    themes = {
        "corporate_blue": {"panel_bg": "002060", "panel_fg": "FFFFFF", "main_bg": "F2F2F2", "main_fg": "000000"},
        "forest_green": {"panel_bg": "274E13", "panel_fg": "FFFFFF", "main_bg": "EBF1E9", "main_fg": "000000"},
        "dark_mode": {"panel_bg": "1E1E1E", "panel_fg": "E0E0E0", "main_bg": "2D2D30", "main_fg": "FFFFFF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    panel_fill = PatternFill(start_color=palette["panel_bg"], end_color=palette["panel_bg"], fill_type="solid")
    main_fill = PatternFill(start_color=palette["main_bg"], end_color=palette["main_bg"], fill_type="solid")
    
    panel_font_title = Font(color=palette["panel_fg"], size=16, bold=True)
    panel_font_kpi_val = Font(color=palette["panel_fg"], size=22, bold=True)
    panel_font_kpi_lbl = Font(color=palette["panel_fg"], size=11, bold=False)

    # 2. Paint the Dashboard Backgrounds
    # In openpyxl, filling columns directly can be buggy in some viewers, so we paint a generous viewport grid.
    for row in range(1, 51):
        for col in range(1, 22):
            cell = ws.cell(row=row, column=col)
            if col <= 3:
                cell.fill = panel_fill
            else:
                cell.fill = main_fill

    # 3. Structure the Columns (A: Padding, B: KPIs, C: Padding, D+: Canvas)
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 3
    for col_idx in range(4, 22):
        ws.column_dimensions[get_column_letter(col_idx)].width = 12

    # 4. Insert Dashboard Title into Left Panel
    title_cell = ws.cell(row=2, column=2, value=title.upper())
    title_cell.font = panel_font_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 40

    # 5. Render KPIs in the Side Panel
    kpis = kwargs.get("kpis", [
        {"label": "Total Orders", "value": "2,400"},
        {"label": "Total Quantity", "value": "11,997"},
        {"label": "Total Revenue", "value": "$649.0K"},
        {"label": "Avg. Rating", "value": "4.0"}
    ])

    current_row = 6
    for kpi in kpis:
        # Label
        lbl_cell = ws.cell(row=current_row, column=2, value=kpi["label"])
        lbl_cell.font = panel_font_kpi_lbl
        lbl_cell.alignment = Alignment(horizontal="center", vertical="bottom")
        ws.row_dimensions[current_row].height = 15
        
        # Value
        val_cell = ws.cell(row=current_row+1, column=2, value=kpi["value"])
        val_cell.font = panel_font_kpi_val
        val_cell.alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[current_row+1].height = 30
        
        current_row += 3  # Spacing between KPIs
