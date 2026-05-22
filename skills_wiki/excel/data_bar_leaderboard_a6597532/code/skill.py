from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter, coordinate_from_string, column_index_from_string

def render(ws, anchor: str, *, title: str = "Sales Agent KPIs", data: list[dict] = None, columns: list[dict] = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a leaderboard table with inline conditional formatting data bars.
    """
    # Provide realistic default data based on the video if none is supplied
    if data is None:
        data = [
            {"name": "Charlie", "calls": 610, "reached": 86, "closed": 67, "value": 45236.03},
            {"name": "Eve", "calls": 722, "reached": 168, "closed": 70, "value": 44841.36},
            {"name": "Bob", "calls": 661, "reached": 73, "closed": 28, "value": 40092.43},
            {"name": "Alice", "calls": 1031, "reached": 56, "closed": 37, "value": 13519.04},
            {"name": "David", "calls": 375, "reached": 120, "closed": 48, "value": 2590.45},
        ]
        
    if columns is None:
        columns = [
            {"header": "Name", "key": "name", "width": 15},
            {"header": "Total Calls", "key": "calls", "width": 12, "bar_color": "FF5B9BD5", "format": "#,##0"},
            {"header": "Calls Reached", "key": "reached", "width": 14, "bar_color": "FFFFC000", "format": "#,##0"},
            {"header": "Deals Closed", "key": "closed", "width": 14, "bar_color": "FFB4A7D6", "format": "#,##0"},
            {"header": "Deal Value ($)", "key": "value", "width": 16, "bar_color": "FF7030A0", "format": "$#,##0"}
        ]

    col_str, row_str = coordinate_from_string(anchor)
    start_col = column_index_from_string(col_str)
    start_row = int(row_str)
    
    # Theme palette fallback (mocking the Aspect theme from the video)
    palette = {
        "primary": "FF5C4084", # Dark purple
        "text_light": "FFFFFFFF",
        "text_dark": "FF000000"
    }
    
    # 1. Render Title
    title_cell = ws.cell(row=start_row, column=start_col, value=title)
    title_cell.font = Font(size=14, bold=True, color=palette["primary"][2:])
    start_row += 2 # Leave a blank row before the table
    
    # 2. Render Headers
    header_font = Font(bold=True, color=palette["text_light"][2:])
    header_fill = PatternFill("solid", fgColor=palette["primary"][2:])
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for c_idx, col_def in enumerate(columns):
        col_letter = get_column_letter(start_col + c_idx)
        ws.column_dimensions[col_letter].width = col_def.get("width", 15)
        
        cell = ws.cell(row=start_row, column=start_col + c_idx, value=col_def["header"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        
    start_row += 1
    
    # 3. Render Data
    for r_idx, row_data in enumerate(data):
        current_row = start_row + r_idx
        for c_idx, col_def in enumerate(columns):
            val = row_data.get(col_def["key"], "")
            cell = ws.cell(row=current_row, column=start_col + c_idx, value=val)
            
            if "format" in col_def:
                cell.number_format = col_def["format"]
            
            # Align numbers right, text left
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")
                
    end_row = start_row + len(data) - 1
    
    # 4. Apply Conditional Formatting (Data Bars)
    for c_idx, col_def in enumerate(columns):
        bar_color = col_def.get("bar_color")
        if bar_color:
            col_letter = get_column_letter(start_col + c_idx)
            range_str = f"{col_letter}{start_row}:{col_letter}{end_row}"
            
            # DataBarRule uses hex colors without the '#' prefix
            rule = DataBarRule(
                start_type='min', 
                end_type='max', 
                color=bar_color, 
                showValue=True
            )
            ws.conditional_formatting.add(range_str, rule)
