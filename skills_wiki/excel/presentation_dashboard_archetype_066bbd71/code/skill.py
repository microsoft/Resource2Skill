from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders an interactive-style presentation dashboard with multiple aligned charts 
    based on aggregated data stored in a hidden background sheet.
    """
    # Clean up default sheets
    for sheet_name in wb.sheetnames:
        del wb[sheet_name]
        
    # Standardize a basic theme palette fallback
    theme_colors = {
        "corporate_blue": {"primary": "003366", "text": "FFFFFF"},
        "dark_mode": {"primary": "333333", "text": "FFFFFF"},
        "emerald": {"primary": "059669", "text": "FFFFFF"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # --- 1. Create Hidden ChartData Sheet ---
    ws_data = wb.create_sheet("ChartData")
    
    # Dummy Aggregated Data: Profit by Market & Product
    summary_profit = [
        ["Country", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle"],
        ["India", 62000, 4800, 21000, 25000],
        ["Philippines", 54000, 7000, 22000, 8000],
        ["United Kingdom", 46000, 5200, 11000, 14000],
        ["United States", 36000, 6300, 22000, 9000]
    ]
    for row in summary_profit:
        ws_data.append(row)
        
    ws_data.append([]) # Spacer row
    trend_start_row = ws_data.max_row + 1
    
    # Dummy Aggregated Data: Trend over months
    summary_trend = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50000, 124000],
        ["Oct", 95000, 228000],
        ["Nov", 65000, 160000],
        ["Dec", 52000, 136000]
    ]
    for row in summary_trend:
        ws_data.append(row)
        
    # Hide the data sheet to preserve the "dashboard" illusion
    ws_data.sheet_state = 'hidden'

    # --- 2. Create Dashboard Sheet ---
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False
    ws_dash.sheet_view.showRowColHeaders = False
    
    # Format Dashboard Header
    ws_dash.merge_cells("B2:O4")
    header_cell = ws_dash["B2"]
    header_cell.value = title
    header_cell.font = Font(size=28, bold=True, color=palette["text"])
    header_cell.fill = PatternFill("solid", fgColor=palette["primary"])
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # --- 3. Build Chart 1: Stacked Bar (Profit by Market) ---
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Product"
    chart1.height = 13.5
    chart1.width = 16
    
    data1 = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    
    ws_dash.add_chart(chart1, "B6")
    
    # --- 4. Build Chart 2: Line (Units Sold) ---
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.height = 6.5
    chart2.width = 13
    
    data2 = Reference(ws_data, min_col=2, min_row=trend_start_row, max_row=trend_start_row+4)
    cats2 = Reference(ws_data, min_col=1, min_row=trend_start_row+1, max_row=trend_start_row+4)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    
    ws_dash.add_chart(chart2, "K6")
    
    # --- 5. Build Chart 3: Line (Profit) ---
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.height = 6.5
    chart3.width = 13
    
    data3 = Reference(ws_data, min_col=3, min_row=trend_start_row, max_row=trend_start_row+4)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats2) # Reuses the same month categories
    
    ws_dash.add_chart(chart3, "K20")
    
    # Set the dashboard as the default active sheet
    wb.active = ws_dash
