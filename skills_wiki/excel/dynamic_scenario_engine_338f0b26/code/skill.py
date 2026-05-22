from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

def render(ws, anchor: str, *, assumptions: list[str] = None, scenarios: list[str] = None, theme: str = "corporate_blue", **kwargs) -> None:
    if assumptions is None:
        assumptions = ["Volume Growth", "Price Increase", "Fixed Costs"]
    if scenarios is None:
        scenarios = ["Base Case", "Optimistic", "Pessimistic"]

    # Basic embedded palette logic mimicking standard helpers
    palettes = {
        "corporate_blue": {"header_bg": "203764", "header_fg": "FFFFFF", "input_bg": "F2F2F2", "live_bg": "D9E1F2", "accent": "4472C4"},
        "emerald": {"header_bg": "004D40", "header_fg": "FFFFFF", "input_bg": "F5F5F5", "live_bg": "B2DFDB", "accent": "00897B"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    header_fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_font = Font(color=palette["header_fg"], bold=True)
    input_fill = PatternFill(start_color=palette["input_bg"], end_color=palette["input_bg"], fill_type="solid")
    live_fill = PatternFill(start_color=palette["live_bg"], end_color=palette["live_bg"], fill_type="solid")
    accent_font = Font(color=palette["accent"], bold=True)
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")

    col_letter, row_str = coordinate_from_string(anchor)
    start_col = column_index_from_string(col_letter)
    start_row = int(row_str)

    current_row = start_row

    # 1. Setup Toggle
    lbl = ws.cell(row=current_row, column=start_col, value="Active Scenario:")
    lbl.font = bold_font
    lbl.alignment = Alignment(horizontal="right")

    toggle_cell = ws.cell(row=current_row, column=start_col + 1, value=1)
    toggle_cell.font = accent_font
    toggle_cell.alignment = center_align
    toggle_cell.fill = input_fill

    border_side = Side(style="thin", color="A6A6A6")
    toggle_cell.border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    toggle_abs = f"${get_column_letter(start_col + 1)}${current_row}"

    # Data Validation for Toggle (Indices: 1, 2, 3...)
    indices_str = ",".join(str(i+1) for i in range(len(scenarios)))
    dv = DataValidation(type="list", formula1=f'"{indices_str}"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(toggle_cell)

    current_row += 2

    # 2. Setup Live Case Header
    live_header = ws.cell(row=current_row, column=start_col, value="Live Assumptions (Powered by CHOOSE)")
    live_header.font = header_font
    live_header.fill = header_fill
    ws.cell(row=current_row, column=start_col+1).fill = header_fill
    current_row += 1

    live_value_rows = []
    for assump in assumptions:
        ws.cell(row=current_row, column=start_col, value=assump)
        live_value_rows.append(current_row)
        current_row += 1

    current_row += 1

    # 3. Setup Scenario Blocks
    scenario_value_coords = {i: [] for i in range(len(scenarios))}

    for s_idx, scenario_name in enumerate(scenarios):
        s_header = ws.cell(row=current_row, column=start_col, value=f"{scenario_name} (Scenario {s_idx + 1})")
        s_header.font = bold_font
        s_header.fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
        ws.cell(row=current_row, column=start_col+1).fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
        current_row += 1

        for a_idx, assump in enumerate(assumptions):
            ws.cell(row=current_row, column=start_col, value=assump)
            # Generate realistic varying placeholder values
            base_val = 100 * (a_idx + 1)
            modifier = 1.0 + (s_idx * 0.2)
            val_cell = ws.cell(row=current_row, column=start_col + 1, value=base_val * modifier)
            val_cell.fill = input_fill
            scenario_value_coords[s_idx].append(val_cell.coordinate)
            current_row += 1

        current_row += 1

    # 4. Wire up Live Case Formulas
    for a_idx, r in enumerate(live_value_rows):
        # Collect references across all scenarios for this specific assumption
        refs = [scenario_value_coords[s_idx][a_idx] for s_idx in range(len(scenarios))]
        refs_str = ", ".join(refs)
        formula = f"=CHOOSE({toggle_abs}, {refs_str})"
        
        live_cell = ws.cell(row=r, column=start_col + 1, value=formula)
        live_cell.fill = live_fill
        live_cell.font = bold_font

    # Formatting adjustments
    ws.column_dimensions[get_column_letter(start_col)].width = 35
    ws.column_dimensions[get_column_letter(start_col + 1)].width = 15
