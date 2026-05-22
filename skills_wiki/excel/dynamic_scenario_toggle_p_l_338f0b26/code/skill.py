from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
import openpyxl

def render_sheet(wb: openpyxl.Workbook, sheet_name: str, *, title: str = "Dynamic P&L Scenarios", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Standard minimal theme palette fallback
    theme_colors = {
        "corporate_blue": {"header_bg": "203764", "header_fg": "FFFFFF", "accent": "8EA9DB"},
        "green": {"header_bg": "385723", "header_fg": "FFFFFF", "accent": "A9D08E"}
    }
    tc = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    header_fill = PatternFill(start_color=tc["header_bg"], end_color=tc["header_bg"], fill_type="solid")
    header_font = Font(color=tc["header_fg"], bold=True)
    input_font = Font(color="0000FF") # Financial modeling standard for hardcoded inputs
    bold_font = Font(bold=True)
    
    # 1. Setup Grid & Column Widths
    ws.column_dimensions['A'].width = 28
    for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
        ws.column_dimensions[col].width = 14

    ws['A1'] = title
    ws['A1'].font = Font(size=14, bold=True, color=tc["header_bg"])

    # 2. Scenario Dropdown Toggle
    ws['G2'] = "Scenario (1-3):"
    ws['G2'].font = bold_font
    ws['G2'].alignment = Alignment(horizontal="right")
    
    ws['H2'] = 1
    ws['H2'].fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    ws['H2'].border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    ws['H2'].alignment = Alignment(horizontal="center")
    
    # Add Data Validation for the toggle
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws['H2'])

    # 3. Structural Headers
    sections = [
        (4, "Income Statement"),
        (11, "Live Assumptions (Driven by Scenario)"),
        (17, "Scenario 1: Base Case"),
        (23, "Scenario 2: Downside"),
        (29, "Scenario 3: Upside")
    ]
    for r, label in sections:
        ws.cell(row=r, column=1, value=label)
        for i, col_char in enumerate(['A', 'B', 'C', 'D', 'E', 'F']):
            cell = ws.cell(row=r, column=i+1)
            cell.fill = header_fill
            cell.font = header_font
            if i > 0:
                cell.value = f"Year {i}"
                cell.alignment = Alignment(horizontal="right")

    # 4. Labels
    assumption_labels = ["Volume (Units)", "Price per Unit", "Cost per Unit", "Fixed Expenses"]
    for idx, label in enumerate(assumption_labels):
        ws.cell(row=12+idx, column=1, value=label)
        ws.cell(row=18+idx, column=1, value=label)
        ws.cell(row=24+idx, column=1, value=label)
        ws.cell(row=30+idx, column=1, value=label)
        
    pl_labels = ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "Net Income"]
    for idx, label in enumerate(pl_labels):
        ws.cell(row=5+idx, column=1, value=label)
        if label in ["Gross Profit", "Net Income"]:
            ws.cell(row=5+idx, column=1).font = bold_font

    # 5. Populate Hardcoded Scenarios
    scenarios = [
        # Base
        {"vol": [1000, 1100, 1210, 1331, 1464], "price": [50]*5, "cost": [20]*5, "fixed": [10000]*5, "start": 18},
        # Downside
        {"vol": [1000, 950, 900, 850, 800], "price": [45]*5, "cost": [22]*5, "fixed": [12000]*5, "start": 24},
        # Upside
        {"vol": [1000, 1200, 1440, 1728, 2073], "price": [55]*5, "cost": [18]*5, "fixed": [10000]*5, "start": 30}
    ]

    for scen in scenarios:
        r = scen["start"]
        for i, col in enumerate(['B', 'C', 'D', 'E', 'F']):
            # Volume
            ws[f'{col}{r}'].value = scen["vol"][i]
            ws[f'{col}{r}'].font = input_font
            ws[f'{col}{r}'].number_format = '#,##0'
            # Price
            ws[f'{col}{r+1}'].value = scen["price"][i]
            ws[f'{col}{r+1}'].font = input_font
            ws[f'{col}{r+1}'].number_format = '"$"#,##0.00'
            # Cost
            ws[f'{col}{r+2}'].value = scen["cost"][i]
            ws[f'{col}{r+2}'].font = input_font
            ws[f'{col}{r+2}'].number_format = '"$"#,##0.00'
            # Fixed
            ws[f'{col}{r+3}'].value = scen["fixed"][i]
            ws[f'{col}{r+3}'].font = input_font
            ws[f'{col}{r+3}'].number_format = '"$"#,##0'

    # 6. Populate Dynamic Live Formulas (The Core Mechanism)
    for col in ['B', 'C', 'D', 'E', 'F']:
        ws[f'{col}12'].value = f'=CHOOSE($H$2, {col}18, {col}24, {col}30)'
        ws[f'{col}12'].number_format = '#,##0'
        
        ws[f'{col}13'].value = f'=CHOOSE($H$2, {col}19, {col}25, {col}31)'
        ws[f'{col}13'].number_format = '"$"#,##0.00'
        
        ws[f'{col}14'].value = f'=CHOOSE($H$2, {col}20, {col}26, {col}32)'
        ws[f'{col}14'].number_format = '"$"#,##0.00'
        
        ws[f'{col}15'].value = f'=CHOOSE($H$2, {col}21, {col}27, {col}33)'
        ws[f'{col}15'].number_format = '"$"#,##0'

    # 7. Populate P&L Formulas based on Live Assumptions
    for col in ['B', 'C', 'D', 'E', 'F']:
        ws[f'{col}5'].value = f'={col}12*{col}13'  # Revenue
        ws[f'{col}6'].value = f'={col}12*{col}14'  # COGS
        ws[f'{col}7'].value = f'={col}5-{col}6'    # Gross Profit
        ws[f'{col}8'].value = f'={col}15'          # OpEx
        ws[f'{col}9'].value = f'={col}7-{col}8'    # Net Income
        
        for r in [5, 6, 7, 8, 9]:
            ws[f'{col}{r}'].number_format = '"$"#,##0'
            
        # P&L Styling Overlays
        ws[f'{col}7'].font = bold_font
        ws[f'{col}9'].font = bold_font
        ws[f'{col}7'].border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
        ws[f'{col}9'].border = Border(top=Side(style='thin'), bottom=Side(style='double'))
