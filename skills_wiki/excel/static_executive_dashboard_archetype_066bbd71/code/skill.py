from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # Attempt to load theme palette, fallback to default corporate colors
    try:
        from _helpers import get_theme
        palette = get_theme(theme)
    except ImportError:
        palette = {"primary": "203764", "text_on_primary": "FFFFFF"}
        
    primary_color = palette.get("primary", "203764").replace("#", "")
    text_color = palette.get("text_on_primary", "FFFFFF").replace("#", "")

    # 1. Setup Sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    ws_data = wb.create_sheet(title="ChartData")
    ws_data.sheet_state = 'hidden'
    
    # 2. Build Dashboard Banner
    ws_dash.merge_cells("A1:R3")
    banner_cell = ws_dash["A1"]
    banner_cell.value = title
    banner_cell.fill = PatternFill("solid", fgColor=primary_color)
    banner_cell.font = Font(name="Calibri", size=24, bold=True, color=text_color)
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add some breathing room below the banner
    ws_dash.row_dimensions[4].height = 15
    
    # 3. Populate Hidden Chart Data
    # Dataset 1: Profit by Market & Product (Simulated Pivot Table)
    data1 = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle", "Sugar"],
        ["India", 62349, 4872, 21028, 25085, 18561],
        ["Philippines", 54618, 7026, 22005, 8313, 14947],
        ["United Kingdom", 46530, 5220, 11497, 14620, 19446],
        ["Malaysia", 46587, 5538, 17536, 20555, 10633],
        ["United States", 36657, 6369, 22260, 9938, 9186]
    ]
    
    for r_idx, row in enumerate(data1, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
            if r_idx > 1 and c_idx > 1:
                cell.number_format = '"$"#,##0'
                
    # Dataset 2: Trend Data (Units Sold and Profit by Month)
    trend_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124813],
        ["Oct", 95622, 228276],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136338]
    ]
    
    r_offset = len(data1) + 3
    for r_idx, row in enumerate(trend_data, 0):
        for c_idx, val in enumerate(row, 1):
            cell = ws_data.cell(row=r_offset + r_idx, column=c_idx, value=val)
            if r_idx > 0:
                if c_idx == 2:
                    cell.number_format = '#,##0'
                elif c_idx == 3:
                    cell.number_format = '"$"#,##0'

    # 4. Create and Anchor Charts
    
    # Chart 1: Stacked Column (Profit by Market & Cookie)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    chart1.height = 12
    chart1.width = 16
    
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=len(data1))
    data_ref1 = Reference(ws_data, min_col=2, max_col=len(data1[0]), min_row=1, max_row=len(data1))
    chart1.add_data(data_ref1, titles_from_data=True)
    chart1.set_categories(cats1)
    
    ws_dash.add_chart(chart1, "B5")
    
    # Chart 2: Smoothed Line (Units Sold)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.height = 6
    chart2.width = 14
    chart2.legend = None
    
    cats2 = Reference(ws_data, min_col=1, min_row=r_offset+1, max_row=r_offset+len(trend_data)-1)
    data_ref2 = Reference(ws_data, min_col=2, min_row=r_offset, max_row=r_offset+len(trend_data)-1)
    chart2.add_data(data_ref2, titles_from_data=True)
    chart2.set_categories(cats2)
    for s in chart2.series:
        s.smooth = True
        
    ws_dash.add_chart(chart2, "K5")
    
    # Chart 3: Smoothed Line (Profit by Month)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.height = 6
    chart3.width = 14
    chart3.legend = None
    
    data_ref3 = Reference(ws_data, min_col=3, min_row=r_offset, max_row=r_offset+len(trend_data)-1)
    chart3.add_data(data_ref3, titles_from_data=True)
    chart3.set_categories(cats2) # Reuses X-axis categories from the previous chart
    for s in chart3.series:
        s.smooth = True
        
    ws_dash.add_chart(chart3, "K17")
