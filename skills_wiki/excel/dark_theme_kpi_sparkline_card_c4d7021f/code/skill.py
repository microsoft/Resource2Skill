from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker
from openpyxl.utils import coordinate_to_tuple

def render(
    ws, 
    anchor: str, 
    *, 
    title: str, 
    current_value: float, 
    value_format: str, 
    percent_change: float, 
    data_sheet_name: str, 
    data_min_col: int, 
    data_min_row: int, 
    data_max_row: int, 
    theme: str = "dark_dashboard", 
    **kwargs
) -> None:
    # Theme palette fallback
    palettes = {
        "dark_dashboard": {
            "bg": "1A1A1D",
            "text_main": "FFFFFF",
            "text_sub": "A0A0A0",
            "pos": "2ECC71",
            "neg": "E74C3C",
            "chart_line": "3498DB"
        }
    }
    colors = palettes.get(theme, palettes["dark_dashboard"])
    
    r, c = coordinate_to_tuple(anchor)
    
    # Paint background block (8 rows by 5 columns)
    bg_fill = PatternFill(start_color=colors["bg"], end_color=colors["bg"], fill_type="solid")
    for row_idx in range(r, r + 8):
        for col_idx in range(c, c + 5):
            ws.cell(row=row_idx, column=col_idx).fill = bg_fill
            
    # Title
    title_cell = ws.cell(row=r, column=c, value=title)
    title_cell.font = Font(color=colors["text_sub"], size=10, bold=True)
    
    # Current Value
    val_cell = ws.cell(row=r + 1, column=c, value=current_value)
    val_cell.font = Font(color=colors["text_main"], size=18, bold=True)
    val_cell.number_format = value_format
    
    # Percent Change
    change_cell = ws.cell(row=r + 1, column=c + 1, value=percent_change)
    change_color = colors["pos"] if percent_change >= 0 else colors["neg"]
    change_cell.font = Font(color=change_color, size=12, bold=True)
    change_cell.number_format = "+0.0%;-0.0%"
    change_cell.alignment = Alignment(vertical="bottom")
    
    # Create Sparkline-style Line Chart
    chart = LineChart()
    data = Reference(ws.parent[data_sheet_name], min_col=data_min_col, min_row=data_min_row, max_row=data_max_row)
    chart.add_data(data)
    
    # Clean up chart elements (strip away axes and legend)
    chart.legend = None
    chart.x_axis.tickLblPos = "none"
    chart.y_axis.tickLblPos = "none"
    chart.x_axis.majorTickMark = "none"
    chart.y_axis.majorTickMark = "none"
    chart.x_axis.minorTickMark = "none"
    chart.y_axis.minorTickMark = "none"
    chart.x_axis.majorGridlines = None
    chart.y_axis.majorGridlines = None
    
    # Hide axis lines completely
    no_fill_line = LineProperties(noFill=True)
    chart.x_axis.spPr = GraphicalProperties(ln=no_fill_line)
    chart.y_axis.spPr = GraphicalProperties(ln=no_fill_line)
    
    # Make chart area and plot area transparent so the painted cells show through
    no_fill_props = GraphicalProperties(noFill=True, ln=no_fill_line)
    chart.graphical_properties = no_fill_props
    chart.plot_area.graphicalProperties = no_fill_props
    
    # Style the data series line and markers
    if chart.series:
        s1 = chart.series[0]
        s1.graphicalProperties.line.solidFill = colors["chart_line"]
        s1.graphicalProperties.line.width = 25000  # EMUs (approx 2pt)
        
        s1.marker = Marker(symbol="circle", size=4)
        s1.marker.graphicalProperties.solidFill = colors["chart_line"]
        s1.marker.graphicalProperties.line.noFill = True
        
    # Set physical dimensions (in cm) and inject into sheet
    chart.width = 12.0
    chart.height = 4.5
    ws.add_chart(chart, ws.cell(row=r + 2, column=c).coordinate)
