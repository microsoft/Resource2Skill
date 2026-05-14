import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Theme
    theme_colors = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF"},
        "dark_slate": {"header_bg": "2F4F4F", "header_fg": "FFFFFF"},
        "emerald": {"header_bg": "2E8B57", "header_fg": "FFFFFF"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # 2. Setup Data Sheet
    ws_data = wb.active
    ws_data.title = "Chart Data"
    
    # Categorical Data for Stacked Column Chart
    market_data = [
        ["Country", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle"],
        ["India", 62000, 23000, 21000, 25000],
        ["Philippines", 54000, 24000, 22000, 8000],
        ["United Kingdom", 46000, 26000, 11000, 14000],
        ["United States", 36000, 32000, 9000, 9000],
    ]
    for row in market_data:
        ws_data.append(row)
        
    ws_data.append([]) # Gap
    
    # Time-series Data for Line Charts
    monthly_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    month_start_row = ws_data.max_row + 1
    for row in monthly_data:
        ws_data.append(row)
        
    # Apply number formatting to data so charts automatically inherit it
    for r in range(2, 6):
        for c in range(2, 6):
            ws_data.cell(row=r, column=c).number_format = '"$"#,##0'
            
    for r in range(month_start_row + 1, month_start_row + 5):
        ws_data.cell(row=r, column=2).number_format = '#,##0'
        ws_data.cell(row=r, column=3).number_format = '"$"#,##0'

    # 3. Setup Dashboard Sheet
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False
    
    # 4. Build Header Banner
    ws_dash.merge_cells("A1:R3")
    title_cell = ws_dash["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["header_fg"])
    title_cell.fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 5. Build Chart 1: Stacked Column (Profit by Market)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    chart1.style = 11
    
    data1 = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=5)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    chart1.height = 12
    chart1.width = 16
    ws_dash.add_chart(chart1, "B5")
    
    # 6. Build Chart 2: Line (Units Sold)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.style = 13
    chart2.legend = None # Clean up unnecessary legend
    
    data2 = Reference(ws_data, min_col=2, min_row=month_start_row, max_col=2, max_row=month_start_row+4)
    cats2 = Reference(ws_data, min_col=1, min_row=month_start_row+1, max_row=month_start_row+4)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.height = 8
    chart2.width = 14
    ws_dash.add_chart(chart2, "K5")
    
    # 7. Build Chart 3: Line (Profit)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.style = 13
    chart3.legend = None
    
    data3 = Reference(ws_data, min_col=3, min_row=month_start_row, max_col=3, max_row=month_start_row+4)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats2)
    chart3.height = 8
    chart3.width = 14
    ws_dash.add_chart(chart3, "K17")
    
    # 8. Hide the data sheet to keep the workbook strictly presentation-focused
    ws_data.sheet_state = 'hidden'
    wb.active = ws_dash
