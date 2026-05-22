def render(ws, data_anchor: str, chart_anchor: str, kpi_name: str, actual: float, target: float, *, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.utils import coordinate_to_tuple
    from openpyxl.chart import DoughnutChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.drawing.line import LineProperties

    # 1. Write Data Backend
    row, col = coordinate_to_tuple(data_anchor)
    ws.cell(row=row, column=col, value=kpi_name)
    ws.cell(row=row, column=col+1, value="Value")
    
    ws.cell(row=row+1, column=col, value="Actual")
    ws.cell(row=row+1, column=col+1, value=actual)
    
    remainder = max(0, target - actual)
    ws.cell(row=row+2, column=col, value="Remainder")
    ws.cell(row=row+2, column=col+1, value=remainder)
    
    # 2. Create Chart
    chart = DoughnutChart()
    chart.title = None
    chart.legend = None
    chart.holeSize = 65  # Expand hole for a modern KPI ring look
    
    # Compact square size for KPI widget
    chart.width = 8
    chart.height = 8
    
    # Map data references
    data = Reference(ws, min_col=col+1, min_row=row+1, max_row=row+2)
    cats = Reference(ws, min_col=col, min_row=row+1, max_row=row+2)
    
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    
    # 3. Format Slices & Chart Area
    palettes = {
        "corporate_blue": {"primary": "1F4E78", "muted": "D9D9D9", "border": "FFFFFF"},
        "emerald_green": {"primary": "385723", "muted": "D9D9D9", "border": "FFFFFF"},
        "crimson_red": {"primary": "C00000", "muted": "D9D9D9", "border": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    # Remove chart area outline for clean embedding
    chart.graphical_properties.line = LineProperties(noFill=True)
    
    if chart.series:
        s = chart.series[0]
        
        # Format "Actual" progress slice
        pt_actual = DataPoint(idx=0)
        pt_actual.graphicalProperties.solidFill = palette["primary"]
        pt_actual.graphicalProperties.line = LineProperties(solidFill=palette["border"])
        
        # Format "Remainder" track slice
        pt_remainder = DataPoint(idx=1)
        pt_remainder.graphicalProperties.solidFill = palette["muted"]
        pt_remainder.graphicalProperties.line = LineProperties(solidFill=palette["border"])
        
        s.dPt.append(pt_actual)
        s.dPt.append(pt_remainder)
        
    # 4. Position Chart
    ws.add_chart(chart, chart_anchor)
