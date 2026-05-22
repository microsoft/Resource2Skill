from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a transparent, minimalist KPI Doughnut Chart component for dashboards.
    """
    # 1. Setup Data
    # For a component, we typically write the source data to a hidden or off-screen range
    data_col = kwargs.get("data_col", 50)  # Default to column AX
    data_row = kwargs.get("data_row", 1)
    
    metric_name = kwargs.get("metric_name", "Customer Satisfaction")
    actual_val = kwargs.get("actual_val", 0.87)
    remain_val = 1.0 - actual_val

    # Write labels
    ws.cell(row=data_row, column=data_col, value=metric_name)
    ws.cell(row=data_row+1, column=data_col, value="Actual")
    ws.cell(row=data_row+2, column=data_col, value="Remaining")

    # Write values and format as percentages
    c_actual = ws.cell(row=data_row+1, column=data_col+1, value=actual_val)
    c_actual.number_format = "0%"
    c_remain = ws.cell(row=data_row+2, column=data_col+1, value=remain_val)
    c_remain.number_format = "0%"

    # 2. Create Chart
    chart = DoughnutChart()
    chart.width = 5.5
    chart.height = 5.5
    chart.holeSize = 65  # Expand the hole for a thin, modern ring

    # 3. Clean UI (No Legend, No Fill, No Border)
    chart.legend = None
    # Remove the white background and border from the chart container
    chart.graphical_properties = GraphicalProperties(noFill=True)
    chart.graphical_properties.line = LineProperties(noFill=True)

    # 4. Attach Data
    data = Reference(ws, min_col=data_col+1, min_row=data_row+1, max_row=data_row+2)
    cats = Reference(ws, min_col=data_col, min_row=data_row+1, max_row=data_row+2)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    # 5. Theme the Slices
    theme_colors = {
        "corporate_blue": {"primary": "2F5597", "bg_accent": "D9E1F2"},
        "modern_dark": {"primary": "4472C4", "bg_accent": "262626"},
        "emerald": {"primary": "385723", "bg_accent": "E2EFDA"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    series = chart.series[0]
    
    # Remove borders between the slices for a seamless ring
    series.graphicalProperties.line = LineProperties(noFill=True)

    # Explicitly color Slice 0 (Actual)
    dp_actual = DataPoint(idx=0)
    dp_actual.graphicalProperties.solidFill = palette["primary"]

    # Explicitly color Slice 1 (Remaining)
    dp_remain = DataPoint(idx=1)
    dp_remain.graphicalProperties.solidFill = palette["bg_accent"]

    series.dPt.append(dp_actual)
    series.dPt.append(dp_remain)

    # 6. Place Chart
    # (In a full Excel dashboard, a text box linked to the 'Actual' cell is often manually overlaid in the center)
    ws.add_chart(chart, anchor)
