from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Fallback dataset if none provided
    data = kwargs.get("data", [
        {"Name": "Alex", "Total Calls": 1031, "Calls Reached": 56, "Deals Closed": 27, "Deal Value": 13519},
        {"Name": "Alice", "Total Calls": 827, "Calls Reached": 128, "Deals Closed": 49, "Deal Value": 41200},
        {"Name": "Bob", "Total Calls": 661, "Calls Reached": 73, "Deals Closed": 28, "Deal Value": 40092},
        {"Name": "Charlie", "Total Calls": 610, "Calls Reached": 86, "Deals Closed": 67, "Deal Value": 45236},
        {"Name": "Diana", "Total Calls": 566, "Calls Reached": 163, "Deals Closed": 26, "Deal Value": 38093},
    ])

    # 1. Define Dashboard Canvas Colors
    dark_bg = "321E4B"    # Dark purple header
    pale_bg = "F2EFF5"    # Pale purple body
    accent = "FFC000"     # Gold/Yellow accent
    text_light = "FFFFFF"

    dark_fill = PatternFill(start_color=dark_bg, end_color=dark_bg, fill_type="solid")
    pale_fill = PatternFill(start_color=pale_bg, end_color=pale_bg, fill_type="solid")

    # Paint Canvas
    for row in range(1, 9):
        for col in range(1, 20):
            ws.cell(row=row, column=col).fill = dark_fill
    for row in range(9, 41):
        for col in range(1, 20):
            ws.cell(row=row, column=col).fill = pale_fill

    # 2. Add Titles
    ws.row_dimensions[2].height = 40
    title_cell = ws.cell(row=2, column=2, value=title)
    title_cell.font = Font(size=32, color=text_light, bold=True)
    
    subtitle_cell = ws.cell(row=3, column=2, value="Evaluating Sales Agent Performance")
    subtitle_cell.font = Font(size=14, color=accent)

    # 3. Generate KPI Card Blocks (Simulating Shapes)
    kpis = [
        {"label": "CALLS", "value": sum(d.get("Total Calls", 0) for d in data)},
        {"label": "REACHED", "value": sum(d.get("Calls Reached", 0) for d in data)},
        {"label": "CLOSED", "value": sum(d.get("Deals Closed", 0) for d in data)},
        {"label": "VALUE", "value": f'${sum(d.get("Deal Value", 0) for d in data):,}'}
    ]
    
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    accent_fill = PatternFill(start_color=accent, end_color=accent, fill_type="solid")
    thin_border = Border(left=Side(style='thin', color='CCCCCC'),
                         right=Side(style='thin', color='CCCCCC'),
                         top=Side(style='thin', color='CCCCCC'),
                         bottom=Side(style='thin', color='CCCCCC'))

    for i, kpi in enumerate(kpis):
        col_offset = 4 + (i * 3) # Spacing cards across columns D, G, J, M
        
        # Color accent strip on the left side of the card
        for r in range(4, 8):
            ws.cell(row=r, column=col_offset).fill = accent_fill
            ws.cell(row=r, column=col_offset).border = thin_border
            
        # Main white body of the card
        for r in range(4, 8):
            cell = ws.cell(row=r, column=col_offset+1)
            cell.fill = white_fill
            cell.border = thin_border

        # Populate Values and Labels
        ws.merge_cells(start_row=4, start_column=col_offset+1, end_row=5, end_column=col_offset+1)
        val_cell = ws.cell(row=4, column=col_offset+1, value=kpi["value"])
        val_cell.font = Font(size=18, bold=True, color=dark_bg)
        val_cell.alignment = Alignment(horizontal="center", vertical="bottom")

        ws.merge_cells(start_row=6, start_column=col_offset+1, end_row=7, end_column=col_offset+1)
        lbl_cell = ws.cell(row=6, column=col_offset+1, value=kpi["label"])
        lbl_cell.font = Font(size=11, color="555555", bold=True)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")

        # Refine widths to shape the card
        ws.column_dimensions[get_column_letter(col_offset)].width = 3
        ws.column_dimensions[get_column_letter(col_offset+1)].width = 16

    # 4. Insert Performance Leaderboard Table
    ws.row_dimensions[9].height = 15 # Spacer
    headers = ["Name", "Total Calls", "Calls Reached", "Deals Closed", "Deal Value"]
    
    for c, h in enumerate(headers, start=2):
        cell = ws.cell(row=10, column=c, value=h)
        cell.font = Font(bold=True, color=text_light)
        cell.fill = dark_fill

    for r_idx, row_data in enumerate(data, start=11):
        ws.cell(row=r_idx, column=2, value=row_data.get("Name", ""))
        ws.cell(row=r_idx, column=3, value=row_data.get("Total Calls", 0))
        ws.cell(row=r_idx, column=4, value=row_data.get("Calls Reached", 0))
        ws.cell(row=r_idx, column=5, value=row_data.get("Deals Closed", 0))
        
        dv_cell = ws.cell(row=r_idx, column=6, value=row_data.get("Deal Value", 0))
        dv_cell.number_format = '"$"#,##0'

    end_row = 10 + len(data)
    
    # 5. In-cell Condtional Formatting Data Bars
    # Calls Reached (Yellow)
    rule_reached = DataBarRule(start_type='min', end_type='max', color="FFFFC000", showValue=True)
    ws.conditional_formatting.add(f"D11:D{end_row}", rule_reached)

    # Deals Closed (Light Purple)
    rule_closed = DataBarRule(start_type='min', end_type='max', color="FFB4A7D6", showValue=True)
    ws.conditional_formatting.add(f"E11:E{end_row}", rule_closed)

    # Deal Value (Dark Purple)
    rule_dv = DataBarRule(start_type='min', end_type='max', color="FF674EA7", showValue=True)
    ws.conditional_formatting.add(f"F11:F{end_row}", rule_dv)

    # Adjust table columns
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 16

    # 6. Add Clean Companion Chart
    chart = BarChart()
    chart.type = "col"
    chart.style = 2
    chart.title = "Total Sales by Agent"
    chart.legend = None
    
    # Remove gridlines for clean aesthetic
    chart.y_axis.majorGridlines = None
    chart.y_axis.title = ""

    data_ref = Reference(ws, min_col=6, min_row=10, max_row=end_row)
    cats_ref = Reference(ws, min_col=2, min_row=11, max_row=end_row)
    
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    
    ws.add_chart(chart, "H10")
