from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Prepare Data Sheet (Hidden) to keep dashboard clean
    data_ws = wb.create_sheet(title=f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'
    
    # Sample Data: Stacked Column (Category & Region)
    cat_data = [
        ["Market", "Chocolate Chip", "Sugar", "Oatmeal Raisin"],
        ["India", 62349, 25085, 21028],
        ["United Kingdom", 46530, 14620, 11497],
        ["United States", 36657, 9938, 5220]
    ]
    for row in cat_data:
        data_ws.append(row)
        
    data_ws.append([]) # Spacer row
    
    # Sample Data: Line Charts (Time Series)
    time_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    for row in time_data:
        data_ws.append(row)
        
    # 2. Setup Dashboard Sheet
    if sheet_name in wb.sheetnames:
        dash_ws = wb[sheet_name]
    else:
        dash_ws = wb.create_sheet(title=sheet_name)
        
    dash_ws.sheet_view.showGridLines = False
    
    # Theme Setup
    theme_colors = {
        "corporate_blue": "2F5597",
        "emerald_green": "385723",
        "dark_mode": "D0D0D0"
    }
    bg_colors = {
        "corporate_blue": "FFFFFF",
        "emerald_green": "FFFFFF",
        "dark_mode": "1E1E1E"
    }
    primary_color = theme_colors.get(theme, "2F5597")
    bg_color = bg_colors.get(theme, "FFFFFF")
    
    # Apply background fill across the visual dashboard area
    bg_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in dash_ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=20):
        for cell in row:
            cell.fill = bg_fill
            
    # Dashboard Title
    dash_ws["C2"] = title
    dash_ws["C2"].font = Font(size=24, bold=True, color=primary_color)
    dash_ws.merge_cells("C2:Q3")
    dash_ws["C2"].alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Create Stacked Bar Chart (Tall Left)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.legend.position = "b"
    bar_chart.y_axis.majorGridlines = None  # Clean look
    
    data_ref = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=4)
    cats_ref = Reference(data_ws, min_col=1, min_row=2, max_row=4)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    
    bar_chart.width = 16
    bar_chart.height = 11.5
    dash_ws.add_chart(bar_chart, "C5")
    
    # 4. Create Line Chart 1 (Short Top-Right)
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.legend = None
    line1.y_axis.majorGridlines = None
    
    data_ref1 = Reference(data_ws, min_col=2, min_row=6, max_row=10)
    cats_ref1 = Reference(data_ws, min_col=1, min_row=7, max_row=10)
    line1.add_data(data_ref1, titles_from_data=True)
    line1.set_categories(cats_ref1)
    
    line1.width = 14
    line1.height = 5.5
    dash_ws.add_chart(line1, "K5")
    
    # 5. Create Line Chart 2 (Short Bottom-Right)
    line2 = LineChart()
    line2.title = "Profit by month"
    line2.legend = None
    line2.y_axis.majorGridlines = None
    
    data_ref2 = Reference(data_ws, min_col=3, min_row=6, max_row=10)
    cats_ref2 = Reference(data_ws, min_col=1, min_row=7, max_row=10)
    line2.add_data(data_ref2, titles_from_data=True)
    line2.set_categories(cats_ref2)
    
    line2.width = 14
    line2.height = 5.5
    dash_ws.add_chart(line2, "K14")

    # Layout spacing adjustments
    dash_ws.column_dimensions['A'].width = 2
    dash_ws.column_dimensions['B'].width = 2
