from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string

def render(ws, anchor: str, *, data: list[dict], theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a leaderboard table with in-cell data bars for numeric columns.
    
    :param data: List of dictionaries representing the rows. First key is typically categorical.
                 Example: [{"Agent": "Alice", "Calls": 1031, "Closed": 37, "Revenue": 13519}, ...]
    """
    if not data:
        return
        
    # Theme palette resolution (fallback to standard if missing)
    theme_colors = {
        "corporate_blue": {
            "primary": "002060",
            "accent1": "4F81BD",
            "accent2": "C0504D",
            "accent3": "9BBB59",
            "accent4": "8064A2",
            "text_light": "FFFFFF"
        }
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    headers = list(data[0].keys())
    anchor_col = column_index_from_string(coordinate_from_string(anchor)[0])
    anchor_row = coordinate_from_string(anchor)[1]
    
    # 1. Render Headers
    header_font = Font(color=palette.get("text_light", "FFFFFF"), bold=True)
    header_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    
    for c_idx, header in enumerate(headers):
        cell = ws.cell(row=anchor_row, column=anchor_col + c_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    # 2. Render Data and Base Formatting
    for r_idx, row_dict in enumerate(data, start=1):
        for c_idx, key in enumerate(headers):
            val = row_dict[key]
            cell = ws.cell(row=anchor_row + r_idx, column=anchor_col + c_idx, value=val)
            
            # Infer basic number formats based on header names
            if isinstance(val, (int, float)):
                if any(kw in str(key).lower() for kw in ["value", "sales", "price", "revenue", "cost", "$"]):
                    cell.number_format = '"$"#,##0'
                elif isinstance(val, float):
                    cell.number_format = '#,##0.00'
                else:
                    cell.number_format = '#,##0'
                    
    # 3. Apply Conditional Formatting Data Bars
    data_start_row = anchor_row + 1
    data_end_row = anchor_row + len(data)
    
    color_keys = ["accent1", "accent2", "accent3", "accent4", "primary"]
    numeric_col_count = 0
    
    for c_idx, key in enumerate(headers):
        first_val = data[0][key]
        # Check if column is numeric to apply data bars
        if isinstance(first_val, (int, float)):
            col_letter = get_column_letter(anchor_col + c_idx)
            range_str = f"{col_letter}{data_start_row}:{col_letter}{data_end_row}"
            
            # DataBarRule color requires 8-character aRGB hex
            bar_color = palette[color_keys[numeric_col_count % len(color_keys)]]
            if len(bar_color) == 6:
                bar_color = f"FF{bar_color}"
                
            rule = DataBarRule(
                start_type="min",
                end_type="max",
                color=bar_color
            )
            ws.conditional_formatting.add(range_str, rule)
            numeric_col_count += 1
            
    # 4. Auto-size columns to fit headers and bars comfortably
    for c_idx, key in enumerate(headers):
        col_letter = get_column_letter(anchor_col + c_idx)
        # Give numeric columns extra padding for the bars to breathe
        is_numeric = isinstance(data[0][key], (int, float))
        padding = 8 if is_numeric else 4
        ws.column_dimensions[col_letter].width = max(len(str(key)) + padding, 12)
