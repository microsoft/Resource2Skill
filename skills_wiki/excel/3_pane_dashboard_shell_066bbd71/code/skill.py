from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a 3-pane dashboard layout onto a new worksheet, including a left-hand 
    control panel and three anchored charts (1 main, 2 secondary).
    """
    # 1. Theme Setup
    palettes = {
        "corporate_blue": {"primary": "003366", "text": "FFFFFF", "panel": "F2F2F2"},
        "modern_dark": {"primary": "333333", "text": "FFFFFF", "panel": "E0E0E0"},
        "forest_green": {"primary": "2E4E3F", "text": "FFFFFF", "panel": "F5F7F5"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])
    
    header_fill = PatternFill("solid", fgColor=colors["primary"])
    header_font = Font(color=colors["text"], size=22, bold=True)
    panel_fill = PatternFill("solid", fgColor=colors["panel"])

    # 2. Sheet Setup
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    data_ws = wb.create_sheet(f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'

    # 3. Dashboard Header
    ws.merge_cells("A1:S3")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.fill = header_fill
    title_cell.font = header_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 4. Control Panel Zone (Left Sidebar)
    ws.merge_cells("A5:D5")
    control_header = ws["A5"]
    control_header.value = "Interactive Controls Drop-Zone"
    control_header.font = Font(bold=True, color=colors["primary"])
    control_header.alignment = Alignment(horizontal="center")
    
    for row in range(5, 26):
        for col in range(1, 5):
            ws.cell(row=row, column=col).fill = panel_fill

    # 5. Populate Dummy Data for Charts
    # Data for Main Chart (Stacked Column: Profit by Market & Product)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62000, 23000, 21000],
        ["Philippines", 54000, 24000, 22000],
        ["United Kingdom", 46000, 26000, 11000],
        ["United States", 36000, 32000, 9000]
    ]
    for r_idx, row in enumerate(market_data, 1):
        for c_idx, val in enumerate(row, 1):
            data_ws.cell(row=r_idx, column=c_idx, value=val)

    # Data for Line Charts (Trends by Month)
    trend_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    for r_idx, row in enumerate(trend_data, 1):
        for c_idx, val in enumerate(row, 7):
            data_ws.cell(row=r_idx, column=c_idx, value=val)

    # 6. Chart 1: Main Stacked Bar Chart (Center)
    main_chart = BarChart()
    main_chart.type = "col"
    main_chart.grouping = "stacked"
    main_chart.overlap = 100
    main_chart.title = "Profit by Market & Cookie Type"
    main_chart.width = 16.5  # cm
    main_chart.height = 11.5 # cm

    data = Reference(data_ws, min_col=2, max_col=4, min_row=1, max_row=5)
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    main_chart.add_data(data, titles_from_data=True)
    main_chart.set_categories(cats)
    ws.add_chart(main_chart, "E5")

    # 7. Chart 2: Top Right Line Chart
    line_top = LineChart()
    line_top.title = "Units Sold Each Month"
    line_top.width = 13.0
    line_top.height = 5.5
    
    data_units = Reference(data_ws, min_col=8, min_row=1, max_row=5)
    cats_months = Reference(data_ws, min_col=7, min_row=2, max_row=5)
    line_top.add_data(data_units, titles_from_data=True)
    line_top.set_categories(cats_months)
    line_top.legend = None
    ws.add_chart(line_top, "M5")

    # 8. Chart 3: Bottom Right Line Chart
    line_bot = LineChart()
    line_bot.title = "Profit By Month"
    line_bot.width = 13.0
    line_bot.height = 5.5
    
    data_profit = Reference(data_ws, min_col=9, min_row=1, max_row=5)
    line_bot.add_data(data_profit, titles_from_data=True)
    line_bot.set_categories(cats_months)
    line_bot.legend = None
    ws.add_chart(line_bot, "M15")
