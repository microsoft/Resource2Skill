def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Theme and Color Setup
    themes = {
        "corporate_blue": {"primary_bg": "203764", "primary_fg": "FFFFFF", "accent": "4472C4"},
        "executive_dark": {"primary_bg": "262626", "primary_fg": "FFFFFF", "accent": "00B050"},
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 2. Setup the Dashboard Sheet
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False  # The key to the "Dashboard" look

    # Create Header Banner
    ws.merge_cells("A1:N3")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["primary_fg"])
    header_cell.fill = PatternFill(fill_type="solid", start_color=palette["primary_bg"], end_color=palette["primary_bg"])
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Setup Hidden Data Sheet
    data_ws = wb.create_sheet(f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden' 

    # Mock Data for Chart 1 (Stacked Bar)
    data_ws.append(["Market", "Fortune Cookie", "Sugar", "Snickerdoodle", "Oatmeal Raisin"])
    data_ws.append(["India", 50000, 20000, 10000, 15000])
    data_ws.append(["Philippines", 40000, 15000, 12000, 14000])
    data_ws.append(["United Kingdom", 30000, 25000, 8000, 12000])
    data_ws.append(["United States", 60000, 30000, 15000, 20000])

    # Mock Data for Charts 2 & 3 (Line Charts)
    data_ws.append([]) # Row 6 empty spacer
    data_ws.append(["Month", "Units Sold", "Profit"]) # Row 7
    months = ["Sep", "Oct", "Nov", "Dec"]
    units = [50601, 95622, 65481, 52970]
    profits = [124812, 228275, 160228, 136337]
    for m, u, p in zip(months, units, profits):
        data_ws.append([m, u, p])

    # 4. Generate Chart 1: Stacked Bar (Profit by Market & Cookie Type)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 10 # Built-in Excel chart style
    
    cats1 = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    data1 = Reference(data_ws, min_col=2, min_row=1, max_col=5, max_row=5)
    bar_chart.add_data(data1, titles_from_data=True)
    bar_chart.set_categories(cats1)
    bar_chart.height = 11.5
    bar_chart.width = 17
    ws.add_chart(bar_chart, "B5")

    # 5. Generate Chart 2: Line (Units Sold Each Month)
    line_chart1 = LineChart()
    line_chart1.title = "Units sold each month"
    line_chart1.style = 13
    cats2 = Reference(data_ws, min_col=1, min_row=8, max_row=11)
    data2 = Reference(data_ws, min_col=2, min_row=7, max_row=11)
    line_chart1.add_data(data2, titles_from_data=True)
    line_chart1.set_categories(cats2)
    line_chart1.height = 7.5
    line_chart1.width = 13
    line_chart1.legend = None # Clean look: no legend needed for single series
    ws.add_chart(line_chart1, "K5")

    # 6. Generate Chart 3: Line (Profit by Month)
    line_chart2 = LineChart()
    line_chart2.title = "Profit by month"
    line_chart2.style = 13
    data3 = Reference(data_ws, min_col=3, min_row=7, max_row=11)
    line_chart2.add_data(data3, titles_from_data=True)
    line_chart2.set_categories(cats2)
    line_chart2.height = 7.5
    line_chart2.width = 13
    line_chart2.legend = None
    ws.add_chart(line_chart2, "K18")
