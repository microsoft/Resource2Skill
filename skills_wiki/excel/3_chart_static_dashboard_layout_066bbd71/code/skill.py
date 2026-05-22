import random
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a 3-chart dashboard layout on the specified sheet.
    Generates hidden mock data if a data sheet doesn't already exist.
    """
    # 1. Theme Setup
    palettes = {
        "corporate_blue": {"primary_bg": "003366", "primary_fg": "FFFFFF"},
        "modern_green": {"primary_bg": "2E7D32", "primary_fg": "FFFFFF"},
        "slate_gray": {"primary_bg": "475569", "primary_fg": "FFFFFF"},
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # 2. Mock Data Setup
    data_ws_name = "ChartData_Hidden"
    if data_ws_name in wb.sheetnames:
        data_ws = wb[data_ws_name]
    else:
        data_ws = wb.create_sheet(data_ws_name)
        data_ws.sheet_state = 'hidden'
        
        # Mock Data for Stacked Bar (Profit by Market & Product)
        markets = ["India", "Malaysia", "Philippines", "UK", "USA"]
        products = ["Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Sugar"]
        
        data_ws.append(["Market"] + products)
        for i, market in enumerate(markets):
            row = [market] + [random.randint(10000, 50000) for _ in products]
            data_ws.append(row)
            # Apply currency format to data so chart axes inherit it
            for col_idx in range(2, 2 + len(products)):
                data_ws.cell(row=i + 2, column=col_idx).number_format = '"$"#,##0'
            
        # Mock Data for Line Charts (Trend by Month)
        months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        data_ws.append([]) # spacer row
        trend_start_row = len(markets) + 3
        
        data_ws.append(["Month", "Units Sold", "Profit"])
        for i, month in enumerate(months):
            units = random.randint(5000, 15000) + (i * 500)
            profit = units * 3.5 + random.randint(-5000, 5000)
            data_ws.append([month, units, profit])
            
            row_idx = trend_start_row + 1 + i
            data_ws.cell(row=row_idx, column=2).number_format = '#,##0'
            data_ws.cell(row=row_idx, column=3).number_format = '"$"#,##0'

    # 3. Dashboard Sheet Setup
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Header Banner
    ws.merge_cells("A1:N3")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["primary_fg"])
    header_cell.fill = PatternFill(start_color=palette["primary_bg"], end_color=palette["primary_bg"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Create Charts
    
    # Chart 1: Stacked Bar (Left)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.height = 14
    bar_chart.width = 16
    
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=1+len(markets))
    data = Reference(data_ws, min_col=2, min_row=1, max_col=1+len(products), max_row=1+len(markets))
    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    
    # Chart 2: Line Chart 1 (Top Right)
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.height = 7
    line1.width = 14
    line1.legend = None # Clean look without legend
    
    cats_trend = Reference(data_ws, min_col=1, min_row=trend_start_row+1, max_row=trend_start_row+len(months))
    data_units = Reference(data_ws, min_col=2, min_row=trend_start_row, max_row=trend_start_row+len(months))
    line1.add_data(data_units, titles_from_data=True)
    line1.set_categories(cats_trend)
    
    # Chart 3: Line Chart 2 (Bottom Right)
    line2 = LineChart()
    line2.title = "Profit by month"
    line2.height = 7
    line2.width = 14
    line2.legend = None
    
    data_profit = Reference(data_ws, min_col=3, min_row=trend_start_row, max_row=trend_start_row+len(months))
    line2.add_data(data_profit, titles_from_data=True)
    line2.set_categories(cats_trend)
    
    # 5. Position Charts in Grid
    ws.add_chart(bar_chart, "B5")
    ws.add_chart(line1, "I5")
    ws.add_chart(line2, "I16")
