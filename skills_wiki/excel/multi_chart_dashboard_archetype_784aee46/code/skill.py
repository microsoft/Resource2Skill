from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import PatternFill, Font, Alignment

def render_workbook(wb: Workbook, *, title: str = "Heroic Insights Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a multi-chart dashboard on a clean presentation sheet, 
    driven by data stored in a separate report data sheet.
    """
    # Color palette fallback based on theme
    palettes = {
        "corporate_blue": {"bg": "1F4E78", "fg": "FFFFFF"},
        "emerald_green": {"bg": "006633", "fg": "FFFFFF"},
        "crimson_red": {"bg": "800000", "fg": "FFFFFF"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])

    # 1. Setup Data Sheet
    ws_data = wb.active
    ws_data.title = "Report Data"
    
    # -- Data Block 1: Monthly Revenue Trend
    ws_data.append(["Month", "Revenue"])
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    revs = [15000, 18000, 22000, 20000, 25000, 28000, 27000, 29000, 31000, 30000, 34000, 36000]
    for m, r in zip(months, revs):
        ws_data.append([m, r])
        
    ws_data.append([]) # Empty row spacer
    
    # -- Data Block 2: Units Sold by Category
    row_offset_cat = ws_data.max_row + 1
    ws_data.append(["Year", "Hoodies", "T-shirts"])
    ws_data.append(["2023", 15201, 17541])
    ws_data.append(["2024", 17000, 19000])
    
    ws_data.append([]) # Empty row spacer
    
    # -- Data Block 3: Top 5 States by Profit
    row_offset_states = ws_data.max_row + 1
    ws_data.append(["State", "Profit"])
    states = ["California", "Texas", "New York", "Florida", "Illinois"]
    profits = [38906, 34420, 32120, 31000, 29000]
    for s, p in zip(states, profits):
        ws_data.append([s, p])

    # 2. Setup Dashboard Sheet
    ws_dash = wb.create_sheet(title="Dashboard", index=0)
    ws_dash.sheet_view.showGridLines = False
    wb.active = 0
    
    # 3. Dashboard Header (mimicking a floating banner shape)
    ws_dash.merge_cells("B2:O4")
    header_cell = ws_dash["B2"]
    header_cell.value = title
    header_cell.font = Font(size=28, bold=True, color=colors["fg"])
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    header_cell.fill = PatternFill("solid", fgColor=colors["bg"])
    
    # 4. Line Chart: Revenue Trend
    chart_rev = LineChart()
    chart_rev.title = "Monthly Revenue Trend"
    chart_rev.style = 13
    chart_rev.y_axis.title = "Revenue (USD)"
    
    data_rev = Reference(ws_data, min_col=2, min_row=1, max_row=1+len(months))
    cats_rev = Reference(ws_data, min_col=1, min_row=2, max_row=1+len(months))
    chart_rev.add_data(data_rev, titles_from_data=True)
    chart_rev.set_categories(cats_rev)
    chart_rev.width = 16
    chart_rev.height = 8
    ws_dash.add_chart(chart_rev, "B6")
    
    # 5. Clustered Column Chart: Units Sold
    chart_units = BarChart()
    chart_units.type = "col" # Vertical bars
    chart_units.style = 10
    chart_units.title = "Units Sold by Category"
    chart_units.y_axis.title = "Units"
    
    data_units = Reference(ws_data, min_col=2, min_row=row_offset_cat, max_row=row_offset_cat+2, max_col=3)
    cats_units = Reference(ws_data, min_col=1, min_row=row_offset_cat+1, max_row=row_offset_cat+2)
    chart_units.add_data(data_units, titles_from_data=True)
    chart_units.set_categories(cats_units)
    chart_units.width = 12
    chart_units.height = 8
    ws_dash.add_chart(chart_units, "K6")
    
    # 6. Horizontal Bar Chart: Top 5 States
    chart_states = BarChart()
    chart_states.type = "bar" # Horizontal bars for ranking
    chart_states.style = 10
    chart_states.title = "Top 5 States by Profit"
    chart_states.x_axis.title = "Profit (USD)"
    
    data_states = Reference(ws_data, min_col=2, min_row=row_offset_states, max_row=row_offset_states+len(states))
    cats_states = Reference(ws_data, min_col=1, min_row=row_offset_states+1, max_row=row_offset_states+len(states))
    chart_states.add_data(data_states, titles_from_data=True)
    chart_states.set_categories(cats_states)
    chart_states.width = 16
    chart_states.height = 8
    ws_dash.add_chart(chart_states, "B21")
