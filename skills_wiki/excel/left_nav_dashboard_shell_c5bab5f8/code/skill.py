from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Generates a two-pane application-style dashboard layout:
    A dark left-hand navigation/KPI panel and a light main content canvas.
    """
    # Standard theme palette fallback
    themes = {
        "corporate_blue": {
            "primary": "1F4E78", 
            "bg_light": "F2F4F7", 
            "text_light": "FFFFFF", 
            "text_dark": "222222",
            "accent": "D9E1E8"
        },
        "emerald_green": {
            "primary": "276749", 
            "bg_light": "E6F4EA", 
            "text_light": "FFFFFF", 
            "text_dark": "1A202C",
            "accent": "C6E4D1"
        }
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # Create or get worksheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    # Disable standard Excel gridlines for a clean UI aesthetic
    ws.sheet_view.showGridLines = False

    # 1. Define Column Layout constraints
    ws.column_dimensions['A'].width = 2   # Left outer margin
    ws.column_dimensions['B'].width = 30  # Slicer / Top-level KPI panel
    ws.column_dimensions['C'].width = 2   # Divider gap

    # Create a uniform snap-to grid for the main content area (Charts/Tables)
    for col_idx in range(4, 27): # Columns D through Z
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 10

    # 2. Setup Pattern Fills
    nav_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    main_fill = PatternFill(start_color=palette["bg_light"], end_color=palette["bg_light"], fill_type="solid")
    divider_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # 3. Apply Background Colors (Painting the canvas)
    # Applying down to row 60 provides plenty of vertical space for a standard 1080p screen
    for row in range(1, 61):
        ws.cell(row=row, column=2).fill = nav_fill
        ws.cell(row=row, column=3).fill = divider_fill
        
        for col in range(4, 27):
            ws.cell(row=row, column=col).fill = main_fill

    # 4. Configure the Header / Dashboard Title
    ws.row_dimensions[2].height = 40
    title_cell = ws['B2']
    title_cell.value = title.upper()
    title_cell.font = Font(name="Arial", size=20, bold=True, color=palette["text_light"])
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 5. Setup Left-Pane KPI Anchors
    ws.row_dimensions[4].height = 20
    kpi_header = ws['B4']
    kpi_header.value = "  OVERALL METRICS"
    kpi_header.font = Font(name="Arial", size=10, bold=True, color=palette["accent"])
    kpi_header.alignment = Alignment(horizontal="left", vertical="bottom")
    
    # Mocking a left-pane KPI card
    ws['B6'].value = "  Total Revenue"
    ws['B6'].font = Font(name="Arial", size=10, color=palette["text_light"])
    ws['B7'].value = "  $649,019"
    ws['B7'].font = Font(name="Arial", size=22, bold=True, color=palette["text_light"])

    # 6. Main Content Area Header
    ws.row_dimensions[3].height = 30
    content_header = ws['D3']
    content_header.value = "Performance Breakdown"
    content_header.font = Font(name="Arial", size=16, bold=True, color=palette["text_dark"])
    content_header.alignment = Alignment(horizontal="left", vertical="center")

    # Mocking a target area for a chart to guide the user/system
    ws.merge_cells('D5:K15')
    chart_anchor = ws['D5']
    chart_anchor.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Apply a subtle border to the chart anchor placeholder
    thin_border = Border(
        left=Side(style='thin', color="CCCCCC"), right=Side(style='thin', color="CCCCCC"),
        top=Side(style='thin', color="CCCCCC"), bottom=Side(style='thin', color="CCCCCC")
    )
    for row in range(5, 16):
        for col in range(4, 12):
            ws.cell(row=row, column=col).border = thin_border
            
    chart_anchor.value = "[ Insert Main Trend Chart Here ]"
    chart_anchor.font = Font(name="Arial", size=12, italic=True, color="888888")
    chart_anchor.alignment = Alignment(horizontal="center", vertical="center")
