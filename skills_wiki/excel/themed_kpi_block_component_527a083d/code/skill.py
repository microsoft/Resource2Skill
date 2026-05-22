from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, region: str = "Asia", label: str = "Revenue", value: float = 369989.0, format_str: str = "$#,##0", percent_val: float = 0.05, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a distinct, colorful KPI block (simulating a shape-based card)
    using cell merges and theme colors.
    """
    # Self-contained theme palette fallback
    themes = {
        "corporate_blue": {"primary": "002060", "text": "FFFFFF", "secondary": "4F81BD"},
        "executive_dark": {"primary": "262626", "text": "FFFFFF", "secondary": "595959"},
        "emerald_green": {"primary": "0F5132", "text": "FFFFFF", "secondary": "198754"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    row, col = coordinate_to_tuple(anchor)

    # 1. Structure & Values
    ws.cell(row=row, column=col, value=region)
    ws.cell(row=row+1, column=col, value=label)
    
    v_cell = ws.cell(row=row+2, column=col, value=value)
    v_cell.number_format = format_str

    p_cell = None
    if percent_val is not None:
        # Merge top rows across 2 columns
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
        ws.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+1)
        
        p_cell = ws.cell(row=row+2, column=col+1, value=percent_val)
        p_cell.number_format = "0%"
    else:
        # Merge all rows if no secondary metric
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
        ws.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+1)
        ws.merge_cells(start_row=row+2, start_column=col, end_row=row+2, end_column=col+1)

    # 2. Formatting Definitions
    main_fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    sec_fill = PatternFill(start_color=palette["secondary"], end_color=palette["secondary"], fill_type="solid")
    
    center_align = Alignment(horizontal="center", vertical="center")
    
    title_font = Font(color=palette["text"], bold=True, size=16)
    label_font = Font(color=palette["text"], italic=True, size=11)
    val_font = Font(color=palette["text"], bold=True, size=14)
    pct_font = Font(color=palette["text"], bold=True, size=12)
    
    thin_border = Border(
        left=Side(style='thin', color="FFFFFF"),
        right=Side(style='thin', color="FFFFFF"),
        top=Side(style='thin', color="FFFFFF"),
        bottom=Side(style='thin', color="FFFFFF")
    )

    # 3. Apply Formats to Block
    for r in range(row, row+3):
        for c in range(col, col+2):
            cell = ws.cell(row=r, column=c)
            cell.fill = main_fill
            cell.alignment = center_align
            cell.border = thin_border

    # Target specific fonts
    ws.cell(row=row, column=col).font = title_font
    ws.cell(row=row+1, column=col).font = label_font
    v_cell.font = val_font
    
    if p_cell:
        p_cell.font = pct_font
        p_cell.fill = sec_fill # Differentiate the secondary metric like the overlay shape

    # 4. Dimension Sizing
    ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions[get_column_letter(col+1)].width = 10
    
    # Pad rows for breathing room
    ws.row_dimensions[row].height = 25
    ws.row_dimensions[row+1].height = 18
    ws.row_dimensions[row+2].height = 25
