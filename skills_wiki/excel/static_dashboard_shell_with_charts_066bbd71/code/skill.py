def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # Standard theme fallback
    theme_palette = {
        "corporate_blue": {"primary_bg": "1F4E78", "primary_fg": "FFFFFF"},
        "modern_dark": {"primary_bg": "262626", "primary_fg": "FFFFFF"},
    }.get(theme, {"primary_bg": "1F4E78", "primary_fg": "FFFFFF"})

    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_data = wb.create_sheet(title="ReportData")

    # --- 1. Populate Hidden Data Sheet ---
    # Data for Stacked Bar (Profit by Region)
    ws_data.append(["Region", "Product A", "Product B", "Product C"])
    ws_data.append(["North America", 125000, 85000, 45000])
    ws_data.append(["Europe", 90000, 75000, 30000])
    ws_data.append(["Asia", 150000, 110000, 60000])

    # Data for Line Chart (Units Sold by Month)
    ws_data.append([]) # spacer row
    ws_data.append(["Month", "Units Sold"])
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    units = [1200, 1350, 1100, 1500, 1800, 1600, 1900, 2100, 2050, 2300, 2500, 2800]
    for m, u in zip(months, units):
        ws_data.append([m, u])

    # --- 2. Build Dashboard Shell ---
    ws_dash.sheet_view.showGridLines = False

    # Header Banner
    ws_dash.merge_cells("A1:O3")
    banner = ws_dash["A1"]
    banner.value = title
    banner.font = Font(size=24, bold=True, color=theme_palette["primary_fg"])
    banner.fill = PatternFill(start_color=theme_palette["primary_bg"], end_color=theme_palette["primary_bg"], fill_type="solid")
    banner.alignment = Alignment(horizontal="center", vertical="center")

    # --- 3. Create and Place Charts ---
    # Stacked Column Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.style = 10
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Region & Product"
    bar_chart.y_axis.title = "Profit ($)"

    bar_data = Reference(ws_data, min_col=2, min_row=1, max_row=4, max_col=4)
    bar_cats = Reference(ws_data, min_col=1, min_row=2, max_row=4)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    bar_chart.width = 16
    bar_chart.height = 8

    ws_dash.add_chart(bar_chart, "B5")

    # Line Chart
    line_chart = LineChart()
    line_chart.style = 13
    line_chart.title = "Units Sold Each Month"
    line_chart.y_axis.title = "Total Units"

    line_data = Reference(ws_data, min_col=2, min_row=6, max_row=18)
    line_cats = Reference(ws_data, min_col=1, min_row=7, max_row=18)
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)
    line_chart.width = 16
    line_chart.height = 8

    ws_dash.add_chart(line_chart, "I5")

    # Hide the data sheet for a clean user experience
    ws_data.sheet_state = 'hidden'
