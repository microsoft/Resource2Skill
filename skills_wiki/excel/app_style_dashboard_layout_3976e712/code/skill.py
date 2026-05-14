def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard South America 2022", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Standard theme palette (fallback)
    palettes = {
        "corporate_blue": {
            "nav_bg": "1A365D",      # Dark blue
            "nav_fg": "FFFFFF",      # White
            "sheet_bg": "F3F4F6",    # Light gray canvas
            "card_bg": "FFFFFF",     # White cards
            "border": "E5E7EB",      # Light gray borders
            "text_main": "111827",   # Dark gray
            "text_muted": "6B7280"   # Gray
        }
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # Create sheet and hide gridlines for the "app" feel
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Define fills and fonts
    nav_fill = PatternFill("solid", fgColor=palette["nav_bg"])
    sheet_fill = PatternFill("solid", fgColor=palette["sheet_bg"])
    card_fill = PatternFill("solid", fgColor=palette["card_bg"])

    nav_font = Font(color=palette["nav_fg"], bold=True, size=16)
    title_font = Font(color=palette["text_main"], bold=True, size=20)
    card_title_font = Font(color=palette["text_muted"], bold=True, size=12)

    center_align = Alignment(horizontal="center", vertical="center")
    
    border_side = Side(border_style="thin", color=palette["border"])

    # Paint the entire canvas with the background color
    for row in ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=15):
        for cell in row:
            cell.fill = sheet_fill

    # 1. Construct Side Navigation Bar (Column A)
    ws.column_dimensions['A'].width = 8
    for row in range(1, 31):
        ws.cell(row=row, column=1).fill = nav_fill

    # Add Navigation Icons with internal hyperlinks
    nav_items = [
        (3, "🏠", f"#'{sheet_name}'!A1"),
        (6, "📊", "#'Inputs'!A1"),
        (9, "📞", "#'Contacts'!A1"),
        (12, "❓", "mailto:support@example.com")
    ]
    
    for row, icon, link in nav_items:
        cell = ws.cell(row=row, column=1, value=icon)
        cell.font = nav_font
        cell.alignment = center_align
        cell.hyperlink = link

    # 2. Define Card Layout (min_col, min_row, max_col, max_row, title)
    cards = [
        (3, 2, 13, 4, title),                          # Top Header Banner
        (3, 6, 5, 10, "Sales"),                        # KPI 1
        (7, 6, 9, 10, "Profit"),                       # KPI 2
        (11, 6, 13, 10, "# of Customers"),             # KPI 3
        (3, 12, 8, 22, "2021-2022 Sales Trend"),       # Main Line Chart Area
        (9, 12, 13, 16, "Sales by Country"),           # Top Right Map Chart Area
        (9, 18, 13, 22, "Customer Satisfaction")       # Bottom Right Radar Chart Area
    ]

    # 3. Draw the Cards
    for min_col, min_row, max_col, max_row, card_title in cards:
        # Apply card fill and outer borders
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Apply borders only to the outer edges of the block
                top = border_side if r == min_row else None
                bottom = border_side if r == max_row else None
                left = border_side if c == min_col else None
                right = border_side if c == max_col else None
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)

        # Set the title of the card
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=min_row, end_column=max_col)
        title_cell = ws.cell(row=min_row, column=min_col, value=card_title)

        if card_title == title:
            # Main dashboard header styling
            title_cell.font = title_font
            title_cell.alignment = center_align
        else:
            # Standard card title styling
            title_cell.font = card_title_font
            title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Clean up column widths for the main dashboard area
    for col in range(2, 15):
        ws.column_dimensions[get_column_letter(col)].width = 12
