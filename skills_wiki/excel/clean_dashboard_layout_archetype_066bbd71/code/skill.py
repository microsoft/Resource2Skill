def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    
    # 1. Prepare backend summary data sheet
    ws_summary = wb.active
    ws_summary.title = "Summary_Data"
    
    summary_rows = [
        # Table 1: Stacked Bar Data
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62000, 4800, 21000],
        ["Philippines", 54000, 7000, 22000],
        ["United Kingdom", 46000, 5200, 11000],
        ["United States", 36000, 6300, 22000],
        [],
        # Table 2: Line Chart Data (Units)
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970],
        [],
        # Table 3: Line Chart Data (Profit)
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    
    for row in summary_rows:
        ws_summary.append(row)
        
    # 2. Setup presentation dashboard sheet
    ws_dash = wb.create_sheet("Dashboard", 0)
    
    # Core Mechanism: Hide gridlines for a clean "software UI" feel
    ws_dash.sheet_view.showGridLines = False
    wb.active = ws_dash
    
    # Standard helper fallback for themes
    theme_palette = {
        "corporate_blue": {"primary": "2B579A", "text": "FFFFFF"},
        "forest_green": {"primary": "27AE60", "text": "FFFFFF"},
        "slate_dark": {"primary": "2C3E50", "text": "FFFFFF"}
    }.get(theme, {"primary": "2B579A", "text": "FFFFFF"})
    
    # 3. Create Branded Title Banner
    ws_dash.merge_cells("A1:P3")
    banner = ws_dash["A1"]
    banner.value = title
    banner.font = Font(size=24, bold=True, color=theme_palette["text"])
    banner.fill = PatternFill("solid", fgColor=theme_palette["primary"])
    banner.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    
    # Add minor padding to the left edge
    ws_dash.column_dimensions['A'].width = 3
    
    # 4. Create Stacked Column Chart
    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Cookie Type"
    c1.width = 16
    c1.height = 11.5
    
    data1 = Reference(ws_summary, min_col=2, min_row=1, max_col=4, max_row=5)
    cats1 = Reference(ws_summary, min_col=1, min_row=2, max_row=5)
    c1.add_data(data1, titles_from_data=True)
    c1.set_categories(cats1)
    ws_dash.add_chart(c1, "B5")
    
    # 5. Create Units Sold Line Chart
    c2 = LineChart()
    c2.title = "Units sold each month"
    c2.width = 14
    c2.height = 8.5
    c2.legend = None  # Strip legend to maximize chart area space
    
    data2 = Reference(ws_summary, min_col=2, min_row=7, max_col=2, max_row=11)
    cats2 = Reference(ws_summary, min_col=1, min_row=8, max_row=11)
    c2.add_data(data2, titles_from_data=True)
    c2.set_categories(cats2)
    ws_dash.add_chart(c2, "J5")
    
    # 6. Create Profit Line Chart
    c3 = LineChart()
    c3.title = "Profit by month"
    c3.width = 14
    c3.height = 8.5
    c3.legend = None
    
    data3 = Reference(ws_summary, min_col=2, min_row=13, max_col=2, max_row=17)
    cats3 = Reference(ws_summary, min_col=1, min_row=14, max_row=17)
    c3.add_data(data3, titles_from_data=True)
    c3.set_categories(cats3)
    ws_dash.add_chart(c3, "J16")
    
    # 7. Hide backend data sheet to enforce dashboard UX
    ws_summary.sheet_state = 'hidden'
