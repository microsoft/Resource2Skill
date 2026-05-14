from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.cell import column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", main_label: str = "Revenue", main_value=369989, secondary_value=0.05, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a standalone KPI Card component using grid cells to simulate a shape-based widget.
    Values can be static numbers or Excel formula strings (e.g. "=B2").
    """
    # Parse anchor cell
    col_letter = "".join([c for c in anchor if c.isalpha()])
    row_num = int("".join([c for c in anchor if c.isdigit()]))
    col_idx = column_index_from_string(col_letter)
    
    # Theme configuration (Fallback to Navy background / White text)
    bg_color = "002060" 
    text_color = "FFFFFF"
    
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    font_title = Font(color=text_color, size=14, bold=True)
    font_label = Font(color=text_color, size=10)
    font_value = Font(color=text_color, size=14, bold=True)
    font_secondary = Font(color=text_color, size=16, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    # Apply unified background fill to the 3x3 block
    for r in range(row_num, row_num + 3):
        for c in range(col_idx, col_idx + 3):
            ws.cell(row=r, column=c).fill = fill
            
    # Top Row: Title Header (Spans all 3 cols)
    ws.merge_cells(start_row=row_num, start_column=col_idx, end_row=row_num, end_column=col_idx+2)
    title_cell = ws.cell(row=row_num, column=col_idx)
    title_cell.value = title
    title_cell.font = font_title
    title_cell.alignment = align_center
    
    # Middle Row Left: Main Metric Label
    label_cell = ws.cell(row=row_num+1, column=col_idx)
    label_cell.value = main_label
    label_cell.font = font_label
    label_cell.alignment = align_left
    
    # Bottom Row Left: Main Metric Value
    val_cell = ws.cell(row=row_num+2, column=col_idx)
    val_cell.value = main_value
    val_cell.font = font_value
    val_cell.alignment = align_left
    val_cell.number_format = '"$"#,##0'
    
    # Middle & Bottom Row Right: Secondary Value (Market Share %)
    ws.merge_cells(start_row=row_num+1, start_column=col_idx+2, end_row=row_num+2, end_column=col_idx+2)
    sec_cell = ws.cell(row=row_num+1, column=col_idx+2)
    sec_cell.value = secondary_value
    sec_cell.font = font_secondary
    sec_cell.alignment = align_right
    sec_cell.number_format = '0%'
    
    # Adjust column widths to shape the card properly
    ws.column_dimensions[get_column_letter(col_idx)].width = 15
    ws.column_dimensions[get_column_letter(col_idx+1)].width = 3  # spacer between metrics
    ws.column_dimensions[get_column_letter(col_idx+2)].width = 10
