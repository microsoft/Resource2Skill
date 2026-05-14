from typing import List, Dict, Any
from openpyxl.workbook import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference

def render_sheet(wb: Workbook, sheet_name: str, *, title: str, theme: str = "viva_green", **kwargs) -> None:
    """
    Renders a two-tone dashboard shell with a left KPI sidebar and a main chart canvas.
    
    Expected kwargs:
    - kpis: List[Dict[str, str]] (e.g., [{"label": "Orders", "value": "2,400"}])
    - trend_data: List[tuple] (e.g., [("Week", "Revenue"), ("W1", 5000), ...])
    - product_data: List[tuple] (e.g., [("Product", "Qty"), ("T-Shirts", 744), ...])
    """
    
    # 1. Theme Setup
    palettes = {
        "viva_green": {
            "sidebar_bg": "2A5336",  # Dark Green
            "sidebar_fg": "FFFFFF",  # White
            "canvas_bg": "E2EFDA",   # Light Green
            "accent": "548235"
        },
        "corporate_blue": {
            "sidebar_bg": "1F4E78",  # Dark Blue
            "sidebar_fg": "FFFFFF",  # White
            "canvas_bg": "F2F2F2",   # Light Gray
            "accent": "2F75B5"
        }
    }
    colors = palettes.get(theme, palettes["corporate_blue"])
    
    fill_sidebar = PatternFill(start_color=colors["sidebar_bg"], end_color=colors["sidebar_bg"], fill_type="solid")
    fill_canvas = PatternFill(start_color=colors["canvas_bg"], end_color=colors["canvas_bg"], fill_type="solid")
    
    font_kpi_val = Font(name="Calibri", size=22, bold=True, color=colors["sidebar_fg"])
    font_kpi_lbl = Font(name="Calibri", size=11, color=colors["sidebar_fg"])
    font_title = Font(name="Calibri", size=24, bold=True, color=colors["sidebar_bg"])

    # 2. Create the Dashboard Sheet
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Configure Layout Widths
    ws.column_dimensions['A'].width = 3   # Left padding
    ws.column_dimensions['B'].width = 18  # KPI column
    ws.column_dimensions['C'].width = 3   # Right padding / divider
    for col in "DEFGHIJKLMNO":
        ws.column_dimensions[col].width = 12

    # Apply Background Fills (Sidebar vs Canvas)
    for row in range(1, 40):
        for col_idx in range(1, 4):  # Cols A:C (Sidebar)
            ws.cell(row=row, column=col_idx).fill = fill_sidebar
        for col_idx in range(4, 16): # Cols D:O (Canvas)
            ws.cell(row=row, column=col_idx).fill = fill_canvas

    # 3. Mount KPIs to Sidebar
    default_kpis = [
        {"label": "Total Orders", "value": "2,400"},
        {"label": "Revenue", "value": "$649.0k"},
        {"label": "Avg Rating", "value": "4.0"},
        {"label": "Days to Deliver", "value": "2.3"}
    ]
    kpis = kwargs.get("kpis", default_kpis)
    
    # Sidebar Title
    logo_cell = ws["B2"]
    logo_cell.value = title.upper()
    logo_cell.font = Font(name="Calibri", size=18, bold=True, color=colors["sidebar_fg"])
    
    start_row = 5
    for kpi in kpis:
        val_cell = ws.cell(row=start_row, column=2)
        lbl_cell = ws.cell(row=start_row + 1, column=2)
        
        val_cell.value = kpi["value"]
        val_cell.font = font_kpi_val
        val_cell.alignment = Alignment(horizontal="left")
        
        lbl_cell.value = kpi["label"]
        lbl_cell.font = font_kpi_lbl
        lbl_cell.alignment = Alignment(horizontal="left")
        
        start_row += 4

    # 4. Prepare Background Data for Charts
    data_ws = wb.create_sheet(title=f"_{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'
    
    default_trend = [
        ("Week", "Revenue"), ("W1", 45000), ("W2", 52000), 
        ("W3", 48000), ("W4", 61000), ("W5", 59000), ("W6", 75000)
    ]
    trend_data = kwargs.get("trend_data", default_trend)
    for row_data in trend_data:
        data_ws.append(row_data)
        
    default_products = [
        ("Product", "Quantity"), ("T-Shirts", 744), ("Jeans", 527), 
        ("Sneakers", 482), ("Hoodies", 318), ("Caps", 215)
    ]
    product_data = kwargs.get("product_data", default_products)
    # Write product data leaving a gap
    for r_idx, row_data in enumerate(product_data, start=1):
        data_ws.cell(row=r_idx, column=4, value=row_data[0])
        data_ws.cell(row=r_idx, column=5, value=row_data[1])

    # 5. Create and Mount Charts to Canvas
    
    # Chart 1: Revenue Trend (Line Chart)
    line_chart = LineChart()
    line_chart.title = "Weekly Revenue Trend"
    line_chart.style = 13
    line_chart.y_axis.title = "Revenue"
    line_chart.x_axis.title = "Week"
    line_chart.legend = None
    
    data_ref = Reference(data_ws, min_col=2, min_row=1, max_row=len(trend_data))
    cats_ref = Reference(data_ws, min_col=1, min_row=2, max_row=len(trend_data))
    line_chart.add_data(data_ref, titles_from_data=True)
    line_chart.set_categories(cats_ref)
    
    # Style the line to match theme
    s1 = line_chart.series[0]
    s1.graphicalProperties.line.solidFill = colors["accent"]
    s1.graphicalProperties.line.width = 30000
    
    line_chart.width = 16
    line_chart.height = 8
    ws.add_chart(line_chart, "E3")

    # Chart 2: Product Popularity (Bar Chart)
    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.style = 10
    bar_chart.title = "Top Products by Quantity"
    bar_chart.legend = None
    
    bar_data_ref = Reference(data_ws, min_col=5, min_row=1, max_row=len(product_data))
    bar_cats_ref = Reference(data_ws, min_col=4, min_row=2, max_row=len(product_data))
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)
    
    # Style the bars
    b1 = bar_chart.series[0]
    b1.graphicalProperties.solidFill = colors["accent"]
    
    bar_chart.width = 16
    bar_chart.height = 10
    ws.add_chart(bar_chart, "E13")
    
    # Freeze the sidebar
    ws.freeze_panes = "D1"
