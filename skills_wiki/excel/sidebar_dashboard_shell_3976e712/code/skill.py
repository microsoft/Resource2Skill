from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.cell import range_boundaries

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a containerized dashboard shell with a navigation sidebar and blank content panels.
    """
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Define layout palette (simulate theme extraction)
    sidebar_color = "1F3864"  # Dark blue sidebar
    canvas_color = "F2F2F2"   # Light gray background
    panel_color = "FFFFFF"    # White content panels
    border_color = "D9D9D9"   # Subtle panel borders
    title_text_color = "595959"

    sidebar_fill = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    canvas_fill = PatternFill(start_color=canvas_color, end_color=canvas_color, fill_type="solid")
    panel_fill = PatternFill(start_color=panel_color, end_color=panel_color, fill_type="solid")

    # 1. Paint the entire canvas background
    for row in ws.iter_rows(min_row=1, max_row=40, min_col=2, max_col=22):
        for cell in row:
            cell.fill = canvas_fill

    # 2. Setup the Sidebar (Column A)
    ws.column_dimensions['A'].width = 8
    for row in range(1, 41):
        ws.cell(row=row, column=1).fill = sidebar_fill

    # 3. Add Dashboard Title
    title_cell = ws.cell(row=2, column=3, value=title)
    title_cell.font = Font(size=20, bold=True, color=sidebar_color)

    # 4. Define Content Panels (Containers)
    # The layout separates the top row into 3 KPI blocks, and the bottom into 2 larger chart blocks
    panels = [
        {"range": "C4:H10", "title": "Sales Overview"},
        {"range": "J4:O10", "title": "Profit Margin"},
        {"range": "Q4:U10", "title": "Customer Satisfaction"},
        {"range": "C12:O26", "title": "Trend Analysis"},
        {"range": "Q12:U26", "title": "Regional Breakdown"}
    ]

    for panel in panels:
        min_col, min_row, max_col, max_row = range_boundaries(panel["range"])
        
        # Apply panel background and outline border
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.fill = panel_fill
                
                # Calculate border edges to create a bounding box
                top_edge = Side(style='thin', color=border_color) if cell.row == min_row else None
                bottom_edge = Side(style='thin', color=border_color) if cell.row == max_row else None
                left_edge = Side(style='thin', color=border_color) if cell.column == min_col else None
                right_edge = Side(style='thin', color=border_color) if cell.column == max_col else None
                
                cell.border = Border(top=top_edge, bottom=bottom_edge, left=left_edge, right=right_edge)

        # Add panel title
        title_c = ws.cell(row=min_row, column=min_col, value=panel["title"])
        title_c.font = Font(bold=True, size=11, color=title_text_color)
        
        # Adjust layout spacing slightly
        ws.row_dimensions[min_row].height = 20
        title_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
