from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Executive Dashboard", theme: str = "corporate_blue", kpis: list = None, **kwargs) -> None:
    """
    Renders a standard dashboard shell with a dark left-hand KPI sidebar 
    and a wide light-themed main canvas for charts.
    """
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    # Standard theme dictionary fallback
    themes = {
        "corporate_blue": {
            "sidebar": "1E3A8A", "canvas": "F3F4F6", 
            "text_light": "FFFFFF", "accent": "60A5FA", "text_dark": "111827"
        },
        "forest_green": {
            "sidebar": "14532D", "canvas": "ECFDF5", 
            "text_light": "FFFFFF", "accent": "34D399", "text_dark": "064E3B"
        },
        "charcoal": {
            "sidebar": "27272A", "canvas": "F4F4F5", 
            "text_light": "F4F4F5", "accent": "FBBF24", "text_dark": "18181B"
        }
    }
    t_colors = themes.get(theme, themes["corporate_blue"])

    fill_sidebar = PatternFill(start_color=t_colors["sidebar"], end_color=t_colors["sidebar"], fill_type="solid")
    fill_canvas = PatternFill(start_color=t_colors["canvas"], end_color=t_colors["canvas"], fill_type="solid")
    
    font_title = Font(name="Calibri", size=24, bold=True, color=t_colors["text_dark"])
    font_kpi_lbl = Font(name="Calibri", size=11, bold=True, color=t_colors["text_light"])
    font_kpi_val = Font(name="Calibri", size=20, bold=True, color=t_colors["accent"])

    # 1. Setup Column Widths for Layout
    ws.column_dimensions['A'].width = 3   # Left padding
    ws.column_dimensions['B'].width = 18  # KPI content
    ws.column_dimensions['C'].width = 3   # Right padding

    for i in range(4, 22):
        ws.column_dimensions[get_column_letter(i)].width = 12

    # 2. Apply Background Fills (Sidebar vs Canvas)
    for row in range(1, 45):
        for col in range(1, 22):
            cell = ws.cell(row=row, column=col)
            if col <= 3:
                cell.fill = fill_sidebar
            else:
                cell.fill = fill_canvas

    # 3. Inject Main Canvas Title
    ws.merge_cells("E2:K3")
    title_cell = ws.cell(row=2, column=5, value=title)
    title_cell.font = font_title
    title_cell.alignment = Alignment(vertical="center")

    # 4. Inject Sidebar KPIs
    if kpis is None:
        kpis = [
            ("TOTAL ORDERS", "2,400"),
            ("TOTAL REVENUE", "$649.0K"),
            ("AVG RATING", "4.0"),
            ("DAYS TO DELIVER", "2.3")
        ]

    start_row = 6
    for lbl, val in kpis:
        # Label
        lbl_cell = ws.cell(row=start_row, column=2, value=lbl)
        lbl_cell.font = font_kpi_lbl
        lbl_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Value
        val_cell = ws.cell(row=start_row + 1, column=2, value=val)
        val_cell.font = font_kpi_val
        val_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Spacer between KPIs
        start_row += 4

    # Hide gridlines for cleaner dashboard look
    ws.sheet_view.showGridLines = False
