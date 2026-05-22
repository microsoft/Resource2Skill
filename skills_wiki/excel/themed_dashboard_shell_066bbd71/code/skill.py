def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, LineChart, Reference

    # Fallback theme dictionary 
    theme_colors = {
        "corporate_blue": "1F4E78",
        "emerald_green": "228B22",
        "slate_gray": "708090"
    }
    primary_color = theme_colors.get(theme, "1F4E78")

    # 1. Setup Dashboard Sheet Canvas
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # 2. Setup Hidden Data Sheet
    data_ws = wb.create_sheet(f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'

    # 3. Populate Chart Data
    # Time Series Data (Rows 1-7)
    data_ws.append(["Month", "Units Sold", "Profit"])
    for m, u, p in zip(
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun"],
        [50000, 95000, 65000, 52000, 75000, 80000],
        [124000, 228000, 160000, 136000, 190000, 210000]
    ):
        data_ws.append([m, u, p])

    # Categorical Data (Rows 8-12)
    products = ["Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Sugar"]
    data_ws.append(["Market"] + products)
    data_ws.append(["India", 62000, 4800, 21000, 25000])
    data_ws.append(["Philippines", 54000, 7000, 22000, 8000])
    data_ws.append(["United Kingdom", 46000, 5200, 11000, 14000])
    data_ws.append(["United States", 36000, 6300, 22000, 9000])

    # 4. Build Title Banner
    ws.merge_cells("A1:M3")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color="FFFFFF")
    header_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 5. Build Main Chart (Stacked Column)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Product"
    bar_chart.height = 12
    bar_chart.width = 16

    # References offset by data block sizes
    bar_data_ref = Reference(data_ws, min_col=2, min_row=8, max_col=5, max_row=12)
    bar_cats_ref = Reference(data_ws, min_col=1, min_row=9, max_row=12)
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)
    ws.add_chart(bar_chart, "C5")

    # 6. Build Secondary Side Charts (Line)
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.height = 6
    line1.width = 12
    line1.legend = None

    line1_data_ref = Reference(data_ws, min_col=2, min_row=1, max_row=7)
    line_cats_ref = Reference(data_ws, min_col=1, min_row=2, max_row=7)
    line1.add_data(line1_data_ref, titles_from_data=True)
    line1.set_categories(line_cats_ref)
    ws.add_chart(line1, "I5")

    line2 = LineChart()
    line2.title = "Profit by month"
    line2.height = 6
    line2.width = 12
    line2.legend = None

    line2_data_ref = Reference(data_ws, min_col=3, min_row=1, max_row=7)
    line2.add_data(line2_data_ref, titles_from_data=True)
    line2.set_categories(line_cats_ref)
    ws.add_chart(line2, "I15")

    # 7. Add Left-Rail Filter/Slicer Region Placeholder
    ws.merge_cells("A5:B20")
    slicer_ph = ws["A5"]
    slicer_ph.value = "Filter Panel\n(Insert Slicers Here)"
    slicer_ph.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    slicer_ph.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    slicer_ph.font = Font(color="595959", italic=True)

    border_side = Side(style="thin", color="D9D9D9")
    slicer_ph.border = Border(top=border_side, left=border_side, right=border_side, bottom=border_side)

    # Clean up structure spacing
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 12
