from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Sets up a clean, gridless dashboard canvas with a header banner and a sidebar area.
    """
    ws = wb.create_sheet(sheet_name)
    
    # 1. Turn off gridlines for a clean "app-like" look
    ws.sheet_view.showGridLines = False
    
    # 2. Theme definitions (fallback provided if external palette is not passed)
    palette = kwargs.get("palette", {
        "primary_bg": "203764",  # Dark Blue
        "primary_fg": "FFFFFF",  # White
        "sidebar_bg": "F2F2F2",  # Light Gray
        "text_muted": "595959"   # Dark Gray
    })
    
    header_fill = PatternFill(start_color=palette["primary_bg"], end_color=palette["primary_bg"], fill_type="solid")
    header_font = Font(name="Calibri", size=28, bold=True, color=palette["primary_fg"])
    sidebar_fill = PatternFill(start_color=palette["sidebar_bg"], end_color=palette["sidebar_bg"], fill_type="solid")
    sidebar_header_font = Font(name="Calibri", size=12, bold=True, color=palette["text_muted"])
    
    # 3. Create Title Banner (Rows 1-3, Columns A-R)
    # Merge a central area for the title, leaving A1:B3 unmerged for a potential logo
    ws.merge_cells("C1:R3")
    title_cell = ws["C1"]
    title_cell.value = title
    title_cell.font = header_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Apply fill to the entire banner strip
    for row in ws["A1:R3"]:
        for cell in row:
            cell.fill = header_fill
            
    # 4. Setup Sidebar (Columns A-B)
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 18
    
    # Fill sidebar background down to row 40
    for row in range(5, 41):
        for col in [1, 2]:
            ws.cell(row=row, column=col).fill = sidebar_fill
            
    # Sidebar Header
    ws.merge_cells("A5:B5")
    sidebar_title = ws["A5"]
    sidebar_title.value = "Filters & Controls"
    sidebar_title.font = sidebar_header_font
    sidebar_title.alignment = Alignment(horizontal="center", vertical="center")
    
    # 5. Setup Main Content Area (Columns C-R)
    # Widen columns to create a spacious grid for dropping in charts
    for col in range(3, 19):
        ws.column_dimensions[get_column_letter(col)].width = 14
