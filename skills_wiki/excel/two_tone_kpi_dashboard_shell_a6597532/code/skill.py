from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, subtitle: str, kpis: list, chart_data: dict, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a Sleek Two-Tone Dashboard Shell.
    
    Expected `kpis` format:
        [{"label": "TOTAL CALLS", "value": 16749, "format": "#,##0"}, ...]
        
    Expected `chart_data` format:
        {
            "headers": ["Month", "Calls Reached", "Deals Closed"],
            "rows": [["Jan", 301, 115], ["Feb", 311, 118], ...]
        }
    """
    ws = wb.create_sheet(sheet_name)
    
    # Hide native gridlines to create a blank canvas
    ws.sheet_view.showGridLines = False
    
    # Structural column spacing
    ws.column_dimensions['A'].width = 4
    for gap_col in ['E', 'I', 'M', 'Q']:
        ws.column_dimensions[gap_col].width = 4
    
    # Theme Extraction (Mocked for standalone execution)
    # In a full framework, load via: palette = get_theme(theme)
    dark_bg = "4B2E83"   # Deep purple header band
    light_bg = "F4F0F9"  # Pale purple body band
    accent = "FFC000"    # Gold accent
    card_bg = "FFFFFF"   # White cards
    text_main = "FFFFFF"
    text_kpi = "4B2E83"
    
    dark_fill = PatternFill("solid", fgColor=dark_bg)
    light_fill = PatternFill("solid", fgColor=light_bg)
    
    # Apply two-tone horizontal bands
    for row in range(1, 10):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = dark_fill
    for row in range(10, 40):
        for col in range(1, 25):
            ws.cell(row=row, column=col).fill = light_fill

    # Render Header Text
    ws["B2"] = title
    ws["B2"].font = Font(size=32, color=text_main, bold=True)
    ws["B3"] = subtitle
    ws["B3"].font = Font(size=14, color=accent)

    # Render Floating KPI Cards (Spanning rows 6-9)
    card_fill = PatternFill("solid", fgColor=card_bg)
    top_accent = Border(top=Side(style="thick", color=accent))
    
    col_idx = 2
    for kpi in kpis:
        # Paint card background and top accent
        for r in range(6, 10):
            for c in range(col_idx, col_idx + 3):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                if r == 6:
                    cell.border = top_accent
                    
        # Metric Value
        ws.merge_cells(start_row=7, start_column=col_idx, end_row=7, end_column=col_idx+2)
        val_cell = ws.cell(row=7, column=col_idx)
        val_cell.value = kpi.get("value", 0)
        val_cell.number_format = kpi.get("format", "General")
        val_cell.font = Font(size=20, color=text_kpi, bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Metric Label
        ws.merge_cells(start_row=8, start_column=col_idx, end_row=8, end_column=col_idx+2)
        lbl_cell = ws.cell(row=8, column=col_idx)
        lbl_cell.value = kpi.get("label", "")
        lbl_cell.font = Font(size=11, color="808080", bold=True)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        col_idx += 4  # Advance 3 columns for card + 1 gap column

    # Write Chart Data invisibly far to the right (Column AD / 30)
    data_c = 30
    data_r = 1
    headers = chart_data.get("headers", [])
    
    for i, h in enumerate(headers):
        ws.cell(row=data_r, column=data_c + i, value=h)
        
    rows = chart_data.get("rows", [])
    for r_idx, row_vals in enumerate(rows, start=1):
        for c_idx, val in enumerate(row_vals):
            ws.cell(row=data_r + r_idx, column=data_c + c_idx, value=val)
            
    # Add cleanly formatted Column Chart
    if rows and headers:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Conversion Pipeline"
        chart.legend.position = "t"
        
        # Remove visual clutter
        chart.y_axis.majorGridlines = None
        
        cats = Reference(ws, min_col=data_c, min_row=data_r+1, max_row=data_r+len(rows))
        data_ref = Reference(ws, min_col=data_c+1, min_row=data_r, max_col=data_c+len(headers)-1, max_row=data_r+len(rows))
        
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        
        ws.add_chart(chart, "B12")
        chart.width = 20
        chart.height = 12
