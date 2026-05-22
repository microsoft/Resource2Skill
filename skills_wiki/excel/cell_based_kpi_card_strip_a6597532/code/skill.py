from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment

def render(ws, anchor: str, *, kpi_data: list = None, theme: str = "purple_gold", **kwargs) -> None:
    if kpi_data is None:
        kpi_data = [
            {"label": "TOTAL CALLS", "value": 16749, "icon": "📞", "format": "#,##0"},
            {"label": "REACHED", "value": 3328, "icon": "🎯", "format": "#,##0"},
            {"label": "CLOSED", "value": 1203, "icon": "🏆", "format": "#,##0"},
            {"label": "DEAL VALUE", "value": 646979, "icon": "💰", "format": "$#,##0"}
        ]

    # Inline palette simulation for self-contained execution
    palettes = {
        "corporate_blue": {"accent": "2B579A", "card": "FFFFFF", "val": "000000", "lbl": "555555"},
        "purple_gold": {"accent": "FFC000", "card": "FFFFFF", "val": "333333", "lbl": "666666"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])
    
    start_row, start_col = coordinate_to_tuple(anchor)
    
    accent_fill = PatternFill("solid", fgColor=palette["accent"])
    card_fill = PatternFill("solid", fgColor=palette["card"])
    
    val_font = Font(color=palette["val"], size=16, bold=True)
    lbl_font = Font(color=palette["lbl"], size=10, bold=True)
    icon_font = Font(color="FFFFFF", size=18)
    
    center_align = Alignment(horizontal="center", vertical="center")
    
    curr_col = start_col
    for kpi in kpi_data:
        c_icon = curr_col
        c_text = curr_col + 1
        c_spacer = curr_col + 2
        
        # 1. Left Icon Bar (Merge 2 rows)
        ws.merge_cells(start_row=start_row, start_column=c_icon, end_row=start_row+1, end_column=c_icon)
        icon_cell = ws.cell(row=start_row, column=c_icon, value=kpi.get("icon", ""))
        icon_cell.font = icon_font
        icon_cell.alignment = center_align
        
        # Apply fill to all underlying merged cells to ensure stable Excel rendering
        ws.cell(row=start_row, column=c_icon).fill = accent_fill
        ws.cell(row=start_row+1, column=c_icon).fill = accent_fill
        
        # 2. Top Right: Value
        val_cell = ws.cell(row=start_row, column=c_text, value=kpi.get("value", ""))
        val_cell.font = val_font
        val_cell.alignment = center_align
        val_cell.fill = card_fill
        if "format" in kpi:
            val_cell.number_format = kpi["format"]
            
        # 3. Bottom Right: Label
        lbl_cell = ws.cell(row=start_row+1, column=c_text, value=kpi.get("label", ""))
        lbl_cell.font = lbl_font
        lbl_cell.alignment = center_align
        lbl_cell.fill = card_fill
        
        # Column sizing
        ws.column_dimensions[get_column_letter(c_icon)].width = 6
        ws.column_dimensions[get_column_letter(c_text)].width = 16
        ws.column_dimensions[get_column_letter(c_spacer)].width = 2
        
        curr_col += 3
        
    # Row sizing for the entire horizontal strip
    ws.row_dimensions[start_row].height = 24
    ws.row_dimensions[start_row+1].height = 16
