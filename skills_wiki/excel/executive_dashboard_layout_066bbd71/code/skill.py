def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    # 1. Define Theme Palette
    themes = {
        "corporate_blue": {"header_bg": "003366", "header_fg": "FFFFFF", "accent": "4F81BD"},
        "modern_dark": {"header_bg": "222222", "header_fg": "E0E0E0", "accent": "00A2E8"},
        "forest_green": {"header_bg": "2E4E3F", "header_fg": "FFFFFF", "accent": "6BAF92"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 2. Create Hidden Calculation Layer
    # Simulates the PivotTables that typically drive an interactive dashboard
    calc_ws = wb.create_sheet("Calc")
    calc_ws.sheet_state = 'hidden'

    # Table 1: Categorical Breakdown
    table1_data = [
        ["Market", "Chocolate Chip", "Sugar", "Oatmeal"],
        ["India", 62349, 25085, 21028],
        ["United Kingdom", 46530, 14620, 22005],
        ["United States", 36657, 9938, 22260],
        ["Malaysia", 46587, 20555, 17536],
        ["Philippines", 54618, 8313, 22005]
    ]
    for row in table1_data:
        calc_ws.append(row)
    
    calc_ws.append([]) # Blank row 7
    
    # Table 2 & 3: Time Series Data
    calc_ws.append(["Month", "Units Sold", "Profit"]) # Row 8
    trend_data = [
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    for row in trend_data:
        calc_ws.append(row)

    # 3. Setup Dashboard Canvas
    dash_ws = wb.active
    dash_ws.title = "Dashboard"
    
    # Turn off gridlines for a clean "App-like" appearance
    dash_ws.sheet_view.showGridLines = False
    
    # 4. Build Full-Width Header
    dash_ws.merge_cells("A1:N3")
    header_cell = dash_ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["header_fg"])
    header_cell.fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Fill remaining header background visually
    for col in range(1, 16):
        for row in range(1, 4):
            dash_ws.cell(row=row, column=col).fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")

    # 5. Build Slicer / Filter Placeholder Panel
    dash_ws.column_dimensions['A'].width = 18
    dash_ws.column_dimensions['B'].width = 4
    filter_header = dash_ws.cell(row=5, column=1, value="Filters / Slicers")
    filter_header.font = Font(bold=True, size=12)
    filter_header.border = Border(bottom=Side(style="thick", color=palette["accent"]))
    dash_ws.cell(row=6, column=1, value="[ Insert Slicers Here ]").font = Font(italic=True, color="7F7F7F")

    # 6. Create Stacked Bar Chart (Categorical)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Cookie Type"
    chart1.height = 10.5
    chart1.width = 14
    
    data1 = Reference(calc_ws, min_col=2, min_row=1, max_col=4, max_row=6)
    cats1 = Reference(calc_ws, min_col=1, min_row=2, max_row=6)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    dash_ws.add_chart(chart1, "C5")

    # 7. Create Line Chart 1 (Primary Metric Trend)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.height = 7.5
    chart2.width = 12
    chart2.legend = None  # Remove legend for cleaner look
    
    data2 = Reference(calc_ws, min_col=2, min_row=8, max_row=12)
    cats2 = Reference(calc_ws, min_col=1, min_row=9, max_row=12)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    dash_ws.add_chart(chart2, "I5")

    # 8. Create Line Chart 2 (Secondary Metric Trend)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.height = 7.5
    chart3.width = 12
    chart3.legend = None 
    
    data3 = Reference(calc_ws, min_col=3, min_row=8, max_row=12)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats2)
    dash_ws.add_chart(chart3, "I21")
