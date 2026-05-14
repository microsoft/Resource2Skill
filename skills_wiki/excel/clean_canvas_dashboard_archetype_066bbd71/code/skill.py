from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font

def render_workbook(wb: Workbook, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a clean, gridless dashboard canvas with a stacked column chart 
    and two line charts, backed by a separate data sheet.
    """
    # 1. Setup Architecture: Backend Data Sheet and Frontend Dashboard
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    ws_data = wb.create_sheet("Data")
    ws_dash = wb.create_sheet("Dashboard")
    
    # 2. Clean Canvas Setup
    ws_dash.sheet_view.showGridLines = False
    
    # Dashboard Header
    ws_dash["B2"] = title
    ws_dash["B2"].font = Font(size=24, bold=True, color="003366")  # Primary theme color placeholder
    
    # 3. Populate Backend Data
    # Matrix data for Stacked Column (Profit by Market & Cookie Type)
    data_mkt_cookie = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["Philippines", 54618, 7026, 22005],
        ["United Kingdom", 46530, 5220, 11497],
        ["United States", 36657, 6368, 22260]
    ]
    
    for r, row in enumerate(data_mkt_cookie, start=1):
        for c, val in enumerate(row, start=1):
            ws_data.cell(row=r, column=c, value=val)
            
    # Time-series data for Line Charts (Monthly Trend Data)
    data_monthly = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    
    start_row_monthly = 10
    for r, row in enumerate(data_monthly, start=start_row_monthly):
        for c, val in enumerate(row, start=1):
            ws_data.cell(row=r, column=c, value=val)
            
    # 4. Generate Charts
    
    # Chart A: Stacked Column (Left Side)
    chart_stacked = BarChart()
    chart_stacked.type = "col"
    chart_stacked.grouping = "stacked"
    chart_stacked.overlap = 100
    chart_stacked.title = "Profit by Market & Cookie Type"
    chart_stacked.height = 14
    chart_stacked.width = 16
    
    cats_stacked = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    data_stacked = Reference(ws_data, min_col=2, max_col=4, min_row=1, max_row=5)
    chart_stacked.add_data(data_stacked, titles_from_data=True)
    chart_stacked.set_categories(cats_stacked)
    
    # Chart B: Line Chart (Top Right)
    chart_line1 = LineChart()
    chart_line1.title = "Units sold each month"
    chart_line1.height = 7
    chart_line1.width = 14
    chart_line1.legend = None  # Clean interface: remove legend
    
    cats_line = Reference(ws_data, min_col=1, min_row=start_row_monthly+1, max_row=start_row_monthly+4)
    data_line1 = Reference(ws_data, min_col=2, max_col=2, min_row=start_row_monthly, max_row=start_row_monthly+4)
    chart_line1.add_data(data_line1, titles_from_data=True)
    chart_line1.set_categories(cats_line)
    
    # Chart C: Line Chart (Bottom Right)
    chart_line2 = LineChart()
    chart_line2.title = "Profit by month"
    chart_line2.height = 7
    chart_line2.width = 14
    chart_line2.legend = None  # Clean interface: remove legend
    
    data_line2 = Reference(ws_data, min_col=3, max_col=3, min_row=start_row_monthly, max_row=start_row_monthly+4)
    chart_line2.add_data(data_line2, titles_from_data=True)
    chart_line2.set_categories(cats_line)
    
    # 5. Position Elements on the Canvas
    ws_dash.add_chart(chart_stacked, "B5")
    ws_dash.add_chart(chart_line1, "K5")
    ws_dash.add_chart(chart_line2, "K19")
