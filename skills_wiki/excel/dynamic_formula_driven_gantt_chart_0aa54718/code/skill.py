import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from datetime import date, timedelta

def render_sheet(wb, sheet_name: str, *, title: str = "Project Tracker", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Define Palette (fallback theme mapping)
    colors = {
        "primary": "674EA7",       # Dark purple (Progress Bar)
        "secondary": "D9D2E9",     # Light purple (Base Bar)
        "success": "D9EAD3",       # Pale green (Completed Row)
        "muted": "EFEFEF",         # Light gray (Weekends)
        "alert": "CC0000",         # Red (Today indicator)
        "text_light": "FFFFFF",
        "border": "FFFFFF"
    }
    
    # Reusable styles
    header_fill = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    header_font = Font(color=colors["text_light"], bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Title setup
    ws["A2"] = title
    ws["A2"].font = Font(size=18, bold=True, color=colors["primary"])
    ws.sheet_view.showGridLines = False
    
    # 1. Setup Data Headers
    headers = ["Task", "Phase", "Owner", "Start", "Duration", "End", "Completed", "Progress"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    # 2. Insert Sample Data
    base_date = date.today()
    data = [
        ["Define project scope", "Planning", "Sarah", base_date, 5, 5],
        ["Stakeholder sign-off", "Planning", "Tom", base_date + timedelta(days=7), 3, 3],
        ["Wireframes & mockups", "Design", "Priya", base_date + timedelta(days=12), 10, 4],
        ["Backend development", "Development", "James", base_date + timedelta(days=15), 15, 0],
        ["QA & testing", "Testing", "Lisa", base_date + timedelta(days=35), 7, 0],
    ]
    
    for i, row_data in enumerate(data, start=6):
        ws.cell(row=i, column=1, value=row_data[0])
        ws.cell(row=i, column=2, value=row_data[1])
        ws.cell(row=i, column=3, value=row_data[2])
        
        # Start Date
        start_cell = ws.cell(row=i, column=4, value=row_data[3])
        start_cell.number_format = "dd-mmm-yy"
        
        # Duration
        ws.cell(row=i, column=5, value=row_data[4])
        
        # End Date (Calculated, skipping weekends)
        ws.cell(row=i, column=6, value=f"=WORKDAY.INTL(D{i}-1, E{i}, 1)")
        ws.cell(row=i, column=6).number_format = "dd-mmm-yy"
        
        # Completed Days
        ws.cell(row=i, column=7, value=row_data[5])
        
        # Progress %
        prog_cell = ws.cell(row=i, column=8, value=f"=G{i}/E{i}")
        prog_cell.number_format = "0%"

    # Format Data as Excel Table
    table_ref = f"A5:H{5 + len(data)}"
    table = Table(displayName="Project", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight14", showFirstColumn=False, showLastColumn=False, 
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    
    # 3. Setup Timeline Matrix (Dynamic Sequence)
    # The SEQUENCE array generates a horizontal list of serial dates
    ws["J5"] = '=SEQUENCE(1, MAX(Project[End])-MIN(Project[Start])+1, MIN(Project[Start]))'
    ws["J5"].font = header_font
    ws["J5"].fill = header_fill
    
    # Pre-format the expected spilled columns for the timeline
    for col in range(10, 60):  # Columns J through BG (50 days capacity)
        cell = ws.cell(row=5, column=col)
        cell.number_format = "dd\nmmm"
        cell.alignment = center_align
        cell.fill = header_fill
        cell.font = header_font
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 4.5
        
    # 4. Setup Conditional Formatting Rules
    tl_range = "J6:BG25"  # Application range for the Gantt Grid
    
    gray_fill = PatternFill(start_color=colors["muted"], end_color=colors["muted"], fill_type="solid")
    dark_fill = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    light_fill = PatternFill(start_color=colors["secondary"], end_color=colors["secondary"], fill_type="solid")
    green_fill = PatternFill(start_color=colors["success"], end_color=colors["success"], fill_type="solid")
    red_font = Font(color=colors["alert"], bold=True)
    
    # White horizontal borders to separate task rows visually
    white_edge = Side(style='thin', color=colors["border"])
    bar_border = Border(top=white_edge, bottom=white_edge)
    
    # Note: Rules are evaluated top-to-bottom in Excel. Add most specific first.
    
    # Rule A: Progress Bar (Dark Purple) - Tasks already completed
    rule_progress = Rule(type="expression", formula=['AND(J$5>=$D6, J$5<=WORKDAY.INTL($D6-1, $G6, 1), $G6>0)'], stopIfTrue=True)
    rule_progress.fill = dark_fill
    rule_progress.border = bar_border
    ws.conditional_formatting.add(tl_range, rule_progress)
    
    # Rule B: Base Bar (Light Purple) - Total duration span
    rule_base = Rule(type="expression", formula=['AND(J$5>=$D6, J$5<=$F6)'], stopIfTrue=True)
    rule_base.fill = light_fill
    rule_base.border = bar_border
    ws.conditional_formatting.add(tl_range, rule_base)
    
    # Rule C: Weekends (Gray vertical bands)
    rule_weekend = Rule(type="expression", formula=['WEEKDAY(J$5, 2)>=6'], stopIfTrue=True)
    rule_weekend.fill = gray_fill
    ws.conditional_formatting.add(tl_range, rule_weekend)
    
    # Rule D: Highlight Row when Task is 100% Complete (Applied to Data Columns)
    rule_complete_row = Rule(type="expression", formula=['$H6=1'], stopIfTrue=True)
    rule_complete_row.fill = green_fill
    ws.conditional_formatting.add("A6:H25", rule_complete_row)
    
    # Rule E: Highlight Today's Date in the Timeline Header
    rule_today = Rule(type="expression", formula=['J$5=TODAY()'], stopIfTrue=True)
    rule_today.font = red_font
    ws.conditional_formatting.add("J5:BG5", rule_today)
    
    # Adjust primary data column widths
    for col in ["A", "B", "C"]:
        ws.column_dimensions[col].width = 20
    for col in ["D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 11
