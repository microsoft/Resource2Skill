from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia Region", main_value: float = 369989, secondary_value: float = 0.05, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a 2x2 cell-based KPI card with a main metric and a secondary 'badge' metric.
    
    :param ws: openpyxl worksheet object.
    :param anchor: Top-left cell coordinate (e.g., 'B2').
    :param title: The title of the KPI card.
    :param main_value: The primary metric (e.g., Revenue).
    :param secondary_value: The secondary metric (e.g., Market Share).
    :param theme: The visual theme palette to use.
    """
    
    # Standard theme palette fallback
    themes = {
        "corporate_blue": {
            "primary": "1F4E78", 
            "secondary": "2F75B5", 
            "text_on_primary": "FFFFFF"
        },
        "executive_dark": {
            "primary": "262626", 
            "secondary": "595959", 
            "text_on_primary": "FFFFFF"
        }
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # Parse anchor
    col_str, row_str = coordinate_from_string(anchor)
    col = column_index_from_string(col_str)
    row = int(row_str)

    # Define cells
    cell_title = ws.cell(row=row, column=col)
    cell_main = ws.cell(row=row+1, column=col)
    cell_badge_top = ws.cell(row=row, column=col+1)
    cell_badge_bottom = ws.cell(row=row+1, column=col+1)

    # Inject values
    cell_title.value = title
    cell_main.value = main_value
    cell_badge_top.value = secondary_value

    # Merge right column for the badge
    ws.merge_cells(start_row=row, start_column=col+1, end_row=row+1, end_column=col+1)

    # Formatting rules
    font_title = Font(color=palette["text_on_primary"], size=11, bold=False)
    font_main = Font(color=palette["text_on_primary"], size=14, bold=True)
    font_badge = Font(color=palette["text_on_primary"], size=14, bold=True)
    
    fill_main = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    fill_badge = PatternFill(start_color=palette["secondary"], end_color=palette["secondary"], fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Apply to left column (Main Info)
    for c in [cell_title, cell_main]:
        c.fill = fill_main
        c.alignment = align_center
    cell_title.font = font_title
    cell_main.font = font_main
    cell_main.number_format = '"$"#,##0'

    # Apply to right column (Badge)
    for c in [cell_badge_top, cell_badge_bottom]:
        c.fill = fill_badge
        c.alignment = align_center
    cell_badge_top.font = font_badge
    cell_badge_top.number_format = '0%'

    # Adjust column widths & row heights to look like a proportional card
    ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions[get_column_letter(col+1)].width = 10
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row+1].height = 25

    # Create a Card Border
    thick_side = Side(style="medium", color=palette["primary"])
    
    cell_title.border = Border(top=thick_side, left=thick_side)
    cell_main.border = Border(bottom=thick_side, left=thick_side)
    cell_badge_top.border = Border(top=thick_side, right=thick_side, bottom=thick_side)
    # Because it's merged, the bottom border goes on the bottom-most cell of the merge
    cell_badge_bottom.border = Border(bottom=thick_side, right=thick_side)
