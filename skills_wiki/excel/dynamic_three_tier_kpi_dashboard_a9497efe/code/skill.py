import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_workbook(wb, *, title: str = "Monthly KPI Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Theme Palette Setup
    themes = {
        "corporate_blue": {"primary": "2F75B5", "accent": "ED7D31", "bg": "F2F2F2", "text": "FFFFFF"},
        "exec_dark": {"primary": "203764", "accent": "FFC000", "bg": "D9D9D9", "text": "FFFFFF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 2. Sheet Initialization
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    ws_staging = wb.create_sheet("Staging")
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False

    # 3. Populate Staging Data
    staging_data = [
        ["Category", "Jan-20", "Feb-20", "Mar-20"],
        ["DSO", 41, 54, 58],
        ["DSO Target", 45, 45, 45],
        ["Gross Margin", 0.40, 0.41, 0.38],
        ["GM Target", 0.38, 0.38, 0.38],
    ]
    for row in staging_data:
        ws_staging.append(row)
        
    for col in range(2, 5):
        ws_staging.cell(row=4, column=col).number_format = "0%"
        ws_staging.cell(row=5, column=col).number_format = "0%"

    # 4. Dashboard Header & Interactivity
    ws_dash.column_dimensions['A'].width = 3
    ws_dash.column_dimensions['B'].width = 16
    ws_dash.column_dimensions['C'].width = 16
    ws_dash.column_dimensions['D'].width = 16
    ws_dash.column_dimensions['E'].width = 16

    ws_dash["B3"] = "For the month of"
    ws_dash["B3"].font = Font(bold=True)
    ws_dash["C3"] = "Jan-20"
    ws_dash["C3"].fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    dv = DataValidation(type="list", formula1='"Jan-20,Feb-20,Mar-20"', allow_blank=False)
    ws_dash.add_data_validation(dv)
    dv.add(ws_dash["C3"])

    ws_dash.merge_cells("B1:E1")
    ws_dash["B1"] = title
    ws_dash["B1"].font = Font(size=20, bold=True, color=palette["text"])
    ws_dash["B1"].fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    ws_dash["B1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 35

    ws_dash.merge_cells("B5:E5")
    ws_dash["B5"] = "Working Capital & Sales KPIs"
    ws_dash["B5"].font = Font(size=14, bold=True, color=palette["text"])
    ws_dash["B5"].fill = PatternFill(start_color=palette["accent"], end_color=palette["accent"], fill_type="solid")
    ws_dash["B5"].alignment = Alignment(horizontal="center", vertical="center")

    # 5. KPI Block Generator
    def build_kpi_block(start_col: int, start_row: int, kpi_title: str, metric_name: str, target_name: str, is_lower_better: bool = True, is_pct: bool = False):
        c1 = get_column_letter(start_col)
        c2 = get_column_letter(start_col + 1)
        
        # Title Ribbon
        ws_dash.merge_cells(f"{c1}{start_row}:{c2}{start_row}")
        title_cell = ws_dash[f"{c1}{start_row}"]
        title_cell.value = kpi_title
        title_cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center")
        title_cell.font = Font(bold=True)
        
        # Value (Dynamic INDEX/MATCH)
        val_row = start_row + 1
        ws_dash.merge_cells(f"{c1}{val_row}:{c2}{val_row}")
        val_cell = ws_dash[f"{c1}{val_row}"]
        val_cell.value = f'=INDEX(Staging!$B$2:$D$5, MATCH("{metric_name}", Staging!$A$2:$A$5, 0), MATCH($C$3, Staging!$B$1:$D$1, 0))'
        val_cell.font = Font(size=24, bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.row_dimensions[val_row].height = 35
        if is_pct:
            val_cell.number_format = "0%"
            
        # Target Subtext
        tgt_row = start_row + 2
        lbl_cell = ws_dash[f"{c1}{tgt_row}"]
        lbl_cell.value = "Vs. Target"
        lbl_cell.font = Font(size=10, italic=True)
        
        tgt_cell = ws_dash[f"{c2}{tgt_row}"]
        tgt_cell.value = f'=INDEX(Staging!$B$2:$D$5, MATCH("{target_name}", Staging!$A$2:$A$5, 0), MATCH($C$3, Staging!$B$1:$D$1, 0))'
        tgt_cell.font = Font(size=10, bold=True)
        tgt_cell.alignment = Alignment(horizontal="right")
        if is_pct:
            tgt_cell.number_format = "0%"
            
        # Draw Borders
        thin = Side(style="thin", color="BFBFBF")
        box_border = Border(top=thin, left=thin, right=thin, bottom=thin)
        for r in range(start_row, start_row + 3):
            for c in range(start_col, start_col + 2):
                ws_dash.cell(row=r, column=c).border = box_border
                
        # Conditional Formatting Array
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        val_range = f"{c1}{val_row}:{c2}{val_row}"
        tgt_ref = f"${c2}${tgt_row}"
        
        if is_lower_better:
            ws_dash.conditional_formatting.add(val_range, CellIsRule(operator='lessThanOrEqual', formula=[tgt_ref], fill=green_fill))
            ws_dash.conditional_formatting.add(val_range, CellIsRule(operator='greaterThan', formula=[tgt_ref], fill=red_fill))
        else:
            ws_dash.conditional_formatting.add(val_range, CellIsRule(operator='greaterThanOrEqual', formula=[tgt_ref], fill=green_fill))
            ws_dash.conditional_formatting.add(val_range, CellIsRule(operator='lessThan', formula=[tgt_ref], fill=red_fill))

    # 6. Render the KPI Blocks
    # Block 1: DSO (Lower is Better)
    build_kpi_block(start_col=2, start_row=6, kpi_title="DSO (Days)", metric_name="DSO", target_name="DSO Target", is_lower_better=True)
    
    # Block 2: Gross Margin (Higher is Better)
    build_kpi_block(start_col=4, start_row=6, kpi_title="Gross Margin", metric_name="Gross Margin", target_name="GM Target", is_lower_better=False, is_pct=True)
