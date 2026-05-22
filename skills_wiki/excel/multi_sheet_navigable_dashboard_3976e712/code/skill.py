def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    # Modern dashboard color palette
    bg_color = "F3F4F6"         # Light gray background
    sidebar_bg = "1E3A8A"       # Dark blue primary
    sidebar_active = "3B82F6"   # Light blue accent
    text_light = "FFFFFF"
    text_dark = "111827"
    text_muted = "6B7280"
    card_bg = "FFFFFF"
    border_color = "E5E7EB"

    sheets = ["Dashboard", "Data Inputs", "Settings"]
    
    # Initialize sheet structure
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    for s in sheets:
        wb.create_sheet(s)

    def draw_card(ws, min_row, min_col, max_row, max_col):
        """Helper to create a unified white 'card' with a subtle border."""
        fill = PatternFill("solid", fgColor=card_bg)
        edge = Side(style='thin', color=border_color)
        
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill
                
                # Apply borders only to the outer edges of the bounding box
                b_top = edge if r == min_row else None
                b_bottom = edge if r == max_row else None
                b_left = edge if c == min_col else None
                b_right = edge if c == max_col else None
                cell.border = Border(top=b_top, bottom=b_bottom, left=b_left, right=b_right)

    # Apply global layout to each sheet
    for sheet_name in sheets:
        ws = wb[sheet_name]
        ws.sheet_view.showGridLines = False
        
        # Base Dimensions
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 3
        for col_letter in "CDEFGH":
            ws.column_dimensions[col_letter].width = 16

        # Apply global light gray background
        light_fill = PatternFill("solid", fgColor=bg_color)
        for r in range(1, 40):
            for c in range(2, 15):
                ws.cell(row=r, column=c).fill = light_fill

        # Paint Sidebar
        sb_fill = PatternFill("solid", fgColor=sidebar_bg)
        for r in range(1, 40):
            ws.cell(row=r, column=1).fill = sb_fill

        # Sidebar Brand / App Title
        logo_cell = ws.cell(row=2, column=1, value="ACME Corp")
        logo_cell.font = Font(color=text_light, bold=True, size=14)
        logo_cell.alignment = Alignment(horizontal="center")

        # Sidebar Navigation Menu
        start_row = 6
        active_fill = PatternFill("solid", fgColor=sidebar_active)
        link_font = Font(color=text_light, underline="single")
        active_font = Font(color=text_light, bold=True)
        
        for idx, s_name in enumerate(sheets):
            r = start_row + (idx * 3)
            cell = ws.cell(row=r, column=1, value=s_name)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            if sheet_name == s_name:
                # Active state styling
                cell.fill = active_fill
                cell.font = active_font
            else:
                # Inactive link styling
                cell.font = link_font
                cell.hyperlink = f"#'{s_name}'!A1"

        # Sheet Header
        title_cell = ws.cell(row=3, column=3, value=f"{title} - {sheet_name}")
        title_cell.font = Font(size=18, bold=True, color=text_dark)

        # Render specific structural cards based on the active view
        if sheet_name == "Dashboard":
            # Top row: KPI Cards
            draw_card(ws, 5, 3, 8, 4)
            ws.cell(row=6, column=3, value=" Revenue").font = Font(color=text_muted, bold=True)
            ws.cell(row=7, column=3, value=" $1.2M").font = Font(size=16, bold=True, color=text_dark)

            draw_card(ws, 5, 5, 8, 6)
            ws.cell(row=6, column=5, value=" Active Users").font = Font(color=text_muted, bold=True)
            ws.cell(row=7, column=5, value=" 45,200").font = Font(size=16, bold=True, color=text_dark)

            draw_card(ws, 5, 7, 8, 8)
            ws.cell(row=6, column=7, value=" Conversion").font = Font(color=text_muted, bold=True)
            ws.cell(row=7, column=7, value=" 4.8%").font = Font(size=16, bold=True, color=text_dark)

            # Main content area: Trend Chart Block
            draw_card(ws, 10, 3, 20, 8)
            ws.cell(row=11, column=3, value=" Monthly Performance Trend").font = Font(color=text_muted, bold=True)
            
        elif sheet_name == "Data Inputs":
            draw_card(ws, 5, 3, 20, 8)
            ws.cell(row=6, column=3, value=" Paste Raw Export Here").font = Font(color=text_muted, bold=True)
            
        elif sheet_name == "Settings":
            draw_card(ws, 5, 3, 10, 8)
            ws.cell(row=6, column=3, value=" Configuration Toggles").font = Font(color=text_muted, bold=True)

    # Focus user on the Dashboard on launch
    wb.active = 0
