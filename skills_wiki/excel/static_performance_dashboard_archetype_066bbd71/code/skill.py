def render_workbook(wb, *, title: str = "Performance Dashboard", data: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    from collections import defaultdict
    import random

    themes = {
        "corporate_blue": {"primary": "003366", "accent": "4F81BD"},
        "midnight": {"primary": "1F497D", "accent": "8064A2"},
        "emerald": {"primary": "006633", "accent": "9BBB59"}
    }
    palette = themes.get(theme, themes["corporate_blue"])
    primary_color = palette["primary"]

    # Generate dummy data if none is provided
    if not data:
        data = []
        markets = ["India", "USA", "UK", "Philippines"]
        products = ["Chocolate Chip", "Fortune Cookie", "Sugar", "Oatmeal Raisin"]
        months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        
        random.seed(42)
        for m_idx, month in enumerate(months):
            for market in markets:
                for product in products:
                    profit = random.randint(1000, 10000)
                    units = random.randint(100, 1000)
                    data.append({
                        "Month": month,
                        "Market": market,
                        "Product": product,
                        "Profit": profit,
                        "Units": units
                    })

    # Python-side Data Aggregations (mimicking Pivot Tables)
    market_product_profit = defaultdict(lambda: defaultdict(int))
    month_units = defaultdict(int)
    month_profit = defaultdict(int)
    
    for row in data:
        market_product_profit[row["Market"]][row["Product"]] += row["Profit"]
        month_units[row["Month"]] += row["Units"]
        month_profit[row["Month"]] += row["Profit"]

    # --- Chart Data Sheet (Hidden) ---
    ws_data = wb.create_sheet("ChartData")
    
    # 1. Market x Product
    markets_sorted = sorted(list(market_product_profit.keys()))
    products_sorted = sorted(list({p for m in markets_sorted for p in market_product_profit[m].keys()}))
    
    ws_data.append(["Market"] + products_sorted)
    for m in markets_sorted:
        row_vals = [m] + [market_product_profit[m][p] for p in products_sorted]
        ws_data.append(row_vals)
    table1_end = ws_data.max_row

    # 2. Monthly Trends
    ws_data.append([]) # Empty row separator
    ws_data.append(["Month", "Units", "Profit"])
    table2_start = ws_data.max_row
    
    months_order = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    for m in months_order:
        ws_data.append([m, month_units.get(m, 0), month_profit.get(m, 0)])
    table2_end = ws_data.max_row
    
    ws_data.sheet_state = 'hidden'

    # --- Dashboard Sheet ---
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    # Header styling
    ws_dash.merge_cells("A1:O2")
    header_cell = ws_dash["A1"]
    header_cell.value = "  " + title
    header_cell.font = Font(color="FFFFFF", size=20, bold=True)
    header_cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_cell.alignment = Alignment(horizontal="left", vertical="center")

    # Chart 1: Profit by Market & Product (Stacked Bar)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Product"
    bar_chart.width = 16
    bar_chart.height = 12
    if bar_chart.legend:
        bar_chart.legend.position = 'b'
    
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=table1_end)
    data1 = Reference(ws_data, min_col=2, max_col=1+len(products_sorted), min_row=1, max_row=table1_end)
    bar_chart.add_data(data1, titles_from_data=True)
    bar_chart.set_categories(cats1)
    ws_dash.add_chart(bar_chart, "B4")

    # Chart 2: Units Sold each Month (Line)
    line1 = LineChart()
    line1.title = "Units Sold each Month"
    line1.width = 14
    line1.height = 6
    line1.legend = None
    
    cats2 = Reference(ws_data, min_col=1, min_row=table2_start+1, max_row=table2_end)
    data2 = Reference(ws_data, min_col=2, min_row=table2_start, max_row=table2_end)
    line1.add_data(data2, titles_from_data=True)
    line1.set_categories(cats2)
    ws_dash.add_chart(line1, "J4")

    # Chart 3: Profit by Month (Line)
    line2 = LineChart()
    line2.title = "Profit by Month"
    line2.width = 14
    line2.height = 6
    line2.legend = None
    
    data3 = Reference(ws_data, min_col=3, min_row=table2_start, max_row=table2_end)
    line2.add_data(data3, titles_from_data=True)
    line2.set_categories(cats2)
    ws_dash.add_chart(line2, "J16")
