from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, title: str, subtitle: str, main_value: str, secondary_value: str = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a stylized KPI Card using cell formatting to simulate a dashboard shape widget.
    
    :param ws: openpyxl worksheet object
    :param anchor: Top-left cell coordinate (e.g., 'B2')
    :param title: Main card heading (e.g., 'Asia')
    :param subtitle: Metric label (e.g., 'Revenue')
    :param main_value: Primary metric value (e.g., '$369,989')
    :param secondary_value: Optional secondary metric (e.g., '5%') for the corner badge
    """
    row, col = coordinate_to_tuple(anchor)
    
    # Theme palette definition
    themes = {
        "corporate_blue": {"bg": "002060", "fg": "FFFFFF", "badge_bg": "4F81BD", "badge_fg": "FFFFFF", "border": "000000"},
        "dark_mode": {"bg": "1E1E1E", "fg": "FFFFFF", "badge_bg": "333333", "badge_fg": "4CAF50", "border": "444444"},
        "light_modern": {"bg": "F3F4F6", "fg": "111827", "badge_bg": "2563EB", "badge_fg": "FFFFFF", "border": "D1D5DB"}
    }
    palette = themes.get(theme, themes["corporate_blue"])
    
    # Define Styles
    bg_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    badge_fill = PatternFill(start_color=palette["badge_bg"], end_color=palette["badge_bg"], fill_type="solid")
    
    title_font = Font(color=palette["fg"], size=14, bold=True)
    subtitle_font = Font(color=palette["fg"], size=11)
    main_val_font = Font(color=palette["fg"], size=16, bold=True)
    badge_font = Font(color=palette["badge_fg"], size=12, bold=True)
    
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Configure grid dimensions to look like a card
    ws.column_dimensions[get_column_letter(col)].width = 12
    ws.column_dimensions[get_column_letter(col+1)].width = 12
    ws.column_dimensions[get_column_letter(col+2)].width = 10
    
    ws.row_dimensions[row].height = 25
    ws.row_dimensions[row+1].height = 18
    ws.row_dimensions[row+2].height = 30
    
    # Set primary cell values
    title_cell = ws.cell(row=row, column=col, value=title)
    subtitle_cell = ws.cell(row=row+1, column=col, value=subtitle)
    main_val_cell = ws.cell(row=row+2, column=col, value=main_value)
    
    # Apply merging based on whether a secondary badge is included
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+2)
    ws.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+2)
    
    if secondary_value:
        ws.merge_cells(start_row=row+2, start_column=col, end_row=row+2, end_column=col+1)
        badge_cell = ws.cell(row=row+2, column=col+2, value=secondary_value)
        badge_cell.fill = badge_fill
        badge_cell.font = badge_font
        badge_cell.alignment = center_align
    else:
        ws.merge_cells(start_row=row+2, start_column=col, end_row=row+2, end_column=col+2)
        
    # Apply baseline formatting to the anchor cells
    title_cell.font = title_font
    title_cell.alignment = center_align
    
    subtitle_cell.font = subtitle_font
    subtitle_cell.alignment = center_align
    
    main_val_cell.font = main_val_font
    main_val_cell.alignment = center_align

    # Apply structural background fill and outer borders
    thin_border = Side(border_style="thin", color=palette["border"])
    
    for r in range(row, row+3):
        for c in range(col, col+3):
            cell = ws.cell(row=r, column=c)
            
            # Skip fill override if this is the secondary badge cell
            if not (r == row+2 and c == col+2 and secondary_value):
                cell.fill = bg_fill
                
            # Apply dynamic outline border to the entire 3x3 widget
            cell.border = Border(
                top=thin_border if r == row else None,
                bottom=thin_border if r == row+2 else None,
                left=thin_border if c == col else None,
                right=thin_border if c == col+2 else None
            )
