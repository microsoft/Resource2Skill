from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a modern, application-like dashboard shell using cell formatting 
    to create a sidebar and floating 'cards' for visual placement.
    """
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        
    # Theme overrides / defaults
    bg_color = "F3F4F6"      # Light gray canvas
    card_bg = "FFFFFF"       # White cards
    sidebar_bg = "1F4E78"    # Dark blue sidebar
    text_main = "333333"     # Dark gray text
    text_sidebar = "FFFFFF"  # White text for sidebar
    border_color = "D1D5DB"  # Soft gray border
    
    # 1. Paint the global canvas background
    fill_bg = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=35, min_col=1, max_col=16):
        for cell in row:
            cell.fill = fill_bg
            
    # 2. Setup the Left Navigation Sidebar (Col A)
    ws.column_dimensions['A'].width = 8
    fill_sidebar = PatternFill(start_color=sidebar_bg, end_color=sidebar_bg, fill_type="solid")
    for row in range(1, 36):
        ws.cell(row=row, column=1).fill = fill_sidebar
        
    # Add fake navigation icons/text in sidebar
    nav_font = Font(color=text_sidebar, bold=True)
    nav_align = Alignment(horizontal="center", vertical="center")
    nav_items = ["Dash", "Data", "Mail", "Help"]
    for i, item in enumerate(nav_items):
        cell = ws.cell(row=5 + i*4, column=1, value=item)
        cell.font = nav_font
        cell.alignment = nav_align

    # 3. Add Dashboard Headers
    ws.column_dimensions['B'].width = 3 # Left gutter
    
    title_cell = ws.cell(row=2, column=3, value=title)
    title_cell.font = Font(color=text_main, size=20, bold=True)
    ws.row_dimensions[2].height = 28
    
    subtitle_cell = ws.cell(row=3, column=3, value="Figures in USD (Millions)")
    subtitle_cell.font = Font(color="7F8C8D", size=11, italic=True)
    
    # 4. Card Drawing Helper
    def draw_card(min_col, min_row, max_col, max_row, card_title):
        fill_card = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
        side_style = Side(style='thin', color=border_color)
        
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_card
                
                # Apply borders only to the outer edges of the range
                cell.border = Border(
                    left=side_style if c == min_col else None,
                    right=side_style if c == max_col else None,
                    top=side_style if r == min_row else None,
                    bottom=side_style if r == max_row else None
                )
        
        # Add Card Title in the top-left of the bounding box
        tc = ws.cell(row=min_row, column=min_col, value=f"  {card_title}")
        tc.font = Font(color=text_main, size=12, bold=True)
        tc.alignment = Alignment(vertical="center")
        ws.row_dimensions[min_row].height = 20
        
    # Standardize card columns widths
    for c in [3, 4, 5, 7, 8, 9, 11, 12, 13, 14]:
        ws.column_dimensions[get_column_letter(c)].width = 11
        
    # Gutters between cards
    ws.column_dimensions['F'].width = 3
    ws.column_dimensions['J'].width = 3
    
    # 5. Draw the Card Layout
    # Top Row (3 KPI Cards)
    draw_card(3, 5, 5, 10, "Sales")
    draw_card(7, 5, 9, 10, "Profit")
    draw_card(11, 5, 14, 10, "Customers")
    
    # Bottom Row (2 Large Chart Cards)
    draw_card(3, 12, 9, 24, "2021-2022 Sales Trend")
    draw_card(11, 12, 14, 24, "Sales by Country")
    
    # 6. Final Polish
    ws.sheet_view.showGridLines = False
