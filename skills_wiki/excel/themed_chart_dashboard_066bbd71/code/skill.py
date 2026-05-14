def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # 1. Theme Palette Fallbacks
    palettes = {
        "corporate_blue": {"primary": "2F5597", "text_light": "FFFFFF"},
        "emerald": {"primary": "005A36", "text_light": "FFFFFF"},
        "dark_slate": {"primary": "2D3436", "text_light": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # 2. Setup Data Sheet
    ws_data = wb.active
    ws_data.title = "ChartData"

    # Mock pre-aggregated multidimensional data
    data1 = [
        ["Market", "Fortune Cookie", "Sugar", "Snickerdoodle"],
        ["India", 62349, 18560, 25085],
        ["Philippines", 54618, 14947, 8313],
        ["United States", 36657, 10633, 20555]
    ]

    # Mock temporal trend data
    data2 = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]

    # Write data blocks to specific coordinate ranges
    for row in data1:
        ws_data.append(row)
        
    for i, row in enumerate(data2):
        ws_data.cell(row=i+1, column=6, value=row[0])
        ws_data.cell(row=i+1, column=7, value=row[1])
        ws_data.cell(row=i+1, column=8, value=row[2])

    ws_data.sheet_state = "hidden"

    # 3. Setup Dashboard Sheet
    ws_dash = wb.create_sheet("Dashboard")
    wb.active = ws_dash
    ws_dash.sheet_view.showGridLines = False

    # 4. Thematic Header
    ws_dash.merge_cells("B2:R3")
    ws_dash.row_dimensions[2].height = 20
    ws_dash.row_dimensions[3].height = 20
    
    header_cell = ws_dash["B2"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["text_light"])
    header_cell.fill = PatternFill("solid", fgColor=palette["primary"])
    header_cell.alignment = Alignment(vertical="center", horizontal="left")

    # 5. Charts Configuration
    
    # Chart 1: Stacked Bar (Left Column, Large)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "stacked"
    chart1.overlap = 100
    chart1.title = "Profit by Market & Product"
    chart1.width = 15.0
    chart1.height = 13.0

    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=4)
    data1_ref = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=4)
    chart1.add_data(data1_ref, titles_from_data=True)
    chart1.set_categories(cats1)
    ws_dash.add_chart(chart1, "B5")

    # Chart 2: Line (Right Column, Top)
    chart2 = LineChart()
    chart2.title = "Units sold each month"
    chart2.width = 12.0
    chart2.height = 6.0
    chart2.legend = None

    cats2 = Reference(ws_data, min_col=6, min_row=2, max_row=5)
    data2_ref = Reference(ws_data, min_col=7, min_row=1, max_row=5)
    chart2.add_data(data2_ref, titles_from_data=True)
    chart2.set_categories(cats2)
    ws_dash.add_chart(chart2, "K5")

    # Chart 3: Line (Right Column, Bottom)
    chart3 = LineChart()
    chart3.title = "Profit by month"
    chart3.width = 12.0
    chart3.height = 6.0
    chart3.legend = None

    data3_ref = Reference(ws_data, min_col=8, min_row=1, max_row=5)
    chart3.add_data(data3_ref, titles_from_data=True)
    chart3.set_categories(cats2)
    ws_dash.add_chart(chart3, "K17")
