from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Fallback palette simulating a dark dashboard theme
    colors = {
        "bg_base": "2A2A2A",     # Deep gray/black canvas
        "bg_card": "3D3D3D",     # Lighter gray for component cards
        "text_main": "FFFFFF",   # White text
        "accent": "FF9900",      # Bright orange accent for active elements
        "inactive": "555555"     # Dimmed gray for inactive elements
    }
    
    fill_base = PatternFill(start_color=colors["bg_base"], end_color=colors["bg_base"], fill_type="solid")
    fill_card = PatternFill(start_color=colors["bg_card"], end_color=colors["bg_card"], fill_type="solid")
    fill_btn_active = PatternFill(start_color=colors["accent"], end_color=colors["accent"], fill_type="solid")
    fill_btn_inactive = PatternFill(start_color=colors["inactive"], end_color=colors["inactive"], fill_type="solid")
    
    font_title = Font(color=colors["text_main"], size=24, bold=True)
    font_btn_active = Font(color=colors["bg_base"], size=12, bold=True)
    font_btn_inactive = Font(color=colors["text_main"], size=12, bold=True)
    
    # 1. Disable Gridlines for a clean application look
    ws.sheet_view.showGridLines = False
    
    # Define the bounded canvas size
    max_col = 20  # Up to column T
    max_row = 40  # Up to row 40
    
    # 2. Paint the base canvas
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row=row, column=col).fill = fill_base
            
    # 3. Add Dashboard Title
    title_cell = ws["B2"]
    title_cell.value = title
    title_cell.font = font_title
    
    # 4. Create "Card" zones for hosting charts/KPIs
    card_zones = [
        ("B6", "H20"),   # Top-left card
        ("J6", "S20"),   # Top-right card
        ("B22", "S38")   # Bottom full-width card
    ]
    
    for start_cell, end_cell in card_zones:
        start_col = ws[start_cell].column
        start_row = ws[start_cell].row
        end_col = ws[end_cell].column
        end_row = ws[end_cell].row
        
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                ws.cell(row=r, column=c).fill = fill_card
                
    # 5. Add Navigation Tabs (Top Right)
    # Active Tab
    ws.merge_cells("P2:Q3")
    btn_active = ws["P2"]
    btn_active.value = "Sales"
    btn_active.fill = fill_btn_active
    btn_active.font = font_btn_active
    btn_active.alignment = Alignment(horizontal="center", vertical="center")
    
    # Inactive Tab (with hyperlink placeholder)
    ws.merge_cells("R2:S3")
    btn_inactive = ws["R2"]
    btn_inactive.value = "Shipping"
    btn_inactive.hyperlink = "#'Shipping'!A1"  # Intra-workbook link
    btn_inactive.fill = fill_btn_inactive
    btn_inactive.font = font_btn_inactive
    btn_inactive.alignment = Alignment(horizontal="center", vertical="center")
    
    # 6. Restrict the scrolling area by hiding unused rows and columns
    # We hide a generous buffer around the canvas to prevent scrolling into white space
    for col_idx in range(max_col + 1, max_col + 50):
        ws.column_dimensions[get_column_letter(col_idx)].hidden = True
        
    for row_idx in range(max_row + 1, max_row + 100):
        ws.row_dimensions[row_idx].hidden = True
