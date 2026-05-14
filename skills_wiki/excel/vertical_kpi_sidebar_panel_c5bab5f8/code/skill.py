from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import column_index_from_string, coordinate_from_string

def render(ws, anchor: str, kpis: list[dict] = None, *, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a vertical KPI sidebar panel, mimicking a modern web app sidebar.
    
    :param ws: openpyxl Worksheet object.
    :param anchor: Top-left cell coordinate (e.g., 'A1').
    :param kpis: List of dicts with 'label', 'value', and 'format'.
    :param theme: Theme string for color palette selection.
    """
    # Self-contained palette definitions (mimics theme loader behavior)
    palettes = {
        "corporate_blue": {"bg": "1A365D", "val": "FFFFFF", "lbl": "A0AEC0"},
        "viva_green": {"bg": "1B4332", "val": "FFFFFF", "lbl": "9AE6B4"}, # Dashboard dark green
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    if kpis is None:
        kpis = [
            {"label": "🛒 Orders", "value": 2400, "format": "#,##0"},
            {"label": "📦 Quantity", "value": 11997, "format": "#,##0"},
            {"label": "💰 Amount", "value": 649019.8, "format": '$#,##0.0,"k"'},
            {"label": "⭐ Avg. Rating", "value": 4.0, "format": "0.0"},
            {"label": "🚚 Days to Deliver", "value": 2.3, "format": "0.0"}
        ]

    col_str, start_row = coordinate_from_string(anchor)
    col_idx = column_index_from_string(col_str)

    # Set sidebar column width to simulate a wide UI panel
    ws.column_dimensions[col_str].width = 24

    bg_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    val_font = Font(name="Calibri", size=20, bold=True, color=palette["val"])
    lbl_font = Font(name="Calibri", size=11, bold=False, color=palette["lbl"])
    center_align = Alignment(horizontal="center", vertical="center")

    # Paint the continuous background fill for the entire sidebar length
    # Adjust the range size depending on the number of KPIs to ensure the panel reaches the screen bottom
    panel_height = max(35, len(kpis) * 5)
    for r in range(start_row, start_row + panel_height):
        ws.cell(row=r, column=col_idx).fill = bg_fill

    current_row = start_row + 2  # Leave top padding

    for kpi in kpis:
        # Render the Metric Value (Large, Bold)
        val_cell = ws.cell(row=current_row, column=col_idx)
        val_cell.value = kpi.get("value", 0)
        if "format" in kpi:
            val_cell.number_format = kpi["format"]
        val_cell.font = val_font
        val_cell.alignment = center_align

        # Render the Metric Label (Small, Muted, with Emojis)
        lbl_cell = ws.cell(row=current_row + 1, column=col_idx)
        lbl_cell.value = kpi.get("label", "")
        lbl_cell.font = lbl_font
        lbl_cell.alignment = center_align

        # Increment by 4 to leave generous whitespace between metric blocks
        current_row += 4
