from openpyxl.styles import PatternFill, Font, Border, Side
from openpyxl.chart import LineChart, DoughnutChart, RadarChart, Reference, Series
from openpyxl.chart.series import DataPoint
from openpyxl.chart.marker import Marker

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Theme Palette (Fallback simulated)
    palette = {
        "corporate_blue": {"primary": "1F3864", "accent": "4472C4", "bg": "F5F6F8", "card": "FFFFFF", "text": "333333", "border": "E0E0E0"}
    }.get(theme, {"primary": "1F3864", "accent": "4472C4", "bg": "F5F6F8", "card": "FFFFFF", "text": "333333", "border": "E0E0E0"})
    
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # 2. Setup Hidden Semantic Data Layer
    data_ws = wb.create_sheet(f"{sheet_name}_Data")
    data_ws.sheet_state = 'hidden'
    
    trend_data = [
        ["Month", "2021", "2022"],
        ["Jan", 201.9, 215.3], ["Feb", 204.2, 217.6], ["Mar", 198.6, 220.1],
        ["Apr", 199.2, 206.4], ["May", 206.4, 204.3], ["Jun", 195.3, 203.0]
    ]
    for r in trend_data: data_ws.append(r) # Rows 1 to 7
    
    kpi_data = [
        ["KPI", "Actual", "Remainder", "Target"],
        ["Sales", 2544, 456, 3000],
        ["Profit", 890, 110, 1000],
        ["Customers", 87, 13, 100]
    ]
    for r in kpi_data: data_ws.append(r) # Rows 8 to 11
    
    radar_data = [
        ["Factor", "Score"],
        ["Speed", 54], ["Quality", 86], ["Hygiene", 91], 
        ["Service", 53], ["Availability", 95]
    ]
    for r in radar_data: data_ws.append(r) # Rows 12 to 17
    
    # 3. Canvas & Layout Construction
    bg_fill = PatternFill("solid", fgColor=palette["bg"])
    for row in ws.iter_rows(min_row=1, max_row=30, min_col=2, max_col=17):
        for cell in row:
            cell.fill = bg_fill
            
    # Sidebar
    sidebar_fill = PatternFill("solid", fgColor=palette["primary"])
    for row in range(1, 31):
        ws.cell(row=row, column=1).fill = sidebar_fill
    ws.column_dimensions['A'].width = 8
    
    ws.cell(row=2, column=3, value=title).font = Font(size=18, bold=True, color=palette["primary"])
    
    # Helper: Grid-Based UI Cards
    def create_card(min_col, min_row, max_col, max_row):
        fill = PatternFill("solid", fgColor=palette["card"])
        side = Side(style="thin", color=palette["border"])
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill
                # Paint borders only on the perimeter to simulate a continuous shape
                top = side if r == min_row else None
                bottom = side if r == max_row else None
                left = side if c == min_col else None
                right = side if c == max_col else None
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)

    # 4. Build Top KPI Cards with Progress Doughnuts
    kpi_anchors = [
        {"col": 3, "kpi_row": 9, "chart_col": "E", "label": "Sales", "val": "$2,544"},
        {"col": 8, "kpi_row": 10, "chart_col": "J", "label": "Profit", "val": "$890"},
        {"col": 13, "kpi_row": 11, "chart_col": "O", "label": "Customers", "val": "87.0"}
    ]
    
    for anchor in kpi_anchors:
        c_idx = anchor["col"]
        create_card(c_idx, 4, c_idx + 3, 8)
        
        # Typography
        lbl_cell = ws.cell(row=5, column=c_idx)
        lbl_cell.value = anchor["label"]
        lbl_cell.font = Font(size=12, bold=True, color="555555")
        
        val_cell = ws.cell(row=6, column=c_idx)
        val_cell.value = anchor["val"]
        val_cell.font = Font(size=16, bold=True, color=palette["primary"])
        
        # Doughnut Chart Injection
        kpi_chart = DoughnutChart()
        kpi_chart.width = 3.5
        kpi_chart.height = 2.5
        kpi_chart.legend = None
        kpi_chart.holeSize = 65
        kpi_chart.graphical_properties.line.noFill = True # Remove external border
        
        data_ref = Reference(data_ws, min_col=2, min_row=anchor["kpi_row"], max_col=3, max_row=anchor["kpi_row"])
        series = Series(data_ref)
        
        # Manual Slice Coloring
        dp_actual = DataPoint(idx=0)
        dp_actual.graphicalProperties.solidFill = palette["primary"]
        dp_rem = DataPoint(idx=1)
        dp_rem.graphicalProperties.solidFill = "E0E0E0"
        series.dp = [dp_actual, dp_rem]
        
        kpi_chart.series.append(series)
        ws.add_chart(kpi_chart, f"{anchor['chart_col']}4")
        
    # 5. Trend Line Chart
    create_card(3, 10, 11, 24)
    line_chart = LineChart()
    line_chart.title = "2021-2022 Sales Trend (in millions)"
    line_chart.width = 13.5
    line_chart.height = 7.5
    line_chart.legend = None
    line_chart.y_axis.majorGridlines = None
    line_chart.x_axis.majorGridlines = None
    line_chart.graphical_properties.line.noFill = True
    
    data = Reference(data_ws, min_col=2, min_row=1, max_col=3, max_row=7)
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=7)
    line_chart.add_data(data, titles_from_data=True)
    line_chart.set_categories(cats)
    
    s1 = line_chart.series[0]
    s1.graphicalProperties.line.solidFill = "CCCCCC"
    s1.graphicalProperties.line.width = 20000
    
    s2 = line_chart.series[1]
    s2.graphicalProperties.line.solidFill = palette["primary"]
    s2.graphicalProperties.line.width = 30000
    
    # Custom Series Marker Design
    marker = Marker(symbol="circle", size=6)
    marker.graphicalProperties.solidFill = palette["card"]
    marker.graphicalProperties.line.solidFill = palette["primary"]
    marker.graphicalProperties.line.width = 15000
    s2.marker = marker
    
    ws.add_chart(line_chart, "C10")
    
    # 6. Satisfaction Radar Chart
    create_card(13, 10, 16, 24)
    radar_chart = RadarChart()
    radar_chart.title = "Customer Satisfaction"
    radar_chart.width = 6.5
    radar_chart.height = 7.5
    radar_chart.legend = None
    radar_chart.graphical_properties.line.noFill = True
    
    r_data = Reference(data_ws, min_col=2, min_row=12, max_row=17)
    r_labels = Reference(data_ws, min_col=1, min_row=13, max_row=17)
    radar_chart.add_data(r_data, titles_from_data=True)
    radar_chart.set_categories(r_labels)
    
    rs = radar_chart.series[0]
    rs.graphicalProperties.line.solidFill = palette["primary"]
    rs.graphicalProperties.line.width = 20000
    rs.marker = marker # Reuse custom marker logic
    
    ws.add_chart(radar_chart, "M10")
