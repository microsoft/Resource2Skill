from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Theme fallback loader
    try:
        from skills_library.excel.components._helpers import load_theme
        palette = load_theme(theme)
    except ImportError:
        palette = {
            "primary": "5B2C6F",     # Deep purple
            "accent": "F1C40F",      # Gold
            "background": "F2EFF5",  # Soft pale purple/grey
            "text_light": "FFFFFF",  # White
            "text_dark": "333333"    # Dark grey
        }
        
    primary_color = palette.get("primary", "5B2C6F")
    accent_color = palette.get("accent", "F1C40F")
    bg_color = palette.get("background", "F2EFF5")
    
    # 1. Background Canvas Fill (Rows 7 to 30)
    bg_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for row in ws.iter_rows(min_row=7, max_row=30, min_col=1, max_col=18):
        for cell in row:
            cell.fill = bg_fill
            
    # 2. Hero Header Fill (Rows 1 to 6)
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=18):
        for cell in row:
            cell.fill = header_fill
            
    # Configure precise column widths for the grid
    for col in range(2, 18):
        letter = get_column_letter(col)
        if col in [5, 9, 13]: # Gutter columns between cards
            ws.column_dimensions[letter].width = 3
        else:
            ws.column_dimensions[letter].width = 11
            
    # Title & Subtitle
    ws["B2"] = title
    ws["B2"].font = Font(color=palette.get("text_light", "FFFFFF"), size=26, bold=True)
    ws["B3"] = kwargs.get("subtitle", "Evaluating Key Performance Indicators")
    ws["B3"].font = Font(color=accent_color, size=14, italic=True)
    
    # 3. Cell-based KPI Cards (Rows 8 to 11)
    kpis = kwargs.get("kpis", [
        {"label": "TOTAL CALLS", "value": "16,749"},
        {"label": "CALLS REACHED", "value": "3,328"},
        {"label": "DEALS CLOSED", "value": "1,203"},
        {"label": "DEAL VALUE", "value": "$646,979"}
    ])
    
    card_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    top_border = Border(top=Side(style="thick", color=accent_color))
    
    col_start = 2 # Start at column 'B'
    for kpi in kpis[:4]:
        # Fill the card background and apply the accent top border
        for r in range(8, 12):
            for c in range(col_start, col_start + 3):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                if r == 8:
                    cell.border = top_border
                    
        # Inject KPI Data
        val_cell = ws.cell(row=9, column=col_start)
        val_cell.value = kpi["value"]
        val_cell.font = Font(size=22, bold=True, color=primary_color)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        lbl_cell = ws.cell(row=10, column=col_start)
        lbl_cell.value = kpi["label"]
        lbl_cell.font = Font(size=10, color=palette.get("text_dark", "333333"))
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Merge rows for exact centering within the card bounds
        ws.merge_cells(start_row=9, start_column=col_start, end_row=9, end_column=col_start+2)
        ws.merge_cells(start_row=10, start_column=col_start, end_row=10, end_column=col_start+2)
        
        col_start += 4
        
    # 4. Lower Dashboard Details Area (Rows 14 to 26)
    # Create a white background panel for a table/chart
    for r in range(14, 27):
        for c in range(2, 8):
            ws.cell(row=r, column=c).fill = card_fill
            
    ws["B15"] = "Agent Performance"
    ws["B15"].font = Font(color=primary_color, bold=True, size=14)
    
    # Render mock table to demonstrate Data Bars
    table_data = [
        ["Agent", "Calls", "Deals", "Conversion"],
        ["Alice", 827, 49, 0.059],
        ["Bob", 661, 28, 0.042],
        ["Charlie", 610, 67, 0.109],
        ["Diana", 566, 26, 0.045],
        ["Evan", 722, 16, 0.022],
    ]
    
    for i, row_data in enumerate(table_data):
        r = 17 + i
        for j, val in enumerate(row_data):
            c = 2 + j
            cell = ws.cell(row=r, column=c, value=val)
            if i == 0:
                cell.font = Font(bold=True, color=palette.get("text_light", "FFFFFF"))
                cell.fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            else:
                if j == 3:
                    cell.number_format = "0.0%"
                    
    # Apply Data Bar conditional formatting to the 'Deals' column
    rule = DataBarRule(start_type="min", end_type="max", color=accent_color)
    ws.conditional_formatting.add("D18:D22", rule)
