from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Financial Model", theme: str = "corporate_blue", **kwargs) -> None:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        
    # Mock palette (in production, use get_theme_palette)
    palettes = {
        "corporate_blue": {"header_bg": "203764", "header_fg": "FFFFFF", "accent": "FFFF00"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])
    
    header_fill = PatternFill(start_color=colors["header_bg"], end_color=colors["header_bg"], fill_type="solid")
    header_font = Font(bold=True, color=colors["header_fg"])
    highlight_fill = PatternFill(start_color=colors["accent"], end_color=colors["accent"], fill_type="solid")
    
    blue_font = Font(color="0000FF") # Financial modeling standard for hardcoded inputs
    bold_font = Font(bold=True)
    
    # 1. Setup Title & Scenario Toggle
    ws['A1'] = title
    ws['A1'].font = Font(size=14, bold=True)
    
    ws['F1'] = "Scenario (1=Base, 2=Down):"
    ws['F1'].font = bold_font
    ws['F1'].alignment = Alignment(horizontal="right")
    
    ws['G1'] = 1
    ws['G1'].fill = highlight_fill
    ws['G1'].border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    ws['G1'].alignment = Alignment(horizontal="center")
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("G1")

    # 2. Table Headers
    headers = ["Line Item", "Unit", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font

    # 3. Income Statement Shell
    ws['A3'] = "Income Statement"
    ws['A3'].font = bold_font
    is_labels = [
        ("Revenue", "$"), 
        ("COGS", "$"), 
        ("Gross Profit", "$"), 
        ("Operating Expenses", "$"), 
        ("Net Profit", "$")
    ]
    for i, (label, unit) in enumerate(is_labels, 4):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=unit)

    # 4. Assumptions Blocks Setup
    blocks = [
        (10, "Live Case Assumptions"),
        (17, "Scenario 1: Base Case"),
        (24, "Scenario 2: Downside Case")
    ]
    assump_labels = [
        ("Number of Orders", "#"), 
        ("Order Growth Rate", "%"), 
        ("Average Order Value", "$"), 
        ("COGS per order", "$"), 
        ("Fixed Opex", "$")
    ]
    
    for start_row, block_title in blocks:
        cell = ws.cell(row=start_row, column=1, value=block_title)
        cell.font = bold_font
        for i, (label, unit) in enumerate(assump_labels, start_row + 1):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=unit)

    # 5. Inject Hardcoded Inputs for Scenarios (Blue Font)
    # Scenario 1 (Base Case)
    ws['C18'].value = 3000
    ws['C18'].font = blue_font
    s1_data = {
        19: [1.0, 0.75, 0.50, 0.35],               # Growth
        20: [39.95, 39.95, 39.95, 39.95, 39.95],   # AOV
        21: [8.75, 8.75, 8.75, 8.75, 8.75],        # COGS
        22: [20000, 30000, 30000, 30000, 30000]    # Opex
    }
    
    # Scenario 2 (Downside Case)
    ws['C25'].value = 2000
    ws['C25'].font = blue_font
    s2_data = {
        26: [0.50, 0.25, 0.10, 0.05],              # Growth
        27: [34.95, 34.95, 34.95, 34.95, 34.95],   # AOV
        28: [9.50, 9.50, 9.50, 9.50, 9.50],        # COGS
        29: [25000, 35000, 35000, 35000, 35000]    # Opex
    }

    def write_scenario(start_col, data_dict):
        for r_idx, values in data_dict.items():
            # offset columns for growth (starts Year 2)
            col_offset = 4 if r_idx in (19, 26) else 3
            for i, val in enumerate(values):
                c = ws.cell(row=r_idx, column=col_offset + i, value=val)
                c.font = blue_font

    write_scenario(3, s1_data)
    write_scenario(3, s2_data)

    # Calculate compounded Orders for Scenarios (Formulas are black)
    for c_idx in range(4, 8):
        prev_col = get_column_letter(c_idx-1)
        curr_col = get_column_letter(c_idx)
        ws.cell(row=18, column=c_idx, value=f"={prev_col}18*(1+{curr_col}19)") # Scen 1 Orders
        ws.cell(row=25, column=c_idx, value=f"={prev_col}25*(1+{curr_col}26)") # Scen 2 Orders

    # 6. Apply Live Case "CHOOSE" Routing formulas
    for c_idx in range(3, 8):
        col_ltr = get_column_letter(c_idx)
        # 11: Orders | 12: Growth | 13: AOV | 14: COGS | 15: Opex
        ws.cell(row=11, column=c_idx, value=f"=CHOOSE($G$1, {col_ltr}18, {col_ltr}25)")
        if c_idx > 3: # Growth only applicable Y2 onward
            ws.cell(row=12, column=c_idx, value=f"=CHOOSE($G$1, {col_ltr}19, {col_ltr}26)")
        ws.cell(row=13, column=c_idx, value=f"=CHOOSE($G$1, {col_ltr}20, {col_ltr}27)")
        ws.cell(row=14, column=c_idx, value=f"=CHOOSE($G$1, {col_ltr}21, {col_ltr}28)")
        ws.cell(row=15, column=c_idx, value=f"=CHOOSE($G$1, {col_ltr}22, {col_ltr}29)")

    # 7. Roll Up Income Statement (Referencing Live Case only)
    for c_idx in range(3, 8):
        col_ltr = get_column_letter(c_idx)
        ws.cell(row=4, column=c_idx, value=f"={col_ltr}11*{col_ltr}13")  # Rev = Orders * AOV
        ws.cell(row=5, column=c_idx, value=f"={col_ltr}11*{col_ltr}14")  # COGS = Orders * COGS/order
        ws.cell(row=6, column=c_idx, value=f"={col_ltr}4-{col_ltr}5")    # GP = Rev - COGS
        ws.cell(row=7, column=c_idx, value=f"={col_ltr}15")              # Opex = Fixed Opex
        ws.cell(row=8, column=c_idx, value=f"={col_ltr}6-{col_ltr}7")    # Net Profit = GP - Opex

    # 8. Formatting & Cleanup
    accounting_fmt = '_($* #,##0_);_($* (#,##0);_($* "-"_);_(@_)'
    currency_fmt = '$#,##0.00'
    number_fmt = '#,##0'
    percent_fmt = '0%'

    # Formatting blocks mapping
    for row in range(4, 9):
        for col in range(3, 8):
            ws.cell(row=row, column=col).number_format = accounting_fmt

    # Bold sub-totals
    for r in [6, 8]:
        for c in range(1, 8):
            ws.cell(row=r, column=c).font = bold_font

    # Format metric blocks (Live, S1, S2)
    for offset in [0, 7, 14]: # Offset for the three blocks
        for c in range(3, 8):
            ws.cell(row=11+offset, column=c).number_format = number_fmt
            ws.cell(row=12+offset, column=c).number_format = percent_fmt
            ws.cell(row=13+offset, column=c).number_format = currency_fmt
            ws.cell(row=14+offset, column=c).number_format = currency_fmt
            ws.cell(row=15+offset, column=c).number_format = accounting_fmt

    # Set Widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 8
    for col in ['C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col].width = 15
