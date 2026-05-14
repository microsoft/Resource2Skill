from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string

def render(ws, anchor: str, *, 
           columns: list[dict] = None, 
           data: list[list] = None, 
           theme: dict = None, 
           **kwargs) -> None:
    """
    Renders a KPI performance table with inline Data Bars for visual benchmarking.
    
    :param columns: List of dicts configuring each column: {"header": str, "format": str, "bar_color": str, "width": int}
    :param data: 2D list of row data matching the columns structure.
    """
    
    # Fallback dashboard data mimicking the video's sales agent performance
    if columns is None:
        columns = [
            {"header": "Agent Name", "format": "General", "bar_color": None, "width": 15},
            {"header": "Total Calls", "format": "#,##0", "bar_color": "FF8064A2", "width": 12},     # Purple
            {"header": "Calls Reached", "format": "#,##0", "bar_color": "FFFFC000", "width": 14},   # Gold/Yellow
            {"header": "Deals Closed", "format": "#,##0", "bar_color": "FFCCC0DA", "width": 14},    # Light Purple
            {"header": "Deal Value", "format": "$#,##0", "bar_color": "FF8064A2", "width": 15},     # Purple
        ]
        
    if data is None:
        data = [
            ["Alice", 1031, 56, 37, 13519.04],
            ["Bob", 661, 73, 28, 40092.43],
            ["Charlie", 610, 86, 67, 45236.03],
            ["David", 375, 120, 48, 2590.45],
            ["Eva", 1057, 63, 17, 37878.60],
            ["Frank", 827, 128, 49, 41200.22],
            ["Grace", 566, 163, 26, 38593.93]
        ]
        
    # Sort data by the primary metric (Deal Value) descending, mimicking the video's presentation
    data = sorted(data, key=lambda x: x[-1], reverse=True)

    # Theme handling
    theme = theme or {}
    header_bg = theme.get("primary_color", "5C2D91").replace("#", "")
    header_fg = theme.get("background_color", "FFFFFF").replace("#", "")

    # Calculate starting indices
    coords = coordinate_from_string(anchor)
    start_col_idx = column_index_from_string(coords[0])
    start_row_idx = coords[1]
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Write and format Headers
    for c_idx, col_def in enumerate(columns):
        col_letter = get_column_letter(start_col_idx + c_idx)
        cell = ws.cell(row=start_row_idx, column=start_col_idx + c_idx, value=col_def["header"])
        cell.font = Font(bold=True, color=header_fg)
        cell.fill = PatternFill(patternType="solid", fgColor=header_bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
        if "width" in col_def:
            ws.column_dimensions[col_letter].width = col_def["width"]

    # 2. Write and format Data
    for r_idx, row_data in enumerate(data):
        current_row = start_row_idx + 1 + r_idx
        for c_idx, val in enumerate(row_data):
            cell = ws.cell(row=current_row, column=start_col_idx + c_idx, value=val)
            col_def = columns[c_idx]
            
            # Number formatting
            if col_def.get("format") and col_def["format"] != "General":
                cell.number_format = col_def["format"]
            
            # Align numbers to the right, text to the left
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left", indent=1)
                
            cell.border = thin_border

    # 3. Apply Conditional Formatting Data Bars
    end_row_idx = start_row_idx + len(data)
    for c_idx, col_def in enumerate(columns):
        bar_color = col_def.get("bar_color")
        if bar_color:
            col_letter = get_column_letter(start_col_idx + c_idx)
            range_str = f"{col_letter}{start_row_idx + 1}:{col_letter}{end_row_idx}"
            
            clean_color = bar_color.replace("#", "")
            rule = DataBarRule(start_type='min', end_type='max', color=clean_color)
            ws.conditional_formatting.add(range_str, rule)
