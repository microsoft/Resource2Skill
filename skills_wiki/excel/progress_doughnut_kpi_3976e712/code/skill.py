from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.marker import DataPoint

def render(ws, anchor: str, *, kpi_name: str = "Sales Completion", actual_val: float = 0.85, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a stylized KPI Doughnut chart with Actual vs Remainder slices.
    The data table is placed at the anchor, and the chart is placed slightly below it.
    """
    row = ws[anchor].row
    col = ws[anchor].column

    # 1. Setup Data for Doughnut (Actual vs Remainder)
    # Clamp actual_val to max 1.0 (100%) to prevent visual distortion
    percent_complete = min(actual_val, 1.0)
    remainder_val = 1.0 - percent_complete
    
    # Write chart data to the worksheet
    ws.cell(row=row, column=col, value="Metric")
    ws.cell(row=row, column=col+1, value="Value")
    
    ws.cell(row=row+1, column=col, value="Complete")
    c_act = ws.cell(row=row+1, column=col+1, value=percent_complete)
    c_act.number_format = "0%"
    
    ws.cell(row=row+2, column=col, value="Remainder")
    c_rem = ws.cell(row=row+2, column=col+1, value=remainder_val)
    c_rem.number_format = "0%"

    # 2. Initialize Doughnut Chart
    chart = DoughnutChart()
    chart.title = kpi_name
    chart.legend = None  # Clean look, no legend for KPIs
    chart.width = 6.0
    chart.height = 6.0
    chart.holeSize = 65  # Matches the thick ring aesthetic in the tutorial

    # 3. Bind Data & Categories
    cats = Reference(ws, min_col=col, min_row=row+1, max_row=row+2)
    data = Reference(ws, min_col=col+1, min_row=row+1, max_row=row+2)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    # 4. Color the Slices
    # openpyxl requires appending DataPoints to series.dPt to format individual slices
    series = chart.series[0]

    # Slice 0: Actual (Primary Theme Color)
    pt_actual = DataPoint(idx=0)
    pt_actual.graphicalProperties.solidFill = "203764" # Dark Blue fallback
    series.dPt.append(pt_actual)

    # Slice 1: Remainder (Muted Gray track)
    pt_remainder = DataPoint(idx=1)
    pt_remainder.graphicalProperties.solidFill = "D9D9D9" # Light Gray fallback
    series.dPt.append(pt_remainder)

    # 5. Place Chart offset from the data so the underlying numbers remain accessible
    chart_anchor = ws.cell(row=row+4, column=col).coordinate
    ws.add_chart(chart, chart_anchor)
