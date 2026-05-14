from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean dashboard shell with a primary layout area and dummy charts, 
    matching the visual presentation of an interactive KPI dashboard.
    """
    ws = wb.create_sheet(sheet_name)

    # 1. Dashboard Canvas Setup
    ws.sheet_view.showGridLines = False

    # Default color fallbacks (simulating a loaded theme palette)
    bg_color = "1F4E78" # Corporate Blue
    fg_color = "FFFFFF" # White text

    # 2. Title Banner
    ws.merge_cells("A1:P3")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=fg_color)
    title_cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Inject Realistic Dummy Data for Layout (placed out of standard view)
    data_col = 20 # Column T
    
    # Data for Primary Stacked Bar Chart
    ws.cell(row=1, column=data_col, value="Market")
    ws.cell(row=1, column=data_col+1, value="Chocolate Chip")
    ws.cell(row=1, column=data_col+2, value="Fortune Cookie")

    markets = ["India", "Malaysia", "Philippines", "United Kingdom", "United States"]
    for i, market in enumerate(markets, start=2):
        ws.cell(row=i, column=data_col, value=market)
        ws.cell(row=i, column=data_col+1, value=60000 - (i*5000))
        ws.cell(row=i, column=data_col+2, value=25000 + (i*2000))

    # Data for Secondary Line Charts
    ws.cell(row=10, column=data_col, value="Month")
    ws.cell(row=10, column=data_col+1, value="Units Sold")
    ws.cell(row=10, column=data_col+2, value="Profit")
    months = ["Sep", "Oct", "Nov", "Dec"]
    for i, month in enumerate(months, start=11):
        ws.cell(row=i, column=data_col, value=month)
        ws.cell(row=i, column=data_col+1, value=50000 + (i % 2) * 40000)
        ws.cell(row=i, column=data_col+2, value=120000 + (i % 3) * 60000)

    # 4. Primary Chart (Stacked Column)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.y_axis.title = "Profit ($)"
    
    # Hide legend to match clean presentation, or leave to explain stack
    bar_chart.legend.position = 'b'

    data_bar = Reference(ws, min_col=data_col+1, max_col=data_col+2, min_row=1, max_row=6)
    cats_bar = Reference(ws, min_col=data_col, min_row=2, max_row=6)
    bar_chart.add_data(data_bar, titles_from_data=True)
    bar_chart.set_categories(cats_bar)

    bar_chart.width = 16.5
    bar_chart.height = 10.5
    ws.add_chart(bar_chart, "B5")

    # 5. Secondary Chart 1 (Line - Units)
    line1 = LineChart()
    line1.title = "Units sold each month"
    data_line1 = Reference(ws, min_col=data_col+1, min_row=10, max_row=14)
    cats_line = Reference(ws, min_col=data_col, min_row=11, max_row=14)
    
    line1.add_data(data_line1, titles_from_data=True)
    line1.set_categories(cats_line)
    line1.legend = None # Remove legend for single series
    line1.width = 12.5
    line1.height = 5
    ws.add_chart(line1, "J5")

    # 6. Secondary Chart 2 (Line - Profit)
    line2 = LineChart()
    line2.title = "Profit by month"
    data_line2 = Reference(ws, min_col=data_col+2, min_row=10, max_row=14)
    
    line2.add_data(data_line2, titles_from_data=True)
    line2.set_categories(cats_line)
    line2.legend = None
    line2.width = 12.5
    line2.height = 5.2
    ws.add_chart(line2, "J11")
