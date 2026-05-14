def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    ws = wb.create_sheet(sheet_name)
    
    # 1. Disable native gridlines for a clean UI
    ws.sheet_view.showGridLines = False

    # 2. Theme Configuration 
    # Fallback dictionary simulating standard palette loading
    palette = {
        "corporate_blue": {"sidebar": "1F3864", "wash": "F2F2F2", "card": "FFFFFF", "text": "262626", "border": "D9D9D9"},
        "modern_dark": {"sidebar": "171717", "wash": "262626", "card": "404040", "text": "FFFFFF", "border": "595959"}
    }.get(theme, {"sidebar": "1F3864", "wash": "F2F2F2", "card": "FFFFFF", "text": "262626", "border": "D9D9D9"})

    fills = {
        "wash": PatternFill("solid", fgColor=palette["wash"]),
        "card": PatternFill("solid", fgColor=palette["card"]),
        "sidebar": PatternFill("solid", fgColor=palette["sidebar"])
    }

    # 3. Apply global background wash (covers A1:R40)
    for row in ws.iter_rows(min_row=1, max_row=40, min_col=1, max_col=18):
        for cell in row:
            cell.fill = fills["wash"]

    # 4. Draw Navigation Sidebar (Col A)
    ws.column_dimensions['A'].width = 8
    for r in range(1, 41):
        ws.cell(row=r, column=1).fill = fills["sidebar"]

    # 5. Dashboard Grid Geometry
    # Setup standard width arrays to form 3 equal columns with spacers
    for col_letter in ['B', 'G', 'L']:
        ws.column_dimensions[col_letter].width = 3  # Spacers
    for col_letter in ['C', 'D', 'E', 'F', 'H', 'I', 'J', 'K', 'M', 'N', 'O', 'P']:
        ws.column_dimensions[col_letter].width = 11 # Card content columns

    # 6. Card Generation Helper
    def make_card(min_col, min_row, max_col, max_row, card_title=None):
        """Punches out a white card area from the wash background and applies subtle borders."""
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fills["card"]

                # Apply perimeter borders
                top = Side(style='thin', color=palette["border"]) if r == min_row else None
                bottom = Side(style='thin', color=palette["border"]) if r == max_row else None
                left = Side(style='thin', color=palette["border"]) if c == min_col else None
                right = Side(style='thin', color=palette["border"]) if c == max_col else None

                if top or bottom or left or right:
                    cell.border = Border(top=top, bottom=bottom, left=left, right=right)

        # Inject Card Header
        if card_title:
            header_cell = ws.cell(row=min_row, column=min_col)
            header_cell.value = "  " + card_title  # Slight indent
            header_cell.font = Font(bold=True, size=12, color=palette["text"])
            header_cell.alignment = Alignment(vertical="center")

    # 7. Construct Layout
    # Top Title Banner
    make_card(3, 2, 16, 4, title)
    ws.cell(row=2, column=3).font = Font(bold=True, size=20, color=palette["text"]) # Upsize main title

    # Three Top-Level KPI Cards
    make_card(3, 6, 6, 10, "Sales Volume")
    make_card(8, 6, 11, 10, "Net Profit")
    make_card(13, 6, 16, 10, "Customer Growth")

    # Two Lower Chart Cards
    make_card(3, 12, 9, 26, "2021-2022 Sales Trend")
    make_card(11, 12, 16, 26, "Regional Performance Breakdown")
