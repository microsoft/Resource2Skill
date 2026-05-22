def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # Standard theme palette routing
    palettes = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF"},
        "modern_dark": {"header_bg": "2B2B2B", "header_fg": "FFFFFF"},
        "forest_green": {"header_bg": "2E7D32", "header_fg": "FFFFFF"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # 1. Create Dashboard Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    # Transform into a clean "Canvas" by removing spreadsheet clutter
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False

    # 2. Render Dominant Header
    ws.merge_cells("A1:R2")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["header_fg"])
    header_cell.fill = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Create Hidden Data Sheet for Charts
    data_sheet_name = f"{sheet_name}_Data"
    if data_sheet_name in wb.sheetnames:
        ws_data = wb[data_sheet_name]
    else:
        ws_data = wb.create_sheet(data_sheet_name)
        # Hide the data sheet to preserve the dashboard illusion
        ws_data.sheet_state = 'hidden'

    # Populate Data for Stacked Bar (Markets vs Products)
    bar_data = [
        ["Market", "Chocolate Chip", "Oatmeal Raisin", "Sugar Cookie"],
        ["India", 62000, 21000, 18000],
        ["Philippines", 54000, 22000, 14000],
        ["United Kingdom", 46000, 11000, 19000],
        ["United States", 36000, 22000, 9000]
    ]
    for row in bar_data:
        ws_data.append(row)

    # Populate Data for Line Charts (Months vs Units/Profit)
    line_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    ws_data.append([]) # Spacer row
    start_line_row = ws_data.max_row + 1
    for row in line_data:
        ws_data.append(row)

    # 4. Render Chart 1: Stacked Bar (Left Column)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 11

    # Map references
    bar_data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=4, max_row=5)
    bar_cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(bar_data_ref, titles_from_data=True)
    bar_chart.set_categories(bar_cats_ref)

    bar_chart.height = 13.5
    bar_chart.width = 18
    ws.add_chart(bar_chart, "B4")

    # 5. Render Chart 2: Top Line Chart (Right Column Top)
    line_chart_units = LineChart()
    line_chart_units.title = "Units sold each month"
    line_chart_units.style = 13
    
    line_u_data = Reference(ws_data, min_col=2, min_row=start_line_row, max_row=start_line_row+4)
    line_cats = Reference(ws_data, min_col=1, min_row=start_line_row+1, max_row=start_line_row+4)
    
    line_chart_units.add_data(line_u_data, titles_from_data=True)
    line_chart_units.set_categories(line_cats)
    line_chart_units.legend = None # Remove legend for cleaner look

    line_chart_units.height = 6.5
    line_chart_units.width = 14
    ws.add_chart(line_chart_units, "L4")

    # 6. Render Chart 3: Bottom Line Chart (Right Column Bottom)
    line_chart_profit = LineChart()
    line_chart_profit.title = "Profit by month"
    line_chart_profit.style = 13
    
    line_p_data = Reference(ws_data, min_col=3, min_row=start_line_row, max_row=start_line_row+4)
    
    line_chart_profit.add_data(line_p_data, titles_from_data=True)
    line_chart_profit.set_categories(line_cats)
    line_chart_profit.legend = None # Remove legend for cleaner look

    line_chart_profit.height = 6.5
    line_chart_profit.width = 14
    ws.add_chart(line_chart_profit, "L11")
