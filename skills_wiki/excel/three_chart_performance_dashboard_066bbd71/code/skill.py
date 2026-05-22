def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # Theme palette fallback
    palettes = {
        "corporate_blue": {"primary": "4F81BD", "text": "333333", "bg": "FFFFFF"},
        "modern_dark": {"primary": "4472C4", "text": "FFFFFF", "bg": "262626"}
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    # Clear default sheets
    for sheet in wb.sheetnames:
        del wb[sheet]

    ws_dash = wb.create_sheet("Dashboard")
    ws_data = wb.create_sheet("Data")
    
    # Hide gridlines on dashboard
    ws_dash.sheet_view.showGridLines = False
    
    # Apply optional background color if not default white
    if palette["bg"] != "FFFFFF":
        fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
        for row in ws_dash.iter_rows(min_row=1, max_row=40, min_col=1, max_col=20):
            for cell in row:
                cell.fill = fill
                
    # Header Title
    ws_dash.merge_cells("B2:P3")
    title_cell = ws_dash["B2"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette["text"])
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # ==========================================
    # DATA SETUP (Hidden Sheet)
    # ==========================================
    
    # Chart 1 Data: Profit by Market & Product
    chart1_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle", "Sugar"],
        ["India", 62349, 4872, 21028, 25085, 18561],
        ["Philippines", 54618, 7026, 22005, 8313, 14947],
        ["United Kingdom", 46530, 5220, 11497, 14620, 19446],
        ["Malaysia", 46587, 5538, 17536, 20555, 10633],
        ["United States", 36657, 6369, 22260, 9938, 9186]
    ]
    for row in chart1_data:
        ws_data.append(row)
        
    ws_data.append([]) # spacer
    start_c2 = ws_data.max_row + 1
    
    # Chart 2 Data: Units Sold Over Time
    chart2_data = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for row in chart2_data:
        ws_data.append(row)
        
    ws_data.append([]) # spacer
    start_c3 = ws_data.max_row + 1
    
    # Chart 3 Data: Profit Over Time
    chart3_data = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for row in chart3_data:
        ws_data.append(row)
        
    # ==========================================
    # CHART SETUP (Dashboard Sheet)
    # ==========================================
    
    # Chart 1: Stacked Column (Left Side)
    c1 = BarChart()
    c1.type = "col"
    c1.style = 10 
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Cookie Type"
    c1.width = 16
    c1.height = 12.5
    
    data_ref1 = Reference(ws_data, min_col=2, min_row=1, max_col=6, max_row=6)
    cats_ref1 = Reference(ws_data, min_col=1, min_row=2, max_row=6)
    c1.add_data(data_ref1, titles_from_data=True)
    c1.set_categories(cats_ref1)
    ws_dash.add_chart(c1, "B5")
    
    # Chart 2: Line - Units Sold (Top Right)
    c2 = LineChart()
    c2.style = 13
    c2.title = "Units sold each month"
    c2.width = 14
    c2.height = 6
    
    data_ref2 = Reference(ws_data, min_col=2, min_row=start_c2, max_col=2, max_row=start_c2+4)
    cats_ref2 = Reference(ws_data, min_col=1, min_row=start_c2+1, max_row=start_c2+4)
    c2.add_data(data_ref2, titles_from_data=True)
    c2.set_categories(cats_ref2)
    c2.legend = None  # Clean look
    ws_dash.add_chart(c2, "J5")
    
    # Chart 3: Line - Profit by Month (Bottom Right)
    c3 = LineChart()
    c3.style = 13
    c3.title = "Profit by month"
    c3.width = 14
    c3.height = 6
    
    data_ref3 = Reference(ws_data, min_col=2, min_row=start_c3, max_col=2, max_row=start_c3+4)
    cats_ref3 = Reference(ws_data, min_col=1, min_row=start_c3+1, max_row=start_c3+4)
    c3.add_data(data_ref3, titles_from_data=True)
    c3.set_categories(cats_ref3)
    c3.legend = None  # Clean look
    ws_dash.add_chart(c3, "J16")
    
    # Hide the data sheet to complete the dashboard illusion
    ws_data.sheet_state = "hidden"
