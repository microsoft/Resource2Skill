def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # Load theme (fallback mechanism)
    try:
        from _helpers import get_theme
        theme_palette = get_theme(theme)
        header_bg = theme_palette.get("primary", "203764")
        header_fg = theme_palette.get("text_light", "FFFFFF")
    except ImportError:
        header_bg = "203764"
        header_fg = "FFFFFF"

    # Create main Dashboard sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
    
    # Disable gridlines to create a clean "canvas" look
    ws.sheet_view.showGridLines = False
    
    # Dashboard Header Setup
    ws.merge_cells("A1:R3")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=header_fg)
    header_cell.fill = PatternFill(start_color=header_bg, end_color=header_bg, fill_type="solid")
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Create hidden data sheet to store chart backing data
    data_sheet_name = f"{sheet_name}_Data"
    if data_sheet_name in wb.sheetnames:
        ws_data = wb[data_sheet_name]
    else:
        ws_data = wb.create_sheet(data_sheet_name)
        ws_data.sheet_state = "hidden"
        
    # --- Data & Chart 1: Profit by Market & Product (Stacked Bar) ---
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin", "Snickerdoodle"],
        ["India", 62349, 4872, 21028, 25085],
        ["United States", 36657, 6368, 22260, 9937],
        ["United Kingdom", 46530, 5220, 11497, 14620],
        ["Philippines", 54618, 7026, 22005, 8313],
    ]
    
    start_row = ws_data.max_row + 1 if ws_data.max_row > 1 else 1
    for r in market_data:
        ws_data.append(r)
        
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.height = 13.5
    bar_chart.width = 16
    
    data_ref = Reference(ws_data, min_col=2, min_row=start_row, max_col=len(market_data[0]), max_row=start_row + len(market_data) - 1)
    cats_ref = Reference(ws_data, min_col=1, min_row=start_row + 1, max_row=start_row + len(market_data) - 1)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    
    ws.add_chart(bar_chart, "B5")
    
    # --- Data & Chart 2 & 3: Trend over time (Line Charts) ---
    start_row = ws_data.max_row + 2
    trend_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    for r in trend_data:
        ws_data.append(r)
        
    # Top Line Chart
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.height = 6.5
    line1.width = 13
    l1_data = Reference(ws_data, min_col=2, min_row=start_row, max_row=start_row + len(trend_data) - 1)
    l1_cats = Reference(ws_data, min_col=1, min_row=start_row + 1, max_row=start_row + len(trend_data) - 1)
    line1.add_data(l1_data, titles_from_data=True)
    line1.set_categories(l1_cats)
    line1.legend = None  # Maximizes chart area
    ws.add_chart(line1, "J5")
    
    # Bottom Line Chart
    line2 = LineChart()
    line2.title = "Profit by month"
    line2.height = 6.5
    line2.width = 13
    l2_data = Reference(ws_data, min_col=3, min_row=start_row, max_row=start_row + len(trend_data) - 1)
    line2.add_data(l2_data, titles_from_data=True)
    line2.set_categories(l1_cats)
    line2.legend = None  # Maximizes chart area
    ws.add_chart(line2, "J16")
