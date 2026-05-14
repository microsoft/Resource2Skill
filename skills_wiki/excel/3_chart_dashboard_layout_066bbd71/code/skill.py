def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, LineChart, Reference

    # Base palette configuration (supports swapping themes)
    theme_colors = {
        "corporate_blue": {"bg": "203764", "fg": "FFFFFF", "panel": "F2F2F2", "border": "D9D9D9"},
        "emerald": {"bg": "0F4C3A", "fg": "FFFFFF", "panel": "EAF3F0", "border": "C8DCD5"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])

    # --- 1. Generate Supporting Data in a Hidden Sheet ---
    data_sheet_name = f"{sheet_name}_Data"
    if data_sheet_name in wb.sheetnames:
        del wb[data_sheet_name]
    ws_data = wb.create_sheet(data_sheet_name)
    ws_data.sheet_state = 'hidden'

    # Primary Chart Data (Profit by Market & Product)
    bar_data = [
        ["Market", "Fortune Cookie", "Sugar", "Snickerdoodle", "Oatmeal Raisin", "White Choc", "Choc Chip"],
        ["India", 4800, 18500, 25000, 21000, 23000, 62000],
        ["Philippines", 7000, 14000, 8000, 22000, 24000, 54000],
        ["UK", 1200, 19000, 14000, 11000, 26000, 46000],
        ["US", 5000, 10000, 20000, 17000, 20000, 46000]
    ]
    for row in bar_data:
        ws_data.append(row)

    # Secondary Charts Data (Trends over time)
    line_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    start_row = 10
    for i, row in enumerate(line_data):
        for j, val in enumerate(row):
            ws_data.cell(row=start_row+i, column=j+1, value=val)

    # --- 2. Setup Dashboard Sheet Canvas ---
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Main Header Banner
    ws.merge_cells("B2:T4")
    header_cell = ws["B2"]
    header_cell.value = title
    header_cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    header_cell.font = Font(color=palette["fg"], size=24, bold=True)
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Left Control Panel Placeholder (for Slicers/Filters)
    ws.merge_cells("B6:C21")
    slicer_area = ws["B6"]
    slicer_area.value = "Filter Controls\n\n(Insert Slicers Here)"
    slicer_area.fill = PatternFill(start_color=palette["panel"], end_color=palette["panel"], fill_type="solid")
    slicer_area.font = Font(color="595959", italic=True)
    slicer_area.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin = Side(style="thin", color=palette["border"])
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    for r in range(6, 22):
        for c in range(2, 4):
            ws.cell(row=r, column=c).border = border

    # --- 3. Construct and Position Charts ---
    # Primary Chart: Stacked Column
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 11
    
    cats = Reference(ws_data, min_col=1, min_row=2, max_row=5)
    data = Reference(ws_data, min_col=2, min_row=1, max_row=5, max_col=7)
    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    bar_chart.width = 14
    bar_chart.height = 8.5
    ws.add_chart(bar_chart, "D6")

    # Secondary Chart 1: Line (Units Sold)
    line1 = LineChart()
    line1.title = "Units sold each month"
    line1.style = 13
    
    cats_l = Reference(ws_data, min_col=1, min_row=11, max_row=14)
    data_l1 = Reference(ws_data, min_col=2, min_row=10, max_row=14)
    line1.add_data(data_l1, titles_from_data=True)
    line1.set_categories(cats_l)
    line1.legend = None
    line1.width = 10
    line1.height = 4.2
    ws.add_chart(line1, "N6")

    # Secondary Chart 2: Line (Profit)
    line2 = LineChart()
    line2.title = "Profit by month"
    line2.style = 13
    
    data_l2 = Reference(ws_data, min_col=3, min_row=10, max_row=14)
    line2.add_data(data_l2, titles_from_data=True)
    line2.set_categories(cats_l)
    line2.legend = None
    line2.width = 10
    line2.height = 4.2
    ws.add_chart(line2, "N14")
