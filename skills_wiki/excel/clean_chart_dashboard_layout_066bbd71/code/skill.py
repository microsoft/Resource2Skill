from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean, presentation-ready dashboard layout with a header banner,
    a reserved side-panel for controls, and a grid of charts.
    """
    # 1. Create Dashboard Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        
    # Disable gridlines for a clean, application-like look
    ws.sheet_view.showGridLines = False
    
    # 2. Define Theme Colors (fallback to a deep corporate blue)
    header_bg_color = kwargs.get("header_bg_color", "203764")
    header_fill = PatternFill("solid", fgColor=header_bg_color)
    header_font = Font(color="FFFFFF", size=28, bold=True)
    
    # 3. Create Header Banner (Rows 1 to 4)
    ws.merge_cells("A1:R4")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = header_font
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for row in ws.iter_rows(min_row=1, max_row=4, min_col=1, max_col=18):
        for cell in row:
            cell.fill = header_fill
            
    # 4. Set Column Widths for Layout Structure
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 18  # Left panel for Slicers/Controls
    ws.column_dimensions['C'].width = 2   # Gutter spacing
    ws.column_dimensions['D'].width = 2
    
    # 5. Add Left Panel Slicer Placeholders
    slicer_font = Font(bold=True, size=11, color="333333")
    slicer_fill = PatternFill("solid", fgColor="F2F2F2")
    
    ws["B6"] = "Filters / Controls"
    ws["B6"].font = Font(bold=True, size=12, color=header_bg_color)
    
    # Simulated Slicer 1
    ws["B8"] = "Date Range"
    ws["B8"].font = slicer_font
    ws["B8"].fill = slicer_fill
    ws["B9"] = "All Periods"
    
    # Simulated Slicer 2
    ws["B12"] = "Market"
    ws["B12"].font = slicer_font
    ws["B12"].fill = slicer_fill
    ws["B13"] = "United States\nIndia\nUnited Kingdom"
    ws["B13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[13].height = 45

    # 6. Setup Backing Data Sheet (Hidden)
    data_sheet_name = f"{sheet_name}_Data"
    if data_sheet_name in wb.sheetnames:
        data_ws = wb[data_sheet_name]
    else:
        data_ws = wb.create_sheet(data_sheet_name)
    data_ws.sheet_state = 'hidden'
    
    # Populate Dummy Data for Charts
    # Stacked Bar Data
    data_ws.append(["Market", "Chocolate Chip", "Oatmeal Raisin", "Sugar Cookie"])
    data_ws.append(["India", 150000, 80000, 45000])
    data_ws.append(["United Kingdom", 120000, 95000, 30000])
    data_ws.append(["United States", 180000, 110000, 60000])
    
    # Line Chart Data (Empty row for spacing)
    data_ws.append([])
    data_ws.append(["Month", "Units Sold", "Profit"])
    data_ws.append(["Sep", 50601, 124812])
    data_ws.append(["Oct", 95622, 228275])
    data_ws.append(["Nov", 65481, 160228])
    data_ws.append(["Dec", 52970, 136337])
    
    # 7. Create Chart 1: Stacked Bar (Profit by Market & Cookie Type)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 10  # Standard clean style
    
    bar_data = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=4)
    bar_cats = Reference(data_ws, min_col=1, min_row=2, max_row=4)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    bar_chart.height = 12
    bar_chart.width = 16
    ws.add_chart(bar_chart, "D6")
    
    # 8. Create Chart 2: Line Chart (Units Sold)
    line_chart_units = LineChart()
    line_chart_units.title = "Units sold each month"
    line_chart_units.style = 13
    
    units_data = Reference(data_ws, min_col=2, min_row=6, max_row=10)
    line_cats = Reference(data_ws, min_col=1, min_row=7, max_row=10)
    line_chart_units.add_data(units_data, titles_from_data=True)
    line_chart_units.set_categories(line_cats)
    line_chart_units.legend = None # Remove legend for single series
    line_chart_units.height = 7.5
    line_chart_units.width = 14
    ws.add_chart(line_chart_units, "L6")
    
    # 9. Create Chart 3: Line Chart (Profit)
    line_chart_profit = LineChart()
    line_chart_profit.title = "Profit by month"
    line_chart_profit.style = 13
    
    profit_data = Reference(data_ws, min_col=3, min_row=6, max_row=10)
    line_chart_profit.add_data(profit_data, titles_from_data=True)
    line_chart_profit.set_categories(line_cats)
    line_chart_profit.legend = None
    line_chart_profit.height = 7.5
    line_chart_profit.width = 14
    ws.add_chart(line_chart_profit, "L21")
