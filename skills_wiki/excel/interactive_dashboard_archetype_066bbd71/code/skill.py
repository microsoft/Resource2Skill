def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Setup Theme Hooks
    theme_colors = {
        "corporate_blue": {"primary": "1F4E78", "bg": "F2F2F2", "text": "FFFFFF", "active": "DDEBF7", "border": "D9D9D9"},
        "emerald": {"primary": "005A36", "bg": "EBF1EE", "text": "FFFFFF", "active": "C3D8CF", "border": "A8C2B5"}
    }
    palette = theme_colors.get(theme, theme_colors["corporate_blue"])
    
    # 2. Setup Sheet Architecture
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    
    ws_calc = wb.create_sheet("Calc")
    ws_data = wb.create_sheet("Data")
    
    # 3. Build Presentation Header
    ws_dash.merge_cells("A1:T3")
    title_cell = ws_dash["A1"]
    title_cell.value = title
    title_cell.font = Font(size=28, bold=True, color=palette["text"])
    title_cell.fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 4. Populate Calculation Data (Subbing for Pivot Tables)
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["United States", 36657, 6368, 22260],
        ["United Kingdom", 46530, 5220, 11497],
        ["Philippines", 54618, 7026, 22005],
    ]
    for r, row in enumerate(market_data, start=1):
        for c, val in enumerate(row, start=1):
            ws_calc.cell(row=r, column=c, value=val)
            
    time_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337],
    ]
    for r, row in enumerate(time_data, start=10):
        for c, val in enumerate(row, start=1):
            ws_calc.cell(row=r, column=c, value=val)
            
    # 5. Build Interactive Controls (Slicer Mockups)
    ws_dash.column_dimensions['A'].width = 2
    ws_dash.column_dimensions['B'].width = 18
    
    ws_dash["B5"] = "Market Filter"
    ws_dash["B5"].font = Font(bold=True, color=palette["primary"])
    
    filters = ["India", "Malaysia", "Philippines", "United Kingdom", "United States"]
    for i, f in enumerate(filters, start=6):
        cell = ws_dash.cell(row=i, column=2, value=f)
        cell.border = Border(
            bottom=Side(style="thin", color=palette["border"]),
            left=Side(style="thin", color=palette["border"]),
            right=Side(style="thin", color=palette["border"]),
            top=Side(style="thin", color=palette["border"]) if i == 6 else None
        )
        if f == "India":
            cell.fill = PatternFill(start_color=palette["active"], end_color=palette["active"], fill_type="solid")
            cell.font = Font(bold=True)
        else:
            cell.fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
            
    # 6. Build and Anchor Charts
    # Chart 1: Stacked Bar for Categorical Breakdowns
    bc = BarChart()
    bc.type = "col"
    bc.grouping = "stacked"
    bc.overlap = 100
    bc.title = "Profit by Market & Cookie Type"
    bc.height = 11
    bc.width = 16
    data = Reference(ws_calc, min_col=2, min_row=1, max_col=4, max_row=5)
    cats = Reference(ws_calc, min_col=1, min_row=2, max_row=5)
    bc.add_data(data, titles_from_data=True)
    bc.set_categories(cats)
    ws_dash.add_chart(bc, "D5")
    
    # Chart 2: Line for Unit Volume
    lc1 = LineChart()
    lc1.title = "Units sold each month"
    lc1.height = 7
    lc1.width = 14
    data_units = Reference(ws_calc, min_col=2, min_row=10, max_row=14)
    cats_time = Reference(ws_calc, min_col=1, min_row=11, max_row=14)
    lc1.add_data(data_units, titles_from_data=True)
    lc1.set_categories(cats_time)
    ws_dash.add_chart(lc1, "N5")
    
    # Chart 3: Line for Profit Trends
    lc2 = LineChart()
    lc2.title = "Profit by month"
    lc2.height = 7
    lc2.width = 14
    data_profit = Reference(ws_calc, min_col=3, min_row=10, max_row=14)
    lc2.add_data(data_profit, titles_from_data=True)
    lc2.set_categories(cats_time)
    ws_dash.add_chart(lc2, "N16")
    
    # 7. Finalize App Feel by hiding engine sheets
    ws_calc.sheet_state = 'hidden'
    ws_data.sheet_state = 'hidden'
