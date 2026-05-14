from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # Basic theme fallback
    themes = {
        "corporate_blue": {"primary": "002060", "text_on_primary": "FFFFFF"},
        "forest_green": {"primary": "2E7D32", "text_on_primary": "FFFFFF"},
        "slate_dark": {"primary": "2F4F4F", "text_on_primary": "FFFFFF"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 1. Create Data Sheet
    ws_data = wb.active
    ws_data.title = "Data_DoNotModify"
    
    # Sample Data - Trend
    trend_data = [
        ["Month", "Units Sold", "Profit"],
        ["Jan", 50601, 124812],
        ["Feb", 52000, 125000],
        ["Mar", 65481, 160228],
        ["Apr", 52970, 136337],
        ["May", 70000, 165000],
        ["Jun", 68000, 155000],
    ]
    
    # Sample Data - Category
    cat_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["Philippines", 54618, 7026, 22005],
        ["United Kingdom", 46530, 5220, 11497],
        ["United States", 36657, 6368, 22260],
    ]
    
    # Write Trend Data
    for row in trend_data:
        ws_data.append(row)
    trend_end_row = len(trend_data)
    
    # Write Category Data
    ws_data.append([]) # Blank spacer row
    cat_start_row = trend_end_row + 2
    for row in cat_data:
        ws_data.append(row)
    cat_end_row = cat_start_row + len(cat_data) - 1

    # 2. Build Charts
    # Chart 1: Units Sold (Line)
    c1 = LineChart()
    c1.title = "Units Sold Each Month"
    c1.style = 13
    c1.y_axis.title = "Units"
    data_ref = Reference(ws_data, min_col=2, min_row=1, max_row=trend_end_row)
    cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=trend_end_row)
    c1.add_data(data_ref, titles_from_data=True)
    c1.set_categories(cats_ref)
    c1.width = 14
    c1.height = 7
    
    # Chart 2: Profit Trend (Line)
    c2 = LineChart()
    c2.title = "Profit by Month"
    c2.style = 13
    data_ref = Reference(ws_data, min_col=3, min_row=1, max_row=trend_end_row)
    c2.add_data(data_ref, titles_from_data=True)
    c2.set_categories(cats_ref)
    c2.width = 14
    c2.height = 7
    
    # Chart 3: Category Profit (Stacked Bar)
    c3 = BarChart()
    c3.type = "col"
    c3.grouping = "stacked"
    c3.overlap = 100
    c3.title = "Profit by Market & Cookie Type"
    c3.style = 11
    data_ref = Reference(ws_data, min_col=2, max_col=4, min_row=cat_start_row, max_row=cat_end_row)
    cats_ref = Reference(ws_data, min_col=1, min_row=cat_start_row+1, max_row=cat_end_row)
    c3.add_data(data_ref, titles_from_data=True)
    c3.set_categories(cats_ref)
    c3.width = 18
    c3.height = 14.5

    # 3. Create Dashboard Presentation Sheet
    ws_dash = wb.create_sheet("Dashboard", 0)
    ws_dash.sheet_view.showGridLines = False  # Critical for dashboard aesthetic
    
    # Add Header Title Band
    ws_dash.merge_cells("A1:Q3")
    header_cell = ws_dash["A1"]
    header_cell.value = f"   {title}"  # Indent title slightly
    header_cell.font = Font(size=24, bold=True, color=palette["text_on_primary"])
    header_cell.alignment = Alignment(vertical="center", horizontal="left")
    
    fill = PatternFill(start_color=palette["primary"], end_color=palette["primary"], fill_type="solid")
    for row in ws_dash.iter_rows(min_row=1, max_row=3, min_col=1, max_col=17):
        for cell in row:
            cell.fill = fill

    # 4. Place Charts on Dashboard
    ws_dash.add_chart(c3, "B5")
    ws_dash.add_chart(c1, "K5")
    ws_dash.add_chart(c2, "K16")
    
    # 5. Hide Data Sheet to prevent user tampering
    ws_data.sheet_state = 'hidden'
