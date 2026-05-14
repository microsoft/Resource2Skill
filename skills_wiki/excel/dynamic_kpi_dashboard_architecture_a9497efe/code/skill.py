from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

def render_workbook(wb, *, title: str = "KPI Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    # 1. Setup Architecture (Presentation, Computation, Data)
    ws_dash = wb.active
    ws_dash.title = "3) Dashboard"
    ws_stage = wb.create_sheet("2) Staging")
    ws_data = wb.create_sheet("1) Data")
    
    # 2. Theme & Palette Definition
    palette = {
        "primary": "2B579A",
        "accent": "F39C12",
        "card_header": "EAECEF",
        "good": "C6EFCE",  # Excel native light green
        "bad": "FFC7CE",   # Excel native light red
        "bg": "F2F4F7"     # Canvas background
    }

    # ---------------------------------------------------------
    # LAYER 1: DATA
    # ---------------------------------------------------------
    headers = ["Metric", "Jan-20", "Feb-20", "Mar-20", "Apr-20", "May-20", "Jun-20"]
    ws_data.append(headers)
    ws_data.append(["Sales", 15000, 18000, 20000, 19000, 22000, 25000])  # Row 2
    ws_data.append(["AR", 5000, 5500, 5000, 4000, 4500, 4000])           # Row 3
    ws_data.append(["Marketing", 2000, 2500, 2200, 2000, 2800, 3000])    # Row 4
    ws_data.append(["New Customers", 100, 120, 110, 105, 140, 160])      # Row 5
    ws_data.append(["COGS", 9000, 10500, 11500, 11000, 13000, 14000])    # Row 6

    # ---------------------------------------------------------
    # LAYER 2: STAGING (Calculations)
    # ---------------------------------------------------------
    ws_stage.append(["KPI", "Jan-20", "Feb-20", "Mar-20", "Apr-20", "May-20", "Jun-20"])
    
    # KPI 1: DSO (Days Sales Outstanding) -> AR / Sales * 30
    ws_stage.append(["DSO"] + [f"=(Data!{chr(65+i)}3/Data!{chr(65+i)}2)*30" for i in range(1, 7)])
    ws_stage.append(["DSO Target"] + [11] * 6)
    ws_stage.append(["DSO Prior", "N/A"] + [f"={chr(66+i)}2" for i in range(5)]) # Offset 1 month
    
    # KPI 2: CAC (Customer Acquisition Cost) -> Marketing / New Customers
    ws_stage.append(["CAC"] + [f"=Data!{chr(65+i)}4/Data!{chr(65+i)}5" for i in range(1, 7)])
    ws_stage.append(["CAC Target"] + [18] * 6)
    ws_stage.append(["CAC Prior", "N/A"] + [f"={chr(66+i)}5" for i in range(5)])
    
    # KPI 3: Gross Margin -> (Sales - COGS) / Sales
    ws_stage.append(["Gross Margin"] + [f"=(Data!{chr(65+i)}2-Data!{chr(65+i)}6)/Data!{chr(65+i)}2" for i in range(1, 7)])
    ws_stage.append(["GM Target"] + [0.42] * 6)
    ws_stage.append(["GM Prior", "N/A"] + [f"={chr(66+i)}8" for i in range(5)])

    # ---------------------------------------------------------
    # LAYER 3: DASHBOARD PRESENTATION
    # ---------------------------------------------------------
    # Wash background color
    dash_bg = PatternFill(fill_type="solid", start_color=palette["bg"])
    for row in ws_dash.iter_rows(min_row=1, max_row=25, min_col=1, max_col=10):
        for cell in row:
            cell.fill = dash_bg

    # Adjust Column Proportions for Card spacing
    ws_dash.column_dimensions["A"].width = 3
    for c in ["B", "C", "D", "E", "F", "G", "H", "I"]:
        ws_dash.column_dimensions[c].width = 12
    ws_dash.sheet_view.showGridLines = False

    # Main Title
    ws_dash.merge_cells("B2:I3")
    t_cell = ws_dash["B2"]
    t_cell.value = title
    t_cell.font = Font(size=18, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(fill_type="solid", start_color=palette["primary"])
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Master Dropdown Control
    ws_dash["B5"] = "Select Month:"
    ws_dash["B5"].font = Font(bold=True)
    ws_dash["C5"].value = "Jun-20"
    ws_dash["C5"].fill = PatternFill(fill_type="solid", start_color="FFFFFF")
    ws_dash["C5"].border = Border(bottom=Side(style="medium", color="D9D9D9"))
    
    dv = DataValidation(type="list", formula1="Staging!$B$1:$G$1", allow_blank=False)
    ws_dash.add_data_validation(dv)
    dv.add(ws_dash["C5"])

    # --- KPI Card Generator Component ---
    def create_kpi_card(ws, start_col, start_row, card_title, staging_start_row, number_format="0.0", lower_is_better=True):
        thin = Side(style="thin", color="D9D9D9")
        box_border = Border(top=thin, left=thin, right=thin, bottom=thin)
        card_fill = PatternFill(fill_type="solid", start_color="FFFFFF")
        
        # Base Card Canvas
        for r in range(start_row, start_row+4):
            for c in range(start_col, start_col+4):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                cell.border = box_border

        # 1. Header Row
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+3)
        hdr_cell = ws.cell(row=start_row, column=start_col)
        hdr_cell.value = card_title
        hdr_cell.font = Font(bold=True, color="333333")
        hdr_cell.fill = PatternFill(fill_type="solid", start_color=palette["card_header"])
        hdr_cell.alignment = Alignment(horizontal="center")
        
        # 2. Main Value Area (Dynamic INDEX/MATCH)
        val_row = start_row + 1
        ws.merge_cells(start_row=val_row, start_column=start_col, end_row=val_row+1, end_column=start_col+3)
        v_cell = ws.cell(row=val_row, column=start_col)
        v_cell.value = f"=INDEX(Staging!$B${staging_start_row}:$G${staging_start_row}, 1, MATCH($C$5, Staging!$B$1:$G$1, 0))"
        v_cell.font = Font(size=24, bold=True)
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell.number_format = number_format
        
        # 3. Sub-metrics (Target & Prior Period)
        sub_row = start_row + 3
        
        # Target Sub-metric
        ws.cell(row=sub_row, column=start_col).value = "Vs. Target"
        ws.cell(row=sub_row, column=start_col).font = Font(size=9, color="555555")
        ws.cell(row=sub_row, column=start_col).alignment = Alignment(horizontal="right")
        
        target_cell = ws.cell(row=sub_row, column=start_col+1)
        target_cell.value = f"=INDEX(Staging!$B${staging_start_row+1}:$G${staging_start_row+1}, 1, MATCH($C$5, Staging!$B$1:$G$1, 0))"
        target_cell.font = Font(size=9, bold=True)
        target_cell.number_format = number_format
        target_cell.alignment = Alignment(horizontal="left")
        
        # Prior Sub-metric
        ws.cell(row=sub_row, column=start_col+2).value = "Vs. Prior"
        ws.cell(row=sub_row, column=start_col+2).font = Font(size=9, color="555555")
        ws.cell(row=sub_row, column=start_col+2).alignment = Alignment(horizontal="right")
        
        prior_cell = ws.cell(row=sub_row, column=start_col+3)
        prior_cell.value = f"=INDEX(Staging!$B${staging_start_row+2}:$G${staging_start_row+2}, 1, MATCH($C$5, Staging!$B$1:$G$1, 0))"
        prior_cell.font = Font(size=9, bold=True)
        prior_cell.number_format = number_format
        prior_cell.alignment = Alignment(horizontal="left")

        # 4. Target Threshold Conditional Formatting
        green_fill = PatternFill(start_color=palette["good"], end_color=palette["good"], fill_type="solid")
        red_fill = PatternFill(start_color=palette["bad"], end_color=palette["bad"], fill_type="solid")
        
        good_op = "lessThanOrEqual" if lower_is_better else "greaterThanOrEqual"
        bad_op = "greaterThan" if lower_is_better else "lessThan"
        
        target_col_letter = get_column_letter(start_col + 1)
        cf_range = f"{v_cell.coordinate}:{ws.cell(row=val_row+1, column=start_col+3).coordinate}"
        
        ws.conditional_formatting.add(cf_range, CellIsRule(operator=good_op, formula=[f"${target_col_letter}${sub_row}"], fill=green_fill))
        ws.conditional_formatting.add(cf_range, CellIsRule(operator=bad_op, formula=[f"${target_col_letter}${sub_row}"], fill=red_fill))

    # --- Render Categories & Cards ---
    
    # Category 1: Efficiency
    ws_dash.merge_cells("B7:I7")
    cat1 = ws_dash["B7"]
    cat1.value = "Working Capital & Efficiency KPIs"
    cat1.font = Font(bold=True, size=12, color="FFFFFF")
    cat1.fill = PatternFill(fill_type="solid", start_color=palette["accent"])
    cat1.alignment = Alignment(horizontal="center")

    create_kpi_card(ws_dash, start_col=2, start_row=9, card_title="DSO (Days Sales Outstanding)", staging_start_row=2, number_format="0", lower_is_better=True)
    create_kpi_card(ws_dash, start_col=6, start_row=9, card_title="CAC (Customer Acquisition Cost)", staging_start_row=5, number_format="$#,##0", lower_is_better=True)

    # Category 2: Profitability
    ws_dash.merge_cells("B15:I15")
    cat2 = ws_dash["B15"]
    cat2.value = "Profitability KPIs"
    cat2.font = Font(bold=True, size=12, color="FFFFFF")
    cat2.fill = PatternFill(fill_type="solid", start_color=palette["accent"])
    cat2.alignment = Alignment(horizontal="center")

    create_kpi_card(ws_dash, start_col=2, start_row=17, card_title="Gross Margin", staging_start_row=8, number_format="0.0%", lower_is_better=False)
