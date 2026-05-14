def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    
    # 1. Setup Backend Data Sheet
    data_ws = wb.create_sheet("Data")
    
    # Dataset 1: Categorical Stacked Bar Data
    data_ws.append(["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"])
    data_ws.append(["India", 62349, 23621, 21028])
    data_ws.append(["United Kingdom", 46530, 26731, 11497])
    data_ws.append(["United States", 36657, 32910, 9938])
    
    # Dataset 2: Time Series 1 (Units)
    data_ws.append([])
    data_ws.append(["Month", "Units Sold"])
    data_ws.append(["Sep", 50601])
    data_ws.append(["Oct", 95622])
    data_ws.append(["Nov", 65481])
    data_ws.append(["Dec", 52970])
    
    # Dataset 3: Time Series 2 (Profit)
    data_ws.append([])
    data_ws.append(["Month", "Profit"])
    data_ws.append(["Sep", 124812])
    data_ws.append(["Oct", 228275])
    data_ws.append(["Nov", 160228])
    data_ws.append(["Dec", 136337])
    
    # Hide the data sheet to focus on the dashboard
    data_ws.sheet_state = 'hidden'
    
    # 2. Setup Dashboard Canvas
    dash_ws = wb.active
    dash_ws.title = "Dashboard"
    dash_ws.sheet_view.showGridLines = False
    
    # Theme configuration
    theme_colors = {
        "corporate_blue": "1F4E78",
        "forest_green": "2E7D32",
        "slate_gray": "455A64"
    }
    header_bg = theme_colors.get(theme, "1F4E78")
    
    # Create Header Banner
    dash_ws.merge_cells("A1:N4")
    banner = dash_ws["A1"]
    banner.value = title
    banner.font = Font(size=28, bold=True, color="FFFFFF")
    banner.fill = PatternFill(start_color=header_bg, fill_type="solid")
    banner.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Build & Place Charts
    
    # Chart 1: Stacked Bar (Profit by Market & Product)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.y_axis.title = "Profit ($)"
    
    bar_data = Reference(data_ws, min_col=2, min_row=1, max_col=4, max_row=4)
    bar_cats = Reference(data_ws, min_col=1, min_row=2, max_row=4)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)
    bar_chart.width = 18
    bar_chart.height = 13
    dash_ws.add_chart(bar_chart, "B6")
    
    # Chart 2: Line (Units Sold)
    line_units = LineChart()
    line_units.title = "Units sold each month"
    line_units.style = 13  # Built-in openpyxl style standard
    
    u_data = Reference(data_ws, min_col=2, min_row=7, max_col=2, max_row=11)
    u_cats = Reference(data_ws, min_col=1, min_row=8, max_row=11)
    line_units.add_data(u_data, titles_from_data=True)
    line_units.set_categories(u_cats)
    line_units.width = 14
    line_units.height = 6.25
    dash_ws.add_chart(line_units, "I6")
    
    # Chart 3: Line (Profit)
    line_profit = LineChart()
    line_profit.title = "Profit by month"
    line_profit.style = 13
    
    p_data = Reference(data_ws, min_col=2, min_row=13, max_col=2, max_row=17)
    p_cats = Reference(data_ws, min_col=1, min_row=14, max_row=17)
    line_profit.add_data(p_data, titles_from_data=True)
    line_profit.set_categories(p_cats)
    line_profit.width = 14
    line_profit.height = 6.25
    dash_ws.add_chart(line_profit, "I18")
