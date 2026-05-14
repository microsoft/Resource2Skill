def render_sheet(wb, sheet_name: str = "Dashboard", *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean, grid-less dashboard presentation sheet with a branded header banner 
    and aligned charts pulling from a hidden background data sheet.
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference

    # Fallback palette resolution
    primary_bg = "003366"  # Corporate Blue
    primary_fg = "FFFFFF"  # White text
    if theme == "dark_mode":
        primary_bg = "1E1E1E"
        primary_fg = "E0E0E0"
    elif theme == "forest_green":
        primary_bg = "2E7D32"
    
    # 1. Initialize Presentation Sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        # Insert at the very front
        ws = wb.create_sheet(sheet_name, 0)
        
    # Disable Excel structural UI for an app-like feel
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False

    # 2. Build Branded Header Banner (Spans rows 1 to 3, columns A to Q)
    banner_fill = PatternFill(start_color=primary_bg, end_color=primary_bg, fill_type="solid")
    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=17):
        for cell in row:
            cell.fill = banner_fill

    ws.merge_cells("A1:Q3")
    banner_cell = ws["A1"]
    banner_cell.value = title
    banner_cell.font = Font(size=24, bold=True, color=primary_fg)
    banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=2)

    # 3. Generate Hidden Data Processing Sheet
    data_ws_name = f"{sheet_name}_Data"
    if data_ws_name in wb.sheetnames:
        data_ws = wb[data_ws_name]
    else:
        data_ws = wb.create_sheet(data_ws_name)
        # Hide the data sheet to prevent clutter
        data_ws.sheet_state = 'hidden'

        # Seed data for Stacked Column (Profit by Market & Cookie Type)
        data_ws.append(["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"])
        data_ws.append(["India", 62349, 4872, 21028])
        data_ws.append(["Philippines", 54618, 7026, 22005])
        data_ws.append(["United Kingdom", 46530, 5220, 11497])
        data_ws.append(["United States", 36657, 6368, 22260])

        # Seed data for Line Chart (Units Sold per Month)
        data_ws.append([]) # Spacer row 6
        data_ws.append(["Month", "Units Sold"]) # Row 7
        data_ws.append(["Sep", 50601])
        data_ws.append(["Oct", 95622])
        data_ws.append(["Nov", 65481])
        data_ws.append(["Dec", 52970])

    # 4. Create Stacked Bar Chart
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 11  # Standard preset style
    bar_chart.height = 10
    bar_chart.width = 16

    bar_data = Reference(data_ws, min_col=2, max_col=4, min_row=1, max_row=5)
    bar_cats = Reference(data_ws, min_col=1, min_row=2, max_row=5)
    bar_chart.add_data(bar_data, titles_from_data=True)
    bar_chart.set_categories(bar_cats)

    ws.add_chart(bar_chart, "B5")

    # 5. Create Trend Line Chart
    line_chart = LineChart()
    line_chart.title = "Units sold each month"
    line_chart.style = 13
    line_chart.legend = None  # Remove legend for cleaner trend view
    line_chart.height = 10
    line_chart.width = 14

    line_data = Reference(data_ws, min_col=2, max_col=2, min_row=7, max_row=11)
    line_cats = Reference(data_ws, min_col=1, min_row=8, max_row=11)
    line_chart.add_data(line_data, titles_from_data=True)
    line_chart.set_categories(line_cats)

    ws.add_chart(line_chart, "J5")
