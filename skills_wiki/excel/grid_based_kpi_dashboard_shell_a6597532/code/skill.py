from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", subtitle: str = "Evaluating Sales Agent Performance", theme: str = "aspect_purple", kpis: list = None, **kwargs) -> None:
    # 1. Setup Theme Registry
    themes = {
        "corporate_blue": {
            "primary": "003366",
            "bg_light": "F2F4F8",
            "accent": "00A859",
            "text_light": "FFFFFF",
            "text_dark": "333333",
            "text_muted": "666666",
            "card_bg": "FFFFFF"
        },
        "aspect_purple": {
            "primary": "4B2E83",
            "bg_light": "F2EFF5",
            "accent": "FFC000",
            "text_light": "FFFFFF",
            "text_dark": "333333",
            "text_muted": "666666",
            "card_bg": "FFFFFF"
        }
    }
    colors = themes.get(theme, themes["aspect_purple"])

    if kpis is None:
        kpis = [
            {"label": "TOTAL CALLS", "value": "16,749"},
            {"label": "REACHED", "value": "3,328"},
            {"label": "CLOSED", "value": "1,203"},
            {"label": "DEAL VALUE", "value": "$646,979"}
        ]

    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Fill entire sheet with background color to simulate the dashboard "canvas"
    bg_fill = PatternFill(start_color=colors["bg_light"], end_color=colors["bg_light"], fill_type="solid")
    for row in range(1, 41):
        for col in range(1, 16):
            ws.cell(row=row, column=col).fill = bg_fill

    # Set column widths to create a grid system with "gutters"
    widths = {
        'A': 3, 'B': 14, 'C': 14, 'D': 3, 'E': 14, 'F': 14, 'G': 3,
        'H': 14, 'I': 14, 'J': 3, 'K': 14, 'L': 14, 'M': 3
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    # 2. Header Block
    header_fill = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    for row in range(1, 5):
        for col in range(1, 14):
            ws.cell(row=row, column=col).fill = header_fill

    title_cell = ws['B2']
    title_cell.value = title
    title_cell.font = Font(name="Arial", size=24, color=colors["text_light"], bold=True)
    
    subtitle_cell = ws['B3']
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(name="Arial", size=14, color=colors["accent"])

    # 3. KPI Cards (Using merged cells and borders instead of floating shapes)
    start_cols = [2, 5, 8, 11] # B, E, H, K
    card_fill = PatternFill(start_color=colors["card_bg"], end_color=colors["card_bg"], fill_type="solid")
    top_border = Border(top=Side(style='thick', color=colors["accent"]))
    
    for i, kpi in enumerate(kpis[:4]):
        sc = start_cols[i]
        ec = sc + 1
        
        # Paint the white card background
        for row in range(6, 9):
            for col in range(sc, ec + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = card_fill
                if row == 6:
                    cell.border = top_border
                    
        # Inject Value
        val_cell = ws.cell(row=7, column=sc)
        ws.merge_cells(start_row=7, start_column=sc, end_row=7, end_column=ec)
        val_cell.value = kpi["value"]
        val_cell.font = Font(name="Arial", size=20, color=colors["text_dark"], bold=True)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Inject Label
        lbl_cell = ws.cell(row=8, column=sc)
        ws.merge_cells(start_row=8, start_column=sc, end_row=8, end_column=ec)
        lbl_cell.value = kpi["label"]
        lbl_cell.font = Font(name="Arial", size=11, color=colors["text_muted"])
        lbl_cell.alignment = Alignment(horizontal="center", vertical="top")

    # 4. Content Area 1: Performance Table with Inline Data Bars
    for row in range(11, 21):
        for col in range(2, 7): # B to F spans the width of the first two KPI cards + gutter
            ws.cell(row=row, column=col).fill = card_fill
            
    ws.cell(row=12, column=2).value = "Sales Agent KPIs"
    ws.cell(row=12, column=2).font = Font(size=14, color=colors["primary"], bold=True)
    
    headers = ["Agent", "Total Calls", "Reached", "Closed", "Deal Value"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=14, column=2+i)
        cell.value = h
        cell.font = Font(bold=True, color=colors["text_muted"])
        cell.border = Border(bottom=Side(style='thin', color="CCCCCC"))
        
    mock_data = [
        ("Alice", 1031, 128, 49, "$13,519"),
        ("Bob", 661, 73, 28, "$40,092"),
        ("Charlie", 610, 86, 67, "$45,236"),
        ("Diana", 566, 163, 26, "$38,093"),
        ("Evan", 722, 168, 16, "$17,105")
    ]
    
    for r, row_data in enumerate(mock_data):
        for c, val in enumerate(row_data):
            ws.cell(row=15+r, column=2+c).value = val
            
    # Apply Data Bar formatting to the 'Closed' column to visualize performance directly in the table
    bar_rule = DataBarRule(start_type='min', end_type='max', color=colors["primary"])
    ws.conditional_formatting.add('E15:E19', bar_rule)

    # 5. Content Area 2: Placeholder Block for Charts
    for row in range(11, 21):
        for col in range(8, 13): # H to L spans the width of the last two KPI cards + gutter
            ws.cell(row=row, column=col).fill = card_fill
            
    ws.cell(row=12, column=8).value = "Monthly Trend Analysis"
    ws.cell(row=12, column=8).font = Font(size=14, color=colors["primary"], bold=True)
    ws.cell(row=15, column=8).value = "[ Insert OpenPyXL Chart Here ]"
    ws.cell(row=15, column=8).font = Font(color=colors["text_muted"], italic=True)
