from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def render_sheet(wb, sheet_name: str, *, title: str, kpis: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Builds a clean, shape-free dashboard layout using cell coloring and borders 
    to emulate a modern web UI.
    """
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    # 1. Dashboard Canvas Prep
    ws.sheet_view.showGridLines = False

    # Color Palette (Standard modern UI fallback)
    bg_color = "F3F4F6"    # Canvas background (light gray)
    sidebar_bg = "1E3A8A"  # Navigation sidebar (dark blue)
    card_bg = "FFFFFF"     # Card background (white)
    border_color = "E5E7EB"# Subtle card border
    text_primary = "111827"
    text_secondary = "6B7280"

    # Apply canvas background across standard viewable area
    canvas_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for r in range(1, 40):
        for c in range(2, 16):
            ws.cell(row=r, column=c).fill = canvas_fill

    # 2. Left Sidebar Navigation Strip
    sidebar_fill = PatternFill(start_color=sidebar_bg, end_color=sidebar_bg, fill_type="solid")
    ws.column_dimensions['A'].width = 8
    for r in range(1, 40):
        ws.cell(row=r, column=1).fill = sidebar_fill

    # Sidebar simulated icons/links
    nav_items = ["Home", "Data", "Docs", "Help"]
    for i, item in enumerate(nav_items):
        cell = ws.cell(row=5 + (i * 3), column=1)
        cell.value = item
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Column Width Setup for Grid (Responsive padding emulation)
    ws.column_dimensions['B'].width = 2   # spacer
    ws.column_dimensions['C'].width = 15  # card 1
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 2   # spacer
    ws.column_dimensions['F'].width = 15  # card 2
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 2   # spacer
    ws.column_dimensions['I'].width = 15  # card 3
    ws.column_dimensions['J'].width = 15

    # 4. Header Section
    ws.merge_cells('C2:J3')
    header = ws['C2']
    header.value = title
    header.font = Font(color=text_primary, size=20, bold=True)
    header.alignment = Alignment(vertical="center")

    # Utility function to draw a "Card UI" module using cells
    def draw_card(start_row, end_row, start_col, end_col, title_text, value_text=""):
        c_fill = PatternFill(start_color=card_bg, end_color=card_bg, fill_type="solid")
        thin_border = Side(style='thin', color=border_color)

        # Fill and Border Application
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = c_fill
                cell.border = Border(
                    left=thin_border if c == start_col else None,
                    right=thin_border if c == end_col else None,
                    top=thin_border if r == start_row else None,
                    bottom=thin_border if r == end_row else None
                )

        # Card Title
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
        title_cell = ws.cell(row=start_row, column=start_col)
        title_cell.value = title_text
        title_cell.font = Font(color=text_secondary, size=11, bold=True)
        title_cell.alignment = Alignment(vertical="center", horizontal="center")

        # Card Value
        if value_text:
            ws.merge_cells(start_row=start_row+1, start_column=start_col, end_row=end_row, end_column=end_col)
            val_cell = ws.cell(row=start_row+1, column=start_col)
            val_cell.value = value_text
            val_cell.font = Font(color=text_primary, size=24, bold=True)
            val_cell.alignment = Alignment(vertical="center", horizontal="center")

    # 5. KPI Cards (Top Row)
    if not kpis:
        kpis = [
            {"title": "Total Sales", "value": "$2,544M"},
            {"title": "Total Profit", "value": "$890M"},
            {"title": "Total Customers", "value": "87.0M"}
        ]

    draw_card(5, 8, 3, 4, kpis[0]["title"], kpis[0]["value"])  # C5:D8
    draw_card(5, 8, 6, 7, kpis[1]["title"], kpis[1]["value"])  # F5:G8
    draw_card(5, 8, 9, 10, kpis[2]["title"], kpis[2]["value"]) # I5:J8

    # 6. Main Chart Panels (Drop zones for visuals)
    draw_card(10, 24, 3, 7, "2021-2022 Sales Trend (in millions)")  # Large panel left
    draw_card(10, 24, 9, 10, "Customer Satisfaction Overview")      # Smaller panel right
