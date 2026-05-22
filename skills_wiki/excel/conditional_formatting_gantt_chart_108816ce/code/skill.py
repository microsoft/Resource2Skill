def render(ws, anchor: str, *, tasks: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    import datetime
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule, DataBarRule
    from openpyxl.utils import coordinate_to_tuple, get_column_letter

    # Standard palette fallback (would normally be fetched via theme loader)
    palette = {
        "primary": "#4F81BD",
        "secondary": "#B8CCE4",
        "text": "#000000",
        "background": "#FFFFFF",
    }
    
    if not tasks:
        base_date = datetime.date.today()
        # Find next Monday for a clean start
        base_date = base_date + datetime.timedelta(days=(7 - base_date.weekday()) % 7)
        tasks = [
            {"name": "Phase 1: Planning", "start": base_date, "duration": 5, "progress": 1.0},
            {"name": "Phase 2: Design", "start": base_date + datetime.timedelta(days=7), "duration": 10, "progress": 0.6},
            {"name": "Phase 3: Development", "start": base_date + datetime.timedelta(days=14), "duration": 15, "progress": 0.2},
            {"name": "Phase 4: Testing", "start": base_date + datetime.timedelta(days=28), "duration": 8, "progress": 0.0},
        ]

    row, col = coordinate_to_tuple(anchor)
    headers = ["Task Name", "Start Date", "Duration", "Progress", "End Date"]
    timeline_start_col = col + len(headers)
    
    min_date = min(t["start"] for t in tasks)
    max_duration = max((t["start"] - min_date).days + t["duration"] for t in tasks)
    num_timeline_cols = max(30, max_duration + 10)

    # 1. Write Headers
    header_fill = PatternFill(start_color=palette["primary"][1:], fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col+i, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    # Generate sequential date headers for the timeline
    for i in range(num_timeline_cols):
        c = ws.cell(row=row, column=timeline_start_col+i, value=min_date + datetime.timedelta(days=i))
        c.number_format = "d-mmm"
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(textRotation=90, horizontal="center", vertical="center")

    # 2. Write Task Data
    for r_idx, task in enumerate(tasks, start=row+1):
        ws.cell(row=r_idx, column=col, value=task["name"])
        
        c_start = ws.cell(row=r_idx, column=col+1, value=task["start"])
        c_start.number_format = "yyyy-mm-dd"
        
        ws.cell(row=r_idx, column=col+2, value=task["duration"])
        
        c_prog = ws.cell(row=r_idx, column=col+3, value=task["progress"])
        c_prog.number_format = "0%"
        
        # End date formula handles weekends natively
        start_cell_ref = f"{get_column_letter(col+1)}{r_idx}"
        dur_cell_ref = f"{get_column_letter(col+2)}{r_idx}"
        c_end = ws.cell(row=r_idx, column=col+4, value=f"=WORKDAY({start_cell_ref}, MAX({dur_cell_ref}-1, 0))")
        c_end.number_format = "yyyy-mm-dd"

    # 3. Column Widths Setup
    ws.column_dimensions[get_column_letter(col)].width = 25
    ws.column_dimensions[get_column_letter(col+1)].width = 12
    ws.column_dimensions[get_column_letter(col+2)].width = 10
    ws.column_dimensions[get_column_letter(col+3)].width = 10
    ws.column_dimensions[get_column_letter(col+4)].width = 12
    for i in range(num_timeline_cols):
        ws.column_dimensions[get_column_letter(timeline_start_col + i)].width = 3.5

    # 4. Draw Grid Borders
    thin_border = Border(left=Side(style='thin', color='D9D9D9'),
                         right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'),
                         bottom=Side(style='thin', color='D9D9D9'))
    for r in range(row, row + len(tasks) + 1):
        for c in range(col, timeline_start_col + num_timeline_cols):
            ws.cell(row=r, column=c).border = thin_border

    # 5. Conditional Formatting: Progress Column Data Bar
    prog_col_ltr = get_column_letter(col+3)
    progress_range = f"{prog_col_ltr}{row+1}:{prog_col_ltr}{row+len(tasks)}"
    
    # Excel requires 8-char hex (FF + 6-char color) for Data Bars
    bar_color = "FF" + palette["primary"][1:]
    progress_rule = DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=bar_color)
    ws.conditional_formatting.add(progress_range, progress_rule)

    # 6. Conditional Formatting: Timeline Gantt Logics
    tl_date_ref = f"{get_column_letter(timeline_start_col)}${row}"
    start_col_ltr = get_column_letter(col+1)
    dur_col_ltr = get_column_letter(col+2)
    end_col_ltr = get_column_letter(col+4)
    
    # Relative row / absolute column bindings for the CF engine
    curr_start_ref = f"${start_col_ltr}{row+1}"
    curr_dur_ref = f"${dur_col_ltr}{row+1}"
    curr_prog_ref = f"${prog_col_ltr}{row+1}"
    curr_end_ref = f"${end_col_ltr}{row+1}"

    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    completed_fill = PatternFill(start_color=palette["primary"][1:], end_color=palette["primary"][1:], fill_type="solid")
    incomplete_fill = PatternFill(start_color=palette["secondary"][1:], end_color=palette["secondary"][1:], fill_type="solid")

    # The order rules are added matters. First evaluated = highest priority.
    # 1. Weekends overwrite everything
    rule_weekend = FormulaRule(formula=[f"WEEKDAY({tl_date_ref}, 2)>5"], fill=gray_fill)
    
    # 2. Completed portions overwrite incomplete
    rule_completed = FormulaRule(
        formula=[f"AND({curr_prog_ref}>0, {tl_date_ref}>={curr_start_ref}, {tl_date_ref}<=WORKDAY({curr_start_ref}, MAX(ROUND({curr_dur_ref}*{curr_prog_ref}, 0)-1, 0)))"],
        fill=completed_fill
    )
    
    # 3. Base duration coverage (incomplete)
    rule_incomplete = FormulaRule(
        formula=[f"AND({tl_date_ref}>={curr_start_ref}, {tl_date_ref}<={curr_end_ref})"],
        fill=incomplete_fill
    )

    timeline_range = f"{get_column_letter(timeline_start_col)}{row+1}:{get_column_letter(timeline_start_col+num_timeline_cols-1)}{row+len(tasks)}"
    
    ws.conditional_formatting.add(timeline_range, rule_weekend)
    ws.conditional_formatting.add(timeline_range, rule_completed)
    ws.conditional_formatting.add(timeline_range, rule_incomplete)
