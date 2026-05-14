def render(ws, anchor: str, *, title: str, label: str, value_ref: str, pct_ref: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a cell-based KPI Card mimicking a linked shape graphic from the dashboard tutorial.
    
    :param ws: openpyxl worksheet
    :param anchor: Top-left cell coordinate (e.g., 'E2')
    :param title: The main title of the card (e.g., 'Asia')
    :param label: The metric label (e.g., 'Revenue')
    :param value_ref: Source cell reference for the primary value (e.g., 'B2')
    :param pct_ref: Source cell reference for the percentage (e.g., 'C2')
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    
    # In a full framework, you would extract these from the loaded theme palette.
    # We fallback to the Navy Blue and White seen in the tutorial.
    bg_color = "1F3864"
    fg_color = "FFFFFF"
    
    col_str, row_str = coordinate_from_string(anchor)
    start_col = column_index_from_string(col_str)
    start_row = int(row_str)
    
    # 1. Map Cell Positions
    title_cell = ws.cell(row=start_row, column=start_col)
    label_cell = ws.cell(row=start_row + 1, column=start_col)
    val_cell = ws.cell(row=start_row + 2, column=start_col)
    pct_cell = ws.cell(row=start_row + 1, column=start_col + 2)
    
    # 2. Merge Block Areas
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 1, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 2, start_column=start_col, end_row=start_row + 2, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 1, start_column=start_col + 2, end_row=start_row + 2, end_column=start_col + 2)
    
    # 3. Inject Content & Live Links
    title_cell.value = title
    label_cell.value = label
    val_cell.value = f"={value_ref}"
    pct_cell.value = f"={pct_ref}"
    
    # 4. Global Card Styling
    card_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    
    for r in range(start_row, start_row + 3):
        for c in range(start_col, start_col + 3):
            # Apply base styling to all cells in the 3x3 grid to ensure merged regions are fully colored
            cell = ws.cell(row=r, column=c)
            cell.fill = card_fill
            cell.alignment = center_align
    
    # 5. Granular Typography & Borders
    title_cell.font = Font(color=fg_color, bold=True, size=14)
    label_cell.font = Font(color=fg_color, size=11)
    val_cell.font = Font(color=fg_color, bold=True, size=14)
    pct_cell.font = Font(color=fg_color, bold=True, size=14)
    
    # Simulate the disconnected circle shape with a thick border around the percentage
    white_border = Border(
        left=Side(style="thick", color=fg_color),
        right=Side(style="thick", color=fg_color),
        top=Side(style="thick", color=fg_color),
        bottom=Side(style="thick", color=fg_color)
    )
    pct_cell.border = white_border
    
    # 6. Formatting
    val_cell.number_format = '"$"#,##0'
    pct_cell.number_format = '0%'
