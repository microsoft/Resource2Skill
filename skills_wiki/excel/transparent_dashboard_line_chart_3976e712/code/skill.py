def render(ws, anchor: str, *, theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    # Theme definitions
    themes = {
        "corporate_blue": {"primary": "002060", "secondary": "C00000"},
        "modern_dark": {"primary": "4F81BD", "secondary": "9BBB59"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # Setup sample trend data isolated away from the dashboard UI (e.g., column AA)
    data = [
        ["Month", "2021", "2022"],
        ["Jan", 201.9, 215.3],
        ["Feb", 204.2, 217.6],
        ["Mar", 198.6, 220.1],
        ["Apr", 199.2, 206.4],
        ["May", 206.4, 204.3],
        ["Jun", 195.1, 203.0]
    ]

    start_col = 27
    start_row = 1
    for row_idx, row_data in enumerate(data, start=start_row):
        for col_idx, value in enumerate(row_data, start=start_col):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Initialize Line Chart
    chart = LineChart()
    chart.title = "Sales Trend (in millions)"
    chart.width = 14
    chart.height = 7
    chart.legend.position = "b"  # Move legend to the bottom

    # Make the chart background and plot area transparent to blend into dashboard UI
    chart.graphical_properties = GraphicalProperties()
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line = LineProperties(noFill=True)
    
    chart.plot_area.graphicalProperties = GraphicalProperties()
    chart.plot_area.graphicalProperties.noFill = True

    # Bind data
    data_ref = Reference(ws, min_col=start_col+1, min_row=start_row, max_col=start_col+2, max_row=start_row+6)
    cats_ref = Reference(ws, min_col=start_col, min_row=start_row+1, max_row=start_row+6)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    # Clean up axes (Remove gridlines, hardcode min/max to zoom into variance)
    chart.y_axis.scaling.min = 180
    chart.y_axis.scaling.max = 230
    chart.y_axis.majorGridlines = None

    # Apply custom line and marker formatting
    colors = [palette["primary"], palette["secondary"]]

    for i, series in enumerate(chart.series):
        color = colors[i % len(colors)]

        # Series line formatting (Width is set in EMUs: 20000 EMU ~= 1.5 pt)
        series.graphicalProperties.line = LineProperties(solidFill=color, w=20000)

        # Custom Marker: Circular, White Fill, Border matches Line Color
        marker = Marker(symbol="circle", size=5)
        marker.graphicalProperties = GraphicalProperties(solidFill="FFFFFF")
        marker.graphicalProperties.line = LineProperties(solidFill=color, w=15000)
        series.marker = marker

    ws.add_chart(chart, anchor)
