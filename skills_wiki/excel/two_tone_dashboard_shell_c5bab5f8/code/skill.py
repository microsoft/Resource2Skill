from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Executive Dashboard", theme: str = "botanical_green", **kwargs) -> None:
    """
    Renders a two-tone dashboard shell with a dark sidebar for KPIs and a light canvas with white widget cards.
    """
    # Self-contained theme resolver
    themes = {
        "corporate_blue": {
            "sidebar": "1F4E78", "canvas": "EAEDED", "card": "FFFFFF", 
            "text_light": "FFFFFF", "text_dark": "333333", "border": "CCCCCC"
        },
        "botanical_green": {
            "sidebar": "2A4B3C", "canvas": "E8F0EA", "card": "FFFFFF", 
            "text_light": "FFFFFF", "text_dark": "1A2F25", "border": "B4C8BB"
        }
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # Create sheet and hide gridlines
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Setup Fills
    sidebar_fill = PatternFill(start_color=palette["sidebar"], fill_type="solid")
    canvas_fill = PatternFill(start_color=palette["canvas"], fill_type="solid")
    card_fill = PatternFill(start_color=palette["card"], fill_type="solid")
    
    # Configure Column Widths
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 2
    ws.column_dimensions['D'].width = 2
    for col in range(5, 16):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # Paint Backgrounds (Sidebar vs Canvas)
    for row in range(1, 40):
        for col in range(1, 16):
            cell = ws.cell(row=row, column=col)
            if col <= 3:
                cell.fill = sidebar_fill
            else:
                cell.fill = canvas_fill
    
    # Render Sidebar Title
    title_cell = ws.cell(row=3, column=2, value=title)
    title_cell.font = Font(name="Arial", size=20, bold=True, color=palette["text_light"])
    title_cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Render Vertical KPI Strip in Sidebar
    kpis = [
        ("Total Orders", "2,400"), 
        ("Total Quantity", "11,997"), 
        ("Total Revenue", "$649.0k"), 
        ("Avg. Rating", "4.0")
    ]
    
    current_row = 7
    for label, val in kpis:
        lbl_cell = ws.cell(row=current_row, column=2, value=label)
        lbl_cell.font = Font(name="Arial", size=10, color=palette["text_light"])
        
        val_cell = ws.cell(row=current_row+1, column=2, value=val)
        val_cell.font = Font(name="Arial", size=16, bold=True, color=palette["text_light"])
        
        current_row += 4

    # Define Widget Cards (start_row, start_col, end_row, end_col, title)
    widgets = [
        (3, 5, 12, 9, "Last 13 Week Trends - Qty & Amount"),
        (3, 10, 12, 14, "How Customers Like to Buy"),
        (14, 5, 28, 9, "Popular Products - Breakdown"),
        (14, 10, 28, 14, "Customer Geographic Distribution")
    ]

    # Render Widget Cards
    for sr, sc, er, ec, w_title in widgets:
        for r in range(sr, er + 1):
            for c in range(sc, ec + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Calculate and apply perimeter borders
                b_left = Side(style='thin', color=palette["border"]) if c == sc else None
                b_right = Side(style='thin', color=palette["border"]) if c == ec else None
                b_top = Side(style='thin', color=palette["border"]) if r == sr else None
                b_bottom = Side(style='thin', color=palette["border"]) if r == er else None
                
                if b_left or b_right or b_top or b_bottom:
                    cell.border = Border(left=b_left, right=b_right, top=b_top, bottom=b_bottom)

        # Card Title Header
        w_title_cell = ws.cell(row=sr+1, column=sc+1, value=w_title)
        w_title_cell.font = Font(name="Arial", size=11, bold=True, color=palette["text_dark"])
