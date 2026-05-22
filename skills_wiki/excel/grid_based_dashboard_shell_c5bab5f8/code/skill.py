from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Creates a modern, grid-based dashboard layout with a dark sidebar and white card drop-zones.
    """
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

    # 1. Clean canvas
    ws.sheet_view.showGridLines = False

    # 2. Theme definitions
    palettes = {
        "corporate_blue": {
            "sidebar_bg": "1A365D",   # Dark navy
            "sidebar_fg": "FFFFFF",   # White text
            "main_bg": "EDF2F7",      # Light gray/blue background
            "card_bg": "FFFFFF",      # Pure white cards
            "card_border": "CBD5E0",  # Subtle gray border
            "text_main": "2D3748",    # Dark gray text
        },
        "forest_green": {
            "sidebar_bg": "276749",
            "sidebar_fg": "FFFFFF",
            "main_bg": "F0FFF4",
            "card_bg": "FFFFFF",
            "card_border": "C6F6D5",
            "text_main": "22543D",
        }
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    sidebar_fill = PatternFill("solid", fgColor=palette["sidebar_bg"])
    main_fill = PatternFill("solid", fgColor=palette["main_bg"])
    card_fill = PatternFill("solid", fgColor=palette["card_bg"])

    # 3. Set structural column widths
    ws.column_dimensions['A'].width = 3   # Sidebar left margin
    ws.column_dimensions['B'].width = 25  # Sidebar content
    ws.column_dimensions['C'].width = 3   # Sidebar right margin
    ws.column_dimensions['D'].width = 3   # Main left margin
    
    # Content grid columns
    for col in range(5, 18):
        ws.column_dimensions[get_column_letter(col)].width = 10

    # 4. Paint Backgrounds (Simulating full screen height)
    max_row_layout = 40
    for row in range(1, max_row_layout + 1):
        # Sidebar
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = sidebar_fill
        # Main stage
        for col in range(4, 18):
            ws.cell(row=row, column=col).fill = main_fill

    # 5. Sidebar Title
    title_cell = ws.cell(row=3, column=2, value=title.upper())
    title_cell.font = Font(color=palette["sidebar_fg"], size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 6. Sidebar KPI Placeholders
    # In a real pipeline, these values would be injected dynamically from data
    kpis = [
        ("TOTAL REVENUE", "$649.0K"),
        ("TOTAL ORDERS", "2,400"),
        ("AVG RATING", "4.0"),
        ("DAYS TO SHIP", "2.3")
    ]
    
    start_row = 7
    for label, val in kpis:
        # KPI Label
        lc = ws.cell(row=start_row, column=2, value=label)
        lc.font = Font(color=palette["sidebar_fg"], size=10, bold=False)
        lc.alignment = Alignment(horizontal="center")
        
        # KPI Value
        vc = ws.cell(row=start_row + 1, column=2, value=val)
        vc.font = Font(color=palette["sidebar_fg"], size=22, bold=True)
        vc.alignment = Alignment(horizontal="center")
        
        start_row += 5

    # 7. Helper to create a "Card" (White box container for charts/tables)
    def create_card(min_col, min_row, max_col, max_row, card_title):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                
                # Apply borders conditionally to form a bounding box
                top = Side(style='thin', color=palette["card_border"]) if r == min_row else None
                bottom = Side(style='thin', color=palette["card_border"]) if r == max_row else None
                left = Side(style='thin', color=palette["card_border"]) if c == min_col else None
                right = Side(style='thin', color=palette["card_border"]) if c == max_col else None
                
                cell.border = Border(top=top, bottom=bottom, left=left, right=right)

        # Apply Card Title
        tc = ws.cell(row=min_row + 1, column=min_col + 1, value=card_title)
        tc.font = Font(color=palette["text_main"], size=12, bold=True)
        tc.alignment = Alignment(vertical="center")

    # 8. Stamp the Card Layouts into the Main Area
    # Top Row Cards
    create_card(min_col=5,  min_row=3, max_col=10, max_row=15, card_title="Last 13 Weeks Trend")
    create_card(min_col=12, min_row=3, max_col=17, max_row=15, card_title="Purchase Demographics")

    # Bottom Row Cards
    create_card(min_col=5,  min_row=17, max_col=10, max_row=30, card_title="Popular Products by Volume")
    create_card(min_col=12, min_row=17, max_col=17, max_row=30, card_title="Shipment Duration Analysis")
