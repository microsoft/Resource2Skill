def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Financial Model", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    
    # Load theme safely assuming _helpers is available in the environment
    try:
        import _helpers
        theme_colors = _helpers.get_theme_colors(theme)
    except ImportError:
        theme_colors = {"header_bg": "4F81BD", "header_fg": "FFFFFF", "accent_bg": "DCE6F1", "accent_fg": "000000"}

    ws = wb.create_sheet(sheet_name)
    
    # Sheet Title & Scenario Toggle
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    
    ws["I2"] = "Active Scenario (1=Base, 2=Downside):"
    ws["I2"].font = Font(bold=True)
    ws["I2"].alignment = Alignment(horizontal="right")
    
    ws["J2"] = 1
    border_thin = Side(style='thin', color='000000')
    ws["J2"].font = Font(bold=True)
    ws["J2"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["J2"].alignment = Alignment(horizontal="center")
    ws["J2"].border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["J2"])
    
    # Column Headers
    cols = ['C', 'D', 'E', 'F', 'G']
    ws["A3"] = "Line Items"
    ws["B3"] = "Unit"
    for i, col in enumerate(cols, start=1):
        ws[f"{col}3"] = f"Year {i}"
        
    # Line Item Labels
    labels = {
        5: ("Revenue", "$"),
        6: ("Cost of Goods Sold", "$"),
        7: ("Gross Profit", "$"),
        8: ("Gross Profit Margin", "%"),
        9: ("Operating Expenses", "$"),
        10: ("Operating Profit", "$"),
        11: ("Operating Profit Margin", "%"),
        12: ("Corporate Tax", "$"),
        13: ("Profit / (Loss)", "$"),
        16: ("Number of Orders", "#"),
        17: ("Order Growth Rate", "%"),
        18: ("Average Order Value", "$"),
        19: ("COGS per order", "$"),
        20: ("Fixed Operating Expenses", "$"),
        21: ("Corporate Tax Rate", "%")
    }
    
    for row_idx, (label, unit) in labels.items():
        ws[f"A{row_idx}"] = label
        ws[f"B{row_idx}"] = unit
        if row_idx >= 16:
            ws[f"A{row_idx+8}"] = label
            ws[f"B{row_idx+8}"] = unit
            ws[f"A{row_idx+16}"] = label
            ws[f"B{row_idx+16}"] = unit

    # Block Headers
    blocks = {4: "Income Statement", 15: "Live Case Assumptions", 23: "Scenario 1 (Base Case)", 31: "Scenario 2 (Downside Case)"}
    for row_idx, text in blocks.items():
        ws[f"A{row_idx}"] = text
            
    # Model Formulas (Columns C through G)
    for col in cols:
        # Income Statement logic connected strictly to Live Case
        ws[f"{col}5"] = f"={col}16*{col}18" # Rev = Orders * AOV
        ws[f"{col}6"] = f"={col}16*{col}19" # COGS = Orders * Unit COGS
        ws[f"{col}7"] = f"={col}5-{col}6"   # GP
        ws[f"{col}8"] = f"=IF({col}5=0, 0, {col}7/{col}5)" # GM %
        ws[f"{col}9"] = f"={col}20" # Fixed OpEx
        ws[f"{col}10"] = f"={col}7-{col}9" # Op Profit
        ws[f"{col}11"] = f"=IF({col}5=0, 0, {col}10/{col}5)" # Op Margin
        ws[f"{col}12"] = f'=IF({col}10<0, "NA", {col}10*{col}21)' # Tax (Show NA if loss)
        ws[f"{col}13"] = f"=IFERROR({col}10-{col}12, {col}10)" # Net Income (Safely ignore text NA)
        
        # Live Case pulling from Scenarios via CHOOSE
        for row in range(16, 22):
            if row == 17 and col == 'C': 
                continue # Year 1 growth rate is intentionally blank
            ws[f"{col}{row}"] = f"=CHOOSE($J$2, {col}{row+8}, {col}{row+16})"

    # Populate Scenario 1 (Base Case) Data
    ws['C24'], ws['C26'], ws['C27'], ws['C28'], ws['C29'] = 3000, 40.0, 10.0, 50000, 0.2
    for i, col in enumerate(cols[1:], start=1):
        prev = cols[i-1]
        ws[f"{col}24"] = f"={prev}24*(1+{col}25)"
        ws[f"{col}25"] = [0.5, 0.4, 0.3, 0.2][i-1]
        ws[f"{col}26"] = 40.0
        ws[f"{col}27"] = 10.0
        ws[f"{col}28"] = 50000 + (i * 5000)
        ws[f"{col}29"] = 0.2

    # Populate Scenario 2 (Downside Case) Data
    ws['C32'], ws['C34'], ws['C35'], ws['C36'], ws['C37'] = 2000, 35.0, 12.0, 40000, 0.2
    for i, col in enumerate(cols[1:], start=1):
        prev = cols[i-1]
        ws[f"{col}32"] = f"={prev}32*(1+{col}33)"
        ws[f"{col}33"] = [0.2, 0.15, 0.1, 0.05][i-1]
        ws[f"{col}34"] = 35.0
        ws[f"{col}35"] = 12.0
        ws[f"{col}36"] = 40000
        ws[f"{col}37"] = 0.2

    # Apply Formatting
    for row in range(5, 38):
        unit = ws[f"B{row}"].value
        if not unit: continue
        
        # Determine number format
        fmt = "#,##0"
        if unit == "$":
            if row in [18, 19, 26, 27, 34, 35]: # Unit costs
                fmt = '"$"#,##0.00'
            else: # Macro financials
                fmt = '_("$"* #,##0_);_("$"* \(#,##0\);_("$"* "-"_);_(@_)'
        elif unit == "%":
            fmt = "0.0%"
            
        for col in cols:
            ws[f"{col}{row}"].number_format = fmt

    # Apply Design Theme
    top_bottom_border = Border(top=border_thin, bottom=border_thin)
    
    for row in [7, 10, 13]: # Margin/Profit lines
        for col in ['A', 'B'] + cols:
            ws[f"{col}{row}"].font = Font(bold=True)
            ws[f"{col}{row}"].border = top_bottom_border

    for r in [4, 15, 23, 31]: # Block Headers
        for col in ['A', 'B'] + cols:
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=theme_colors.get("accent_bg", "DCE6F1"))
            ws[f"{col}{r}"].font = Font(bold=True, color=theme_colors.get("accent_fg", "000000"))

    for col in ['A', 'B'] + cols: # Column Headers
        ws[f"{col}3"].fill = PatternFill("solid", fgColor=theme_colors.get("header_bg", "4F81BD"))
        ws[f"{col}3"].font = Font(bold=True, color=theme_colors.get("header_fg", "FFFFFF"))

    # Layout dimensions
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 8
    for col in cols:
        ws.column_dimensions[col].width = 15
