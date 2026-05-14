from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_sheet(wb, sheet_name: str, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a clean, presentation-ready dashboard layout containing a header, 
    a left-side control panel, and a grid of charts fed by hidden aggregated data.
    """
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    # Theme palette fallback
    themes = {
        "corporate_blue": {"header_bg": "203764", "header_fg": "FFFFFF", "panel_bg": "F2F2F2"},
        "light_green": {"header_bg": "375623", "header_fg": "FFFFFF", "panel_bg": "E2EFDA"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    # 1. Header Layout
    ws.merge_cells("A1:P2")
    header_cell = ws["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=palette["header_fg"])
    header_cell.alignment = Alignment(vertical="center", horizontal="left")
    header_fill = PatternFill("solid", fgColor=palette["header_bg"])
    
    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=16):
        for cell in row:
            cell.fill = header_fill

    # 2. Control Panel Area (Visual placeholder for Slicers)
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 25
    panel_fill = PatternFill("solid", fgColor=palette["panel_bg"])
    
    ws["B4"].value = "Filters / Controls"
    ws["B4"].font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, max_row=35, min_col=2, max_col=2):
        for cell in row:
            if cell.row != 4:
                cell.fill = panel_fill

    # 3. Mock Aggregated Data for Charts (Stored in AA onwards)
    # Stacked Column Data: Profit by Market & Cookie
    market_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62000, 23000, 21000],
        ["United Kingdom", 46000, 24000, 22000],
        ["United States", 36000, 32000, 9000]
    ]
    for r_idx, row_data in enumerate(market_data, start=1):
        for c_idx, val in enumerate(row_data, start=27): # Column 27 is 'AA'
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Line Chart Data: Monthly Units & Profit
    month_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    for r_idx, row_data in enumerate(month_data, start=10):
        for c_idx, val in enumerate(row_data, start=27):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # 4. Insert Charts
    # Chart 1: Stacked Column (Left Side)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Cookie Type"
    bar_chart.style = 10
    bar_chart.height = 14
    bar_chart.width = 16
    
    data_ref = Reference(ws, min_col=28, min_row=1, max_col=30, max_row=4)
    cats_ref = Reference(ws, min_col=27, min_row=2, max_row=4)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    ws.add_chart(bar_chart, "D4")

    # Chart 2: Line Chart - Units Sold (Right Side, Top)
    line_chart1 = LineChart()
    line_chart1.title = "Units sold each month"
    line_chart1.style = 13
    line_chart1.height = 9
    line_chart1.width = 14
    
    data_ref2 = Reference(ws, min_col=28, min_row=10, max_row=14)
    cats_ref2 = Reference(ws, min_col=27, min_row=11, max_row=14)
    line_chart1.add_data(data_ref2, titles_from_data=True)
    line_chart1.set_categories(cats_ref2)
    ws.add_chart(line_chart1, "J4")

    # Chart 3: Line Chart - Profit by Month (Right Side, Bottom)
    line_chart2 = LineChart()
    line_chart2.title = "Profit by month"
    line_chart2.style = 13
    line_chart2.height = 9
    line_chart2.width = 14
    
    data_ref3 = Reference(ws, min_col=29, min_row=10, max_row=14)
    line_chart2.add_data(data_ref3, titles_from_data=True)
    line_chart2.set_categories(cats_ref2)
    ws.add_chart(line_chart2, "J19")

    # 5. Hide the helper data columns
    for col in ['AA', 'AB', 'AC', 'AD']:
        ws.column_dimensions[col].hidden = True
