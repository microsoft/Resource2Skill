from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, title: str = "Asia", value: float = 369989, sub_value: float = 0.05, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a stylized cell-based KPI tile, replicating the look of shape-based KPIs.
    """
    # Map theme to palette (fallback to the video's Navy/Dark theme)
    palettes = {
        "corporate_blue": {"bg": "002060", "fg": "FFFFFF"},
        "dark_mode": {"bg": "222222", "fg": "FFFFFF"},
        "emerald": {"bg": "005830", "fg": "FFFFFF"}
    }
    colors = palettes.get(theme, palettes["corporate_blue"])
    bg_color = colors["bg"]
    fg_color = colors["fg"]
    
    # Parse anchor
    start_row, start_col = coordinate_to_tuple(anchor)
    end_row = start_row + 2
    end_col = start_col + 1
    
    # Set values in the anchor column
    ws.cell(row=start_row, column=start_col, value=title)
    ws.cell(row=start_row+1, column=start_col, value=value)
    ws.cell(row=start_row+2, column=start_col, value=sub_value)
    
    # Merge cells for each row in the KPI card to create a single readable block
    for r in range(start_row, end_row + 1):
        ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=end_col)
    
    # Define styles
    title_font = Font(name="Calibri", size=12, color=fg_color)
    value_font = Font(name="Calibri", size=18, bold=True, color=fg_color)
    sub_font = Font(name="Calibri", size=12, italic=True, color=fg_color)
    
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    
    border_side = Side(border_style="medium", color=bg_color)
    card_border = Border(top=border_side, bottom=border_side, left=border_side, right=border_side)
    
    # Apply number formatting
    ws.cell(row=start_row+1, column=start_col).number_format = '"$"#,##0'
    ws.cell(row=start_row+2, column=start_col).number_format = '0%'
    
    # Apply styling to all cells in the simulated shape block
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.alignment = center_align
            cell.border = card_border
            
            # Apply specific fonts to the visible left column
            if c == start_col:
                if r == start_row:
                    cell.font = title_font
                elif r == start_row + 1:
                    cell.font = value_font
                elif r == start_row + 2:
                    cell.font = sub_font
                
    # Adjust dimensions to make the tile feel like a spacious shape
    ws.column_dimensions[get_column_letter(start_col)].width = 10
    ws.column_dimensions[get_column_letter(start_col+1)].width = 10
    ws.row_dimensions[start_row].height = 18
    ws.row_dimensions[start_row+1].height = 28  # Taller row for the large KPI value
    ws.row_dimensions[start_row+2].height = 18
