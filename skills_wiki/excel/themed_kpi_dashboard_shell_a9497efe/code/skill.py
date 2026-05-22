from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str, theme: str = "corporate_blue", kpi_data: dict = None, **kwargs) -> None:
    """
    Renders a KPI Dashboard with grouped metric cards, large typography, and conditional formatting.
    """
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    
    # Theme palette fallback
    colors = {
        "primary": "4F81BD",       # Group Header BG
        "primary_fg": "FFFFFF",    # Group Header Text
        "card_bg": "F2F2F2",       # KPI Title/Context BG
        "border": "D9D9D9",        # KPI Card Border
        "success_bg": "C6EFCE",    # Good KPI
        "success_fg": "006100",
        "danger_bg": "FFC7CE",     # Bad KPI
        "danger_fg": "9C0006"
    }

    # Default realistic data if none provided
    if not kpi_data:
        kpi_data = {
            "Working Capital Efficiency": [
                {"name": "DSO (Days Sales Outstanding)", "value": 41, "target": 45, "prior": 53, "good_is_up": False, "format": "0"},
                {"name": "DPO (Days Payables Outstanding)", "value": 90, "target": 90, "prior": 89, "good_is_up": True, "format": "0"},
                {"name": "Non-Current AR %", "value": 0.12, "target": 0.03, "prior": 0.11, "good_is_up": False, "format": "0%"}
            ],
            "Sales KPIs": [
                {"name": "CAC (Customer Acq Cost)", "value": 17725, "target": 15000, "prior": 18236, "good_is_up": False, "format": "$#,##0"},
                {"name": "Sales vs. Budget %", "value": 1.27, "target": 1.00, "prior": 0.98, "good_is_up": True, "format": "0%"},
                {"name": "Gross Margin", "value": 0.26, "target": 0.38, "prior": 0.26, "good_is_up": True, "format": "0%"}
            ],
            "Cost KPIs": [
                {"name": "OPEX Actual vs. Budget", "value": 1.05, "target": 1.00, "prior": 1.07, "good_is_up": False, "format": "0%"},
                {"name": "Cost Per Full Time Employee", "value": 12965, "target": 12500, "prior": 13200, "good_is_up": False, "format": "$#,##0"}
            ]
        }

    # Styles
    thin_border = Border(
        left=Side(style='thin', color=colors["border"]),
        right=Side(style='thin', color=colors["border"]),
        top=Side(style='thin', color=colors["border"]),
        bottom=Side(style='thin', color=colors["border"])
    )
    group_header_fill = PatternFill(start_color=colors["primary"], end_color=colors["primary"], fill_type="solid")
    card_header_fill = PatternFill(start_color=colors["card_bg"], end_color=colors["card_bg"], fill_type="solid")
    
    # Dashboard Header
    ws["B2"] = title
    ws["B2"].font = Font(size=20, bold=True, color=colors["primary"])
    
    current_row = 4
    
    # Iterate through KPI Groups
    for group_name, kpis in kpi_data.items():
        # Group Header Bar (spans columns B through G, assuming max 3 KPIs per row, 2 cols each)
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=7)
        hdr_cell = ws.cell(row=current_row, column=2, value=group_name)
        hdr_cell.font = Font(size=14, bold=True, color=colors["primary_fg"])
        hdr_cell.fill = group_header_fill
        hdr_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        current_row += 2
        col = 2
        
        # Iterate through KPIs in the group
        for kpi in kpis:
            # Wrap to next row if we exceed 3 KPIs (6 columns)
            if col > 7:
                col = 2
                current_row += 4
                
            # 1. KPI Title (Merged over 2 columns)
            ws.merge_cells(start_row=current_row, start_column=col, end_row=current_row, end_column=col+1)
            title_cell = ws.cell(row=current_row, column=col, value=kpi["name"])
            title_cell.font = Font(bold=True, color="333333")
            title_cell.fill = card_header_fill
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # 2. KPI Value (Merged over 2 columns)
            ws.merge_cells(start_row=current_row+1, start_column=col, end_row=current_row+1, end_column=col+1)
            val_cell = ws.cell(row=current_row+1, column=col, value=kpi["value"])
            val_cell.font = Font(size=24, bold=True)
            val_cell.number_format = kpi["format"]
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Apply Conditional Formatting to Value
            green_fill = PatternFill(start_color=colors["success_bg"], end_color=colors["success_bg"], fill_type="solid")
            red_fill = PatternFill(start_color=colors["danger_bg"], end_color=colors["danger_bg"], fill_type="solid")
            green_font = Font(size=24, bold=True, color=colors["success_fg"])
            red_font = Font(size=24, bold=True, color=colors["danger_fg"])
            
            target_str = str(kpi["target"])
            if kpi["good_is_up"]:
                good_op, bad_op = 'greaterThanOrEqual', 'lessThan'
            else:
                good_op, bad_op = 'lessThanOrEqual', 'greaterThan'
                
            rule_good = CellIsRule(operator=good_op, formula=[target_str], stopIfTrue=True, fill=green_fill, font=green_font)
            rule_bad = CellIsRule(operator=bad_op, formula=[target_str], stopIfTrue=True, fill=red_fill, font=red_font)
            
            coord = val_cell.coordinate
            ws.conditional_formatting.add(f"{coord}:{coord}", rule_good)
            ws.conditional_formatting.add(f"{coord}:{coord}", rule_bad)

            # 3. Context Row (Vs. Target | Vs. Prior)
            ctx_target = ws.cell(row=current_row+2, column=col, value=f"Vs. Target: {kpi['target']}")
            ctx_prior = ws.cell(row=current_row+2, column=col+1, value=f"Vs. Prior: {kpi['prior']}")
            
            for ctx_cell in [ctx_target, ctx_prior]:
                ctx_cell.font = Font(size=9, color="555555")
                ctx_cell.fill = card_header_fill
                ctx_cell.alignment = Alignment(horizontal="center", vertical="center")
                
            # Apply borders to the 3x2 KPI Card block
            for r in range(current_row, current_row + 3):
                for c in range(col, col + 2):
                    ws.cell(row=r, column=c).border = thin_border
                    
            col += 2
            
        current_row += 4

    # Adjust column widths for clean rendering
    ws.column_dimensions["A"].width = 2
    for i in range(2, 8):
        ws.column_dimensions[get_column_letter(i)].width = 18
