from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import coordinate_to_tuple, get_column_letter

def render(ws, anchor: str, *, theme: str = "corporate_blue", metric_name: str = "YTD IT Expenditure", actual_val: float = 555700000, plan_val: float = 611300000, **kwargs) -> None:
    # Theme palette fallback
    palettes = {
        "corporate_blue": {
            "bg_dark": "1E3A8A", # Deep blue background
            "text_main": "FFFFFF", # White text
            "text_muted": "93C5FD", # Light blue text
            "favorable": "22C55E", # Green
            "adverse": "EF4444" # Red
        }
    }
    palette = palettes.get(theme, palettes["corporate_blue"])

    row, col = coordinate_to_tuple(anchor)
    
    # 1. Write backing data below the visual KPI strip
    data_row = row + 10
    actual_cell = f"{get_column_letter(col)}{data_row}"
    plan_cell = f"{get_column_letter(col+1)}{data_row}"
    ws[actual_cell] = actual_val
    ws[plan_cell] = plan_val
    
    # Hide the data row to keep the dashboard clean
    ws.row_dimensions[data_row].hidden = True

    # 2. Layout the visual KPI Strip
    title_cell = f"{get_column_letter(col)}{row}"
    var_title_cell = f"{get_column_letter(col+1)}{row}"
    
    actual_disp_cell = f"{get_column_letter(col)}{row+1}"
    var_disp_cell = f"{get_column_letter(col+1)}{row+1}"
    var_pct_cell = f"{get_column_letter(col+2)}{row+1}"
    
    ws[title_cell] = metric_name
    ws[var_title_cell] = "Plan v Actual Var."
    
    # Formulas linking to backing data
    ws[actual_disp_cell] = f"={actual_cell}"
    ws[var_disp_cell] = f"={plan_cell}-{actual_cell}"
    
    # Dynamic text formula: Unicode arrow + absolute percentage
    # In cost contexts: Actual < Plan is Favorable (Positive variance = ▼ down arrow indicating cost drop)
    ws[var_pct_cell] = f'=IF({var_disp_cell}<0, "▲ ", "▼ ") & TEXT(ABS({var_disp_cell}/{plan_cell}), "0.0%")'
    
    # 3. Styling and Formatting
    bg_fill = PatternFill("solid", fgColor=palette["bg_dark"])
    main_font = Font(color=palette["text_main"], bold=True, size=18)
    muted_font = Font(color=palette["text_muted"], bold=True, size=11)
    var_font = Font(color=palette["text_main"], bold=True, size=14)
    
    # Apply baseline styling to the 3x2 grid
    for r in range(row, row+2):
        for c in range(col, col+3):
            cell = ws.cell(row=r, column=c)
            cell.fill = bg_fill
            cell.alignment = Alignment(vertical="center", horizontal="left")
            
    ws[title_cell].font = muted_font
    ws[var_title_cell].font = muted_font
    
    ws[actual_disp_cell].font = main_font
    ws[actual_disp_cell].number_format = '$#,##0.0,, "M"'
    
    ws[var_disp_cell].font = var_font
    ws[var_disp_cell].number_format = '$#,##0.0,, "M"'
    
    ws[var_pct_cell].font = var_font
    
    # Adjust column widths for visual breathing room
    ws.column_dimensions[get_column_letter(col)].width = 22
    ws.column_dimensions[get_column_letter(col+1)].width = 18
    ws.column_dimensions[get_column_letter(col+2)].width = 12

    # 4. Conditional Formatting for the Arrow/Pct cell
    green_font = Font(color=palette["favorable"], bold=True, size=14)
    red_font = Font(color=palette["adverse"], bold=True, size=14)
    
    var_disp_abs = f"${get_column_letter(col+1)}${row+1}"
    
    # Rule 1: Variance < 0 -> Overspend (Adverse) -> Red Font
    ws.conditional_formatting.add(
        var_pct_cell,
        FormulaRule(formula=[f"{var_disp_abs}<0"], font=red_font)
    )
    
    # Rule 2: Variance >= 0 -> Underspend (Favorable) -> Green Font
    ws.conditional_formatting.add(
        var_pct_cell,
        FormulaRule(formula=[f"{var_disp_abs}>=0"], font=green_font)
    )
