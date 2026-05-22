def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo

    # 1. Theme Extraction
    primary_color = theme.get("primary", "003366") if isinstance(theme, dict) else "003366"
    text_color = theme.get("text", "FFFFFF") if isinstance(theme, dict) else "FFFFFF"

    # 2. Setup Dashboard Sheet (Presentation Layer)
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False

    # Dashboard Header
    ws_dash.merge_cells("A1:P2")
    header_cell = ws_dash["A1"]
    header_cell.value = title
    header_cell.font = Font(size=24, bold=True, color=text_color)
    header_cell.fill = PatternFill("solid", fgColor=primary_color)
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Setup Calc Sheet (Calculation Layer - Hidden)
    ws_calc = wb.create_sheet("Calc")
    ws_calc.sheet_state = "hidden"

    # Data for Chart 1: Stacked Bar (Profit by Market & Cookie)
    calc_data_1 = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Oatmeal Raisin"],
        ["India", 62349, 4872, 21028],
        ["Philippines", 54618, 7026, 22005],
        ["United Kingdom", 46530, 5220, 11497],
        ["United States", 36657, 6368, 22260]
    ]
    for row in calc_data_1:
        ws_calc.append(row)
    
    ws_calc.append([]) # Spacer at row 6

    # Data for Chart 2: Line (Units Sold)
    calc_data_2 = [
        ["Month", "Units Sold"],
        ["Sep", 50601],
        ["Oct", 95622],
        ["Nov", 65481],
        ["Dec", 52970]
    ]
    for row in calc_data_2:
        ws_calc.append(row)
        
    ws_calc.append([]) # Spacer at row 12

    # Data for Chart 3: Line (Profit)
    calc_data_3 = [
        ["Month", "Profit"],
        ["Sep", 124812],
        ["Oct", 228275],
        ["Nov", 160228],
        ["Dec", 136337]
    ]
    for row in calc_data_3:
        ws_calc.append(row)

    # 4. Build and Position Charts
    # Chart 1: Stacked Column
    c1 = BarChart()
    c1.type = "col"
    c1.style = 10
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "Profit by Market & Cookie Type"
    c1.height = 14
    c1.width = 15
    
    data1 = Reference(ws_calc, min_col=2, min_row=1, max_col=4, max_row=5)
    cats1 = Reference(ws_calc, min_col=1, min_row=2, max_row=5)
    c1.add_data(data1, titles_from_data=True)
    c1.set_categories(cats1)
    ws_dash.add_chart(c1, "B4")

    # Chart 2: Line (Units Sold)
    c2 = LineChart()
    c2.title = "Units sold each month"
    c2.style = 13
    c2.height = 7
    c2.width = 13
    
    data2 = Reference(ws_calc, min_col=2, min_row=7, max_col=2, max_row=11)
    cats2 = Reference(ws_calc, min_col=1, min_row=8, max_row=11)
    c2.add_data(data2, titles_from_data=True)
    c2.set_categories(cats2)
    c2.legend = None  # Hide legend for clean look
    ws_dash.add_chart(c2, "J4")

    # Chart 3: Line (Profit)
    c3 = LineChart()
    c3.title = "Profit by month"
    c3.style = 13
    c3.height = 7
    c3.width = 13
    
    data3 = Reference(ws_calc, min_col=2, min_row=13, max_col=2, max_row=17)
    cats3 = Reference(ws_calc, min_col=1, min_row=14, max_row=17)
    c3.add_data(data3, titles_from_data=True)
    c3.set_categories(cats3)
    c3.legend = None  # Hide legend for clean look
    ws_dash.add_chart(c3, "J15")

    # 5. Setup Raw Data Sheet (Data Layer)
    ws_data = wb.create_sheet("Data")
    headers = ["Country", "Product", "Units Sold", "Revenue", "Cost", "Profit", "Date"]
    ws_data.append(headers)
    
    sample_data = [
        ["India", "Chocolate Chip", 1725, 8625, 3450, 5175, "11/1/2019"],
        ["India", "Chocolate Chip", 2152, 10760, 4304, 6456, "12/1/2019"],
        ["Philippines", "Fortune Cookie", 1404, 7020, 2808, 4212, "11/1/2019"],
        ["United States", "Oatmeal Raisin", 2470, 12350, 4940, 7410, "12/1/2019"]
    ]
    for row in sample_data:
        ws_data.append(row)
        
    tab = Table(displayName="RawData", ref=f"A1:G5")
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws_data.add_table(tab)
