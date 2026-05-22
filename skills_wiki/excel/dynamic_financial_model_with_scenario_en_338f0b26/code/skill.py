from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "Financial Model", theme: str = "corporate_blue", **kwargs) -> None:
    ws = wb.create_sheet(sheet_name)
    
    # Styles
    # Assuming corporate_blue defaults; in a real engine these might be pulled from a theme dict
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    input_font = Font(color="0000FF") # Standard financial modeling color for hardcodes
    toggle_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid") # Yellow to draw attention
    
    top_border = Border(top=Side(style='thin'))
    double_bottom = Border(bottom=Side(style='double'), top=Side(style='thin'))
    
    # Title and Scenario Toggle
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["G2"] = "Active Scenario:"
    ws["G2"].alignment = Alignment(horizontal="right")
    ws["H2"] = 1
    ws["H2"].fill = toggle_fill
    ws["H2"].font = input_font
    ws["H2"].alignment = Alignment(horizontal="center")
    
    # Add Data Validation for Scenario Toggle (1=Base, 2=Best, 3=Worst)
    dv = DataValidation(type="list", formula1='"1,2,3"', allowBlank=False)
    dv.error ='Your entry is not in the list'
    dv.errorTitle = 'Invalid Entry'
    dv.prompt = 'Please select scenario 1, 2, or 3'
    dv.promptTitle = 'Scenario Select'
    ws.add_data_validation(dv)
    dv.add(ws["H2"])

    # Setup Columns and Headers
    ws.column_dimensions["A"].width = 25
    headers = ["Figures in USD", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        if col_idx > 1:
            ws.column_dimensions[get_column_letter(col_idx)].width = 12

    # --- Section 1: Income Statement ---
    pnl_labels = [
        "Revenue", "Cost of Goods Sold", "Gross Profit", 
        "Operating Expenses", "Operating Profit", 
        "Corporate Tax", "Net Income"
    ]
    for i, label in enumerate(pnl_labels, start=5):
        ws.cell(row=i, column=1, value=label)
        
    for c in range(2, 7):
        col = get_column_letter(c)
        # Revenue = Orders * AOV
        ws[f"{col}5"] = f"={col}20*{col}21"
        # COGS = Orders * Cost Per Order
        ws[f"{col}6"] = f"={col}20*{col}22"
        # Gross Profit
        ws[f"{col}7"] = f"={col}5-{col}6"
        ws[f"{col}7"].font = bold_font
        ws[f"{col}7"].border = top_border
        # Opex (Fixed)
        ws[f"{col}8"] = f"={col}23"
        # Operating Profit
        ws[f"{col}9"] = f"={col}7-{col}8"
        ws[f"{col}9"].font = bold_font
        ws[f"{col}9"].border = top_border
        # Tax
        ws[f"{col}10"] = f"=MAX(0, {col}9*{col}24)"
        # Net Income
        ws[f"{col}11"] = f"={col}9-{col}10"
        ws[f"{col}11"].font = bold_font
        ws[f"{col}11"].border = double_bottom

        # Number formatting for P&L
        for r in range(5, 12):
            ws[f"{col}{r}"].number_format = '#,##0'

    # --- Section 2: Live Assumptions (Driven by Engine) ---
    ws["A18"] = "Live Case Assumptions"
    ws["A18"].font = bold_font
    ws["A18"].border = double_bottom
    
    assump_labels = ["Number of Orders", "Average Order Value", "COGS per Order", "Operating Expenses", "Tax Rate"]
    for i, label in enumerate(assump_labels, start=20):
        ws.cell(row=i, column=1, value=label)

    for c in range(2, 7):
        col = get_column_letter(c)
        # CHOOSE Formula Engine: Links $H$2 to the respective rows in scenarios below
        ws[f"{col}20"] = f"=CHOOSE($H$2, {col}30, {col}38, {col}46)"
        ws[f"{col}21"] = f"=CHOOSE($H$2, {col}31, {col}39, {col}47)"
        ws[f"{col}22"] = f"=CHOOSE($H$2, {col}32, {col}40, {col}48)"
        ws[f"{col}23"] = f"=CHOOSE($H$2, {col}33, {col}41, {col}49)"
        ws[f"{col}24"] = f"=CHOOSE($H$2, {col}34, {col}42, {col}50)"
        
        ws[f"{col}20"].number_format = '#,##0'
        ws[f"{col}21"].number_format = '$#,##0.00'
        ws[f"{col}22"].number_format = '$#,##0.00'
        ws[f"{col}23"].number_format = '$#,##0'
        ws[f"{col}24"].number_format = '0.0%'

    # --- Section 3: Scenario Data Blocks (Hard-coded Inputs in Blue) ---
    def build_scenario_block(start_row, title, orders, aov, cogs, opex, tax):
        ws.cell(row=start_row, column=1, value=title).font = bold_font
        ws.cell(row=start_row, column=1).border = double_bottom
        for i, label in enumerate(assump_labels, start=start_row+2):
            ws.cell(row=i, column=1, value=label)
            
        for c_idx in range(2, 7):
            growth_mult = 1.0 + (0.2 * (c_idx-2)) # Simple dummy growth logic across years
            col = get_column_letter(c_idx)
            
            c_orders = ws.cell(row=start_row+2, column=c_idx, value=orders * growth_mult)
            c_aov = ws.cell(row=start_row+3, column=c_idx, value=aov)
            c_cogs = ws.cell(row=start_row+4, column=c_idx, value=cogs)
            c_opex = ws.cell(row=start_row+5, column=c_idx, value=opex)
            c_tax = ws.cell(row=start_row+6, column=c_idx, value=tax)
            
            # Format as inputs (Blue)
            for cell in [c_orders, c_aov, c_cogs, c_opex, c_tax]:
                cell.font = input_font
            
            c_orders.number_format = '#,##0'
            c_aov.number_format = '$#,##0.00'
            c_cogs.number_format = '$#,##0.00'
            c_opex.number_format = '$#,##0'
            c_tax.number_format = '0.0%'

    # Scenario 1: Base Case
    build_scenario_block(28, "Scenario 1: Base Case", 3000, 39.95, 8.75, 50000, 0.20)
    
    # Scenario 2: Best Case
    build_scenario_block(36, "Scenario 2: Best Case", 4000, 44.95, 8.00, 55000, 0.20)
    
    # Scenario 3: Worst Case
    build_scenario_block(44, "Scenario 3: Worst Case", 2000, 34.95, 9.50, 45000, 0.25)
