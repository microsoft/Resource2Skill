from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def _create_card(ws, min_col: int, min_row: int, max_col: int, max_row: int, title: str = "", title_color: str = "333333") -> None:
    """Helper to draw a white 'card' over a range of cells with a subtle border."""
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Side(border_style="thin", color="E5E7EB")
    
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = white_fill
            
            # Apply borders only to the outer edges of the card range
            t = thin_border if row == min_row else None
            b = thin_border if row == max_row else None
            l = thin_border if col == min_col else None
            r = thin_border if col == max_col else None
            
            if any([t, b, l, r]):
                cell.border = Border(top=t, bottom=b, left=l, right=r)

    # Inject and style the card title
    if title:
        title_cell = ws.cell(row=min_row, column=min_col)
        title_cell.value = title
        title_cell.font = Font(bold=True, size=12, color=title_color)
        title_cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.row_dimensions[min_row].height = 25

def render_sheet(wb, sheet_name: str, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]
        
    sidebar_links = kwargs.get("sidebar_links", ["Dashboard", "Inputs", "Contacts"])
    
    # Modern web-app palette (Tailwind-inspired defaults)
    bg_color = "F3F4F6"       # Gray 100 
    sidebar_color = "1F2937"  # Gray 800
    
    bg_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    sidebar_fill = PatternFill(start_color=sidebar_color, end_color=sidebar_color, fill_type="solid")
    
    # 1. Clean the canvas
    ws.sheet_view.showGridLines = False
    
    for row in range(1, 40):
        for col in range(1, 22):
            ws.cell(row=row, column=col).fill = bg_fill
            
    # 2. Build the Navigation Sidebar (Col A)
    ws.column_dimensions['A'].width = 18
    for row in range(1, 40):
        ws.cell(row=row, column=1).fill = sidebar_fill
        
    ws.column_dimensions['B'].width = 4  # Spacer
    
    # 3. Add Sidebar Hyperlinks
    for i, link_name in enumerate(sidebar_links):
        # Ensure target sheet exists so hyperlinks resolve natively
        if link_name not in wb.sheetnames:
            wb.create_sheet(link_name)
            
        cell = ws.cell(row=4 + (i * 3), column=1, value=link_name)
        cell.hyperlink = f"#'{link_name}'!A1"
        cell.font = Font(color="FFFFFF", underline="none", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # 4. Layout the Cards (Simulating shape objects)
    # Title Header Card
    _create_card(ws, 3, 2, 19, 4, title, title_color=sidebar_color)
    ws.cell(row=2, column=3).font = Font(bold=True, size=16, color=sidebar_color) # Override size for main title
    
    # Top KPI Row (3 Cards)
    _create_card(ws, min_col=3,  min_row=6, max_col=7,  max_row=11, title="Sales")
    _create_card(ws, min_col=9,  min_row=6, max_col=13, max_row=11, title="Profit")
    _create_card(ws, min_col=15, min_row=6, max_col=19, max_row=11, title="# of Customers")
    
    # Main Chart Row (1 Wide, 1 Narrow)
    _create_card(ws, min_col=3,  min_row=13, max_col=13, max_row=26, title="Sales Trend")
    _create_card(ws, min_col=15, min_row=13, max_col=19, max_row=26, title="Customer Satisfaction")
