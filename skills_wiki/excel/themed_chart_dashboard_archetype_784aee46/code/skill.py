def render_workbook(wb, *, title: str = "Sales Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.chart import LineChart, BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    # Try standard palette injection, fallback to corporate defaults
    try:
        from _helpers import get_theme_palette
        palette = get_theme_palette(theme)
    except ImportError:
        palette = {
            "header_bg": "002060", 
            "header_fg": "FFFFFF", 
            "bg": "F2F2F2", 
            "border": "CCCCCC"
        }

    # 1. Setup Backend Data Sheet
    ws_data = wb.active
    ws_data.title = "ReportData"
    # Hide the data sheet to focus user attention on the dashboard
    ws_data.sheet_state = 'hidden'

    # Populate Data for Chart 1: Monthly Trend
    ws_data.append(["Month", "Revenue"])
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    rev = [15000, 18000, 22000, 19000, 25000, 28000, 24000, 30000, 32000, 35000, 40000, 45000]
    for m, r in zip(months, rev):
        ws_data.append([m, r])

    # Populate Data for Chart 2: Units by Category
    ws_data.append([])
    ws_data.append(["Year", "Hoodies", "T-Shirts"])
    ws_data.append(["2023", 7500, 15000])
    ws_data.append(["2024", 9000, 17500])

    # Populate Data for Chart 3: Top States
    ws_data.append([])
    ws_data.append(["State", "Profit"])
    states = ["California", "Texas", "New York", "Florida", "Illinois"]
    profits = [38000, 34000, 32000, 31000, 29000]
    for s, p in zip(states, profits):
        ws_data.append([s, p])

    # 2. Setup Dashboard Canvas
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False

    # Set background fill for the dashboard canvas viewport
    bg_fill = PatternFill(start_color=palette.get("bg", "F2F2F2"), fill_type="solid")
    for row in ws_dash.iter_rows(min_row=1, max_row=40, min_col=1, max_col=20):
        for cell in row:
            cell.fill = bg_fill

    # Title Banner (Spans B2:P4)
    ws_dash.merge_cells("B2:P4")
    title_cell = ws_dash["B2"]
    title_cell.value = title
    title_cell.font = Font(size=24, bold=True, color=palette.get("header_fg", "FFFFFF"))
    title_cell.fill = PatternFill(start_color=palette.get("header_bg", "002060"), fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Create Charts
    # Chart 1: Line Chart
    c1 = LineChart()
    c1.title = "Monthly Revenue Trend"
    c1.style = 13
    c1.y_axis.title = "Revenue (USD)"
    data1 = Reference(ws_data, min_col=2, min_row=1, max_row=13)
    cats1 = Reference(ws_data, min_col=1, min_row=2, max_row=13)
    c1.add_data(data1, titles_from_data=True)
    c1.set_categories(cats1)
    c1.width = 16
    c1.height = 8

    # Chart 2: Clustered Column Chart
    c2 = BarChart()
    c2.type = "col"
    c2.style = 10
    c2.title = "Units Sold by Category"
    data2 = Reference(ws_data, min_col=2, max_col=3, min_row=16, max_row=18)
    cats2 = Reference(ws_data, min_col=1, min_row=17, max_row=18)
    c2.add_data(data2, titles_from_data=True)
    c2.set_categories(cats2)
    c2.width = 14
    c2.height = 8

    # Chart 3: Horizontal Bar Chart
    c3 = BarChart()
    c3.type = "bar"
    c3.style = 11
    c3.title = "Top 5 States by Profit"
    data3 = Reference(ws_data, min_col=2, min_row=21, max_row=26)
    cats3 = Reference(ws_data, min_col=1, min_row=22, max_row=26)
    c3.add_data(data3, titles_from_data=True)
    c3.set_categories(cats3)
    c3.width = 14
    c3.height = 8

    # 4. Place Charts onto the Canvas Grid
    ws_dash.add_chart(c1, "B6")
    ws_dash.add_chart(c2, "J6")
    ws_dash.add_chart(c3, "F22")

    # 5. Mock Visual Slicer/Filter Panel (For visual completeness)
    ws_dash.merge_cells("B22:D22")
    filter_hdr = ws_dash["B22"]
    filter_hdr.value = "Active Filters"
    filter_hdr.font = Font(bold=True, color=palette.get("header_fg", "FFFFFF"))
    filter_hdr.fill = PatternFill(start_color=palette.get("header_bg", "002060"), fill_type="solid")
    filter_hdr.alignment = Alignment(horizontal="center")
    
    filters = ["Year: All", "Category: All", "State: All"]
    border_side = Side(border_style="thin", color=palette.get("border", "CCCCCC"))
    box_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    for i, f in enumerate(filters):
        row_idx = 23 + i
        ws_dash.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        cell = ws_dash.cell(row=row_idx, column=2)
        cell.value = f
        cell.alignment = Alignment(horizontal="center")
        
        # Apply borders to the simulated slicer items
        for col_idx in range(2, 5):
            ws_dash.cell(row=row_idx, column=col_idx).border = box_border
