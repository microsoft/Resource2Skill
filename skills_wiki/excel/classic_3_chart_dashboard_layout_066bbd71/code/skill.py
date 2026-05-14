from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

def render_workbook(wb, *, title: str = "Performance Dashboard", theme: str = "corporate_blue", **kwargs) -> None:
    """
    Constructs a complete 3-chart dashboard with a left-hand control panel area.
    Simulates the clean layout and structure of a PivotTable-driven interactive dashboard.
    """
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    
    # 1. Dashboard Page Setup
    ws_dash.sheet_view.showGridLines = False
    
    # Setup Header
    ws_dash.merge_cells("A1:U3")
    header = ws_dash["A1"]
    header.value = title
    header.font = Font(size=24, bold=True, color="FFFFFF")
    header.alignment = Alignment(horizontal="center", vertical="center")
    
    # Theme color fallback
    bg_color = "2F5597" if theme == "corporate_blue" else "4472C4"
    header.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    # 2. Hidden Data Sheet for Charts
    ws_data = wb.create_sheet("Data")
    ws_data.sheet_state = 'hidden'
    
    # Bar Chart Data (Profit by Market & Product)
    bar_data = [
        ["Market", "Chocolate Chip", "Fortune Cookie", "Sugar", "Snickerdoodle"],
        ["India", 62349, 4872, 18560, 25085],
        ["Philippines", 54618, 7026, 8313, 22005],
        ["United Kingdom", 46530, 5220, 14620, 11497],
        ["Malaysia", 46587, 5537, 20555, 17536],
        ["United States", 36657, 6368, 9937, 22260]
    ]
    for row in bar_data:
        ws_data.append(row)
        
    # Line Chart Data (Trend over Months)
    line_data = [
        ["Month", "Units Sold", "Profit"],
        ["Sep", 50601, 124812],
        ["Oct", 95622, 228275],
        ["Nov", 65481, 160228],
        ["Dec", 52970, 136337]
    ]
    
    # Leave a gap in the data sheet
    ws_data.append([])
    start_line_row = ws_data.max_row + 1
    for row in line_data:
        ws_data.append(row)

    # 3. Main Stacked Bar Chart (Left Area)
    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.grouping = "stacked"
    bar_chart.overlap = 100
    bar_chart.title = "Profit by Market & Product"
    bar_chart.y_axis.title = "Profit ($)"
    
    cats = Reference(ws_data, min_col=1, min_row=2, max_row=6)
    data = Reference(ws_data, min_col=2, min_row=1, max_col=5, max_row=6)
    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    
    # Size and position
    bar_chart.width = 18
    bar_chart.height = 11.5
    ws_dash.add_chart(bar_chart, "D5")

    # 4. Secondary Line Charts (Right Area)
    # Line 1: Units Sold Trend
    line_units = LineChart()
    line_units.title = "Units sold each month"
    line_units.legend = None
    
    l_cats = Reference(ws_data, min_col=1, min_row=start_line_row+1, max_row=start_line_row+4)
    l_data_u = Reference(ws_data, min_col=2, min_row=start_line_row, max_row=start_line_row+4)
    line_units.add_data(l_data_u, titles_from_data=True)
    line_units.set_categories(l_cats)
    
    line_units.width = 12
    line_units.height = 5.5
    ws_dash.add_chart(line_units, "N5")

    # Line 2: Profit Trend
    line_profit = LineChart()
    line_profit.title = "Profit by month"
    line_profit.legend = None
    
    l_data_p = Reference(ws_data, min_col=3, min_row=start_line_row, max_row=start_line_row+4)
    line_profit.add_data(l_data_p, titles_from_data=True)
    line_profit.set_categories(l_cats)
    
    line_profit.width = 12
    line_profit.height = 5.5
    ws_dash.add_chart(line_profit, "N15")

    # 5. Control Panel / Slicer Mockup (Far Left Column)
    slicer_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    hdr_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    
    controls = [
        ("Timeframe", ["Q1", "Q2", "Q3", "Q4"]),
        ("Market", ["India", "Philippines", "UK", "US"]),
        ("Product", ["Chocolate Chip", "Fortune Cookie", "Sugar"])
    ]
    
    row_idx = 5
    for c_title, items in controls:
        # Pseudo-Slicer Header
        hdr_cell = ws_dash.cell(row=row_idx, column=2, value=c_title)
        hdr_cell.font = Font(bold=True)
        hdr_cell.fill = hdr_fill
        
        # Pseudo-Slicer Options
        for i, item in enumerate(items, 1):
            cell = ws_dash.cell(row=row_idx+i, column=2, value=item)
            cell.fill = slicer_fill
            
        row_idx += len(items) + 2
        
    # Set spacing for the control panel column
    ws_dash.column_dimensions['B'].width = 18
    ws_dash.column_dimensions['C'].width = 2
