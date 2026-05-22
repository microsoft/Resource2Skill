def render_sheet(wb, sheet_name: str, *, title: str = "Dynamic Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    # Theme abstraction fallback
    themes = {
        "corporate_blue": {"header_bg": "D9E1F2", "header_fg": "000000", "input_fg": "0000FF", "highlight_bg": "FFF2CC"},
        "executive_grey": {"header_bg": "D9D9D9", "header_fg": "000000", "input_fg": "000080", "highlight_bg": "E2EFDA"}
    }
    palette = themes.get(theme, themes["corporate_blue"])

    ws = wb.create_sheet(sheet_name)
    
    bold = Font(bold=True)
    input_font = Font(color=palette["input_fg"])
    bg_header = PatternFill(start_color=palette["header_bg"], end_color=palette["header_bg"], fill_type="solid")
    border_subtotal = Border(top=Side(style='thin'), bottom=Side(style='double'))
    
    # 1. Title and Scenario Toggle
    ws['A1'] = title
    ws['A1'].font = Font(size=14, bold=True)
    
    ws['I2'] = "Active Scenario:"
    ws['I2'].font = bold
    ws['I2'].alignment = Alignment(horizontal="right")
    
    # Scenario Dropdown Setup
    ws['J2'] = 1
    ws['J2'].font = Font(bold=True, color=palette["input_fg"])
    ws['J2'].fill = PatternFill(start_color=palette["highlight_bg"], end_color=palette["highlight_bg"], fill_type="solid")
    ws['J2'].border = Border(outline=True, top=Side(style="thin"), bottom=Side(style="thin"), left=Side(style="thin"), right=Side(style="thin"))
    ws['J2'].alignment = Alignment(horizontal="center")
    
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws['J2'])
    ws['K2'] = "<- 1 = Upper Case, 2 = Lower Case"
    ws['K2'].font = Font(italic=True, color="595959")
    
    # 2. Section Headers
    headers = ["Figures in USD", "Unit", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.font = bold
        c.fill = bg_header

    for r_idx, label in [(12, "Assumptions: Live Case"), (20, "Assumptions: Scenario 1 (Upper Case)"), (28, "Assumptions: Scenario 2 (Lower Case)")]:
        c = ws.cell(row=r_idx, column=1, value=label)
        c.font = bold
        c.fill = bg_header
        for col_idx in range(2, 8):
            ws.cell(row=r_idx, column=col_idx).fill = bg_header
    
    # 3. Setup Row Labels
    labels = ["Number of Orders", "Order Growth Rate", "Average Order Value", "COGS per order", "Operating Expenses"]
    units = ["#", "%", "$", "$", "$"]
    
    for start_row in [13, 21, 29]:
        for i, (lbl, unit) in enumerate(zip(labels, units)):
            ws.cell(row=start_row + i, column=1, value=lbl)
            ws.cell(row=start_row + i, column=2, value=unit).alignment = Alignment(horizontal="center")
            
    is_labels = ["Revenue", "Cost of Goods Sold", "Gross Profit", "Operating Expenses", "Operating Profit", "Margin %"]
    is_units = ["$", "$", "$", "$", "$", "%"]
    for i, (lbl, unit) in enumerate(zip(is_labels, is_units)):
        ws.cell(row=5 + i, column=1, value=lbl)
        ws.cell(row=5 + i, column=2, value=unit).alignment = Alignment(horizontal="center")

    # 4. Populate Scenarios (Upper = 21, Lower = 29)
    # Upper Case inputs
    ws['C21'] = 3000; ws['C21'].font = input_font
    for i, gr in enumerate([1.0, 0.75, 0.50, 0.35], 4):
        c = ws.cell(row=22, column=i, value=gr); c.font = input_font
    for i in range(3, 8): # Cols C to G
        if i > 3:
            ws.cell(row=21, column=i, value=f"={get_column_letter(i-1)}21*(1+{get_column_letter(i)}22)")
        ws.cell(row=23, column=i, value=39.95).font = input_font
        ws.cell(row=24, column=i, value=8.75).font = input_font
        ws.cell(row=25, column=i, value=20000 + (i-3)*5000).font = input_font
        
    # Lower Case inputs
    ws['C29'] = 2000; ws['C29'].font = input_font
    for i, gr in enumerate([0.50, 0.25, 0.15, 0.10], 4):
        c = ws.cell(row=30, column=i, value=gr); c.font = input_font
    for i in range(3, 8):
        if i > 3:
            ws.cell(row=29, column=i, value=f"={get_column_letter(i-1)}29*(1+{get_column_letter(i)}30)")
        ws.cell(row=31, column=i, value=34.95).font = input_font
        ws.cell(row=32, column=i, value=8.00).font = input_font
        ws.cell(row=33, column=i, value=20000).font = input_font

    # 5. Connect the Live Case using CHOOSE
    for row_offset in range(5):
        live_r = 13 + row_offset
        s1_r = 21 + row_offset
        s2_r = 29 + row_offset
        
        start_col = 4 if row_offset == 1 else 3 # Growth rate starts in Year 2
        for col_idx in range(start_col, 8):
            col_let = get_column_letter(col_idx)
            ws.cell(row=live_r, column=col_idx, value=f"=CHOOSE($J$2, {col_let}{s1_r}, {col_let}{s2_r})")
            
    # 6. Build the Income Statement
    for col_idx in range(3, 8):
        col = get_column_letter(col_idx)
        ws[f'{col}5'] = f"={col}13*{col}15" # Revenue
        ws[f'{col}6'] = f"={col}13*{col}16" # COGS
        ws[f'{col}7'] = f"={col}5-{col}6"   # Gross Profit
        ws[f'{col}8'] = f"={col}17"         # Opex
        ws[f'{col}9'] = f"={col}7-{col}8"   # Op Profit
        ws[f'{col}10'] = f"=IF({col}5>0, {col}9/{col}5, 0)" # Margin

    # 7. Formatting & Column Widths
    for r in [7, 9]: # Subtotal styling
        for col_idx in range(1, 8):
            c = ws.cell(row=r, column=col_idx)
            c.font = bold
            if col_idx > 2:
                c.border = border_subtotal

    for row_offset in range(5):
        is_pct = (row_offset == 1)
        for block_start in [13, 21, 29]:
            r = block_start + row_offset
            for col_idx in range(3, 8):
                if is_pct and col_idx == 3: continue
                # Specific formats: float for AOV/COGS, int for orders/opex
                fmt = "0%" if is_pct else "#,##0.00" if row_offset in [2, 3] else "#,##0"
                ws.cell(row=r, column=col_idx).number_format = fmt

    for r in range(5, 10):
        for col_idx in range(3, 8):
            ws.cell(row=r, column=col_idx).number_format = "#,##0"
    for col_idx in range(3, 8):
        ws.cell(row=10, column=col_idx).number_format = "0%"

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 8
    for i in range(3, 9):
        ws.column_dimensions[get_column_letter(i)].width = 14
        
    # Group scenario rows to allow collapsing
    ws.row_dimensions.group(20, 34, hidden=False)
