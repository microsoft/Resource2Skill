from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def render_sheet(wb, sheet_name: str, *, title: str = "SALES & PROFIT DASHBOARD", kpis: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a structured, grid-based dashboard shell layout using cells as containers.
    """
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    if kpis is None:
        kpis = [
            {"label": "TOTAL SALES", "value": "$19,288,888"},
            {"label": "TOTAL PROFIT", "value": "$2,477,961"},
            {"label": "CUSTOMER COUNT", "value": "2,319"}
        ]

    # Modern Light-Mode Dashboard Palette
    canvas_fill = PatternFill("solid", fgColor="F3F6F9")
    card_fill = PatternFill("solid", fgColor="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A365D")
    
    font_header = Font(name="Segoe UI", size=20, bold=True, color="FFFFFF")
    font_kpi_lbl = Font(name="Segoe UI", size=10, bold=True, color="718096")
    font_kpi_val = Font(name="Segoe UI", size=18, bold=True, color="2D3748")
    
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # 1. Paint Canvas Area
    for row in ws.iter_rows(min_row=1, max_row=35, min_col=1, max_col=13):
        for cell in row:
            cell.fill = canvas_fill

    # 2. Configure Underlying Grid Columns
    # Alternating layout: 1 spacer (width 2), 3 data cols (width 10 each)
    col_widths = [2, 10, 10, 10, 2, 10, 10, 10, 2, 10, 10, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 3. Header Bar
    ws.merge_cells("B2:L3")
    header_cell = ws["B2"]
    header_cell.value = title
    header_cell.font = font_header
    header_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    for r in range(2, 4):
        for c in range(2, 13):
            ws.cell(row=r, column=c).fill = header_fill

    # Helper function to construct structured bordered cards
    def create_card(start_row: int, start_col: int, end_row: int, end_col: int):
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = card_fill
                # Apply boundary borders to individual cells to create a unified block border
                left = thin_border.left if c == start_col else None
                right = thin_border.right if c == end_col else None
                top = thin_border.top if r == start_row else None
                bottom = thin_border.bottom if r == end_row else None
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)
        return ws.cell(row=start_row, column=start_col)

    # 4. Inject KPI Cards
    # Anchoring to our 3-column data blocks (Cols B-D, F-H, J-L)
    kpi_anchors = [
        (2, 4),   # Left Block
        (6, 8),   # Center Block
        (10, 12)  # Right Block
    ]
    
    kpi_row = 5
    for idx, kpi in enumerate(kpis[:3]):
        start_c, end_c = kpi_anchors[idx]
        
        # Upper half of card for the muted label
        lbl_cell = create_card(kpi_row, start_c, kpi_row, end_c)
        lbl_cell.value = kpi["label"]
        lbl_cell.font = font_kpi_lbl
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Lower half for the prominent metric
        val_cell = create_card(kpi_row + 1, start_c, kpi_row + 2, end_c)
        val_cell.value = kpi["value"]
        val_cell.font = font_kpi_val
        val_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 5. Chart Placeholder Zones
    # Main wide chart zone (Spans Left & Center blocks)
    main_chart = create_card(9, 2, 24, 8)
    main_chart.value = "[ Main Primary Chart Area ]"
    main_chart.font = font_kpi_lbl
    main_chart.alignment = Alignment(horizontal="center", vertical="center")
    
    # Top right zone (Spans Right block only)
    tr_chart = create_card(9, 10, 16, 12)
    tr_chart.value = "[ Top Right Chart ]"
    tr_chart.font = font_kpi_lbl
    tr_chart.alignment = Alignment(horizontal="center", vertical="center")
    
    # Bottom right zone (Spans Right block only)
    br_chart = create_card(17, 10, 24, 12)
    br_chart.value = "[ Bottom Right Chart ]"
    br_chart.font = font_kpi_lbl
    br_chart.alignment = Alignment(horizontal="center", vertical="center")
