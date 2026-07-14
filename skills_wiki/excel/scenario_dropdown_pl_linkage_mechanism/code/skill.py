from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill


def build_scenario_dropdown_pl(path="scenario_dropdown_pl.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"

    ws.append(["Driver", "Base", "Bull", "Bear"])
    rows = [
        ["Revenue", 24000000, 32000000, 19000000],
        ["Gross Margin", 0.58, 0.63, 0.49],
        ["Opex", 11800000, 14200000, 10200000],
    ]
    for row in rows:
        ws.append(row)

    ws["F1"] = "Scenario"
    ws["G1"] = 1
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["G1"])
    ws["F2"] = "1=Base, 2=Bull, 3=Bear"

    outputs = ["Revenue", "COGS", "Opex", "EBITDA"]
    for i, label in enumerate(outputs, 4):
        ws.cell(i, 6, label)
    ws["G4"] = "=CHOOSE($G$1,$B$2,$C$2,$D$2)"
    ws["G5"] = "=G4*(1-CHOOSE($G$1,$B$3,$C$3,$D$3))"
    ws["G6"] = "=CHOOSE($G$1,$B$4,$C$4,$D$4)"
    ws["G7"] = "=G4-G5-G6"

    for cell in ("F1", "G1"):
        ws[cell].font = Font(bold=True)
        ws[cell].fill = PatternFill("solid", fgColor="D9EAF7")
    wb.save(path)
    return path
