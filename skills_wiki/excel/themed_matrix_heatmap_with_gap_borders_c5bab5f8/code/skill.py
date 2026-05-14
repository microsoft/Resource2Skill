from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils import get_column_letter

def render(ws, anchor: str, *, title: str = "Purchase Patterns by Channel", row_labels: list = None, col_labels: list = None, data_matrix: list = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a heatmap matrix with conditional formatting and thick "gap" borders.
    """
    # Default realistic data if none provided
    if row_labels is None:
        row_labels = ["App", "Instagram", "Partner App", "Target.com", "Website"]
    if col_labels is None:
        col_labels = ["Female", "Male", "Other", "Unknown"]
    if data_matrix is None:
        data_matrix = [
            [0.190, 0.123, 0.011, 0.031],
            [0.050, 0.043, 0.002, 0.018],
            [0.063, 0.046, 0.002, 0.004],
            [0.096, 0.068, 0.005, 0.013],
            [0.125, 0.091, 0.006, 0.012]
        ]

    start_row, start_col = coordinate_to_tuple(anchor)
    end_col = start_col + len(col_labels)

    # 1. Write Component Title
    title_cell = ws.cell(row=start_row, column=start_col)
    title_cell.value = title
    title_cell.font = Font(bold=True, size=12, color="333333")
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)

    # 2. Write Column Labels
    header_font = Font(bold=True, color="555555")
    header_align = Alignment(horizontal="center", vertical="center")
    for j, c_label in enumerate(col_labels):
        c_cell = ws.cell(row=start_row + 1, column=start_col + 1 + j)
        c_cell.value = c_label
        c_cell.font = header_font
        c_cell.alignment = header_align

    # 3. Setup Gap Borders and Write Row Labels / Data Matrix
    # The thick white border creates the "gap" effect between colored heatmap cells
    gap_border = Border(
        left=Side(style='thick', color='FFFFFF'),
        right=Side(style='thick', color='FFFFFF'),
        top=Side(style='thick', color='FFFFFF'),
        bottom=Side(style='thick', color='FFFFFF')
    )
    data_font = Font(color="000000")
    
    for i, r_label in enumerate(row_labels):
        # Write Row Label
        r_cell = ws.cell(row=start_row + 2 + i, column=start_col)
        r_cell.value = r_label
        r_cell.font = header_font
        r_cell.alignment = Alignment(horizontal="right", vertical="center")

        # Write Matrix Data Row
        for j, val in enumerate(data_matrix[i]):
            d_cell = ws.cell(row=start_row + 2 + i, column=start_col + 1 + j)
            d_cell.value = val
            d_cell.number_format = "0.0%"
            d_cell.font = data_font
            d_cell.alignment = header_align
            d_cell.border = gap_border

    # 4. Apply Conditional Formatting (Color Scale)
    first_data_coord = ws.cell(row=start_row + 2, column=start_col + 1).coordinate
    last_data_coord = ws.cell(row=start_row + 1 + len(row_labels), column=end_col).coordinate
    matrix_range = f"{first_data_coord}:{last_data_coord}"

    # In a full framework, this pulls from theme tokens. Fallback to a clean Blue tint.
    theme_accent = "1F4E78" if theme == "corporate_blue" else "D92B30"
    base_tint = "F2F6FA" if theme == "corporate_blue" else "FDF2F2"

    scale_rule = ColorScaleRule(
        start_type='min', start_color=base_tint,
        end_type='max', end_color=theme_accent
    )
    ws.conditional_formatting.add(matrix_range, scale_rule)

    # 5. Clean up column widths for uniform square-like appearance
    ws.column_dimensions[get_column_letter(start_col)].width = 15  # Label col
    for col_idx in range(start_col + 1, end_col + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 10
