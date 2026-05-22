from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def render(ws, anchor: str, *, title: str, value: str, percent: str = None, theme: str = "corporate_blue", **kwargs) -> None:
    """
    Renders a grid-aligned KPI card that mimics a floating shape widget.
    
    :param ws: The worksheet to render on.
    :param anchor: Top-left cell coordinate (e.g., "B2").
    :param title: The title of the KPI (e.g., "Revenue").
    :param value: The main value or formula (e.g., "$7,708,632" or "=Data!B2").
    :param percent: Optional secondary value/formula for the badge (e.g., "5%").
    """
    col_str, row_str = coordinate_from_string(anchor)
    c_idx = column_index_from_string(col_str)
    r_idx = int(row_str)
    
    # 1. Resolve Theme Colors (Fallback to corporate blue/navy scheme)
    # In a full framework, these would be loaded via the theme configuration
    bg_color = "002060"     # Primary dark navy
    text_color = "FFFFFF"   # White text
    badge_bg = "004080"     # Slightly lighter blue for the badge
    
    fill_main = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    fill_badge = PatternFill(start_color=badge_bg, end_color=badge_bg, fill_type="solid")
    
    font_title = Font(name="Calibri", size=12, color=text_color, bold=False)
    font_value = Font(name="Calibri", size=18, color=text_color, bold=True)
    font_percent = Font(name="Calibri", size=12, color=text_color, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Border styles
    thin_white = Side(border_style="thin", color="FFFFFF")
    thick_navy = Side(border_style="medium", color=bg_color)
    
    # 2. Layout the KPI Card (2 rows x 2 columns)
    
    # Row 1: Title (Merged across both columns)
    ws.merge_cells(start_row=r_idx, start_column=c_idx, end_row=r_idx, end_column=c_idx+1)
    title_cell = ws.cell(row=r_idx, column=c_idx)
    title_cell.value = title
    title_cell.fill = fill_main
    title_cell.font = font_title
    title_cell.alignment = align_center
    
    # Apply fill to the adjacent merged cell so borders/fills render correctly
    ws.cell(row=r_idx, column=c_idx+1).fill = fill_main
    
    # Row 2: Value and Percent
    value_cell = ws.cell(row=r_idx+1, column=c_idx)
    value_cell.value = value
    value_cell.fill = fill_main
    value_cell.font = font_value
    value_cell.alignment = align_center
    
    percent_cell = ws.cell(row=r_idx+1, column=c_idx+1)
    if percent is not None:
        percent_cell.value = percent
        percent_cell.fill = fill_badge
        percent_cell.font = font_percent
        percent_cell.alignment = align_center
        # White border creates the illusion of a floating, distinct badge over the card
        percent_cell.border = Border(top=thin_white, left=thin_white, right=thin_white, bottom=thin_white)
    else:
        percent_cell.fill = fill_main
        
    # 3. Size the Grid Cells
    ws.row_dimensions[r_idx].height = 20
    ws.row_dimensions[r_idx+1].height = 35
    ws.column_dimensions[get_column_letter(c_idx)].width = 18
    ws.column_dimensions[get_column_letter(c_idx+1)].width = 10
    
    # 4. Apply Outer Card Border
    for r in range(r_idx, r_idx+2):
        for c in range(c_idx, c_idx+2):
            cell = ws.cell(row=r, column=c)
            current_border = cell.border
            
            # Keep existing border if it has a style (e.g. the badge's thin white border), 
            # else apply the thick outer border to the outer edges of the 2x2 block
            b_top = current_border.top if (current_border.top and current_border.top.style) else (thick_navy if r == r_idx else Side(style=None))
            b_bottom = current_border.bottom if (current_border.bottom and current_border.bottom.style) else (thick_navy if r == r_idx+1 else Side(style=None))
            b_left = current_border.left if (current_border.left and current_border.left.style) else (thick_navy if c == c_idx else Side(style=None))
            b_right = current_border.right if (current_border.right and current_border.right.style) else (thick_navy if c == c_idx+1 else Side(style=None))
            
            cell.border = Border(top=b_top, bottom=b_bottom, left=b_left, right=b_right)
