def render_sheet(wb, sheet_name: str, *, title: str = "Scenario Model", theme: str = "corporate_blue", **kwargs) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(sheet_name)
    
    # Theme palette simulation
    theme_palette = {
        "corporate_blue": {"header_bg": "1F4E78", "header_fg": "FFFFFF", "input_fg": "0070C0", "highlight_bg": "FFF2CC"},
        "modern_green": {"header_bg": "2E7D32", "header_fg": "FFFFFF", "input_fg": "1565C0", "highlight_bg": "E8F5E9"}
    }
    palette = theme_palette.get(theme, theme_palette["corporate_blue"])
    
    # Styles
    header_fill = PatternFill("solid", fgColor=palette["header_bg"])
    header_font = Font(color=palette["header_fg"], bold=True)
    input_font = Font(color=palette["input_fg"])
    toggle_fill = PatternFill("solid", fgColor=palette["highlight_bg"])
    bold_font = Font(bold=True)
    thin_border = Border(
        top=Side(border_style="thin", color="000000"),
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )
    top_thin_bottom_double = Border(
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="double", color="000000")
    )

    # 1. Setup Scenario Toggle (F2) with Data Validation
    ws['E2'] = "Scenario Selector:"
    ws['E2'].font = bold_font
    ws['E2'].alignment = Alignment(horizontal="right")
    
    ws['F2'] = 1
    ws['F2'].fill = toggle_fill
    ws['F2'].border = thin_border
    ws['F2'].alignment = Alignment(horizontal="center")
    
    dv = DataValidation(type="list", formula1='"1,2"', allowBlank=False)
    dv.error = "Please select Scenario 1 or 2"
    ws.add_data_validation(dv)
    dv.add(ws['F2'])

    # 2. Setup Income Statement Headers
    ws['B4'] = title
    ws['B4'].font = bold_font
    headers = ["Year 1", "Year 2", "Year 3"]
    for col_idx, val in enumerate(headers, start=3):
        cell = ws.cell(row=4, column=col_idx)
        cell.value = val
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 3. Write Specific Scenarios (Inputs - Blue Font)
    labels = ["Revenue Growth", "COGS Margin", "Operating Expenses"]
    
    # Scenario 1 (Upper Case)
    ws['B16'] = "Scenario 1: Upper Case"
    ws['B16'].font = bold_font
    scen1_data = [[0.50, 0.40, 0.30], [0.30, 0.30, 0.30], [20000, 25000, 30000]]
    for i, label in enumerate(labels):
        ws.cell(row=17+i, column=2).value = label
        for j, val in enumerate(scen1_data[i]):
            cell = ws.cell(row=17+i, column=3+j)
            cell.value = val
            cell.font = input_font
            cell.number_format = "0%" if i < 2 else "#,##0"

    # Scenario 2 (Lower Case)
    ws['B21'] = "Scenario 2: Lower Case"
    ws['B21'].font = bold_font
    scen2_data = [[0.10, 0.05, 0.05], [0.40, 0.40, 0.40], [20000, 20000, 20000]]
    for i, label in enumerate(labels):
        ws.cell(row=22+i, column=2).value = label
        for j, val in enumerate(scen2_data[i]):
            cell = ws.cell(row=22+i, column=3+j)
            cell.value = val
            cell.font = input_font
            cell.number_format = "0%" if i < 2 else "#,##0"

    # 4. Write Live Case utilizing CHOOSE mechanism
    ws['B11'] = "Live Case Assumptions"
    ws['B11'].font = bold_font
    for i, label in enumerate(labels):
        ws.cell(row=12+i, column=2).value = label
        for j, col_letter in enumerate(['C', 'D', 'E']):
            cell = ws.cell(row=12+i, column=3+j)
            # Core Mechanism: CHOOSE($F$2, Scen1, Scen2)
            cell.formula = f"=CHOOSE($F$2, {col_letter}{17+i}, {col_letter}{22+i})"
            cell.number_format = "0%" if i < 2 else "#,##0"

    # 5. Connect Income Statement to Live Case
    ws['B5'] = "Revenue"
    ws['C5'].formula = "=100000 * (1+C12)" # Base Yr0 = 100k
    ws['D5'].formula = "=C5 * (1+D12)"
    ws['E5'].formula = "=D5 * (1+E12)"
    
    ws['B6'] = "Cost of Goods Sold"
    for col in ['C', 'D', 'E']:
        ws[f'{col}6'].formula = f"={col}5 * {col}13"
        
    ws['B7'] = "Gross Profit"
    for col in ['C', 'D', 'E']:
        ws[f'{col}7'].formula = f"={col}5 - {col}6"
        ws[f'{col}7'].font = bold_font
        ws[f'{col}7'].border = Border(top=Side(border_style="thin", color="000000"))
        
    ws['B8'] = "Operating Expenses"
    for col in ['C', 'D', 'E']:
        ws[f'{col}8'].formula = f"={col}14"
        
    ws['B9'] = "Operating Profit"
    for col in ['C', 'D', 'E']:
        ws[f'{col}9'].formula = f"={col}7 - {col}8"
        ws[f'{col}9'].font = bold_font
        ws[f'{col}9'].border = top_thin_bottom_double

    # Format P&L block
    for row in range(5, 10):
        for col in ['C', 'D', 'E']:
            ws[f'{col}{row}'].number_format = "#,##0"

    # Polish column widths
    ws.column_dimensions['B'].width = 25
    for col in ['C', 'D', 'E']:
        ws.column_dimensions[col].width = 14
